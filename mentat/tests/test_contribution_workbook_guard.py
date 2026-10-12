"""The current workbook must agree with its analytical receipt source."""
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'app'))
import contribution_workbook_guard as guard


class PublishedWorkbookGuardTests(unittest.TestCase):
    def setUp(self):
        verifier=patch.object(guard.source_guard,'verify',return_value={'full_financial_universe_preserved':True,'proof_token':'test-proof'})
        self.verifier=verifier.start();self.addCleanup(verifier.stop)
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.path=Path(self.tmp.name)/'Example Industry Spending Overview.xlsx'
        book=openpyxl.Workbook();sheet=book.active;sheet.title='Spending Overview'
        sheet.append(['Example report'])
        for row in range(2,5):sheet.append([])
        sheet.append(['Cycle','Scope','Total Giving'])
        sheet.append(['2023-2024','Total Giving',100.25])
        sheet.append(['2023-2024','Employee Giving',100.25])
        book.save(self.path);book.close()
        self.source={'id':'Example','industry':'Example','available':True,'workbook':self.path,
            'workbook_sha256':'test-source','expected_cycle_totals':{'2023-2024':100.25}}

    def test_current_cached_totals_match_without_summing_breakdowns(self):
        with patch.object(guard.contributions,'_sources',return_value=[self.source]):
            result=guard.verify_sources(self.path.parent)
        self.assertEqual(result[0]['published_cycle_cents'],{'2023-2024':10025})
        self.assertEqual(result[0]['source_proof_token'],'test-proof')
        self.verifier.assert_called_once_with(self.source)

    def test_source_financial_failure_is_rejected_or_isolated_by_catalog_mode(self):
        self.verifier.side_effect=ValueError('Financial identities changed')
        with patch.object(guard.contributions,'_sources',return_value=[self.source]):
            with self.assertRaisesRegex(ValueError,'Financial identities changed'):
                guard.verify_sources(self.path.parent)
            self.assertEqual(guard.verify_sources(self.path.parent,strict=False),[
                {'dataset':'Example','matched':False,'error':'Example: Financial identities changed'}])

    def test_altered_total_is_rejected_even_when_source_notes_are_unchanged(self):
        with patch.object(guard.contributions,'_sources',return_value=[self.source]):
            guard.verify_sources(self.path.parent)
            book=openpyxl.load_workbook(self.path);book.active['C6']=200.50;book.save(self.path);book.close()
            with self.assertRaisesRegex(ValueError,'differ from the pinned contribution source'):
                guard.verify_sources(self.path.parent)

    def test_duplicate_total_row_is_not_silently_overwritten(self):
        book=openpyxl.load_workbook(self.path);book.active.append(['2023-2024','Total Giving',100.25]);book.save(self.path);book.close()
        with self.assertRaisesRegex(ValueError,'more than one'):
            guard.cycle_totals(guard.contributions._stamp(self.path))


if __name__=='__main__':unittest.main()
