"""A stale or missing transfer index publishes no zeros (Paul's ruling 2026-09-25, F2).

A cycle whose transfer index is not current publishes NO transfer-derived numbers: no
Transfer Paths, Allocation Estimates or Allocation Coverage rows, and its conditional
estimates are blank (None / an empty CSV cell), never 0.00. A visible note says
"Not traced: the transfer index for <cycle> is <status>; rebuild with <command>", and the
status is named per cycle in metadata.transfer_index_status, a transfer_index_<cycle>
quality check and the audit manifest. Everything else in the report is unchanged, and a
CURRENT index produces exactly what the code produced before the ruling.
"""
import csv
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

MENTAT = Path(__file__).resolve().parents[1]
TOOLS = MENTAT / 'tools'
sys.path.insert(0, str(TOOLS))
import enrich_downstream  # noqa: E402
from enrich_downstream import enrich  # noqa: E402

TRANSFER_TABLES = ('transfer_paths', 'allocation_estimates', 'allocation_coverage')
ESTIMATES = ('estimated_non_opposition', 'estimated_opposition', 'direct_plus_estimated_non_opposition')
AUDIT_NAMES = ('transfer_paths', 'allocation_estimates', 'allocation_coverage', 'attribution_summary', 'funding_cohorts')

# The pre-ruling enrich_downstream.py (blob cf388500776d81796edecb5d20c5cf0c6bf8d0bd, commit
# 3cec98524) run on fixture('indexed'): the payload with the three additive status fields
# removed, and the audit CSVs it publishes. A current index must reproduce these bytes.
# F14 (Paul approved option (a), 2026-09-26) reshaped transfer_paths to one row per transfer
# and cycle by design (tests/test_transfer_paths_contract.py): the payload is compared without
# it (pre-ruling payload f43aadb8... hashed that way is 71d594b2...), and its CSV golden is the
# new shape's. The four other CSVs keep their pre-ruling bytes.
GOLDEN_PAYLOAD_SHA256 = '71d594b295beafff41f3fb67bf69e23208b7dcbfa9e93a28636a3d66fc94fa1e'
GOLDEN_CSV_SHA256 = {
    'transfer_paths': 'e5fa9f8099c96322b91645dd6edc0455abfa0aa3f064cc63474acbd5ce3b6c41',
    'allocation_estimates': '168b28c341e8ef8ffa912e6414f32329ebfb1369609c37925444ae2bfe35047a',
    'allocation_coverage': 'b168abc622ec85e8aa37250e68511e950426be63bf515990a1b5b201946e9f64',
    'attribution_summary': '11ccae265549b04fda877cfc58effb4949abcdc72718d9eb3a2337d2e4cf0668',
    'funding_cohorts': '7df7ac32030cf4de78afca933a6bbeae67f1827a771c309cbd4854c5eee9a2dc',
}


def fixture(status_2024):
    """Two cycles with identical activity: 2022 is current; 2024 has `status_2024`.

    status_2024 is a transfer_coverage status ('indexed', 'stale_transfer_index',
    'missing_transfer_source'), 'unresolved' (indexed but unresolved_all_sources) or None
    (no coverage row at all). A root super PAC (C00000001) funds candidate activity and
    transfers to C00000002, which opposes H2; C00000009 is H9's principal committee.
    """
    coverage = [dict(cycle=2022, status='indexed', counts={})]
    if status_2024 == 'unresolved':
        coverage.append(dict(cycle=2024, status='indexed', counts={'malformed_rows': 2}, unresolved_all_sources=True))
    elif status_2024 == 'stale_transfer_index':
        coverage.append(dict(cycle=2024, status=status_2024, unresolved_all_sources=True))
    elif status_2024 is not None:
        coverage.append(dict(cycle=2024, status=status_2024, counts={}))
    payload = {'metadata': {'industry': 'Fixture', 'transfer_coverage': coverage},
               'cycles': [{'cycle': 2022}, {'cycle': 2024}], 'committee_map': [], 'candidate_master': [],
               'raw_rows': [], 'downstream_activity': [], 'candidate_committee_links': [], 'transfer_edges': []}
    for cycle in (2022, 2024):
        payload['raw_rows'] += [
            dict(cycle=cycle, committee_id='C00000001', amount=100, company='Acme', include_in_analysis='Yes',
                 receipt_role='contribution', date=f'{cycle}-03-01'),
            dict(cycle=cycle, committee_id='C00000009', amount=50, company='Acme', include_in_analysis='Yes',
                 receipt_role='contribution', date=f'{cycle}-04-01')]
        payload['candidate_committee_links'].append(
            dict(cycle=cycle, committee_id='C00000009', candidate_id='H9', designation='P'))
        payload['candidate_master'] += [dict(cycle=cycle, candidate_id=f'H{n}', candidate_name=f'Candidate {n}',
                                             candidate_party='DEM') for n in (1, 2, 9)]
        payload['transfer_edges'].append(dict(cycle=cycle, source='C00000001', destination='C00000002', amount=60))
        payload['downstream_activity'] += [
            dict(cycle=cycle, source_committee_id='C00000001', candidate_id='H1', support_oppose='Support',
                 preliminary_included_amount=40),
            dict(cycle=cycle, source_committee_id='C00000002', candidate_id='H2', support_oppose='Oppose',
                 preliminary_included_amount=30),
            dict(cycle=cycle, source_committee_id='C00000001', candidate_id='H9', support_oppose='Support',
                 preliminary_included_amount=10)]
    return payload


def strip_additive(result):
    """The payload as the pre-ruling code wrote it: the three status fields removed."""
    value = json.loads(json.dumps(result))
    value['metadata'].pop('transfer_index_status', None)
    value['metadata'].pop('transfer_index_notes', None)
    value['quality_checks'] = [r for r in value['quality_checks'] if not str(r.get('id', '')).startswith('transfer_index_')]
    return value


def payload_sha256(result):
    value = strip_additive(result)
    value.pop('transfer_paths', None)
    return hashlib.sha256(json.dumps(value, separators=(',', ':')).encode()).hexdigest()


def zero_valued_transfer_numbers(result, cycle):
    """Every transfer-derived number a report would print for `cycle` that is 0 or 0.00."""
    zeros = []
    for table in TRANSFER_TABLES:
        for row in result[table]:
            if row.get('cycle') == cycle:
                zeros += [(table, key) for key, value in row.items() if key not in ('cycle', 'depth')
                          and isinstance(value, (int, float)) and not isinstance(value, bool) and value == 0]
    for row in result['attribution_summary']:
        if row['cycle'] == cycle:
            zeros += [('attribution_summary', key) for key in ESTIMATES if row[key] == 0]
    return zeros


def run_audit(payload, root, workbook_name='Fixture Industry Spending Overview.xlsx'):
    """Publish the audit bundle through the real CLI, exactly as build_industry_report.ps1 does."""
    payload_path = root / 'payload.json'
    payload_path.write_text(json.dumps(payload), encoding='utf-8')
    workbook = root / workbook_name
    if not workbook.exists():
        workbook.write_bytes(b'fixture workbook bytes')
    audit = root / (workbook.stem + '.audit')
    subprocess.run([sys.executable, str(TOOLS / 'enrich_downstream.py'), str(payload_path), str(root),
                    '--audit-only', '--audit-directory', str(audit), '--workbook', str(workbook)],
                   check=True, capture_output=True)
    return audit


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream))


class NotTracedPublishesNothingTransferDerived(unittest.TestCase):
    def enrich(self, payload):
        with tempfile.TemporaryDirectory() as directory:
            result = enrich(payload, Path(directory))
            return result, Path(directory).resolve()

    def assert_not_traced(self, result, cycle, status, directory):
        for table in TRANSFER_TABLES:
            self.assertEqual([r for r in result[table] if r.get('cycle') == cycle], [], table)
        rows = [r for r in result['attribution_summary'] if r['cycle'] == cycle]
        self.assertEqual([r['candidate_id'] for r in rows], ['H9'])  # direct funding still publishes
        self.assertEqual(rows[0]['direct'], 50)
        self.assertEqual([rows[0][key] for key in ESTIMATES], [None, None, None])
        self.assertEqual(zero_valued_transfer_numbers(result, cycle), [])
        self.assertEqual(result['metadata']['transfer_index_status'], {'2022': 'current', str(cycle): status})
        command = (f'python "{TOOLS.resolve() / "fec_transfer_graph.py"}" --directory "{directory}" --cycles {cycle}')
        note = result['metadata']['transfer_index_notes'][str(cycle)]
        self.assertTrue(note.startswith(f'Not traced: the transfer index for {cycle} is {status}; rebuild with {command}'), note)
        self.assertEqual(list(result['metadata']['transfer_index_notes']), [str(cycle)])
        check = next(r for r in result['quality_checks'] if r['id'] == f'transfer_index_{cycle}')
        self.assertEqual((check['status'], check['passed'], check['severity'], check['cycle']), (status, False, 'warn', cycle))
        self.assertTrue(check['detail'].startswith(note))
        # The plain `warn` is replaced by the named status; the current cycle keeps its check.
        self.assertFalse(any(r['id'] == f'transfer_source_{cycle}' for r in result['quality_checks']))
        self.assertTrue(any(r['id'] == 'transfer_source_2022' for r in result['quality_checks']))
        current = next(r for r in result['quality_checks'] if r['id'] == 'transfer_index_2022')
        self.assertEqual((current['status'], current['passed'], current['severity']), ('current', True, 'info'))

    def assert_2022_still_traced(self, result):
        estimates = {(r['candidate_id'], r['direction']): r['estimated_amount']
                     for r in result['allocation_estimates'] if r['cycle'] == 2022}
        self.assertAlmostEqual(estimates['H1', 'Support'], 100 * 40 / 110)
        self.assertAlmostEqual(estimates['H2', 'Oppose'], 100 * 60 / 110)
        self.assertTrue(any(r['cycle'] == 2022 for r in result['transfer_paths']))

    def test_a_stale_index_publishes_no_zero_valued_transfer_rows_and_names_its_status(self):
        result, directory = self.enrich(fixture('stale_transfer_index'))
        self.assert_not_traced(result, 2024, 'stale', directory)
        self.assert_2022_still_traced(result)

    def test_a_missing_index_is_named_missing_whether_or_not_r_wrote_a_row(self):
        for status in ('missing_transfer_source', None):
            with self.subTest(status=status):
                result, directory = self.enrich(fixture(status))
                self.assert_not_traced(result, 2024, 'missing', directory)
                self.assert_2022_still_traced(result)

    def test_an_index_that_cannot_place_its_source_is_named_unresolved(self):
        result, directory = self.enrich(fixture('unresolved'))
        self.assert_not_traced(result, 2024, 'unresolved', directory)

    def test_negative_control_without_the_status_gate_the_stale_cycle_publishes_zeros(self):
        # The pre-ruling behaviour, reproduced by pretending every index is current: the
        # stale cycle's estimates come out as 0.00 and its roots as "unresolved" coverage.
        # The assertions above would fail on exactly this output.
        with mock.patch.object(enrich_downstream, 'transfer_index_status', return_value='current'):
            result, _ = self.enrich(fixture('stale_transfer_index'))
        self.assertIn(('attribution_summary', 'estimated_non_opposition'), zero_valued_transfer_numbers(result, 2024))
        self.assertTrue(any(r['cycle'] == 2024 for r in result['allocation_coverage']))
        with self.assertRaises(AssertionError):
            self.assert_not_traced(result, 2024, 'stale', Path('.'))

    def test_the_audit_bundle_keeps_blanks_blank_and_carries_the_status(self):
        result, _ = self.enrich(fixture('stale_transfer_index'))
        with tempfile.TemporaryDirectory() as temporary:
            audit = run_audit(result, Path(temporary))
            manifest = json.loads((audit / 'audit-manifest.json').read_text(encoding='utf-8'))
            self.assertEqual(manifest['transfer_index_status'], {'2022': 'current', '2024': 'stale'})
            self.assertEqual(list(manifest['transfer_index_notes']), ['2024'])
            summary = [r for r in read_csv(audit / 'attribution_summary.csv') if r['cycle'] == '2024']
            self.assertEqual([[r[key] for key in ESTIMATES] for r in summary], [['', '', '']])
            self.assertEqual(float(summary[0]['direct']), 50)
            for name in TRANSFER_TABLES:
                self.assertFalse(any(r.get('cycle') == '2024' for r in read_csv(audit / f'{name}.csv')), name)
            checks = {r['id']: r for r in read_csv(audit / 'quality_checks.csv')}
            self.assertEqual((checks['transfer_index_2024']['status'], checks['transfer_index_2024']['passed']),
                             ('stale', 'False'))
            self.assertTrue(checks['transfer_index_2024']['detail'].startswith('Not traced: the transfer index for 2024 is stale; rebuild with python '))

    def test_a_per_cycle_bundle_names_only_its_own_cycle(self):
        result, _ = self.enrich(fixture('stale_transfer_index'))
        result['cycles'] = [{'cycle': 2022}]
        with tempfile.TemporaryDirectory() as temporary:
            manifest = json.loads((run_audit(result, Path(temporary)) / 'audit-manifest.json').read_text(encoding='utf-8'))
        self.assertEqual((manifest['transfer_index_status'], manifest['transfer_index_notes']), ({'2022': 'current'}, {}))

    def test_a_payload_enriched_before_the_status_existed_writes_the_manifest_it_always_did(self):
        result, _ = self.enrich(fixture('indexed'))
        legacy = strip_additive(result)
        with tempfile.TemporaryDirectory() as temporary:
            manifest = json.loads((run_audit(legacy, Path(temporary)) / 'audit-manifest.json').read_text(encoding='utf-8'))
        self.assertEqual(list(manifest), ['version', 'workbook_name', 'workbook_sha256', 'payload_path',
                                          'payload_sha256', 'csv_sha256'])


class CurrentIndexIsUnchanged(unittest.TestCase):
    def test_a_current_index_reproduces_the_pre_ruling_payload_and_audit_csvs_byte_for_byte(self):
        with tempfile.TemporaryDirectory() as directory:
            result = enrich(fixture('indexed'), Path(directory))
        self.assertEqual(result['metadata']['transfer_index_status'], {'2022': 'current', '2024': 'current'})
        self.assertEqual(result['metadata']['transfer_index_notes'], {})
        self.assertEqual(payload_sha256(result), GOLDEN_PAYLOAD_SHA256)
        with tempfile.TemporaryDirectory() as temporary:
            audit = run_audit(result, Path(temporary))
            actual = {name: hashlib.sha256((audit / f'{name}.csv').read_bytes()).hexdigest() for name in AUDIT_NAMES}
        self.assertEqual(actual, GOLDEN_CSV_SHA256)


@unittest.skipUnless(shutil.which('node') and (Path(os.environ.get('MENTAT_NODE_MODULES') or Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules') / '@oai/artifact-tool').is_dir(),
                     'The attribution workbook needs node and @oai/artifact-tool')
class AttributionWorkbookKeepsBlanksBlank(unittest.TestCase):
    """build_attribution_report.py -> attribution_workbook_builder.mjs -> verify_attribution_report.py."""

    def build(self, payload):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        with tempfile.TemporaryDirectory() as directory:
            result = enrich(payload, Path(directory))
        run_audit(result, root)
        subprocess.run([sys.executable, str(TOOLS / 'build_attribution_report.py'), str(root), '--node', shutil.which('node')],
                       check=True, capture_output=True, cwd=str(TOOLS))
        return root

    def test_untraced_rows_are_blank_and_carry_no_formula_while_traced_rows_keep_theirs(self):
        import openpyxl
        sys.path.insert(0, str(TOOLS))
        import verify_attribution_report
        root = self.build(fixture('stale_transfer_index'))
        self.assertEqual(verify_attribution_report.verify(root)['status'], 'passed')
        book = openpyxl.load_workbook(root / 'Attributed Industry Spending.xlsx')
        try:
            sheet = book['Candidate Funding']
            rows = {(row[1].value, row[2].value): row for row in sheet.iter_rows(min_row=6)}
            untraced = rows[2024, 'H9']
            self.assertEqual(untraced[5].value, 50)
            self.assertEqual([untraced[i].value in (None, '') for i in (6, 7, 8)], [True, True, True])
            traced = rows[2022, 'H9']
            self.assertTrue(str(traced[8].value).startswith('='), traced[8].value)
            self.assertIn('Blank is not zero', sheet['A2'].value)
        finally:
            book.close()

    def test_negative_control_a_blank_that_came_back_as_zero_is_refused(self):
        import openpyxl
        import verify_attribution_report
        payload = fixture('stale_transfer_index')
        payload['cycles'] = [{'cycle': 2024}]  # one untraced cycle only: no formula cells to lose on save
        root = self.build(payload)
        path = root / 'Attributed Industry Spending.xlsx'
        book = openpyxl.load_workbook(path)
        book['Candidate Funding'].cell(6, 7).value = 0
        book.save(path)
        with self.assertRaisesRegex(ValueError, 'must stay blank'):
            verify_attribution_report.verify(root)


if __name__ == '__main__':
    unittest.main()
