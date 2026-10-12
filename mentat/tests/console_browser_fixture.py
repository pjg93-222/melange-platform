"""Temporary local-only report server for test_console_browser.cjs."""
import json
import sqlite3
from contextlib import closing
from pathlib import Path
import sys
import tempfile
import threading

import openpyxl
from graph_fixture import add_graph_sheets

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'app'))
import server as app
from test_validate_industry_coverage import create_complete_database, write_receipt
import validate_industry_coverage as coverage_gate


class FixtureConsole(app.Console):
    def __init__(self, root):
        self.root = root
        self.report_directory = root
        self.token = 'offline-fixture-token'
        self.report_cache = {}
        self.processes = {}
        self.lock = threading.Lock()

    def runs(self):
        return [{'id':'fixture','name':'Fictional credit union sample','state':'completed',
                 'created_at':1,'directory':self.root}]

    def summary(self, include_archived=False):
        archived = 'fixture' in app.read_json(self.root / 'run-archives.json').get('runs', {})
        return {'runs':[{'id':'fixture','name':'Fictional credit union sample','state':'completed',
                         'archived':archived,'can_archive':True}] if include_archived or not archived else [],
                'coverage':self.coverage_catalog(self.runs()),
                'archived_count':int(archived),'total_run_count':1,
                'collection_directory':str(self.root),'report_directory':str(self.report_directory),
                'key_configured':False,'active_key_count':0,'stored_key_count':0,'evidence_collection_available':True}


def main():
    with tempfile.TemporaryDirectory(prefix='mentat-chart-test-') as temporary:
        root = Path(temporary)
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'Candidate Funding'
        for row in [['Credit union sample report'],['Fictional values for UI tests'],[],[],
                    ['Company','Amount','Other Amount','Cycle','Date'],
                    ['Harborlight Credit Union',10,100,2026,'2026-01-01'],['Harborlight Credit Union',30,120,2026,'2026-01-02'],
                    ['Civic Horizon Bank',60,200,2024,'2026-01-03'],['Civic Horizon Bank',None,300,2024,'2026-01-04'],
                    ['<img src=x onerror=alert(1)>',-5,0,2026,'2026-01-05']]:
            sheet.append(row)
        receipts = book.create_sheet('Source Freshness')
        receipts.append(['FEC Dataset','Cycle','Downloaded At','Checked At','Verified Data Through'])
        receipts.append(['cm',2026,'2026-08-01T00:00:00Z','2026-09-08T00:00:00Z','2026-07-31'])
        receipts.append(['oth',2026,None,'2026-09-08T00:00:00Z',None])
        regions = book.create_sheet('Regional sample')
        for row in [['Region','Volume','Cost'], ['North',10,20], ['South',15,30], ['East',25,60]]:
            regions.append(row)
        history = book.create_sheet('Annual sample')
        for row in [['Date', 'Amount'], ['2022-01-01', 100], ['2023-01-01', 120], ['2024-01-01', 140], ['2025-01-01', 160]]:
            history.append(row)
        add_graph_sheets(book)
        if '--multi-candidate' in sys.argv:
            book['Downstream Activity']['A3']='101 source rows condensed to 5 committee-candidate-activity groups'
            book['Downstream Activity'].append(['2025-2026','C00000003','H0CA00002','Candidate Two','R','Support',250])
        report = root / 'Credit Union Sample.xlsx'
        book.save(report)
        book.close()
        database, fanout, coverage = root / 'mentat.db', root / 'fanout.json', root / 'coverage.json'
        create_complete_database(database)
        with closing(sqlite3.connect(database)) as connection, connection:
            connection.execute("UPDATE companies SET industry='Credit Unions'")
        assertions = [dict(industry='Credit Unions',company=company,cycle=cycle,kind=kind,
                           coverage_start=f'{cycle-1}-01-01',verified_data_through=f'{cycle}-12-31' if cycle == 2022 else '2026-06-30',
                           verification_evidence='Fictional dated source evidence')
                      for company in ('Alpha','Beta') for cycle in (2022,2026) for kind in ('employer','corporate')]
        write_receipt(fanout, database, fec_date_evidence=assertions)
        coverage.write_text(json.dumps(coverage_gate.validate(database,'Credit Unions',(2022,2024,2026),fanout)))
        (root / 'status.json').write_text(json.dumps({'stage':'completed','cycles':[2022,2024,2026],'industries':[
            {'industry':'Credit Unions','status':'succeeded','output_workbook':str(report),
             'database':str(database),'coverage_receipt':str(coverage)}]}))
        server = app.ThreadingHTTPServer(('127.0.0.1',0), app.Handler)
        server.app = FixtureConsole(root)
        worker = threading.Thread(target=server.serve_forever,daemon=True)
        worker.start()
        print(f'http://127.0.0.1:{server.server_port}',flush=True)
        try:
            if '--serve' in sys.argv:
                threading.Event().wait()
            else:
                sys.stdin.readline()
        finally:
            server.shutdown()
            server.server_close()
            worker.join()


if __name__ == '__main__':
    main()
