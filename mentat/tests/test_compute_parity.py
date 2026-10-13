"""Content-level parity harness (tools/compute_parity.py) on tiny synthetic outputs, and the anchors job end to end: the same job run
twice (variance) and once more from a copy at another root (the container's mounts), compared with a mask derived from the first two."""
import json
import os
from pathlib import Path
import shutil
import sqlite3
import sys
import tempfile
import time
import unittest
import zipfile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools'))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import compute_job as cj  # noqa: E402
import compute_parity as parity  # noqa: E402
from compute_world import make_layout, move_layout  # noqa: E402


def write_run(folder, stamp, root, order=(1, 2, 3), text_newline='\n'):
    """One synthetic job output: every file kind, a run stamp (noise) and a root path (the mount) in it."""
    folder = Path(folder)
    folder.mkdir(parents=True)
    (folder / 'snapshot.json').write_text(json.dumps({'built_at_utc': stamp, 'levels': {'banks': [1.5, 2.5], 'cu': 7}, 'input': root + '/reports/Banking.xlsx',
                                                     'note': 'read ' + root + '/state ok'}), encoding='utf-8')
    (folder / 'table.csv').write_text('id,when,amount\n1,%s,10.5\n2,%s,20\n' % (stamp, stamp), encoding='utf-8')
    db = sqlite3.connect(folder / 'store.sqlite')
    db.execute('CREATE TABLE runs(id INTEGER PRIMARY KEY, finished_at TEXT, amount REAL)')
    db.execute('CREATE TABLE loose(name TEXT, n INTEGER)')
    for i in order:
        db.execute('INSERT INTO runs VALUES (?,?,?)', (i, stamp, i * 1.25))
        db.execute('INSERT INTO loose VALUES (?,?)', ('row%d' % i, i))
    db.commit()
    db.close()
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Summary'
    sheet['A1'], sheet['B1'], sheet['A2'], sheet['B2'] = 'Generated', stamp, 'Total', '=SUM(C1:C3)'
    for row in (1, 2, 3):
        sheet.cell(row=row, column=3, value=row * 100)
    book.create_sheet('Detail')['A1'] = 'x'
    book.save(folder / 'book.xlsx')
    with zipfile.ZipFile(folder / 'bundle.zip', 'w') as archive:
        archive.writestr(zipfile.ZipInfo('data.json', (2026, 1, 1 + int(stamp[-3:-1]) % 20, 0, 0, 0)), '{"rows": [1, 2]}')
        archive.writestr(zipfile.ZipInfo('meta.json', (2026, 1, 1, 0, 0, 0)), json.dumps({'stamp': stamp}))
    (folder / 'notes.txt').write_bytes(('line one\nline two\n').replace('\n', text_newline).encode())
    (folder / 'blob.bin').write_bytes(b'\x00\x01\x02 constant')


class Temp(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)

    def runs(self):
        write_run(self.base / 'a1', '2026-10-12T03:30:01Z', 'C:/Users/pgoff/Mentat')
        write_run(self.base / 'a2', '2026-10-12T03:30:09Z', 'C:/Users/pgoff/Mentat')
        write_run(self.base / 'b', '2026-10-13T03:30:05Z', '/mentat', text_newline='\r\n')
        return self.base / 'a1', self.base / 'a2', self.base / 'b'


MAPS = [('C:/Users/pgoff/Mentat', '/mentat')]


class HarnessTests(Temp):
    def test_variance_masks_exactly_the_locations_that_differ_between_two_reference_runs(self):
        a1, a2, _ = self.runs()
        mask = parity.variance(a1, a2)
        self.assertEqual(mask['files']['snapshot.json'], ['/built_at_utc'])
        self.assertEqual(mask['files']['table.csv'], ['r2/c2', 'r3/c2'])
        self.assertEqual(mask['files']['store.sqlite'], ['table:runs/col:finished_at'])
        self.assertEqual(mask['files']['book.xlsx'], ['Summary!B1', 'Summary!B1@cached'])
        self.assertNotIn('Summary!A1', mask['files']['book.xlsx'])
        self.assertEqual(mask['files']['bundle.zip'], ['meta.json'])
        self.assertNotIn('blob.bin', mask['files'])
        self.assertNotIn('notes.txt', mask['files'])
        self.assertEqual(mask['unstable_files'], [])

    def test_the_container_run_matches_outside_the_mask_with_the_mount_declared_and_each_class_is_named(self):
        a1, a2, b = self.runs()
        report = parity.compare(a1, b, parity.variance(a1, a2), MAPS)
        self.assertTrue(report['ok'], report)
        classes = {rel: v['class'] for rel, v in report['files'].items()}
        self.assertEqual(classes, {'snapshot.json': 'PATH_ONLY', 'table.csv': 'CSV_CANONICAL', 'store.sqlite': 'SQLITE_CANONICAL', 'book.xlsx': 'OOXML_CANONICAL',
                                   'bundle.zip': 'ZIP_CANONICAL', 'notes.txt': 'EOL_ONLY', 'blob.bin': 'BYTE_EQUAL'})

    def test_without_the_mask_or_the_mount_the_same_runs_are_different(self):
        a1, a2, b = self.runs()
        unmasked = parity.compare(a1, b, None, MAPS)
        self.assertFalse(unmasked['ok'])
        self.assertEqual({rel for rel, v in unmasked['files'].items() if v['class'] == 'DIFFERENT'}, {'snapshot.json', 'table.csv', 'store.sqlite', 'book.xlsx', 'bundle.zip'})
        self.assertEqual(parity.compare(a1, b, parity.variance(a1, a2), [])['files']['snapshot.json']['class'], 'DIFFERENT')

    def test_a_row_order_difference_is_not_a_difference_and_a_changed_value_is(self):
        a1, a2, _ = self.runs()
        write_run(self.base / 'shuffled', '2026-10-12T03:30:01Z', 'C:/Users/pgoff/Mentat', order=(3, 1, 2))
        mask = parity.variance(a1, a2)
        self.assertTrue(parity.compare(a1, self.base / 'shuffled', mask, MAPS)['ok'])
        db = sqlite3.connect(self.base / 'shuffled' / 'store.sqlite')
        db.execute('UPDATE loose SET n = 99 WHERE name = "row2"')
        db.commit()
        db.close()
        report = parity.compare(a1, self.base / 'shuffled', mask, MAPS)
        self.assertEqual((report['ok'], report['files']['store.sqlite']['where']), (False, ['table:loose/col:n']))

    def test_negative_controls_each_name_the_location_that_changed(self):
        a1, a2, b = self.runs()
        mask = parity.variance(a1, a2)
        snapshot = json.loads((b / 'snapshot.json').read_text())
        snapshot['levels']['banks'][1] = 2.6
        (b / 'snapshot.json').write_text(json.dumps(snapshot), encoding='utf-8')
        book = openpyxl.load_workbook(b / 'book.xlsx')
        book['Summary']['C2'] = 201
        book['Detail']['A1'] = 'y'
        book.save(b / 'book.xlsx')
        with zipfile.ZipFile(b / 'bundle.zip', 'w') as archive:
            archive.writestr('data.json', '{"rows": [1, 3]}')
            archive.writestr('meta.json', '{}')
        (b / 'table.csv').write_text('id,when,amount\n1,x,10.5\n2,x,21\n', encoding='utf-8')
        report = parity.compare(a1, b, mask, MAPS)
        self.assertFalse(report['ok'])
        self.assertEqual(report['files']['snapshot.json']['where'], ['/levels/banks/1'])
        self.assertEqual(report['files']['book.xlsx']['where'], ['Detail!A1', 'Detail!A1@cached', 'Summary!C2', 'Summary!C2@cached'])
        self.assertEqual(report['files']['bundle.zip']['where'], ['data.json'])
        self.assertEqual(report['files']['table.csv']['where'], ['r3/c3'])

    def test_a_file_on_one_side_only_fails_the_census_unless_the_reference_runs_disagree_about_it(self):
        a1, a2, b = self.runs()
        (b / 'extra.json').write_text('{}', encoding='utf-8')
        (b / 'blob.bin').unlink()
        report = parity.compare(a1, b, parity.variance(a1, a2), MAPS)
        self.assertEqual((report['ok'], report['only_in_a'], report['only_in_b']), (False, ['blob.bin'], ['extra.json']))
        (a2 / 'blob.bin').unlink()
        self.assertEqual(parity.compare(a1, b, parity.variance(a1, a2), MAPS)['only_in_a'], [])            # unstable in the reference itself

    def test_a_mask_belongs_to_the_reference_run_it_came_from(self):
        a1, a2, b = self.runs()
        mask = parity.variance(a1, a2)
        with self.assertRaises(ValueError):
            parity.compare(b, a1, mask, MAPS)                                    # the container run cannot be the reference of a Windows-derived mask
        (a1 / 'blob.bin').write_bytes(b'tampered')
        with self.assertRaises(ValueError):
            parity.compare(a1, b, mask, MAPS)
        write_run(self.base / 'a3', '2026-10-12T03:30:02Z', 'C:/Users/pgoff/Mentat')
        self.assertTrue(parity.compare(a2, self.base / 'a3', mask, MAPS)['ok'])             # A2 is as good a reference as A1

    def test_ignore_leaves_run_specific_files_out_of_the_census_and_the_digest(self):
        a1, a2, b = self.runs()
        for folder in (a1, a2, b):
            (folder / ('receipt-%s.json' % folder.name)).write_text('{}', encoding='utf-8')
        report = parity.compare(a1, b, parity.variance(a1, a2, ['receipt-*.json']), MAPS, ['receipt-*.json'])
        self.assertTrue(report['ok'], report)

    def test_the_command_line_runs_variance_then_compare_and_exits_by_the_verdict(self):
        a1, a2, b = self.runs()
        mask = self.base / 'mask.json'
        self.assertEqual(parity.main(['variance', str(a1), str(a2), '--out', str(mask)]), 0)
        self.assertEqual(parity.main(['compare', str(a1), str(b), '--mask', str(mask), '--map', 'C:/Users/pgoff/Mentat=/mentat', '--report', str(self.base / 'report.json')]), 0)
        self.assertTrue(json.loads((self.base / 'report.json').read_text())['ok'])
        (b / 'blob.bin').write_bytes(b'other')
        self.assertEqual(parity.main(['compare', str(a1), str(b), '--mask', str(mask), '--map', 'C:/Users/pgoff/Mentat=/mentat']), 1)


class AnchorsParityTests(Temp):
    """The anchors job, run for real: A1 and A2 on one frozen layout, B from a copy at another root (what the container's mounts are)."""

    def run_anchors(self, name, env_extra, extra, output):
        env = {**{k: v for k, v in os.environ.items() if not k.startswith(('FEC_API_KEY', 'MENTAT_', 'RF_'))}, **env_extra}
        args = [*extra[:-1], str(output)]
        from unittest.mock import patch
        with patch.object(cj.glassbox_emit, 'EVENTS_DIR', self.base / 'gb'):
            self.assertEqual(cj.run('anchors', name, args, env), 0)

    def test_the_container_layout_builds_the_same_anchors_outside_the_run_noise(self):
        env, extra = make_layout(self.base / 'win')
        self.run_anchors('a1', env, extra, self.base / 'a1')
        time.sleep(1.1)                                                  # built_at_utc has one-second resolution: two reference runs must not share a second
        self.run_anchors('a2', env, extra, self.base / 'a2')
        # elapsed_seconds of two runs of a tiny fixture can coincide to the hundredth; the method measures the noise two runs show, so guarantee that it shows
        receipt = self.base / 'a2' / 'giving-anchors-v1.receipt.json'
        record = json.loads(receipt.read_text(encoding='utf-8'))
        record['elapsed_seconds'] += 1000
        receipt.write_text(json.dumps(record, indent=1), encoding='utf-8')
        env_b, extra_b = move_layout(self.base / 'win', self.base / 'box')
        self.run_anchors('b', env_b, extra_b, self.base / 'b')
        mask = parity.variance(self.base / 'a1', self.base / 'a2')
        self.assertIn('/built_at_utc', mask['files']['giving-anchors-v1.json'])
        self.assertIn('/snapshot_sha256', mask['files']['giving-anchors-v1.receipt.json'])
        report = parity.compare(self.base / 'a1', self.base / 'b', mask, [(str(self.base / 'win').replace('\\', '/'), str(self.base / 'box').replace('\\', '/'))])
        self.assertTrue(report['ok'], report)
        self.assertEqual(report['files']['giving-anchors-v1.receipt.json']['class'], 'PATH_ONLY')       # the receipt names the inputs by their mount
        self.assertEqual(report['files']['giving-anchors-v1.json']['class'], 'JSON_CANONICAL')
        # negative control: one anchor changed in the container run is found, with its location
        snapshot = json.loads((self.base / 'b' / 'giving-anchors-v1.json').read_text(encoding='utf-8'))
        cycles = snapshot['industries']['commercial_banks']['general_giving']['cycles']
        key = sorted(cycles)[0]
        cycles[key] += 1
        (self.base / 'b' / 'giving-anchors-v1.json').write_text(json.dumps(snapshot, indent=1), encoding='utf-8')
        failed = parity.compare(self.base / 'a1', self.base / 'b', mask, [(str(self.base / 'win').replace('\\', '/'), str(self.base / 'box').replace('\\', '/'))])
        self.assertFalse(failed['ok'])
        self.assertEqual(failed['files']['giving-anchors-v1.json']['where'], ['/industries/commercial_banks/general_giving/cycles/' + key])


if __name__ == '__main__':
    unittest.main()
