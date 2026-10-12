"""FIX 6: Graph follows transfers from the audited per-cycle book when the served one carries none.

The served Telecom Overview is an all-cycles build with no Transfer Paths sheet, so Graph said "no saved
network, and this workbook has no Transfer Paths sheet" for 2025-26. The audited per-cycle book of the same
industry (reports/<Industry> Industry Spending Overview <yyyy-yyyy>.xlsx) has the sheet, capped at 10,000
rows, and its .audit/transfer_paths.csv holds every row. Resolution order, each source said on screen:
  1. the served workbook's own Transfer Paths sheet;
  2. its pinned full source, when that sheet is capped;
  3. the audited book's audit CSV, else its pinned tidy store, else its own complete sheet.
Three fixtures (served without the sheet + audited CSV; a capped sheet + a full source; no source at all),
negative controls for each refusal, the cache dependency, and a read-only check on the real Telecom data.
"""
from decimal import Decimal
from pathlib import Path
import csv
import json
import re
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'app'))
import graph_data  # noqa: E402
from graph_data import graph_view  # noqa: E402
from report_insights import digest, stamp  # noqa: E402
from graph_fixture import add_graph_sheets  # noqa: E402

HEADER = 'cycle,source,destination,amount,record_count,first_date,last_date,root,depth,path,status'
BUILD = 'ab' * 32
# (cycle, source, destination, amount, root, depth, status): one edge repeated under a second root, a candidate
# gift that is never expanded, and a cycle edge back to the funded committee.
ROWS = [
    ('2026', 'C00000003', 'C00000009', 100, 'C00000003', 1, 'followed'),
    ('2026', 'C00000003', 'C00000009', 100, 'C00000004', 2, 'followed'),
    ('2026', 'C00000009', 'C00000011', 40, 'C00000003', 2, 'depth_limit'),
    ('2026', 'C00000003', 'C00000010', 5, 'C00000003', 1, 'candidate_committee'),
    ('2026', 'C00000011', 'C00000003', 1, 'C00000003', 3, 'cycle'),
]
UNIQUE = [('2026', 'C00000003', 'C00000009', Decimal(100)), ('2026', 'C00000009', 'C00000011', Decimal(40)),
          ('2026', 'C00000011', 'C00000003', Decimal(1))]
BASIS = ('Transfers for 2025-2026 were read from the audited 2025-2026 workbook (build abababab…), '
         'not the served workbook')


def csv_text(rows):
    return HEADER + '\n' + ''.join(
        f'{c},{s},{d},{a},1,2025-01-01,2026-01-01,{r},{depth},{r} > {d},{status}\n'
        for c, s, d, a, r, depth, status in rows)


def edges(found):
    return sorted((e['cycle'], e['source'], e['destination'], e['amount']) for e in found)


class Fixture(unittest.TestCase):
    """A served Overview with no Transfer Paths sheet, and a reports folder for the audited books."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        root = Path(self.tmp.name)
        self.reports = root / 'reports'
        self.reports.mkdir()
        (root / 'served').mkdir()
        self.path = root / 'served' / 'Fixture Industry Spending Overview.xlsx'
        self.served()
        patcher = patch.object(graph_data, 'audited_reports_directory', return_value=self.reports)
        patcher.start()
        self.addCleanup(patcher.stop)

    def served(self, sheet=None):
        book = openpyxl.Workbook()
        add_graph_sheets(book)
        book['Downstream Activity'].append(['2025-2026', 'C00000009', 'H0KY06104', 'BARR, GARLAND ANDY', 'R', 'Support', 106708])
        book['Downstream Activity']['A3'] = '101 source rows condensed to 5 committee-candidate-activity groups'
        if sheet is not None:
            add_sheet(book, *sheet)
        book.save(self.path)
        book.close()

    def audited(self, rows=ROWS, sheet=None, csv_file=True, label='2025-2026', build=BUILD):
        """reports/Fixture Industry Spending Overview <label>.xlsx and its .audit bundle, as audit CSVs are bound."""
        book = openpyxl.Workbook()
        book.active['A1'] = 'Source DB SHA-256: ' + build
        if sheet is not None:
            add_sheet(book, *sheet)
        path = self.reports / f'Fixture Industry Spending Overview {label}.xlsx'
        book.save(path)
        book.close()
        audit = path.with_suffix('.audit')
        audit.mkdir(exist_ok=True)
        files = {}
        if csv_file:
            (audit / 'transfer_paths.csv').write_text(csv_text(rows), encoding='utf-8-sig')
            files['transfer_paths.csv'] = digest(audit / 'transfer_paths.csv')
        (audit / 'audit-manifest.json').write_text(json.dumps(
            {'version': 1, 'workbook_name': path.name, 'workbook_sha256': digest(path), 'csv_sha256': files}), encoding='utf-8')
        return path

    def rebind(self, path):
        """Rewrite the audit manifest after the CSV changed, as a rebuild would."""
        audit = path.with_suffix('.audit')
        manifest = json.loads((audit / 'audit-manifest.json').read_text(encoding='utf-8'))
        manifest['csv_sha256'] = {'transfer_paths.csv': digest(audit / 'transfer_paths.csv')}
        (audit / 'audit-manifest.json').write_text(json.dumps(manifest), encoding='utf-8')

    def activity(self, **filters):
        with patch('downstream.inventory', return_value={}):
            return graph_view([self.path], {'mode': 'activity', **filters})


def add_sheet(book, note, rows=()):
    sheet = book.create_sheet('Transfer Paths')
    for row in [['Routes'], [], [note], [],
                ['Cycle', 'Source Committee', 'Destination Committee', 'Reported Net Transfer', 'Traversal Status'], *rows]:
        sheet.append(row)


class ServedWithoutTheSheet(Fixture):
    """(a) The served book has no Transfer Paths sheet; the audited book's CSV supplies the cycle."""

    def test_the_audit_csv_supplies_each_transfer_once_and_never_a_candidate_gift(self):
        self.audited()
        found = graph_data.report_transfers(stamp(self.path), audited=True)
        self.assertEqual(edges(found), UNIQUE)

    def test_the_basis_says_it_is_the_audited_book_not_the_served_one(self):
        self.audited()
        basis = graph_data.transfer_basis(stamp(self.path), audited=True)
        self.assertEqual([b['cycles'] for b in basis], [['2026']])
        self.assertEqual(basis[0]['text'], BASIS + ', from transfer_paths.csv (5 rows).')

    def test_a_capped_audited_sheet_is_named_with_both_counts(self):
        self.audited(sheet=('5 records; 2 shown. Full records are in the accompanying audit CSV files.',
                            [(2026, 'C00000003', 'C00000009', 100, 'followed')]))
        text = graph_data.transfer_basis(stamp(self.path), audited=True)[0]['text']
        self.assertEqual(text, BASIS + '. Its Transfer Paths sheet shows the first 2 of 5 rows; '
                                       'all 5 were read from transfer_paths.csv.')

    def test_the_cycle_is_covered_and_the_graph_traces_through_it(self):
        self.audited()
        value = self.activity(recipient='C00000003')
        self.assertTrue(value['candidate_available'], value['candidate_note'])
        trace = value['transfer_trace']
        self.assertTrue(trace['root_traced'])
        self.assertEqual(trace['committees'], [])
        self.assertNotIn(graph_data.NO_TRANSFER_PATHS, value['candidate_note'])
        # The route C00000003 > C00000009 reaches the spending the workbook records for C00000009.
        self.assertEqual(next(r['amount'] for r in value['candidates'] if r['id'] == 'H0KY06104'), 106708)
        self.assertEqual(sorted((e['source'], e['destination'], e['amount']) for e in value['downstream_edges']),
                         [('C00000003', 'C00000009', 100), ('C00000009', 'C00000011', 40), ('C00000011', 'C00000003', 1)])
        self.assertEqual(graph_data.not_traced_problems(value), [])

    def test_the_trace_note_and_the_method_note_both_say_what_was_used(self):
        self.audited()
        value = self.activity(recipient='C00000003')
        said = BASIS + ', from transfer_paths.csv (5 rows).'
        self.assertEqual(value['transfer_trace']['note'], said)
        self.assertTrue(value['candidate_note'].endswith(said), value['candidate_note'])
        self.assertNotIn('not traced', value['transfer_trace']['note'])

    def test_a_view_of_another_cycle_does_not_claim_this_cycles_source(self):
        # Two served cycles, an audited book for 2025-2026 only: 2023-2024 is still not traced, and says so.
        book = openpyxl.load_workbook(self.path)
        book.copy_worksheet(book['2025-2026']).title = '2023-2024'
        book.save(self.path)
        book.close()
        self.audited()
        self.assertEqual(graph_data.audited_sources(self.path).keys(), {'2026'})
        other = self.activity(recipient='C00000003', cycle='2024')
        self.assertNotIn('audited', other['transfer_trace']['note'])
        self.assertIn(graph_data.NO_TRANSFER_PATHS, other['transfer_trace']['note'])
        self.assertFalse(other['transfer_trace']['root_traced'])
        this = self.activity(recipient='C00000003', cycle='2026')
        self.assertIn(BASIS, this['transfer_trace']['note'])
        self.assertTrue(this['transfer_trace']['root_traced'])
        both = self.activity(recipient='C00000003')   # both cycles: one untraced, one read from the audited book
        self.assertTrue(both['transfer_trace']['note'].startswith('Transfers from this committee are not traced for 2024'))
        self.assertTrue(both['transfer_trace']['note'].endswith(BASIS + ', from transfer_paths.csv (5 rows).'))
        for answer in (other, this, both):
            self.assertEqual(graph_data.not_traced_problems(answer), [])

    def test_trace_coverage_counts_the_audited_cycle_as_covered(self):
        self.audited()
        covered = graph_data.trace_coverage([stamp(self.path)], [{'cycles': ['2026']}], {})
        self.assertEqual(covered, {'2026': (True, None)})

    def test_the_audit_csv_beats_the_audited_books_capped_sheet(self):
        # The sheet lists one edge; the CSV lists the three. The full source wins.
        self.audited(sheet=('5 records; 1 shown. Full records are in the accompanying audit CSV files.',
                            [(2026, 'C00000003', 'C00000009', 100, 'followed')]))
        self.assertEqual(len(graph_data.report_transfers(stamp(self.path), audited=True)), 3)

    def test_a_complete_audited_sheet_serves_when_it_has_no_csv_or_store(self):
        self.audited(csv_file=False, sheet=('2 records; 2 shown. Complete destination data.',
                                            [(2026, 'C00000003', 'C00000009', 100, 'followed'),
                                             (2026, 'C00000009', 'C00000011', 40, 'followed')]))
        self.assertEqual(edges(graph_data.report_transfers(stamp(self.path), audited=True)), UNIQUE[:2])
        self.assertEqual(graph_data.transfer_basis(stamp(self.path), audited=True)[0]['text'],
                         BASIS + ", from its Transfer Paths sheet (2 rows).")

    def test_a_capped_audited_sheet_with_no_full_source_is_not_a_source(self):
        self.audited(csv_file=False, sheet=('5 records; 1 shown. First records only.',
                                            [(2026, 'C00000003', 'C00000009', 100, 'followed')]))
        self.assertEqual(graph_data.audited_sources(self.path), {})
        self.assertEqual(graph_data.report_transfers(stamp(self.path), audited=True), [])


class TidyStore(Fixture):
    def store(self, rows=ROWS):
        path = Path(self.tmp.name) / 'tidy-20252026-fixture.sqlite'
        path.unlink(missing_ok=True)
        database = sqlite3.connect(path)
        database.execute('CREATE TABLE transfer_path_detail (cycle_label TEXT, cycle INTEGER, funded_committee TEXT, '
                         'source_committee TEXT, destination_committee TEXT, reported_net_transfer REAL, depth INTEGER, '
                         'representative_path TEXT, traversal_status TEXT, via_carryover INTEGER)')
        database.executemany('INSERT INTO transfer_path_detail VALUES (?,?,?,?,?,?,?,?,?,0)',
                             [('2025-2026', int(c), r, s, d, a, depth, f'{r} > {d}', status)
                              for c, s, d, a, r, depth, status in rows])
        database.commit()
        database.close()
        manifest = Path(self.tmp.name) / 'manifest-20252026-fixture.json'
        manifest.write_text(json.dumps({'stage_sha256': digest(path)}), encoding='utf-8')
        return path, manifest

    def test_the_pinned_tidy_store_serves_when_the_audited_book_has_no_csv(self):
        self.audited(csv_file=False)
        store, manifest = self.store()
        with patch.object(graph_data, 'audited_store', return_value=(store, manifest)):
            self.assertEqual(graph_data.audited_sources(self.path)['2026']['kind'], 'store')
            graph_data._pinned_detail_store.cache_clear()
            found = graph_data.report_transfers(stamp(self.path), audited=True)
            text = graph_data.transfer_basis(stamp(self.path), audited=True)[0]['text']
        self.assertEqual(edges(found), UNIQUE)
        self.assertEqual(text, BASIS + ', from tidy-20252026-fixture.sqlite (5 rows).')

    def test_the_csv_is_preferred_to_the_store(self):
        self.audited()
        store, manifest = self.store()
        with patch.object(graph_data, 'audited_store', return_value=(store, manifest)):
            self.assertEqual(graph_data.audited_sources(self.path)['2026']['kind'], 'csv')

    def test_negative_control_a_store_that_drifted_from_its_fingerprint_is_refused(self):
        self.audited(csv_file=False)
        store, manifest = self.store()
        manifest.write_text(json.dumps({'stage_sha256': '0' * 64}), encoding='utf-8')
        with patch.object(graph_data, 'audited_store', return_value=(store, manifest)):
            graph_data._pinned_detail_store.cache_clear()
            with self.assertRaisesRegex(ValueError, 'fingerprint'):
                graph_data.report_transfers(stamp(self.path), audited=True)


class CappedSheetWithAFullSource(Fixture):
    """(b) A capped Transfer Paths sheet: its pinned full source first, then the audited book."""

    NOTE = '3 records; 1 shown. First records only.'
    ONE = (2026, 'C00000003', 'C00000009', 100, 'followed')

    def capped(self):
        self.served(sheet=(self.NOTE, [self.ONE]))

    def pinned(self):
        payload, manifest = Path(self.tmp.name) / 'full.json', Path(self.tmp.name) / 'manifest.json'
        edge = {'cycle': 2026, 'source': 'C00000003', 'destination': 'C00000009', 'amount': 100, 'status': 'followed'}
        payload.write_text(json.dumps({'transfer_paths': [
            edge, {**edge, 'root': 'another'}, {**edge, 'destination': 'C00000010', 'status': 'candidate_committee'}]}))
        manifest.write_text(json.dumps({'database_sha256': 'pinned', 'data_sha256': digest(payload)}))
        return payload, manifest

    def test_the_pinned_full_source_is_named_with_the_sheets_cap(self):
        self.capped()
        with patch.object(graph_data, 'transfer_source', return_value=self.pinned()), \
                patch('contribution_data._workbook_source', return_value='pinned'):
            found = graph_data.report_transfers(stamp(self.path), audited=True)
            basis = graph_data.transfer_basis(stamp(self.path), audited=True)
        self.assertEqual(edges(found), [('2026', 'C00000003', 'C00000009', Decimal(100))])
        # The payload's raw row count is not known (native_scan collapses repeats first): no "all N" claimed.
        self.assertEqual(basis, [{'cycles': ['2026'], 'text': 'The Transfer Paths sheet shows the first 1 of 3 rows; '
                                                              'the full list was read from full.json.'}])

    def test_the_pinned_source_is_preferred_to_an_audited_book(self):
        self.capped()
        self.audited()
        with patch.object(graph_data, 'transfer_source', return_value=self.pinned()), \
                patch('contribution_data._workbook_source', return_value='pinned'):
            self.assertEqual(len(graph_data.report_transfers(stamp(self.path), audited=True)), 1)
            self.assertNotIn('audited', graph_data.transfer_basis(stamp(self.path), audited=True)[0]['text'])

    def test_with_no_pinned_source_the_audited_csv_replaces_the_partial_display(self):
        self.capped()
        self.audited()
        found = graph_data.report_transfers(stamp(self.path), audited=True)
        self.assertEqual(edges(found), UNIQUE)
        self.assertTrue(graph_data.transfer_basis(stamp(self.path), audited=True)[0]['text'].startswith(BASIS))

    def test_negative_control_a_partial_display_with_no_full_source_is_still_refused(self):
        self.capped()
        with self.assertRaisesRegex(ValueError, 'partial display'):
            graph_data.report_transfers(stamp(self.path), audited=True)

    def test_negative_control_a_complete_served_sheet_is_the_served_books_own_and_says_nothing(self):
        self.served(sheet=('1 records; 1 shown. Complete destination data.', [self.ONE]))
        self.audited()   # present, and ignored: the served sheet comes first
        self.assertEqual(edges(graph_data.report_transfers(stamp(self.path), audited=True)),
                         [('2026', 'C00000003', 'C00000009', Decimal(100))])
        self.assertEqual(graph_data.transfer_basis(stamp(self.path), audited=True), [])
        value = self.activity(recipient='C00000003')
        self.assertEqual(value['transfer_trace']['note'], '')


class NoSource(Fixture):
    """(c) Nothing supplies the cycle: Graph still says there is no Transfer Paths sheet."""

    def test_no_audited_book_keeps_the_no_transfer_paths_sentence(self):
        self.assertEqual(graph_data.audited_sources(self.path), {})
        self.assertEqual(graph_data.report_transfers(stamp(self.path), audited=True), [])
        self.assertEqual(graph_data.transfer_basis(stamp(self.path), audited=True), [])
        value = self.activity(recipient='C00000003')
        self.assertFalse(value['transfer_trace']['root_traced'])
        self.assertEqual(value['transfer_trace']['note'], 'Transfers from this committee are not traced: '
                                                         + graph_data.NO_TRANSFER_PATHS + '.')
        self.assertEqual(graph_data.trace_coverage([stamp(self.path)], [{'cycles': ['2026']}], {}),
                         {'2026': (False, None)})

    def test_an_audited_book_of_another_cycle_does_not_cover_this_one(self):
        self.audited(label='2023-2024', rows=[('2024', 'C00000003', 'C00000009', 100, 'C00000003', 1, 'followed')])
        self.assertEqual(graph_data.audited_sources(self.path), {})
        self.assertEqual(self.activity(recipient='C00000003')['transfer_trace']['note'],
                         'Transfers from this committee are not traced: ' + graph_data.NO_TRANSFER_PATHS + '.')

    def test_a_workbook_that_is_not_a_published_overview_has_no_audited_book(self):
        other = self.path.with_name('Giving.xlsx')
        self.path.rename(other)
        self.audited()
        self.assertEqual(graph_data.audited_sources(other), {})

    def test_an_audited_book_with_neither_csv_store_nor_sheet_is_no_source(self):
        self.audited(csv_file=False)
        self.assertEqual(graph_data.audited_sources(self.path), {})


class Refusals(Fixture):
    """Every refusal reads as the Graph's unavailable note, never as a partial network."""

    def refused(self, pattern):
        with self.assertRaisesRegex(ValueError, pattern):
            graph_data.report_transfers(stamp(self.path), audited=True)
        value = self.activity(recipient='C00000003')
        self.assertFalse(value['candidate_available'])
        self.assertRegex(value['candidate_note'], pattern)

    def test_a_csv_that_does_not_match_its_manifest_is_refused(self):
        path = self.audited()
        with (path.with_suffix('.audit') / 'transfer_paths.csv').open('a', encoding='utf-8') as stream:
            stream.write('2026,C00000003,C00000012,9,1,x,y,C00000003,1,p,followed\n')
        self.refused('audited transfer paths file is missing or changed')

    def test_a_csv_with_no_audit_manifest_is_refused(self):
        path = self.audited()
        (path.with_suffix('.audit') / 'audit-manifest.json').unlink()
        self.refused('no audit manifest')

    def test_a_manifest_bound_to_another_workbook_is_refused(self):
        path = self.audited()
        manifest = path.with_suffix('.audit') / 'audit-manifest.json'
        manifest.write_text(json.dumps({**json.loads(manifest.read_text()), 'workbook_sha256': '0' * 64}), encoding='utf-8')
        self.refused('do not match their audited workbook')

    def test_one_transfer_with_two_amounts_is_refused(self):
        self.audited(rows=[*ROWS, ('2026', 'C00000003', 'C00000009', 101, 'C00000005', 1, 'followed')])
        self.refused('disagree on one transfer amount')

    def test_a_candidate_gift_with_another_amount_is_not_a_disagreement(self):
        self.audited(rows=[*ROWS, ('2026', 'C00000003', 'C00000009', 999, 'C00000005', 1, 'candidate_committee')])
        self.assertEqual(edges(graph_data.report_transfers(stamp(self.path), audited=True)), UNIQUE)

    def test_an_invalid_committee_identifier_is_refused(self):
        self.audited(rows=[*ROWS, ('2026', 'C00000003', 'not-an-id', 5, 'C00000003', 1, 'followed')])
        self.refused('invalid committee or cycle identifiers')

    def test_a_csv_that_carries_another_cycle_is_refused(self):
        self.audited(rows=[*ROWS, ('2024', 'C00000003', 'C00000012', 5, 'C00000003', 1, 'followed')])
        self.refused('include another cycle')


class AuditedBooksAreOptIn(Fixture):
    """P1-1: only the Graph view reads the audited per-cycle books. report_transfers(stamp) is what the flow
    (flow_data.industry_flow) and the drawer sidecar warmers call, and it stays the served build's transfers."""

    def test_a_served_book_without_the_sheet_has_no_transfers_unless_the_graph_asks(self):
        self.audited()
        self.assertEqual(graph_data.report_transfers(stamp(self.path)), [])
        self.assertEqual(graph_data.transfer_basis(stamp(self.path)), [])
        self.assertEqual(edges(graph_data.report_transfers(stamp(self.path), audited=True)), UNIQUE)
        self.assertEqual(graph_data.trace_coverage([stamp(self.path)], [{'cycles': ['2026']}], {}),
                         {'2026': (True, None)})

    def test_a_capped_sheet_with_no_pinned_source_stays_refused_for_everyone_but_the_graph(self):
        self.served(sheet=('3 records; 1 shown. First records only.', [(2026, 'C00000003', 'C00000009', 100, 'followed')]))
        self.audited()
        with self.assertRaisesRegex(ValueError, 'partial display'):
            graph_data.report_transfers(stamp(self.path))
        self.assertEqual(edges(graph_data.report_transfers(stamp(self.path), audited=True)), UNIQUE)

    def test_the_two_variants_are_cached_apart_in_either_order(self):
        self.audited()
        for order in ((True, False), (False, True)):
            graph_data.report_transfers.cache_clear()
            found = {flag: edges(graph_data.report_transfers(stamp(self.path), audited=flag)) for flag in order}
            self.assertEqual(found, {True: UNIQUE, False: []})

    def test_an_audited_file_changing_does_not_replace_the_default_variants_key(self):
        path = self.audited()
        graph_data.report_transfers(stamp(self.path))
        before = graph_data.report_transfers.cache_info().misses
        (path.with_suffix('.audit') / 'transfer_paths.csv').write_text(csv_text(ROWS[:1]), encoding='utf-8-sig')
        self.rebind(path)
        graph_data.report_transfers(stamp(self.path))
        self.assertEqual(graph_data.report_transfers.cache_info().misses, before)

    def test_the_flow_and_the_sidecar_warmers_do_not_ask_for_the_audited_books(self):
        app = Path(graph_data.__file__).parent
        for name in ('flow_data.py', 'drawer_sidecars.py'):
            calls = re.findall(r'report_transfers\([^)]*\)', (app / name).read_text(encoding='utf-8'))
            self.assertTrue(calls, name)
            self.assertFalse([c for c in calls if 'audited' in c], name)


class CacheDependencies(Fixture):
    def test_every_audited_file_is_a_dependency_even_before_it_exists(self):
        book = self.reports / 'Fixture Industry Spending Overview 2025-2026.xlsx'
        wanted = {book, book.with_suffix('.audit') / 'audit-manifest.json', book.with_suffix('.audit') / 'transfer_paths.csv'}
        self.assertFalse(book.exists())
        self.assertTrue(wanted <= set(map(Path, graph_data.transfer_files(self.path))))
        self.audited()
        self.assertTrue(wanted <= set(map(Path, graph_data.transfer_files(self.path))))

    def test_a_changed_csv_replaces_the_cached_transfers_and_the_cached_view(self):
        path = self.audited()
        first = self.activity(recipient='C00000003')
        self.assertEqual(sorted(e['amount'] for e in first['downstream_edges']), [1, 40, 100])
        (path.with_suffix('.audit') / 'transfer_paths.csv').write_text(
            csv_text([('2026', 'C00000003', 'C00000009', 250, 'C00000003', 1, 'followed')]), encoding='utf-8-sig')
        self.rebind(path)
        second = self.activity(recipient='C00000003')
        self.assertEqual([e['amount'] for e in second['downstream_edges']], [250])
        self.assertEqual(len(graph_data.report_transfers(stamp(self.path), audited=True)), 1)

    def test_an_audited_book_that_appears_later_replaces_a_cached_no_source_view(self):
        first = self.activity(recipient='C00000003')
        self.assertFalse(first['transfer_trace']['root_traced'])
        self.audited()
        second = self.activity(recipient='C00000003')
        self.assertTrue(second['transfer_trace']['root_traced'])
        self.assertIn(BASIS, second['candidate_note'])


REAL_SERVED = Path('C:/Users/pgoff/OneDrive/Desktop/Mentat Complete Files/Telecom Industry Spending Overview.xlsx')
REAL_REPORTS = Path(__file__).resolve().parents[2] / 'reports'
REAL_CSV = REAL_REPORTS / 'Telecom Industry Spending Overview 2025-2026.audit' / 'transfer_paths.csv'


@unittest.skipUnless(REAL_SERVED.is_file() and REAL_CSV.is_file(), 'the real Telecom books are not on this machine')
class RealTelecom(unittest.TestCase):
    """Read-only: the served all-cycles book has no Transfer Paths sheet; 2026 comes from the audited book's CSV."""

    def test_2026_is_covered_and_resolution_returns_the_csvs_unique_edges(self):
        self.assertEqual(graph_data.audited_reports_directory().resolve(), REAL_REPORTS.resolve())
        self.assertFalse(graph_data.has_transfer_paths(stamp(REAL_SERVED)))
        sources = graph_data.audited_sources(REAL_SERVED)
        self.assertIn('2026', sources)
        self.assertEqual((sources['2026']['kind'], sources['2026']['path'].resolve()), ('csv', REAL_CSV.resolve()))
        covered = graph_data.trace_coverage([stamp(REAL_SERVED)], [{'cycles': ['2026']}], {})
        self.assertTrue(covered['2026'][0])
        found, sentence = graph_data.audited_transfers(sources['2026'])
        unique, rows = {}, 0
        with REAL_CSV.open(encoding='utf-8-sig', newline='') as stream:
            for row in csv.DictReader(stream):
                rows += 1
                if row['status'] == 'candidate_committee':
                    continue
                unique[row['cycle'], row['source'], row['destination']] = Decimal(row['amount'])
        self.assertEqual(rows, 230075)
        self.assertEqual(len(unique), 4515)
        self.assertEqual(len(found), len(unique))
        self.assertEqual({(e['cycle'], e['source'], e['destination']): e['amount'] for e in found}, unique)
        self.assertRegex(sentence, r'^Transfers for 2025-2026 were read from the audited 2025-2026 workbook '
                                   r'\(build 5bd16e96…\), not the served workbook\. Its Transfer Paths sheet shows '
                                   r'the first 10,000 of 230,075 rows; all 230,075 were read from transfer_paths\.csv\.$')


if __name__ == '__main__':
    unittest.main()
