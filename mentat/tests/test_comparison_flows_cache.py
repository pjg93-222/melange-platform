"""/api/comparison-flows through response_cache: the same bytes as the uncached build, from memory,
from disk and in a restarted process; issues, moving reports, changed content and changed code are
never answered from a stored entry.

A child interpreter of this file stands in for a restarted console (output to real files, waited
on with a timeout and tree-killed: brain DE-026).
"""
import gzip
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

APP = Path(__file__).resolve().parents[1] / 'app'
sys.path.insert(0, str(APP))
if os.environ.get('MENTAT_CACHE_TEST_OVERLAY'):
    sys.path.insert(0, os.environ['MENTAT_CACHE_TEST_OVERLAY'])
import response_cache as rc  # noqa: E402
import comparison_data  # noqa: E402
import report_insights  # noqa: E402
from graph_fixture import add_graph_sheets  # noqa: E402

NO_WINDOW = getattr(subprocess, 'CREATE_NO_WINDOW', 0)
NAMESPACE = comparison_data.FLOWS_NAMESPACE


def kill_tree(process):
    if process.poll() is None:
        killer = subprocess.Popen(['taskkill', '/PID', str(process.pid), '/T', '/F'], stdout=subprocess.DEVNULL,
                                  stderr=subprocess.DEVNULL, creationflags=NO_WINDOW)
        try:
            killer.wait(timeout=15)
        except subprocess.TimeoutExpired:
            killer.kill()
    try:
        process.wait(timeout=15)
    except subprocess.TimeoutExpired:
        process.kill()


class FlowsCacheTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)
        self.root = self.dir / 'response-cache'
        saved = rc.ROOT
        self.addCleanup(setattr, rc, 'ROOT', saved)
        self.addCleanup(rc.clear_memory)
        rc.configure(self.root)
        rc.clear_memory()
        self.path = self.dir / 'flows.xlsx'
        book = openpyxl.Workbook()
        add_graph_sheets(book)
        book['Downstream Activity']['A3'] = '101 source rows condensed to 5 committee-candidate-activity groups'
        book['Downstream Activity'].append(['2025-2026', 'C00000003', 'H0CA00002', 'Candidate Two', 'R', 'Support', 250])
        book.save(self.path)
        book.close()

    def stored(self):
        folder = self.root / NAMESPACE
        return sorted(p.name for p in folder.glob('*.json')) if folder.exists() else []

    def edit(self, change):
        book = openpyxl.load_workbook(self.path)
        change(book)
        book.save(self.path)
        book.close()

    def test_every_path_answers_the_bytes_of_the_uncached_build(self):
        fresh = rc.json_body(comparison_data._flows([self.path], 'Sample'))
        built = comparison_data.flows_entry([self.path], 'Sample')
        self.assertEqual((built.hit, built.body), ('built', fresh))
        memory = comparison_data.flows_entry([self.path], 'Sample')
        self.assertEqual((memory.hit, memory.body), ('memory', fresh))
        self.assertEqual(rc.json_body(comparison_data.flows([self.path], 'Sample')), fresh)
        rc.clear_memory()  # what a restart loses
        disk = comparison_data.flows_entry([self.path], 'Sample')
        self.assertEqual((disk.hit, disk.body), ('disk', fresh))
        self.assertEqual(disk.gzip, gzip.compress(fresh, compresslevel=1, mtime=0))  # Handler.send's own gzip
        rc.clear_memory()
        self.assertEqual(rc.json_body(comparison_data.flows([self.path], 'Sample')), fresh)
        self.assertEqual(len(self.stored()), 1)

    def test_each_caller_owns_its_object(self):
        first = comparison_data.flows([self.path], 'Sample')
        first['rows'].clear()
        self.assertTrue(comparison_data.flows([self.path], 'Sample')['rows'])

    def test_a_restarted_process_serves_flows_from_disk(self):
        built = comparison_data.flows_entry([self.path], 'Sample')
        out = self.dir / 'child.json'
        log = open(self.dir / 'child.log', 'wb')
        self.addCleanup(log.close)
        process = subprocess.Popen([sys.executable, str(Path(__file__).resolve()), 'child', str(self.root),
                                    str(self.path), str(out)], stdout=log, stderr=subprocess.STDOUT,
                                   cwd=str(Path(__file__).resolve().parent), creationflags=NO_WINDOW)
        self.addCleanup(kill_tree, process)
        process.wait(timeout=180)
        self.assertTrue(out.exists(), (self.dir / 'child.log').read_text(errors='replace'))
        result = json.loads(out.read_text())
        self.assertEqual((result['hit'], result['computed']), ('disk', 0))
        self.assertEqual(result['sha256'], hashlib.sha256(built.body).hexdigest())
        self.assertEqual(result['object_sha256'], hashlib.sha256(built.body).hexdigest())

    def test_answers_with_issues_are_recomputed_and_never_stored(self):
        self.edit(lambda book: book['Downstream Activity'].cell(3, 1, 'condensed to 500 groups'))
        with patch.object(comparison_data, '_flows', wraps=comparison_data._flows) as build:
            first = comparison_data.flows([self.path], 'Sample')
            comparison_data.flows([self.path], 'Sample')
        self.assertTrue(first['issues'])
        self.assertEqual(build.call_count, 2)
        self.assertEqual(self.stored(), [])

    def test_a_report_that_moves_during_the_build_is_answered_not_stored(self):
        real = comparison_data._flows

        def moving(paths, industry):
            info = self.path.stat()
            os.utime(self.path, ns=(info.st_atime_ns, info.st_mtime_ns + 5_000_000_000))
            return real(paths, industry)
        with patch.object(comparison_data, '_flows', side_effect=moving):
            entry = comparison_data.flows_entry([self.path], 'Sample')
        self.assertEqual(entry.hit, 'built')
        self.assertEqual(json.loads(entry.body)['industry'], 'Sample')
        self.assertEqual(self.stored(), [])
        self.assertEqual(comparison_data.flows_entry([self.path], 'Sample').hit, 'built')  # the new stamp, once
        self.assertEqual(len(self.stored()), 1)

    def test_a_changed_report_is_a_new_key(self):
        before = comparison_data.flows([self.path], 'Sample')
        self.edit(lambda book: book['Downstream Activity'].append(['2025-2026', 'C00000004', 'H0CA00002', 'Candidate Two',
                                                                    'R', 'Oppose', 5]))
        self.edit(lambda book: book['Downstream Activity'].cell(3, 1, '102 source rows condensed to 6 committee-candidate-activity groups'))
        after = comparison_data.flows_entry([self.path], 'Sample')
        self.assertEqual(after.hit, 'built')
        self.assertNotEqual(json.loads(after.body)['rows'], before['rows'])

    def test_content_hash_is_part_of_the_key(self):
        comparison_data.flows_entry([self.path], 'Sample')
        rc.clear_memory()
        report_insights.digest_stamp.cache_clear()
        with patch.object(comparison_data, 'digest_stamp', return_value='0' * 64):
            entry = comparison_data.flows_entry([self.path], 'Sample')
        self.assertEqual(entry.hit, 'built')  # same stamps, other content: never the stored answer

    def test_code_that_changed_under_the_process_keeps_flows_off_disk(self):
        with patch.object(rc, 'code_identity', return_value=None):
            first = comparison_data.flows_entry([self.path], 'Sample')
            second = comparison_data.flows_entry([self.path], 'Sample')
        self.assertEqual((first.hit, second.hit), ('built', 'memory'))  # this process's own code
        self.assertEqual(self.stored(), [])

    def test_overlapping_reports_are_still_refused(self):
        comparison_data.flows([self.path], 'Sample')  # [A] is stored; [A, A] is another question
        with self.assertRaisesRegex(ValueError, 'overlap'):
            comparison_data.flows([self.path, self.path], 'Sample')
        self.assertEqual(len(self.stored()), 1)

    def test_the_industry_label_is_part_of_the_key(self):
        comparison_data.flows([self.path], 'Sample')
        other = comparison_data.flows_entry([self.path], 'Other label')
        self.assertEqual((other.hit, json.loads(other.body)['industry']), ('built', 'Other label'))

    def test_warm_hook_builds_once_and_reports(self):
        first = comparison_data.warm_flows([self.path], 'Sample')
        rc.clear_memory()
        second = comparison_data.warm_flows([self.path], 'Sample')
        self.assertEqual((first['hit'], second['hit']), ('built', 'disk'))
        self.assertEqual(first['etag'], second['etag'])
        self.assertGreater(first['body_bytes'], first['gzip_bytes'])


def child(args):
    root, path, out = map(Path, args)
    rc.configure(root)
    computed = []
    real = comparison_data._flows

    def counting(paths, industry):
        computed.append(1)
        return real(paths, industry)
    comparison_data._flows = counting
    entry = comparison_data.flows_entry([path], 'Sample')
    value = comparison_data.flows([path], 'Sample')
    out.write_text(json.dumps({'hit': entry.hit, 'computed': len(computed),
                               'sha256': hashlib.sha256(entry.body).hexdigest(),
                               'object_sha256': hashlib.sha256(rc.json_body(value)).hexdigest()}))


if __name__ == '__main__':
    if len(sys.argv) > 1 and sys.argv[1] == 'child':
        child(sys.argv[2:])
    else:
        unittest.main()
