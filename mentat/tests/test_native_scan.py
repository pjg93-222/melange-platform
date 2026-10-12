"""Native decoding must leave downstream accounting unchanged."""
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import os
import json
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'app'))
import native_scan
import graph_data
import contribution_workbook_guard as workbook_guard
import contribution_data
from report_insights import stamp
from graph_fixture import add_graph_sheets


@unittest.skipUnless(native_scan.executable(), 'Build with tools/build_native.ps1 first')
class NativeScanTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / 'report.xlsx'
        book = openpyxl.Workbook()
        add_graph_sheets(book)
        book.save(self.path)
        book.close()
        self.addCleanup(native_scan.shutdown)

    def both(self, reader):
        results = []
        for engine in ('off', 'on'):
            reader.cache_clear()
            with patch.dict(os.environ, MENTAT_NATIVE=engine):
                results.append(reader(stamp(self.path)))
        self.assertEqual(*results)
        return results[1]

    def test_reference_parity_and_source_bytes(self):
        before = self.path.read_bytes()
        self.both(graph_data.report_facts)
        self.both(graph_data.report_activity)
        self.assertEqual(before, self.path.read_bytes())

    def test_direction_groups_match_individual_views(self):
        body = graph_data.graph_view([self.path], {'mode':'activity'})
        for direction in graph_data.ACTIVITY_DIRECTIONS:
            one = graph_data.graph_view([self.path], {'mode':'activity','direction':direction})
            self.assertEqual(body['candidates_by_direction'][direction], one['candidates'])

    def test_view_cache_does_not_hide_duplicate_report_inputs(self):
        graph_data.graph_view([self.path],{})
        with self.assertRaisesRegex(ValueError,'overlap'):
            graph_data.graph_view([self.path,self.path],{})

    def test_signed_outlier_and_missing_amount_are_not_sanitized(self):
        book = openpyxl.load_workbook(self.path)
        sheet = book['Downstream Activity']
        amount_col = next(c.column for c in sheet[5] if graph_data.key(c.value)=='net_included_amount')
        sheet.cell(6, amount_col, -9876543210.25)
        book.save(self.path)
        book.close()
        self.assertTrue(any(r['amount'] < -9_000_000_000 for r in self.both(graph_data.report_activity)))
        book = openpyxl.load_workbook(self.path)
        book['Downstream Activity'].cell(6,amount_col).value = None
        book.save(self.path)
        book.close()
        for engine in ('off','on'):
            graph_data.report_activity.cache_clear()
            with patch.dict(os.environ, MENTAT_NATIVE=engine), self.assertRaises(ValueError):
                graph_data.report_activity(stamp(self.path))

    def test_projection_unicode_formulas_rich_strings_and_limits(self):
        path = Path(self.tmp.name) / 'strings.xlsx'
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('xl/workbook.xml','<workbook xmlns:r="r"><sheets><sheet name="Test" r:id="s1"/></sheets></workbook>')
            archive.writestr('xl/_rels/workbook.xml.rels','<Relationships><Relationship Id="s1" Target="worksheets/sheet1.xml"/></Relationships>')
            archive.writestr('xl/sharedStrings.xml','<sst><si><r><t>Zoë &amp; </t></r><r><t>東京</t></r><rPh><t>phonetic</t></rPh></si></sst>')
            archive.writestr('xl/worksheets/sheet1.xml','''<worksheet><sheetData>
              <row r="1"><c r="A1" t="inlineStr"><is><t>Name</t></is></c><c r="B1" t="inlineStr"><is><t>Amount</t></is></c></row>
              <row r="2"><c r="A2" t="s"><v>0</v></c><c r="B2"><f>1-2.5</f><v>-1.5</v></c></row>
              <row r="3"><c r="A3" t="inlineStr"><is><t>missing amount</t></is></c><c r="B3"/></row>
              </sheetData></worksheet>''')
        result = native_scan.tables(stamp(path),[{'name':'Test','header':1}])
        self.assertEqual(result['tables']['Test']['rows'],[['Zoë & 東京','-1.5'],['missing amount',None]])
        projected = native_scan.tables(stamp(path),[{'name':'Test','header':1,'columns':['amount']}])
        self.assertEqual(projected['tables']['Test']['rows'],[['-1.5']])
        with self.assertRaisesRegex(ValueError,'limit'):
            native_scan.tables(stamp(path),[{'name':'Test','header':1}],max_rows=1)

    def test_worker_reuse_and_errors_not_empty_results(self):
        request = [{'name':'Downstream Activity','header':5}]
        for _ in range(3):
            self.assertTrue(native_scan.tables(stamp(self.path),request)['tables'])
        self.assertEqual(len(native_scan._all),1)
        bad = Path(self.tmp.name) / 'bad.xlsx'
        bad.write_bytes(b'not a ZIP')
        with self.assertRaises(ValueError):native_scan.tables(stamp(bad),request)
        self.assertEqual(len(native_scan._all),0)

    def test_bounded_overview_scan_and_full_transfer_payload(self):
        result = native_scan.tables(stamp(self.path), [{'name':'Downstream Activity','header':5,'end_row':6}])
        self.assertEqual(len(result['tables']['Downstream Activity']['rows']), 1)
        payload = Path(self.tmp.name) / 'source.json'
        edge = dict(cycle=2024,source='C00000001',destination='C00000002',amount=125.25,status='expanded')
        payload.write_text(json.dumps({'ignored':[{'large':'unneeded'}], 'transfer_paths':[edge,{**edge,'root':'different'},
            {**edge,'destination':'C00000003','status':'candidate_committee'}]}),encoding='utf-8')
        result = native_scan.tables(stamp(payload), [], payload_table='transfer_paths')
        self.assertEqual(result['tables']['transfer_paths']['rows'], [['2024','C00000001','C00000002','125.25']])
        payload.write_text(json.dumps({'transfer_paths':[edge,{**edge,'amount':99}]}),encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'Conflicting'):
            native_scan.tables(stamp(payload), [], payload_table='transfer_paths')

    def test_workbook_guard_adapts_stamp_order_and_matches_reference(self):
        book = openpyxl.load_workbook(self.path)
        overview = book.create_sheet('Spending Overview')
        overview.append([])
        for column, value in enumerate(['Cycle','Scope','Total Giving'],1): overview.cell(5,column,value)
        for column, value in enumerate(['2023-2024','Total Giving',123.45],1): overview.cell(6,column,value)
        book.save(self.path); book.close()
        totals=[]
        for mode in ('off','on'):
            workbook_guard.cycle_totals.cache_clear()
            with patch.dict(os.environ,MENTAT_NATIVE=mode):
                totals.append(workbook_guard.cycle_totals(contribution_data._stamp(self.path)))
        self.assertEqual(totals,[{'2023-2024':12345}]*2)

    def test_parallel_reads_and_absent_binary_fallback(self):
        with ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(lambda _: native_scan.tables(stamp(self.path),[{'name':'Downstream Activity','header':5}]),range(4)))
        native = [r for r in results if r is not None]
        self.assertTrue(native)
        self.assertTrue(all(r==native[0] for r in native))
        with patch.dict(os.environ,MENTAT_NATIVE_SCANNER=str(Path(self.tmp.name)/'missing.exe')):
            self.assertIsNone(native_scan.tables(stamp(self.path),[]))


if __name__ == '__main__':unittest.main()
