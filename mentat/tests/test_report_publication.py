"""Owned generation publication and read boundaries, using isolated fixtures only."""
from concurrent.futures import ThreadPoolExecutor
import hashlib
import http.client
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import threading
import types
import unittest
from unittest.mock import patch

import openpyxl

PROJECT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(PROJECT / 'app'), str(PROJECT / 'tools'), str(PROJECT / 'tests')]
import report_publication as publication
import contribution_data as contributions
import graph_data
import report_insights
import comparison_data
import response_cache
import source_memo
import server as app
from graph_fixture import add_graph_sheets, GIVING


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value), encoding='utf-8')


class PublicationFixture(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix='mentat-publication-test-')
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.reports, self.stage = self.root / 'reports', self.root / 'stage'
        self.reports.mkdir(); self.stage.mkdir()
        self.environment = patch.dict(os.environ, {'MENTAT_NATIVE': '0', 'MENTAT_DRAWER_SIDECARS': '0',
            'MENTAT_DATA_ROOT': '', 'MENTAT_REPORT_CACHE_ROOT': '', 'MENTAT_RECEIPT_ROOTS': ''})
        self.environment.start(); self.addCleanup(self.environment.stop)
        project = patch.object(contributions, 'PROJECT', self.root)
        project.start(); self.addCleanup(project.stop)
        source_memo.uninstall(); self.addCleanup(source_memo.uninstall)
        response_cache.clear_memory()
        saved = response_cache.ROOT; response_cache.ROOT = None
        self.addCleanup(setattr, response_cache, 'ROOT', saved)
        self.addCleanup(response_cache.clear_memory)
        self.database = self.root / 'source-marker.db'
        self.database.write_bytes(b'synthetic discovery marker only; never queried as a financial database')
        self.source_hash = contributions._hash(self.database)
        self.overview = self.stage / 'Fixture Industry Spending Overview.xlsx'
        self.data = self.stage / 'Fixture Industry Spending Data.xlsx'
        self.target = self.reports / self.overview.name
        self.paths = publication.locations(self.target)
        self.make_books('first')
        cache = self.root / '.cache/industry-report/fixture'
        cache.mkdir(parents=True)
        payload = cache / ('industry_data-' + self.source_hash[:16] + '.json')
        write_json(payload, {'metadata': {'industry': 'Fixture', 'database_path': str(self.database),
            'expected_cycle_totals': {'2025-2026': sum(row[4] for row in GIVING)}}})
        write_json(cache / ('manifest-' + self.source_hash[:16] + '.json'),
                   {'database_sha256': self.source_hash, 'data_sha256': contributions._hash(payload)})
        self.console = object.__new__(app.Console)
        self.console.report_directory = self.reports
        self.console.token = ''

    def make_books(self, label):
        book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Spending Overview'
        sheet['A1'] = label; sheet['A3'] = 'Source DB SHA-256: ' + self.source_hash
        sheet.append([]); sheet['A5'] = 'Cycle'; sheet['B5'] = 'Scope'; sheet['C5'] = 'Total Giving'
        sheet.append(['2025-2026', 'Total Giving', sum(row[4] for row in GIVING)])
        add_graph_sheets(book); book.save(self.overview); book.close()
        book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Data'
        sheet.append(['Generation', 'Value']); sheet.append([label, 12])
        sheet['A3'] = 'Source DB SHA-256: ' + self.source_hash
        book.save(self.data); book.close()

    def publish(self, **kwargs):
        return publication.publish(self.overview, self.data, self.reports, **kwargs)

    def pending(self):
        write_json(self.paths['intent'], {'schema_version': 1, 'generation': 'f' * 32})


class GenerationProtocol(PublicationFixture):
    def test_fresh_server_import_and_real_publication_cli_do_not_shadow_each_other(self):
        result = subprocess.run([sys.executable, str(PROJECT / 'tools/publish_report_generation.py'),
            '--overview', str(self.overview), '--data', str(self.data)],
            env=dict(os.environ, MENTAT_REPORT_DIRECTORY=str(self.reports)), capture_output=True, text=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(publication.capture(self.target)['generation'], json.loads(result.stdout)['generation'])
        result = subprocess.run([sys.executable, '-c', 'import sys;sys.path.insert(0,sys.argv[1]);import server;print(server.report_publication.VERSION)',
            str(PROJECT / 'app')], env=dict(os.environ), capture_output=True, text=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stdout.strip(), '1')

    def test_legacy_and_committed_are_explicit_and_loss_cannot_downgrade(self):
        shutil.copyfile(self.overview, self.target)
        self.assertEqual(publication.capture(self.target)['state'], 'legacy')
        manifest = self.publish()
        self.assertEqual(publication.capture(self.target)['generation'], manifest['generation'])
        self.assertEqual(publication.capture(self.paths['data'])['state'], 'committed')
        self.assertTrue(self.paths['required'].is_file())
        self.assertFalse(self.paths['intent'].exists())
        self.paths['manifest'].unlink()
        with self.assertRaisesRegex(ValueError, 'publication is incomplete'):
            publication.capture(self.target)

    def test_concurrent_readers_refuse_split_pair_and_second_writer_cannot_interleave(self):
        self.publish(); self.make_books('second')
        at_split, resume = threading.Event(), threading.Event()
        def checkpoint(phase, manifest):
            if phase == 'overview':
                at_split.set()
                if not resume.wait(30): raise RuntimeError('fixture barrier timed out')
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(self.publish, checkpoint=checkpoint)
            if not at_split.wait(30):
                if future.done(): future.result()  # surface writer failure instead of a misleading barrier timeout
                self.fail("Publisher did not reach the fixture barrier within 30 seconds.")
            try:
                for _ in range(10):
                    with self.assertRaises(ValueError): publication.capture(self.target)
                self.assertEqual(self.console.published_reports(), [])
                with self.assertRaisesRegex(ValueError, 'Another publisher owns'):
                    self.publish(wait_seconds=0)
            finally:
                resume.set()
            manifest = future.result(timeout=10)
        self.assertEqual(publication.capture(self.target)['generation'], manifest['generation'])

    def test_process_death_after_first_replacement_is_recoverable_only_by_new_owner(self):
        old = self.publish(); self.make_books('after process death')
        child = self.root / 'crash_publisher.py'
        child.write_text("import os,sys\nfrom pathlib import Path\nsys.path.insert(0,sys.argv[1])\nimport report_publication as p\np.publish(Path(sys.argv[2]),Path(sys.argv[3]),Path(sys.argv[4]),checkpoint=lambda phase,manifest: os._exit(73) if phase=='overview' else None)\n")
        result = subprocess.run([sys.executable, str(child), str(PROJECT / 'app'), str(self.overview), str(self.data), str(self.reports)],
                                env=dict(os.environ), capture_output=True, timeout=20)
        self.assertEqual(result.returncode, 73, result.stderr.decode(errors='replace'))
        incomplete = json.loads(self.paths['intent'].read_text())
        self.assertNotEqual(incomplete['generation'], old['generation'])
        with self.assertRaises(ValueError): publication.capture(self.target)
        recovered = self.publish(wait_seconds=2)
        self.assertEqual(publication.capture(self.target)['generation'], recovered['generation'])
        history = list(self.reports.glob('*.publication.pending.json.recovered-*.json'))
        self.assertEqual(len(history), 1)
        self.assertEqual(json.loads(history[0].read_text())['previous_intent']['generation'], incomplete['generation'])

    def test_staged_change_and_manifest_filename_tamper_are_refused(self):
        def change(phase, manifest):
            if phase == 'intent': self.data.write_bytes(b'changed after source fingerprints were taken')
        with self.assertRaisesRegex(ValueError, 'changed while being copied'):
            self.publish(checkpoint=change)
        with self.assertRaises(ValueError): publication.capture(self.target)
        self.make_books('restored'); self.publish()
        manifest = json.loads(self.paths['manifest'].read_text())
        manifest['data']['name'] = '../outside.xlsx'
        write_json(self.paths['manifest'], manifest)
        with self.assertRaises(ValueError): publication.capture(self.target)

    def test_read_boundary_notices_intent_after_actual_or_cached_read(self):
        self.publish()
        def read_then_intent():
            content = self.target.read_bytes()
            self.pending()
            return content
        with self.assertRaises(ValueError): publication.read([self.target], read_then_intent)

    def test_overview_only_generation_does_not_expose_old_data(self):
        self.publish()
        publication.publish(self.overview, None, self.reports)
        self.assertFalse(publication.capture(self.target)['data_published'])
        with self.assertRaises(ValueError): publication.capture(self.paths['data'])
        source = contributions._sources(self.reports)[0]
        self.assertNotIn('data_workbook', source)

    def test_audit_hashes_are_verified_before_final_artifacts_change(self):
        self.publish()
        before = self.target.read_bytes()
        audit = self.overview.with_suffix('.audit'); audit.mkdir()
        (audit / 'paths.csv').write_text('cycle,amount\n2026,2\n')
        write_json(audit / 'audit-manifest.json', {'workbook_name': self.overview.name,
            'workbook_sha256': contributions._hash(self.overview), 'csv_sha256': {'paths.csv': 'a' * 64}})
        with self.assertRaisesRegex(ValueError, 'audit CSV differs'):
            self.publish(audit_directory=audit)
        self.assertEqual(self.target.read_bytes(), before)
        self.assertEqual(publication.capture(self.target)['state'], 'committed')

    def test_audit_junction_escape_is_refused(self):
        outside = self.root / 'outside'; outside.mkdir()
        link = self.paths['audit'].parent
        self.assertTrue(outside.resolve().is_relative_to(self.root.resolve()))
        self.assertTrue(link.absolute().is_relative_to(self.root.resolve()))
        if os.name == 'nt':
            powershell = shutil.which('powershell.exe')
            if not powershell: self.skipTest('PowerShell unavailable for isolated junction fixture')
            subprocess.run([powershell, '-NoProfile', '-NonInteractive', '-Command',
                'New-Item -ItemType Junction -Path $env:REPORT_TEST_LINK -Target $env:REPORT_TEST_TARGET | Out-Null'],
                env=dict(os.environ, REPORT_TEST_LINK=str(link), REPORT_TEST_TARGET=str(outside)),
                capture_output=True, check=True, timeout=10,
                creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        else: link.symlink_to(outside, target_is_directory=True)
        with self.assertRaises(ValueError): self.publish()
        self.assertEqual(list(outside.iterdir()), [])


class ConsumerBoundaries(PublicationFixture):
    def test_marker_invalidates_source_memo_and_cached_worksheet_graph_comparison(self):
        self.publish(); source_memo.install()
        self.assertTrue(contributions._sources(self.reports)[0]['available'])
        report_insights.worksheet_entry(self.target, 'Spending Overview', 5)
        graph_data.graph_view([self.target], {})
        comparison_data.flows_entry([self.target], 'Fixture')
        before = self.target.stat()
        self.pending()
        self.assertEqual((self.target.stat().st_size, self.target.stat().st_mtime_ns), (before.st_size, before.st_mtime_ns))
        self.assertFalse(contributions._sources(self.reports)[0]['available'])
        for reader in (lambda: report_insights.worksheet_entry(self.target, 'Spending Overview', 5),
                       lambda: graph_data.graph_view([self.target], {}),
                       lambda: comparison_data.flows_entry([self.target], 'Fixture')):
            with self.assertRaises(ValueError): reader()

    def test_contribution_http_boundary_refuses_intent_created_by_reader(self):
        self.publish()
        sources = contributions._sources(self.reports)
        handler = object.__new__(app.Handler)
        handler.server = types.SimpleNamespace(app=self.console)
        def reader(*args):
            self.pending()
            return {'rows': [{'amount': 1}]}
        with patch.object(handler, '_contribution_result', side_effect=reader):
            with self.assertRaises(ValueError):
                handler.contribution_result(None, {'dataset': 'Fixture'}, sources)
        with self.assertRaises(ValueError): contributions._token(sources[0])

    def test_workbook_hash_from_before_committed_capture_cannot_mix_with_new_book(self):
        shutil.copyfile(self.overview, self.target)
        saved = self.root / 'new.xlsx'; self.make_books('changed before capture'); shutil.copyfile(self.overview, saved)
        real = publication.capture; changed = []
        def capture(path):
            if not changed:
                shutil.copyfile(saved, self.target); changed.append(True)
            return real(path)
        with patch.object(publication, 'capture', side_effect=capture):
            self.assertFalse(contributions._sources(self.reports)[0]['available'])

    def metadata_get(self, route, mutation=None):
        tools = self.root / 'tools'; tools.mkdir(exist_ok=True)
        (tools / 'cross_reference_engine.py').write_text('# synthetic catalogue-only rule marker\n')
        httpd = app.ThreadingHTTPServer(('127.0.0.1', 0), app.Handler); httpd.app = self.console
        worker = threading.Thread(target=httpd.serve_forever, daemon=True); worker.start()
        try:
            with patch.object(app.contribution_workbook_guard, 'verify_sources', return_value=[]):
                def get():
                    client = http.client.HTTPConnection('127.0.0.1', httpd.server_port, timeout=10)
                    try:
                        client.request('GET', route); response = client.getresponse()
                        return response.status, json.loads(response.read())
                    finally: client.close()
                if mutation is None: return get()
                module, name = mutation; original = getattr(module, name)
                def changed(*args, **kwargs):
                    value = original(*args, **kwargs); self.pending(); return value
                with patch.object(module, name, side_effect=changed): return get()
        finally:
            httpd.shutdown(); httpd.server_close(); worker.join(timeout=5)

    def test_metadata_http_refuses_intent_after_scan_on_all_four_routes(self):
        import analysis_data
        targets = [('/api/contributions/catalog', app.contribution_query.semantics, 'catalog'),
                   ('/api/insights/catalog', analysis_data, 'catalog'),
                   ('/api/contributions/session', app.snapshot_registry, 'registry'),
                   ('/api/snapshots', app.snapshot_registry, 'registry')]
        for route, module, name in targets:
            with self.subTest(route=route):
                self.publish()
                status, body = self.metadata_get(route, (module, name))
                self.assertEqual(status, 400, body)
                self.assertIn('publication is incomplete', body['error'])

    def test_metadata_http_keeps_legacy_committed_and_pending_availability_honest(self):
        routes = [('/api/contributions/catalog', 'datasets'), ('/api/insights/catalog', 'datasets'),
                  ('/api/contributions/session', 'industries'), ('/api/snapshots', 'snapshots')]
        shutil.copyfile(self.overview, self.target)
        shutil.copyfile(self.data, self.paths['data'])
        for state in ('legacy', 'committed', 'pending'):
            if state == 'committed': self.publish()
            if state == 'pending': self.pending()
            for route, key in routes:
                with self.subTest(state=state, route=route):
                    status, body = self.metadata_get(route)
                    self.assertEqual(status, 200, body)
                    self.assertEqual(len(body[key]), 1)
                    self.assertEqual(body[key][0]['available'], state != 'pending')

    def test_metadata_refuses_old_scan_after_a_new_committed_generation(self):
        self.publish(); sources = contributions._sources(self.reports)
        self.publish()  # identical workbook bytes, different committed generation
        handler = object.__new__(app.Handler); handler.server = types.SimpleNamespace(app=self.console)
        calls = []
        with patch.object(contributions, '_sources', return_value=sources):
            with self.assertRaisesRegex(ValueError, 'publication is incomplete'):
                handler.source_metadata(lambda scanned: calls.append(scanned))
        self.assertEqual(calls, [])

    def test_graph_envelope_refuses_old_registry_echo_and_late_pending_intent(self):
        self.publish()
        original = self.console.graph
        def replaced(*args, **kwargs):
            self.make_books('new generation after registry'); self.publish()
            return original(*args, **kwargs)
        with patch.object(self.console, 'graph', side_effect=replaced):
            status, body = self.metadata_get('/api/graph?id=published-reports&industry=Fixture&cycle=2026')
        self.assertEqual(status, 400, body)
        self.assertIn('publication is incomplete', body['error'])
        original_view = app.graph_view
        def late(*args, **kwargs):
            value = original_view(*args, **kwargs); self.pending(); return value
        with patch.object(app, 'graph_view', side_effect=late):
            status, body = self.metadata_get('/api/graph?id=published-reports&industry=Fixture&cycle=2026')
        self.assertEqual(status, 400, body)
        self.assertIn('publication is incomplete', body['error'])

    def test_graph_refuses_new_data_generation_even_when_overview_bytes_are_identical(self):
        self.publish(); original_hash = contributions._hash(self.target)
        original = self.console.graph
        def changed(*args, **kwargs):
            book = openpyxl.load_workbook(self.data); book.active['A1'] = 'changed Data only'
            book.save(self.data); book.close(); self.publish()
            self.assertEqual(contributions._hash(self.target), original_hash)
            return original(*args, **kwargs)
        with patch.object(self.console, 'graph', side_effect=changed):
            status, body = self.metadata_get('/api/graph?id=published-reports&industry=Fixture&cycle=2026')
        self.assertEqual(status, 400, body)
        self.assertIn('publication is incomplete', body['error'])

    def test_graph_and_comparison_do_not_gate_an_unrelated_pending_industry(self):
        self.publish()
        other = self.stage / 'Other Industry Spending Overview.xlsx'
        shutil.copyfile(self.overview, other)
        publication.publish(other, None, self.reports)
        cache = self.root / '.cache/industry-report/other'
        payload = cache / ('industry_data-' + self.source_hash[:16] + '.json')
        write_json(payload, {'metadata': {'industry': 'Other', 'database_path': str(self.database),
            'expected_cycle_totals': {'2025-2026': sum(row[4] for row in GIVING)}}})
        write_json(cache / ('manifest-' + self.source_hash[:16] + '.json'),
                   {'database_sha256': self.source_hash, 'data_sha256': contributions._hash(payload)})
        self.assertTrue(next(v for v in contributions._sources(self.reports) if v['industry'] == 'Other')['available'])
        paths = publication.locations(self.reports / other.name)
        write_json(paths['intent'], {'schema_version': 1, 'generation': 'f' * 32})
        for route in ('/api/graph', '/api/comparison', '/api/comparison-flows'):
            with self.subTest(route=route):
                status, body = self.metadata_get(route + '?id=published-reports&industry=Fixture&cycle=2026')
                self.assertEqual(status, 200, body)
                self.assertTrue(body['reports'])
                self.assertTrue(all(v['name'] == self.target.name for v in body['reports']))

    def test_producer_to_catalogue_actual_http_graph_and_worksheet(self):
        manifest = self.publish()
        httpd = app.ThreadingHTTPServer(('127.0.0.1', 0), app.Handler)
        httpd.app = self.console
        worker = threading.Thread(target=httpd.serve_forever, daemon=True); worker.start()
        try:
            def get(path):
                client = http.client.HTTPConnection('127.0.0.1', httpd.server_port, timeout=10)
                try:
                    client.request('GET', path); reply = client.getresponse()
                    return reply.status, json.loads(reply.read())
                finally: client.close()
            status, catalog = get('/api/reports?id=published-reports')
            self.assertEqual(status, 200)
            overview = next(v for v in catalog['reports'] if v['kind'] == 'overview')
            self.assertEqual(overview['publication']['generation'], manifest['generation'])
            status, graph = get('/api/graph?id=published-reports&industry=Fixture&cycle=2026')
            self.assertEqual(status, 200, graph)
            self.assertEqual(graph['totals']['giving'], sum(row[4] for row in GIVING))
            status, sheet = get('/api/report-data?id=published-reports&index=' + str(overview['index']) + '&sheet=Spending%20Overview&header=5')
            self.assertEqual(status, 200, sheet)
            self.assertEqual(sheet['rows'][0][2], graph['totals']['giving'])
            self.pending()
            status, refused = get('/api/report-data?id=published-reports&index=' + str(overview['index']) + '&sheet=Spending%20Overview&header=5')
            self.assertEqual(status, 400)
        finally:
            httpd.shutdown(); httpd.server_close(); worker.join(timeout=5)


if __name__ == '__main__': unittest.main()
