"""Unit tests for the tidy workbook writer.

The fixture store is built from the contract's own section-4 example rows, so a
schema drift between `STAGING_CONTRACT.md` and the writer fails here first.
"""
from __future__ import annotations

import datetime
from contextlib import closing
import json
import sqlite3
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
import build_tidy_workbook as writer
# The console module that reads a split sheet's parts back (SheetNameTest).
sys.path.append(str(Path(__file__).resolve().parents[1] / 'app'))
import report_insights

CONTRIBUTION = {
    "source_id": "4021620231727827649", "transaction_id": "SA11AI.12345",
    "back_reference_transaction_id": None, "cycle_label": "2021-2022", "cycle": 2022,
    "source_cycle": 2022, "cycle_alignment": "OK", "company": "Example Credit Union",
    "contributor": "DOE, JANE", "person_key": "Example Credit Union|DOE|JANE|C01",
    "identity_basis": "Exact name + employer", "employee_tier": "Employee",
    "employer": "EXAMPLE CREDIT UNION", "occupation": "MEMBER SERVICE", "city": "VIENNA",
    "state": "VA", "zip5": "22180", "zipcode": "221804907", "date": "2021-06-28",
    "amount": 25.0, "analysis_amount": 25.0, "include_in_analysis": 1,
    "inclusion_reason": "Countable receipt", "company_match_status": "accepted",
    "company_match_reason": "verified_roster_alias", "matched_alias": "EXAMPLE CREDIT UNION",
    "recipient": "EXAMPLE CONDUIT", "committee_id": "C00000001",
    "recipient_key": "FEC:C00000001", "recipient_category": "Candidate",
    "category_source": "master", "party": "Democratic", "party_source": "master",
    "committee_type_code": "O", "committee_type": "Independent expenditure-only (Super PAC)",
    "committee_designation_code": "U", "committee_designation": "Unauthorized",
    "record_type": "contribution", "source_table": "contributions",
    "funding_source_type": "individual", "funding_direction": "inbound",
    "receipt_role": "original", "receipt_type_full": "Earmarked contribution",
    "filing_form": "F3X", "line_number": "11AI", "memo_text": None, "is_memo_leg": 0,
    "is_conduit_leg": 0, "superseded_by_amendment": 0, "duplicate_filing_version": 0,
    "duplicate_memo": 0, "duplicate_of_sub_id": None, "duplicate_reason": None,
    "recipient_state": "MA", "represented_state": None,
    "source_url": "https://docquery.fec.gov/cgi-bin/fecimg/?000000000000000000",
}

GIVING = {"cycle_label": "2021-2022", "cycle": 2022, "company": "Example Credit Union",
          "employee_tier": "Employee", "recipient_category": "Candidate",
          "party": "Democratic", "giving": 25.0, "gift_count": 1,
          "distinct_contributors": 1, "distinct_recipients": 1}

RECIPIENT = {"cycle_label": "2021-2022", "cycle": 2022, "company": "Example Credit Union",
             "committee_id": "C00000001", "employee_tier": "Employee",
             "recipient": "EXAMPLE CONDUIT", "recipient_category": "Candidate",
             "category_source": "master", "party": "Democratic", "party_source": "master",
             "committee_type_code": "O", "committee_type": "Super PAC",
             "committee_designation_code": "U", "committee_designation": "Unauthorized",
             "sponsor_candidate_id": None, "sponsor_candidate": None,
             "sponsor_office": None, "sponsor_state": None, "sponsor_district": None,
             "sponsor_count": 0, "connected_organization": None,
             "committee_relationship": "Conduit", "ownership_status": "Unaffiliated",
             "giving": 25.0, "gift_count": 1, "distinct_contributors": 1,
             "first_gift_date": "2021-06-28", "last_gift_date": "2021-06-28",
             "fec_committee_url": "https://www.fec.gov/data/committee/C00000001/"}

CONTRIBUTOR = {"cycle_label": "2021-2022", "cycle": 2022, "company": "Example Credit Union",
               "person_key": "Example Credit Union|DOE|JANE|C01", "party": "Democratic",
               "contributor": "DOE, JANE", "employee_tier": "Employee",
               "identity_basis": "Exact name + employer", "city": "VIENNA", "state": "VA",
               "zip5": "22180", "employer": "EXAMPLE CREDIT UNION",
               "occupation": "MEMBER SERVICE", "distinct_addresses": 1, "giving": 25.0,
               "gift_count": 1}

LOCATION = {"cycle_label": "2021-2022", "cycle": 2022, "company": "Example Credit Union",
            "state": "VA", "city": "VIENNA", "zip5": "22180", "giving": 25.0,
            "gift_count": 1, "distinct_contributors": 1}

CANDIDATE = {"candidate_id": "H0XX00001", "candidate": "ROE, RICHARD",
             "party": "Democratic", "party_affiliation": "DEM", "office_code": "H",
             "office": "House", "state": "XX", "district": "01",
             "incumbent_status": "Incumbent", "principal_committee_id": "C00000002",
             "fec_candidate_url": "https://www.fec.gov/data/candidate/H0XX00001/",
             "source_url": "https://www.fec.gov/files/bulk-downloads/2022/cn22.zip"}

DATABASE_SHA = "a" * 64

METADATA = {"schema_version": "tidy-1", "industry": "Example", "industry_slug": "example",
            "cycle_label": "", "supported_cycles": "[2022, 2024]",
            "cycle_labels": '["2021-2022", "2023-2024"]',
            "database_path": "C:/nowhere/mentat.db", "database_sha256": DATABASE_SHA,
            "generated_at": "2026-09-14T00:00:00+00:00", "exporter_version": "test",
            "expected_cycle_totals": '{"2021-2022": 25.0, "2023-2024": 25.0}',
            "accounting_note": "Analysis Amount counts included receipts only."}


def second_cycle(row):
    """The same observation moved into 2023-2024 so the split has two groups."""
    copy = dict(row)
    for key, value in row.items():
        if key == "cycle_label":
            copy[key] = "2023-2024"
        elif key == "cycle":
            copy[key] = 2024
        elif key == "source_id":
            copy[key] = "4021620231727827650"
        elif key in ("date", "first_gift_date", "last_gift_date"):
            copy[key] = "2023-06-28"
    return copy


def create(connection, table, rows):
    columns = list(rows[0])
    connection.execute(f'CREATE TABLE {table} ({", ".join(columns)})')
    connection.executemany(
        f'INSERT INTO {table} VALUES ({", ".join("?" * len(columns))})',
        [[row[c] for c in columns] for row in rows])


def build_store(path, giving_override=None, recipient_override=None,
                relationship_override=None):
    connection = sqlite3.connect(path)
    with connection:
        create(connection, "metadata",
               [{"key": k, "value": v} for k, v in METADATA.items()])
        create(connection, "contributions", [CONTRIBUTION, second_cycle(CONTRIBUTION)])
        giving = [GIVING, second_cycle(GIVING)]
        if giving_override is not None:
            giving[0] = dict(giving[0], giving=giving_override)
        create(connection, "giving_facts", giving)
        recipients = [RECIPIENT, second_cycle(RECIPIENT)]
        if recipient_override is not None:
            recipients[0] = dict(recipients[0], recipient=recipient_override)
        if relationship_override is not None:
            recipients[0] = dict(recipients[0],
                                 committee_relationship=relationship_override)
        create(connection, "recipient_facts", recipients)
        create(connection, "contributor_facts", [CONTRIBUTOR, second_cycle(CONTRIBUTOR)])
        create(connection, "location_facts", [LOCATION, second_cycle(LOCATION)])
        create(connection, "candidates", [CANDIDATE])
        create(connection, "data_quality", [{"check_id": "roster_coverage",
                                             "source": "r_guard", "severity": "fail",
                                             "passed": 1, "detail": "all rows collected",
                                             "metric": "10"}])
    connection.close()


class WriterTest(unittest.TestCase):
    def setUp(self):
        self.directory = Path(tempfile.mkdtemp(prefix="tidy-test-"))
        self.store = self.directory / "tidy-test.sqlite"
        self.out = self.directory / "Example Industry Spending Data.xlsx"
        build_store(self.store)

    def run_writer(self, *extra):
        return writer.main(["--staging", str(self.store), "--out", str(self.out), *extra])

    def book(self):
        return openpyxl.load_workbook(self.out, read_only=False, data_only=True)

    def test_sheets_and_headers(self):
        self.assertEqual(0, self.run_writer())
        book = self.book()
        self.assertEqual(
            ["About", "Contributions", "Giving Facts", "Recipient Facts",
             "Contributor Facts", "Location Facts", "Candidates", "Data Quality"],
            book.sheetnames)
        for name in book.sheetnames:
            sheet = book[name]
            headers = [c.value for c in sheet[1]]
            self.assertTrue(all(isinstance(h, str) and h.strip() for h in headers), name)
            self.assertEqual(len(headers), len(set(headers)), name)
            self.assertEqual("A2", sheet.freeze_panes, name)
            self.assertEqual([], list(sheet.merged_cells.ranges), name)

    def test_network_activity_is_complete_and_company_associations_remain_separate(self):
        activity=[dict(cycle_label='2021-2022',cycle=2022,company='Industry Total',
            source_committee_id='PAC',candidate_id=f'H{i:08}',bridge_key=f'network-{i}',
            downstream_activity=float(i),support_amount=float(i),oppose_amount=0.,other_activity_amount=0.)
            for i in range(1,106)]
        associations=[dict(row,company=company,bridge_key=f'{company}-{row["candidate_id"]}')
            for company in ['A','B'] for row in activity]
        with closing(sqlite3.connect(self.store)) as connection,connection:
            create(connection,'beneficiary_activity',activity)
            create(connection,'beneficiary_bridge',associations)
        self.assertEqual(0,self.run_writer())
        book=self.book()
        network=list(book['Network Candidate Activity'].values)
        companies=list(book['Beneficiary Bridge'].values)
        self.assertEqual(len(network)-1,105)
        self.assertEqual(len(companies)-1,210)
        col=network[0].index('Downstream Activity')
        self.assertEqual(sum(row[col] for row in network[1:]),5565)
        col=companies[0].index('Downstream Activity')
        self.assertEqual(sum(row[col] for row in companies[1:]),11130)

    def test_downstream_uncertainty_survives_workbook_export(self):
        with closing(sqlite3.connect(self.store)) as connection,connection:
            create(connection,'downstream_activity',[dict(cycle_label='2021-2022',cycle=2022,
                source_committee_id='PAC',candidate_id='H1',activity_type='Independent expenditure',
                support_oppose='Support',transaction_type='24E',unverified_version_rows=7,
                unresolved_version_rows=2,unresolved_version_amount=12.34,repeated_itemization_rows=3)])
        self.assertEqual(0,self.run_writer())
        rows=list(self.book()['Downstream Activity'].values)
        record=dict(zip(rows[0],rows[1]))
        self.assertEqual(record['Unverified Version Rows'],7)
        self.assertEqual(record['Unresolved Version Amount'],12.34)

    def test_table_refs_span_every_column_and_row(self):
        self.assertEqual(0, self.run_writer())
        book = self.book()
        for name in book.sheetnames:
            sheet = book[name]
            tables = list(sheet.tables.values())
            self.assertEqual(1, len(tables), name)
            reference = tables[0].ref
            start, end = reference.split(":")
            self.assertEqual("A1", start, name)
            self.assertEqual(sheet.cell(sheet.max_row, sheet.max_column).coordinate,
                             end, name)
            self.assertEqual([c.name for c in tables[0].tableColumns],
                             [c.value for c in sheet[1]], name)

    def test_types_and_formats(self):
        self.assertEqual(0, self.run_writer())
        sheet = self.book()["Contributions"]
        headers = [c.value for c in sheet[1]]
        row = {h: sheet.cell(2, i + 1) for i, h in enumerate(headers)}
        self.assertIsInstance(row["Include in Analysis"].value, bool)
        self.assertIsInstance(row["Memo Leg"].value, bool)
        self.assertIsInstance(row["Date"].value, datetime.datetime)
        self.assertEqual(writer.DATE_FORMAT, row["Date"].number_format)
        self.assertIsInstance(row["Source ID"].value, str)
        self.assertIsInstance(row["Contributor ZIP"].value, str)
        self.assertIsInstance(row["Line Number"].value, str)
        self.assertEqual(writer.MONEY_FORMAT, row["Amount"].number_format)
        self.assertIsNone(row["Memo Text"].value)
        self.assertIsNone(row["Back Reference Transaction ID"].value)

    def test_downstream_ledgers_and_hashes_survive_staging_and_workbook(self):
        from enrich_downstream import stage
        payload = {'downstream_reach':[dict(cycle=2024,view='Company',company='Example',
            source_committee_id='C00836221',representative_root='C00835959',
            representative_path='C00835959 > C00836221')],
            'downstream_measurement_ledger':[dict(cycle=2024,view='Company',company='Example',
                candidate_id='H0KY06104',before_amount=0,after_amount=106708,delta=106708,
                basis='Same source rows; synthetic test')],
            'downstream_sources':[dict(cycle=2024,source='fixture',path='source-fixture',
                sha256='a'*64,status='fixture only',policy='count_all')]}
        stage(payload,self.store)
        self.assertEqual(0,self.run_writer())
        book=self.book()
        for name in ['Downstream Reach','Downstream Measurement','Downstream Sources']:
            self.assertEqual(book[name].max_row,2)
        sheet=book['Downstream Measurement']
        row=dict(zip([c.value for c in sheet[1]],[c.value for c in sheet[2]]))
        self.assertEqual(row['Before Amount'],0)
        self.assertEqual(row['After Amount'],106708)
        self.assertEqual(row['Candidate ID'],'H0KY06104')
        self.assertIn('a'*64,[c.value for c in book['Downstream Sources'][2]])

    def test_split_forces_per_cycle_sheets(self):
        self.assertEqual(0, self.run_writer("--split-rows", "1"))
        book = self.book()
        # S1 (Paul, 2026-09-26): a part is `<Sheet> <election year>`. Before it,
        # this test pinned "Contributions 2021-2022" and "Contributions 2023-2024".
        self.assertIn("Contributions 2022", book.sheetnames)
        self.assertIn("Contributions 2024", book.sheetnames)
        self.assertNotIn("Contributions", book.sheetnames)
        self.assertNotIn("Contributions 2021-2022", book.sheetnames)
        for name in book.sheetnames:
            self.assertFalse(name.strip().replace("-", "").isdigit(), name)
        self.assertEqual(2, book["Contributions 2022"].max_row)
        tables = list(book["Contributions 2022"].tables)
        self.assertEqual(["tbl_contributions_2022"], tables)
        receipt = json.loads(Path(str(self.out) + ".receipt.json").read_text())
        self.assertIn("Contributions", receipt["split_sheets"])
        self.assertEqual(["Contributions 2022", "Contributions 2024"],
                         receipt["split_sheets"]["Contributions"])

    def test_network_candidate_activity_splits_under_its_short_name(self):
        # The crash S1 fixes: the old part name was 36 characters where Excel
        # allows 31, so every all-cycles build that split this sheet died in
        # open_sheet. Its parts are now `Network Candidates <election year>`.
        self.assertGreater(len("Network Candidate Activity 2021-2022"), writer.MAX_SHEET_NAME)
        activity = [dict(cycle_label=label, cycle=int(label[-4:]), company="Industry Total",
                         source_committee_id="PAC", candidate_id="H00000001",
                         bridge_key=f"network-{label}", downstream_activity=10.0)
                    for label in ("2021-2022", "2023-2024")]
        with closing(sqlite3.connect(self.store)) as connection, connection:
            create(connection, "beneficiary_activity", activity)
        self.assertEqual(0, self.run_writer("--split-rows", "1"))
        book = self.book()
        self.assertIn("Network Candidates 2022", book.sheetnames)
        self.assertIn("Network Candidates 2024", book.sheetnames)
        self.assertNotIn("Network Candidate Activity", book.sheetnames)
        for name in book.sheetnames:
            self.assertEqual([], writer.sheet_name_problems(name), name)
        self.assertEqual(2, book["Network Candidates 2024"].max_row)
        receipt = json.loads(Path(str(self.out) + ".receipt.json").read_text())
        self.assertEqual(["Network Candidates 2022", "Network Candidates 2024"],
                         receipt["split_sheets"]["Network Candidate Activity"])

    def test_checks_pass_and_about_lineage(self):
        self.assertEqual(0, self.run_writer())
        book = self.book()
        about = {r[0].value: r[1].value for r in book["About"].iter_rows(min_row=2)}
        self.assertEqual(f"Source DB SHA-256: {DATABASE_SHA}", about["Lineage"])
        self.assertEqual("2", about["Rows: Contributions"])
        self.assertEqual(str(book["About"].max_row - 1), about["Rows: About"])
        absent = json.loads(about["Absent tables"])
        self.assertIn("committees", absent)
        self.assertIn("downstream_activity", absent)
        blanks = json.loads(about["Blank columns"])
        # A header is blank when the store lacks the column OR every row is NULL.
        self.assertIn("Contributions / Memo Text", blanks)
        self.assertIn("Recipient Facts / Sponsor Candidate", blanks)
        self.assertNotIn("Contributions / Amount", blanks)
        self.assertNotIn("Giving Facts / Giving", blanks)
        quality = book["Data Quality"]
        headers = [c.value for c in quality[1]]
        rows = [dict(zip(headers, [c.value for c in r]))
                for r in quality.iter_rows(min_row=2)]
        identifiers = {r["Check"] for r in rows}
        for check in ("giving_facts_vs_contributions", "recipient_facts_vs_contributions",
                      "contributor_facts_vs_contributions",
                      "location_facts_vs_contributions", "expected_cycle_totals",
                      "overview_cycle_totals", "database_amount_and_rows", "unique_keys",
                      "no_multivalue_cells", "header_vocabulary", "lineage_literal"):
            self.assertIn(check, identifiers, check)
        self.assertTrue(all(r["Passed"] for r in rows if r["Severity"] == "fail"))
        self.assertTrue(any(r["Source"] == "r_guard" for r in rows))
        self.assertTrue(all(isinstance(r["Passed"], bool) for r in rows))

    def test_skipped_checks_state_their_reason(self):
        self.assertEqual(0, self.run_writer())
        book = self.book()
        quality = book["Data Quality"]
        headers = [c.value for c in quality[1]]
        rows = [dict(zip(headers, [c.value for c in r]))
                for r in quality.iter_rows(min_row=2)]
        skipped = {r["Check"]: r["Detail"] for r in rows if r["Severity"] == "info"}
        self.assertIn("overview_cycle_totals", skipped)
        self.assertTrue(skipped["overview_cycle_totals"].startswith("SKIPPED: "))
        self.assertIn("database_amount_and_rows", skipped)
        self.assertIn("--database not supplied", skipped["database_amount_and_rows"])
        self.assertEqual(0, self.run_writer("--database", str(self.directory / "gone.db")))
        book = self.book()
        quality = book["Data Quality"]
        headers = [c.value for c in quality[1]]
        rows = [dict(zip(headers, [c.value for c in r]))
                for r in quality.iter_rows(min_row=2)]
        detail = next(r["Detail"] for r in rows
                      if r["Check"] == "database_amount_and_rows")
        self.assertIn("source database missing", detail)

    def test_committee_name_with_semicolon_does_not_fail_w9(self):
        """A real FEC name carrying ';' must publish.

        "MADISON PAC; THE" is one committee, not a packed list. Under the old
        exempt-list form of W9 the first live Credit Union export refused on exactly
        this value (Downstream Activity / Source Committee, 78 rows) and on
        "JFC TRANSFER: ...; 2021 AGGREGATE" (Contributions / Receipt Type, 6 rows).
        Free-text name columns are not checked at all now, so both publish.
        """
        # "Recipient" is free text and must not be checked.
        self.assertNotIn("Recipient", writer.MULTIVALUE_CHECKED_HEADERS)
        # "Source Committee" is the column that actually refused on 2026-09-15. The
        # fixture store has no downstream_activity table, so the guarantee is asserted
        # against the rule itself rather than end to end.
        self.assertNotIn("Source Committee", writer.MULTIVALUE_CHECKED_HEADERS)
        self.assertNotIn("Receipt Type", writer.MULTIVALUE_CHECKED_HEADERS)

        self.store.unlink()
        build_store(self.store, recipient_override="MADISON PAC; THE")
        self.assertEqual(0, self.run_writer())
        receipt = json.loads(Path(str(self.out) + ".receipt.json").read_text())
        self.assertTrue(receipt["ok"])
        multivalue = [c for c in receipt["checks"]
                      if c["check_id"] == "no_multivalue_cells"]
        self.assertTrue(multivalue)
        for check in multivalue:
            self.assertTrue(check["passed"], check["detail"])
        # and the value really did survive into the sheet
        sheet = self.book()["Recipient Facts"]
        headers = [c.value for c in sheet[1]]
        column = headers.index("Recipient") + 1
        values = {sheet.cell(row=r, column=column).value
                  for r in range(2, sheet.max_row + 1)}
        self.assertIn("MADISON PAC; THE", values)

    def test_packed_value_in_categorical_column_fails_w9(self):
        """The guard must still bite where a packed list can actually arise.

        "Committee Relationship" is a categorical dimension our own derivation writes,
        so a '; '-joined value there is a real defect and must block publication.
        """
        self.assertIn("Committee Relationship", writer.MULTIVALUE_CHECKED_HEADERS)
        self.store.unlink()
        build_store(self.store, relationship_override="Conduit; Principal campaign committee")
        self.assertEqual(2, self.run_writer())
        self.assertFalse(self.out.exists())
        receipt = json.loads(Path(str(self.out) + ".receipt.json").read_text())
        self.assertFalse(receipt["ok"])
        failed = [c for c in receipt["checks"]
                  if c["check_id"] == "no_multivalue_cells" and not c["passed"]]
        self.assertEqual(1, len(failed), receipt["checks"])
        self.assertIn("Committee Relationship", failed[0]["detail"])

    def test_corrupt_fact_row_blocks_publication(self):
        self.store.unlink()
        build_store(self.store, giving_override=26.0)
        self.assertEqual(2, self.run_writer())
        self.assertFalse(self.out.exists())
        self.assertEqual([], list(self.out.parent.glob("*.pending-*")))
        receipt = json.loads(Path(str(self.out) + ".receipt.json").read_text())
        self.assertFalse(receipt["ok"])
        self.assertIsNone(receipt["sha256"])
        failures = [c for c in receipt["checks"]
                    if c["severity"] == "fail" and not c["passed"]]
        self.assertTrue(any(c["check_id"] == "giving_facts_vs_contributions"
                            for c in failures))

    def test_receipt_shape(self):
        self.assertEqual(0, self.run_writer())
        receipt = json.loads(Path(str(self.out) + ".receipt.json").read_text())
        for key in ("expected_worksheet_count", "sha256", "rows_per_sheet", "checks",
                    "split_sheets", "schema_version", "writer_version"):
            self.assertIn(key, receipt)
        self.assertEqual(writer.SCHEMA_VERSION, receipt["schema_version"])
        self.assertEqual(8, receipt["expected_worksheet_count"])
        self.assertEqual(writer.file_sha256(self.out), receipt["sha256"])


class HelperTest(unittest.TestCase):
    def test_title_header_builds_flag_headers(self):
        self.assertEqual("Has Support Activity", writer.title_header("has_support"))
        self.assertEqual("Has Oppose Activity", writer.title_header("has_oppose"))
        self.assertEqual("Has Contribution to Nonaffiliated Committee Activity",
                         writer.title_header("has_contribution_to_nonaffiliated_committee"))
        self.assertEqual("Committee ID", writer.title_header("committee_id"))

    def test_convert(self):
        self.assertIs(True, writer.convert(1, "B"))
        self.assertIs(False, writer.convert(0, "B"))
        self.assertIsNone(writer.convert(None, "B"))
        self.assertEqual(datetime.date(2021, 6, 28), writer.convert("2021-06-28", "D"))
        self.assertIsNone(writer.convert("", "T"))
        self.assertEqual("007", writer.convert("007", "ID"))

    def test_year_columns_are_not_thousands_formatted(self):
        self.assertEqual(writer.YEAR_FORMAT, writer.number_format("I", "Cycle End Year"))
        self.assertEqual(writer.INTEGER_FORMAT, writer.number_format("I", "Gift Count"))


class FormulaGuardTest(unittest.TestCase):
    """Text from an FEC filing must reach the workbook as text, never as a formula.

    openpyxl turns a string that begins with '=' into a live formula; the writer
    stores every string that begins with = + - @ tab or CR as an explicit string.
    """

    COLUMNS = [("employer", "Employer", "T"), ("occupation", "Occupation", "T"),
               ("amount", "Amount", "R")]
    HOSTILE = ['=HYPERLINK("http://x","y")', "+1+1", "@SUM(A1)", "-2+3",
               "\t=1+1", "\r=1+1"]

    def written(self, rows):
        """Rows through the real append_rows, saved and reloaded formulas-intact."""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "guard.xlsx"
            book = openpyxl.Workbook(write_only=True)
            sheet = writer.open_sheet(book, "Guard", self.COLUMNS)
            writer.append_rows(sheet, self.COLUMNS, rows)
            book.save(path)
            with zipfile.ZipFile(path) as archive:
                xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
            loaded = openpyxl.load_workbook(path, data_only=False)
            return [[cell for cell in row] for row in loaded["Guard"].iter_rows(min_row=2)], xml

    def test_hostile_text_stays_text_and_numbers_stay_numbers(self):
        rows = [(value, "x", 1.0) for value in self.HOSTILE]
        rows += [("Acme Corp", "Legit & Sons", -5), ("-", "", -0.25)]
        cells, xml = self.written(rows)
        self.assertNotIn("<f>", xml)
        for row, value in zip(cells, self.HOSTILE):
            self.assertEqual("s", row[0].data_type, repr(value))
            # XML normalizes a bare CR to LF on the way back, so compare stripped.
            self.assertEqual(value.strip(), row[0].value.strip(), repr(value))
        legit = cells[len(self.HOSTILE)]
        self.assertEqual(("Acme Corp", "Legit & Sons"), (legit[0].value, legit[1].value))
        # A real negative number is not a string and keeps its numeric cell.
        self.assertEqual("n", legit[2].data_type)
        self.assertEqual(-5, legit[2].value)
        self.assertEqual(writer.MONEY_FORMAT, legit[2].number_format)
        self.assertEqual("-", cells[-1][0].value)
        self.assertEqual(-0.25, cells[-1][2].value)

    def test_the_check_can_fail(self):
        # Negative control: with the guard off, openpyxl writes the '=' value as a formula.
        original = writer.FORMULA_LEADERS
        writer.FORMULA_LEADERS = ()
        try:
            cells, xml = self.written([(self.HOSTILE[0], "x", 1.0)])
        finally:
            writer.FORMULA_LEADERS = original
        self.assertEqual("f", cells[0][0].data_type)
        self.assertIn("<f>", xml)

    def test_employer_and_occupation_survive_the_full_writer(self):
        with tempfile.TemporaryDirectory(prefix="tidy-guard-") as directory:
            store = Path(directory) / "tidy.sqlite"
            out = Path(directory) / "Guard Industry Spending Data.xlsx"
            build_store(store)
            with closing(sqlite3.connect(store)) as connection, connection:
                for table in ("contributions", "contributor_facts"):
                    connection.execute(f"UPDATE {table} SET employer=?, occupation=?",
                                       (self.HOSTILE[0], self.HOSTILE[2]))
            self.assertEqual(0, writer.main(["--staging", str(store), "--out", str(out)]))
            book = openpyxl.load_workbook(out, data_only=False)
            for name in ("Contributions", "Contributor Facts"):
                headers = [c.value for c in book[name][1]]
                row = {h: book[name].cell(2, i + 1) for i, h in enumerate(headers)}
                for header, value in (("Employer", self.HOSTILE[0]),
                                      ("Occupation", self.HOSTILE[2])):
                    self.assertEqual("s", row[header].data_type, (name, header))
                    self.assertEqual(value, row[header].value, (name, header))
            book.close()
            with zipfile.ZipFile(out) as archive:
                self.assertFalse([n for n in archive.namelist()
                                  if n.startswith("xl/worksheets/")
                                  and b"<f>" in archive.read(n)])


class SheetNameTest(unittest.TestCase):
    """S1 (Paul, 2026-09-26): every sheet name the writer can emit is one Excel accepts.

    A sheet over `--split-rows` becomes one sheet per cycle named `<split base>
    <election year>`. The old `<Sheet> <cycle_label>` made "Network Candidate
    Activity 2021-2022", 36 characters, and every all-cycles build that split that
    sheet crashed. The checks walk the registry itself, so a sheet added later is
    covered without anyone listing it here.
    """

    FORBIDDEN = set("[]:*?/\\")
    # An election year is always four digits, so every part name of a sheet has
    # the same length; the largest even four-digit year stands in for the longest.
    CYCLES = ("2021-2022", "2025-2026", "9997-9998")

    def test_every_splittable_sheet_name_fits_excel_with_its_longest_year(self):
        # No registry sheet is exempt: a dynamic table takes whatever columns its
        # store has, so any of them can arrive with a cycle_label and split.
        for sheet in writer.SHEETS:
            parts = [writer.split_sheet_name(sheet, label) for label in self.CYCLES]
            longest = max(parts, key=len)
            with self.subTest(sheet=sheet.name, longest=longest):
                self.assertEqual({len(sheet.split_base) + 5}, {len(part) for part in parts})
                for name in (sheet.name, longest):
                    self.assertLessEqual(len(name), 31, name)
                    self.assertFalse(self.FORBIDDEN & set(name), name)
                    self.assertEqual([], writer.sheet_name_problems(name), name)

    def test_a_part_is_named_for_its_cycle_s_election_year(self):
        self.assertEqual("2022", writer.election_year("2021-2022"))
        self.assertEqual("2026", writer.election_year("2025-2026"))
        self.assertEqual("2024", writer.election_year("2024"))
        for label in (None, "", "NA-NA", "2022-2023", "2020-2022", "2023", "22", "Unknown"):
            with self.subTest(label=label):
                self.assertIsNone(writer.election_year(label))
        sheets = {sheet.table: sheet for sheet in writer.SHEETS}
        self.assertEqual("Contributions 2022",
                         writer.split_sheet_name(sheets["contributions"], "2021-2022"))
        self.assertEqual("Network Candidates 2022",
                         writer.split_sheet_name(sheets["beneficiary_activity"], "2021-2022"))
        # A label that is not an FEC cycle (the exporter's sprintf writes "NA-NA"
        # for a receipt with no date) keeps the pre-S1 spelling, so a store that
        # published before still publishes; a NULL label cannot be split at all.
        self.assertEqual("Contributions NA-NA",
                         writer.split_sheet_name(sheets["contributions"], "NA-NA"))
        with self.assertRaises(ValueError):
            writer.split_sheet_name(sheets["contributions"], None)

    def test_the_name_checks_can_fail(self):
        # Negative controls: each rule refuses the name it exists to refuse, and
        # the spelling S1 retired is refused for its length.
        self.assertTrue(writer.sheet_name_problems("Network Candidate Activity 2021-2022"))
        self.assertTrue(writer.sheet_name_problems("x" * 32))
        for character in sorted(self.FORBIDDEN):
            self.assertTrue(writer.sheet_name_problems(f"Transfer{character}Paths"), character)
        for name in ("", "   ", "'Quoted", "Quoted'", "History"):
            self.assertTrue(writer.sheet_name_problems(name), repr(name))

    def test_the_console_reads_every_part_back_to_its_sheet(self):
        for sheet in writer.SHEETS:
            self.assertEqual(sheet.name, report_insights.split_sheet_base(sheet.name))
            for label in ("2021-2022", "2023-2024", "2025-2026"):
                with self.subTest(sheet=sheet.name, cycle=label):
                    new = writer.split_sheet_name(sheet, label)
                    self.assertEqual(sheet.name, report_insights.split_sheet_base(new))
                    # Books published before S1 keep `<Sheet> <cycle_label>`.
                    self.assertEqual(sheet.name,
                                     report_insights.split_sheet_base(f"{sheet.name} {label}"))


class RawSourcePolicyTests(unittest.TestCase):
    def test_all_observations_are_reconciled_once_with_report_priority(self):
        with tempfile.TemporaryDirectory() as directory:
            database=Path(directory)/'source.db'
            with closing(sqlite3.connect(database)) as connection, connection:
                for table in ('contributions','corporate_giving'):
                    connection.execute(f'CREATE TABLE {table}(sub_id TEXT, industry TEXT, company TEXT, contribution_receipt_date TEXT, contribution_receipt_amount REAL, company_match_status TEXT)')
                connection.executemany('INSERT INTO contributions VALUES(?,?,?,?,?,?)',[
                    ('one','Fixture','Uncertain','2024-01-01',100,'review'),
                    ('two','Fixture','Retained','2024-01-01',50,'rejected')])
                connection.execute('INSERT INTO corporate_giving VALUES(?,?,?,?,?,?)',
                                   ('one','Fixture','Accepted','2024-01-01',100,'accepted'))
            self.assertEqual(writer.source_facts(database,'Fixture',{2024},True),
                             {('2023-2024','Accepted'):(100,1),('2023-2024','Retained'):(50,1)})
            self.assertEqual(writer.source_facts(database,'Fixture',{2024}),{('2023-2024','Accepted'):(100,1)})


if __name__ == "__main__":
    unittest.main()
