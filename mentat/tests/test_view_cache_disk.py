"""ViewCache's disk tier: Graph views survive a restart byte for byte; changed inputs, changed
content under an unchanged stamp and changed code are misses; a failed activity read stays off disk.

A child interpreter of this file stands in for a restarted console (output to real files,
waited on with a timeout and tree-killed: brain DE-026). downstream.ROOT points at an empty
folder so activity views never load the real PAC evidence inventory.
"""
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

import openpyxl

APP = Path(__file__).resolve().parents[1] / 'app'
sys.path.insert(0, str(APP))
if os.environ.get('MENTAT_CACHE_TEST_OVERLAY'):
    sys.path.insert(0, os.environ['MENTAT_CACHE_TEST_OVERLAY'])
import response_cache as rc  # noqa: E402
import view_cache  # noqa: E402
from view_cache import ViewCache, file_signature  # noqa: E402
import graph_data  # noqa: E402
import downstream  # noqa: E402
import report_insights  # noqa: E402
from graph_fixture import add_graph_sheets  # noqa: E402

NO_WINDOW = getattr(subprocess, 'CREATE_NO_WINDOW', 0)
NAMESPACE = 'graph-views'
FILTERS = [{}, {'latest': '1'}, {'company': 'Civic Horizon Bank'}, {'cycle': '2026'}, {'funding': 'corporate'},
           {'tier': 'Chief'}, {'recipient': 'C00000003'}, {'mode': 'activity'},
           {'mode': 'activity', 'direction': 'Oppose'}, {'mode': 'activity', 'company': 'Harborlight Credit Union'}]


def encoded(value):
    return view_cache._encode(value)


def kill_tree(process):
    if process.poll() is None:
        killer = subprocess.Popen(['taskkill', '/PID', str(process.pid), '/T', '/F'], stdout=subprocess.DEVNULL,
                                  stderr=subprocess.DEVNULL, creationflags=NO_WINDOW)
        try:
            killer.wait(timeout=15)
        except subprocess.TimeoutExpired:
            killer.kill()
    try:
        process.wait(timeout=15)
    except subprocess.TimeoutExpired:
        process.kill()


def make_workbook(path):
    book = openpyxl.Workbook()
    add_graph_sheets(book)
    book.save(path)
    book.close()


class DiskCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)
        self.root = self.dir / 'response-cache'
        self.evidence = self.dir / 'pac-evidence'
        self.evidence.mkdir()
        saved = rc.ROOT
        self.addCleanup(setattr, rc, 'ROOT', saved)
        self.addCleanup(rc.clear_memory)
        self.addCleanup(graph_data._views.clear)
        rc.configure(self.root)
        rc.clear_memory()
        graph_data._views.clear()
        patcher = patch.object(downstream, 'ROOT', self.evidence)
        patcher.start()
        self.addCleanup(patcher.stop)

    def stored(self):
        folder = self.root / NAMESPACE
        return sorted(p.name for p in folder.glob('*.json')) if folder.exists() else []


class GraphViewDiskTests(DiskCase):
    def setUp(self):
        super().setUp()
        self.path = self.dir / 'Giving.xlsx'
        make_workbook(self.path)

    def edit(self, change):
        book = openpyxl.load_workbook(self.path)
        change(book)
        book.save(self.path)
        book.close()

    def test_views_read_back_from_disk_are_byte_identical_to_fresh_builds(self):
        for filters in FILTERS:
            with self.subTest(filters=filters):
                graph_data._views.clear()
                built = graph_data.graph_view([self.path], filters)
                fresh = graph_data._graph_view([self.path], filters)  # no cache anywhere
                self.assertEqual(encoded(built), encoded(fresh))
                self.assertEqual(graph_data._views.info()['disk_builds'], 1)
                graph_data._views.clear()  # what a restart loses
                again = graph_data.graph_view([self.path], filters)
                self.assertEqual(graph_data._views.info()['disk_hits'], 1)
                self.assertEqual(encoded(again), encoded(fresh))
                again['totals'] = 'mutated by one caller'  # each caller owns its copy
                self.assertEqual(encoded(graph_data.graph_view([self.path], filters)), encoded(fresh))
        self.assertEqual(len(self.stored()), len(FILTERS))

    def test_a_restarted_process_serves_the_view_from_disk(self):
        built = graph_data.graph_view([self.path], {'latest': '1'})
        out = self.dir / 'child.json'
        log = open(self.dir / 'child.log', 'wb')
        self.addCleanup(log.close)
        process = subprocess.Popen([sys.executable, str(Path(__file__).resolve()), 'child', str(self.root),
                                    str(self.evidence), str(self.path), str(out)], stdout=log, stderr=subprocess.STDOUT,
                                   cwd=str(Path(__file__).resolve().parent), creationflags=NO_WINDOW)
        self.addCleanup(kill_tree, process)
        process.wait(timeout=180)
        self.assertTrue(out.exists(), (self.dir / 'child.log').read_text(errors='replace'))
        result = json.loads(out.read_text())
        self.assertEqual((result['disk_hits'], result['built']), (1, 0))
        self.assertEqual(result['sha256'], hashlib.sha256(encoded(built)).hexdigest())

    def test_a_changed_workbook_is_a_new_key(self):
        first = graph_data.graph_view([self.path], {})
        self.edit(lambda book: book['2025-2026'].append(['Recipient', 'Civic Horizon Bank', 'Chief', 7, 'Super PAC',
                                                          'FEC:C00000003', 'Yes']))
        self.edit(lambda book: book['2025-2026'].append(['Master', 'Civic Horizon Bank', 'Chief', 7, 'Super PAC',
                                                          None, 'Yes']))
        graph_data._views.clear()
        second = graph_data.graph_view([self.path], {})
        self.assertEqual(graph_data._views.info()['disk_hits'], 0)
        self.assertEqual(second['totals']['giving'], first['totals']['giving'] + 7)
        self.assertEqual(len(self.stored()), 2)

    def test_an_unavailable_activity_view_stays_in_memory_only(self):
        self.edit(lambda book: book['Downstream Activity'].cell(3, 1, '100 source rows condensed to 8 groups'))
        value = graph_data.graph_view([self.path], {'mode': 'activity'})
        self.assertFalse(value['candidate_available'])
        self.assertEqual(self.stored(), [])
        self.assertEqual(graph_data._views.info()['entries'], 1)  # kept in memory, as before
        self.assertEqual(rc.stats(NAMESPACE)['namespaces'][NAMESPACE]['kept_out'], 1)

    def test_code_that_changed_under_the_process_keeps_views_off_disk(self):
        with patch.object(rc, 'code_identity', return_value=None):
            value = graph_data.graph_view([self.path], {})
        self.assertEqual(value['totals']['giving'], 1585)
        self.assertEqual(self.stored(), [])

    def test_no_disk_tier_until_configured(self):
        rc.ROOT = None
        graph_data.graph_view([self.path], {})
        self.assertFalse(self.root.exists())
        self.assertEqual(graph_data._views.info()['disk_builds'], 0)  # never routed through the disk tier

    def test_the_code_identity_covers_graph_data_and_what_it_imports(self):
        # In a fresh interpreter that has imported graph_data only: nothing has put mentat/tools on
        # sys.path yet, and the closure must be the same as in a console that has.
        out = self.dir / 'closure.json'
        log = open(self.dir / 'closure.log', 'wb')
        self.addCleanup(log.close)
        process = subprocess.Popen([sys.executable, str(Path(__file__).resolve()), 'closure', str(out)], stdout=log,
                                   stderr=subprocess.STDOUT, cwd=str(Path(__file__).resolve().parent),
                                   creationflags=NO_WINDOW)
        self.addCleanup(kill_tree, process)
        process.wait(timeout=120)
        self.assertTrue(out.exists(), (self.dir / 'closure.log').read_text(errors='replace'))
        fresh = json.loads(out.read_text())
        self.assertNotIn('contribution_data', fresh['loaded'])
        for found in (fresh['closure'], sorted(rc.module_closure('graph_data'))):
            for name in ('graph_data', 'view_cache', 'response_cache', 'report_insights', 'native_scan', 'downstream',
                         'contribution_data', 'cross_reference_engine'):  # the last is mentat/tools, loaded by contribution_data
                self.assertIn(name, found)
            self.assertNotIn('server', found)
        self.assertEqual(fresh['closure'], sorted(rc.module_closure('graph_data')))


class ContentHashTests(DiskCase):
    def test_same_stamp_rewrite_is_a_miss_after_a_restart(self):
        aux = self.dir / 'input.txt'
        aux.write_text('first')

        def build():
            return {'content': aux.read_text()}
        tiers = {build.__module__: ('test-views', view_cache.graph_view_persistable, lambda: {'test': True})}
        with patch.dict(view_cache.DISK_TIERS, tiers):
            key = ('question', file_signature([aux]))
            self.assertEqual(ViewCache().get(key, build), {'content': 'first'})
            info = aux.stat()
            aux.write_text('other')  # same size
            os.utime(aux, ns=(info.st_atime_ns, info.st_mtime_ns))  # a stamp a same-size rewrite can keep
            self.assertEqual(file_signature([aux]), key[1])  # the stat signature cannot tell
            report_insights.digest_stamp.cache_clear()  # a restarted process has no hash memo
            self.assertEqual(ViewCache().get(key, build), {'content': 'other'})

    def test_builders_outside_disk_tiers_stay_memory_only(self):
        cache = ViewCache()
        self.assertEqual(cache.get(('x', 1), lambda: {'a': 1}), {'a': 1})
        self.assertEqual(cache.get(('x', 1), lambda: {'a': 2}), {'a': 1})
        self.assertFalse(self.root.exists())

    def test_file_signature_is_still_a_tuple(self):
        signature = file_signature([self.dir / 'missing.txt'])
        self.assertIsInstance(signature, tuple)
        self.assertEqual(signature, ((str((self.dir / 'missing.txt').absolute()), None, None, None),))
        self.assertEqual(hash(signature), hash(tuple(signature)))


def child(args):
    root, evidence, path, out = map(Path, args)
    rc.configure(root)
    downstream.ROOT = evidence
    built = []
    real = graph_data._graph_view

    def counting(paths, filters):
        built.append(1)
        return real(paths, filters)
    graph_data._graph_view = counting
    value = graph_data.graph_view([path], {'latest': '1'})
    out.write_text(json.dumps({'disk_hits': graph_data._views.info()['disk_hits'], 'built': len(built),
                               'sha256': hashlib.sha256(encoded(value)).hexdigest()}))


def closure(out):
    Path(out).write_text(json.dumps({'closure': sorted(rc.module_closure('graph_data')),
                                     'loaded': sorted(name for name in ('contribution_data', 'cross_reference_engine')
                                                      if name in sys.modules)}))


if __name__ == '__main__':
    if len(sys.argv) > 1 and sys.argv[1] == 'child':
        child(sys.argv[2:])
    elif len(sys.argv) > 1 and sys.argv[1] == 'closure':
        closure(sys.argv[2])
    else:
        unittest.main()
