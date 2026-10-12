"""/api/report-data through response_cache (report_insights.worksheet_entry, perf lane 2026-10-08).

The route used to re-parse the whole workbook on every lru_cache miss (4-38 s for a Banking sheet), let
the sheet list and each sheet push each other out of a six-entry cache, never cached a refusal (the
over-50,000-row error cost the same read again on every click), built twice for two simultaneous
requests, and re-encoded and re-compressed the answer on every hit.

These tests go through the HTTP route, so they hold for the cache and for the bytes a client gets. Each
one that guards a new behavior fails on the code before this change (checked against a copy of the
pre-change tree: see the lane report).
"""
from concurrent.futures import ThreadPoolExecutor
import gzip
import http.client
import json
from pathlib import Path
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

import openpyxl

APP = Path(__file__).resolve().parents[1] / 'app'
sys.path.insert(0, str(APP))
import response_cache as rc  # noqa: E402
import report_insights as ri  # noqa: E402
import server as app  # noqa: E402

COMMAND = 'python "C:/Mentat/tools/fec_transfer_graph.py" --cycles 2024'


def write_book(path, sheets):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    book.save(path)
    book.close()


class ReportDataRoute(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.report = self.root / 'Example.xlsx'
        saved = rc.ROOT
        self.addCleanup(setattr, rc, 'ROOT', saved)
        self.addCleanup(rc.clear_memory)
        rc.ROOT = None  # memory tier only, unless a test turns the disk tier on
        rc.clear_memory()
        self.run_directory = self.root / 'run'
        self.run_directory.mkdir()
        self.serve(self.report)

    def serve(self, report):
        (self.run_directory / 'status.json').write_text(json.dumps(
            {'industries': [{'status': 'succeeded', 'output_workbook': str(report)}]}))
        console = object.__new__(app.Console)
        console.report_directory = self.root
        console.find = lambda identifier: {'directory': self.run_directory} if identifier == 'test' else None
        server = app.ThreadingHTTPServer(('127.0.0.1', 0), app.Handler)
        server.app = console
        worker = threading.Thread(target=server.serve_forever, daemon=True)
        worker.start()

        def stop():
            server.shutdown()
            server.server_close()
            worker.join()
        self.addCleanup(stop)
        self.port = server.server_port

    def get(self, sheet=None, header=None, headers=None):
        query = 'id=test&index=0' + ('&sheet=' + sheet.replace(' ', '%20') if sheet else '') + \
            (f'&header={header}' if header else '')
        connection = http.client.HTTPConnection('127.0.0.1', self.port, timeout=60)
        try:
            connection.request('GET', '/api/report-data?' + query, headers=headers or {})
            response = connection.getresponse()
            return response.status, dict(response.getheaders()), response.read()
        finally:
            connection.close()

    def data(self, sheet=None, header=None):
        status, _, body = self.get(sheet, header)
        return status, json.loads(body)

    def spy(self):
        """Count every openpyxl workbook load while still performing it."""
        return patch.object(ri.openpyxl, 'load_workbook', wraps=openpyxl.load_workbook)


class RepeatedRequests(ReportDataRoute):
    def test_the_second_request_does_not_read_the_workbook_and_can_revalidate(self):
        write_book(self.report, {'Data': [['Name', 'Value'], ['A', 4]]})
        with self.spy() as loads:
            first = self.get('Data')
            second = self.get('Data')
        self.assertEqual((first[0], second[0], loads.call_count), (200, 200, 1))
        self.assertEqual(first[2], second[2])
        self.assertEqual(json.loads(first[2])['rows'], [['A', 4]])
        tag = first[1]['ETag']
        self.assertEqual(second[1]['ETag'], tag)
        status, headers, body = self.get('Data', headers={'If-None-Match': tag})
        self.assertEqual((status, body), (304, b''))
        self.assertEqual(headers['ETag'], tag)

    def test_the_bytes_are_the_ones_the_route_always_wrote_and_gzip_is_pre_encoded(self):
        rows = [['Name', 'Value']] + [[f'Company {i}', i] for i in range(400)]
        write_book(self.report, {'Data': rows})
        expected = rc.json_body(ri.worksheet_data.__wrapped__(ri.stamp(self.report), 'Data', 0))
        status, headers, body = self.get('Data')
        self.assertEqual((status, body), (200, expected))
        self.assertGreater(len(expected), 2048)
        status, headers, body = self.get('Data', headers={'Accept-Encoding': 'gzip'})
        self.assertEqual(headers['Content-Encoding'], 'gzip')
        self.assertEqual(body, gzip.compress(expected, compresslevel=1, mtime=0))  # Handler.send's own gzip

    def test_the_sheet_list_and_the_sheets_do_not_push_each_other_out(self):
        names = [f'Sheet {i}' for i in range(8)]  # more than the six entries the old cache held
        write_book(self.report, {name: [['Name', 'Value'], [name, i]] for i, name in enumerate(names)})
        with self.spy() as loads:
            for _ in range(2):
                self.assertEqual(self.data()[1]['sheets'], names)
                for i, name in enumerate(names):
                    self.assertEqual(self.data(name)[1]['rows'], [[name, i]])
        self.assertEqual(loads.call_count, len(names))  # one read per sheet, none on the second pass

    def test_simultaneous_requests_for_one_sheet_build_it_once(self):
        write_book(self.report, {'Data': [['Name', 'Value'], ['A', 4]]})
        real = openpyxl.load_workbook

        def slow(*args, **kwargs):
            time.sleep(.5)
            return real(*args, **kwargs)
        with patch.object(ri.openpyxl, 'load_workbook', side_effect=slow) as loads, ThreadPoolExecutor(4) as pool:
            results = list(pool.map(lambda _: self.get('Data'), range(4)))
        self.assertEqual({status for status, _, _ in results}, {200})
        self.assertEqual(len({body for _, _, body in results}), 1)
        self.assertEqual(loads.call_count, 1)

    def test_a_restart_is_answered_from_disk(self):
        write_book(self.report, {'Data': [['Name', 'Value'], ['A', 4]]})
        rc.configure(self.root / 'response-cache')
        built = self.get('Data')
        rc.clear_memory()  # what a restart loses
        with patch.object(ri.openpyxl, 'load_workbook', side_effect=AssertionError('read again')):
            again = self.get('Data')
        self.assertEqual((again[0], again[2]), (200, built[2]))
        self.assertEqual(ri.worksheet_entry(self.report, 'Data').hit, 'memory')  # remembered from the disk hit
        self.assertEqual(len(list((self.root / 'response-cache' / ri.WORKSHEET_NAMESPACE).glob('*.json'))), 1)


class Invalidation(ReportDataRoute):
    def test_a_changed_workbook_is_a_new_answer(self):
        write_book(self.report, {'Data': [['Name', 'Value'], ['A', 4]]})
        before = self.data('Data')[1]
        write_book(self.report, {'Data': [['Name', 'Value'], ['B', 88]]})
        after = self.data('Data')[1]
        self.assertEqual((before['rows'], after['rows']), ([['A', 4]], [['B', 88]]))
        self.assertNotEqual(before['sha256'], after['sha256'])
        self.assertEqual(self.data()[1]['sha256'], after['sha256'])  # the sheet list moves with it

    def test_the_audit_manifest_is_a_dependency_of_the_not_traced_note(self):
        overview = self.root / 'Fixture Industry Spending Overview.xlsx'
        data_book = self.root / 'Fixture Industry Spending Data.xlsx'
        sheets = {'Beneficiary Summary': [['Beneficiary Summary'], ['note'], ['1 records; 1 shown.'], [],
                                          ['Cycle', 'Candidate ID', 'Support Activity'], [2024, 'H0CA00001', 100]]}
        write_book(overview, sheets)
        write_book(data_book, sheets)
        self.serve(data_book)

        def manifest(statuses):
            audit = overview.with_suffix('.audit')
            audit.mkdir(exist_ok=True)
            (audit / 'audit-manifest.json').write_text(json.dumps({
                'version': 1, 'workbook_name': overview.name, 'workbook_sha256': ri.digest(overview), 'csv_sha256': {},
                'transfer_index_status': statuses,
                'transfer_index_notes': {year: f'Not traced: the transfer index for {year} is {status}; rebuild with {COMMAND}'
                                         for year, status in statuses.items() if status != 'current'}}), encoding='utf-8')

        def note():
            status, body = self.data('Beneficiary Summary', header=5)
            self.assertEqual(status, 200, body)
            return body['note']
        self.assertNotIn('not traced', note().lower())  # no bundle yet
        manifest({'2024': 'stale'})
        self.assertIn('Direct only; transfers not traced for 2024', note())
        self.assertIn('Direct only; transfers not traced for 2024', note())  # and now from the cache
        manifest({'2024': 'current'})  # the book itself did not change
        self.assertNotIn('not traced', note().lower())


class Refusals(ReportDataRoute):
    def test_an_over_limit_refusal_is_remembered_and_answers_the_same_400(self):
        write_book(self.report, {'Data': [['Name', 'Value']] + [[f'R{i}', i] for i in range(6)]})
        with patch.object(ri, 'MAX_ROWS', 2), self.spy() as loads:
            first = self.get('Data')
            second = self.get('Data')
            third = self.get('Data', headers={'Accept-Encoding': 'gzip'})
        self.assertEqual((first[0], second[0], third[0]), (400, 400, 400))
        message = json.loads(first[2])['error']
        self.assertIn('exceeds 2 rows', message)
        self.assertIn('no partial chart', message)
        self.assertEqual(second[2], first[2])
        self.assertEqual(json.loads(third[2]), {'error': message})
        self.assertEqual(loads.call_count, 1)

    def test_other_refusals_are_remembered_too(self):
        write_book(self.report, {'Data': [['Name', 'Value'], ['A', 4]]})
        with self.spy() as loads:
            for _ in range(2):
                self.assertEqual(self.data('Missing'), (400, {'error': 'Choose a worksheet from this report.'}))
                self.assertEqual(self.data('Data', header=101), (400, {'error': 'Header row must be between 1 and 100.'}))
        self.assertEqual(loads.call_count, 2)  # one read per refused question

    def test_a_refusal_ends_with_the_workbook_that_caused_it(self):
        write_book(self.report, {'Data': [['Name', 'Value']] + [[f'R{i}', i] for i in range(6)]})
        with patch.object(ri, 'MAX_ROWS', 2):
            self.assertEqual(self.data('Data')[0], 400)
        write_book(self.report, {'Data': [['Name', 'Value'], ['A', 4]]})
        with patch.object(ri, 'MAX_ROWS', 2):
            status, body = self.data('Data')
        self.assertEqual((status, body['rows']), (200, [['A', 4]]))

    def test_a_transient_error_is_not_remembered(self):
        write_book(self.report, {'Data': [['Name', 'Value'], ['A', 4]]})
        real = openpyxl.load_workbook
        calls = []

        def locked_once(*args, **kwargs):
            calls.append(1)
            if len(calls) == 1:
                raise PermissionError('The report is open in Excel.')
            return real(*args, **kwargs)
        with patch.object(ri.openpyxl, 'load_workbook', side_effect=locked_once):
            self.assertEqual(self.data('Data'), (400, {'error': 'The report is open in Excel.'}))
            status, body = self.data('Data')
        self.assertEqual((status, body['rows']), (200, [['A', 4]]))
        self.assertEqual(len(calls), 2)

    def test_a_cached_refusal_is_the_same_exception_for_direct_callers(self):
        write_book(self.report, {'Data': [['Name', 'Value']] + [[f'R{i}', i] for i in range(6)]})
        with patch.object(ri, 'MAX_ROWS', 2):
            messages = []
            for _ in range(2):
                with self.assertRaises(ri.WorksheetRefused) as caught:
                    ri.worksheet_entry(self.report, 'Data')
                messages.append(str(caught.exception))
        self.assertEqual(messages[0], messages[1])
        self.assertIsInstance(caught.exception, ValueError)  # the route's existing handler answers 400


class PrewarmHook(unittest.TestCase):
    def test_after_the_warm_hook_the_route_answers_from_the_cache(self):
        import prewarm
        with tempfile.TemporaryDirectory() as folder:
            report = Path(folder) / 'Example Industry Spending Overview.xlsx'
            write_book(report, {'Data': [['Name', 'Value'], ['A', 4]],
                                'Big': [['Name', 'Value']] + [[f'R{i}', i] for i in range(6)]})
            row = {'index': 7, 'name': report.name, 'industry': 'Example', 'kind': 'overview'}
            stub = type('Stub', (), {'published_reports': lambda self: [row, {**row, 'industry': 'Other'}],
                                     'report_path': lambda self, identifier, index: report})()
            backend = type('Backend', (), {
                'report_directory': Path(folder), '_console_stub': lambda self: stub,
                'snapshot_registry': type('Registry', (), {'registry': staticmethod(lambda directory: {})}),
                'industry_scope': type('Scope', (), {'label_or_value': staticmethod(lambda registry, value: value)})})()
            saved = rc.ROOT
            self.addCleanup(setattr, rc, 'ROOT', saved)
            self.addCleanup(rc.clear_memory)
            rc.ROOT = None
            rc.clear_memory()
            self.assertIn('worksheets', prewarm.AppBackend.artifact_hooks(backend))
            self.assertIn('worksheets', prewarm.ARTIFACT_KINDS)
            with patch.object(ri, 'MAX_ROWS', 2):
                info = prewarm.AppBackend.warm_worksheets(backend, 'Example')
                self.assertEqual((info['sheets'], info['refused']), (1, 1))  # Big is refused, cached and skipped
                with patch.object(ri.openpyxl, 'load_workbook', side_effect=AssertionError('read again')):
                    self.assertEqual(ri.worksheet_entry(report).hit, 'memory')
                    self.assertEqual(ri.worksheet_entry(report, 'Data').hit, 'memory')
                    with self.assertRaises(ri.WorksheetRefused):
                        ri.worksheet_entry(report, 'Big')


if __name__ == '__main__':
    unittest.main()
