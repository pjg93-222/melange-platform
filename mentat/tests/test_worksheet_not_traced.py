"""F2 and F4 (Paul, 2026-09-25): a worksheet view says when its figures are not traced.

With a stale transfer index, Transfer Paths, Allocation Estimates, Allocation
Coverage and Candidate Funding publish nothing transfer-derived for that cycle,
and Beneficiary Summary, Downstream Reach and Downstream Measurement fall back to
directly funded committees only. report_insights.worksheet_data puts the
sentence in the note Chart Studio shows beside the data. Negative controls: a
current index, a legacy bundle, no bundle, and a sheet that is not transfer-derived.
"""
from pathlib import Path
import json
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'app'))
from report_insights import digest, stamp, worksheet_data  # noqa: E402

COMMAND = 'python "C:/Mentat/tools/fec_transfer_graph.py" --cycles 2024'


def workbook(path):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, headers, row in [
            ('Beneficiary Summary', ['Cycle', 'Candidate ID', 'Support Activity'], [2024, 'H0CA00001', 100]),
            ('Allocation Estimates', ['Cycle', 'Funded Committee', 'Conditional Estimate'], [2022, 'C00000001', 5]),
            ('2023-2024', ['Company', 'Amount'], ['Acme', 10])]:
        sheet = book.create_sheet(name)
        for line in [[name], ['note'], ['1 records; 1 shown.'], [], headers, row]:
            sheet.append(line)
    book.save(path)
    book.close()


def bundle(path, statuses):
    audit = path.with_suffix('.audit')
    audit.mkdir(exist_ok=True)
    manifest = {'version': 1, 'workbook_name': path.name, 'workbook_sha256': digest(path), 'csv_sha256': {}}
    if statuses is not None:
        manifest['transfer_index_status'] = statuses
        manifest['transfer_index_notes'] = {year: f'Not traced: the transfer index for {year} is {status}; rebuild with {COMMAND}'
                                            for year, status in statuses.items() if status != 'current'}
    (audit / 'audit-manifest.json').write_text(json.dumps(manifest), encoding='utf-8')


class WorksheetNotes(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / 'Fixture Industry Spending Overview.xlsx'
        workbook(self.path)
        worksheet_data.cache_clear()
        self.addCleanup(worksheet_data.cache_clear)

    def note(self, sheet):
        worksheet_data.cache_clear()
        return worksheet_data(stamp(self.path), sheet, 5)['note']

    def test_a_fallback_sheet_is_labelled_direct_only_for_the_stale_cycle(self):
        bundle(self.path, {'2022': 'current', '2024': 'stale'})
        note = self.note('Beneficiary Summary')
        self.assertIn('Direct only; transfers not traced for 2024: the transfer index for 2024 is stale.', note)
        self.assertNotIn('2022', note)

    def test_a_transfer_derived_sheet_says_not_traced_and_that_blank_is_not_zero(self):
        bundle(self.path, {'2024': 'missing'})
        note = self.note('Allocation Estimates')
        self.assertIn('Not traced: the transfer index for 2024 is missing: 2024 has no transfer-derived rows '
                      'or values here; blank is not zero.', note)
        self.assertNotIn('rebuild with', note)

    def test_the_companion_data_workbook_reads_its_overview_s_bundle(self):
        bundle(self.path, {'2024': 'stale'})
        data = self.path.with_name('Fixture Industry Spending Data.xlsx')
        workbook(data)
        worksheet_data.cache_clear()
        self.assertIn('Direct only; transfers not traced for 2024',
                      worksheet_data(stamp(data), 'Beneficiary Summary', 5)['note'])

    def test_negative_controls_say_nothing(self):
        for statuses in ({'2024': 'current'}, None):
            with self.subTest(statuses=statuses):
                bundle(self.path, statuses)
                self.assertNotIn('not traced', self.note('Beneficiary Summary').lower())
                self.assertNotIn('Not traced', self.note('Allocation Estimates'))
        bundle(self.path, {'2024': 'stale'})
        self.assertNotIn('not traced', self.note('2023-2024').lower())

    def test_no_bundle_says_nothing(self):
        self.assertNotIn('not traced', self.note('Beneficiary Summary').lower())

    def test_a_split_part_of_the_data_workbook_carries_its_sheet_s_note(self):
        # S1 (Paul, 2026-09-26): the Data workbook names a per-cycle part
        # `<Sheet> <election year>`; books published before keep
        # `<Sheet> <cycle label>` until rebuilt. Either spelling is its sheet.
        bundle(self.path, {'2024': 'stale'})
        data = self.path.with_name('Fixture Industry Spending Data.xlsx')
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name in ('Allocation Estimates 2024', 'Allocation Estimates 2023-2024',
                     'Downstream Measurement 2024', 'Downstream Measurement 2023-2024',
                     'Giving Facts 2024'):
            sheet = book.create_sheet(name)
            sheet.append(['Cycle', 'Amount'])
            sheet.append([2024, 5])
        book.save(data)
        book.close()

        def note(name):
            worksheet_data.cache_clear()
            return worksheet_data(stamp(data), name, 1)['note']

        for name in ('Allocation Estimates 2024', 'Allocation Estimates 2023-2024'):
            self.assertIn('Not traced: the transfer index for 2024 is stale: 2024 has no '
                          'transfer-derived rows or values here; blank is not zero.', note(name), name)
        for name in ('Downstream Measurement 2024', 'Downstream Measurement 2023-2024'):
            self.assertIn('Direct only; transfers not traced for 2024', note(name), name)
        self.assertNotIn('not traced', note('Giving Facts 2024').lower())


if __name__ == '__main__':
    unittest.main()
