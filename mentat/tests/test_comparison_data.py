"""Independent subject partitions and person IDs, against report-shaped fixtures."""
import sys
import tempfile
from pathlib import Path
import unittest
from unittest.mock import patch
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'app'))
from comparison_data import subjects, flows, report_facts
from graph_fixture import add_graph_sheets


class ComparisonTests(unittest.TestCase):
    def fixture(self, tmp, second=True):
        path=Path(tmp)/'flows.xlsx';book=openpyxl.Workbook();add_graph_sheets(book)
        if second:
            book['Downstream Activity']['A3']='101 source rows condensed to 5 committee-candidate-activity groups'
            book['Downstream Activity'].append(['2025-2026','C00000003','H0CA00002','Candidate Two','R','Support',250])
        book.save(path);book.close();return path

    def test_source_candidate_relationships_keep_channels_and_refunds_separate(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=self.fixture(tmp);before=path.read_bytes();result=flows([path],'Sample')
            rows=result['rows'];candidate='Candidate:H0CA00001'
            company=next(r for r in rows if r['source']=='company:Harborlight Credit Union' and r['destination']==candidate)
            self.assertEqual(company['amount'],100)
            self.assertEqual(company['channel'],'Company-associated personal receipts')
            refund=next(r for r in rows if r['source']=='company:Civic Horizon Bank' and r['destination']==candidate)
            self.assertEqual(refund['amount'],-10)
            pac=[r for r in rows if r['source']=='committee:C00000003' and r['destination']==candidate]
            self.assertEqual({r['channel']:r['amount'] for r in pac},{'Independent spending · support':777,'Independent spending · opposition':999,'Committee contribution / other':111})
            self.assertEqual(sum(r['amount'] for r in rows if r['source']=='committee:C00000003' and r['destination']=='Candidate:H0CA00002'),250)
            self.assertEqual(len([s for s in result['sources'] if s['id']=='committee:C00000003']),1)
            self.assertTrue(any(d['category']=='Super PAC' and d['id']=='Recipient:C00000003' for d in result['destinations']))
            self.assertEqual(path.read_bytes(),before)
            with self.assertRaisesRegex(ValueError,'overlap'):flows([path,path],'Sample')

    def test_partial_activity_retains_receipts_but_never_publishes_partial_pac_totals(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=self.fixture(tmp);book=openpyxl.load_workbook(path);book['Downstream Activity']['A3']='condensed to 500 groups';book.save(path);book.close()
            result=flows([path],'Sample')
            self.assertTrue(result['issues'])
            self.assertEqual(result['unavailable_activity_cycles'],['2026'])
            self.assertTrue(any(s['id']=='committee:C00000003' for s in result['sources']))
            self.assertTrue(result['rows'])
            self.assertTrue(all(r['source'].startswith('company:') for r in result['rows']))

    def test_report_changed_between_receipts_and_activity_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=self.fixture(tmp)
            def changed(_):
                path.write_bytes(path.read_bytes()+b'changed')
                return []
            with patch('comparison_data.report_activity',side_effect=changed):
                with self.assertRaisesRegex(ValueError,'report changed'):flows([path],'Sample')

    def test_subjects_do_not_double_count_summary_families(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'sample.xlsx'
            book=openpyxl.Workbook();add_graph_sheets(book)
            sheet=book['2025-2026'];sheet.cell(4,8,'Person Key')
            sheet.append(['Employee','Harborlight Credit Union','Chief',100,None,None,'Yes','PERSON:1'])
            people=book.create_sheet('Employee List')
            for row in [[],[],[],[],['Person Key','Person','Company'],['PERSON:1','Alex Example','Harborlight Credit Union']]:people.append(row)
            book.save(path);book.close();before=path.read_bytes()
            result=subjects([path],'Credit Unions')
            industry=next(s for s in result['subjects'] if s['kind']=='Industry')
            self.assertEqual(industry['periods']['2026']['giving'],1585)
            person=next(s for s in result['subjects'] if s['kind']=='Person')
            self.assertEqual(person['label'],'Alex Example')
            self.assertEqual(person['periods']['2026'],{'giving':100})
            self.assertEqual(path.read_bytes(),before)
            with self.assertRaisesRegex(ValueError,'overlap'):subjects([path,path],'Credit Unions')

    def test_subject_report_changed_during_comparison_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self.fixture(tmp)
            def changed(version):
                result = report_facts(version)
                path.write_bytes(path.read_bytes() + b'changed')
                return result
            with patch('comparison_data.report_facts', side_effect=changed):
                with self.assertRaisesRegex(ValueError, 'report changed'):
                    subjects([path], 'Sample')


if __name__=='__main__':unittest.main()
