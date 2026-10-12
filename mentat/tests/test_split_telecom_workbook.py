import importlib.util
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table


TOOLS = Path(__file__).resolve().parents[1] / "tools"
SPEC = importlib.util.spec_from_file_location(
    "split_telecom_workbook",
    TOOLS / "split_telecom_workbook.py",
)
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


def test_make_tables_excel_safe_normalizes_strict_excel_rules():
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    summary.append(["Company", "Total"])
    summary.append(["Example", 100])
    summary.add_table(Table(displayName="SummaryTable", ref="A1:B2"))
    summary.auto_filter.ref = "A1:B2"

    overview = book.create_sheet("Overview")
    overview.append([None, None, None])
    overview.append([None, "D/R", "D/R"])
    overview.append([None, 40, 60])
    overview_table = Table(displayName="OverviewTable", ref="B2:C3")
    overview.add_table(overview_table)
    overview_table._initialise_columns()
    overview_table.tableColumns[0].id = 2
    overview_table.tableColumns[1].id = 3

    empty = book.create_sheet("Empty")
    empty.append(["Recipient", "Total"])
    empty.add_table(Table(displayName="EmptyTable", ref="A1:B1"))

    MODULE.make_tables_excel_safe(book)

    assert summary.auto_filter.ref is None
    assert [column.id for column in overview_table.tableColumns] == [1, 2]
    assert [column.name for column in overview_table.tableColumns] == [
        "D/R",
        "D/R (2)",
    ]
    assert overview["B2"].value == "D/R"
    assert overview["C2"].value == "D/R (2)"
    assert empty.tables["EmptyTable"].ref == "A1:B2"
