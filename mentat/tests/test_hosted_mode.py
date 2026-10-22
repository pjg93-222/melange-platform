"""Hosted, read-only, Access-gated Mentat console (Lane M).

Default mode must be what it always was: a 127.0.0.1 bind, the loopback Host and Origin allowlists, the
per-process session token, no new routes, no new headers. Enforced mode (Cloudflare Access) must be read-only,
must refuse everything that exports, downloads, uploads or reads a path, and must never hand out the console's
own token. MENTAT_STATE_WRITE_ROOT must move every cache the console writes; receipt paths written on Windows
must resolve on Linux.

Nothing here touches the real data root, the real state folder or a real port: servers listen on 127.0.0.1:0
over stand-in consoles in temp dirs, and the start-up tests run server.py on an OS-chosen port with every
path in a temp dir. RSA keys are generated at run time with `cryptography` and signed in memory; only the
public JWKS is written (to a temp dir). The enforced-mode tests skip without `cryptography`.
"""
import base64
import contextlib
from decimal import Decimal
import hashlib
import http.client
import inspect
import io
import json
import os
from pathlib import Path
import re
import socket
import subprocess
import sys
import tempfile
import textwrap
import threading
import time
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch
from urllib.parse import quote

try:
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.asymmetric import padding, rsa
    HAVE_CRYPTO = True
except ImportError:
    HAVE_CRYPTO = False

try:
    import openpyxl
except ImportError:
    openpyxl = None

from test_contribution_http import app

import contribution_data
import contribution_source_guard
import evidence_csv
import receipt_path as receipt_paths
import response_cache
import source_memo

gate = app.access_gate
needs_crypto = unittest.skipUnless(HAVE_CRYPTO, 'cryptography is not installed')
needs_openpyxl = unittest.skipUnless(openpyxl, 'openpyxl is not installed')

AUD = 'mentat-test-aud'
ISS = 'https://example-team.cloudflareaccess.com'
HOST = 'mentat.example.test'
EMAIL = 'analyst@example.test'
REAL_TOKEN = 'real-console-token-never-shown'
ANALYTICAL = {'/api/contributions/query', '/api/contributions/records', '/api/contributions/geography',
              '/api/insights/query', '/api/insights/records', '/api/forecast/project'}
SERVER_PY = Path(app.__file__).resolve()
# deploy/access-gate: fec/access_gate.py. The three apps carry byte-identical copies; change one, recopy, update this.
GATE_SHA256 = 'dc7ef05085044601884ce037f21c1b44d774e7dffba3a0276a3ea5423227e58d'


def b64u(raw):
    return base64.urlsafe_b64encode(raw).rstrip(b'=').decode('ascii')


PERSON_SHEETS = ('Employee List', 'Receipts', 'Names')
AGGREGATE_SHEETS = ('Giving Summary', 'Tier Totals')


def synthetic_report(path):
    """A small report workbook built at test time, fake values only: two aggregate sheets and three person-level
    ones, each with a title block above its header row, as the published reports have."""
    book = openpyxl.Workbook()
    del book['Sheet']
    for name, header, rows in (
            ('Giving Summary', ['Company', 'Cycle', 'Total giving'], [['Example Co', '2024', 100], ['Other Co', '2024', 50]]),
            ('Tier Totals', ['Contributor Tier', 'Distinct Contributors', 'Amount'], [['Employee', 3, 10], ['Chief', 1, 5]]),
            ('Employee List', ['Person', 'Person Key', 'Titles', 'Total'], [['Test Person A', 'k1', 'Engineer', 5], ['Test Person B', 'k2', 'Manager', 7]]),
            ('Receipts', ['Date', 'Contributor', 'Street 1', 'City', 'ZIP', 'Employer', 'Occupation', 'Amount'],
             [['2024-01-02', 'Test Person A', '1 Test St', 'Testville', '00000', 'Example Co', 'Engineer', 5]] * 2),
            ('Names', ['Name', 'Amount'], [['Test Person A', 5], ['Test Person B', 7]])):
        sheet = book.create_sheet(name)
        sheet.append([name + ' (synthetic)'])
        sheet.append([])
        sheet.append(header)
        for row in rows:
            sheet.append(row)
    book.save(path)


def comparison_body(directory):
    """A real /api/comparison body (comparison_data.subjects) over a report-shaped workbook that has one employee,
    fake values only."""
    import comparison_data
    from graph_fixture import add_graph_sheets
    path = Path(directory) / 'comparison-report.xlsx'
    book = openpyxl.Workbook()
    add_graph_sheets(book)
    book['2025-2026'].cell(4, 8, 'Person Key')
    book['2025-2026'].append(['Employee', 'Harborlight Credit Union', 'Chief', 100, None, None, 'Yes', 'PERSON:1'])
    people = book.create_sheet('Employee List')
    for row in [[], [], [], [], ['Person Key', 'Person', 'Company'], ['PERSON:1', 'Alex Example', 'Harborlight Credit Union']]:
        people.append(row)
    book.save(path)
    return json.loads(json.dumps(comparison_data.subjects([path], 'Credit Unions')))


def header_workbook(path, headers):
    """One sheet per header (S0, S1, ...), each with a title block above a header row of Company, that header and
    Amount, and one fake data row."""
    book = openpyxl.Workbook()
    del book['Sheet']
    for number, header in enumerate(headers):
        sheet = book.create_sheet(f'S{number}')
        for row in ([f'S{number} (synthetic)'], [], ['Company', header, 'Amount'], ['Example Co', 'x', 1]):
            sheet.append(row)
    book.save(path)


SUMMARY_HEADER = ['Cycle', 'Section', 'Rank', 'Company', 'Name / Recipient', 'Relationship', 'Total / Activity', 'Reference Key']
# (Name / Recipient and Reference Key are columns 4 and 7.) Rows of a '* Summary' or '<Company> View' sheet, fake values
# only: two contributors, the three recipient and aggregate sections, a section nobody knows, a blank one and a
# near-miss spelling of a contributor's.
SUMMARY_ROWS = [
    ['2024', 'Top Contributor', 1, 'Example Co', 'Zed Quill', 'Person Key', 500, 'key-zed'],
    ['2024', 'Top Contributor', 2, 'Example Co', 'Ida Marsh', 'Person Key', 400, 'key-ida'],
    ['2024', 'Top Direct Recipient', 1, 'Example Co', 'Example PAC', 'Committee', 300, 'key-pac'],
    ['2024', 'Top Downstream Beneficiary', 1, 'Example Co', 'Sample Candidate', 'Candidate', 200, 'key-cand'],
    ['2024', 'Company Total', None, 'Example Co', 'Example Co', 'Company', 1400, 'key-co'],
    ['2024', 'Top Mystery Section', 1, 'Example Co', 'Unknown Kind', 'Other', 100, 'key-new'],
    ['2024', None, 1, 'Example Co', 'Blank Section', 'Other', 50, 'key-blank'],
    ['2024', 'top contributor', 1, 'Example Co', 'Lower Case', 'Other', 25, 'key-lower']]
SUMMARY_KEPT = (2, 3, 4)                       # the indexes of SUMMARY_ROWS a hosted viewer who is not listed still sees whole
SUMMARY_HIDDEN = ('Zed Quill', 'key-zed', 'Ida Marsh', 'key-ida', 'Unknown Kind', 'key-new', 'Blank Section', 'key-blank',
                  'Lower Case', 'key-lower')
SUMMARY_SHEETS = {   # sheet name -> (header, rows)
    'Example Co View': (SUMMARY_HEADER, SUMMARY_ROWS),
    '2024 Summary': (SUMMARY_HEADER, SUMMARY_ROWS[:3]),
    'Twice View': (['Section', 'Name / Recipient', 'Name / Recipient', 'Reference Key', 'Total'],
                   [['Top Contributor', 'Zed Quill', 'Zed Q', 'key-zed', 5], ['Company Total', 'Example Co', 'Example Co', 'key-co', 9]]),
    'Plain View': (['Section', 'Company', 'Total'], [['Company Total', 'Example Co', 9]]),
    'Employer View': ([*SUMMARY_HEADER, 'Employer'], [[*row, 'Example Co'] for row in SUMMARY_ROWS]),
    'Two Sections View': (['Section', 'Name / Recipient', 'Section', 'Total'], [['Top Contributor', 'Zed Quill', 'Top Contributor', 5]]),
    'No Section Sheet': (['Company', 'Name / Recipient', 'Reference Key', 'Total'], [['Example Co', 'Zed Quill', 'key-zed', 5]])}


def summary_workbook(path):
    """SUMMARY_SHEETS as a workbook, each sheet with a title block above its header row (row 3; data from row 4)."""
    book = openpyxl.Workbook()
    del book['Sheet']
    for name, (header, rows) in SUMMARY_SHEETS.items():
        sheet = book.create_sheet(name)
        for row in ([name + ' (synthetic)'], [], header, *rows):
            sheet.append(row)
    book.save(path)


def blanked(row, *columns):
    return [None if index in columns else cell for index, cell in enumerate(row)]


class Idp:
    """A throw-away Access: one RSA key, its JWKS file in `directory`, and a way to sign tokens."""

    def __init__(self, directory):
        self.key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        numbers = self.key.public_key().public_numbers()
        big = lambda n: n.to_bytes((n.bit_length() + 7) // 8, 'big')
        self.jwks = Path(directory) / 'certs.json'
        self.jwks.write_text(json.dumps({'keys': [{'kty': 'RSA', 'kid': 'kid-a', 'use': 'sig', 'alg': 'RS256',
                                                   'n': b64u(big(numbers.n)), 'e': b64u(big(numbers.e))}]}))

    def token(self, **over):
        now = int(time.time())
        claims = {'aud': [AUD], 'iss': ISS, 'exp': now + 3600, 'nbf': now - 10, 'email': EMAIL, **over}
        encode = lambda value: b64u(json.dumps(value).encode())
        signing = encode({'alg': 'RS256', 'kid': 'kid-a', 'typ': 'JWT'}) + '.' + encode(claims)
        signature = self.key.sign(signing.encode('ascii'), padding.PKCS1v15(), hashes.SHA256())
        return signing + '.' + b64u(signature)

    def env(self, **over):
        return {'ACCESS_ENFORCE': '1', 'ALLOWED_HOSTS': HOST, 'ACCESS_AUD': AUD, 'ACCESS_ISSUER': ISS,
                'ACCESS_JWKS_FILE': str(self.jwks), **over}


class Served:
    """A Handler class on 127.0.0.1:0 in a thread, over a stand-in console."""

    def __init__(self, handler, console):
        self.server = app.ThreadingHTTPServer(('127.0.0.1', 0), handler)
        self.server.app = console
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.port = self.server.server_port
        self.stderr = ''

    def close(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=5)

    def request(self, path, *, method='GET', data=None, headers=None, raw=None):
        """(status, headers, body); body is parsed JSON, or bytes. self.stderr is what the server wrote."""
        connection = http.client.HTTPConnection('127.0.0.1', self.port, timeout=10)
        body = raw if raw is not None else json.dumps(data).encode() if data is not None else None
        sent = {'Host': f'127.0.0.1:{self.port}'}
        if method == 'POST':
            sent['Content-Type'] = 'application/json'
        sent.update(headers or {})
        captured = io.StringIO()
        try:
            with contextlib.redirect_stderr(captured):
                connection.request(method, path, body=body, headers=sent)
                response = connection.getresponse()
                content = response.read()
            parsed = json.loads(content) if response.getheader('Content-Type', '').startswith('application/json') else content
            return response.status, dict(response.getheaders()), parsed
        finally:
            connection.close()
            self.stderr = captured.getvalue()


def free_port():
    with socket.socket() as probe:
        probe.bind(('127.0.0.1', 0))
        return probe.getsockname()[1]


def keys_of(value):
    """Every dict key at any depth."""
    if isinstance(value, dict):
        for key, item in value.items():
            yield key
            yield from keys_of(item)
    elif isinstance(value, list):
        for item in value:
            yield from keys_of(item)


def flow_body():
    """A /api/flow-shaped body, fake values only. Its first link carries the real evidence record
    (downstream.filing_record through flow_data._link_out); its second, and a nested copy, carry the person-level
    keys a record would hold if _link_out's list of keys grew."""
    from decimal import Decimal
    import downstream
    import flow_data
    record = downstream.filing_record({
        'sub_id': '1', 'transaction_id': 'T1', 'image_number': '202401019999', 'contribution_receipt_amount': 250,
        'contribution_receipt_date': '2024-01-02', 'contributor_name': 'TEST PERSON A', 'contributor_city': 'TESTVILLE',
        'contributor_state': 'ZZ', 'contributor_zip': '00000', 'contributor_employer': 'EXAMPLE CO',
        'expenditure_description': 'Test purpose'})
    link = flow_data._link_out({'source': 'donor:A', 'target': 'committee:C00000001', 'value': Decimal('250'),
                                'kind': 'receipt', 'rows': 1, 'evidence': [record], 'committee': 'C00000001'})
    wider = {**record, 'contributor': 'TEST PERSON A', 'street_1': '1 Test St', 'street_2': 'Apt 2', 'zip5': '00000',
             'raw_zip': '00000', 'occupation': 'Engineer', 'reported_memo_text': 'memo', 'reported_filing_form': 'F3'}
    return {'version': 1,
            'nodes': [{'id': 'committee:C00000001', 'label': 'Example PAC', 'kind': 'committee',
                       'committee_id': 'C00000001', 'state': 'ZZ', 'value_in': 250.0, 'value_out': 250.0}],
            'links': [link, {**link, 'evidence': [wider]}], 'routes': [{'committee': 'C00000001', 'value': 250.0}],
            'totals': {'in': 250.0}, 'coverage': {'through': '2024-12-31'}, 'deep': {'records': [wider]}}


def strings_of(value):
    """Every dict key and string value at any depth."""
    if isinstance(value, dict):
        for key, item in value.items():
            yield key
            yield from strings_of(item)
    elif isinstance(value, list):
        for item in value:
            yield from strings_of(item)
    elif isinstance(value, str):
        yield value


def individuals_flow_body():
    """A /api/flow-shaped body with named individuals among its donors, fake values only: two individuals (one typed
    IND, one untyped "LAST, FIRST", the two kinds flow_data draws as people), a company, a PAC and the unitemized node
    (none of them people), an "Other" bucket that lists named donors, and a refunded sink whose text names a
    contributor. Every receipt link carries a real evidence record whose purpose names a payee."""
    from decimal import Decimal
    import downstream
    import flow_data

    def receipt(source, value, name):
        record = downstream.filing_record({
            'sub_id': '1', 'transaction_id': 'T1', 'image_number': '202401019999', 'contribution_receipt_amount': value,
            'contribution_receipt_date': '2024-01-02', 'contributor_name': name, 'contributor_city': 'TESTVILLE',
            'contributor_state': 'ZZ', 'contributor_zip': '00000', 'expenditure_description': f'Paid {name} for consulting'})
        return flow_data._link_out({'source': source, 'target': 'committee:C00000001', 'value': Decimal(str(value)),
                                    'kind': 'receipt', 'rows': 1, 'evidence': [record]})

    donor = lambda ident, label, value, **fields: {'id': 'donor:' + ident, 'stage': 0, 'label': label, 'kind': 'donor',
                                                   'contribution': True, 'value_in': 0.0, 'value_out': value, **fields}
    return {
        'version': 1,
        'nodes': [
            {'id': 'committee:C00000001', 'stage': 1, 'label': 'Example PAC', 'kind': 'committee', 'root': True,
             'committee_id': 'C00000001', 'name_as_filed': 'EXAMPLE PAC', 'state': 'ZZ', 'value_in': 1000.37,
             'value_out': 1000.37},
            donor('TEST PERSON A', 'Test Person A', 100.1, name_as_filed='TEST PERSON A', entity='IND', held=5.5,
                  held_kinds=['stock']),
            donor('PERSON JANE', 'Jane Person', 200.2, name_as_filed='PERSON, JANE'),
            donor('EXAMPLE CORP', 'Example Corp', 300.03, name_as_filed='EXAMPLE CORP', entity='ORG'),
            donor('SAMPLE PAC', 'Sample PAC', 150.0, name_as_filed='SAMPLE PAC', entity='PAC'),
            donor('unitemized', 'Unitemized contributions', 50.04, unitemized=True),
            {'id': 'other:0', 'stage': 0, 'label': 'All other contributors', 'kind': 'bucket', 'bucket': 'other',
             'member_count': 2, 'value_in': 0.0, 'value_out': 200.0,
             'members': [{'id': 'donor:TEST PERSON C', 'label': 'Test Person C', 'kind': 'donor', 'value': 120.0},
                         {'id': 'donor:TEST PERSON D', 'label': 'Test Person D', 'kind': 'donor', 'value': 80.0}]},
            {'id': 'balance:refunded:1', 'stage': 1, 'label': 'Refunded to donors', 'kind': 'balance',
             'balance': 'refunded', 'value_in': 50.0, 'value_out': 0.0,
             'details': ["$50.00 in negative receipts from 1 contributor beyond their receipts in this window: "
                         "Test Person E $50.00 (line 17)"]}],
        'links': [receipt('donor:TEST PERSON A', 100.1, 'TEST PERSON A'), receipt('donor:PERSON JANE', 200.2, 'PERSON, JANE'),
                  receipt('donor:EXAMPLE CORP', 300.03, 'EXAMPLE CORP'), receipt('donor:SAMPLE PAC', 150.0, 'SAMPLE PAC'),
                  {'source': 'donor:unitemized', 'target': 'committee:C00000001', 'value': 50.04, 'kind': 'receipt', 'rows': 0},
                  receipt('other:0', 200.0, 'TEST PERSON C'),
                  {'source': 'committee:C00000001', 'target': 'balance:refunded:1', 'value': 50.0, 'kind': 'operating',
                   'rows': 0}],
        'routes': [{'committee': 'C00000001', 'value': 50.0}], 'totals': {'donors': 1000.37, 'refunds_to_donors': 50.0},
        'coverage': {'through': '2024-12-31'},
        'notes': [{'key': 'refunds_to_donors', 'text': 'Refunds to donors $50.00: 1 contributor whose receipts in this '
                                                       'window net below zero (negative entries as filed); not drawn as donors'}]}


CACHED_BODY = {'sheet': 'Sources', 'rows': [['Company', '/srv/mentat/report.xlsx', 5], ['Other', 'C:\\Users\\x\\r.xlsx', 6]],
               'file': '/srv/y', 'note': "[Errno 2] No such file or directory: 'x'", 'ok': 'fine/relative', 'n': 3}
GEOGRAPHY_RESULT = {'points': [], 'layers': {}, 'coverage': {}, 'warnings': ['ZIP locations are unavailable.'],
                    'errors': {'points': "[Errno 2] No such file or directory: '/srv/geo.json'",
                               'states': "HTTPSConnectionPool(host='www2.census.gov', port=443)"}}


class AccessGateCopy(unittest.TestCase):
    def test_the_copy_is_the_canonical_gate_byte_for_byte(self):
        self.assertEqual(hashlib.sha256(Path(gate.__file__).read_bytes()).hexdigest(), GATE_SHA256)


class ServerPaths(unittest.TestCase):
    def test_absolute_paths_are_omitted_and_the_rest_is_kept(self):
        body = {'a': 1, 'b': 'C:\\x\\y', 'c': 'c:/x', 'd': '\\\\host\\share', 'e': '/srv/x', 'f': 'relative/x',
                'g': ['/srv/y', 'ok', {'h': '/z', 'i': 2}], 'j': {'k': 'D:\\q'}}
        self.assertEqual(app.without_server_paths(body), {'a': 1, 'f': 'relative/x', 'g': ['ok', {'i': 2}], 'j': {}})

    def test_a_worksheet_row_keeps_its_columns_when_a_cell_is_a_path(self):
        body = {'rows': [['A', '/srv/x', 5], ['B', 'C:\\x', 6]], 'file': '/srv/y', 'ok': 'a/b'}
        self.assertEqual(app.without_server_paths(body, True), {'rows': [['A', '', 5], ['B', '', 6]], 'ok': 'a/b'})

    def test_os_error_text_is_one_sentence(self):
        body = {'a': "[Errno 2] No such file or directory: 'x'", 'b': ['[WinError 3] The system cannot find the path', 'ok'],
                'c': {'d': 'FileNotFoundError: [Errno 2] gone', 'n': 1}, 'e': 'a [text] in brackets'}
        self.assertEqual(app.without_os_error_text(body), {
            'a': app.HOSTED_IO_ERROR, 'b': [app.HOSTED_IO_ERROR, 'ok'], 'c': {'d': app.HOSTED_IO_ERROR, 'n': 1},
            'e': 'a [text] in brackets'})

    def test_only_an_explicit_organization_or_committee_keeps_its_flow_donor_name(self):
        donor = lambda entity: {'kind': 'donor', 'id': 'donor:X', 'name_as_filed': 'ACME HARDWARE', 'entity': entity}
        for entity in ('ORG', 'COM', 'PAC', 'PTY', 'CCM', 'CAN', 'ORG_PAC', 'org'):
            self.assertFalse(app.is_individual_donor(donor(entity)), entity)
        for entity in ('IND', '', None, 'XYZ', 'OTHER'):    # a person, no type, a type nothing here knows
            self.assertTrue(app.is_individual_donor(donor(entity)), entity)
        self.assertFalse(app.is_individual_donor({**donor(''), 'unitemized': True}))
        self.assertFalse(app.is_individual_donor({**donor(''), 'kind': 'committee', 'id': 'committee:C00000001'}))

    def test_a_person_subject_query_is_read_from_the_subject_and_kind_keys_only(self):
        person = json.dumps(['Person', ['Harborlight Credit Union', 'PERSON:1']])
        for query in ({'kind': ['Person']}, {'kind': [' person ']}, {'subject': [person]}, {'subjects': ['x', person]},
                      {'anything': ['["person", ["c", "p"]]']}):
            self.assertTrue(app.names_person_subject(query), query)
        company = json.dumps(['Company', 'Harborlight Credit Union'])
        for query in ({}, {'id': ['Person'], 'industry': ['Person']}, {'kind': ['Company']}, {'measure': ['personal']},
                      {'subject': [company]}):
            self.assertFalse(app.names_person_subject(query), query)

    def test_only_a_json_cached_body_is_scrubbed_when_hosted(self):
        class Fake:
            access_cfg = SimpleNamespace(enforced=True)

            def __init__(self):
                self.headers, self.wfile = {}, io.BytesIO()

            send_response = send_header = lambda self, *_: None
            end_headers = lambda self: None

        entry = response_cache.from_value(CACHED_BODY)
        text, data = Fake(), Fake()
        app.Handler.send_entry(text, entry, content_type='text/plain')
        app.Handler.send_entry(data, entry)
        self.assertEqual(text.wfile.getvalue(), entry.body)
        self.assertEqual(json.loads(data.wfile.getvalue())['rows'], [['Company', '', 5], ['Other', '', 6]])
        local = Fake()
        local.access_cfg = SimpleNamespace(enforced=False)
        app.Handler.send_entry(local, entry)
        self.assertEqual(local.wfile.getvalue(), entry.body, 'a local console sends the entry untouched')


class DefaultModeIsUnchanged(unittest.TestCase):
    """No Access environment: every local contract holds, and nothing hosted leaks in."""

    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.console = SimpleNamespace(
            token='local-token', read_only=False, report_directory=Path(cls.temp.name),
            configured_root=cls.temp.name, root=Path(cls.temp.name),
            evidence_files=Mock(return_value={'files': []}),
            import_evidence_path=Mock(return_value={'id': 'job', 'state': 'reading'}),
            restart=Mock(return_value={'restarting': True}))

        class Listed(app.Handler):
            # ALLOWED_HOSTS is set but nothing is enforced: it must add no host and no origin.
            access_cfg = gate.AccessConfig(allowed_hosts=gate.LOOPBACK | {HOST})

        cls.plain, cls.listed = Served(app.Handler, cls.console), Served(Listed, cls.console)

    @classmethod
    def tearDownClass(cls):
        cls.plain.close()
        cls.listed.close()
        cls.temp.cleanup()

    def setUp(self):
        self.console.read_only = False
        for method in (self.console.evidence_files, self.console.import_evidence_path, self.console.restart):
            method.reset_mock()

    def post(self, served, path, data=None, **headers):
        return served.request(path, method='POST', data={} if data is None else data,
                              headers={'X-Mentat-Token': self.console.token, **headers})

    def test_nothing_is_enforced_and_the_bind_is_loopback(self):
        self.assertFalse(app.Handler.access_cfg.enforced)
        self.assertEqual(app.Handler.access_cfg.bind_host, '127.0.0.1')
        self.assertFalse(gate.AccessConfig.from_env({}).enforced)

        class Spy:
            request_queue_size = 0

            def __init__(self, address, handler, bind_and_activate=True):
                Spy.address = address

            server_bind = server_activate = server_close = lambda self: None

        with patch.object(app, 'ThreadingHTTPServer', Spy):
            app.console_server(8123)
            self.assertEqual(Spy.address, ('127.0.0.1', 8123))
            app.console_server(8123, host='10.0.0.5')
            self.assertEqual(Spy.address, ('10.0.0.5', 8123))

    def test_host_must_be_loopback_with_this_port(self):
        for served in (self.plain, self.listed):
            status, _, body = served.request('/api/session?light=1', headers={'Host': 'elsewhere.invalid'})
            self.assertEqual((status, body), (403, {'error': 'Local access only.'}))
            status, _, _ = served.request('/api/session?light=1', headers={'Host': '127.0.0.1:1'})
            self.assertEqual(status, 403)
        status, _, body = self.listed.request('/api/session?light=1', headers={'Host': HOST})
        self.assertEqual((status, body), (403, {'error': 'Local access only.'}), 'ALLOWED_HOSTS alone adds nothing')
        status, _, body = self.plain.request('/api/session?light=1')
        self.assertEqual((status, body['token'], body['read_only']), (200, 'local-token', False))
        status, _, body = self.plain.request('/api/session?light=1', headers={'Host': f'localhost:{self.plain.port}'})
        self.assertEqual(status, 200)

    def test_origin_must_be_loopback(self):
        with patch('giving_forecast.project', return_value={'ok': True}) as project:
            for served, origin in ((self.plain, 'https://elsewhere.invalid'), (self.listed, f'https://{HOST}'),
                                   (self.listed, f'http://{HOST}')):
                status, _, body = self.post(served, '/api/forecast/project', Origin=origin)
                self.assertEqual((status, body), (403, {'error': 'Local access only.'}), origin)
            project.assert_not_called()
            for origin in (f'http://127.0.0.1:{self.plain.port}', f'http://localhost:{self.plain.port}'):
                status, _, body = self.post(self.plain, '/api/forecast/project', Origin=origin)
                self.assertEqual((status, body), (200, {'ok': True}), origin)
            self.assertEqual(self.post(self.plain, '/api/forecast/project')[0], 200, 'no Origin header is still fine')

    def test_token_flow_is_intact(self):
        status, headers, body = self.plain.request('/api/session?light=1')
        self.assertEqual((status, body['token']), (200, 'local-token'))
        self.assertEqual(headers['Cache-Control'], 'no-store')
        with patch('giving_forecast.project', return_value={'ok': True}):
            self.assertEqual(self.post(self.plain, '/api/forecast/project')[0], 200)
            status, _, body = self.plain.request('/api/forecast/project', method='POST', data={},
                                                 headers={'X-Mentat-Token': 'hosted-read-only'})
            self.assertEqual((status, body), (403, {'error': 'Refresh the app before continuing.'}))
            status, _, _ = self.plain.request('/api/forecast/project', method='POST', data={})
            self.assertEqual(status, 403)

    def test_health_routes_but_no_hosted_headers_or_log_lines(self):
        status, _, body = self.plain.request('/healthz')
        self.assertEqual((status, body), (200, {'status': 'ok'}))
        status, _, body = self.plain.request('/version')
        self.assertEqual((status, body['app'], set(body)), (200, 'mentat-console', {'app', 'git_sha', 'build_time', 'snapshot_id'}))
        status, headers, _ = self.plain.request('/app.js')
        self.assertEqual(status, 200)
        self.assertEqual(headers['Cache-Control'], 'private, no-cache')
        for name in ('Referrer-Policy', 'X-Frame-Options', 'Cross-Origin-Resource-Policy'):
            self.assertNotIn(name, headers)
        self.assertEqual(self.plain.stderr, '', 'a local console logs nothing')

    def test_local_health_routes_keep_the_loopback_host_check(self):
        for host in ('evil.example', 'localhost'):
            status, _, body = self.plain.request('/healthz', headers={'Host': host})
            self.assertEqual((status, body), (403, {'error': 'Local access only.'}), host)
        status, _, _ = self.plain.request('/healthz', method='POST', data={}, headers={'Host': f'localhost:{self.plain.port}'})
        self.assertEqual(status, 403, 'a POST to a health route still needs the token')

    def test_other_verbs_and_the_local_answers_are_what_they_were(self):
        for method in ('PUT', 'DELETE', 'HEAD', 'OPTIONS'):
            self.assertEqual(self.plain.request('/api/runs', method=method)[0], 501, method)
        status, _, body = self.plain.request('/api/health')
        self.assertEqual(status, 200)
        self.assertEqual((body['project'], body['report_directory'], body['resolved_data_root']),
                         (str(app.PROJECT), self.temp.name, str(Path(self.temp.name))))
        status, _, body = self.plain.request('/api/session?light=1')
        self.assertEqual(body['report_directory'], self.temp.name)

    def test_local_donor_rows_are_not_gated(self):
        with patch.object(app.contribution_workbook_guard, 'verify_sources', return_value=[]), \
                patch.object(app.contribution_query, 'records', return_value={'rows': [{'source_id': 'a'}]}):
            status, _, body = self.post(self.plain, '/api/contributions/records', {'dataset': 'Example'})
        self.assertEqual((status, body['rows']), (200, [{'source_id': 'a'}]))

    def test_local_read_only_allowlist_is_what_it_was(self):
        self.console.read_only = True
        status, _, body = self.post(self.plain, '/api/evidence-files', {'path': '.'})
        self.assertEqual((status, body), (200, {'files': []}))
        status, _, body = self.post(self.plain, '/api/evidence-csv', {'path': 'C:/somewhere/file.csv'})
        self.assertEqual((status, body['state']), (200, 'reading'))
        status, _, body = self.post(self.plain, '/api/restart')
        self.assertEqual(status, 403)
        self.assertIn('read-only while the original console', body['error'])
        self.console.restart.assert_not_called()

    def test_local_get_routes_are_not_allowlisted(self):
        for route in ('/api/prime-status', '/api/evidence'):
            status, _, body = self.plain.request(route)
            self.assertEqual(status, 200, route)
            self.assertNotEqual(body, {'error': app.HOSTED_REFUSED}, route)
        self.assertEqual(self.plain.request('/api/not-a-route')[:3:2], (404, {'error': 'Not found.'}))

    def test_local_person_level_queries_are_not_refused(self):
        body = {'dataset': 'Example', 'dimensions': ['contributor', 'city'], 'filters': [{'field': 'zip5', 'op': 'eq', 'value': '0'}]}
        with patch.object(app.contribution_workbook_guard, 'verify_sources', return_value=[]), \
                patch.object(app.contribution_query, 'query', return_value={'rows': []}) as reader:
            status, _, answer = self.post(self.plain, '/api/contributions/query', body)
        self.assertEqual((status, answer['rows']), (200, []))
        reader.assert_called_once()

    def test_local_errors_keep_their_text_and_bodies_keep_their_paths(self):
        error = OSError(2, 'No such file or directory', '/srv/mentat/x.sqlite')
        with patch.object(self.console, 'summary', Mock(side_effect=error), create=True):
            status, _, body = self.plain.request('/api/runs')
        self.assertEqual(status, 400)
        self.assertIn('No such file or directory', body['error'])
        self.assertIn('x.sqlite', body['error'])
        catalog = {'path': '/srv/x', 'files': ['C:\\x']}
        with patch.object(self.console, 'report_catalog', Mock(return_value=catalog), create=True):
            self.assertEqual(self.plain.request('/api/reports?id=a')[2], catalog)

    def test_local_flow_keeps_its_person_level_keys(self):
        body = flow_body()
        self.assertTrue(set(keys_of(body)) & app.FLOW_PERSON_KEYS)
        with patch('flow_data.flow', return_value=body):
            status, _, answer = self.plain.request('/api/flow?root=industry&cycle=2024')
        self.assertEqual((status, answer), (200, json.loads(json.dumps(body))))

    def test_local_flow_keeps_its_individual_contributors(self):
        body = individuals_flow_body()
        with patch('flow_data.flow', return_value=body):
            status, _, answer = self.plain.request('/api/flow?root=industry&cycle=2024')
        self.assertEqual((status, answer), (200, json.loads(json.dumps(body))))
        self.assertIn('TEST PERSON A', json.dumps(answer))

    @needs_openpyxl
    def test_local_comparison_keeps_its_person_subjects(self):
        body = comparison_body(self.temp.name)
        self.assertIn('Person', {subject['kind'] for subject in body['subjects']})
        with patch.object(self.console, 'comparison', Mock(return_value=body), create=True):
            for query in ('', '&kind=Person'):
                status, _, answer = self.plain.request('/api/comparison?id=published-reports' + query)
                self.assertEqual((status, answer), (200, body), query)
        self.assertIn('Alex Example', json.dumps(answer))

    def test_local_json_bodies_keep_their_os_error_text(self):
        catalog = {'note': "[Errno 2] No such file or directory: 'x'", 'rows': ['[WinError 3] gone']}
        with patch.object(self.console, 'report_catalog', Mock(return_value=catalog), create=True):
            self.assertEqual(self.plain.request('/api/reports?id=a')[2], catalog)

    def test_local_cached_bodies_and_geography_errors_are_sent_as_they_are(self):
        entry = response_cache.from_value(CACHED_BODY)
        with patch.object(app, 'worksheet_entry', return_value=entry), \
                patch.object(self.console, 'report_path', Mock(return_value=Path('r.xlsx')), create=True):
            status, headers, answer = self.plain.request('/api/report-data?id=r&index=0')
        self.assertEqual((status, answer, headers['ETag']), (200, CACHED_BODY, entry.etag))
        with patch.object(app.contribution_geo, 'geography', return_value=dict(GEOGRAPHY_RESULT)):
            status, _, answer = self.post(self.plain, '/api/contributions/geography', {'zip5': ['02134']})
        self.assertEqual((status, answer['errors']), (200, GEOGRAPHY_RESULT['errors']))

    @needs_openpyxl
    def test_local_worksheets_have_no_person_level_rule(self):
        path = Path(self.temp.name) / 'synthetic-report.xlsx'
        synthetic_report(path)
        with patch.object(self.console, 'report_path', Mock(return_value=path), create=True):
            for name in PERSON_SHEETS:
                status, _, body = self.plain.request(f'/api/report-data?id=r&index=0&sheet={quote(name)}')
                self.assertEqual((status, body['row_count']), (200, 2), name)

    @needs_openpyxl
    def test_local_summary_sheets_keep_their_contributor_names(self):
        path = Path(self.temp.name) / 'summary-report.xlsx'
        summary_workbook(path)
        with patch.object(self.console, 'report_path', Mock(return_value=path), create=True):
            for name in SUMMARY_SHEETS:
                status, _, body = self.plain.request(f'/api/report-data?id=r&index=0&sheet={quote(name)}')
                self.assertEqual((status, body['rows']), (200, SUMMARY_SHEETS[name][1]), name)


@needs_crypto
class HostedMode(unittest.TestCase):
    """Access enforced on a loopback bind (ACCESS_ENFORCE=1), as the container will run it."""

    MUTATORS = ('upload', 'create_roster', 'collect_evidence', 'launch', 'resume', 'acknowledge_quality',
                'archive_run', 'stop', 'set_workers', 'publish_pocket', 'add_to_briefing', 'remove_from_briefing',
                'order_briefing', 'save_view', 'delete_view', 'restart', 'import_evidence_path',
                'import_evidence_csv', 'evidence_files', 'start_evidence_geocode')

    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.idp = Idp(cls.temp.name)
        cls.mutators = {name: Mock() for name in cls.MUTATORS}
        cls.console = SimpleNamespace(
            token=REAL_TOKEN, read_only=True, report_directory=Path(cls.temp.name),
            configured_root=cls.temp.name, root=Path(cls.temp.name),
            summary=Mock(return_value={'runs': []}), report_path=Mock(), evidence_store=Mock(), **cls.mutators)
        config = gate.AccessConfig.from_env(cls.idp.env())

        class Hosted(app.Handler):
            access_cfg = config

        cls.served = Served(Hosted, cls.console)

    @classmethod
    def tearDownClass(cls):
        cls.served.close()
        cls.temp.cleanup()

    def get(self, path, *, token=True, host=HOST, **headers):
        sent = {'Host': host, **headers}
        if token:
            sent['Cf-Access-Jwt-Assertion'] = self.idp.token() if token is True else token
        return self.served.request(path, headers=sent)

    def post(self, path, data=None, *, marker='hosted-read-only', host=HOST, jwt=True, **headers):
        sent = {'Host': host, **headers}
        if marker is not None:
            sent['X-Mentat-Token'] = marker
        if jwt:
            sent['Cf-Access-Jwt-Assertion'] = self.idp.token()
        return self.served.request(path, method='POST', data={} if data is None else data, headers=sent)

    def setUp(self):
        for method in (*self.mutators.values(), self.console.report_path, self.console.summary,
                       self.console.evidence_store):
            method.reset_mock()

    def assert_untouched(self):
        for name, method in self.mutators.items():
            self.assertFalse(method.called, name)
        self.assertFalse(self.console.report_path.called)
        self.assertEqual(self.console.evidence_store.mock_calls, [])

    def test_the_config_is_the_environments(self):
        config = gate.AccessConfig.from_env(self.idp.env())
        self.assertTrue(config.enforced)
        self.assertEqual(config.bind_host, '127.0.0.1')
        self.assertEqual(config.allowed_hosts, gate.LOOPBACK | {HOST})

    def test_no_valid_jwt_gets_nothing(self):
        for route in ('/', '/api/session?light=1', '/app.js', '/api/runs', '/api/report?id=a&index=0'):
            status, headers, body = self.get(route, token=False)
            self.assertEqual((status, body), (403, b'forbidden\n'), route)
        for name, bad in (('wrong audience', self.idp.token(aud=['someone-else'])),
                          ('expired', self.idp.token(exp=int(time.time()) - 3600)),
                          ('no identity', self.idp.token(email=None)),
                          ('garbage', 'not.a.jwt')):
            self.assertEqual(self.get('/api/session?light=1', token=bad)[0], 403, name)
        status, _, _ = self.post('/api/forecast/project', jwt=False)
        self.assertEqual(status, 403)
        self.assert_untouched()

    def test_wrong_host_is_refused_even_with_a_valid_jwt(self):
        self.assertEqual(self.get('/api/session?light=1', host='evil.example')[0], 403)
        self.assertEqual(self.post('/api/forecast/project', host='evil.example')[0], 403)

    def test_a_read_only_get_is_served_with_security_headers(self):
        status, headers, body = self.get('/api/session?light=1')
        self.assertEqual((status, body['read_only']), (200, True))
        status, headers, _ = self.get('/app.js')
        self.assertEqual(status, 200)
        self.assertEqual(headers['Referrer-Policy'], 'no-referrer')
        self.assertEqual(headers['X-Frame-Options'], 'DENY')
        self.assertEqual(headers['Cross-Origin-Resource-Policy'], 'same-origin')
        self.assertEqual(headers['X-Content-Type-Options'], 'nosniff')
        self.assertIn("frame-ancestors 'none'", headers['Content-Security-Policy'])
        self.assertEqual(headers['Cache-Control'], 'private, no-cache', 'the route keeps its own caching rule')
        # one copy of each header, not the route's and the gate's
        raw = http.client.HTTPConnection('127.0.0.1', self.served.port)
        with contextlib.redirect_stderr(io.StringIO()):
            raw.request('GET', '/app.js', headers={'Host': HOST, 'Cf-Access-Jwt-Assertion': self.idp.token()})
            response = raw.getresponse()
            response.read()
        names = [name.lower() for name, _ in response.getheaders()]
        raw.close()
        for name in ('x-content-type-options', 'cache-control', 'referrer-policy', 'x-frame-options'):
            self.assertEqual(names.count(name), 1, name)

    def test_the_consoles_token_is_never_handed_out(self):
        for route in ('/api/session?light=1', '/api/session', '/api/contributions/session'):
            status, headers, body = self.get(route)
            self.assertEqual(status, 200, route)
            self.assertEqual(body['token'], 'hosted-read-only', route)
            self.assertNotIn(REAL_TOKEN, json.dumps(body))
        self.assertNotIn(REAL_TOKEN, self.served.stderr)

    def test_a_post_needs_the_marker_and_an_allowed_origin(self):
        with patch('giving_forecast.project', return_value={'ok': True}) as project:
            self.assertEqual(self.post('/api/forecast/project'), (200, unittest.mock.ANY, {'ok': True}))
            for name, kwargs in (('no marker', {'marker': None}), ('the console token', {'marker': REAL_TOKEN}),
                                 ('a guess', {'marker': 'x'}),
                                 ('a foreign origin', {'Origin': 'https://elsewhere.invalid'}),
                                 ('http origin', {'Origin': f'http://{HOST}'})):
                status, _, _ = self.post('/api/forecast/project', **kwargs)
                self.assertEqual(status, 403, name)
            self.assertEqual(project.call_count, 1)
            self.assertEqual(self.post('/api/forecast/project', Origin=f'https://{HOST}')[0], 200)
            self.assertEqual(self.post('/api/forecast/project', Origin=f'https://{HOST.upper()}')[0], 200)

    def test_every_post_route_but_the_analytical_reads_is_refused(self):
        literals = set(re.findall(r"""['"](/api/[a-z0-9/-]+)['"]""", inspect.getsource(app.Handler.do_POST)))
        refused = sorted(literals - ANALYTICAL)
        for needed in ('/api/evidence-csv', '/api/evidence-files', '/api/export/pdf', '/api/upload', '/api/restart',
                       '/api/stop', '/api/keys', '/api/start', '/api/evidence-geocode', '/api/studio/views'):
            self.assertIn(needed, refused)
        bodies = {'/api/evidence-csv': {'path': 'C:/Windows/win.ini'}, '/api/evidence-files': {'path': 'C:/'},
                  '/api/upload': {'filename': 'x.xlsx', 'content': 'AAAA'}}
        for route in refused:
            with self.subTest(route=route):
                status, _, body = self.post(route, bodies.get(route))
                self.assertEqual(status, 403, body)
        # the raw CSV upload form of /api/evidence-csv, and a path with a query string
        self.assertEqual(self.served.request('/api/evidence-csv', method='POST', raw=b'a,b\n1,2\n', headers={
            'Host': HOST, 'X-Mentat-Token': 'hosted-read-only', 'Content-Type': 'text/csv',
            'X-Evidence-Name': 'x.csv', 'Cf-Access-Jwt-Assertion': self.idp.token()})[0], 403)
        self.assertEqual(self.post('/api/evidence-csv?x=1', {'path': 'C:/'})[0], 403)
        self.assert_untouched()

    def test_every_verb_is_gated_not_only_get_and_post(self):
        token = self.idp.token()
        for method in ('HEAD', 'PUT', 'DELETE', 'PATCH', 'OPTIONS', 'TRACE'):
            with self.subTest(method=method):
                status, _, _ = self.served.request('/api/runs', method=method, headers={'Host': HOST})
                self.assertEqual(status, 403, 'refused by the gate, not left to http.server as 501')
                # past the gate nothing handles it
                status, _, _ = self.served.request('/api/runs', method=method, headers={
                    'Host': HOST, 'Cf-Access-Jwt-Assertion': token})
                self.assertEqual(status, 501)
        self.assert_untouched()
        for name in ('do_GET', 'do_POST'):
            self.assertNotIn('access_check', inspect.getsource(getattr(app.Handler, name)), name)
        self.assertIs(app.Handler.parse_request, gate.AccessMixin.parse_request)

    def records(self, route, email=EMAIL):
        return self.served.request(route, method='POST', data={'dataset': 'Example'}, headers={
            'Host': HOST, 'X-Mentat-Token': 'hosted-read-only', 'Cf-Access-Jwt-Assertion': self.idp.token(email=email)})

    def test_donor_rows_are_refused_unless_the_email_is_listed(self):
        with patch.object(app.contribution_workbook_guard, 'verify_sources', return_value=[]), \
                patch.object(app.contribution_query, 'records', return_value={'rows': [{'source_id': 'a'}]}) as contribution, \
                patch.object(app.contribution_query, 'query', return_value={'rows': []}) as aggregate, \
                patch('analysis_data.records', return_value={'rows': [{'source_id': 'b'}]}) as insights, \
                patch.dict(os.environ):
            os.environ.pop('MENTAT_DONOR_ROW_EMAILS', None)
            readers = {'/api/contributions/records': contribution, '/api/insights/records': insights}
            for route, reader in readers.items():          # default: nobody
                status, _, body = self.records(route)
                self.assertEqual((status, body), (403, {'error': app.HOSTED_DONOR_ROWS_REFUSED}), route)
                reader.assert_not_called()
            os.environ['MENTAT_DONOR_ROW_EMAILS'] = ' Someone@Example.test , ' + EMAIL.upper()   # case and spaces
            for route, reader in readers.items():
                status, _, body = self.records(route)
                self.assertEqual((status, len(body['rows'])), (200, 1), route)
                self.assertEqual(reader.call_count, 1)
                status, _, body = self.records(route, email='stranger@example.test')
                self.assertEqual((status, body), (403, {'error': app.HOSTED_DONOR_ROWS_REFUSED}), route)
                self.assertEqual(reader.call_count, 1)
            self.assertEqual(self.records('/api/contributions/query', email='stranger@example.test')[0], 200,
                             'only the two records routes carry donor rows')
            aggregate.assert_called_once()
            os.environ['MENTAT_DONOR_ROW_EMAILS'] = ''
            self.assertEqual(self.records('/api/contributions/records')[0], 403)
        self.assert_untouched()

    def test_health_and_session_carry_no_server_paths(self):
        windows = 'C:\\Users\\pgoff\\Mentat\\mentat\\.cache\\x\\mentat.db'
        summary = {'runs': [], 'collection_directory': self.temp.name, 'report_directory': self.temp.name,
                   'coverage': [{'industry': 'Banking', 'receipt': windows, 'database': '/srv/mentat/db.sqlite'}],
                   'read_only': True}

        def strings(value):
            if isinstance(value, dict):
                value = list(value.values())
            for item in value if isinstance(value, list) else [value]:
                if isinstance(item, (dict, list)):
                    yield from strings(item)
                elif isinstance(item, str):
                    yield item

        with patch.object(self.console, 'summary', Mock(return_value=summary)):
            for route in ('/api/health', '/api/session?light=1', '/api/session', '/api/runs'):
                status, _, body = self.get(route)
                self.assertEqual(status, 200, route)
                self.assertEqual([s for s in strings(body) if app._SERVER_PATH.match(s)], [], route)
                for key in ('project', 'data_root', 'resolved_data_root', 'report_directory', 'collection_directory'):
                    self.assertNotIn(key, body, route)
        status, _, health = self.get('/api/health')
        self.assertEqual((health['application'], health['read_only']), ('MENTAT', True))
        self.assertIn('code_sha256', health)
        with patch.object(self.console, 'summary', Mock(return_value=summary)):
            status, _, session = self.get('/api/session')
        self.assertEqual(session['coverage'], [{'industry': 'Banking'}])
        self.assertEqual(session['token'], 'hosted-read-only')

    def test_the_workbook_download_is_refused(self):
        status, _, body = self.get('/api/report?id=run&index=0')
        self.assertEqual((status, body), (403, {'error': app.HOSTED_REFUSED}))
        self.assert_untouched()

    def test_healthz_and_version(self):
        # no token, no proxy headers, loopback peer: this is the container's own health check
        status, _, body = self.served.request('/healthz', headers={'Host': f'127.0.0.1:{self.served.port}'})
        self.assertEqual((status, body), (200, {'status': 'ok'}))
        status, _, body = self.served.request('/version', headers={'Host': 'localhost'})
        self.assertEqual((status, body['app']), (200, 'mentat-console'))
        self.assertEqual(set(body), {'app', 'git_sha', 'build_time', 'snapshot_id'})
        # the same request seen through the tunnel (Cf- headers), a foreign Host, or on any other route
        self.assertEqual(self.served.request('/healthz', headers={'Host': HOST, 'Cf-Ray': '1'})[0], 403)
        self.assertEqual(self.served.request('/healthz', headers={'Host': 'evil.example'})[0], 403)
        self.assertEqual(self.served.request('/api/runs', headers={'Host': f'127.0.0.1:{self.served.port}'})[0], 403)
        self.assertEqual(self.get('/healthz')[0], 200)

    def test_one_audit_line_per_request_without_query_or_address(self):
        self.get('/api/session?light=1&secret=abc')
        line = json.loads(self.served.stderr.strip().splitlines()[-1])
        self.assertEqual({k: line[k] for k in ('app', 'user', 'method', 'route', 'status')},
                         {'app': 'mentat-console', 'user': EMAIL, 'method': 'GET', 'route': '/api/session', 'status': 200})
        self.assertNotIn('secret', self.served.stderr)
        self.assertNotIn('127.0.0.1', self.served.stderr)
        # The gate answers a refused request and only then writes its line, so the capture that ends with the
        # response can end first: capture around the whole request and wait (5 s at most) for the line.
        captured = io.StringIO()
        connection = http.client.HTTPConnection('127.0.0.1', self.served.port, timeout=10)
        try:
            with contextlib.redirect_stderr(captured):
                connection.request('GET', '/api/runs', headers={'Host': HOST})
                connection.getresponse().read()
                deadline = time.monotonic() + 5
                while not captured.getvalue().endswith('\n') and time.monotonic() < deadline:
                    time.sleep(0.01)
        finally:
            connection.close()
        self.assertTrue(captured.getvalue().endswith('\n'), 'no audit line for the refused request within 5 s')
        denied = json.loads(captured.getvalue().strip().splitlines()[-1])
        self.assertEqual((denied['user'], denied['status'], denied['reason']), ('-', 403, 'no token'))

    def test_every_get_route_is_allowlisted_or_refused(self):
        literals = set(re.findall(r"""['"](/[a-z0-9/_-]*)['"]""", inspect.getsource(app.Handler.do_GET)))
        prefixes = {item for item in literals if item.endswith('/') and item != '/'}
        routes = literals - prefixes - set(gate.HEALTH_ROUTES)   # /healthz and /version are answered before the gate
        self.assertEqual(prefixes, {'/geo/v2/', '/basemap/', '/api/public-data/'},
                         'the two file trees and the refused public-data API are the only prefixes')
        self.assertEqual(app.HOSTED_GET_ROUTES - routes, set(), 'an allowlisted route do_GET no longer has')
        refused = sorted(routes - app.HOSTED_GET_ROUTES)
        for needed in ('/api/report', '/api/run', '/api/evidence', '/api/evidence-csv', '/api/evidence-job',
                       '/api/evidence-stats', '/api/evidence-geo', '/api/evidence-geocode', '/api/evidence-network',
                       '/api/publish/pocket', '/api/prime-status', '/api/studio/briefing', '/api/studio/briefing/item',
                       '/api/studio/views', '/api/public-data/catalogue', '/api/public-data/rows',
                       '/api/public-data/download', '/api/public-data/provenance'):
            self.assertIn(needed, refused)
        for route in (*refused, '/api/not-a-route', '/api/evidence-csv?job=x', '/api/studio/views?id=a',
                      '/api/public-data/rows?table=x&page=0', '/api/public-data/not-a-route'):
            with self.subTest(route=route):
                status, _, body = self.get(route)
                self.assertEqual((status, body), (403, {'error': app.HOSTED_REFUSED}))
        # the page, every static file and the two file trees pass the gate (their own answers are 200 or 404)
        for route in (*app.STATIC_ASSETS, '/geo/v2/albers/none.json', '/basemap/none'):
            with self.subTest(route=route):
                status, _, body = self.get(route)
                self.assertNotEqual((status, body), (403, {'error': app.HOSTED_REFUSED}))
        self.assert_untouched()

    def query(self, route, body, email=EMAIL):
        return self.served.request(route, method='POST', data={'dataset': 'Example', **body}, headers={
            'Host': HOST, 'X-Mentat-Token': 'hosted-read-only', 'Cf-Access-Jwt-Assertion': self.idp.token(email=email)})

    def test_the_person_level_field_set_is_computed_from_the_semantics(self):
        semantics = app.contribution_query.semantics
        self.assertEqual(app.PERSON_LEVEL_FIELDS,
                         {'contributor', 'street_1', 'street_2', 'city', 'zip5', 'employer', 'occupation', 'date'}
                         | set(semantics.AUDIT) | set(semantics.RAW))
        self.assertGreater(len(semantics.RAW), 40)
        self.assertEqual(app.AGGREGATE_FIELDS & app.PERSON_LEVEL_FIELDS, set())
        self.assertEqual(app.AGGREGATE_FIELDS | app.PERSON_LEVEL_FIELDS, set(semantics.DIMENSIONS),
                         'every dimension is one or the other; a new one is person-level until listed')

    def test_a_person_level_query_is_refused_unless_the_email_is_listed(self):
        person = ({'dimensions': ['contributor']}, {'dimensions': ['company', 'city'], 'measures': ['amount']},
                  {'dimensions': ['company'], 'filters': [{'field': 'zip5', 'op': 'eq', 'value': '00000'}]},
                  {'dimensions': ['company'], 'aggregate_filters': [{'field': 'source_id', 'op': 'gte', 'value': 1}]},
                  {'dimensions': ['company'], 'sort': [{'field': 'employer', 'direction': 'asc'}]},
                  {'dimensions': ['date']}, {'dimensions': ['street_1', 'street_2']}, {'dimensions': ['occupation']},
                  {'dimensions': ['reasons']}, {'filters': [{'field': 'raw_zip', 'op': 'eq', 'value': '0'}]},
                  {'dimensions': ['reported_memo_text']}, {'filters': [{'field': 'reported_filing_form', 'op': 'eq', 'value': 'F3'}]},
                  {'dimensions': 'contributor'}, {'filters': ['city']}, {'filters': {'field': 'employer'}})
        aggregate = {'dimensions': ['company', 'cycle', 'state'], 'measures': ['amount', 'records', 'contributors'],
                     'filters': [{'field': 'cycle', 'op': 'eq', 'value': '2024'}, {'field': 'amount', 'op': 'gte', 'value': 100},
                                 {'field': 'recipient_category', 'op': 'eq', 'value': 'Candidate'}],
                     'aggregate_filters': [{'field': 'amount', 'op': 'gte', 'value': 1}],
                     'sort': [{'field': 'amount', 'direction': 'desc'}]}
        with patch.object(app.contribution_workbook_guard, 'verify_sources', return_value=[]), \
                patch.object(app.contribution_query, 'query', return_value={'rows': []}) as contribution, \
                patch('analysis_data.query', return_value={'rows': []}) as insights, patch.dict(os.environ):
            os.environ.pop('MENTAT_DONOR_ROW_EMAILS', None)
            readers = {'/api/contributions/query': contribution, '/api/insights/query': insights}
            for route, reader in readers.items():
                for body in person:
                    status, _, answer = self.query(route, body)
                    self.assertEqual((status, answer), (403, {'error': app.HOSTED_PERSON_LEVEL_REFUSED}), (route, body))
                self.assertEqual(self.query(route, aggregate)[0], 200, route)
                self.assertEqual(reader.call_count, 1, 'only the aggregate query reached the reader')
            os.environ['MENTAT_DONOR_ROW_EMAILS'] = ' Someone@Example.test , ' + EMAIL.upper()
            for route, reader in readers.items():
                before = reader.call_count
                for body in person:
                    self.assertEqual(self.query(route, body)[0], 200, (route, body))
                self.assertEqual(reader.call_count, before + len(person))
                status, _, answer = self.query(route, person[0], email='stranger@example.test')
                self.assertEqual((status, answer), (403, {'error': app.HOSTED_PERSON_LEVEL_REFUSED}))
            # the other analytical POSTs name no field, so they are not touched by the rule
            with patch('giving_forecast.project', return_value={'ok': True}):
                self.assertEqual(self.post('/api/forecast/project', {'industry': 'x'})[0], 200)
        self.assert_untouched()

    @needs_openpyxl
    def test_a_person_level_worksheet_is_refused_unless_the_email_is_listed(self):
        path = Path(self.temp.name) / 'synthetic-report.xlsx'
        synthetic_report(path)
        broken = Path(self.temp.name) / 'broken-report.xlsx'
        broken.write_bytes(b'this is not a workbook')

        def sheet(name, *, header=0, **headers):
            return self.get(f'/api/report-data?id=r&index=0&sheet={quote(name)}' + (f'&header={header}' if header else ''),
                            **headers)

        with patch.object(self.console, 'report_path', Mock(return_value=path)), patch.dict(os.environ):
            os.environ.pop('MENTAT_DONOR_ROW_EMAILS', None)
            self.assertEqual(self.get('/api/report-data?id=r&index=0')[2]['sheets'], [*AGGREGATE_SHEETS, *PERSON_SHEETS],
                             'the sheet list is names only')
            for name in AGGREGATE_SHEETS:
                status, _, body = sheet(name)
                self.assertEqual((status, body['sheet'], body['row_count']), (200, name, 2), name)
            for name in PERSON_SHEETS:
                self.assertEqual(sheet(name)[:3:2], (403, {'error': app.HOSTED_PERSON_LEVEL_REFUSED}), name)
            # asking for another header row does not get past the sheet's own header row
            for header in (5, 6, 50, 100):
                self.assertEqual(sheet('Employee List', header=header)[0], 403, header)
            # an aggregate sheet stays available at its own header row and at any other row that names no person
            self.assertEqual((sheet('Giving Summary', header=3)[0], sheet('Giving Summary', header=1)[0]), (200, 200))
            self.assertEqual(sheet('Nonexistent Sheet')[0], 403, 'fail closed on a sheet that cannot be read')
            os.environ['MENTAT_DONOR_ROW_EMAILS'] = EMAIL
            for name in PERSON_SHEETS:
                status, _, body = sheet(name)
                self.assertEqual((status, body['row_count']), (200, 2), name)
            os.environ.pop('MENTAT_DONOR_ROW_EMAILS')
            self.console.report_path.return_value = broken
            self.assertEqual(sheet('Giving Summary')[:3:2], (403, {'error': app.HOSTED_PERSON_LEVEL_REFUSED}),
                             'fail closed when the headers cannot be read')

    def test_hosted_errors_carry_no_os_text_or_server_path(self):
        error = OSError(2, 'No such file or directory', '/srv/mentat/secret.sqlite')
        with patch.object(self.console, 'summary', Mock(side_effect=error)):
            status, _, body = self.get('/api/runs')
        self.assertEqual((status, body), (400, {'error': app.HOSTED_IO_ERROR}))
        with patch('giving_forecast.project', side_effect=error):
            status, _, body = self.post('/api/forecast/project')
        self.assertEqual((status, body), (400, {'error': app.HOSTED_IO_ERROR}))
        with patch.object(self.console, 'summary', Mock(side_effect=PermissionError('prewarm sandbox refused open: /srv/x'))):
            self.assertEqual(self.get('/api/runs')[2], {'error': app.HOSTED_IO_ERROR})
        with patch.object(self.console, 'summary', Mock(side_effect=ValueError('Choose a run.'))):
            self.assertEqual(self.get('/api/runs')[:3:2], (400, {'error': 'Choose a run.'}), 'a reader sentence is unchanged')
        # every JSON body is scrubbed in send(), not only the routes that used to call public()
        catalog = {'name': 'ok', 'path': '/srv/x', 'files': ['C:\\x', 'y'], 'nested': [{'p': '\\\\host\\s', 'n': 1}]}
        with patch.object(self.console, 'report_catalog', Mock(return_value=catalog), create=True):
            status, _, body = self.get('/api/reports?id=a')
        self.assertEqual((status, body), (200, {'name': 'ok', 'files': ['y'], 'nested': [{'n': 1}]}))

    def flow(self, email=EMAIL):
        return self.served.request('/api/flow?root=industry&cycle=2024', headers={
            'Host': HOST, 'Cf-Access-Jwt-Assertion': self.idp.token(email=email)})

    def test_flow_has_no_person_level_key_unless_the_email_is_listed(self):
        body = flow_body()
        self.assertTrue(set(keys_of(body)) & app.FLOW_PERSON_KEYS, 'the fixture must carry what is to be stripped')
        # amounts, committee ids and the aggregates (state) stay; the evidence keeps its amount (purpose is free text
        # that can name a payee: it goes)
        kept = {'amount': 250.0}
        kept_wide = {**kept, 'state': 'ZZ'}
        with patch('flow_data.flow', return_value=body), patch.dict(os.environ):
            os.environ.pop('MENTAT_DONOR_ROW_EMAILS', None)
            for email in (EMAIL, 'stranger@example.test'):
                status, _, answer = self.flow(email)
                self.assertEqual(status, 200)
                self.assertEqual(set(keys_of(answer)) & app.FLOW_PERSON_KEYS, set(), email)
                self.assertEqual([link['evidence'] for link in answer['links']], [[kept], [kept_wide]])
                self.assertEqual((answer['links'][0]['value'], answer['links'][0]['committee']), (250.0, 'C00000001'))
                self.assertEqual(answer['nodes'], body['nodes'])
                self.assertEqual((answer['totals'], answer['coverage'], answer['routes']),
                                 (body['totals'], body['coverage'], body['routes']))
                self.assertEqual(answer['deep'], {'records': [kept_wide]})
            os.environ['MENTAT_DONOR_ROW_EMAILS'] = ' Someone@Example.test , ' + EMAIL.upper()
            status, _, answer = self.flow()
            self.assertEqual((status, answer), (200, json.loads(json.dumps(body))), 'a listed account gets the body as built')
            self.assertEqual(self.flow('stranger@example.test')[2]['links'][0]['evidence'], [kept])

    def test_flow_names_no_individual_contributor_unless_the_email_is_listed(self):
        body = individuals_flow_body()
        before = json.dumps(body)
        into = lambda flow: sum(Decimal(str(link['value'])) for link in flow['links']
                                if link['target'] == 'committee:C00000001')
        with patch('flow_data.flow', return_value=body), patch.dict(os.environ):
            os.environ.pop('MENTAT_DONOR_ROW_EMAILS', None)
            for email in (EMAIL, 'stranger@example.test'):
                status, _, answer = self.flow(email)
                self.assertEqual(status, 200)
                # a recursive walk of every key and string: no individual's id, label, name as filed, payee or refund text
                leaks = [text for text in strings_of(answer) if any(piece in text.upper() for piece in (
                    'TEST PERSON', 'PERSON, JANE', 'JANE PERSON', 'CONSULTING')) or text in ('purpose', 'held')]
                self.assertEqual(leaks, [], email)
                self.assertEqual({text for text in strings_of(answer) if text.startswith('donor:')},
                                 {app.INDIVIDUALS_ID, 'donor:EXAMPLE CORP', 'donor:SAMPLE PAC', 'donor:unitemized'})
                nodes = {node['id']: node for node in answer['nodes']}
                self.assertEqual(list(nodes), ['committee:C00000001', app.INDIVIDUALS_ID, 'donor:EXAMPLE CORP',
                                               'donor:SAMPLE PAC', 'donor:unitemized', 'other:0', 'balance:refunded:1'])
                # one aggregate node: no per-person field, amounts summed to the cent
                self.assertEqual(nodes[app.INDIVIDUALS_ID], {
                    'id': app.INDIVIDUALS_ID, 'stage': 0, 'label': 'Individual contributors (2)', 'kind': 'donor',
                    'contribution': True, 'value_in': 0.0, 'value_out': 300.3})
                # committees, companies, PACs and the unitemized node are as built
                original = {node['id']: node for node in body['nodes']}
                for ident in ('committee:C00000001', 'donor:EXAMPLE CORP', 'donor:SAMPLE PAC', 'donor:unitemized',
                              'balance:refunded:1'):
                    expected = dict(original[ident], **({'details': [app.REFUNDS_HIDDEN]} if 'details' in original[ident] else {}))
                    self.assertEqual(nodes[ident], expected, ident)
                # the Other bucket keeps its label and count, not its members
                self.assertEqual(nodes['other:0'], {k: v for k, v in original['other:0'].items() if k != 'members'})
                # links are repointed and summed: what flows into the committee is what was built, and rows add up
                self.assertEqual(into(answer), into(body))
                self.assertEqual(into(answer), Decimal('1000.37'))
                link = next(link for link in answer['links'] if link['source'] == app.INDIVIDUALS_ID)
                self.assertEqual((link['target'], link['kind'], link['value'], link['rows']),
                                 ('committee:C00000001', 'receipt', 300.3, 2))
                self.assertEqual(link['evidence'], [{'amount': 100.1}, {'amount': 200.2}])
                self.assertEqual(len(answer['links']), len(body['links']) - 1)
                # everything outside nodes and links is as built
                self.assertEqual({k: v for k, v in answer.items() if k not in ('nodes', 'links')},
                                 {k: v for k, v in body.items() if k not in ('nodes', 'links')})
            self.assertEqual(json.dumps(body), before, 'the builder\'s own body is not changed')
            os.environ['MENTAT_DONOR_ROW_EMAILS'] = ' Someone@Example.test , ' + EMAIL.upper()
            status, _, answer = self.flow()
            self.assertEqual(status, 200)
            self.assertEqual(json.dumps(answer, ensure_ascii=False, separators=(',', ':')),
                             json.dumps(body, ensure_ascii=False, separators=(',', ':')),
                             'a listed account gets the body as built, byte for byte')

    def test_flow_aggregates_donors_of_no_known_type_unless_the_email_is_listed(self):
        body = individuals_flow_body()
        # an untyped name without a comma and a type nothing here knows are people until a filing says otherwise;
        # each explicit organization or committee type keeps its name
        added = [('ACME HARDWARE', '', 70.0), ('MYSTERY HOLDINGS', 'XYZ', 30.0), ('NAMED ORG', 'ORG', 11.0),
                 ('NAMED COM', 'COM', 12.0), ('NAMED PTY', 'PTY', 13.0), ('NAMED CCM', 'CCM', 14.0),
                 ('NAMED CAN', 'CAN', 15.0)]
        for name, entity, value in added:
            body['nodes'].append({'id': 'donor:' + name, 'stage': 0, 'label': name.title(), 'kind': 'donor',
                                  'contribution': True, 'value_in': 0.0, 'value_out': value, 'name_as_filed': name,
                                  'entity': entity})
            body['links'].append({'source': 'donor:' + name, 'target': 'committee:C00000001', 'value': value,
                                  'kind': 'receipt', 'rows': 1})
        with patch('flow_data.flow', return_value=body), patch.dict(os.environ):
            os.environ.pop('MENTAT_DONOR_ROW_EMAILS', None)
            for email in (EMAIL, 'stranger@example.test'):
                status, _, answer = self.flow(email)
                self.assertEqual(status, 200)
                text = json.dumps(answer).upper()
                for hidden in ('ACME', 'MYSTERY'):
                    self.assertNotIn(hidden, text, email)
                nodes = {node['id']: node for node in answer['nodes']}
                for name, entity, _ in added[2:]:
                    self.assertEqual(nodes['donor:' + name]['label'], name.title(), name)
                # the two typed people and the two of unknown type, in one node
                self.assertEqual((nodes[app.INDIVIDUALS_ID]['label'], nodes[app.INDIVIDUALS_ID]['value_out']),
                                 ('Individual contributors (4)', 400.3), email)
                self.assertEqual([link['value'] for link in answer['links'] if link['source'] == app.INDIVIDUALS_ID],
                                 [400.3])
            os.environ['MENTAT_DONOR_ROW_EMAILS'] = EMAIL
            status, _, answer = self.flow()
            self.assertEqual((status, answer), (200, json.loads(json.dumps(body))), 'a listed account gets it as built')

    def comparison(self, query='', email=EMAIL):
        return self.served.request('/api/comparison?id=published-reports' + query, headers={
            'Host': HOST, 'Cf-Access-Jwt-Assertion': self.idp.token(email=email)})

    @needs_openpyxl
    def test_comparison_has_no_person_subject_unless_the_email_is_listed(self):
        body = comparison_body(self.temp.name)
        self.assertEqual({subject['kind'] for subject in body['subjects']},
                         {'Industry', 'Company', 'Recipient', 'Candidate', 'Person'}, 'the fixture must carry a Person')
        person = next(subject for subject in body['subjects'] if subject['kind'] == 'Person')
        kept = [subject for subject in body['subjects'] if subject['kind'] != 'Person']
        with patch.object(self.console, 'comparison', Mock(return_value=body), create=True) as reader, \
                patch.dict(os.environ):
            os.environ.pop('MENTAT_DONOR_ROW_EMAILS', None)
            for email in (EMAIL, 'stranger@example.test'):
                status, _, answer = self.comparison(email=email)
                self.assertEqual(status, 200)
                # a recursive walk of every key and string: no employee's name, person id or Person subject id
                leaks = [text for text in strings_of(answer) if text == 'Person' or any(piece in text for piece in (
                    'Alex', 'PERSON:1', '"Person"'))]
                self.assertEqual(leaks, [], email)
                # company, industry, recipient and candidate subjects, and everything else, are as built
                self.assertEqual(answer['subjects'], kept)
                self.assertEqual({key: value for key, value in answer.items() if key != 'subjects'},
                                 {key: value for key, value in body.items() if key != 'subjects'})
                self.assertTrue(any(subject['kind'] == 'Company' for subject in answer['subjects']))
            # a request that selects a Person subject is refused before anything is read
            reader.reset_mock()
            for query in ('&kind=Person', '&kind=person', '&subject=' + quote(person['id']),
                          '&subjects=' + quote(person['id'])):
                status, _, answer = self.comparison(query)
                self.assertEqual((status, answer), (403, {'error': app.HOSTED_PERSON_LEVEL_REFUSED}), query)
            reader.assert_not_called()
            self.assertEqual(self.comparison('&kind=Company')[0], 200)
            os.environ['MENTAT_DONOR_ROW_EMAILS'] = ' Someone@Example.test , ' + EMAIL.upper()
            status, _, answer = self.comparison()
            self.assertEqual((status, answer), (200, body), 'a listed account gets the body as built')
            self.assertEqual(self.comparison('&kind=Person')[0], 200)
            self.assertEqual(self.comparison('', 'stranger@example.test')[2]['subjects'], kept)

    @needs_openpyxl
    def test_a_header_with_any_person_word_is_refused_unless_the_email_is_listed(self):
        path = Path(self.temp.name) / 'header-report.xlsx'
        refused = ['Employee Name', 'Email', 'E-mail', 'Payee', 'Given Name', 'Contact', 'Phone', 'Telephone', 'Surname',
                   'First Name', 'Last Name', 'Donor', 'Contributor', 'Mailing Address', 'Street 1', 'City', 'ZIP5',
                   'Employer', 'Occupation', 'Job Title', 'Person Key', 'EmployeeName', 'Company Name', 'Employee Emails',
                   # on the real '* Summary' and '<Company> View' sheets, the Top Contributor rows' contributor; these
                   # sheets have no Section column, so they are refused (with one it is redacted: see the summary tests)
                   'Name / Recipient']
        kept = ['Contributor Tier', 'Employee Tier', 'Employee', 'Employee %', 'Employee D', 'Employee R %',
                'Contributors', 'Distinct Contributors', 'Total Giving', 'Candidate / Officeholder', 'Reference Key']
        header_workbook(path, refused + kept)
        sheet = lambda number, **headers: self.get(f'/api/report-data?id=r&index=0&sheet=S{number}', **headers)
        with patch.object(self.console, 'report_path', Mock(return_value=path)), patch.dict(os.environ):
            os.environ.pop('MENTAT_DONOR_ROW_EMAILS', None)
            for number, header in enumerate(refused):
                self.assertEqual(sheet(number)[:3:2], (403, {'error': app.HOSTED_PERSON_LEVEL_REFUSED}), header)
            for number, header in enumerate(kept, len(refused)):
                status, _, body = sheet(number)
                self.assertEqual((status, body['row_count']), (200, 1), header)
            os.environ['MENTAT_DONOR_ROW_EMAILS'] = EMAIL
            for number, header in enumerate(refused):
                self.assertEqual(sheet(number)[0], 200, header)

    @needs_openpyxl
    def test_summary_sheets_are_redacted_unless_the_email_is_listed(self):
        path = Path(self.temp.name) / 'summary-report.xlsx'
        summary_workbook(path)
        sheet = lambda name, header=0, email=EMAIL: self.get(
            f'/api/report-data?id=r&index=0&sheet={quote(name)}' + (f'&header={header}' if header else ''),
            token=self.idp.token(email=email))
        # Name / Recipient and Reference Key (columns 4 and 7) are blank on every row whose Section is not one of the
        # three recipient and aggregate ones; the rest of each row is as built
        expected = [row if index in SUMMARY_KEPT else blanked(row, 4, 7) for index, row in enumerate(SUMMARY_ROWS)]
        with patch.object(self.console, 'report_path', Mock(return_value=path)), patch.dict(os.environ):
            os.environ.pop('MENTAT_DONOR_ROW_EMAILS', None)
            built = app.worksheet_entry(path, 'Example Co View')
            for email in (EMAIL, 'stranger@example.test'):
                status, headers, body = sheet('Example Co View', email=email)
                self.assertEqual(status, 200, email)
                self.assertEqual(body['rows'], expected)
                self.assertEqual((body['row_count'], body['row_numbers']), (8, list(range(4, 12))))
                text = json.dumps(body)
                for hidden in SUMMARY_HIDDEN:
                    self.assertNotIn(hidden, text, hidden)
                for kept in ('Example PAC', 'key-pac', 'Sample Candidate', 'key-cand', 'Example Co', 'key-co'):
                    self.assertIn(kept, text, kept)
                # the column profile follows the rows sent, not the ones redacted; the others are as built
                columns = {column['name']: column for column in body['columns']}
                built_columns = {column['name']: column for column in json.loads(built.body)['columns']}
                for name in ('Name / Recipient', 'Reference Key'):
                    self.assertEqual((columns[name]['missing'], columns[name]['distinct']), (5, 3), name)
                    self.assertGreater(built_columns[name]['distinct'], 3, name)
                for name in set(columns) - {'Name / Recipient', 'Reference Key'}:
                    self.assertEqual(columns[name], built_columns[name], name)
                self.assertEqual({key: value for key, value in body.items() if key not in ('rows', 'columns')},
                                 {key: value for key, value in json.loads(built.body).items() if key not in ('rows', 'columns')})
                self.assertEqual(headers['ETag'], response_cache.from_value(body).etag, 'the tag follows the sent bytes')
                self.assertNotEqual(headers['ETag'], built.etag)
            # the sheet's own header row asked for by number is the same sheet
            self.assertEqual(sheet('Example Co View', header=3)[2]['rows'], expected)
            # another '* Summary', and a sheet whose Name / Recipient is twice there (the second is 'Name / Recipient (2)')
            self.assertEqual(sheet('2024 Summary')[2]['rows'], [blanked(row, 4, 7) if index < 2 else row
                                                                for index, row in enumerate(SUMMARY_ROWS[:3])])
            status, _, body = sheet('Twice View')
            self.assertEqual((status, body['rows']), (200, [['Top Contributor', None, None, None, 5],
                                                            ['Company Total', 'Example Co', 'Example Co', 'key-co', 9]]))
            self.assertEqual([column['name'] for column in body['columns']],
                             ['Section', 'Name / Recipient', 'Name / Recipient (2)', 'Reference Key', 'Total'])
            # a sheet that names no person is not touched
            self.assertEqual(sheet('Plain View')[2]['rows'], [['Company Total', 'Example Co', 9]])
            # read at a row that is not its header, nothing says which cells are a contributor's: refused, not sent as is
            # (row 4 is a contributor's row, which names a person in its own right; row 6 a recipient's, which does not)
            for header in (4, 6):
                self.assertEqual(sheet('Example Co View', header=header)[:3:2],
                                 (403, {'error': app.HOSTED_PERSON_LEVEL_REFUSED}), header)
            # a listed account gets every sheet as built
            os.environ['MENTAT_DONOR_ROW_EMAILS'] = ' Someone@Example.test , ' + EMAIL.upper()
            status, headers, body = sheet('Example Co View')
            self.assertEqual((status, body, headers['ETag']), (200, json.loads(built.body), built.etag))
            self.assertEqual(body['rows'], SUMMARY_ROWS)
            self.assertEqual(sheet('Example Co View', header=6)[0], 200)
            self.assertEqual(sheet('Example Co View', email='stranger@example.test')[2]['rows'], expected,
                             'the list names one account, not every account')

    @needs_openpyxl
    def test_a_summary_sheet_with_another_person_column_is_still_refused(self):
        path = Path(self.temp.name) / 'summary-report.xlsx'
        summary_workbook(path)
        refused = ('Employer View', 'Two Sections View', 'No Section Sheet')
        with patch.object(self.console, 'report_path', Mock(return_value=path)), patch.dict(os.environ):
            os.environ.pop('MENTAT_DONOR_ROW_EMAILS', None)
            for name in refused:
                status, _, body = self.get(f'/api/report-data?id=r&index=0&sheet={quote(name)}')
                self.assertEqual((status, body), (403, {'error': app.HOSTED_PERSON_LEVEL_REFUSED}), name)
            os.environ['MENTAT_DONOR_ROW_EMAILS'] = EMAIL
            for name in refused:
                self.assertEqual(self.get(f'/api/report-data?id=r&index=0&sheet={quote(name)}')[0], 200, name)

    @needs_openpyxl
    def test_a_worksheet_is_allowed_redacted_or_refused_by_its_header_row(self):
        import report_insights
        path = Path(self.temp.name) / 'summary-report.xlsx'
        summary_workbook(path)
        outcome = lambda name, header=0: report_insights.person_level_sheet(report_insights.stamp(path), name, header)
        self.assertEqual({name: outcome(name) for name in SUMMARY_SHEETS},
                         {'Example Co View': report_insights.REDACT, '2024 Summary': report_insights.REDACT,
                          'Twice View': report_insights.REDACT, 'Plain View': report_insights.ALLOW,
                          'Employer View': report_insights.REFUSE, 'Two Sections View': report_insights.REFUSE,
                          'No Section Sheet': report_insights.REFUSE})
        # the most restrictive of the sheet's own header row and the one asked for
        self.assertEqual(outcome('Plain View', 3), report_insights.ALLOW)
        self.assertEqual(outcome('Example Co View', 6), report_insights.REDACT)    # a row that names no person
        self.assertEqual(outcome('Example Co View', 4), report_insights.REFUSE)    # a row that does
        self.assertEqual(outcome('Plain View', 4), report_insights.ALLOW)
        self.assertEqual(outcome('No Such Sheet'), report_insights.REFUSE, 'fail closed')

    @needs_openpyxl
    def test_a_body_with_no_single_section_column_cannot_be_redacted(self):
        body = lambda *names: {'columns': [{'name': name} for name in names], 'rows': [['Top Contributor'] * len(names)]}
        for names in (('Name / Recipient', 'Total'), ('Section', 'Name / Recipient', 'Section')):
            self.assertIsNone(app.without_contributor_names(body(*names)), names)
        given = {'columns': [{'name': 'Section'}, {'name': 'Name / Recipient'}, {'name': 'Reference Key'}],
                 'rows': [['Top Contributor', 'Zed Quill', 'key-zed'], ['Company Total', 'Example Co', 'key-co']]}
        before = json.dumps(given)
        self.assertEqual(app.without_contributor_names(given)['rows'],
                         [['Top Contributor', None, None], ['Company Total', 'Example Co', 'key-co']])
        self.assertEqual(json.dumps(given), before, 'the body it is given is not changed')

    def test_json_bodies_carry_no_os_error_text_when_hosted(self):
        catalog = {'name': 'ok', 'note': "[Errno 2] No such file or directory: 'x'", 'rows': ['[WinError 3] gone', 'fine']}
        with patch.object(self.console, 'report_catalog', Mock(return_value=catalog), create=True):
            status, _, body = self.get('/api/reports?id=a')
        self.assertEqual((status, body), (200, {'name': 'ok', 'note': app.HOSTED_IO_ERROR,
                                                'rows': [app.HOSTED_IO_ERROR, 'fine']}))

    def test_a_cached_body_is_scrubbed_when_hosted(self):
        entry = response_cache.from_value(CACHED_BODY)
        scrubbed = {'sheet': 'Sources', 'rows': [['Company', '', 5], ['Other', '', 6]], 'note': app.HOSTED_IO_ERROR,
                    'ok': 'fine/relative', 'n': 3}
        with patch.object(app, 'worksheet_entry', return_value=entry):
            status, headers, body = self.get('/api/report-data?id=r&index=0')
            self.assertEqual((status, body), (200, scrubbed))
            self.assertEqual(headers['ETag'], response_cache.from_value(scrubbed).etag, 'the tag follows the sent bytes')
            self.assertNotEqual(headers['ETag'], entry.etag)
        clean = response_cache.from_value({'rows': [['a', 1]], 'ok': 'fine/relative'})
        with patch.object(app, 'worksheet_entry', return_value=clean):
            status, headers, body = self.get('/api/report-data?id=r&index=0')
        self.assertEqual((status, body, headers['ETag']), (200, {'rows': [['a', 1]], 'ok': 'fine/relative'}, clean.etag))

    def test_geography_errors_are_one_sentence_when_hosted(self):
        result = json.loads(json.dumps(GEOGRAPHY_RESULT))
        with patch.object(app.contribution_geo, 'geography', return_value=result):
            status, _, body = self.post('/api/contributions/geography', {'zip5': ['02134']})
        self.assertEqual(status, 200)
        self.assertEqual(body['errors'], {'points': app.HOSTED_IO_ERROR, 'states': app.HOSTED_IO_ERROR})
        self.assertEqual(body['warnings'], GEOGRAPHY_RESULT['warnings'])
        self.assertNotIn('Errno', json.dumps(body))
        self.assertEqual(result['errors'], GEOGRAPHY_RESULT['errors'], 'the reader\'s own answer is not changed')


class StateWriteRoot(unittest.TestCase):
    """MENTAT_STATE_WRITE_ROOT: the caches the console writes move; the inputs stay put."""

    PAIRS = (('contribution_query_store', 'ROOT'), ('analysis_history', 'ROOT'),
             ('contribution_source_guard', 'CACHE_ROOT'), ('analysis_data', 'ROOT'), ('response_cache', 'ROOT'),
             ('analysis_data', 'HINTS_PATH'), ('access_stats', 'PATH'), ('contribution_geo', 'CACHE'),
             ('glassbox_emit', 'EVENTS_DIR'),
             ('contribution_data', '_materialize'), ('contribution_data', 'CACHE_ROOT'),
             ('contribution_source_guard', 'ACCEPTED_ROOT'))

    def setUp(self):
        import importlib
        import drawer_sidecars
        import prewarm
        self.prewarm, self.drawer = prewarm, drawer_sidecars
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name) / 'write-root'
        self.shipped = Path(self.temp.name) / 'shipped-state'
        modules = {name: importlib.import_module(name) for name, _ in self.PAIRS}
        self.before = {pair: getattr(modules[pair[0]], pair[1]) for pair in self.PAIRS}
        self.modules = modules
        previous = drawer_sidecars.configure()
        project_state = Path(app.PROJECT) / 'state'
        self.state_listing = sorted(str(p) for p in project_state.rglob('*')) if project_state.exists() else []

        def restore():
            for (name, attribute), value in self.before.items():
                setattr(modules[name], attribute, value)
            drawer_sidecars.configure(**previous)
            contribution_data._jobs.pop('tok', None)
            self.temp.cleanup()

        self.addCleanup(restore)
        (self.shipped / 'contribution-source-proofs').mkdir(parents=True)
        (self.shipped / 'contribution-source-proofs' / 'a.json').write_text('{"proof": 1}')
        (self.shipped / 'geography' / 'census-v1').mkdir(parents=True)
        (self.shipped / 'geography' / 'census-v1' / 'b.json').write_text('{"census": 1}')

    def redirect(self):
        with patch.object(self.prewarm, 'DEFAULT_STATE_ROOT', self.shipped):
            return app.redirect_state_writes(self.root)

    def test_writers_move_and_inputs_stay(self):
        self.assertEqual(self.redirect(), self.root.resolve())
        root = self.root.resolve()
        for (name, attribute), original in self.before.items():
            now = getattr(self.modules[name], attribute)
            if attribute in ('CACHE_ROOT', 'ACCEPTED_ROOT') and (name, attribute) in self.prewarm.READ_ONLY_STATE:
                self.assertEqual(now, original, f'{name}.{attribute} is a shipped input')
            elif attribute != '_materialize':
                self.assertTrue(Path(now).is_relative_to(root), f'{name}.{attribute} -> {now}')
        self.assertEqual(self.modules['contribution_source_guard'].CACHE_ROOT, root / 'contribution-source-proofs')
        self.assertEqual(self.modules['response_cache'].ROOT, root / 'response-cache')
        self.assertTrue(self.drawer.state_root().is_relative_to(root))
        self.assertFalse(self.drawer._settings['emit'])

    def test_what_the_console_writes_lands_under_the_root_and_nothing_under_state(self):
        self.redirect()
        store = self.modules['contribution_query_store']
        store._atomic(store.ROOT / 'responses' / 'x.json', b'1')
        self.assertTrue((self.root / 'contribution-query-cache' / 'responses' / 'x.json').is_file())
        geo = self.modules['contribution_geo']
        self.assertEqual(geo._cached('k', lambda: {'a': 1})['data'], {'a': 1})
        self.assertEqual(len(list((self.root / 'geography' / 'census-v1').glob('*.json'))) , 2, 'seeded file + new one')
        self.drawer._write(self.drawer.state_root() / 'native' / 's.bin', {'k': 1}, b'body')
        self.assertTrue((self.root / 'drawer-sidecars' / 'native' / 's.bin').is_file())
        project_state = Path(app.PROJECT) / 'state'
        after = sorted(str(p) for p in project_state.rglob('*')) if project_state.exists() else []
        self.assertEqual(after, self.state_listing)

    def test_the_shipped_caches_are_seeded_once_and_not_touched(self):
        self.redirect()
        self.assertEqual((self.root / 'contribution-source-proofs' / 'a.json').read_text(), '{"proof": 1}')
        self.assertEqual((self.root / 'geography' / 'census-v1' / 'b.json').read_text(), '{"census": 1}')
        (self.root / 'contribution-source-proofs' / 'a.json').write_text('changed here')
        self.redirect()
        self.assertEqual((self.root / 'contribution-source-proofs' / 'a.json').read_text(), 'changed here')
        self.assertEqual((self.shipped / 'contribution-source-proofs' / 'a.json').read_text(), '{"proof": 1}')

    def test_hosted_installs_the_write_guard_and_a_local_redirect_does_not(self):
        reports = Path(self.temp.name) / 'reports'
        with patch.object(self.prewarm, 'DEFAULT_STATE_ROOT', self.shipped), \
                patch.object(self.prewarm, 'install_write_guard') as guard:
            app.redirect_state_writes(self.root)
            guard.assert_not_called()
            app.redirect_state_writes(self.root, reports)
        guard.assert_called_once_with(guarded=(app.PROJECT, reports), allowed=(self.root.resolve(), tempfile.gettempdir()))

    def test_a_missing_index_is_refused_not_built(self):
        self.redirect()
        built = Path(self.temp.name) / 'index.sqlite'
        contribution_data._materialize({'database': 'x'}, 'tok', built)
        job = contribution_data._jobs['tok']
        self.assertEqual((job['status'], job['message']), ('error', self.prewarm.INDEX_REFUSED))
        self.assertFalse(built.exists())


WRITE_GUARD_PROBE = textwrap.dedent("""
    import json, os, sys
    from pathlib import Path
    app_dir, shipped, root, reports, attempts = sys.argv[1:6]
    sys.path.insert(0, app_dir)
    import server, prewarm
    prewarm.DEFAULT_STATE_ROOT = Path(shipped)
    server.redirect_state_writes(root, reports)
    out = {}
    for name, path in json.loads(attempts).items():
        try:
            Path(path).write_text('x')
            out[name] = 'written'
        except PermissionError:
            out[name] = 'refused'
    try:
        os.makedirs(Path(reports) / 'new')
        out['mkdir'] = 'made'
    except PermissionError:
        out['mkdir'] = 'refused'
    out['read'] = len(Path(server.__file__).read_text(encoding='utf-8')) > 0
    print(json.dumps(out))
""")


class HostedWriteGuard(unittest.TestCase):
    """What a hosted start does after redirect_state_writes, in its own process: the guard is process-wide."""

    def test_a_write_outside_the_write_root_is_refused(self):
        probe = Path(app.PROJECT) / '_write_guard_probe.tmp'
        self.addCleanup(lambda: probe.unlink(missing_ok=True))
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            for name in ('tmp', 'write-root', 'reports', 'shipped'):
                (base / name).mkdir()
            attempts = {'write root': base / 'write-root' / 'x.txt', 'temp dir': base / 'tmp' / 'x.txt',
                        'report folder': base / 'reports' / 'x.txt', 'checkout': probe}
            env = {**os.environ, 'TMP': str(base / 'tmp'), 'TEMP': str(base / 'tmp'), 'TMPDIR': str(base / 'tmp'),
                   'PYTHONDONTWRITEBYTECODE': '1'}
            done = subprocess.run([sys.executable, '-B', '-c', WRITE_GUARD_PROBE, str(SERVER_PY.parent),
                                   str(base / 'shipped'), str(base / 'write-root'), str(base / 'reports'),
                                   json.dumps({key: str(value) for key, value in attempts.items()})],
                                  env=env, capture_output=True, text=True, timeout=120)
            self.assertEqual(done.returncode, 0, done.stderr)
            self.assertEqual(json.loads(done.stdout.strip().splitlines()[-1]),
                             {'write root': 'written', 'temp dir': 'written', 'report folder': 'refused',
                              'checkout': 'refused', 'mkdir': 'refused', 'read': True})
        self.assertFalse(probe.exists())


class ReceiptPaths(unittest.TestCase):
    """Receipts and accepted-source pointers hold Windows paths; Linux must still find the files."""

    def as_linux(self, text):
        with patch.object(receipt_paths, 'os', SimpleNamespace(name='posix', environ={})):
            return receipt_paths.receipt_path(text)

    def relative_to_project(self, text):
        return self.as_linux(text).relative_to(receipt_paths.PROJECT).as_posix()

    def test_a_windows_path_is_rebased_onto_this_checkout(self):
        for text, expected in (
                ('C:/Users/pgoff/Mentat/mentat/.cache/report-repairs/banking-20260910/data/mentat.db',
                 '.cache/report-repairs/banking-20260910/data/mentat.db'),
                ('C:\\Users\\pgoff\\Mentat\\mentat\\state\\accepted-report-sources\\credit_union\\2022-2024-2026\\x.db',
                 'state/accepted-report-sources/credit_union/2022-2024-2026/x.db'),
                ('c:\\users\\PGOFF\\MENTAT\\MENTAT\\.cache\\y.db', '.cache/y.db')):
            self.assertEqual(self.relative_to_project(text), expected)
        outside = self.as_linux('C:/Users/pgoff/Mentat/tmp/production-all/runs/x/data/mentat.db')
        self.assertEqual(outside.relative_to(receipt_paths.PROJECT.parent).as_posix(), 'tmp/production-all/runs/x/data/mentat.db')

    def test_everything_else_is_left_alone(self):
        for text in ('/srv/mentat/.cache/x.db', 'relative/x.db', 'D:/elsewhere/x.db', ''):
            self.assertEqual(self.as_linux(text), Path(text), text)

    def test_windows_keeps_the_path_it_was_given(self):
        if os.name != 'nt':
            self.skipTest('Windows only')
        self.assertEqual(receipt_paths.receipt_path('C:/Users/pgoff/Mentat/mentat/.cache/x'), Path('C:/Users/pgoff/Mentat/mentat/.cache/x'))

    def test_the_readers_use_it(self):
        for module in (contribution_data, contribution_source_guard, source_memo):
            self.assertIs(module.receipt_path, receipt_paths.receipt_path, module.__name__)
        # contribution_data._sources and source_memo._discover are too heavy to drive here: guard the call sites.
        for function, needles in ((contribution_data._sources, ("receipt_path(metadata['database_path'])",
                                                                "receipt_path(hydration['output'])")),
                                  (source_memo._discover, ("receipt_path(data._payload_metadata(payload)['database_path'])",
                                                           "receipt_path(hydration['output'])"))):
            source = inspect.getsource(function)
            for needle in needles:
                self.assertIn(needle, source, function.__name__)

    def test_verify_finds_the_snapshot_a_windows_pointer_names(self):
        guard = contribution_source_guard
        pointer = {'snapshot': 'C:\\Users\\pgoff\\Mentat\\mentat\\state\\accepted-report-sources\\x\\2022-2024-2026\\s.db'}
        source = {'hydration_manifest': 'h', 'industry': 'X', 'database': 'recovered.db', 'database_sha256': 'a',
                  'original_database_sha256': 'b'}
        fake_stamp = lambda path: guard.LOADED_ALGORITHM_STAMP if path == guard.__file__ else (str(path), 1, 1, 1)
        with patch.object(guard, '_stamp', side_effect=fake_stamp), patch.object(guard, '_read', return_value=pointer), \
                patch.object(guard, '_quiescent'), patch.object(guard, '_verify_cached', return_value='proof') as verified, \
                patch.object(receipt_paths, 'os', SimpleNamespace(name='posix', environ={})):
            self.assertEqual(guard.verify(source), 'proof')
        expected = (guard.ACCEPTED_ROOT / 'x' / '2022-2024-2026' / 's.db').resolve()
        self.assertEqual(Path(verified.call_args.args[2][0]), expected)

    def test_an_accepted_pointer_written_on_windows_matches_its_snapshot_on_linux(self):
        guard = contribution_source_guard
        with tempfile.TemporaryDirectory() as temp:
            temp = Path(temp)
            database = temp / 'recovered.db'
            database.write_bytes(b'db')
            pointer = temp / 'current.json'
            pointer.write_text(json.dumps({
                'status': 'verified', 'industry': 'X', 'cycles': list(guard.CYCLES), 'database_sha256': 'original',
                'snapshot': 'C:\\Users\\pgoff\\Mentat\\mentat\\state\\accepted-report-sources\\x\\2022-2024-2026\\s.db'}))
            snapshot = guard.ACCEPTED_ROOT / 'x' / '2022-2024-2026' / 's.db'
            stamp = lambda path: (str(path.resolve()), 1, 1, 1)
            arguments = (stamp(database), stamp(pointer), stamp(snapshot), 'not-its-hash', 'original', 'X', 'algorithm',
                         str(temp))
            # The pointer names the snapshot, so verification moves on to the next check (the database's hash).
            with patch.object(receipt_paths, 'os', SimpleNamespace(name='posix', environ={})):
                with self.assertRaisesRegex(ValueError, 'recovered database changed'):
                    guard._verify_cached.__wrapped__(*arguments)
            # Named some other snapshot, it is refused at the pointer.
            pointer.write_text(pointer.read_text().replace('\\\\s.db', '\\\\other.db'))
            with patch.object(receipt_paths, 'os', SimpleNamespace(name='posix', environ={})):
                with self.assertRaisesRegex(ValueError, 'does not match the original published source scope'):
                    guard._verify_cached.__wrapped__(*arguments)

    def test_a_stored_windows_source_path_still_names_its_industry(self):
        record = {'name': 'upload.csv', 'catalog': {'row_count': 0, 'committees': []},
                  'source_path': 'C:\\Users\\pgoff\\Downloads\\Banking Industry Spending Overview.xlsx'}
        self.assertEqual(evidence_csv.stored_catalog(record)['industry'], 'Banking')


class KeyStoreImport(unittest.TestCase):
    def test_wintypes_is_not_needed_to_import_the_key_store(self):
        tools = Path(app.PROJECT) / 'tools'
        code = ('import sys; sys.path.insert(0, %r); import fec_key_store; '
                'sys.exit(1 if "ctypes.wintypes" in sys.modules else 0)' % str(tools))
        done = subprocess.run([sys.executable, '-B', '-c', code], capture_output=True, text=True, timeout=60)
        self.assertEqual(done.returncode, 0, done.stderr)

    @unittest.skipUnless(os.name == 'nt', 'DPAPI is Windows only')
    def test_the_dpapi_round_trip_still_works(self):
        sys.path.insert(0, str(Path(app.PROJECT) / 'tools'))
        import fec_key_store
        self.assertEqual(fec_key_store.crypt(fec_key_store.crypt(b'not-a-key'), decrypt=True), b'not-a-key')


class HostedStartup(unittest.TestCase):
    """server.py main(), run for real on an OS-chosen port with every path in a temp dir."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.reports = Path(self.temp.name) / 'reports'
        self.reports.mkdir()
        self.port = free_port()

    def command(self, *extra):
        return [sys.executable, '-B', str(SERVER_PY), '--port', str(self.port), '--data-root',
                str(Path(self.temp.name) / 'data'), *extra]

    def environment(self, **over):
        env = {k: v for k, v in os.environ.items()
               if not k.startswith(('ACCESS_', 'MENTAT_')) and k not in ('BIND_HOST', 'ALLOWED_HOSTS', 'GIT_SHA')}
        env.update(MENTAT_REPORT_DIRECTORY=str(self.reports), PYTHONDONTWRITEBYTECODE='1',
                   GLASSBOX_EVENTS_DIR=str(Path(self.temp.name) / 'glassbox'))
        env.update(over)
        return env

    def run_to_exit(self, env, *extra):
        return subprocess.run(self.command(*extra), env=env, capture_output=True, text=True, timeout=120)

    def test_a_non_loopback_bind_without_the_access_environment_does_not_start(self):
        done = self.run_to_exit(self.environment(BIND_HOST='0.0.0.0'), '--read-only')
        self.assertEqual(done.returncode, 2, done.stderr)
        self.assertIn('ACCESS_AUD', done.stderr)

    def run_main(self, bound, env, *, gate_check=None):
        """server.main() in this process with everything past the bind stubbed (so that, without the bind check, it
        runs on to serve_forever and touches nothing). Returns (the SystemExit code or None, stderr, the stand-ins)."""
        server = SimpleNamespace(server_address=(bound, 1), serve_forever=Mock(), server_close=Mock())
        argv = ['server.py', '--read-only', '--data-root', str(Path(self.temp.name) / 'data')]
        stderr, code = io.StringIO(), None
        with patch.dict(os.environ, env, clear=True), patch.object(sys, 'argv', argv), \
                patch.object(app, 'console_server', return_value=server), patch.object(app, 'Console') as console, \
                patch.object(app, 'redirect_state_writes'), patch('source_memo.install'), patch.object(app.os, 'chdir'), \
                patch.object(app.Handler, 'access_cfg', app.Handler.access_cfg), \
                patch.object(app.access_gate, 'assert_bind_consistent', gate_check or app.access_gate.assert_bind_consistent), \
                contextlib.redirect_stderr(stderr), contextlib.redirect_stdout(io.StringIO()):
            try:
                app.main()
            except SystemExit as stop:
                code = stop.code
        return code, stderr.getvalue(), SimpleNamespace(server=server, console=console)

    def test_main_does_not_serve_when_the_address_bound_is_not_the_gates(self):
        # not enforced, yet bound to every interface: the gate's own check, before a console exists
        code, stderr, stubs = self.run_main('0.0.0.0', self.environment())
        self.assertEqual(code, 2, stderr)
        self.assertIn("access gate: bound to '0.0.0.0' but the access gate is not enforced", stderr)
        stubs.console.assert_not_called()
        stubs.server.serve_forever.assert_not_called()
        stubs.server.server_close.assert_called_once()
        # the same stand-in bound where the gate expects it serves
        code, stderr, stubs = self.run_main('127.0.0.1', self.environment())
        self.assertEqual((code, stderr), (None, ''))
        stubs.server.serve_forever.assert_called_once()

    @needs_crypto
    def test_a_hosted_main_asks_the_gate_about_the_address_it_bound(self):
        idp = Idp(self.temp.name)
        env = self.environment(**idp.env(MENTAT_STATE_WRITE_ROOT=str(Path(self.temp.name) / 'state-write')))
        refuse = Mock(side_effect=gate.AccessConfigError('bound to the wrong address'))
        code, stderr, stubs = self.run_main('127.0.0.1', env, gate_check=refuse)
        self.assertEqual((code, stderr), (2, 'access gate: bound to the wrong address\n'))
        refuse.assert_called_once()
        config, bound = refuse.call_args.args
        self.assertEqual((config.enforced, bound), (True, '127.0.0.1'))
        stubs.console.assert_not_called()
        stubs.server.serve_forever.assert_not_called()

    def test_an_enforced_console_must_be_read_only(self):
        env = self.environment(ACCESS_ENFORCE='1', ACCESS_AUD=AUD, ACCESS_ISSUER=ISS,
                               ACCESS_JWKS_FILE=str(Path(self.temp.name) / 'certs.json'))
        done = self.run_to_exit(env)
        self.assertEqual(done.returncode, 2, done.stderr)
        self.assertIn('--read-only', done.stderr)
        self.assertNotIn('Mentat app', done.stdout)

    @needs_crypto
    def test_a_hosted_console_without_a_state_write_root_does_not_start(self):
        idp = Idp(self.temp.name)
        for name, extra in (('unset', {}), ('blank', {'MENTAT_STATE_WRITE_ROOT': '  '})):
            done = self.run_to_exit(self.environment(**idp.env(**extra)), '--read-only')
            self.assertEqual(done.returncode, 2, (name, done.stderr))
            self.assertIn('MENTAT_STATE_WRITE_ROOT', done.stderr)
            self.assertNotIn('Mentat app', done.stdout)

    @needs_crypto
    def test_a_hosted_console_serves_only_behind_access(self):
        idp = Idp(self.temp.name)
        env = self.environment(**idp.env(MENTAT_STATE_WRITE_ROOT=str(Path(self.temp.name) / 'state-write'),
                                         GIT_SHA='abc1234'))
        process = subprocess.Popen(self.command('--read-only'), env=env, stdout=subprocess.PIPE,
                                   stderr=subprocess.PIPE, text=True)
        try:
            def call(path, headers):
                connection = http.client.HTTPConnection('127.0.0.1', self.port, timeout=10)
                try:
                    connection.request('GET', path, headers=headers)
                    response = connection.getresponse()
                    return response.status, response.read()
                finally:
                    connection.close()

            deadline = time.time() + 120
            while True:
                self.assertIsNone(process.poll(), 'the console exited during start-up')
                try:
                    if call('/healthz', {'Host': f'127.0.0.1:{self.port}'})[0] == 200:
                        break
                except OSError:
                    pass
                self.assertLess(time.time(), deadline, 'the console did not come up')
                time.sleep(.25)
            self.assertEqual(call('/api/session?light=1', {'Host': HOST})[0], 403)
            status, body = call('/api/session?light=1', {'Host': HOST, 'Cf-Access-Jwt-Assertion': idp.token()})
            self.assertEqual((status, json.loads(body)['token'], json.loads(body)['read_only']), (200, 'hosted-read-only', True))
            status, body = call('/version', {'Host': f'127.0.0.1:{self.port}'})
            self.assertEqual((status, json.loads(body)['git_sha']), (200, 'abc1234'))
        finally:
            process.terminate()
            try:
                out, err = process.communicate(timeout=20)
            except subprocess.TimeoutExpired:
                process.kill()
                out, err = process.communicate()
        self.assertIn('hosted, read-only', out)
        self.assertNotIn('MENTAT_STATE_WRITE_ROOT', err)
        lines = [json.loads(line) for line in err.splitlines() if line.startswith('{')]
        self.assertTrue(any(line['route'] == '/api/session' and line['status'] == 200 and line['user'] == EMAIL
                            for line in lines), err)
        self.assertTrue(any(line.get('reason') == 'no token' for line in lines), err)


if __name__ == '__main__':
    unittest.main()
