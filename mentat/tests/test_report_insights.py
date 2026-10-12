"""Offline freshness, worksheet scanning and local HTTP regressions."""
import csv
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'app'))
import report_insights as insights
import server as app


class ReportInsightsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.report = self.root / 'Example.xlsx'

    def book(self, sheets):
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, rows in sheets.items():
            sheet = book.create_sheet(name)
            for row in rows:
                sheet.append(row)
        book.save(self.report)
        book.close()

    def test_sheet_list_reads_metadata_without_materializing_cells(self):
        self.book({'Data': [['Company','Amount'],['A',12]], 'History': [['Year','Value']]})
        with patch.object(insights.openpyxl, 'load_workbook', side_effect=AssertionError('cell scan not needed')):
            result = insights.worksheet_data(insights.stamp(self.report))
        self.assertEqual(result['sheets'], ['Data','History'])
        self.assertEqual(result['sha256'], hashlib.sha256(self.report.read_bytes()).hexdigest())

    def test_missing_receipts_do_not_use_mtime_as_coverage(self):
        self.book({'Data': [['Company','Amount'],['A',12]]})
        os.utime(self.report, (1800000000,1800000000))
        result = insights.report_freshness(self.report)
        self.assertEqual(result['freshness'], 'unknown')
        self.assertEqual(result['sources'], [])
        self.assertTrue(result['report_modified_at'].endswith('+00:00'))
        self.assertIn('No report-bound', result['issues'][0])

    def test_retrieval_check_remote_mtime_and_verified_date_are_distinct(self):
        self.book({'Source Freshness': [
            ['Source','Cycle','Downloaded At','Checked At','Server Last Modified','Data Through','Data Through Verified'],
            ['cm',2026,'2026-08-01T10:00:00Z','2026-09-08T10:00:00Z','Mon, 07 Sep 2026 05:00:00 GMT','2026-07-31',True],
            ['oth',2026,'bad date','2026-09-08T10:00:00Z','Mon, 07 Sep 2026 05:00:00 GMT','2026-08-31',False]]})
        result = insights.report_freshness(self.report)
        first, second = result['sources']
        self.assertEqual(first['retrieved_at'], '2026-08-01T10:00:00+00:00')
        self.assertEqual(first['checked_at'], '2026-09-08T10:00:00+00:00')
        self.assertEqual(first['server_last_modified'], '2026-09-07T05:00:00+00:00')
        self.assertEqual(first['verified_data_through'], '2026-07-31')
        self.assertIsNone(first['file_modified_at'])
        self.assertEqual(first['freshness'], 'verified')
        self.assertIsNone(second['verified_data_through'])
        self.assertIsNone(second['retrieved_at'])
        self.assertEqual(second['freshness'], 'unknown')
        self.assertEqual(result['freshness'], 'unknown')

    def test_checked_at_and_max_transaction_date_are_not_retrieval_or_coverage(self):
        self.book({'Source Freshness': [['Source','Checked At','Max Transaction Date','Verified Data Through'],
                                       ['cm','2026-09-08T01:00:00Z','2026-09-01','2026-02-30']]})
        row = insights.report_freshness(self.report)['sources'][0]
        self.assertIsNone(row['retrieved_at'])
        self.assertIsNone(row['verified_data_through'])
        self.assertEqual(row['freshness'], 'unknown')

    def test_excel_date_cell_and_malformed_receipt_remain_readable(self):
        self.book({'Source Freshness': [['Source','Verified Data Through Date'],['cm',datetime(2026,7,31)]]})
        result = insights.report_freshness(self.report)
        self.assertEqual(result['sources'][0]['verified_data_through'],'2026-07-31')
        audit = self.report.with_suffix('.audit')
        audit.mkdir()
        (audit / 'audit-manifest.json').write_text('not JSON')
        result = insights.report_freshness(self.report)
        self.assertTrue(result['issues'])
        self.assertEqual(result['sources'][0]['verified_data_through'],'2026-07-31')
        self.report.write_bytes(b'not a workbook')
        result = insights.report_freshness(self.report)
        self.assertEqual(result['freshness'],'unknown')
        self.assertEqual(result['sources'],[])

    def receipt_csv(self):
        directory = self.root / 'Example.audit'
        directory.mkdir(exist_ok=True)
        receipt = directory / 'source_freshness.csv'
        receipt.write_text('source,cycle,downloaded_at,verified_data_through\ncm,2026,2026-08-01T00:00:00Z,2026-07-31\n', encoding='utf-8')
        return receipt

    def test_source_receipts_use_hash_bound_csv_and_invalidate_without_workbook_change(self):
        receipt = self.receipt_csv()
        self.book({'Source Receipts': [['Report title'],[],[],[],['Industry','Audit File','SHA-256'],
                                      ['Test',str(receipt), insights.digest(receipt)]]})
        first = insights.report_freshness(self.report)
        self.assertEqual(first['sources'][0]['verified_data_through'],'2026-07-31')
        receipt.write_text(receipt.read_text().replace('2026-07-31','2026-08-31'))
        changed = insights.report_freshness(self.report)
        self.assertEqual(changed['sources'][0]['freshness'], 'unknown')
        self.assertIn('checksum',changed['issues'][0])
        receipt.unlink()
        self.assertEqual(insights.report_freshness(self.report)['sources'][0]['freshness'], 'unknown')

    def test_companion_audit_must_match_workbook_and_csv(self):
        receipt = self.receipt_csv()
        self.book({'Data': [['Company','Amount'],['A',12]]})
        manifest = {'workbook_name':self.report.name, 'workbook_sha256':insights.digest(self.report),
                    'csv_sha256':{'source_freshness.csv':insights.digest(receipt)}}
        (receipt.parent / 'audit-manifest.json').write_text(json.dumps(manifest))
        self.assertEqual(insights.report_freshness(self.report)['verified_source_count'], 1)
        self.book({'Data': [['Company','Amount'],['B',99]]})
        result = insights.report_freshness(self.report)
        self.assertEqual(result['sources'], [])
        self.assertIn('different workbook',result['issues'][0])

    def test_report_reference_cannot_read_outside_report_directory_or_network(self):
        for reference in ('../source_freshness.csv','//server/share/source_freshness.csv','https://example.com/source_freshness.csv'):
            with self.subTest(reference=reference):
                with self.assertRaises(ValueError):
                    insights.local_reference(reference,self.root)

    def test_pinned_source_enrichment_requires_the_exact_report_source(self):
        source = self.root / 'inputs/fec_committee_master/cm26'
        source.mkdir(parents=True)
        source_file = source / 'cm.txt'
        source_file.write_text('old snapshot')
        signature = insights.digest(source_file)
        receipt = {'family':'cm','cycle':2026,'file':'cm.txt','sha256':signature,'downloaded_at':'2026-08-01T00:00:00Z',
                   'checked_at':'2026-09-08T00:00:00Z'}
        (source / 'source-receipt.json').write_text(json.dumps(receipt))
        self.book({'Source Freshness': [['FEC Dataset','Cycle','SHA-256','Checked At'],['cm',2026,signature,'2026-08-02T00:00:00Z']]})
        root = self.root / 'inputs'
        value = insights.report_freshness(self.report,root)['sources'][0]
        self.assertIsNotNone(value['file_modified_at'])
        self.assertEqual(value['checked_at'],'2026-08-02T00:00:00+00:00')
        self.assertEqual(value['freshness'],'unknown')
        source_file.write_text('new snapshot')
        changed = insights.report_freshness(self.report,root)
        self.assertIsNone(changed['sources'][0]['file_modified_at'])
        self.assertIn('no longer matches',changed['issues'][0])
        receipt['sha256'] = insights.digest(source_file)
        (source / 'source-receipt.json').write_text(json.dumps(receipt))
        self.assertIsNone(insights.report_freshness(self.report,root)['sources'][0]['retrieved_at'])

    def test_column_scan_types_duplicates_missing_formulas_and_immutability(self):
        self.book({'Funding': [['Funding by company'],['Explanatory text'],[],[],
                              ['Company','Amount','Amount','Cycle','Date','Uncalculated'],
                              ['A',10,'$1,250.50',2026,'2026-01-01','=B6*2'],
                              ['B',None,'($50.00)',2026,'2026-01-02',None]]})
        before = self.report.read_bytes()
        result = insights.worksheet_data(insights.stamp(self.report),'Funding')
        self.assertEqual(result['header_row'],5)
        self.assertEqual(result['row_numbers'],[6,7])
        self.assertEqual([c['type'] for c in result['columns']],['category','number','number','category','date','category'])
        self.assertEqual(result['columns'][2]['name'],'Amount (2)')
        self.assertEqual(result['columns'][1]['missing'],1)
        self.assertIsNone(result['rows'][0][-1])
        self.assertEqual(self.report.read_bytes(),before)

    def test_row_one_and_manual_subsection_header_and_limits(self):
        self.book({'Data': [['Name','Value'],['A',4],['B',5],['New group','New value'],['C',9]]})
        self.assertEqual(insights.worksheet_data(insights.stamp(self.report),'Data')['header_row'],1)
        self.assertEqual(insights.worksheet_data(insights.stamp(self.report),'Data',4)['rows'],[['C',9]])
        with self.assertRaises(ValueError): insights.worksheet_data(insights.stamp(self.report),'Data',101)
        with self.assertRaises(ValueError): insights.worksheet_data(insights.stamp(self.report),'Missing')
        insights.worksheet_data.cache_clear()
        with patch.object(insights,'MAX_ROWS',2), self.assertRaisesRegex(ValueError,'no partial chart'):
            insights.worksheet_data(insights.stamp(self.report),'Data')

    def test_changed_workbook_invalidates_column_cache(self):
        self.book({'Data': [['Name','Value'],['A',4]]})
        first = insights.worksheet_data(insights.stamp(self.report),'Data')
        self.book({'Data': [['Name','Value'],['B',88]]})
        second = insights.worksheet_data(insights.stamp(self.report),'Data')
        self.assertNotEqual(first['sha256'],second['sha256'])
        self.assertEqual(second['rows'],[['B',88]])

    def test_overlapping_summary_views_remain_selectable_with_explicit_grain(self):
        self.book({'Spending Overview':[['Cycle','Scope','Total Giving'],['2023-2024','Total Giving',100],['2023-2024','Corporate',40],['2023-2024','Employee',60],['2023-2024','Candidate',100]],
                   'Recipient List':[['Cycle','View','Total Giving'],['2023-2024','Industry',100],['2023-2024','Company',100]]})
        overview=insights.worksheet_data(insights.stamp(self.report),'Spending Overview')
        self.assertEqual(len(overview['rows']),4)
        self.assertEqual(overview['row_numbers'],[2,3,4,5])
        self.assertEqual(overview['accounting']['independent_dimensions'],[1])
        self.assertIsNone(overview['accounting']['row_role_column'])
        recipient=insights.worksheet_data(insights.stamp(self.report),'Recipient List')
        self.assertEqual(recipient['rows'],[['2023-2024','Industry',100],['2023-2024','Company',100]])
        self.assertEqual(recipient['accounting']['independent_dimensions'],[1])

    def test_tidy_data_book_cells_type_and_summability_as_the_writer_intends(self):
        # The companion Data workbook writes REAL Excel types, not text: dates
        # as datetimes and flags as booleans. numeric() returns None for a bool
        # (report_insights.py:215), so a flag must classify as a dimension and
        # never as a measure a reader could sum; cell_value isoformats a
        # datetime, so a date column must classify 'date' and stay chartable on
        # a time axis. Both are silent failures if they regress - a summed flag
        # column produces a plausible number.
        self.book({'Giving Facts': [
            ['Election Cycle','Date','Include in Analysis','Giving',
             'Distinct Contributors'],
            ['2023-2024',datetime(2024,1,1),True,1250.5,12],
            ['2023-2024',datetime(2024,2,1),False,900.25,9]]})
        result = insights.worksheet_data(insights.stamp(self.report),'Giving Facts')
        self.assertEqual(result['header_row'],1)
        columns = {c['name']:c for c in result['columns']}
        self.assertEqual(columns['Date']['type'],'date')
        self.assertEqual(columns['Include in Analysis']['type'],'category')
        self.assertEqual(columns['Giving']['type'],'number')
        self.assertEqual(columns['Giving']['unit'],'currency')
        self.assertTrue(columns['Giving'].get('accounting',{}).get('safe_to_sum',True))
        self.assertFalse(
            columns['Distinct Contributors']['accounting']['safe_to_sum'])

    def test_all_text_sheet_takes_row_one_as_its_header(self):
        # Every dimension sheet in the Data book (Committees, Candidates,
        # Committee Sponsors) is entirely text. header_row picks the row with
        # the most DISTINCT text cells and breaks ties on the earliest row, so
        # a header of unique names always wins - but only while the headers
        # stay unique. A duplicated header would hand row 1 away to a data row
        # and shift every column by one.
        self.book({'Committees': [
            ['Committee ID','Committee','Committee Type','Mapping Source'],
            ['C00000001','CONDUIT ONE','Super PAC','FEC master'],
            ['C00000002','LEADERSHIP TWO','PAC - qualified','Multiple']]})
        result = insights.worksheet_data(insights.stamp(self.report),'Committees')
        self.assertEqual(result['header_row'],1)
        self.assertEqual([c['name'] for c in result['columns']],
                         ['Committee ID','Committee','Committee Type','Mapping Source'])
        self.assertEqual(result['row_numbers'],[2,3])

    def test_per_cycle_split_sheets_keep_their_sheet_level_summability(self):
        # The Data workbook splits any sheet over MENTAT_DATA_SPLIT_ROWS into
        # `<Sheet> <cycle_label>` parts. The sheet-level non-summable list is
        # matched by EXACT name, so before the base-name match those parts were
        # offered as safe to sum - the association-path double count, split per
        # cycle and harder to see than the unsplit version it replaced.
        self.book({'Transfer Paths 2021-2022':[['Funded Committee','Reported Net Transfer'],
                                               ['C00000001',100],['C00000002',100]],
                   'Giving Facts 2023-2024':[['Company','Giving'],['A',100],['B',100]]})
        split = insights.worksheet_data(insights.stamp(self.report),'Transfer Paths 2021-2022')
        self.assertFalse(split['accounting']['safe_to_sum'])
        self.assertIn('association paths',split['accounting']['reason'])
        # A split sheet that was never on the list stays summable: the change
        # must not make every per-cycle sheet non-summable.
        facts = insights.worksheet_data(insights.stamp(self.report),'Giving Facts 2023-2024')
        self.assertTrue(facts['accounting']['safe_to_sum'])
        giving = next(c for c in facts['columns'] if c['name'] == 'Giving')
        self.assertEqual(giving['unit'],'currency')

    def test_split_parts_keep_their_sheet_s_summability_under_either_name(self):
        # S1 (Paul, 2026-09-26): the writer now names a part `<Sheet> <election
        # year>`; books published before keep `<Sheet> <cycle_label>` until they
        # are rebuilt, and both keep the sheet-level non-summable flag.
        rows = [['Candidate','Associated Funding'],['A',100],['B',100]]
        self.book({'Transfer Paths 2022':[['Funded Committee','Reported Net Transfer'],
                                          ['C00000001',100],['C00000002',100]],
                   'Beneficiary Bridge 2024':rows,'Beneficiary Bridge 2023-2024':rows,
                   'Funding Cohorts 2026':rows,
                   'Giving Facts 2024':[['Company','Giving'],['A',100],['B',100]]})
        for name in ('Transfer Paths 2022','Beneficiary Bridge 2024','Beneficiary Bridge 2023-2024',
                     'Funding Cohorts 2026'):
            with self.subTest(name=name):
                result = insights.worksheet_data(insights.stamp(self.report),name)
                self.assertFalse(result['accounting']['safe_to_sum'])
                self.assertIn('association paths',result['accounting']['reason'])
        facts = insights.worksheet_data(insights.stamp(self.report),'Giving Facts 2024')
        self.assertTrue(facts['accounting']['safe_to_sum'])

    def test_split_sheet_base_names_the_sheet_a_part_belongs_to(self):
        for name, base in (('Transfer Paths 2022','Transfer Paths'),
                           ('Transfer Paths 2021-2022','Transfer Paths'),
                           ('Network Candidates 2022','Network Candidate Activity'),
                           ('Network Candidates 2025-2026','Network Candidate Activity'),
                           ('Transfer Paths','Transfer Paths'),
                           ('Network Candidates','Network Candidates'),
                           ('2023-2024','2023-2024'),('2024','2024'),
                           ('Transfer Paths 22','Transfer Paths 22'),
                           ('Recipient List','Recipient List'),('',''),(None,'')):
            with self.subTest(name=name):
                self.assertEqual(base,insights.split_sheet_base(name))

    def test_raw_amount_and_transfer_context_cannot_silently_be_summed(self):
        self.book({'2023-2024':[['Source ID','Amount','Analysis Amount'],['a',100,100],['b',100,0]],
                   'Beneficiary Bridge':[['Candidate','Associated Funding'],['A',100],['B',100]]})
        result=insights.worksheet_data(insights.stamp(self.report),'2023-2024')
        self.assertFalse(result['columns'][1]['accounting']['safe_to_sum'])
        self.assertTrue(result['columns'][2]['accounting']['safe_to_sum'])
        self.assertFalse(insights.worksheet_data(insights.stamp(self.report),'Beneficiary Bridge')['accounting']['safe_to_sum'])

    def test_http_report_data_rejects_arbitrary_paths_incomplete_reports_and_hosts(self):
        self.book({'Data': [['Name','Value'],['A',4]]})
        directory = self.root / 'run'
        directory.mkdir()
        (directory / 'status.json').write_text(json.dumps({'industries':[
            {'status':'succeeded','output_workbook':str(self.report)},
            {'status':'running','output_workbook':str(self.report)}]}))
        console = object.__new__(app.Console)
        console.report_directory = self.root
        console.find = lambda identifier: {'directory':directory} if identifier == 'test' else None
        server = app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
        server.app = console
        worker = threading.Thread(target=server.serve_forever,daemon=True)
        worker.start()
        origin = f'http://127.0.0.1:{server.server_port}'
        try:
            with urlopen(origin+'/api/report-data?id=test&index=0&sheet=Data') as response:
                self.assertEqual(json.load(response)['rows'],[['A',4]])
                self.assertIn("default-src 'self'",response.headers['Content-Security-Policy'])
            for query in ('id=../test&index=0','id=test&index=1','id=test&index=-1','id=test&index=999','path='+str(self.report)):
                with self.subTest(query=query), self.assertRaises(HTTPError) as error:
                    urlopen(origin+'/api/report-data?'+query)
                self.assertEqual(error.exception.code,400)
            with self.assertRaises(HTTPError) as error:
                urlopen(Request(origin+'/api/report-data?id=test&index=0',headers={'Host':'untrusted.example'}))
            self.assertEqual(error.exception.code,403)
        finally:
            server.shutdown(); server.server_close(); worker.join()


if __name__ == '__main__':
    unittest.main()
