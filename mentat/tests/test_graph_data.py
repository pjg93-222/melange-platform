"""Guided Graph regressions: scope, unique joins, directions and report bindings."""
from pathlib import Path
import json
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import openpyxl

sys.path.insert(0,str(Path(__file__).resolve().parents[1] / 'app'))
import graph_data
from graph_data import graph_view, report_facts, singleflight_cache
from report_insights import digest, stamp
import server as app
from graph_fixture import add_graph_sheets


class GraphTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.path=self.root/'Giving.xlsx'
        book=openpyxl.Workbook()
        add_graph_sheets(book)
        book.save(self.path)
        book.close()

    def view(self,**filters):
        return graph_view([self.path],filters)

    def edit(self,fn):
        book=openpyxl.load_workbook(self.path)
        fn(book)
        book.save(self.path)
        book.close()

    def test_report_activity_keeps_optional_filing_records(self):
        records = [{'amount':700,'source_url':'https://docquery.fec.gov/cgi-bin/fecimg/?202401010000000001'},
                   {'amount':77,'source_url':'https://docquery.fec.gov/cgi-bin/fecimg/?202401010000000002'}]
        def add_records(book):
            sheet = book['Downstream Activity']
            sheet['H5'] = 'Filing Records'
            sheet['H6'] = json.dumps(records)
        self.edit(add_records)
        activity = graph_data.report_activity(stamp(self.path), include_filings=True)
        match = next(row for row in activity if row['committee'] == 'C00000003'
                     and row['direction'] == 'Support')
        self.assertEqual(match['filings'], records)

    def test_industry_totals_use_one_fact_partition(self):
        before=self.path.read_bytes()
        value=self.view()
        self.assertEqual(value['totals']['giving'],1585)
        self.assertEqual(value['totals']['direct'],90)
        self.assertEqual(value['totals']['unresolved_candidate_links'],20)
        self.assertEqual(value['totals']['organizations'],1475)
        self.assertEqual(sum(r['amount'] for r in value['organizations']),1475)
        self.assertEqual({r['id']:r['amount'] for r in value['levels']},
                         {'Chief':370,'Middle-Management':300,'Employees':415})
        self.assertEqual({r['id']:r['amount'] for r in value['categories']},
                         {'Super PAC':1200,'JFC':200,'Candidate':110,'Trade Association':50,'Leadership PAC':25})
        self.assertEqual(sum(r['amount'] for r in value['companies']),1585)
        self.assertEqual(self.path.read_bytes(),before)

    def test_levels_by_company_split_each_company_without_corporate(self):
        # "By company" under Giving by employee level (board M4): the same
        # non-Corporate rows as `levels`, per company, each company's levels in
        # the legend's order (Employees, Middle-Management, Chief) with a
        # missing level at 0, the largest company first.
        value=self.view()
        self.assertEqual([(r['company'],r['amount'],[(l['id'],l['amount']) for l in r['levels']])
                          for r in value['levels_by_company']],
                         [('Harborlight Credit Union',625,[('Employees',25),('Middle-Management',300),('Chief',300)]),
                          ('Civic Horizon Bank',460,[('Employees',390),('Middle-Management',0),('Chief',70)])])
        # Corporate is left out: Harborlight's $500 corporate receipt is in
        # companies[] (1125) but in no level, and the per-company splits add up
        # to the levels themselves.
        self.assertEqual({r['id']:r['amount'] for r in value['companies']}['Harborlight Credit Union'],1125)
        self.assertEqual(sum(r['amount'] for r in value['levels_by_company']),
                         sum(r['amount'] for r in value['levels']))
        for company in value['levels_by_company']:
            self.assertEqual(sum(l['amount'] for l in company['levels']),company['amount'])
        self.assertEqual(self.view(funding='corporate')['levels_by_company'],[])
        # A selected company is its own single split.
        self.assertEqual(self.view(company='Civic Horizon Bank')['levels_by_company'],
                         [{'company':'Civic Horizon Bank','amount':460,'levels':[
                             {'id':'Employees','label':'Employees','amount':390},
                             {'id':'Middle-Management','label':'Middle-Management','amount':0},
                             {'id':'Chief','label':'Chief','amount':70}]}])

    def test_workbook_routes_reach_barr_without_snapshot_or_duplicate_paths(self):
        def extend(book):
            sheet=book.create_sheet('Transfer Paths')
            for row in [['Routes'],[],['3 records; 3 shown. Complete destination data.'],[],
                        ['Cycle','Source Committee','Destination Committee','Reported Net Transfer','Traversal Status'],
                        [2026,'C00000003','C00000009',100,'followed'],
                        [2026,'C00000003','C00000009',100,'followed'],
                        [2026,'C00000009','C00000003',1,'cycle']]:
                sheet.append(row)
            book['Downstream Activity'].append(['2025-2026','C00000009','H0KY06104','BARR, GARLAND ANDY','R','Support',106708])
            book['Downstream Activity']['A3']='101 source rows condensed to 5 committee-candidate-activity groups'
        self.edit(extend)
        with patch('downstream.inventory',return_value={}):
            value=self.view(mode='activity',category='Super PAC')
        self.assertTrue(value['candidate_available'],value['candidate_note'])
        self.assertEqual(next(r['amount'] for r in value['candidates'] if r['id']=='H0KY06104'),106708)
        self.assertEqual(len(value['downstream_edges']),2)
        # A capped route display must never masquerade as complete evidence.
        self.edit(lambda b:setattr(b['Transfer Paths']['A3'],'value','10001 records; 3 shown. Full records are in the accompanying audit CSV files.'))
        value=self.view(mode='activity')
        self.assertFalse(value['candidate_available'])
        self.assertIn('partial display',value['candidate_note'])

    def test_company_levels_and_corporate_scope_are_distinct(self):
        self.assertEqual(self.view(company='Harborlight Credit Union',tier='Chief')['totals']['giving'],300)
        self.assertEqual(self.view(company='Harborlight Credit Union',tier='Chief')['totals']['direct'],100)
        self.assertEqual(self.view(tier='Middle-Management')['totals']['giving'],300)
        self.assertEqual(self.view(tier='Employees')['totals']['giving'],415)
        self.assertEqual(self.view(funding='corporate')['totals']['giving'],500)
        self.assertEqual(self.view(funding='corporate')['levels'],[])

    def test_partial_transfer_preview_uses_only_matching_full_payload_and_revalidates(self):
        def preview(book):
            sheet=book.create_sheet('Transfer Paths')
            for row in [['Routes'],[],['3 records; 1 shown. First records only.'],[],
                        ['Cycle','Source Committee','Destination Committee','Reported Net Transfer','Traversal Status'],
                        [2026,'C00000003','C00000009',100,'followed']]:sheet.append(row)
        self.edit(preview)
        payload=self.root/'full.json';manifest=self.root/'manifest.json'
        edge={'cycle':2026,'source':'C00000003','destination':'C00000009','amount':100,'status':'followed'}
        payload.write_text(json.dumps({'transfer_paths':[edge,{**edge,'root':'another'},
            {**edge,'destination':'C00000010','status':'candidate_committee'}]}))
        manifest.write_text(json.dumps({'database_sha256':'pinned','data_sha256':digest(payload)}))
        with patch.object(graph_data,'transfer_source',return_value=(payload,manifest)), patch('contribution_data._workbook_source',return_value='pinned'):
            native=graph_data.report_transfers(stamp(self.path))
            graph_data.report_transfers.cache_clear()
            with patch.object(graph_data.native_scan,'tables',return_value=None):
                reference=graph_data.report_transfers(stamp(self.path))
            self.assertEqual(native,reference)
            self.assertEqual(len(native),1)
            self.assertEqual(native[0]['amount'],100)
            payload.write_text(json.dumps({'transfer_paths':[dict(edge,amount=999)]}))
            with self.assertRaisesRegex(ValueError,'fingerprint'):
                graph_data.report_transfers(stamp(self.path))
        self.assertEqual(self.view(funding='personal')['totals']['giving'],1085)
        with self.assertRaises(ValueError):self.view(tier='Chief',funding='corporate')

    def test_jfc_leadership_and_unauthorized_links_are_not_direct_candidates(self):
        for category in ['JFC','Leadership PAC','Super PAC']:
            result=self.view(category=category)
            self.assertEqual(result['candidates'],[])
            self.assertEqual(result['totals']['direct'],0)
        self.assertEqual(self.view()['candidates'][0]['amount'],90)
        self.edit(lambda b:b['PAC & Candidate Map'].cell(6,6,'H0CA00001; H0CA00002'))
        self.assertEqual(self.view()['candidates'],[])
        self.assertEqual(self.view()['totals']['unresolved_candidate_links'],110)

    def test_pac_activity_is_counted_once_for_industry_and_directions_stay_separate(self):
        self.assertEqual(self.view(mode='activity')['candidates'][0]['amount'],810)
        self.assertEqual(self.view(mode='activity',category='Super PAC')['candidates'][0]['amount'],777)
        self.assertEqual(self.view(mode='activity',direction='Oppose')['candidates'][0]['amount'],999)
        self.assertEqual(self.view(mode='activity',direction='Contribution / Other')['candidates'][0]['amount'],111)
        a=self.view(company='Harborlight Credit Union',tier='Middle-Management',mode='activity')
        b=self.view(company='Civic Horizon Bank',tier='Employees',mode='activity')
        self.assertEqual(a['candidates'][0]['amount'],777)
        self.assertEqual(b['candidates'][0]['amount'],777)
        self.assertIn('not money attributed',a['candidate_note'])

    def test_direction_totals_answer_every_side_from_one_read(self):
        # P5. The drawer asks for one side and has to label the others, so
        # every direction is summarised from the same eligible rows the
        # selected side was cut from - including the empty ones.
        # Paul, 2026-09-23, on the drawer's "Contributions / Other" segment:
        # "It should be two separate segments". The sides are now Contribution
        # and Other (the recorded bucket split by FEC transaction code) beside
        # Support and Oppose, then 'Contribution / Other' for amounts whose
        # source records no code, and Unknown. This fixture's Downstream
        # Activity sheet has no Transaction Type column - a workbook written
        # before it existed - so its 111 cannot be split and stays on the
        # recorded side; test_contribution_other_splits_by_transaction_type
        # below pins the split itself.
        value = self.view(mode='activity')
        self.assertEqual([r['direction'] for r in value['direction_totals']],
                         ['Contribution','Other','Support','Oppose','Contribution / Other','Unknown'])
        self.assertEqual({r['direction']:(r['candidates'],r['amount']) for r in value['direction_totals']},
                         {'Contribution':(0,0),'Other':(0,0),'Support':(1,810),'Oppose':(1,999),
                          'Contribution / Other':(1,111),'Unknown':(0,0)})
        # The selected side still ranks exactly what it always did.
        self.assertEqual(value['candidates'][0]['amount'],810)
        # One PAC narrows every side, not just the one being shown.
        one = self.view(mode='activity', recipient='C00000003')
        self.assertEqual({r['direction']:(r['candidates'],r['amount']) for r in one['direction_totals']},
                         {'Contribution':(0,0),'Other':(0,0),'Support':(1,777),'Oppose':(1,999),
                          'Contribution / Other':(1,111),'Unknown':(0,0)})
        # A side with no rows is reported as zero, never omitted.
        jfc = self.view(mode='activity', recipient='C00000002')
        self.assertEqual({r['direction']:(r['candidates'],r['amount']) for r in jfc['direction_totals']},
                         {'Contribution':(0,0),'Other':(0,0),'Support':(1,33),'Oppose':(0,0),
                          'Contribution / Other':(0,0),'Unknown':(0,0)})
        # The recorded bucket, each candidate once: here all of it is on the
        # unsplit side, so nothing is on more than one side.
        self.assertEqual(value['contribution_other_union'],
                         {'direction':'Contribution / Other','sides':['Contribution','Other','Contribution / Other'],
                          'candidates':1,'on_more_than_one_side':0,'amount':111,'unidentified_amount':0})
        # Only the activity view has sides to report.
        self.assertNotIn('direction_totals', self.view())
        self.assertNotIn('direction_totals', self.view(mode='estimates'))
        self.assertNotIn('contribution_other_union', self.view())
        self.assertNotIn('contribution_other_union', self.view(mode='estimates'))

    def test_contribution_other_splits_by_transaction_type(self):
        # Paul, 2026-09-23, on the drawer's "Contributions / Other" segment:
        # "It should be two separate segments". The workbook records one
        # direction for every non-independent-expenditure code, and carries
        # the code beside it (Transaction Type). Contribution is the codes
        # whose recorded Activity Type names a contribution (24K, 24P, 24Z);
        # Other is every other recorded code; an amount with no code stays
        # 'Contribution / Other'. The halves add up to the recorded bucket.
        def coded(book):
            sheet = book['Downstream Activity']
            sheet['H5'] = 'Transaction Type'
            for row, code in zip(range(6, 10), ['24E', '24A', '24K', '24E']):
                sheet.cell(row, 8, code)
            sheet.append(['2025-2026','C00000003','H0CA00001','Candidate One','D','Contribution / Other',40,'24C'])
            sheet.append(['2025-2026','C00000003','H0CA00002','Candidate Two','R','Contribution / Other',7,None])
            sheet.append(['2025-2026','C00000003','H0CA00003','Candidate Three','R','Contribution / Other',5,'24z'])
            sheet['A3'] = '100 source rows condensed to 7 committee-candidate-activity groups'
        self.edit(coded)
        value = self.view(mode='activity', recipient='C00000003')
        totals = {r['direction']:r for r in value['direction_totals']}
        self.assertEqual([r['direction'] for r in value['direction_totals']],
                         ['Contribution','Other','Support','Oppose','Contribution / Other','Unknown'])
        self.assertEqual({name:(r['candidates'],r['amount'],r['transaction_types']) for name,r in totals.items()},
                         {'Contribution':(2,116,['24K','24Z']),'Other':(1,40,['24C']),
                          'Support':(1,777,[]),'Oppose':(1,999,[]),
                          'Contribution / Other':(1,7,[]),'Unknown':(0,0,[])})
        # Nothing is re-attributed: the three sides add up to the recorded
        # bucket, 111 + 40 + 7 + 5; Candidate One is on both halves, so the
        # counts need not add (2 + 1 + 1 = 4 for 3 candidates).
        self.assertEqual(sum(totals[name]['amount'] for name in ('Contribution','Other','Contribution / Other')),163)
        self.assertEqual({r['id'] for name in ('Contribution','Other','Contribution / Other')
                          for r in value['candidates_by_direction'][name]},{'H0CA00001','H0CA00002','H0CA00003'})
        self.assertEqual(value['candidates_by_direction']['Other'],
                         [{'id':'H0CA00001','label':'Candidate One','amount':40,'party':'D'}])
        # Round-3 verify (D1): the segments' counts are each side's own, so
        # they need not add up to the candidates behind them. The response
        # says by how much, for the drawer to show under the segments: the
        # bucket's distinct candidates (3), those on more than one side (1,
        # Candidate One), and its amount, which the sides add up to.
        union = value['contribution_other_union']
        self.assertEqual(union,{'direction':'Contribution / Other','sides':['Contribution','Other','Contribution / Other'],
                                'candidates':3,'on_more_than_one_side':1,'amount':163,'unidentified_amount':0})
        counts = [totals[name]['candidates'] for name in union['sides']]
        self.assertEqual(sum(counts) - union['on_more_than_one_side'], union['candidates'])
        self.assertEqual(union['amount'], sum(totals[name]['amount'] for name in union['sides']))
        # Whichever side is asked for, the reconciliation is the same.
        self.assertEqual(self.view(mode='activity',recipient='C00000003',direction='Other')['contribution_other_union'],union)
        # The explorer follows the chosen side: each PAC-candidate tie carries it.
        ties = {(t['direction'],t['candidate_id']):t['amount'] for t in value['pac_candidate_relationships']
                if t['direction'] in ('Contribution','Other','Contribution / Other')}
        self.assertEqual(ties,{('Contribution','H0CA00001'):111,('Contribution','H0CA00003'):5,
                               ('Other','H0CA00001'):40,('Contribution / Other','H0CA00002'):7})
        # Each side is a request of its own, for the page's Direction select.
        self.assertEqual(self.view(mode='activity',direction='Other')['candidates'][0]['amount'],40)
        self.assertEqual([r['amount'] for r in self.view(mode='activity',direction='Contribution')['candidates']],[111,5])
        # The workbook read itself is unchanged: one recorded row per
        # committee, candidate and direction, its codes kept beside it.
        row = next(r for r in graph_data.report_activity(stamp(self.path))
                   if r['candidate_id'] == 'H0CA00001' and r['direction'] == 'Contribution / Other')
        self.assertEqual(row['amount'],151)
        self.assertEqual({code:part['amount'] for code,part in row['transaction_parts'].items()},{'24K':111,'24C':40})
        # The reference reader agrees with whichever reader ran above.
        graph_data.report_activity.cache_clear()
        with patch.object(graph_data.native_scan,'tables',return_value=None):
            reference = graph_data.report_activity(stamp(self.path))
        self.assertEqual({code:part['amount'] for r in reference if r['candidate_id'] == 'H0CA00001'
                          and r['direction'] == 'Contribution / Other' for code,part in r['transaction_parts'].items()},
                         {'24K':111,'24C':40})

    def test_filing_relationship_note_counts_each_observation_once_after_the_split(self):
        # r4 code review F1 (major): the drawer split ("It should be two
        # separate segments", Paul, 2026-09-23) turns one recorded
        # 'Contribution / Other' row into one row per side, each carrying the
        # recorded row's unverified/unresolved observation counts. The note
        # summed them over the sides, so a row split in two counted its
        # observations twice. It must count each observation once.
        def flagged(book):
            sheet = book['Downstream Activity']
            sheet['H5'] = 'Transaction Type'
            sheet['I5'] = 'Unverified Version Rows'
            sheet['J5'] = 'Unresolved Version Rows'
            sheet['H8'] = '24K'
            sheet['I8'] = 3
            sheet['J8'] = 1
            # The same committee, candidate and direction again, as a 24C
            # coordinated expenditure: one recorded group, two sides.
            sheet.append(['2025-2026','C00000003','H0CA00001','Candidate One','D','Contribution / Other',40,'24C',0,0])
            sheet['A3'] = '100 source rows condensed to 5 committee-candidate-activity groups'
        self.edit(flagged)
        recorded = graph_data.report_activity(stamp(self.path))
        row = next(r for r in recorded if r['direction'] == 'Contribution / Other')
        self.assertEqual((row['unverified_version_rows'], row['unresolved_version_rows']), (3, 1))
        self.assertEqual({code:part['amount'] for code,part in row['transaction_parts'].items()}, {'24K':111,'24C':40})
        # The negative control: the split really does put the one recorded
        # row on two sides, each still carrying 3 and 1 - summing over the
        # sides is what gave 6 and 2.
        sides = [r for r in graph_data.activity_sides(recorded) if r['candidate_id'] == 'H0CA00001'
                 and r['committee'] == 'C00000003' and r['direction'] in ('Contribution','Other')]
        self.assertEqual(sorted(r['direction'] for r in sides), ['Contribution','Other'])
        self.assertEqual(sum(r['unverified_version_rows'] for r in sides), 6)
        # The note counts the recorded observations once, whichever side is
        # asked for and whether or not the view is narrowed to the PAC.
        expected = sum(r['unverified_version_rows'] for r in recorded), sum(r['unresolved_version_rows'] for r in recorded)
        self.assertEqual(expected, (3, 1))
        for filters in ({}, {'direction':'Contribution'}, {'direction':'Other'}, {'recipient':'C00000003'}):
            with self.subTest(filters=filters):
                note = self.view(mode='activity', **filters)['candidate_note']
                self.assertIn('Filing relationships: 3 unverified observations retained; 1 unresolved observations in the saved source.', note)
                self.assertNotIn('6 unverified', note)
        # The amounts are the split's, unchanged by this: 111 + 40 on the sides.
        totals = {r['direction']:r['amount'] for r in self.view(mode='activity')['direction_totals']}
        self.assertEqual((totals['Contribution'], totals['Other']), (111, 40))

    def test_estimates_without_an_audit_bundle_say_so_in_words(self):
        # r4 code review (observed): estimates on a report written without
        # its audit bundle answered with the raw OS error and a local path.
        # F4 (Paul approved 2026-09-25): the words name what is missing and
        # what makes it (test_graph_not_traced.py has the F2 cases).
        value = self.view(mode='estimates')
        self.assertFalse(value['candidate_available'])
        self.assertEqual(value['candidates'], [])
        self.assertEqual(value['candidate_note'], 'No estimate bundle for this industry yet; rebuild the workbook.')
        self.assertNotIn('Errno', value['candidate_note'])

    def test_direction_totals_are_empty_when_the_downstream_sheet_is_refused(self):
        self.edit(lambda b:b['Downstream Activity'].cell(3,1,'100 source rows condensed to 8 committee-candidate-activity groups'))
        value = self.view(mode='activity')
        self.assertFalse(value['candidate_available'])
        self.assertEqual(value['direction_totals'],[])
        self.assertIsNone(value['contribution_other_union'])

    def test_recipient_filter_and_negative_values_reconcile(self):
        value=self.view(recipient='C00000003')
        self.assertEqual(value['totals']['giving'],1200)
        self.assertEqual({r['id']:r['amount'] for r in value['companies']},{'Harborlight Credit Union':800,'Civic Horizon Bank':400})
        value=self.view(company='Civic Horizon Bank',tier='Employees',category='Candidate')
        self.assertEqual(value['totals']['giving'],-10)
        self.assertEqual(value['candidates'][0]['amount'],-10)

    def test_partial_activity_is_unavailable_not_a_smaller_total(self):
        self.edit(lambda b:b['Downstream Activity'].cell(3,1,'100 source rows condensed to 8 committee-candidate-activity groups'))
        value=self.view(mode='activity')
        self.assertFalse(value['candidate_available'])
        self.assertEqual(value['candidates'],[])
        self.assertIn('partial display',value['candidate_note'])

    def test_broken_fact_partition_and_missing_amount_refuse_partial_graph(self):
        self.edit(lambda b:b['2025-2026'].cell(5,4,101))
        with self.assertRaisesRegex(ValueError,'reconcile'):self.view()
        self.edit(lambda b:setattr(b['2025-2026'].cell(5,4),'value',None))
        with self.assertRaisesRegex(ValueError,'missing or nonnumeric'):self.view()

    def test_unknown_level_is_not_reclassified_as_employee(self):
        def rename(book):
            for row in book['2025-2026'].iter_rows(min_row=5):
                if row[2].value=='Chief':row[2].value='Not classified'
        self.edit(rename)
        self.assertTrue(self.view()['options']['unknown_tiers'])
        self.assertEqual(self.view(tier='Unknown')['totals']['giving'],370)
        self.assertEqual(self.view(tier='Employees')['totals']['giving'],415)

    def test_overlapping_cycles_and_invalid_filters_rejected(self):
        with self.assertRaisesRegex(ValueError,'overlap'):graph_view([self.path,self.path],{})
        for filters in ({'company':'not in roster'},{'cycle':'2024'},{'tier':'CEO'},{'mode':'guess'},{'funding':'guess'}):
            with self.subTest(filters=filters),self.assertRaises(ValueError):self.view(**filters)

    def test_estimates_require_whole_industry_and_verified_full_csv(self):
        audit=self.path.with_suffix('.audit');audit.mkdir()
        csv=audit/'allocation_estimates.csv'
        csv.write_text('cycle,root,candidate_id,candidate_name,direction,estimated_amount\n2026,C00000003,H0CA00001,Candidate One,Support,123\n2026,C00000003,H0CA00001,Candidate One,Oppose,456\n')
        manifest={'workbook_name':self.path.name,'workbook_sha256':digest(self.path),
                  'csv_sha256':{csv.name:digest(csv)}}
        (audit/'audit-manifest.json').write_text(json.dumps(manifest))
        self.assertEqual(self.view(mode='estimates')['candidates'][0]['amount'],123)
        self.assertEqual(self.view(mode='estimates',direction='Oppose')['candidates'][0]['amount'],456)
        limited=self.view(company='Harborlight Credit Union',mode='estimates')
        self.assertFalse(limited['candidate_available'])
        self.assertEqual(limited['candidates'],[])
        csv.write_text(csv.read_text().replace('123','888'))
        changed=self.view(mode='estimates')
        self.assertFalse(changed['candidate_available'])
        self.assertEqual(changed['candidates'],[])

    def test_raw_receipt_layout_uses_only_included_rows(self):
        def convert(book):
            del book['2025-2026']
            s=book.create_sheet('2025-2026')
            for row in [['Raw'],[],[],['Company','Employee Tier','Amount','Recipient Category','Committee ID','Include in Analysis'],
                        ['Harborlight Credit Union','Chief',12,'Candidate','C00000001','Yes'],['Harborlight Credit Union','Chief',200,'Candidate','C00000001','No']]:s.append(row)
        self.edit(convert)
        self.assertEqual(self.view()['totals']['direct'],12)

    def test_report_changed_during_downstream_read_is_rejected(self):
        version = stamp(self.path)
        def changed(input_version):
            self.assertEqual(input_version, version)
            self.path.write_bytes(self.path.read_bytes() + b'changed')
            return []
        with patch('graph_data.report_activity', side_effect=changed):
            with self.assertRaisesRegex(ValueError, 'report changed'):
                self.view(mode='activity')

    def test_concurrent_requests_share_a_cold_read_and_other_files_can_load(self):
        from concurrent.futures import ThreadPoolExecutor, TimeoutError
        started, release, second_started = threading.Event(), threading.Event(), threading.Event()
        calls = []
        @singleflight_cache(maxsize=2)
        def load(version):
            calls.append(version)
            if version == 'first':
                started.set()
                if not release.wait(5):
                    raise RuntimeError('test release timed out')
            return object()
        def second_request():
            second_started.set()
            return load('first')
        with ThreadPoolExecutor(max_workers=3) as workers:
            first = workers.submit(load, 'first')
            self.assertTrue(started.wait(5))
            second = workers.submit(second_request)
            self.assertTrue(second_started.wait(5))
            try:
                with self.assertRaises(TimeoutError):
                    second.result(timeout=.05)
                self.assertIsNotNone(workers.submit(load, 'different').result(timeout=5))
            finally:
                release.set()
            self.assertIs(first.result(timeout=5), second.result(timeout=5))
        self.assertEqual(calls.count('first'), 1)
        self.assertEqual(calls.count('different'), 1)

    def test_failed_read_releases_the_next_request(self):
        calls = []
        @singleflight_cache(maxsize=1)
        def load(version):
            calls.append(version)
            if len(calls) == 1:
                raise ValueError('temporarily unavailable')
            return version
        with self.assertRaisesRegex(ValueError, 'temporarily unavailable'):
            load('report')
        self.assertEqual(load('report'), 'report')
        self.assertEqual(calls, ['report', 'report'])

    def test_a_stalled_read_stops_blocking_the_requests_behind_it(self):
        from concurrent.futures import ThreadPoolExecutor
        started, release = threading.Event(), threading.Event()
        calls = []
        @singleflight_cache(maxsize=2)
        def load(version):
            calls.append(version)
            if len(calls) == 1:
                started.set()
                if not release.wait(5):
                    raise RuntimeError('test release timed out')
            return version + ' read'
        # reading is closure state; only the caller that added a stamp may remove it.
        tokens = dict(zip(load.__code__.co_freevars, load.__closure__))['reading'].cell_contents
        with patch('graph_data.SINGLEFLIGHT_WAIT_SECONDS', .2), ThreadPoolExecutor(max_workers=2) as workers:
            stalled = workers.submit(load, 'stalled')
            try:
                self.assertTrue(started.wait(5))
                self.assertEqual(workers.submit(load, 'stalled').result(timeout=5), 'stalled read')
                self.assertFalse(stalled.done())
                self.assertEqual(tokens, {'stalled'})
            finally:
                release.set()
            self.assertEqual(stalled.result(timeout=5), 'stalled read')
            self.assertEqual(tokens, set())
            self.assertEqual(load('stalled'), 'stalled read')
        self.assertEqual(calls, ['stalled', 'stalled'])

    def test_http_catalog_and_graph_use_authoritative_report_folder_only(self):
        renamed = self.root / 'Test 2025-2026 Industry Spending Overview.xlsx'
        self.path.rename(renamed)
        self.path = renamed
        directory=self.root/'run';directory.mkdir()
        (directory/'status.json').write_text(json.dumps({'industries':[
            {'industry':'Test 2025-2026','status':'succeeded','output_workbook':str(self.path)},
            {'industry':'Pending','status':'running','output_workbook':str(self.path)}]}))
        console=object.__new__(app.Console)
        console.report_directory = self.root
        console.find=lambda identifier:{'directory':directory} if identifier=='fixture' else None
        server=app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler);server.app=console
        worker=threading.Thread(target=server.serve_forever,daemon=True);worker.start()
        origin=f'http://127.0.0.1:{server.server_port}'
        try:
            with urlopen(origin+'/api/graph-catalog?id=fixture') as response:
                self.assertEqual(json.load(response)['industries'], [])
            with urlopen(origin+'/api/graph-catalog?id=published-reports') as response:
                self.assertEqual(json.load(response)['industries'],[{'name':'Test','indices':[console.published_reports()[0]['index']]}])
            with urlopen(origin+'/api/comparison?id=published-reports&industry=Test') as response:
                comparison = json.load(response)
                self.assertEqual(next(s for s in comparison['subjects'] if s['kind']=='Industry')['periods']['2026']['giving'],1585)
            with urlopen(origin+'/api/graph?id=published-reports&industry=Test&latest=1') as response:
                self.assertEqual(json.load(response)['filters']['cycle'],'2026')
            with urlopen(origin+'/api/graph?id=published-reports&industry=Test&company=Harborlight%20Credit%20Union&tier=Chief') as response:
                body=json.load(response)
                self.assertEqual(body['totals']['direct'],100)
                self.assertEqual(body['reports'][0]['freshness']['freshness'],'unknown')
                self.assertEqual(body['reports'][0]['freshness']['sources'],[])
            for query in ['id=published-reports&industry=Pending','id=published-reports&industry=../../outside','id=fixture&industry=Test','id=other&industry=Test']:
                with self.assertRaises(HTTPError) as error:urlopen(origin+'/api/graph?'+query)
                self.assertEqual(error.exception.code,400)
            with self.assertRaises(HTTPError) as error:
                urlopen(Request(origin+'/api/graph-catalog?id=fixture',headers={'Host':'untrusted.example'}))
            self.assertEqual(error.exception.code,403)
        finally:
            server.shutdown();server.server_close();worker.join()


if __name__=='__main__':unittest.main()
