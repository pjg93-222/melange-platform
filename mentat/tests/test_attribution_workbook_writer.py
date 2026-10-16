"""Portable attribution workbooks preserve traced estimates, blanks and presentation.

Only synthetic data in temporary directories is used; no report directory is read.
Run: python -m unittest discover -s mentat/tests -p test_attribution_workbook_writer.py
"""
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

import openpyxl

BUILDER = Path(__file__).resolve().parents[1] / 'tools' / 'attribution_workbook_builder.mjs'
NODE = os.environ.get('MENTAT_NODE_PATH') or shutil.which('node')
SHEETS = ['Candidate Funding', 'Root Coverage', 'Data Quality', 'Source Receipts', 'Methodology']
MONEY = '$#,##0.00;[Red]($#,##0.00);$0.00'


@unittest.skipUnless(NODE, 'node is not available')
class AttributionWorkbookWriter(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def run_builder(self, payload, writer='openpyxl'):
        source, output = self.root / 'synthetic.json', self.root / 'synthetic.xlsx'
        source.write_text(json.dumps(payload), encoding='utf-8')
        env = {**os.environ, 'MENTAT_PYTHON': sys.executable,
               'MENTAT_NODE_MODULES': str(self.root / 'no-artifact-modules')}
        env.pop('MENTAT_XLSX_WRITER', None)
        if writer is not None:
            env['MENTAT_XLSX_WRITER'] = writer
        result = subprocess.run([NODE, str(BUILDER), str(source), str(output)], env=env,
                                capture_output=True, text=True, encoding='utf-8', timeout=60)
        return result, output

    def build(self, payload):
        result, output = self.run_builder(payload)
        self.assertEqual(result.returncode, 0, result.stderr)
        book = openpyxl.load_workbook(output, data_only=False)
        self.addCleanup(book.close)
        self.assertFalse(output.with_name(output.name + '.spec.ndjson').exists())
        return book

    def assert_structure(self, book):
        self.assertEqual(book.sheetnames, SHEETS)
        for name in SHEETS:
            sheet = book[name]
            self.assertEqual(sheet.freeze_panes, 'A6', name)
            self.assertEqual(len(sheet.tables), 1, name)
            table = next(iter(sheet.tables.values()))
            self.assertEqual(table.tableStyleInfo.name, 'TableStyleMedium2', name)
            self.assertEqual(sheet['A1'].fill.fgColor.rgb, 'FF123145', name)
            self.assertEqual(sheet['A1'].font.color.rgb, 'FFFFFFFF', name)
            self.assertTrue(sheet['A1'].font.bold, name)
            self.assertEqual(sheet['A1'].font.sz, 19, name)
            self.assertEqual(sheet.row_dimensions[1].height, 36, name)
            self.assertEqual(sheet.row_dimensions[5].height, 42, name)
            self.assertTrue(sheet['A5'].alignment.wrap_text, name)
        funding = book['Candidate Funding']
        self.assertEqual([funding.column_dimensions[c].width for c in 'ABCDEFGHI'],
                         [25, 10, 16, 37, 9, 25, 25, 25, 30])
        self.assertEqual(book['Methodology'].column_dimensions['B'].width, 120)
        self.assertEqual(book['Data Quality'].column_dimensions['E'].width, 100)
        self.assertEqual(book['Source Receipts'].column_dimensions['C'].width, 65)
        self.assertIn('A1:I1', funding.merged_cells)
        self.assertIn('A2:I3', funding.merged_cells)

    def test_mixed_traced_and_untraced_rows_keep_formulas_and_unknowns_distinct(self):
        rows = []
        # Adjacent traced rows, a gap, and a traced zero test formula segment boundaries.
        for index, estimate in enumerate((25, 15, None, 0, None)):
            rows.append(dict(industry='Synthetic', cycle=2024, candidate_id=f'SYN{index}',
                             candidate_name=f'Fixture {index}', party='X', direct=100 + index,
                             estimated_non_opposition=estimate,
                             estimated_opposition=None if estimate is None else 0,
                             direct_plus_estimated_non_opposition=None if estimate is None else 999))
        payload = dict(rows=rows,
                       coverage=[dict(industry='Synthetic', cycle='2024', root='SYNROOT',
                                      model_funding=100, estimated_candidate_activity=40,
                                      unresolved=60, unresolved_reasons='Synthetic gap')],
                       warnings=[dict(industry='Synthetic', id='transfer_index_2024', severity='warn',
                                      passed='False', detail='Synthetic untraced cycle')],
                       sources=[dict(industry='Synthetic', file='fixture.csv', sha256='a' * 64)])
        book = self.build(payload)
        self.assert_structure(book)
        funding = book['Candidate Funding']
        self.assertEqual(funding.tables['CandidateFundingTable'].ref, 'A5:I10')
        for row in (6, 7, 9):
            self.assertEqual(funding[f'I{row}'].value, f'=F{row}+G{row}')
            self.assertEqual(funding[f'I{row}'].data_type, 'f')
        for row in (8, 10):
            self.assertEqual([funding[f'{column}{row}'].value for column in 'GHI'], [None] * 3)
            self.assertNotEqual(funding[f'I{row}'].data_type, 'f')
            self.assertEqual(funding[f'F{row}'].value, 100 + row - 6)
        self.assertEqual(funding['G9'].value, 0)
        self.assertIn('Blank is not zero', funding['A2'].value)
        for row in range(6, 11):
            self.assertEqual(funding[f'I{row}'].number_format, MONEY)
        self.assertEqual(book['Root Coverage']['B6'].value, 2024)
        self.assertEqual(book['Root Coverage']['D6'].number_format, MONEY)
        self.assertIs(book['Data Quality']['D6'].value, False)
        self.assertTrue(book['Data Quality']['E6'].alignment.wrap_text)
        self.assertEqual(book['Data Quality'].row_dimensions[6].height, 70)
        self.assertEqual(book['Source Receipts']['C6'].value, 'a' * 64)

    def test_empty_payload_has_valid_placeholder_tables_and_formatted_blank_cells(self):
        book = self.build(dict(rows=[], coverage=[], warnings=[], sources=[]))
        self.assert_structure(book)
        for name, last_column in zip(SHEETS[:4], 'IGEC'):
            sheet = book[name]
            table = next(iter(sheet.tables.values()))
            self.assertEqual(table.ref, f'A5:{last_column}6')
            self.assertTrue(all(cell.value is None for cell in sheet[6]), name)
        self.assertEqual(book['Candidate Funding']['I6'].number_format, MONEY)
        self.assertEqual(book['Root Coverage']['F6'].number_format, MONEY)
        self.assertTrue(book['Data Quality']['E6'].alignment.wrap_text)
        self.assertEqual(book['Data Quality'].row_dimensions[6].height, 70)
        self.assertFalse(any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row))

    def test_unknown_writer_fails_and_default_still_resolves_artifact_tool(self):
        empty = dict(rows=[], coverage=[], warnings=[], sources=[])
        result, output = self.run_builder(empty, 'bogus')
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('MENTAT_XLSX_WRITER must be artifact or openpyxl, got bogus', result.stderr)
        self.assertFalse(output.exists())
        # An absent explicit module root proves the default did not silently choose openpyxl.
        result, output = self.run_builder(empty, None)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('@oai/artifact-tool', result.stderr)
        self.assertFalse(output.exists())


if __name__ == '__main__':
    unittest.main()
