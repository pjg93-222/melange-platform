"""FIX 4: the report's own collection_completeness check, read for display only.

The parser takes the check's Detail text apart. It never recomputes a count, so no
receipt or money code is involved; the cells are carried to Chart Studio, Graph and
Insights, which say "short" beside any chart whose company x cycle scope they touch.
"""
from pathlib import Path
import csv
import http.client
import importlib.util
import json
import sys
import tempfile
import threading
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'app'))
import collection_gaps
from graph_data import graph_view
from graph_fixture import add_graph_sheets

# The Detail of the audited Telecom reference, copied from its Data Quality sheet.
REAL_DETAIL = (
    "walks stored materially fewer rows than FEC reported - the cell looks populated but is short, "
    "recollect before trusting it: Lumen Technologies/2024 contributions 14133 of 28266; "
    "Comcast/2024 contributions 17869 of 29991; Comcast/2022 contributions 26441 of 34014; "
    "Lumen Technologies/2022 contributions 7388 of 13779; Verizon/2024 contributions 14783 of 19170; "
    "Verizon/2022 contributions 23415 of 27056; Comcast/2026 contributions 11753 of 15269; "
    "Comcast/2022 corporate 1369 of 4365; Comcast/2024 corporate 1299 of 4225; "
    "Verizon/2026 contributions 31654 of 33964; Altice USA/2024 contributions 1697 of 3394; "
    "AT&T/2022 corporate 7064 of 8432 ... and 19 more")
SERVED = Path('C:/Users/pgoff/OneDrive/Desktop/Mentat Complete Files/Telecom Industry Spending Overview.xlsx')
REFERENCE = Path(__file__).resolve().parents[2] / 'reports' / 'Telecom Industry Spending Overview 2025-2026.xlsx'


def by_cell(parsed):
    return {(c['company'], c['cycle'], c['destination']): c for c in parsed['cells']}


def quality_book(path, detail=REAL_DETAIL, passed=False, graph=False):
    book = openpyxl.Workbook()
    if graph:
        add_graph_sheets(book)
    sheet = book.create_sheet('Data Quality')
    for row in [['Data Quality'], [], [], [], ['Check', 'Severity', 'Passed', 'Detail'],
                ['roster_coverage', 'fail', True, 'all 10 roster companies have collected rows'],
                ['collection_completeness', 'warn', passed, detail]]:
        sheet.append(row)
    book.save(path)
    book.close()


def payload_file(path, detail=REAL_DETAIL, passed=False):
    """A pinned build payload: the metadata object first, as the generator writes it."""
    path.write_text(json.dumps({'metadata': {'industry': 'Telecom', 'export_guards': [
        {'id': 'roster_coverage', 'severity': 'fail', 'passed': True, 'detail': 'all 10 roster companies', 'metric': 10},
        {'id': 'collection_completeness', 'severity': 'warn', 'passed': passed, 'detail': detail, 'metric': 76249}]},
        'raw_rows': []}), encoding='utf-8')
    return path


def quality_csv(path, detail=REAL_DETAIL, passed='False'):
    folder = path.with_suffix('.audit')
    folder.mkdir()
    with (folder / 'quality_checks.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(['id', 'severity', 'passed', 'detail', 'metric'])
        writer.writerow(['roster_coverage', 'fail', 'True', 'all 10 roster companies have collected rows', '10'])
        writer.writerow(['collection_completeness', 'warn', passed, detail, '74442'])


class ParseDetailTests(unittest.TestCase):
    def test_the_real_detail_itemizes_twelve_cells_and_counts_nineteen_more(self):
        parsed = collection_gaps.parse_detail(REAL_DETAIL)
        self.assertEqual(len(parsed['cells']), 12)
        self.assertEqual(parsed['unitemized'], 19)
        cells = by_cell(parsed)
        comcast = cells['Comcast', '2025-2026', 'contributions']
        self.assertEqual((comcast['held'], comcast['reported'], comcast['percent']), (11753, 15269, 77))
        self.assertEqual(cells['Verizon', '2025-2026', 'contributions']['percent'], 93)
        self.assertEqual(cells['Lumen Technologies', '2023-2024', 'contributions']['percent'], 50)
        self.assertEqual(cells['Comcast', '2021-2022', 'corporate']['percent'], 31)
        self.assertEqual(cells['AT&T', '2021-2022', 'corporate']['percent'], 84)
        self.assertEqual(cells['Altice USA', '2023-2024', 'contributions']['percent'], 50)

    def test_the_served_builds_remainder_of_twenty_one_reads_like_the_audited_nineteen(self):
        served = collection_gaps.parse_detail(REAL_DETAIL.replace('and 19 more', 'and 21 more'))
        self.assertEqual(served['unitemized'], 21)
        self.assertEqual(served['cells'], collection_gaps.parse_detail(REAL_DETAIL)['cells'])

    def test_every_cell_keeps_the_exact_words_of_its_destination_and_counts(self):
        for cell in collection_gaps.parse_detail(REAL_DETAIL)['cells']:
            self.assertEqual(set(cell), {'company', 'cycle', 'destination', 'held', 'reported', 'percent'})
            self.assertIn(cell['destination'], ('contributions', 'corporate'))
            self.assertLess(cell['held'], cell['reported'])

    def test_a_detail_that_itemizes_everything_has_no_unitemized_remainder(self):
        parsed = collection_gaps.parse_detail('x: Comcast/2026 contributions 1 of 3; Verizon/2024 corporate 2 of 4')
        self.assertEqual(parsed['unitemized'], 0)
        self.assertEqual([(c['company'], c['cycle'], c['percent']) for c in parsed['cells']],
                         [('Comcast', '2025-2026', 33), ('Verizon', '2023-2024', 50)])

    def test_a_detail_that_lists_every_short_cell_reads_cell_for_cell_like_the_truncated_one(self):
        # export_guards.R lists every short cell now; a build made before that prints the first twelve and
        # "... and N more". Both forms keep the Detail's format and order, so the full one extends the short one.
        head = REAL_DETAIL.rpartition(' ... and 19 more')[0]
        extra = [f'Extra Co {n}/2024 contributions {n * 100} of {n * 200}' for n in range(1, 20)]
        full = collection_gaps.parse_detail(head + '; ' + '; '.join(extra))
        truncated = collection_gaps.parse_detail(REAL_DETAIL)
        self.assertEqual((len(full['cells']), full['unitemized']), (31, 0))
        self.assertEqual(full['cells'][:12], truncated['cells'])
        self.assertEqual([c['company'] for c in full['cells'][12:]], [f'Extra Co {n}' for n in range(1, 20)])
        self.assertEqual({c['percent'] for c in full['cells'][12:]}, {50})
        self.assertEqual((len(truncated['cells']), truncated['unitemized']), (12, 19))

    def test_nothing_to_itemize_is_empty_not_an_error(self):
        for text in (None, '', 'every company x cycle cell walks stored what FEC reported'):
            self.assertEqual(collection_gaps.parse_detail(text), {'cells': [], 'unitemized': 0})

    def test_a_cell_FEC_reports_nothing_for_cannot_be_a_shortfall(self):
        self.assertEqual(collection_gaps.parse_detail('x: Comcast/2026 contributions 0 of 0')['cells'], [])

    def test_percent_rounds_half_up_like_math_round(self):
        for held, reported, expected in ((1, 8, 13), (1, 3, 33), (2, 3, 67), (14133, 28266, 50), (99, 100, 99)):
            cell = collection_gaps.parse_detail(f'x: A/2026 contributions {held} of {reported}')['cells'][0]
            self.assertEqual(cell['percent'], expected, (held, reported))

    @unittest.skipUnless(REFERENCE.with_suffix('.audit').joinpath('quality_checks.csv').is_file(),
                         'the audited Telecom reference is not on this machine')
    def test_the_audited_reference_csv_reads_as_the_pasted_detail_does(self):
        with REFERENCE.with_suffix('.audit').joinpath('quality_checks.csv').open(encoding='utf-8-sig', newline='') as stream:
            detail = next(r['detail'] for r in csv.DictReader(stream) if r['id'] == 'collection_completeness')
        self.assertEqual(collection_gaps.parse_detail(detail), collection_gaps.parse_detail(REAL_DETAIL))


class WorkbookGapTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.path = self.root / 'Telecom Industry Spending Overview 2025-2026.xlsx'

    def test_the_workbooks_own_data_quality_sheet_is_read(self):
        quality_book(self.path)
        gaps = collection_gaps.workbook_gaps(self.path)
        self.assertTrue(gaps['failed'])
        self.assertEqual(gaps['source'], 'Data Quality sheet')
        self.assertEqual(gaps['unitemized'], 19)
        self.assertEqual(by_cell(gaps)['Comcast', '2025-2026', 'contributions']['percent'], 77)

    def test_a_passing_check_names_no_cells(self):
        quality_book(self.path, detail='60 company x cycle walks stored what FEC reported', passed=True)
        gaps = collection_gaps.workbook_gaps(self.path)
        self.assertEqual((gaps['failed'], gaps['cells'], gaps['unitemized']), (False, [], 0))

    def test_without_the_sheet_the_sibling_audit_csv_answers(self):
        book = openpyxl.Workbook()
        book.save(self.path)
        book.close()
        quality_csv(self.path)
        gaps = collection_gaps.workbook_gaps(self.path)
        self.assertTrue(gaps['failed'])
        self.assertEqual(gaps['source'], 'quality_checks.csv')
        self.assertEqual((len(gaps['cells']), gaps['unitemized']), (12, 19))

    def test_a_workbook_that_cannot_be_opened_still_falls_back_to_the_csv(self):
        self.path.write_bytes(b'not a workbook')
        quality_csv(self.path)
        self.assertEqual(collection_gaps.workbook_gaps(self.path)['source'], 'quality_checks.csv')

    def test_no_sheet_and_no_csv_says_nothing_is_known_not_that_nothing_is_short(self):
        book = openpyxl.Workbook()
        book.save(self.path)
        book.close()
        gaps = collection_gaps.workbook_gaps(self.path)
        self.assertEqual((gaps['failed'], gaps['cells'], gaps['unitemized'], gaps['source']), (False, [], 0, None))
        self.assertEqual(collection_gaps.workbook_gaps(self.root / 'absent.xlsx')['source'], None)

    def test_with_neither_sheet_nor_csv_the_pinned_payloads_export_guard_answers(self):
        book = openpyxl.Workbook()
        book.save(self.path)
        book.close()
        payload = payload_file(self.root / 'industry_data-34ef41793b3ab4c9.json', REAL_DETAIL.replace('and 19 more', 'and 21 more'))
        with patch.object(collection_gaps, '_payload_path', return_value=payload):
            gaps = collection_gaps.workbook_gaps(self.path)
        self.assertEqual((gaps['failed'], gaps['source'], gaps['unitemized']), (True, 'pinned payload export_guards', 21))
        self.assertEqual(by_cell(gaps)['Comcast', '2025-2026', 'contributions']['percent'], 77)

    def test_the_sheet_outranks_the_csv_and_the_csv_outranks_the_payload(self):
        payload = payload_file(self.root / 'industry_data-1.json', REAL_DETAIL.replace('and 19 more', 'and 21 more'))
        quality_book(self.path)
        quality_csv(self.path, REAL_DETAIL.replace('and 19 more', 'and 20 more'))
        with patch.object(collection_gaps, '_payload_path', return_value=payload):
            self.assertEqual((collection_gaps.workbook_gaps(self.path)['source'], collection_gaps.workbook_gaps(self.path)['unitemized']),
                             ('Data Quality sheet', 19))
            book = openpyxl.Workbook()
            book.save(self.path)
            book.close()
            self.assertEqual((collection_gaps.workbook_gaps(self.path)['source'], collection_gaps.workbook_gaps(self.path)['unitemized']),
                             ('quality_checks.csv', 20))
            (self.path.with_suffix('.audit') / 'quality_checks.csv').unlink()
            self.assertEqual((collection_gaps.workbook_gaps(self.path)['source'], collection_gaps.workbook_gaps(self.path)['unitemized']),
                             ('pinned payload export_guards', 21))

    def test_a_payload_without_the_check_or_unreadable_says_nothing_is_known(self):
        book = openpyxl.Workbook()
        book.save(self.path)
        book.close()
        bare = self.root / 'bare.json'
        bare.write_text(json.dumps({'metadata': {'export_guards': []}}), encoding='utf-8')
        for payload in (bare, self.root / 'absent.json'):
            with patch.object(collection_gaps, '_payload_path', return_value=payload):
                self.assertEqual(collection_gaps.workbook_gaps(self.path)['source'], None)
        broken = self.root / 'broken.json'
        broken.write_text('{"raw_rows": []}', encoding='utf-8')
        with patch.object(collection_gaps, '_payload_path', return_value=broken):
            self.assertEqual(collection_gaps.workbook_gaps(self.path)['source'], None)

    def test_a_regenerated_payload_is_read_again(self):
        book = openpyxl.Workbook()
        book.save(self.path)
        book.close()
        payload = payload_file(self.root / 'industry_data-2.json')
        with patch.object(collection_gaps, '_payload_path', return_value=payload):
            self.assertTrue(collection_gaps.workbook_gaps(self.path)['failed'])
            payload_file(payload, 'every cell walks stored what FEC reported', passed=True)
            self.assertFalse(collection_gaps.workbook_gaps(self.path)['failed'])

    def test_a_workbook_that_is_not_a_published_report_has_no_pinned_payload(self):
        self.assertIsNone(collection_gaps._pinned(self.root / 'Notes.xlsx'))
        self.assertIsNone(collection_gaps._pinned(self.path))

    @unittest.skipUnless(SERVED.is_file(), 'the served Telecom workbook is not on this machine')
    def test_the_served_telecom_build_answers_from_its_own_pinned_guard(self):
        gaps = collection_gaps.workbook_gaps(SERVED)
        if gaps['source'] is None:
            self.skipTest('the served Telecom workbook has no pinned payload guard here')
        self.assertEqual(gaps['source'], 'pinned payload export_guards')
        self.assertEqual((gaps['failed'], len(gaps['cells']), gaps['unitemized']), (True, 12, 21))
        cells = by_cell(gaps)
        self.assertEqual(cells['Comcast', '2025-2026', 'contributions']['percent'], 77)
        self.assertEqual(cells['Verizon', '2025-2026', 'contributions']['percent'], 93)

    def test_a_rebuilt_workbook_is_read_again(self):
        quality_book(self.path)
        self.assertTrue(collection_gaps.workbook_gaps(self.path)['failed'])
        quality_book(self.path, detail='60 company x cycle walks stored what FEC reported', passed=True)
        self.assertFalse(collection_gaps.workbook_gaps(self.path)['failed'])

    def test_callers_get_their_own_copy(self):
        quality_book(self.path)
        collection_gaps.workbook_gaps(self.path)['cells'].clear()
        self.assertEqual(len(collection_gaps.workbook_gaps(self.path)['cells']), 12)


class ScopeAndMergeTests(unittest.TestCase):
    def test_scoped_keeps_only_the_cycles_a_view_can_show(self):
        gaps = {'failed': True, 'cells': collection_gaps.parse_detail(REAL_DETAIL)['cells'], 'unitemized': 19, 'source': 'x'}
        only = collection_gaps.scoped(gaps, [2026])
        self.assertEqual([(c['company'], c['cycle']) for c in only['cells']],
                         [('Comcast', '2025-2026'), ('Verizon', '2025-2026')])
        self.assertEqual((only['failed'], only['unitemized']), (True, 19))
        self.assertEqual(len(collection_gaps.scoped(gaps, None)['cells']), 12)

    def test_a_check_that_names_only_other_cycles_is_not_a_failure_for_this_one(self):
        # P2-2: the cells the scope removes were the only ones, and none were left unitemized.
        cells = [c for c in collection_gaps.parse_detail(REAL_DETAIL)['cells'] if c['cycle'] == '2023-2024']
        gaps = {'failed': True, 'cells': cells, 'unitemized': 0, 'source': 'x'}
        away = collection_gaps.scoped(gaps, [2026])
        self.assertEqual((away['failed'], away['cells'], away['unitemized'], away['source']), (False, [], 0, 'x'))
        # Still a failure when a cell is in scope, when the check left some unitemized, or when it itemized nothing.
        self.assertEqual(collection_gaps.scoped(gaps, [2024]), gaps)
        self.assertTrue(collection_gaps.scoped({**gaps, 'unitemized': 3}, [2026])['failed'])
        self.assertTrue(collection_gaps.scoped({**gaps, 'cells': []}, [2026])['failed'])

    def test_merge_lists_each_cell_once_and_never_adds_the_remainders(self):
        cells = collection_gaps.parse_detail(REAL_DETAIL)['cells']
        one = {'failed': True, 'cells': cells[:7], 'unitemized': 19, 'source': 'a'}
        two = {'failed': True, 'cells': cells[5:], 'unitemized': 19, 'source': 'b'}
        clean = {'failed': False, 'cells': [], 'unitemized': 0, 'source': None}
        merged = collection_gaps.merge([one, two, clean])
        self.assertEqual(len(merged['cells']), 12)
        self.assertEqual((merged['failed'], merged['unitemized']), (True, 19))
        self.assertEqual(collection_gaps.merge([clean]), clean)
        self.assertEqual(collection_gaps.merge([])['failed'], False)


class PayloadTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_the_graph_view_carries_the_gaps_for_the_cycles_it_shows(self):
        path = self.root / 'Giving.xlsx'
        quality_book(path, graph=True)
        value = graph_view([path], {})
        gaps = value['collection_gaps']
        self.assertTrue(gaps['failed'])
        self.assertEqual(value['options']['cycles'], ['2026'])
        self.assertEqual({(c['company'], c['cycle']) for c in gaps['cells']},
                         {('Comcast', '2025-2026'), ('Verizon', '2025-2026')})
        self.assertEqual(gaps['unitemized'], 19)

    def test_a_graph_view_over_a_clean_report_says_so(self):
        path = self.root / 'Giving.xlsx'
        quality_book(path, detail='60 company x cycle walks stored what FEC reported', passed=True, graph=True)
        self.assertEqual(graph_view([path], {})['collection_gaps']['cells'], [])
        self.assertFalse(graph_view([path], {})['collection_gaps']['failed'])

    def catalog(self, name):
        path = self.root / name
        quality_book(path)
        return {'datasets': [{'id': name[:-5], 'available': True, 'provenance': {'workbook': name}},
                             {'id': 'No provenance', 'available': False}]}

    def test_a_cycle_report_is_catalogued_with_its_own_cycle_only(self):
        catalog = collection_gaps.attach(self.root, self.catalog('Telecom Industry Spending Overview 2025-2026.xlsx'))
        gaps = catalog['datasets'][0]['collection_gaps']
        self.assertEqual({(c['company'], c['cycle']) for c in gaps['cells']},
                         {('Comcast', '2025-2026'), ('Verizon', '2025-2026')})
        self.assertEqual(gaps['unitemized'], 19)
        self.assertNotIn('collection_gaps', catalog['datasets'][1])

    OTHER_CYCLES_ONLY = ('walks stored materially fewer rows than FEC reported - recollect before trusting it: '
                         'Comcast/2024 contributions 17869 of 29991; Verizon/2024 contributions 14783 of 19170')

    def test_a_cycle_report_whose_check_names_only_other_cycles_is_catalogued_as_not_failing(self):
        path = self.root / 'Telecom Industry Spending Overview 2025-2026.xlsx'
        quality_book(path, detail=self.OTHER_CYCLES_ONLY)
        catalog = {'datasets': [{'id': 'a', 'available': True, 'provenance': {'workbook': path.name}}]}
        gaps = collection_gaps.attach(self.root, catalog)['datasets'][0]['collection_gaps']
        self.assertEqual((gaps['failed'], gaps['cells'], gaps['unitemized']), (False, [], 0))

    def test_a_graph_view_whose_check_names_only_other_cycles_carries_no_failure(self):
        path = self.root / 'Giving.xlsx'
        quality_book(path, detail=self.OTHER_CYCLES_ONLY, graph=True)
        gaps = graph_view([path], {})['collection_gaps']
        self.assertEqual((gaps['failed'], gaps['cells'], gaps['unitemized']), (False, [], 0))

    def test_an_industry_report_is_catalogued_with_every_cycle_it_names(self):
        catalog = collection_gaps.attach(self.root, self.catalog('Telecom Industry Spending Overview.xlsx'))
        self.assertEqual(len(catalog['datasets'][0]['collection_gaps']['cells']), 12)

    def test_an_unexpected_error_leaves_the_catalogue_whole(self):
        catalog = {'datasets': [{'id': 'a', 'provenance': {'workbook': 'Telecom Industry Spending Overview.xlsx'}}]}
        with patch.object(collection_gaps, 'workbook_gaps', side_effect=RuntimeError('boom')):
            self.assertIs(collection_gaps.attach(self.root, catalog), catalog)
        self.assertNotIn('collection_gaps', catalog['datasets'][0])

    def test_an_unreadable_or_foreign_workbook_name_leaves_the_dataset_alone(self):
        catalog = {'datasets': [{'id': 'a', 'provenance': {'workbook': 'Notes.xlsx'}},
                                {'id': 'b', 'provenance': {'workbook': 'Telecom Industry Spending Overview.xlsx'}}]}
        collection_gaps.attach(self.root, catalog)
        self.assertNotIn('collection_gaps', catalog['datasets'][0])
        self.assertEqual(catalog['datasets'][1]['collection_gaps'], collection_gaps.EMPTY)


class CatalogRouteTests(unittest.TestCase):
    """Both catalogue routes (Chart Studio's and Insights') add the gaps; no new round trip."""

    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name)
        quality_book(cls.root / 'Telecom Industry Spending Overview 2023-2024.xlsx')
        spec = importlib.util.spec_from_file_location('collection_gaps_server', Path(collection_gaps.__file__).with_name('server.py'))
        cls.app = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.app)
        cls.console = SimpleNamespace(token='t', read_only=True, report_directory=cls.root,
                                      configured_root=str(cls.root), root=cls.root)
        cls.server = cls.app.ThreadingHTTPServer(('127.0.0.1', 0), cls.app.Handler)
        cls.server.app = cls.console
        cls.worker = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.worker.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.worker.join(timeout=5)
        cls.temp.cleanup()

    def get(self, path):
        connection = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=8)
        try:
            connection.request('GET', path, headers={'Host': f'127.0.0.1:{self.server.server_port}'})
            response = connection.getresponse()
            return response.status, json.loads(response.read())
        finally:
            connection.close()

    def dataset(self):
        return {'id': 'Telecom 2023-2024', 'available': True,
                'provenance': {'workbook': 'Telecom Industry Spending Overview 2023-2024.xlsx'}}

    def test_chart_studio_catalogue_route_carries_the_cycle_gaps(self):
        with patch.object(self.app.contribution_data, 'catalog', return_value={'datasets': [self.dataset()]}), \
                patch.object(self.app.contribution_workbook_guard, 'verify_sources', return_value=[]):
            status, body = self.get('/api/contributions/catalog')
        self.assertEqual(status, 200)
        cells = body['datasets'][0]['collection_gaps']['cells']
        self.assertEqual(sorted((c['company'], c['destination'], c['percent']) for c in cells),
                         [('Altice USA', 'contributions', 50), ('Comcast', 'contributions', 60),
                          ('Comcast', 'corporate', 31), ('Lumen Technologies', 'contributions', 50),
                          ('Verizon', 'contributions', 77)])
        self.assertEqual({c['cycle'] for c in cells}, {'2023-2024'})
        self.assertEqual(body['datasets'][0]['collection_gaps']['unitemized'], 19)

    def test_insights_catalogue_route_carries_the_same_gaps(self):
        import analysis_data
        with patch.object(analysis_data, 'catalog', return_value={'datasets': [self.dataset()], 'panels': []}):
            status, body = self.get('/api/insights/catalog')
        self.assertEqual(status, 200)
        self.assertEqual({c['cycle'] for c in body['datasets'][0]['collection_gaps']['cells']}, {'2023-2024'})


if __name__ == '__main__':
    unittest.main()
