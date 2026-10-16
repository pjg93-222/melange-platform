"""The actual Data writer carries truthful receipt-basis metadata without changing cells."""
from contextlib import closing
import importlib.util,json,os,re,sqlite3,sys,tempfile,unittest
from pathlib import Path
import openpyxl
PRODUCT=Path(os.environ.get('MENTAT_ACCOUNTING_COPY_PRODUCT',Path(__file__).resolve().parents[1]))
EXPORTER=Path(os.environ.get('MENTAT_ACCOUNTING_COPY_EXPORTER',PRODUCT/'tools/export_crypto_spending_data.R'))
sys.path.insert(0,str(PRODUCT/'tests'))
import test_tidy_workbook as fixture

def accounting_note(path):
    source=path.read_text(encoding='utf-8-sig')
    block=re.search(r'accounting_note = paste\(\s*(.*?)\s*\),\s*blank_columns',source,re.S)
    if block is None:raise ValueError('Exact exporter accounting-note block missing')
    parts=re.findall(r'"(?:[^"\\]|\\.)*"',block.group(1))
    return ' '.join(json.loads(part) for part in parts)

class ExportAccountingCopyTests(unittest.TestCase):
    def test_actual_writer_changes_only_note_and_keeps_included_memo_conduit_money(self):
        note=accounting_note(EXPORTER)
        self.assertIn('at the named recipient; excluded rows are 0.',note)
        self.assertIn('Valid conduit and nonduplicated memo receipts can be included.',note)
        self.assertTrue(note.endswith('Neither column certifies unique giving.'))
        self.assertNotIn('0 for memo, conduit',note)
        with tempfile.TemporaryDirectory(prefix='mentat-accounting-copy-') as name:
            folder=Path(name);store=folder/'fixture.sqlite';fixture.build_store(store)
            old='Analysis Amount counts included itemized receipts at the named recipient and is 0 for memo, conduit, superseded, duplicate and out-of-cycle legs; Amount is the figure as filed. Neither is verified unique giving.'
            with closing(sqlite3.connect(store)) as db,db:
                db.execute("UPDATE metadata SET value=? WHERE key='accounting_note'",(old,))
                db.execute("UPDATE contributions SET is_conduit_leg=1 WHERE cycle=2022")
                db.execute("UPDATE contributions SET is_memo_leg=1,duplicate_memo=0 WHERE cycle=2024")
            first=folder/'before.xlsx';second=folder/'after.xlsx'
            self.assertEqual(fixture.writer.main(['--staging',str(store),'--out',str(first)]),0)
            with closing(sqlite3.connect(store)) as db,db:db.execute("UPDATE metadata SET value=? WHERE key='accounting_note'",(note,))
            self.assertEqual(fixture.writer.main(['--staging',str(store),'--out',str(second)]),0)
            books=[openpyxl.load_workbook(p,read_only=True,data_only=True) for p in (first,second)]
            try:
                self.assertEqual(books[0].sheetnames,books[1].sheetnames)
                for name in books[0].sheetnames:
                    before=list(books[0][name].values);after=list(books[1][name].values)
                    if name=='About':
                        a=dict(before[1:]);b=dict(after[1:]);self.assertEqual(a.pop('Accounting Note'),old);self.assertEqual(b.pop('Accounting Note'),note);self.assertEqual(a,b)
                    else:self.assertEqual(before,after,name)
                rows=list(books[1]['Contributions'].values);records=[dict(zip(rows[0],r)) for r in rows[1:]]
                self.assertEqual(len(records),2)
                self.assertEqual([(r['Amount'],r['Analysis Amount'],r['Include in Analysis']) for r in records],[(25,25,1),(25,25,1)])
                self.assertEqual(records[0]['Conduit Leg'],1);self.assertEqual(records[1]['Memo Leg'],1);self.assertEqual(records[1]['Duplicate Memo'],0)
            finally:
                for book in books:book.close()

if __name__=='__main__':unittest.main()
