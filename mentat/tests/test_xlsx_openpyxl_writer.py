"""MENTAT_XLSX_WRITER=openpyxl (D-178): the artifact-tool stand-in for the Overview workbook builder.

xlsx_spec_writer.py is tested on hand-written operation lists; openpyxl_workbook_shim.mjs is tested through Node (skipped
when there is none). The artifact-tool side-by-side proof lives in the parity runs, because it needs the proprietary
runtime; these tests need neither it nor any payload.

Run from inside mentat/tests: python -m unittest test_xlsx_openpyxl_writer
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

import openpyxl

HERE = Path(__file__).resolve().parent
TOOLS = HERE.parent / 'tools'
sys.path.insert(0, str(TOOLS))
import xlsx_spec_writer as W  # noqa: E402

NODE = os.environ.get('MENTAT_NODE_PATH') or shutil.which('node')


def values(s, r, c, rows):
    return {'s': s, 'op': 'values', 'r': r, 'c': c, 'rows': rows}


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)

    def write(self, sheets, ops):
        spec, out = self.tmp / 'spec.ndjson', self.tmp / 'out.xlsx'
        lines = [{'op': 'book', 'sheets': sheets}] + ops
        spec.write_text('\n'.join(json.dumps(line) for line in lines) + '\n', encoding='utf-8')
        W.replay(str(spec), str(out))
        return out

    def raw(self, out, part='xl/worksheets/sheet1.xml'):
        with ZipFile(out) as z:
            return z.read(part).decode('utf-8')


class CellTests(Base):
    def test_values_formulas_dates_and_text(self):
        out = self.write(['S'], [
            values(0, 1, 1, [['Name', 'Amount', 'Check'], ['Acme', 10.5, '=B2*2'], [None, '', 'x']]),
            {'s': 0, 'op': 'formulas', 'r': 4, 'c': 1, 'rows': [['=SUM(B2:B3)', '', 7]]},
            values(0, 6, 1, [['2024-03-01T12:30:00Z', '2024-03-01T12:30:00', '2024-03-01T12:30:00-05:00', '=A1']]),
        ])
        ws = openpyxl.load_workbook(out)['S']
        self.assertEqual([c.value for c in ws[1]][:3], ['Name', 'Amount', 'Check'])
        self.assertEqual(ws['C2'].value, '=B2*2')
        self.assertIsNone(ws['A3'].value)
        self.assertEqual((ws['A4'].value, ws['C4'].value), ('=SUM(B2:B3)', 7))
        self.assertEqual(ws['A6'].value, datetime(2024, 3, 1, 12, 30))           # zoned: a UTC date cell
        self.assertEqual(ws['C6'].value, datetime(2024, 3, 1, 17, 30))
        self.assertEqual(ws['B6'].value, '2024-03-01T12:30:00')                  # no zone: text
        self.assertEqual(ws['D6'].value, '=A1')
        self.assertEqual(ws['A6'].number_format, 'General')

    def test_floats_are_the_same_double_and_long_text_is_kept(self):
        long = 'x' * 40000
        out = self.write(['S'], [values(0, 1, 1, [[0.25555555555555554, 3.0, -0.0, long]])])
        xml = self.raw(out)
        self.assertIn('<v>0.25555555555555554</v>', xml)
        self.assertIn('<v>3</v>', xml)
        self.assertEqual(openpyxl.load_workbook(out)['S']['D1'].value, long)

    def test_fill_down_moves_relative_references(self):
        out = self.write(['S'], [
            values(0, 5, 1, [['=IF(R5="Yes",H5,0)'], [None], [None]]),
            {'s': 0, 'op': 'fillDown', 'ref': 'A5:A7'},
        ])
        ws = openpyxl.load_workbook(out)['S']
        self.assertEqual([ws['A%d' % r].value for r in (5, 6, 7)], ['=IF(R5="Yes",H5,0)', '=IF(R6="Yes",H6,0)', '=IF(R7="Yes",H7,0)'])


class FormatTests(Base):
    def test_formats_merge_by_property_and_keep_earlier_ones(self):
        out = self.write(['S'], [
            values(0, 1, 1, [['Title', 'x'], [1, 2]]),
            {'s': 0, 'op': 'fmt', 'ref': 'A1:B1', 'p': {'fill': '#17365D'}},
            {'s': 0, 'op': 'fmt', 'ref': 'A1:B1', 'p': {'font': {'bold': True, 'color': '#FFFFFF', 'size': 18}}},
            {'s': 0, 'op': 'fmt', 'ref': 'A1:B1', 'p': {'v': 'center', 'h': 'left', 'wrap': True}},
            {'s': 0, 'op': 'fmt', 'ref': 'A2:B2', 'p': {'numberFormat': '$#,##0.00;[Red]($#,##0.00);-'}},
            {'s': 0, 'op': 'rowHeight', 'r1': 1, 'r2': 1, 'h': 30},
            {'s': 0, 'op': 'colWidth', 'c1': 1, 'c2': 2, 'w': 21},
        ])
        ws = openpyxl.load_workbook(out)['S']
        top = ws['B1']
        self.assertEqual(top.fill.fgColor.rgb, 'FF17365D')
        self.assertEqual((top.font.b, top.font.sz, top.font.color.rgb), (True, 18, 'FFFFFFFF'))
        self.assertEqual((top.alignment.vertical, top.alignment.horizontal, top.alignment.wrap_text), ('center', 'left', True))
        self.assertEqual(ws['A2'].number_format, '$#,##0.00;[Red]($#,##0.00);-')
        self.assertEqual((ws.row_dimensions[1].height, ws.column_dimensions['A'].width, ws.column_dimensions['B'].width), (30, 21, 21))

    def test_an_unknown_font_property_or_operation_is_an_error_not_a_dropped_format(self):
        with self.assertRaisesRegex(ValueError, 'unsupported font property: underline'):
            self.write(['S'], [values(0, 1, 1, [['x']]), {'s': 0, 'op': 'fmt', 'ref': 'A1', 'p': {'font': {'underline': True}}}])
        with self.assertRaisesRegex(ValueError, 'unknown spec operation'):
            self.write(['S'], [{'s': 0, 'op': 'sparkline'}])


class StructureTests(Base):
    def test_table_merge_freeze_grid_and_conditional_formats(self):
        out = self.write(['S'], [
            values(0, 1, 1, [['Title'], ['sub'], ['note'], ['Company', 'Total', 'Check'], ['A', 1, 'OK'], ['B', 2, 'CHECK']]),
            {'s': 0, 'op': 'merge', 'ref': 'A1:C1'},
            {'s': 0, 'op': 'cf', 'ref': 'C5:C6', 'type': 'expression', 'formula': '=$C5<>"OK"',
             'fill': '#FFF2CC', 'font': {'bold': True, 'color': '#C00000'}},
            {'s': 0, 'op': 'cf', 'ref': 'A5:A6', 'type': 'beginsWith', 'text': 'Conservative split',
             'fill': '#FFF2CC', 'font': {'color': '#7F6000'}},
            {'s': 0, 'op': 'table', 'ref': 'A4:C6', 'name': 'T_One', 'style': 'TableStyleMedium2', 'headers': True, 'filter': True},
            {'s': 0, 'op': 'freeze', 'rows': 4, 'cols': 1},
            {'s': 0, 'op': 'grid', 'show': False},
        ])
        ws = openpyxl.load_workbook(out)['S']
        self.assertEqual([str(r) for r in ws.merged_cells.ranges], ['A1:C1'])
        self.assertEqual((ws.freeze_panes, ws.sheet_view.showGridLines), ('B5', False))
        table = ws.tables['T_One']
        self.assertEqual((table.ref, table.tableStyleInfo.name, table.autoFilter.ref), ('A4:C6', 'TableStyleMedium2', 'A4:C6'))
        self.assertEqual([c.name for c in table.tableColumns], ['Company', 'Total', 'Check'])
        xml = self.raw(out)
        self.assertIn('<formula>$C5&lt;&gt;"OK"</formula>', xml)
        begins = xml[xml.index('type="beginsWith"'):].split('>')[0]
        self.assertIn('text="Conservative split"', begins)
        self.assertNotIn('<formula>', xml.split('type="beginsWith"')[1].split('</conditionalFormatting>')[0])

    def test_sheet_order_and_a_single_sheet_without_freeze(self):
        out = self.write(['B', 'A'], [values(1, 1, 1, [['x']]), {'s': 0, 'op': 'freeze', 'rows': 0, 'cols': 0}])
        wb = openpyxl.load_workbook(out)
        self.assertEqual(wb.sheetnames, ['B', 'A'])
        self.assertIsNone(wb['B'].freeze_panes)


SHIM_PROBE = r"""
import { Workbook, SpreadsheetFile } from %(shim)s;
const wb = Workbook.create();
const s = wb.worksheets.add("Data");
s.showGridLines = false;
s.getRange("A1:C1").merge();
s.getRange("A1").values = [["Title"]];
s.getRange("A1:C1").format = { fill: "#17365D", font: { bold: true, color: "#FFFFFF", size: 18 }, verticalAlignment: "center" };
s.getRange("A1:C1").format.rowHeight = 30;
s.getRange("A3:C5").values = [["Company", "Date", "Amount"], ["Acme", new Date("2024-12-31T00:00:00"), 10.5], ["Beta", new Date("2023-01-17T20:02:36Z"), 3]];
s.getRange("D3").values = [["Next"]];
s.getRange("D4").formulas = [["=B4+1"]];
s.getRange("D4:D5").fillDown();
s.getRange("C4:C5").format.numberFormat = "#,##0.00";
s.getRange("B4:B5").format.numberFormat = "yyyy-mm-dd";
s.getRange("A3:D5").conditionalFormats.addCustom('=$A3<>"OK"', { fill: "#FFF2CC", font: { bold: true, color: "#C00000" } });
const t = s.tables.add("A3:D5", true, "DataTable");
t.style = "TableStyleMedium2"; t.showHeaders = true; t.showFilterButton = true;
s.freezePanes.freezeRows(3);
s.freezePanes.freezeColumns(1);
for (const attempt of [() => s.getRange("A1").format = { borders: {} }, () => s.getRange("A1").conditionalFormats.add("top10", {}),
                       () => s.getRange("A1").frobnicate(), () => wb.worksheets.add("Data"), () => wb.worksheets.getItem("Nope")]) {
  try { attempt(); console.log("NO ERROR"); } catch (error) { console.log("error:" + error.constructor.name); }
}
wb.inspect = async () => ({ ndjson: "" });
await (await SpreadsheetFile.exportXlsx(wb)).save(process.argv[2]);
"""


@unittest.skipUnless(NODE, 'node is not available')
class ShimTests(Base):
    def test_the_builders_calls_become_a_workbook(self):
        script, out = self.tmp / 'probe.mjs', self.tmp / 'probe.xlsx'
        script.write_text(SHIM_PROBE % {'shim': json.dumps((TOOLS / 'openpyxl_workbook_shim.mjs').as_uri())}, encoding='utf-8')
        env = {**os.environ, 'MENTAT_PYTHON': sys.executable}
        run = subprocess.run([NODE, str(script), str(out)], capture_output=True, text=True, encoding='utf-8', env=env)
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertEqual([line for line in run.stdout.split() if line.startswith('error:')],
                         ['error:Error', 'error:Error', 'error:TypeError', 'error:Error', 'error:Error'])
        self.assertNotIn('NO ERROR', run.stdout)
        self.assertFalse((self.tmp / 'probe.xlsx.spec.ndjson').exists())
        ws = openpyxl.load_workbook(out)['Data']
        self.assertEqual([str(r) for r in ws.merged_cells.ranges], ['A1:C1'])
        self.assertEqual((ws.freeze_panes, ws.sheet_view.showGridLines, ws.row_dimensions[1].height), ('B4', False, 30))
        self.assertEqual((ws['A4'].value, ws['C4'].value, ws['C4'].number_format), ('Acme', 10.5, '#,##0.00'))
        self.assertEqual([ws['D4'].value, ws['D5'].value], ['=B4+1', '=B5+1'])
        self.assertEqual(ws['B4'].value, datetime(2024, 12, 31))         # a Date at local midnight is that calendar day
        self.assertEqual(ws['B5'].value, datetime(2023, 1, 17, 20, 2, 36))
        self.assertEqual(ws.tables['DataTable'].ref, 'A3:D5')
        self.assertEqual([c.name for c in ws.tables['DataTable'].tableColumns], ['Company', 'Date', 'Amount', 'Next'])

    def test_a_failing_writer_fails_the_save_and_keeps_the_spec(self):
        script, out = self.tmp / 'probe.mjs', self.tmp / 'probe.xlsx'
        script.write_text(SHIM_PROBE % {'shim': json.dumps((TOOLS / 'openpyxl_workbook_shim.mjs').as_uri())}, encoding='utf-8')
        env = {**os.environ, 'MENTAT_PYTHON': str(self.tmp / 'no-such-python')}
        run = subprocess.run([NODE, str(script), str(out)], capture_output=True, text=True, encoding='utf-8', env=env)
        self.assertNotEqual(run.returncode, 0)
        self.assertIn('openpyxl workbook writer failed', run.stderr)
        self.assertTrue((self.tmp / 'probe.xlsx.spec.ndjson').exists())


class BuilderSwitchTests(unittest.TestCase):
    """The builder is a top-level script, so its switch is pinned as text (the other builder tests do the same)."""
    SOURCE = (TOOLS / 'crypto_workbook_builder.mjs').read_text(encoding='utf-8')

    def test_the_default_stays_artifact_tool_and_an_unknown_writer_is_refused(self):
        self.assertIn('process.env.MENTAT_XLSX_WRITER ?? ""', self.SOURCE)
        self.assertIn('|| "artifact"', self.SOURCE)
        self.assertIn('MENTAT_XLSX_WRITER must be artifact or openpyxl', self.SOURCE)
        self.assertIn('await import("./openpyxl_workbook_shim.mjs")', self.SOURCE)
        self.assertIn('await import(pathToFileURL(resolveArtifactToolEntry()).href)', self.SOURCE)

    def test_previews_are_skipped_not_faked_under_openpyxl(self):
        self.assertIn('MENTAT_RENDER_PREVIEWS === "1" && xlsxWriter !== "artifact"', self.SOURCE)


if __name__ == '__main__':
    unittest.main()
