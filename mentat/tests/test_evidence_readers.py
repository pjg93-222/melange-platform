"""Any-file-type / any-path evidence import (S3b): xlsx, delimited text (csv/tsv/txt
with sniffed delimiter and encoding fallback), JSON, and FEC bulk (zip or bare
pipe-delimited txt), plus Store.import_path(path=...) and the browse() file-picker
API. Written against the S3b CONTRACT in this lane's HANDOFF.md before Architect A
(server) and executor B (client) land their implementation; run first against a
pristine worktree of HEAD to record the RED baseline (see receipts/S3b_before_python.txt).

This file extends, but does not modify, test_evidence_csv_stream.py's S3 suite: it
imports evidence_csv/server exactly the same way, but keeps its own local copies of
the message constants and the schema-by-example helper (contracts/README.md does not
mandate sharing assertion code between same-language suites -- the R/cjs suites for
other contracts in this packet each keep their own local assertions too), so a
producer that silently changed a message or a shape shows up here as a real
assertion failure, not as an accidental pass through borrowed code.
"""
from __future__ import annotations

import datetime
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import unittest
import urllib.parse
import zipfile
from contextlib import contextmanager
from datetime import date as date_cls
from pathlib import Path
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'app'))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
import evidence_csv as csv_data  # noqa: E402
import server as app  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parents[2]
CONTRACTS = REPO_ROOT / 'audits' / 'mentat_failings_20260915' / 'contracts'
MENTAT_INPUTS = Path(__file__).resolve().parents[1] / 'inputs'


def load_contract(name):
    return json.loads((CONTRACTS / f'{name}.json').read_text(encoding='utf-8'))


# ---------------------------------------------------------------------------
# Exact messages, copied verbatim from the S3b CONTRACT (not read off
# evidence_csv.py -- see the module docstring for why).
# ---------------------------------------------------------------------------
M_BINARY = 'This file contains binary data and could not be read as text.'
M_JSON_SIZE = 'Choose a JSON file no larger than 25 MB, or export it as CSV.'
M_JSON_SHAPE = 'The JSON must be an openFEC export ({"results": [...]}) or a list of objects.'
M_HEADER = ('This looks like an FEC bulk file without a header row and its name is not a known bulk type. '
            'Give its header file too.')
M_XLSX_READ = 'The workbook could not be opened. It may be corrupt or password-protected.'
WINDOWS_1252_WARNING = 'Decoded as Windows-1252 because the file is not valid UTF-8.'
HEADER_BUNDLED_WARNING = 'Header taken from bundled.'


def m_type(suffix):
    return f'This file type ({suffix or "no extension"}) is not supported. Choose CSV, TSV, TXT, XLSX, JSON or an FEC bulk ZIP.'


def m_zip(members):
    return f'The ZIP must hold exactly one .txt or .csv file; it holds: {", ".join(members[:10])}.'


def m_sheet(name, sheets):
    return f'Worksheet {name!r} is not in this workbook. Sheets: {", ".join(sheets)}.'


def m_path(path, reason):
    return f'Cannot open {path}: {reason}'


def m_read(row_number):
    return f'The CSV could not be read at row {row_number:,}. Check its quoted fields and delimiter.'


# ---------------------------------------------------------------------------
# Fixture plumbing shared with test_evidence_csv_stream.py's own local copies
# (kept local here too, not imported, for the same reason as the message
# constants above).
# ---------------------------------------------------------------------------
def rmtree_retry(path, attempts=15, delay=0.2):
    path = Path(path)
    if not path.exists():
        return
    for attempt in range(attempts):
        try:
            shutil.rmtree(path)
            return
        except (PermissionError, OSError):
            if attempt == attempts - 1:
                raise
            time.sleep(delay)


@contextmanager
def temp_data_root():
    root = Path(tempfile.mkdtemp(prefix='mentat-evidence-readers-'))
    try:
        yield root
    finally:
        rmtree_retry(root)


@contextmanager
def temp_fixture_dir():
    root = Path(tempfile.mkdtemp(prefix='mentat-evidence-readers-fixtures-'))
    try:
        yield root
    finally:
        rmtree_retry(root)


def start_console_server(root, *, read_only=True):
    with patch.object(app, 'report_directory', return_value=root), patch.object(app, 'environment_keys', return_value=[]):
        console = app.Console(root / 'state', read_only=read_only)
    server = app.ThreadingHTTPServer(('127.0.0.1', 0), app.Handler)
    server.app = console
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    return console, server, worker


def stop_console_server(server, worker):
    server.shutdown()
    server.server_close()
    worker.join()


# ---------------------------------------------------------------------------
# Fixture builders
# ---------------------------------------------------------------------------
def write_xlsx(path, sheets, active=None):
    """sheets: {name: [[cell, cell, ...], ...]} (first row is the header row).
    active: name of the sheet openpyxl should leave selected (first sheet by
    default, matching a plain save)."""
    workbook = openpyxl.Workbook()
    default_sheet = workbook.active
    names = list(sheets)
    for index, name in enumerate(names):
        sheet = default_sheet if index == 0 else workbook.create_sheet()
        sheet.title = name
        for row in sheets[name]:
            sheet.append(row)
    if active and active in names:
        workbook.active = names.index(active)
    workbook.save(path)


def two_sheet_workbook(path):
    write_xlsx(path, {
        'Sheet1': [
            ['committee_id', 'committee_name', 'contribution_receipt_amount', 'contribution_receipt_date',
             'two_year_transaction_period'],
            ['C00000001', 'Example PAC', 10.0, '2026-01-05', '2026'],
            ['C00000001', 'Example PAC', 20.0, '2026-01-06', '2026'],
        ],
        'Detail': [
            ['committee_id', 'committee_name', 'contribution_receipt_amount', 'contribution_receipt_date',
             'two_year_transaction_period'],
            ['C00000002', 'Detail PAC', 30.0, '2026-02-05', '2026'],
        ],
    })


def cells_xlsx(path):
    """One row exercising every cell-to-text rule the contract lists: None -> '',
    an int-valued float -> its integer text, a date -> 'YYYY-MM-DD', a bool ->
    'TRUE'/'FALSE'."""
    write_xlsx(path, {'Sheet1': [
        ['committee_id', 'note', 'amount', 'when', 'flag'],
        ['C00000001', None, 12.0, date_cls(2026, 3, 4), True],
    ]})


def wide_xlsx_201_columns(path):
    header = ['committee_id'] + [f'col{i}' for i in range(200)]
    row = ['C00000001'] + ['x'] * 200
    write_xlsx(path, {'Sheet1': [header, row]})


FEC_INDIV_HEADER = ('CMTE_ID,AMNDT_IND,RPT_TP,TRANSACTION_PGI,IMAGE_NUM,TRANSACTION_TP,ENTITY_TP,NAME,CITY,STATE,'
                    'ZIP_CODE,EMPLOYER,OCCUPATION,TRANSACTION_DT,TRANSACTION_AMT,OTHER_ID,TRAN_ID,FILE_NUM,'
                    'MEMO_CD,MEMO_TEXT,SUB_ID').split(',')


def fec_indiv_row(index):
    # NAME deliberately has no comma (real FEC bulk data does: "Doe, Jane"):
    # a comma in an otherwise-pipe-only file is exactly the case where
    # csv.Sniffer's own count-based fallback (evidence_csv.sniff_dialect) can
    # mis-pick ',' over '|', which only matters for the two tests that route
    # through the sniffed delimited_reader (an unrecognised stem, no
    # header_path); bulk_reader itself (a recognised stem, or header_path
    # given) never sniffs at all -- it hardcodes delimiter='|' -- so this
    # fixture is safe to reuse everywhere.
    values = {name: '' for name in FEC_INDIV_HEADER}
    values.update({
        'CMTE_ID': 'C00000001', 'AMNDT_IND': 'N', 'RPT_TP': 'Q1', 'TRANSACTION_PGI': 'P2026',
        'IMAGE_NUM': f'2026030100000{index}', 'TRANSACTION_TP': '15', 'ENTITY_TP': 'IND',
        'NAME': f'DOE JANE {index}', 'CITY': 'WASHINGTON', 'STATE': 'DC', 'ZIP_CODE': '20002',
        'EMPLOYER': 'EXAMPLE EMPLOYER', 'OCCUPATION': 'ANALYST',
        'TRANSACTION_DT': '0%d152026' % (1 + index % 9), 'TRANSACTION_AMT': f'{index}00.00',
        'OTHER_ID': '', 'TRAN_ID': f'SA-{index}', 'FILE_NUM': '1234567', 'MEMO_CD': '', 'MEMO_TEXT': '',
        'SUB_ID': f'400000000000{index}',
    })
    return '|'.join(values[name] for name in FEC_INDIV_HEADER)


def fec_indiv_body(row_count=5):
    return '\r\n'.join(fec_indiv_row(index) for index in range(1, row_count + 1)) + '\r\n'


def write_fec_bulk_zip(path, member_name, body):
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(member_name, body)


def bundled_indiv_header_expected():
    """The contract falls back to mentat/inputs/**/indiv_header_file.csv before the
    bundled table; this repo's live mentat/inputs/ has no such file (confirmed by a
    directory search), so 'bundled' is the expected header_source today. Written as
    a live filesystem check, not a hardcoded assumption, so a future inputs file
    changes this test's expectation instead of silently going stale."""
    if any(MENTAT_INPUTS.glob('**/indiv_header_file.csv')):
        return 'inputs'
    return 'bundled'


# ---------------------------------------------------------------------------
# Contract assertion helpers: schema-by-example, kept local to this file (see
# module docstring). Mirrors test_evidence_csv_stream.py's helper in spirit; the
# S3b examples living inside evidence-csv-import.json carry MORE keys than the
# S3-era ones (format/source_kind/source_path/sheet/sheets/encoding/delimiter/
# header_source), so they land in their own frozenset-keyed group and never
# collide with the S3 suite's own CATALOG_EXAMPLES/JOB_EXAMPLES groups.
# ---------------------------------------------------------------------------
def _type_family(value):
    if isinstance(value, bool):
        return 'bool'
    if isinstance(value, (int, float)):
        return 'number'
    if isinstance(value, str):
        return 'string'
    if value is None:
        return 'null'
    if isinstance(value, list):
        return 'array'
    if isinstance(value, dict):
        return 'object'
    raise TypeError(f'unsupported example value type: {type(value)!r}')


def assert_shape(value, examples, enums=None, path='$'):
    enums = enums or {}
    if not isinstance(value, dict):
        raise AssertionError(f'{path}: expected an object, got {type(value).__name__}')
    keyset = set(examples[0].keys())
    for example in examples[1:]:
        if set(example.keys()) != keyset:
            raise AssertionError(f'{path}: contract examples disagree on keys: '
                                  f'{sorted(example.keys())} vs {sorted(keyset)}')
    if set(value.keys()) != keyset:
        raise AssertionError(f'{path}: key set {sorted(value.keys())} != contract key set {sorted(keyset)}')
    for key in sorted(keyset):
        field_path = f'{path}.{key}'
        allowed = {_type_family(example[key]) for example in examples}
        actual = _type_family(value[key])
        if actual not in allowed:
            raise AssertionError(f'{field_path}: type {actual!r} not in contract types {sorted(allowed)}')
        if key in enums and value[key] not in enums[key]:
            raise AssertionError(f'{field_path}: value {value[key]!r} not in allowed set {sorted(enums[key])}')


def _group_by_keyset(examples):
    groups = {}
    for example in examples:
        groups.setdefault(frozenset(example.keys()), []).append(example)
    return groups


IMPORT_CONTRACT = load_contract('evidence-csv-import')
BROWSE_CONTRACT = load_contract('evidence-files-browse')
_import_groups = _group_by_keyset(IMPORT_CONTRACT['examples'])
_browse_groups = _group_by_keyset(BROWSE_CONTRACT['examples'])

JOB_KEYS_S3B = frozenset({'id', 'name', 'state', 'bytes', 'bytes_read', 'rows_read', 'message', 'started_at',
                          'finished_at', 'catalog', 'format', 'source_kind', 'source_path', 'sheet', 'sheets', 'sheets_read',
                          'encoding', 'delimiter', 'header_source'})
LIMITS_KEYS_S3B = frozenset({'max_bytes', 'max_columns', 'max_rows', 'max_bytes_message', 'max_columns_message',
                             'max_rows_message', 'accepted_suffixes', 'max_json_bytes'})
SOURCES_LISTING_KEYS = frozenset({'sources', 'limits', 'session_only'})
CATALOG_KEYS_S3B = frozenset({'id', 'name', 'sha256', 'row_count', 'committees', 'cycles', 'supported_rows',
                              'unsupported_rows', 'unassigned_rows', 'session_only', 'imported_at', 'bytes',
                              'warnings', 'format', 'source_kind', 'source_path', 'sheet', 'sheets', 'sheets_read', 'encoding',
                              'delimiter', 'header_source',
                              # S7
                              'categories', 'excluded_rows', 'catalog_version', 'state', 'reread', 'stale',
                              'rows_read',
                              # S6: the industry a Mentat workbook file name declares
                              'industry'})
COMMITTEE_KEYS = frozenset({'id', 'name', 'rows', 'supported_rows', 'excluded_rows', 'category',
                            'category_source'})

JOB_EXAMPLES_S3B = _import_groups[JOB_KEYS_S3B]
LIMITS_EXAMPLES_S3B = _import_groups[LIMITS_KEYS_S3B]
CATALOG_EXAMPLES_S3B = [example['catalog'] for example in JOB_EXAMPLES_S3B if example['catalog'] is not None]
COMMITTEE_EXAMPLES = CATALOG_EXAMPLES_S3B[0]['committees']
# v5 (console redesign build 06, board B19): in the sources listing, each committee of
# an import that is not being re-read also carries the figure its pane leads with, for the
# PAC picker rail. The listing example's committees show both states (a figure, and none).
LISTED_FIGURE_KEYS = ('receipts', 'receipts_channel', 'receipts_by_cycle')
LISTED_COMMITTEE_EXAMPLES = next(example['sources'][0]['committees']
                                 for example in _import_groups[SOURCES_LISTING_KEYS] if example['sources'])

ENTRY_KEYS = frozenset({'name', 'path', 'kind', 'size', 'modified', 'supported'})
SHORTCUT_KEYS = frozenset({'name', 'path'})
LISTING_KEYS = frozenset({'path', 'parent', 'roots', 'shortcuts', 'entries', 'truncated'})
BROWSE_ERROR_KEYS = frozenset({'error'})
LISTING_EXAMPLES = _browse_groups[LISTING_KEYS]
BROWSE_ERROR_EXAMPLES = _browse_groups[BROWSE_ERROR_KEYS]
ENTRY_EXAMPLES = [entry for example in LISTING_EXAMPLES for entry in example['entries']]
SHORTCUT_EXAMPLES = [shortcut for example in LISTING_EXAMPLES for shortcut in example['shortcuts']]

def category_warning(text):
    """True for the per-import sentence S7 built to replace the static committee-type note."""
    return (text.startswith('Committee categories ')
            or text.startswith('The FEC committee master is not available')
            or text.startswith('The FEC committee master file '))


CATALOG_WARNINGS_S3B = {
    'Imported observations are not verified coverage or donor-attributed PAC spending. Duplicate and amended '
    'rows are preserved; totals may overlap.',
    'Cycle aggregates have no inferred transaction date. Missing or ambiguous amounts, dates, or types remain '
    'unsupported.',
    WINDOWS_1252_WARNING,
    HEADER_BUNDLED_WARNING,
    'Header taken from inputs.',
    'Header taken from header_path.',
}


def assert_committee(value):
    assert_shape(value, COMMITTEE_EXAMPLES, enums={'category': set(csv_data.CATEGORIES) | {''}})


def assert_catalog_s3b(value):
    assert_shape(value, CATALOG_EXAMPLES_S3B,
                 enums={'format': {'delimited', 'xlsx', 'json', 'fec_bulk'},
                        'source_kind': {'upload', 'path'},
                        'state': {'ready', 'reading'},
                        'encoding': {'utf-8', 'utf-8-sig', 'utf-16', 'cp1252', ''},
                        'delimiter': {',', ';', '\t', '|', ''},
                        'header_source': {'row', 'header_path', 'inputs', 'bundled', ''}})
    if not isinstance(value['sheets'], list):
        raise AssertionError(f"catalog.sheets must be an array, got {type(value['sheets']).__name__}")
    for committee in value['committees']:
        assert_committee(committee)
    names = [item['category'] for item in value['categories']]
    if names != [name for name in csv_data.CATEGORIES if name in set(names)]:
        raise AssertionError(f'catalog.categories out of vocabulary order: {names}')
    for warning in value['warnings']:
        if warning not in CATALOG_WARNINGS_S3B and not category_warning(warning):
            raise AssertionError(f'unexpected catalog warning: {warning!r}')


def assert_job_s3b(value):
    assert_shape(value, JOB_EXAMPLES_S3B, enums={'state': {'reading', 'ready', 'failed'}})
    ready = value['state'] == 'ready'
    has_catalog = value['catalog'] is not None
    if ready != has_catalog:
        raise AssertionError("JOB.catalog must be non-null iff state == 'ready' "
                              f"(state={value['state']!r}, catalog is {'not ' if has_catalog else ''}None)")
    if has_catalog:
        assert_catalog_s3b(value['catalog'])
    if value['state'] != 'failed' and value['message'] != '':
        raise AssertionError(f"JOB.message must be '' unless failed, got {value['message']!r}")


def assert_limits_s3b(value):
    assert_shape(value, LIMITS_EXAMPLES_S3B)
    if not isinstance(value['accepted_suffixes'], list):
        raise AssertionError('limits.accepted_suffixes must be an array')


def assert_listed_committee(value):
    assert_shape(value, LISTED_COMMITTEE_EXAMPLES, enums={'category': set(csv_data.CATEGORIES) | {''}})


def without_listed_figures(source):
    """The listed import as its CATALOG: every committee without the rail's figure."""
    return {**source, 'committees': [{key: item for key, item in committee.items() if key not in LISTED_FIGURE_KEYS}
                                     for committee in source['committees']]}


def assert_sources_listing_s3b(value):
    assert_shape(value, _import_groups[SOURCES_LISTING_KEYS], enums={'session_only': {False}})
    assert_limits_s3b(value['limits'])
    for source in value['sources']:
        if source.get('state') == 'reading':
            # v5: an import being re-read lists its committees exactly as its CATALOG does.
            assert_catalog_s3b(source)
            continue
        for committee in source['committees']:
            assert_listed_committee(committee)
        assert_catalog_s3b(without_listed_figures(source))


def assert_entry(value):
    assert_shape(value, ENTRY_EXAMPLES, enums={'kind': {'dir', 'file'}})


def assert_shortcut(value):
    assert_shape(value, SHORTCUT_EXAMPLES)


def assert_listing(value):
    assert_shape(value, LISTING_EXAMPLES)
    for shortcut in value['shortcuts']:
        assert_shortcut(shortcut)
    seen_file = False
    previous_key = None
    for entry in value['entries']:
        assert_entry(entry)
        if entry['kind'] == 'file':
            seen_file = True
        elif seen_file:
            raise AssertionError('browse() must list directories before files: '
                                  f'a dir ({entry["name"]!r}) followed a file')
        sort_key = (entry['kind'] != 'dir', entry['name'].lower())
        if previous_key is not None and sort_key < previous_key:
            raise AssertionError(f'browse() entries are not sorted (dirs first, then name '
                                  f'case-insensitive): {entry["name"]!r} out of order')
        previous_key = sort_key


def assert_browse_error(value):
    assert_shape(value, BROWSE_ERROR_EXAMPLES)
    if not value['error'].startswith('Cannot open '):
        raise AssertionError(f"a browse error must read 'Cannot open <path>: <reason>', got {value['error']!r}")


IMPORT_DISPATCH_S3B = {'catalog': assert_catalog_s3b, 'sources_listing': assert_sources_listing_s3b}
BROWSE_DISPATCH = {'listing': assert_listing, 'error': assert_browse_error}


class EvidenceReadersTests(unittest.TestCase):
    # -- xlsx ----------------------------------------------------------------------
    def test_xlsx_first_sheet_default(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            workbook_path = fixtures / 'sheets.xlsx'
            two_sheet_workbook(workbook_path)
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(workbook_path))['id'])
            self.assertEqual(job['state'], 'ready')
            catalog = job['catalog']
            self.assertEqual(catalog['format'], 'xlsx')
            self.assertEqual(catalog['sheets'], ['Sheet1', 'Detail'])
            self.assertEqual(catalog['sheet'], 'Sheet1')
            # Sheet1 and Detail share one header, so both are read as one table.
            self.assertEqual(catalog['sheets_read'], ['Sheet1', 'Detail'])
            self.assertEqual(catalog['row_count'], 3)
            self.assertEqual({c['id'] for c in catalog['committees']}, {'C00000001', 'C00000002'})
            self.assertIn('Read 2 worksheets with the same columns: Sheet1, Detail.', catalog['warnings'])

    def test_xlsx_named_sheet_distinct_id(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            workbook_path = fixtures / 'sheets.xlsx'
            two_sheet_workbook(workbook_path)
            store = csv_data.Store(root)
            default_job = store.wait(store.import_path(str(workbook_path))['id'])
            detail_job = store.wait(store.import_path(str(workbook_path), sheet='Detail')['id'])
            self.assertEqual(default_job['state'], 'ready')
            self.assertEqual(detail_job['state'], 'ready')
            self.assertNotEqual(default_job['id'], detail_job['id'],
                                 'importing the same bytes under a different sheet must mint a distinct id')
            self.assertEqual(detail_job['catalog']['sheet'], 'Detail')
            self.assertEqual({c['id'] for c in detail_job['catalog']['committees']}, {'C00000002'})

    def test_xlsx_unknown_sheet_message(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            workbook_path = fixtures / 'sheets.xlsx'
            two_sheet_workbook(workbook_path)
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(workbook_path), sheet='Nope')['id'])
            self.assertEqual(job['state'], 'failed')
            self.assertEqual(job['message'], m_sheet('Nope', ['Sheet1', 'Detail']))

    def test_xlsx_cells_to_text(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            workbook_path = fixtures / 'cells.xlsx'
            cells_xlsx(workbook_path)
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(workbook_path))['id'])
            self.assertEqual(job['state'], 'ready')
            selection = store.select(job['id'], committee='C00000001')
            self.assertEqual(len(selection['rows']), 1)
            raw = selection['rows'][0]['raw']
            self.assertEqual(raw['note'], '')
            self.assertEqual(raw['amount'], '12')
            self.assertEqual(raw['when'], '2026-03-04')
            self.assertEqual(raw['flag'], 'TRUE')

    def test_xlsx_201_columns_rejected(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            workbook_path = fixtures / 'wide201.xlsx'
            wide_xlsx_201_columns(workbook_path)
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(workbook_path))['id'])
            self.assertEqual(job['state'], 'failed')
            self.assertEqual(job['message'], csv_data.columns_message())

    def test_xlsx_corrupt_bytes(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            corrupt_path = fixtures / 'corrupt.xlsx'
            corrupt_path.write_bytes(os.urandom(256))
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(corrupt_path))['id'])
            self.assertEqual(job['state'], 'failed')
            self.assertEqual(job['message'], M_XLSX_READ)

    # -- delimited text (tsv/txt, delimiter sniffing, encoding fallback) ------------
    def test_tsv_tab_delimiter(self):
        text = 'committee_id\tcommittee_name\tcontribution_receipt_amount\tcontribution_receipt_date\ttwo_year_transaction_period\r\n' \
               'C00000001\tExample PAC\t10.00\t2026-01-05\t2026\r\n'
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'evidence.tsv'
            path.write_text(text, encoding='utf-8')
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'ready')
            self.assertEqual(job['catalog']['format'], 'delimited')
            self.assertEqual(job['catalog']['delimiter'], '\t')
            self.assertEqual(job['catalog']['row_count'], 1)

    def test_txt_pipe_with_header(self):
        text = 'committee_id|committee_name|contribution_receipt_amount|contribution_receipt_date|two_year_transaction_period\r\n' \
               'C00000001|Example PAC|10.00|2026-01-05|2026\r\n'
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'evidence-pipe.txt'
            path.write_text(text, encoding='utf-8')
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'ready')
            self.assertEqual(job['catalog']['delimiter'], '|')

    def test_txt_semicolon(self):
        text = 'committee_id;committee_name;contribution_receipt_amount;contribution_receipt_date;two_year_transaction_period\r\n' \
               'C00000001;Example PAC;10.00;2026-01-05;2026\r\n'
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'evidence-semi.txt'
            path.write_text(text, encoding='utf-8')
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'ready')
            self.assertEqual(job['catalog']['delimiter'], ';')

    def test_csv_windows_1252_fallback_warns(self):
        text = 'committee_id,committee_name,contribution_receipt_amount,contribution_receipt_date,two_year_transaction_period\r\n' \
               'C00000001,Caf\xe9 PAC,10.00,2026-01-05,2026\r\n'  # 'Café PAC' -- 0xE9 is invalid standalone UTF-8
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'evidence-1252.csv'
            path.write_bytes(text.encode('cp1252'))
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'ready')
            self.assertEqual(job['catalog']['encoding'], 'cp1252')
            self.assertIn(WINDOWS_1252_WARNING, job['catalog']['warnings'])

    def test_csv_utf16_bom(self):
        text = ('committee_id,committee_name,contribution_receipt_amount,contribution_receipt_date,'
                'two_year_transaction_period\r\nC00000001,Example PAC,10.00,2026-01-05,2026\r\n')
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'evidence-utf16.csv'
            path.write_bytes(text.encode('utf-16'))  # includes the LE/BE BOM
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'ready')
            self.assertEqual(job['catalog']['encoding'], 'utf-16')
            self.assertEqual(job['catalog']['row_count'], 1)

    def test_csv_nul_bytes_rejected(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'evidence-nul.csv'
            path.write_bytes(b'committee_id,amount\r\nC00000001,\x0010.00\r\n')
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'failed')
            self.assertEqual(job['message'], M_BINARY)

    def test_truncated_delimited_names_row(self):
        text = ('committee_id,committee_name,contribution_receipt_amount,contribution_receipt_date\r\n'
                'C00000001,Example PAC,10.00,2026-01-01\r\n'
                'C00000002,Example PAC,10.00,"Doe, Jane')
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'truncated.csv'
            path.write_text(text, encoding='utf-8')
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'failed')
            self.assertRegex(job['message'], r'^The CSV could not be read at row [\d,]+\. Check its quoted fields and delimiter\.$')

    # -- json ------------------------------------------------------------------------
    def test_json_results_shape(self):
        payload = {'results': [
            {'committee_id': 'C00000001', 'committee_name': 'Example PAC',
             'contribution_receipt_amount': 10.0, 'contribution_receipt_date': '2026-01-05',
             'two_year_transaction_period': '2026'},
        ]}
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'results.json'
            path.write_text(json.dumps(payload), encoding='utf-8')
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'ready')
            self.assertEqual(job['catalog']['format'], 'json')
            self.assertEqual(job['catalog']['row_count'], 1)

    def test_json_list_shape(self):
        payload = [
            {'committee_id': 'C00000001', 'committee_name': 'Example PAC',
             'contribution_receipt_amount': 10.0, 'contribution_receipt_date': '2026-01-05',
             'two_year_transaction_period': '2026'},
        ]
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'list.json'
            path.write_text(json.dumps(payload), encoding='utf-8')
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'ready')
            self.assertEqual(job['catalog']['row_count'], 1)

    def test_json_wrong_shape(self):
        payload = {'committee_id': 'C00000001'}  # neither {"results": [...]} nor a bare list
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'wrong.json'
            path.write_text(json.dumps(payload), encoding='utf-8')
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'failed')
            self.assertEqual(job['message'], M_JSON_SHAPE)

    def test_json_over_cap(self):
        # The message names the CURRENT cap (evidence_csv.json_size_message()),
        # the same pattern test_evidence_csv_stream.py uses for bytes_message()
        # -- M_JSON_SIZE (this file's hardcoded copy of the DEFAULT 25 MB text)
        # would go stale the moment the cap is narrowed for the test's own duration.
        original_cap = csv_data.MAX_JSON_BYTES
        csv_data.MAX_JSON_BYTES = 1000
        try:
            payload = {'results': [{'committee_id': 'C00000001', 'filler': 'x' * 2000}]}
            with temp_fixture_dir() as fixtures, temp_data_root() as root:
                path = fixtures / 'big.json'
                path.write_text(json.dumps(payload), encoding='utf-8')
                store = csv_data.Store(root)
                job = store.wait(store.import_path(str(path))['id'])
                self.assertEqual(job['state'], 'failed')
                self.assertEqual(job['message'], csv_data.json_size_message())
        finally:
            csv_data.MAX_JSON_BYTES = original_cap

    def test_json_nested_values_are_strings(self):
        payload = {'results': [
            {'committee_id': 'C00000001', 'committee_name': 'Example PAC',
             'contribution_receipt_amount': 10.0, 'contribution_receipt_date': '2026-01-05',
             'two_year_transaction_period': '2026', 'contributor_employer': {'name': 'Acme', 'sector': 'Tech'}},
        ]}
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'nested.json'
            path.write_text(json.dumps(payload), encoding='utf-8')
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'ready')
            selection = store.select(job['id'], committee='C00000001')
            nested_text = selection['rows'][0]['raw']['contributor_employer']
            self.assertIsInstance(nested_text, str)
            self.assertEqual(json.loads(nested_text), {'name': 'Acme', 'sector': 'Tech'})

    # -- FEC bulk ----------------------------------------------------------------------
    def test_fec_bulk_zip_indiv_bundled_header(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            zip_path = fixtures / 'indiv26.zip'
            write_fec_bulk_zip(zip_path, 'indiv.txt', fec_indiv_body(5))
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(zip_path))['id'])
            self.assertEqual(job['state'], 'ready')
            catalog = job['catalog']
            self.assertEqual(catalog['format'], 'fec_bulk')
            self.assertEqual(catalog['delimiter'], '|')
            self.assertEqual(catalog['row_count'], 5)
            self.assertEqual(catalog['header_source'], bundled_indiv_header_expected())
            self.assertIn('C00000001', {c['id'] for c in catalog['committees']})

    def test_fec_bulk_txt_with_header_path(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            header_path = fixtures / 'my_header.csv'
            header_path.write_text(','.join(FEC_INDIV_HEADER), encoding='utf-8')
            txt_path = fixtures / 'unrecognized_stem.txt'
            txt_path.write_text(fec_indiv_body(3), encoding='utf-8')
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(txt_path), header_path=str(header_path))['id'])
            self.assertEqual(job['state'], 'ready')
            catalog = job['catalog']
            self.assertEqual(catalog['format'], 'fec_bulk')
            self.assertEqual(catalog['header_source'], 'header_path')
            self.assertEqual(catalog['row_count'], 3)

    def test_fec_bulk_zip_two_members_rejected(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            zip_path = fixtures / 'indiv-two.zip'
            with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as archive:
                archive.writestr('indiv-a.txt', fec_indiv_body(2))
                archive.writestr('indiv-b.txt', fec_indiv_body(2))
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(zip_path))['id'])
            self.assertEqual(job['state'], 'failed')
            self.assertEqual(job['message'], m_zip(['indiv-a.txt', 'indiv-b.txt']))

    def test_fec_bulk_zip_member_over_cap(self):
        # MAX_BYTES has to sit strictly between the ZIP's own (compressed) size
        # on disk and the member's UNCOMPRESSED size, or import_path()'s own
        # whole-file stat check (see test_import_path_over_cap) fires first and
        # this never reaches the per-member check inside bulk_reader. 20,000
        # highly repetitive pipe rows measured ~11.5 KB compressed / ~2.9 MB
        # uncompressed; 1,000,000 sits cleanly between the two.
        original_max = csv_data.MAX_BYTES
        csv_data.MAX_BYTES = 1_000_000
        try:
            with temp_fixture_dir() as fixtures, temp_data_root() as root:
                zip_path = fixtures / 'indiv-big.zip'
                write_fec_bulk_zip(zip_path, 'indiv.txt', fec_indiv_body(20000))
                self.assertLess(zip_path.stat().st_size, csv_data.MAX_BYTES,
                                 'fixture no longer compresses under MAX_BYTES; the whole-file '
                                 'check would fire before the per-member one this test targets')
                store = csv_data.Store(root)
                job = store.wait(store.import_path(str(zip_path))['id'])
                self.assertEqual(job['state'], 'failed')
                self.assertEqual(job['message'], csv_data.bytes_message())
        finally:
            csv_data.MAX_BYTES = original_max

    def test_unknown_pipe_txt_without_header_rejected(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'mystery.txt'
            path.write_text(fec_indiv_body(2), encoding='utf-8')  # unrecognised stem, no header row, no header_path
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'failed')
            self.assertEqual(job['message'], M_HEADER)

    def test_unknown_pipe_txt_with_header_row_is_delimited(self):
        text = 'committee_id|committee_name|contribution_receipt_amount|contribution_receipt_date|two_year_transaction_period\r\n' \
               'C00000001|Example PAC|10.00|2026-01-05|2026\r\n'
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'mystery-with-header.txt'
            path.write_text(text, encoding='utf-8')
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(path))['id'])
            self.assertEqual(job['state'], 'ready')
            self.assertEqual(job['catalog']['format'], 'delimited')

    # -- unsupported / id compatibility ------------------------------------------------
    def test_unsupported_suffix(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            store = csv_data.Store(root)
            pdf_path = fixtures / 'notes.pdf'
            pdf_path.write_bytes(b'%PDF-1.4 not really a pdf')
            job = store.wait(store.import_path(str(pdf_path))['id'])
            self.assertEqual(job['state'], 'failed')
            self.assertEqual(job['message'], m_type('.pdf'))

            no_ext_path = fixtures / 'README'
            no_ext_path.write_text('just some notes', encoding='utf-8')
            job2 = store.wait(store.import_path(str(no_ext_path))['id'])
            self.assertEqual(job2['state'], 'failed')
            self.assertEqual(job2['message'], m_type(''))

    def test_csv_id_unchanged_from_s3(self):
        # The S3b id formula is 'csv-' + sha256(bytes + b'\x00' + sheet.encode())[:32];
        # "sheet '' for non-xlsx (existing CSV ids unchanged)" is read here as the
        # authoritative outcome -- for sheet == '' the NUL+sheet suffix contributes
        # nothing, so the id must equal the plain S3-era sha256(bytes) id
        # (test_evidence_csv_stream.py:test_dedup_same_bytes_same_id uses exactly
        # this formula and must keep passing once S3b lands).
        text = ('committee_id,committee_name,contribution_receipt_amount,contribution_receipt_date,'
                'two_year_transaction_period\r\nC00000001,Example PAC,10.00,2026-01-05,2026\r\n')
        expected_id = 'csv-' + hashlib.sha256(text.encode('utf-8')).hexdigest()[:32]
        with temp_data_root() as root:
            store = csv_data.Store(root)
            job = store.wait(store.ingest_text('same-as-s3.csv', text)['id'])
            self.assertEqual(job['id'], expected_id)

    # -- Store.import_path -------------------------------------------------------------
    # Python API: "Store.import_path(path, sheet='', header_path='') -> JOB (async
    # ingest; M_PATH cases; ...; unsupported suffix -> M_TYPE)" -- read literally,
    # import_path() always returns a JOB and every failure mode named there (a bad
    # path, an unsupported suffix, an unknown sheet, a corrupt workbook) lands as
    # job['state'] == 'failed' with the message, the same as the pre-existing
    # M_COLUMNS/M_READ/etc. failures in test_evidence_csv_stream.py -- never a
    # raised exception from import_path() itself.
    def test_import_path_xlsx(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            workbook_path = fixtures / 'plain.xlsx'
            write_xlsx(workbook_path, {'Sheet1': [
                ['committee_id', 'committee_name', 'contribution_receipt_amount', 'contribution_receipt_date',
                 'two_year_transaction_period'],
                ['C00000001', 'Example PAC', 10.0, '2026-01-05', '2026'],
            ]})
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(workbook_path))['id'])
            self.assertEqual(job['state'], 'ready')
            self.assertEqual(job['catalog']['source_kind'], 'path')
            self.assertEqual(job['catalog']['source_path'], str(workbook_path))

    def test_import_path_missing(self):
        with temp_data_root() as root:
            store = csv_data.Store(root)
            missing = str(Path(root) / 'does-not-exist.csv')
            job = store.wait(store.import_path(missing)['id'])
            self.assertEqual(job['state'], 'failed')
            self.assertEqual(job['message'], m_path(missing, 'no such file'))

    def test_import_path_folder(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            folder = fixtures / 'a-folder.csv'
            folder.mkdir()
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(folder))['id'])
            self.assertEqual(job['state'], 'failed')
            self.assertEqual(job['message'], m_path(str(folder), 'it is a folder'))

    def test_import_path_over_cap(self):
        original_max = csv_data.MAX_BYTES
        csv_data.MAX_BYTES = 2_000_000
        try:
            with temp_fixture_dir() as fixtures, temp_data_root() as root:
                big_path = fixtures / 'big.csv'
                with big_path.open('wb') as handle:
                    handle.truncate(csv_data.MAX_BYTES + 1)  # sparse; no real write needed
                store = csv_data.Store(root)
                job = store.wait(store.import_path(str(big_path))['id'])
                self.assertEqual(job['state'], 'failed')
                self.assertEqual(job['message'], m_path(str(big_path), f'larger than {csv_data.MAX_BYTES // 1_000_000} MB'))
        finally:
            csv_data.MAX_BYTES = original_max

    def test_import_path_copies_source(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            workbook_path = fixtures / 'copyme.xlsx'
            write_xlsx(workbook_path, {'Sheet1': [
                ['committee_id', 'committee_name', 'contribution_receipt_amount', 'contribution_receipt_date',
                 'two_year_transaction_period'],
                ['C00000001', 'Example PAC', 10.0, '2026-01-05', '2026'],
            ]})
            original_bytes = workbook_path.read_bytes()
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(workbook_path))['id'])
            self.assertEqual(job['state'], 'ready')
            stored_source = store.directory(job['id']) / 'source.xlsx'
            self.assertTrue(stored_source.is_file(), f'expected {stored_source} to exist')
            self.assertEqual(stored_source.read_bytes(), original_bytes)

    # -- browse() ------------------------------------------------------------------------
    def test_browse_root_lists_roots_and_shortcuts(self):
        result = csv_data.browse('')
        assert_listing(result)
        self.assertEqual(result['path'], '')
        self.assertEqual(result['entries'], [])
        self.assertGreater(len(result['roots']), 0)

    def test_browse_dir_sorted_dirs_first_and_supported_flags(self):
        with temp_fixture_dir() as fixtures:
            (fixtures / 'a-subdir').mkdir()
            (fixtures / 'z-subdir').mkdir()
            (fixtures / 'data.csv').write_text('committee_id\r\nC00000001\r\n', encoding='utf-8')
            (fixtures / 'notes.pdf').write_bytes(b'not a real pdf')
            result = csv_data.browse(str(fixtures))
            # assert_listing() already enforces "dirs first, then name case-insensitive"
            # structurally (see assert_listing's ordering check above); a violation here
            # would fail there, not just in this test's own assertions below.
            assert_listing(result)
            self.assertEqual(result['path'], str(fixtures))
            kinds = [(entry['name'], entry['kind'], entry['supported']) for entry in result['entries']]
            self.assertEqual([name for name, kind, _ in kinds if kind == 'dir'], ['a-subdir', 'z-subdir'])
            self.assertIn(('data.csv', 'file', True), kinds)
            self.assertIn(('notes.pdf', 'file', False), kinds)

    def test_browse_truncated_at_2000(self):
        with temp_fixture_dir() as fixtures:
            for index in range(2005):
                (fixtures / f'f{index:05d}.csv').write_text('', encoding='utf-8')
            result = csv_data.browse(str(fixtures))
            assert_listing(result)
            self.assertEqual(len(result['entries']), 2000)
            self.assertTrue(result['truncated'])

    def test_browse_unreadable(self):
        with temp_data_root() as root:
            missing = str(Path(root) / 'nowhere')
            with self.assertRaises(ValueError) as ctx:
                csv_data.browse(missing)
            self.assertEqual(str(ctx.exception), m_path(missing, 'no such file'))

        locked = r'C:\Windows\System32\config'
        with self.assertRaises(ValueError) as ctx:
            csv_data.browse(locked)
        self.assertTrue(str(ctx.exception).startswith('Cannot open '),
                         f'expected a bounded M_PATH message, got {ctx.exception!r}')

    # -- HTTP surface ------------------------------------------------------------------
    def test_http_path_route_and_browse(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            path = fixtures / 'http-path.csv'
            path.write_text('committee_id,contribution_receipt_amount,contribution_receipt_date,'
                             'two_year_transaction_period\r\nC00000001,10.00,2026-01-05,2026\r\n', encoding='utf-8')
            console, server, worker = start_console_server(root, read_only=False)
            origin = f'http://127.0.0.1:{server.server_port}'
            try:
                payload = json.dumps({'path': str(path)}).encode()
                headers = {'Content-Type': 'application/json', 'X-Mentat-Token': console.token,
                           'Content-Length': str(len(payload))}
                with urlopen(Request(origin + '/api/evidence-csv', data=payload, headers=headers)) as response:
                    self.assertEqual(response.status, 200)
                    job = json.load(response)
                deadline = time.time() + 15
                while job['state'] == 'reading' and time.time() < deadline:
                    time.sleep(0.05)
                    with urlopen(origin + '/api/evidence-csv?job=' + job['id']) as response:
                        job = json.load(response)
                self.assertEqual(job['state'], 'ready')

                browse_payload = json.dumps({'path': str(fixtures)}).encode()
                browse_headers = {'Content-Type': 'application/json', 'X-Mentat-Token': console.token,
                                  'Content-Length': str(len(browse_payload))}
                with urlopen(Request(origin + '/api/evidence-files', data=browse_payload, headers=browse_headers)) as response:
                    self.assertEqual(response.status, 200)
                    listing = json.load(response)
                assert_listing(listing)

                no_token = Request(origin + '/api/evidence-files', data=browse_payload,
                                    headers={'Content-Type': 'application/json', 'Content-Length': str(len(browse_payload))})
                with self.assertRaises(HTTPError) as ctx:
                    urlopen(no_token)
                self.assertEqual(ctx.exception.code, 403)

                with urlopen(origin + '/api/evidence-csv') as response:
                    body = json.load(response)
                self.assertIn('accepted_suffixes', body['limits'])
            finally:
                stop_console_server(server, worker)

    def test_http_streaming_with_sheet_header(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            workbook_path = fixtures / 'stream-sheets.xlsx'
            two_sheet_workbook(workbook_path)
            console, server, worker = start_console_server(root, read_only=True)
            origin = f'http://127.0.0.1:{server.server_port}'
            try:
                payload = workbook_path.read_bytes()
                headers = {'Content-Type': 'application/octet-stream', 'X-Mentat-Token': console.token,
                           'X-Evidence-Name': urllib.parse.quote('stream-sheets.xlsx'),
                           'X-Evidence-Sheet': urllib.parse.quote('Detail'),
                           'Content-Length': str(len(payload))}
                with urlopen(Request(origin + '/api/evidence-csv', data=payload, headers=headers, method='POST')) as response:
                    self.assertEqual(response.status, 200)
                    job = json.load(response)
                deadline = time.time() + 15
                while job['state'] == 'reading' and time.time() < deadline:
                    time.sleep(0.05)
                    with urlopen(origin + '/api/evidence-csv?job=' + job['id']) as response:
                        job = json.load(response)
                self.assertEqual(job['state'], 'ready')
                self.assertEqual(job['catalog']['sheet'], 'Detail')
            finally:
                stop_console_server(server, worker)

    # -- contract --------------------------------------------------------------------
    def test_contract_examples_against_producers(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            store = csv_data.Store(root)
            assert_limits_s3b(csv_data.limits())

            workbook_path = fixtures / 'contract.xlsx'
            two_sheet_workbook(workbook_path)
            job = store.wait(store.import_path(str(workbook_path), sheet='Detail')['id'])
            assert_job_s3b(job)
            self.assertEqual(job['state'], 'ready')

            listing = {'sources': store.sources(), 'limits': csv_data.limits(), 'session_only': False}
            assert_sources_listing_s3b(listing)

            browse_result = csv_data.browse('')
            assert_listing(browse_result)
            dir_result = csv_data.browse(str(fixtures))
            assert_listing(dir_result)

        for control in IMPORT_CONTRACT['negative_controls']:
            with self.subTest(contract='evidence-csv-import', why=control['why']):
                if control['shape'] not in IMPORT_DISPATCH_S3B:
                    continue  # owned by the S3-era shapes (catalog/sources_listing/error); asserted there.
                with self.assertRaises(AssertionError):
                    IMPORT_DISPATCH_S3B[control['shape']](control['payload'])
        for control in BROWSE_CONTRACT['negative_controls']:
            with self.subTest(contract='evidence-files-browse', why=control['why']):
                with self.assertRaises(AssertionError):
                    BROWSE_DISPATCH[control['shape']](control['payload'])



# ---------------------------------------------------------------------------
# S3c: a Mentat industry workbook (title rows above the header, header on row 4,
# several cycle sheets sharing one header, aggregate sheets alongside). Paul's
# 'Credit Union Industry Spending Overview.xlsx' failed on 21c9cab6d because the
# reader took the first sheet's first row as the header.
# ---------------------------------------------------------------------------
class MentatWorkbookSelection(unittest.TestCase):
    COLUMNS = ['Source ID', 'Cycle', 'FEC Filing Cycle', 'Cycle Alignment', 'Company', 'Contributor', 'Employee Tier',
               'Amount', 'Analysis Amount', 'Party', 'Recipient Category', 'Category Source', 'Party Source', 'Recipient',
               'Committee ID', 'Recipient Key', 'Date', 'City', 'State', 'ZIP', 'Employer', 'Occupation', 'Record Type']
    CYCLES = {
        '2023-2024': [('C00007880', "AMERICA'S CREDIT UNIONS PAC", 10.0, datetime.datetime(2024, 12, 31)),
                      ('C00007880', "AMERICA'S CREDIT UNIONS PAC", 15.5, datetime.datetime(2024, 6, 1)),
                      ('C00919084', 'TALARICO FOR TEXAS', 25.0, datetime.datetime(2024, 1, 2))],
        '2025-2026': [('C00919084', 'TALARICO FOR TEXAS', 30.77, datetime.datetime(2026, 3, 3))],
    }

    def build(self, folder, *, with_data=True):
        book = openpyxl.Workbook()
        overview = book.active
        overview.title = 'Master Overview'
        overview.append(['Credit Union Master Overview'])
        overview.append(['Mentat master-workbook layout'])
        overview.append([])
        overview.append(['Company', 'Rank', 'Total Giving'])
        overview.append(['Alpha CU', 1, 100.0])
        if with_data:
            for cycle_name, rows in self.CYCLES.items():
                sheet = book.create_sheet(cycle_name)
                sheet.append([f'{cycle_name} Credit Union Contributions'])
                sheet.append(['Contribution-date cycle.'])
                sheet.append(['Source DB SHA-256: fixture'])
                sheet.append(self.COLUMNS)
                for index, (committee, recipient, amount, when) in enumerate(rows, 1):
                    sheet.append([f'40{index}', cycle_name, cycle_name[-4:], 'OK', 'Alpha CU', 'DOE, JANE', 'Employee',
                                  amount, amount, 'D', 'Candidate', 'fec', '', recipient, committee, f'FEC:{committee}',
                                  when, 'CITY', 'ST', '00000', 'ALPHA CU', 'ANALYST', 'CONTRIBUTION'])
            downstream = book.create_sheet('Downstream Activity')
            downstream.append(['Downstream Committee-to-Candidate Activity'])
            downstream.append(['Same-cycle activity.'])
            downstream.append(['885,132 exact source rows'])
            downstream.append([])
            downstream.append(['Cycle', 'Source Committee', 'Source Committee ID', 'Candidate / Officeholder',
                               'Support / Oppose', 'Transaction Type', 'Net Included Amount'])
            downstream.append(['2021-2022', 'ABRAHAM LINCOLN PAC', 'C00631051', 'ARENHOLZ, ASHLEY', 'Contribution / Other',
                               '24K', 10000])
        target = Path(folder) / 'Credit Union Industry Spending Overview.xlsx'
        book.save(target)
        return target

    def test_contribution_sheets_chosen_and_totalled(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            target = self.build(fixtures)
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(target))['id'], timeout=120)
            self.assertEqual(job['state'], 'ready', job.get('message'))
            self.assertEqual(job['format'], 'xlsx')
            self.assertEqual(job['sheet'], '2023-2024')
            self.assertEqual(job['sheets_read'], ['2023-2024', '2025-2026'])
            self.assertEqual(job['sheets'], ['Master Overview', '2023-2024', '2025-2026', 'Downstream Activity'])
            catalog = job['catalog']
            self.assertEqual(catalog['row_count'], 4)
            self.assertEqual(catalog['supported_rows'], 4, catalog)
            self.assertIn('Read 2 worksheets with the same columns: 2023-2024, 2025-2026.', catalog['warnings'])
            names = {row['id']: row['name'] for row in catalog['committees']}
            self.assertEqual(names['C00919084'], 'TALARICO FOR TEXAS')
            self.assertEqual(catalog['cycles'], ['2026', '2024'])
            page = store.select(job['id'], 'C00007880', '2024')
            self.assertEqual(page['total_rows'], 2)
            self.assertEqual(page['totals'], [{'channel': 'Receipts', 'amount': 25.5, 'row_count': 2}])
            self.assertEqual(page['rows'][0]['date'], '2024-12-31')
            self.assertEqual(page['rows'][0]['status'], 'supported')

    def test_explicit_sheet_reads_the_aggregate_table(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            target = self.build(fixtures)
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(target), sheet='Downstream Activity')['id'], timeout=120)
            self.assertEqual(job['state'], 'ready', job.get('message'))
            self.assertEqual(job['sheet'], 'Downstream Activity')
            self.assertEqual(job['sheets_read'], ['Downstream Activity'])
            page = store.select(job['id'], 'C00631051')
            self.assertEqual(page['total_rows'], 1)
            self.assertTrue(page['rows'][0]['channel'].startswith('Reported cycle aggregate'), page['rows'][0])

    def test_workbook_without_a_committee_column_names_its_sheets(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            target = self.build(fixtures, with_data=False)
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(target))['id'], timeout=120)
            self.assertEqual(job['state'], 'failed')
            self.assertEqual(job['message'], 'No worksheet in this workbook has a Committee ID column. Sheets: Master Overview.')
            self.assertEqual(job['sheets'], ['Master Overview'])


# ---------------------------------------------------------------------------
# S3f: the FEC bulk TRANSACTION_TP -> channel ruling (inputs/fec_transaction_types.json).
# ---------------------------------------------------------------------------
class BulkCodeMapping(unittest.TestCase):
    HEADER = ('CMTE_ID,AMNDT_IND,RPT_TP,TRANSACTION_PGI,IMAGE_NUM,TRANSACTION_TP,ENTITY_TP,NAME,CITY,STATE,'
              'ZIP_CODE,EMPLOYER,OCCUPATION,TRANSACTION_DT,TRANSACTION_AMT,OTHER_ID,TRAN_ID,FILE_NUM,MEMO_CD,'
              'MEMO_TEXT,SUB_ID').split(',')

    def row(self, code, amount, *, rpt='Q1', memo='', index=1):
        values = {name: '' for name in self.HEADER}
        values.update({'CMTE_ID': 'C00000001', 'AMNDT_IND': 'N', 'RPT_TP': rpt, 'TRANSACTION_PGI': 'P2026',
                       'IMAGE_NUM': f'2026030100000{index}', 'TRANSACTION_TP': code, 'ENTITY_TP': 'IND',
                       'NAME': f'DOE JANE {index}', 'TRANSACTION_DT': '02152026', 'TRANSACTION_AMT': str(amount),
                       'TRAN_ID': f'SA-{index}', 'FILE_NUM': '1234567', 'MEMO_CD': memo, 'SUB_ID': f'4000000000{index:02d}'})
        return '|'.join(values[name] for name in self.HEADER)

    def build(self, folder):
        rows = [self.row('15', 100, index=1), self.row('18K', 250, index=2), self.row('22Y', 30, index=3),
                self.row('24K', 500, index=4), self.row('24E', 1000, index=5), self.row('24A', 700, index=6),
                self.row('24F', 90, index=7), self.row('15J', 40, index=8), self.row('15', 60, memo='X', index=9),
                self.row('24E', 300, rpt='24', index=10), self.row('99Z', 5, index=11)]
        target = Path(folder) / 'indiv26.txt'
        target.write_text('\r\n'.join(rows) + '\r\n', encoding='utf-8')
        return target

    def test_codes_map_and_exclusions_hold(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            store = csv_data.Store(root)
            job = store.wait(store.import_path(str(self.build(fixtures)))['id'], timeout=120)
            self.assertEqual(job['state'], 'ready', job.get('message'))
            self.assertEqual(job['format'], 'fec_bulk')
            catalog = job['catalog']
            self.assertEqual(catalog['row_count'], 11)
            self.assertEqual(catalog['supported_rows'], 6, catalog)
            self.assertTrue(any('Disbursements here are contributions' in note for note in catalog['warnings']))
            self.assertTrue(any('not been verified' in note for note in catalog['warnings']))
            page = store.select(job['id'], 'C00000001', '', 0, 100)
            totals = {item['channel']: (item['amount'], item['row_count']) for item in page['totals']}
            self.assertEqual(totals['Receipts'], (350.0, 2))
            self.assertEqual(totals['Disbursements'], (530.0, 2))
            self.assertEqual(totals['Independent spending · support'], (1000.0, 1))
            self.assertEqual(totals['Independent spending · oppose'], (700.0, 1))
            by_code = {row['raw']['TRANSACTION_TP'] + ':' + row['raw']['TRAN_ID']: row for row in page['rows']}
            self.assertEqual(by_code['24F:SA-7']['status'], 'unsupported')
            self.assertIn("Transaction type code 24F is excluded by ruling: Form 7 communication cost for a candidate: "
                          "the corporation's or union's own spending, not the PAC's.", by_code['24F:SA-7']['issues'])
            self.assertNotIn('The CSV does not identify a supported receipt or activity type.', by_code['24F:SA-7']['issues'])
            self.assertIn('Memo entry; not counted in totals.', by_code['15J:SA-8']['issues'])
            self.assertIn('Memo entry; not counted in totals.', by_code['15:SA-9']['issues'])
            self.assertEqual(by_code['15:SA-9']['channel'], 'Receipts')
            self.assertIn('Notice report (24/48-hour); the periodic filing carries this amount.', by_code['24E:SA-10']['issues'])
            self.assertEqual(by_code['24E:SA-10']['direction'], 'Support')
            self.assertIn('Transaction type code 99Z is not in the mapping table (inputs/fec_transaction_types.json).', by_code['99Z:SA-11']['issues'])
            self.assertNotIn('The CSV does not identify a supported receipt or activity type.', by_code['99Z:SA-11']['issues'])

    def test_missing_table_maps_nothing(self):
        with temp_fixture_dir() as fixtures, temp_data_root() as root:
            csv_data.transaction_types.cache_clear()
            with patch.object(csv_data, 'INPUTS', Path(fixtures) / 'no-inputs'):
                try:
                    store = csv_data.Store(root)
                    job = store.wait(store.import_path(str(self.build(fixtures)))['id'], timeout=120)
                    self.assertEqual(job['state'], 'ready', job.get('message'))
                    self.assertEqual(job['catalog']['supported_rows'], 0)
                    self.assertTrue(any('is missing' in note for note in job['catalog']['warnings']))
                finally:
                    csv_data.transaction_types.cache_clear()

    def test_workbook_transaction_type_column_is_untouched(self):
        # A Mentat 'Downstream Activity' aggregate row carries transaction_type (not
        # transaction_tp) and keeps its reported_aggregate classification.
        columns = ['cycle', 'source_committee', 'source_committee_id', 'support_oppose', 'transaction_type', 'net_included_amount']
        raw = dict(zip(['Cycle', 'Source Committee', 'Source Committee ID', 'Support / Oppose', 'Transaction Type', 'Net Included Amount'],
                       ['2021-2022', 'ABRAHAM LINCOLN PAC', 'C00631051', 'Contribution / Other', '24K', '10000']))
        row = csv_data.normalize(raw, 6, columns)
        self.assertEqual(row['kind'], 'reported_aggregate')
        self.assertEqual(row['status'], 'supported')


if __name__ == '__main__':
    unittest.main()
