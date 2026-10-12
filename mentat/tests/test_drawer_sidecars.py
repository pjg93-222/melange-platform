"""Graph drawer sidecars (perf.md H8): the drawer's cold reads kept on disk, its answers unchanged.

The snapshot index must describe exactly the units json.loads reads. The lazy inventory must select,
order and label them exactly as the old reader did, and raise exactly its errors. A changed input is
never served an old sidecar; a damaged sidecar is rebuilt, never trusted; two processes build one
sidecar once; a stale lock is broken and a live one waited for; the drawer's answer is byte for byte
the old reader's. Every sidecar here is written under a temporary state root.
"""
import hashlib
import json
import os
from pathlib import Path
import random
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

import openpyxl

APP = Path(__file__).resolve().parents[1] / 'app'
sys.path.insert(0, str(APP))
import drawer_sidecars  # noqa: E402
import downstream  # noqa: E402
import graph_data  # noqa: E402
import native_scan  # noqa: E402
from report_insights import stamp  # noqa: E402
from graph_fixture import add_graph_sheets  # noqa: E402

NO_WINDOW = getattr(subprocess, 'CREATE_NO_WINDOW', 0)
ACTIVITY_REQUEST = [{'name': 'Downstream Activity', 'header': 5,
                     'columns': ['cycle', 'source_committee_id', 'candidate_id', 'candidate_officeholder', 'candidate_party',
                                 'support_oppose', 'net_included_amount', 'unverified_version_rows',
                                 'unresolved_version_rows', 'transaction_type']}]


def saved_unit(committee, cycle=2026, coverage='2026-06-30', spend=0, destinations=(), receipts=1, name=None):
    return {'committee_id': committee, 'cycle': cycle,
            'receipts': [{'sub_id': f'r{committee}{i}', 'contribution_receipt_amount': 10 + i,
                          'committee': {'name': name or f'{committee} PAC'},
                          'memo_text': 'a "quoted" [bracket] {brace} \\ backslash, comma: colon'} for i in range(receipts)],
            'disbursements': [{'sub_id': f'd{committee}{d}', 'line_number': '22', 'recipient_committee_id': d,
                               'disbursement_amount': a, 'disbursement_date': '2026-02-01'} for d, a in destinations],
            'expenditures': [{'sub_id': f'e{committee}', 'candidate_id': 'H0KY06104', 'candidate_name': 'BARR, ANDY',
                              'expenditure_amount': spend, 'support_oppose_indicator': 'S',
                              'expenditure_date': '2026-03-01'}] if spend else [],
            'totals': {'disbursements': sum(a for _, a in destinations) + spend, 'coverage_start_date': '2025-01-01',
                       'coverage_end_date': coverage}}


def outcome(function):
    """('ok', dict) or ('error', type, message): how the old and new readers are compared."""
    try:
        value = function()
    except Exception as error:  # noqa: BLE001
        return ('error', type(error).__name__, str(error))
    return ('ok', dict(value))


def fixture_workbook(path):
    book = openpyxl.Workbook()
    add_graph_sheets(book)
    book.save(path)
    book.close()
    return path


def run_child(arguments, timeout=90):
    """A child with its streams on real files, waited on with a timeout (brain DE-026)."""
    with tempfile.TemporaryFile() as out, tempfile.TemporaryFile() as err:
        child = subprocess.Popen(arguments, stdin=subprocess.DEVNULL, stdout=out, stderr=err, creationflags=NO_WINDOW)
        try:
            child.wait(timeout=timeout)
        finally:
            if child.poll() is None:
                child.kill()
                child.wait(timeout=15)
        out.seek(0)
        err.seek(0)
        return child.returncode, out.read().decode('utf-8', 'replace'), err.read().decode('utf-8', 'replace')


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.home = Path(self.tmp.name)
        self.inputs = self.home / 'inputs'
        self.snapshots = self.inputs / 'pac-evidence'
        self.snapshots.mkdir(parents=True)
        self.state = self.home / 'state'
        previous = drawer_sidecars.configure(state_root=self.state, persist_under=[self.inputs], emit=False)
        self.addCleanup(lambda: drawer_sidecars.configure(**previous))
        downstream.cache_clear()
        self.addCleanup(downstream.cache_clear)

    def write(self, folder, units, **dump):
        path = self.snapshots / folder / 'snapshots.json'
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(json.dumps(units, **dump).encode('utf-8'))
        return path

    def stamps(self):
        files = sorted(self.snapshots.glob('*/snapshots.json'))
        return tuple((str(p), p.stat().st_mtime_ns, p.stat().st_size) for p in files)

    def old(self):
        downstream._read_inventory.cache_clear()
        return downstream._read_inventory(self.stamps())

    def kept(self, folder, pattern='*.json'):
        return sorted((self.state / folder).glob(pattern))


class SnapshotIndexTests(Base):
    def test_index_units_are_exactly_what_json_loads_reads(self):
        rng = random.Random(20260924)
        units = [saved_unit(f'C{n:08d}', cycle=rng.choice([2024, 2026, '2026']),
                            spend=rng.choice([0, 7, 12345678901234567890]),
                            destinations=[(f'C{rng.randrange(10 ** 8):08d}', rng.choice([1, 2.5, 1e-7, -0.0, 3.14159e+300]))],
                            receipts=rng.randrange(0, 4)) for n in range(12)]
        units[3]['totals']['coverage_end_date'] = None
        units[4]['notes'] = {'nested': [[], {}, [1, [2, [3]]], 'tab\tnewline\n', '\u00e9\u6771\U0001f600']}
        for style in ({}, {'indent': 1}, {'indent': '\t', 'separators': (', ', ' : ')}):
            path = self.write('style', units, **style)
            raw = path.read_bytes()
            expected = json.loads(raw)
            for chunk in (1, 3, 64, 4096, 1 << 20):
                with self.subTest(style=style, chunk=chunk), patch.object(drawer_sidecars, 'CHUNK', chunk):
                    info = path.stat()
                    fingerprint, index = drawer_sidecars._build_index(path, info.st_mtime_ns, info.st_size)
                    self.assertEqual(fingerprint, hashlib.sha256(raw).hexdigest())
                    self.assertEqual(len(index), len(expected))
                    for entry, unit in zip(index, expected):
                        piece = raw[entry['offset']:entry['offset'] + entry['length']]
                        self.assertEqual(json.loads(piece), unit)
                        self.assertEqual(entry['sha256'], hashlib.sha256(piece).hexdigest())
                        self.assertEqual((entry['cycle'], entry['committee_id']), (unit['cycle'], unit['committee_id']))
                        self.assertEqual(entry['totals'], {'disbursements': True,
                                                           'coverage_end_date': unit['totals']['coverage_end_date']})
                        self.assertEqual(entry['kinds'], {k: drawer_sidecars._kind(v) for k, v in unit.items()})

    def test_every_unusual_file_answers_or_raises_exactly_as_the_old_reader(self):
        good = saved_unit('C00000001', spend=7)
        text = json.dumps([good]).encode()
        cases = {
            'bom': b'\xef\xbb\xbf' + text,
            'raw utf-8': json.dumps([dict(good, note='Zo\u00eb')], ensure_ascii=False).encode('utf-8'),
            'trailing comma': text[:-1] + b',]',
            'data after the array': text + b' []',
            'an object': json.dumps({'units': [good]}).encode(),
            'a scalar unit': json.dumps([good, 5]).encode(),
            'a repeated field': text.replace(b'"cycle": 2026', b'"cycle": 2024, "cycle": 2026'),
            'an equal repeat': json.dumps([good, good]).encode(),
            'a conflicting repeat': json.dumps([good, dict(good, totals={**good['totals'], 'disbursements': 1})]).encode(),
            'a missing schedule': json.dumps([{k: v for k, v in good.items() if k != 'expenditures'}]).encode(),
            'a schedule that is not a list': json.dumps([dict(good, receipts={})]).encode(),
            'totals without disbursements': json.dumps([dict(good, totals={'coverage_end_date': '2026-06-30'})]).encode(),
            'no cycle': json.dumps([{k: v for k, v in good.items() if k != 'cycle'}]).encode(),
            'a numeric committee': json.dumps([dict(good, committee_id=7)]).encode(),
            'truncated': text[:-5],
            'empty': b'',
            'an empty list': b'[]',
            'whitespace everywhere': b' \r\n\t[ \n' + json.dumps(good, indent=2).encode() + b'\n ] \n',
        }
        for name, data in cases.items():
            with self.subTest(case=name):
                for old in self.snapshots.glob('*/snapshots.json'):
                    old.unlink()
                path = self.snapshots / 'case' / 'snapshots.json'
                path.parent.mkdir(exist_ok=True)
                path.write_bytes(data)
                downstream.cache_clear()
                new = outcome(lambda: downstream.inventory(self.snapshots))
                downstream.cache_clear()
                self.assertEqual(new, outcome(lambda: downstream._read_inventory(self.stamps())))

    def test_selection_order_and_labels_match_the_old_reader(self):
        first = self.write('b-first', [saved_unit('C00000001', coverage='2026-03-31', spend=2),
                                       saved_unit('C00000003', coverage='2026-06-30', spend=4),
                                       saved_unit('C00000002', spend=3)])
        second = self.write('a-second', [saved_unit('C00000001', coverage='2026-06-30', spend=7),
                                         saved_unit('C00000003', coverage='2026-06-30', spend=5)])
        # C00000003 ties on coverage; the newer file wins although its path sorts first.
        now = time.time_ns()
        os.utime(first, ns=(now, now - 60 * 10 ** 9))
        os.utime(second, ns=(now, now))
        new = downstream.inventory(self.snapshots)
        self.assertIsInstance(new, downstream.SavedUnits)
        old = self.old()
        self.assertEqual(len(new), len(old))
        self.assertIsNone(new.get(('2026', 'C00000009')))
        self.assertIn(('2026', 'C00000001'), new)
        self.assertEqual(new[('2026', 'C00000003')]['expenditures'][0]['expenditure_amount'], 5)
        self.assertEqual(new[('2026', 'C00000001')]['expenditures'][0]['expenditure_amount'], 7)
        self.assertEqual(list(new), list(old))
        self.assertEqual(dict(new), old)

    def test_a_traversal_reads_only_the_units_it_reaches(self):
        self.write('one', [saved_unit('C00000001', destinations=[('C00000002', 5)]), saved_unit('C00000002', spend=5),
                           saved_unit('C00000009', spend=9), saved_unit('C00000008', cycle=2024, spend=1)])
        reads = []
        original = drawer_sidecars.read_unit

        def counting(*arguments):
            reads.append(arguments[3])
            return original(*arguments)
        with patch.object(downstream, 'ROOT', self.snapshots), patch.object(drawer_sidecars, 'read_unit', counting):
            found = downstream.network([('2026', 'C00000001')])
            rows = downstream.activity(found)
            ties = downstream.relationships(rows, [('2026', 'C00000001')], found)
        self.assertEqual(len(reads), 2)
        self.assertEqual(sorted(found['units']), [('2026', 'C00000001'), ('2026', 'C00000002')])
        before = downstream.network([('2026', 'C00000001')], self.old())
        old_rows = downstream.activity(before)
        self.assertEqual(json.dumps([rows, ties, downstream.sources(found)], default=str),
                         json.dumps([old_rows, downstream.relationships(old_rows, [('2026', 'C00000001')], before),
                                     downstream.sources(before)], default=str))

    def test_inputs_outside_the_kept_roots_are_never_written(self):
        elsewhere = self.home / 'elsewhere'
        (elsewhere / 'x').mkdir(parents=True)
        (elsewhere / 'x' / 'snapshots.json').write_text(json.dumps([saved_unit('C00000001', spend=7)]))
        self.assertEqual(dict(downstream.inventory(elsewhere)),
                         downstream._read_inventory(tuple((str(p), p.stat().st_mtime_ns, p.stat().st_size)
                                                          for p in sorted(elsewhere.glob('*/snapshots.json')))))
        self.write('kept', [saved_unit('C00000001', spend=7)])
        with patch.dict(os.environ, MENTAT_DRAWER_SIDECARS='0'):
            self.assertEqual(dict(downstream.inventory(self.snapshots)), self.old())
        self.assertFalse(self.state.exists())


class KeptIndexTests(Base):
    def test_index_is_kept_and_read_back_without_walking(self):
        self.write('kept', [saved_unit('C00000001', spend=7), saved_unit('C00000002', spend=8)])
        first = dict(downstream.inventory(self.snapshots))
        self.assertEqual(len(self.kept('snapshot-index')), 1)
        downstream.cache_clear()   # a restarted console
        with patch.object(drawer_sidecars, '_walk', side_effect=AssertionError('the file was walked again')):
            lazy = downstream.inventory(self.snapshots)
            self.assertIsInstance(lazy, downstream.SavedUnits)   # the index, not the old reader's fallback
            again = dict(lazy)
        self.assertEqual(again, first)
        self.assertEqual(again, self.old())

    def test_changed_content_under_the_same_size_and_mtime_rebuilds_the_index(self):
        path = self.write('change', [saved_unit('C00000001', spend=7)])
        info = path.stat()
        self.assertEqual(downstream.inventory(self.snapshots)[('2026', 'C00000001')]['expenditures'][0]['expenditure_amount'], 7)
        path.write_bytes(path.read_bytes().replace(b'"expenditure_amount": 7', b'"expenditure_amount": 8'))
        os.utime(path, ns=(info.st_atime_ns, info.st_mtime_ns))
        self.assertEqual((path.stat().st_size, path.stat().st_mtime_ns), (info.st_size, info.st_mtime_ns))
        downstream.cache_clear()   # a restarted console
        drawer_sidecars.forget(contents=True)
        lazy = downstream.inventory(self.snapshots)
        self.assertIsInstance(lazy, downstream.SavedUnits)
        new = dict(lazy)
        self.assertEqual(new[('2026', 'C00000001')]['expenditures'][0]['expenditure_amount'], 8)
        self.assertEqual(new, self.old())
        self.assertEqual(len(self.kept('snapshot-index')), 1)   # the old version is gone

    def test_damaged_index_is_rebuilt_not_trusted(self):
        self.write('damaged', [saved_unit('C00000001', spend=7), saved_unit('C00000002', spend=9)])
        expected = self.old()
        dict(downstream.inventory(self.snapshots))
        kept = self.kept('snapshot-index')[0]
        original = kept.read_bytes()
        header, body = original.split(b'\n', 1)
        self.assertIn(b'"offset":1,', body)
        damages = {'a body byte': header + b'\n' + body.replace(b'"offset":1,', b'"offset":2,', 1),
                   'truncated': original[:-10], 'a broken header': b'{not json\n' + body, 'empty': b''}
        for name, data in damages.items():
            with self.subTest(damage=name):
                kept.write_bytes(data)
                downstream.cache_clear()
                self.assertEqual(dict(downstream.inventory(self.snapshots)), expected)
                self.assertEqual(kept.read_bytes().split(b'\n', 1)[1], body)


@unittest.skipUnless(native_scan.executable(), 'Build with tools/build_native.ps1 first')
class NativeSidecarTests(Base):
    def setUp(self):
        super().setUp()
        self.addCleanup(native_scan.shutdown)
        self.payload = self.inputs / 'industry_data-0123456789abcdef.json'

    def write_payload(self, amount):
        edge = dict(cycle=2024, source='C00000001', destination='C00000002', amount=amount, status='expanded')
        self.payload.write_text(json.dumps({'metadata': {}, 'ignored': [{'large': 'x' * 1000}],
                                            'transfer_paths': [edge, dict(edge, destination='C00000003', amount='5.00'),
                                                               dict(edge, destination='C00000004', status='candidate_committee')]}),
                                encoding='utf-8')
        return stamp(self.payload)

    def fresh(self, version, requests=(), payload_table='transfer_paths'):
        with patch.dict(os.environ, MENTAT_DRAWER_SIDECARS='0'):
            return native_scan.tables(version, list(requests), payload_table=payload_table)

    def test_payload_sidecar_is_built_once_served_after_and_keyed_by_content(self):
        version = self.write_payload('125.25')
        expected = self.fresh(version)
        self.assertEqual(native_scan.tables(version, [], payload_table='transfer_paths'), expected)
        self.assertEqual(len(self.kept('native', 'payload-*.json')), 1)
        with patch.object(native_scan, '_scan', side_effect=AssertionError('scanned again')):
            self.assertEqual(native_scan.tables(version, [], payload_table='transfer_paths'), expected)
        changed = self.write_payload('125.35')   # the same size: only the content says it changed
        self.assertEqual(self.payload.stat().st_size, version[3])
        self.assertNotEqual(self.fresh(changed), expected)
        self.assertEqual(native_scan.tables(changed, [], payload_table='transfer_paths'), self.fresh(changed))
        self.assertEqual(len(self.kept('native', 'payload-*.json')), 1)   # the old version is gone

    def test_damaged_payload_sidecar_is_rescanned_not_served(self):
        version = self.write_payload('125.25')
        expected = self.fresh(version)
        native_scan.tables(version, [], payload_table='transfer_paths')
        kept = self.kept('native', 'payload-*.json')[0]
        original = kept.read_bytes()
        header, body = original.split(b'\n', 1)
        self.assertIn(b'125.25', body)
        for data in (header + b'\n' + body.replace(b'125.25', b'125.35'), original[:len(original) // 2], b''):
            with self.subTest(size=len(data)):
                kept.write_bytes(data)
                scans = []
                real = native_scan._scan

                def counting(*arguments):
                    scans.append(1)
                    return real(*arguments)
                with patch.object(native_scan, '_scan', counting):
                    self.assertEqual(native_scan.tables(version, [], payload_table='transfer_paths'), expected)
                self.assertEqual(scans, [1])
                self.assertEqual(kept.read_bytes().split(b'\n', 1)[1], body)

    def test_downstream_activity_scan_is_kept_and_other_scans_are_not(self):
        workbook = fixture_workbook(self.inputs / 'Fixture Industry Spending Overview.xlsx')
        version = stamp(workbook)
        expected = self.fresh(version, ACTIVITY_REQUEST, None)
        self.assertEqual(native_scan.tables(version, ACTIVITY_REQUEST), expected)
        self.assertEqual(len(self.kept('native', 'workbook-*.json')), 1)
        with patch.object(native_scan, '_scan', side_effect=AssertionError('scanned again')):
            self.assertEqual(native_scan.tables(version, ACTIVITY_REQUEST), expected)
        native_scan.tables(version, [{'name': 'PAC & Candidate Map', 'header': 5, 'columns': ['cycle', 'committee_id']}])
        native_scan.tables(version, [])
        # The evidence maps' projection with every filing record, and an all-column read, are not kept.
        native_scan.tables(version, [{**ACTIVITY_REQUEST[0], 'columns': [*ACTIVITY_REQUEST[0]['columns'], 'filing_records']}])
        native_scan.tables(version, [{'name': 'Downstream Activity', 'header': 5}])
        self.assertEqual(len(self.kept('native')), 1)
        graph_data.report_activity.cache_clear()
        through_sidecar = graph_data.report_activity(version)
        graph_data.report_activity.cache_clear()
        with patch.dict(os.environ, MENTAT_DRAWER_SIDECARS='0'):
            self.assertEqual(graph_data.report_activity(version), through_sidecar)
        graph_data.report_activity.cache_clear()

    def test_scan_errors_pass_through_unchanged_and_bookkeeping_failures_fall_back(self):
        self.payload.write_text(json.dumps({'metadata': {}, 'other': []}), encoding='utf-8')
        version = stamp(self.payload)
        with patch.dict(os.environ, MENTAT_DRAWER_SIDECARS='0'), self.assertRaises(ValueError) as unkept:
            native_scan.tables(version, [], payload_table='transfer_paths')
        with self.assertRaises(ValueError) as kept:
            native_scan.tables(version, [], payload_table='transfer_paths')
        self.assertIs(type(kept.exception), type(unkept.exception))
        self.assertEqual(str(kept.exception), str(unkept.exception))
        self.assertEqual(self.kept('native'), [])   # an error is never kept
        good = self.write_payload('125.25')
        with patch.object(drawer_sidecars, '_native_result', side_effect=RuntimeError('bookkeeping broke')):
            self.assertEqual(native_scan.tables(good, [], payload_table='transfer_paths'), self.fresh(good))


CHILD = r'''
import json, os, sys, time
sys.path.insert(0, sys.argv[1])
import drawer_sidecars
drawer_sidecars.configure(state_root=sys.argv[2], emit=False)
log, start = sys.argv[3], float(sys.argv[4])
target = drawer_sidecars.state_root() / 'race.json'
def build():
    with open(log, 'a') as handle:
        handle.write(f'{os.getpid()}\n')
    time.sleep(1.5)
    return {'value': 42}
def load():
    try:
        return json.loads(target.read_text())
    except (OSError, ValueError):
        return None
def store(value):
    temporary = target.with_name(f'race.{os.getpid()}.tmp')
    temporary.write_text(json.dumps(value))
    os.replace(temporary, target)
    return 1
while time.time() < start:
    time.sleep(0.005)
value, how = drawer_sidecars._single_flight('race', load, build, store)
print(json.dumps({'value': value, 'how': how}))
'''


class LockTests(Base):
    def test_two_processes_build_one_sidecar_once(self):
        log = self.home / 'builds.log'
        start = time.time() + 1.5
        here = str(Path(drawer_sidecars.__file__).resolve().parent)
        results = [None, None]

        def launch(slot):
            results[slot] = run_child([sys.executable, '-c', CHILD, here, str(self.state), str(log), str(start)])
        threads = [threading.Thread(target=launch, args=(slot,)) for slot in range(2)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(120)
        values = []
        for code, out, err in results:
            self.assertEqual(code, 0, err)
            values.append(json.loads(out))
        self.assertEqual([v['value'] for v in values], [{'value': 42}] * 2)
        self.assertEqual(sorted(v['how'] for v in values), ['built', 'hit'])
        self.assertEqual(len(log.read_text().splitlines()), 1)

    def test_threads_share_one_build(self):
        builds, box, results = [], {}, []

        def build():
            builds.append(threading.get_ident())
            time.sleep(0.5)
            return {'n': 1}

        def store(value):
            box['value'] = value
            return 1
        threads = [threading.Thread(target=lambda: results.append(
            drawer_sidecars._single_flight('threads', lambda: box.get('value'), build, store))) for _ in range(8)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(60)
        self.assertEqual(len(builds), 1)
        self.assertEqual([value for value, _ in results], [{'n': 1}] * 8)

    def test_a_dead_left_behind_or_recycled_lock_is_broken_and_a_live_one_is_waited_for_at_any_age(self):
        path = drawer_sidecars._lock_path('stale')
        path.parent.mkdir(parents=True, exist_ok=True)
        gone = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(0.5)'], creationflags=NO_WINDOW)
        gone_started = drawer_sidecars.process_started(gone.pid)
        gone.wait(timeout=30)
        me = drawer_sidecars.process_started(os.getpid())
        builds = []

        def build():
            builds.append(1)
            return {'ok': True}
        with patch.object(drawer_sidecars, 'LOCK_WAIT_SECONDS', 3):
            for name, info in {'dead': {'pid': gone.pid, 'started': gone_started, 'created': time.time()},
                               'left behind by this process': {'pid': os.getpid(), 'started': me,
                                                               'created': time.time() - 31 * 60},
                               'recycled': {'pid': os.getpid(), 'started': 12345, 'created': time.time()}}.items():
                with self.subTest(lock=name):
                    path.write_text(json.dumps({**info, 'token': name}))
                    self.assertEqual(drawer_sidecars._single_flight('stale', lambda: None, build, lambda value: 1),
                                     ({'ok': True}, 'built'))
                    self.assertFalse(path.exists())
            # A live holder in another process, its lock older than 30 minutes (review H2): waited for.
            holder = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'], creationflags=NO_WINDOW)
            self.addCleanup(lambda: (holder.kill(), holder.wait(timeout=15)))
            deadline = time.monotonic() + 30
            while not drawer_sidecars.process_started(holder.pid) and time.monotonic() < deadline:
                time.sleep(0.02)
            path.write_text(json.dumps({'pid': holder.pid, 'started': drawer_sidecars.process_started(holder.pid),
                                        'created': time.time() - 31 * 60, 'token': 'live'}))
            moment = time.time() - 31 * 60
            os.utime(path, (moment, moment))
            box = {}

            def finish():
                box['value'] = {'from': 'the other builder'}
                path.unlink()
            timer = threading.Timer(1.0, finish)
            timer.start()
            began = time.monotonic()
            self.assertEqual(drawer_sidecars._single_flight('stale', lambda: box.get('value'), build, lambda value: 1),
                             ({'from': 'the other builder'}, 'hit'))
            self.assertGreaterEqual(time.monotonic() - began, 0.9)
            timer.join()
        self.assertEqual(len(builds), 3)

    def test_the_liveness_probe_tells_running_from_exited_without_signalling(self):
        child = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'], creationflags=NO_WINDOW)
        try:
            self.assertIsInstance(drawer_sidecars.process_started(child.pid), int)
            time.sleep(0.3)
            self.assertIsNone(child.poll())   # still running: the probe did not signal it
        finally:
            child.kill()
            child.wait(timeout=15)
        self.assertIsNone(drawer_sidecars.process_started(child.pid))


class DrawerAnswerTests(Base):
    def test_drawer_answers_are_byte_identical_to_the_old_reader(self):
        workbook = fixture_workbook(self.inputs / 'Fixture Industry Spending Overview.xlsx')
        self.write('saved', [saved_unit('C00000003', cycle='2026', destinations=[('C00000007', 30)], spend=11,
                                        name='Members First Super PAC'),
                             saved_unit('C00000007', cycle='2026', spend=30),
                             saved_unit('C00000002', cycle=2024, spend=1)])
        requests = [{'mode': 'activity', 'recipient': 'C00000003', 'cycle': '2026', 'direction': direction}
                    for direction in ('Support', 'Oppose', 'Contribution / Other')]
        requests += [{'mode': 'activity'}, {'mode': 'activity', 'recipient': 'C00000002'},
                     {'mode': 'activity', 'company': 'Harborlight Credit Union', 'direction': 'Oppose'}]
        with patch.object(downstream, 'ROOT', self.snapshots):
            for filters in requests:
                with self.subTest(filters=filters):
                    graph_data._views.clear()
                    downstream.cache_clear()
                    reads, real = [], drawer_sidecars.read_unit
                    with patch.object(drawer_sidecars, 'read_unit', lambda *a: reads.append(a) or real(*a)):
                        new = graph_data.graph_view([workbook], dict(filters))
                    graph_data._views.clear()
                    downstream.cache_clear()
                    with patch.object(downstream, '_selection', return_value=None):
                        old = graph_data.graph_view([workbook], dict(filters))
                    encode = lambda value: json.dumps(value, ensure_ascii=False, separators=(',', ':')).encode()
                    self.assertEqual(encode(new), encode(old))
                    if filters.get('recipient') == 'C00000003':
                        self.assertEqual(len(new['downstream_sources']), 2)
                        self.assertEqual(len(reads), 2)   # the index answered, not the old reader
        graph_data._views.clear()

    def test_build_sidecars_reports_each_step(self):
        workbook = fixture_workbook(self.inputs / 'Fixture Industry Spending Overview.xlsx')
        self.write('saved', [saved_unit('C00000003', spend=1)])
        with patch.object(downstream, 'ROOT', self.snapshots):
            summary = drawer_sidecars.build_sidecars(str(workbook))
        json.dumps(summary)
        self.assertTrue(summary['ok'], summary)
        self.assertEqual([s['step'] for s in summary['steps']], ['snapshot indexes', 'downstream activity', 'transfer paths'])
        self.assertEqual(summary['steps'][0]['result'], {'files': 1, 'indexed': True, 'units': 1})
        self.assertEqual(len(self.kept('snapshot-index')), 1)
        if native_scan.executable():
            self.assertEqual(len(self.kept('native', 'workbook-*.json')), 1)
        self.assertEqual(graph_data.report_activity.cache_info().currsize, 0)


if __name__ == '__main__':
    unittest.main()
