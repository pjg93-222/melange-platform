"""Offline upload, launch isolation, telemetry, and HTTP boundary checks."""
import base64
import importlib.util
import json
import os
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch, MagicMock
from urllib.request import Request, urlopen
from urllib.error import HTTPError

APP = Path(__file__).resolve().parents[1] / "app"
spec = importlib.util.spec_from_file_location("collection_console", APP / "server.py")
app = importlib.util.module_from_spec(spec)
spec.loader.exec_module(app)


class ConsoleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        report_patch = patch.object(app,'report_directory',return_value=self.root)
        report_patch.start()
        self.addCleanup(report_patch.stop)
        self.key_patch = patch('fec_key_store.key_path', return_value=self.root / 'test-keys.dpapi')
        self.key_patch.start()
        self.addCleanup(self.key_patch.stop)
        self.environment_patch = patch.object(app, 'environment_keys', return_value=[])
        self.environment_patch.start()
        self.addCleanup(self.environment_patch.stop)
        original_read = app.read_json
        source_patch = patch.object(app, 'read_json', side_effect=lambda path:
            {} if Path(path).name == 'active_fec_sources.json' else original_read(path))
        source_patch.start()
        self.addCleanup(source_patch.stop)
        self.console = app.Console(self.root)

    def tearDown(self):
        self.temp.cleanup()

    def upload(self, text="canonical_name,alias,industry\nAcme,Acme Inc,Test\nAcme,ACME CORP,Test\nOther,Other LLC,Test\n"):
        return self.console.upload({"filename":"test.csv", "content":base64.b64encode(text.encode()).decode()})

    def test_repeated_canonical_rows_group_aliases(self):
        value = self.upload()
        self.assertEqual(value["state"], "ready")
        industry = value["readiness"]["readiness"][0]
        self.assertEqual(industry["company_count"], 2)
        self.assertEqual(industry["companies"][0]["variant_count"], 3)

    def test_completed_stage_with_failed_industry_is_failed(self):
        uploaded = self.upload()
        directory = self.root / uploaded['id'] / 'runs' / 'sample'
        directory.mkdir(parents=True)
        app.atomic_json(directory / 'manifest.json', {'run_id':'sample'})
        app.atomic_json(directory / 'status.json', {'stage':'completed', 'industries':[{'industry':'Test','status':'failed'}]})
        self.assertEqual(self.console.find(uploaded['id'])['state'], 'failed')
        self.assertEqual(self.console.report_catalog(uploaded['id'])['reports'], [])

    def test_evidence_collection_validates_and_preserves_configuration(self):
        data = {'company':'Acme', 'aliases':['Acme Inc'], 'committees':'C00000001, C00000002', 'cycles':'2024,2026'}
        with patch.object(app, 'PROJECT', self.root), patch.object(app, 'load_keys', return_value=['synthetic']), patch.object(app.threading, 'Thread') as thread:
            for change in ({'committees':'../bad'}, {'cycles':'2025'}, {'company':''}):
                with self.assertRaises(ValueError):
                    self.console.collect_evidence({**data, **change})
            thread.assert_not_called()
            result = self.console.collect_evidence(data)
            path = self.root / '.cache/pac-evidence' / result['id'] / 'collection.json'
            saved = app.read_json(path)
            self.assertEqual(saved['aliases'], ['Acme Inc'])
            self.assertEqual(saved['cycles'], ['2024', '2026'])
            with patch.object(app.subprocess, 'Popen') as popen:
                popen.return_value.wait.return_value = 0
                popen.return_value.pid = 123
                app.atomic_json(path.parent / 'evidence.json', {'version':1})
                thread.call_args.kwargs['target']()
                self.assertEqual(app.read_json(path)['state'], 'completed')
                self.assertNotIn('synthetic', str(popen.call_args))
                (path.parent / 'evidence.json').unlink()
                thread.call_args.kwargs['target']()
                self.assertEqual(app.read_json(path)['state'], 'failed')

    def test_manual_sector_roster_persists_names_and_commas_without_collecting(self):
        value = self.console.create_roster({'sector': 'Energy', 'companies': [
            {'name': 'Acme', 'aliases': ['Acme, Inc.', 'ACME ENERGY', 'ACME ENERGY']},
            {'name': 'Other', 'aliases': []}]})
        self.assertEqual(value['state'], 'ready')
        self.assertEqual(self.console.processes, {})
        roster = value['readiness']['readiness'][0]
        self.assertEqual(roster['industry'], 'Energy')
        self.assertEqual(roster['company_count'], 2)
        self.assertEqual(roster['companies'][0]['variant_count'], 3)
        self.assertIn('Acme, Inc.', [v['variant'] for v in roster['companies'][0]['variants']])
        self.assertEqual(app.read_json(self.root / value['id'] / 'upload.json'), value)

    def test_manual_alias_conflicts_and_bad_input_cannot_start(self):
        value = self.console.create_roster({'sector': 'Energy', 'companies': [
            {'name': 'One', 'aliases': ['SHARED']}, {'name': 'Two', 'aliases': ['SHARED']}]})
        self.assertEqual(value['state'], 'needs review')
        with self.assertRaises(ValueError):
            self.console.launch({'id': value['id'], 'cycles': [2026]})
        for data in ({'sector': '', 'companies': []}, {'sector': 'Energy', 'companies': 'oops'},
                     {'sector': 'Energy', 'companies': [{'name': 'Acme'}, {'name': ' ACME '}]},
                     {'sector': 'Energy', 'companies': [{'name': 'Acme', 'aliases': 'wrong'}]}):
            with self.assertRaises(ValueError): self.console.create_roster(data)

    def test_sector_upload_header_is_supported_and_conflicts_are_rejected(self):
        value = self.upload('canonical_name,alias,sector\nAcme,Acme Inc,Energy\n')
        self.assertEqual(value['readiness']['readiness'][0]['industry'], 'Energy')
        with self.assertRaisesRegex(ValueError, 'must agree'):
            self.upload('canonical_name,sector,industry\nAcme,Energy,Crypto\n')

    def test_cross_company_alias_collision_is_not_ready(self):
        value = self.upload("canonical_name,alias\nOne,SHARED\nTwo,SHARED\n")
        self.assertEqual(value["state"], "needs review")
        with self.assertRaises(ValueError):
            self.console.launch({"id":value["id"],"cycles":[2026]})

    def test_crypto_keeps_one_company_per_row(self):
        value = self.upload("canonical_name,alias,industry\nOne,One LLC,Crypto\nOne,One Inc,Crypto\nTwo,Two Inc,Crypto\n")
        self.assertEqual(value["readiness"]["readiness"][0]["company_count"], 2)

    def test_empty_name_rejected(self):
        with self.assertRaisesRegex(ValueError, "canonical_name is required"):
            self.upload("canonical_name,alias\n,hello\n")

    def test_existing_workbook_roundtrip_preserves_roster(self):
        first = self.upload()
        content = (self.root / first["id"] / "roster.xlsx").read_bytes()
        second = self.console.upload({"filename":"roster.xlsx","content":base64.b64encode(content).decode()})
        self.assertEqual(first["readiness"]["readiness"], second["readiness"]["readiness"])

    def test_launch_pins_uploaded_roster_and_keeps_key_out_of_arguments_and_disk(self):
        value = self.upload()
        runner = MagicMock()
        runner.pid = 987654
        runner.poll.return_value = None
        # The FEC bulk snapshots are not kept in git, and this test is about key
        # handling and the pinned roster rather than about staging them.
        with patch.object(app.subprocess,"Popen", return_value=runner) as launch, \
                patch.object(app.threading,"Thread"), patch.object(app, "copy_fec_sources"):
            self.console.launch({"id":value["id"],"cycles":[2026],"api_key":"private-test-key"})
        command = launch.call_args.args[0]
        self.assertNotIn("private-test-key", repr(command))
        self.assertEqual(launch.call_args.kwargs["env"]["FEC_API_KEYS"], "private-test-key")
        self.assertNotIn(b'private-test-key', (self.root / 'test-keys.dpapi').read_bytes())
        self.assertEqual(app.load_keys(), ['private-test-key'])
        directory = self.root / value["id"]
        self.assertNotIn("private-test-key", (directory / "upload.json").read_text())
        runtime = directory / "runtime/mentat"
        self.assertEqual((runtime / "inputs/Company Name.xlsx").read_bytes(), (directory / "roster.xlsx").read_bytes())
        self.assertEqual(Path(launch.call_args.kwargs["env"]["MENTAT_INPUT_DIR"]), runtime / "inputs")
        with self.assertRaises(ValueError):
            self.console.launch({"id":value["id"],"cycles":[2026],"api_key":"private-test-key"})

    def test_invalid_cycle_and_identifier_cannot_launch(self):
        with self.assertRaises(ValueError):
            self.console.launch({"id":"../outside", "cycles":[2026]})
        value = self.upload()
        with self.assertRaises(ValueError):
            self.console.launch({"id":value["id"], "cycles":[2028]})

    def test_background_launch_reserves_run_before_copying(self):
        value = self.upload()
        with patch.object(app.threading, 'Thread') as worker, patch.object(app.subprocess, 'Popen') as process:
            result = self.console.launch({'id':value['id'], 'cycles':[2026], 'api_key':'offline-key'}, background=True)
            self.assertEqual(result['state'], 'preparing')
            self.assertEqual(app.read_json(self.root/value['id']/'upload.json')['state'], 'preparing')
            self.assertIn(value['id'], self.console.preparing)
            process.assert_not_called()
            worker.return_value.start.assert_called_once()
            with self.assertRaises(ValueError):self.console.launch({'id':value['id'], 'cycles':[2026]})

    def test_preparation_failure_has_terminal_receipt(self):
        value = self.upload()
        self.console.preparing.add(value['id'])
        with patch.object(self.console, 'start_prepared_process', side_effect=OSError('private diagnostic')):
            self.console.prepare_launch(value['id'], self.root/value['id'], [2026], ['offline-key'], True)
        manifest = app.read_json(self.root/value['id']/'upload.json')
        self.assertEqual(manifest['state'], 'launch failed')
        self.assertIsInstance(manifest['finished_at'], float)
        self.assertNotIn(value['id'], self.console.preparing)
        self.assertNotIn('private diagnostic', str(manifest))

    def test_finish_preserves_stop_state_and_frozen_timestamp_without_keys(self):
        value = self.upload()
        directory = self.root/value['id']
        manifest = app.read_json(directory/'upload.json')
        manifest.update(state='stopped', stop_requested='now', finished_at=1234)
        app.atomic_json(directory/'upload.json', manifest)
        process = MagicMock();process.wait.return_value = 1
        with patch.object(app, 'load_keys', side_effect=AssertionError('Finishing must not load credentials')):
            self.console.finish(process, directory)
        result = app.read_json(directory/'upload.json')
        self.assertEqual(result['state'], 'stopped')
        self.assertEqual(result['finished_at'], 1234)

    # Unlike its neighbours this one stages the real project and runs the real R
    # ingest against it, so it wants the bulk snapshots themselves, not a stub.
    @unittest.skipUnless(all((app.PROJECT / "inputs" / name).is_dir()
                             for name in ("fec_beneficiary", "fec_committee_master")),
                         "R ingest stages the FEC bulk sources; install them with tools/refresh_fec_sources.py")
    @unittest.skipUnless(Path("C:/Program Files/R/R-4.6.0/bin/Rscript.exe").is_file(), "R integration requires installed R 4.6")
    def test_runner_plan_and_offline_R_ingest_use_uploaded_aliases(self):
        value = self.upload()
        runner = MagicMock()
        runner.pid = 987654
        with patch.object(app.subprocess, "Popen", return_value=runner) as launch, patch.object(app.threading,"Thread"):
            self.console.launch({"id":value["id"],"cycles":[2026],"api_key":"offline-plan-placeholder"})
        kwargs = launch.call_args.kwargs
        hashing = subprocess.run([launch.call_args.args[0][0], '-NoProfile', '-ExecutionPolicy', 'Bypass', '-Command',
                                  "(Get-FileHash -LiteralPath '" + str(runtime_path := self.root / value['id'] / 'roster.xlsx').replace("'", "''") + "').Algorithm"],
                                 env=kwargs['env'], capture_output=True, text=True, timeout=30)
        self.assertEqual(hashing.returncode, 0, hashing.stderr)
        self.assertEqual(hashing.stdout.strip(), 'SHA256')
        plan = subprocess.run(launch.call_args.args[0] + ["-PlanOnly"], cwd=kwargs["cwd"],
                              env=kwargs["env"], capture_output=True, text=True, timeout=60)
        self.assertEqual(plan.returncode, 0, plan.stdout + plan.stderr)
        self.assertEqual(json.loads(plan.stdout)["discovered_industries"], ["Test"])
        runtime = Path(kwargs["cwd"])
        data_root = self.root / "offline-ingest"
        result = subprocess.run([self.console.rscript, str(runtime / "mentat.R"),
                                 "--industry=Test", "--cycles=2026", "--skip-scrape", "--defer-report-files",
                                 "--data-root=" + str(data_root)], cwd=runtime, env=kwargs["env"],
                                capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        connection = sqlite3.connect(data_root / "mentat.db")
        try:
            companies = connection.execute("SELECT DISTINCT company FROM companies ORDER BY company").fetchall()
            aliases = connection.execute("SELECT variant FROM companies WHERE company='Acme'").fetchall()
            self.assertEqual(companies, [("Acme",), ("Other",)])
            self.assertEqual({row[0] for row in aliases}, {"Acme", "Acme Inc", "ACME CORP"})
        finally:
            connection.close()

    def test_live_read_uses_database_without_modifying_it(self):
        database = self.root / "mentat.db"
        with sqlite3.connect(database) as connection:
            connection.execute("CREATE TABLE api_work_progress(company,cycle,kind,rows_seen,records_total,fetched_complete,stored_complete,updated_at)")
            connection.execute("INSERT INTO api_work_progress VALUES ('Acme',2026,'employer',10,20,0,0,'now')")
            connection.execute("CREATE TABLE contributions(sub_id,contributor_name,company,matched_variant,contributor_employer,committee_name,contribution_receipt_date,contribution_receipt_amount,company_match_status)")
            connection.execute("INSERT INTO contributions VALUES ('123','Sample','Acme','ACME INC','Acme Inc','Committee','2026-01-01',100,'accepted')")
        connection.close()
        before = database.read_bytes()
        view = app.lane_preview(database)
        self.assertTrue(view["available"])
        self.assertEqual(view["progress"][0]["rows_seen"], 10)
        self.assertEqual(view["records"][0]["sub_id"], "123")
        self.assertEqual(database.read_bytes(), before)

    def test_missing_database_has_no_fake_zero(self):
        self.assertEqual(app.lane_preview(self.root / "absent.db"), {"available":False})

    def test_report_inspection_detects_missing_directions_and_changed_file(self):
        import openpyxl
        path = self.root / "report.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Beneficiary Summary"
        sheet.cell(5, 1, "Downstream Activity")
        book.save(path)
        first = self.console.inspect_report(path)
        self.assertFalse(first['direction_columns'])
        sheet.cell(5, 2, "Support Activity")
        sheet.cell(5, 3, "Oppose Activity")
        book.save(path)
        book.close()
        second = self.console.inspect_report(path)
        self.assertTrue(second['direction_columns'])
        self.assertNotEqual(first['sha256'], second['sha256'])

    def test_attributed_report_warnings_use_named_columns(self):
        import openpyxl
        path = self.root / 'attributed.xlsx'
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'Candidate Funding'
        sheet.cell(5, 1, 'Conditional Non-opposition')
        sheet.cell(5, 2, 'Conditional Opposition')
        quality = book.create_sheet('Data Quality')
        for column, value in enumerate(('Industry', 'Check', 'Severity', 'Passed', 'Detail'), 1):
            quality.cell(5, column, value)
        for column, value in enumerate(('Fixture', 'coverage', 'warn', False, 'Unverified source'), 1):
            quality.cell(6, column, value)
        book.save(path)
        book.close()
        result = self.console.inspect_report(path)
        self.assertTrue(result['direction_columns'])
        self.assertEqual(result['warnings'], [{'check': 'coverage', 'severity': 'warn', 'detail': 'Unverified source'}])

    def test_saved_run_reports_failed_coverage_separately_from_completion(self):
        upload = self.upload()
        directory = self.root / upload['id'] / 'runs' / 'test-run'
        directory.mkdir(parents=True)
        coverage = directory / 'coverage.json'
        app.atomic_json(coverage, {'status':'failed','checks':{'company_cycle_grid':False}})
        receipt = directory / 'fanout.json'
        app.atomic_json(receipt, {'task_count':3,'succeeded':2,'failed':1})
        report = self.root / 'report.xlsx'
        import openpyxl
        book = openpyxl.Workbook()
        book.save(report)
        book.close()
        expected = self.console.inspect_report(report)['sha256'].upper()
        app.atomic_json(directory / 'manifest.json', {'run_id':'test-run'})
        app.atomic_json(directory / 'status.json', {'stage':'completed','industries':[
            {'industry':'Test','status':'succeeded','coverage_receipt':str(coverage),
             'fanout_receipt':str(receipt),'output_workbook':str(report),'workbook_sha256':expected}]})
        result = self.console.snapshot(upload['id'])
        self.assertEqual(result['assurance'][0]['collection'],'failed')
        self.assertTrue(result['reports'][0]['quality']['matches_run'])
        self.assertEqual(result['completed_collection'], {'total':3,'completed':3,'failed':1,'running':0})

    def test_chart_report_list_does_not_inspect_workbooks_or_worker_databases(self):
        upload = self.upload()
        directory = self.root / upload['id'] / 'runs' / 'list-run'
        directory.mkdir(parents=True)
        report = directory / 'sample.xlsx'
        report.write_bytes(b'not scanned by the catalog')
        app.atomic_json(directory / 'manifest.json', {'run_id':'list-run'})
        app.atomic_json(directory / 'status.json', {'stage':'collecting', 'industries':[
            {'status':'failed','output_workbook':str(report)},
            {'status':'succeeded','output_workbook':str(report)},
            {'status':'succeeded','output_workbook':str(directory/'missing.xlsx')}]})
        with patch.object(self.console, 'inspect_report', side_effect=AssertionError('must not scan')):
            self.assertEqual(self.console.report_catalog(upload['id'])['reports'], [{'index':1,'name':'sample.xlsx','source_directory':str(self.root)}])

    def test_shared_worker_monitor_reads_current_state_with_stable_lane_identity(self):
        upload = self.upload()
        directory = self.root / upload['id'] / 'runs' / 'shared-run'
        lane = directory / 'industries/test/fanout/lanes/lane-01'
        lane.mkdir(parents=True)
        app.atomic_json(directory / 'manifest.json', {'run_id':'shared-run'})
        app.atomic_json(directory / 'status.json', {'stage':'collecting','industries':[]})
        app.atomic_json(lane / 'status.old.json', {'status':'succeeded','current_company':'Old'})
        path = lane / 'status.work-1-current.json'
        app.atomic_json(path, {'status':'running','current_company':'Current'})
        manifest = {'industry':'Test','scheduling':{'strategy':'company_work_sharing'},
                    'lanes':[{'lane':1,'state':str(path),'database':str(lane/'mentat.db')}]}
        manifest_path = lane.parent.parent / 'manifest.json'
        app.atomic_json(manifest_path, manifest)
        first = self.console.snapshot(upload['id'])
        self.assertEqual(len(first['lanes']), 1)
        self.assertEqual(first['lanes'][0]['state']['current_company'], 'Current')
        # The discovered directory and saved status can name the same physical
        # run through different path spellings (including Windows app aliases).
        (directory / 'alias-segment').mkdir()
        run = self.console.find(upload['id'])
        alias_run = {**run, 'directory':directory / 'alias-segment' / '..'}
        with patch.object(self.console, 'find', return_value=alias_run):
            aliased = self.console.snapshot(upload['id'], first['selected_lane'])
        self.assertEqual(aliased['selected_lane'], first['selected_lane'])
        self.assertEqual(aliased['lanes'], first['lanes'])
        next_path = lane / 'status.work-1-next.json'
        app.atomic_json(next_path, {'status':'running','current_company':'Next'})
        manifest['lanes'][0]['state'] = str(next_path)
        app.atomic_json(manifest_path, manifest)
        second = self.console.snapshot(upload['id'], first['selected_lane'])
        self.assertEqual(second['selected_lane'], first['selected_lane'])
        self.assertEqual(second['lanes'][0]['state']['current_company'], 'Next')

    def test_worker_status_outside_run_is_rejected_before_read(self):
        upload = self.upload()
        directory = self.root / upload['id'] / 'runs' / 'shared-run'
        fanout = directory / 'industries/test/fanout'
        fanout.mkdir(parents=True)
        app.atomic_json(directory / 'manifest.json', {'run_id':'shared-run'})
        app.atomic_json(directory / 'status.json', {'stage':'collecting','industries':[]})
        outside = self.root / 'outside-status.json'
        app.atomic_json(outside, {'status':'running'})
        app.atomic_json(fanout / 'manifest.json', {
            'industry':'Test','scheduling':{'strategy':'company_work_sharing'},
            'lanes':[{'lane':1,'state':str(outside),'database':str(self.root/'outside.db')}]})
        original = app.read_json
        def guarded_read(path):
            if Path(path).resolve() == outside.resolve():
                self.fail('An outside status must not be read.')
            return original(path)
        with patch.object(app, 'read_json', side_effect=guarded_read):
            with self.assertRaisesRegex(ValueError, 'outside its collection run'):
                self.console.snapshot(upload['id'])

    def test_shared_worker_history_counts_every_invocation_and_newer_fanout_status_wins(self):
        # A shared worker writes one status file per invocation and the fanout
        # manifest names only the current one; the company table needs them all.
        upload = self.upload()
        directory = self.root / upload['id'] / 'runs' / 'shared-run'
        fanout = directory / 'industries/test/fanout'
        lane_one, lane_two = fanout / 'lanes/lane-01', fanout / 'lanes/lane-02'
        lane_one.mkdir(parents=True)
        lane_two.mkdir(parents=True)
        app.atomic_json(directory / 'manifest.json', {'run_id':'shared-run'})
        stale = {'industry':'Test','run_id':'fan-1','total':4,'completed':2,'running':1,'failed':0,
                 'updated_at':'2026-09-17T19:55:59.336431+00:00'}
        app.atomic_json(directory / 'status.json', {'stage':'collecting','industries':[],'fanout':stale})
        task = lambda company, cycle, status, rows: {'task_id':f'{company}-{cycle}'.lower(), 'company':company,
                                                    'cycle':cycle, 'status':status, 'rows':rows}
        first = lane_one / 'status.work-1-a.json'
        app.atomic_json(first, {'status':'failed','tasks':[task('Acme',2022,'succeeded',5),
                                                           task('Acme',2024,'failed',0),
                                                           task('Acme',2026,'succeeded',7)]})
        retry = lane_one / 'status.recovery-1.json'
        app.atomic_json(retry, {'status':'succeeded','tasks':[task('Acme',2024,'succeeded',3)]})
        other = lane_two / 'status.work-1-b.json'
        app.atomic_json(other, {'status':'failed','tasks':[task('Other',2022,'failed',0)]})
        os.utime(first, (1_000_000, 1_000_000))
        os.utime(retry, (2_000_000, 2_000_000))
        # lane-02's current status file never landed; its earlier receipt still counts.
        app.atomic_json(fanout / 'manifest.json', {
            'industry':'Test','run_id':'fan-1','cycles':[2022,2024,2026],'scheduling':{'strategy':'company_work_sharing'},
            'lanes':[{'lane':1,'state':str(retry),'database':str(lane_one/'mentat.db')},
                     {'lane':2,'state':str(lane_two/'status.split-3-missing.json'),'database':str(lane_two/'mentat.db')}]})
        fresh = {**stale, 'completed':4, 'running':0, 'failed':1, 'updated_at':'2026-09-17T19:56:00.073477+00:00'}
        app.atomic_json(fanout / 'status.json', fresh)
        result = self.console.snapshot(upload['id'])
        tasks = {(row['company'], row['cycle']): (row['status'], row['rows'], row['industry']) for row in result['company_tasks']}
        self.assertEqual(tasks, {('Acme',2022):('succeeded',5,'Test'), ('Acme',2024):('succeeded',3,'Test'),
                                 ('Acme',2026):('succeeded',7,'Test'), ('Other',2022):('failed',0,'Test')})
        self.assertEqual(result['cycles'], [2022, 2024, 2026])
        self.assertEqual(result['status']['fanout'], fresh)
        # The fanout's own file wins only when it is newer and names the same run.
        app.atomic_json(fanout / 'status.json', {**fresh, 'updated_at':'2026-09-17T19:00:00+00:00'})
        self.assertEqual(self.console.snapshot(upload['id'])['status']['fanout'], stale)
        app.atomic_json(fanout / 'status.json', {**fresh, 'run_id':'another-fanout'})
        self.assertEqual(self.console.snapshot(upload['id'])['status']['fanout'], stale)

    def test_stopped_run_keeps_the_collectors_elapsed_time(self):
        upload = self.upload()
        directory = self.root / upload['id'] / 'runs' / 'stopped-run'
        directory.mkdir(parents=True)
        app.atomic_json(directory / 'manifest.json', {'run_id':'stopped-run'})
        app.atomic_json(directory / 'status.json', {'stage':'collecting','elapsed_seconds':11614.7,'industries':[]})
        manifest_path = self.root / upload['id'] / 'upload.json'
        app.atomic_json(manifest_path, {**app.read_json(manifest_path), 'state':'stopped', 'started_at':time.time() - 500000})
        result = self.console.snapshot(upload['id'])
        self.assertEqual(result['state'], 'stopped')
        self.assertEqual(result['elapsed_seconds'], 11614.7)

    def test_evidence_endpoint_rejects_estimate_from_old_model_code(self):
        import hashlib
        from pac_cash_model import write_estimates
        from test_pac_cash_model import snapshot
        directory=self.root/'.cache/pac-evidence/fixture';directory.mkdir(parents=True)
        (self.root/'tools').mkdir()
        (self.root/'tools/pac_cash_model.py').write_bytes((app.PROJECT/'tools/pac_cash_model.py').read_bytes())
        source=json.dumps([snapshot()]).encode()
        (directory/'snapshots.json').write_bytes(source)
        app.atomic_json(directory/'evidence.json',dict(version=1,company='Acme',aliases=['Acme Inc'],
            source_snapshot_sha256=hashlib.sha256(source).hexdigest()))
        current=write_estimates(directory)
        server=app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler);server.app=self.console
        worker=threading.Thread(target=server.serve_forever,daemon=True);worker.start()
        try:
            with patch.object(app,'PROJECT',self.root):
                url=f'http://127.0.0.1:{server.server_port}/api/evidence?id=fixture'
                with urlopen(url) as response:self.assertIn('cash_estimates',json.load(response))
                app.atomic_json(directory/'cash-estimates.json',dict(current,method_code_sha256='old-model'))
                with urlopen(url) as response:
                    result=json.load(response)
                    self.assertNotIn('cash_estimates',result)
                    self.assertEqual(result['company'],'Acme')
        finally:
            server.shutdown();server.server_close();worker.join()

    def test_http_rejects_untrusted_posts_and_hosts(self):
        server = app.ThreadingHTTPServer(("127.0.0.1",0),app.Handler)
        server.app = self.console
        worker = threading.Thread(target=server.serve_forever,daemon=True)
        worker.start()
        origin = f"http://127.0.0.1:{server.server_port}"
        try:
            with urlopen(origin + "/api/session") as response:
                self.assertEqual(json.load(response)["token"],self.console.token)
            for headers in ({}, {"X-Mentat-Token":self.console.token,"Origin":"https://untrusted.example"}):
                with self.assertRaises(HTTPError) as error:
                    urlopen(Request(origin + "/api/start", data=b"{}", headers=headers))
                self.assertEqual(error.exception.code,403)
            with self.assertRaises(HTTPError) as error:
                urlopen(Request(origin + "/api/session",headers={"Host":"untrusted.example"}))
            self.assertEqual(error.exception.code,403)
        finally:
            server.shutdown()
            server.server_close()
            worker.join()

    def test_review_is_bound_to_workbook_and_warning_and_preserves_evidence(self):
        directory = self.root / 'fixture'
        directory.mkdir()
        workbook = self.root / 'report.xlsx'
        workbook.write_bytes(b'fixture')
        app.atomic_json(directory / 'status.json', {'industries': [
            {'industry': 'Test', 'status': 'succeeded', 'output_workbook': str(workbook)}]})
        run = {'id': 'test', 'name': 'Fixture', 'state': 'completed', 'directory': directory}
        warning = {'check': 'coverage', 'severity': 'warn', 'detail': 'Old source'}
        quality = {'sha256': 'first', 'warnings': [warning]}
        with patch.object(self.console, 'find', return_value=run), patch.object(self.console, 'inspect_report', return_value=quality):
            first = self.console.snapshot('test')['reports'][0]['quality']['warnings'][0]
            self.console.acknowledge_quality({'id': 'test', 'warning_ids': [first['id']]})
            reviewed = self.console.snapshot('test')['reports'][0]['quality']['warnings'][0]
            self.assertTrue(reviewed['reviewed_at'])
            self.assertEqual(reviewed['detail'], 'Old source')
            self.assertEqual(workbook.read_bytes(), b'fixture')
            # Repeated polling does not mutate cached warnings or their IDs.
            self.assertEqual(first['id'], reviewed['id'])
            self.assertEqual(quality['warnings'], [warning])
            quality['sha256'] = 'new-workbook'
            changed = self.console.snapshot('test')['reports'][0]['quality']['warnings'][0]
            self.assertIsNone(changed['reviewed_at'])
            with self.assertRaisesRegex(ValueError, 'changed'):
                self.console.acknowledge_quality({'id': 'test', 'warning_ids': [first['id']]})
            quality['sha256'] = 'first'
            quality['warnings'] = [{**warning, 'detail': 'New finding'}]
            self.assertIsNone(self.console.snapshot('test')['reports'][0]['quality']['warnings'][0]['reviewed_at'])

    def test_failed_check_cannot_be_acknowledged_away(self):
        with patch.object(self.console, 'snapshot', return_value={'reports': [
            {'quality': {'warnings': [{'id': 'hard-fail', 'severity': 'fail'}]}}]}):
            with self.assertRaises(ValueError):
                self.console.acknowledge_quality({'id': 'test', 'warning_ids': ['hard-fail']})

    def resumable_run(self, cycles=(2024,)):
        """A saved roster whose most recent run stopped with lane state on disk."""
        uploaded = self.upload()
        directory = self.root / uploaded['id']
        run_id = '20260917T101500-4242'
        run = directory / 'runs' / run_id
        fanout = run / 'industries' / 'test' / 'fanout'
        fanout.mkdir(parents=True)
        app.atomic_json(run / 'manifest.json', {'run_id': run_id})
        app.atomic_json(run / 'status.json', {'run_id': run_id, 'stage': 'collecting'})
        manifest = app.read_json(directory / 'upload.json')
        manifest.update(state='stopped', cycles=list(cycles), stop_requested=True,
                        stop_mode='force', exit_code=3, finished_at=1.0, pid=4242)
        app.atomic_json(directory / 'upload.json', manifest)
        return uploaded['id'], directory, run, fanout

    def test_resume_refuses_wrong_state_missing_lane_state_or_active_collection(self):
        identifier, directory, run, fanout = self.resumable_run()
        # Resume applies to a terminal run only; nothing else is "the last run".
        for state in ('ready', 'preparing', 'running', 'completed'):
            manifest = app.read_json(directory / 'upload.json')
            manifest['state'] = state
            app.atomic_json(directory / 'upload.json', manifest)
            with self.assertRaisesRegex(ValueError, 'Nothing to resume for this roster'):
                self.console.resume({'id': identifier})
        manifest = app.read_json(directory / 'upload.json')
        manifest['state'] = 'stopped'
        app.atomic_json(directory / 'upload.json', manifest)
        # Stopped, but no fanout manifest: the fanout has no lane snapshots to reload.
        with self.assertRaisesRegex(ValueError, 'Nothing to resume for this roster'):
            self.console.resume({'id': identifier})
        app.atomic_json(fanout / 'manifest.json', {'run_id': run.name, 'lane_count': 1})
        self.assertTrue(self.console.find(identifier)['state'] == 'stopped')
        self.assertTrue(self.console.snapshot(identifier)['resumable'])
        # Resumable, but the runtime is still owned by a live collection.
        self.console.processes['other'] = MagicMock(**{'poll.return_value': None})
        with self.assertRaisesRegex(ValueError, 'already running'):
            self.console.resume({'id': identifier})
        self.console.processes.clear()
        with self.assertRaises(app.RunNotFound):
            self.console.resume({'id': 'f' * 32})
        with self.assertRaises(app.RunNotFound):
            self.console.resume({'id': 'not-an-id'})
        # No refused resume may leave the roster mid-launch or edit its controls.
        self.assertEqual(app.read_json(directory / 'upload.json')['state'], 'stopped')
        self.assertEqual(self.console.preparing, set())
        self.assertEqual(app.read_json(directory / 'control.json'), {})

    def test_resume_clears_every_stop_flag_before_relaunching_the_saved_run(self):
        identifier, directory, run, fanout = self.resumable_run()
        app.atomic_json(fanout / 'manifest.json', {'run_id': run.name, 'lane_count': 1})
        stale = '2026-09-17T10:20:00+00:00'
        app.atomic_json(directory / 'control.json', {'stop': True, 'desired_workers': 0, 'updated_at': stale})
        app.atomic_json(run / 'control.json', {'stop': True, 'desired_workers': 6, 'updated_at': stale})
        app.atomic_json(fanout / 'control.json', {'stop': True, 'desired_workers': 0, 'updated_at': stale})
        (directory / 'launcher.log').write_bytes(b'stopped attempt\n')
        self.assertEqual(self.console.resume_control_targets(identifier, run.name),
                         [directory / 'control.json', run / 'control.json', fanout / 'control.json'])
        self.console.rscript = sys.executable
        commands, observed = [], {}

        def record(command, **kwargs):
            commands.append(list(command))
            # Read the controls the relaunched orchestrator will read, at the
            # moment it is launched: a stop flag left here stops it in seconds.
            observed.update({str(path): app.read_json(path) for path in
                             (directory / 'control.json', run / 'control.json', fanout / 'control.json')})
            return MagicMock(pid=4321)

        with patch.object(app, 'load_keys', return_value=['stored-key']), \
                patch.object(app, 'r_library_path', return_value=''), \
                patch.object(app, 'copy_fec_sources'), \
                patch.object(app, 'process_identity', return_value='fixture'), \
                patch.object(app.shutil, 'copytree'), patch.object(app.shutil, 'copy2'), \
                patch.object(app.threading, 'Thread'), \
                patch.object(app.subprocess, 'Popen', side_effect=record):
            value = self.console.resume({'id': identifier}, background=False)
        self.assertEqual(value, {'id': identifier, 'state': 'starting', 'resumed_from': run.name})
        command = commands[0]
        self.assertEqual(command[command.index('-ResumeRunId') + 1], run.name)
        self.assertNotIn('stored-key', ' '.join(command))
        # Appended, never truncated: the stopped attempt's launcher log survives.
        self.assertTrue((directory / 'launcher.log').read_bytes().startswith(b'stopped attempt\n'))
        # Cleared before the launch, and only a worker count the operator chose
        # is carried over - the zero stop() wrote is dropped, not saved.
        for path, workers in ((directory / 'control.json', None), (run / 'control.json', 6),
                              (fanout / 'control.json', None)):
            control = observed[str(path)]
            self.assertEqual(control, app.read_json(path))
            self.assertIs(control['stop'], False)
            self.assertEqual(control.get('desired_workers'), workers)
            self.assertNotEqual(control['updated_at'], stale)
        saved = app.read_json(directory / 'upload.json')
        self.assertEqual(saved['resumed_from'], run.name)
        self.assertEqual(saved['resume_count'], 1)
        self.assertEqual(saved['state'], 'starting')
        for key in ('stop_requested', 'stop_mode', 'exit_code', 'finished_at', 'launch_error'):
            self.assertNotIn(key, saved)
        self.assertEqual(self.console.preparing, set())

    def orphaned_run(self, state='starting', pid=None, identity=None, lane_state=True):
        """A saved roster left live-looking by a console that was killed.

        The orchestrator wrote its status while it collected and then died with
        the console that owned it, so no terminal state was ever recorded. The
        console PID stays this live process on purpose: the console that
        launched a run being alive says nothing about the orchestrator, so it
        must never be what keeps a dead collection looking live.
        """
        uploaded = self.upload()
        directory = self.root / uploaded['id']
        run_id = '20260917T124148-57020'
        run = directory / 'runs' / run_id
        fanout = run / 'industries' / 'test' / 'fanout'
        fanout.mkdir(parents=True)
        app.atomic_json(run / 'manifest.json', {'run_id': run_id})
        app.atomic_json(run / 'status.json', {'run_id': run_id, 'stage': 'collecting'})
        if lane_state:
            app.atomic_json(fanout / 'manifest.json', {'run_id': run_id, 'lane_count': 1})
        manifest = app.read_json(directory / 'upload.json')
        manifest.update(state=state, cycles=[2024], started_at=1.0, console_pid=os.getpid())
        if pid is not None:
            manifest.update(pid=pid, process_identity=identity)
        app.atomic_json(directory / 'upload.json', manifest)
        return uploaded['id'], directory, run

    def dead_pid(self):
        """A PID with no live process: one this test started and watched exit.

        The finished Popen object is kept for the life of the test, so Windows
        cannot hand the PID to an unrelated process and make the fixture lie.
        """
        self.finished = subprocess.Popen([sys.executable, '-c', 'pass'])
        self.finished.wait()
        self.assertFalse(app.process_alive(self.finished.pid))
        return self.finished.pid

    def summary_row(self, identifier):
        """The row /api/runs hands the page for one roster."""
        return next(row for row in self.console.summary()['runs'] if row['id'] == identifier)

    def test_dead_collection_reports_stopped_and_resumable_without_rewriting_it(self):
        pid = self.dead_pid()
        identifier, directory, run = self.orphaned_run(
            state='starting', pid=pid, identity={'pid': pid, 'created': 1, 'executable': 'orchestrator.exe'})
        before = (directory / 'upload.json').read_bytes()
        row = self.summary_row(identifier)
        self.assertEqual(row['state'], 'stopped')
        self.assertTrue(row['resumable'])
        # Computed, not recorded: a GET leaves the saved manifest byte for byte.
        self.assertEqual((directory / 'upload.json').read_bytes(), before)
        self.assertEqual(app.read_json(directory / 'upload.json')['state'], 'starting')
        # find() feeds resume() and snapshot() from the same computed list.
        self.assertEqual(self.console.find(identifier)['state'], 'stopped')
        self.assertTrue(self.console.snapshot(identifier)['resumable'])
        self.assertEqual(self.console.resume_target(self.console.find(identifier)), run)

    def test_live_collection_is_never_demoted(self):
        identifier, directory, run = self.orphaned_run(
            state='starting', pid=os.getpid(), identity=app.process_identity(os.getpid()))
        row = self.summary_row(identifier)
        self.assertEqual(row['state'], 'starting')
        self.assertFalse(row['resumable'])
        self.assertEqual(app.read_json(directory / 'upload.json')['state'], 'starting')

    @unittest.skipUnless(os.name == 'nt', 'Windows process identity')
    def test_recycled_pid_cannot_keep_a_dead_collection_looking_live(self):
        identifier, directory, run = self.orphaned_run(
            state='running', pid=os.getpid(),
            identity={**app.process_identity(os.getpid()), 'created': 1})
        self.assertEqual(self.summary_row(identifier)['state'], 'stopped')

    def test_run_this_console_still_owns_is_never_demoted(self):
        identifier, directory, run = self.orphaned_run(state='starting', pid=self.dead_pid())
        self.console.preparing.add(identifier)
        self.assertEqual(self.summary_row(identifier)['state'], 'starting')
        self.console.preparing.clear()
        self.console.processes[identifier] = MagicMock(**{'poll.return_value': None})
        self.assertEqual(self.summary_row(identifier)['state'], 'starting')
        self.console.processes[identifier].poll.return_value = 1
        self.assertEqual(self.summary_row(identifier)['state'], 'stopped')

    def test_preparing_without_a_pid_is_not_demoted(self):
        identifier, directory, run = self.orphaned_run(state='preparing')
        status = run / 'status.json'
        os.utime(status, (time.time(), time.time() - app.ORPHANED_RUN_GRACE_SECONDS - 3600))
        row = self.summary_row(identifier)
        self.assertEqual(row['state'], 'preparing')
        self.assertFalse(row['resumable'])

    def test_collecting_without_a_pid_is_demoted_only_after_the_grace_elapses(self):
        self.assertGreaterEqual(app.ORPHANED_RUN_GRACE_SECONDS, 15 * 60)
        identifier, directory, run = self.orphaned_run(state='starting')
        before = (directory / 'upload.json').read_bytes()
        status = run / 'status.json'
        for age in (0, app.ORPHANED_RUN_GRACE_SECONDS - 60):
            os.utime(status, (time.time(), time.time() - age))
            row = self.summary_row(identifier)
            self.assertEqual(row['state'], 'starting', f'age {age}s is inside the grace')
            self.assertFalse(row['resumable'])
        os.utime(status, (time.time(), time.time() - app.ORPHANED_RUN_GRACE_SECONDS - 60))
        row = self.summary_row(identifier)
        self.assertEqual(row['state'], 'stopped')
        self.assertTrue(row['resumable'])
        self.assertEqual((directory / 'upload.json').read_bytes(), before)

    def test_demoted_run_passes_the_resume_state_guard(self):
        pid = self.dead_pid()
        identifier, directory, run = self.orphaned_run(
            state='starting', pid=pid, identity={'pid': pid, 'created': 1, 'executable': 'orchestrator.exe'})
        self.console.rscript = sys.executable
        with patch.object(app, 'load_keys', return_value=['stored-key']), \
                patch.object(app, 'r_library_path', return_value=''), \
                patch.object(app, 'copy_fec_sources'), \
                patch.object(app, 'process_identity', return_value='fixture'), \
                patch.object(app.shutil, 'copytree'), patch.object(app.shutil, 'copy2'), \
                patch.object(app.threading, 'Thread'), \
                patch.object(app.subprocess, 'Popen', return_value=MagicMock(pid=4321)) as launch:
            value = self.console.resume({'id': identifier}, background=False)
        self.assertEqual(value, {'id': identifier, 'state': 'starting', 'resumed_from': run.name})
        command = launch.call_args.args[0]
        self.assertEqual(command[command.index('-ResumeRunId') + 1], run.name)
        saved = app.read_json(directory / 'upload.json')
        self.assertEqual(saved['resumed_from'], run.name)
        self.assertEqual(saved['pid'], 4321)
        self.assertEqual(saved['resume_count'], 1)


if __name__ == "__main__":
    unittest.main()
