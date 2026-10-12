"""Independently reconcile the consolidated workbook to its bound audit CSVs."""
import argparse
from collections import defaultdict
import csv
import hashlib
import json
from pathlib import Path

import openpyxl
from verify_repaired_reports import near


def signature(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def verify(directory):
    expected = {}
    roots = defaultdict(lambda: [0., 0., 0.])
    for audit in directory.glob('*.audit'):
        manifest = json.loads((audit / 'audit-manifest.json').read_text(encoding='utf-8'))
        if signature(directory / manifest['workbook_name']) != manifest['workbook_sha256']:
            raise ValueError('Source workbook checksum failed')
        for name, digest in manifest['csv_sha256'].items():
            if signature(audit / name) != digest:
                raise ValueError('Source CSV checksum failed')
        with (audit / 'attribution_summary.csv').open(encoding='utf-8-sig', newline='') as stream:
            for row in csv.DictReader(stream):
                key = (row['industry'], int(row['cycle']), row['candidate_id'])
                if key in expected:
                    raise ValueError('Duplicate candidate across source reports')
                # Blank is not traced (enrich_downstream.py, transfer index status).
                expected[key] = [None if row[name] == '' else float(row[name]) for name in ('direct',
                    'estimated_non_opposition', 'estimated_opposition', 'direct_plus_estimated_non_opposition')]
        with (audit / 'allocation_coverage.csv').open(encoding='utf-8-sig', newline='') as stream:
            for row in csv.DictReader(stream):
                key = (int(row['cycle']), row['root'])
                for i, name in enumerate(('model_funding', 'estimated_candidate_activity', 'unresolved')):
                    roots[key][i] += float(row[name])
    path = directory / 'Attributed Industry Spending.xlsx'
    before = signature(path)
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    seen = set()
    ordering = []
    actual_roots = defaultdict(lambda: [0., 0., 0.])
    try:
        for row in book['Candidate Funding'].iter_rows(min_row=6, values_only=True):
            key = (row[0], int(row[1]), row[2])
            if key in seen or key not in expected:
                raise ValueError('Duplicate or unexpected candidate in workbook')
            seen.add(key)
            if expected[key][1] is None:
                # Not traced: the three conditional cells stay blank. near() reads
                # None as 0, so a blank that came back as 0 would pass it.
                if any(value not in (None, '') for value in row[6:9]):
                    raise ValueError('Untraced conditional estimates must stay blank in the workbook')
                near(row[5], expected[key][0], 'Candidate workbook versus full CSV')
                ordering.append((key[0], key[1], -row[5], key[2]))
                continue
            if any(value in (None, '') for value in row[6:8]):
                raise ValueError('Traced conditional estimates are blank in the workbook')
            for actual, value in zip(row[5:9], expected[key]):
                near(actual, value, 'Candidate workbook versus full CSV')
            near(row[8], row[5] + row[6], 'Opposition excluded from combined estimate')
            ordering.append((key[0], key[1], -row[8], key[2]))
        if seen != set(expected) or ordering != sorted(ordering):
            raise ValueError('Missing or incorrectly sorted candidate rows')
        for row in book['Root Coverage'].iter_rows(min_row=6, values_only=True):
            near(row[3], row[4] + row[5], 'Root conservation')
            for i, value in enumerate(row[3:6]):
                actual_roots[int(row[1]), row[2]][i] += value
        if set(actual_roots) != set(roots):
            raise ValueError('Missing root coverage')
        for key in roots:
            for actual, value in zip(actual_roots[key], roots[key]):
                near(actual, value, 'Root workbook versus full CSV')
        for sheet in book:
            for row in sheet:
                if any(cell.data_type == 'e' for cell in row):
                    raise ValueError('Cached Excel error')
    finally:
        book.close()
    if signature(path) != before:
        raise ValueError('Workbook changed during audit')
    return {'status': 'passed', 'file': str(path.resolve()), 'sha256': before,
        'candidate_rows': len(seen), 'industries': sorted({key[0] for key in seen}),
        'industry_cycles': len({key[:2] for key in seen})}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('directory', type=Path)
    parser.add_argument('--receipt', type=Path, required=True)
    args = parser.parse_args()
    result = verify(args.directory)
    args.receipt.write_text(json.dumps(result, indent=2), encoding='utf-8')
    print(json.dumps(result))
