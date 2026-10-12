import gc
import re
import sys
import time
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils.cell import get_column_letter, range_boundaries


SOURCE = (
    Path(sys.argv[1]).resolve()
    if __name__ == "__main__" and len(sys.argv) >= 3
    else None
)
OUTPUT_DIRECTORY = (
    Path(sys.argv[2]).resolve()
    if __name__ == "__main__" and len(sys.argv) >= 3
    else None
)

CYCLES = (
    ("21-22", "2021-2022", 2022),
    ("23-24", "2023-2024", 2024),
    ("25-26", "2025-2026", 2026),
)
COMPANIES = (
    "AT&T",
    "Altice USA",
    "Charter Communications",
    "Comcast",
    "Cox Communications",
    "EchoStar",
    "Frontier Communications",
    "Lumen Technologies",
    "T-Mobile",
    "Verizon",
)


def safe_table_name(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_]", "_", value)
    if not value or value[0].isdigit():
        value = f"T_{value}"
    return value[:250]


def make_tables_excel_safe(book):
    """Normalize table metadata to the stricter rules enforced by Excel."""
    next_table_id = 1
    for sheet in book.worksheets:
        tables = list(sheet.tables.values())
        if not tables:
            continue

        # A worksheet-level AutoFilter cannot overlap an Excel table's own
        # AutoFilter. The table retains its filter controls.
        sheet.auto_filter.ref = None

        for table in tables:
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)

            # Excel removes tables whose range contains only a header row.
            # Preserve the empty table as a header plus one blank data row.
            if max_row == min_row:
                max_row += 1
                table.ref = (
                    f"{get_column_letter(min_col)}{min_row}:"
                    f"{get_column_letter(max_col)}{max_row}"
                )

            column_count = max_col - min_col + 1
            if len(table.tableColumns) != column_count:
                table.tableColumns = []
                table._initialise_columns()

            # Excel table-column IDs must be 1..N, and every displayed header
            # must be unique within a table. The legacy overview used worksheet
            # column numbers (2..84) and repeated labels, which Excel repaired.
            seen = {}
            for offset, column in enumerate(table.tableColumns, start=1):
                cell = sheet.cell(min_row, min_col + offset - 1)
                base = str(cell.value or f"Column {offset}").strip()
                occurrence = seen.get(base.casefold(), 0) + 1
                seen[base.casefold()] = occurrence
                header = base if occurrence == 1 else f"{base} ({occurrence})"
                cell.value = header
                column.id = offset
                column.name = header

            table.id = next_table_id
            next_table_id += 1


def add_summary_sheet(book, title: str, cycle_label: str, cycle: int, index: int):
    sheet = book.create_sheet(title, index)
    headers = (
        "Company",
        "Cycle",
        "Total Giving",
        "Democratic",
        "Republican",
        "Other / Unclassified",
        "Democratic %",
        "Republican %",
        "Cycle Change %",
    )
    sheet.append(headers)
    for company_index, company in enumerate(COMPANIES, start=2):
        sheet.cell(company_index, 1, company)
        sheet.cell(company_index, 2, cycle_label)
        company_ref = f"$A{company_index}"
        cycle_ref = f"$B{company_index}"
        total = (
            f'=SUMIFS(Facts!$G:$G,Facts!$B:$B,{cycle_ref},'
            f'Facts!$C:$C,{company_ref})'
        )
        democratic = (
            f'=SUMIFS(Facts!$G:$G,Facts!$B:$B,{cycle_ref},'
            f'Facts!$C:$C,{company_ref},Facts!$F:$F,"D")'
        )
        republican = (
            f'=SUMIFS(Facts!$G:$G,Facts!$B:$B,{cycle_ref},'
            f'Facts!$C:$C,{company_ref},Facts!$F:$F,"R")'
        )
        other = f"=C{company_index}-D{company_index}-E{company_index}"
        democratic_pct = (
            f'=IF((D{company_index}+E{company_index})=0,0,'
            f'D{company_index}/(D{company_index}+E{company_index}))'
        )
        republican_pct = (
            f'=IF((D{company_index}+E{company_index})=0,0,'
            f'E{company_index}/(D{company_index}+E{company_index}))'
        )
        previous_cycle = cycle - 2
        change = (
            f'=IF(SUMIFS(Facts!$G:$G,Facts!$B:$B,'
            f'"{previous_cycle - 1}-{previous_cycle}",Facts!$C:$C,{company_ref})=0,'
            f'0,C{company_index}/SUMIFS(Facts!$G:$G,Facts!$B:$B,'
            f'"{previous_cycle - 1}-{previous_cycle}",Facts!$C:$C,{company_ref})-1)'
        )
        for column, formula in enumerate(
            (
                total,
                democratic,
                republican,
                other,
                democratic_pct,
                republican_pct,
                change,
            ),
            start=3,
        ):
            sheet.cell(company_index, column, formula)

    table = Table(
        displayName=safe_table_name(f"Tbl_{title}_{cycle_label}"),
        ref=f"A1:I{len(COMPANIES) + 1}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    navy = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = navy
        cell.font = Font(color="FFFFFF", bold=True)
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=6):
        for cell in row:
            cell.number_format = '$#,##0;[Red]($#,##0);-'
    for row in sheet.iter_rows(min_row=2, min_col=7, max_col=9):
        for cell in row:
            cell.number_format = '0.0%;[Red](0.0%);-'
    widths = (24, 14, 16, 16, 16, 20, 16, 16, 16)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return sheet


def filter_facts(sheet, cycle_labels):
    cycle_labels = set(cycle_labels)
    keep_rows = []
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 2).value in cycle_labels:
            keep_rows.append(row)
    if not keep_rows:
        raise RuntimeError(f"Facts contains no rows for {sorted(cycle_labels)}")

    # Delete excluded rows in descending contiguous runs. Later-cycle workbooks
    # retain the immediately preceding cycle so their change formulas have a
    # valid Excel source while still staying well below the combined file size.
    keep_row_set = set(keep_rows)
    excluded = [
        row for row in range(2, sheet.max_row + 1) if row not in keep_row_set
    ]
    runs = []
    for row in excluded:
        if not runs or row != runs[-1][1] + 1:
            runs.append([row, row])
        else:
            runs[-1][1] = row
    for first, last in reversed(runs):
        sheet.delete_rows(first, last - first + 1)

    final_row = sheet.max_row
    for table in sheet.tables.values():
        table.ref = f"A1:L{final_row}"
    return final_row


def split_cycle(short_label: str, cycle_label: str, cycle: int):
    if SOURCE is None or OUTPUT_DIRECTORY is None:
        raise RuntimeError("split_cycle requires source and output arguments")
    started = time.perf_counter()
    OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)
    print(f"LOAD {cycle_label}", flush=True)
    book = openpyxl.load_workbook(
        SOURCE,
        read_only=False,
        data_only=False,
        keep_links=False,
    )
    suffix = f" {short_label}"
    keep = {
        "IND TELECOM OVERVIEW",
        f"IND TELECOM Raw{suffix}",
        f"IND TELECOM Recip{suffix}",
        f"IND TELECOM Contrib{suffix}",
        "Facts",
        "PAC Facts",
    }
    for company in COMPANIES:
        prefix = {
            "Charter Communications": "Charter Comm",
            "Cox Communications": "Cox Communic",
            "Frontier Communications": "Frontier Com",
            "Lumen Technologies": "Lumen Techno",
        }.get(company, company)
        keep.update(
            {
                f"{prefix} Raw{suffix}",
                f"{prefix} Recip{suffix}",
                f"{prefix} Contrib{suffix}",
            }
        )
    missing = sorted(keep - set(book.sheetnames))
    if missing:
        raise RuntimeError(f"Missing cycle sheets for {cycle_label}: {missing}")
    for sheet in list(book.worksheets):
        if sheet.title not in keep:
            book.remove(sheet)

    overview = book["IND TELECOM OVERVIEW"]
    # Excel worksheet titles are limited to 31 characters. The full report
    # name belongs in the filename; this shorter sheet title is Excel-safe.
    overview.title = "Telecom Spending Overview"
    selected_rows = {2026: (2, 14), 2024: (16, 28), 2022: (30, 42)}[cycle]
    for row in range(1, overview.max_row + 1):
        overview.row_dimensions[row].hidden = not (
            selected_rows[0] <= row <= selected_rows[1]
        )
    overview.sheet_view.showGridLines = False
    overview.freeze_panes = "B3"

    employee = book[f"IND TELECOM Contrib{suffix}"]
    employee.title = "Employee List"
    recipient = book[f"IND TELECOM Recip{suffix}"]
    recipient.title = "Top Recipients"
    raw = book[f"IND TELECOM Raw{suffix}"]
    raw.title = f"Contributions {cycle_label}"

    fact_cycles = {
        "2021-2022": ("2021-2022",),
        "2023-2024": ("2021-2022", "2023-2024"),
        "2025-2026": ("2023-2024", "2025-2026"),
    }[cycle_label]
    fact_rows = filter_facts(book["Facts"], fact_cycles)
    master = add_summary_sheet(book, "Master Overview", cycle_label, cycle, 0)
    company_list = add_summary_sheet(book, "Company List", cycle_label, cycle, 2)
    company_list.sheet_state = "visible"
    book.active = book.index(master)
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    book.calculation.calcMode = "auto"
    make_tables_excel_safe(book)

    target = OUTPUT_DIRECTORY / (
        f"Telecom Spending Overview {cycle_label}.xlsx"
    )
    pending = target.with_name(f"{target.name}.pending")
    if pending.exists():
        pending.unlink()
    print(
        f"SAVE {cycle_label} sheets={len(book.sheetnames)} fact_rows={fact_rows}",
        flush=True,
    )
    book.save(pending)
    pending.replace(target)
    book.close()
    gc.collect()
    print(
        f"DONE {cycle_label} path={target} bytes={target.stat().st_size} "
        f"elapsed={time.perf_counter() - started:.1f}s",
        flush=True,
    )


if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit(
            "usage: split_telecom_workbook.py SOURCE.xlsx OUTPUT_DIRECTORY "
            "[2021-2022 2023-2024 2025-2026]"
        )
    requested_cycles = set(sys.argv[3:])
    for cycle_spec in CYCLES:
        if requested_cycles and cycle_spec[1] not in requested_cycles:
            continue
        split_cycle(*cycle_spec)
