"""Read-only evidence audit of exported Mentat workbooks and source payloads."""
import collections
import datetime
import hashlib
import json
from pathlib import Path
import re
import sqlite3
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[2]


def sha(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def rows(book, name, header_row=5):
    if name not in book.sheetnames:
        return []
    iterator = book[name].iter_rows(min_row=header_row, values_only=True)
    headers = next(iterator, ())
    return [dict(zip(headers, row), _row=index) for index, row in enumerate(iterator, header_row+1)
            if any(value is not None for value in row)]


def amount(value):
    return float(value) if isinstance(value, (int,float)) else 0


def workbook_audit(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = {'path':str(path),'sha256':sha(path),'modified_at':datetime.datetime.fromtimestamp(path.stat().st_mtime).isoformat(),
              'sheets':{sheet.title:{'rows':sheet.max_row,'columns':sheet.max_column} for sheet in book}}
    try:
        summary = rows(book,'Beneficiary Summary')
        bridge = rows(book,'Beneficiary Bridge')
        downstream = rows(book,'Downstream Activity')
        recipient = rows(book,'Recipient List')
        if summary:
            result['beneficiary_has_direction_columns'] = 'Support Activity' in summary[0] and 'Oppose Activity' in summary[0]
        activity = collections.defaultdict(lambda:collections.defaultdict(float))
        bad_dates=[]
        for row in downstream:
            key=(row['Cycle'],row['Source Committee ID'],row['Candidate ID'])
            activity[key][row['Support / Oppose']]+=amount(row['Net Included Amount'])
            for field in ('First Activity Date','Last Activity Date'):
                value=row.get(field)
                # Prior-cycle transaction dates can legitimately occur in a filing
                # cycle. Flag impossible century parsing, not ordinary 2020 carryover.
                if value and str(value)[:4].isdigit() and int(str(value)[:4])<1900:
                    bad_dates.append({'cell':f'{field}:{row["_row"]}','value':str(value),'committee':row['Source Committee ID']})
        result['downstream_invalid_year_dates']={'count':len(bad_dates),'examples':bad_dates[:4]}
        by_candidate=collections.defaultdict(lambda:collections.defaultdict(float))
        for row in bridge:
            if row.get('View')!='Industry':continue
            key=(row['Cycle'],row['Candidate ID'])
            committee_id = row.get('Source Committee ID')
            for direction,value in activity[(row['Cycle'],committee_id,row['Candidate ID'])].items():
                by_candidate[key][direction]+=value
        affected=[]
        for row in summary:
            if row.get('View')!='Industry':continue
            key=(row['Cycle'],row['Candidate ID'])
            oppose=by_candidate[key]['Oppose']
            if oppose>0.01:
                affected.append({'row':row['_row'],'cycle':row['Cycle'],'candidate':row['Candidate / Officeholder'],
                                 'party':row['Party'],'downstream':row['Downstream Activity'],
                                 'reconstructed_support':round(by_candidate[key]['Support'],2),'reconstructed_oppose':round(oppose,2),
                                 'democratic_activity':row['Democratic Activity'],'republican_activity':row['Republican Activity']})
        result['industry_beneficiaries_with_opposition']={'count':len(affected),
            'oppose_dollars':round(sum(row['reconstructed_oppose'] for row in affected),2),
            'examples':sorted(affected,key=lambda row:row['reconstructed_oppose'],reverse=True)[:6]}
        # ActBlue and WinRed are their own Conduit category (2026-10-09). A book built
        # before that files them under Carey or Super PAC; earmark_conduits_not_conduit counts those rows.
        result['earmark_conduits']=[{key:row.get(key) for key in ('_row','Cycle','Recipient','Committee ID','Recipient Category','Total Giving','Category Source')}
            for row in recipient if row.get('View')=='Industry' and row.get('Committee ID') in ('C00401224','C00694323')]
        result['earmark_conduits_not_conduit']=sum(1 for row in result['earmark_conduits'] if row['Recipient Category']!='Conduit')
        cycle_results={}
        for name in ('2021-2022','2023-2024','2025-2026'):
            data=rows(book,name,4)
            if not data:continue
            totals=collections.defaultdict(float)
            errors=[]; out_of_order=[]; prior=None; ids=collections.Counter()
            for row in data:
                totals['analysis']+=amount(row.get('Analysis Amount'))
                totals['raw']+=amount(row.get('Amount'))
                if row.get('Include in Analysis')=='Yes':totals['included_raw']+=amount(row.get('Amount'))
                totals['unclassified']+=amount(row.get('Analysis Amount')) if row.get('Recipient Category')=='Unclassified' else 0
                if row.get('Company Match Status')!='accepted':totals['nonaccepted_rows']+=1
                if row.get('Recipient Category')=='Super PAC':totals['superpac']+=amount(row.get('Analysis Amount'))
                value=row.get('Date')
                if isinstance(value,datetime.datetime):
                    if prior and value>prior:out_of_order.append(row['_row'])
                    prior=value
                ids[row.get('Source ID')]+=1
                for key,value in row.items():
                    if isinstance(value,str) and value in ('#REF!','#VALUE!','#DIV/0!','#NAME?','#N/A','#NUM!'):
                        errors.append({'row':row['_row'],'column':str(key),'error':value})
            cycle_results[name]={'rows':len(data),'totals':{key:round(value,2) for key,value in totals.items()},
                'formula_errors':errors[:10],'formula_error_count':len(errors),'date_sort_violations':out_of_order[:10],
                'duplicate_source_ids':sum(count-1 for key,count in ids.items() if key and count>1)}
        result['cycles']=cycle_results
        company=rows(book,'Company List')
        result['company_list_headers']=list(company[0]) if company else []
        result['company_list_totals_by_cycle']={cycle:round(sum(amount(row.get('Total Giving')) for row in company if row.get('Cycle')==cycle),2) for cycle in cycle_results}
    finally:
        book.close()
    return result


def payload_audit(path):
    data=json.loads(path.read_text(encoding='utf-8-sig'))
    meta=data['metadata']
    result={'path':str(path),'metadata':meta}
    raw=data['raw_rows']
    by_type=collections.defaultdict(lambda:{'committees':set(),'amount':0,'rows':0})
    cmap={(row['cycle'],row['committee_id']):row for row in data['committee_map']}
    for row in raw:
        if row['include_in_analysis']!='Yes':continue
        mapped=cmap.get((row['cycle'],row['committee_id']),{})
        kind=mapped.get('master_committee_type')
        if kind in ('O','U','V','W') or row['committee_id'] in ('C00401224','C00694323'):
            key=f'{kind}|{row["recipient_category"]}|{row["category_source"]}'
            by_type[key]['committees'].add(row['committee_id']);by_type[key]['amount']+=row['amount'];by_type[key]['rows']+=1
    result['pac_taxonomy']={key:{**value,'amount':round(value['amount'],2),'committees':sorted(value['committees'])} for key,value in by_type.items()}
    included_ids={row['committee_id'] for row in raw if row['include_in_analysis']=='Yes'}
    downstream={(row['cycle'],row['source_committee_id']) for row in data['downstream_activity'] if row['included_rows']>0}
    result['superpacs_without_downstream']= [{key:row[key] for key in ('cycle','committee_id','committee_name','recipient_category','master_committee_type','preliminary_crypto_funding')}
        for row in data['committee_map'] if row['recipient_category']=='Super PAC' and row['preliminary_crypto_funding']>0 and (row['cycle'],row['committee_id']) not in downstream]
    # Compare ranking all offices before filtering with ranking Congress alone.
    funded={(r['cycle'],r['committee_id']) for r in raw if r['include_in_analysis']=='Yes' and r['committee_id']}
    candidates={}
    for row in data['downstream_activity']:
        if (row['cycle'],row['source_committee_id']) not in funded:continue
        key=(row['cycle'],row['candidate_id'])
        record=candidates.setdefault(key,{'candidate':row['candidate_name'],'office':row['office'],'amount':0})
        record['amount']+=amount(row['preliminary_included_amount'])
    result['congress_rank_filter']=[]
    for cycle in sorted({key[0] for key in candidates}):
        ordered=sorted([(key[1],value) for key,value in candidates.items() if key[0]==cycle],key=lambda item:(-item[1]['amount'],item[1]['candidate']))
        top_all=[key for key,value in ordered[:100] if value['office'] in ('H','S')]
        top_congress=[key for key,value in ordered if value['office'] in ('H','S')][:100]
        result['congress_rank_filter'].append({'cycle':cycle,'congress_in_first_100_all_offices':len(top_all),
            'congress_available_up_to_100':len(top_congress),'omitted_by_filter_order':[key for key in top_congress if key not in top_all]})
    database=Path(meta['database_path'])
    if database.exists():
        connection=sqlite3.connect(database.as_uri()+'?mode=ro',uri=True)
        try:
            result['api_progress']=connection.execute('SELECT cycle, COUNT(*), SUM(fetched_complete), SUM(stored_complete), MAX(updated_at) FROM api_work_progress GROUP BY cycle').fetchall()
            result['source_max_dates']=connection.execute('SELECT two_year_transaction_period,COUNT(*),MAX(contribution_receipt_date),MAX(scraped_at) FROM contributions GROUP BY two_year_transaction_period').fetchall()
            result['db_tables']=[row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")]
        except sqlite3.Error as error:result['db_error']=str(error)
        finally:connection.close()
    return result


def attributed_audit(path):
    book=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        sheet=book.worksheets[0]
        data=rows(book,sheet.title,4)
        totals=collections.defaultdict(lambda:collections.defaultdict(float))
        errors=[];duplicates=collections.Counter();sort_errors=[];previous={};negative=[];shares=[]
        for row in data:
            group=(row['Industry'],row['Cycle'])
            direct=amount(row['DIRECT (attributed)']);estimated=amount(row['PASSTHROUGH (est.)']);total=amount(row['TOTAL INDUSTRY SPEND'])
            for label,value in [('direct',direct),('estimated',estimated),('total',total)]:totals['|'.join(group)][label]+=value
            if abs(total-direct-estimated)>.011:errors.append(row['_row'])
            if group in previous and total>previous[group]+.011:sort_errors.append(row['_row'])
            previous[group]=total
            duplicates[(*group,row['Candidate ID'])]+=1
            if direct<0 or estimated<0:negative.append(row['_row'])
            share=row.get('its pro-member share')
            if isinstance(share,(float,int)) and not 0<=share<=1:shares.append({'row':row['_row'],'share':share})
        return {'path':str(path),'sha256':sha(path),'sheet':sheet.title,'rows':len(data),'methodology_cell_A2':sheet.cell(2,1).value,
                'total_arithmetic_errors':errors,'within_group_sort_errors':sort_errors,'duplicate_member_cycle_industry_keys':sum(v-1 for v in duplicates.values() if v>1),
                'negative_direct_or_estimate_rows':negative[:20],'out_of_range_largest_committee_shares':shares[:20],
                'totals_by_industry_cycle':{key:{k:round(v,2) for k,v in value.items()} for key,value in totals.items()}}
    finally:book.close()


def main():
    output=Path(sys.argv[1]);output.parent.mkdir(parents=True,exist_ok=True)
    targets=[ROOT/'reports'/f'{name} Industry Spending Overview.xlsx' for name in ('Crypto','Fintech','Credit Union')]
    targets.append(ROOT/'mentat/.cache/industry-report/crypto/stage-all_cycles/Crypto Industry Spending Overview.xlsx')
    evidence={'audited_at':datetime.datetime.now().isoformat(),'workbooks':[],'payloads':[]}
    for path in targets:
        if path.exists():
            result=workbook_audit(path);evidence['workbooks'].append(result)
            print(json.dumps({'file':path.name,'parent':str(path.parent),'direction_columns':result.get('beneficiary_has_direction_columns'),
                'opposition':result.get('industry_beneficiaries_with_opposition'), 'date_errors':result.get('downstream_invalid_year_dates'),
                'cycles':result.get('cycles'),'company_totals':result.get('company_list_totals_by_cycle')},default=str),flush=True)
    for industry in ('crypto','fintech','credit_union'):
        path=ROOT/'mentat/.cache/industry-report'/industry/'industry_data.json'
        if path.exists():
            result=payload_audit(path);evidence['payloads'].append(result)
            print(json.dumps({'payload':industry,'metadata':result['metadata'],'progress':result.get('api_progress'),
                'unresolved_superpacs':result['superpacs_without_downstream'],'pac_taxonomy':result['pac_taxonomy']},default=str),flush=True)
    attributed=ROOT/'reports/Attributed Industry Spending.xlsx'
    if attributed.exists():evidence['attributed']=attributed_audit(attributed)
    output.write_text(json.dumps(evidence,indent=2,default=str),encoding='utf-8')
    print('Evidence written: '+str(output))


if __name__=='__main__':main()
