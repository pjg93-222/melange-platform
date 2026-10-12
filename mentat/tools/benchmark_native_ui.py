"""Offline, synthetic workload: exact-result parity before reporting timings."""
from pathlib import Path
import argparse
import json
import os
import statistics
import sys
import time
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'app'))
import graph_data
import native_scan
from report_insights import stamp


def fixture(path,count):
    book=openpyxl.Workbook(write_only=True)
    receipts=book.create_sheet('2025-2026')
    for row in [['SYNTHETIC PERFORMANCE FIXTURE'],[],[],['Company','Employee Tier','Amount','Include in Analysis','Recipient Category','Recipient Key','Person Key','Contributor']]:receipts.append(row)
    for index in range(count):
        receipts.append([f'Fixture Company {index%25}','Employee',-100.25 if index%37==0 else 100.25,
                         'Yes','Super PAC',f'FEC:C{index%100:08}',f'P{index}',f'Person {index}'])
    activity=book.create_sheet('Downstream Activity')
    for row in [['SYNTHETIC PERFORMANCE FIXTURE'],[],[f'{count} source rows condensed to {count} committee-candidate-activity groups'],[],
                ['Cycle','Source Committee ID','Candidate ID','Candidate / Officeholder','Candidate Party','Support / Oppose','Net Included Amount']]:activity.append(row)
    for index in range(count):activity.append(['2025-2026',f'C{index%100:08}',f'H{index%1000:08}',f'Candidate {index%1000}','D','Support' if index%2 else 'Oppose',-1.25 if index%37==0 else 100.25])
    book.save(path)


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--rows',type=int,default=50000)
    parser.add_argument('--repeats',type=int,default=3)
    args=parser.parse_args()
    if not native_scan.executable():raise SystemExit('Build the scanner before benchmarking.')
    folder=ROOT/'.cache/ui-performance-20260920'
    folder.mkdir(parents=True,exist_ok=True)
    path=folder/'synthetic-workbook.xlsx'
    fixture(path,args.rows)
    file_stamp=stamp(path)
    timings={}
    for label,reader in [('receipt_scan',graph_data.report_facts),('activity_scan',graph_data.report_activity)]:
        reference=None
        for engine in ('off','on'):
            os.environ['MENTAT_NATIVE']=engine
            samples=[]
            for _ in range(args.repeats):
                reader.cache_clear()
                start=time.perf_counter()
                result=reader(file_stamp)
                samples.append(round((time.perf_counter()-start)*1000,3))
                if reference is None:reference=result
                assert result==reference, 'Native/reference result drift'
            timings[f'{label}_{"python" if engine=="off" else "rust"}']={'median_ms':statistics.median(samples),'samples_ms':samples}
            print(label,engine,timings[f'{label}_{"python" if engine=="off" else "rust"}'],flush=True)
    filters={}
    expected=graph_data._graph_view([path],filters)
    graph_data._views.clear()
    graph_data.graph_view([path],filters)
    for label,reader in [('view_reaggregation',graph_data._graph_view),('view_cached',graph_data.graph_view)]:
        samples=[]
        for _ in range(7):
            start=time.perf_counter()
            result=reader([path],filters)
            samples.append(round((time.perf_counter()-start)*1000,3))
            assert result==expected
        timings[label]={'median_ms':statistics.median(samples),'samples_ms':samples}
    output={'synthetic':True,'rows_per_sheet':args.rows,'xlsx_bytes':path.stat().st_size,
            'parity':'exact Python object equality for both readers and complete Graph payload',
            'method':'Uncached workbook reads, persistent native process; warm OS filesystem. View timings use already-decoded facts. Not end-to-end report generation.',
            'timings':timings}
    (folder/'results.json').write_text(json.dumps(output,indent=2),encoding='utf-8')
    print(json.dumps(output,indent=2),flush=True)
    native_scan.shutdown()


if __name__=='__main__':main()
