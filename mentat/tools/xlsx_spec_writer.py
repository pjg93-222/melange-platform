"""Write the .xlsx that crypto_workbook_builder.mjs describes, with openpyxl (MENTAT_XLSX_WRITER=openpyxl, D-178).

openpyxl_workbook_shim.mjs records what the builder asks the spreadsheet library to do (cell values, formulas, formats,
tables, merges, ...) as one JSON object per line; this program replays those operations, in order, onto an openpyxl
workbook. Nothing here decides what goes in a cell: the builder does, and writes the same cells whichever library
saves them.

    python xlsx_spec_writer.py spec.ndjson out.xlsx

Each artifact-tool behaviour that the content depends on is reproduced and named where it is done:
  * a string that is a zoned ISO-8601 timestamp becomes a UTC date cell, other text stays text;
  * text beginning with '=' is a formula;
  * text longer than Excel's 32,767 characters is kept whole (openpyxl would cut it);
  * numbers are written as the shortest text that reads back as the same double (openpyxl prints 16 digits);
  * fillDown copies the first cell's formula down with its relative references moved.
Not reproduced: PNG previews (MENTAT_RENDER_PREVIEWS) and the in-library recalculation, which openpyxl cannot do.
"""
from __future__ import annotations

import json
import math
import re
import sys
from copy import copy
from datetime import datetime, timezone

import openpyxl
import openpyxl.cell._writer as _cell_writer
from openpyxl.formatting.rule import Rule
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Color, Font, PatternFill
from openpyxl.styles.cell_style import StyleArray
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo

STRING_MAX = 32767
ISO_ZONED = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(\.\d+)?(Z|[+-]\d{2}:\d{2})$")
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_safe_string = _cell_writer.safe_string


def exact_number(value):
    """A number as JavaScript prints it; artifact-tool wrote that, and openpyxl's own '%.16g' loses the last digit."""
    if isinstance(value, float) and math.isfinite(value):
        if value == 0:
            return "0"
        return str(int(value)) if value == int(value) and abs(value) < 1e21 else repr(value)
    return _safe_string(value)


def coerce(value):
    if isinstance(value, str) and ISO_ZONED.match(value):
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(timezone.utc).replace(tzinfo=None)
        except ValueError:
            pass
    return value


def argb(color):
    return "FF" + color.lstrip("#").upper()


def assign(cell, value):
    """artifact-tool's range.values semantics for one cell."""
    value = coerce(value)
    if isinstance(value, str):
        value = ILLEGAL.sub("", value)
        if len(value) > STRING_MAX and not value.startswith("="):
            cell._value, cell.data_type = value, "s"
            return
    cell.value = value
    if isinstance(value, datetime):
        cell.number_format = "General"


class Styler:
    """Apply a format to a cell as property merges, cached by (current style, format) so a million cells cost a lookup."""

    def __init__(self, wb):
        self.wb = wb
        self.scratch = wb.create_sheet("__scratch__").cell(1, 1)
        self.cache = {}

    def done(self):
        self.wb.remove(self.scratch.parent)

    def style_of(self, cell):
        return tuple(cell._style) if cell._style is not None else ()

    def merged(self, old, key, props):
        hit = self.cache.get((old, key))
        if hit is None:
            probe = self.scratch
            probe._style = StyleArray(old) if old else StyleArray()
            if "fill" in props:
                probe.fill = PatternFill("solid", fgColor=argb(props["fill"]))
            if "font" in props:
                font = copy(probe.font)
                spec = props["font"]
                unknown = set(spec) - {"name", "size", "bold", "italic", "color"}
                if unknown:
                    raise ValueError("unsupported font property: %s" % ", ".join(sorted(unknown)))
                if "name" in spec:
                    font.name = spec["name"]
                if "size" in spec:
                    font.sz = spec["size"]
                if "bold" in spec:
                    font.b = bool(spec["bold"])
                if "italic" in spec:
                    font.i = bool(spec["italic"])
                if "color" in spec:
                    font.color = Color(rgb=argb(spec["color"]))
                probe.font = font
            if "numberFormat" in props:
                probe.number_format = props["numberFormat"]
            if props.keys() & {"wrap", "v", "h"}:
                align = copy(probe.alignment)
                if "wrap" in props:
                    align.wrap_text = bool(props["wrap"])
                if "v" in props:
                    align.vertical = props["v"]
                if "h" in props:
                    align.horizontal = props["h"]
                probe.alignment = align
            hit = self.cache[(old, key)] = tuple(probe._style)
        return hit

    def apply(self, cell, key, props):
        cell._style = StyleArray(self.merged(self.style_of(cell), key, props))


def cells_in(ws, ref):
    c1, r1, c2, r2 = range_boundaries(ref)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            yield ws.cell(row=r, column=c)


def dxf(spec):
    fill = PatternFill("solid", bgColor=argb(spec["fill"])) if spec.get("fill") else None
    font = None
    if spec.get("font"):
        font = Font(b=spec["font"].get("bold") or None, color=argb(spec["font"]["color"]) if spec["font"].get("color") else None)
    return DifferentialStyle(font=font, fill=fill)


def replay(spec_path, output):
    wb = openpyxl.Workbook(iso_dates=True)        # date cells are written t="d", as artifact-tool writes them
    wb.remove(wb.active)
    styler, sheets, tables, merges, formats = None, [], [], [], 0
    with open(spec_path, encoding="utf-8") as handle:
        for line in handle:
            op = json.loads(line)
            kind = op["op"]
            if kind == "book":
                sheets = [wb.create_sheet(name) for name in op["sheets"]]
                styler = Styler(wb)
                continue
            ws = sheets[op["s"]]
            if kind in ("values", "formulas"):
                for dr, row in enumerate(op["rows"]):
                    for dc, value in enumerate(row):
                        if value is None or value == "":
                            continue
                        assign(ws.cell(row=op["r"] + dr, column=op["c"] + dc), value)
            elif kind == "fillDown":
                c1, r1, c2, r2 = range_boundaries(op["ref"])
                for c in range(c1, c2 + 1):
                    first = ws.cell(row=r1, column=c)
                    for r in range(r1 + 1, r2 + 1):
                        cell = ws.cell(row=r, column=c)
                        if isinstance(first.value, str) and first.value.startswith("="):
                            cell.value = Translator(first.value, origin=first.coordinate).translate_formula(cell.coordinate)
                        else:
                            cell.value = first.value
            elif kind == "fmt":
                formats += 1
                for cell in cells_in(ws, op["ref"]):
                    styler.apply(cell, formats, op["p"])
            elif kind == "rowHeight":
                for r in range(op["r1"], op["r2"] + 1):
                    ws.row_dimensions[r].height = op["h"]
            elif kind == "colWidth":
                for c in range(op["c1"], op["c2"] + 1):
                    ws.column_dimensions[get_column_letter(c)].width = op["w"]
            elif kind == "merge":
                merges.append((ws, op["ref"]))
            elif kind == "cf":
                if op["type"] == "expression":
                    rule = Rule(type="expression", formula=[op["formula"].lstrip("=")], dxf=dxf(op))
                else:       # beginsWith: artifact-tool writes the rule with its operator and text and no formula
                    rule = Rule(type="beginsWith", operator="beginsWith", text=op["text"], dxf=dxf(op))
                ws.conditional_formatting.add(op["ref"], rule)
            elif kind == "table":
                table = Table(displayName=op["name"], name=op["name"], ref=op["ref"], headerRowCount=1 if op["headers"] else 0,
                              tableStyleInfo=TableStyleInfo(name=op["style"], showFirstColumn=False, showLastColumn=False,
                                                            showRowStripes=True, showColumnStripes=False))
                ws.add_table(table)
            elif kind == "freeze":
                ws.freeze_panes = "%s%d" % (get_column_letter(op["cols"] + 1), op["rows"] + 1) if op["rows"] or op["cols"] else None
            elif kind == "grid":
                ws.sheet_view.showGridLines = bool(op["show"])
            else:
                raise ValueError("unknown spec operation %r" % kind)
    styler.done()
    for ws, ref in merges:      # last: merging empties the covered cells, so nothing may be written there afterwards
        ws.merge_cells(ref)
    _cell_writer.safe_string = exact_number
    try:
        wb.save(output)
    finally:
        _cell_writer.safe_string = _safe_string
    return {"sheets": len(sheets), "formats": formats, "tables": sum(len(ws._tables) for ws in sheets)}


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) != 2:
        print("usage: xlsx_spec_writer.py spec.ndjson out.xlsx", file=sys.stderr)
        return 2
    print(json.dumps(replay(*argv)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
