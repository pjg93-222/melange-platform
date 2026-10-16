"""Read-only report receipts and worksheet data. No network or collection calls."""
from __future__ import annotations

import csv
from datetime import date, datetime, timezone
from email.utils import parsedate_to_datetime
from functools import lru_cache
import hashlib
import json
import math
from pathlib import Path
import re
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.cell.text import Text
from openpyxl.packaging.manifest import Manifest
from openpyxl.reader.excel import _find_workbook_part, _validate_archive
from openpyxl.reader.workbook import WorkbookParser
from openpyxl.styles.stylesheet import apply_stylesheet
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet._reader import WorkSheetParser
from openpyxl.xml.constants import ARC_CONTENT_TYPES, SHARED_STRINGS, SHEET_MAIN_NS
from openpyxl.xml.functions import fromstring, iterparse

import response_cache
import report_publication

MAX_ROWS = 50000
MAX_COLUMNS = 200
RECEIPT_SHEETS = ('Source Freshness', 'Source Receipts')
RECEIPT_LIMIT = 2000
# records() reads rows 1-30 to find its header, then the header row and up to
# RECEIPT_LIMIT + 1 more; openpyxl parses one row past a max_row to stop there.
_LAST_RECEIPT_ROW = 30 + RECEIPT_LIMIT + 1
_SHARED_STRING = '{%s}si' % SHEET_MAIN_NS
WORKSHEET_NAMESPACE = 'report-data'


class WorksheetRefused(ValueError):
    """A refusal that follows from the workbook and the question alone (too many rows, an unknown
    sheet, no header), never from a locked file or a read error. worksheet_entry stores these."""


def stamp(path):
    stat = Path(path).stat()
    return (str(Path(path).resolve()), stat.st_mtime_ns, stat.st_ctime_ns, stat.st_size)


@lru_cache(maxsize=128)
def digest_stamp(key):
    with Path(key[0]).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def digest(path):
    return digest_stamp(stamp(path))


def modified(path):
    return datetime.fromtimestamp(Path(path).stat().st_mtime, timezone.utc).isoformat()


def key(value):
    normalized = re.sub(r'[^a-z0-9]+', '_', str(value or '').lower()).strip('_')
    return {'sha_256':'sha256', 'verified_data_through_date':'verified_data_through'}.get(normalized, normalized)


def timestamp(value, day=False):
    if not value:
        return None
    try:
        if day:
            if isinstance(value, datetime):
                return value.date().isoformat()
            if isinstance(value, date):
                return value.isoformat()
            return date.fromisoformat(str(value)).isoformat()
        try:
            result = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        except ValueError:
            result = parsedate_to_datetime(str(value))
        # Do not guess a timezone for legacy dates.
        return result.isoformat() if result.tzinfo else str(value) + ' (timezone unspecified)'
    except (ValueError, TypeError, OverflowError):
        return None


def header_row(sheet):
    """Prefer the earliest wide text row over title blocks and numeric data."""
    candidates = []
    for number, row in enumerate(sheet.iter_rows(max_row=30, max_col=MAX_COLUMNS, values_only=True), 1):
        filled = [v for v in row if v is not None and str(v).strip()]
        if filled and all(isinstance(v, str) for v in filled):
            candidates.append((len(set(filled)), -number))
    return -max(candidates)[1] if candidates else 1


# A header is person-level (a person, or where one gives from) if it has any of these words: the report workbooks'
# Employee List (Person, Person Key, Titles, Zips, Cities / States), Cycle Exceptions, Contributions and Contributor
# Detail (Contributor, Street 1, City, ZIP, Employer, Occupation), and the names, e-mail, phone, address, payee and
# contact columns another workbook may carry. Fail closed: a header not on AGGREGATE_HEADERS and with one of these
# words is refused.
PERSON_HEADER = re.compile(r'\b(?:names?|persons?|people|contributors?|donors?|employees?|payees?|contacts?|e ?mails?'
                           r'|phones?|telephones?|address(?:es)?|streets?|cit(?:y|ies)|zips?|employers?|occupations?'
                           r'|titles?|given|surnames?)\b')
# The headers with one of those words that the real report workbooks (Mentat Complete Files, 2026-10-12) carry on
# aggregate columns only, as header_row() reads them, lower-cased with everything but letters spaced out
# ('Employee R %' is 'employee r'): the four-way tier label (Employee, Middle Management, Chief, Corporate) and the
# per-tier giving columns of Master Overview, Company List and Recipient List (numbers), and the contributor counts
# (whole numbers). 'Name / Recipient' is not here: on the Top Contributor rows of every '* Summary' and '<Company> View'
# sheet it holds the contributor's name.
AGGREGATE_HEADERS = frozenset({'contributor tier', 'employee tier', 'employee', 'employee d', 'employee r',
                               'contributors', 'distinct contributors'})
# What the hosted console does with a worksheet for a viewer not listed in MENTAT_DONOR_ROW_EMAILS. REDACT: the sheet's
# only person-level header is 'Name / Recipient' and it has one 'Section' column, which is the shape of every '* Summary'
# and '<Company> View' sheet (Mentat Complete Files, 2026-10-12). Their Top Contributor rows name an individual in
# 'Name / Recipient' and carry the name key in 'Reference Key'; server.without_contributor_names blanks both on every row
# whose Section is not a recipient or aggregate one. Any other person-level header: REFUSE.
ALLOW, REDACT, REFUSE = 'allow', 'redact', 'refuse'
REDACTED_HEADERS = frozenset({'name recipient', 'reference key'})   # header_text() of the columns that are blanked
SECTION_HEADER = 'section'


def header_text(value):
    return ' '.join(re.findall(r'[a-z]+', re.sub(r'(?<=[a-z])(?=[A-Z])', ' ', str(value)).lower()))


def person_header(value):
    text = header_text(value)
    return text not in AGGREGATE_HEADERS and bool(PERSON_HEADER.search(text))


@lru_cache(maxsize=64)
def _header_candidates(file_stamp, sheet_name):
    """(the row header_row() chooses, the sheet's first 100 rows): the rows a header can be."""
    book = openpyxl.load_workbook(file_stamp[0], read_only=True, data_only=True)
    try:
        sheet = book[sheet_name]
        return header_row(sheet), tuple(sheet.iter_rows(max_row=100, max_col=MAX_COLUMNS, values_only=True))
    finally:
        book.close()


def _header_outcome(cells):
    texts = [header_text(cell) for cell in cells if isinstance(cell, str)]
    people = {text for text in texts if person_header(text)}
    if not people:
        return ALLOW
    return REDACT if people <= {'name recipient'} and texts.count(SECTION_HEADER) == 1 else REFUSE


def person_level_sheet(file_stamp, sheet_name, header=0):
    """ALLOW, REDACT or REFUSE for a sheet read by a viewer who may not see person-level fields: the most
    restrictive outcome over the sheet's own header row (header_row's choice) and the one asked for (worksheet_data
    takes 1-100). Fail closed: a sheet that cannot be read is REFUSE."""
    try:
        chosen, rows = _header_candidates(file_stamp, sheet_name)
        return max((_header_outcome(rows[number - 1]) for number in {chosen, header} if 1 <= number <= len(rows)),
                   key=(ALLOW, REDACT, REFUSE).index, default=ALLOW)
    except Exception:  # noqa: BLE001 - unreadable means refused
        return REFUSE


def records(sheet, limit=RECEIPT_LIMIT):
    start = header_row(sheet)
    rows = sheet.iter_rows(min_row=start, max_row=start + limit + 1, max_col=MAX_COLUMNS, values_only=True)
    headers = [key(v) for v in next(rows, ())]
    result = []
    for row in rows:
        value = {h: v for h, v in zip(headers, row) if h}
        if any(v is not None and v != '' for v in value.values()):
            result.append(value)
        if len(result) > limit:
            raise ValueError('Receipt sheet exceeds the local inspection limit.')
    return result


def _workbook_receipts(path):
    """openpyxl's whole read-only, data-only load: the reference behaviour."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return {name: records(book[name]) for name in RECEIPT_SHEETS
                if name in book.sheetnames}
    finally:
        book.close()


class _StringsUsed(dict):
    """Pre-scan stand-in for the shared-string table; records each index read."""
    def __missing__(self, index):
        self[index] = ''
        return ''


class _StringSubset(dict):
    """The referenced shared strings by table index. An index outside the table
    fails as openpyxl's list does."""
    def __missing__(self, index):
        raise IndexError('list index out of range')


def _streamed_receipts(path):
    """The receipt sheets exactly as _workbook_receipts reads them, parsing only
    those sheets and the shared strings they use (Banking parsed all 123,088
    strings for a 15-row sheet). Rows and cells go through openpyxl's own
    read-only worksheet and cell parser, with openpyxl's workbook, sheet, epoch
    and number-format (date/duration) handling. Returns None for a shape this
    does not mirror, and the caller then uses the whole load."""
    archive = _validate_archive(path)
    try:
        package = Manifest.from_tree(fromstring(archive.read(ARC_CONTENT_TYPES)))
        parser = WorkbookParser(archive, _find_workbook_part(package).PartName[1:])
        parser.parse()
        book = parser.wb
        book._data_only, book._read_only, book._archive = True, True, archive
        present = set(archive.namelist())
        found = {}
        for sheet, rel in parser.find_sheets():
            if rel.target in present and sheet.name in RECEIPT_SHEETS:
                found.setdefault(sheet.name, []).append((rel.target, 'chartsheet' in rel.Type))
        if any(len(entries) != 1 or entries[0][1] for entries in found.values()):
            return None  # a duplicate or a chart sheet under a receipt name
        wanted = {name: found[name][0][0] for name in RECEIPT_SHEETS if name in found}
        used, numeric = _StringsUsed(), False
        for target in wanted.values():
            with archive.open(target) as source:
                for number, cells in WorkSheetParser(source, used, data_only=True).parse():
                    numeric = numeric or any(cell['data_type'] == 'n' and cell['value'] is not None
                                             for cell in cells)
                    if number > _LAST_RECEIPT_ROW:
                        break
        if any(index < 0 for index in used):
            return None  # a negative index reads from the end of the whole table
        strings = _StringSubset()
        part = package.find(SHARED_STRINGS)
        if used and part is not None:
            ceiling, index = max(used), 0
            with archive.open(part.PartName[1:]) as source:
                # openpyxl.reader.strings.read_string_table, stopping at the last index used
                for _, node in iterparse(source):
                    if node.tag == _SHARED_STRING:
                        if index in used:
                            strings[index] = Text.from_tree(node).content.replace('x005F_', '')
                        node.clear()
                        index += 1
                        if index > ceiling:
                            break
        if numeric:
            apply_stylesheet(archive, book)  # date and duration formats, as load_workbook does
        return {name: records(ReadOnlyWorksheet(book, name, target, strings))
                for name, target in wanted.items()}
    finally:
        archive.close()


@lru_cache(maxsize=32)
def embedded_receipts(file_stamp):
    # Anything the streamed read refuses or fails on is answered by the whole
    # load, so an unusual or damaged workbook returns or raises exactly as before.
    try:
        result = _streamed_receipts(file_stamp[0])
    except Exception:
        result = None
    return _workbook_receipts(file_stamp[0]) if result is None else result


def local_reference(value, root):
    """References may only read inside their report/source directory, never UNC."""
    text = str(value or '')
    if not text or text.startswith(('\\\\', '//')) or '://' in text:
        raise ValueError('Receipt reference is not a local file.')
    path = Path(text)
    path = (path if path.is_absolute() else root / path).resolve()
    if not path.is_relative_to(root.resolve()) or str(path).startswith(('\\\\', '//')):
        raise ValueError('Receipt reference is outside the report directory.')
    return path


def csv_receipts(path, expected):
    if not expected or digest(path).lower() != str(expected).lower():
        raise ValueError('Source receipt checksum does not match the report.')
    if path.stat().st_size > 2_000_000:
        raise ValueError('Source receipt exceeds the local inspection limit.')
    with path.open(encoding='utf-8-sig', newline='') as stream:
        return [{key(k): v for k, v in row.items()} for row in csv.DictReader(stream)]


def source_view(row, origin):
    row = {key(k): v for k, v in row.items()}
    # Only an explicit verified date or date+boolean assertion can establish coverage.
    through = timestamp(row.get('verified_data_through'), day=True)
    if not through and str(row.get('data_through_verified', '')).lower() == 'true':
        through = timestamp(row.get('data_through'), day=True)
    return {'source': row.get('source') or row.get('fec_dataset') or row.get('family') or 'Unnamed source',
            'cycle': row.get('cycle'), 'industry': row.get('industry'), 'receipt': origin,
            'sha256': row.get('sha256'), 'retrieved_at': timestamp(row.get('retrieved_at') or row.get('downloaded_at')),
            'checked_at': timestamp(row.get('checked_at')),
            'file_modified_at': None,
            'server_last_modified': timestamp(row.get('server_last_modified') or row.get('last_modified')),
            'verified_data_through': through, 'freshness': 'verified' if through else 'unknown',
            'verification_evidence': row.get('verification_evidence'),
            'note': 'Data-through date asserted by this receipt.' if through else 'No verified data-through date in this receipt.'}


def report_freshness(path, source_root=None):
    path = Path(path)
    result = {'report_modified_at': modified(path), 'sources': [], 'issues': [], 'freshness': 'unknown',
              'note': 'Source dates do not establish complete report coverage. Retrieval, checks and file modification are separate events.'}
    sources = []
    try:
        embedded = embedded_receipts(stamp(path))
        # Prefer the exact hash-bound audit bundle over the worksheet's displayed subset.
        audit = path.with_suffix('.audit')
        manifest_path = audit / 'audit-manifest.json'
        if manifest_path.is_file():
            try:
                manifest = json.loads(manifest_path.read_text(encoding='utf-8-sig'))
                if (manifest.get('workbook_name') != path.name or
                        str(manifest.get('workbook_sha256', '')).lower() != digest(path)):
                    raise ValueError('Companion audit bundle belongs to a different workbook revision.')
                rows = csv_receipts(audit / 'source_freshness.csv', manifest.get('csv_sha256', {}).get('source_freshness.csv'))
                sources.extend(source_view(row, str(audit / 'source_freshness.csv')) for row in rows)
            except (OSError, ValueError, AttributeError) as error:
                result['issues'].append(str(error))
        if not sources:
            sources.extend(source_view(row, 'Workbook / Source Freshness') for row in embedded.get('Source Freshness', []))
        for reference in embedded.get('Source Receipts', []):
            if Path(str(reference.get('audit_file', ''))).name != 'source_freshness.csv':
                continue
            try:
                receipt = local_reference(reference['audit_file'], path.parent)
                rows = csv_receipts(receipt, reference.get('sha256'))
                sources.extend(source_view({**row, 'industry': reference.get('industry')}, str(receipt)) for row in rows)
            except (OSError, ValueError) as error:
                result['issues'].append(str(error))
                sources.append({'source': reference.get('industry') or 'Audit source', 'freshness': 'unknown',
                                'note': 'Referenced receipt is missing, unreadable or does not match its recorded hash.'})
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as error:
        result['issues'].append('Workbook receipts unavailable: ' + str(error))
    # A newer active source must never make an old report look current. Only enrich
    # from the run's pinned source root, and require both source and receipt hashes.
    if source_root and not str(source_root).startswith(('\\\\', '//')):
        root = Path(source_root)
        candidates = [*root.glob('fec_beneficiary/*/source-receipt.json'),
                      *root.glob('fec_committee_master/*/source-receipt.json')]
        for receipt_path in candidates:
            try:
                raw = json.loads(receipt_path.read_text(encoding='utf-8-sig'))
                matching = [s for s in sources if s.get('sha256') and s['sha256'] == raw.get('sha256')
                            and str(s.get('cycle')) == str(raw.get('cycle')) and s['source'] == raw.get('family')]
                if not matching:
                    continue
                source_file = local_reference(raw.get('file'), receipt_path.parent)
                if digest(source_file) != raw.get('sha256'):
                    raise ValueError('Pinned source file no longer matches its receipt: ' + receipt_path.parent.name)
                extra = source_view(raw, str(receipt_path))
                for source in matching:
                    source['file_modified_at'] = modified(source_file)
                    # These are current file metadata; retain the report's original check time.
                    source['retrieved_at'] = source.get('retrieved_at') or extra['retrieved_at']
                    source['local_receipt'] = str(receipt_path)
            except (OSError, ValueError, AttributeError) as error:
                result['issues'].append(str(error))
    result['sources'] = sources
    if not sources:
        result['issues'].append('No report-bound source freshness receipts available.')
    result['verified_source_count'] = sum(s['freshness'] == 'verified' for s in sources)
    return result


def cell_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, str) and value.startswith('='):
        return None
    return value


def numeric(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value if math.isfinite(value) else None
    text = str(value).strip()
    if not re.fullmatch(r'\(?-?\$?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?%?\)?', text):
        return None
    negative = text.startswith('(') and text.endswith(')')
    try:
        number = float(text.strip('()').replace('$', '').replace(',', '').rstrip('%'))
        number = -number if negative else number
        return number / 100 if text.endswith('%') else number
    except ValueError:
        return None


def profile_columns(headers, rows):
    columns = []
    for index, name in enumerate(headers):
        values = [r[index] for r in rows if r[index] is not None and r[index] != '']
        is_id = bool(re.search(r'(^|[ _])(id|zip|code|cycle|year|district)([ _]|$)', name, re.I))
        numbers = sum(numeric(v) is not None for v in values)
        dates = sum(bool(re.fullmatch(r'\d{4}-\d{2}-\d{2}(?:T.*)?', str(v))) for v in values)
        kind = 'number' if values and numbers == len(values) and not is_id else 'date' if values and dates == len(values) else 'category'
        columns.append({'index': index, 'name': name, 'type': kind, 'missing': len(rows) - len(values),
                        'numeric_values': numbers, 'distinct': len(set(str(v) for v in values))})
    return columns


# F2 and F4 (Paul, 2026-09-25). A cycle whose transfer index was not current
# published no transfer-derived rows or values, and three sheets fell back to
# directly funded committees only - an undercount, not a zero. A worksheet view
# of any of them says so in its note, which Chart Studio shows with the data.
TRANSFER_DERIVED_SHEETS = frozenset({'Transfer Paths', 'Allocation Estimates', 'Allocation Coverage',
                                     'Candidate Funding'})
DIRECT_ONLY_SHEETS = frozenset({'Beneficiary Summary', 'Downstream Reach', 'Downstream Measurement'})


# S1 (Paul, 2026-09-26). The Data workbook splits a sheet over its row
# threshold into one sheet per cycle named `<Sheet> <election year>`
# ('Transfer Paths 2022'). Books published before keep `<Sheet> <cycle label>`
# ('Transfer Paths 2021-2022') until they are rebuilt, so both are read. The one
# sheet whose old part name broke Excel's 31 characters splits under a shorter
# base, which maps back to it. tools/build_tidy_workbook.split_sheet_name is the
# writer's side; tests/test_tidy_workbook.py holds the two to each other.
SPLIT_SHEET_SUFFIX = re.compile(r'\s+(?:\d{4}-)?\d{4}$')
SPLIT_SHEET_BASES = {'Network Candidates': 'Network Candidate Activity'}


def split_sheet_base(sheet_name):
    """The sheet a per-cycle part belongs to; any other name comes back as it is.

    A bare cycle label ('2023-2024', an Overview cycle sheet) is not a part:
    nothing would be left of it, so it is not stripped (graph_data.py parses
    those sheets separately).
    """
    name = str(sheet_name or '').strip()
    base = SPLIT_SHEET_SUFFIX.sub('', name).strip()
    if not base or base == name:
        return name
    return SPLIT_SHEET_BASES.get(base, base)


def _audit_book(path):
    """The book whose audit bundle names the not-traced cycles: a Data book's Overview sibling when it exists."""
    path = Path(path)
    overview = path.with_name(re.sub(r' Industry Spending Data\.xlsx$', ' Industry Spending Overview.xlsx', path.name))
    return overview if overview.is_file() else path


def untraced_sheet_notes(path, sheet_name):
    """The not-traced sentences for one worksheet, from its workbook's audit bundle (F2).

    A companion Data workbook shares its Overview's bundle (one build writes
    both). No bundle, a legacy bundle or a current index says nothing.
    """
    base = split_sheet_base(sheet_name)
    if base not in TRANSFER_DERIVED_SHEETS | DIRECT_ONLY_SHEETS:
        return []
    import graph_data  # graph_data imports this module, so it is read on use
    rows = [row for row in graph_data.transfer_index(_audit_book(path)).values()
            if row['status'] != 'current']
    if base in DIRECT_ONLY_SHEETS:
        return [f"Direct only; transfers not traced for {row['cycle']}: the transfer index for {row['cycle']} "
                f"is {row['status']}." for row in rows]
    return [f"{row['note']}: {row['cycle']} has no transfer-derived rows or values here; blank is not zero."
            for row in rows]


def _notes_identity(path, sheet_name):
    """What untraced_sheet_notes reads for this sheet, as a cache dependency: None for a sheet with no
    such note; otherwise the book carrying the audit bundle and the bundle's manifest (None when it has
    none), each by stamp and content. The manifest is bound to its book by hash, so the book counts too."""
    if split_sheet_base(sheet_name) not in TRANSFER_DERIVED_SHEETS | DIRECT_ONLY_SHEETS:
        return None
    import graph_data
    def state(file):
        version = stamp(file)
        return [list(version), digest_stamp(version)]
    book = _audit_book(path)
    manifest = graph_data.audit_manifest_path(book)
    return {'book': state(book), 'manifest': state(manifest) if manifest.is_file() else None}


def worksheet_accounting(sheet_name, headers, rows, row_numbers):
    """Expose all worksheet rows with explicit overlap guards, never a hidden filter."""
    notes = []
    accounting = {'safe_to_sum': True, 'grain': 'worksheet row',
                  'row_role_column': None, 'included_row_roles': [],
                  'independent_dimensions': [], 'basis': 'saved worksheet values'}
    for name, chosen in (('View', 'Industry'), ('Scope', 'Total Giving'), ('Section', 'Company Total')):
        if name not in headers:
            continue
        index = headers.index(name)
        roles = {str(row[index]) for row in rows}
        if chosen in roles and len(roles) > 1:
            accounting['independent_dimensions'].append(index)
            notes.append(f'{name} contains overlapping published summaries. Choose any value on Filters, or keep {name} in the grouping; all worksheet rows remain selectable.')
    if 'Source ID' in headers:
        accounting.update(grain='FEC source observation', basis='included itemized receipts at the named recipient; not verified unique giving')
        notes.append('Use Analysis Amount for included receipts. Raw Amount can include excluded or repeated source observations.')
    # Match the BASE sheet name. The companion Data workbook splits any sheet
    # over its row threshold into per-cycle parts, so the sheet that used to be
    # called exactly 'Transfer Paths' can arrive as 'Transfer Paths 2022' (or,
    # in a book published before S1 of 2026-09-26, 'Transfer Paths 2021-2022').
    # Against the exact-name list those parts fell through and their sums were
    # offered as safe - the same association-path double count, now split three
    # ways and harder to spot. split_sheet_base reads both spellings and leaves
    # a bare cycle label alone (graph_data.py parses those separately).
    base_sheet_name = split_sheet_base(sheet_name) or sheet_name
    if base_sheet_name in ('Beneficiary Bridge', 'Beneficiary Summary', 'Funding Cohorts', 'Allocation Estimates'):
        accounting.update(safe_to_sum=False, reason='This sheet repeats funding or committee activity across association paths. Use the contribution dataset and a named accounting basis.')
    elif base_sheet_name == 'Transfer Paths':
        # F14 (2026-09-26): one row per transfer and cycle, no longer one per funded
        # root; still not a total, because money moving A -> B -> C is in two rows.
        accounting.update(safe_to_sum=False, reason='Each row is one committee-to-committee transfer on the funded roots\' association paths; money that moves through several committees appears at every transfer, so a sum is not a total. Use the contribution dataset and a named accounting basis.')
    columns = profile_columns(headers, rows)
    for column in columns:
        name = column['name']
        if name == 'Amount' and ('Analysis Amount' in headers or 'Include in Analysis' in headers):
            column['accounting'] = {'safe_to_sum':False, 'reason':'Raw observations include excluded or repeated records. Choose Analysis Amount or the contribution dataset.'}
        elif name == 'Analysis Amount':
            column['accounting'] = {'safe_to_sum':True, 'basis':'included_receipts'}
        elif re.search(r'%|percent|coverage|rank|average|median|distinct|contributors',name,re.I):
            column['accounting'] = {'safe_to_sum':False, 'reason':'This is a percentage, rank, average, or precomputed distinct count; summing it does not produce a valid total.'}
        if re.search(r'amount|giving|spending|activity|funding',name,re.I) and column['type']=='number':
            column['unit']='currency'
    return rows,row_numbers,columns,accounting,notes


@lru_cache(maxsize=6)
def worksheet_data(file_stamp, sheet_name='', header=0):
    if not sheet_name:
        # Listing worksheets must not load the potentially huge shared-string
        # table. Only materialize cells when the user chooses a worksheet.
        try:
            with zipfile.ZipFile(file_stamp[0]) as archive:
                tree = ET.fromstring(archive.read('xl/workbook.xml'))
            sheets = [node.get('name') for node in tree.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet')]
            return {'sheets': sheets, 'sha256': digest_stamp(file_stamp)}
        except ET.ParseError as error:
            raise WorksheetRefused('Workbook metadata is malformed.') from error
    book = openpyxl.load_workbook(file_stamp[0], read_only=True, data_only=True)
    try:
        if sheet_name not in book.sheetnames:
            raise WorksheetRefused('Choose a worksheet from this report.')
        sheet = book[sheet_name]
        start = header or header_row(sheet)
        if not 1 <= start <= 100:
            raise WorksheetRefused('Header row must be between 1 and 100.')
        if sheet.max_column and sheet.max_column > MAX_COLUMNS:
            raise WorksheetRefused(f'This worksheet exceeds {MAX_COLUMNS} columns.')
        iterator = sheet.iter_rows(min_row=start, max_col=min(sheet.max_column or MAX_COLUMNS, MAX_COLUMNS), values_only=True)
        raw_headers = list(next(iterator, ()))
        while raw_headers and raw_headers[-1] is None:
            raw_headers.pop()
        if not raw_headers:
            raise WorksheetRefused('No column headers found. Choose another header row.')
        used = set()
        headers = []
        for i, value in enumerate(raw_headers):
            name = str(value).strip() if value is not None else f'Column {i + 1}'
            original = name
            suffix = 2
            while name in used:
                name = f'{original} ({suffix})'
                suffix += 1
            used.add(name)
            headers.append(name)
        rows, row_numbers = [], []
        for number, row in enumerate(iterator, start + 1):
            if number > start + MAX_ROWS:
                raise WorksheetRefused(f'Worksheet exceeds {MAX_ROWS:,} rows. Choose a smaller report subsection; no partial chart was created.')
            values = [cell_value(v) for v in row[:len(headers)]]
            if any(v is not None and v != '' for v in values):
                rows.append(values)
                row_numbers.append(number)
        rows,row_numbers,columns,accounting,notes=worksheet_accounting(sheet_name,headers,rows,row_numbers)
        notes = notes + untraced_sheet_notes(file_stamp[0], sheet_name)
        return {'sheet': sheet_name, 'header_row': start, 'columns': columns,
                'rows': rows, 'row_numbers': row_numbers, 'row_count': len(rows), 'sha256': digest_stamp(file_stamp),
                'accounting': accounting,
                'note': ' '.join(notes + ['Worksheet formulas use their last saved results; empty caches remain missing. The contribution dataset provides the explicit accounting basis and transaction geography.'])}
    finally:
        book.close()


def worksheet_entry(path, sheet_name='', header=0):
    return report_publication.read([path], lambda: _worksheet_entry(path, sheet_name, header))


def _worksheet_entry(path, sheet_name='', header=0):
    """/api/report-data as a response_cache entry: exactly the bytes Handler.send writes for a fresh
    worksheet_data answer (response_cache.json_body), and their gzip, built once per workbook version.

    Key: the workbook's resolved path, the sheet and the header row. Dependencies: the workbook's stamp
    and content SHA-256, what untraced_sheet_notes reads for the sheet (_notes_identity: the audit
    manifest and the book it is bound to), and the code identity of this module and every application
    module it imports, with the openpyxl that reads the cells. The sheet list (no sheet) and each
    sheet's data are separate keys in one byte-bounded cache, so they do not push each other out.

    A refusal that follows from the workbook (WorksheetRefused: over 50,000 rows, an unknown sheet) is
    stored the same way as an answer, and raised again as the same WorksheetRefused. Any other error
    (a locked or unreadable file, a broken archive) propagates and stores nothing. Nothing is stored
    when the workbook or the manifest moved while the answer was built; that answer is served once."""
    version = stamp(path)
    key_parts = {'report': version[0], 'sheet': sheet_name, 'header': header}

    def reads():
        return stamp(path), _notes_identity(path, sheet_name)

    try:
        before = reads()
        content = digest_stamp(version)
        steady = before[0] == version and reads() == before
        code = response_cache.code_identity(__name__, extra={'openpyxl': openpyxl.__version__}) if steady else None
    except OSError:
        before, content, steady, code = (None, None), None, False, None
    deps = {'stamp': list(version), 'content_sha256': content, 'notes': before[1], 'code': code or 'this process only'}
    moved = []

    def compute():
        # lru_cache's __wrapped__: the uncached read; this cache replaces worksheet_data's own.
        try:
            value = worksheet_data.__wrapped__(version, sheet_name, header)
        except WorksheetRefused as error:
            value = {'error': str(error)}
        try:
            moved.append(reads() != before)
        except OSError:
            moved.append(True)
        return value

    entry = response_cache.get_or_compute(WORKSHEET_NAMESPACE, key_parts, deps, compute,
                                          persist=lambda value: not moved[-1], memory=steady, disk=code is not None)
    # Every refusal body is small and starts {"error": - no worksheet answer does (they start {"sheet).
    if entry.body_bytes < 1024 and entry.body.startswith(b'{"error":'):
        raise WorksheetRefused(json.loads(entry.body)['error'])
    return entry
