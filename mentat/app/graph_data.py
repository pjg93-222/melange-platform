"""Guided report graphs. Read saved report facts without changing classifications."""
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from functools import lru_cache, wraps
import csv
import json
from pathlib import Path
import posixpath
import re
import sqlite3
import threading
import time
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import collection_gaps
import downstream
import evidence_csv
import native_scan
from view_cache import ViewCache, file_signature

from report_insights import digest, digest_stamp, key, local_reference, stamp

TIERS = {'Chief': 'Chief', 'Middle Management': 'Middle-Management', 'Employee': 'Employees'}
# The order the employee-level legend reads in, largest level first on every
# saved report so far; levels_by_company keeps each company's levels in it.
LEVEL_ORDER = (TIERS['Employee'], TIERS['Middle Management'], TIERS['Chief'])
CATEGORIES = ['Super PAC', 'Trade Association', 'JFC', 'Company PAC', 'Leadership PAC',
              'Carey', 'Other PAC', 'Candidate', 'Party Committee', 'Unclassified', 'Conduit']
AUTHORIZED = {'Principal campaign committee', 'Authorized candidate committee'}
MAX_SOURCE_ROWS = 1000000
# The drawer's sides, in the order the console offers them. direction_totals
# is emitted for every one of them, including the empty ones: a drawer that
# showed only the non-empty sides could not say that Support is zero, which is
# the answer the reader came for.
#
# Paul, 2026-09-23, on the drawer's "Contributions / Other" segment: "It
# should be two separate segments". The workbook records one direction,
# 'Contribution / Other', for every FEC transaction code that is not an
# independent expenditure (R/fec_beneficiary.R fec_downstream_direction), and
# the same rows carry that code (Transaction Type; a saved snapshot's
# disbursement_type). So the recorded bucket is split by its code, never by
# guess: Contribution is the codes whose recorded Activity Type names a
# contribution (fec_downstream_relationship: 24K "Contribution to
# nonaffiliated committee", 24P "Contribution to possible federal candidate",
# 24Z "In-kind contribution to registered filer"); Other is every other
# recorded code (24C coordinated party expenditure, 24H, 24R, ...). A row
# whose source records no code cannot be split, so it stays what it was
# recorded as, 'Contribution / Other' - and that is also what an old
# permalink and the saved estimates (which carry no code) still ask for.
COMBINED_DIRECTION = 'Contribution / Other'
CONTRIBUTION_CODES = frozenset({'24K', '24P', '24Z'})
ACTIVITY_DIRECTIONS = ('Contribution', 'Other', 'Support', 'Oppose', COMBINED_DIRECTION, 'Unknown')
# The sides the recorded 'Contribution / Other' bucket is drawn as, for
# recorded_bucket below.
RECORDED_SIDES = ('Contribution', 'Other', COMBINED_DIRECTION)
SINGLEFLIGHT_WAIT_SECONDS = 120


def singleflight_cache(maxsize):
    """Share an in-progress file read without blocking reads of other files.

    A caller that has waited SINGLEFLIGHT_WAIT_SECONDS for the read in progress
    stops sharing it and reads the file itself, so one stalled read delays the
    requests behind it instead of blocking them forever.
    """
    def decorate(load):
        cached = lru_cache(maxsize=maxsize)(load)
        condition = threading.Condition()
        reading = set()

        @wraps(load)
        def read(file_stamp):
            deadline = time.monotonic() + SINGLEFLIGHT_WAIT_SECONDS
            leader = True
            with condition:
                while file_stamp in reading:
                    remaining = deadline - time.monotonic()
                    if remaining <= 0 or not condition.wait(remaining):
                        # The read in progress is stalled. Read independently, as
                        # plain lru_cache would, and leave its token to its owner.
                        leader = False
                        break
                if leader:
                    reading.add(file_stamp)
            try:
                return cached(file_stamp)
            finally:
                if leader:
                    with condition:
                        reading.remove(file_stamp)
                        condition.notify_all()

        read.cache_clear = cached.cache_clear
        read.cache_info = cached.cache_info
        return read
    return decorate


def cycle(value):
    match = re.fullmatch(r'(?:\d{4}[-–])?(\d{4})', str(value or '').strip())
    return match[1] if match else ''


def amount(value):
    try:
        result = Decimal(str(value))
        if result.is_finite():
            return result
    except (InvalidOperation, ValueError):
        pass
    raise ValueError('A required report amount is missing or nonnumeric; no partial graph was created.')


def money(value):
    return float(round(value, 2))


def table_rows(sheet, header=5, required=()):
    iterator = sheet.iter_rows(min_row=header, values_only=True)
    headers = [key(v) for v in next(iterator, ())]
    if not set(required) <= set(headers):
        raise ValueError(f'{sheet.title}: required report columns are missing.')
    for number, row in enumerate(iterator, 1):
        if number > MAX_SOURCE_ROWS:
            raise ValueError(f'{sheet.title}: report exceeds the local graph reading limit.')
        record = {h: v for h, v in zip(headers, row) if h}
        if any(v is not None and v != '' for v in record.values()):
            yield record


def activity_xml_rows(path, metadata, sheet_name='Downstream Activity', required=None, optional=False,
                      optional_fields=()):
    """Stream selected activity columns; ignore styles, dates and hyperlinks."""
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    relationship = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    required = set(required or {'cycle','source_committee_id','candidate_id','candidate_officeholder',
                'candidate_party','support_oppose','net_included_amount'})
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        sheet = next((s for s in workbook.iter(ns+'sheet') if s.get('name') == sheet_name), None)
        if sheet is None:
            if optional:
                return
            raise ValueError('This report has no saved downstream activity.')
        relations = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        relation = next((r for r in relations if r.get('Id') == sheet.get(relationship)), None)
        if relation is None or relation.get('TargetMode') == 'External':
            raise ValueError('Downstream worksheet is not stored in this report.')
        target = relation.get('Target','')
        target = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
        strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            with archive.open('xl/sharedStrings.xml') as stream:
                for _,element in ET.iterparse(stream,events=('end',)):
                    if element.tag == ns+'si':
                        strings.append(''.join(t.text or '' for t in element.iter(ns+'t')))
                        element.clear()
        def value(cell):
            if cell.get('t') == 'inlineStr':
                return ''.join(t.text or '' for t in cell.iter(ns+'t'))
            cached = cell.find(ns+'v')
            if cached is None or cached.text is None:
                return None
            return strings[int(cached.text)] if cell.get('t') == 's' else cached.text
        columns = None
        with archive.open(target) as stream:
            for _,element in ET.iterparse(stream,events=('end',)):
                if element.tag != ns+'row':
                    continue
                number = int(element.get('r','0'))
                if number > MAX_SOURCE_ROWS+5:
                    raise ValueError('Downstream report exceeds the local graph reading limit.')
                if number == 3:
                    metadata['note'] = next((value(c) for c in element if c.get('r') == 'A3'), '')
                elif number == 5:
                    headers = {c.get('r','').rstrip('0123456789'):key(value(c)) for c in element if c.tag == ns+'c'}
                    if not required <= set(headers.values()):
                        raise ValueError(f'{sheet_name}: required report columns are missing.')
                    columns = {col:field for col,field in headers.items() if field in required |
                               {'unverified_version_rows','unresolved_version_rows'} | set(optional_fields)}
                elif number > 5:
                    if columns is None:
                        raise ValueError('Downstream Activity: header row is missing.')
                    row = {columns[col]:value(c) for c in element
                           if c.tag == ns+'c' and (col := c.get('r','').rstrip('0123456789')) in columns}
                    if any(v is not None and v != '' for v in row.values()):
                        yield row
                element.clear()


def report_transfers(file_stamp, audited=False):
    """The workbook's transfers: its own sheet, else its pinned full source. `audited` also lets the audited
    per-cycle books fill in; only the Graph view asks (the flow and the sidecar warmers read the served build)."""
    return _transfers(file_stamp, audited)[0]


def transfer_basis(file_stamp, audited=False):
    """[{'cycles': [end years], 'text': a sentence}] saying where the transfers that were not the workbook's own
    complete Transfer Paths sheet were read from; [] when they all were. The Graph's trace note carries the text."""
    return _transfers(file_stamp, audited)[1]


def _transfers(file_stamp, audited):
    return _report_transfers((file_stamp, file_signature(transfer_files(file_stamp[0], audited)), audited))


# FIX 6 (P2): the transfers behind a workbook, in this order. (1) Its own Transfer Paths sheet. (2) The full source
# pinned to it, when that sheet is capped (transfer_source). (3) The audited per-cycle book of the same industry and
# cycle, <Industry> Industry Spending Overview <yyyy-yyyy>.xlsx in the reports folder: its audit CSV, else its pinned
# tidy store (both full), else its own sheet when that is complete. A source that is not the workbook's own is always
# said (transfer_basis); a full source is never presented as the capped sheet it replaces.
TRANSFER_CSV = 'transfer_paths.csv'
SHEET_RECORDS = re.compile(r'([\d,]+)\s+records;')
SHEET_SHOWN = re.compile(r'([\d,]+)\s+records;\s*([\d,]+)\s+shown')


def sheet_transfers(path):
    """(edges, note, rows): the workbook's own Transfer Paths sheet. ([], '', 0) when it has none."""
    metadata = {}
    result = []
    required = {'cycle','source_committee','destination_committee','reported_net_transfer','traversal_status'}
    count = 0
    # F14 (2026-09-26): the sheet lists each transfer once per cycle (earlier books
    # repeated it once per funded root; both shapes carry these five columns). A
    # transfer seen twice must agree on its amount, as in the full-source read below.
    seen = {}
    for row in activity_xml_rows(path, metadata, TRANSFER_PATHS, required, optional=True):
        if not row.get('source_committee'):
            continue
        count += 1
        # Candidate gifts are already in activity. Never expand a terminal gift.
        if row.get('traversal_status') == 'candidate_committee':
            continue
        edge = {'cycle':cycle(row['cycle']), 'source':str(row['source_committee']),
                'destination':str(row['destination_committee']),
                'amount':amount(row['reported_net_transfer'])}
        identity = (edge['cycle'], edge['source'], edge['destination'])
        if identity in seen:
            if seen[identity] != edge['amount']:
                raise ValueError('Transfer Paths rows disagree on one transfer amount.')
            continue
        seen[identity] = edge['amount']
        result.append(edge)
    return result, str(metadata.get('note') or ''), count


def sheet_note(path):
    """The count line of a workbook's Transfer Paths sheet ('230,075 records; 10,000 shown. ...'), '' without one."""
    metadata = {}
    try:
        for _ in activity_xml_rows(path, metadata, TRANSFER_PATHS, {'cycle'}, optional=True):
            break
    except ValueError:
        return ''
    return str(metadata.get('note') or '')


def sheet_is_partial(note, count):
    expected = SHEET_RECORDS.search(note)
    return 'first ' in note.lower() or bool(expected and int(expected[1].replace(',','')) != count)


def sheet_shown(note):
    """(shown, total) as the sheet's count line states them, or None."""
    found = SHEET_SHOWN.search(note)
    return (int(found[2].replace(',', '')), int(found[1].replace(',', ''))) if found else None


def sheet_capped(note):
    """Whether the count line says the sheet lists fewer rows than the workbook holds (or does not say)."""
    counts = sheet_shown(note)
    return 'first ' in note.lower() or counts is None or counts[0] != counts[1]


def sheet_sentence(note, rows, name, whose):
    """'The Transfer Paths sheet shows the first 10,000 of 230,075 rows; all 230,075 were read from <file>.'"""
    counts = sheet_shown(note)
    if counts is None:
        return f'{whose} Transfer Paths sheet is a partial display; the full list was read from {name}.'
    shown, total = counts
    read = ('the full list was read' if rows is None else f'all {total:,} were read' if rows == total
            else f'{rows:,} rows were read')
    return f'{whose} Transfer Paths sheet shows the first {shown:,} of {total:,} rows; {read} from {name}.'


def unique_transfers(rows):
    """({(cycle, source, destination): amount}, rows read) for a full source: a transfer once per cycle, a
    candidate gift never expanded, committee ids checked, and one amount per transfer."""
    unique, count = {}, 0
    for row in rows:
        count += 1
        if row.get('status') == 'candidate_committee':
            continue
        identity = (cycle(row['cycle']), str(row['source']), str(row['destination']))
        if not identity[0] or not all(re.fullmatch(r'C\d{8}', value) for value in identity[1:]):
            raise ValueError('The full transfer source has invalid committee or cycle identifiers.')
        value = amount(row['amount'])
        if identity in unique and unique[identity] != value:
            raise ValueError('Full report paths disagree on one transfer amount.')
        unique[identity] = value
    return unique, count


def edge_list(unique):
    return [dict(cycle=k[0], source=k[1], destination=k[2], amount=v) for k,v in sorted(unique.items())]


@singleflight_cache(maxsize=4)
def _report_transfers(version):
    """(edges, basis) for one workbook; `dependencies` is the signature of every file they may come from, and
    `use_audited` whether the audited per-cycle books may fill in (the Graph view) or not (everything else)."""
    file_stamp, dependencies, use_audited = version
    path = file_stamp[0]
    result, note, count = sheet_transfers(path)
    if not has_transfer_paths(file_stamp):
        return read_audited(file_stamp, audited_sources(path), dependencies) if use_audited else ([], [])
    if not sheet_is_partial(note, count):
        return result, []
    payload, manifest = transfer_source(path)
    if not payload or not manifest.is_file() or not payload.is_file():
        # No pinned full source: the audited books may supply every cycle of this workbook, else the display
        # stays partial and is refused rather than shown as the whole.
        audited = audited_sources(path) if use_audited else {}
        if audited and set(audited) == {cycle(label) for label in graph_schema(file_stamp)}:
            return read_audited(file_stamp, audited, dependencies)
        raise ValueError('Transfer Paths is a partial display; its pinned full report source is unavailable.')
    import contribution_data as source_data
    proof = json.loads(manifest.read_text(encoding='utf-8-sig'))
    if (proof.get('database_sha256') != source_data._workbook_source(file_stamp)
            or not proof.get('data_sha256') or digest(payload) != proof['data_sha256']):
        raise ValueError('Full transfer source does not match the published report fingerprint.')
    payload_stamp = stamp(payload)
    native = native_scan.tables(payload_stamp, [], payload_table='transfer_paths')
    if native is not None:
        table = native['tables']['transfer_paths']
        rows = (dict(zip(table['headers'], row)) for row in table['rows'])
    else:
        rows = source_data._payload_rows(payload, 'transfer_paths')
    unique, _ = unique_transfers(rows)
    if (stamp(payload) != payload_stamp or stamp(path) != file_stamp
            or file_signature(transfer_files(path, use_audited)) != dependencies):
        raise ValueError('Full report source changed while reading transfers.')
    # The payload's raw row count is not known here (native_scan collapses repeats and candidate gifts first),
    # so the sentence does not claim one.
    return edge_list(unique), [{'cycles': sorted({k[0] for k in unique}),
                                'text': sheet_sentence(note, None, payload.name, 'The')}]


report_transfers.cache_clear = _report_transfers.cache_clear
report_transfers.cache_info = _report_transfers.cache_info


def audited_reports_directory():
    """Where the audited per-cycle books live: the repository's reports folder, as tools/audit_collection_outputs.py
    and the collection-gap tests name it. The console's report folder holds the served books, never these."""
    import contribution_data as source_data
    return source_data.PROJECT.parent / 'reports'


def audited_book(path, label):
    """The audited per-cycle book of the Overview `path` names for one cycle label ('2025-2026'), whether or not
    it exists yet; None for a workbook that is not a published Overview or is that book itself."""
    from report_identity import identity
    info = identity(Path(path).name)
    if info is None or info['kind'] != 'overview':
        return None
    book = audited_reports_directory() / f"{info['industry']} Industry Spending Overview {label}.xlsx"
    return None if book.resolve() == Path(path).resolve() else book


def audited_store(book, label):
    """(tidy staging store, manifest) pinned to an audited per-cycle book's Source DB, or (None, None)."""
    import contribution_data as source_data
    from report_identity import identity, lineage_paths
    try:
        fingerprint = source_data._workbook_source(stamp(book))
    except (ValueError, OSError, KeyError, zipfile.BadZipFile):
        return None, None
    cache = source_data.PROJECT / '.cache/industry-report' / source_data._slug(identity(Path(book).name)['industry'])
    manifest, payload = lineage_paths(cache, fingerprint, label)
    return tidy_beside(payload), manifest


def audited_source(path, label):
    """How the audited per-cycle book supplies one cycle's transfers, or None. A full source (its audit CSV, then
    its pinned tidy store) comes before the book's Transfer Paths sheet, which counts only when it is complete."""
    book = audited_book(path, label)
    if book is None or not book.is_file():
        return None
    found = {'cycle': cycle(label), 'label': label, 'book': book}
    csv_path = audit_manifest_path(book).with_name(TRANSFER_CSV)
    if csv_path.is_file():
        return {**found, 'kind': 'csv', 'path': csv_path}
    store, manifest = audited_store(book, label)
    if store is not None and store.is_file() and manifest.is_file():
        return {**found, 'kind': 'store', 'path': store, 'manifest': manifest}
    if has_transfer_paths(stamp(book)) and not sheet_capped(sheet_note(book)):
        return {**found, 'kind': 'sheet', 'path': book}
    return None


def audited_sources(path):
    """{cycle end year: source} for the cycles of an Overview that an audited per-cycle book can supply. Whether
    they are used is the caller's: a workbook's own complete Transfer Paths sheet comes first."""
    sources = (audited_source(path, label.replace('–', '-')) for label in graph_schema(stamp(path)))
    return {source['cycle']: source for source in sources if source}


def transfer_files(path, audited=True):
    """Every file this workbook's transfers may be read from, present or not, for the dependency signatures of
    the caches that hold transfers or Graph views: a changed CSV, or a book that appears later, is a new key.
    With audited=False, the pinned full source alone (what the flow and the sidecar warmers read)."""
    files = [p for p in transfer_source(path) if p is not None]
    if not audited:
        return files
    try:
        labels = graph_schema(stamp(path))
    except (OSError, ValueError, KeyError, zipfile.BadZipFile):
        return files
    for label in labels:
        label = label.replace('–', '-')
        book = audited_book(path, label)
        if book is None:
            continue
        audit = audit_manifest_path(book)
        files += [book, audit, audit.with_name(TRANSFER_CSV)]
        if book.is_file():
            files += [p for p in audited_store(book, label) if p is not None]
    return files


def audit_csv_rows(book, path):
    """The rows of an audit CSV, once its bundle manifest has bound it to its workbook and its own hash."""
    try:
        manifest = audit_manifest(book)
    except ValueError as error:
        raise ValueError('The audited transfer paths do not match their audited workbook.') from error
    if manifest is None:
        raise ValueError('The audited transfer paths have no audit manifest to check them against.')
    if (manifest.get('csv_sha256') or {}).get(path.name) != digest(path):
        raise ValueError('The audited transfer paths file is missing or changed.')
    with path.open(encoding='utf-8-sig', newline='') as stream:
        for number, row in enumerate(csv.DictReader(stream), 1):
            if number > MAX_SOURCE_ROWS:
                raise ValueError('The audited transfer paths exceed the local graph reading limit.')
            yield row


def store_rows(store, manifest, year):
    """One cycle's rows of the tidy store's transfer_path_detail, once the store matches its published fingerprint."""
    version = _pinned_detail_store(stamp(store), stamp(manifest))
    connection = sqlite3.connect(Path(version[0]).resolve().as_uri() + '?mode=ro', uri=True)
    try:
        connection.execute('PRAGMA query_only=ON')
        tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if 'transfer_path_detail' not in tables:
            raise ValueError('This build predates the transfer path detail table (F14).')
        for row in connection.execute(
                'SELECT "cycle", "source_committee", "destination_committee", "reported_net_transfer", '
                '"traversal_status" FROM "transfer_path_detail" WHERE "cycle" = ?', (int(year),)):
            yield dict(zip(('cycle', 'source', 'destination', 'amount', 'status'), row))
    finally:
        connection.close()


def audited_build(book):
    """The audited book's Source DB fingerprint, shortened as a report's source note prints it."""
    import contribution_data as source_data
    try:
        return source_data._workbook_source(stamp(book))[:8] + '…'
    except (ValueError, OSError, KeyError, zipfile.BadZipFile):
        return 'unrecorded'


def audited_transfers(source):
    """([edge], sentence) read from one audited source, the way the pinned full source is read."""
    kind, book, label, year = source['kind'], source['book'], source['label'], source['cycle']
    name = source['path'].name
    if kind == 'sheet':
        edges, note, rows = sheet_transfers(book)
        if sheet_is_partial(note, rows):
            raise ValueError(f'The audited {label} Transfer Paths sheet is a partial display and has no full source.')
        unique = {(e['cycle'], e['source'], e['destination']): e['amount'] for e in edges}
        where = f', from its Transfer Paths sheet ({rows:,} rows).'
    else:
        unique, rows = unique_transfers(audit_csv_rows(book, source['path']) if kind == 'csv'
                                        else store_rows(source['path'], source['manifest'], year))
        note = sheet_note(book)
        where = f', from {name} ({rows:,} rows).'
    if any(k[0] != year for k in unique):
        raise ValueError(f'The audited {label} transfer paths include another cycle.')
    head = f'Transfers for {label} were read from the audited {label} workbook (build {audited_build(book)}), not the served workbook'
    if kind != 'sheet' and note and sheet_capped(note):
        return edge_list(unique), head + '. ' + sheet_sentence(note, rows, name, 'Its')
    return edge_list(unique), head + where


def read_audited(file_stamp, sources, dependencies):
    """(edges, basis) from the audited sources, cycle by cycle; ([], []) when there are none."""
    edges, basis = [], []
    for year in sorted(sources):
        found, sentence = audited_transfers(sources[year])
        edges += found
        basis.append({'cycles': [year], 'text': sentence})
    if sources and (stamp(file_stamp[0]) != file_stamp or file_signature(transfer_files(file_stamp[0])) != dependencies):
        raise ValueError('An audited transfer source changed while reading transfers.')
    return edges, basis


def transfer_source(path):
    """The full export pinned to this workbook, never a newer collection."""
    path = Path(path)
    if not path.name.endswith(' Industry Spending Overview.xlsx'):
        return None, None
    import contribution_data as source_data
    try:
        fingerprint = source_data._workbook_source(stamp(path))
    except (ValueError, OSError, KeyError, zipfile.BadZipFile):
        return None, None
    industry = path.name.removesuffix(' Industry Spending Overview.xlsx')
    folder = source_data.PROJECT / '.cache/industry-report' / source_data._slug(industry)
    return (folder / f'industry_data-{fingerprint[:16]}.json',
            folder / f'manifest-{fingerprint[:16]}.json')


# F14 (Paul approved option (a), 2026-09-26): the Transfer Paths sheet lists each
# transfer once per cycle. The full expansion - every funded root x committee it
# reaches x transfer out of it - is the transfer_path_detail table of the tidy
# staging store the same build pinned (its manifest's stage_sha256), never a sheet.
TRANSFER_PATH_DETAIL_COLUMNS = ('cycle_label', 'cycle', 'funded_committee', 'source_committee',
                                'destination_committee', 'reported_net_transfer', 'depth',
                                'representative_path', 'traversal_status', 'via_carryover')


def tidy_beside(payload):
    """The tidy staging store a build writes beside its pinned payload."""
    return payload.with_name('tidy-' + payload.name.removeprefix('industry_data-').removesuffix('.json') + '.sqlite')


def transfer_detail_source(path):
    """(tidy staging store, manifest) pinned to this workbook, or (None, None)."""
    payload, manifest = transfer_source(path)
    if payload is None:
        return None, None
    return tidy_beside(payload), manifest


@lru_cache(maxsize=8)
def _pinned_detail_store(store_stamp, manifest_stamp):
    proof = json.loads(Path(manifest_stamp[0]).read_text(encoding='utf-8-sig'))
    if not proof.get('stage_sha256') or digest(Path(store_stamp[0])) != proof['stage_sha256']:
        raise ValueError('The transfer path detail store does not match the published report fingerprint.')
    return store_stamp


def transfer_path_detail(path, year, root=None, source=None, destination=None, limit=1000):
    """Routes from the full root x committee x transfer expansion for one workbook and cycle.

    Filter by funded root, or by transfer (source and/or destination). Read-only,
    from the store the workbook's build pinned; a store that does not match that
    fingerprint, or a build older than F14, is refused rather than read.
    """
    store, manifest = transfer_detail_source(path)
    if store is None or manifest is None or not store.is_file() or not manifest.is_file():
        raise ValueError('This workbook has no pinned transfer path detail store.')
    version = _pinned_detail_store(stamp(store), stamp(manifest))
    clauses, values = ['"cycle" = ?'], [int(cycle(year) or 0)]
    for column, value in (('funded_committee', root), ('source_committee', source),
                          ('destination_committee', destination)):
        if value:
            clauses.append(f'"{column}" = ?')
            values.append(str(value))
    fields = ', '.join(f'"{name}"' for name in TRANSFER_PATH_DETAIL_COLUMNS)
    connection = sqlite3.connect(Path(version[0]).resolve().as_uri() + '?mode=ro', uri=True)
    try:
        connection.execute('PRAGMA query_only=ON')
        tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if 'transfer_path_detail' not in tables:
            raise ValueError('This build predates the transfer path detail table (F14).')
        rows = connection.execute(
            f'SELECT {fields} FROM "transfer_path_detail" WHERE {" AND ".join(clauses)} '
            f'ORDER BY "funded_committee", "depth", "source_committee", "destination_committee" LIMIT ?',
            (*values, int(limit))).fetchall()
    finally:
        connection.close()
    return [dict(zip(TRANSFER_PATH_DETAIL_COLUMNS, row)) for row in rows]


@lru_cache(maxsize=32)
def workbook_sheets(file_stamp):
    """Every sheet name, in workbook order, from the workbook's own index."""
    try:
        # Workbook metadata is sufficient. Loading openpyxl here materialized the
        # shared-string table for every report just to list its sheet names.
        with zipfile.ZipFile(file_stamp[0]) as archive:
            tree = ET.fromstring(archive.read('xl/workbook.xml'))
            ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
            return tuple(s.get('name', '') for s in tree.iter(ns+'sheet'))
    except ET.ParseError as error:
        raise ValueError('Workbook sheet metadata is malformed.') from error


@lru_cache(maxsize=32)
def graph_schema(file_stamp):
    return [name for name in workbook_sheets(file_stamp) if re.fullmatch(r'\d{4}[-–]\d{4}', name)]


TRANSFER_PATHS = 'Transfer Paths'


def has_transfer_paths(file_stamp):
    """server.report_quality's `transfer_paths` flag: 'Transfer Paths' in the sheet names.

    report_quality answers it by loading the workbook in openpyxl (measured
    2026-09-25: 21 s for the 40 MB Banking book, 1.7 s for Crypto), too slow for
    a drawer, so Graph reads the same fact from the sheet index. The Graph tests
    hold the two to the same answer.
    """
    return TRANSFER_PATHS in workbook_sheets(file_stamp)


def committee_id(value):
    match = re.fullmatch(r'(?:FEC:)?(C\d{8})', str(value or ''))
    return match[1] if match else ''


@singleflight_cache(maxsize=8)
def report_facts(file_stamp):
    """Recipient facts are a complete partition; Master/Employee facts repeat them."""
    cycles = graph_schema(file_stamp)
    projection = ('company','employee_tier','amount','include_in_analysis','recipient_category',
                  'fact_type','person_key','contributor','recipient_key','committee_id','recipient')
    scanned = native_scan.tables(file_stamp, [
        {'name':'Employee List','header':5,'columns':['person_key','person','company']},
        {'name':'PAC & Candidate Map','header':5},
        {'name':'Recipient List','header':5,'columns':['cycle','recipient_key','recipient']},
        {'name':'Company List','header':5,'columns':['company']},
        *[{'name':name,'header':4,'columns':projection} for name in cycles]])
    book = None if scanned else openpyxl.load_workbook(file_stamp[0], read_only=True, data_only=True)
    sheetnames = scanned['sheets'] if scanned else book.sheetnames
    def rows(name, header=5, required=()):
        if scanned is None:
            yield from table_rows(book[name], header, required)
            return
        table = scanned['tables'][name]
        if not set(required) <= set(table['headers']):
            raise ValueError(f'{name}: required report columns are missing.')
        for row in table['rows']:
            yield dict(zip(table['headers'], row))
    try:
        maps, names, companies = {}, {}, set()
        people, person_names = [], {}
        if 'Employee List' in sheetnames:
            for row in rows('Employee List', required=('person_key', 'person', 'company')):
                if row.get('person_key') and row.get('person'):
                    person_names[(str(row['company']), str(row['person_key']))] = str(row['person'])
        if 'PAC & Candidate Map' in sheetnames:
            for row in rows('PAC & Candidate Map', required=('cycle', 'committee_id')):
                lookup = (cycle(row.get('cycle')), str(row.get('committee_id') or ''))
                if lookup in maps and maps[lookup] != row:
                    raise ValueError('Conflicting candidate links in this report; direct rankings are unavailable.')
                maps[lookup] = row
        if 'Recipient List' in sheetnames:
            for row in rows('Recipient List', required=('cycle', 'recipient_key', 'recipient')):
                names.setdefault((cycle(row['cycle']), row['recipient_key']), row['recipient'])
        if 'Company List' in sheetnames:
            companies.update(str(r['company']) for r in rows('Company List', required=('company',)) if r.get('company'))
        facts, masters, partitions = {}, defaultdict(Decimal), defaultdict(Decimal)
        sheets = graph_schema(file_stamp)
        if not sheets:
            raise ValueError('This report has no supported cycle receipt tables. Use Advanced for custom worksheet graphs.')
        for sheet_name in sheets:
            year = cycle(sheet_name)
            fact_mode = False
            for row in rows(sheet_name, 4, ('company','employee_tier','amount','include_in_analysis')):
                if str(row.get('include_in_analysis', '')).lower() != 'yes':
                    continue
                company = str(row.get('company') or 'Unknown company')
                companies.add(company)
                tier = str(row.get('employee_tier') or 'Unknown')
                category = str(row.get('recipient_category') or 'Unclassified')
                scope = (year, company, tier, category)
                fact_type = row.get('fact_type')
                if (fact_type == 'Employee' or not fact_type) and row.get('person_key'):
                    person = str(row['person_key'])
                    people.append({'cycle':year, 'company':company, 'tier':tier, 'person_id':person,
                                   'person':person_names.get((company, person), str(row.get('contributor') or person)),
                                   'amount':amount(row['amount'])})
                if fact_type:
                    fact_mode = True
                    if fact_type not in ('Master','Employee','Recipient'):
                        raise ValueError('Unknown report fact type; no partial graph was created.')
                    if fact_type == 'Master':
                        masters[scope] += amount(row['amount'])
                    if fact_type != 'Recipient':
                        continue
                value = amount(row['amount'])
                partitions[scope] += value
                recipient = str(row.get('recipient_key') or row.get('committee_id') or row.get('recipient') or 'Unknown recipient')
                identifier = committee_id(row.get('committee_id')) or committee_id(recipient)
                mapping = maps.get((year, identifier), {})
                candidate = str(mapping.get('candidate_id_s') or '')
                linked = (category == 'Candidate' and mapping.get('recipient_category') == 'Candidate'
                          and mapping.get('relationship') in AUTHORIZED and re.fullmatch(r'[HSP][A-Z0-9]{8}', candidate))
                # A workbook built before 2026-09-08 files ActBlue and WinRed under Super PAC. They are named
                # here, on the recipient fact, and only as a label: `scope` above still holds the workbook's
                # own category, which is what the Master rows (they carry no committee) reconcile against.
                shown = evidence_csv.conduit_category(identifier, category)
                identity = (year, company, tier, shown, recipient)
                if identity not in facts:
                    facts[identity] = {'cycle':year, 'company':company, 'tier':tier, 'category':shown,
                        'recipient_id':identifier or recipient,
                        'recipient':str(mapping.get('recipient_committee') or names.get((year,recipient)) or row.get('recipient') or recipient),
                        'candidate_id':candidate if linked else '',
                        'candidate':str(mapping.get('direct_candidate_sponsor') or candidate) if linked else '',
                        'party':str(mapping.get('direct_party') or '') if linked else '', 'amount':Decimal(0)}
                facts[identity]['amount'] += value
            if fact_mode:
                keys = {s for s in masters.keys() | partitions.keys() if s[0] == year}
                if any(abs(masters[s] - partitions[s]) > Decimal('.01') for s in keys):
                    raise ValueError('Recipient facts do not reconcile to the report totals; no partial graph was created.')
        if stamp(file_stamp[0]) != file_stamp:
            raise ValueError('The workbook changed while it was being read. Refresh Graph.')
        return {'facts':list(facts.values()), 'people':people, 'companies':sorted(companies), 'cycles':sorted(cycle(s) for s in sheets),
                'committees':[{'cycle':year, 'id':identifier, 'label':str(row.get('recipient_committee') or identifier),
                               'category':evidence_csv.conduit_category(identifier, str(row.get('recipient_category') or 'Unclassified'))}
                              for (year,identifier),row in maps.items()],
                'sha256':digest_stamp(file_stamp), 'report':Path(file_stamp[0]).name}
    finally:
        if book is not None:
            book.close()


@singleflight_cache(maxsize=8)
def _report_activity(options):
    file_stamp, include_filings = options
    try:
        metadata = {}
        groups = defaultdict(Decimal)
        unverified, unresolved = defaultdict(int), defaultdict(int)
        filings = defaultdict(list)
        labels = {}
        count = 0
        required = ('cycle','source_committee_id','candidate_id','candidate_officeholder',
                    'candidate_party','support_oppose','net_included_amount')
        # Transaction Type is optional: a workbook written before the column
        # existed still reads, and its rows carry no code (see activity_sides).
        parts = defaultdict(lambda: defaultdict(lambda: [Decimal(0), []]))
        scanned = native_scan.tables(file_stamp, [{'name':'Downstream Activity','header':5,
            'columns':[*required,'unverified_version_rows','unresolved_version_rows','transaction_type',
                       *(['filing_records'] if include_filings else [])]}])
        if scanned:
            table = scanned['tables'].get('Downstream Activity')
            if table is None:
                raise ValueError('This report has no saved downstream activity.')
            if not set(required) <= set(table['headers']):
                raise ValueError('Downstream Activity: required report columns are missing.')
            metadata['note'] = table['note']
            source_rows = (dict(zip(table['headers'], row)) for row in table['rows'])
        else:
            source_rows = activity_xml_rows(file_stamp[0], metadata,
                                            optional_fields=('transaction_type',
                                                             *(('filing_records',) if include_filings else ())))
        for row in source_rows:
            count += 1
            identity = (cycle(row['cycle']), str(row['source_committee_id']), str(row['candidate_id'] or ''), str(row['support_oppose'] or 'Unknown'))
            value = amount(row['net_included_amount'])
            groups[identity] += value
            # The same amount again under its FEC transaction code, so the one
            # recorded bucket can be split by code without a second read. The
            # group itself (its identity and amount) is unchanged.
            part = parts[identity][transaction_code(row.get('transaction_type'))]
            part[0] += value
            unverified[identity] += int(row.get('unverified_version_rows') or 0)
            unresolved[identity] += int(row.get('unresolved_version_rows') or 0)
            if include_filings and row.get('filing_records'):
                try:
                    records = json.loads(row['filing_records'])
                except (TypeError, ValueError):
                    records = []
                if isinstance(records, list):
                    records = [record for record in records if isinstance(record, dict)]
                    filings[identity].extend(records)
                    part[1].extend(records)
            labels[identity] = (str(row.get('candidate_officeholder') or identity[2] or 'Unidentified candidate'), str(row.get('candidate_party') or ''))
        note = str(metadata.get('note') or '')
        expected = re.search(r'condensed to ([\d,]+)', note)
        if expected and int(expected[1].replace(',','')) != count:
            raise ValueError('Downstream worksheet is a partial display; full activity rankings are unavailable.')
        if stamp(file_stamp[0]) != file_stamp:
            raise ValueError('The workbook changed while it was being read. Refresh Graph.')
        return [{'cycle':k[0], 'committee':k[1], 'candidate_id':k[2], 'direction':k[3],
                 'candidate':labels[k][0], 'party':labels[k][1], 'amount':v,
                 'filings':filings[k],
                 'transaction_parts':{code:{'amount':total, 'filings':records}
                                      for code,(total,records) in parts[k].items()},
                 'unverified_version_rows':unverified[k], 'unresolved_version_rows':unresolved[k]} for k,v in groups.items()]
    except (ET.ParseError, IndexError, KeyError, zipfile.BadZipFile) as error:
        raise ValueError('Saved downstream activity could not be read.') from error


def report_activity(file_stamp, include_filings=False):
    return _report_activity((file_stamp, include_filings))


def transaction_code(value):
    """An FEC transaction code as recorded ('24K'), or '' when none is."""
    return str(value or '').strip().upper()


def activity_side(code):
    """The drawer side a recorded 'Contribution / Other' amount belongs to."""
    if code in CONTRIBUTION_CODES:
        return 'Contribution'
    return 'Other' if code else COMBINED_DIRECTION


def activity_sides(rows):
    """Paul, 2026-09-23: "It should be two separate segments".

    Each recorded 'Contribution / Other' row becomes one row per side its
    transaction codes fall on - Contribution, Other, and, for amounts whose
    source recorded no code, 'Contribution / Other' itself - each carrying
    that side's own amount and filings. Nothing is re-attributed: the sides
    of one row add up to its recorded amount to the cent, and every other
    row (Support, Oppose, Unknown, estimates) passes through as it is.

    unverified_version_rows and unresolved_version_rows count the recorded
    row's filing observations, not a side's: the reader does not record them
    per code, so each side carries its row's counts as they are, and a total
    of them is taken over the recorded rows, before this split (graph_view),
    never over the sides - a row split in two would count its observations
    twice (r4 code review F1).
    """
    result = []
    for row in rows:
        parts = row.get('transaction_parts')
        if row.get('direction') != COMBINED_DIRECTION or not parts:
            result.append(row)
            continue
        sides = {}
        for code, part in parts.items():
            side = sides.setdefault(activity_side(code), {'amount':Decimal(0), 'filings':[], 'codes':set()})
            side['amount'] += part['amount']
            side['filings'].extend(part['filings'])
            if code:
                side['codes'].add(code)
        if sum((side['amount'] for side in sides.values()), Decimal(0)) != row['amount']:
            raise ValueError('Saved activity by transaction type does not add up to its recorded amount.')
        for name in ACTIVITY_DIRECTIONS:
            if name in sides:
                side = sides[name]
                result.append({**row, 'direction':name, 'amount':side['amount'], 'filings':side['filings'],
                               'transaction_types':sorted(side['codes'])})
    return result


def recorded_bucket(side_groups, buckets):
    """The recorded 'Contribution / Other' bucket, each candidate once.

    Paul, 2026-09-23: "It should be two separate segments". A candidate can
    have both a contribution (24K) and other activity (24C) from the same
    committee, so it is in both segments' counts, and the counts add up to
    more than the candidates behind them (ActBlue, Banking 2025-2026: 1,000
    + 45 for 1,005 candidates). This is the reconciliation the drawer shows
    under them: the distinct identified candidates across the sides, how many
    are on more than one side, and the bucket's amount, which the sides'
    amounts add up to to the cent. side_groups are the sides' candidate
    groups (candidates_by_direction), so a candidate is on a side exactly
    when that side's count includes it.
    """
    sides = defaultdict(int)
    for name in RECORDED_SIDES:
        for candidate in side_groups.get(name) or ():
            if candidate['id'] != 'unidentified-candidate':
                sides[candidate['id']] += 1
    rows = [r for name in RECORDED_SIDES for r in buckets.get(name) or ()]
    return {'direction':COMBINED_DIRECTION, 'sides':list(RECORDED_SIDES), 'candidates':len(sides),
            'on_more_than_one_side':sum(1 for count in sides.values() if count > 1),
            'amount':money(sum((r['amount'] for r in rows), Decimal(0))),
            'unidentified_amount':money(sum((r['amount'] for r in rows if not r['candidate_id']), Decimal(0)))}


report_activity.cache_clear = _report_activity.cache_clear
report_activity.cache_info = _report_activity.cache_info


# F4 (Paul approved 2026-09-25): a surface says when money is not traced. No
# industry had an estimate bundle when this was written, and the Estimates view
# answered with an empty panel; it now says what is missing and what makes it.
# Lane F's attributed Flow answers 409 with the same sentence.
NO_ESTIMATE_BUNDLE = 'No estimate bundle for this industry yet; rebuild the workbook.'


def audit_manifest_path(path):
    path = Path(path)
    return path.with_suffix('.audit') / 'audit-manifest.json'


def audit_manifest(path):
    """The workbook's audit-bundle manifest, bound to this workbook; None when it has no bundle."""
    path = Path(path)
    try:
        manifest = json.loads(audit_manifest_path(path).read_text(encoding='utf-8-sig'))
    except FileNotFoundError:
        return None
    if (not isinstance(manifest, dict) or manifest.get('workbook_name') != path.name
            or manifest.get('workbook_sha256') != digest(path)):
        raise ValueError('The estimate audit bundle does not match this workbook.')
    return manifest


def not_traced(year, status, recorded=''):
    """('Not traced: the transfer index for <cycle> is <status>', its rebuild command).

    F2 (lane A, tools/enrich_downstream.py not_traced_note) records one note per
    cycle that is not current: 'Not traced: the transfer index for 2024 is
    stale; rebuild with python ... --cycles 2024'. The words are shown; the
    command goes in a tooltip. A current index has neither.
    """
    if status == 'current':
        return '', ''
    text = f'Not traced: the transfer index for {year} is {status}'
    recorded = str(recorded or '')
    command = recorded.split('; rebuild with ', 1)[1].strip() if '; rebuild with ' in recorded else ''
    return text, command


def transfer_index(path):
    """{cycle: {cycle, status, note, command}} as the workbook's audit bundle names them (F2).

    enrich_downstream.py writes transfer_index_status ({"2024": "stale", ...};
    current | stale | missing | unresolved) and transfer_index_notes into the
    audit manifest. Any status but current published no transfer-derived number
    for that cycle. A workbook with no bundle, a bundle written before the
    status existed, or one that does not match the workbook names nothing here,
    so nothing is claimed either way.
    """
    try:
        manifest = audit_manifest(path)
    except (OSError, ValueError):
        return {}
    if not manifest:
        return {}
    statuses = manifest.get('transfer_index_status') or {}
    notes = manifest.get('transfer_index_notes') or {}
    if not isinstance(statuses, dict) or not isinstance(notes, dict):
        return {}
    result = {}
    for year, status in statuses.items():
        key_ = cycle(year)
        if not key_ or not isinstance(status, str) or not status:
            continue
        text, command = not_traced(key_, status, notes.get(str(year)))
        result[key_] = {'cycle': key_, 'status': status, 'note': text, 'command': command}
    return result


def rebuild_title(rows):
    """The tooltip for not-traced rows: each distinct rebuild command, on its own line."""
    commands = sorted({row['command'] for row in rows if row.get('command')})
    return '\n'.join('Rebuild with: ' + command for command in commands)


def estimate_receipt(path):
    """Require the full hash-bound CSV, never the 10,000-row worksheet preview."""
    path = Path(path)
    manifest_path = audit_manifest_path(path)
    # The reader-facing note, not the OS error with a local path in it (r4 code
    # review: estimates on a report written without its audit bundle answered
    # "[Errno 2] No such file or directory: ...").
    manifest = audit_manifest(path)
    if manifest is None:
        raise ValueError(NO_ESTIMATE_BUNDLE)
    receipt = local_reference('allocation_estimates.csv', manifest_path.parent)
    if manifest.get('csv_sha256', {}).get(receipt.name) != digest(receipt):
        raise ValueError('The estimate receipt is missing or changed.')
    return stamp(receipt)


@lru_cache(maxsize=4)
def estimate_rows(file_stamp):
    groups = defaultdict(Decimal)
    labels = {}
    with Path(file_stamp[0]).open(encoding='utf-8-sig',newline='') as stream:
        for i, row in enumerate(csv.DictReader(stream),1):
            if i > MAX_SOURCE_ROWS:
                raise ValueError('Estimate receipt exceeds the local graph reading limit.')
            identity = (cycle(row.get('cycle')), row.get('root',''), row.get('candidate_id',''), row.get('direction','Unknown'))
            if not identity[2]:
                raise ValueError('Estimate receipt lacks candidate identifiers.')
            groups[identity] += amount(row.get('estimated_amount'))
            labels[identity] = row.get('candidate_name') or identity[2]
    return [{'cycle':k[0],'committee':k[1],'candidate_id':k[2],'direction':k[3],
             'candidate':labels[k],'party':'','amount':v} for k,v in groups.items()]


def grouped(rows, id_field, label_field=None):
    totals, labels = defaultdict(Decimal), {}
    for row in rows:
        identifier = row[id_field]
        if not identifier and id_field == 'candidate_id':
            identifier = 'unidentified-candidate'
        if not identifier:
            continue
        totals[identifier] += row['amount']
        labels[identifier] = str(row.get(label_field or id_field) or identifier)
    return [{'id':identifier,'label':labels[identifier],'amount':money(value)}
            for identifier,value in sorted(totals.items(), key=lambda item:(-item[1],str(item[0]))) if value != 0]


def levels_by_company(rows):
    """Each company's own non-Corporate giving split by level, for the "By
    company" bars under Giving by employee level (board M4). The rows are the
    ones `levels` is grouped from, so Corporate never enters; a company's
    amount is its non-Corporate total, unlike companies[].amount, which counts
    Corporate giving when the scope is all giving. Every company lists the
    legend's levels in one order, a missing one at 0, then any other level it
    has (Unknown), so its shares always read in the same place."""
    totals = defaultdict(lambda: defaultdict(Decimal))
    for row in rows:
        totals[row['company']][row['level']] += row['amount']
    result = []
    for company, levels in totals.items():
        order = [*LEVEL_ORDER, *sorted(level for level in levels if level not in LEVEL_ORDER)]
        result.append({'company':company, 'amount':money(sum(levels.values(), Decimal(0))),
                       'levels':[{'id':level,'label':level,'amount':money(levels.get(level, Decimal(0)))}
                                 for level in order]})
    return sorted(result, key=lambda item:(-item['amount'], item['company']))


# F4 (Paul approved 2026-09-25). A committee with no saved snapshot is followed
# only through its workbook's Transfer Paths sheet (downstream.network's
# report_edges). With no such sheet - or with F2's stale or missing transfer
# index, which publishes none of that cycle's paths - nothing past that
# committee is traced, and Graph says so in words instead of a generic note.
# FIX 6: "no such sheet" now means no source at all - neither the sheet, nor its
# pinned full source, nor an audited per-cycle book (report_transfers).
NO_TRANSFER_PATHS = 'no saved network, and this workbook has no Transfer Paths sheet'


def trace_coverage(stamps, reports, index):
    """{cycle: (some source covers its transfers, F2's row for the cycle or None)}: the covering workbook's own
    Transfer Paths sheet, or an audited per-cycle book for the cycle (audited_sources)."""
    covered = {}
    for version, report in zip(stamps, reports):
        sheet = has_transfer_paths(version)
        audited = {} if sheet else audited_sources(version[0])
        for year in report['cycles']:
            covered[year] = (sheet or year in audited, index.get(year))
    return covered


def transfer_trace(network, roots, covered, basis=()):
    """Which committees' transfers this view could not follow, and the words for it.

    roots are the view's funded committee/cycles; network['missing'] the ones
    reached without a saved snapshot. `basis` are the sentences saying where the
    transfers it did follow were read from (transfer_basis), after the untraced
    ones. Returns {'committees': [...], 'root_traced': False when every root is
    untraced, 'note': the sentences, 'command': the rebuild commands for a tooltip}.
    """
    rows = []
    for year, committee in network['missing']:
        if year not in covered:
            continue
        sheet, index = covered[year]
        if not sheet:
            reason, status, command = NO_TRANSFER_PATHS, None, ''
        elif index and index['status'] != 'current':
            reason = f"no saved network, and the transfer index for {year} is {index['status']}"
            status, command = index['status'], index['command']
        else:
            continue
        rows.append({'cycle': year, 'committee': committee, 'root': (year, committee) in roots,
                     'status': status, 'reason': reason, 'command': command})
    committees = sorted({committee for _, committee in roots})
    single = len(committees) == 1
    sentences = []
    at_root = defaultdict(list)
    for row in rows:
        if row['root']:
            at_root[row['reason']].append(row)
    for reason, group in sorted(at_root.items()):
        if single:
            untraced = sorted({row['cycle'] for row in group})
            scope = '' if untraced == sorted({year for year, _ in roots}) else ' for ' + ' and '.join(untraced)
            sentences.append(f'Transfers from this committee are not traced{scope}: {reason}.')
        else:
            count = len({row['committee'] for row in group})
            sentences.append(f'Transfers from {count:,} of the {len(committees):,} funded committees '
                             f'are not traced: {reason}.')
    further = defaultdict(set)
    for row in rows:
        if not row['root']:
            further[row['reason']].add(row['committee'])
    for reason, group in sorted(further.items()):
        count = len(group)
        sentences.append(f"Transfers from {count:,} other {'committee' if count == 1 else 'committees'} "
                         f"in this network are not traced: {reason}.")
    sentences.extend(basis)
    untraced_roots = {(row['cycle'], row['committee']) for row in rows if row['root']}
    return {'committees': rows,
            'root_traced': not (single and bool(roots) and set(roots) <= untraced_roots),
            'note': ' '.join(sentences), 'command': rebuild_title(rows)}


def not_traced_problems(payload):
    """Every way a Graph answer breaks contracts/graph-not-traced.json, as plain
    sentences; [] when it holds. Applies to candidate views (activity, estimates);
    a direct answer carries none of the fields and has nothing to check."""
    problems = []
    if not isinstance(payload, dict):
        return ['the answer is not an object']
    mode = (payload.get('filters') or {}).get('mode')
    if mode == 'direct':
        return [f'a direct answer carries {field}' for field in ('transfer_trace', 'transfer_index',
                                                                  'candidate_note_title') if field in payload]
    for field in ('transfer_trace', 'transfer_index', 'candidate_note_title'):
        if field not in payload:
            problems.append(f'the answer has no {field}')
    title = payload.get('candidate_note_title') or ''
    note = payload.get('candidate_note') or ''
    rows = payload.get('transfer_index') or []
    for row in rows:
        if not isinstance(row, dict) or not cycle(row.get('cycle')) or not row.get('status'):
            problems.append('a transfer_index row states no cycle and status')
            continue
        expected = not_traced(row['cycle'], row['status'])[0]
        if row['status'] == 'current' and (row.get('note') or row.get('command')):
            problems.append(f"a current index carries a note for {row['cycle']}")
        elif row['status'] != 'current' and row.get('note') != expected:
            problems.append(f"the {row['status']} index for {row['cycle']} is not named Not traced")
    if mode == 'estimates':
        selected = (payload.get('filters') or {}).get('cycle') or ''
        scope = [selected] if selected else list((payload.get('options') or {}).get('cycles') or [])
        for row in rows:
            if isinstance(row, dict) and row.get('status') != 'current' and row.get('cycle') in scope:
                if row.get('note') and row['note'] not in note:
                    problems.append(f"the answer does not say {row['cycle']} is not traced")
                if row.get('command') and row['command'] not in title:
                    problems.append(f"the tooltip does not carry the rebuild command for {row['cycle']}")
    trace = payload.get('transfer_trace')
    if mode == 'activity' and trace is not None:
        if not isinstance(trace, dict) or not isinstance(trace.get('committees'), list):
            return problems + ['transfer_trace must list its committees']
        for row in trace['committees']:
            reason = str(row.get('reason') or '') if isinstance(row, dict) else ''
            if reason != NO_TRANSFER_PATHS and not re.fullmatch(
                    r'no saved network, and the transfer index for \d{4} is \w+', reason):
                problems.append(f"the untraced committee {row.get('committee') if isinstance(row, dict) else row} "
                                f"states no reason")
            elif row.get('command') and row['command'] not in title:
                problems.append(f"the tooltip does not carry the rebuild command for {row.get('committee')}")
        said = str(trace.get('note') or '')
        if bool(trace['committees']) != ('not traced' in said):
            problems.append('transfer_trace.note does not say what the committees list holds')
        if trace.get('root_traced') is False and not said.startswith('Transfers from this committee are not traced'):
            problems.append('root_traced is false but the note does not say so')
        if said and said not in note:
            problems.append('the candidate note does not carry the trace sentence')
    return problems


_views = ViewCache()


def graph_view(paths, filters):
    paths = tuple(map(Path, paths))
    # Saved estimate bundles have their own manifest/hash protocol. Keep that
    # gate on every estimate read rather than treating a view cache as proof.
    if filters.get('mode') == 'estimates':
        return _graph_view(paths, filters)
    def dependencies():
        snapshots = list(downstream.ROOT.glob('*/snapshots.json')) if filters.get('mode') == 'activity' else []
        full_sources = [p for path in paths for p in transfer_files(path)] if filters.get('mode') == 'activity' else []
        bundles = [audit_manifest_path(path) for path in paths] if filters.get('mode') == 'activity' else []
        return file_signature([*paths, *snapshots, *full_sources, *bundles, *map(collection_gaps.quality_csv, paths)])
    before = dependencies()
    selection = tuple((k,str(filters.get(k) or '')) for k in
                      ('company','tier','cycle','category','recipient','latest','funding','mode','direction'))
    def build():
        value = _graph_view(paths, filters)
        if dependencies() != before:
            raise ValueError('A source changed while preparing Graph. Refresh.')
        return value
    # Keep order and multiplicity: [A, A] is invalid overlapping evidence,
    # not the same request as [A], even though both have the same file set.
    return _views.get((tuple(str(path.resolve()) for path in paths), before, selection), build)


def _graph_view(paths, filters):
    stamps = [stamp(path) for path in paths]
    reports = [report_facts(version) for version in stamps]
    years = [year for report in reports for year in report['cycles']]
    if len(years) != len(set(years)):
        raise ValueError('These reports overlap election cycles. Choose a run with one report per industry/cycle.')
    facts = [r for report in reports for r in report['facts']]
    companies = sorted({c for report in reports for c in report['companies']})
    selected = {k:str(filters.get(k) or '') for k in ('company','tier','cycle','category','recipient')}
    if filters.get('latest') == '1' and years:
        selected['cycle'] = max(years)
    funding = str(filters.get('funding') or 'all')
    mode = str(filters.get('mode') or 'direct')
    direction = str(filters.get('direction') or 'Support')
    if funding not in ('all','personal','corporate') or mode not in ('direct','activity','estimates'):
        raise ValueError('Choose a valid giving scope and candidate view.')
    if selected['company'] and selected['company'] not in companies:
        raise ValueError('Choose a company in this industry.')
    if selected['cycle'] and selected['cycle'] not in years:
        raise ValueError('Choose an election cycle in this report.')
    if selected['tier'] and selected['tier'] not in TIERS.values() and selected['tier'] != 'Unknown':
        raise ValueError('Choose Chief, Middle-Management or Employees.')
    if selected['tier'] and funding == 'corporate':
        raise ValueError('Employee levels apply to personal giving, not corporate giving.')
    scope = [r for r in facts if (not selected['company'] or r['company'] == selected['company'])
             and (not selected['cycle'] or r['cycle'] == selected['cycle'])
             and (funding == 'all' or (r['tier'] == 'Corporate') == (funding == 'corporate'))
             and (not selected['tier'] or (r['tier'] != 'Corporate' and TIERS.get(r['tier'],'Unknown') == selected['tier']))]
    categories = sorted({r['category'] for r in scope})
    recipients = {r['recipient_id']:r['recipient'] for r in scope if not selected['category'] or r['category'] == selected['category']}
    filtered = [r for r in scope if (not selected['category'] or r['category'] == selected['category'])
                and (not selected['recipient'] or r['recipient_id'] == selected['recipient'])]
    direct = [r for r in filtered if r['candidate_id']]
    organizations = [r for r in filtered if r['category'] != 'Candidate']
    employee_rows = [{**r, 'level':TIERS.get(r['tier'], 'Unknown')}
                     for r in filtered if r['tier'] != 'Corporate']
    unresolved = sum((r['amount'] for r in filtered if r['category'] == 'Candidate' and not r['candidate_id']), Decimal(0))
    candidate_rows = direct
    buckets = None
    available, note = True, 'Direct contributions to uniquely linked principal or authorized candidate committees.'
    ties = []
    trace, index_rows, note_title = None, [], ''
    if mode != 'direct':
        candidate_rows = []
        try:
            if mode == 'estimates' and (selected['company'] or selected['tier'] or funding != 'all'):
                raise ValueError('Saved estimates cover the whole industry. Select all companies, all levels and all giving to view them.')
            funded = defaultdict(Decimal)
            for r in filtered:
                funded[r['cycle'],r['recipient_id']] += r['amount']
            eligible = {k for k,v in funded.items() if v > 0}
            if direction not in ACTIVITY_DIRECTIONS:
                raise ValueError('Choose a recorded activity direction.')
            buckets = {name: [] for name in ACTIVITY_DIRECTIONS}
            network_keys = eligible
            activity_rows = None
            if mode == 'activity':
                activity_rows, network = downstream.merge_activity(
                    [r for version in stamps for r in report_activity(version)], eligible,
                    report_edges=[r for version in stamps for r in report_transfers(version, audited=True)])
                # Split after the merge: merge_activity compares whole recorded
                # partitions between a snapshot and the workbook, and those
                # still read 'Contribution / Other' on both sides. The recorded
                # rows are kept for the filing-relationship counts below: a
                # split row carries its row's counts on each side.
                recorded_rows = activity_rows
                activity_rows = activity_sides(activity_rows)
                network_keys = network['keys']
            index = {}
            if mode == 'activity':
                for path in paths:
                    index.update(transfer_index(path))
                index_rows = sorted(index.values(), key=lambda row: row['cycle'])
                reached = {key[0] for key in network['keys']}
                trace = transfer_trace(network, eligible, trace_coverage(stamps, reports, index),
                                       [b['text'] for version in stamps for b in transfer_basis(version, audited=True)
                                        if reached & set(b['cycles'])])
                note_title = trace['command']
                rows = activity_rows
            else:
                rows = [r for path in paths for r in estimate_rows(estimate_receipt(path))]
                for path in paths:
                    index.update(transfer_index(path))
                index_rows = sorted(index.values(), key=lambda row: row['cycle'])
                # F2: a cycle whose transfer index was not current published no
                # estimates. Say so; its absence is never shown as $0.
                scope_cycles = [selected['cycle']] if selected['cycle'] else sorted(set(years))
                untraced = [index[year] for year in scope_cycles if year in index and index[year]['status'] != 'current']
                note_title = rebuild_title(untraced)
                if untraced:
                    gone = {row['cycle'] for row in untraced}
                    rows = [r for r in rows if r['cycle'] not in gone]
                    if gone >= set(scope_cycles):
                        raise ValueError(' '.join(row['note'] + '.' for row in untraced))
            for r in rows:
                if (r['cycle'],r['committee']) in network_keys and r['amount'] != 0 and r['direction'] in buckets:
                    buckets[r['direction']].append(r)
            candidate_rows = buckets[direction]
            note = ('Saved observed activity across reachable committees, counted once per committee/cycle. Routes and activity are limited to saved report coverage. This is not money attributed to the selected company or employee level; company views can share the same PAC activity.'
                    if mode == 'activity' else 'Saved conditional industry estimates, not traced company dollars. Support, opposition and contribution/other activity remain separate.')
            if mode == 'activity' and network['units']:
                note = downstream.NOTE + f" {len(network['units'])} saved committee/cycles; {len(network['missing'])} without a saved snapshot. Snapshots supply whole schedule partitions; workbook activity is used only for committees without snapshots."
                if network.get('measurement_comparisons'):
                    note += ' Snapshot and workbook measurements differ; their amounts are shown separately below and are not reconciled.'
            if mode == 'activity' and trace['note']:
                note += ' ' + trace['note']
            if mode == 'estimates':
                partial = [index[year] for year in scope_cycles if year in index and index[year]['status'] != 'current']
                if partial:
                    note += ' ' + ' '.join(row['note'] + '.' for row in partial)
            if mode == 'activity' and not activity_rows and network['missing']:
                raise ValueError('Downstream activity is unavailable: no saved activity for the selected committee/cycles. Missing evidence is not zero spending.'
                                 + (' ' + trace['note'] if trace['note'] else ''))
            if mode == 'activity':
                ties = downstream.relationships(activity_rows, eligible, network)
                # Each observation once: summed over the recorded rows, not the
                # sides they were split into (r4 code review F1).
                uncertain = sum(r.get('unverified_version_rows',0) for r in recorded_rows)
                unresolved_versions = sum(r.get('unresolved_version_rows',0) for r in recorded_rows)
                if uncertain or unresolved_versions:
                    note += f' Filing relationships: {uncertain} unverified observations retained; {unresolved_versions} unresolved observations in the saved source. Amounts are provisional, not certified unique spending.'
        except (OSError, ValueError, KeyError, AttributeError) as error:
            # The partial buckets go with it: a refused read has no sides to
            # report, and four rows of '0 candidates' would be read as a
            # measurement rather than as the absence of one.
            available, note, buckets = False, str(error), None
    # Activity is what the committees that received this industry's money did, not
    # what the industry gave. Each identified candidate carries the industry's own
    # direct amount for the same filters (the direct view's row, 0 when it has
    # none), on the ranking and on every side group alike, so a reader never takes
    # the reached amount for the giving. A recipient category or a recipient
    # committee narrows the direct rows to none (a PAC is not a candidate
    # committee): that 0 would say the industry gave nothing, so those views carry
    # no direct amount rather than a false one.
    own = ({row['id']: row['amount'] for row in grouped(direct,'candidate_id','candidate')}
           if mode == 'activity' and not selected['category'] and not selected['recipient'] else None)
    def candidate_groups(rows):
        result = grouped(rows,'candidate_id','candidate')
        parties = defaultdict(set)
        for row in rows:
            if row.get('party'):
                parties[row['candidate_id']].add(row['party'])
        for candidate in result:
            candidate['party'] = ' / '.join(sorted(parties[candidate['id']]))
            if own is not None and candidate['id'] != 'unidentified-candidate':
                candidate['direct'] = own.get(candidate['id'], 0.0)
        return result
    chart_candidates = candidate_groups(candidate_rows) if available else []
    side_groups = ({name:candidate_groups(rows) for name,rows in buckets.items()}
                   if mode == 'activity' and buckets is not None else None)
    if any(stamp(path) != version for path,version in zip(paths,stamps)):
        raise ValueError('A report changed while building the graph. Refresh Graph.')
    return {'filters':{**selected,'funding':funding,'mode':mode,'direction':direction},
            'options':{'companies':companies,'cycles':sorted(set(years)), 'categories':categories,
                       'recipients':[{'id':i,'label':n} for i,n in sorted(recipients.items(),key=lambda p:p[1])],
                       'unknown_tiers':any(r['tier'] not in TIERS and r['tier'] != 'Corporate' for r in facts)},
            'totals':{'giving':money(sum((r['amount'] for r in filtered),Decimal(0))),
                      'direct':money(sum((r['amount'] for r in direct),Decimal(0))),
                      'organizations':money(sum((r['amount'] for r in organizations),Decimal(0))),
                      'unresolved_candidate_links':money(unresolved),
                      'recipients':len({r['recipient_id'] for r in filtered}), 'receipt_groups':len(filtered)},
            'categories':grouped(filtered,'category'), 'organizations':grouped(organizations,'recipient_id','recipient'),
            'companies':grouped(filtered,'company'),
            'levels':grouped(employee_rows,'level'),
            'levels_by_company':levels_by_company(employee_rows),
            'candidates':chart_candidates, 'candidate_available':available, 'candidate_note':note,
            # F4: the not-traced facts behind the notes, and the tooltip that
            # holds a rebuild command. Candidate views only.
            **({'candidate_note_title':note_title, 'transfer_index':index_rows,
                'transfer_trace':trace} if mode != 'direct' else {}),
            'downstream_edges':network['edges'] if mode == 'activity' and available else [],
            'pac_candidate_relationships':ties,
            'downstream_sources':downstream.sources(network) if mode == 'activity' and available else [],
            'measurement_comparisons':network.get('measurement_comparisons',[]) if mode == 'activity' and available else [],
            'candidates_by_direction':side_groups if side_groups is not None else {},
            **({'direction_totals':[
                {'direction':name,
                 'candidates':len([r for r in grouped(buckets.get(name) or [],'candidate_id','candidate') if r['id'] != 'unidentified-candidate']),
                 'unidentified_amount':money(sum((r['amount'] for r in buckets.get(name) or [] if not r['candidate_id']),Decimal(0))),
                 'amount':money(sum((r['amount'] for r in buckets.get(name) or []),Decimal(0))),
                 # The recorded FEC codes behind a split side, so its label can
                 # say what it holds (graph.js drawDirections).
                 'transaction_types':sorted({code for r in buckets.get(name) or [] for code in r.get('transaction_types') or ()})}
                for name in ACTIVITY_DIRECTIONS],
                # The two segments' counts reconciled (recorded_bucket).
                'contribution_other_union':recorded_bucket(side_groups, buckets)}
              if mode == 'activity' and buckets is not None
              else {'direction_totals':[], 'contribution_other_union':None} if mode == 'activity' else {}),
            'candidate_details':[{'candidate_id':r['candidate_id'],'candidate':r['candidate'],
                                 'company':r.get('company',''), 'tier':TIERS.get(r.get('tier'),r.get('tier','')),
                                 'recipient':r.get('recipient',r.get('committee','')), 'cycle':r['cycle'], 'amount':money(r['amount'])}
                                 for r in candidate_rows] if available else [],
            'reports':[{'name':r['report'],'sha256':r['sha256']} for r in reports],
            # FIX 4: what each report's own Data Quality check says the collection fell short of, for the
            # cycles shown here; the page names a cell only when its company x cycle scope touches it.
            'collection_gaps':collection_gaps.scoped(collection_gaps.merge(map(collection_gaps.workbook_gaps, paths)), years),
            'note':'Included net giving from saved report receipts; corporate and personal giving remain distinct. Totals may include unverified attribution or receipt roles and are not certified unique giving. Levels follow existing occupation/title rules, not a verified company hierarchy. Employees includes unrecognized titles. Amounts retain the saved report’s rules; code changes do not rewrite historical workbooks.'}
