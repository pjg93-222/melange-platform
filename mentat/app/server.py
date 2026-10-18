"""Local Mentat upload and collection console; no collector runs on import."""
from __future__ import annotations

import argparse
import base64
import csv
import datetime
from functools import lru_cache
from contextlib import contextmanager
import gzip
import hashlib
import importlib
import io
import json
import os
from pathlib import Path
import re
import secrets
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs
import uuid
import zipfile

PROJECT = Path(__file__).resolve().parents[1]
# Capture the loaded server version once. Reading the file on each health
# request would incorrectly certify an old process after an on-disk update.
SERVER_CODE_SHA256 = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(PROJECT / "tools"))
# report_insights and graph_data import openpyxl (0.4 s warm, seconds on a cold
# disk), which no request needs until a workbook is opened. They are imported
# on first use; the names stay module attributes so tests can patch them.
def report_freshness(*args, **kwargs):
    from report_insights import report_freshness as implementation
    return implementation(*args, **kwargs)


def worksheet_data(*args, **kwargs):
    from report_insights import worksheet_data as implementation
    return implementation(*args, **kwargs)


def worksheet_entry(*args, **kwargs):
    from report_insights import worksheet_entry as implementation
    return implementation(*args, **kwargs)


def stamp(*args, **kwargs):
    from report_insights import stamp as implementation
    return implementation(*args, **kwargs)


def graph_schema(*args, **kwargs):
    from graph_data import graph_schema as implementation
    return implementation(*args, **kwargs)


def graph_view(*args, **kwargs):
    from graph_data import graph_view as implementation
    return implementation(*args, **kwargs)


# comparison_data imports graph_data and report_insights (openpyxl) at its top;
# same treatment: imported on first use, names kept as module attributes.
def comparison_subjects(*args, **kwargs):
    from comparison_data import subjects as implementation
    return implementation(*args, **kwargs)


def comparison_flows(*args, **kwargs):
    from comparison_data import flows as implementation
    return implementation(*args, **kwargs)


from cross_reference_data import latest_cross_reference, cross_reference_for_industry
import contribution_data
import contribution_query
import contribution_geo
import contribution_workbook_guard
import collection_gaps
import snapshot_registry
import industry_scope
from runtime_paths import collector_environment
from api_request_budget import budget_path, legacy_budget_paths

# Bump the server revision with schema changes so existing desktop launchers
# detect that the running process needs the updated contribution reader.
CONTRIBUTION_STUDIO_REVISION = '2026-09-12-typed-fields-chart-families'
import evidence_csv
import evidence_geo
import geocode
from report_directory import report_directory, require_report_path
import report_publication
from validate_industry_coverage import workbook_roster
from collection_coverage import inspect as inspect_coverage, cycles_from
from fec_key_store import load_keys, save_keys, environment_keys, status as key_status
import mentat_env
from runtime_inputs import copy_fec_sources
import access_gate

# Files this console serves. Hoisted out of the request handler so the build
# identity below can hash exactly the set the browser can ask for.
# "/" is answered by console_page.page(), which renders app/templates/ at
# request time. The entry stays because index.html - the golden render those
# templates produce - must keep its place in the asset hash below.
STATIC_ASSETS = {"/": ("index.html", "text/html; charset=utf-8"), "/app.js": ("app.js", "text/javascript"),
                 '/insights.js': ('insights.js', 'text/javascript'),
                 '/insights.css': ('insights.css', 'text/css'),
                 '/giving-forecast.js': ('giving-forecast.js', 'text/javascript'),
                 '/giving-forecast.css': ('giving-forecast.css', 'text/css'),
                 '/pac-giving-forecast.js': ('pac-giving-forecast.js', 'text/javascript'),
                 '/pac-giving-forecast.css': ('pac-giving-forecast.css', 'text/css'),
                 '/contribution-studio': ('contribution-studio.html', 'text/html; charset=utf-8'),
                 '/contribution-studio.css': ('contribution-studio.css', 'text/css'),
                 '/contribution-studio.js': ('contribution-studio.js', 'text/javascript'),
                 '/contribution-plot.js': ('contribution-plot.js', 'text/javascript'),
                 '/contribution-charts.js': ('contribution-charts.js', 'text/javascript'),
                 '/contribution-map.js': ('contribution-map.js', 'text/javascript'),
                 '/mentat-icon.svg': ('mentat-icon.svg', 'image/svg+xml'),
                 "/style.css": ("style.css", "text/css"), "/template.csv": ("template.csv", "text/csv"),
                 '/workspace.css': ('workspace.css', 'text/css'),
                 '/workspace-polish.css': ('workspace-polish.css', 'text/css'),
                 '/evidence.js': ('evidence.js', 'text/javascript'),
                 # P4b. The interactive map. No script tag loads it: evidence.js
                 # appends <script src="/evidence-map.js"> the first time a map
                 # opens (CSP script-src 'self' allows a same-origin src), so a
                 # console that never opens a map pays nothing for it.
                 '/evidence-map.js': ('evidence-map.js', 'text/javascript'),
                 # Lane F (boards #44-#46): the Sankey layout and renderer, the two views that
                 # open it (PAC evidence Flow, Graph's Follow the money) and their stylesheet.
                 '/flow-layout.js': ('flow-layout.js', 'text/javascript'),
                 '/flow-chart.js': ('flow-chart.js', 'text/javascript'),
                 '/evidence-flow.js': ('evidence-flow.js', 'text/javascript'),
                 '/graph-flow.js': ('graph-flow.js', 'text/javascript'),
                 '/studio-flow.js': ('studio-flow.js', 'text/javascript'),
                 '/candidate-spending.js': ('candidate-spending.js', 'text/javascript'),
                 '/flow.css': ('flow.css', 'text/css'),
                 # Lane P (boards #47-#55): maps whose marks cards are its layers. The
                 # model (map-layers.js), the Albers SVG engine and Layer control
                 # (map-engine.js), the Chart Studio wiring (map-studio.js), the
                 # offline street renderer (map-streets.mjs, an ES module the engine
                 # imports only when Streets is chosen) and their stylesheet. The v2
                 # geography index is listed so a rebuilt tree moves the asset hash;
                 # the tree itself and the street files have routes of their own.
                 '/map-layers.js': ('map-layers.js', 'text/javascript'),
                 '/map-engine.js': ('map-engine.js', 'text/javascript'),
                 '/map-studio.js': ('map-studio.js', 'text/javascript'),
                 '/map-streets.mjs': ('map-streets.mjs', 'text/javascript'),
                 '/map-evidence.js': ('map-evidence.js', 'text/javascript'),
                 '/map.css': ('map.css', 'text/css'),
                 '/geo/v2/index.json': ('geo/v2/index.json', 'application/json; charset=utf-8'),
                 '/charts.js': ('charts.js', 'text/javascript'), '/chart-math.js': ('chart-math.js', 'text/javascript'),
                 '/graph.js': ('graph.js', 'text/javascript'),
                 '/graph-explorer.js': ('graph-explorer.js', 'text/javascript'),
                 '/scope.js': ('scope.js', 'text/javascript'),
                 '/ui-state.js': ('ui-state.js', 'text/javascript'),
                 '/coverage-view.js': ('coverage-view.js', 'text/javascript'),
                 '/cross-reference.js': ('cross-reference.js', 'text/javascript'),
                 '/cross-reference.css': ('cross-reference.css', 'text/css'),
                 '/comparison.js': ('comparison.js', 'text/javascript'),
                 '/grouped-chart.js': ('grouped-chart.js', 'text/javascript'),
                 '/funding-comparison.js': ('funding-comparison.js', 'text/javascript'),
                 '/comparison.css': ('comparison.css', 'text/css'),
                 '/chart-export.js': ('chart-export.js', 'text/javascript'),
                 '/figure-pdf.js': ('figure-pdf.js', 'text/javascript'),
                 '/worksheet-comparison.js': ('worksheet-comparison.js', 'text/javascript'),
                 '/graph.css': ('graph.css', 'text/css'),
                 '/acu-theme.css': ('acu-theme.css', 'text/css'),
                 '/chart-studio-core.js': ('chart-studio-core.js', 'text/javascript'),
                 '/chart-studio-contributions.js': ('chart-studio-contributions.js', 'text/javascript'),
                 '/chart-studio.js': ('chart-studio.js', 'text/javascript'),
                 '/chart-studio.css': ('chart-studio.css', 'text/css'),
                 '/theme.js': ('theme.js', 'text/javascript'),
                 '/chart-briefing.js': ('chart-briefing.js', 'text/javascript'),
                 '/chart-briefing.css': ('chart-briefing.css', 'text/css'),
                 '/vendor/pptxgen.bundle.js': ('vendor/pptxgen.bundle.js', 'text/javascript'),
                 # S8: the PAC evidence map draws from geography built once by
                 # tools/build_district_geo.py and shipped like any other asset.
                 # Without these entries an unregistered path falls through to
                 # {"error": "Not found."} and the map 404s.
                 # The charset form is deliberate: send() JSON-ENCODES its value when the
                 # content type is exactly 'application/json', so a file already read
                 # as bytes must not use that exact string (measured: TypeError,
                 # 'Object of type bytes is not JSON serializable', and the connection
                 # dropped). The compressible test below matches the prefix, so these
                 # still gzip.
                 '/geo/index.json': ('geo/index.json', 'application/json; charset=utf-8'),
                 '/geo/cd119.json': ('geo/cd119.json', 'application/json; charset=utf-8'),
                 '/geo/cd118.json': ('geo/cd118.json', 'application/json; charset=utf-8'),
                 '/geo/cd116.json': ('geo/cd116.json', 'application/json; charset=utf-8'),
                 '/geo/zcta_cd119.json': ('geo/zcta_cd119.json', 'application/json; charset=utf-8'),
                 '/geo/zcta_cd118.json': ('geo/zcta_cd118.json', 'application/json; charset=utf-8'),
                 # P4b: ZIP centre points, already projected onto the same
                 # 975x610 canvas as the district paths, so the page places a
                 # circle without re-deriving the projection.
                 '/geo/zcta_points.json': ('geo/zcta_points.json', 'application/json; charset=utf-8'),
                 # S8b: the census-tract crosswalk that turns a geocoded street
                 # address into a district. The batch geocoder answers a 2020
                 # census block, never a district (measured 2026-09-18), so the
                 # bridge ships like the ZIP crosswalk beside it.
                 '/geo/tract_cd119.json': ('geo/tract_cd119.json', 'application/json; charset=utf-8'),
                 '/geo/tract_cd118.json': ('geo/tract_cd118.json', 'application/json; charset=utf-8')}

# Only packaged legislative files are public; arbitrary geo paths never fall
# through to a filesystem handler. State bundles are fetched on selection.
for _geo_path in (PROJECT / 'app' / 'geo' / 'sld').glob('*/*.json'):
    _geo_name = _geo_path.relative_to(PROJECT / 'app').as_posix()
    STATIC_ASSETS['/' + _geo_name] = (_geo_name, 'application/json; charset=utf-8')
for _geo_name in ('geo/sld/index.json', 'geo/zcta_sld.json'):
    STATIC_ASSETS['/' + _geo_name] = (_geo_name, 'application/json; charset=utf-8')

# Build identity. A console window outlives edits to this tree, so the page and
# the launcher both compare what this process loaded with what is on disk now.
# File names are folded into the tree hash, so a rename changes it too.


@lru_cache(maxsize=512)
def file_digest(identity):
    """SHA-256 of one file, cached by (path, size, mtime_ns)."""
    try:
        return hashlib.sha256(Path(identity[0]).read_bytes()).digest()
    except OSError:
        return b'file unavailable'


def tree_sha256(paths):
    paths = [Path(path) for path in paths]
    # Geography now contains repeated state filenames in separate directories.
    # Include the relative path so two layers cannot exchange bytes unnoticed;
    # sorting also makes the fingerprint independent of set iteration order.
    root = Path(os.path.commonpath([str(path.parent) for path in paths])) if paths else None
    digest = hashlib.sha256()
    for path in sorted(paths, key=lambda path: path.relative_to(root).as_posix()):
        try:
            info = path.stat()
            identity = (str(path), info.st_size, info.st_mtime_ns)
        except OSError:
            identity = (str(path), -1, -1)
        digest.update(path.relative_to(root).as_posix().encode('utf-8'))
        digest.update(b'\0')
        digest.update(file_digest(identity))
        digest.update(b'\0')
    return digest.hexdigest()


def code_files():
    """Every Python module this console loads from its own folder."""
    return sorted(Path(__file__).resolve().parent.glob('*.py'), key=lambda path: path.name)


def asset_files():
    # The page is rendered from app/templates/, so a partial edited there has to
    # move this hash exactly as an edit to index.html used to; without it the
    # page would never notice that the build on disk had moved past this
    # process. tree_sha256 includes relative directories as well as filenames.
    import console_page
    directory = Path(__file__).resolve().parent
    return sorted({directory / name for name, _ in STATIC_ASSETS.values()}
                  | set(console_page.template_files()), key=lambda path: path.name)


GIT_SHA_COMMAND = ['git', 'rev-parse', '--short', 'HEAD']
GIT_SHA_TIMEOUT_SECONDS = 10
# Nine hex digits is what `git rev-parse --short HEAD` answers in this repo
# today, so the two readers below agree on sight. The width is safe to fix
# here because the value is display-only: D-029's restart gate compares
# code_sha256 and assets_sha256, never the git sha.
GIT_SHA_DISPLAY_LENGTH = 9


def git_sha_from_gitdir():
    """Short revision read off Git's own files, with no child process at all.

    On this OneDrive-backed tree a `git` child can wedge on placeholder I/O and
    then refuse to die (measured 2026-09-18: two orphaned git.exe alive for an
    hour), so the ordinary answer is read straight out of the repository.
    Returns None when anything is missing or malformed, and the caller then
    falls back to Git itself.
    """
    try:
        # PROJECT is the app's parent folder, not the checkout root, so
        # PROJECT/.git does not exist. `git rev-parse` answered anyway because
        # it searches upward from its working directory; this reader searches
        # the same way so that both paths speak about the same checkout.
        pointer = None
        for candidate in (PROJECT, *PROJECT.parents):
            if (candidate / '.git').exists():
                pointer = candidate / '.git'
                break
        if pointer is None:
            return None
        if pointer.is_dir():
            gitdir = pointer
        elif pointer.is_file():
            # A linked worktree keeps `.git` as a FILE holding `gitdir: <path>`,
            # and that path may be relative to the folder holding the file.
            marker = pointer.read_text(encoding='utf-8', errors='ignore').strip()
            if not marker.startswith('gitdir:'):
                return None
            gitdir = Path(os.path.normpath(str(pointer.parent / marker[len('gitdir:'):].strip())))
        else:
            return None
        head = (gitdir / 'HEAD').read_text(encoding='utf-8', errors='ignore').strip()
    except OSError:
        return None
    if re.fullmatch(r'[0-9a-f]{40}', head):
        # A detached HEAD holds the commit itself rather than a ref.
        return head[:GIT_SHA_DISPLAY_LENGTH]
    if not head.startswith('ref:'):
        return None
    reference = head[len('ref:'):].strip()
    # Only a plain ref path is followed, so a malformed HEAD cannot send this
    # reader up out of the repository.
    if not re.fullmatch(r'refs/[\w.-]+(?:/[\w.-]+)*', reference) or '..' in reference:
        return None
    # A worktree's own refs folder holds only the refs it owns; a shared branch
    # lives in the common directory instead, loose or packed.
    common = gitdir
    try:
        if (gitdir / 'commondir').is_file():
            relative = (gitdir / 'commondir').read_text(encoding='utf-8', errors='ignore').strip()
            if relative:
                common = Path(os.path.normpath(str(gitdir / relative)))
    except OSError:
        return None
    for candidate in (gitdir / reference, common / reference):
        try:
            value = candidate.read_text(encoding='utf-8', errors='ignore').strip()
        except OSError:
            continue
        if re.fullmatch(r'[0-9a-f]{40}', value):
            return value[:GIT_SHA_DISPLAY_LENGTH]
    try:
        packed = (common / 'packed-refs').read_text(encoding='utf-8', errors='ignore')
    except OSError:
        return None
    for line in packed.splitlines():
        # '#' opens the header line and '^' a tag's peeled commit, not a ref.
        if not line or line[0] in '#^':
            continue
        parts = line.split()
        if len(parts) == 2 and parts[1] == reference and re.fullmatch(r'[0-9a-f]{40}', parts[0]):
            return parts[0][:GIT_SHA_DISPLAY_LENGTH]
    return None


def git_sha():
    """Short revision from Git itself, or 'unknown' when Git cannot answer.

    stdout goes to a temporary FILE, never a pipe. subprocess.run answers a
    timeout by re-entering communicate() to drain, and communicate() then
    blocks forever joining the pipe-reader threads of a child that will not
    die: that is how one health call hung the whole console. A real file
    handle creates no reader threads, so kill() is the entire cleanup and this
    returns on time even when the child is wedged.
    """
    try:
        with tempfile.TemporaryFile() as sink:
            child = subprocess.Popen(GIT_SHA_COMMAND, cwd=str(PROJECT), stdin=subprocess.DEVNULL,
                                     stdout=sink, stderr=subprocess.DEVNULL,
                                     creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
            try:
                code = child.wait(timeout=GIT_SHA_TIMEOUT_SECONDS)
            except subprocess.TimeoutExpired:
                child.kill()
                try:
                    # Reap the child if the kill lands; never wait on it again.
                    child.wait(timeout=1)
                except subprocess.TimeoutExpired:
                    pass
                return 'unknown'
            sink.seek(0)
            value = sink.read().decode('utf-8', errors='ignore').strip()
    except (OSError, ValueError, subprocess.SubprocessError):
        return 'unknown'
    return value if code == 0 and re.fullmatch(r'[0-9a-f]{4,40}', value) else 'unknown'


CODE_TREE_SHA256 = tree_sha256(code_files())
ASSET_TREE_SHA256 = tree_sha256(asset_files())


_git_sha_lock = threading.Lock()
_git_sha_value = None
_git_sha_started = False


def git_sha_lookup():
    """Where the build stamp's revision comes from: Git's files, then Git."""
    return git_sha_from_gitdir() or git_sha()


def resolve_git_sha():
    """Store the revision once, off the request path, whatever happens."""
    global _git_sha_value
    try:
        value = git_sha_lookup()
    except Exception:
        value = None
    with _git_sha_lock:
        _git_sha_value = value or 'unknown'


def git_sha_cached():
    """The revision if it is already known, otherwise 'unknown' at once.

    A request must never wait on this lookup. lru_cache did wait: it does not
    serialise concurrent callers, so one page load raced two request threads
    into the git child, and when that child wedged /api/session never
    answered and the console window hung at startup. One daemon thread does
    the work instead and later requests pick up the answer it left.
    """
    global _git_sha_started
    if _git_sha_value is not None:
        return _git_sha_value
    with _git_sha_lock:
        if _git_sha_value is not None:
            return _git_sha_value
        if not _git_sha_started:
            _git_sha_started = True
            threading.Thread(target=resolve_git_sha, name='git-sha', daemon=True).start()
    return 'unknown'


def reset_git_sha_cache():
    """Forget the resolved revision so a test can drive the resolver again."""
    global _git_sha_value, _git_sha_started
    with _git_sha_lock:
        _git_sha_value = None
        _git_sha_started = False
STARTED_AT = time.time()


def console_busy(app):
    """Whether this console owns work a restart would interrupt.

    A console that cannot answer is reported busy: never stop work on a guess.
    """
    reader = getattr(app, 'busy', None)
    if not callable(reader):
        return False
    try:
        return bool(reader())
    except Exception:
        return True


def build_status(app, port):
    """What this process loaded, beside what is on disk right now."""
    return {'git_sha': git_sha_cached(), 'code_sha256': CODE_TREE_SHA256, 'assets_sha256': ASSET_TREE_SHA256,
            'on_disk_code_sha256_now': tree_sha256(code_files()),
            'on_disk_assets_sha256_now': tree_sha256(asset_files()),
            'started_at': STARTED_AT, 'port': int(port), 'pid': os.getpid(),
            'busy': console_busy(app)}


@lru_cache(maxsize=64)
def static_asset(file_stamp):
    content = Path(file_stamp[0]).read_bytes()
    return content, 'W/"' + hashlib.sha256(content).hexdigest() + '"'


# Lane P (boards #47-#55). Two trees the maps read that are too large for
# STATIC_ASSETS, each behind one route that answers only the names its pattern
# allows (a path is matched, never joined blindly, so `..` cannot leave it):
#   /geo/v2/<form>/...   lane G's layer geography, mentat/app/geo/v2 (866 files,
#                        205 MB: listing them would make the build-identity hash
#                        read them all at start-up and stat them on every health
#                        check; geo/v2/index.json, which carries every file's
#                        sha256, is in STATIC_ASSETS instead)
#   /basemap/...         lane E's offline street basemap, mentat/.cache/basemap:
#                        MapLibre GL 6.11.2, pmtiles 4.5.0, the generated style,
#                        two Noto Sans glyph stacks, the v4 light sprite, and the
#                        tile archive, which is answered by HTTP Range
#                        (send_ranged) because pmtiles reads it a few KB at a time.
#
# The console's one security policy. worker-src and connect-src are stated
# rather than left to default-src (lane P): the street renderer runs a
# same-origin module worker and fetches the style, glyphs and tiles. Both stay
# 'self', and nothing else in the policy moves.
CONSOLE_CSP = ("default-src 'self'; style-src 'self'; script-src 'self'; img-src 'self' blob:; "
               "worker-src 'self'; connect-src 'self'; frame-ancestors 'none'")
# Hosted mode: Cloudflare Access is enforced (access_gate.py, ACCESS_* environment). Access authenticates; the
# console's per-process token no longer does, so it is never handed out. The page still insists on a session
# token before it will POST, so it is given this fixed marker instead, and a POST must send it back: a cross-site
# form or fetch cannot add the header without a CORS preflight, which this server never answers.
HOSTED_SESSION_TOKEN = 'hosted-read-only'
HOSTED_REFUSED = 'This route is not available on the hosted read-only console.'
# Donor-level rows (the two records routes) are served to an Access identity only if MENTAT_DONOR_ROW_EMAILS
# (comma-separated, case-insensitive; default empty, so nobody) lists its email. Pending Paul's governance ruling.
DONOR_ROW_PATHS = ('/api/contributions/records', '/api/insights/records')
HOSTED_DONOR_ROWS_REFUSED = 'Donor-level rows are not available to this account on the hosted console.'
HOSTED_PERSON_LEVEL_REFUSED = 'Person-level fields are not available to this account on the hosted console.'
HOSTED_IO_ERROR = 'The server could not read that data.'
# Hosted GET is an allowlist (do_POST's is analytical_paths): the page, its static and map files (STATIC_ASSETS,
# /geo/v2/, /basemap/), /healthz and /version (access_gate.HEALTH_ROUTES, answered first) and exactly these routes.
# Anything else, including every evidence, run, publish and studio route and the raw workbook download, is 403.
HOSTED_GET_ROUTES = frozenset({
    '/', '/api/health', '/api/session', '/api/contributions/session', '/api/runs',
    '/api/insights/catalog', '/api/forecast/pac', '/api/forecast/catalog', '/api/contributions/catalog',
    '/api/contributions/geography/catalog', '/api/graph-catalog', '/api/reports', '/api/report-sources',
    '/api/pac-flow-catalog', '/api/snapshots', '/api/graph', '/api/comparison', '/api/comparison-flows',
    '/api/flow', '/api/cross-reference', '/api/report-data'})
# Person-level fields: every dimension the contribution index can group or filter by except these coarse ones. Derived
# from contribution_semantics.DIMENSIONS (the contributor, street, city, ZIP, employer, occupation and date dimensions,
# plus every AUDIT and RAW field), so a field added there is person-level until it is listed here on purpose.
# ponytail: no small-cell suppression (minimum group size) on the aggregate fields; that threshold is Paul's call.
AGGREGATE_FIELDS = frozenset({'industry', 'cycle', 'year', 'company', 'state', 'recipient', 'committee_id',
                              'recipient_category', 'party', 'funding_source', 'decision_status', 'receipt_role',
                              *contribution_data.DESTINATIONS, *contribution_query.semantics.DERIVED})
PERSON_LEVEL_FIELDS = frozenset(contribution_query.semantics.DIMENSIONS) - AGGREGATE_FIELDS
_SERVER_PATH = re.compile(r'[A-Za-z]:[\\/]|\\\\|/')


def donor_rows_allowed(email):
    listed = {item.strip().lower() for item in os.environ.get('MENTAT_DONOR_ROW_EMAILS', '').split(',') if item.strip()}
    return bool(email) and email.strip().lower() in listed


def names_person_level_field(body):
    """True when a query body names a person-level field as a dimension or measure or in a filter, aggregate
    filter or sort. An item that is not an object is read as the field name itself (fail closed)."""
    named = []
    for key, value in body.items():
        items = value if isinstance(value, list) else [value]
        if key in ('dimensions', 'measures'):
            named += items
        elif key in ('filters', 'aggregate_filters', 'sort'):
            named += [item.get('field') if isinstance(item, dict) else item for item in items]
    return any(isinstance(name, str) and name in PERSON_LEVEL_FIELDS for name in named)


def without_server_paths(value, keep_places=False):
    """`value` minus every string that is an absolute server path: a dict entry or list item holding one is omitted.
    keep_places: a list item becomes '' instead, so a worksheet row keeps its columns."""
    drop = lambda item: isinstance(item, str) and _SERVER_PATH.match(item)
    if isinstance(value, dict):
        return {key: without_server_paths(item, keep_places) for key, item in value.items() if not drop(item)}
    if isinstance(value, list):
        if keep_places:
            return ['' if drop(item) else without_server_paths(item, True) for item in value]
        return [without_server_paths(item) for item in value if not drop(item)]
    return value


_OS_ERROR_TEXT = re.compile(r'\[(?:Errno|WinError) -?\d+\]')


def without_os_error_text(value):
    """`value` with every string that carries OSError text ("[Errno 2] ...", "[WinError 3] ...") replaced by HOSTED_IO_ERROR."""
    if isinstance(value, str):
        return HOSTED_IO_ERROR if _OS_ERROR_TEXT.search(value) else value
    if isinstance(value, dict):
        return {key: without_os_error_text(item) for key, item in value.items()}
    if isinstance(value, list):
        return [without_os_error_text(item) for item in value]
    return value


# Hosted /api/flow: its records (links[].evidence[], downstream.filing_record through flow_data._link_out) name a
# counterparty, a date and the filing they were read from. An account not in MENTAT_DONOR_ROW_EMAILS gets none of
# M1's person-level query fields in that body (PERSON_LEVEL_FIELDS: contributor, street, city, ZIP, employer,
# occupation, date, every AUDIT and RAW field, source_url among them), nor the record's own names for them:
# counterparty (the contributor), zip, and the row pointers sub_id and transaction_id (source_id and
# reported_transaction_id there), nor purpose (free text that can name a payee). 'state' and the committee and
# candidate ids are aggregates and stay.
FLOW_PERSON_KEYS = PERSON_LEVEL_FIELDS | {'counterparty', 'zip', 'sub_id', 'transaction_id', 'purpose'}


def without_person_keys(value):
    """`value` minus every key in FLOW_PERSON_KEYS, at any depth."""
    if isinstance(value, dict):
        return {key: without_person_keys(item) for key, item in value.items() if key not in FLOW_PERSON_KEYS}
    if isinstance(value, list):
        return [without_person_keys(item) for item in value]
    return value


# Hosted /api/flow, same viewers: no individual contributor's identity at all (Claude's default pending Paul's
# governance ruling). Every donor node that is not explicitly an organization or committee becomes the one aggregate
# node below; the "Other" bucket keeps its count but not its member names; the refunded sink's text names no one.
INDIVIDUALS_ID = 'donor:individuals'
REFUNDS_HIDDEN = 'Refunds to individual contributors are not itemized on the hosted console.'


def is_individual_donor(node):
    """A donor node that is not explicitly an organization or committee: FEC entity type IND, and also no type, or a
    type this does not know, whatever the name looks like (flow_data draws an untyped name without a comma as an
    organization, and a filer's "Acme Hardware" can be a sole proprietor). Only ORG, CAN and
    flow_data.COMMITTEE_ENTITIES (PAC, CCM, PTY, COM, ORG_PAC) keep their names; the unitemized node is no one."""
    from flow_data import COMMITTEE_ENTITIES
    return (node.get('kind') == 'donor' and str(node.get('id')).startswith('donor:') and not node.get('unitemized')
            and str(node.get('entity') or '').upper() not in COMMITTEE_ENTITIES | {'ORG', 'CAN'})


def without_individual_donors(body):
    """A /api/flow body with its individual donor nodes merged into one aggregate node (amounts summed, the links
    repointed and summed, every per-person field dropped), the "Other" bucket's member list and the refunded sink's
    details replaced by counts and a generic sentence. Totals, committees and everything else are as built."""
    from flow_data import EVIDENCE_LIMIT
    people = [node for node in body.get('nodes', []) if is_individual_donor(node)]
    ids = {node['id'] for node in people}
    nodes = []
    for node in body.get('nodes', []):
        if node['id'] in ids:
            if node is people[0]:
                nodes.append({'id': INDIVIDUALS_ID, 'stage': node.get('stage', 0),
                              'label': f'Individual contributors ({len(people)})', 'kind': 'donor',
                              'contribution': all(item.get('contribution', True) for item in people),
                              'value_in': round(sum(item.get('value_in', 0) for item in people), 2),
                              'value_out': round(sum(item.get('value_out', 0) for item in people), 2)})
            continue
        node = dict(node)
        if node.get('bucket') == 'other' and any(str(member.get('id')).startswith('donor:')
                                                 for member in node.get('members', [])):
            node.pop('members')                      # member_count stays
        if node.get('balance') == 'refunded' and 'details' in node:
            node['details'] = [REFUNDS_HIDDEN]
        nodes.append(node)
    links, repointed = [], {}
    for link in body.get('links', []):
        if link['source'] not in ids and link['target'] not in ids:
            links.append(link)
            continue
        link = {**link, 'source': INDIVIDUALS_ID if link['source'] in ids else link['source'],
                'target': INDIVIDUALS_ID if link['target'] in ids else link['target']}
        key = (link['source'], link['target'], link['kind'])
        if key not in repointed:
            repointed[key] = link
            links.append(link)
            continue
        # ponytail: group()'s merge without its 'committee' rule (a donor receipt link has none); copy it if one ever does
        merged = repointed[key]
        merged['value'] = round(merged['value'] + link['value'], 2)
        merged['rows'] = merged.get('rows', 0) + link.get('rows', 0)
        if merged.get('evidence') or link.get('evidence'):
            merged['evidence'] = (merged.get('evidence', []) + link.get('evidence', []))[:EVIDENCE_LIMIT]
        if link.get('date_unverified'):
            merged['date_unverified'] = True
    return {**body, 'nodes': nodes, 'links': links}


# Hosted /api/comparison, same viewers: its Person subjects (comparison_data.subjects) are the report's employees, each
# with a name, a company and per-cycle totals. They are dropped, and a query that selects one is refused.
def without_person_subjects(body):
    return {**body, 'subjects': [subject for subject in body.get('subjects', []) if subject.get('kind') != 'Person']}


def names_person_subject(query):
    """True when a /api/comparison query (parse_qs) other than its run and industry says 'Person' (kind=Person) or
    gives a subject id of that kind (comparison_data.subjects: '["Person", [company, person id]]')."""
    return any(re.fullmatch(r'\s*person\s*', value, re.I) or re.match(r'\s*\[\s*"person"', value, re.I)
               for key, values in query.items() if key not in ('id', 'industry') for value in values)


# Hosted /api/report-data, same viewers, for a sheet report_insights.person_level_sheet calls REDACT (every '* Summary' and
# '<Company> View' sheet): Claude's default pending Paul's governance ruling is to show it with the contributors' names and
# name keys blanked, not to refuse it. A row keeps its Name / Recipient and Reference Key only when its Section is one of
# these, the recipient and aggregate sections the real workbooks carry (the fourth, 'Top Contributor', names an individual);
# any other value, a blank one included, is blanked, so a section added later is redacted until it is listed here.
RECIPIENT_SECTIONS = frozenset({'Top Direct Recipient', 'Top Downstream Beneficiary', 'Company Total'})


def without_contributor_names(body):
    """A worksheet_data body with every Name / Recipient and Reference Key cell blanked on the rows whose Section is
    not in RECIPIENT_SECTIONS, and those columns' profile recomputed. None (the caller refuses) when the body has no
    single Section column: the header row asked for was not the sheet's own, so no cell can be told from a contributor's."""
    import report_insights as insights
    names = [column.get('name') for column in body.get('columns', [])]
    texts = [insights.header_text(name) for name in names]
    if texts.count(insights.SECTION_HEADER) != 1:
        return None
    section = texts.index(insights.SECTION_HEADER)
    blank = [index for index, text in enumerate(texts) if text in insights.REDACTED_HEADERS]
    rows = [row if str(row[section]).strip() in RECIPIENT_SECTIONS else
            [None if index in blank else cell for index, cell in enumerate(row)] for row in body.get('rows', [])]
    profile = insights.profile_columns(names, rows)
    columns = [{**column, **{key: profile[index][key] for key in ('type', 'missing', 'numeric_values', 'distinct')}}
               if index in blank else column for index, column in enumerate(body['columns'])]
    return {**body, 'rows': rows, 'columns': columns}


GEO_V2_DIR = Path(__file__).resolve().parent / 'geo' / 'v2'
GEO_V2_ROUTE = re.compile(r'/geo/v2/((?:albers|wgs84|lookup)(?:/[A-Za-z0-9_@.-]+){1,4}\.(?:json|geojson))')
BASEMAP_DIR = PROJECT / '.cache' / 'basemap'
BASEMAP_FILES = {
    '/basemap/vendor/maplibre-gl.mjs': ('vendor/maplibre-gl-6.11.2/maplibre-gl.mjs', 'text/javascript'),
    '/basemap/vendor/maplibre-gl-worker.mjs': ('vendor/maplibre-gl-6.11.2/maplibre-gl-worker.mjs', 'text/javascript'),
    '/basemap/vendor/maplibre-gl-shared.mjs': ('vendor/maplibre-gl-6.11.2/maplibre-gl-shared.mjs', 'text/javascript'),
    '/basemap/vendor/maplibre-gl.css': ('vendor/maplibre-gl-6.11.2/maplibre-gl.css', 'text/css'),
    '/basemap/vendor/pmtiles.js': ('vendor/pmtiles-4.5.0/pmtiles.js', 'text/javascript'),
    '/basemap/style/light.json': ('style/light.json', 'application/json; charset=utf-8'),
    '/basemap/style/dark.json': ('style/dark.json', 'application/json; charset=utf-8'),
    '/basemap/sprites/v4/light.json': ('sprites/v4/light.json', 'application/json; charset=utf-8'),
    '/basemap/sprites/v4/light.png': ('sprites/v4/light.png', 'image/png'),
    '/basemap/sprites/v4/light@2x.json': ('sprites/v4/light@2x.json', 'application/json; charset=utf-8'),
    '/basemap/sprites/v4/light@2x.png': ('sprites/v4/light@2x.png', 'image/png'),
    '/basemap/ATTRIBUTION.txt': ('licenses/ATTRIBUTION.txt', 'text/plain; charset=utf-8'),
}
BASEMAP_FONT_ROUTE = re.compile(r'/basemap/fonts/(Noto Sans (?:Regular|Medium))/(\d{1,5}-\d{1,5})\.pbf')
BASEMAP_TILES_ROUTE = re.compile(r'/basemap/tiles/([a-z0-9][a-z0-9_-]{0,63}\.pmtiles)')
_BYTE_RANGE = re.compile(r'\s*bytes\s*=\s*(\d*)\s*-\s*(\d*)\s*')


def byte_range(header, size):
    """The one byte range a Range header asks of a `size`-byte file.

    (first, last) inclusive for a satisfiable single range; 'unsatisfiable'
    for a well-formed single range wholly past the end (answered 416 with
    `Content-Range: bytes */size`, which pmtiles reads to learn the size);
    None when there is no Range header, or one this server does not serve
    (several ranges, another unit, a malformed or backwards range), in which
    case the whole file is sent with 200, as RFC 9110 permits.
    """
    if not header:
        return None
    match = _BYTE_RANGE.fullmatch(str(header))
    if not match:
        return None
    first, last = match.group(1), match.group(2)
    if not first and not last:
        return None
    if not first:
        suffix = int(last)
        if suffix == 0 or size == 0:
            return 'unsatisfiable'
        return max(0, size - suffix), size - 1
    start = int(first)
    if last and int(last) < start:
        return None
    if start >= size:
        return 'unsatisfiable'
    return start, min(int(last), size - 1) if last else size - 1


def accepts_gzip(header):
    for entry in (header or '').split(','):
        parts = [part.strip() for part in entry.lower().split(';')]
        if parts[0] != 'gzip':
            continue
        try:
            return all(float(part[2:]) > 0 for part in parts[1:] if part.startswith('q='))
        except ValueError:
            return False
    return False


def read_json(path):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except (OSError, ValueError):
        return {}


@lru_cache(maxsize=4096)
def _cached_json(identity):
    return read_json(identity[0])


def read_json_cached(path):
    """read_json keyed by (path, size, mtime_ns): a run's status file is parsed
    again only when it changes. /api/runs re-read every status file on every
    page load, which cost 3 s on a data root with a history of runs."""
    try:
        info = Path(path).stat()
    except OSError:
        return {}
    return dict(_cached_json((str(path), info.st_size, info.st_mtime_ns)))


def newer_timestamp(candidate, current):
    """True when ISO-8601 `candidate` is strictly later than `current`; False
    when either is missing or unreadable, so an unparseable stamp never wins."""
    try:
        return datetime.datetime.fromisoformat(str(candidate)) > datetime.datetime.fromisoformat(str(current))
    except (TypeError, ValueError):
        return False


def recorded_tasks(task_files, root):
    """The newest recorded attempt of every collection task, one row per task.

    `task_files` holds (industry, status path) pairs. Files are read oldest
    first, so a retried task (a recovery or split invocation) replaces its
    earlier failure. Each file is parsed again only when it changes, and a
    file that resolves outside the run `root` is never read."""
    stamped = []
    for industry, path in task_files:
        try:
            path = Path(path).resolve()
            path.relative_to(root)
            stamped.append((path.stat().st_mtime_ns, str(path), industry, path))
        except (OSError, ValueError):
            continue
    tasks = {}
    for _, _, industry, path in sorted(stamped):
        for task in read_json_cached(path).get('tasks') or []:
            if isinstance(task, dict) and task.get('company'):
                identity = task.get('task_id') or f"{task.get('company')}|{task.get('cycle')}"
                tasks[(industry, str(identity))] = {**task, 'industry': industry}
    return list(tasks.values())


def r_library_path(environment, project=PROJECT, rscript=""):
    """Choose an existing R library for this child, without changing global settings."""
    def existing(candidate):
        try:
            return str(candidate.resolve()) if candidate.is_dir() else None
        except (OSError, ValueError):
            return None
    configured = next((value for key,value in environment.items() if key.upper() == 'R_LIBS_USER'), '')
    for value in configured.split(os.pathsep):
        if value.strip():
            candidate = Path(os.path.expandvars(value.strip().strip('"'))).expanduser()
            found = existing(candidate)
            if found:
                return found
    version = re.search(r'R-(\d+\.\d+)(?:\.\d+)?', str(rscript))
    version = version.group(1) if version else '4.6'
    local = Path(environment.get('LOCALAPPDATA') or Path.home() / 'AppData' / 'Local')
    user_library = local / 'R' / 'win-library' / version
    found = existing(user_library)
    if found:
        return found
    bundled = Path(project) / '.cache' / 'audit-r-library'
    try:
        if all((bundled / package / 'DESCRIPTION').is_file() for package in ('digest','DBI','RSQLite','data.table')):
            return str(bundled.resolve())
    except OSError:
        pass
    return None


def run_state(status, fallback):
    """Account for older runners that marked failed industries completed."""
    stage = status.get('stage')
    if stage == 'completed':
        return 'failed' if any(row.get('status') != 'succeeded' for row in status.get('industries', [])) else 'completed'
    return stage if stage in ('failed', 'stopped') else fallback


def run_has_lane_state(directory):
    """Whether a run directory holds fanout manifests a resume could continue.

    The manifest is what the fanout's --resume reloads. Without one there are
    no lane snapshots to reuse, so a resume would have nothing to carry forward
    and would only re-collect what the stopped attempt already collected.
    """
    if not directory:
        return False
    return next(Path(directory).glob('industries/*/fanout/manifest.json'), None) is not None


def stale_resume_status(directory, meta):
    """A resumed run's saved status still describes the attempt that stopped.

    Reporting that terminal stage back would drop the run into History seconds
    after the operator resumed it, and would offer Resume again for a run that
    is already being relaunched. Only the relaunched orchestrator's own rewrite
    makes the file current, so it is trusted from the moment it is rewritten.
    """
    if not meta.get('resumed_from'):
        return False
    try:
        return (Path(directory) / 'status.json').stat().st_mtime < (meta.get('started_at') or 0)
    except OSError:
        return True


ORPHANED_RUN_GRACE_SECONDS = 15 * 60


def run_status_is_stale(directory, grace=ORPHANED_RUN_GRACE_SECONDS):
    """Whether a run stopped writing its status file longer than a grace ago.

    Only consulted where the manifest names no process to ask about, so the
    file's own age is the only evidence there is. A run with no directory or no
    status file yet has produced no evidence at all and is never called stale:
    a collection seconds from its first write must not be read as dead.
    """
    if not directory:
        return False
    try:
        return time.time() - (Path(directory) / 'status.json').stat().st_mtime > grace
    except OSError:
        return False


def process_alive(pid):
    if not isinstance(pid, int) or pid <= 0:
        return False
    if os.name == 'nt':
        import ctypes
        kernel = ctypes.WinDLL('kernel32', use_last_error=True)
        kernel.OpenProcess.restype = ctypes.c_void_p
        kernel.GetExitCodeProcess.argtypes = [ctypes.c_void_p, ctypes.POINTER(ctypes.c_ulong)]
        kernel.CloseHandle.argtypes = [ctypes.c_void_p]
        handle = kernel.OpenProcess(0x1000, False, pid)
        if not handle:
            return False
        try:
            code = ctypes.c_ulong()
            return bool(kernel.GetExitCodeProcess(handle, ctypes.byref(code))) and code.value == 259
        finally:
            kernel.CloseHandle(handle)
    try:
        os.kill(pid, 0)
        return True
    except OSError:
        return False


@contextmanager
def console_state_lock(root):
    """Serialize collection claims and manifest changes across consoles."""
    with (Path(root) / 'collection-control.lock').open('a+b') as stream:
        if stream.tell() == 0:
            stream.write(b'0'); stream.flush()
        stream.seek(0)
        deadline = time.monotonic() + 10
        while True:
            try:
                if os.name == 'nt':
                    import msvcrt
                    msvcrt.locking(stream.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl
                    fcntl.flock(stream, fcntl.LOCK_EX | fcntl.LOCK_NB)
                break
            except OSError:
                if time.monotonic() >= deadline:
                    raise ValueError('Another console is updating collection state. Try again shortly.')
                time.sleep(.025)
        try:
            yield
        finally:
            stream.seek(0)
            if os.name == 'nt':
                msvcrt.locking(stream.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(stream, fcntl.LOCK_UN)


@contextmanager
def process_identity_handle(pid):
    """Hold the Windows process object open while checking and using its PID.

    An open handle prevents PID reuse between verification and termination.
    Inaccessible or missing identity leaves only the graceful stop path.
    """
    if os.name != 'nt' or type(pid) is not int or pid <= 0:
        yield None
        return
    import ctypes
    from ctypes import wintypes
    kernel = ctypes.WinDLL('kernel32', use_last_error=True)
    kernel.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
    kernel.OpenProcess.restype = wintypes.HANDLE
    kernel.CloseHandle.argtypes = [wintypes.HANDLE]
    kernel.GetProcessTimes.argtypes = [wintypes.HANDLE] + [ctypes.POINTER(wintypes.FILETIME)] * 4
    kernel.GetExitCodeProcess.argtypes = [wintypes.HANDLE, ctypes.POINTER(wintypes.DWORD)]
    kernel.QueryFullProcessImageNameW.argtypes = [wintypes.HANDLE, wintypes.DWORD, wintypes.LPWSTR, ctypes.POINTER(wintypes.DWORD)]
    handle = kernel.OpenProcess(0x1000, False, pid)
    if not handle:
        yield None
        return
    try:
        created, ended, kernel_time, user_time = (wintypes.FILETIME() for _ in range(4))
        code = wintypes.DWORD(); size = wintypes.DWORD(32768)
        executable = ctypes.create_unicode_buffer(size.value)
        if (kernel.GetExitCodeProcess(handle, ctypes.byref(code)) and code.value == 259
                and kernel.GetProcessTimes(handle, ctypes.byref(created), ctypes.byref(ended), ctypes.byref(kernel_time), ctypes.byref(user_time))
                and kernel.QueryFullProcessImageNameW(handle, 0, executable, ctypes.byref(size))):
            yield {'pid':pid, 'created':(created.dwHighDateTime << 32) | created.dwLowDateTime,
                   'executable':os.path.normcase(executable.value)}
        else:
            yield None
    finally:
        kernel.CloseHandle(handle)


def process_identity(pid):
    with process_identity_handle(pid) as identity:
        return identity


def process_confirmed_dead(manifest):
    """Whether the collection process a saved run names is provably gone.

    Only a positive answer is ever acted on, so every ambiguity answers False:
    no PID recorded, a process this session cannot open, a manifest written
    before identities were kept. Where an identity was kept it is compared
    through one open handle, the way stop() verifies before it kills, so a
    recycled PID that now belongs to something else cannot keep a dead run
    looking alive.
    """
    pid = manifest.get('pid')
    if type(pid) is not int or pid <= 0:
        return False
    if not process_alive(pid):
        return True
    saved = manifest.get('process_identity')
    if not isinstance(saved, dict):
        return False
    with process_identity_handle(pid) as identity:
        return identity is not None and identity != saved


def replace_with_retry(temporary, path, attempts=10):
    """os.replace onto a destination another process has open, on Windows.

    A reader holding the destination open makes os.replace raise PermissionError
    (WinError 5) for as long as it holds it - microseconds, but a console listing
    saved views while a worker rewrites index.json was enough to lose two runs in
    twenty. Windows offers no way to wait on that handle, so back off and try
    again: 20 ms doubling to a second, about four seconds in all. A hold that
    never clears is a real failure and is raised, never swallowed.
    """
    delay = 0.02
    for attempt in range(attempts):
        try:
            os.replace(temporary, path)
            return
        except PermissionError:
            if attempt == attempts - 1:
                raise
            time.sleep(delay)
            delay = min(delay * 2, 1.0)


def atomic_json(path, value):
    path = Path(path)
    temporary = path.with_name(path.name + '.' + uuid.uuid4().hex + '.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    replace_with_retry(temporary, path)


def normalize_upload(filename, content, target):
    """Convert a table or legacy roster to the existing collector's input format."""
    import openpyxl
    suffix = Path(filename).suffix.lower()
    rows = []
    if suffix in (".csv", ".tsv"):
        text = content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text), delimiter="\t" if suffix == ".tsv" else ","))
    elif suffix == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30_000_000:
                raise ValueError("The expanded workbook must be under 30 MB.")
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            rows = [list(row) for row in book.active.iter_rows(values_only=True)]
            header = [str(cell or "").strip().lower() for cell in rows[0]] if rows else []
            if "canonical_name" not in header:
                target.write_bytes(content)
                return workbook_roster(target)
        finally:
            book.close()
    else:
        raise ValueError("Upload a CSV, TSV, or XLSX file.")
    if not rows:
        raise ValueError("The file is empty.")
    headers = [str(value or "").strip().lower() for value in rows[0]]
    if "canonical_name" not in headers:
        raise ValueError("The table needs a canonical_name column and an optional alias column.")
    if len(headers) != len(set(headers)):
        raise ValueError("Each column needs a unique header.")
    if len(rows) > 10001:
        raise ValueError("Upload up to 10,000 rows at a time.")
    groups = {}
    for number, row in enumerate(rows[1:], 2):
        if not any(str(value or "").strip() for value in row):
            continue
        item = dict(zip(headers, [str(value or "").strip() for value in row]))
        name = item.get("canonical_name", "")
        if item.get("sector") and item.get("industry") and item['sector'] != item['industry']:
            raise ValueError(f"Row {number}: sector and industry must agree when both are provided.")
        industry = item.get("sector") or item.get("industry") or "Uploaded names"
        if not name:
            raise ValueError(f"Row {number}: canonical_name is required.")
        if any(len(value) > 250 for value in item.values()):
            raise ValueError(f"Row {number}: names and aliases must be under 250 characters.")
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9 &()._-]{0,30}", industry):
            raise ValueError(f"Row {number}: use a short industry name with letters, numbers, spaces, or &()._-.")
        company = groups.setdefault(industry, {}).setdefault(name, {"aliases": [], "exclusions": []})
        for alias in (name, item.get("alias", "")):
            if alias and alias not in company["aliases"]:
                company["aliases"].append(alias)
        exclusion = item.get("exclusion", "")
        if exclusion and exclusion not in company["exclusions"]:
            company["exclusions"].append(exclusion)
    if not groups:
        raise ValueError("The file contains no company names.")
    book = openpyxl.Workbook()
    book.remove(book.active)
    for industry, companies in groups.items():
        sheet = book.create_sheet(industry)
        for company in companies.values():
            values = company["aliases"] + ["NOT " + value for value in company["exclusions"]]
            if industry == "Crypto":
                sheet.append(values)
            else:
                for value in values:
                    sheet.append([value])
                sheet.append([None])
        # Uploaded text must remain text, even if it begins with '='.
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    cell.data_type = "s"
    book.save(target)
    book.close()
    return workbook_roster(target)


def lane_preview(database):
    """Bounded, read-only queries against a single lane; rows remain provisional."""
    path = Path(database)
    if not path.is_file():
        return {"available": False}
    connection = None
    try:
        connection = sqlite3.connect(path.resolve().as_uri() + "?mode=ro", uri=True, timeout=0.15)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA query_only=ON")
        progress = [dict(row) for row in connection.execute(
            "SELECT company,cycle,kind,rows_seen,records_total,fetched_complete,stored_complete,updated_at "
            "FROM api_work_progress ORDER BY updated_at DESC LIMIT 50")]
        columns = {row[1] for row in connection.execute('PRAGMA table_info(contributions)')}
        optional = ','.join(name if name in columns else f'NULL AS {name}'
                            for name in ('contributor_occupation', 'employee_tier'))
        records = [dict(row) for row in connection.execute(
            "SELECT sub_id,contributor_name,company,matched_variant,contributor_employer,"
            "committee_name,contribution_receipt_date,contribution_receipt_amount,company_match_status," + optional + " "
            "FROM contributions ORDER BY rowid DESC LIMIT 25")]
        return {"available": True, "progress": progress, "records": records}
    except sqlite3.Error:
        return {"available": False, "message": "Waiting for a readable collection checkpoint."}
    finally:
        if connection:
            connection.close()


def report_quality(path):
    """Inspect exported headers, not runner exit status, for the direction fix."""
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        warnings = []
        if 'Data Quality' in book.sheetnames:
            headers = next(book['Data Quality'].iter_rows(min_row=5, max_row=5, values_only=True), ())
            for row in book['Data Quality'].iter_rows(min_row=6, values_only=True):
                values = dict(zip(headers, row))
                if values.get('Check') and values.get('Passed') is False:
                    warnings.append({'check': values['Check'], 'severity': values.get('Severity'),
                                     'detail': values.get('Detail')})
        quality = {'warnings': warnings, 'transfer_paths': 'Transfer Paths' in book.sheetnames}
        if 'Candidate Funding' in book.sheetnames:
            headers = next(book['Candidate Funding'].iter_rows(min_row=5, max_row=5, values_only=True), ())
            return {**quality, 'direction_columns': all(name in headers for name in
                ('Conditional Non-opposition', 'Conditional Opposition'))}
        if "Beneficiary Summary" not in book.sheetnames:
            return {**quality, "direction_columns": None}
        headers = next(book["Beneficiary Summary"].iter_rows(min_row=5, max_row=5, values_only=True), ())
        return {**quality, "direction_columns": all(name in headers for name in ("Support Activity", "Oppose Activity"))}
    finally:
        book.close()


def view_industries(value):
    """Distinct industry labels a saved view addresses, first-seen order.

    Read from the three places the save gate validates: the view source,
    each sheet's source, and each sheet's series sources.
    """
    found = {}
    for source in ([value.get('source')]
                   + [sheet.get('source') for sheet in value.get('sheets') or []
                      if isinstance(sheet, dict)]
                   + [item for sheet in value.get('sheets') or [] if isinstance(sheet, dict)
                      for item in sheet.get('sources') or []]):
        if isinstance(source, dict) and isinstance(source.get('industry'), str) and source['industry']:
            found.setdefault(source['industry'], None)
    return list(found)


def validate_view_industries(data, source, sheets):
    """The industry shape a saved view must carry. Membership is NOT checked.

    A saved view has to reload after its snapshot is retired, so the registry
    is deliberately not consulted here; the client reports an unaddressable
    source at open time. What is checked is the shape DE-021 got wrong: a
    field agreed as a string that a producer sends as something else, or
    puts where save_view's record whitelist below silently drops it.
    """
    named = 'Every series source in a saved view needs an industry.'
    if 'industry' in data:
        raise ValueError('A saved view carries its industry inside its data source.')

    def carries(value):
        return isinstance(value, dict) and isinstance(value.get('industry'), str) and bool(value['industry'].strip())

    if 'industry' in source and not carries(source):
        raise ValueError(named)
    for sheet in sheets:
        nested = sheet.get('source')
        if isinstance(nested, dict) and 'industry' in nested and not carries(nested):
            raise ValueError(named)
        if 'sources' in sheet:
            series = sheet.get('sources')
            if not isinstance(series, list) or not all(carries(item) for item in series):
                raise ValueError(named)
            if any(not isinstance(item[key], str)
                   for item in series for key in ('dataset', 'label') if key in item):
                raise ValueError(named)
    return data


def validate_view_board(data, sheets):
    """The board a saved view MAY carry (studio_api.json v2): 1, 2 or 4 tiles.

    Returns the board to store, or None when the view has none. A board is
    written only when the reader chose 2 or 4 tiles, so a v1 view is stored
    exactly as it always was. Everything here is shape: `tiles` is the integer
    1, 2 or 4; `slots` holds exactly that many entries, each a sheet index in
    this view or null for an empty tile, and no sheet sits in two tiles. A
    dangling index would draw a tile from a sheet that is not in the view, so
    it is refused rather than stored and explained later.

    chart-studio-core.js:normalizeBoard applies the same rules in the browser.
    An explicit null is "no board" there and here, not an error.
    """
    if data.get('board') is None:
        return None
    board = data['board']
    refusal = "The saved view's board does not match its sheets."
    if not isinstance(board, dict):
        raise ValueError(refusal)
    tiles = board.get('tiles')
    if isinstance(tiles, bool) or not isinstance(tiles, int) or tiles not in (1, 2, 4):
        raise ValueError(refusal)
    slots = board.get('slots')
    if not isinstance(slots, list) or len(slots) != tiles:
        raise ValueError(refusal)
    placed = set()
    for slot in slots:
        if slot is None:
            continue
        if isinstance(slot, bool) or not isinstance(slot, int) or not 0 <= slot < len(sheets):
            raise ValueError(refusal)
        if slot in placed:
            raise ValueError(refusal)
        placed.add(slot)
    return {'tiles': tiles, 'slots': list(slots)}


# Marks cards (studio_api.json v3, lane M, boards #50-#52). The same rules as
# chart-studio-core.js cardFault / marksCardsFault: a card set that opens in
# the browser is one that saves, and a malformed one is refused before it can
# be stored, the way a malformed board is.
VIEW_MARK_TYPES = frozenset((
    'automatic', 'bar', 'line', 'area', 'square', 'circle', 'shape', 'pie', 'doughnut', 'text', 'box',
    'histogram', 'waterfall', 'pareto', 'radar', 'polygon', 'density', 'map', 'heatmap', 'highlight',
    'stacked', 'dot', 'series-dot', 'treemap', 'bubbles', 'gantt', 'bullet', 'combination', 'dual-line',
    'filled-map'))
_GLYPH_BASES = ('circle', 'square', 'plus', 'times', 'asterisk', 'diamond', 'triangle-up', 'triangle-down',
                'triangle-left', 'triangle-right', 'star')
VIEW_GLYPHS = frozenset(_GLYPH_BASES + tuple(glyph + '-open' for glyph in _GLYPH_BASES))
VIEW_STYLE_COLOURS = frozenset(('automatic', 'none', 'ink', 'white') + tuple('c%d' % n for n in range(1, 13)))
VIEW_AGGREGATIONS = ('sum', 'mean', 'count', 'median', 'min', 'max', 'countd')
VIEW_ENCODINGS = ('color', 'size', 'label', 'detail', 'tooltip', 'shape', 'path', 'angle')
VIEW_CARD_KEYS = frozenset(('id', 'name', 'measure', 'type', 'encodings', 'opacity', 'border', 'halo',
                            'visible', 'selectable', 'zoomExtent', 'source', 'geography', 'map'))

# Lane F (board #46, studio_api.json v4): Flow is a stored card type whose four options
# travel as `flow`, on a Flow card only (chart-studio-core.js flowFault is the same rule).
VIEW_MARK_TYPES = VIEW_MARK_TYPES | {'flow'}
VIEW_CARD_KEYS = VIEW_CARD_KEYS | {'flow'}
VIEW_FLOW_RIBBONS = ('source', 'target', 'outcome', 'gradient')
VIEW_FLOW_ORDERS = ('value', 'alphabetical')


def _flow_ok(card):
    """A Flow card carries ribbons, order, opacity and groupAfter; no other card has them."""
    if card.get('type') != 'flow':
        return 'flow' not in card
    options = card.get('flow')
    if not isinstance(options, dict) or set(options) != {'ribbons', 'order', 'opacity', 'groupAfter'}:
        return False
    if options['ribbons'] not in VIEW_FLOW_RIBBONS or options['order'] not in VIEW_FLOW_ORDERS:
        return False
    if not _whole(options['opacity']) or not 0 <= options['opacity'] <= 100:
        return False
    group = options['groupAfter']
    return _whole(group) and (group == 0 or 2 <= group <= 50)
# Map layers (lane P, boards #47-#55; studio_api.json v3, its map examples). A
# card that names a geography is a layer of a map sheet: its level, its own
# source and its map options, checked by the same rules as map-layers.js
# layerFault and chart-studio-core.js cardFault. The sheet's own `map` holds the
# Map Layers pane (background, washout, base layers, data layer, map options).
VIEW_MAP_LEVELS = frozenset(('state', 'county', 'cd', 'sldu', 'sldl', 'place', 'zcta', 'tract',
                             'hex10', 'hex25', 'hex50', 'point'))
VIEW_MAP_MARKS = frozenset(('automatic', 'filled-map', 'circle', 'shape', 'density', 'text', 'line'))
VIEW_MAP_BASES = frozenset(('receipt', 'original_observed', 'original_supported', 'raw', 'pac_activity'))
VIEW_MAP_RAMPS = frozenset(('teal', 'purple', 'mauve', 'gold', 'heat'))
VIEW_MAP_BACKGROUNDS = frozenset(('boundaries', 'light', 'dark', 'streets', 'none'))
VIEW_MAP_BASE_LAYERS = frozenset(('coastlines', 'state-borders', 'state-names', 'county-borders', 'county-names',
                                  'cd', 'sld', 'zcta', 'cities', 'water', 'parks', 'streets'))
VIEW_MAP_DATA_BY = frozenset(('state', 'county', 'zcta', 'tract'))
VIEW_MAP_OPTIONS = {'scale': bool, 'units': ('us', 'metric'), 'search': bool, 'toolbar': bool,
                    'panZoom': bool, 'layerControl': bool}
VIEW_MAP_KEYS = frozenset(('layers', 'congress', 'background', 'washout', 'base', 'data', 'options'))
# Lane P (2026-09-26, #55's Opposition): a PAC Spent layer may take one side of the spending
# (`stance`: the server's per-place support or oppose sum, never the two netted), and a Filled
# layer may draw with the hatch (`map.pattern`); map-layers.js sourceFault/layerFault, same rules.
VIEW_MAP_STANCES = ('support', 'oppose')
VIEW_MAP_PATTERNS = ('hatch',)


def _whole(value):
    """A JSON number JavaScript's Number.isInteger accepts (85 and 85.0 alike)."""
    if isinstance(value, bool):
        return False
    return isinstance(value, int) or (isinstance(value, float) and value.is_integer())


def _pill_ok(value, glyphs=False):
    if not isinstance(value, dict):
        return False
    if set(value) - ({'column', 'aggregation', 'glyphs'} if glyphs else {'column', 'aggregation'}):
        return False
    if not _whole(value.get('column')) or value['column'] < 0 or value.get('aggregation') not in VIEW_AGGREGATIONS:
        return False
    if glyphs and value.get('glyphs') is not None:
        chosen = value['glyphs']
        if not isinstance(chosen, dict) or len(chosen) > 64:
            return False
        if any(len(member) > 200 or glyph not in VIEW_GLYPHS for member, glyph in chosen.items()):
            return False
    return True


def _shelf_key(value):
    """chart-studio-core.js pill() read as column|aggregation, or None."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, float):
        return '%d|sum' % value if value.is_integer() else None
    if isinstance(value, int):
        return '%d|sum' % value
    if isinstance(value, str):
        return value + '|sum' if re.fullmatch(r'-?\d+', value) else None
    if not isinstance(value, dict):
        return None
    column = value.get('column')
    if 'column' in value and column is None:
        column = 0  # Number(null) is 0 in pill()
    if isinstance(column, bool):
        column = int(column)
    elif isinstance(column, float) and column.is_integer():
        column = int(column)
    elif isinstance(column, str) and re.fullmatch(r'-?\d+', column):
        column = int(column)
    if not isinstance(column, int):
        return None
    aggregation = value.get('aggregation') if value.get('aggregation') in VIEW_AGGREGATIONS else 'sum'
    return '%d|%s' % (column, aggregation)


def _card_ok(card):
    if not isinstance(card, dict) or set(card) - VIEW_CARD_KEYS:
        return False
    if not isinstance(card.get('id'), str) or not 1 <= len(card['id']) <= 64:
        return False
    if not isinstance(card.get('name'), str) or len(card['name']) > 120:
        return False
    if card.get('type') not in VIEW_MARK_TYPES:
        return False
    if not _flow_ok(card):
        return False
    encodings = card.get('encodings')
    if not isinstance(encodings, dict) or set(encodings) != set(VIEW_ENCODINGS):
        return False
    for key in ('color', 'size', 'label', 'path', 'angle'):
        if encodings[key] is not None and not _pill_ok(encodings[key]):
            return False
    for key in ('detail', 'tooltip'):
        items = encodings[key]
        if not isinstance(items, list) or len(items) > 24 or not all(_pill_ok(item) for item in items):
            return False
    shape = encodings['shape']
    if shape is not None:
        if isinstance(shape, dict) and 'glyph' in shape:
            if set(shape) != {'glyph'} or shape['glyph'] not in VIEW_GLYPHS:
                return False
        elif not _pill_ok(shape, glyphs=True):
            return False
    if not _whole(card.get('opacity')) or not 0 <= card['opacity'] <= 100:
        return False
    if card.get('border') not in VIEW_STYLE_COLOURS or card.get('halo') not in VIEW_STYLE_COLOURS:
        return False
    if any(not isinstance(card.get(key), bool) for key in ('visible', 'selectable', 'zoomExtent')):
        return False
    if 'measure' in card and not _pill_ok(card['measure']):
        return False
    if 'source' in card and not _map_source_ok(card['source']):
        return False
    if 'geography' in card and (card['geography'] not in VIEW_MAP_LEVELS or card.get('type') not in VIEW_MAP_MARKS):
        return False
    if 'map' in card and ('geography' not in card or not _card_map_ok(card['map'])):
        return False
    if 'map' in card and 'pattern' in card['map'] and not _card_fills(card):
        return False  # a hatch fills shapes, so it belongs to a Filled layer
    return True


def _text_ok(value, most):
    return isinstance(value, str) and 1 <= len(value.strip()) <= most


def _map_source_ok(source):
    """A map layer's own source: an industry snapshot, addressed by its label
    as D-030's `industry` parameter, or one PAC evidence pane (the
    /api/evidence-geo request: source, id, committee, cycle, direction)."""
    if not isinstance(source, dict):
        return False
    kind = source.get('kind')
    if kind == 'industry':
        return (not set(source) - {'kind', 'industry', 'basis', 'label'} and _text_ok(source.get('industry'), 120)
                and ('basis' not in source or source['basis'] in VIEW_MAP_BASES)
                and ('label' not in source or _text_ok(source['label'], 120)))
    if kind == 'pac':
        return (not set(source) - {'kind', 'source', 'id', 'committee', 'cycle', 'direction', 'name', 'industry', 'stance'}
                and source.get('source') in ('snapshot', 'csv')
                and isinstance(source.get('id'), str) and re.fullmatch(r'[a-z0-9][a-z0-9_-]{0,99}', source['id']) is not None
                and isinstance(source.get('committee'), str) and re.fullmatch(r'C\d{8}', source['committee']) is not None
                and isinstance(source.get('cycle'), str) and re.fullmatch(r'\d{4}', source['cycle']) is not None
                and source.get('direction') in ('raised', 'spent')
                and all(key not in source or _text_ok(source[key], 120) for key in ('name', 'industry'))
                and ('stance' not in source
                     or (source['stance'] in VIEW_MAP_STANCES and source['direction'] == 'spent')))
    return False


def _card_map_ok(options):
    """A layer's own map options: density radius and intensity, its ramp, and its pattern."""
    if not isinstance(options, dict) or set(options) - {'radius', 'intensity', 'ramp', 'pattern'}:
        return False
    if 'pattern' in options and options['pattern'] not in VIEW_MAP_PATTERNS:
        return False
    if 'radius' in options and (not _whole(options['radius']) or not 2 <= options['radius'] <= 80):
        return False
    if 'intensity' in options and (not _whole(options['intensity']) or not 1 <= options['intensity'] <= 100):
        return False
    return 'ramp' not in options or options['ramp'] in VIEW_MAP_RAMPS


def _card_fills(card):
    """map-layers.js markOf(card) is 'filled-map': Filled, or Automatic at any level but points."""
    return card.get('type') == 'filled-map' or (card.get('type') == 'automatic' and card.get('geography') != 'point')


def _sheet_map_ok(settings):
    """A sheet's `map`: today's boundary choices (layers, congress) plus the Map
    Layers pane of a map whose cards are its layers. Nothing in it is a layer."""
    if not isinstance(settings, dict) or set(settings) - VIEW_MAP_KEYS:
        return False
    if 'layers' in settings and (not isinstance(settings['layers'], list)
                                 or not all(isinstance(item, str) for item in settings['layers'])):
        return False
    if 'congress' in settings and not isinstance(settings['congress'], str):
        return False
    if 'background' in settings and settings['background'] not in VIEW_MAP_BACKGROUNDS:
        return False
    if 'washout' in settings and (not _whole(settings['washout']) or not 0 <= settings['washout'] <= 100):
        return False
    if 'base' in settings:
        base = settings['base']
        if (not isinstance(base, list) or len(base) > len(VIEW_MAP_BASE_LAYERS) or len(set(map(str, base))) != len(base)
                or not all(item in VIEW_MAP_BASE_LAYERS for item in base)):
            return False
    if 'data' in settings and settings['data'] is not None:
        data = settings['data']
        if (not isinstance(data, dict) or set(data) != {'measure', 'by', 'using'} or not _pill_ok(data['measure'])
                or data['by'] not in VIEW_MAP_DATA_BY or data['using'] not in VIEW_MAP_RAMPS):
            return False
    if 'options' in settings:
        options = settings['options']
        if not isinstance(options, dict) or set(options) - set(VIEW_MAP_OPTIONS):
            return False
        for key, rule in VIEW_MAP_OPTIONS.items():
            if key in options and not (isinstance(options[key], bool) if rule is bool else options[key] in rule):
                return False
    return True


def validate_view_marks(sheets):
    """The marks cards a saved sheet MAY carry (studio_api.json v3).

    Absent is v2: the sheet's `marks` is its one card, stored as it always was.
    Present, `marksCards` is 1 to 12 well-formed cards with unique ids; two or
    more are a dual axis, where every card names a distinct measure that sits on
    the sheet's Rows (a card for a measure the shelves do not hold would draw an
    axis from nothing, so it is refused like a dangling board slot).

    Lane P: a sheet whose every card names a geography is a map, and its cards
    are its layers (draw order = card order), each with its own source, so the
    dual-axis rule does not apply to it; a card set mixing layers and axis
    cards is refused. A sheet's `map` settings must be well formed.
    """
    refusal = "The saved view's marks cards are malformed."
    for sheet in sheets:
        if 'map' in sheet and not _sheet_map_ok(sheet['map']):
            raise ValueError("The saved view's map settings are malformed.")
        if 'syncAxes' in sheet and not isinstance(sheet['syncAxes'], bool):
            raise ValueError(refusal)
        if 'marksCards' not in sheet:
            continue
        cards = sheet['marksCards']
        if not isinstance(cards, list) or not 1 <= len(cards) <= 12 or not all(_card_ok(card) for card in cards):
            raise ValueError(refusal)
        if len({card['id'] for card in cards}) != len(cards):
            raise ValueError(refusal)
        if len(cards) >= 2 and any(card['type'] == 'flow' for card in cards):
            raise ValueError(refusal)  # lane F: a flow draws alone, sharing no axis and no map
        layers = ['geography' in card for card in cards]
        if any(layers) and not all(layers):
            raise ValueError(refusal)
        if len(cards) >= 2 and not all(layers):
            rows = sheet.get('rows') if isinstance(sheet.get('rows'), list) else []
            on_rows = {key for key in (_shelf_key(item) for item in rows) if key}
            measures = ['%d|%s' % (card['measure']['column'], card['measure']['aggregation'])
                        if 'measure' in card else '' for card in cards]
            if not all(measures) or len(set(measures)) != len(measures) or not set(measures) <= on_rows:
                raise ValueError(refusal)
    return sheets


class Conflict(Exception):
    """A well-formed request the current state cannot accept yet."""


class RunNotFound(Exception):
    """A stop/workers request named a run id with no matching upload."""


class Console:
    def __init__(self, root, keys_file="", rscript="", read_only=False):
        self.configured_root = os.path.abspath(root)
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.report_directory = report_directory(PROJECT)
        self.keys_file = keys_file
        self.read_only = read_only
        # The pocket publisher reads this console over HTTP, so the console
        # needs the port it was given. main() sets it before serve_forever.
        self.port = 0
        self.rscript = rscript or mentat_env.seam("MENTAT_RSCRIPT")
        self.token = secrets.token_urlsafe(32)
        self.lock = threading.Lock()
        self.processes = {}
        self.preparing = set()
        self.jobs = set()
        self.report_cache = {}
        configured = environment_keys()
        if self.keys_file:
            configured.append(Path(self.keys_file).read_text(encoding="utf-8-sig"))
        if configured and not read_only:
            save_keys(configured)

    def busy(self):
        """True while this console owns work that stopping it would interrupt."""
        if self.preparing or self.jobs:
            return True
        return any(process.poll() is None for process in list(self.processes.values()))

    def restart(self, port):
        """Hand this console's port back to the launcher, which replaces it.

        The launcher stops this process only after it confirms the identity,
        the code hash and an idle busy flag, then binds the same port again.
        """
        script = Path(__file__).resolve().parent / 'Start-MentatApp.ps1'
        if not script.is_file():
            raise ValueError('The launcher script is unavailable. Close this window and open MENTAT again.')
        shell = shutil.which('powershell') or str(Path(os.environ.get('SystemRoot', 'C:/Windows'))
                                                  / 'System32/WindowsPowerShell/v1.0/powershell.exe')
        command = [shell, '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass', '-File', str(script),
                   '-NoBrowser', '-Port', str(int(port)), '-DataRoot', self.configured_root]
        # Identity is project path plus data root; a restart that dropped the
        # data root would open a second console on another port instead.
        if self.keys_file:
            command += ['-KeysFile', self.keys_file]
        if self.rscript:
            command += ['-RScriptPath', self.rscript]
        # The replacement keeps this process's interpreter. Without it the launcher
        # chooses again (Mentat's venv, the Codex runtime, PATH), and a console started
        # on a named Python could come back on another one.
        command += ['-PythonPath', sys.executable]
        # DETACHED_PROCESS would leave powershell.exe without a console: it
        # exits 0 without running the launcher, so the restart silently fails.
        flags = subprocess.CREATE_NO_WINDOW | subprocess.CREATE_NEW_PROCESS_GROUP if os.name == 'nt' else 0
        try:
            subprocess.Popen(command, cwd=str(PROJECT), stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                             stderr=subprocess.DEVNULL, close_fds=True, creationflags=flags)
        except OSError:
            raise ValueError('The launcher could not be started. Close this window and open MENTAT again.')
        return {'restarting': True, 'port': int(port), 'pid': os.getpid()}

    def inspect_report(self, path):
        stat = path.stat()
        key = (str(path), stat.st_mtime_ns, stat.st_size)
        if key not in self.report_cache:
            try:
                value = report_quality(path)
                with path.open("rb") as stream:
                    value["sha256"] = hashlib.file_digest(stream, "sha256").hexdigest()
                self.report_cache[key] = value
            except (OSError, ValueError, KeyError, zipfile.BadZipFile):
                return {"direction_columns": None, "warning": "Workbook inspection unavailable."}
        return self.report_cache[key]

    def acknowledge_quality(self, data):
        identifier = str(data.get('id', ''))
        requested = data.get('warning_ids', [])
        if not isinstance(requested, list) or not requested or not all(isinstance(i, str) for i in requested):
            raise ValueError('Choose the report warnings currently shown.')
        with self.lock:
            current = self.snapshot(identifier)
            eligible = {w['id'] for report in current['reports']
                        for w in report.get('quality', {}).get('warnings', [])
                        if str(w.get('severity', '')).lower() != 'fail'}
            if not set(requested) <= eligible:
                raise ValueError('The reports changed. Refresh and review their current warnings.')
            path = self.root / 'quality-reviews.json'
            reviews = read_json(path)
            stamp = datetime.datetime.now(datetime.timezone.utc).isoformat()
            for warning_id in requested:
                reviews[warning_id] = {'reviewed_at': stamp, 'run_id': identifier}
            atomic_json(path, reviews)
        return {'reviewed': len(set(requested))}

    def upload(self, data):
        content = base64.b64decode(data.get("content", ""), validate=True)
        if not content or len(content) > 5_000_000:
            raise ValueError("Upload a nonempty file smaller than 5 MB.")
        identifier = uuid.uuid4().hex
        directory = self.root / identifier
        directory.mkdir()
        readiness = normalize_upload(str(data.get("filename", "")), content, directory / "roster.xlsx")
        manifest = {"id": identifier, "name": Path(str(data["filename"])).name,
                    "created_at": time.time(), "state": "ready" if readiness["status"] == "succeeded" else "needs review",
                    "readiness": readiness}
        atomic_json(directory / "upload.json", manifest)
        return manifest

    def create_roster(self, data):
        """Use the same persisted readiness/launch path as file uploads."""
        sector = data.get('sector')
        companies = data.get('companies')
        if not isinstance(sector, str) or not sector.strip():
            raise ValueError('Enter a sector for this collection.')
        if not isinstance(companies, list) or not 1 <= len(companies) <= 1000:
            raise ValueError('Enter between 1 and 1,000 companies.')
        output = io.StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(('canonical_name', 'alias', 'sector'))
        seen = set()
        for index, company in enumerate(companies, 1):
            if not isinstance(company, dict):
                raise ValueError(f'Company {index}: enter a company and its alternative names.')
            name, aliases = company.get('name'), company.get('aliases', [])
            if not isinstance(name, str) or not name.strip():
                raise ValueError(f'Company {index}: a canonical name is required.')
            name = name.strip()
            normalized = re.sub(r'\s+', ' ', name).casefold()
            if normalized in seen:
                raise ValueError(f'{name} appears more than once. Combine its alternative names in one company.')
            seen.add(normalized)
            if not isinstance(aliases, list) or len(aliases) > 100 or not all(isinstance(a, str) for a in aliases):
                raise ValueError(f'{name}: enter up to 100 alternative names, one per line.')
            for alias in dict.fromkeys(['', *[a.strip() for a in aliases if a.strip()]]):
                writer.writerow((name, alias, sector.strip()))
        return self.upload({'filename': sector.strip() + ' roster.csv',
                            'content': base64.b64encode(output.getvalue().encode('utf-8')).decode('ascii')})

    def launch(self, data, background=False):
        identifier = data.get("id", "")
        if not re.fullmatch(r"[a-f0-9]{32}", identifier):
            raise ValueError("Choose a validated upload.")
        directory = self.root / identifier
        cycles = data.get("cycles", [])
        if not cycles or any(type(c) is not int or c not in (2022, 2024, 2026) for c in cycles):
            raise ValueError("Select at least one supported election cycle.")
        with self.lock, console_state_lock(self.root):
            manifest = read_json(directory / "upload.json")
            if manifest.get("state") != "ready":
                raise ValueError("This upload is not ready, or has already been started.")
            archives = read_json(self.root / 'run-archives.json')
            if identifier in (archives.get('runs', {}) if isinstance(archives, dict) else {}):
                raise ValueError('Restore this saved roster from Archived runs before starting collection.')
            if self.preparing or any(process.poll() is None for process in self.processes.values()):
                raise ValueError("An app collection is already running. Wait for it to finish.")
            if any(run['state'] in ('preparing', 'starting', 'running') and
                   process_alive(run.get('upload', {}).get('pid') or run.get('upload', {}).get('console_pid'))
                   for run in self.runs()):
                raise ValueError("A saved collection is still active. Wait for it to finish.")
            if not Path(self.rscript).is_file():
                raise ValueError("Rscript is unavailable. Set -RScriptPath in the launcher.")
            api_key = str(data.get("api_key", "")).strip()
            if api_key:
                save_keys([api_key])
            stored_keys = load_keys()
            if not stored_keys:
                raise ValueError("Enter an FEC API key before starting collection.")
            manifest.update(state="preparing", cycles=cycles, started_at=time.time(), console_pid=os.getpid())
            atomic_json(directory / "upload.json", manifest)
            self.preparing.add(identifier)
        if background:
            threading.Thread(target=self.prepare_launch, args=(identifier, directory, cycles, stored_keys, True), daemon=True).start()
            return {"id":identifier, "state":"preparing"}
        self.prepare_launch(identifier, directory, cycles, stored_keys)
        return {"id":identifier, "state":"starting"}

    def prepare_launch(self, identifier, directory, cycles, stored_keys, background=False, resume_run_id=""):
        try:
            self.start_prepared_process(identifier, directory, cycles, stored_keys, resume_run_id)
        except Exception as error:
            with self.lock, console_state_lock(self.root):
                manifest = read_json(directory / "upload.json")
                manifest.update(state="launch failed", finished_at=time.time(),
                                launch_error=f'Runtime preparation or launch failed ({type(error).__name__}).')
                atomic_json(directory / "upload.json", manifest)
            if not background:
                raise
        finally:
            self.preparing.discard(identifier)

    def start_prepared_process(self, identifier, directory, cycles, stored_keys, resume_run_id=""):
        runtime = directory / "runtime" / "mentat"
        runtime.mkdir(parents=True, exist_ok=True)
        # Pin the current runner and inputs together; R resolves its roster relative to its entry point.
        active_sources = read_json(PROJECT / "inputs/active_fec_sources.json")
        for folder in ("R", "tools", "inputs"):
            shutil.copytree(PROJECT / folder, runtime / folder, dirs_exist_ok=True,
                            ignore=shutil.ignore_patterns("__pycache__", "node_modules", ".cache", "*.bak*",
                                    *(("fec_beneficiary", "fec_committee_master") if folder == "inputs" else ())))
        shutil.copy2(PROJECT / "mentat.R", runtime / "mentat.R")
        # The path-only resolver is shared by the console and copied collectors;
        # no model data or other personal app assets are staged into this folder.
        (runtime / 'app').mkdir(exist_ok=True)
        shutil.copy2(Path(__file__).parent / 'runtime_paths.py', runtime / 'app' / 'runtime_paths.py')
        shutil.copy2(directory / "roster.xlsx", runtime / "inputs" / "Company Name.xlsx")
        source_root = Path(active_sources['path']) if active_sources.get('path') else PROJECT / 'inputs'
        copy_fec_sources(source_root, runtime / 'inputs', cycles)
        command = ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                   str(runtime / "tools/parallel/Start-MentatResilientIndustryReports.ps1"),
                   "-Cycles", "|".join(str(c) for c in sorted(set(cycles))),
                   "-RunRoot", str(directory / "runs"), "-OutputDirectory", str(self.report_directory),
                   "-RawDataDirectory", str(directory / "raw"), "-PythonPath", sys.executable,
                   "-RScriptPath", self.rscript, "-NoCache"]
        if resume_run_id:
            # Continue the existing run folder. The orchestrator keeps the run
            # id, appends to its logs and events, and skips initialization for
            # every industry whose saved lane state a fanout resume can reuse.
            # The runtime above is re-staged either way, so a code fix landed
            # since the stopped attempt is the code the resume runs.
            command += ["-ResumeRunId", resume_run_id]
        environment = collector_environment(os.environ, runtime)
        library = r_library_path(environment, PROJECT, self.rscript)
        if library:
            command.extend(['-RLibraryPath', library])
            environment['R_LIBS_USER'] = library
        if os.name == 'nt':
            # Python can inherit PowerShell 7's module path without the native
            # shell's compatibility adjustment. Windows PowerShell must load
            # its own Utility module first (including Get-FileHash).
            inherited_modules = ''
            for name in list(environment):
                if name.upper() == 'PSMODULEPATH':
                    inherited_modules = environment.pop(name)
            native_modules = Path(os.environ.get('SystemRoot', 'C:/Windows')) / 'System32/WindowsPowerShell/v1.0/Modules'
            environment['PSModulePath'] = str(native_modules) + (';' + inherited_modules if inherited_modules else '')
        environment["MENTAT_INPUT_DIR"] = str(runtime / "inputs")
        environment["MENTAT_CUSTOM_ROSTER"] = "1"
        environment["MENTAT_REPORT_DIRECTORY"] = str(self.report_directory)
        environment["MENTAT_CROSS_REFERENCE_ROOT"] = str(PROJECT / "state" / "cross-reference")
        environment["MENTAT_ACCEPTED_SOURCE_ROOT"] = str(PROJECT / "state" / "accepted-report-sources")
        environment["MENTAT_API_BUDGET_DB"] = str(budget_path(PROJECT, environment=environment))
        environment["MENTAT_API_BUDGET_IMPORTS"] = json.dumps(
            [str(path) for path in legacy_budget_paths(PROJECT, environment=environment)])
        environment["MENTAT_FEC_SOURCE_ROOT"] = str(runtime / 'inputs')
        for name in list(environment):
            if name == 'FEC_API_KEY' or re.fullmatch(r'FEC_API_KEY_\d+', name):
                environment.pop(name)
        environment["FEC_API_KEYS"] = "|".join(stored_keys)
        with self.lock, console_state_lock(self.root):
            stored_keys = load_keys()
            if not stored_keys:
                raise ValueError('No enabled FEC keys remain. Save an enabled key before launching again.')
            environment['FEC_API_KEYS'] = '|'.join(stored_keys)
            manifest = read_json(directory / "upload.json")
            if manifest.get("stop_requested"):
                manifest.update(state="stopped", finished_at=time.time())
                atomic_json(directory / "upload.json", manifest)
                return
            manifest.update(state="starting")
            atomic_json(directory / "upload.json", manifest)
            try:
                with (directory / "launcher.log").open("ab" if resume_run_id else "wb") as log:
                    process = subprocess.Popen(command, cwd=runtime, env=environment, stdout=log,
                                               stderr=subprocess.STDOUT,
                                               creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            except OSError:
                manifest["state"] = "launch failed"
                atomic_json(directory / "upload.json", manifest)
                raise ValueError("The collection process could not start. Check the local launcher log.")
            self.processes[identifier] = process
            manifest["pid"] = process.pid
            manifest['process_identity'] = process_identity(process.pid)
            atomic_json(directory / "upload.json", manifest)
            threading.Thread(target=self.finish, args=(process, directory), daemon=True).start()

    def finish(self, process, directory):
        code = process.wait()
        with self.lock, console_state_lock(self.root):
            manifest = read_json(directory / "upload.json")
            manifest.update(state="stopped" if manifest.get("stop_requested") else "completed" if code == 0 else "failed", exit_code=code, finished_at=manifest.get("finished_at") or time.time())
            atomic_json(directory / "upload.json", manifest)

    def control_targets(self, identifier):
        """Run-level control.json plus every per-industry fanout dir the run has already reached."""
        directory = self.root / identifier
        targets = [directory / "control.json"]
        for fanout in sorted(directory.glob("runs/*/industries/*/fanout")):
            if fanout.is_dir():
                targets.append(fanout / "control.json")
        return targets

    def resume_control_targets(self, identifier, run_id):
        """Every control file a resume of one saved run has to clear.

        stop() writes {"desired_workers": 0, "stop": true} to the run-level
        control file and to every fanout control file the run had reached. A
        relaunched orchestrator reads those same files, so a resume that left
        them alone would stop itself within three seconds of starting.
        """
        directory = self.root / identifier
        targets = [directory / "control.json"]
        run = directory / "runs" / run_id
        if (run / "control.json").is_file():
            targets.append(run / "control.json")
        for fanout in sorted(run.glob("industries/*/fanout")):
            if fanout.is_dir():
                targets.append(fanout / "control.json")
        return targets

    def resume_target(self, run):
        """The run directory a stopped or failed collection could continue.

        None means this run has nothing to resume: it is not an uploaded
        roster, it never reached a terminal state a resume applies to, or it
        holds no fanout manifest for the lane snapshots to be reloaded from.
        """
        if not isinstance(run, dict) or not run.get('upload'):
            return None
        if run.get('state') not in ('stopped', 'failed'):
            return None
        directory = run.get('directory')
        return directory if run_has_lane_state(directory) else None

    def resume(self, data, background=True):
        """Continue a stopped or failed collection in its own run folder.

        Every verified company-cycle task the stopped attempt finished stays
        finished: the fanout reloads its lane databases and re-runs only what
        did not succeed. Merge semantics are untouched - the industry source
        database is still written only after every task in the plan passes.
        """
        identifier = str(data.get("id", ""))
        if not re.fullmatch(r"[a-f0-9]{32}", identifier):
            raise RunNotFound("Run not found.")
        directory = self.root / identifier
        with self.lock, console_state_lock(self.root):
            run = self.find(identifier)
            if run is None or not run.get('upload'):
                raise RunNotFound("Run not found.")
            target = self.resume_target(run)
            if target is None:
                raise ValueError("Nothing to resume for this roster.")
            archives = read_json(self.root / 'run-archives.json')
            if identifier in (archives.get('runs', {}) if isinstance(archives, dict) else {}):
                raise ValueError('Restore this saved roster from Archived runs before resuming collection.')
            if self.preparing or any(process.poll() is None for process in self.processes.values()):
                raise ValueError("An app collection is already running. Wait for it to finish.")
            if any(other['state'] in ('preparing', 'starting', 'running') and
                   process_alive(other.get('upload', {}).get('pid') or other.get('upload', {}).get('console_pid'))
                   for other in self.runs()):
                raise ValueError("A saved collection is still active. Wait for it to finish.")
            if not Path(self.rscript).is_file():
                raise ValueError("Rscript is unavailable. Set -RScriptPath in the launcher.")
            stored_keys = load_keys()
            if not stored_keys:
                raise ValueError("Enter an FEC API key before resuming collection.")
            manifest = read_json(directory / "upload.json")
            cycles = [cycle for cycle in manifest.get("cycles", []) if type(cycle) is int]
            if not cycles:
                raise ValueError("Nothing to resume for this roster.")
            run_id = Path(target).name
            stamp = datetime.datetime.now(datetime.timezone.utc).isoformat()
            # Clear the stop flags before the orchestrator can read them. A
            # worker count the operator chose is kept; the zero stop() wrote is
            # dropped rather than saved, so the fanout picks its own start.
            for control_path in self.resume_control_targets(identifier, run_id):
                previous = read_json(control_path)
                previous = previous if isinstance(previous, dict) else {}
                workers = previous.get("desired_workers")
                control = {"stop": False}
                if type(workers) is int and workers > 0:
                    control["desired_workers"] = workers
                control["updated_at"] = stamp
                atomic_json(control_path, control)
            manifest.update(state="preparing", cycles=cycles, resumed_from=run_id,
                            resume_count=int(manifest.get("resume_count") or 0) + 1,
                            started_at=time.time(), console_pid=os.getpid())
            # A stop flag or an exit code left in the manifest would make
            # start_prepared_process refuse the launch, or finish() record the
            # relaunch as stopped the moment it exits.
            for key in ("stop_requested", "stop_mode", "finished_at", "exit_code",
                        "launch_error", "pid", "process_identity"):
                manifest.pop(key, None)
            atomic_json(directory / "upload.json", manifest)
            self.preparing.add(identifier)
        if background:
            threading.Thread(target=self.prepare_launch,
                             args=(identifier, directory, cycles, stored_keys, True),
                             kwargs={"resume_run_id": run_id}, daemon=True).start()
            return {"id": identifier, "state": "preparing", "resumed_from": run_id}
        self.prepare_launch(identifier, directory, cycles, stored_keys, resume_run_id=run_id)
        return {"id": identifier, "state": "starting", "resumed_from": run_id}

    def stop(self, data):
        identifier = str(data.get("id", ""))
        force = bool(data.get("force", True))
        if not re.fullmatch(r"[a-f0-9]{32}", identifier):
            raise RunNotFound("Run not found.")
        directory = self.root / identifier
        with self.lock, console_state_lock(self.root):
            manifest = read_json(directory / "upload.json")
            if not manifest:
                raise RunNotFound("Run not found.")
            if manifest.get('state') not in ('preparing', 'starting', 'running'):
                raise ValueError('This collection is not active; no process was stopped.')
            stamp = datetime.datetime.now(datetime.timezone.utc).isoformat()
            control = {"desired_workers": 0, "stop": True, "updated_at": stamp}
            files = 0
            for target in self.control_targets(identifier):
                atomic_json(target, control)
                files += 1
            # Write the stop flag everywhere before any kill attempt, so a process that
            # survives taskkill - or a server crash mid-kill - still sees stop:true.
            manifest.update(stop_requested=stamp, stop_mode="force" if force else "graceful")
            atomic_json(directory / "upload.json", manifest)
            if manifest.get('state') == 'preparing' and not manifest.get('pid'):
                return {"ok":True, "mode":"preparation cancellation requested", "files":files}
            result = {"ok": True, "mode": "graceful", "files": files}
            if force:
                process = self.processes.get(identifier)
                pid = process.pid if process is not None else manifest.get("pid")
                result["mode"] = "force"
                result["taskkill_returncode"] = None
                if not pid or not manifest.get('process_identity'):
                    result["ok"] = False
                    result['mode'] = 'graceful'
                    result["message"] = "Graceful stop requested. This run has no verified process identity, so force-stop was withheld."
                    return result
                with process_identity_handle(pid) as identity:
                    if identity is None or identity != manifest['process_identity']:
                        result.update(ok=False, mode='graceful', message='Graceful stop requested. The original process identity no longer matches; no process was killed.')
                        return result
                    kill = subprocess.run(["taskkill", "/PID", str(pid), "/T", "/F"], capture_output=True)
                result["taskkill_returncode"] = kill.returncode
                if kill.returncode == 0:
                    # Mirror finish()'s state/exit_code/finished_at shape so snapshot()/runs()
                    # stop presenting this run as live and Start controls unblock right away.
                    manifest.update(state="stopped", exit_code=process.poll() if process is not None else None,
                                    finished_at=time.time())
                    atomic_json(directory / "upload.json", manifest)
                else:
                    result["ok"] = False
                    result["message"] = "taskkill failed: " + kill.stderr.decode("utf-8", "replace").strip()
        return result

    def set_workers(self, data):
        workers = data.get("workers")
        if type(workers) is not int or workers < 1 or workers > 64:
            raise ValueError("Choose a worker count between 1 and 64.")
        identifier = str(data.get("id", ""))
        if not re.fullmatch(r"[a-f0-9]{32}", identifier):
            raise RunNotFound("Run not found.")
        directory = self.root / identifier
        with self.lock:
            manifest = read_json(directory / "upload.json")
            if not manifest:
                raise RunNotFound("Run not found.")
            stamp = datetime.datetime.now(datetime.timezone.utc).isoformat()
            files = 0
            for target in self.control_targets(identifier):
                control = read_json(target)
                control = control if isinstance(control, dict) else {}
                control["desired_workers"] = workers
                control.setdefault("stop", False)
                control["updated_at"] = stamp
                atomic_json(target, control)
                files += 1
        return {"ok": True, "desired_workers": workers, "files": files}

    def collect_evidence(self, data):
        company = str(data.get('company', '')).strip()
        aliases = [str(v).strip() for v in data.get('aliases', [])] if isinstance(data.get('aliases'), list) else []
        committees = list(dict.fromkeys(re.split(r'[\s,]+', str(data.get('committees', '')).strip())))
        cycles = list(dict.fromkeys(re.split(r'[\s,]+', str(data.get('cycles', '')).strip())))
        if not company or len(company) > 250 or len(aliases) > 50 or any(len(v) > 250 or '|' in v for v in aliases):
            raise ValueError('Enter a company and up to 50 aliases, each under 250 characters.')
        if not committees or any(not re.fullmatch(r'C\d{8}', v) for v in committees):
            raise ValueError('Enter comma-separated FEC committee IDs such as C00000001.')
        if not cycles or any(not re.fullmatch(r'\d{4}', v) or int(v) < 1980 or int(v) > 2100 or int(v) % 2 for v in cycles):
            raise ValueError('Enter even election years, separated by commas.')
        if len(committees) * len(cycles) > 30:
            raise ValueError('Choose up to 30 committee/cycle combinations.')
        if not load_keys():
            raise ValueError('Save an enabled FEC key in Loading before collecting evidence.')
        identifier = 'snapshot-' + uuid.uuid4().hex
        root = PROJECT / '.cache/pac-evidence' / identifier
        root.mkdir(parents=True)
        job = {'id':identifier, 'company':company, 'aliases':aliases, 'committees':committees,
               'cycles':cycles, 'state':'collecting', 'pid':os.getpid(), 'started_at':time.time()}
        atomic_json(root / 'collection.json', job)
        def collect():
            try:
                with (root / 'collection.log').open('w', encoding='utf-8') as log:
                    process = subprocess.Popen([sys.executable, str(PROJECT / 'tools/pac_evidence.py'),
                        '--company=' + company, '--aliases=' + '|'.join(aliases), '--committees', ','.join(committees),
                        '--cycles', ','.join(cycles), '--output', str(root)], stdout=log, stderr=subprocess.STDOUT,
                        creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
                    job['pid'] = process.pid
                    atomic_json(root / 'collection.json', job)
                    code = process.wait()
                ready = code == 0 and read_json(root / 'evidence.json').get('version') == 1
                job['state'] = 'completed' if ready else 'failed'
                if ready:
                    job['message'] = 'Snapshot ready.'
                else:
                    # Name the failure in the status line. "Did not finalize a
                    # snapshot" told the reader nothing and left the reason in a
                    # file nobody opened.
                    try:
                        said = [line.strip() for line in (root / 'collection.log').read_text(
                            encoding='utf-8', errors='replace').splitlines() if line.strip()]
                    except OSError:
                        said = []
                    job['message'] = f'Collection failed: {said[-1] if said else "no output"}. Full log: {root / "collection.log"}'
            except OSError as error:
                job['state'] = 'failed'
                job['message'] = f'Could not start the evidence collector: {error}'
            job['finished_at'] = time.time()
            atomic_json(root / 'collection.json', job)
        def run_collect():
            try:
                collect()
            finally:
                self.jobs.discard(identifier)
        self.jobs.add(identifier)
        threading.Thread(target=run_collect, daemon=True).start()
        return {'id':identifier, 'state':'collecting'}

    # Chart Studio saved views. They live under the data root beside the runs,
    # so a second console opened on the same data root lists the same views;
    # that is what "save here, open on the mirror console" means. A read-only
    # preview refuses the writes at the route, before the body is parsed.

    def studio_directory(self, *parts):
        return self.root.joinpath('studio', *parts)

    def list_views(self):
        """Saved views for the picker, newest first, without their sheets."""
        directory = self.studio_directory('views')
        values = []
        if directory.is_dir():
            for path in directory.glob('*.json'):
                value = read_json(path)
                if value.get('id') != path.stem or not isinstance(value.get('sheets'), list):
                    continue
                values.append({'id': value['id'], 'name': value.get('name', ''),
                               'saved_at': value.get('saved_at', 0),
                               'sourceKey': value.get('sourceKey', ''),
                               'source': value.get('source', {}),
                               'industries': view_industries(value),
                               'sheet_count': len(value['sheets'])})
        values.sort(key=lambda item: (-(item['saved_at'] or 0), item['id']))
        return {'views': values}

    def get_view(self, identifier):
        """One saved view with its sheets, or None when this console has no such view."""
        if not re.fullmatch(r'[a-f0-9]{32}', identifier or ''):
            raise ValueError('Unknown saved view.')
        value = read_json(self.studio_directory('views', identifier + '.json'))
        if value.get('id') != identifier or not isinstance(value.get('sheets'), list):
            return None
        view = {'id': value['id'], 'name': value.get('name', ''), 'saved_at': value.get('saved_at', 0),
                'sourceKey': value.get('sourceKey', ''), 'source': value.get('source', {}),
                'sheets': value['sheets'], 'active': value.get('active', 0)}
        # Echoed exactly as saved, and absent when it was never saved.
        if 'board' in value:
            view['board'] = value['board']
        return view

    def save_view(self, data):
        name = str(data.get('name', '')).strip()
        if not 1 <= len(name) <= 80:
            raise ValueError('View name must be 1 to 80 characters.')
        source_key, source = data.get('sourceKey'), data.get('source')
        if not isinstance(source_key, str) or not isinstance(source, dict):
            raise ValueError('Saved view is missing its data source.')
        sheets = data.get('sheets')
        if not isinstance(sheets, list) or not 1 <= len(sheets) <= 24:
            raise ValueError('A view holds 1 to 24 sheets.')
        if any(not isinstance(sheet, dict) or not str(sheet.get('name', '')).strip() for sheet in sheets):
            raise ValueError('Every sheet in a saved view needs a name.')
        active = data.get('active')
        if not isinstance(active, int) or isinstance(active, bool):
            raise ValueError('Active sheet must be a sheet number.')
        validate_view_industries(data, source, sheets)
        board = validate_view_board(data, sheets)
        validate_view_marks(sheets)
        identifier = uuid.uuid4().hex
        value = {'id': identifier, 'name': name, 'saved_at': time.time(), 'sourceKey': source_key,
                 'source': source, 'sheets': sheets, 'active': active}
        # The record whitelist above is the whole record: a key it does not name
        # is dropped in silence. The board is added only when there is one, so a
        # view saved without one is the same bytes on disk it has always been.
        if board is not None:
            value['board'] = board
        directory = self.studio_directory('views')
        with self.lock:
            directory.mkdir(parents=True, exist_ok=True)
            atomic_json(directory / (identifier + '.json'), value)
        return {'id': identifier, 'saved_at': value['saved_at']}

    def delete_view(self, data):
        """Remove a saved view. Deleting one that is already gone is success."""
        identifier = str(data.get('id', ''))
        if not re.fullmatch(r'[a-f0-9]{32}', identifier):
            raise ValueError('Unknown saved view.')
        with self.lock:
            try:
                (self.studio_directory('views') / (identifier + '.json')).unlink()
            except FileNotFoundError:
                pass
        return {'deleted': True}

    # The briefing tray: one shared tray per data root (Paul's v1 ruling), so
    # every console on this data root shows the same items in the same order.

    def briefing(self):
        """The ordered tray. Missing or damaged index files read as empty."""
        value = read_json(self.studio_directory('briefing', 'index.json'))
        items = value.get('items')
        return {'items': [item for item in items if isinstance(item, dict)] if isinstance(items, list) else []}

    def briefing_item(self, identifier):
        """The stored SVG for one item, or None when this console has no such item."""
        if not re.fullmatch(r'[a-f0-9]{32}', identifier or ''):
            raise ValueError('Unknown briefing item.')
        if not any(item.get('id') == identifier for item in self.briefing()['items']):
            return None
        try:
            return self.studio_directory('briefing', identifier + '.svg').read_text(encoding='utf-8')
        except OSError:
            return None

    def add_to_briefing(self, data):
        svg = data.get('svg')
        if not isinstance(svg, str) or not svg.startswith('<svg'):
            raise ValueError('Briefing item is missing its chart.')
        if len(svg.encode('utf-8')) > 6 * 1024 * 1024:
            raise ValueError('Chart is over 6 MB. Narrow the view first.')

        def dimension(value):
            try:
                return max(0, int(round(float(value))))
            except (TypeError, ValueError):
                return 0

        identifier = uuid.uuid4().hex
        item = {'id': identifier, 'title': str(data.get('title', '')).strip(),
                'scope': str(data.get('scope', '')).strip(),
                # The short scope chip every reader-facing surface prints first.
                # Dropping it here left the drawer, the printed sheet, the HTML
                # bundle and the slide notes falling back to the full provenance
                # string the moment the tray was read back off this console.
                'chip': str(data.get('chip', '')).strip(),
                'note': str(data.get('note', '')).strip(),
                'basis': str(data.get('basis', '')).strip(), 'sheetId': str(data.get('sheetId', '')).strip(),
                'sheetName': str(data.get('sheetName', '')).strip(),
                'source': data.get('source') if isinstance(data.get('source'), dict) else {},
                'provenance': data.get('provenance') if isinstance(data.get('provenance'), dict) else {},
                'width': dimension(data.get('width')), 'height': dimension(data.get('height')),
                'added_at': time.time(), 'origin': str(data.get('origin', '')).strip()}
        directory = self.studio_directory('briefing')
        with self.lock:
            directory.mkdir(parents=True, exist_ok=True)
            items = self.briefing()['items']
            if len(items) >= 40:
                raise Conflict('Briefing is full (40 charts). Remove one first.')
            temporary = directory / (identifier + '.svg.' + uuid.uuid4().hex + '.tmp')
            temporary.write_text(svg, encoding='utf-8')
            replace_with_retry(temporary, directory / (identifier + '.svg'))
            items.append(item)
            atomic_json(directory / 'index.json', {'items': items})
        return {'id': identifier, 'added_at': item['added_at']}

    def remove_from_briefing(self, data):
        """Remove one item. Removing one that is already gone is success."""
        identifier = str(data.get('id', ''))
        if not re.fullmatch(r'[a-f0-9]{32}', identifier):
            raise ValueError('Unknown briefing item.')
        directory = self.studio_directory('briefing')
        with self.lock:
            directory.mkdir(parents=True, exist_ok=True)
            items = [item for item in self.briefing()['items'] if item.get('id') != identifier]
            atomic_json(directory / 'index.json', {'items': items})
            try:
                (directory / (identifier + '.svg')).unlink()
            except FileNotFoundError:
                pass
        return {'removed': True}

    def order_briefing(self, data):
        """Reorder the tray. The list has to be a permutation of what is there."""
        ids = data.get('ids')
        if not isinstance(ids, list) or any(not isinstance(value, str) for value in ids):
            raise ValueError('Briefing order must be a list of item ids.')
        directory = self.studio_directory('briefing')
        with self.lock:
            directory.mkdir(parents=True, exist_ok=True)
            items = self.briefing()['items']
            index = {item.get('id'): item for item in items}
            if len(ids) != len(items) or len(set(ids)) != len(ids) or any(value not in index for value in ids):
                raise ValueError('Briefing order must list every item once.')
            atomic_json(directory / 'index.json', {'items': [index[value] for value in ids]})
        return {'ok': True}

    def publish_pocket(self, data):
        """Publish the phone snapshot: read this console, then build the page.

        Modelled on collect_evidence. One publication at a time per console;
        the job state lives at <data root>/publish/<id>/job.json so a reader
        that arrives after the answer was sent still sees the outcome.
        """
        port = int(getattr(self, 'port', 0) or 0)
        if not port:
            raise ValueError('Restart the console, then publish again.')
        datasets = data.get('datasets')
        if not isinstance(datasets, list) or not datasets:
            raise ValueError('Choose at least one report.')
        datasets = list(dict.fromkeys(str(value).strip() for value in datasets))
        if any(not value or ',' in value for value in datasets):
            raise ValueError('Choose reports from the catalog.')
        known = {str(entry.get('id')) for entry in
                 contribution_query.semantics.catalog(self.report_directory).get('datasets', [])}
        unknown = [value for value in datasets if value not in known]
        if unknown:
            raise ValueError('Not in this report: ' + ', '.join(unknown) + '.')
        identifier = 'publish-' + uuid.uuid4().hex
        with self.lock:
            if any(str(job).startswith('publish-') for job in self.jobs):
                raise Conflict('A phone snapshot is already publishing. Wait for it to finish.')
            self.jobs.add(identifier)
        root = self.root / 'publish' / identifier
        root.mkdir(parents=True, exist_ok=True)
        out_directory = Path(self.report_directory) / 'pocket'
        tools = PROJECT / 'tools' / 'pocket'
        flags = subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0
        log_path = root / 'publish.log'
        job = {'id': identifier, 'status': 'queued', 'message': 'Preparing the snapshot.',
               'datasets': datasets, 'output_path': str(out_directory / 'mentat-pocket.html'),
               'console_pid': os.getpid(), 'output_files': [], 'selfcheck': None,
               'selfcheck_message': '', 'started_at': time.time(), 'finished_at': None}

        def record():
            """Write the job state, retrying a reader's brief hold on the file.

            os.replace fails while another process has the destination open,
            and a poller reading job.json is such a process. A state write lost
            that way would strand the job in 'running'. A copy is written: the
            reply below shares this dict with the worker thread.
            """
            for attempt in range(5):
                try:
                    return atomic_json(root / 'job.json', dict(job))
                except OSError:
                    time.sleep(0.05)

        record()

        def step(script, arguments, log):
            process = subprocess.Popen([sys.executable, str(tools / script)] + arguments, stdout=log,
                                       stderr=subprocess.STDOUT, creationflags=flags)
            job['pid'] = process.pid
            record()
            return process.wait()

        def run_selfcheck(log):
            """Check what was published. Advisory: the verdict is recorded, not enforced."""
            log.write('\n--- selfcheck ---\n')
            log.flush()
            try:
                finished = subprocess.run([sys.executable, str(tools / 'selfcheck.py'),
                                           '--dir', str(out_directory)], capture_output=True,
                                          text=True, creationflags=flags, timeout=600)
            except (OSError, subprocess.SubprocessError) as error:
                log.write('could not run: ' + str(error) + '\n')
                return 'fail', 'The published page could not be checked: ' + str(error)
            output = (finished.stdout or '') + (finished.stderr or '')
            log.write(output)
            summary = next((line.strip() for line in reversed(output.splitlines())
                            if line.startswith('ALL CHECKS PASSED') or 'CHECK(S) FAILED' in line),
                           'The self-check produced no summary. See ' + str(log_path))
            return ('pass' if finished.returncode == 0 else 'fail'), summary

        def publish():
            failure = ''
            try:
                out_directory.mkdir(parents=True, exist_ok=True)
                # The page links these as siblings. A snapshot published without
                # them is a page a phone cannot install, so they ship with it.
                for name in ('icon-180.png', 'manifest.webmanifest'):
                    shutil.copyfile(tools / name, out_directory / name)
                job.update(status='running', message='Reading the published report through this console.')
                record()
                with log_path.open('w', encoding='utf-8') as log:
                    if step('extract_pocket_data.py', ['--port', str(port), '--out-dir', str(out_directory),
                                                       '--datasets', ','.join(datasets)], log):
                        failure = 'The snapshot reader did not finish. See ' + str(log_path)
                    else:
                        job.update(message='Building the page.')
                        record()
                        if step('build_page.py', ['--data', str(out_directory / 'pocket-data.json'),
                                                  '--out', job['output_path']], log):
                            failure = 'The page could not be built from the snapshot. See ' + str(log_path)
                        else:
                            job.update(message='Checking the published page.')
                            record()
                            job['selfcheck'], job['selfcheck_message'] = run_selfcheck(log)
            except OSError as error:
                failure = 'The publisher could not be started: ' + str(error)
            job['output_files'] = [str(out_directory / name) for name in
                                   ('mentat-pocket.html', 'pocket-data.json', 'pocket-integrity.json',
                                    'icon-180.png', 'manifest.webmanifest')
                                   if (out_directory / name).is_file()]
            if not failure and not Path(job['output_path']).is_file():
                failure = 'The page was not written. See ' + str(log_path)
            job.update(status='failed' if failure else 'done',
                       message=failure or ('Snapshot published to ' + job['output_path']),
                       finished_at=time.time())
            record()

        def guarded():
            try:
                publish()
            finally:
                self.jobs.discard(identifier)

        threading.Thread(target=guarded, daemon=True).start()
        return {'job': dict(job)}

    def publication(self, identifier):
        """The stored state of one publication, or None when there is no such job.

        A console that dies mid-publish leaves 'running' on disk forever, so a
        state whose owning console is gone is reported failed. The owner is
        checked, not the child: between the child exiting and this thread
        writing 'done' the child pid is already dead and the job is fine.
        """
        if not re.fullmatch(r'publish-[a-f0-9]{32}', identifier or ''):
            raise ValueError('That is not a phone snapshot publication.')
        path = self.root / 'publish' / identifier / 'job.json'
        if not path.parent.is_dir():
            return None
        # A reader and the worker's os.replace collide on Windows: the read
        # fails with PermissionError while the state file is being swapped,
        # and read_json reports that as an empty object. Measured at 8.9% of
        # reads under contention, so a state that will not read is retried
        # rather than reported missing. A job that never existed has no
        # directory and is answered on the first pass.
        for attempt in range(6):
            job = read_json(path)
            if job.get('id') == identifier:
                break
            time.sleep(0.05)
        else:
            return None
        if job.get('status') in ('queued', 'running') and not process_alive(job.get('console_pid')):
            job['status'] = 'failed'
            job['message'] = 'The publishing console stopped.'
        return job

    def orphaned_state(self, identifier, state, meta, directory):
        """Report a live-looking run whose collection process is gone as stopped.

        finish() records an exit only for a process THIS console holds a handle
        for, so a console that is killed leaves its run showing whatever state
        the orchestrator last wrote - "starting", hours after every process it
        named has exited. That presents a dead collection as live, and because
        resume_target() applies only to a terminal state it also locks the run
        out of Resume, which is precisely the run Resume exists for. The
        correction is computed, never written: the manifest belongs to whichever
        console owns the roster next, and a read path leaves it untouched.
        """
        if state not in ('preparing', 'starting', 'running'):
            return state
        # Work this console owns is never second-guessed. A launch thread
        # between claiming the roster and writing the PID is alive by
        # definition, and its manifest has no PID to check yet.
        if identifier in getattr(self, 'preparing', set()):
            return state
        process = getattr(self, 'processes', {}).get(identifier)
        if process is not None and process.poll() is None:
            return state
        if process_confirmed_dead(meta):
            return 'stopped'
        # No PID names no process to ask about. Preparing may be seconds from
        # writing one, so it stands whatever its age; a run that claimed to be
        # collecting and then wrote nothing for the whole grace period does not.
        if state != 'preparing' and not meta.get('pid') and run_status_is_stale(directory):
            return 'stopped'
        return state

    def runs(self):
        values = []
        for path in self.root.glob("*/upload.json"):
            meta = read_json_cached(path)
            if not meta:
                continue
            run_paths = sorted((path.parent / "runs").glob("*/manifest.json"))
            state = meta["state"]
            if (run_paths and state not in ('stopped', 'failed', 'launch failed')
                    and not stale_resume_status(run_paths[-1].parent, meta)):
                state = run_state(read_json_cached(run_paths[-1].parent / 'status.json'), state)
            state = self.orphaned_state(meta["id"], state, meta,
                                        run_paths[-1].parent if run_paths else None)
            values.append({"id": meta["id"], "name": meta["name"], "state": state,
                           "created_at": meta["created_at"], "directory": run_paths[-1].parent if run_paths else None,
                           "upload": meta})
        existing = PROJECT / ".cache/resilient-industry-runs"
        for path in sorted(existing.glob("*/manifest.json"), reverse=True)[:15]:
            values.append({"id": "existing-" + path.parent.name, "name": path.parent.name,
                           "state": run_state(read_json_cached(path.parent / 'status.json'), 'existing run'), "created_at": path.stat().st_mtime, "directory": path.parent})
        for path in (PROJECT / '.cache/report-repairs').glob('*/manifest.json'):
            meta = read_json(path)
            values.append({'id': 'repair-' + path.parent.name,
                           'name': meta.get('name', 'Rebuilt reports'), 'state': 'completed',
                           'created_at': path.stat().st_mtime, 'directory': path.parent})
        return sorted(values, key=lambda item: item["created_at"], reverse=True)

    def can_archive_run(self, run):
        if run.get('state') not in ('ready', 'needs review', 'completed', 'failed', 'stopped', 'launch failed'):
            return False
        identifier = run['id']
        if identifier in getattr(self, 'preparing', set()):
            return False
        process = getattr(self, 'processes', {}).get(identifier)
        if process is not None and process.poll() is None:
            return False
        return not process_alive(run.get('upload', {}).get('pid'))

    def archive_run(self, data):
        identifier, archived = data.get('id'), data.get('archived')
        if not isinstance(identifier, str) or not isinstance(archived, bool):
            raise ValueError('Choose a saved run and whether to archive or restore it.')
        with self.lock:
            run = self.find(identifier)
            if run is None:
                raise RunNotFound('Run not found.')
            if archived and not self.can_archive_run(run):
                raise ValueError('Active, queued, or unfinished collections cannot be archived. Wait for the run to finish or stop.')
            path = self.root / 'run-archives.json'
            metadata = read_json(path)
            entries = metadata.get('runs', {}) if isinstance(metadata, dict) else {}
            entries = entries if isinstance(entries, dict) else {}
            if archived:
                entries[identifier] = {'archived_at':datetime.datetime.now(datetime.timezone.utc).isoformat()}
            else:
                entries.pop(identifier, None)
            atomic_json(path, {'version':1, 'runs':entries})
        return {'id':identifier, 'archived':archived}

    @property
    def evidence_store(self):
        """CSV imports kept under the data root, not in this process.

        The old dict lived in process memory, so a stale-build restart, an idle
        stop or the launcher landing on a different port silently dropped every
        import the reader had chosen. The data root is shared by every console
        the launcher starts, so a file imported once stays listed.
        """
        with self.lock:
            store = getattr(self, 'csv_store', None)
            if store is None:
                store = self.csv_store = evidence_csv.Store(
                    self.root, evict=not getattr(self, 'read_only', False))
        return store

    def import_evidence_csv(self, data):
        """Legacy JSON import: read the text through the store, then answer."""
        store = self.evidence_store
        job = store.wait(store.ingest_text(data.get('name'), data.get('text'))['id'])
        if job['state'] != 'ready':
            raise ValueError(job['message'] or evidence_csv.M_GONE(job['id']))
        return job['catalog']

    def import_evidence_path(self, data):
        """Read a file the reader chose by path on this computer.

        Asynchronous like the streaming upload, and answered with the same JOB
        the page already polls, so a large workbook on a slow disk is never
        read inside the request thread.
        """
        return self.evidence_store.import_path(data.get('path'), sheet=data.get('sheet', ''),
                                               header_path=data.get('header_path', ''))

    def evidence_files(self, data):
        """One folder of this computer, for the console's own file picker."""
        return evidence_csv.browse(data.get('path', ''))

    def evidence_csv_job(self, identifier):
        return self.evidence_store.job(identifier)

    def evidence_csv_view(self, identifier='', committee='', cycle='', offset=0, limit=200):
        store = self.evidence_store
        if not identifier:
            return {'sources':store.sources(), 'limits':evidence_csv.limits(), 'session_only':False}
        return store.select(identifier, committee, cycle, *evidence_csv.page_args(offset, limit))

    def evidence_csv_summary(self, identifier='', committee='', cycle=''):
        """One committee's figures for a pane header, without its rows."""
        return self.evidence_store.summary(identifier, committee, cycle)

    def evidence_geography(self, source='', identifier='', committee='', cycle='', direction='spent'):
        """One pane's map: where this committee's money came from, or went.

        The import store is opened only for an import, so a snapshot map never
        touches the CSV database, and the snapshot root is this console's own
        collection cache - the same directory /api/evidence reads.
        """
        return evidence_geo.cached_geography(source, identifier, committee, cycle, direction,
                                      store=self.evidence_store if source == 'csv' else None,
                                      snapshot_root=PROJECT / '.cache/pac-evidence',
                                      data_root=self.root)

    def evidence_stats(self, identifier='', committee='', cycle=''):
        """The one pane figure the CSV summary does not carry: distinct contributors.

        /api/evidence-csv?summary=1 ranks the top ten counterparties and counts
        rows, so the pane could print "6 contributors" only by counting the ten
        it was shown. This counts them where the rows are, over exactly the
        supported rows the map places, so the pane and the map agree.
        """
        return evidence_geo.counterparty_count(self.evidence_store, identifier, committee, cycle)

    def evidence_geocode(self, identifier=''):
        """What this import's geocoding has done, without scanning the import.

        Polled twice a second while a job runs, so it reads the job file and
        nothing else: the row counts a reader cares about (matched, by ZIP,
        still to do) come from /api/evidence-geo, which has to read the rows
        anyway.
        """
        if not evidence_csv.IDENTIFIER.fullmatch(str(identifier or '')):
            raise ValueError(evidence_csv.M_GONE(identifier))
        return geocode.status(self.root, identifier)

    def start_evidence_geocode(self, data):
        """Ruling 10: one job per import, and ONLY from an explicit Start.

        The read-only preview is refused here as well as at the handler, because
        it shares its data root with the console that owns these imports and a
        job started from it would write that console's cache.
        """
        if getattr(self, 'read_only', False):
            raise ValueError(geocode.M_READ_ONLY)
        identifier = str(data.get('id') or '')
        if not evidence_csv.IDENTIFIER.fullmatch(identifier):
            raise ValueError(evidence_csv.M_GONE(identifier))
        if str(data.get('action') or 'start').strip().lower() == 'stop':
            return geocode.stop(self.root, identifier)
        addresses = geocode.addresses_for(self.evidence_store, identifier,
                                          str(data.get('committee') or '').strip().upper(),
                                          str(data.get('cycle') or '').strip())
        return geocode.start(self.root, identifier, addresses)

    def summary(self, include_archived=False):
        runs = self.runs()
        try:
            credentials = {**key_status(), 'key_status_available':True}
        except (RuntimeError, OSError, ValueError):
            # Report analysis does not require opening the encrypted key store.
            # Collection still uses load_keys and its normal authorization checks.
            credentials = {'active_key_count':None, 'stored_key_count':None, 'encrypted':True,
                           'key_status_available':False,
                           'key_status_message':'Saved FEC keys are unavailable to this Windows session. Report analysis remains available.'}
        metadata = read_json(self.root / 'run-archives.json')
        archives = metadata.get('runs', {}) if isinstance(metadata, dict) else {}
        archives = archives if isinstance(archives, dict) else {}
        visible = [{**{key:value for key,value in run.items() if key not in ('directory', 'upload')},
                    'archived':run['id'] in archives, 'can_archive':self.can_archive_run(run),
                    'resumable':self.resume_target(run) is not None,
                    'resumed_from':run.get('upload', {}).get('resumed_from')}
                   for run in runs if include_archived or run['id'] not in archives]
        return {"runs":visible, 'archived_count':sum(run['id'] in archives for run in runs),
                'total_run_count':len(runs), 'collection_directory':str(self.root),
                'report_directory':str(self.report_directory),
                "coverage": self.coverage_catalog(runs), "key_configured": bool(credentials.get('active_key_count', 0)),
                "read_only":getattr(self, 'read_only', False), "api_version":2,
                "evidence_collection_available":True, **credentials}

    def collection_coverage(self, run, status=None, excluded_scopes=()):
        directory = run.get('directory')
        if not directory:
            return []
        status = status if status is not None else read_json(directory / 'status.json')
        manifest = read_json(directory / 'manifest.json')
        cycles = cycles_from(status.get('cycles') or manifest.get('cycles') or run.get('upload', {}).get('cycles'))
        rows = []
        for industry in status.get('industries', []):
            receipt_path = industry.get('coverage_receipt')
            receipt = read_json(receipt_path) if receipt_path else {}
            if not isinstance(receipt, dict):
                receipt = {}
            selected_cycles = cycles_from(industry.get('cycles')) or cycles or cycles_from(receipt.get('cycles'))
            # A legacy run lacking even its requested cycles must visibly keep
            # every supported cycle unknown instead of disappearing from Loading.
            selected_cycles = selected_cycles or [2022, 2024, 2026]
            name = str(industry.get('industry') or 'Unknown sector')
            if all((name, cycle) in excluded_scopes for cycle in selected_cycles):
                continue
            database = industry.get('database') or receipt.get('database')
            values = inspect_coverage(database, receipt_path, name, selected_cycles)
            for value in values:
                value.update(run_id=run['id'], run_name=run['name'], receipt=receipt_path,
                             database=database)
            rows.extend(values)
        return rows

    def coverage_catalog(self, runs):
        # Select the latest saved collection for each sector/cycle, including
        # failures. Never make an older verified run appear to verify a new one.
        selected = {}
        for run in runs:
            for row in self.collection_coverage(run, excluded_scopes=selected):
                selected.setdefault((row['industry'], row['cycle']), row)
        return list(selected.values())

    def find(self, identifier):
        return next((run for run in self.runs() if run["id"] == identifier), None)

    def report_path(self, identifier, index):
        if identifier == 'published-reports':
            item = next((row for row in self.published_reports() if row['index'] == index), None)
            if item is None:
                raise ValueError('Report not available in the authoritative folder. Refresh the report list.')
            path = require_report_path(self.report_directory / item['name'], self.report_directory)
            report_publication.capture(path)
            return path
        run = self.find(identifier)
        if not run or not run['directory']:
            raise ValueError('Report not available.')
        industries = read_json(run['directory'] / 'status.json').get('industries', [])
        if index < 0 or index >= len(industries) or industries[index].get('status') != 'succeeded':
            raise ValueError('Report is not complete.')
        path = Path(industries[index].get('output_workbook') or '__absent__')
        path = require_report_path(path, self.report_directory)
        if str(path).startswith(('\\\\', '//')) or not path.is_file():
            raise ValueError('Local report not available.')
        report_publication.capture(path)
        return path

    def published_reports(self):
        """Discover only final workbooks in the configured folder, never fall back."""
        if not self.report_directory.is_dir():
            raise ValueError('The authoritative report folder is unavailable; no alternate source will be used.')
        rows = []
        identifiers = set()
        for path in sorted(self.report_directory.glob('*.xlsx'), key=lambda p:p.name.casefold()):
            if path.name.startswith('~$'):
                continue
            try:
                resolved = require_report_path(path, self.report_directory)
            except ValueError:
                continue
            if not resolved.is_file():
                continue
            try:
                publication = report_publication.capture(resolved)
            except ValueError:
                continue
            index = int(hashlib.sha256(path.name.encode('utf-8')).hexdigest()[:12],16)
            if index in identifiers:
                raise ValueError('Report filename identifiers conflict; no ambiguous report will be opened.')
            identifiers.add(index)
            # The companion `<Industry> Industry Spending Data.xlsx` is the tall,
            # Tableau-shaped view of the SAME export. Stripping either suffix
            # files both books under one industry instead of inventing an
            # industry literally called "Credit Union Industry Spending Data";
            # `kind` is what lets a caller tell the two apart without parsing
            # the filename again.
            from report_identity import identity
            report_identity = identity(path.name)
            if report_publication.capture(resolved) != publication:
                continue
            stem = report_identity['label'] if report_identity else path.stem
            kind = report_identity['kind'] if report_identity else 'overview'
            rows.append({'index':index,'name':path.name,'industry':stem,'kind':kind,
                         'source_directory':str(self.report_directory), 'publication': publication})
        return rows

    def report_sources(self):
        return {'runs':[{'id':'published-reports','name':self.report_directory.name,
                         'state':'completed','source_directory':str(self.report_directory)}],
                'report_directory':str(self.report_directory)}

    def report_catalog(self, identifier):
        """Chart source names only; do not scan reports or live worker databases."""
        if identifier == 'published-reports':
            return {'reports':self.published_reports(),'report_directory':str(self.report_directory)}
        run = self.find(identifier)
        if not run or not run['directory']:
            return {'reports': []}
        reports = []
        for index, item in enumerate(read_json(run['directory'] / 'status.json').get('industries', [])):
            if item.get('status') != 'succeeded':
                continue
            try:
                path = self.report_path(identifier,index)
            except (OSError,ValueError):
                continue
            reports.append({'index':index,'name':path.name,'source_directory':str(self.report_directory)})
        return {'reports':reports,'report_directory':str(self.report_directory),
                'message':'Only reports in the authoritative folder are available. Select '+self.report_directory.name+' in the report viewer.'}

    def graph_catalog(self, identifier, reports=None):
        if identifier != 'published-reports':
            return {'industries':[],'message':'Select the authoritative report folder: '+self.report_directory.name+'.'}
        # `reports` lets one request reuse a folder scan it already paid for.
        industries = self.published_reports() if reports is None else reports
        groups, unavailable = {}, []
        for industry in industries:
            index = industry['index']
            try:
                path = require_report_path(self.report_directory / industry['name'], self.report_directory)
                if not graph_schema(stamp(path)):
                    continue
            except (OSError, ValueError, KeyError, zipfile.BadZipFile):
                unavailable.append(str(industry.get('industry') or 'Unnamed report'))
                continue
            from report_identity import identity
            report_identity = identity(path.name)
            name = (report_identity['label'] if report_identity and report_identity['cycle_label'] else
                    re.sub(r'\s+\d{4}[-–]\d{4}$', '', str(industry.get('industry') or path.stem)))
            groups.setdefault(name, []).append(index)
        return {'industries':[{'name':name,'indices':indices} for name,indices in sorted(groups.items())],
                'unavailable_reports':unavailable,
                'report_directory':str(self.report_directory),
                'message':'No supported reports are available in the authoritative folder.' if not groups else ''}

    def graph(self, identifier, industry, filters, registry=None):
        # The folder is read once here and shared with the catalogue and the
        # stamping below, and the caller's registry is reused when it has one.
        reports = self.published_reports() if identifier == 'published-reports' else None
        selected = next((g for g in self.graph_catalog(identifier, reports)['industries'] if g['name'] == industry), None)
        if not selected:
            raise ValueError('Choose an industry with completed reports.')
        paths = [self.report_path(identifier,i) for i in selected['indices']]
        def build():
            result = graph_view(paths,filters)
            # A report is stamped through its own published industry, not through
            # the grouped catalogue name: one group can hold several workbooks.
            published = {row['name']: row['industry'] for row in reports or []}
            snapshots = (registry or snapshot_registry.registry(self.report_directory))['snapshots']
            stamps = {item['industry']: item['build_stamp'] for item in snapshots}
            by_workbook = {item['workbook']['name']: item for item in snapshots if item.get('workbook')}
            for report, path in zip(result['reports'], paths):
                scanned = by_workbook.get(path.name)
                if scanned is not None and scanned['workbook']['sha256'] != report['sha256']:
                    raise ValueError(report_publication.MESSAGE)
                report['freshness'] = report_freshness(path)
                report['source_directory'] = str(self.report_directory)
                report['build_stamp'] = stamps.get(published.get(path.name))
            result['report_directory'] = str(self.report_directory)
            result['industry'] = industry
            return result
        return report_publication.read(paths, build)

    def comparison(self, identifier, industry, relationships=False, reports=None):
        if reports is None and identifier == 'published-reports':
            reports = self.published_reports()
        selected = next((g for g in self.graph_catalog(identifier, reports)['industries'] if g['name'] == industry), None)
        if not selected:
            raise ValueError('Choose an industry with completed reports.')
        build = comparison_flows if relationships else comparison_subjects
        return build([self.report_path(identifier, i) for i in selected['indices']], industry)

    def comparison_entry(self, identifier, industry, reports=None):
        """/api/comparison-flows as a response_cache entry (comparison_data.flows_entry): the same
        report selection as comparison(), answered from memory or disk while nothing changed."""
        if reports is None and identifier == 'published-reports':
            reports = self.published_reports()
        selected = next((g for g in self.graph_catalog(identifier, reports)['industries'] if g['name'] == industry), None)
        if not selected:
            raise ValueError('Choose an industry with completed reports.')
        from comparison_data import flows_entry
        return flows_entry([self.report_path(identifier, i) for i in selected['indices']], industry)

    def snapshot(self, identifier, lane_id=""):
        run = self.find(identifier)
        if not run:
            raise ValueError("Run not found.")
        directory = run["directory"]
        result = {"id": identifier, "name": run["name"], "state": run["state"],
                  "readiness": run.get("upload", {}).get("readiness"), "lanes": [], "events": [], "reports": [], "assurance": [], "fec_coverage": []}
        manifest = run.get('upload', {})
        result['launch_error'] = manifest.get('launch_error')
        result['resumable'] = self.resume_target(run) is not None
        result['resumed_from'] = manifest.get('resumed_from')
        if isinstance(manifest.get('started_at'), (int, float)):
            result['elapsed_seconds'] = max(0, (manifest.get('finished_at') or time.time()) - manifest['started_at'])
        if not directory:
            return result
        status = read_json(directory / "status.json")
        result["status"] = status
        result['fec_coverage'] = self.collection_coverage(run, status)
        # Every terminal state keeps the collector's own elapsed time when the
        # manifest never recorded a finish: a stopped run is not still running,
        # so wall time since its start would grow for as long as it is kept.
        if run['state'] in ('completed', 'failed', 'stopped', 'launch failed') and not manifest.get('finished_at') and isinstance(status.get('elapsed_seconds'), (int, float)):
            result['elapsed_seconds'] = status['elapsed_seconds']
        result["control"] = read_json(directory / "control.json")
        result["age_seconds"] = round(time.time() - (directory / "status.json").stat().st_mtime) if status else None
        if 'elapsed_seconds' not in result and isinstance(status.get('elapsed_seconds'), (int, float)):
            terminal = run['state'] in ('completed', 'failed', 'stopped') or status.get('stage') in ('completed', 'failed', 'stopped')
            result['elapsed_seconds'] = status['elapsed_seconds'] + (0 if terminal else max(0, result['age_seconds'] or 0))
        lanes, fanout_dirs = [], {}
        # (industry, status path) for every task receipt a worker has written,
        # current or historical; company_tasks below is built from them.
        task_files, fanout_cycles = [], []
        resolved_directory = directory.resolve()
        for fanout_manifest in directory.glob("industries/*/fanout/manifest.json"):
            fanout = read_json(fanout_manifest)
            fanout_dirs[fanout.get("industry")] = fanout_manifest.parent
            fanout_cycles += fanout.get('cycles') or []
            # The run's status.json holds a copy of the fanout's own status that
            # the launcher refreshes on its own schedule; a run stopped between
            # the two writes keeps a stale copy (27 of 30 done, one running).
            # The fanout's file wins when it is the newer record of the same run.
            own, copied = read_json(fanout_manifest.parent / 'status.json'), status.get('fanout')
            if (isinstance(own, dict) and isinstance(copied, dict) and own.get('industry') == copied.get('industry')
                    and own.get('run_id') == copied.get('run_id') and newer_timestamp(own.get('updated_at'), copied.get('updated_at'))):
                status['fanout'] = own
            sharing = fanout.get('scheduling', {}).get('strategy') == 'company_work_sharing'
            fanout_directory = fanout_manifest.parent.resolve()
            lane_databases = {Path(lane['state']).resolve():Path(lane['database'])
                              for lane in fanout.get('lanes', []) if lane.get('state') and lane.get('database')}
            # Shared workers publish their current receipt in the manifest. This
            # avoids rereading every historical company attempt on every poll.
            paths = [Path(lane['state']) for lane in fanout.get('lanes', []) if lane.get('state')] if sharing else sorted(fanout_manifest.parent.glob("**/status*.json"))
            for path in paths:
                # Windows can discover this run through a LocalAppData alias
                # while the worker saved its canonical package-cache path.
                path = path.resolve()
                if path.parent == fanout_directory:
                    continue
                try:
                    key = str((path.parent if sharing else path).relative_to(resolved_directory))
                except ValueError:
                    raise ValueError('Saved worker status is outside its collection run.') from None
                # A shared worker's lane folder keeps one status file per
                # invocation (work, recovery, split); the manifest names only
                # the current one, so the earlier attempts are read from disk.
                task_files += [(fanout.get("industry"), item) for item in
                               (sorted(path.parent.glob("status*.json")) if sharing else [path])]
                state = read_json(path)
                if not state:
                    continue
                db = lane_databases.get(path, path.parent / 'mentat.db')
                lanes.append({"id": key, "industry": fanout.get("industry"), "label": str(path.parent.name),
                              "state": state, "database": str(db)})
        result["lanes"] = [{key: value for key, value in lane.items() if key != "database"} for lane in lanes]
        result['company_tasks'] = recorded_tasks(task_files, resolved_directory)
        result['cycles'] = cycles_from(status.get('cycles') or manifest.get('cycles') or fanout_cycles)
        selected = next((lane for lane in lanes if lane["id"] == lane_id), lanes[0] if lanes else None)
        if selected:
            result["selected_lane"] = selected["id"]
            result["preview"] = lane_preview(selected["database"])
        events = directory / "events.jsonl"
        if events.is_file():
            with events.open("rb") as stream:
                stream.seek(max(0, events.stat().st_size - 24000))
                for line in stream.read().decode("utf-8-sig", errors="replace").splitlines()[-35:]:
                    try:
                        result["events"].append(json.loads(line))
                    except ValueError:
                        pass
        finished_collections = []
        source_root = status.get('source_root')
        for index, industry in enumerate(status.get("industries", [])):
            workers = read_json(Path(fanout_dirs.get(industry.get("industry")) or "__absent__") / "workers.json")
            shared = read_json(Path(fanout_dirs.get(industry.get("industry")) or "__absent__") / "request-workers.json")
            if shared and time.time() - shared.get('updated_at_epoch', 0) < 15 and not shared.get('stopped'):
                workers = {**workers, 'database_writers':workers.get('active_workers'),
                           **shared, 'request_pool':shared}
            if isinstance(workers, dict) and workers:
                industry["workers"] = workers
            receipt = read_json(industry.get("fanout_receipt") or "__absent__")
            if receipt.get("task_count") is not None:
                finished_collections.append(receipt)
            coverage_path = Path(industry.get("coverage_receipt") or "__absent__")
            coverage = read_json(coverage_path)
            if not isinstance(coverage, dict):
                coverage = {}
            check = {"industry": industry.get("industry"), "collection": "unverified"}
            if coverage:
                checks = coverage.get("checks", {})
                check["collection"] = "passed" if coverage.get("status") == "succeeded" and isinstance(checks, dict) and checks and all(value is True for value in checks.values()) else "failed"
                check["recorded_checks"] = len(checks) if isinstance(checks, dict) else 0
            result["assurance"].append(check)
            output = None
            if industry.get('status') == 'succeeded':
                try:
                    output = self.report_path(identifier,index)
                except (OSError,ValueError):
                    result['report_source_notice'] = 'Historical report paths are unavailable here. Current workbooks are in '+str(self.report_directory)+'.'
            if output is not None:
                quality = self.inspect_report(output)
                expected = industry.get("workbook_sha256")
                quality = {**quality, "matches_run": quality["sha256"].lower() == expected.lower()
                           if expected and quality.get("sha256") else None}
                result["reports"].append({"name": output.name, "index": index, "quality": quality,
                                          "freshness": report_freshness(output),
                                          "source_directory":str(self.report_directory)})
        if (status.get("stage") in ("completed", "failed") and finished_collections
                and len(finished_collections) == len(status.get("industries", []))):
            result["completed_collection"] = {
                "total": sum(item["task_count"] for item in finished_collections),
                "completed": sum(item.get("succeeded", 0) + item.get("failed", 0) for item in finished_collections),
                "failed": sum(item.get("failed", 0) for item in finished_collections), "running": 0}
        inputs = self.root / identifier / "runtime/mentat/inputs" if run.get("upload") else PROJECT / "inputs"
        if status.get('source_root'):
            inputs = Path(status['source_root'])
        result["source_snapshots"] = [
            {"name": label, "modified_at": datetime.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
             "age_days": max(0, int((time.time() - path.stat().st_mtime) / 86400))}
            for path, label in ((inputs / "fec_beneficiary/pas226/itpas2.txt", "2025–26 downstream activity"),
                                (inputs / "fec_committee_master/cm26/cm.txt", "2025–26 committee register")) if path.is_file()
        ]
        reviews = read_json(self.root / 'quality-reviews.json')
        for report in result['reports']:
            quality = report['quality']
            annotated = []
            for warning in quality.get('warnings', []):
                fingerprint = hashlib.sha256(json.dumps(
                    [identifier, report['name'], quality.get('sha256'), warning],
                    sort_keys=True, ensure_ascii=False).encode()).hexdigest()
                annotated.append({**warning, 'id': fingerprint,
                                  'reviewed_at': reviews.get(fingerprint, {}).get('reviewed_at')})
            report['quality'] = {**quality, 'warnings': annotated}
        return result


# Access recording for the value-ranked prewarm (RACF, prewarm.py "value ranking"). main() sets this to a
# prewarm.AccessRecorder only on the console on the user's own data root that may write; everywhere else
# it stays None and record_access() does nothing.
ACCESS_RECORDER = None
# graph.js marks its idle prefetch of the likely next Graph views with this request header (value '1'). Such
# a read is answered exactly like a click on the same URL (same body, same ETag, same 304), but it is the
# page's own guess, not a user access, so /api/graph does not call record_access for it (review H4,
# 2026-09-24): recording it would rank the warmer's priming by what the page guessed. A custom X- header,
# because browsers forbid scripts to set Sec-Purpose; nothing varies on it, so the browser's cache entry
# for the URL is shared by the prefetch and the click.
PREFETCH_HEADER = 'X-Mentat-Prefetch'


def prefetch_request(handler):
    """Whether this request is graph.js's idle prefetch (PREFETCH_HEADER: 1), not a user access."""
    return (handler.headers.get(PREFETCH_HEADER) or '').strip() == '1'


def record_access(handler, kind, body, started):
    """After a response is written: one record of (kind, parameters or body, handler seconds, bytes sent).
    The recorder only enqueues, so this never delays or fails the response."""
    recorder = ACCESS_RECORDER
    if recorder is None:
        return
    try:
        recorder.submit(kind, body, cost_s=round(time.perf_counter() - started, 4),
                        size=getattr(handler, '_sent_bytes', None))
    except Exception:
        pass


class Handler(access_gate.AccessMixin, BaseHTTPRequestHandler):
    # The page loads about 35 small files. Under HTTP/1.0 each was a new
    # connection, and with Nagle on, the header write and the body write
    # became two segments that tripped Windows' 200 ms delayed-ACK stall,
    # so the static waterfall alone took 5 s. Keep-alive plus TCP_NODELAY
    # removes both; every response already carries a Content-Length.
    protocol_version = 'HTTP/1.1'
    disable_nagle_algorithm = True
    access_app = 'mentat-console'
    access_cfg = access_gate.AccessConfig()   # not enforced; main() installs the environment's
    access_frame_options = 'DENY'             # CONSOLE_CSP already carries frame-ancestors 'none' (D-061)

    def log_message(self, *_):
        pass

    def log_request(self, code='-', size='-'):
        """Hosted: one audit line per answered request (who, method, route without query, status).
        Everywhere else, silent as ever."""
        if self.access_cfg.enforced:
            sys.stderr.write(access_gate.audit_line(self.access_app, self.access_user, self.command, self.path,
                                                    code if isinstance(code, int) else 0) + '\n')

    def end_headers(self):
        """Hosted: add the gate's security headers that this response has not already set."""
        if self.access_cfg.enforced:
            sent = {line.split(b':', 1)[0].strip().lower() for line in getattr(self, '_headers_buffer', ())}
            for name, value in access_gate.security_headers(self.access_cfg, self.access_frame_options):
                if name.lower().encode('latin-1') not in sent:
                    self.send_header(name, value)
        super().end_headers()

    def send(self, value, status=200, content_type="application/json", *, cache_control="no-store", etag=None):
        if self.access_cfg.enforced and isinstance(value, (dict, list)):
            value = without_os_error_text(without_server_paths(value))   # hosted: no server path or OSError text in any JSON body
        content = json.dumps(value, ensure_ascii=False, separators=(',', ':')).encode() if content_type == "application/json" else value
        compressible = content_type.startswith(('application/json', 'text/', 'image/svg+xml'))
        compressed = compressible and status != 304 and len(content) >= 2048 and accepts_gzip(self.headers.get('Accept-Encoding'))
        if compressed:
            content = gzip.compress(content, compresslevel=1, mtime=0)
        self._sent_bytes = len(content)
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        if status != 304:
            self.send_header("Content-Length", str(len(content)))
        self.send_header("Cache-Control", cache_control)
        if compressible:
            self.send_header('Vary', 'Accept-Encoding')
        if compressed:
            self.send_header('Content-Encoding', 'gzip')
        if etag:
            self.send_header('ETag', etag)
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Content-Security-Policy", CONSOLE_CSP)
        self.end_headers()
        self.wfile.write(content)

    def send_entry(self, entry, content_type="application/json"):
        """A response_cache entry: 304 when If-None-Match names its ETag, else its stored gzip
        (Accept-Encoding permitting, same rule and 2048-byte floor as send()) or its body. The body
        bytes are exactly send()'s; the headers differ only in Cache-Control (private, no-cache: the
        browser keeps the answer and revalidates every use) and the strong ETag."""
        import response_cache
        if self.access_cfg.enforced and content_type.startswith('application/json'):
            # Hosted: a cached JSON body gets send()'s path scrub and the OSError-text scrub too, then is keyed by
            # its own bytes (the ETag follows the scrubbed body). ponytail: a full parse and encode on every hosted
            # request; scrub once when the entry is stored if a hosted entry ever gets hot.
            entry = response_cache.from_value(without_os_error_text(without_server_paths(json.loads(entry.body), True)))
        status, headers, body = response_cache.respond(entry, self.headers.get('If-None-Match'),
                                                       gzip_ok=accepts_gzip(self.headers.get('Accept-Encoding')),
                                                       content_type=content_type)
        self._sent_bytes = len(body or b'')
        self.send_response(status)
        for name, value in headers:
            self.send_header(name, value)
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Content-Security-Policy", CONSOLE_CSP)
        self.end_headers()
        if body:
            self.wfile.write(body)

    # Lane P: the two trees behind GEO_V2_ROUTE and the BASEMAP_* routes. A file
    # is validated by its size and modification time ("<size>-<mtime>"), so a
    # 2 MB layer or a 112 MB archive is never hashed to answer a request.
    @staticmethod
    def file_tag(info):
        return '"%x-%x"' % (info.st_size, info.st_mtime_ns)

    def send_file(self, path, mime):
        """One whitelisted file: 304 when If-None-Match names it, else the bytes."""
        try:
            info = path.stat()
        except OSError:
            return self.send({"error": "Not found."}, 404)
        if not path.is_file():
            return self.send({"error": "Not found."}, 404)
        tag = self.file_tag(info)
        if tag in [value.strip() for value in self.headers.get('If-None-Match', '').split(',')]:
            return self.send(b'', 304, content_type=mime, cache_control='private, no-cache', etag=tag)
        return self.send(path.read_bytes(), content_type=mime, cache_control='private, no-cache', etag=tag)

    def send_ranged(self, path, mime='application/octet-stream'):
        """HTTP byte serving (RFC 9110 section 14) for the street tile archive.

        pmtiles asks for a few KB at a time with `Range: bytes=a-b` and refuses
        a 200 whose Content-Length exceeds what it asked for, which is all a
        plain http.server can send. A single satisfiable range is 206 with
        Content-Range; a range past the end is 416 with `bytes */size`; no
        Range (or one this server does not serve, or an If-Range naming another
        version) is the whole file with 200. The body is streamed in 64 KB
        pieces, never read whole, and never compressed.
        """
        try:
            info = path.stat()
        except OSError:
            return self.send({"error": "Not found."}, 404)
        if not path.is_file():
            return self.send({"error": "Not found."}, 404)
        size, tag = info.st_size, self.file_tag(info)
        wanted = byte_range(self.headers.get('Range'), size)
        if_range = self.headers.get('If-Range')
        if wanted is not None and if_range and if_range.strip() != tag:
            wanted = None
        if wanted == 'unsatisfiable':
            self._sent_bytes = 0
            self.send_response(416)
            self.send_header('Content-Range', 'bytes */%d' % size)
            self.send_header('Accept-Ranges', 'bytes')
            self.send_header('Content-Length', '0')
            self.send_header('Cache-Control', 'no-store')
            self.send_header('X-Content-Type-Options', 'nosniff')
            self.send_header('Content-Security-Policy', CONSOLE_CSP)
            self.end_headers()
            return None
        first, last = wanted if wanted is not None else (0, size - 1)
        length = max(0, last - first + 1)
        self.send_response(206 if wanted is not None else 200)
        self.send_header('Content-Type', mime)
        self.send_header('Content-Length', str(length))
        self.send_header('Accept-Ranges', 'bytes')
        if wanted is not None:
            self.send_header('Content-Range', 'bytes %d-%d/%d' % (first, last, size))
        self.send_header('ETag', tag)
        self.send_header('Cache-Control', 'private, no-cache')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.send_header('Content-Security-Policy', CONSOLE_CSP)
        self.end_headers()
        self._sent_bytes = length
        with open(path, 'rb') as stream:
            stream.seek(first)
            left = length
            while left > 0:
                chunk = stream.read(min(65536, left))
                if not chunk:
                    break
                self.wfile.write(chunk)
                left -= len(chunk)
        return None

    def send_map_file(self, raw_path):
        """/geo/v2/... and /basemap/...: a whitelisted name, or 404."""
        from urllib.parse import unquote
        if raw_path.startswith('/geo/v2/'):
            match = GEO_V2_ROUTE.fullmatch(raw_path)
            if not match or '..' in match.group(1).split('/'):
                return self.send({"error": "Not found."}, 404)
            return self.send_file(GEO_V2_DIR / match.group(1), 'application/json; charset=utf-8')
        if raw_path in BASEMAP_FILES:
            relative, mime = BASEMAP_FILES[raw_path]
            return self.send_file(BASEMAP_DIR / relative, mime)
        font = BASEMAP_FONT_ROUTE.fullmatch(unquote(raw_path))
        if font:
            return self.send_file(BASEMAP_DIR / 'fonts' / font.group(1) / (font.group(2) + '.pbf'),
                                  'application/x-protobuf')
        tiles = BASEMAP_TILES_ROUTE.fullmatch(raw_path)
        if tiles:
            return self.send_ranged(BASEMAP_DIR / 'tiles' / tiles.group(1))
        return self.send({"error": "Not found."}, 404)

    def discard_body(self):
        """Read and drop an unread request body before an early refusal.

        Closing a socket that still has inbound data queued makes Windows send
        RST, and the client can lose the 4xx it is already reading.

        Only a body this server would otherwise have accepted is drained, and
        only briefly. A client that declares a length it never sends must not
        hold this thread: ThreadingHTTPServer gives it a thread per connection.
        A missing, malformed or oversized Content-Length is left alone, because
        there is then no body this server would have read either.
        """
        try:
            length = int(self.headers.get('Content-Length', '0'))
        except (TypeError, ValueError):
            return
        if length < 1 or length > 7_000_000:
            return
        previous = self.connection.gettimeout()
        self.connection.settimeout(2)
        try:
            while length > 0:
                chunk = self.rfile.read(min(length, 65_536))
                if not chunk:
                    return
                length -= len(chunk)
        except OSError:
            return
        finally:
            try:
                self.connection.settimeout(previous)
            except OSError:
                pass

    def valid_host(self):
        host = self.headers.get("Host")
        return (host in (f"127.0.0.1:{self.server.server_port}", f"localhost:{self.server.server_port}")
                or (self.access_cfg.enforced and access_gate.host_ok(host, self.access_cfg)))

    def valid_origin(self, origin):
        port = self.server.server_port
        return (origin in (f"http://localhost:{port}", f"http://127.0.0.1:{port}")
                or (self.access_cfg.enforced
                    and origin.lower() in {'https://' + host for host in self.access_cfg.allowed_hosts - access_gate.LOOPBACK}))

    def session_token(self):
        """What the page is told to send back: this console's token, or HOSTED_SESSION_TOKEN once hosted."""
        return HOSTED_SESSION_TOKEN if self.access_cfg.enforced else self.server.app.token

    def scope_body(self, data, sources=None):
        """Resolve this body's industry, fill the legacy dataset, drop the key.

        The readers select by `dataset`, and neither of their validators knows
        `industry`, so it is removed before the reader sees it rather than left
        for a future stricter validator to refuse.
        """
        entry = industry_scope.from_body(self.server.app.report_directory, data, sources)
        if entry and 'dataset' not in data:
            data['dataset'] = entry['dataset']
        data.pop('industry', None)
        return entry

    def source_metadata(self, reader, industry=None):
        """Keep available scanned sources and their metadata in one generation.

        Pending/unavailable entries remain listed unavailable. They cannot
        authorize facts, and are deliberately not upgraded by this boundary.
        """
        sources = contribution_data._sources(self.server.app.report_directory)
        available = [item for item in sources if item.get('available') and item.get('workbook')]
        if industry is not None:
            registry = snapshot_registry.registry(self.server.app.report_directory, sources)
            selected = snapshot_registry.find(registry, industry)
            available = [item for item in available if selected and item['id'] == selected['dataset']]
        def build():
            for item in available:
                scanned = item.get('publication')
                if scanned is not None and report_publication.capture(item['workbook']) != scanned:
                    raise ValueError(report_publication.MESSAGE)
            return reader(sources)
        return report_publication.read([item['workbook'] for item in available], build)

    def contribution_result(self, reader, data, sources=None):
        directory = self.server.app.report_directory
        sources = contribution_data._sources(directory) if sources is None else sources
        paths = [item['workbook'] for item in sources if item['id'] == data.get('dataset')]
        return report_publication.read(paths, lambda: self._contribution_result(reader, data, sources))

    def _contribution_result(self, reader, data, sources=None):
        directory = self.server.app.report_directory
        checked = contribution_workbook_guard.verify_sources(directory, data.get('dataset'), sources=sources)
        result = reader(directory, data, sources)
        pinned = next((item for item in checked if item['dataset'] == data.get('dataset')), None)
        if pinned:
            provenance = result.get('provenance', {})
            # The guard and the reader share one scan, so their hashes agree by construction; a report
            # replaced since that scan is caught by hashing its file again.
            scanned = next((item for item in sources or () if item['id'] == pinned['dataset']), None)
            if (any(pinned.get(key) is not None and provenance.get(key) != pinned[key]
                    for key in ('workbook_sha256', 'database_sha256'))
                    or (scanned and contribution_workbook_guard.changed_since_scan(scanned))):
                raise ValueError('The published report or contribution source changed while the view was being prepared. Refresh the source before continuing.')
            provenance.update(source_financial_universe_verified=pinned.get('source_financial_universe_verified'),
                              source_proof_token=pinned.get('source_proof_token'))
        return result

    def error_text(self, error):
        """The reader's own sentence; hosted, an OSError (an errno, a file name) is one fixed sentence instead."""
        return HOSTED_IO_ERROR if self.access_cfg.enforced and isinstance(error, OSError) else str(error)

    def do_GET(self):
        if not self.valid_host():
            return self.send({"error": "Local access only."}, 403)
        parsed = urlparse(self.path)
        query = parse_qs(parsed.query)
        app = self.server.app
        try:
            if self.access_cfg.enforced:
                if parsed.path in access_gate.HEALTH_ROUTES:
                    return self.send({'status': 'ok'} if parsed.path == '/healthz'
                                     else access_gate.version_payload(self.access_app))
                if not (parsed.path in HOSTED_GET_ROUTES or parsed.path in STATIC_ASSETS
                        or parsed.path.startswith(('/geo/v2/', '/basemap/'))):
                    return self.send({'error': HOSTED_REFUSED}, 403)
            if parsed.path == '/api/health':
                return self.send({'application': 'MENTAT', 'project': str(PROJECT),
                                  'data_root': app.configured_root,
                                  'resolved_data_root': str(app.root),
                                  'server_code_sha256': SERVER_CODE_SHA256,
                                  'contribution_studio_revision': CONTRIBUTION_STUDIO_REVISION,
                                  'read_only': bool(getattr(app, 'read_only', False)),
                                  'report_directory':str(app.report_directory),
                                  **build_status(app, self.server.server_port)})
            if parsed.path == '/api/prime-status':
                # Read-only on every console: this process's startup prime (prewarm.prime_status), which
                # Start-MentatApp.ps1 waits on before it opens the app window.
                import prewarm
                return self.send(prewarm.prime_status())
            if parsed.path == "/api/session":
                if query.get('light', ['0'])[0] == '1':
                    build = build_status(app, self.server.server_port)
                    # hosted: the page switches off what the hosted allowlist refuses (the candidate-spending card and panel)
                    return self.send({'token':self.session_token(), 'read_only':bool(getattr(app, 'read_only', False)),
                                      'hosted': bool(self.access_cfg.enforced),
                                      'evidence_collection_available':True, 'api_version':2,
                                      'report_directory':str(app.report_directory), 'build': build, **build})
                return self.send({"token": self.session_token(), 'hosted': bool(self.access_cfg.enforced),
                                  **app.summary(query.get('include_archived', ['0'])[0] in ('1','true'))})
            if parsed.path == '/api/contributions/session':
                def build(sources):
                    registry = snapshot_registry.registry(app.report_directory, sources)
                    entry = industry_scope.from_query(registry, parsed.query)
                    return {'token': self.session_token(), 'read_only': bool(getattr(app, 'read_only', False)),
                            'industry': entry['industry'] if entry else None,
                            'industries': [{key: item[key] for key in ('industry','slug','available','build_stamp')}
                                           for item in registry['snapshots']]}
                return self.send(self.source_metadata(build))
            if parsed.path == '/api/insights/catalog':
                import analysis_data
                import response_cache
                # FIX 4: each dataset names what its report's own Data Quality check says was collected short.
                return self.send_entry(self.source_metadata(lambda sources: response_cache.from_value(
                    collection_gaps.attach(app.report_directory, analysis_data.catalog(app.report_directory, sources)))))
            if parsed.path == '/api/forecast/pac':
                import pac_giving_forecast
                return self.send(pac_giving_forecast.catalog())
            if parsed.path == '/api/forecast/catalog':
                import giving_forecast
                return self.send(giving_forecast.catalog())
            if parsed.path == '/api/contributions/catalog':
                # One scan of the report folder for this request: the semantic
                # catalogue and the registry annotation below read the same list.
                def build(sources):
                    catalog = contribution_query.semantics.catalog(app.report_directory, sources)
                    checks = {item['dataset']: item for item in contribution_workbook_guard.verify_sources(app.report_directory, strict=False, sources=sources)}
                    for dataset in catalog['datasets']:
                        check = checks.get(dataset['id'])
                        if check and not check['matched']:
                            dataset.update(available=False, message=check['error'])
                        elif check and any(check.get(key) is not None and dataset.get('provenance', {}).get(key) != check[key]
                                           for key in ('workbook_sha256', 'database_sha256')):
                            dataset.update(available=False, message='The report changed during source discovery. Refresh sources.')
                    registry = snapshot_registry.registry(app.report_directory, sources)
                    entry = industry_scope.from_query(registry, parsed.query)
                    if entry:
                        catalog['datasets'] = [item for item in catalog['datasets'] if item.get('id') == entry['dataset']]
                    stamps = {item['dataset']: item['build_stamp'] for item in registry['snapshots']}
                    for dataset in catalog['datasets']:
                        dataset['build_stamp'] = stamps.get(dataset.get('id'))
                    catalog['industry'] = entry['industry'] if entry else None
                    import response_cache
                    return response_cache.from_value(collection_gaps.attach(app.report_directory, catalog))
                return self.send_entry(self.source_metadata(build))
            if parsed.path == '/api/contributions/geography/catalog':
                # geo_catalog() is industry-independent; the parameter is accepted
                # and echoed so one client call carries one scope everywhere.
                entry = industry_scope.from_query(app.report_directory, parsed.query)
                return self.send({**contribution_geo.geo_catalog(),
                                  'industry': entry['industry'] if entry else None})
            if parsed.path == '/api/snapshots':
                def build(sources):
                    registry = snapshot_registry.registry(app.report_directory, sources)
                    entry = industry_scope.from_query(registry, parsed.query)
                    return {**registry, 'snapshots': [entry]} if entry else registry
                return self.send(self.source_metadata(build))
            if parsed.path == "/api/runs":
                return self.send(app.summary(query.get('include_archived', ['0'])[0] in ('1','true')))
            if parsed.path == '/api/evidence-csv':
                if query.get('job', [''])[0]:
                    return self.send(app.evidence_csv_job(query.get('job', [''])[0]))
                # A pane asks for its committee's figures, not its rows: same id
                # and committee, no paging, and the same refusals (an unknown
                # committee is still M_COMMITTEE 400).
                if query.get('summary', [''])[0] and query.get('id', [''])[0]:
                    return self.send(app.evidence_csv_summary(query.get('id', [''])[0],
                        query.get('committee', [''])[0], query.get('cycle', [''])[0]))
                # Page arguments are validated by the reader, which answers with
                # the reader's own message. A bare int() leaked Python's own text
                # ("invalid literal for int()") into the status line.
                return self.send(app.evidence_csv_view(query.get('id', [''])[0], query.get('committee', [''])[0],
                    query.get('cycle', [''])[0], query.get('offset', ['0'])[0], query.get('limit', ['200'])[0]))
            if parsed.path == '/api/evidence-stats':
                # One number, for the pane's stat strip. Every refusal is the
                # reader's own sentence, answered 400 by the handler below.
                return self.send(app.evidence_stats(
                    query.get('id', [''])[0], query.get('committee', [''])[0],
                    query.get('cycle', [''])[0]))
            if parsed.path == '/api/evidence-geocode':
                # Read-only on purpose: starting and stopping are POSTs, so a
                # GET can never begin sending addresses anywhere.
                return self.send(app.evidence_geocode(query.get('id', [''])[0]))
            if parsed.path == '/api/evidence-geo':
                # One pane, one cycle, one direction. Every refusal is the
                # reader's own sentence, raised as ValueError and answered 400
                # by the handler below, like every other evidence route.
                return self.send(app.evidence_geography(
                    query.get('source', [''])[0], query.get('id', [''])[0],
                    query.get('committee', [''])[0], query.get('cycle', [''])[0],
                    query.get('direction', ['spent'])[0]))
            if parsed.path == '/api/evidence':
                identifier = query.get('id', [''])[0]
                root = PROJECT / '.cache/pac-evidence'
                if not identifier:
                    return self.send({'snapshots': [
                        {'id': p.parent.name, 'company': value.get('company'), 'retrieved_at': value.get('retrieved_at'),
                         'cycles': sorted({row.get('cycle') for row in value.get('committee_finances', []) if row.get('cycle')}),
                         'committees': [{'id': row.get('committee_id'), 'cycle': str(row.get('cycle'))}
                                        for row in value.get('committee_finances', []) if row.get('committee_id')]}
                        for p in sorted(root.glob('*/evidence.json')) if not p.parent.name.startswith('direct-')
                        if (value := read_json(p)).get('version') == 1]})
                if not re.fullmatch(r'[a-z0-9][a-z0-9_-]{0,99}', identifier):
                    raise ValueError('Invalid evidence snapshot.')
                value = read_json(root / identifier / 'evidence.json')
                if value.get('version') != 1:
                    raise ValueError('Evidence snapshot unavailable.')
                # S7: a snapshot names its committees but never says what kind of
                # committee each one is, so the picker read 'FAIRSHAKE - C00835959'
                # with nothing to tell a Super PAC from a candidate committee. The
                # label is read from the FEC committee master here, at read time,
                # rather than written into the stored snapshot: it is a label, not
                # evidence, and a master that cannot be read must never cost the
                # reader the snapshot itself.
                try:
                    for committee in value.get('committee_finances') or []:
                        committee['category'], committee['category_source'] = (
                            evidence_csv.classify_committee(committee.get('committee_id', '')))
                except Exception:
                    pass
                estimates = read_json(root / identifier / 'cash-estimates.json')
                evidence_hash = hashlib.sha256((root / identifier / 'evidence.json').read_bytes()).hexdigest()
                if (estimates.get('source_snapshot_sha256') == value.get('source_snapshot_sha256')
                        and estimates.get('evidence_sha256') == evidence_hash
                        and estimates.get('method_code_sha256') == hashlib.sha256((PROJECT / 'tools/pac_cash_model.py').read_bytes()).hexdigest()):
                    value['cash_estimates'] = estimates
                return self.send(value)
            if parsed.path == '/api/evidence-network':
                from downstream import snapshot_view
                return self.send(snapshot_view(query.get('id', [''])[0],
                    query.get('committee', [''])[0], query.get('cycle', [''])[0],
                    PROJECT / '.cache/pac-evidence'))
            if parsed.path == '/api/pac-flow-catalog':
                from pac_flow_access import catalog
                return self.send(catalog(app, query.get('industry', [''])[0], query.get('q', [''])[0]))
            if parsed.path == '/api/candidate-spending':
                # A campaign's own Schedule B spending, from saved snapshots only (candidate_spending.py): no fetch here.
                import candidate_spending
                return self.send(*candidate_spending.respond(query))
            if parsed.path == '/api/flow':
                # Lane F (boards #44-#46): where a committee's (or an industry's) money came
                # from and where it went, built from the same downstream measurement as
                # /api/evidence-network (flow_data.py). A malformed request is 400 like every
                # route; a flow that disagrees with the evidence network, or an attributed
                # basis without an estimate bundle, is refused 409 with its reason.
                import flow_data
                run = query.get('run', ['published-reports'])[0] or 'published-reports'

                def workbooks(label):
                    # The industry's workbooks, chosen exactly as /api/graph chooses them.
                    reports = app.published_reports() if run == 'published-reports' else None
                    chosen = next((group for group in app.graph_catalog(run, reports)['industries']
                                   if group['name'] == label), None)
                    if not chosen:
                        raise ValueError('Choose an industry with completed reports.')
                    return [app.report_path(run, index) for index in chosen['indices']]
                try:
                    body = flow_data.flow(query, report_directory=app.report_directory,
                                          root=PROJECT / '.cache/pac-evidence', workbooks=workbooks,
                                          allow_direct_collect=not getattr(app, 'read_only', False))
                except flow_data.FlowRefused as error:
                    return self.send(error.body(), 409)
                if self.access_cfg.enforced and not donor_rows_allowed(self.access_user):
                    body = without_person_keys(without_individual_donors(body))
                return self.send(body)
            if parsed.path == '/api/evidence-job':
                identifier = query.get('id', [''])[0]
                if not re.fullmatch(r'snapshot-[a-f0-9]{32}', identifier):
                    raise ValueError('Invalid evidence collection.')
                root = PROJECT / '.cache/pac-evidence' / identifier
                job = read_json(root / 'collection.json')
                if not job:
                    raise ValueError('Evidence collection unavailable.')
                if job.get('state') == 'collecting' and not process_alive(job.get('pid')):
                    job['state'] = 'completed' if read_json(root / 'evidence.json').get('version') == 1 else 'failed'
                    job['message'] = 'Snapshot ready.' if job['state'] == 'completed' else 'Collection interrupted; saved responses are retained.'
                return self.send(job)
            if parsed.path == "/api/run":
                return self.send(app.snapshot(query.get("id", [""])[0], query.get("lane", [""])[0]))
            if parsed.path == "/api/report":
                path = app.report_path(query.get('id', [''])[0], int(query.get('index', ['-1'])[0]))
                return self.send(report_publication.read([path], path.read_bytes), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            if parsed.path == '/api/reports':
                return self.send(app.report_catalog(query.get('id', [''])[0]))
            if parsed.path == '/api/report-sources':
                return self.send(app.report_sources())
            if parsed.path == '/api/report-data':
                path = app.report_path(query.get('id', [''])[0], int(query.get('index', ['-1'])[0]))
                sheet, header = query.get('sheet', [''])[0], int(query.get('header', ['0'])[0])
                if sheet and self.access_cfg.enforced and not donor_rows_allowed(self.access_user):
                    import report_insights, response_cache
                    outcome = report_insights.person_level_sheet(report_insights.stamp(path), sheet, header)
                    if outcome != report_insights.ALLOW:
                        redacted = outcome == report_insights.REDACT and without_contributor_names(
                            json.loads(worksheet_entry(path, sheet, header).body))
                        if not redacted:   # REFUSE, or a REDACT sheet read at a header row that is not its own
                            return self.send({'error': HOSTED_PERSON_LEVEL_REFUSED}, 403)
                        return self.send_entry(response_cache.from_value(redacted))
                return self.send_entry(worksheet_entry(path, sheet, header))
            if parsed.path == '/api/graph-catalog':
                return self.send(app.graph_catalog(query.get('id',[''])[0]))
            if parsed.path == '/api/cross-reference':
                # industry is the only key this view accepts; every other one is
                # still refused, including a path or a run id.
                selector = parse_qs(parsed.query, keep_blank_values=True)
                if set(selector) - {'industry'}:
                    raise ValueError('The cross-reference view accepts no path or run parameters.')
                root = PROJECT / 'state' / 'cross-reference'
                if 'industry' not in selector:
                    return self.send(latest_cross_reference(root))
                scope = industry_scope.resolve(app.report_directory, industry=selector['industry'][0])
                return self.send(cross_reference_for_industry(root, scope['industry']))
            if parsed.path in ('/api/comparison', '/api/comparison-flows', '/api/graph'):
                # A slug resolves to its label; a catalogue name the registry does
                # not publish is handed on untouched, so these routes keep serving
                # every workbook graph_catalog() offers. One registry per request,
                # and none at all when no industry was named.
                wanted = query.get('industry',[''])[0]
                if parsed.path == '/api/graph':
                    import response_cache
                    started = time.perf_counter()
                    params = {**{k:v[0] for k,v in query.items()}, 'industry': wanted}
                    def build(sources):
                        registry = snapshot_registry.registry(app.report_directory, sources) if wanted else None
                        label = industry_scope.label_or_value(registry, wanted) if registry else wanted
                        params['industry'] = label
                        return response_cache.from_value(app.graph(query.get('id',[''])[0], label,
                                                                  dict(params), registry))
                    self.send_entry(self.source_metadata(build, wanted))
                    if prefetch_request(self):
                        return   # graph.js's idle prefetch: answered the same, never recorded (H4)
                    return record_access(self, 'graph_drawer' if params.get('mode') == 'activity' else 'graph_view',
                                         params, started)
                registry = snapshot_registry.registry(app.report_directory) if wanted else None
                label = industry_scope.label_or_value(registry, wanted) if registry else wanted
                if parsed.path == '/api/comparison-flows':
                    started = time.perf_counter()
                    self.send_entry(app.comparison_entry(query.get('id',[''])[0], label))
                    return record_access(self, 'comparison_flows', {'id': query.get('id',[''])[0], 'industry': label},
                                         started)
                if self.access_cfg.enforced and not donor_rows_allowed(self.access_user):
                    if names_person_subject(query):
                        return self.send({'error': HOSTED_PERSON_LEVEL_REFUSED}, 403)
                    return self.send(without_person_subjects(app.comparison(query.get('id',[''])[0], label, False)))
                return self.send(app.comparison(query.get('id',[''])[0], label, False))
            if parsed.path == '/api/publish/pocket':
                job = app.publication(query.get('id', [''])[0])
                if job is None:
                    return self.send({'error': 'Publication not on this console.'}, 404)
                return self.send(job)
            if parsed.path == '/api/studio/briefing':
                return self.send(app.briefing())
            if parsed.path == '/api/studio/briefing/item':
                svg = app.briefing_item(query.get('id', [''])[0])
                if svg is None:
                    return self.send({'error': 'Briefing item not on this console.'}, 404)
                return self.send(svg.encode('utf-8'), content_type='text/plain; charset=utf-8')
            if parsed.path == '/api/studio/views':
                identifier = query.get('id', [''])[0]
                if not identifier:
                    return self.send(app.list_views())
                view = app.get_view(identifier)
                if view is None:
                    return self.send({'error': 'Saved view not on this console.'}, 404)
                return self.send({'view': view})
            if parsed.path == '/':
                # The page is rendered from app/templates/, with exactly the
                # headers it has always carried: same content type, same
                # revalidation, same 304 on a matching ETag.
                import console_page
                before = console_page.stamp()
                content, etag = console_page.page()
                if console_page.stamp() != before:
                    # The static route's double-stat guard, read over the whole
                    # template set: a page half from one edit and half from the
                    # next is refused rather than served.
                    return self.send({'error':'The interface changed while loading. Refresh the page.'}, 409)
                unchanged = etag in [tag.strip() for tag in self.headers.get('If-None-Match', '').split(',')]
                return self.send(b'' if unchanged else content, 304 if unchanged else 200,
                                 content_type='text/html; charset=utf-8',
                                 cache_control='private, no-cache', etag=etag)
            # Lane P: map layer geography and the offline street basemap, each
            # a whitelisted route (send_map_file); the tile archive by Range.
            if parsed.path.startswith(('/geo/v2/', '/basemap/')) and parsed.path not in STATIC_ASSETS:
                return self.send_map_file(parsed.path)
            assets = STATIC_ASSETS
            if parsed.path in assets:
                filename, mime = assets[parsed.path]
                path = Path(__file__).parent / filename
                if not path.is_file():
                    return self.send({"error": "Not found."}, 404)
                info = path.stat()
                content, etag = static_asset((str(path), info.st_size, info.st_mtime_ns, info.st_ctime_ns))
                current = path.stat()
                if (current.st_size, current.st_mtime_ns, current.st_ctime_ns) != (info.st_size, info.st_mtime_ns, info.st_ctime_ns):
                    return self.send({'error':'The interface changed while loading. Refresh the page.'}, 409)
                unchanged = etag in [tag.strip() for tag in self.headers.get('If-None-Match', '').split(',')]
                return self.send(b'' if unchanged else content, 304 if unchanged else 200,
                                 content_type=mime, cache_control='private, no-cache', etag=etag)
            return self.send({"error": "Not found."}, 404)
        except (ValueError, OSError, KeyError, zipfile.BadZipFile) as error:
            return self.send({"error": self.error_text(error)}, 400)

    def do_POST(self):
        analytical_paths = {'/api/contributions/query', '/api/contributions/records', '/api/contributions/geography',
                            '/api/insights/query', '/api/insights/records', '/api/forecast/project'}
        hosted = self.access_cfg.enforced
        read_only = hosted or getattr(self.server.app, 'read_only', False)
        if read_only and self.path.startswith('/api/studio/briefing'):
            self.discard_body()
            return self.send({'error': 'Read-only preview. Briefing changes are not saved.'}, 403)
        if read_only and self.path.startswith('/api/studio/views'):
            self.discard_body()
            return self.send({'error': 'Read-only preview. Save views from the main console.'}, 403)
        if read_only and self.path.startswith('/api/evidence-geocode'):
            self.discard_body()
            return self.send({'error': geocode.M_READ_ONLY}, 403)
        # /api/export/pdf (figure_pdf.py) writes only a temporary file, so the read-only preview serves it. A hosted
        # console serves none of the three: /api/evidence-files lists folders, /api/evidence-csv reads any path it is
        # given and takes uploads, and /api/export/pdf is a download.
        if read_only and self.path not in (analytical_paths if hosted else
                                           analytical_paths | {'/api/evidence-csv', '/api/evidence-files', '/api/export/pdf'}):
            self.discard_body()
            return self.send({"error": HOSTED_REFUSED if hosted else
                              "This preview is read-only while the original console manages collection."}, 403)
        if hosted and self.path in DONOR_ROW_PATHS and not donor_rows_allowed(self.access_user):
            self.discard_body()
            return self.send({'error': HOSTED_DONOR_ROWS_REFUSED}, 403)
        if not self.valid_host() or not secrets.compare_digest(self.headers.get("X-Mentat-Token", ""), self.session_token()):
            self.discard_body()
            return self.send({"error": "Refresh the app before continuing."}, 403)
        origin = self.headers.get("Origin")
        if origin and not self.valid_origin(origin):
            self.discard_body()
            return self.send({"error": "Local access only."}, 403)
        try:
            # A chosen CSV streams straight to the data-root store: the request
            # thread copies it in chunks and never holds the decoded file, so
            # the generic body gate below does not apply to it. Its own ceiling
            # is the reader's MAX_BYTES, which the page pre-flights from
            # /api/evidence-csv so a refusal is shown before anything uploads.
            if self.path == '/api/evidence-csv' and 'application/json' not in self.headers.get('Content-Type', '').lower():
                store = self.server.app.evidence_store
                try:
                    expected = int(self.headers.get('Content-Length', '0'))
                except ValueError:
                    raise ValueError('The upload did not state its size.') from None
                if expected > evidence_csv.MAX_BYTES:
                    self.close_connection = True
                    return self.send({'error': evidence_csv.M_BYTES}, 400)
                upload = store.begin_upload(
                    evidence_csv.decode_name(self.headers.get('X-Evidence-Name', '')),
                    sheet=evidence_csv.decode_name(self.headers.get('X-Evidence-Sheet', '')))
                try:
                    received = 0
                    while received < expected:
                        chunk = self.rfile.read(min(evidence_csv.CHUNK_BYTES, expected - received))
                        if not chunk:
                            raise ValueError(evidence_csv.M_SHORT(received, expected))
                        received += len(chunk)
                        upload.write(chunk)
                except BaseException:
                    upload.abort()
                    raise
                return self.send(upload.finish())
            if self.path == '/api/export/pdf':
                # One figure to one PDF (figure-pdf.js). It reads its own body: a figure with a map
                # raster can pass the 7 MB cap below, and it answers a download, not JSON.
                import figure_pdf
                return figure_pdf.respond(self, STATIC_ASSETS, Path(__file__).parent, CONSOLE_CSP)
            length = int(self.headers.get("Content-Length", "0"))
            if length < 1 or length > 7_000_000:
                self.discard_body()
                raise ValueError("Request is empty or too large.")
            if self.path == '/api/studio/views' and length > 512 * 1024:
                self.discard_body()
                raise ValueError('That view is too large to save. Remove a sheet and save again.')
            data = json.loads(self.rfile.read(length))
            if not isinstance(data, dict):
                raise ValueError("Expected an upload or run request.")
            if hosted and not donor_rows_allowed(self.access_user) and names_person_level_field(data):
                return self.send({'error': HOSTED_PERSON_LEVEL_REFUSED}, 403)
            if self.path == '/api/forecast/project':
                import giving_forecast
                return self.send(giving_forecast.project(data))
            # The body as the client sent it (scope_body and the readers change `data`), for record_access.
            started = time.perf_counter()
            sent = (json.loads(json.dumps(data)) if ACCESS_RECORDER is not None and
                    self.path in ('/api/insights/query', '/api/contributions/query') else None)
            if self.path in ('/api/insights/query', '/api/insights/records'):
                import analysis_data
                # One scan of the report folder for this request: scope, guard and reader share it.
                sources = contribution_data._sources(self.server.app.report_directory)
                entry = self.scope_body(data, sources)
                reader = analysis_data.query if self.path.endswith('/query') else analysis_data.records
                result = self.contribution_result(reader, data, sources)
                # The records come straight off the stored facts, which still file ActBlue and WinRed under
                # their old category; analysis_data's code hash keys every Insights analysis, so the label
                # is applied to the answer here rather than there.
                evidence_csv.relabel_conduit_rows(result.get('rows'))
                result.update(industry_scope.echo(entry))
                self.send(result)
                if sent is not None and not (isinstance(result, dict) and result.get('status') in ('building', 'queued', 'preparing')):
                    sent.pop('retry', None)
                    record_access(self, 'insights_query', sent, started)
                return
            if self.path in ('/api/contributions/query', '/api/contributions/records'):
                sources = contribution_data._sources(self.server.app.report_directory)
                entry = self.scope_body(data, sources)
                reader = contribution_query.query if self.path.endswith('/query') else contribution_query.records
                result = self.contribution_result(reader, data, sources)
                if isinstance(result, dict):
                    # Same label as the Insights route above, whichever store the records came from.
                    if self.path.endswith('/records'):
                        evidence_csv.relabel_conduit_rows(result.get('rows'))
                    result.update(industry_scope.echo(entry))
                self.send(result)
                if sent is not None and not (isinstance(result, dict) and result.get('status') in ('building', 'queued', 'preparing')):
                    record_access(self, 'contribution_query', sent, started)
                return
            if self.path == '/api/contributions/geography':
                entry = self.scope_body(data)
                result = contribution_geo.geography(data)
                if hosted and isinstance(result, dict) and result.get('errors'):
                    # geography() files str(error) of a failed Census read (a URL, a file, an errno): one sentence
                    result = {**result, 'errors': {key: HOSTED_IO_ERROR for key in result['errors']}}
                if isinstance(result, dict):
                    result.update(industry_scope.echo(entry))
                return self.send(result)
            if self.path == '/api/restart':
                if console_busy(self.server.app):
                    return self.send({'error': 'Collection is running. Restart after it finishes.'}, 409)
                # A console already running the build that is on disk has nothing
                # to restart. Answering 'restarting' anyway sent the launcher to
                # reuse the process it already had - correctly - and left the
                # reader watching a poll for a new pid that could never arrive,
                # and two minutes later being told to close the tab.
                # One snapshot, compared against itself, so the block handed back
                # with current:true cannot carry on-disk values from a different
                # moment than the ones the decision was made on.
                build = build_status(self.server.app, self.server.server_port)
                if (build['code_sha256'] == build['on_disk_code_sha256_now']
                        and build['assets_sha256'] == build['on_disk_assets_sha256_now']):
                    return self.send({'restarting': False, 'current': True, 'build': build})
                return self.send(self.server.app.restart(self.server.server_port))
            if self.path == "/api/upload":
                return self.send(self.server.app.upload(data))
            if self.path == '/api/roster':
                return self.send(self.server.app.create_roster(data))
            if self.path == '/api/evidence-collect':
                return self.send(self.server.app.collect_evidence(data))
            if self.path == "/api/keys":
                with self.server.app.lock:
                    count = save_keys([str(data.get("keys", ""))])
                return self.send(key_status())
            if self.path == "/api/start":
                return self.send(self.server.app.launch(data, background=True))
            if self.path == "/api/resume":
                return self.send(self.server.app.resume(data))
            if self.path == '/api/quality/acknowledge':
                return self.send(self.server.app.acknowledge_quality(data))
            if self.path == '/api/run-archive':
                return self.send(self.server.app.archive_run(data))
            if self.path == '/api/evidence-csv':
                # A body carrying a path names a file on this computer and is
                # read off disk by the store, asynchronously, exactly like a
                # streamed upload. The legacy {name, text} body is unchanged.
                if 'path' in data:
                    job = self.server.app.import_evidence_path(data)
                    # The store records a refusal it reached before reading a
                    # byte (no such file, a folder, too large, a type it does
                    # not open) as a failed JOB, so the page can poll it like
                    # any other failure. Over HTTP that same refusal is the
                    # answer to this request, so it is sent as an error rather
                    # than as a job the caller would poll to learn it never
                    # started. Only a refusal can be failed here: a reader that
                    # has started is still 'reading' when this returns.
                    if job.get('state') == 'failed':
                        return self.send({'error': job.get('message') or evidence_csv.M_GONE(job.get('id', ''))}, 400)
                    return self.send(job)
                return self.send(self.server.app.import_evidence_csv(data))
            if self.path == '/api/evidence-geocode':
                # S8b ruling 10. The ONLY door that sends contributor addresses
                # to the US Census Geocoder, and it opens on an explicit Start.
                return self.send(self.server.app.start_evidence_geocode(data))
            if self.path == '/api/evidence-files':
                # Read-only, and authenticated by the same token as every other
                # POST: do_GET checks only the Host header, so a folder listing
                # has no business being answered there.
                return self.send(self.server.app.evidence_files(data))
            if self.path == "/api/stop":
                return self.send(self.server.app.stop(data))
            if self.path == "/api/workers":
                return self.send(self.server.app.set_workers(data))
            if self.path == '/api/publish/pocket':
                return self.send(self.server.app.publish_pocket(data))
            if self.path == '/api/studio/briefing/add':
                return self.send(self.server.app.add_to_briefing(data))
            if self.path == '/api/studio/briefing/remove':
                return self.send(self.server.app.remove_from_briefing(data))
            if self.path == '/api/studio/briefing/order':
                return self.send(self.server.app.order_briefing(data))
            if self.path == '/api/studio/views':
                return self.send(self.server.app.save_view(data))
            if self.path == '/api/studio/views/delete':
                return self.send(self.server.app.delete_view(data))
            return self.send({"error": "Not found."}, 404)
        except Conflict as error:
            return self.send({'error': str(error)}, 409)
        except RunNotFound as error:
            return self.send({"error": str(error)}, 404)
        except RuntimeError as error:
            if self.path in analytical_paths:
                return self.send({'error': str(error)}, 400)
            return self.send({"error": "The encrypted key store could not be accessed for this Windows user."}, 400)
        except (ValueError, OSError, KeyError, zipfile.BadZipFile) as error:
            return self.send({"error": self.error_text(error)}, 400)


# The stdlib listen backlog is 5: while query threads hold the GIL the accept
# loop falls behind, and Windows refuses the sixth waiting connection outright,
# so a page load lost stylesheets and scripts (ERR_CONNECTION_REFUSED) whenever
# analysis requests were running.
CONSOLE_LISTEN_BACKLOG = 128


def redirect_state_writes(root, report_directory=None):
    """MENTAT_STATE_WRITE_ROOT: every cache the console writes while it serves moves under `root`, so the
    checkout's state/ and the report folder can be mounted read-only. The inputs stay where they are:
    prewarm.READ_ONLY_STATE (analytical indexes, accepted sources) and the rest of state/. Two small caches
    that are also shipped inputs, the source proofs and the Census geography, are seeded into `root` once.
    A missing analytical index is refused, as prewarm's own sandbox does, instead of built.
    With a `report_directory` (hosted), a write to the checkout or the report folder that was not redirected is
    refused too (prewarm.install_write_guard, process-wide: only `root` and the temp directory may be written)."""
    import access_stats, analysis_data, contribution_data, drawer_sidecars, glassbox_emit, prewarm
    root = Path(root).resolve()
    for module_name, attribute, relative in prewarm.STATE_REDIRECTS:
        setattr(importlib.import_module(module_name), attribute, root / relative)
    analysis_data.HINTS_PATH = root / 'prewarm' / 'hints.json'
    access_stats.configure(root / 'prewarm' / 'access.sqlite')
    contribution_geo.CACHE = root / 'geography' / 'census-v1'
    drawer_sidecars.configure(state_root=root / 'drawer-sidecars', emit=False)
    glassbox_emit.EVENTS_DIR = root / 'glassbox'
    for relative in ('contribution-source-proofs', 'geography/census-v1'):
        shipped, target = prewarm.DEFAULT_STATE_ROOT / relative, root / relative
        if shipped.is_dir() and not target.exists():
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copytree(shipped, target)

    def refuse_build(source, token, destination):
        contribution_data._progress(token, status='error', failed_at=time.time(), message=prewarm.INDEX_REFUSED)

    contribution_data._materialize = refuse_build
    if report_directory is not None:
        prewarm.install_write_guard(guarded=(PROJECT, report_directory), allowed=(root, tempfile.gettempdir()))
    return root


def console_server(port, handler=None, host='127.0.0.1'):
    """The console's listener: ThreadingHTTPServer with a deeper listen backlog.
    Built through the module's ThreadingHTTPServer name, bound after the backlog
    is set, so a test that replaces that name still gets its stand-in."""
    server = ThreadingHTTPServer((host, port), handler or Handler, bind_and_activate=False)
    server.request_queue_size = CONSOLE_LISTEN_BACKLOG
    try:
        server.server_bind()
        server.server_activate()
    except BaseException:
        server.server_close()
        raise
    return server


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--port", type=int, default=8802)
    parser.add_argument("--data-root", default=str(Path(os.environ.get("LOCALAPPDATA", Path.home())) / "Mentat" / "console"))
    parser.add_argument("--keys-file", default="")
    parser.add_argument("--rscript", default="")
    parser.add_argument("--read-only", action="store_true", help="Inspect reports without allowing mutations or collection")
    options = parser.parse_args()
    access = access_gate.require_config_or_exit()
    if access.enforced and not options.read_only:
        sys.stderr.write("access gate: a hosted console (Cloudflare Access enforced) must run with --read-only\n")
        sys.exit(2)
    write_root = os.environ.get('MENTAT_STATE_WRITE_ROOT', '').strip()
    if access.enforced and not write_root:
        sys.stderr.write("access gate: a hosted console (Cloudflare Access enforced) must set MENTAT_STATE_WRITE_ROOT\n")
        sys.exit(2)
    Handler.access_cfg = access
    data_root = Path(options.data_root).resolve()
    keys_file = str(Path(options.keys_file).resolve()) if options.keys_file else ''
    rscript = str(Path(options.rscript).resolve()) if options.rscript else ''
    reports = report_directory(PROJECT)
    os.chdir(reports)
    server = console_server(options.port, host=access.bind_host)
    try:
        access_gate.assert_bind_consistent(access, server.server_address[0])    # the address really bound, before anything serves
    except access_gate.AccessConfigError as error:
        server.server_close()
        sys.stderr.write(f"access gate: {error}\n")
        sys.exit(2)
    server.app = Console(data_root, keys_file, rscript, options.read_only)
    server.app.port = options.port or server.server_port
    # One source discovery per stamp change, not per request (perf lane 2026-09-24).
    import source_memo
    source_memo.install()
    import analysis_data
    import prewarm
    # Only the console on the user's own data root warms the shared caches and records
    # warm-up hints: consoles that test suites start on a temporary data root must not
    # warm mentat/state. MENTAT_PREWARM=1 forces warming on, MENTAT_PREWARM=0 turns it off.
    choice = os.environ.get('MENTAT_PREWARM', '').strip()
    own_root = data_root == (Path(os.environ.get("LOCALAPPDATA", Path.home())) / "Mentat" / "console").resolve()
    warm = choice == '1' or (choice != '0' and own_root)
    # MENTAT_STATE_WRITE_ROOT (see redirect_state_writes). Neither a hosted console nor one with redirected state
    # warms: the warmer's own state and its calls back into the console are not redirected or authorised.
    if write_root:
        redirect_state_writes(write_root, server.app.report_directory if access.enforced else None)
    warm = warm and not (access.enforced or write_root)
    # Encoded Graph views and comparison flows persist across restarts in mentat/state/response-cache,
    # content-verified by every new process (perf lane 2026-09-24, cache). Like the warm caches, only
    # the console on the user's own data root writes there; set before the startup warm below, which
    # then refills memory from disk.
    if own_root and not write_root:
        import response_cache
        response_cache.configure()
    analysis_data.HINTS_ENABLED = warm and not options.read_only
    # Access recording (RACF value ranking) follows the same gate as the hints: only the console that
    # warms and may write. Without access_stats the recorder is None and the ranking uses hints only.
    global ACCESS_RECORDER
    ACCESS_RECORDER = prewarm.AccessRecorder.start() if warm and not options.read_only else None
    if warm:
        prewarm.start_startup_warm(read_only=options.read_only, app=server.app)
        prewarm.launch_warmer(data_root=data_root, read_only=options.read_only, console_port=server.app.port)
        # Lane P1: after a cold /api/flow, the snapshot's other cycles for the same committee are
        # prepared in the background (flow_data.warm_cycles; MENTAT_FLOW_WARM=0 turns it off).
        import flow_data
        flow_data.WARM_CYCLES = True
    print(f"Mentat app: http://127.0.0.1:{options.port}" if not access.enforced else
          f"Mentat app (hosted, read-only): bind {access.bind_host or '*'}:{options.port}", flush=True)
    server.serve_forever()


if __name__ == "__main__":
    main()
