"""fec/xlsx_openpyxl_authors.py (D-178): the openpyxl stand-ins for the three @oai/artifact-tool helpers, no Node needed."""
from __future__ import annotations

import hashlib
import json
import os
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock
from zipfile import ZipFile

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "fec"))

import xlsx_openpyxl_authors as A                          # noqa: E402
from stance.workbook_cell_sync import merge_cells          # noqa: E402

from ops.compute import xlsx_content_diff as X              # noqa: E402


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)

    def write(self, name, payload):
        path = self.tmp / name
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        return path

    def cells(self, path, sheet):
        return {c.coordinate: c.value for row in openpyxl.load_workbook(path)[sheet].iter_rows() for c in row if c.value is not None}


class SwitchTests(unittest.TestCase):
    def test_default_is_artifact_and_openpyxl_is_selectable(self):
        with mock.patch.dict(os.environ, clear=False):
            os.environ.pop("MELANGE_XLSX_WRITER", None)
            self.assertEqual(A.xlsx_writer(), "artifact")
            os.environ["MELANGE_XLSX_WRITER"] = " OpenPyXL "
            self.assertEqual(A.xlsx_writer(), "openpyxl")

    def test_an_unknown_writer_is_an_error_not_a_silent_fallback(self):
        with mock.patch.dict(os.environ, {"MELANGE_XLSX_WRITER": "excel"}):
            with self.assertRaises(ValueError):
                A.xlsx_writer()


class CoerceTests(unittest.TestCase):
    def test_zoned_timestamps_become_utc_dates_and_everything_else_stays_text(self):
        self.assertEqual(A._coerce("2023-01-17T20:02:36Z"), datetime(2023, 1, 17, 20, 2, 36))
        self.assertEqual(A._coerce("2023-01-17T20:02:36-05:00"), datetime(2023, 1, 18, 1, 2, 36))
        self.assertEqual(A._coerce("2023-01-17T20:02:36.123Z"), datetime(2023, 1, 17, 20, 2, 36, 123000))
        for text in ("2023-01-17T20:02:36", "2023-01-17", "007", "=A1", "2023-13-45T99:99:99Z", "", "trail "):
            self.assertEqual(A._coerce(text), text)


class CellSyncTests(Base):
    def changes(self):
        return [{"sheet": "Rank & Allocation", "cell": "G3", "value": "Champion"},
                {"sheet": "Rank & Allocation", "cell": "D4", "value": "='Blended Model'!D4"},
                {"sheet": "Rank & Allocation", "cell": "G5", "value": None},
                {"sheet": "Blended Model", "cell": "D4", "value": 0.25},
                {"sheet": "Blended Model", "cell": "A9", "value": "2023-01-17T20:02:36Z"}]

    def test_authored_cells_and_a_hash_bound_proof(self):
        output = self.tmp / "authored.xlsx"
        plan = self.write("plan.json", {"output": str(output), "changes": self.changes()})
        proof = A.author_cell_sync(plan)
        self.assertEqual((proof["schema"], proof["status"], proof["engine"]), ("workbook-cell-author-v1", "PASS", "openpyxl"))
        self.assertIsNone(proof["scores_and_ranks_verified"])
        self.assertEqual(proof["recalc"], "not-run")
        self.assertEqual((proof["plan_sha256"], proof["helper_sha256"], proof["output_sha256"]),
                         (sha(plan), sha(A.AUTHOR_MODULE), sha(output)))
        self.assertEqual(json.loads(Path(str(output) + ".report.json").read_text()), proof)
        self.assertEqual(self.cells(output, "Rank & Allocation"), {"G3": "Champion", "D4": "='Blended Model'!D4"})
        self.assertEqual(self.cells(output, "Blended Model")["D4"], 0.25)

    def test_merge_into_a_package_keeps_the_original_cells_and_styles(self):
        source = self.tmp / "source.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Rank & Allocation"
        ws["A3"], ws["G3"], ws["G6"] = "Member", "Old", "Kept"
        wb.create_sheet("Blended Model")["D4"] = 1
        wb.save(source)
        output = self.tmp / "authored.xlsx"
        plan = self.write("plan.json", {"output": str(output), "changes": self.changes()})
        A.author_cell_sync(plan)
        merged = self.tmp / "merged.xlsx"
        merge_cells(source, output, self.changes(), merged)
        merged_cells = self.cells(merged, "Rank & Allocation")
        self.assertEqual(merged_cells, {"A3": "Member", "G3": "Champion", "D4": "='Blended Model'!D4", "G6": "Kept"})
        expected = self.tmp / "expected.xlsx"
        wb["Rank & Allocation"]["G3"], wb["Rank & Allocation"]["D4"] = "Champion", "='Blended Model'!D4"
        wb["Blended Model"]["D4"] = 0.25
        wb["Blended Model"]["A9"] = datetime(2023, 1, 17, 20, 2, 36)
        wb["Blended Model"]["A9"].number_format = "General"          # a date cell with no number format, as authored
        wb.iso_dates = True
        wb.save(expected)
        self.assertEqual(X.compare_files(str(expected), str(merged))["counts"], {})


def honors_plan(output):
    def view(name, width):
        title = ["Melange — %s" % name] + [None] * (width - 1)
        return {"name": name, "values": [title, ["note"] + [None] * (width - 1), ["Generation", "g1"] + [None] * (width - 2),
                                         ["Member", "Tier"] + ["c%d" % i for i in range(width - 2)],
                                         ["Smith", "Tier 1 · Champion"] + [1.5] * (width - 2),
                                         ["Jones", "—"] + [2] * (width - 2)]}
    return {"generation_id": "g1", "output": str(output),
            "views": [view("Field Recognition", 8), view("Recognition Detail", 14), view("Overall Standing", 24)]}


class HonorsTests(Base):
    def test_three_views_styled_and_receipt_bound(self):
        output = self.tmp / "recognition_views.xlsx"
        plan = self.write("plan.json", honors_plan(output))
        receipt = A.author_honors_views(plan)
        self.assertEqual((receipt["schema"], receipt["status"], receipt["generation_id"]), ("honors-workbook-author-v1", "PASS", "g1"))
        self.assertEqual((receipt["plan_sha256"], receipt["helper_sha256"], receipt["output_sha256"]),
                         (sha(plan), sha(A.AUTHOR_MODULE), sha(output)))
        wb = openpyxl.load_workbook(output)
        self.assertEqual(wb.sheetnames, ["Field Recognition", "Recognition Detail", "Overall Standing"])
        overall, detail, field = wb["Overall Standing"], wb["Recognition Detail"], wb["Field Recognition"]
        self.assertEqual(sorted(str(m) for m in field.merged_cells.ranges), ["A1:F1", "A2:F2"])
        self.assertEqual((field.freeze_panes, detail.freeze_panes), ("E5", "C5"))
        self.assertEqual(field["A1"].fill.fgColor.rgb, "FF163F4B")
        self.assertEqual((field["A1"].font.sz, field["A1"].font.b), (17, True))
        self.assertEqual((detail["E5"].number_format, detail["I5"].number_format), ("0", "0.00"))
        self.assertEqual((overall["G5"].number_format, overall["H5"].number_format, overall["V5"].number_format),
                         ("0.00", "0", "0.00"))
        self.assertEqual((overall.column_dimensions["W"].width, overall.column_dimensions["A"].width), (38, 28))
        self.assertEqual(overall["B5"].value, "Tier 1 · Champion")

    def test_contents_round_trip_the_plan(self):
        output = self.tmp / "recognition_views.xlsx"
        plan = honors_plan(output)
        A.author_honors_views(self.write("plan.json", plan))
        wb = openpyxl.load_workbook(output)
        for view in plan["views"]:
            got = [[c.value for c in row] for row in wb[view["name"]].iter_rows()]
            self.assertEqual(got, view["values"])


def sentinel_plan(output, ledger_error=False):
    header = ["Bill", "Title", "Date", "Note"] + ["c%d" % i for i in range(23)] + ["Why"]
    row = lambda n: ["B%d" % n, "Title %d" % n, "2023-01-17T20:02:36Z", "x"] + [None] * 23 + ["#REF! broken" if ledger_error and n == 2 else "ok"]

    def interaction(sheet, bioguide, role, congress):
        return {"sheet": sheet, "_bioguide": bioguide,
                "values": ["Sec", role, "Member", "Name", "Title", "Sponsor", 3, "Cmte", "2023-01-17T20:02:36Z", congress, 0.5, "http://x", "sum"]}
    return {"schema": "sentinel-bill-workbook-plan-v1", "ledger_sheet": "Sentinel Bill Ledger", "output_workbook": str(output),
            "ledger_rows": [header, row(1), row(2)],
            "accepted_interactions": [interaction("Co-Sponsor Interactions", "A1", "r", 118),
                                      interaction("Co-Sponsor Interactions", "A1", "R", 118),      # same key, case-folded: dropped
                                      interaction("Co-Sponsor Interactions", "A2", "r", 118),
                                      interaction("Original Co-Sponsor Interaction", "A1", "r", 117)]}


class SentinelTests(Base):
    def test_ledger_interactions_dedup_and_report(self):
        output = self.tmp / "out" / "Sentinel Bill Ledger.xlsx"
        output.parent.mkdir()
        plan, report_path = self.write("plan.json", sentinel_plan(output)), self.tmp / "report.json"
        report = A.author_sentinel(plan, report_path)
        self.assertEqual(json.loads(report_path.read_text()), report)
        self.assertEqual((report["schema"], report["status"], report["engine"]), ("sentinel-bill-workbook-artifact-report-v1", "PASS", "openpyxl"))
        self.assertEqual((report["input_plan_sha256"], report["output_sha256"]), (sha(plan), sha(output)))
        self.assertEqual(report["appended"], {"Original Co-Sponsor Interaction": 1, "Co-Sponsor Interactions": 2})
        self.assertEqual(report["ledger_range"], "A1:AB3")
        self.assertEqual([i["sheet"] for i in report["inspections"]],
                         ["Sentinel Bill Ledger", "Original Co-Sponsor Interaction", "Co-Sponsor Interactions"])
        wb = openpyxl.load_workbook(output)
        self.assertEqual(wb.sheetnames, ["Sentinel Bill Ledger", "Original Co-Sponsor Interaction", "Co-Sponsor Interactions"])
        ledger, inter = wb["Sentinel Bill Ledger"], wb["Co-Sponsor Interactions"]
        self.assertEqual((ledger.freeze_panes, ledger.sheet_view.showGridLines), ("A2", False))
        self.assertEqual(ledger["A1"].fill.fgColor.rgb, "FF17365D")
        self.assertEqual(ledger["C2"].value, datetime(2023, 1, 17, 20, 2, 36))
        self.assertEqual([c.value for c in inter[1]][:3], ["Section", "Role", "Member"])
        self.assertEqual(inter.max_row, 3)
        self.assertEqual((inter.column_dimensions["M"].width, inter.row_dimensions[2].height), (72, 18))

    def test_a_formula_error_in_the_generated_ranges_stops_the_install(self):
        output = self.tmp / "Sentinel Bill Ledger.xlsx"
        with self.assertRaisesRegex(RuntimeError, "formula error"):
            A.author_sentinel(self.write("plan.json", sentinel_plan(output, ledger_error=True)), self.tmp / "report.json")
        self.assertFalse(output.exists())

    def test_a_plan_of_the_wrong_shape_is_refused(self):
        with self.assertRaisesRegex(ValueError, "invalid plan"):
            A.author_sentinel(self.write("plan.json", {"schema": "other"}), self.tmp / "report.json")


class ExactNumberTests(Base):
    def test_floats_are_written_as_the_same_double_not_sixteen_digits(self):
        values = [0.25555555555555554, 0.1140845070422535, 3.0, -0.0, 1e22, 5e-324, 123456789012345680.0]
        output = self.tmp / "authored.xlsx"
        changes = [{"sheet": "S", "cell": "A%d" % (i + 1), "value": v} for i, v in enumerate(values)]
        A.author_cell_sync(self.write("plan.json", {"output": str(output), "changes": changes}))
        with ZipFile(output) as z:
            xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        texts = [part.split("</v>")[0] for part in xml.split("<v>")[1:]]
        self.assertEqual([float(t) for t in texts], [v if v else 0.0 for v in values])
        self.assertEqual(texts[:4], ["0.25555555555555554", "0.1140845070422535", "3", "0"])

    def test_the_patch_is_scoped_to_the_authors(self):
        import openpyxl.cell._writer as w
        before = w.safe_string
        with A.exact_numbers():
            self.assertIsNot(w.safe_string, before)
        self.assertIs(w.safe_string, before)


class PackageTests(Base):
    def test_authored_workbooks_hold_no_shared_string_table_the_mergers_do_not_expect(self):
        output = self.tmp / "authored.xlsx"
        A.author_cell_sync(self.write("plan.json", {"output": str(output), "changes": CellSyncTests.changes(self)}))
        with ZipFile(output) as z:
            self.assertIn("xl/worksheets/sheet1.xml", z.namelist())
            self.assertIn(b't="inlineStr"', z.read("xl/worksheets/sheet1.xml"))


if __name__ == "__main__":
    unittest.main()
