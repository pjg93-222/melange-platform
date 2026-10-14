"""Tests for ops/compute/chain_parity.py (the Melange chain parity harness).

Synthetic trees under tmp_path only: no chain run, no Docker, no network, no live tree. The workbook tests build
their files with openpyxl (pinned in requirements.txt) and skip cleanly where it is missing.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import os
import re
import shutil
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "compute"))
import chain_parity as cp  # noqa: E402

STAMP_OLD = dt.datetime(2020, 1, 1, 12, 0, 0)
STAMP_NEW = dt.datetime(2026, 10, 11, 9, 30, 0)


# ------------------------------------------------------------------------------------ synthetic trees

def ranks_bytes(n: int = 533, swap=None, scores: dict | None = None) -> bytes:
    names = [f"Member {i:03d}" for i in range(n)]
    rank = {m: i + 1 for i, m in enumerate(names)}
    if swap:
        x, y = (names[i] for i in swap)
        rank[x], rank[y] = rank[y], rank[x]
    score = {m: 100.0 - i * 0.25 for i, m in enumerate(names)}
    score.update(scores or {})
    return json.dumps({"score": score, "rank": rank}, separators=(",", ":")).encode("ascii")


def base_outputs() -> dict:
    return {"ranks_python.json": ranks_bytes(), "dashboard_data.json": b'{"members":533,"built":"fixed"}',
            "bipartisan.json": b'{"pairs":[1,2,3]}', "out/log.txt": b"run finished\n"}


def write_tree(root: Path, files: dict) -> None:
    for rel, data in files.items():
        if data is not None:
            (root / rel).parent.mkdir(parents=True, exist_ok=True)
            (root / rel).write_bytes(data)  # bytes, so CRLF and exact content survive


def make_case(tmp_path: Path, a1=None, a2=None, b=None, c=None, frozen_extra=None) -> SimpleNamespace:
    """Frozen tree -> manifest -> three (optionally four) identical copies, each then 'run' by writing the
    base outputs plus its overrides (a None value leaves that file out of the root)."""
    frozen = tmp_path / "frozen"
    write_tree(frozen, {"inputs/seed.txt": b"seed\n", "inputs/cfg.json": b'{"k":1}', **(frozen_extra or {})})
    manifest = tmp_path / "frozen.manifest.json"
    assert cp.main(["snapshot", "--root", str(frozen), "--out", str(manifest)]) == 0
    roots = {}
    for label, over in (("A1", a1), ("A2", a2), ("B", b), ("C", c)):
        if over is None and label == "C":
            continue
        root = roots[label] = tmp_path / label
        shutil.copytree(frozen, root)
        write_tree(root, {**base_outputs(), **(over or {})})
        for rel, data in (over or {}).items():
            if data is None and (root / rel).exists():
                (root / rel).unlink()
    return SimpleNamespace(manifest=manifest, roots=roots, tmp=tmp_path)


def compare(case, *extra: str):
    args = ["compare", "--frozen", str(case.manifest), "--a1", str(case.roots["A1"]), "--a2", str(case.roots["A2"]),
            "--b", str(case.roots["B"])]
    if "C" in case.roots:
        args += ["--c", str(case.roots["C"])]
    report = case.tmp / "report.json"
    rc = cp.main([*args, "--report", str(report), *extra])
    return rc, json.loads(report.read_text(encoding="utf-8"))


def gate(rep, name):
    return next(g for g in rep["gates"] if g["name"] == name)


def failed(rep):
    return [g["name"] for g in rep["gates"] if g["status"] == "FAIL"]


def detail(rep, name):
    return "\n".join(gate(rep, name)["detail"])


# ------------------------------------------------------------------------------------- faithful trio

def test_faithful_trio_passes_with_exit_0(tmp_path):
    rc, rep = compare(make_case(tmp_path))
    assert rc == 0 and rep["result"] == "PASS" and rep["exit_code"] == 0
    assert [g["name"] for g in rep["gates"]] == ["G0", "G1", "G2", "G3", "G4"]
    assert failed(rep) == []
    sha = rep["byte_artifacts"]["ranks_python.json"]
    assert sha["A1"] == sha["A2"] == sha["B"] == hashlib.sha256(ranks_bytes()).hexdigest()
    assert set(rep["byte_artifacts"]) == set(cp.BYTE_ARTIFACTS)
    assert rep["classes"]["B"] == {"out/log.txt": {"class": "BYTE_EQUAL", "detail": ""}}
    assert rep["needs_acceptance"]["pending"] == [] and rep["movers"]["A1_vs_B"]["zero"] is True
    assert rep["counts"]["changed"]["B"] == {"added": 4, "removed": 0, "modified": 0}


def test_optional_c_root_is_gated_like_b(tmp_path):
    rc, rep = compare(make_case(tmp_path, c={}))
    assert rc == 0 and "C" in rep["byte_artifacts"]["bipartisan.json"] and "A1_vs_C" in rep["movers"]
    rc, rep = compare(make_case(tmp_path / "bad", c={"bipartisan.json": b'{"pairs":[1,2,4]}'}))
    assert rc == 1 and failed(rep) == ["G2"] and "bipartisan.json: C != A1" in detail(rep, "G2")


# ------------------------------------------------------------------------------------------- G2 byte

def test_one_byte_in_dashboard_data_fails_the_byte_gate(tmp_path):
    rc, rep = compare(make_case(tmp_path, b={"dashboard_data.json": b'{"members":533,"built":"fixeD"}'}))
    assert rc == 1 and failed(rep) == ["G2"]
    assert "dashboard_data.json: B != A1" in detail(rep, "G2")


def test_a1_a2_byte_difference_is_a_noise_floor_failure(tmp_path):
    rc, rep = compare(make_case(tmp_path, a2={"bipartisan.json": b'{"pairs":[1,2,9]}'}))
    assert rc == 1 and failed(rep) == ["G2"]
    assert "noise floor violated" in detail(rep, "G2") and "no mask can rescue a byte gate" in detail(rep, "G2")


def test_a_missing_byte_artifact_fails_never_skips(tmp_path):
    gone = {"bipartisan.json": None}
    rc, rep = compare(make_case(tmp_path, a1=gone, a2=gone, b=gone))
    assert rc == 1 and failed(rep) == ["G2"]
    assert "bipartisan.json: missing in A1, A2, B" in detail(rep, "G2") and "never a skip" in detail(rep, "G2")


def test_an_empty_byte_artifact_fails_instead_of_passing_vacuously(tmp_path):
    empty = {"dashboard_data.json": b""}
    rc, rep = compare(make_case(tmp_path, a1=empty, a2=empty, b=empty))
    assert rc == 1 and failed(rep) == ["G2"] and "dashboard_data.json: empty in A1, A2, B" in detail(rep, "G2")


# ------------------------------------------------------------------------------------------ G3 ranks

def test_rank_swap_in_b_fails_the_rank_gate_and_names_the_movers(tmp_path):
    rc, rep = compare(make_case(tmp_path, b={"ranks_python.json": ranks_bytes(swap=(10, 11))}))
    assert rc == 1 and failed(rep) == ["G2", "G3"]
    d = rep["movers"]["A1_vs_B"]
    assert d["zero"] is False and d["added"] == [] and d["dropped"] == []
    assert {(m["member"], m["before"], m["after"]) for m in d["rank_movers"]} == {
        ("Member 010", 11, 12), ("Member 011", 12, 11)}
    assert "Member 010" in detail(rep, "G3") and rep["movers"]["A1_vs_A2"]["zero"] is True


def test_member_count_532_fails_even_when_every_root_agrees(tmp_path):
    short = {"ranks_python.json": ranks_bytes(n=532)}
    rc, rep = compare(make_case(tmp_path, a1=short, a2=short, b=short))
    assert rc == 1 and failed(rep) == ["G3"]
    assert "A1: 532 members, expected 533" in detail(rep, "G3")
    rc, _ = compare(make_case(tmp_path / "ok", a1=short, a2=short, b=short), "--expect-members", "532")
    assert rc == 0


def test_dropped_member_in_b_is_listed(tmp_path):
    rc, rep = compare(make_case(tmp_path, b={"ranks_python.json": ranks_bytes(n=532)}))
    assert rc == 1 and rep["movers"]["A1_vs_B"]["dropped"] == ["Member 532"]


def test_rank_delta_tolerance_boundary_and_nan():
    base = json.loads(ranks_bytes(n=3))

    def moved(**scores):
        return cp.rank_delta(base, json.loads(ranks_bytes(n=3, scores=scores)))["score_movers"]

    assert moved(**{"Member 001": base["score"]["Member 001"] + 1e-13}) == []
    assert [m["member"] for m in moved(**{"Member 001": base["score"]["Member 001"] + 1e-11})] == ["Member 001"]
    assert [m["member"] for m in moved(**{"Member 002": float("nan")})] == ["Member 002"]
    assert cp.rank_delta(base, base)["zero"] is True


def test_unreadable_ranks_file_fails_the_gate_instead_of_crashing(tmp_path):
    junk = {"ranks_python.json": b"not json"}
    rc, rep = compare(make_case(tmp_path, a1=junk, a2=junk, b=junk))
    assert rc == 1 and "G3" in failed(rep) and "ranks_python.json" in detail(rep, "G3")


# ------------------------------------------------------------------------------------- G0 and G1 sets

def test_a_path_changed_only_in_b_fails_g1(tmp_path):
    rc, rep = compare(make_case(tmp_path, b={"out/extra.txt": b"x"}))
    assert rc == 1 and "G1" in failed(rep)
    assert "changed on B but not on A1: out/extra.txt" in detail(rep, "G1")
    assert rep["gates"][1]["data"]["B"]["only_B"] == ["out/extra.txt"]


def test_a_path_changed_only_in_a1_fails_g1(tmp_path):
    rc, rep = compare(make_case(tmp_path, a1={"out/extra.txt": b"x"}, a2={"out/extra.txt": b"x"}))
    assert rc == 1 and "G1" in failed(rep) and "changed on A1 but not on B: out/extra.txt" in detail(rep, "G1")


def test_nothing_ran_fails_g0(tmp_path):
    rc, rep = compare(make_case(tmp_path, frozen_extra=base_outputs()))
    assert rc == 1 and "G0" in failed(rep)
    assert all(f"{lbl}: no path changed versus the frozen manifest" in detail(rep, "G0") for lbl in ("A1", "A2", "B"))


def test_a1_a2_changed_path_sets_must_match(tmp_path):
    rc, rep = compare(make_case(tmp_path, a2={"out/only_a2.txt": b"x"}))
    assert rc == 1 and failed(rep) == ["G0"]
    assert "noise floor" in detail(rep, "G0") and "out/only_a2.txt" in detail(rep, "G0")


def test_the_same_directory_twice_is_not_a_comparison(tmp_path):
    case = make_case(tmp_path)
    report = tmp_path / "r.json"
    rc = cp.main(["compare", "--frozen", str(case.manifest), "--a1", str(case.roots["A1"]), "--a2", str(case.roots["A1"]),
                  "--b", str(case.roots["B"]), "--report", str(report)])
    assert rc == 1 and "not distinct" in detail(json.loads(report.read_text("utf-8")), "G0")


def test_usage_errors_return_2(tmp_path, capsys):
    case = make_case(tmp_path)
    assert cp.main([]) == 2
    assert cp.main(["compare", "--frozen", str(tmp_path / "nope.json"), "--a1", "x", "--a2", "y", "--b", "z"]) == 2
    assert cp.main(["compare", "--frozen", str(case.manifest), "--a1", str(tmp_path / "missing"),
                    "--a2", str(case.roots["A2"]), "--b", str(case.roots["B"])]) == 2
    assert cp.main(["compare", "--frozen", str(case.manifest), "--a1", str(case.roots["A1"]), "--a2", str(case.roots["A2"]),
                    "--b", str(case.roots["B"]), "--map", "no-equals-sign"]) == 2
    assert "error:" in capsys.readouterr().err


# --------------------------------------------------------------------------- G4 json / mask semantics

def js(ts: str, x: int = 1) -> bytes:
    return json.dumps({"built_at": ts, "x": x, "rows": [{"t": ts, "n": 1}]}, indent=1).encode("ascii")


def test_timestamp_difference_that_also_varies_between_a1_and_a2_is_json_canonical(tmp_path):
    rc, rep = compare(make_case(tmp_path, a1={"out/state.json": js("t1")}, a2={"out/state.json": js("t2")},
                                b={"out/state.json": js("t3")}))
    assert rc == 0 and rep["classes"]["B"]["out/state.json"]["class"] == "JSON_CANONICAL"
    assert rep["masks"]["out/state.json"] == ["/built_at", "/rows/0/t"]


def test_the_same_difference_only_between_a1_and_b_is_different(tmp_path):
    rc, rep = compare(make_case(tmp_path, a1={"out/state.json": js("t1")}, a2={"out/state.json": js("t1")},
                                b={"out/state.json": js("t2")}))
    assert rc == 1 and failed(rep) == ["G4"]
    assert rep["classes"]["B"]["out/state.json"]["class"] == "DIFFERENT"
    assert "outside the A1/A2 mask" in rep["classes"]["B"]["out/state.json"]["detail"] and "out/state.json" not in rep["masks"]


def test_the_mask_is_never_widened_by_what_b_contains(tmp_path):
    rc, rep = compare(make_case(tmp_path, a1={"out/state.json": js("t1")}, a2={"out/state.json": js("t2")},
                                b={"out/state.json": js("t3", x=2)}))
    assert rc == 1 and rep["classes"]["B"]["out/state.json"]["class"] == "DIFFERENT"
    assert "/x" in rep["classes"]["B"]["out/state.json"]["detail"]


def test_structural_differences_are_never_masked(tmp_path):
    arr = lambda *v: json.dumps({"k": list(v)}).encode()  # noqa: E731
    rc, rep = compare(make_case(tmp_path, a1={"out/s.json": arr(1, 2)}, a2={"out/s.json": arr(1, 2, 3)},
                                b={"out/s.json": arr(1, 2, 4)}))
    assert rc == 1 and rep["classes"]["B"]["out/s.json"]["class"] == "DIFFERENT"


def test_json_comparison_is_type_exact(tmp_path):
    rc, rep = compare(make_case(tmp_path, a1={"out/n.json": b'{"n":1}'}, a2={"out/n.json": b'{"n":1}'},
                                b={"out/n.json": b'{"n":1.0}'}))
    assert rc == 1 and rep["classes"]["B"]["out/n.json"]["class"] == "DIFFERENT"


def test_ndjson_timestamp_noise_is_masked_and_key_order_alone_is_canonical(tmp_path):
    nd = lambda t: ("\n".join(json.dumps({"ts": t, "id": i}) for i in range(3)) + "\n").encode()  # noqa: E731
    reordered = b'{"id":0,"ts":"t1"}\n{"id":1,"ts":"t1"}\n{"id":2,"ts":"t1"}\n'
    rc, rep = compare(make_case(tmp_path, a1={"out/ev.ndjson": nd("t1"), "out/k.json": b'{"a":1,"b":2}'},
                                a2={"out/ev.ndjson": nd("t2"), "out/k.json": b'{"a":1,"b":2}'},
                                b={"out/ev.ndjson": nd("t3"), "out/k.json": b'{"b":2,"a":1}'}))
    assert rc == 0 and {r["class"] for r in rep["classes"]["B"].values()} == {"BYTE_EQUAL", "JSON_CANONICAL"}
    rc, rep = compare(make_case(tmp_path / "re", a1={"out/ev.ndjson": nd("t1")}, a2={"out/ev.ndjson": nd("t1")},
                                b={"out/ev.ndjson": reordered}))
    assert rc == 0 and rep["classes"]["B"]["out/ev.ndjson"]["class"] == "JSON_CANONICAL"


# --------------------------------------------------------------------------- G4 eol / zip / path-only

def test_crlf_only_difference_is_eol_only_and_needs_acceptance(tmp_path):
    case = make_case(tmp_path, a1={"out/notes.txt": b"a\nb\n"}, a2={"out/notes.txt": b"a\nb\n"},
                     b={"out/notes.txt": b"a\r\nb\r\n"})
    rc, rep = compare(case)
    assert rc == 3 and rep["result"] == "NEEDS_ACCEPTANCE" and failed(rep) == []
    assert rep["classes"]["B"]["out/notes.txt"]["class"] == "EOL_ONLY"
    assert rep["needs_acceptance"]["pending"] == ["out/notes.txt"]
    rc, rep = compare(case, "--accept", "out/notes.txt")
    assert rc == 0 and rep["result"] == "PASS" and rep["needs_acceptance"]["accepted"] == ["out/notes.txt"]
    rc, rep = compare(case, "--accept", "out\\notes.txt", "out/never-needed.txt")
    assert rc == 0 and rep["needs_acceptance"]["unused_accepts"] == ["out/never-needed.txt"]


def test_acceptance_never_rescues_a_failed_gate(tmp_path):
    rc, rep = compare(make_case(tmp_path, a1={"out/notes.txt": b"a\n"}, a2={"out/notes.txt": b"a\n"},
                                b={"out/notes.txt": b"a\r\n", "dashboard_data.json": b"{}"}), "--accept", "out/notes.txt")
    assert rc == 1 and rep["result"] == "FAIL" and "G2" in failed(rep)


def make_zip(entries: dict, date_time=(2020, 1, 1, 0, 0, 0), compression=zipfile.ZIP_DEFLATED) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, data in entries.items():
            info = zipfile.ZipInfo(name, date_time)
            info.compress_type = compression
            z.writestr(info, data)
    return buf.getvalue()


def test_zip_with_other_mtimes_and_compression_is_zip_content(tmp_path):
    payload = {"m/a.txt": b"alpha" * 50, "m/b.bin": bytes(range(256))}
    one = make_zip(payload)
    two = make_zip(payload, date_time=(2024, 6, 7, 8, 9, 10), compression=zipfile.ZIP_STORED)
    assert one != two
    case = make_case(tmp_path, a1={"out/bundle.zip": one}, a2={"out/bundle.zip": one}, b={"out/bundle.zip": two})
    rc, rep = compare(case)
    assert rc == 3 and rep["classes"]["B"]["out/bundle.zip"]["class"] == "ZIP_CONTENT"
    assert compare(case, "--accept", "out/bundle.zip")[0] == 0


def test_zip_with_a_different_payload_or_member_set_is_different(tmp_path):
    one = make_zip({"m/a.txt": b"alpha"})
    for i, other in enumerate((make_zip({"m/a.txt": b"alphA"}), make_zip({"m/a.txt": b"alpha", "m/extra.txt": b""}))):
        case = make_case(tmp_path / f"case{i}", a1={"out/bundle.zip": one},
                         a2={"out/bundle.zip": one}, b={"out/bundle.zip": other})
        rc, rep = compare(case)
        assert rc == 1 and rep["classes"]["B"]["out/bundle.zip"]["class"] == "DIFFERENT"


WIN = b"root=C:\\melange-data\\audio\\a.json\nbase=C:\\melange-data\n"
BOX = b"root=/melange-data/audio/a.json\nbase=/melange-data\n"
MAP = ["--map", "C:\\melange-data=/melange-data"]


def test_path_only_needs_the_map_and_then_acceptance(tmp_path):
    case = make_case(tmp_path, a1={"out/paths.txt": WIN}, a2={"out/paths.txt": WIN}, b={"out/paths.txt": BOX})
    rc, rep = compare(case)  # no --map: nothing can explain the difference
    assert rc == 1 and rep["classes"]["B"]["out/paths.txt"]["class"] == "DIFFERENT"
    rc, rep = compare(case, *MAP)
    assert rc == 3 and rep["classes"]["B"]["out/paths.txt"]["class"] == "PATH_ONLY"
    assert compare(case, *MAP, "--accept", "out/paths.txt")[0] == 0


def test_path_only_handles_json_escaped_backslashes_and_wins_over_json_canonical(tmp_path):
    win = b'{"p":"C:\\\\melange-data\\\\audio\\\\a.json","q":1}'
    box = b'{"p":"/melange-data/audio/a.json","q":1}'
    rc, rep = compare(make_case(tmp_path, a1={"out/p.json": win}, a2={"out/p.json": win}, b={"out/p.json": box}), *MAP)
    assert rc == 3 and rep["classes"]["B"]["out/p.json"]["class"] == "PATH_ONLY"


@pytest.mark.parametrize("box", [
    b"root=/melange-data/audio/b.json\nbase=/melange-data\n",       # a different file name behind the same prefix
    b"root=/melange-data-old/audio/a.json\nbase=/melange-data\n",   # a longer lookalike prefix
    b"root=/opt/melange-data/audio/a.json\nbase=/melange-data\n",   # the prefix in the middle of another path
])
def test_path_only_does_not_hide_real_differences(tmp_path, box):
    rc, rep = compare(make_case(tmp_path, a1={"out/paths.txt": WIN}, a2={"out/paths.txt": WIN},
                                b={"out/paths.txt": box}), *MAP)
    assert rc == 1 and rep["classes"]["B"]["out/paths.txt"]["class"] == "DIFFERENT"


def test_exclude_keeps_bytecode_out_of_the_census(tmp_path):
    case = make_case(tmp_path, a1={"pkg/__pycache__/m.pyc": b"1"}, a2={"pkg/__pycache__/m.pyc": b"1"},
                     b={"pkg/__pycache__/m.pyc": b"2"})
    assert compare(case)[0] == 1
    assert compare(case, "--exclude", "__pycache__")[0] == 0


# ------------------------------------------------------------------------------------------ workbooks

def build_book(path: Path, *, stamp=STAMP_OLD, cells=None, merged=None, name=None, table=None, dv=0,
               sheets=("Data", "Other")) -> bytes:
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for s in sheets:
        wb.create_sheet(s)
    ws = wb["Data"]
    for coord, v in (cells or {"A1": 1, "A2": 2, "A3": "=A1+A2"}).items():
        ws[coord] = v
    if merged:
        ws.merge_cells(merged)
    if name:
        wb.defined_names["Total"] = DefinedName("Total", attr_text=name)
    if table:
        ws["F1"], ws["F2"], ws["F3"] = "h", 1, 2
        ws.add_table(Table(displayName="T1", ref=table))
    for i in range(dv):  # a validation with no cell range is dropped on save, so give it one
        rule = DataValidation(type="whole", operator="between", formula1="0", formula2="9")
        rule.add(f"H{i + 1}")
        ws.add_data_validation(rule)
    wb.properties.created = wb.properties.modified = stamp
    wb.save(path)
    return path.read_bytes()


def with_cached(data: bytes, value) -> bytes:
    """Give the formula cell a cached <v> (openpyxl writes none), as Excel would."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                body, n = re.subn(rb"<v\s*/>|<v></v>", f"<v>{value}</v>".encode(), body)     # openpyxl spells it <v></v> with lxml, <v /> without
                assert n
            dst.writestr(info, body)
    return out.getvalue()


def test_workbook_differing_only_in_docprops_is_ooxml_canonical(tmp_path):
    old, new = build_book(tmp_path / "o.xlsx"), build_book(tmp_path / "n.xlsx", stamp=STAMP_NEW)
    assert old != new
    rc, rep = compare(make_case(tmp_path / "c", a1={"out/w.xlsx": old}, a2={"out/w.xlsx": old}, b={"out/w.xlsx": new}))
    assert rc == 0 and rep["classes"]["B"]["out/w.xlsx"]["class"] == "OOXML_CANONICAL"


def test_workbook_with_a_changed_formula_is_different(tmp_path):
    old = build_book(tmp_path / "o.xlsx")
    new = build_book(tmp_path / "n.xlsx", cells={"A1": 1, "A2": 2, "A3": "=A1-A2"})
    rc, rep = compare(make_case(tmp_path / "c", a1={"out/w.xlsx": old}, a2={"out/w.xlsx": old}, b={"out/w.xlsx": new}))
    assert rc == 1 and rep["classes"]["B"]["out/w.xlsx"]["class"] == "DIFFERENT"
    assert "Data!A3" in rep["classes"]["B"]["out/w.xlsx"]["detail"]


def test_workbook_with_a_changed_cached_value_is_different_and_equal_caches_are_canonical(tmp_path):
    base = build_book(tmp_path / "o.xlsx")
    three, four = with_cached(base, 3), with_cached(base, 4)
    rc, rep = compare(make_case(tmp_path / "bad", a1={"out/w.xlsx": three}, a2={"out/w.xlsx": three}, b={"out/w.xlsx": four}))
    assert rc == 1 and rep["classes"]["B"]["out/w.xlsx"]["class"] == "DIFFERENT"
    assert cp.compare_workbooks(io.BytesIO(three), io.BytesIO(four))["cell_diffs"][0]["mode"] == "value"
    again = with_cached(build_book(tmp_path / "n.xlsx", stamp=STAMP_NEW), 3)
    rc, rep = compare(make_case(tmp_path / "ok", a1={"out/w.xlsx": three}, a2={"out/w.xlsx": three}, b={"out/w.xlsx": again}))
    assert rc == 0 and rep["classes"]["B"]["out/w.xlsx"]["class"] == "OOXML_CANONICAL"


def test_compare_workbooks_lists_capped_cells_with_sheet_and_coordinate(tmp_path):
    a = build_book(tmp_path / "a.xlsx", cells={f"B{i}": i for i in range(1, 31)})
    b = build_book(tmp_path / "b.xlsx", cells={f"B{i}": i + 1 for i in range(1, 31)})
    res = cp.compare_workbooks(tmp_path / "a.xlsx", tmp_path / "b.xlsx", max_cells=5)
    assert (a != b) and res["equal"] is False and res["cell_diff_count"] == 60  # 30 cells in each of two load modes
    assert len(res["cell_diffs"]) == 5 and all({"sheet", "cell", "mode", "a", "b"} <= set(d) for d in res["cell_diffs"])
    assert res["cell_diffs"][0]["sheet"] == "Data" and res["cell_diffs"][0]["cell"] == "B1"
    assert cp.compare_workbooks(tmp_path / "a.xlsx", tmp_path / "a.xlsx")["equal"] is True


@pytest.mark.parametrize("label, kwargs", [
    ("sheet names (in order)", {"sheets": ("Other", "Data")}),
    ("merged ranges", {"merged": "D1:E1"}),
    ("defined names", {"name": "Data!$A$1"}),
    ("tables", {"table": "F1:F3"}),
    ("data validation count", {"dv": 1}),
])
def test_compare_workbooks_flags_each_structural_dimension(tmp_path, label, kwargs):
    build_book(tmp_path / "a.xlsx")
    build_book(tmp_path / "b.xlsx", **kwargs)
    res = cp.compare_workbooks(tmp_path / "a.xlsx", tmp_path / "b.xlsx")
    assert res["equal"] is False and any(label in line for line in res["differences"]), res


# ------------------------------------------------------------------------------- freeze / verify / census

def make_src(tmp_path: Path) -> Path:
    src = tmp_path / "src"
    write_tree(src, {"fec/a.py": b"print(1)\n", "fec/sub/b.py": b"x = 2\n", "fec/__pycache__/a.cpython-312.pyc": b"\0\1",
                     "data/big.csv": b"1,2\n", "data/skip.tmp": b"t", "ranks_python.json": b"{}", "notes/readme.md": b"hi"})
    for p in src.rglob("*"):
        if p.is_file():
            os.utime(p, (1_000_000_000, 1_000_000_000))
    return src


FROZEN_FILES = {"fec/a.py", "fec/sub/b.py", "data/big.csv"}


def freeze_args(src, dest, *more):
    return ["freeze", "--src", str(src), "--dest", str(dest), "--include", "fec", "data/*", "--exclude", "__pycache__", "*.tmp", *more]


def test_freeze_verify_round_trip(tmp_path, capsys):
    src, dest = make_src(tmp_path), tmp_path / "dest"
    assert cp.main(freeze_args(src, dest)) == 0
    assert {p.relative_to(dest).as_posix() for p in dest.rglob("*") if p.is_file()} == FROZEN_FILES
    manifest = tmp_path / "dest.manifest.json"
    man = json.loads(manifest.read_text(encoding="utf-8"))
    assert set(man) == FROZEN_FILES and man["fec/a.py"] == {"sha256": hashlib.sha256(b"print(1)\n").hexdigest(), "size": 9}
    assert list(man) == sorted(man) and manifest.read_bytes().endswith(b"\n") and b"\r" not in manifest.read_bytes()
    assert all(abs((dest / rel).stat().st_mtime - 1_000_000_000) < 2 for rel in FROZEN_FILES)  # copy2 kept mtimes
    capsys.readouterr()
    assert cp.main(["verify", "--root", str(dest), "--manifest", str(manifest)]) == 0
    assert "byte-identical" in capsys.readouterr().out


def test_verify_catches_modified_extra_and_missing_files(tmp_path, capsys):
    src, dest = make_src(tmp_path), tmp_path / "dest"
    assert cp.main(freeze_args(src, dest)) == 0
    verify = ["verify", "--root", str(dest), "--manifest", str(tmp_path / "dest.manifest.json")]
    (dest / "fec/a.py").write_bytes(b"print(2)\n")
    assert cp.main(verify) == 1 and "modified: fec/a.py" in capsys.readouterr().out
    (dest / "fec/a.py").write_bytes(b"print(1)\n")
    (dest / "stray.txt").write_bytes(b"x")
    assert cp.main(verify) == 1 and "extra: stray.txt" in capsys.readouterr().out
    (dest / "stray.txt").unlink()
    (dest / "data/big.csv").unlink()
    assert cp.main(verify) == 1 and "missing: data/big.csv" in capsys.readouterr().out


def test_freeze_refuses_a_non_empty_dest_and_an_existing_manifest(tmp_path, capsys):
    src = make_src(tmp_path)
    dest = tmp_path / "dest"
    write_tree(dest, {"old.txt": b"keep"})
    assert cp.main(freeze_args(src, dest)) == 2 and "not empty" in capsys.readouterr().err
    assert [p.name for p in dest.iterdir()] == ["old.txt"] and not (tmp_path / "dest.manifest.json").exists()
    fresh = tmp_path / "fresh"
    write_tree(tmp_path, {"fresh.manifest.json": b"{}"})
    assert cp.main(freeze_args(src, fresh)) == 2 and not fresh.exists()
    assert cp.main(freeze_args(src, src / "inside")) == 2  # dest inside src
    assert cp.main(["freeze", "--src", str(src), "--dest", str(tmp_path / "e"), "--include", "nothing-matches"]) == 2


def test_freeze_from_a_list_file(tmp_path, capsys):
    src, dest = make_src(tmp_path), tmp_path / "dest"
    listing = tmp_path / "read_set.txt"
    listing.write_text("# traced read set\n\nfec/a.py\nnotes\\readme.md\n", encoding="utf-8")
    assert cp.main(["freeze", "--src", str(src), "--dest", str(dest), "--list", str(listing)]) == 0
    assert set(json.loads((tmp_path / "dest.manifest.json").read_text("utf-8"))) == {"fec/a.py", "notes/readme.md"}
    for bad in ("fec/missing.py", "../outside.txt", "C:/abs.txt"):
        listing.write_text(bad + "\n", encoding="utf-8")
        assert cp.main(["freeze", "--src", str(src), "--dest", str(tmp_path / "x"), "--list", str(listing),
                        "--manifest", str(tmp_path / "x.json")]) == 2, bad


def test_snapshot_and_census(tmp_path, capsys):
    root = tmp_path / "t"
    write_tree(root, {"keep.txt": b"k", "gone.txt": b"g", "chg.txt": b"1", "skip.log": b"s"})
    before, after = tmp_path / "m1.json", tmp_path / "m2.json"
    assert cp.main(["snapshot", "--root", str(root), "--out", str(before), "--exclude", "*.log"]) == 0
    assert set(json.loads(before.read_text("utf-8"))) == {"keep.txt", "gone.txt", "chg.txt"}
    (root / "gone.txt").unlink()
    (root / "chg.txt").write_bytes(b"2")
    (root / "new.txt").write_bytes(b"n")
    assert cp.main(["snapshot", "--root", str(root), "--out", str(after), "--exclude", "*.log"]) == 0
    capsys.readouterr()
    assert cp.main(["census", "--before", str(before), "--after", str(after)]) == 0
    out = capsys.readouterr().out
    assert "+ new.txt" in out and "- gone.txt" in out and "~ chg.txt" in out and "added 1, removed 1, changed 1" in out


# ------------------------------------------------------------------------------------------- command

MANIFEST_IMAGE = "ghcr.io/pjg93-222/compute-melange@sha256:" + "ab" * 32


@pytest.fixture
def clean_env(monkeypatch):
    for name, _ in cp.PASSTHROUGH:
        monkeypatch.delenv(name, raising=False)
    monkeypatch.setattr(cp, "current_image", lambda name=cp.IMAGE_NAME: MANIFEST_IMAGE)      # the verified current manifest
    monkeypatch.setattr(cp.subprocess, "run", lambda *a, **k: pytest.fail("printing a command must never execute it"))


def command_lines(capsys, *argv) -> dict:
    assert cp.main(["command", *argv]) == 0
    return dict(line.split("=", 1) for line in capsys.readouterr().out.splitlines() if "=" in line and not line.startswith("#"))


def test_command_prints_the_same_today_on_both_sides(tmp_path, capsys, clean_env):
    win = command_lines(capsys, "--side", "windows", "--root", str(tmp_path), "--today", "2026-10-12", "--local-dir", str(tmp_path / "Melange"), "--glassbox-dir", str(tmp_path / "g"))
    box = command_lines(capsys, "--side", "container", "--root", str(tmp_path), "--today", "2026-10-12")
    assert win["MELANGE_TODAY"] == box["MELANGE_TODAY"] == "2026-10-12"
    assert win["MONEYBALL_ROOT"] == os.path.abspath(tmp_path) and "MELANGE_TREE_DIR" not in win
    assert box["MELANGE_TREE_DIR"] == os.path.abspath(tmp_path) and "MONEYBALL_ROOT" not in box
    assert box["IMAGE_COMPUTE_MELANGE"] == MANIFEST_IMAGE and "IMAGE_COMPUTE_MELANGE" not in win       # no flag, no env: the manifest's
    assert os.path.join("fec", "run_pipeline.py") in win["run"]
    assert os.path.basename(sys.executable) in win["run"]
    assert "docker compose -f ops/compute/melange.compose.yaml --profile chain run --rm chain" in box["run"]


def test_command_container_passthrough_from_flags_and_environment(tmp_path, capsys, clean_env, monkeypatch):
    monkeypatch.setenv("MELANGE_DATA_DIR", "/from/env")
    monkeypatch.setenv("MELANGE_DRAFT_DIR", "/env/draft")
    box = command_lines(capsys, "--side", "container", "--root", str(tmp_path), "--today", "2026-10-12",
                        "--draft-dir", "/from/flag", "--image", "compute-melange:test")
    assert box["MELANGE_DRAFT_DIR"] == "/from/flag" and box["MELANGE_DATA_DIR"] == "/from/env"
    assert box["IMAGE_COMPUTE_MELANGE"] == "compute-melange:test"
    assert "MELANGE_LOCAL_DIR" not in box and "MELANGE_GLASSBOX_DIR" not in box
    win = command_lines(capsys, "--side", "windows", "--root", str(tmp_path), "--today", "2026-10-12", "--local-dir", str(tmp_path / "Melange"), "--glassbox-dir", str(tmp_path / "g"))
    assert "MELANGE_DATA_DIR" not in win


def test_command_container_image_from_the_environment_beats_the_manifest(tmp_path, capsys, clean_env, monkeypatch):
    monkeypatch.setattr(cp, "current_image", lambda name=cp.IMAGE_NAME: pytest.fail("an explicit image never consults the manifest"))
    monkeypatch.setenv("IMAGE_COMPUTE_MELANGE", "ghcr.io/pjg93-222/compute-melange@sha256:" + "cd" * 32)
    box = command_lines(capsys, "--side", "container", "--root", str(tmp_path), "--today", "2026-10-12")
    assert box["IMAGE_COMPUTE_MELANGE"].endswith("cd" * 32)


def test_command_container_without_a_verified_manifest_or_an_explicit_image_refuses_and_prints_nothing(tmp_path, capsys, clean_env, monkeypatch):
    def refuse(name=cp.IMAGE_NAME):
        raise cp.Usage("no verified image for compute-melange: not running")
    monkeypatch.setattr(cp, "current_image", refuse)
    assert cp.main(["command", "--side", "container", "--root", str(tmp_path), "--today", "2026-10-12"]) == 2
    captured = capsys.readouterr()
    assert captured.out == "" and "no verified image" in captured.err
    assert cp.main(["command", "--side", "windows", "--root", str(tmp_path), "--today", "2026-10-12", "--local-dir", str(tmp_path / "Melange"), "--glassbox-dir", str(tmp_path / "g")]) == 0     # the Windows side needs no image


@pytest.mark.parametrize("returncode, stdout", [(1, ""), (1, MANIFEST_IMAGE + "\n"), (0, "\n")])
def test_current_image_fails_closed_on_a_nonzero_exit_or_no_output(monkeypatch, returncode, stdout):
    monkeypatch.setattr(cp.subprocess, "run", lambda *a, **k: SimpleNamespace(returncode=returncode, stdout=stdout, stderr="current: NO_CURRENT"))
    with pytest.raises(cp.Usage, match="no verified image for compute-melange"):
        cp.current_image()


def test_current_image_asks_ops_release_lib_current_for_the_image_by_name(monkeypatch):
    seen = {}

    def fake_run(argv, **kwargs):
        seen["argv"] = argv
        return SimpleNamespace(returncode=0, stdout=MANIFEST_IMAGE + "\n", stderr="")

    monkeypatch.setattr(cp.subprocess, "run", fake_run)
    assert cp.current_image() == MANIFEST_IMAGE
    assert seen["argv"][0] == sys.executable and seen["argv"][-2:] == ["image", "compute-melange"]
    assert Path(seen["argv"][1]).as_posix().endswith("ops/release_lib/current.py")


@pytest.mark.parametrize("side", ["windows", "container"])
def test_command_refuses_without_a_valid_today(tmp_path, capsys, clean_env, side):
    assert cp.main(["command", "--side", side, "--root", str(tmp_path)]) == 2
    captured = capsys.readouterr()
    assert captured.out == "" and "--today" in captured.err and "MELANGE_TODAY" in captured.err
    assert cp.main(["command", "--side", side, "--root", str(tmp_path), "--today", "2026-13-40"]) == 2
    assert cp.main(["command", "--side", side, "--root", str(tmp_path), "--today", "10/12/2026"]) == 2
    assert capsys.readouterr().out == ""


def test_command_execute_plumbing_with_a_fake_runner(tmp_path, monkeypatch, capsys):
    seen = {}

    def fake_run(argv, cwd, env, timeout):
        seen.update(argv=argv, cwd=cwd, env=env, timeout=timeout)
        return SimpleNamespace(returncode=7)

    monkeypatch.setattr(cp.subprocess, "run", fake_run)
    rc = cp.main(["command", "--side", "windows", "--root", str(tmp_path), "--today", "2026-10-12", "--execute", "--timeout", "30", "--local-dir", str(tmp_path / "Melange"), "--glassbox-dir", str(tmp_path / "g")])
    assert rc == 7 and seen["timeout"] == 30 and seen["cwd"] == os.path.abspath(tmp_path)
    assert seen["env"]["MONEYBALL_ROOT"] == os.path.abspath(tmp_path) and seen["env"]["MELANGE_TODAY"] == "2026-10-12"
    assert seen["argv"][0] == sys.executable and "running it now" in capsys.readouterr().out


# ----------------------------------------------------------------------------- relocated, clean-environment runs (P7)

def test_windows_side_refuses_without_a_local_dir_and_prints_nothing(tmp_path, capsys, clean_env):
    assert cp.main(["command", "--side", "windows", "--root", str(tmp_path), "--today", "2026-10-12"]) == 2
    captured = capsys.readouterr()
    assert captured.out == "" and "--local-dir" in captured.err


def test_windows_side_refuses_without_a_glassbox_dir_and_prints_nothing(tmp_path, capsys, clean_env):
    assert cp.main(["command", "--side", "windows", "--root", str(tmp_path), "--today", "2026-10-12", "--local-dir", str(tmp_path / "Melange")]) == 2
    captured = capsys.readouterr()
    assert captured.out == "" and "--glassbox-dir" in captured.err


def test_windows_side_refuses_a_local_dir_not_named_melange_or_the_live_store(tmp_path, capsys, clean_env, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "live"))
    for local in (tmp_path / "loc", tmp_path / "live" / "Melange"):   # a sibling LOCALAPPDATA\Melange; the live store itself
        assert cp.main(["command", "--side", "windows", "--root", str(tmp_path), "--today", "2026-10-12", "--local-dir", str(local)]) == 2
        assert "named Melange" in capsys.readouterr().err


def test_from_step_warns_that_a_resumed_run_is_no_parity_proof(tmp_path, capsys, clean_env):
    assert cp.main(["command", "--side", "container", "--root", str(tmp_path), "--today", "2026-10-12", "--from", "x.py"]) == 0
    assert "DE-045" in capsys.readouterr().err
    assert cp.main(["command", "--side", "container", "--root", str(tmp_path), "--today", "2026-10-12"]) == 0
    assert "DE-045" not in capsys.readouterr().err


def test_windows_side_redirects_every_root_and_rebases_stored_paths(tmp_path, capsys, clean_env):
    win = command_lines(capsys, "--side", "windows", "--root", str(tmp_path / "t"), "--today", "2026-10-12",
                        "--local-dir", str(tmp_path / "l" / "Melange"), "--data-dir", str(tmp_path / "d"),
                        "--draft-dir", str(tmp_path / "dr"), "--glassbox-dir", str(tmp_path / "g"))
    assert win["MELANGE_REBASE_STORED"] == "1" and win["PYTHONUTF8"] == "1" and win["MELANGE_XLSX_WRITER"] == "openpyxl"
    assert win["MELANGE_LOCAL_ROOT"] == os.path.abspath(tmp_path / "l" / "Melange") and win["LOCALAPPDATA"] == os.path.abspath(tmp_path / "l")
    assert win["MELANGE_DATA_ROOT"] == os.path.abspath(tmp_path / "d") and win["MONEYBALL_DRAFT"] == os.path.abspath(tmp_path / "dr")
    assert win["GLASSBOX_EVENTS_DIR"] == os.path.abspath(tmp_path / "g")


def test_windows_side_executes_with_only_the_allowlisted_environment(tmp_path, monkeypatch):
    seen = {}
    monkeypatch.setattr(cp.subprocess, "run", lambda argv, cwd, env, timeout: seen.update(env=env) or SimpleNamespace(returncode=0))
    for name, value in (("COSPONSOR_FORECAST_REFRESH", "1"), ("ANTHROPIC_API_KEY", "secret"), ("MELANGE_LOCAL_ROOT", "C:/live"), ("PATH", "/bin")):
        monkeypatch.setenv(name, value)
    assert cp.main(["command", "--side", "windows", "--root", str(tmp_path), "--today", "2026-10-12", "--execute",
                    "--local-dir", str(tmp_path / "Melange"), "--glassbox-dir", str(tmp_path / "g")]) == 0
    env = seen["env"]
    assert "COSPONSOR_FORECAST_REFRESH" not in env and "ANTHROPIC_API_KEY" not in env
    assert env["MELANGE_LOCAL_ROOT"] == os.path.abspath(tmp_path / "Melange") and env["PATH"] == "/bin"
    assert all(k in cp.ENV_ALLOW or k.startswith(("MELANGE_", "MONEYBALL_", "PYTHON")) or k in ("LOCALAPPDATA", "GLASSBOX_EVENTS_DIR") for k in env), sorted(env)


def test_from_step_is_appended_for_both_sides(tmp_path, capsys, clean_env):
    win = command_lines(capsys, "--side", "windows", "--root", str(tmp_path), "--today", "2026-10-12", "--local-dir", str(tmp_path / "Melange"),
                        "--glassbox-dir", str(tmp_path / "g"), "--from", "build_vote_category")
    assert win["run"].endswith("run_pipeline.py --from build_vote_category")
    box = command_lines(capsys, "--side", "container", "--root", str(tmp_path), "--today", "2026-10-12", "--from", "build_vote_category")
    assert box["run"].endswith("run --rm chain python /melange/fec/run_pipeline.py --from build_vote_category")
    plain = command_lines(capsys, "--side", "container", "--root", str(tmp_path), "--today", "2026-10-12")
    assert plain["run"].endswith("run --rm chain")


def test_audio_pilot_dir_adds_the_parity_override_and_its_mount_point(tmp_path, capsys, clean_env, monkeypatch):
    pilot = tmp_path / "pilot"
    box = command_lines(capsys, "--side", "container", "--root", str(tmp_path / "t"), "--today", "2026-10-12", "--audio-pilot-dir", str(pilot))
    assert box["MELANGE_AUDIO_PILOT_DIR"] == str(pilot)
    assert "-f ops/compute/melange.compose.yaml -f ops/compute/melange.parity.compose.yaml --profile chain" in box["run"]
    plain = command_lines(capsys, "--side", "container", "--root", str(tmp_path / "t"), "--today", "2026-10-12")
    assert "MELANGE_AUDIO_PILOT_DIR" not in plain and "parity.compose" not in plain["run"]
    ran = {}
    monkeypatch.setattr(cp.subprocess, "run", lambda argv, cwd, env, timeout: ran.update(argv=argv) or SimpleNamespace(returncode=0))
    (tmp_path / "t").mkdir()
    assert cp.main(["command", "--side", "container", "--root", str(tmp_path / "t"), "--today", "2026-10-12", "--execute",
                    "--audio-pilot-dir", str(pilot)]) == 0
    assert (tmp_path / "t" / "audio_pilot").is_dir() and "ops/compute/melange.parity.compose.yaml" in ran["argv"]


def test_walk_files_does_not_enter_a_directory_link(tmp_path):
    import subprocess
    target = tmp_path / "live" / "audio_pilot"
    write_tree(target, {"a.wav": b"x" * 10})
    root = tmp_path / "run"
    write_tree(root, {"keep.txt": b"k"})
    link = root / "audio_pilot"
    try:
        os.symlink(target, link, target_is_directory=True)
    except (OSError, NotImplementedError):
        if not (os.name == "nt" and subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(target)], capture_output=True).returncode == 0):
            pytest.skip("cannot create a symlink or junction here")
    assert os.path.exists(link / "a.wav")                       # the link works ...
    assert cp.walk_files(root) == ["keep.txt"]                  # ... and is still not entered
    assert list(cp.snapshot_tree(root)) == ["keep.txt"]


def test_at_file_arguments_expand_with_comments_and_blank_lines(tmp_path):
    src = make_src(tmp_path)
    args = tmp_path / "freeze.args"
    args.write_text("# read set\n--include\n\nfec/a.py\n# a comment\nnotes/readme.md\n", encoding="utf-8")
    assert cp.main(["freeze", "--src", str(src), "--dest", str(tmp_path / "dest"), "@" + str(args)]) == 0
    assert set(json.loads((tmp_path / "dest.manifest.json").read_text("utf-8"))) == {"fec/a.py", "notes/readme.md"}


def test_the_shipped_local_read_set_names_the_two_directories_the_traced_set_missed(tmp_path):
    shipped = Path(cp.__file__).with_name("parity_readset_local_extra.args")
    local = tmp_path / "local"
    write_tree(local, {"staffer_network/normalized/lda-source-native/generation-20260830T231200Z/a.bin": b"a",
                       "staffer_network/normalized/lda-source-native/generation-OTHER/b.bin": b"b",
                       "demographics/lodes/x/y.csv": b"y", "demographics/other.csv": b"o", "elsewhere.txt": b"e"})
    assert cp.main(["freeze", "--src", str(local), "--dest", str(tmp_path / "dest"), "@" + str(shipped)]) == 0
    assert set(json.loads((tmp_path / "dest.manifest.json").read_text("utf-8"))) == {
        "staffer_network/normalized/lda-source-native/generation-20260830T231200Z/a.bin", "demographics/lodes/x/y.csv"}


def test_compare_store_writes_one_snapshot_per_root(tmp_path):
    case = make_case(tmp_path)
    rc, _ = compare(case, "--store", str(tmp_path / "store"))
    assert rc == 0
    assert sorted(p.name for p in (tmp_path / "store").iterdir()) == ["A1.manifest.json", "A2.manifest.json", "B.manifest.json"]
    assert "ranks_python.json" in json.loads((tmp_path / "store" / "B.manifest.json").read_text("utf-8"))


def test_workbook_style_only_difference_is_canonical_but_a_cached_value_is_not(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.styles import Font
    plain = build_book(tmp_path / "p.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "p.xlsx")
    wb["Data"]["A1"].font = Font(bold=True)
    wb.save(tmp_path / "s.xlsx")
    styled = (tmp_path / "s.xlsx").read_bytes()
    rc, rep = compare(make_case(tmp_path / "c", a1={"out/w.xlsx": plain}, a2={"out/w.xlsx": plain}, b={"out/w.xlsx": styled}))
    assert rc == 0 and rep["classes"]["B"]["out/w.xlsx"]["class"] == "OOXML_CANONICAL"      # xlsx_content_diff: styles reported, not gated
    rc, rep = compare(make_case(tmp_path / "d", a1={"out/w.xlsx": with_cached(plain, 3)}, a2={"out/w.xlsx": with_cached(plain, 3)},
                                b={"out/w.xlsx": with_cached(plain, 4)}))
    assert rc == 1 and "cached_value" in rep["classes"]["B"]["out/w.xlsx"]["detail"]
