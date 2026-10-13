"""ops.compute.xlsx_content_diff on synthetic workbooks: what counts as the same content, and what does not."""
from __future__ import annotations

import contextlib
import io
import json
import re
import tempfile
import unittest
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

from ops.compute import xlsx_content_diff as X


def build(path, edit=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Amount", "Check"])
    ws.append(["Acme", 10.5, "=B2*2"])
    ws.append(["Beta", 3, "=B3*2"])
    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "hello"
    if edit:
        edit(wb)
    wb.save(path)
    return path


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)

    def pair(self, edit_a=None, edit_b=None, **kw):
        a, b = build(self.tmp / "a.xlsx", edit_a), build(self.tmp / "b.xlsx", edit_b)
        return X.compare_files(str(a), str(b), **kw)

    def verdict(self, result, **kw):
        return X.verdict(result, **kw)


class ContentTests(Base):
    def test_same_content_is_identical_whatever_the_package_metadata(self):
        a = build(self.tmp / "a.xlsx", lambda wb: setattr(wb.properties, "creator", "alice"))
        b = build(self.tmp / "b.xlsx", lambda wb: setattr(wb.properties, "creator", "bob"))
        result = X.compare_files(str(a), str(b))
        self.assertEqual(self.verdict(result), "IDENTICAL")
        self.assertEqual(result["counts"], {})
        self.assertEqual(result["cells_compared"], 3 * 3 + 1)

    def test_value_formula_and_kind_differences_are_told_apart(self):
        def edit(wb):
            ws = wb["Data"]
            ws["A2"] = "Acme Corp"      # value
            ws["C2"] = "=B2*3"          # formula
            ws["B3"] = "3"              # number -> string: kind
        result = self.pair(None, edit)
        self.assertEqual(result["counts"], {"cell_value": 1, "cell_formula": 1, "cell_kind": 1})
        self.assertEqual(self.verdict(result), "DIFF")

    def test_missing_cells_and_sheets_and_order_and_visibility(self):
        def edit(wb):
            wb["Data"]["D1"] = "extra"
            wb.create_sheet("Extra")
            wb.move_sheet("Notes", offset=-1)
            wb["Notes"].sheet_state = "hidden"
        result = self.pair(None, edit)
        counts = result["counts"]
        self.assertEqual(counts["cell_only_in_b"], 1)
        self.assertEqual(counts["sheet_missing"], 1)
        self.assertEqual(counts["sheet_order"], 1)
        self.assertEqual(counts["sheet_state"], 1)

    def test_number_spelling_and_blank_versus_empty_string_do_not_matter(self):
        def one(wb):
            wb["Data"]["B3"] = 3
            wb["Data"]["E1"] = None

        def other(wb):
            wb["Data"]["B3"] = 3.0
            wb["Data"]["E1"] = ""
        self.assertEqual(self.verdict(self.pair(one, other)), "IDENTICAL")

    def test_numbers_are_the_same_only_when_they_are_the_same_double(self):
        def with_value(name, text):
            source, out = build(self.tmp / ("src_" + name)), self.tmp / name
            with ZipFile(source) as src, ZipFile(out, "w", ZIP_DEFLATED) as dst:
                for info in src.infolist():
                    data = src.read(info.filename)
                    if info.filename == "xl/worksheets/sheet1.xml":
                        self.assertIn(b"<v>3</v>", data)
                        data = data.replace(b"<v>3</v>", b"<v>%s</v>" % text.encode())
                    dst.writestr(info, data)
            return str(out)
        exact, again, short = (with_value("a.xlsx", "0.25555555555555554"), with_value("b.xlsx", "0.25555555555555554"),
                               with_value("c.xlsx", "0.2555555555555555"))
        self.assertEqual(X.compare_files(exact, again)["counts"], {})
        self.assertEqual(X.compare_files(exact, short)["counts"], {"cell_value": 1})

    def test_formula_whitespace_and_case_are_normalised_but_structured_column_names_are_not(self):
        def spaced(wb):
            wb["Data"]["C2"] = "= sum( B2 , 1 )"
            wb["Data"]["D2"] = "=Table1[Total Giving]"

        def tight(wb):
            wb["Data"]["C2"] = "=SUM(B2,1)"
            wb["Data"]["D2"] = "=Table1[TotalGiving]"
        result = self.pair(spaced, tight)
        self.assertEqual(result["counts"], {"cell_formula": 1})
        self.assertEqual(result["examples"]["cell_formula"], ["Data!D2"])     # C2 differs only in spacing and case

    def test_quoted_text_and_sheet_names_are_compared_exactly(self):
        def upper(wb):
            wb["Data"]["C2"] = '=IF(A2="OK",1,0)'

        def lower(wb):
            wb["Data"]["C2"] = '=IF(A2="ok",1,0)'
        self.assertEqual(self.pair(upper, lower)["counts"], {"cell_formula": 1})

    def test_defined_names_tables_merges_freeze_conditional_formats(self):
        def edit(wb):
            ws = wb["Data"]
            wb.defined_names["Total"] = DefinedName("Total", attr_text="Data!$B$2:$B$3")
            ws.add_table(Table(displayName="DataTable", ref="A1:C3", tableStyleInfo=TableStyleInfo(name="TableStyleMedium2")))
            ws.merge_cells("E1:F1")
            ws.freeze_panes = "A2"
            ws.conditional_formatting.add("C2:C3", FormulaRule(formula=['$C2<>"OK"'], fill=PatternFill("solid", bgColor="FFF2CC")))
        result = self.pair(None, edit)
        self.assertEqual(result["counts"], {"defined_name": 1, "table": 1, "merge": 1, "freeze_pane": 1,
                                            "conditional_format": 1})
        self.assertEqual(self.verdict(self.pair(edit, edit)), "IDENTICAL")

    def test_table_definitions_compare_name_ref_and_columns(self):
        def table(ref, name):
            return lambda wb: wb["Data"].add_table(Table(displayName=name, ref=ref))
        self.assertEqual(self.verdict(self.pair(table("A1:C3", "T"), table("A1:C3", "T"))), "IDENTICAL")
        self.assertEqual(self.pair(table("A1:C3", "T"), table("A1:C2", "T"))["counts"], {"table": 2})
        self.assertEqual(self.pair(table("A1:C3", "T"), table("A1:C3", "U"))["counts"], {"table": 2})

    def test_a_one_cell_range_has_one_spelling(self):
        self.assertEqual((X.norm_sqref("O6:O6"), X.norm_sqref("O6"), X.norm_sqref("A1:B2 C3:C3")), ("O6", "O6", "A1:B2 C3"))

        def rule(ref):
            return lambda wb: wb["Data"].conditional_formatting.add(
                ref, FormulaRule(formula=['$C2<>"OK"'], fill=PatternFill("solid", bgColor="FFF2CC")))
        self.assertEqual(self.pair(rule("C2"), rule("C2:C2"))["counts"], {})
        self.assertEqual(self.pair(rule("C2"), rule("C2:C3"))["counts"], {"conditional_format": 2})

    def test_conditional_format_rule_content_is_compared(self):
        def rule(text):
            return lambda wb: wb["Data"].conditional_formatting.add(
                "B2:B3", CellIsRule(operator="greaterThan", formula=[text], font=Font(bold=True)))
        self.assertEqual(self.verdict(self.pair(rule("5"), rule("5"))), "IDENTICAL")
        self.assertEqual(self.pair(rule("5"), rule("6"))["counts"], {"conditional_format": 2})


class StyleAndCachedTests(Base):
    def edit_styles(self, wb):
        ws = wb["Data"]
        ws["A1"].font = Font(bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="17365D")
        ws["B2"].number_format = '$#,##0.00'
        ws["A3"].alignment = Alignment(wrap_text=True)
        ws.column_dimensions["A"].width = 30
        ws.row_dimensions[1].height = 30

    def test_style_differences_are_reported_but_gate_only_on_request(self):
        result = self.pair(None, self.edit_styles)
        self.assertEqual(self.verdict(result), "IDENTICAL")
        self.assertEqual(result["counts"], {"cell_style": 3, "column_width": 1, "row_height": 1})
        self.assertEqual(self.verdict(result, gate_styles=True), "DIFF")
        self.assertEqual(self.verdict(self.pair(self.edit_styles, self.edit_styles), gate_styles=True), "IDENTICAL")

    def test_no_styles_skips_the_group(self):
        a, b = build(self.tmp / "a.xlsx"), build(self.tmp / "b.xlsx", self.edit_styles)
        self.assertEqual(X.compare_files(str(a), str(b), styles=False)["counts"], {})

    def test_style_on_an_empty_cell_is_invisible(self):
        def edit(wb):
            wb["Data"]["H9"].fill = PatternFill("solid", fgColor="FF0000")
        self.assertEqual(self.pair(None, edit)["counts"], {})

    def test_unstyled_cell_is_the_workbook_default_and_theme_text_colour_is_no_colour(self):
        def edit(wb):
            wb["Data"]["A2"].font = Font(name="Calibri", size=11)       # the default font, spelled without its theme colour
        self.assertEqual(self.pair(None, edit)["counts"], {})
        def red(wb):
            wb["Data"]["A2"].font = Font(name="Calibri", size=11, color="FF0000")
        self.assertEqual(self.pair(None, red)["counts"], {"cell_style": 1})

    def test_cached_values_are_reported_not_gated(self):
        a = build(self.tmp / "a.xlsx")
        b = build(self.tmp / "b.xlsx")
        patched = self.tmp / "c.xlsx"
        with ZipFile(b) as src, ZipFile(patched, "w", ZIP_DEFLATED) as out:
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"<f>B2*2</f><v></v>", b"<f>B2*2</f><v>21</v>").replace(b"<f>B2*2</f><v />", b"<f>B2*2</f><v>21</v>")
                out.writestr(info, data)
        result = X.compare_files(str(a), str(patched))
        self.assertEqual(result["counts"], {"cached_value": 1})
        self.assertEqual(self.verdict(result), "IDENTICAL")
        self.assertEqual(self.verdict(result, gate_cached=True), "DIFF")


class ShapeTests(Base):
    def test_shared_formulas_expand_to_the_same_content_as_explicit_ones(self):
        explicit = build(self.tmp / "a.xlsx")
        shared = self.tmp / "b.xlsx"
        with ZipFile(explicit) as src, ZipFile(shared, "w", ZIP_DEFLATED) as out:
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    text = data.decode("utf-8")
                    self.assertIn("<f>B3*2</f>", text)
                    text = text.replace("<f>B2*2</f>", '<f t="shared" ref="C2:C3" si="0">B2*2</f>')
                    text = text.replace("<f>B3*2</f>", '<f t="shared" si="0"/>')
                    data = text.encode("utf-8")
                out.writestr(info, data)
        self.assertEqual(self.verdict(X.compare_files(str(explicit), str(shared))), "IDENTICAL")

    def test_shared_and_inline_strings_are_the_same_content(self):
        inline = build(self.tmp / "a.xlsx")      # openpyxl 3.1 writes t="inlineStr"
        shared = self.tmp / "b.xlsx"
        with ZipFile(inline) as src, ZipFile(shared, "w", ZIP_DEFLATED) as out:
            strings = []

            def to_shared(match):
                text = match[3]
                strings.append(text)
                return '<c r="%s"%s t="s"><v>%d</v></c>' % (match[1], match[2], len(strings) - 1)
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename.startswith("xl/worksheets/sheet"):
                    data = re.sub(r'<c r="([A-Z]+\d+)"([^>]*?) t="inlineStr"><is><t>(.*?)</t></is></c>', to_shared,
                                  data.decode("utf-8")).encode("utf-8")
                out.writestr(info, data)
            body = "".join("<si><t>%s</t></si>" % t for t in strings)
            out.writestr("xl/sharedStrings.xml", '<sst xmlns="%s">%s</sst>' % (X.NS, body))
        self.assertTrue(strings)
        self.assertEqual(self.verdict(X.compare_files(str(inline), str(shared))), "IDENTICAL")


class CliTests(Base):
    def run_cli(self, *args):
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = X.main(list(args))
        return code, out.getvalue(), err.getvalue()

    def test_exit_codes_and_a_report_with_no_values(self):
        a = build(self.tmp / "a.xlsx")
        b = build(self.tmp / "b.xlsx", lambda wb: wb["Data"].__setitem__("A2", "SECRETNAME"))
        code, out, _ = self.run_cli(str(a), str(a))
        self.assertEqual((code, out.splitlines()[0]), (0, "VERDICT IDENTICAL"))
        code, out, _ = self.run_cli(str(a), str(b), "--examples", "3", "--json", str(self.tmp / "r.json"))
        self.assertEqual((code, out.splitlines()[0]), (1, "VERDICT DIFF"))
        self.assertIn("cell_value=1", out)
        self.assertIn("Data!A2", out)
        self.assertNotIn("SECRETNAME", out)
        self.assertEqual(json.loads((self.tmp / "r.json").read_text())["counts"], {"cell_value": 1})
        _, with_values, _ = self.run_cli(str(a), str(b), "--examples", "3", "--show-values")
        self.assertIn("SECRETNAME", with_values)

    def test_unreadable_input_is_exit_two(self):
        (self.tmp / "bad.xlsx").write_bytes(b"not a zip")
        code, _, err = self.run_cli(str(self.tmp / "bad.xlsx"), str(build(self.tmp / "a.xlsx")))
        self.assertEqual(code, 2)
        self.assertIn("UNREADABLE", err)


if __name__ == "__main__":
    unittest.main()
