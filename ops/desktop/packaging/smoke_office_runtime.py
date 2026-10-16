"""Real private UNO recalculation of a synthetic workbook; never a parity claim."""
import argparse
import hashlib
import json
import os
from pathlib import Path
import sys
import tempfile


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--office', type=Path, required=True)
    parser.add_argument('--compute', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    options = parser.parse_args()
    sys.path.insert(0, str(options.compute.resolve()))
    import xl_engine
    from openpyxl import Workbook, load_workbook
    office = options.office.resolve()/'program'
    os.environ.update(MELANGE_OFFICE_CALC_LOCAL='1', MELANGE_OFFICE_PYTHON=str(office/'python.exe'),
                      SOFFICE=str(office/'soffice.com'))
    with tempfile.TemporaryDirectory(prefix='desktop-office-smoke-') as folder:
        root=Path(folder)
        original, output, receipt = root/'original.xlsx', root/'view.xlsx', root/'receipt.json'
        workbook=Workbook()
        workbook.active['A1']=17
        workbook.active['A2']=25
        workbook.active['A3']='=SUM(A1:A2)'
        workbook.save(original)
        digest=hashlib.sha256(original.read_bytes()).hexdigest()
        xl_engine.recalc_copy(original, expected_sha256=digest, out=output, receipt=receipt,
                              engine_name='libreoffice', timeout=90)
        cached=load_workbook(output, data_only=True)
        value=cached.active['A3'].value
        cached.close()
        report={'sourcePreserved':hashlib.sha256(original.read_bytes()).hexdigest()==digest,
                'formulaValue':value,'syntheticAcceptanceOnly':True,'receipt':json.loads(receipt.read_text())}
        assert report['sourcePreserved'] and value==42
    options.output.write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps({'sourcePreserved':True,'formulaValue':42,'syntheticAcceptanceOnly':True}))


if __name__=='__main__':
    main()
