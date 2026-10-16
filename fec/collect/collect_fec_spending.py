# -*- coding: utf-8 -*-
"""Standalone ingest for Mentat FEC industry spending through 2026-06-30.

Inputs are six industry workbooks in the configured FEC spending raw
directory and a read-only legislators.csv candidate-id crosswalk. Outputs are
member attribution, industry-cycle receipts, a normalized committee map, and a
fail-closed source manifest in the output directory.
Decision log (2026-07-31): keep this collector data-plane-independent; both
--raw and --out must be outside MB_ROOT.
Decision log (2026-07-31): --legislators is a deliberate read-only exception to
the data-plane boundary and is never modified by this collector.
Decision log (2026-07-31): direct candidate dollars, leadership-PAC sponsor
dollars, and unattributed dollars are separate metrics and are never mixed.
Decision log (2026-07-31): Downstream Activity remains excluded. Per the user
directive to ratify wiring, Beneficiary Summary and Beneficiary Bridge are now
consumed as committee-path EXPOSURE, never receipts or candidate-summable
allocations; Bridge trade-association tagging is best-effort via the existing
committee map, and blank or missing map categories are not counted as trade.
Decision log (2026-07-31): refunds are real observations; negative amounts are
retained and reported rather than rejected.
Decision log (2026-07-31): associated bucket values are refund-netted and may
be negative; consumers clamp them only at the point of metric use.
Decision log (2026-07-31): Credit Arms ships three stacked views distinguished
by Fact Type; only Recipient is consumed because it reconciles exactly and
carries Committee IDs, while summing Employee/Recipient/Master would
triple-count the same dollars.
Decision log (2026-10-07): Telecom ships a COMPACT 12-column cycle sheet
(Fact Type, Company, Employee Tier, Amount, Analysis Amount, Party, Recipient
Category, Recipient Key, Gift Count, Include in Analysis, Person Key, Logic
Check) with no Cycle and no Committee ID column and no dates. The adapter
derives Cycle from the sheet name and Committee ID from the Recipient Key
(FEC:C########); only Fact Type Recipient is consumed, exactly as for Credit
Arms, and Telecom is therefore cycle_only. Full-schema workbooks are untouched.
"""

import os as _o, sys as _s; _s.path.insert(0, _o.path.dirname(_o.path.dirname(_o.path.abspath(__file__)))); from paths import ROOT as MB_ROOT

import argparse
import csv
import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import json
import os
import re
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# MB_ROOT is imported for repository convention and as an isolation boundary.
# This standalone collector never reads or writes Moneyball data-plane artifacts.
try:
    from moneyball_core import dump_atomic
except Exception:
    # Worktree data planes can be empty, so retain temp-sibling + os.replace
    # semantics when the kernel import is unavailable.
    def dump_atomic(obj, path, **kw):
        temp_path = path + ".tmp." + str(os.getpid())
        try:
            with open(temp_path, "w", encoding="utf-8") as handle:
                json.dump(obj, handle, **kw)
            os.replace(temp_path, path)
        except BaseException:
            try:
                os.remove(temp_path)
            except OSError:
                pass
            raise


VINTAGE = "2026-06-30"
DETAIL_SHEETS = ("2021-2022", "2023-2024", "2025-2026")
EXCEPTIONS_SHEET = "Cycle Exceptions"
MAP_SHEET = "PAC & Candidate Map"
COMPANY_SHEET = "Company List"
SUMMARY_SHEET = "Beneficiary Summary"
BRIDGE_SHEET = "Beneficiary Bridge"
OVERVIEW_SHEETS = ("Spending Overview", "Crypto Spending Overview")
VIEW_COLUMN = "Fact Type"
VIEW_KEEP = "Recipient"
CONDUIT_COMMITTEE_IDS = {"C00401224", "C00694323"}

WORKBOOKS = {
    "credit_union": {
        "filename": "Credit Union Industry Spending Overview.xlsx",
        "side": "ours",
    },
    "community_banks": {
        "filename": "Community Banks Industry Spending Overview.xlsx",
        "side": "opposition",
    },
    "fintech": {
        "filename": "Fintech Industry Spending Overview.xlsx",
        "side": "opposition",
    },
    "crypto": {
        "filename": "Crypto Industry Spending Overview.xlsx",
        "side": "opposition",
    },
    "credit_arms": {
        "filename": "Credit Arms Industry Spending Overview.xlsx",
        "side": "opposition",
    },
    "telecom": {
        "filename": "Telecom Industry Spending Overview.xlsx",
        "side": "opposition",
    },
}

REQUIRED_DETAIL_COLUMNS = (
    "Cycle",
    "Company",
    "Amount",
    "Analysis Amount",
    "Recipient Category",
    "Committee ID",
    "Include in Analysis",
)
# Compact cycle-sheet schema (Telecom): the full-schema Cycle and Committee ID
# columns are absent, so they are derived (see _detail_headers).
COMPACT_DETAIL_COLUMNS = (
    VIEW_COLUMN,
    "Company",
    "Amount",
    "Analysis Amount",
    "Recipient Category",
    "Recipient Key",
    "Include in Analysis",
)
RECIPIENT_KEY_PREFIX = "FEC:"
OPTIONAL_DETAIL_COLUMNS = (
    "Date",
    "Employee Tier",
    "Party",
    "Recipient",
    "Contributor",
)
REQUIRED_MAP_COLUMNS = (
    "Cycle",
    "Recipient Committee",
    "Committee ID",
    "Recipient Category",
    "Candidate ID(s)",
    "Relationship",
)

MEMBER_FIELDS = [
    "bioguide",
    "candidate_ids",
    "industry",
    "side",
    "cycle",
    "month",
    "direct_dollars",
    "sponsor_dollars",
    "n_contributions",
    "first_date",
    "last_date",
]
TOTAL_FIELDS = [
    "industry",
    "side",
    "cycle",
    "recipient_category",
    "total_dollars",
    "n_rows",
    "attributed_direct",
    "attributed_sponsor",
    "unattributed",
]
MAP_FIELDS = [
    "industry",
    "cycle",
    "committee_id",
    "recipient_committee",
    "recipient_category",
    "candidate_ids",
    "bioguide",
    "relationship",
]
ASSOCIATED_FIELDS = [
    "bioguide",
    "candidate_id",
    "candidate_name",
    "party",
    "office",
    "state",
    "district",
    "industry",
    "side",
    "cycle",
    "associated_dollars",
    "assoc_support_dollars",
    "assoc_oppose_dollars",
    "assoc_mixed_dollars",
    "assoc_conduit_dollars",
    "assoc_relational_dollars",
    "assoc_noncond_trade_dollars",
]

FEC_ID_RE = re.compile(r"[HSP]\d[A-Z0-9]{2}\d{5}")
CENT = Decimal("0.01")
ZERO = Decimal("0")


class IngestError(Exception):
    """Expected ingest failure that must exit 2 without certified outputs."""


def _write_csv_atomic(path, fieldnames, rows):
    temp_path = path + ".tmp." + str(os.getpid())
    try:
        with open(temp_path, "w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)
        os.replace(temp_path, path)
    except BaseException:
        try:
            os.remove(temp_path)
        except OSError:
            pass
        raise


def _file_md5(path):
    digest = hashlib.md5()
    with open(path, "rb") as handle:
        while True:
            chunk = handle.read(1024 * 1024)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def _outside_mb_root(path_name, path):
    resolved_path = os.path.realpath(os.path.abspath(path))
    root_real = os.path.realpath(MB_ROOT)
    try:
        under_root = os.path.normcase(
            os.path.commonpath([resolved_path, root_real])
        ) == os.path.normcase(root_real)
    except ValueError:
        under_root = False
    if under_root:
        raise IngestError(
            f"{path_name} path {resolved_path} is inside MB_ROOT; this collector "
            "is deliberately data-plane-independent"
        )
    return resolved_path


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _row_has_data(row):
    return any(_text(value) for value in row)


def _headers(ws, row_number, required, workbook_name):
    row = next(
        ws.iter_rows(
            min_row=row_number,
            max_row=row_number,
            values_only=True,
        ),
        (),
    )
    positions = {}
    duplicates = set()
    for index, value in enumerate(row):
        name = _text(value)
        if not name:
            continue
        if name in positions:
            duplicates.add(name)
        else:
            positions[name] = index
    for column in required:
        if column not in positions:
            raise IngestError(
                f"{workbook_name}/{ws.title}: required column {column!r} missing"
            )
        if column in duplicates:
            raise IngestError(
                f"{workbook_name}/{ws.title}: required column {column!r} duplicated"
            )
    return positions


def _cell(row, positions, column):
    index = positions[column]
    return row[index] if index < len(row) else None


def _detail_headers(ws, row_number, workbook_name):
    """Return (positions, compact) for a detail or exception sheet.

    A sheet that carries Cycle or Committee ID is read exactly as before. A
    sheet that carries neither is the compact schema: it must supply
    COMPACT_DETAIL_COLUMNS, and "Cycle" and "Committee ID" are mapped to two
    virtual columns appended by _expand_compact_row.
    """
    probe = _headers(ws, row_number, (), workbook_name)
    if "Cycle" in probe or "Committee ID" in probe:
        return _headers(ws, row_number, REQUIRED_DETAIL_COLUMNS, workbook_name), False
    positions = _headers(ws, row_number, COMPACT_DETAIL_COLUMNS, workbook_name)
    width = max(positions.values()) + 1
    positions["Cycle"] = width
    positions["Committee ID"] = width + 1
    return positions, True


def _expand_compact_row(row, positions, sheet_name, workbook_name, row_number):
    """Append the derived Cycle (sheet name) and Committee ID (Recipient Key)."""
    width = positions["Cycle"]
    padded = list(row[:width]) + [None] * (width - len(row))
    key = _text(padded[positions["Recipient Key"]])
    if not key:
        committee_id = ""
    elif key.upper().startswith(RECIPIENT_KEY_PREFIX):
        committee_id = key[len(RECIPIENT_KEY_PREFIX):].strip().upper()
    else:
        raise IngestError(
            f"{workbook_name}/{sheet_name} row {row_number}: Recipient Key "
            f"{key!r} is not an {RECIPIENT_KEY_PREFIX}<committee id> key"
        )
    return tuple(padded) + (sheet_name, committee_id)


def _decimal(value, workbook_name, sheet_name, row_number, column):
    text = _text(value)
    if not text:
        return ZERO
    try:
        amount = Decimal(text.replace(",", "").replace("$", ""))
    except InvalidOperation as exc:
        raise IngestError(
            f"{workbook_name}/{sheet_name} row {row_number}: invalid "
            f"{column} value {text!r}"
        ) from exc
    if not amount.is_finite():
        raise IngestError(
            f"{workbook_name}/{sheet_name} row {row_number}: non-finite "
            f"{column} value {text!r}"
        )
    return amount


def _include_flag(value, workbook_name, sheet_name, row_number):
    if isinstance(value, bool):
        return value
    token = _text(value).casefold()
    if token in {"yes", "y", "true", "1"}:
        return True
    if token in {"no", "n", "false", "0"}:
        return False
    raise IngestError(
        f"{workbook_name}/{sheet_name} row {row_number}: invalid "
        f"Include in Analysis value {_text(value)!r}"
    )


def _date_value(value, workbook_name, sheet_name, row_number):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = _text(value)
    if not text:
        raise IngestError(
            f"{workbook_name}/{sheet_name} row {row_number}: Date is blank"
        )
    parsers = (
        lambda token: datetime.date.fromisoformat(token[:10]),
        lambda token: datetime.datetime.strptime(token, "%m/%d/%Y").date(),
    )
    for parser in parsers:
        try:
            return parser(text)
        except ValueError:
            pass
    raise IngestError(
        f"{workbook_name}/{sheet_name} row {row_number}: invalid Date {text!r}"
    )


def _candidate_ids(value):
    return sorted(set(FEC_ID_RE.findall(_text(value).upper())))


def _money_decimal(value):
    rounded = value.quantize(CENT, rounding=ROUND_HALF_UP)
    return Decimal("0.00") if rounded == 0 else rounded


def _money_text(value):
    return format(_money_decimal(value), ".2f")


def _money_number(value):
    return float(_money_decimal(value))


def _load_legislators(path):
    fec_to_bioguide = {}
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = set(reader.fieldnames or [])
        for required in ("bioguide_id", "fec_ids"):
            if required not in fields:
                raise IngestError(
                    f"legislators.csv required column {required!r} missing: {path}"
                )
        for row_number, row in enumerate(reader, start=2):
            bioguide = _text(row.get("bioguide_id"))
            if not bioguide:
                continue
            for candidate_id in _candidate_ids(row.get("fec_ids")):
                prior = fec_to_bioguide.get(candidate_id)
                if prior and prior != bioguide:
                    raise IngestError(
                        f"legislators.csv row {row_number}: candidate id "
                        f"{candidate_id} maps to both {prior} and {bioguide}"
                    )
                fec_to_bioguide[candidate_id] = bioguide
    if not fec_to_bioguide:
        raise IngestError(
            f"legislators.csv yielded 0 fec-id-to-bioguide pairs: {path}"
        )
    return fec_to_bioguide


def _merge_label(existing, incoming):
    if not existing:
        return incoming
    if not incoming or incoming == existing:
        return existing
    return "; ".join(sorted(set(existing.split("; ") + [incoming])))


def _load_candidate_map(ws, workbook_name, industry, fec_to_bioguide):
    positions = _headers(ws, 5, REQUIRED_MAP_COLUMNS, workbook_name)
    entries = {}
    blank_committee_rows = 0
    for row_number, row in enumerate(
        ws.iter_rows(min_row=6, values_only=True), start=6
    ):
        if not _row_has_data(row):
            continue
        cycle = _text(_cell(row, positions, "Cycle"))
        committee_id = _text(_cell(row, positions, "Committee ID")).upper()
        if not committee_id:
            blank_committee_rows += 1
            continue
        key = (cycle, committee_id)
        entry = entries.get(key)
        if entry is None:
            entry = {
                "industry": industry,
                "cycle": cycle,
                "committee_id": committee_id,
                "recipient_committee": "",
                "recipient_category": "",
                "candidate_ids_set": set(),
                "relationship": "",
            }
            entries[key] = entry
        entry["recipient_committee"] = _merge_label(
            entry["recipient_committee"],
            _text(_cell(row, positions, "Recipient Committee")),
        )
        entry["recipient_category"] = _merge_label(
            entry["recipient_category"],
            _text(_cell(row, positions, "Recipient Category")),
        )
        entry["relationship"] = _merge_label(
            entry["relationship"],
            _text(_cell(row, positions, "Relationship")),
        )
        entry["candidate_ids_set"].update(
            _candidate_ids(_cell(row, positions, "Candidate ID(s)"))
        )

    output_rows = []
    for key in sorted(entries):
        entry = entries[key]
        mapped_bioguides = {
            fec_to_bioguide[candidate_id]
            for candidate_id in entry["candidate_ids_set"]
            if candidate_id in fec_to_bioguide
        }
        entry["mapped_bioguides"] = mapped_bioguides
        entry["bioguide"] = (
            next(iter(mapped_bioguides)) if len(mapped_bioguides) == 1 else ""
        )
        output_rows.append(
            {
                "industry": industry,
                "cycle": entry["cycle"],
                "committee_id": entry["committee_id"],
                "recipient_committee": entry["recipient_committee"],
                "recipient_category": entry["recipient_category"],
                "candidate_ids": ";".join(sorted(entry["candidate_ids_set"])),
                "bioguide": entry["bioguide"],
                "relationship": entry["relationship"],
            }
        )
    return entries, output_rows, blank_committee_rows


def _load_company_roster(ws, workbook_name):
    positions = _headers(ws, 5, ("Company",), workbook_name)
    companies = set()
    for row in ws.iter_rows(min_row=6, values_only=True):
        if not _row_has_data(row):
            continue
        company = _text(_cell(row, positions, "Company"))
        if company:
            companies.add(company)
    return sorted(companies, key=lambda value: (value.casefold(), value))


def _associated_column(positions, workbook_name, sheet_name):
    matches = [
        column
        for column in positions
        if column.casefold().startswith("associated")
        and column.casefold().endswith("funding")
    ]
    if len(matches) != 1:
        observed = ", ".join(repr(value) for value in sorted(matches)) or "none"
        raise IngestError(
            f"{workbook_name}/{sheet_name}: expected exactly one Associated "
            f"* Funding column; observed matches: {observed}"
        )
    return matches[0]


def _committee_category_lookups(committee_map):
    by_cycle = {}
    by_committee = {}
    cycles_by_committee = {}
    for (cycle, committee_id), entry in committee_map.items():
        categories = {
            value.strip().casefold()
            for value in entry["recipient_category"].split(";")
            if value.strip()
        }
        by_cycle[(cycle, committee_id)] = categories
        by_committee.setdefault(committee_id, set()).update(categories)
        cycles_by_committee.setdefault(committee_id, set()).add(cycle)
    multi_cycle_category = {
        committee_id
        for committee_id, cycles in cycles_by_committee.items()
        if len({frozenset(by_cycle[(cycle, committee_id)]) for cycle in cycles})
        > 1
    }
    return by_cycle, by_committee, multi_cycle_category


def _new_bridge_aggregate():
    return {
        "total": ZERO,
        "support": ZERO,
        "oppose": ZERO,
        "mixed": ZERO,
        "conduit": ZERO,
        "relational": ZERO,
        "noncond_trade": ZERO,
    }


def _load_associated_views(
    workbook,
    workbook_name,
    industry,
    side,
    fec_to_bioguide,
    committee_map,
):
    summary_ws = workbook[SUMMARY_SHEET]
    summary_positions = _headers(
        summary_ws,
        5,
        (
            "Cycle",
            "View",
            "Company",
            "Candidate / Officeholder",
            "Candidate ID",
            "Party",
            "Office",
            "State",
            "District",
        ),
        workbook_name,
    )
    summary_amount_column = _associated_column(
        summary_positions, workbook_name, SUMMARY_SHEET
    )
    summary_aggregates = {}
    summary_total = ZERO
    n_industry_view_rows = 0
    candidate_ids = set()
    sitting_bioguides = set()
    nonmember_candidate_ids = set()
    company_values = set()

    for row_number, row in enumerate(
        summary_ws.iter_rows(min_row=6, values_only=True), start=6
    ):
        if not _row_has_data(row):
            continue
        if _text(_cell(row, summary_positions, "View")).casefold() != "industry":
            continue
        n_industry_view_rows += 1
        candidate_id = _text(
            _cell(row, summary_positions, "Candidate ID")
        ).upper()
        if not candidate_id:
            raise IngestError(
                f"{workbook_name}/{SUMMARY_SHEET} row {row_number}: "
                "Industry-view Candidate ID is blank"
            )
        cycle = _text(_cell(row, summary_positions, "Cycle"))
        amount = _decimal(
            _cell(row, summary_positions, summary_amount_column),
            workbook_name,
            SUMMARY_SHEET,
            row_number,
            summary_amount_column,
        )
        summary_total += amount
        candidate_ids.add(candidate_id)
        company = _text(_cell(row, summary_positions, "Company"))
        if company:
            company_values.add(company)
        bioguide = fec_to_bioguide.get(candidate_id, "")
        if bioguide:
            sitting_bioguides.add(bioguide)
        else:
            nonmember_candidate_ids.add(candidate_id)

        key = (candidate_id, cycle)
        aggregate = summary_aggregates.get(key)
        if aggregate is None:
            aggregate = {
                "candidate_id": candidate_id,
                "cycle": cycle,
                "bioguide": bioguide,
                "candidate_name": "",
                "party": "",
                "office": "",
                "state": "",
                "district": "",
                "associated": ZERO,
            }
            summary_aggregates[key] = aggregate
        aggregate["associated"] += amount
        for target, column in (
            ("candidate_name", "Candidate / Officeholder"),
            ("party", "Party"),
            ("office", "Office"),
            ("state", "State"),
            ("district", "District"),
        ):
            aggregate[target] = _merge_label(
                aggregate[target],
                _text(_cell(row, summary_positions, column)),
            )

    if n_industry_view_rows == 0:
        raise IngestError(
            f"{workbook_name}/{SUMMARY_SHEET}: 0 Industry-view rows"
        )
    if summary_total == 0:
        raise IngestError(
            f"{workbook_name}: {SUMMARY_SHEET} has Industry-view rows but "
            "their total associated sum is 0.00"
        )

    bridge_ws = workbook[BRIDGE_SHEET]
    bridge_positions = _headers(
        bridge_ws,
        5,
        (
            "Cycle",
            "View",
            "Source Committee ID",
            "Candidate ID",
            "Directions",
            "Activity Relationships",
        ),
        workbook_name,
    )
    bridge_amount_column = _associated_column(
        bridge_positions, workbook_name, BRIDGE_SHEET
    )
    (
        categories_by_cycle,
        categories_by_committee,
        committees_with_multi_cycle_category,
    ) = _committee_category_lookups(committee_map)
    bridge_aggregates = {}
    bridge_total = ZERO
    n_mixed_direction_rows = 0
    n_negative_bridge_rows = 0
    conduit_with_direction_dollars = ZERO
    trade_category_cycle_fallbacks = 0
    for row_number, row in enumerate(
        bridge_ws.iter_rows(min_row=6, values_only=True), start=6
    ):
        if not _row_has_data(row):
            continue
        if _text(_cell(row, bridge_positions, "View")).casefold() != "industry":
            continue
        candidate_id = _text(
            _cell(row, bridge_positions, "Candidate ID")
        ).upper()
        cycle = _text(_cell(row, bridge_positions, "Cycle"))
        source_committee_id = _text(
            _cell(row, bridge_positions, "Source Committee ID")
        ).upper()
        amount = _decimal(
            _cell(row, bridge_positions, bridge_amount_column),
            workbook_name,
            BRIDGE_SHEET,
            row_number,
            bridge_amount_column,
        )
        bridge_total += amount
        if amount < 0:
            n_negative_bridge_rows += 1
        key = (candidate_id, cycle)
        aggregate = bridge_aggregates.get(key)
        if aggregate is None:
            aggregate = _new_bridge_aggregate()
            bridge_aggregates[key] = aggregate
        aggregate["total"] += amount
        direction_text = (
            _text(_cell(row, bridge_positions, "Directions"))
            + " "
            + _text(_cell(row, bridge_positions, "Activity Relationships"))
        ).casefold()
        has_oppose = "oppose" in direction_text
        has_support = "support" in direction_text
        if has_oppose and has_support:
            n_mixed_direction_rows += 1

        if source_committee_id in CONDUIT_COMMITTEE_IDS:
            aggregate["conduit"] += amount
            if has_oppose or has_support:
                conduit_with_direction_dollars += amount
        elif has_oppose and has_support:
            aggregate["mixed"] += amount
        elif has_oppose:
            aggregate["oppose"] += amount
        elif has_support:
            aggregate["support"] += amount
        else:
            aggregate["relational"] += amount

        if source_committee_id not in CONDUIT_COMMITTEE_IDS:
            categories = categories_by_cycle.get(
                (cycle, source_committee_id)
            )
            if categories is None:
                categories = categories_by_committee.get(
                    source_committee_id,
                    set(),
                )
                if source_committee_id in categories_by_committee:
                    trade_category_cycle_fallbacks += 1
            if "trade association" in categories:
                aggregate["noncond_trade"] += amount

    bridge_keys = set(summary_aggregates) | set(bridge_aggregates)
    for candidate_id, cycle in sorted(bridge_keys):
        summary_aggregate = summary_aggregates.get((candidate_id, cycle))
        bridge_aggregate = bridge_aggregates.get((candidate_id, cycle))
        summary_amount = (
            summary_aggregate["associated"] if summary_aggregate else ZERO
        )
        bridge_amount = (
            bridge_aggregate["total"] if bridge_aggregate else ZERO
        )
        delta = bridge_amount - summary_amount
        if abs(delta) > Decimal("0.01"):
            raise IngestError(
                f"{workbook_name}: cycle {cycle}, candidate {candidate_id!r} "
                f"has {SUMMARY_SHEET} associated "
                f"{_money_text(summary_amount)} and {BRIDGE_SHEET} associated "
                f"{_money_text(bridge_amount)}; delta {_money_text(delta)}"
            )

    bridge_delta = bridge_total - summary_total
    if abs(bridge_delta) > Decimal("0.01"):
        raise IngestError(
            f"{workbook_name}: {BRIDGE_SHEET} Industry-view sum "
            f"{_money_text(bridge_total)} does not reconcile to {SUMMARY_SHEET} "
            f"Industry-view sum {_money_text(summary_total)}; delta "
            f"{_money_text(bridge_delta)}"
        )

    output_rows = []
    n_negative_bucket_values = 0
    for key in sorted(summary_aggregates):
        aggregate = summary_aggregates[key]
        bridge = bridge_aggregates.get(key, _new_bridge_aggregate())
        bucket_names = (
            "conduit",
            "mixed",
            "oppose",
            "support",
            "relational",
        )
        bucket_total = sum(
            (bridge[bucket] for bucket in bucket_names),
            ZERO,
        )
        bucket_delta = bucket_total - aggregate["associated"]
        if abs(bucket_delta) > Decimal("0.01"):
            raise IngestError(
                f"{workbook_name}: candidate {aggregate['candidate_id']!r}, "
                f"industry {industry}, cycle {aggregate['cycle']} has bucket "
                f"sum {_money_text(bucket_total)} and associated_dollars "
                f"{_money_text(aggregate['associated'])}; delta "
                f"{_money_text(bucket_delta)}"
            )
        n_negative_bucket_values += sum(
            bridge[bucket] < 0 for bucket in bucket_names
        )
        output_rows.append(
            {
                "bioguide": aggregate["bioguide"],
                "candidate_id": aggregate["candidate_id"],
                "candidate_name": aggregate["candidate_name"],
                "party": aggregate["party"],
                "office": aggregate["office"],
                "state": aggregate["state"],
                "district": aggregate["district"],
                "industry": industry,
                "side": side,
                "cycle": aggregate["cycle"],
                "associated_dollars": _money_text(aggregate["associated"]),
                "assoc_support_dollars": _money_text(
                    bridge["support"]
                ),
                "assoc_oppose_dollars": _money_text(
                    bridge["oppose"]
                ),
                "assoc_mixed_dollars": _money_text(
                    bridge["mixed"]
                ),
                "assoc_conduit_dollars": _money_text(
                    bridge["conduit"]
                ),
                "assoc_relational_dollars": _money_text(
                    bridge["relational"]
                ),
                "assoc_noncond_trade_dollars": _money_text(
                    bridge["noncond_trade"]
                ),
            }
        )
    stats = {
        "assoc_sum": summary_total,
        "n_industry_view_rows": n_industry_view_rows,
        "candidate_ids": candidate_ids,
        "sitting_bioguides": sitting_bioguides,
        "nonmember_candidate_ids": nonmember_candidate_ids,
        "company_values": company_values,
        "bridge_reconciled": True,
        "bridge_keys_checked": len(bridge_keys),
        "n_mixed_direction_rows": n_mixed_direction_rows,
        "n_negative_bridge_rows": n_negative_bridge_rows,
        "n_negative_bucket_values": n_negative_bucket_values,
        "conduit_with_direction_dollars": (
            conduit_with_direction_dollars
        ),
        "trade_category_cycle_fallbacks": (
            trade_category_cycle_fallbacks
        ),
        "committees_with_multi_cycle_category": (
            committees_with_multi_cycle_category
        ),
    }
    return output_rows, stats


def _new_receipt():
    return {
        "total_rows": 0,
        "excluded_exception_rows": 0,
        "rows_include_no": 0,
        "committees": set(),
        "candidate_ids": set(),
        "bioguides": set(),
        "direct_dollars": ZERO,
        "sponsor_dollars": ZERO,
        "unattributed_dollars": ZERO,
        "unattributed_by_category": {},
        "ambiguous_multi_candidate_rows": 0,
        "map_rows_blank_committee_id": 0,
        "eligible_rows_blank_committee_id": 0,
        "eligible_rows_unmapped_committee": 0,
        "eligible_unmapped_committee_dollars": ZERO,
        "assoc_sum": ZERO,
        "assoc_authoritative_total": ZERO,
        "assoc_authoritative_available": True,
        "n_industry_view_rows": 0,
        "assoc_candidate_ids": set(),
        "assoc_sitting_bioguides": set(),
        "assoc_nonmember_candidate_ids": set(),
        "assoc_company_values": set(),
        "bridge_reconciled": True,
        "bridge_keys_checked": 0,
        "n_mixed_direction_rows": 0,
        "n_negative_bridge_rows": 0,
        "n_negative_bucket_values": 0,
        "conduit_with_direction_dollars": ZERO,
        "trade_category_cycle_fallbacks": 0,
        "committees_with_multi_cycle_category": set(),
        "negative_count": 0,
        "negative_sum": ZERO,
        "dates": set(),
        "months": set(),
    }


def _new_member_aggregate():
    return {
        "candidate_ids": set(),
        "direct_dollars": ZERO,
        "sponsor_dollars": ZERO,
        "n_contributions": 0,
        "dates": set(),
    }


def _new_total_aggregate(industry, side, cycle, category):
    return {
        "industry": industry,
        "side": side,
        "cycle": cycle,
        "recipient_category": category,
        "total_dollars": ZERO,
        "n_rows": 0,
        "attributed_direct": ZERO,
        "attributed_sponsor": ZERO,
        "unattributed": ZERO,
    }


def _count_exception_rows(ws, workbook_name):
    positions, compact = _detail_headers(ws, 4, workbook_name)
    all_rows = 0
    kept_rows = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not _row_has_data(row):
            continue
        all_rows += 1
        # Compact exception rows carry the single Fact Type "Exception", so every
        # one of them is a kept (unstacked) row.
        if compact or VIEW_COLUMN not in positions or _text(
            _cell(row, positions, VIEW_COLUMN)
        ).casefold() == VIEW_KEEP.casefold():
            kept_rows += 1
    return kept_rows, all_rows


def _load_reconciliation(
    workbook,
    workbook_name,
    computed_by_cycle,
):
    overview_name = next(
        (name for name in OVERVIEW_SHEETS if name in workbook.sheetnames),
        None,
    )
    if overview_name is None:
        raise IngestError(
            f"{workbook_name}: no overview sheet "
            f"({' or '.join(OVERVIEW_SHEETS)}); cannot reconcile against "
            f"Total Giving"
        )
    ws = workbook[overview_name]
    positions = _headers(
        ws,
        5,
        ("Cycle", "Scope", "Total Giving"),
        workbook_name,
    )
    authoritative = {}
    for row_number, row in enumerate(
        ws.iter_rows(min_row=6, values_only=True), start=6
    ):
        if not _row_has_data(row):
            continue
        if _text(_cell(row, positions, "Scope")).casefold() != "total giving":
            continue
        cycle = _text(_cell(row, positions, "Cycle"))
        total = _decimal(
            _cell(row, positions, "Total Giving"),
            workbook_name,
            overview_name,
            row_number,
            "Total Giving",
        )
        if cycle in authoritative and authoritative[cycle] != total:
            raise IngestError(
                f"{workbook_name}/{overview_name}: conflicting Total Giving "
                f"rows for cycle {cycle!r}"
            )
        authoritative[cycle] = total
    for cycle in sorted(set(computed_by_cycle) - set(authoritative)):
        raise IngestError(
            f"{workbook_name}/{overview_name}: computed cycle {cycle!r} has "
            f"no authoritative Total Giving counterpart; computed total "
            f"{_money_text(computed_by_cycle[cycle])}"
        )

    reconciliation = {}
    for cycle in sorted(authoritative):
        workbook_total = authoritative[cycle]
        computed_total = computed_by_cycle.get(cycle, ZERO)
        delta = computed_total - workbook_total
        ok = abs(delta) <= Decimal("0.01")
        reconciliation[cycle] = {
            "workbook_total": _money_number(workbook_total),
            "computed_total": _money_number(computed_total),
            "delta": _money_number(delta),
            "ok": ok,
        }
        if not ok:
            raise IngestError(
                f"{workbook_name}/{overview_name}: reconciliation delta for "
                f"{cycle} is {_money_text(delta)}, exceeding $0.01"
            )
    return reconciliation, sum(authoritative.values(), ZERO)


def _process_workbook(path, industry, side, fec_to_bioguide):
    workbook_name = os.path.basename(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in DETAIL_SHEETS + (
            EXCEPTIONS_SHEET,
            MAP_SHEET,
            COMPANY_SHEET,
            SUMMARY_SHEET,
            BRIDGE_SHEET,
        ):
            if sheet_name not in workbook.sheetnames:
                raise IngestError(
                    f"{workbook_name}: required sheet {sheet_name!r} missing"
                )

        detail_positions = {}
        compact_sheets = {}
        optional_columns = {}
        date_presence = []
        for sheet_name in DETAIL_SHEETS:
            positions, compact_sheets[sheet_name] = _detail_headers(
                workbook[sheet_name],
                4,
                workbook_name,
            )
            detail_positions[sheet_name] = positions
            optional_columns[sheet_name] = [
                column for column in OPTIONAL_DETAIL_COLUMNS if column in positions
            ]
            date_presence.append("Date" in positions)
        if all(date_presence):
            date_status = "dated"
        elif not any(date_presence):
            date_status = "cycle_only"
        else:
            raise IngestError(
                f"{workbook_name}: Date column is inconsistently present across "
                "detail sheets"
            )
        view_column_present = any(
            VIEW_COLUMN in positions for positions in detail_positions.values()
        )
        view_rows_skipped = 0
        view_row_counts = {}

        committee_map, map_rows, blank_committee_rows = _load_candidate_map(
            workbook[MAP_SHEET],
            workbook_name,
            industry,
            fec_to_bioguide,
        )
        roster = _load_company_roster(workbook[COMPANY_SHEET], workbook_name)
        receipt = _new_receipt()
        receipt["map_rows_blank_committee_id"] = blank_committee_rows
        receipt["candidate_ids"].update(
            candidate_id
            for entry in committee_map.values()
            for candidate_id in entry["candidate_ids_set"]
        )
        (
            receipt["excluded_exception_rows"],
            excluded_exception_rows_all_views,
        ) = _count_exception_rows(
            workbook[EXCEPTIONS_SHEET], workbook_name
        )
        associated_rows, associated_stats = _load_associated_views(
            workbook,
            workbook_name,
            industry,
            side,
            fec_to_bioguide,
            committee_map,
        )
        receipt["assoc_sum"] = associated_stats["assoc_sum"]
        receipt["n_industry_view_rows"] = associated_stats[
            "n_industry_view_rows"
        ]
        receipt["assoc_candidate_ids"].update(
            associated_stats["candidate_ids"]
        )
        receipt["assoc_sitting_bioguides"].update(
            associated_stats["sitting_bioguides"]
        )
        receipt["assoc_nonmember_candidate_ids"].update(
            associated_stats["nonmember_candidate_ids"]
        )
        receipt["assoc_company_values"].update(
            associated_stats["company_values"]
        )
        receipt["bridge_reconciled"] = associated_stats[
            "bridge_reconciled"
        ]
        receipt["bridge_keys_checked"] = associated_stats[
            "bridge_keys_checked"
        ]
        receipt["n_mixed_direction_rows"] = associated_stats[
            "n_mixed_direction_rows"
        ]
        receipt["n_negative_bridge_rows"] = associated_stats[
            "n_negative_bridge_rows"
        ]
        receipt["n_negative_bucket_values"] = associated_stats[
            "n_negative_bucket_values"
        ]
        receipt["conduit_with_direction_dollars"] = associated_stats[
            "conduit_with_direction_dollars"
        ]
        receipt["trade_category_cycle_fallbacks"] = associated_stats[
            "trade_category_cycle_fallbacks"
        ]
        receipt["committees_with_multi_cycle_category"].update(
            associated_stats["committees_with_multi_cycle_category"]
        )

        member_aggregates = {}
        total_aggregates = {}
        detail_row_counts = {}
        computed_by_cycle = {}

        for sheet_name in DETAIL_SHEETS:
            ws = workbook[sheet_name]
            positions = detail_positions[sheet_name]
            data_rows = 0
            sheet_view_values = set()
            for row_number, row in enumerate(
                ws.iter_rows(min_row=5, values_only=True), start=5
            ):
                if not _row_has_data(row):
                    continue
                if VIEW_COLUMN in positions:
                    fact_type = _text(_cell(row, positions, VIEW_COLUMN))
                    sheet_view_values.add(fact_type)
                    view_row_counts[fact_type] = (
                        view_row_counts.get(fact_type, 0) + 1
                    )
                    if fact_type.casefold() != VIEW_KEEP.casefold():
                        view_rows_skipped += 1
                        continue
                if compact_sheets[sheet_name]:
                    row = _expand_compact_row(
                        row, positions, sheet_name, workbook_name, row_number
                    )
                data_rows += 1
                receipt["total_rows"] += 1

                cycle = _text(_cell(row, positions, "Cycle"))
                if not cycle:
                    raise IngestError(
                        f"{workbook_name}/{sheet_name} row {row_number}: "
                        "Cycle is blank"
                    )
                source_amount = _decimal(
                    _cell(row, positions, "Amount"),
                    workbook_name,
                    sheet_name,
                    row_number,
                    "Amount",
                )
                analysis_amount = _decimal(
                    _cell(row, positions, "Analysis Amount"),
                    workbook_name,
                    sheet_name,
                    row_number,
                    "Analysis Amount",
                )
                include = _include_flag(
                    _cell(row, positions, "Include in Analysis"),
                    workbook_name,
                    sheet_name,
                    row_number,
                )
                if not include:
                    receipt["rows_include_no"] += 1

                transaction_date = None
                if date_status == "dated":
                    transaction_date = _date_value(
                        _cell(row, positions, "Date"),
                        workbook_name,
                        sheet_name,
                        row_number,
                    )
                    date_text = transaction_date.isoformat()
                    receipt["dates"].add(date_text)
                    receipt["months"].add(date_text[:7])

                if source_amount < 0:
                    receipt["negative_count"] += 1
                    receipt["negative_sum"] += source_amount

                committee_id = _text(
                    _cell(row, positions, "Committee ID")
                ).upper()
                if committee_id:
                    receipt["committees"].add(committee_id)
                category = (
                    _text(_cell(row, positions, "Recipient Category"))
                    or "Unclassified"
                )
                total_key = (industry, cycle, category)
                total = total_aggregates.get(total_key)
                if total is None:
                    total = _new_total_aggregate(
                        industry, side, cycle, category
                    )
                    total_aggregates[total_key] = total
                total["n_rows"] += 1

                normalized_amount = analysis_amount if include else ZERO
                total["total_dollars"] += normalized_amount
                computed_by_cycle[cycle] = (
                    computed_by_cycle.get(cycle, ZERO) + normalized_amount
                )
                if not include:
                    continue

                category_key = category.casefold()
                eligible = category_key in {"candidate", "leadership pac"}
                if eligible and not committee_id:
                    receipt["eligible_rows_blank_committee_id"] += 1
                elif (
                    eligible
                    and (cycle, committee_id) not in committee_map
                ):
                    receipt["eligible_rows_unmapped_committee"] += 1
                    receipt["eligible_unmapped_committee_dollars"] += (
                        normalized_amount
                    )
                map_entry = (
                    committee_map.get((cycle, committee_id))
                    if committee_id
                    else None
                )
                mapped_bioguides = (
                    map_entry["mapped_bioguides"] if map_entry else set()
                )
                if eligible and len(mapped_bioguides) > 1:
                    receipt["ambiguous_multi_candidate_rows"] += 1

                if eligible and len(mapped_bioguides) == 1:
                    bioguide = next(iter(mapped_bioguides))
                    month = (
                        transaction_date.strftime("%Y-%m")
                        if transaction_date is not None
                        else ""
                    )
                    member_key = (bioguide, industry, cycle, month)
                    member = member_aggregates.get(member_key)
                    if member is None:
                        member = _new_member_aggregate()
                        member_aggregates[member_key] = member
                    member["candidate_ids"].update(
                        candidate_id
                        for candidate_id in map_entry["candidate_ids_set"]
                        if fec_to_bioguide.get(candidate_id) == bioguide
                    )
                    member["n_contributions"] += 1
                    if transaction_date is not None:
                        member["dates"].add(transaction_date.isoformat())
                    receipt["bioguides"].add(bioguide)
                    if category_key == "candidate":
                        member["direct_dollars"] += normalized_amount
                        total["attributed_direct"] += normalized_amount
                        receipt["direct_dollars"] += normalized_amount
                    else:
                        member["sponsor_dollars"] += normalized_amount
                        total["attributed_sponsor"] += normalized_amount
                        receipt["sponsor_dollars"] += normalized_amount
                else:
                    total["unattributed"] += normalized_amount
                    receipt["unattributed_dollars"] += normalized_amount
                    by_category = receipt["unattributed_by_category"]
                    by_category[category] = (
                        by_category.get(category, ZERO) + normalized_amount
                    )

            if VIEW_COLUMN in positions and data_rows == 0:
                observed = ", ".join(
                    repr(value) for value in sorted(sheet_view_values)
                ) or "none"
                raise IngestError(
                    f"{workbook_name}/{sheet_name}: 0 rows matched "
                    f"{VIEW_COLUMN}={VIEW_KEEP!r}; observed Fact Type values: "
                    f"{observed}"
                )
            if data_rows == 0:
                raise IngestError(
                    f"{workbook_name}/{sheet_name}: detail sheet has 0 data rows"
                )
            detail_row_counts[sheet_name] = data_rows

        computed_total_all_cycles = sum(computed_by_cycle.values(), ZERO)
        if (
            sum(detail_row_counts.values()) > 0
            and computed_total_all_cycles == 0
        ):
            raise IngestError(
                f"{workbook_name}: processed detail rows have a computed "
                "total of 0.00 across all cycles"
            )
        reconciliation, authoritative_total = _load_reconciliation(
            workbook,
            workbook_name,
            computed_by_cycle,
        )
        if authoritative_total is None:
            receipt["assoc_authoritative_available"] = False
        else:
            receipt["assoc_authoritative_total"] = authoritative_total
        workbook_manifest = {
            "present": True,
            "processed": True,
            "path": path,
            "sheet_count": len(workbook.sheetnames),
            "detail_row_counts": detail_row_counts,
            "date_status": date_status,
            "optional_detail_columns": optional_columns,
            "view_column": VIEW_COLUMN if view_column_present else None,
            "roster": roster,
            "reconciliation": reconciliation,
        }
        if any(compact_sheets.values()):
            workbook_manifest["detail_schema"] = {
                "kind": "compact_cycle_sheet",
                "cycle": "sheet name",
                "committee_id": "Recipient Key without the FEC: prefix",
                "compact_sheets": sorted(
                    name for name, flag in compact_sheets.items() if flag
                ),
            }
        if view_column_present:
            workbook_manifest.update(
                {
                    "view_kept": VIEW_KEEP,
                    "view_rows_skipped": view_rows_skipped,
                    "view_row_counts": dict(sorted(view_row_counts.items())),
                    "excluded_exception_rows_all_views": (
                        excluded_exception_rows_all_views
                    ),
                }
            )
        return (
            workbook_manifest,
            member_aggregates,
            total_aggregates,
            map_rows,
            associated_rows,
            receipt,
        )
    finally:
        workbook.close()


def _merge_receipts(receipts):
    merged = _new_receipt()
    for receipt in receipts:
        for key in (
            "total_rows",
            "excluded_exception_rows",
            "rows_include_no",
            "ambiguous_multi_candidate_rows",
            "map_rows_blank_committee_id",
            "eligible_rows_blank_committee_id",
            "eligible_rows_unmapped_committee",
            "n_industry_view_rows",
            "bridge_keys_checked",
            "n_mixed_direction_rows",
            "n_negative_bridge_rows",
            "n_negative_bucket_values",
            "trade_category_cycle_fallbacks",
            "negative_count",
        ):
            merged[key] += receipt[key]
        for key in (
            "direct_dollars",
            "sponsor_dollars",
            "unattributed_dollars",
            "eligible_unmapped_committee_dollars",
            "assoc_sum",
            "assoc_authoritative_total",
            "conduit_with_direction_dollars",
            "negative_sum",
        ):
            merged[key] += receipt[key]
        for key in (
            "committees",
            "candidate_ids",
            "bioguides",
            "dates",
            "months",
            "assoc_candidate_ids",
            "assoc_sitting_bioguides",
            "assoc_nonmember_candidate_ids",
            "assoc_company_values",
            "committees_with_multi_cycle_category",
        ):
            merged[key].update(receipt[key])
        merged["assoc_authoritative_available"] = (
            merged["assoc_authoritative_available"]
            and receipt["assoc_authoritative_available"]
        )
        merged["bridge_reconciled"] = (
            merged["bridge_reconciled"] and receipt["bridge_reconciled"]
        )
        for category, amount in receipt["unattributed_by_category"].items():
            merged["unattributed_by_category"][category] = (
                merged["unattributed_by_category"].get(category, ZERO) + amount
            )
    return merged


def _finalize_receipt(receipt):
    dates = sorted(receipt["dates"])
    if (
        receipt["assoc_authoritative_available"]
        and receipt["assoc_authoritative_total"] != 0
    ):
        fanout_ratio = float(
            receipt["assoc_sum"] / receipt["assoc_authoritative_total"]
        )
    else:
        fanout_ratio = None
    return {
        "total_rows": receipt["total_rows"],
        "excluded_exception_rows": receipt["excluded_exception_rows"],
        "rows_include_no": receipt["rows_include_no"],
        "distinct_committees": len(receipt["committees"]),
        "distinct_candidate_ids": len(receipt["candidate_ids"]),
        "distinct_bioguides": len(receipt["bioguides"]),
        "direct_dollars": _money_number(receipt["direct_dollars"]),
        "sponsor_dollars": _money_number(receipt["sponsor_dollars"]),
        "unattributed_dollars": _money_number(receipt["unattributed_dollars"]),
        "unattributed_by_category": {
            category: _money_number(amount)
            for category, amount in sorted(
                receipt["unattributed_by_category"].items()
            )
        },
        "ambiguous_multi_candidate_rows": receipt[
            "ambiguous_multi_candidate_rows"
        ],
        "map_rows_blank_committee_id": receipt[
            "map_rows_blank_committee_id"
        ],
        "eligible_rows_blank_committee_id": receipt[
            "eligible_rows_blank_committee_id"
        ],
        "eligible_rows_unmapped_committee": receipt[
            "eligible_rows_unmapped_committee"
        ],
        "eligible_unmapped_committee_dollars": _money_number(
            receipt["eligible_unmapped_committee_dollars"]
        ),
        "assoc_sum": _money_number(receipt["assoc_sum"]),
        "fanout_ratio": fanout_ratio,
        "n_industry_view_rows": receipt["n_industry_view_rows"],
        "n_candidates": len(receipt["assoc_candidate_ids"]),
        "n_sitting_members": len(receipt["assoc_sitting_bioguides"]),
        "n_nonmember_candidates": len(
            receipt["assoc_nonmember_candidate_ids"]
        ),
        "company_values": sorted(receipt["assoc_company_values"]),
        "bridge_reconciled": receipt["bridge_reconciled"],
        "bridge_keys_checked": receipt["bridge_keys_checked"],
        "n_mixed_direction_rows": receipt["n_mixed_direction_rows"],
        "n_negative_bridge_rows": receipt["n_negative_bridge_rows"],
        "n_negative_bucket_values": receipt[
            "n_negative_bucket_values"
        ],
        "conduit_with_direction_dollars": _money_number(
            receipt["conduit_with_direction_dollars"]
        ),
        "trade_category_cycle_fallbacks": receipt[
            "trade_category_cycle_fallbacks"
        ],
        "committees_with_multi_cycle_category": len(
            receipt["committees_with_multi_cycle_category"]
        ),
        "negative_amount_rows": {
            "count": receipt["negative_count"],
            "sum": _money_number(receipt["negative_sum"]),
        },
        "date_min": dates[0] if dates else None,
        "date_max": dates[-1] if dates else None,
        "months_covered": sorted(receipt["months"]),
    }


def _member_output_rows(aggregates):
    rows = []
    for key in sorted(aggregates):
        bioguide, industry, cycle, month = key
        aggregate = aggregates[key]
        dates = sorted(aggregate["dates"])
        rows.append(
            {
                "bioguide": bioguide,
                "candidate_ids": ";".join(
                    sorted(aggregate["candidate_ids"])
                ),
                "industry": industry,
                "side": WORKBOOKS[industry]["side"],
                "cycle": cycle,
                "month": month,
                "direct_dollars": _money_text(
                    aggregate["direct_dollars"]
                ),
                "sponsor_dollars": _money_text(
                    aggregate["sponsor_dollars"]
                ),
                "n_contributions": aggregate["n_contributions"],
                "first_date": dates[0] if dates else "",
                "last_date": dates[-1] if dates else "",
            }
        )
    return rows


def _total_output_rows(aggregates):
    rows = []
    for key in sorted(aggregates):
        aggregate = aggregates[key]
        rows.append(
            {
                "industry": aggregate["industry"],
                "side": aggregate["side"],
                "cycle": aggregate["cycle"],
                "recipient_category": aggregate["recipient_category"],
                "total_dollars": _money_text(aggregate["total_dollars"]),
                "n_rows": aggregate["n_rows"],
                "attributed_direct": _money_text(
                    aggregate["attributed_direct"]
                ),
                "attributed_sponsor": _money_text(
                    aggregate["attributed_sponsor"]
                ),
                "unattributed": _money_text(aggregate["unattributed"]),
            }
        )
    return rows


def _absent_workbook_inventory(path):
    return {
        "present": False,
        "processed": False,
        "path": path,
        "md5": None,
        "sheet_count": None,
        "detail_row_counts": {},
        "date_status": None,
        "roster": [],
        "reconciliation": {"status": "unavailable"},
    }


def _build_parser():
    local_appdata = os.environ.get("LOCALAPPDATA", "%LOCALAPPDATA%")
    default_raw = os.environ.get(
        "FEC_SPENDING_RAW",
        os.path.join(local_appdata, "Melange", "fec_spending", "raw"),
    )
    default_out = os.environ.get(
        "FEC_SPENDING_OUT",
        os.path.join(local_appdata, "Melange", "fec_spending", "derived"),
    )
    default_legislators = (
        r"C:\Users\pgoff\OneDrive\Desktop\Melange Systems\fec\legislators.csv"
    )
    parser = argparse.ArgumentParser(
        description="Ingest and attribute Mentat FEC industry spending."
    )
    parser.add_argument("--raw", default=default_raw, help="FEC workbook directory")
    parser.add_argument("--out", default=default_out, help="Derived output directory")
    parser.add_argument(
        "--legislators",
        default=default_legislators,
        help="Read-only legislators.csv candidate-id crosswalk",
    )
    return parser


def _run(args):
    raw_dir = _outside_mb_root("raw", args.raw)
    out_dir = _outside_mb_root("out", args.out)
    legislators_path = os.path.abspath(args.legislators)

    if not os.path.isdir(raw_dir):
        raise IngestError(f"raw directory not found: {raw_dir}")

    workbook_paths = {}
    missing = []
    initial_inventory = {}
    for industry, specification in WORKBOOKS.items():
        path = os.path.join(raw_dir, specification["filename"])
        workbook_paths[industry] = path
        if not os.path.isfile(path):
            missing.append(specification["filename"])
            initial_inventory[industry] = _absent_workbook_inventory(path)
            print(f"WAITING: {specification['filename']} not present in raw dir")
    if missing:
        raise IngestError("required workbooks missing: " + ", ".join(missing))
    if not os.path.isfile(legislators_path):
        raise IngestError(f"legislators.csv not found: {legislators_path}")

    fec_to_bioguide = _load_legislators(legislators_path)
    workbook_manifest = {}
    all_member_aggregates = {}
    all_total_aggregates = {}
    all_map_rows = []
    all_associated_rows = []
    industry_receipts = {}

    for industry, specification in WORKBOOKS.items():
        path = workbook_paths[industry]
        starting_md5 = _file_md5(path)
        try:
            (
                inventory,
                member_aggregates,
                total_aggregates,
                map_rows,
                associated_rows,
                receipt,
            ) = _process_workbook(
                path,
                industry,
                specification["side"],
                fec_to_bioguide,
            )
        except Exception as exc:
            if _file_md5(path) != starting_md5:
                raise IngestError(
                    f"input changed mid-run: {path} (while processing: {exc})"
                ) from exc
            raise
        if _file_md5(path) != starting_md5:
            raise IngestError(f"input changed mid-run: {path}")
        inventory["md5"] = starting_md5
        workbook_manifest[industry] = inventory
        all_member_aggregates.update(member_aggregates)
        all_total_aggregates.update(total_aggregates)
        all_map_rows.extend(map_rows)
        all_associated_rows.extend(associated_rows)
        industry_receipts[industry] = receipt

    total_attributed = sum(
        (
            receipt["direct_dollars"] + receipt["sponsor_dollars"]
            for receipt in industry_receipts.values()
        ),
        ZERO,
    )
    if total_attributed == 0:
        raise IngestError(
            "total attributed direct plus sponsor dollars is 0 across all "
            "industries; refusing to publish silent-zero artifacts"
        )

    global_receipt = _merge_receipts(industry_receipts.values())
    member_rows = _member_output_rows(all_member_aggregates)
    total_rows = _total_output_rows(all_total_aggregates)
    map_rows = sorted(
        all_map_rows,
        key=lambda row: (
            row["industry"],
            row["cycle"],
            row["committee_id"],
            row["recipient_committee"],
        ),
    )
    associated_rows = sorted(
        all_associated_rows,
        key=lambda row: (
            row["bioguide"],
            row["candidate_id"],
            row["industry"],
            row["cycle"],
        ),
    )

    generated = datetime.datetime.now(datetime.timezone.utc).isoformat()
    run_id = hashlib.sha1(
        (generated + raw_dir + out_dir).encode("utf-8")
    ).hexdigest()[:12]
    manifest_payload = {
        "vintage": VINTAGE,
        "generated": generated,
        "run_id": run_id,
        "raw_dir": raw_dir,
        "out_dir": out_dir,
        "workbooks": workbook_manifest,
        "receipts": {
            "global": _finalize_receipt(global_receipt),
            "by_industry": {
                industry: _finalize_receipt(industry_receipts[industry])
                for industry in WORKBOOKS
            },
        },
    }

    # This is intentionally the first filesystem mutation in the ingest.
    os.makedirs(out_dir, exist_ok=True)
    member_output = os.path.join(out_dir, "fec_member_attribution.csv")
    totals_output = os.path.join(out_dir, "fec_industry_cycle_totals.csv")
    map_output = os.path.join(out_dir, "fec_committee_candidate_map.csv")
    associated_output = os.path.join(out_dir, "fec_member_associated.csv")
    _write_csv_atomic(member_output, MEMBER_FIELDS, member_rows)
    _write_csv_atomic(totals_output, TOTAL_FIELDS, total_rows)
    _write_csv_atomic(map_output, MAP_FIELDS, map_rows)
    _write_csv_atomic(associated_output, ASSOCIATED_FIELDS, associated_rows)
    produced_outputs = [
        member_output,
        totals_output,
        map_output,
        associated_output,
    ]

    # The manifest is the final write: its presence certifies that every earlier
    # output for this run reached its atomic replacement point.
    manifest_payload["outputs"] = {
        os.path.basename(path): _file_md5(path) for path in produced_outputs
    }
    dump_atomic(
        manifest_payload,
        os.path.join(out_dir, "fec_spending_manifest.json"),
        indent=1,
        ensure_ascii=False,
        allow_nan=False,
    )
    print(
        f"fec spending ingest: {global_receipt['total_rows']} rows -> "
        f"{len(member_rows)} member-month rows | direct "
        f"{_money_text(global_receipt['direct_dollars'])} | sponsor "
        f"{_money_text(global_receipt['sponsor_dollars'])} | unattributed "
        f"{_money_text(global_receipt['unattributed_dollars'])}"
    )
    return 0


def main():
    # Windows limits csv.field_size_limit to a 32-bit C long.
    limit = sys.maxsize
    while True:
        try:
            csv.field_size_limit(limit)
            break
        except OverflowError:
            limit //= 2
    args = _build_parser().parse_args()
    try:
        return _run(args)
    except IngestError as exc:
        print(f"ERROR: {exc}")
        return 2
    except Exception as exc:
        print(f"ERROR: FEC spending ingest failed: {exc}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
