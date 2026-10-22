"""Project Mentat's tidy industry workbooks onto the member roster (read-only).

Decision record: audits/decisions_20261011/B_money_wiring.md (MW-1..MW-6, PM-A4).
Plan: audits/mentat_money_wiring_20261006/MENTAT_TO_MELANGE_WIRING.md.

Reads each `<Industry> Industry Spending Data.xlsx` (Mentat's published tidy
contract) and writes data/money/mentat_money_projection.json. Nothing in Mentat
is opened for writing; the store named in each About sheet is only hashed.

Rules (each one is a test in fec/tests/collect/test_import_mentat_money.py):
- Side map (MW-1) is data: data/money/mentat_side_map.json. A book the map does
  not name REFUSES the import. Context books never reach `cu` or `ind`.
- Channels stay separate raw fields (MW-2); weights are applied in the kernel
  accessor (moneyball_core.money_table), never here. Amounts are Mentat's
  `Giving` (included analysis amount) as given, never re-derived.
- Joins: a principal/authorized candidate committee (FEC ccl link) -> candidate
  -> bioguide through fec/legislators.csv fec_ids (refresh_fec_money.candidate_map).
  The member's own leadership PAC (FEC ccl link only, never `provisional`) is
  its own channel. Conduits (ActBlue, WinRed) never join: their dollars are the
  book's unattributed share. JFC rows never join. Downstream / Network /
  Beneficiary tables are not read at all.
- A Corporate-tier row from a company whose PAC is one of the 44 FEC-leg
  committees is dropped (the FEC leg already counts it).
- Freshness (MW-4), per book: admissible iff the current-cycle pas2 Source
  Freshness date and the store's last write are on or after D_last (latest
  quarterly deadline with deadline + MENTAT_FRESH_GRACE_DAYS <= today) AND the
  About-sheet store SHA-256 equals the sha256 of the store it names. A stale or
  mismatched book is dropped with a status; it never holds the run.
- Deterministic: no timestamps in the body (D_last is a rule input, recorded).
"""
from pathlib import Path
import sys, os
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import argparse, hashlib, json, math, re
from datetime import date, datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
from paths import ROOT
from moneyball_core import dump_atomic, compet_roster_hash
from collect.refresh_fec_money import committee_policy, candidate_map

SCHEMA = 'mentat-money-projection-v1'
SIDES = ('cu', 'ind', 'context')
TIER_CHANNEL = {'Corporate': 'corporate', 'Chief': 'chief',
                'Middle Management': 'middle', 'Employee': 'employee'}
CHANNELS = ('corporate', 'chief', 'middle', 'employee', 'own_lpac')
CONDUITS = {'C00401224': 'ActBlue', 'C00694323': 'WinRed'}   # D-132
CANDIDATE_LINKS = {'Principal campaign committee', 'Authorized candidate committee'}
OWN_LPAC_LINK = 'Leadership PAC linked to candidate'          # FEC ccl leadership list
_UNKEYED = object()   # add() default: "no key", distinct from a None/blank key
DEADLINES = ((1, 31), (4, 15), (7, 15), (10, 15))             # 52 U.S.C. 30104(a) quarterly; re-verify on the day relied on
# PM-A4 (decision B, 2026-10-11): Mentat's counting is frozen and unchanged; Melange records the
# H3 bitcoin double count and never consumes it (FAIRSHAKE is a super PAC; Crypto is context only).
KNOWN_CROSS_CYCLE_DUPLICATES = [{
    'book': 'Crypto', 'company': 'Coinbase', 'committee_id': 'C00835959',
    'sub_ids': ['4030420261386079658', '4030420261386079660'], 'gift_sub_id': '4020420251141540904',
    'amount': 11175000, 'cycle_counted_twice': '2025-2026',
    'disclosure': ("of which $11,175,000 is FAIRSHAKE's 2025 resale of Coinbase's Dec-2024 bitcoin gift, "
                   "counted in 2023-24")}]


def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for block in iter(lambda: f.read(1 << 20), b''):
            h.update(block)
    return h.hexdigest()


def current_cycle_label(today):
    end = today.year + (today.year % 2)
    return '%d-%d' % (end - 1, end)


def last_deadline(today, grace_days):
    """Latest quarterly FEC deadline d with d + grace <= today."""
    found = [date(y, m, d) for y in (today.year - 1, today.year) for m, d in DEADLINES
             if date(y, m, d) + timedelta(days=grace_days) <= today]
    return max(found)


def parse_day(text):
    """Source Freshness dates are RFC-1123 (`Sun, 06 Sep 2026 15:55:17 GMT`) or ISO; None if absent/unparseable."""
    if not text:
        return None
    try:
        return parsedate_to_datetime(str(text)).date()
    except (TypeError, ValueError):
        try:
            return datetime.fromisoformat(str(text)[:19]).date()
        except ValueError:
            return None


def book_status(about, freshness, d_last, cycle, store_sha=None, store_day=None):
    """MW-4 admissibility for one book -> (status, pas2_day). Missing evidence counts against the book."""
    pas2 = parse_day(next((r.get('Server Last Modified') for r in freshness
                           if r.get('Election Cycle') == cycle and r.get('FEC Dataset') == 'pas2'), None))
    claimed = (about.get('Source DB SHA-256') or '').strip().lower()
    if not claimed or store_sha is None or store_sha != claimed:
        return 'mentat_hash_mismatch', pas2
    if pas2 is None or pas2 < d_last or store_day is None or store_day < d_last:
        return 'mentat_stale', pas2
    return 'ok', pas2


def load_side_map(path):
    side = json.loads(Path(path).read_text(encoding='utf-8'))
    bad = {k: v for k, v in side.items() if v not in SIDES}
    if bad:
        raise ValueError('side map: unknown side(s) %r' % bad)
    return side


def pac44_companies(groups, affiliations_path, supplement_path):
    """Companies whose PAC is one of the 44 FEC-leg committees (Mentat's verified affiliations + Melange supplement)."""
    out = {}
    for r in json.loads(Path(affiliations_path).read_text(encoding='utf-8')).get('records', []):
        if r.get('entity_id') in groups:
            out.setdefault(r['company'], set()).add(r['entity_id'])
    for company, ids in json.loads(Path(supplement_path).read_text(encoding='utf-8')).get('companies', {}).items():
        unknown = set(ids) - set(groups)
        if unknown:
            raise ValueError('company PAC supplement names committees outside the 44: %r' % sorted(unknown))
        out.setdefault(company, set()).update(ids)
    return {k: sorted(v) for k, v in out.items()}


class GivingInvalid(ValueError):
    """A Recipient Facts `Giving` cell that is not a finite number: the book is dropped, the run is not."""


def project_book(rf_rows, sponsor_rows, cand_map, roster_bios, pac44, stats=None):
    """One book's Recipient Facts -> (members, accounting). members: {bio: {cycle: {channel: $}}}.
    A blank Giving counts as $0 (tallied in stats['n_blank_giving']); a non-numeric one raises GivingInvalid."""
    links = {}
    for r in sponsor_rows:
        text = ' '.join(str(v) for v in r.values() if v is not None).lower()
        if r.get('Link Source') != 'FEC ccl' or 'provisional' in text:
            continue
        links.setdefault((r.get('Election Cycle'), r.get('Committee ID'), r.get('Relationship')), set()).add(r.get('Candidate ID'))
    members, acct = {}, {}

    def add(bucket, cycle, amount, key=_UNKEYED):
        keyed = key is not _UNKEYED   # a None/blank id is still a keyed bucket: never mix float and dict, never a None dict key
        b = acct.setdefault(cycle, {}).setdefault(bucket, {} if keyed else 0.0)
        if keyed:
            key = '(blank)' if key in (None, '') else key
            b[key] = round(b.get(key, 0.0) + amount, 2)
        else:
            acct[cycle][bucket] = round(b + amount, 2)

    for n, r in enumerate(rf_rows, 1):
        cycle, cid, raw = r.get('Election Cycle'), r.get('Committee ID'), r.get('Giving')
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            amount = 0.0
            if stats is not None:
                stats['n_blank_giving'] = stats.get('n_blank_giving', 0) + 1
        else:
            try:
                amount = float(raw)
                if not math.isfinite(amount):
                    raise ValueError
            except (TypeError, ValueError):
                raise GivingInvalid('Recipient Facts data row %d (cycle %r, committee %r): non-numeric Giving %r' % (n, cycle, cid, raw))
        tier, category = r.get('Contributor Tier'), r.get('Recipient Category')
        if tier not in TIER_CHANNEL:
            raise ValueError('Recipient Facts: unknown contributor tier %r' % tier)
        add('book_total', cycle, amount)
        if cid in CONDUITS or category == 'Conduit':
            add('unattributed_conduit', cycle, amount); continue
        if category == 'JFC' or r.get('Committee Designation Code') == 'J':
            add('excluded_jfc', cycle, amount); continue
        if category == 'Candidate' and r.get('Committee Designation Code') in ('P', 'A') \
                and r.get('Committee Relationship') in CANDIDATE_LINKS:
            channel = TIER_CHANNEL[tier]
            if channel == 'corporate' and r.get('Company') in pac44:
                add('dropped_fec_leg_duplicate', cycle, amount); continue
            cands = set()
            for rel in CANDIDATE_LINKS:
                cands |= links.get((cycle, cid, rel), set())
        elif category == 'Leadership PAC' and r.get('Committee Relationship') == OWN_LPAC_LINK:
            channel, cands = 'own_lpac', links.get((cycle, cid, OWN_LPAC_LINK), set())
        else:
            add('other_recipients', cycle, amount); continue
        bios = {cand_map.get(c) for c in cands}
        if not cands:
            add('unlinked_committee', cycle, amount, cid); continue
        if len(bios) != 1:
            add('ambiguous_committee', cycle, amount, cid); continue
        bio = bios.pop()
        if bio is None or bio not in roster_bios:
            add('unmatched_candidate', cycle, amount, sorted(cands)[0]); continue
        slot = members.setdefault(bio, {}).setdefault(cycle, {})
        slot[channel] = round(slot.get(channel, 0.0) + amount, 2)
        add('joined_' + channel, cycle, amount)
    return members, acct


def sheet_rows(wb, base):
    """Dict rows of a tidy sheet, following the writer's per-cycle split (`<base> <year>`)."""
    names = [n for n in wb.sheetnames if n == base or re.fullmatch(re.escape(base) + r' \d{4}', n)]
    for name in names:
        it = wb[name].iter_rows(values_only=True)
        header = next(it, None) or ()
        for row in it:
            if any(v is not None for v in row):
                yield dict(zip(header, row))


def discover(dirs):
    found = {}
    for d in dirs:
        for p in sorted(Path(d).glob('* Industry Spending Data.xlsx')) if d and Path(d).is_dir() else []:
            found.setdefault(p.name[:-len(' Industry Spending Data.xlsx')], []).append(p)
    return found


def admit_books(found, side):
    unknown = sorted(set(found) - set(side))
    if unknown:
        raise ValueError('Mentat book(s) not in the side map: %r (rule MW-1: an unknown book refuses the import)' % unknown)
    return found


def build(root, today, grace, dirs, mentat_inputs):
    import openpyxl
    root = Path(root)
    side_path = root / 'data/money/mentat_side_map.json'
    supplement = root / 'data/money/mentat_company_pacs.json'
    side = load_side_map(side_path)
    groups = committee_policy(root / 'fec/attic/full_congestion.py')
    roster = json.loads((root / 'data/engine/congestion.json').read_text(encoding='utf-8'))
    roster_bios = {r['bioguide'] for r in roster if r.get('bioguide')}
    cand_map = candidate_map(root / 'fec/legislators.csv')
    affiliations = Path(mentat_inputs) / 'verified_company_affiliations.json'
    pac44 = pac44_companies(groups, affiliations, supplement)
    cycle, d_last = current_cycle_label(today), last_deadline(today, grace)
    found = admit_books(discover(dirs), side)
    books, members = {}, {}
    for book in sorted(side):
        paths = found.get(book, [])
        entry = {'side': side[book]}
        books[book] = entry
        if len(paths) > 1:
            entry.update(status='mentat_ambiguous', paths=[str(p) for p in paths]); continue
        if not paths:
            entry['status'] = 'mentat_missing'; continue
        path = paths[0]
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            about = {r.get('Key'): r.get('Value') for r in sheet_rows(wb, 'About')}
            if about.get('Industry') != book:
                raise ValueError('%s: About Industry %r != file book %r' % (path, about.get('Industry'), book))
            freshness = list(sheet_rows(wb, 'Source Freshness'))
            store = Path(about.get('Source DB Path') or '')
            store_ok = store.is_file()
            store_sha = sha256(store) if store_ok else None
            store_day = datetime.fromtimestamp(store.stat().st_mtime, timezone.utc).date() if store_ok else None
            status, pas2 = book_status(about, freshness, d_last, cycle, store_sha, store_day)
            entry.update(status=status, workbook=str(path), workbook_sha256=sha256(path),
                         store_path=str(store), store_sha256=store_sha, store_last_write=str(store_day) if store_day else None,
                         pas2_last_modified=str(pas2) if pas2 else None, generated_at=about.get('Generated At'),
                         cycles=json.loads(about.get('Cycles') or '[]'))
            if status != 'ok':
                continue
            stats = {'n_blank_giving': 0}
            try:
                book_members, acct = project_book(sheet_rows(wb, 'Recipient Facts'), list(sheet_rows(wb, 'Committee Sponsors')),
                                                  cand_map, roster_bios, pac44, stats)
            except GivingInvalid as e:
                entry.update(status='mentat_giving_invalid', giving_error=str(e)); continue
            entry['n_blank_giving'] = stats['n_blank_giving']
        finally:
            wb.close()
        entry['accounting'] = acct
        for bio, cycles in book_members.items():
            members.setdefault(bio, {})[book] = cycles
    header = {'schema': SCHEMA, 'current_cycle': cycle, 'd_last': str(d_last), 'fresh_grace_days': grace,
              'side_map': side, 'side_map_sha256': sha256(side_path),
              'company_pacs_sha256': sha256(supplement), 'affiliations_sha256': sha256(affiliations),
              'fec_leg_company_pacs': pac44, 'roster_sha256': compet_roster_hash(roster),
              'channels': list(CHANNELS), 'books': books,
              'known_cross_cycle_duplicates': KNOWN_CROSS_CYCLE_DUPLICATES,
              'not_read': 'Downstream Activity, Network Candidate Activity, Beneficiary Bridge, Allocation Estimates (MW-2 weight 0)'}
    return {'header': header, 'members': {b: members[b] for b in sorted(members)}}


DEFAULT_DIRS = (r'C:\Users\pgoff\OneDrive\Desktop\Mentat Complete Files', r'C:\melange-data\mentat_tidy_staging')


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument('--root', type=Path, default=Path(ROOT))
    p.add_argument('--out', type=Path, default=None, help='default <root>/data/money/mentat_money_projection.json')
    p.add_argument('--today', default=None, help='YYYY-MM-DD (tests / replays); default = UTC today')
    p.add_argument('--mentat-inputs', default=r'C:\Users\pgoff\Mentat\mentat\inputs')
    a = p.parse_args(argv)
    today = date.fromisoformat(a.today) if a.today else datetime.now(timezone.utc).date()
    grace = int(os.environ.get('MENTAT_FRESH_GRACE_DAYS', '14'))   # MW-4 grace (arbitrary, decision B); deadlines are statutory
    dirs = [d for d in os.environ.get('MENTAT_TIDY_DIRS', os.pathsep.join(DEFAULT_DIRS)).split(os.pathsep) if d]
    doc = build(a.root, today, grace, dirs, a.mentat_inputs)
    out = a.out or a.root / 'data/money/mentat_money_projection.json'
    dump_atomic(doc, str(out), indent=1, sort_keys=True)
    st = {b: e['status'] for b, e in doc['header']['books'].items()}
    print('mentat projection: %d members, books %s, D_last %s -> %s' % (len(doc['members']), st, doc['header']['d_last'], out))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
