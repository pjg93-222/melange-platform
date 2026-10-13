# -*- coding: utf-8 -*-
"""openpyxl authors for the three chain steps that wrote workbooks through @oai/artifact-tool (D-178).

MELANGE_XLSX_WRITER selects the author for all three: ``artifact`` (default, unchanged) runs the Node helpers;
``openpyxl`` runs the functions below, in process, with no Node and no proprietary runtime.

Each function stands in for one helper and leaves the same hash-bound receipt behind, so the callers' packaging
and verification code does not change:

  author_cell_sync     workbook_cell_sync.cjs          sync_stance_g (step 39)
  author_honors_views  honors_workbook.cjs             honors_workbook publish (steps 28, 40)
  author_sentinel      sync_sentinel_bill_ledger_artifact.cjs   sync_sentinel_bill_ledger (step 14)

What artifact-tool did that openpyxl cannot, and how each is handled (see ops/compute/xlsx_content_diff.py for the
parity proof):
  * recalculation + the 533-member score/rank check in workbook_cell_sync.cjs -- openpyxl has no formula engine, so
    the proof records ``scores_and_ranks_verified: null`` and ``recalc: not-run``; the check moves to the
    LibreOffice ``office-calc`` job (design F1: Python, not the workbook, is the rank authority).
  * autofitColumns() -- approximated from the text length (widths are presentation, reported not gated).
  * PNG previews (render()) -- not produced.
  * cached results -- artifact-tool left a computed value beside the formulas it authored; openpyxl leaves none (the
    workbook holds no caches today and Excel/LibreOffice compute them on open), so the comparator reports them apart.
Floats are written as the shortest text that reads back as the same double (openpyxl's own '%.16g' drops the last
digit of about one double in ten), so numbers match artifact-tool's bit for bit.
"""
from __future__ import annotations

import hashlib
import json
import math
import os
import re
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import openpyxl.cell._writer as _cell_writer
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

WRITERS = ('artifact', 'openpyxl')
AUTHOR_MODULE = Path(__file__).resolve()      # the receipts bind this file where the Node helper's hash used to go


def xlsx_writer():
    """MELANGE_XLSX_WRITER=artifact|openpyxl; anything else is an error, never a silent fallback."""
    value = (os.environ.get('MELANGE_XLSX_WRITER') or 'artifact').strip().lower()
    if value not in WRITERS:
        raise ValueError('MELANGE_XLSX_WRITER must be one of %s, got %r' % ('|'.join(WRITERS), value))
    return value


def _sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


_safe_string = _cell_writer.safe_string


def _exact_number(value):
    """A number as JavaScript prints it: openpyxl's own '%.16g' drops the 17th digit of about one double in ten,
    and artifact-tool wrote the shortest text that reads back as the same double."""
    if isinstance(value, float) and math.isfinite(value):
        if value == 0:
            return '0'
        return str(int(value)) if value == int(value) and abs(value) < 1e21 else repr(value)
    return _safe_string(value)


@contextmanager
def exact_numbers():
    """Within the block openpyxl writes every float exactly (the sheet XML is serialised at save time)."""
    _cell_writer.safe_string = _exact_number
    try:
        yield
    finally:
        _cell_writer.safe_string = _safe_string


ISO_ZONED = re.compile(r'^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(\.\d+)?(Z|[+-]\d{2}:\d{2})$')
STRING_MAX = 32767      # Excel's cell-text ceiling; openpyxl truncates there, artifact-tool does not


def _coerce(value):
    """What artifact-tool's range.values does to a string: a zoned ISO-8601 timestamp becomes a UTC date cell."""
    if isinstance(value, str) and ISO_ZONED.match(value):
        try:
            return datetime.fromisoformat(value.replace('Z', '+00:00')).astimezone(timezone.utc).replace(tzinfo=None)
        except ValueError:
            pass
    return value


def _assign(cell, value):
    """artifact-tool's range.values semantics: '=' text is a formula, a zoned timestamp is a date cell with no
    number format, and long text is kept whole (the two cells over Excel's limit are in the ledger today)."""
    value = _coerce(value)
    if isinstance(value, str) and len(value) > STRING_MAX and not value.startswith('='):
        cell._value, cell.data_type = value, 's'
    else:
        cell.value = value
    if isinstance(value, datetime):
        cell.number_format = 'General'
    return cell


def _put(ws, row, column, value):
    cell = ws.cell(row=row, column=column)
    return cell if value is None else _assign(cell, value)


def _workbook():
    wb = openpyxl.Workbook(iso_dates=True)       # date cells are written t="d", as artifact-tool writes them
    wb.remove(wb.active)
    return wb


def _argb(color):
    return 'FF' + color.lstrip('#').upper()


def _fill(color):
    return PatternFill('solid', fgColor=_argb(color))


def _font(**kw):
    kw.setdefault('name', 'Calibri')
    kw.setdefault('size', 11)
    if 'color' in kw:
        kw['color'] = _argb(kw['color'])
    return Font(**kw)


def _paint(ws, rows, columns, *, font=None, fill=None, alignment=None, border=None, number_format=None):
    for r in rows:
        for c in columns:
            cell = ws.cell(row=r, column=c)
            if font is not None:
                cell.font = font
            if fill is not None:
                cell.fill = fill
            if alignment is not None:
                cell.alignment = alignment
            if border is not None:
                cell.border = border
            if number_format is not None:
                cell.number_format = number_format


def _column_index(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


# --------------------------------------------------------------------------- sync_stance_g (workbook_cell_sync)

def author_cell_sync(plan_path):
    """Write exactly the changed cells; merge_cells() copies those cells into the package and nothing else.

    The Node helper also rebuilt four whole views to recalculate and compare 533 scores and ranks. That check needs a
    formula engine, so it is recorded as not run rather than faked.
    """
    plan_path = Path(plan_path)
    plan = json.loads(plan_path.read_text(encoding='utf-8'))
    wb = _workbook()
    sheets = {}
    for change in plan['changes']:
        ws = sheets.get(change['sheet']) or sheets.setdefault(change['sheet'], wb.create_sheet(change['sheet']))
        cell = ws[change['cell']]
        if change['value'] is not None:
            _assign(cell, change['value'])       # a '=' string IS a formula here, as in the workbook the change came from
    output = Path(plan['output'])
    with exact_numbers():
        wb.save(output)
    proof = {'schema': 'workbook-cell-author-v1', 'status': 'PASS', 'engine': 'openpyxl',
             'recalc': 'not-run', 'scores_and_ranks_verified': None,
             'authored': len(plan['changes']), 'output_sha256': _sha(output),
             'plan_sha256': _sha(plan_path), 'helper_sha256': _sha(AUTHOR_MODULE)}
    Path(str(output) + '.report.json').write_text(json.dumps(proof), encoding='utf-8')
    return proof


# --------------------------------------------------------------------------- honors_workbook

def author_honors_views(plan_path):
    """The three recognition views with the styling honors_workbook.cjs gave them."""
    plan_path = Path(plan_path)
    plan = json.loads(plan_path.read_text(encoding='utf-8'))
    wb = _workbook()
    for view in plan['views']:
        name, values = view['name'], view['values']
        ws = wb.create_sheet(name)
        height, width = len(values), len(values[0])
        columns, last = range(1, width + 1), get_column_letter(width)
        base = _font()
        for r, row in enumerate(values, 1):
            for c, value in enumerate(row, 1):
                _put(ws, r, c, value).font = base
        ws.merge_cells('A1:F1')
        ws.merge_cells('A2:F2')
        _paint(ws, [1], columns, fill=_fill('#163F4B'), font=_font(size=17, bold=True, color='#FFFFFF'))
        _paint(ws, [2], columns, font=_font(color='#425563'), alignment=Alignment(wrap_text=True))
        _paint(ws, [3], columns, font=_font(size=9, color='#64748B'))
        _paint(ws, [4], columns, fill=_fill('#DCECEF'), font=_font(bold=True), alignment=Alignment(wrap_text=True))
        _paint(ws, range(5, height + 1), columns, alignment=Alignment(wrap_text=True))
        for r, h in {1: 32, 2: 42, 3: 22, 4: 44}.items():
            ws.row_dimensions[r].height = h
        for r in range(5, height + 1):
            ws.row_dimensions[r].height = 30
        for c in columns:
            ws.column_dimensions[get_column_letter(c)].width = 23
        ws.column_dimensions['A'].width = 28
        detail = name == 'Recognition Detail'
        if not detail:
            for letter, w in {'B': 25, 'C': 13, 'D': 26, 'E': 24}.items():
                ws.column_dimensions[letter].width = w
        body = range(5, height + 1)
        if detail:
            ws.column_dimensions['N'].width = 38
            _paint(ws, body, range(5, 9), number_format='0')
            _paint(ws, body, range(9, 11), number_format='0.00')
        if name == 'Overall Standing':
            ws.column_dimensions['W'].width = 38
            ws.column_dimensions['X'].width = 38
            _paint(ws, body, [7], number_format='0.00')
            _paint(ws, body, range(8, 22), number_format='0')
            _paint(ws, body, [22], number_format='0.00')
        ws.freeze_panes = '%s5' % get_column_letter((2 if detail else 4) + 1)
        ws.sheet_view.tabSelected = False
    output = Path(plan['output'])
    with exact_numbers():
        wb.save(output)
    receipt = {'schema': 'honors-workbook-author-v1', 'status': 'PASS', 'engine': 'openpyxl',
               'generation_id': plan['generation_id'], 'plan_sha256': _sha(plan_path),
               'helper_sha256': _sha(AUTHOR_MODULE), 'output_sha256': _sha(output)}
    Path(str(output) + '.receipt.json').write_text(json.dumps(receipt), encoding='utf-8')
    return receipt


# --------------------------------------------------------------------------- sync_sentinel_bill_ledger

ERROR_PATTERN = re.compile(r'#REF!|#DIV/0!|#VALUE!|#NAME\?|#N/A')
INTERACTION_HEADERS = ['Section', 'Role', 'Member', 'Name', 'Title', 'Sponsor', 'CosponsorCount', 'Committee',
                       'LastAction', 'Congress', 'DecayWeight', 'DetailLink', 'bill_summary']
LEDGER_WIDE_COLUMNS = ('A', 'B', 'D', 'I', 'J', 'N', 'O', 'T', 'U', 'V', 'W', 'X', 'Z', 'AA')
THIN_EDGE = Side(style='thin', color=_argb('#D9E2F3'))


def _js_string(value):
    """String(value ?? "") as JavaScript spells it, for the interaction de-duplication key."""
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value)


def _interaction_key(record):
    values = record['values']
    return '\x1f'.join([_js_string(record.get('_bioguide')), _js_string(values[0]), _js_string(values[1]),
                        _js_string(values[2]), _js_string(values[3]), _js_string(values[9])]).lower()


def _autofit(ws, columns, rows):
    """artifact-tool measured the rendered text; this approximates that from the character count (presentation only)."""
    for c in columns:
        longest = max((len(str(cell.value)) for cell in ws[get_column_letter(c)][:rows] if cell.value is not None), default=0)
        ws.column_dimensions[get_column_letter(c)].width = round(min(max(0.88 * longest + 0.6, 6.0), 72.0), 2)


def author_sentinel(plan_path, report_path):
    """The bill ledger plus the interaction tabs, with the formatting sync_sentinel_bill_ledger_artifact.cjs applied."""
    plan_path, report_path = Path(plan_path), Path(report_path)
    plan = json.loads(plan_path.read_bytes().decode('utf-8'))      # bytes, not read_text: the plan hash is of the file
    if plan.get('schema') != 'sentinel-bill-workbook-plan-v1' or not isinstance(plan.get('ledger_rows'), list):
        raise ValueError('invalid plan')
    wb = _workbook()
    rows = plan['ledger_rows']
    ledger = wb.create_sheet(plan['ledger_sheet'])
    width = len(rows[0])
    ledger_range = 'A1:%s%d' % (get_column_letter(width), len(rows))
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row, 1):
            _put(ledger, r, c, value)
    ledger.sheet_view.showGridLines = False
    ledger.freeze_panes = 'A2'
    header_border = Border(left=THIN_EDGE, right=THIN_EDGE, top=THIN_EDGE, bottom=THIN_EDGE)
    header = dict(fill=_fill('#17365D'), border=header_border, alignment=Alignment(wrap_text=True),
                  font=_font(bold=True, color='#FFFFFF'))
    _paint(ledger, [1], range(1, width + 1), **header)
    _autofit(ledger, range(1, width + 1), 200)
    for letter in LEDGER_WIDE_COLUMNS:
        ledger.column_dimensions[letter].width = 72 if letter == 'X' else 24
    for r in range(2, len(rows) + 1):
        ledger['X%d' % r].alignment = Alignment(wrap_text=True)

    grouped = {}
    for record in plan['accepted_interactions']:
        grouped.setdefault(record['sheet'], []).append(record)
    names = ['Original Co-Sponsor Interaction', 'Co-Sponsor Interactions']
    if 'Sponsor Interactions' in grouped:
        names.append('Sponsor Interactions')
    appended, inspections = {}, []
    for name in names:
        ws = wb.create_sheet(name)
        seen, additions = set(), []
        for record in grouped.get(name, []):
            key = _interaction_key(record)
            if key not in seen:
                seen.add(key)
                additions.append(record['values'])
        for c, text in enumerate(INTERACTION_HEADERS, 1):
            ws.cell(row=1, column=c, value=text)
        for r, values in enumerate(additions, 2):
            for c, value in enumerate(values, 1):
                _put(ws, r, c, value)
        last = 1 + len(additions)
        _paint(ws, [1], range(1, 14), fill=_fill('#2F75B5'), border=header_border, alignment=Alignment(wrap_text=True),
               font=_font(bold=True, color='#FFFFFF'))
        ws.freeze_panes = 'A2'
        _autofit(ws, range(1, 14), 200)
        for letter, w in {'A': 28, 'C': 24, 'E': 34, 'F': 34, 'M': 72}.items():
            ws.column_dimensions[letter].width = w
        _paint(ws, range(2, last + 1), range(1, 14), alignment=Alignment(vertical='top', wrap_text=False))
        for r in range(2, last + 1):
            ws.row_dimensions[r].height = 18
        appended[name] = len(additions)
    scanned = [(plan['ledger_sheet'], ledger_range)] + [(n, 'A1:M%d' % (1 + appended[n])) for n in names]
    errors = 0
    for sheet, _ in scanned:
        for row in wb[sheet].iter_rows(values_only=True):
            errors += sum(1 for v in row if isinstance(v, str) and ERROR_PATTERN.search(v))
        digest = hashlib.sha256(json.dumps([list(r) for r in wb[sheet].iter_rows(values_only=True)],
                                           ensure_ascii=False, default=str).encode('utf-8')).hexdigest()
        inspections.append({'sheet': sheet, 'range': dict(scanned)[sheet], 'sha256': digest})
    if errors:
        raise RuntimeError('mutated bill ranges contain %d formula error matches' % errors)
    output = Path(plan['output_workbook']).resolve()
    with exact_numbers():
        wb.save(output)
    report = {'schema': 'sentinel-bill-workbook-artifact-report-v1', 'status': 'PASS', 'engine': 'openpyxl',
              'input_plan_sha256': _sha(plan_path), 'output_sha256': _sha(output),
              'ledger_range': ledger_range, 'appended': appended, 'targeted_formula_error_count': 0,
              'formula_error_scan_scope': 'all generated bill ranges', 'inspections': inspections, 'previews': [],
              'output': str(output)}
    report_path.write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8')
    return report
