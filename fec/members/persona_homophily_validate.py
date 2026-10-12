# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT
"""PERSONA-HOMOPHILY VALIDATION GATE (user-directed 2026-07-03: "lets do it, if the validity changes
after all the data comes back lets make that change"). Tests whether interpersonal-persona
compatibility predicts real cosponsorship ties, cross-party (party confound removed). Four
compatibility models (similarity / IPC-complementarity / communion-match / both-warm). If ANY clears
|r|>=0.05 on the settled corpus, persona-homophily edge-weighting in cosponsor_route.py is EARNED and
can be proposed for ratification. Until then it stays OFF. Zero API cost. Re-run after each corpus
refresh: `python fec/members/persona_homophily_validate.py`.
Result 2026-07-03 (partial corpus, 494 covered members): all models |r|<=0.015 -> UNEARNED, held."""
import glob, json, math, os, sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
import openpyxl
import moneyball_core as MC

B = MB_ROOT
ACP = B + r"\data\members\agency_communion.json"
if not os.path.exists(ACP):
    print("data/members/agency_communion.json missing — run agency_communion_llm.py first"); sys.exit(1)
AC = {m: d for m, d in json.load(open(ACP))['members'].items() if d.get('n', 0) >= 15}
if len(AC) < 30:
    print("only %d covered members — too few to validate" % len(AC)); sys.exit(0)
ag = [d['agency'] for d in AC.values()]; co = [d['communion'] for d in AC.values()]
ma, sa = sum(ag)/len(ag), (sum((x-sum(ag)/len(ag))**2 for x in ag)/len(ag))**.5 or 1
mc, sc = sum(co)/len(co), (sum((x-sum(co)/len(co))**2 for x in co)/len(co))**.5 or 1
Z = {m: ((d['agency']-ma)/sa, (d['communion']-mc)/sc) for m, d in AC.items()}
party = MC.party()
net = max([p for p in glob.glob(B+r"\congress_cosponsor_network_*.xlsx") if 'AUG' not in os.path.basename(p)], key=os.path.getmtime)
ne = openpyxl.load_workbook(net, data_only=True)['Edges']; b2n = MC.bio2name()
W = {}
for r in range(2, ne.max_row+1):
    a, b, w = ne.cell(r,1).value, ne.cell(r,2).value, ne.cell(r,3).value
    a, b = b2n.get(a), b2n.get(b)
    if a in Z and b in Z and a != b and isinstance(w, (int, float)): W[(min(a,b),max(a,b))] = w
MODELS = {
    'similarity (match both)': lambda ax,cx,ay,cy: -(abs(ax-ay)+abs(cx-cy)),
    'complementarity (IPC)':   lambda ax,cx,ay,cy: -abs(cx-cy) - (ax*ay),
    'communion-match only':    lambda ax,cx,ay,cy: -abs(cx-cy),
    'both-warm (additive)':    lambda ax,cx,ay,cy: cx+cy,
}
def corr(pairs):
    xs=[p[0] for p in pairs]; ys=[p[1] for p in pairs]; n=len(xs)
    if n<50: return float('nan')
    mx,my=sum(xs)/n,sum(ys)/n
    sx=math.sqrt(sum((x-mx)**2 for x in xs)); sy=math.sqrt(sum((y-my)**2 for y in ys))
    return sum((x-mx)*(y-my) for x,y in zip(xs,ys))/max(1e-12,sx*sy)
mems=list(Z); cross=[]
for i in range(len(mems)):
    for j in range(i+1,len(mems)):
        a,b=mems[i],mems[j]
        if party.get(a)==party.get(b): continue
        ax,cx=Z[a]; ay,cy=Z[b]; cross.append((ax,cx,ay,cy,W.get((min(a,b),max(a,b)),0.0)))
print("covered members: %d | cross-party dyads: %d" % (len(Z), len(cross)))
earned=False
for name,f in MODELS.items():
    r=corr([(f(ax,cx,ay,cy),w) for ax,cx,ay,cy,w in cross])
    hit=abs(r)>=0.05
    earned=earned or hit
    print("  %-26s r=%+.3f %s" % (name, r, '<-- EARNED' if hit else ''))
print("\nVERDICT:", "SIGNAL — propose persona-homophily routing for ratification" if earned
      else "UNEARNED — persona-homophily stays OFF (re-run after next corpus refresh)")
