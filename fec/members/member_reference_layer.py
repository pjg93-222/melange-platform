# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root
import html as _html
"""SIGNED MEMBER-REFERENCE LAYER (user-directed 2026-07-04: "see who they are referencing... in press
articles and transcripts or re-tweet their colleagues... either negatively or positively" + "not
formalities, like actual substantive mentions").

Extraction: every reference to ANOTHER roster member inside a member's OWN words — floor speech,
hearings, interviews (speeches table), press headlines, and X mentions/RTs. Attribution uses the
QA2-hardened rules (accent transliteration, first-name compatibility, no unique-surname guessing
without a title). SELF-references excluded.

FORMALITY DISCIPLINE (the user's explicit rule): chamber courtesy is NOISE. Two stages:
  1. Free regex prefilter drops pure ritual ("I yield to...", "I thank the gentleman from...",
     recognition/colloquy boilerplate) when the window carries no evaluative content.
  2. Fable classifies every surviving mention: positive / negative / neutral-substantive / FORMALITY
     — and formality is DISCARDED. Only substantive references enter the layer.
Cost control: per-dyad-per-source cap (newest first) for the scored pass; raw counts kept for all.
Modes:
  python fec/members/member_reference_layer.py --probe          # free: extraction + prefilter volume + cost estimate
  python fec/members/member_reference_layer.py --score [--go]   # Fable valence pass (cost-gated), writes member_references.json
  python fec/members/member_reference_layer.py --rebuild        # FREE: re-extract + re-aggregate from the existing cache only
                                                    # (new/uncached mentions count in n_raw, never spend)
  python fec/members/member_reference_layer.py --verify         # write member_reference_verify.csv — every edge with its
                                                    # evidence quotes, for the user's person/valence eyeball pass
Output: member_references.json {edges:[{a,b,n_raw,n_scored,pos,neg,neu,net,sources{},examples[]}]}
2026-07-04 resolver upgrade (user: "make sure the person is being mentioned"): nickname classes
(kernel NICK_SETS: bernie~bernard), chamber-scoped title pools ("senator johnson" resolves in the
SENATE surname pool), single-alternation bare-name scan; examples now attach to EVERY scored class
(pos/neg preferred, neutral fills) so the desk can show the phrase on every reference row."""
import csv, hashlib, io, json, os, re, sqlite3, sys, time, collections, unicodedata
import moneyball_core as MBC   # kernel: name_variants (nickname classes), dump_atomic later
from moneyball_core import dump_atomic
from press_stream import REFRESH_MENTIONS_PATH, load_tone_rows
from members import reference_filters as RF   # pure admission rules: no data, no side effects
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
B = MB_ROOT
DYAD_CAP = int(os.environ.get('REF_DYAD_CAP', '4'))     # scored contexts per (a->b, source)
# ADMISSION FILTERS (Paul-approved 2026-09-19; measured in audits/laundry_list_closeout_20260918
# FINDINGS.md B6/B4/B5). A filtered window is out before n_raw and before the scoring queue, so
# --score never pays for one, and the census lands in the artifact's 'admission_filters' block.
# Each DEFAULT is the admitted-set rule; rollback is a dial flip, never a revert.
REF_PRESS_OWN_DOMAIN = os.environ.get('REF_PRESS_OWN_DOMAIN', '1') == '1'   # PRESS pass: a headline is the speaker's own words only on their own .house.gov/.senate.gov domain ('0' = any domain, the old behaviour)
REF_ROSTER_FILTER = os.environ.get('REF_ROSTER_FILTER', '1') == '1'         # roll-call rosters and signature blocks name colleagues without referencing them ('0' = counted as references, the old behaviour)
REF_IDENTICAL_ONCE = os.environ.get('REF_IDENTICAL_ONCE', '1') == '1'       # per edge, windows identical after whitespace normalisation count once in n_scored/pos/neg/neu ('0' = once per copy, the old behaviour)
AF_STATE = {'REF_PRESS_OWN_DOMAIN': REF_PRESS_OWN_DOMAIN, 'REF_ROSTER_FILTER': REF_ROSTER_FILTER,
            'REF_IDENTICAL_ONCE': REF_IDENTICAL_ONCE}
FILTERED = collections.Counter()   # what each admission filter excluded, reported in the artifact
MODEL = 'claude-fable-5'
CACHE = B + r'\data\members\member_reference_cache_fable.jsonl'

C = {m['member']: m for m in json.load(open(B + r'\data\engine\congestion.json'))}

def norm(s):
    s = unicodedata.normalize('NFD', s); s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z ]", '', s.lower()).strip()

# reference patterns: title+surname (resolved only if unique IN THE TITLE'S CHAMBER) OR full
# first+last (nickname-variant aware: 'bernie sanders' -> Bernard Sanders)
SUR = collections.defaultdict(list)
for m in C: SUR[norm(m).split()[-1]].append(m)
SUR_SEN = collections.defaultdict(list); SUR_HOU = collections.defaultdict(list)
for m in C:
    (SUR_SEN if (C[m].get('chamber') == 'Senate') else SUR_HOU)[norm(m).split()[-1]].append(m)
FULLSET = {}
for m in C:
    p = norm(m).split()
    if len(p) >= 2:
        for v in MBC.name_variants(p[0]): FULLSET.setdefault((v, p[-1]), m)
# one-pass bare-name scan: single alternation over every (variant-first, last) pair
BARE2MEM = {fn + ' ' + ln: mem for (fn, ln), mem in FULLSET.items()}
BARE_RE = re.compile(r'\b(?:' + '|'.join(sorted((re.escape(fn) + r'\s+' + re.escape(ln)
                     for fn, ln in FULLSET), key=len, reverse=True)) + r')\b')
TITLE_RE = r'(?:senator|sen\.|representative|rep\.|congressman|congresswoman|speaker|leader|chairman|chairwoman|ranking member)'
SURNAMES = sorted({norm(m).split()[-1] for m in C}, key=len, reverse=True)
REF_RE = re.compile(r'\b(' + TITLE_RE + r')\s+([a-z]+(?:\s[a-z]+)?)|\b([a-z]+)\s([a-z\'\-]+)\b', re.I)

# stage-1 ritual prefilter: window is PURE formality if it matches ritual AND has no evaluative token
RITUAL = re.compile(r'\b(i yield|yield(s|ing)? (to|back)|i thank (the|my|senator|rep)|thank the (gentleman|gentlewoman|gentlelady|chair)|'
                    r'(gentleman|gentlewoman|gentlelady) from|i recognize|is recognized|i join .{0,40} in (recogniz|celebrat|honor|wish)|'
                    r'colloquy|unanimous consent|i ask .{0,30} to join|pursuant to|the clerk|quorum|i reserve)\b', re.I)
EVALUATIVE = re.compile(r'\b(wrong|right|great|dangerous|disastrous|brilliant|leadership|fail|failed|courage|proud|disgrace|'
                        r'attack|defend|oppose|support|fight|betray|hero|corrupt|honest|lie|lying|truth|gut|destroy|save|'
                        r'protect|reckless|smart|effective|partner|friend|mentor|extremis|radical|principled|bipartisan)\b', re.I)

def resolve_ref(text_l, speaker):
    """Yield (member, char_pos) references found in lowercase text. Conservative: title+surname
    unique WITHIN THE TITLE'S CHAMBER ('senator johnson' -> Ron Johnson even though the House has
    more Johnsons), or full first+last (nickname-variant compatible). Never bare-surname guessing."""
    out = []
    # title + name form (chamber-scoped pool when the title implies one)
    for m in re.finditer(r'\b(' + TITLE_RE + r')\s+([a-zÀ-ɏ\'\-]+(?:\s[a-zÀ-ɏ\'\-]+)?)', text_l):
        ttl = m.group(1)
        raw_name=m.group(2)
        possessive_surname=raw_name.split()[0].endswith("'s")
        cand = norm(re.sub(r"'s\b",'',raw_name))
        parts = cand.split()
        if not parts: continue   # accent/punctuation-only capture normalizes to empty
        pool = (SUR_SEN if ttl.startswith('sen') else
                SUR_HOU if ttl.startswith(('rep', 'congress')) else SUR)
        hit = None
        if len(parts) == 1:
            cs = pool.get(parts[0], [])
            if len(cs) == 1: hit = cs[0]
        else:
            hit = FULLSET.get((parts[0], parts[-1]))
            # The two-word capture also consumes 'Warren on' or 'Scott has'.
            # Resolve the titled surname only when the next token is an explicit
            # connective/verb, never by guessing an unknown person's first name.
            if not hit and (possessive_surname or parts[1] in {'on','for','and','has','is','was','will','would','who','said',
                                        'the','to','of','with','in','as','at','about','from','that',
                                        'this','because','joins','joined','introduced','announced','supports'}):
                cs = pool.get(parts[0], [])
                if len(cs) == 1: hit = cs[0]
            if not hit:
                cs = pool.get(parts[-1], [])
                if len(cs) == 1:
                    rf = norm(cs[0]).split()[0]
                    if rf.startswith(parts[0]) or parts[0].startswith(rf) or (MBC.name_variants(parts[0]) & MBC.name_variants(rf)):
                        hit = cs[0]
        if hit and hit != speaker: out.append((hit, m.start()))
    # bare full-name form — one alternation pass over all (variant-first, last) pairs
    for m in BARE_RE.finditer(text_l):
        mem = BARE2MEM.get(re.sub(r'\s+', ' ', m.group(0)))
        if mem and mem != speaker: out.append((mem, m.start()))
    # dedupe by (member, pos//80)
    seen = set(); res = []
    for mem, pos in out:
        k = (mem, pos // 80)
        if k not in seen: seen.add(k); res.append((mem, pos))
    return res

def extract():
    """All (speaker -> referenced, source, date, window) mentions passing the ritual prefilter."""
    refs = []
    con = sqlite3.connect(B + r'\social_media.db', timeout=120)
    for spk, date, src, text in con.execute("SELECT member, date, source, text FROM speeches"):
        if spk not in C or not text: continue
        tl = ' ' + text.lower() + ' '
        for mem, pos in resolve_ref(tl, spk):
            w = text[max(0, pos - 170):pos + 200]
            if RITUAL.search(w) and not EVALUATIVE.search(w): continue   # pure formality -> drop free
            refs.append({'a': spk, 'b': mem, 'src': src, 'date': date or '', 'w': re.sub(r'\s+', ' ', w).strip()})
    con.close()
    # press headlines: member A's article naming member B
    for r in load_tone_rows():
        spk = r.get('member'); t = (r.get('title') or '')
        if spk not in C or not t: continue
        tl = ' ' + t.lower() + ' '
        for mem, pos in resolve_ref(tl, spk):
            refs.append({'a': spk, 'b': mem, 'src': 'PRESS', 'date': r.get('date') or '', 'w': t[:300],
                         'url': r.get('url') or ''})   # url: the own-domain gate below reads it (PRESS only)
    # Typed press-body relationships from the refresh.  These are reference candidates only:
    # they still pass the same ritual/Fable substantive-valence discipline below and never become
    # a stance row or attributed speech turn.
    if os.path.exists(REFRESH_MENTIONS_PATH):
        for line in io.open(REFRESH_MENTIONS_PATH,encoding='utf-8'):
            try: r=json.loads(line)
            except ValueError: continue
            spk=r.get('host_member'); mem=r.get('mentioned_member')
            if spk not in C or mem not in C or spk==mem: continue
            window=(r.get('evidence') or r.get('title') or '')[:1200]
            if RITUAL.search(window) and not EVALUATIVE.search(window): continue
            refs.append({'a':spk,'b':mem,'src':'PRESS_OFFICIAL_MENTION',
                         'date':r.get('date') or '','w':window})
    # X original-post mentions (framed text; plain RTs counted separately as endorsement events)
    try:
        import openpyxl
        con2 = sqlite3.connect(B + r'\social_media.db', timeout=120)
        h2m = {}
        bio2name = {C[m]['bioguide']: m for m in C if C[m].get('bioguide')}
        for bio, tw in con2.execute("SELECT bioguide,twitter FROM members WHERE twitter IS NOT NULL AND twitter!=''"):
            if bio in bio2name: h2m[tw.lower().lstrip('@')] = bio2name[bio]
        con2.close()
        ws = openpyxl.load_workbook(B + r'\Quorum Processed.xlsx', read_only=True)['Members']
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        im, ic, irt, idt = hdr.index('member'), hdr.index('Content'), hdr.index('is_rt'), hdr.index('dt') if 'dt' in hdr else None
        for row in ws.iter_rows(min_row=2, values_only=True):
            spk = row[im]; t = str(row[ic] or '')
            if spk not in C or not t or str(row[irt]).lower() in ('1', 'true'): continue
            for h in re.findall(r'@(\w+)', t):
                mem = h2m.get(h.lower())
                if mem and mem != spk:
                    refs.append({'a': spk, 'b': mem, 'src': 'X', 'date': str(row[idt] or '') if idt is not None else '', 'w': t[:300]})
    except Exception as e:
        print('  (X mention pass skipped: %s)' % str(e)[:80])
    # Admission: a third-party press headline and a roll-call/signature-block window are not
    # references. Dropped here, so they never reach n_raw, the dyad cap or the scoring queue.
    refs, _excl = RF.apply_admission_filters(refs, frozenset(SURNAMES),
                                             press_own_domain=REF_PRESS_OWN_DOMAIN,
                                             roster_filter=REF_ROSTER_FILTER)
    FILTERED.update(_excl)
    return refs

SYSTEM = """You judge how a member of Congress references a COLLEAGUE in their own words. Given a text window
where SPEAKER mentions TARGET, classify the reference:
  * positive  — praise, endorsement, credit, alliance, defending the target ("Senator X's leadership on this
    bill", "I'm proud to partner with...", co-authorship framed with approval)
  * negative  — criticism, attack, blame, mockery, opposing the target's position or character
  * neutral-substantive — a real policy interaction without clear valence ("Senator X's amendment changes
    the threshold to 3%") — substantive but neither praise nor attack
  * formality — chamber courtesy/procedure only: yields, thanks-the-gentleman, recognition, colloquy setup,
    ceremonial co-mentions with NO evaluative or substantive content. When in doubt between formality and
    neutral-substantive, choose formality.
Judge only the SPEAKER's expressed attitude toward the TARGET in this window."""
SCHEMA = {"type": "object",
          "properties": {"cls": {"type": "string", "enum": ["positive", "negative", "neutral-substantive", "formality"]},
                         "reason": {"type": "string"}},
          "required": ["cls", "reason"], "additionalProperties": False}

if '--probe' in sys.argv:
    t0 = time.time()
    refs = extract()
    by_src = collections.Counter(r['src'] for r in refs)
    dyads = collections.Counter((r['a'], r['b']) for r in refs)
    # scored set under the dyad cap
    bykey = collections.defaultdict(list)
    for r in refs: bykey[(r['a'], r['b'], r['src'])].append(r)
    n_scored = sum(min(DYAD_CAP, len(v)) for v in bykey.values())
    print('substantive-candidate references: %d (post ritual-prefilter) in %.0fs' % (len(refs), time.time() - t0))
    print('  admission filters: %s' % json.dumps(RF.admission_filters_block(AF_STATE, FILTERED), sort_keys=True))
    print('    (identical_window_repeat is counted where the scored set is built: --score / --rebuild)')
    print('  by source: %s' % dict(by_src))
    print('  directed dyads: %d | dyads with 3+ refs: %d' % (len(dyads), sum(1 for v in dyads.values() if v >= 3)))
    print('  to score under DYAD_CAP=%d: %d contexts | est Fable cost ~$%.2f' % (DYAD_CAP, n_scored, n_scored * 0.006))
    for r in refs[:6]:
        print('  eg [%s] %s -> %s | %s...' % (r['src'], r['a'].split()[-1], r['b'].split()[-1], r['w'][:90]))
    sys.exit(0)

if '--score' in sys.argv or '--rebuild' in sys.argv:
    REBUILD = '--rebuild' in sys.argv   # free path: aggregate from the existing cache, never call the API
    lock = None
    _policy_path=os.path.join(B,'data/engine/sentinel_admission_policy.json')
    _policy=json.load(open(_policy_path,encoding='utf-8')) if os.path.exists(_policy_path) else {}
    _transport=_policy.get('reference_transport','api')
    _new_model=_policy.get('reference_model',MODEL)
    if not REBUILD:
        if _transport=='api' and not os.environ.get('ANTHROPIC_API_KEY'):
            _kf = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'anthropic_key.txt')
            if os.path.exists(_kf): os.environ['ANTHROPIC_API_KEY'] = open(_kf).read().strip()
        import threading
        from concurrent.futures import ThreadPoolExecutor, as_completed
        lock = threading.Lock()
        if _transport=='api':
            import anthropic
            client = anthropic.Anthropic()
    refs = extract()
    bykey = collections.defaultdict(list)
    for r in refs: bykey[(r['a'], r['b'], r['src'])].append(r)
    target_all = []
    for k, v in bykey.items():
        v.sort(key=lambda r: r['date'], reverse=True)
        target_all.extend(v[:DYAD_CAP])
    # A repeat of a window already targeted on this edge carries the same label by construction
    # (the cache key hashes a, b, w[:400]), so it is never queued and never paid for.
    target_all, _dup_target = RF.collapse_identical(target_all, REF_IDENTICAL_ONCE)
    cache = {}
    _cache_models=set()
    if os.path.exists(CACHE):
        for ln in open(CACHE, encoding='utf-8'):
            # first-wins: the append-only cache can hold duplicate keys with conflicting LLM 'cls'
            # labels; keep the earliest so classification is deterministic (not append-order dependent).
            try:
                d = json.loads(ln)
                _cache_models.add(d.get('model',MODEL))
                if d['k'] not in cache: cache[d['k']] = d['out']
            except Exception: pass
    from members.reference_cache import key, bind_legacy, validate
    legacy_ambiguous = bind_legacy(cache, refs, os.path.join(B,'data/members/member_reference_cache_baseline_v2.json'))
    # Cache-only rebuilds must not let a new, uncached window consume the cap and
    # evict an older classified window.  Keep n_raw on the full fresh corpus,
    # report uncached windows from the intended newest-first target, but
    # aggregate the newest DYAD_CAP *cached* windows per dyad/source until a paid
    # scoring pass is explicitly authorized.
    if REBUILD:
        todo_all = []
        for v in bykey.values():
            todo_all.extend([r for r in v if key(r) in cache][:DYAD_CAP])
    else:
        todo_all = target_all
    pending = {key(r):r for r in target_all if key(r) not in cache or key(r) in legacy_ambiguous}
    batch_limit = max(1, int(_policy.get('reference_batch_limit',128)))
    todo = [] if REBUILD else sorted(pending.values(),key=lambda r:r['date'],reverse=True)[:batch_limit]
    deferred = max(0,len(pending)-len(todo))
    if REBUILD:
        _unc = sum(1 for r in target_all if key(r) not in cache)
        _target_keys = {key(r) for r in target_all}
        _fallback = sum(1 for r in todo_all if key(r) not in _target_keys)
        print('REBUILD (cache-only, $0): target set %d | cached target %d | uncached target '
              '(counted in n_raw only) %d | cached fallback contexts %d'
              % (len(target_all), len(target_all) - _unc, _unc, _fallback))
    else:
        print('scored set: %d | cached %d | to classify %d | est ~$%.2f' % (len(todo_all), len(todo_all) - len(todo), len(todo), len(todo) * 0.006))
    if todo and '--go' not in sys.argv: print('COST GATE: add --go to spend.'); sys.exit(0)
    def work(r):
        try:
            if _transport=='claude-cli':
                from classifier_runtime import classify_json
                value,provenance=classify_json(SYSTEM,"SPEAKER: %s\nTARGET: %s\nWindow: %s\n(reason under 10 words)" % (r['a'],r['b'],r['w'][:400]),SCHEMA,_new_model)
                validate(value)
                return r,value,None,provenance
            resp = client.messages.create(model=_new_model, max_tokens=400,
                output_config={"format": {"type": "json_schema", "schema": SCHEMA}},
                system=[{"type": "text", "text": SYSTEM, "cache_control": {"type": "ephemeral"}}],
                messages=[{"role": "user", "content": "SPEAKER: %s\nTARGET: %s\nWindow: %s\n(reason under 10 words)" % (r['a'], r['b'], r['w'][:400])}])
            return r, validate(json.loads(next(b.text for b in resp.content if b.type == 'text'))), None, {'model':resp.model,'response_id':resp.id,'transport':'api'}
        except Exception as e:
            return r, None, (str(e)[:100] or type(e).__name__), {}
    done = errs = 0
    error_details=[]
    if todo:
        cf = io.open(CACHE, 'a', encoding='utf-8')
        with ThreadPoolExecutor(max_workers=int(os.environ.get('REF_WORKERS', '2' if _transport=='claude-cli' else '8'))) as ex:
            for f in as_completed([ex.submit(work, r) for r in todo]):
                r, out, err, provenance = f.result()
                with lock:
                    done += 1
                    if err or not out:
                        errs += 1
                        error_details.append({'speaker':r['a'],'target':r['b'],'error':err or 'empty classification'})
                        print('Reference classification error: '+str(err),flush=True)
                    else:
                        cache[key(r)] = out;_cache_models.add(_new_model)
                        cf.write(json.dumps({'k':key(r),'out':out,'model':_new_model,'transport':_transport,'classifier':provenance},ensure_ascii=True)+'\n')
                        cf.flush()
                    if done % 500 == 0 or done == len(todo): print('  %d/%d (%d errors)' % (done, len(todo), errs), flush=True)
        cf.close()
    if not REBUILD:
        dump_atomic({'schema':'member-reference-classification-v1','status':'FAIL' if errs else 'PASS',
                     'classified':done-errs,'errors':errs,'error_details':error_details,'model':_new_model,'transport':_transport,
                     'deferred_contexts':deferred,'batch_limit':batch_limit,
                     'legacy_ambiguous_contexts_at_start':len(legacy_ambiguous),
                     'cache_key_schema':'complete-classifier-window-v2',
                     'checked_at':__import__('datetime').datetime.now(__import__('datetime').timezone.utc).isoformat()},
                    os.path.join(B,'data/members/member_reference_classification_receipt.json'))
        if errs:
            print('Reference classification failed; prior network retained; successful cache entries saved.')
            sys.exit(1)
    # A bounded refresh keeps the newest classified fallback per dyad/source.
    # Unclassified targets remain visible in n_raw and the backlog receipt.
    todo_all = [r for group in bykey.values() for r in [x for x in group if key(x) in cache][:DYAD_CAP]]
    # Identical windows on one edge are one piece of evidence: they contribute once to n_scored
    # and to the pos/neg/neu tallies. n_raw below keeps counting every raw mention.
    todo_all, _dup_scored = RF.collapse_identical(todo_all, REF_IDENTICAL_ONCE)
    FILTERED['identical_window_repeat'] = _dup_scored
    edges = collections.defaultdict(lambda: {'n_raw': 0, 'n_scored': 0, 'pos': 0, 'neg': 0, 'neu': 0, 'form': 0,
                                             'sources': collections.Counter(), 'examples': [], '_neu_ex': []})
    for r in refs: edges[(r['a'], r['b'])]['n_raw'] += 1; edges[(r['a'], r['b'])]['sources'][r['src']] += 1
    for r in todo_all:
        o = cache.get(key(r))
        if not o: continue
        e = edges[(r['a'], r['b'])]; e['n_scored'] += 1
        c = o['cls']
        if c == 'positive': e['pos'] += 1
        elif c == 'negative': e['neg'] += 1
        elif c == 'neutral-substantive': e['neu'] += 1
        else: e['form'] += 1
        # evidence for EVERY class the layer keeps (user 2026-07-04: the phrase must be showable):
        # pos/neg preferred, neutral-substantive fills remaining slots after the loop
        if c in ('positive', 'negative') and len(e['examples']) < 2:
            e['examples'].append({'cls': c, 'w': _html.unescape(r['w'])[:200], 'src': r['src']})
        elif c == 'neutral-substantive' and len(e['_neu_ex']) < 2:
            e['_neu_ex'].append({'cls': c, 'w': _html.unescape(r['w'])[:200], 'src': r['src']})
    for e in edges.values():
        while len(e['examples']) < 2 and e['_neu_ex']: e['examples'].append(e['_neu_ex'].pop(0))
    out_edges = []
    for (a, b), e in edges.items():
        sub = e['pos'] + e['neg'] + e['neu']
        if sub == 0: continue                      # formality-only dyads never enter the layer
        out_edges.append({'a': a, 'b': b, 'n_raw': e['n_raw'], 'n_scored': e['n_scored'],
                          'pos': e['pos'], 'neg': e['neg'], 'neu': e['neu'], 'formality_dropped': e['form'],
                          'net': round((e['pos'] - e['neg']) / max(1, sub), 3),
                          'sources': dict(e['sources']), 'examples': e['examples']})
    out_edges.sort(key=lambda x: -(x['pos'] + x['neg'] + x['neu']))
    from moneyball_core import dump_atomic
    dump_atomic({'edges': out_edges, 'dyad_cap': DYAD_CAP, 'model': ', '.join(sorted(_cache_models)) or MODEL,
                 'admission_filters': RF.admission_filters_block(AF_STATE, FILTERED),
                 'provenance': 'substantive member references (formality discarded per user 2026-07-04); speeches+press+X; Fable-valenced'},
                B + r'\member_references.json')
    npos = sum(1 for e in out_edges if e['net'] > 0.2); nneg = sum(1 for e in out_edges if e['net'] < -0.2)
    print('member_references.json: %d substantive directed edges | positive %d | negative %d' % (len(out_edges), npos, nneg))
    sys.exit(0)

if '--verify' in sys.argv:
    # the user's eyeball sheet: every edge with its evidence quotes (person-check + valence-check)
    d = json.load(open(B + r'\member_references.json', encoding='utf-8'))
    outp = B + r'\data\members\member_reference_verify.csv'
    with io.open(outp, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow(['speaker', 'about', 'net', 'pos', 'neg', 'neu', 'n_raw', 'sources',
                    'evidence_1', 'evidence_2', 'wrong_person_or_wrong_valence?'])
        for e in sorted(d['edges'], key=lambda x: -x['n_scored'])[:500]:
            ex = e.get('examples') or []
            w.writerow([e['a'], e['b'], e['net'], e['pos'], e['neg'], e['neu'], e['n_raw'],
                        ' '.join('%s:%d' % kv for kv in sorted((e.get('sources') or {}).items())),
                        (ex[0]['w'] if ex else ''), (ex[1]['w'] if len(ex) > 1 else ''), ''])
    print('wrote %s (%d edges, capped 500)' % (outp, min(500, len(d['edges']))))
    sys.exit(0)

print('usage: member_reference_layer.py [--probe | --score [--go] | --rebuild | --verify]')
