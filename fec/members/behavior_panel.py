# -*- coding: utf-8 -*-
import os as _o, sys as _s; _s.path.insert(0, _o.path.dirname(_o.path.dirname(_o.path.abspath(__file__)))); from paths import ROOT as MB_ROOT
"""BEHAVIORAL CRITERION PANEL (2026-07-04 'Full Reset': user rejected gold labels for latent
constructs — the research-correct instrument is the NOMOLOGICAL NETWORK (Cronbach & Meehl 1955):
trait measures are validated by predicting theory-specified BEHAVIORS. This module assembles the
per-member behavioral outcomes from data already on disk. The trait->behavior sign predictions are
pre-registered from the literature (research agent report) BEFORE any correlation is computed;
validate_nomological.py runs the one-shot test.
Outcomes assembled (per member):
  n_floor_speeches, floor_words          CREC volume (E+: verbal output; RKH speech-frequency)
  n_interviews, n_interview_programs     broadcast media-seeking (E+; media-seeking literature)
  n_tweets                               posting volume (E+ at mass level)
  n_sponsored, n_cosponsored             legislative initiative (C+/agency+; RKH sponsorship)
  cosponsor_degree                       network embedding (communion+/E+)
  bipartisanship                         cross-party tie share (A+/communion+)
  committee_leader                       chair/ranking role (agency+/dominance; status attainment)
  nominate_extremity                     |DW-NOM dim1| (O- per conservatism link is DIRECTIONAL not extremity — kept as covariate)
  votes_cast                             roll-call participation volume (C+)
  vote_share_last                        electoral margin (covariate: safe seats change behavior)
Read-only on every source; writes behavior_panel.json."""
import csv, io, json, os, sqlite3, sys, collections, glob
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
import openpyxl
B = MB_ROOT
C = {m['member']: m for m in json.load(open(B + r'\data\engine\congestion.json'))}
P = collections.defaultdict(dict)

# speech + interview volumes
con = sqlite3.connect(B + r'\social_media.db', timeout=30)
con.execute('PRAGMA query_only=1')
for m, n, w in con.execute("select member, count(*), sum(nwords) from speeches where source='CREC' group by member"):
    if m in C: P[m]['n_floor_speeches'] = n; P[m]['floor_words'] = w or 0
for m, n in con.execute("select member, count(*) from speeches where source='INTERVIEW' group by member"):
    if m in C: P[m]['n_interviews'] = n
for m, n in con.execute("select member, count(distinct package) from speeches where source='INTERVIEW' group by member"):
    if m in C: P[m]['n_interview_programs'] = n
# CHRG is deliberately absent from the behavioral panel. The legacy collector assigned some
# witness and wrong-chamber testimony to current members; corpus_integrity_audit.py inventories the
# quarantined rows. Do not restore hearing participation until CHRG is rebuilt from
# chamber/tenure-constrained speaker attribution and independently verified.

# tweet volume (labeled corpus)
tw = collections.Counter()
tp = B + r'\data\stance\tweet_stance_final.csv'
if os.path.exists(tp):
    for r in csv.DictReader(io.open(tp, encoding='utf-8-sig', newline='')):
        if r.get('member') in C: tw[r['member']] += 1
for m, n in tw.items(): P[m]['n_tweets'] = n

# sponsorship / cosponsorship from the bills archive
bp = max(glob.glob(B + r'\congress_bills_archive_*.xlsx'), key=os.path.getmtime)
ws = openpyxl.load_workbook(bp, read_only=True)['All Bills']
rows = ws.iter_rows(min_row=1, values_only=True)
hdr = [str(h or '') for h in next(rows)]
iR, iM, iN, iCg = hdr.index('Role'), hdr.index('Member'), hdr.index('Name'), hdr.index('Congress')
# CU_IDENTITY_FIX: name-keyed archive (Sen. Menendez's Senate bills -> Rep. M001226); shared check.
from members.member_identity import load_billstatus, archive_row_unverified
_billstatus = load_billstatus(B); identity_rejected = collections.Counter()
sp = collections.Counter(); co = collections.Counter()
for r in rows:
    m = r[iM]; role = str(r[iR] or '').lower()
    if m not in C: continue
    if archive_row_unverified(_billstatus, r[iN], r[iCg], C[m].get('bioguide')):
        identity_rejected[m] += 1; continue
    if 'sponsor' == role or role.startswith('sponsor'): sp[m] += 1
    elif 'cosponsor' in role: co[m] += 1
print('identity rows unverified by BILLSTATUS (CU_IDENTITY_FIX): %d' % sum(identity_rejected.values()), dict(identity_rejected))
for m in C:
    if sp[m] or co[m]: P[m]['n_sponsored'] = sp[m]; P[m]['n_cosponsored'] = co[m]

# cosponsor network degree (non-AUG file = pure cosponsorship)
np_ = max([q for q in glob.glob(B + r'\congress_cosponsor_network_*.xlsx') if 'AUG' not in os.path.basename(q)],
          key=os.path.getmtime)
b2n = {info.get('bioguide'): m for m, info in C.items() if info.get('bioguide')}
ne = openpyxl.load_workbook(np_, read_only=True, data_only=True)['Edges']
deg = collections.Counter()
it = ne.iter_rows(min_row=2, values_only=True)
for r in it:
    a, b = b2n.get(r[0]), b2n.get(r[1])
    if a and b and a != b: deg[a] += 1; deg[b] += 1
for m, d in deg.items(): P[m]['cosponsor_degree'] = d

# bipartisanship (pipeline artifact)
bip = json.load(open(B + r'\bipartisan.json')) if os.path.exists(B + r'\bipartisan.json') else {}
for m, v in bip.items():
    if m in C: P[m]['bipartisanship'] = round(float(v), 4)

# committee leadership from the desk bundle (chair/ranking/vice on any committee)
mdb = B + r'\member_desk_bundle.json'
if os.path.exists(mdb):
    for rec in json.load(open(mdb, encoding='utf-8')).get('members', []):
        m = rec.get('name')
        if m not in C: continue
        roles = [c.get('role') or '' for c in (rec.get('committees') or [])]
        P[m]['committee_leader'] = 1 if any(r for r in roles) else 0

# voteview: votes cast + ideology extremity + party
for cn in (119,):
    p = os.path.join(B, 'votes', 'HS%d_members.csv' % cn)
    if not os.path.exists(p): continue
    for r in csv.DictReader(open(p, encoding='utf-8')):
        m = b2n.get((r.get('bioguide_id') or '').strip())
        if not m: continue
        # Voteview HS119 can list a bioguide twice (party-switch pseudo-record party_code=328, or a
        # bicameral split). Keep the row with the MOST votes so a low-vote phantom copy can't clobber
        # the real member (e.g. K000401 Kiley: real 386-vote row vs a 109-vote 328 phantom).
        try: v = int(float(r.get('nominate_number_of_votes') or 0))
        except ValueError: continue
        if v <= P[m].get('votes_cast', -1): continue
        P[m]['votes_cast'] = v
        try: P[m]['nominate_extremity'] = round(abs(float(r.get('nominate_dim1') or 0)), 3)
        except ValueError: pass

# electoral margin covariate
eh = json.load(open(B + r'\data\money\electoral_history.json', encoding='utf-8')) if os.path.exists(B + r'\data\money\electoral_history.json') else {}
for m, d in eh.items():
    if m in C and isinstance(d, dict):
        # guard the pct range so a corrupt electoral_history row (e.g. a -100 or 119 MEDSL artifact)
        # can't leak into the vote_share_last validation criterion.
        gen = [g for g in (d.get('general') or []) if isinstance(g, dict) and isinstance(g.get('pct'), (int, float)) and 0 <= g['pct'] <= 100]
        if gen:
            try: P[m]['vote_share_last'] = round(float(gen[-1]['pct']), 2)
            except (ValueError, TypeError, KeyError): pass

from moneyball_core import dump_atomic
OUTS = ['n_floor_speeches', 'floor_words', 'n_interviews', 'n_interview_programs', 'n_tweets',
        'n_sponsored', 'n_cosponsored', 'cosponsor_degree', 'bipartisanship', 'committee_leader',
        'votes_cast', 'nominate_extremity', 'vote_share_last']
cov = {o: sum(1 for m in P if o in P[m]) for o in OUTS}
dump_atomic({'members': dict(P), 'outcomes': OUTS, 'coverage': cov,
             'note': 'behavioral criterion panel for nomological validation; assembled read-only from '
                     'existing artifacts; sign predictions pre-registered separately before testing; '
                     'legacy CHRG hearing counts quarantined after attribution-integrity failure'},
            B + r'\behavior_panel.json')
print('behavior panel: %d members' % len(P))
for o in OUTS: print('  %-22s coverage %d' % (o, cov[o]))
