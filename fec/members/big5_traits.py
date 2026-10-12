# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
"""BIG FIVE TRAIT ESTIMATION FROM MEMBER TEXT — Ramey, Klingler & Hollibaugh (PSRM 2016)
"Measuring Elite Personality Using Speech", synthesized (user 2026-07-02: "copy what these articles
did bar for bar in a synthesized manner, it does not have to be domain specific").

Their pipeline, kept step for step:
  1. Aggregate each legislator's text over a fixed span (theirs: floor speeches by year; ours: the
     member's full tweet/post corpus — Birkenmaier & Lechner (Pol.Analysis 2025) validate social
     posts as personality-cue carriers on par with speeches).
  2. Score the text on closed-vocabulary LIWC/MRC-style category rates (per-1000-word usage).
  3. Standardize the category rates ACROSS THE CORPUS and combine with standardized weights into
     the Big Five (their step: Mairesse et al. 2007 pre-trained standardized SMOreg over the same
     features; the model file is LIWC-license-bound, so we use the published Pennebaker & King
     (1999) / Yarkoni (2010) / Mairesse Table-3 correlates as the standardized linear weights —
     same feature families, same signs, linear instead of SVR).
  4. Scores are z-normalized RELATIVE TO THIS CORPUS (their explicit framing: estimates are only
     meaningful relative to the corpus used — congressional distribution, not population norms).
  5. Confidence: they show CI width blows up below ~1k words/member (Fig 3, median 11k/yr).  We
     carry tokens and flag lowconf below LOWCONF_TOKENS (default 1500); consumers shrink lowconf
     trait z-scores toward 0 (uncertainty -> neutral, never a fabricated personality).
Emotional Stability (ES) is reported Ramey-style; Neuroticism N = -ES.
Output: big5_traits.json  {member: {O,C,E,A,ES,N (z-scores), tokens, nposts, lowconf}}"""
import sys, os, json, re, collections, math
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
import openpyxl,pandas as pd
B=MB_ROOT
LOWCONF_TOKENS=int(os.environ.get('LOWCONF_TOKENS','1500'))
C={m['member']:m for m in json.load(open(B+r'\data\engine\congestion.json'))}; members=list(C)
# ---------------- corpus: posts (ALL topics; RTs EXCLUDED — an RT is someone else's words) + floor
# speeches / hearing transcripts (speeches table, fetch_congressional_record.py — the papers' native
# medium; Ramey used the Congressional Record itself) ----------------
# validation hooks (validate_big5.py MTMM/test-retest axes; guarded, default = full corpus, no change):
#   BIG5_SOURCES  csv subset of {social, CREC, INTERVIEW, CHRG} — genre-restricted scoring
#   BIG5_BEFORE / BIG5_AFTER  ISO date — time-sliced SPEECHES (social rows carry no usable date here)
_SRC=set(s.strip() for s in os.environ.get('BIG5_SOURCES','').split(',') if s.strip()) or None
# Legacy CHRG rows are not safe default personality evidence: historical surname resolution attached
# some witness turns to same-surname members, and prepared committee statements are outside the
# user-approved ordinary-questioning scope.  A caller may still request CHRG explicitly for an audit.
_SPEECH_SRC=_SRC or {'CREC','INTERVIEW'}
_BEF=os.environ.get('BIG5_BEFORE',''); _AFT=os.environ.get('BIG5_AFTER','')
if _SRC or _BEF or _AFT: print('CORPUS FILTER: sources=%s before=%s after=%s'%(sorted(_SRC) if _SRC else 'all',_BEF or '-',_AFT or '-'))
texts=collections.defaultdict(list)
def clean_social_text(value):
    """Return usable authored text, never pandas' missing-value sentinel as corpus language."""
    if value is None or (isinstance(value,float) and math.isnan(value)): return None
    text=str(value).strip()
    return None if not text or text.lower() in {'nan','none','<na>','nat'} else text
if _SRC is None or 'social' in _SRC:
    ws=openpyxl.load_workbook(B+r'\Quorum Processed.xlsx',read_only=True)['Members']
    hdr=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
    im,ic,irt=hdr.index('member'),hdr.index('Content'),hdr.index('is_rt')
    for row in ws.iter_rows(min_row=2,values_only=True):
        m,t=row[im],clean_social_text(row[ic])
        if m in C and t and str(row[irt]).lower() not in ('1','true'): texts[m].append(t)
    sm=pd.read_excel(B+r'\Social Master.xlsx',sheet_name='Corpus',dtype=str)
    for _,r in sm[sm['platform']=='bluesky'].iterrows():
        t=clean_social_text(r.get('text'))
        if r['member'] in C and t and str(r.get('is_rt','')).lower() not in ('1','true'): texts[r['member']].append(t)
import sqlite3,hashlib as _hl
nsp=0; ndup=0; _seen=set()
try:
    con=sqlite3.connect(B+r'\social_media.db')
    for m,t,src,dt in con.execute('SELECT member,text,source,date FROM speeches'):
        if src not in _SPEECH_SRC: continue
        d=str(dt or '')[:10]
        if _BEF and (not d or d>=_BEF): continue
        if _AFT and (not d or d<_AFT): continue
        if m in C and t:
            _h=(m,_hl.md5(t.encode('utf-8','replace')).hexdigest())
            if _h in _seen: ndup+=1; continue   # QA 2026-07-03: re-aired soundbites + crash-refetch dups must not inflate rates
            _seen.add(_h)
            texts[m].append(t); nsp+=1
except sqlite3.OperationalError as _db_exc:
    raise RuntimeError('speech corpus read failed (%s); every trait z is standardized within THIS corpus, so a missing table or a mid-read lock must not silently reshape all 533 scores' % _db_exc) from _db_exc
print('corpus: speeches %d turns folded in (%d duplicates skipped)'%(nsp,ndup))
# ---------------- closed-vocabulary categories (LIWC-analog; word lists are the published exemplars) ----------------
CAT={
 'i':"i me my mine myself im ive ill id",
 'we':"we us our ours ourselves were weve well",
 'you':"you your yours youre",
 'article':"a an the",
 'negate':"no not never none cannot cant dont wont isnt wasnt arent didnt doesnt wouldnt couldnt shouldnt nothing neither nor",
 'posemo':"good great love happy hope proud thank thanks celebrate honor honored best wonderful glad support win wins success benefit benefits opportunity grateful appreciate excellent congratulations congrats delighted terrific amazing",
 'negemo':"bad wrong hate hurt nasty terrible awful fail failed failure crisis threat worry worried fear afraid angry anger sad loss lose damage harm danger attack disaster outrage shameful disgrace broken devastating reckless",
 'anx':"worried anxious afraid fear fearful nervous risk risky threat threatened threatens uncertain uncertainty alarm alarming danger dangerous crisis vulnerable insecure unstable",
 'anger':"angry anger outrage outrageous furious fight fighting attack attacks hate disgusting shameful corrupt corruption betray betrayal rigged radical disgrace hostile",
 'social':"talk talked share friend friends family families community communities neighbor neighbors together meeting met join joined joining team people folks everyone partnership colleagues constituents neighbors town",
 'family':"family families mother father mom dad son daughter kids children parents grandmother grandfather wife husband brother sister grandparents",
 'achieve':"work working achieve achieved achievement accomplish accomplished succeed success successful win won effort efforts goal goals deliver delivered results earn earned lead leader leadership fought secured secure passed",
 'work':"job jobs work working business businesses economy economic employment worker workers industry wages career manufacturing farmers",
 'insight':"think thinking believe know understand consider considered because reason reasons means thus therefore however although question questions idea ideas analysis examine evidence solution solutions",
 'certain':"always never must certainly absolutely definitely clear clearly undeniable undoubtedly every all guarantee fact facts truth proven commit committed",
 'tentat':"maybe perhaps might may could seems appears possibly somewhat unclear likely probably potential potentially",
 # expansion 2026-07-03: classic LIWC-family categories (Mairesse feature space) — Empath-overlaid via
 # fec/liwc/empath_categories.dic; consumed only through FITTED weights (mairesse_smoreg.json), so with
 # the published-correlates fallback they are computed but unweighted (harmless).
 'money':"money cost costs price prices dollar dollars fund funds funding budget tax taxes pay paid spending afford debt",
 'relig':"god faith church prayer pray blessed bless religious christian jewish muslim worship holy",
 'swear':"damn hell crap",
 'leisure':"vacation game games sports fun play movie music concert hobby travel weekend",
 'body':"body hands hand head heart blood eyes feet skin",
 'health':"health healthy sick illness disease doctor hospital medicine medical care treatment patient patients",
 'death':"death die died dying dead killed kill funeral grief mourn",
 'sad':"sad sadness cry crying tears grief heartbroken sorrow disappointed loss",
}
CATW={k:set(v.split()) for k,v in CAT.items()}
CATPRE={k:tuple() for k in CAT}   # wildcard prefixes (LIWC '.dic' entries like 'happi*')
# ---- EXACT-ARTICLE PARITY HOOKS (user-ratified 2026-07-03: "yes to LIWC/SMOreg") ----
# Drop the licensed files into fec/liwc/ and they take over; absent, the published-exemplar lists
# and correlate weights below stay in force (current behavior, zero output change).
# * fec/liwc/*.dic — LIWC 2007/2015 dictionary ('%'-delimited header of "id<TAB>name" rows, then
#   "word<TAB>id id..." rows; trailing '*' = prefix wildcard). LIWC2015 category NAMES map directly
#   onto our features (i,we,you,article,negate,posemo,negemo,anx,anger,social,family,achieve,work,
#   insight,certain,tentat); unmatched LIWC categories are ignored. Note: full-dictionary wildcard
#   matching makes the corpus pass considerably slower (minutes, not seconds) — fine for weekly runs.
# * fec/liwc/mairesse_smoreg.json — {trait:{feature:weight}} replaces the linear weights W (their
#   SMOreg model exported as weights).
import glob as _g
_dicf=sorted(_g.glob(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),'liwc','*.dic')))
if _dicf:
    _hdr={}; _inh=False; _exact={c:set() for c in CAT}; _pre={c:[] for c in CAT}; _used=set()
    for _ln in open(_dicf[0],encoding='utf-8',errors='replace'):
        _ln=_ln.strip()
        if _ln=='%': _inh=not _inh; continue
        if not _ln: continue
        _p=_ln.split('\t') if '\t' in _ln else _ln.split()
        if _inh:
            if len(_p)>=2 and _p[0].isdigit(): _hdr[_p[0]]=_p[1].strip().lower()
        else:
            _w=_p[0].lower()
            for _i in (_q for _q in _p[1:] if _q.isdigit()):
                _c=_hdr.get(_i)
                if _c in CAT:
                    _used.add(_c)
                    (_pre[_c].append(_w[:-1]) if _w.endswith('*') else _exact[_c].add(_w))
    for _c in _used:
        CATW[_c]=_exact[_c]; CATPRE[_c]=tuple(sorted(set(_pre[_c])))
    print('LIWC dictionary ACTIVE: %s -> %d/%d categories replaced (%s)'%(
          os.path.basename(_dicf[0]),len(_used),len(CAT),', '.join(sorted(_used))))
# ---- psycholinguistic norms (research 2026-07-03: openness lives in ABSTRACTNESS + lexical
# diversity — Mairesse/JAIR07, Frontiers 2022). Brysbaert et al. concreteness norms (free, 40k
# words) in fec/liwc/norms/; guarded like the Empath dic.
# LEXICAL DIVERSITY = MATTR (moving-average type-token ratio, window 50). QA wave-3 2026-07-03:
# the Guiraud index (types/sqrt n) was r=0.70 with token count — a VOLUME PROXY, not diversity
# (it feeds openness, the very trait we're fixing). MATTR averages TTR over fixed-size windows so
# it is genuinely length-invariant (verified r~0 vs log tokens after the swap).
CONC={}
_cnf=os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),'liwc','norms','concreteness_brysbaert.txt')
if os.path.exists(_cnf):
    for _ln in open(_cnf,encoding='utf-8',errors='replace'):
        _p=_ln.rstrip('\n').split('\t')
        if len(_p)>=3 and _p[1]=='0':
            try: CONC[_p[0].lower()]=float(_p[2])
            except ValueError: pass
    print('concreteness norms ACTIVE: %d single words (Brysbaert)'%len(CONC))
# ---- B&L persona cue features (agency_communion_llm.py): first-class trait features when the
# artifact exists. Uncovered members get the covered-population MEAN (z=0 after standardization —
# neutral, never an invented persona). Coverage grows to the full roster with the --all run.
ACF={}
_acp=os.path.join(B,'data/members/agency_communion.json')
if os.path.exists(_acp):
    with open(_acp,encoding='utf-8') as _af:
        _ac=json.load(_af)
    if not isinstance(_ac.get('members'),dict):
        raise ValueError('data/members/agency_communion.json: required members mapping is missing')
    ACF={}
    for m,d in _ac['members'].items():
        if not isinstance(d,dict):
            raise ValueError('data/members/agency_communion.json: malformed member record for %s'%m)
        if (d.get('n') or 0)>=15:
            if not all(isinstance(d.get(k),(int,float)) for k in ('agency','communion')):
                raise ValueError('data/members/agency_communion.json: covered member %s lacks numeric features'%m)
            ACF[m]=d
    if ACF: print('persona cue features ACTIVE: %d members covered (uncovered -> neutral mean)'%len(ACF))
FEATS=list(CAT)+['sixltr','exclam','number','mattr']+(['concrete'] if CONC else [])+(['llm_agency','llm_communion'] if ACF else [])
# ---------------- standardized trait weights (Pennebaker–King / Yarkoni / Mairesse correlates) ----------------
# Signs are the replicated literature directions; magnitudes are the relative correlate strengths.
W={
 'E':{'social':.30,'posemo':.20,'we':.15,'exclam':.10,'tentat':-.15,'negemo':-.10},          # extraverts: social+positive, low hedging
 'ES':{'anx':-.30,'negemo':-.25,'i':-.20,'negate':-.15,'posemo':.10},                        # stability = inverse neuroticism markers
 'A':{'posemo':.25,'family':.20,'social':.15,'we':.15,'anger':-.25},                         # agreeable: warm/communal, low hostility
 'C':{'achieve':.35,'work':.20,'certain':.10,'negate':-.20,'negemo':-.15},                   # conscientious: achievement talk, low negations
 'O':{'sixltr':.30,'insight':.25,'article':.20,'i':-.15,'certain':-.10},                     # open: lexical complexity + cognitive words
}
_smof=os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),'liwc','mairesse_smoreg.json')
if os.path.exists(_smof):
    W=json.load(open(_smof,encoding='utf-8'))
    print('Mairesse SMOreg weights ACTIVE (liwc/mairesse_smoreg.json) — replacing published-correlate linear weights')
WORD=re.compile(r"[a-z']+")
# chamber RITUAL is not lexical behavior — floor-management boilerplate (heavier for the majority that
# runs procedure) contaminates formality features; strip the fixed phrases before category counting.
RITUAL=re.compile(r'\b(i yield (back|to|the floor|myself)[^.]*|mr\.? (president|speaker)|madam (president|speaker)|'
 r'i ask unanimous consent[^.]*|the (gentleman|gentlewoman|gentlelady) from [a-z ]+|'
 r'i (rise|reserve the balance of my time)|pursuant to [^.]*|the clerk will [^.]*|'
 r'i suggest the absence of a quorum|the balance of my time)\b',re.I)
MATTR_W=50
def rates(txts):
    cnt=collections.Counter(); n=0; six=0; excl=0; num=0
    conc=0.0; ncc=0
    win=collections.deque(); wc=collections.Counter(); ttr_sum=0.0; ttr_n=0   # MATTR streaming state
    for t in txts:
        t=re.sub(r'https?://\S+|\bt\.co/\S+|@\w+','',t.lower())   # URLs + @handles are not lexical behavior
        t=RITUAL.sub(' ',t)
        excl+=t.count('!'); num+=len(re.findall(r'\d',t))
        for w in WORD.findall(t.replace("'","")):
            n+=1
            if len(w)>=6: six+=1
            win.append(w); wc[w]+=1
            if len(win)>MATTR_W:
                old=win.popleft(); wc[old]-=1
                if wc[old]==0: del wc[old]
            if len(win)==MATTR_W: ttr_sum+=len(wc)/MATTR_W; ttr_n+=1
            if CONC:
                cv=CONC.get(w)
                if cv is not None: conc+=cv; ncc+=1
            for c,s in CATW.items():
                if w in s or (CATPRE[c] and w.startswith(CATPRE[c])): cnt[c]+=1
    if n==0: return None,0
    r={c:1000.0*cnt[c]/n for c in CAT}; r['sixltr']=1000.0*six/n; r['exclam']=1000.0*excl/n; r['number']=1000.0*num/n
    r['mattr']=(ttr_sum/ttr_n) if ttr_n else (len(wc)/max(1,len(win)))   # length-invariant diversity; short corpora -> single-window TTR
    if CONC: r['concrete']=(conc/ncc) if ncc else 2.5  # mean Brysbaert concreteness (openness loads NEGATIVE)
    return r,n
# split-half hook (validate_big5.py reliability axis): BIG5_HALF=1|2 scores alternating documents
# only; BIG5_OUT redirects the artifact so halves never clobber production. Default: unchanged.
_HALF=os.environ.get('BIG5_HALF','')
if _HALF in ('1','2'):
    texts={m:v[int(_HALF)-1::2] for m,v in texts.items()}
    print('SPLIT-HALF mode: half %s (alternating documents)'%_HALF)
R={}; TOK={}
for m in members:
    r,n=rates(texts.get(m,[]))
    if r: R[m]=r; TOK[m]=n
if ACF:
    _ma=sum(d['agency'] for d in ACF.values())/len(ACF); _mc=sum(d['communion'] for d in ACF.values())/len(ACF)
    for m in R:
        d=ACF.get(m)
        R[m]['llm_agency']=d['agency'] if d else _ma
        R[m]['llm_communion']=d['communion'] if d else _mc
# corpus-standardize each feature (Ramey step: standardized features -> standardized model)
mu={f:sum(R[m][f] for m in R)/len(R) for f in FEATS}
sd={f:(sum((R[m][f]-mu[f])**2 for m in R)/max(1,len(R)-1))**.5 or 1.0 for f in FEATS}
Z={m:{f:(R[m][f]-mu[f])/sd[f] for f in FEATS} for m in R}
raw={m:{t:sum(w*Z[m][f] for f,w in W[t].items()) for t in W} for m in R}
tmu={t:sum(raw[m][t] for m in raw)/len(raw) for t in W}
tsd={t:(sum((raw[m][t]-tmu[t])**2 for m in raw)/max(1,len(raw)-1))**.5 or 1.0 for t in W}
out={}
for m in members:
    if m in raw:
        tr={t:round((raw[m][t]-tmu[t])/tsd[t],3) for t in W}
        out[m]={**tr,'N':round(-tr['ES'],3),'tokens':TOK[m],'nposts':len(texts.get(m,[])),'lowconf':TOK[m]<LOWCONF_TOKENS}
    else:
        out[m]={'O':0.0,'C':0.0,'E':0.0,'A':0.0,'ES':0.0,'N':0.0,'tokens':0,'nposts':0,'lowconf':True}
from moneyball_core import dump_atomic   # atomic write: this runs UNATTENDED in weekly_refresh
dump_atomic({'traits':out,'method':'Ramey-synth closed-vocab standardized linear; Empath lexicons + essays-fitted SMOreg weights when fec/liwc/ files present; z relative to THIS corpus',
           'featz':{m:{f:round(Z[m][f],4) for f in FEATS} for m in R},   # standardized features (fit_smoreg_rkh.py trains on these — scorer and fitter can never drift)
           'lowconf_tokens':LOWCONF_TOKENS},os.environ.get('BIG5_OUT') or B+r'\data\members\big5_traits.json')
nm=sum(1 for m in out if out[m]['tokens']>0); lc=sum(1 for m in out if out[m]['lowconf'])
tk=sorted(TOK.values())
print('BIG5: %d/%d members with text | median tokens %d | lowconf(<%d tok): %d'%(nm,len(members),tk[len(tk)//2] if tk else 0,LOWCONF_TOKENS,lc))
for t in ('O','C','E','A','ES'):
    top=sorted((m for m in out if not out[m]['lowconf']),key=lambda m:-out[m][t])[:3]
    bot=sorted((m for m in out if not out[m]['lowconf']),key=lambda m:out[m][t])[:3]
    print('  %-2s high: %-45s low: %s'%(t,', '.join(top),', '.join(bot)))
# ---- Ramey Table-1 style external validation: trait vs party (their DV: Common Space conservatism;
# expected signs there: O −, C +, E −, A −, ES +). Party = roster proxy for ideology. ----
import csv as _csv
party={}
for cn in (119,118,117):
    p=B+r'\votes\HS%d_members.csv'%cn
    if os.path.exists(p):
        for row in _csv.DictReader(open(p,encoding='utf-8')):
            b=row.get('bioguide_id'); pc=str(row.get('party_code'))
            if b and b not in party: party[b]={'100':0,'200':1}.get(pc)
pr={m:party.get(C[m].get('bioguide','')) for m in members}
EXP={'O':-1,'C':1,'E':-1,'A':-1,'ES':1}
print('validation (point-biserial r vs Republican, weighted members only; Ramey-expected sign in []):')
val=[m for m in out if not out[m]['lowconf'] and pr.get(m) is not None]
for t in ('O','C','E','A','ES'):
    xs=[out[m][t] for m in val]; ys=[pr[m] for m in val]
    mx=sum(xs)/len(xs); my=sum(ys)/len(ys)
    sx=(sum((x-mx)**2 for x in xs))**.5; sy=(sum((y-my)**2 for y in ys))**.5
    r=sum((x-mx)*(y-my) for x,y in zip(xs,ys))/max(1e-9,sx*sy)
    ok='OK' if r*EXP[t]>0 else 'REVERSED'
    print('  %-2s r=%+.3f [%+d expected]  %s'%(t,r,EXP[t],ok))
print('wrote %s'%(os.environ.get('BIG5_OUT') or 'big5_traits.json'))
