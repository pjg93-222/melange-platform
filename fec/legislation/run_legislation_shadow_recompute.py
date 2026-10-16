# -*- coding: utf-8 -*-
"""Run the legislation refresh twice in a sealed, hash-bound shadow workspace.

This orchestrator never promotes artifacts.  It validates the immutable
production baseline, separates strict discovery from every final coverage
gate, verifies the accepted-generation -> direction/base-network/roll-call
hash chain, and then executes an explicit isolation plan twice.

Legacy consumers that cannot accept an isolated output root are not executed:
their missing or unsafe plan step is a blocking error.  A technical
``staged_acceptance_receipt.json`` is emitted only when coverage, producer
currency, baseline immutability, attribution, and both value and byte
idempotence all pass.
"""
from __future__ import annotations

import os as _bootstrap_os, sys as _bootstrap_sys; _bootstrap_sys.path.insert(0, _bootstrap_os.path.dirname(_bootstrap_os.path.dirname(_bootstrap_os.path.abspath(__file__))))

import argparse
import csv
from dataclasses import dataclass
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import string
import subprocess
import sys
import tempfile
from typing import Any, Iterable, Mapping, Sequence

from legislation.assemble_approved_legislation_generation import (
    ROLE_METHOD_PROFILE,
    ROLE_METHOD_PROFILE_ID,
)


HERE = Path(__file__).resolve().parent.parent
ROOT = HERE.parent
SHADOW_ROOT = (ROOT / "audits" / "legfix5" / "shadow_recompute").resolve()
REPO_TEMP_ROOT = (ROOT / "tmp").resolve()
REGISTRY = (
    ROOT
    / "source_data"
    / "legislation"
    / "acu_legislation_search_registry.json"
)

AUDIT_SCHEMA = "legislation-shadow-recompute-audit-v1"
RECEIPT_SCHEMA = "legislation-shadow-recompute-acceptance-v1"
PLAN_SCHEMA = "legislation-shadow-consumer-plan-v1"
BASELINE_SCHEMA = "legislation-production-baseline-v1"
ACCEPTED_MANIFEST_SCHEMA = "accepted-legislation-generation-manifest-v1"
REVIEW_MANIFEST_SCHEMA = "acu-legislation-membership-review-v3"
STEP_LEDGER_SCHEMA = "legislation-shadow-step-receipts-v1"
STEP_LEDGER_NAME = Path("intermediate") / "shadow_step_receipts.json"
HEX_SHA_RE = re.compile(r"^[0-9a-f]{64}$", re.IGNORECASE)

DEFAULT_STEP_TIMEOUT_SECONDS = 6 * 60 * 60
REQUIRED_BASELINE_ROLES = (
    "legislation_archive_workbook",
    "archive_leg",
    "legislation_routing_audit",
    "legacy_rollcall_aggregate",
    "base_network",
    "augmented_network",
    "bipartisan",
    "behavior_panel",
    "behavioral_dispositions",
    "cooperation",
    "receptivity",
    "category_stance",
    "say_do",
    "blended_workbook",
    "pillar_snapshot",
    "ranks",
    "electoral_plays",
    "per_topic_alloc",
    "per_topic_coalition",
    "budget_frontier",
    "budget_frontier_electoral",
    "optimize_result",
    "im_allocation",
    "salience_advisory",
    "dashboard_bundle",
    "dashboard_html",
    "member_desk",
    "member_desk_html",
    "dashboard_exclusion_reconciliation",
    "combined_html",
)

BASELINE_ROLE_BY_BASENAME = {
    "archive_leg.json": "archive_leg",
    "legislation_routing_audit.json": "legislation_routing_audit",
    "vote_category_lean.json": "legacy_rollcall_aggregate",
    "bipartisan.json": "bipartisan",
    "behavior_panel.json": "behavior_panel",
    "behavioral_dispositions.json": "behavioral_dispositions",
    "coop_v2_candidate.json": "cooperation",
    "member_receptivity.json": "receptivity",
    "category_stance.json": "category_stance",
    "say_do_gap.json": "say_do",
    "moneyball main version - blended.xlsx": "blended_workbook",
    "data/engine/pillar_inputs_snapshot.json": "pillar_snapshot",
    "ranks_python.json": "ranks",
    "data/money/electoral_plays.json": "electoral_plays",
    "per_topic_alloc.json": "per_topic_alloc",
    "data/lobbying/per_topic_coalition.json": "per_topic_coalition",
    "data/money/budget_frontier.json": "budget_frontier",
    "data/money/budget_frontier_electoral.json": "budget_frontier_electoral",
    "optimize_result.json": "optimize_result",
    "im_allocation.json": "im_allocation",
    "salience_advisory.json": "salience_advisory",
    "dashboard_data.json": "dashboard_bundle",
    "moneyball dashboard.html": "dashboard_html",
    "member_desk_bundle.json": "member_desk",
    "cu-member-desk.html": "member_desk_html",
    "dashboard_exclusion_reconciliation.json": (
        "dashboard_exclusion_reconciliation"
    ),
    "moneyball.html": "combined_html",
}

# This is the corrected dependency order from the downstream contract.  The
# second say-do pass and its dependents are deliberate: the current production
# chain otherwise retains a one-generation lag.
REQUIRED_STEP_ORDER = (
    "materialize_archive_and_legislative_workbook",
    "materialize_rollcall_aggregate",
    "materialize_base_network_and_prebehavior_bipartisan",
    "recompute_behavior_panel",
    "recompute_behavioral_dispositions",
    "recompute_cooperation",
    "recompute_receptivity",
    "recompute_augmented_network",
    "recompute_category_stance",
    "recompute_say_do",
    "recompute_legislative_raw_normalized_sublayer",
    "recompute_pillars_and_ranks",
    "recompute_coalitions_flips_budget_optimizer_allocation",
    "recompute_dashboard_and_member_desk",
    "recompute_say_do_second_pass",
    "recompute_second_pass_dependents",
)

ALLOWED_ATTRIBUTION_DRIVERS = frozenset(
    {
        "accepted_legislation",
        "relationship_roles",
        "directions",
        "decay",
        "rollcalls",
        "approved_computation_repair",
    }
)

# An attribution only "explains" a delta if a person wrote it.  The consumer
# plan builder auto-fills an identical rationale onto every baseline role, so a
# present rationale is not evidence that anyone looked.  Attribution rows now
# carry their authorship, and a machine-written rationale on a CHANGED
# production artifact does not satisfy the explained requirement.
HUMAN_ATTRIBUTION_AUTHORS = frozenset({"human", "person", "reviewer"})
MACHINE_RATIONALE_TEMPLATES = frozenset(
    {
        (
            "this isolated consumer is downstream of the accepted "
            "legislation generation."
        ),
        "declared isolated intermediate artifact.",
    }
)
ACCEPTED_ATTRIBUTION_STATUSES = frozenset(
    {"unchanged", "attributed", "staged_intermediate"}
)

# ---------------------------------------------------------------------------
# Rank drift.
#
# The only baseline comparison used to be ``new sha256 != old sha256`` on
# ranks_python.json, which says a byte changed and nothing else - not how many
# members moved, not how far, not who.  The integration architecture requires
# that "Large or unexplained movement blocks promotion", so movement has to be
# measured before it can be judged.
# ---------------------------------------------------------------------------
RANK_DRIFT_SCHEMA = "legislation-shadow-rank-drift-v1"
RANK_DRIFT_ROLE = "ranks"
RANK_DRIFT_ARTIFACT_NAME = "rank_drift.json"
RANK_DRIFT_TOP_MOVER_LIMIT = 25
# Fail closed by default: any movement at all needs an explicit human
# allowance on the command line.  No magic tolerance is invented here.
DEFAULT_MAX_RANK_MOVERS = 0
DEFAULT_MAX_ABS_RANK_DELTA = 0

# Each acceptance gate declares how its value came to exist.  "measured" gates
# are computed from this run's artifacts; "enforced_by_precondition" gates are
# upheld by a named check that raises before the gate is assembled, so their
# True is a restatement of an earlier abort, not an independent measurement.
GATE_PROVENANCE: dict[str, tuple[str, str]] = {
    "strict_discovery": ("enforced_by_precondition", "validate_coverage"),
    "final_merged_coverage": ("enforced_by_precondition", "validate_coverage"),
    "membership_review": ("enforced_by_precondition", "validate_coverage"),
    "relationships": ("enforced_by_precondition", "validate_coverage"),
    "coverage_reconciliation": (
        "enforced_by_precondition",
        "validate_coverage",
    ),
    "accepted_generation_current": (
        "enforced_by_precondition",
        "validate_producers",
    ),
    "direction_current": ("enforced_by_precondition", "validate_producers"),
    "base_network_current": ("enforced_by_precondition", "validate_producers"),
    "rollcall_current": ("enforced_by_precondition", "validate_producers"),
    "production_baseline_immutable": (
        "enforced_by_precondition",
        "validate_baseline_unchanged",
    ),
    "all_required_consumers_isolated": (
        "enforced_by_precondition",
        "load_plan",
    ),
    "old_new_delta_attribution": ("measured", "validate_attribution"),
    "rank_drift_within_threshold": ("measured", "compute_rank_drift"),
    "cross_run_value_idempotence": ("measured", "idempotence_audit"),
    "cross_run_byte_idempotence": ("measured", "idempotence_audit"),
    "complete_baseline_role_output_closure": ("measured", "execute_run"),
}
APPROVED_PLAN_PLACEHOLDERS = frozenset(
    {
        "run_dir",
        "workspace_root",
        "input_root",
        "accepted_input_dir",
        "direction_input_dir",
        "base_network_input_dir",
        "rollcall_input_dir",
    }
)
NON_AUTHORITATIVE_PATH_TOKENS = (
    "/billstatus_bulk_enrichment/review_sidecars/",
    "/billstatus_bulk_enrichment/structured_relevance_sidecars/",
    "/partial_review_prefill/",
    "/partial_review_queue/",
)
NON_AUTHORITATIVE_NAME_TOKENS = (
    "review_sidecar",
    "structured_relevance",
    "partial_review_prefill",
    "partial_review_queue",
    "membership_review_proposals",
    "staged_review_queue",
)
NON_AUTHORITATIVE_TRUE_FLAGS = (
    "review_prefill_only",
    "proposal_only",
    "partial_discovery_only",
)
NON_AUTHORITATIVE_FALSE_FLAGS = (
    "downstream_consumption_allowed",
    "decision_import_allowed",
    "scoring_input_allowed",
    "membership_disposition_authority",
    "ratification_authority",
    "scoring_authority",
    "usable_as_membership_review_ledger",
)


@dataclass(frozen=True)
class ImmutableInputContract:
    """A non-legislation input that must be frozen into each shadow run."""

    path: str
    consumers: tuple[str, ...]
    required: bool = True
    transient_root_compatibility: bool = False


# These are the raw/background reads of the complete 16-stage chain.  The five
# explicitly optional files are absent in the audited production source; every
# source that currently carries a model/display layer is required so a legacy
# ``if exists`` branch cannot silently turn that layer off.
IMMUTABLE_INPUT_CONTRACTS = (
    ImmutableInputContract(
        "data/engine/congestion.json",
        ("behavior", "network", "stance", "dependents", "dashboard"),
    ),
    ImmutableInputContract(
        "social_media.db", ("behavior", "member_desk")
    ),
    ImmutableInputContract(
        "social_media.db-wal", ("behavior", "member_desk"), required=False
    ),
    ImmutableInputContract(
        "data/stance/tweet_stance_final.csv", ("behavior", "member_desk")
    ),
    ImmutableInputContract(
        "votes/HS119_members.csv", ("behavior", "network", "dependents")
    ),
    ImmutableInputContract(
        "votes/HS118_members.csv", ("network", "dependents")
    ),
    ImmutableInputContract(
        "votes/HS117_members.csv", ("network", "dependents")
    ),
    ImmutableInputContract(
        "data/money/electoral_history.json", ("behavior", "member_desk")
    ),
    ImmutableInputContract(
        "data/members/lep_criterion.json", ("behavior", "say_do", "member_desk")
    ),
    ImmutableInputContract(
        "data/lobbying/interaction_features.json", ("behavior", "receptivity")
    ),
    ImmutableInputContract(
        "data/members/big5_traits.json", ("receptivity", "member_desk")
    ),
    ImmutableInputContract(
        "data/money/affective_n.json", ("receptivity", "dashboard")
    ),
    ImmutableInputContract(
        "data/money/electoral.json", ("receptivity", "dependents")
    ),
    ImmutableInputContract(
        "data/members/agency_communion.json", ("receptivity", "member_desk")
    ),
    ImmutableInputContract(
        "fec/retiring_members.txt", ("receptivity",)
    ),
    ImmutableInputContract(
        "data/stance/press_engagement.json", ("augmented_network", "dashboard")
    ),
    ImmutableInputContract(
        "data/stance/press_hcoi.json", ("augmented_network", "dashboard")
    ),
    ImmutableInputContract(
        "data/lobbying/social_interactions.json", ("augmented_network", "dashboard")
    ),
    ImmutableInputContract(
        "data/stance/frame_alignment.json", ("dependents",)
    ),
    ImmutableInputContract(
        "data/members/member_desk_src.json", ("say_do", "member_desk")
    ),
    ImmutableInputContract(
        "fec/model_bio.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "member_desk_manual.json", ("member_desk",), required=False
    ),
    ImmutableInputContract(
        "data/stance/press_date_overrides.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/stance/press_stance_overrides.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/lobbying/district_industry.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/lobbying/district_industry_4digit.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/district/district_employers.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/district/district_financial.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/district/district_hmda.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/district/district_community.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "member_references.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/members/member_personal_texture.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/members/big5_blue_validated.json", ("member_desk",)
    ),
    ImmutableInputContract("data/members/big5_blue_emb.json", ("member_desk",)),
    ImmutableInputContract("data/members/big5_v2.json", ("member_desk",)),
    ImmutableInputContract(
        "data/members/big5_llm_opus.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/stance/big5_senate_expert_gold.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/stance/big5_external_gold_validation.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/stance/big5_e_gold_anchored.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/stance/big5_gold_anchored.json", ("member_desk",)
    ),
    ImmutableInputContract("data/network/kb_benchmark.json", ("member_desk",)),
    ImmutableInputContract(
        "data/lobbying/signed_coalitions.json", ("dashboard", "member_desk")
    ),
    ImmutableInputContract(
        "data/members/behavioral_roles_v2_candidate.json", ("dashboard", "member_desk")
    ),
    ImmutableInputContract(
        "communication_style.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/members/profile_basis_v2.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "data/stance/interview_antagonism.json", ("behavior", "member_desk")
    ),
    ImmutableInputContract(
        "data/stance/hearing_antagonism.json", ("dashboard", "member_desk")
    ),
    ImmutableInputContract(
        "member_narratives_v2.json", ("member_desk",), required=False
    ),
    ImmutableInputContract(
        "data/members/staff_directory.json", ("member_desk",)
    ),
    ImmutableInputContract(
        "staff_state.json", ("member_desk",), required=False
    ),
    ImmutableInputContract(
        "staff_interactions.json", ("member_desk",), required=False
    ),
    ImmutableInputContract(
        "data/members/formal_visibility_roles.json", ("pillars_ranks", "member_desk")
    ),
    ImmutableInputContract("data/members/confidence.json", ("dashboard", "member_desk")),
    ImmutableInputContract("data/network/co_appearance.json", ("member_desk",)),
    ImmutableInputContract(
        "Social Master.xlsx", ("press_social", "member_desk")
    ),
    ImmutableInputContract(
        "data/stance/interview_stance_final.csv", ("press_social", "member_desk")
    ),
    ImmutableInputContract(
        "data/stance/interviews_cu.csv", ("press_social", "member_desk")
    ),
    ImmutableInputContract(
        "data/stance/hearing_stance_final.csv", ("press_social", "member_desk")
    ),
    ImmutableInputContract(
        "data/stance/hearings_cu.csv", ("press_social", "member_desk")
    ),
    ImmutableInputContract(
        "data/legislation/bill_overrides.csv", ("archive", "member_desk")
    ),
    ImmutableInputContract(
        "DEMOGRAPHIC_POTENTIAL_METHOD.md", ("member_desk",)
    ),
    ImmutableInputContract("fec/legislators.csv", ("member_desk",)),
    ImmutableInputContract("tier_overrides.json", ("dashboard",)),
)

CANONICAL_MANUAL_EVIDENCE = "manual_bill_evidence.csv"
TRANSIENT_BASELINE_INPUTS = (
    ImmutableInputContract(
        "member_desk_bundle.json",
        ("say_do_first_pass",),
        transient_root_compatibility=True,
    ),
)
PILLAR_PREREQUISITE_STEP = "recompute_pillars_and_ranks"
NETWORK_PREREQUISITE_STEP = (
    "recompute_coalitions_flips_budget_optimizer_allocation"
)
SAY_DO_FIRST_STEP = "recompute_say_do"


class ShadowError(RuntimeError):
    """A fail-closed shadow recomputation contract violation."""


def non_authoritative_markers(value: Mapping[str, Any]) -> list[str]:
    markers = [
        name
        for name in NON_AUTHORITATIVE_TRUE_FLAGS
        if value.get(name) is True
    ]
    markers.extend(
        name
        for name in NON_AUTHORITATIVE_FALSE_FLAGS
        if value.get(name) is False
    )
    if str(value.get("ratification_state") or "").strip().upper() == (
        "NOT_RATIFIED"
    ):
        markers.append("ratification_state=NOT_RATIFIED")
    return sorted(markers)


def non_authoritative_path(value: str | Path) -> bool:
    normalized = "/" + str(value).replace("\\", "/").casefold().lstrip("/")
    return any(
        token in normalized for token in NON_AUTHORITATIVE_PATH_TOKENS
    ) or any(token in normalized for token in NON_AUTHORITATIVE_NAME_TOKENS)


def authoritative_review_gate(manifest: Mapping[str, Any]) -> bool:
    return bool(
        manifest.get("schema") == REVIEW_MANIFEST_SCHEMA
        and manifest.get("promotion_gate_passed") is True
        and not non_authoritative_markers(manifest)
        and manifest.get("discovery_records_are_never_deleted") is True
        and manifest.get("automated_semantic_exclusion") is False
        and manifest.get(
            "official_text_signals_require_review_before_inclusion"
        )
        is True
        and manifest.get(
            "resolved_decisions_require_reproducible_official_evidence"
        )
        is True
    )


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(
        description="Run two isolated legislation shadow recomputations."
    )
    result.add_argument("--coverage-run", required=True, type=Path)
    result.add_argument("--accepted-generation", required=True, type=Path)
    result.add_argument("--direction", required=True, type=Path)
    result.add_argument("--base-network", required=True, type=Path)
    result.add_argument("--rollcall", required=True, type=Path)
    result.add_argument("--baseline-manifest", required=True, type=Path)
    result.add_argument("--consumer-plan", required=True, type=Path)
    result.add_argument("--out", required=True, type=Path)
    result.add_argument("--strict-discovery", type=Path)
    result.add_argument("--final-merged", type=Path)
    result.add_argument("--review", type=Path)
    result.add_argument("--relationships", type=Path)
    result.add_argument("--reconciliation", type=Path)
    result.add_argument(
        "--max-rank-movers",
        type=int,
        default=None,
        help=(
            "Maximum number of members whose rank may change before the "
            "rank-drift gate blocks.  Omitted means zero: any movement needs "
            "an explicit human allowance."
        ),
    )
    result.add_argument(
        "--max-abs-rank-delta",
        type=int,
        default=None,
        help=(
            "Maximum absolute rank delta any single member may show before "
            "the rank-drift gate blocks.  Omitted means zero."
        ),
    )
    return result


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def receipt(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "path": str(path.resolve()),
        "size": stat.st_size,
        "sha256": sha256_file(path),
    }


def read_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise ShadowError(f"required JSON does not exist: {path}")
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise ShadowError(f"invalid JSON in {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise ShadowError(f"JSON root must be an object: {path}")
    return value


def write_json(path: Path, value: Any) -> None:
    path.write_text(
        json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False) + "\n",
        encoding="utf-8",
        newline="\n",
    )


def write_ndjson(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.write_text(
        "".join(
            json.dumps(
                row,
                sort_keys=True,
                ensure_ascii=False,
                separators=(",", ":"),
            )
            + "\n"
            for row in rows
        ),
        encoding="utf-8",
        newline="\n",
    )


def _is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def allowed_shadow_roots() -> tuple[Path, ...]:
    roots = {SHADOW_ROOT, REPO_TEMP_ROOT, Path(tempfile.gettempdir()).resolve()}
    for name in ("TEMP", "TMP", "TMPDIR"):
        if os.environ.get(name):
            roots.add(Path(os.environ[name]).resolve())
    return tuple(sorted(roots, key=str))


def validate_out(raw: Path) -> Path:
    out = raw.resolve()
    if out.suffix:
        raise ShadowError("--out must be a directory")
    if not any(_is_relative_to(out, root) for root in allowed_shadow_roots()):
        raise ShadowError(
            "--out must resolve under audits/legfix5/shadow_recompute or temp; "
            f"received {out}"
        )
    if out.exists() and any(out.iterdir()):
        raise ShadowError(
            f"shadow output must be new or empty; existing content found at {out}"
        )
    return out


def recursive_sha_values(value: Any) -> set[str]:
    result: set[str] = set()
    if isinstance(value, dict):
        for key, item in value.items():
            if (
                isinstance(item, str)
                and "sha256" in str(key).casefold()
                and HEX_SHA_RE.fullmatch(item)
            ):
                result.add(item.casefold())
            result.update(recursive_sha_values(item))
    elif isinstance(value, list):
        for item in value:
            result.update(recursive_sha_values(item))
    return result


def explicit_pass(manifest: dict[str, Any]) -> bool:
    if manifest.get("acceptance_passed") is True:
        return True
    if manifest.get("promotion_gate_passed") is True:
        return True
    if (
        manifest.get("generation_complete") is True
        and manifest.get("upstream_acceptance_passed") is True
        and manifest.get("promotion_ready") is True
    ):
        return True
    status = str(manifest.get("status") or "").strip().casefold()
    return status in {"pass", "passed", "accepted"}


def resolve_manifest(
    value: Path,
    label: str,
    schema_tokens: Sequence[str],
) -> tuple[Path, dict[str, Any]]:
    path = value.resolve()
    if path.is_file():
        return path, read_json(path)
    if not path.is_dir():
        raise ShadowError(f"{label} path does not exist: {path}")
    direct = [path / "manifest.json"]
    direct.extend(sorted(path.glob("*.manifest.json")))
    candidates: list[tuple[Path, dict[str, Any]]] = []
    for candidate in direct:
        if not candidate.is_file():
            continue
        manifest = read_json(candidate)
        schema = str(manifest.get("schema") or "").casefold()
        if not schema_tokens or all(token in schema for token in schema_tokens):
            candidates.append((candidate, manifest))
    if len(candidates) == 1:
        return candidates[0]
    if not candidates and (path / "manifest.json").is_file():
        manifest_path = path / "manifest.json"
        return manifest_path, read_json(manifest_path)
    raise ShadowError(
        f"{label} manifest discovery is "
        f"{'ambiguous' if candidates else 'empty'} at {path}"
    )


def resolve_explicit_or_candidates(
    explicit: Path | None,
    run_dir: Path,
    candidates: Sequence[str],
    label: str,
) -> tuple[Path, dict[str, Any]]:
    if explicit is not None:
        path = explicit.resolve()
        return path, read_json(path)
    found = [run_dir / candidate for candidate in candidates]
    found = [path for path in found if path.is_file()]
    if len(found) != 1:
        raise ShadowError(
            f"{label} manifest discovery expected one file, found "
            f"{len(found)} under {run_dir}: {[str(path) for path in found]}"
        )
    return found[0], read_json(found[0])


def query_is_exact(query: dict[str, Any]) -> bool:
    validation = query.get("validation") or {}
    return bool(
        validation.get("exact_result_total")
        and validation.get("unique_result_total")
        and validation.get("every_page_exact")
    )


def validate_coverage(args: argparse.Namespace) -> dict[str, Any]:
    run_dir = args.coverage_run.resolve()
    if not run_dir.is_dir():
        raise ShadowError(f"coverage run does not exist: {run_dir}")
    registry_hash = sha256_file(REGISTRY)
    strict_path, strict = resolve_explicit_or_candidates(
        args.strict_discovery,
        run_dir,
        ("manifest.json",),
        "strict discovery",
    )
    selected = list(strict.get("selected_query_ids") or [])
    queries = list(strict.get("queries") or [])
    completed_ids = [query.get("query_id") for query in queries]
    strict_pass = bool(
        strict.get("status") == "complete"
        and strict.get("registry_sha256") == registry_hash
        and len(selected) == len(set(selected)) == 62
        and len(queries) == 62
        and set(completed_ids) == set(selected)
        and len(completed_ids) == len(set(completed_ids))
        and all(query_is_exact(query) for query in queries)
    )
    if not strict_pass:
        raise ShadowError(
            "strict_discovery_gate_failed: current 62-query Congress.gov "
            "Advanced Search manifest is not complete/exact"
        )

    final_explicit = args.final_merged
    dedicated_final_candidates = (
        "final_merged/manifest.json",
        "merged/manifest.json",
        "final_coverage/manifest.json",
    )
    if (
        final_explicit is None
        and not any((run_dir / name).is_file() for name in dedicated_final_candidates)
    ):
        raise ShadowError(
            "final_merged_gate_failed: a dedicated historical-merged "
            "generation manifest is required; strict discovery cannot "
            "self-certify final recall coverage"
        )
    final_path, final = resolve_explicit_or_candidates(
        final_explicit,
        run_dir,
        dedicated_final_candidates,
        "final merged coverage",
    )
    if final_path == strict_path:
        raise ShadowError(
            "final_merged_gate_failed: strict discovery and final merged "
            "coverage must be distinct manifested generations"
        )
    final_status = str(final.get("status") or "").casefold()
    final_count = (
        final.get("record_count")
        or final.get("unique_record_count")
        or (final.get("union") or {}).get("unique_record_count")
        or (final.get("union") or {}).get("record_count")
    )
    gap_count = (
        final.get("coverage_gap_count")
        if "coverage_gap_count" in final
        else final.get("unresolved_gap_count")
    )
    coverage_evidence = final.get("coverage_evidence")
    evidence_pass = isinstance(coverage_evidence, dict)
    if evidence_pass:
        for key in (
            "missed_candidates_ledger",
            "exclusion_ledger",
            "scope_and_exclusion_contract",
        ):
            item = coverage_evidence.get(key)
            if not isinstance(item, dict):
                evidence_pass = False
                break
            evidence_path = Path(str(item.get("path") or ""))
            if not evidence_path.is_absolute():
                evidence_path = final_path.parent / evidence_path
            evidence_path = evidence_path.resolve()
            if not (
                evidence_path.is_relative_to(final_path.parent.resolve())
                and evidence_path.is_file()
                and item.get("size") == evidence_path.stat().st_size
                and str(item.get("sha256") or "").casefold()
                == sha256_file(evidence_path)
            ):
                evidence_pass = False
                break
    evidence_pass = bool(
        evidence_pass
        and coverage_evidence.get("unresolved_gap_count") == 0
        and coverage_evidence.get("recovered_candidates_remain_visible") is True
        and coverage_evidence.get("semantic_exclusions_are_never_automatic")
        is True
    )
    final_pass = bool(
        final_status in {"complete", "pass", "passed"}
        and isinstance(final_count, int)
        and not isinstance(final_count, bool)
        and final_count > 0
        and isinstance(gap_count, int)
        and not isinstance(gap_count, bool)
        and gap_count == 0
        and evidence_pass
        and (
            explicit_pass(final)
            or final_status == "complete"
        )
    )
    if not final_pass:
        raise ShadowError(
            "final_merged_gate_failed: merged lossless corpus is not complete"
        )

    review_path, review = resolve_explicit_or_candidates(
        args.review,
        run_dir,
        ("candidate_membership_review/manifest.json",),
        "membership review",
    )
    review_pass = bool(
        authoritative_review_gate(review)
        and int(review.get("unresolved_count") or 0) == 0
        and int(review.get("weak_provenance_count") or 0) == 0
    )
    if not review_pass:
        raise ShadowError(
            "membership_review_gate_failed: unresolved or unaccepted memberships remain"
        )

    relationship_path, relationships = resolve_explicit_or_candidates(
        args.relationships,
        run_dir,
        ("relationships/manifest.json",),
        "relationship enrichment",
    )
    relationship_status = str(relationships.get("status") or "").casefold()
    relationship_pass = bool(
        relationship_status in {"complete", "pass", "passed"}
        and (
            explicit_pass(relationships)
            or relationships.get("reconciliation_passed") is True
            or relationships.get("all_relationships_reconciled") is True
        )
    )
    if not relationship_pass:
        raise ShadowError(
            "relationship_gate_failed: sponsor/original/later/withdrawn "
            "reconciliation is not PASS"
        )

    reconciliation_path, reconciliation = resolve_explicit_or_candidates(
        args.reconciliation,
        run_dir,
        ("reconciliation/reconciliation_report.json",),
        "coverage reconciliation",
    )
    reconciliation_pass = bool(
        explicit_pass(reconciliation)
        or (reconciliation.get("acceptance") or {}).get("passed") is True
    )
    if not reconciliation_pass:
        raise ShadowError(
            "reconciliation_gate_failed: final coverage acceptance is not PASS"
        )
    return {
        "registry_sha256": registry_hash,
        "strict_discovery": {
            "passed": True,
            "manifest": receipt(strict_path),
            "query_count": len(queries),
        },
        "final_merged": {
            "passed": True,
            "manifest": receipt(final_path),
            "record_count": final_count,
            "gap_count": gap_count or 0,
        },
        "membership_review": {
            "passed": True,
            "manifest": receipt(review_path),
        },
        "relationships": {
            "passed": True,
            "manifest": receipt(relationship_path),
        },
        "reconciliation": {
            "passed": True,
            "manifest": receipt(reconciliation_path),
        },
        "_manifest_paths": {
            "strict": strict_path,
            "final": final_path,
            "review": review_path,
            "relationships": relationship_path,
            "reconciliation": reconciliation_path,
        },
    }


def output_receipts(
    manifest_path: Path, manifest: dict[str, Any], label: str
) -> dict[str, dict[str, Any]]:
    raw = manifest.get("outputs")
    if (
        label == "accepted_generation"
        and (not isinstance(raw, dict) or not raw)
        and manifest.get("canonical_ledger_file")
        and manifest.get("canonical_ledger_sha256")
    ):
        raw = {
            "canonical_ledger": {
                "path": manifest["canonical_ledger_file"],
                "sha256": manifest["canonical_ledger_sha256"],
            }
        }
    if not isinstance(raw, dict) or not raw:
        raise ShadowError(f"{label} manifest has no output receipts")
    outputs: dict[str, dict[str, Any]] = {}
    for name, item in raw.items():
        if not isinstance(item, dict):
            continue
        if non_authoritative_path(str(name)):
            raise ShadowError(
                f"{label} declares a non-authoritative sidecar output: {name}"
            )
        sha = str(item.get("sha256") or "").casefold()
        raw_path = item.get("path")
        if not HEX_SHA_RE.fullmatch(sha) or not raw_path:
            continue
        path = Path(str(raw_path))
        if not path.is_absolute():
            path = manifest_path.parent / path
        path = path.resolve()
        if non_authoritative_path(path):
            raise ShadowError(
                f"{label} output is a non-authoritative "
                f"review/relevance/prefill/queue sidecar: {path}"
            )
        if not path.is_file():
            raise ShadowError(f"{label} output is missing: {path}")
        if sha256_file(path) != sha:
            raise ShadowError(f"{label} output hash mismatch: {path}")
        expected_size = item.get("size")
        if expected_size is not None and path.stat().st_size != int(expected_size):
            raise ShadowError(f"{label} output size mismatch: {path}")
        outputs[str(name)] = {**receipt(path), "logical_name": str(name)}
    if not outputs:
        raise ShadowError(f"{label} manifest has no valid output receipts")
    return outputs


def accepted_generation_gate(manifest: Mapping[str, Any]) -> bool:
    return bool(
        manifest.get("schema") == ACCEPTED_MANIFEST_SCHEMA
        and manifest.get("generation_complete") is True
        and manifest.get("upstream_acceptance_passed") is True
        and manifest.get("promotion_ready") is True
        and manifest.get("role_weight_policy_reconciled") is True
        and manifest.get("staged_only") is True
        and manifest.get("production_artifacts_written") is False
        and not non_authoritative_markers(manifest)
        and bool(manifest.get("selected_role_weight_policy"))
        and manifest.get("selected_method_profile")
        == ROLE_METHOD_PROFILE_ID
        and manifest.get("consumer_role_methods") == ROLE_METHOD_PROFILE
    )


def validate_producers(
    args: argparse.Namespace, coverage: dict[str, Any]
) -> dict[str, Any]:
    producer_specs = (
        ("accepted_generation", args.accepted_generation, ("accepted", "generation")),
        ("direction", args.direction, ("direction",)),
        ("base_network", args.base_network, ("base", "network")),
        ("rollcall", args.rollcall, ("rollcall",)),
    )
    producers: dict[str, Any] = {}
    for label, path, tokens in producer_specs:
        manifest_path, manifest = resolve_manifest(path, label, tokens)
        if non_authoritative_path(manifest_path):
            raise ShadowError(
                f"{label}_gate_failed: producer manifest is in a "
                "non-authoritative review/relevance/prefill/queue namespace"
            )
        if label == "accepted_generation" and not accepted_generation_gate(
            manifest
        ):
            raise ShadowError(
                "accepted_generation_gate_failed: generation, upstream "
                "acceptance, promotion readiness, and exact consumer-specific "
                "legacy role-method reconciliation must all be explicit"
            )
        if not explicit_pass(manifest):
            raise ShadowError(f"{label}_gate_failed: producer manifest is not PASS")
        outputs = output_receipts(manifest_path, manifest, label)
        producers[label] = {
            "manifest_path": manifest_path,
            "manifest": manifest,
            "manifest_receipt": receipt(manifest_path),
            "outputs": outputs,
        }

    accepted = producers["accepted_generation"]
    required_coverage_hashes = {
        *(
            sha256_file(path)
            for path in coverage["_manifest_paths"].values()
        ),
    }
    accepted_hashes = recursive_sha_values(accepted["manifest"])
    missing_bindings = sorted(required_coverage_hashes - accepted_hashes)
    if missing_bindings:
        raise ShadowError(
            "accepted_generation_currentness_failed: manifest is not hash-bound "
            f"to current coverage gates ({len(missing_bindings)} missing hashes)"
        )

    accepted_output_hashes = {
        item["sha256"] for item in accepted["outputs"].values()
    }
    if not accepted_output_hashes:
        raise ShadowError("accepted generation has no hashed outputs")

    direction = producers["direction"]
    direction_input_hashes = recursive_sha_values(
        direction["manifest"].get("inputs")
    )
    if not (direction_input_hashes & accepted_output_hashes):
        raise ShadowError(
            "direction_currentness_failed: no accepted-generation output hash "
            "is bound into the direction manifest"
        )

    base = producers["base_network"]
    base_input_hashes = recursive_sha_values(base["manifest"].get("inputs"))
    if not (base_input_hashes & accepted_output_hashes):
        raise ShadowError(
            "base_network_currentness_failed: no accepted relationship ledger "
            "hash is bound into the base-network manifest"
        )

    direction_output_hashes = {
        item["sha256"] for item in direction["outputs"].values()
    }
    rollcall = producers["rollcall"]
    rollcall_input_hashes = recursive_sha_values(
        rollcall["manifest"].get("inputs")
    )
    if not (rollcall_input_hashes & accepted_output_hashes):
        raise ShadowError(
            "rollcall_currentness_failed: accepted ledger hash is absent"
        )
    if not (rollcall_input_hashes & direction_output_hashes):
        raise ShadowError(
            "rollcall_currentness_failed: scoring direction hash is absent"
        )
    return producers


def resolve_baseline_path(
    raw_path: Any, baseline_path: Path, baseline: dict[str, Any]
) -> Path:
    path = Path(str(raw_path))
    if path.is_absolute():
        return path.resolve()
    workspace = baseline.get("workspace_root") or baseline.get("source_root")
    base = Path(str(workspace)).resolve() if workspace else baseline_path.parent
    return (base / path).resolve()


def inferred_baseline_role(
    entry: dict[str, Any], baseline: dict[str, Any]
) -> str | None:
    explicit = str(entry.get("role") or "").strip()
    if explicit:
        return explicit
    source = Path(str(entry.get("path") or entry.get("source") or "")).name
    name = source.casefold()
    direct = next(
        (
            role
            for relative, role in BASELINE_ROLE_BY_BASENAME.items()
            if name == Path(relative).name.casefold()
        ),
        None,
    )
    if direct:
        return direct
    selection = baseline.get("selection") or {}
    selected_names = {
        "legislation_archive_workbook": selection.get(
            "archive_selected_by_mtime"
        ),
        "base_network": selection.get("base_network_selected_by_mtime"),
        "augmented_network": selection.get(
            "engine_network_selected_augmented_first"
        ),
    }
    for role, selected in selected_names.items():
        if selected and Path(str(selected)).name.casefold() == name:
            return role
    return None


def load_baseline(path: Path, out: Path) -> dict[str, dict[str, Any]]:
    baseline_path = path.resolve()
    baseline = read_json(baseline_path)
    if baseline.get("schema") != BASELINE_SCHEMA:
        raise ShadowError(
            f"baseline schema must be {BASELINE_SCHEMA!r}"
        )
    raw_files = baseline.get("files")
    entries: list[dict[str, Any]] = []
    if isinstance(raw_files, list):
        entries = [item for item in raw_files if isinstance(item, dict)]
    elif isinstance(raw_files, dict):
        for role, item in raw_files.items():
            if isinstance(item, dict):
                entries.append({"role": role, **item})
    else:
        raise ShadowError("baseline files must be a list or object")
    legacy_scope_complete = bool(
        baseline.get("immutable_existing_copy_policy") is True
        and baseline.get("missing_named_artifacts") == []
        and baseline.get("file_count") == len(entries)
    )
    if baseline.get("scope_complete") is not True and not legacy_scope_complete:
        raise ShadowError(
            "baseline scope is not explicitly complete and its frozen-copy "
            "completeness contract does not pass"
        )
    by_role: dict[str, dict[str, Any]] = {}
    for index, entry in enumerate(entries):
        inferred_role = inferred_baseline_role(entry, baseline)
        role = inferred_role or (
            f"__immutable_guard__:{index:03d}:"
            f"{Path(str(entry.get('source') or entry.get('path') or 'file')).name}"
        )
        if role in by_role:
            raise ShadowError(f"duplicate baseline role: {role}")
        raw_path = entry.get("path") or entry.get("source")
        expected_hash = str(entry.get("sha256") or "").casefold()
        if inferred_role and inferred_role not in REQUIRED_BASELINE_ROLES:
            raise ShadowError(f"invalid baseline role: {inferred_role}")
        if not raw_path:
            raise ShadowError(f"invalid baseline entry: {entry!r}")
        if not HEX_SHA_RE.fullmatch(expected_hash):
            raise ShadowError(f"invalid baseline sha256 for role {role}")
        file_path = resolve_baseline_path(raw_path, baseline_path, baseline)
        if _is_relative_to(file_path, out):
            raise ShadowError(f"baseline role {role} points inside shadow output")
        by_role[role] = {
            "role": role,
            "path": file_path,
            "sha256": expected_hash,
            "size": entry.get("size"),
        }
        raw_copy = entry.get("copy")
        if raw_copy:
            copy_path = resolve_baseline_path(raw_copy, baseline_path, baseline)
            copy_role = f"__frozen_copy__:{index:03d}:{copy_path.name}"
            if copy_role in by_role:
                raise ShadowError(f"duplicate frozen-copy guard: {copy_role}")
            by_role[copy_role] = {
                "role": copy_role,
                "path": copy_path,
                "sha256": expected_hash,
                "size": entry.get("size"),
            }
    missing = sorted(set(REQUIRED_BASELINE_ROLES) - set(by_role))
    if missing:
        raise ShadowError(
            f"baseline is incomplete; missing roles: {', '.join(missing)}"
        )
    validate_baseline_unchanged(by_role, "initial")
    return by_role


def validate_baseline_unchanged(
    baseline: dict[str, dict[str, Any]], phase: str
) -> None:
    for role, entry in baseline.items():
        path = entry["path"]
        if not path.is_file():
            raise ShadowError(
                f"production_baseline_mutated:{phase}:{role}:missing"
            )
        actual = sha256_file(path)
        if actual != entry["sha256"]:
            raise ShadowError(
                f"production_baseline_mutated:{phase}:{role}:"
                f"expected={entry['sha256']} actual={actual}"
            )
        if entry["size"] is not None and path.stat().st_size != int(
            entry["size"]
        ):
            raise ShadowError(
                f"production_baseline_mutated:{phase}:{role}:size"
            )


def safe_relative_output(value: Any) -> Path:
    path = Path(str(value))
    if path.is_absolute() or ".." in path.parts or not path.parts:
        raise ShadowError(f"plan output must be a safe relative path: {value!r}")
    return path


def validate_plan_placeholders(values: Iterable[str], label: str) -> None:
    formatter = string.Formatter()
    for value in values:
        try:
            fields = [
                field
                for _, field, _, _ in formatter.parse(value)
                if field
            ]
        except ValueError as exc:
            raise ShadowError(
                f"{label} has invalid placeholder syntax: {value!r}"
            ) from exc
        unknown = sorted(
            field for field in fields if field not in APPROVED_PLAN_PLACEHOLDERS
        )
        if unknown:
            raise ShadowError(
                f"{label} has unapproved placeholder(s): {unknown!r}"
            )


def load_plan(
    path: Path, baseline: dict[str, dict[str, Any]]
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    plan_path = path.resolve()
    plan = read_json(plan_path)
    if plan.get("schema") != PLAN_SCHEMA:
        raise ShadowError(f"consumer plan schema must be {PLAN_SCHEMA!r}")
    steps = plan.get("steps")
    if not isinstance(steps, list) or not all(
        isinstance(step, dict) for step in steps
    ):
        raise ShadowError("consumer plan steps must be a list of objects")
    names = [step.get("name") for step in steps]
    if tuple(names) != REQUIRED_STEP_ORDER:
        missing = [
            name for name in REQUIRED_STEP_ORDER if name not in set(names)
        ]
        raise ShadowError(
            "required_consumer_cannot_be_isolated: corrected dependency order "
            f"is incomplete or reordered; missing={missing!r}"
        )
    mapped_roles: list[str] = []
    normalized: list[dict[str, Any]] = []
    for step in steps:
        name = str(step["name"])
        if (
            step.get("isolated") is not True
            or step.get("writes_only_under_run_dir") is not True
        ):
            raise ShadowError(
                f"required_consumer_cannot_be_isolated:{name}:"
                "isolation contract absent"
            )
        command = step.get("command")
        if (
            not isinstance(command, list)
            or not command
            or not all(isinstance(token, str) and token for token in command)
        ):
            raise ShadowError(
                f"required_consumer_cannot_be_isolated:{name}:"
                "explicit argv command absent"
            )
        validate_plan_placeholders(command, f"consumer {name} command")
        raw_environment = step.get("environment", {})
        if not isinstance(raw_environment, dict) or not all(
            isinstance(key, str)
            and bool(key)
            and isinstance(value, str)
            and bool(value)
            for key, value in (
                raw_environment.items()
                if isinstance(raw_environment, dict)
                else ()
            )
        ):
            raise ShadowError(
                f"required_consumer_cannot_be_isolated:{name}:"
                "environment keys and values must be nonempty strings"
            )
        validate_plan_placeholders(
            raw_environment.values(), f"consumer {name} environment"
        )
        command_bound = any("{run_dir}" in token for token in command)
        root_bound = raw_environment.get("MONEYBALL_ROOT") == "{run_dir}"
        if not (command_bound or root_bound):
            raise ShadowError(
                f"required_consumer_cannot_be_isolated:{name}:"
                "no explicit {run_dir} argv or MONEYBALL_ROOT binding"
            )
        for key, value in raw_environment.items():
            if key.endswith(("_OUT", "_ROOT")) and not value.startswith(
                "{run_dir}"
            ):
                raise ShadowError(
                    f"required_consumer_cannot_be_isolated:{name}:"
                    f"unsafe environment path override {key}={value!r}"
                )
        raw_outputs = step.get("outputs")
        if not isinstance(raw_outputs, list) or not raw_outputs:
            raise ShadowError(f"consumer {name} has no declared outputs")
        outputs = []
        for raw in raw_outputs:
            if not isinstance(raw, dict):
                raise ShadowError(f"consumer {name} output is not an object")
            relative = safe_relative_output(raw.get("path"))
            role = raw.get("baseline_role")
            intermediate = raw.get("intermediate") is True
            if bool(role) == intermediate:
                raise ShadowError(
                    f"consumer {name}/{relative} must declare exactly one of "
                    "baseline_role or intermediate=true"
                )
            if role:
                if role not in baseline:
                    raise ShadowError(
                        f"consumer {name} references unknown baseline role {role}"
                    )
                mapped_roles.append(str(role))
            outputs.append({**raw, "path": relative})
        normalized.append(
            {
                **step,
                "environment": dict(raw_environment),
                "outputs": outputs,
            }
        )
    if sorted(mapped_roles) != sorted(REQUIRED_BASELINE_ROLES):
        missing = sorted(set(REQUIRED_BASELINE_ROLES) - set(mapped_roles))
        duplicates = sorted(
            role for role in set(mapped_roles) if mapped_roles.count(role) > 1
        )
        raise ShadowError(
            "old_new_attribution_scope_incomplete: every production baseline "
            f"role must map exactly once; missing={missing!r} duplicates={duplicates!r}"
        )
    return {
        "path": plan_path,
        "manifest": plan,
        "receipt": receipt(plan_path),
    }, normalized


def _safe_run_child(run_dir: Path, relative: str | Path) -> Path:
    raw = Path(relative)
    if raw.is_absolute() or ".." in raw.parts or not raw.parts:
        raise ShadowError(f"unsafe immutable input path: {relative!r}")
    result = (run_dir / raw).resolve()
    if not _is_relative_to(result, run_dir.resolve()):
        raise ShadowError(f"immutable input escapes shadow run: {relative}")
    return result


def _copy_immutable(
    source: Path,
    destination: Path,
    *,
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    before = sha256_file(source)
    if expected_sha256 is not None and before != expected_sha256:
        raise ShadowError(
            f"immutable source hash mismatch: {source}: "
            f"expected={expected_sha256} observed={before}"
        )
    if destination.exists():
        raise ShadowError(
            f"immutable input destination already exists: {destination}"
        )
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    copied = sha256_file(destination)
    after = sha256_file(source)
    if before != copied or before != after:
        raise ShadowError(
            f"immutable input changed or copied incorrectly: {source}"
        )
    return {
        "source_path": str(source.resolve()),
        "source_size": source.stat().st_size,
        "source_sha256": before,
        "materialized_path": str(destination),
        "materialized_size": destination.stat().st_size,
        "materialized_sha256": copied,
    }


def _accepted_manual_source(
    producers: Mapping[str, Any],
) -> tuple[Path, str]:
    accepted = producers.get("accepted_generation") or {}
    outputs = accepted.get("outputs")
    if not isinstance(outputs, Mapping):
        raise ShadowError(
            "accepted generation has no outputs for manual evidence"
        )
    matches = []
    for logical_name, item in outputs.items():
        if not isinstance(item, Mapping):
            continue
        path = Path(str(item.get("path") or "")).resolve()
        normalized = str(logical_name).casefold().replace("-", "_")
        if (
            path.name.casefold() == CANONICAL_MANUAL_EVIDENCE.casefold()
            or normalized
            in {"manual_bill_evidence", "manual_bill_evidence.csv"}
        ):
            matches.append((path, str(item.get("sha256") or "").casefold()))
    if len(matches) != 1:
        raise ShadowError(
            "accepted generation must declare exactly one hash-bound "
            f"{CANONICAL_MANUAL_EVIDENCE}; found {len(matches)}"
        )
    source, expected = matches[0]
    if not source.is_file() or not HEX_SHA_RE.fullmatch(expected):
        raise ShadowError("accepted manual evidence receipt is invalid")
    return source, expected


def stage_immutable_consumer_inputs(
    run_dir: Path,
    *,
    producers: Mapping[str, Any],
    baseline: Mapping[str, Mapping[str, Any]],
    workspace_root: Path = ROOT,
    contracts: Sequence[ImmutableInputContract] = (
        IMMUTABLE_INPUT_CONTRACTS
    ),
) -> dict[str, Any]:
    """Freeze all non-legislation raw reads before the first consumer runs."""
    paths = [contract.path for contract in contracts]
    duplicates = sorted(
        path for path in set(paths) if paths.count(path) != 1
    )
    if duplicates:
        raise ShadowError(
            f"immutable input contract contains duplicates: {duplicates}"
        )
    missing = [
        contract.path
        for contract in contracts
        if contract.required
        and not (workspace_root / contract.path).resolve().is_file()
    ]
    if missing:
        raise ShadowError(
            "required immutable consumer inputs are missing: "
            + ", ".join(sorted(missing))
        )

    records: list[dict[str, Any]] = []
    absent_optional: list[str] = []
    for contract in contracts:
        source = (workspace_root / contract.path).resolve()
        if not source.is_file():
            absent_optional.append(contract.path)
            records.append(
                {
                    "source_group": "immutable_nonlegislation",
                    "logical_name": contract.path,
                    "required": False,
                    "consumers": list(contract.consumers),
                    "status": "ABSENT_OPTIONAL",
                }
            )
            continue
        destination = _safe_run_child(run_dir, contract.path)
        copied = _copy_immutable(source, destination)
        records.append(
            {
                "source_group": "immutable_nonlegislation",
                "logical_name": contract.path,
                "required": contract.required,
                "consumers": list(contract.consumers),
                "status": "STAGED",
                **{
                    **copied,
                    "materialized_path": str(
                        destination.relative_to(run_dir)
                    ).replace("\\", "/"),
                },
            }
        )

    manual_source, manual_sha = _accepted_manual_source(producers)
    manual_destination = _safe_run_child(
        run_dir, CANONICAL_MANUAL_EVIDENCE
    )
    manual_copy = _copy_immutable(
        manual_source,
        manual_destination,
        expected_sha256=manual_sha,
    )
    records.append(
        {
            "source_group": "accepted_generation",
            "logical_name": "manual_bill_evidence",
            "required": True,
            "generation_bound": True,
            "consumers": ["workbook", "archive", "member_desk"],
            "status": "STAGED_CANONICAL_ROOT",
            **{
                **manual_copy,
                "materialized_path": CANONICAL_MANUAL_EVIDENCE,
            },
        }
    )

    member_desk = baseline.get("member_desk")
    if not isinstance(member_desk, Mapping):
        raise ShadowError(
            "immutable baseline member_desk role is unavailable"
        )
    member_source = Path(str(member_desk.get("path") or "")).resolve()
    member_hash = str(member_desk.get("sha256") or "").casefold()
    if not member_source.is_file() or not HEX_SHA_RE.fullmatch(member_hash):
        raise ShadowError("baseline member_desk receipt is invalid")
    canonical_member = _safe_run_child(
        run_dir, "inputs/immutable/member_desk_bundle.json"
    )
    member_copy = _copy_immutable(
        member_source,
        canonical_member,
        expected_sha256=member_hash,
    )
    records.append(
        {
            "source_group": "production_baseline",
            "logical_name": "member_desk_bundle_first_pass",
            "baseline_role": "member_desk",
            "required": True,
            "transient_root_compatibility": True,
            "consumers": ["say_do_first_pass"],
            "status": "STAGED_CANONICAL",
            **{
                **member_copy,
                "materialized_path": str(
                    canonical_member.relative_to(run_dir)
                ).replace("\\", "/"),
                "compatibility_path": "member_desk_bundle.json",
            },
        }
    )
    return {
        "schema": "legislation-shadow-immutable-input-closure-v1",
        "status": "PASS",
        "production_written": False,
        "required_count": sum(
            1 for contract in contracts if contract.required
        )
        + 2,
        "staged_count": sum(
            1 for row in records if str(row.get("status")).startswith("STAGED")
        ),
        "absent_optional": sorted(absent_optional),
        "records": records,
        "transient_inputs": {
            "member_desk_bundle.json": {
                "canonical_path": str(
                    canonical_member.relative_to(run_dir)
                ).replace("\\", "/"),
                "sha256": member_hash,
            }
        },
    }


def materialize_inputs(
    run_dir: Path,
    coverage: dict[str, Any],
    producers: dict[str, Any],
    baseline: Mapping[str, Mapping[str, Any]],
    *,
    workspace_root: Path = ROOT,
) -> dict[str, Any]:
    input_root = run_dir / "inputs"
    input_root.mkdir(parents=True)
    rows: list[dict[str, Any]] = []
    for label, path in sorted(coverage["_manifest_paths"].items()):
        destination_dir = input_root / "coverage"
        destination_dir.mkdir(exist_ok=True)
        destination = destination_dir / f"{label}.json"
        shutil.copyfile(path, destination)
        rows.append(
            {
                "source_group": "coverage",
                "logical_name": label,
                "source_sha256": sha256_file(path),
                "materialized_path": str(destination.relative_to(run_dir)),
                "materialized_sha256": sha256_file(destination),
            }
        )
    for label, producer in sorted(producers.items()):
        destination_dir = input_root / label
        destination_dir.mkdir()
        manifest_destination = destination_dir / "manifest.json"
        shutil.copyfile(producer["manifest_path"], manifest_destination)
        rows.append(
            {
                "source_group": label,
                "logical_name": "manifest",
                "source_sha256": sha256_file(producer["manifest_path"]),
                "materialized_path": str(
                    manifest_destination.relative_to(run_dir)
                ),
                "materialized_sha256": sha256_file(manifest_destination),
            }
        )
        for index, (logical_name, item) in enumerate(
            sorted(producer["outputs"].items())
        ):
            source = Path(item["path"])
            destination = (
                destination_dir
                / f"{index:03d}_{source.name}"
            )
            shutil.copyfile(source, destination)
            copied_hash = sha256_file(destination)
            if copied_hash != item["sha256"]:
                raise ShadowError(
                    f"materialized input hash mismatch: {label}/{logical_name}"
                )
            rows.append(
                {
                    "source_group": label,
                    "logical_name": logical_name,
                    "source_sha256": item["sha256"],
                    "materialized_path": str(destination.relative_to(run_dir)),
                    "materialized_sha256": copied_hash,
                }
            )
    immutable = stage_immutable_consumer_inputs(
        run_dir,
        producers=producers,
        baseline=baseline,
        workspace_root=workspace_root,
    )
    rows.extend(immutable["records"])
    write_json(
        run_dir / "input_materialization.json",
        {
            "schema": "legislation-shadow-input-materialization-v1",
            "status": "PASS",
            "staged_only": True,
            "production_written": False,
            "immutable_input_closure": {
                key: value
                for key, value in immutable.items()
                if key not in {"records", "transient_inputs"}
            },
            "records": rows,
        },
    )
    return {
        "input_root": input_root,
        "records": rows,
        "immutable": immutable,
        "receipt": receipt(run_dir / "input_materialization.json"),
    }


def _run_relative_receipt(path: Path, run_dir: Path) -> dict[str, Any]:
    resolved = path.resolve()
    if not resolved.is_file() or not _is_relative_to(resolved, run_dir):
        raise ShadowError(f"staged prerequisite artifact missing: {resolved}")
    return {
        "path": str(resolved.relative_to(run_dir)).replace("\\", "/"),
        "size": resolved.stat().st_size,
        "sha256": sha256_file(resolved),
    }


def _ledger_step_output(
    ledger: Mapping[str, Any], step_name: str, relative: str
) -> dict[str, Any]:
    normalized = relative.replace("\\", "/")
    matches = []
    for step in ledger.get("completed_steps") or []:
        if not isinstance(step, Mapping) or step.get("name") != step_name:
            continue
        outputs = step.get("outputs")
        if isinstance(outputs, Mapping):
            item = outputs.get(normalized)
            if isinstance(item, Mapping):
                matches.append(item)
    if len(matches) != 1:
        raise ShadowError(
            f"step ledger must declare exactly one {step_name}/{normalized}; "
            f"found {len(matches)}"
        )
    sha = str(matches[0].get("sha256") or "").casefold()
    if not HEX_SHA_RE.fullmatch(sha):
        raise ShadowError(
            f"step ledger output lacks SHA-256: {step_name}/{normalized}"
        )
    return {"sha256": sha, "size": matches[0].get("size")}


def _validated_generation_id(ledger: Mapping[str, Any]) -> str:
    generation_id = str(ledger.get("generation_id") or "").casefold()
    if not (
        ledger.get("schema") == STEP_LEDGER_SCHEMA
        and ledger.get("status") == "RUNNING"
        and ledger.get("staged_only") is True
        and ledger.get("production_written") is False
        and HEX_SHA_RE.fullmatch(generation_id)
    ):
        raise ShadowError("step ledger is not a running accepted generation")
    return generation_id


def _write_ledger_prefix_snapshot(
    run_dir: Path,
    ledger: Mapping[str, Any],
    *,
    consumer_step: str,
) -> dict[str, Any]:
    """Freeze the completed-step prefix before the mutable ledger advances."""
    snapshot = _safe_run_child(
        run_dir,
        Path("intermediate")
        / "step_ledger_prefixes"
        / f"{consumer_step}.json",
    )
    if snapshot.exists():
        raise ShadowError(
            f"step-ledger prefix snapshot already exists: {consumer_step}"
        )
    snapshot.parent.mkdir(parents=True, exist_ok=True)
    write_json(snapshot, dict(ledger))
    return _run_relative_receipt(snapshot, run_dir)


def build_pillars_ranks_prerequisites(
    run_dir: Path, ledger: Mapping[str, Any]
) -> Path:
    """Create the first-pass pillar/rank prerequisite at its consumption edge."""
    generation_id = _validated_generation_id(ledger)
    snapshot_name = "data/engine/pillar_inputs_snapshot.json"
    workbook = _safe_run_child(
        run_dir, "Moneyball Main Version - Blended.xlsx"
    )
    snapshot = _safe_run_child(run_dir, snapshot_name)
    roles = _safe_run_child(run_dir, "data/members/formal_visibility_roles.json")
    workbook_receipt_path = _safe_run_child(
        run_dir, "legislation_workbook.receipt.json"
    )
    for path in (workbook, snapshot, roles, workbook_receipt_path):
        if not path.is_file():
            raise ShadowError(
                f"pillar/rank prerequisite source is missing: {path.name}"
            )
    ledger_workbook = _ledger_step_output(
        ledger,
        "recompute_legislative_raw_normalized_sublayer",
        workbook.name,
    )
    ledger_receipt = _ledger_step_output(
        ledger,
        "recompute_legislative_raw_normalized_sublayer",
        workbook_receipt_path.name,
    )
    if (
        ledger_workbook["sha256"] != sha256_file(workbook)
        or ledger_receipt["sha256"] != sha256_file(workbook_receipt_path)
    ):
        raise ShadowError(
            "pillar/rank workbook or receipt is stale for the step ledger"
        )
    workbook_receipt_value = read_json(workbook_receipt_path)
    receipt_outputs = workbook_receipt_value.get("outputs")
    if not (
        workbook_receipt_value.get("schema")
        == "legislation-workbook-recompute-receipt-v1"
        and workbook_receipt_value.get("status") == "PASS"
        and workbook_receipt_value.get("staged_only") is True
        and workbook_receipt_value.get("production_written") is False
        and workbook_receipt_value.get("generation_id") == generation_id
        and isinstance(receipt_outputs, Mapping)
        and (receipt_outputs.get(workbook.name) or {}).get("sha256")
        == sha256_file(workbook)
        and (receipt_outputs.get(snapshot_name) or {}).get("sha256")
        == sha256_file(snapshot)
    ):
        raise ShadowError(
            "legislative workbook receipt does not bind pillar inputs"
        )
    artifacts = {
        workbook.name: _run_relative_receipt(workbook, run_dir),
        snapshot_name: _run_relative_receipt(snapshot, run_dir),
        # Keyed by the repo-relative name like snapshot_name, not Path.name: the two were
        # identical while this artifact lived at the root, and diverged when Phase 2 moved
        # it under data/members/.
        "data/members/formal_visibility_roles.json": _run_relative_receipt(roles, run_dir),
    }
    value = {
        "schema": "legislation-pillars-ranks-prerequisites-v1",
        "status": "PASS",
        "staged_only": True,
        "production_written": False,
        "mainframe_written": False,
        "generation_id": generation_id,
        "artifacts": artifacts,
        "legislative_workbook_receipt": _run_relative_receipt(
            workbook_receipt_path, run_dir
        ),
        "source_step_ledger": _write_ledger_prefix_snapshot(
            run_dir,
            ledger,
            consumer_step=PILLAR_PREREQUISITE_STEP,
        ),
    }
    output = _safe_run_child(run_dir, "pillars_ranks_prerequisites.json")
    write_json(output, value)
    return output


NETWORK_FIRST_SOURCE_STEPS = {
    "Moneyball Main Version - Blended.xlsx": (
        "recompute_legislative_raw_normalized_sublayer"
    ),
    "ranks_python.json": "recompute_pillars_and_ranks",
    "category_stance.json": "recompute_category_stance",
    "congress_cosponsor_network_AUG_shadow.xlsx": (
        "recompute_augmented_network"
    ),
    "member_receptivity.json": "recompute_receptivity",
    "bipartisan.json": (
        "materialize_base_network_and_prebehavior_bipartisan"
    ),
}


def _write_network_binding(
    run_dir: Path,
    *,
    logical_name: str,
    receipt_name: str,
    expected_schema: str | None,
    artifact_path: str,
    artifact_sha256: str,
    source_step: str,
    generation_id: str,
    source_receipt: dict[str, Any] | None = None,
) -> dict[str, Any]:
    binding = _safe_run_child(
        run_dir,
        Path("intermediate")
        / "network_first_pass_bindings"
        / f"{logical_name}.receipt.json",
    )
    binding.parent.mkdir(parents=True, exist_ok=True)
    value = {
        "schema": expected_schema
        or "legislation-first-pass-artifact-binding-v1",
        "status": "PASS",
        "receipt_role": "accepted-generation-binding-adapter",
        "generation_id": generation_id,
        "staged_only": True,
        "production_written": False,
        "mainframe_written": False,
        "source_step": source_step,
        "artifact": {
            "logical_name": logical_name,
            "path": artifact_path,
            "sha256": artifact_sha256,
        },
        "source_receipt": source_receipt,
    }
    write_json(binding, value)
    return {
        **_run_relative_receipt(binding, run_dir),
        "generation_id": generation_id,
        "receipt_name": receipt_name,
    }


def build_network_dependents_prerequisites(
    run_dir: Path, ledger: Mapping[str, Any]
) -> Path:
    """Create and validate the first-pass network prerequisite manifest."""
    from legislation import run_legislation_network_dependents as network

    generation_id = _validated_generation_id(ledger)
    artifacts: dict[str, dict[str, Any]] = {}
    upstream: dict[str, dict[str, Any]] = {}
    say_do_receipt_path = _safe_run_child(
        run_dir, "intermediate/say_do_first_pass.receipt.json"
    )
    say_do_value = (
        read_json(say_do_receipt_path)
        if say_do_receipt_path.is_file()
        else {}
    )
    for logical_name, contract in sorted(network.REQUIRED_ARTIFACTS.items()):
        artifact = _safe_run_child(run_dir, contract.path)
        if not artifact.is_file():
            raise ShadowError(
                f"network prerequisite artifact is missing: {contract.path}"
            )
        artifact_hash = sha256_file(artifact)
        item: dict[str, Any] = {
            "path": contract.path,
            "sha256": artifact_hash,
            "generation_bound": contract.generation_bound,
        }
        if contract.generation_bound:
            receipt_name = contract.upstream_receipt
            if not receipt_name:
                raise ShadowError(
                    f"generation-bound network input lacks receipt: {logical_name}"
                )
            if logical_name == "say_do_first_pass":
                ledger_receipt = _ledger_step_output(
                    ledger,
                    SAY_DO_FIRST_STEP,
                    "intermediate/say_do_first_pass.receipt.json",
                )
                if not (
                    ledger_receipt["sha256"]
                    == sha256_file(say_do_receipt_path)
                    and say_do_value.get("schema")
                    == "legislation-say-do-pass-receipt-v1"
                    and say_do_value.get("status") == "PASS"
                    and say_do_value.get("pass") == "first"
                    and say_do_value.get("generation_id") == generation_id
                    and (say_do_value.get("root_output") or {}).get("sha256")
                    == artifact_hash
                ):
                    raise ShadowError(
                        "first-pass say/do artifact is stale or unbound"
                    )
                source_step = SAY_DO_FIRST_STEP
                source_receipt = _run_relative_receipt(
                    say_do_receipt_path, run_dir
                )
            else:
                source_step = NETWORK_FIRST_SOURCE_STEPS.get(contract.path)
                if not source_step:
                    raise ShadowError(
                        f"network input has no source step: {contract.path}"
                    )
                ledger_artifact = _ledger_step_output(
                    ledger, source_step, contract.path
                )
                if ledger_artifact["sha256"] != artifact_hash:
                    raise ShadowError(
                        f"network artifact is stale for step ledger: {logical_name}"
                    )
                source_receipt = None
            bound = _write_network_binding(
                run_dir,
                logical_name=logical_name,
                receipt_name=receipt_name,
                expected_schema=network.EXPECTED_RECEIPT_SCHEMAS.get(
                    receipt_name
                ),
                artifact_path=contract.path,
                artifact_sha256=artifact_hash,
                source_step=source_step,
                generation_id=generation_id,
                source_receipt=source_receipt,
            )
            existing = upstream.get(receipt_name)
            entry = {
                key: value
                for key, value in bound.items()
                if key != "receipt_name"
            }
            if existing is not None and existing != entry:
                raise ShadowError(
                    f"network upstream receipt collision: {receipt_name}"
                )
            upstream[receipt_name] = entry
            item.update(
                {
                    "generation_id": generation_id,
                    "upstream_receipt": receipt_name,
                }
            )
        artifacts[logical_name] = item
    value = {
        "schema": network.PREREQUISITES_SCHEMA,
        "status": "PASS",
        "phase": "first",
        "staged_only": True,
        "production_written": False,
        "mainframe_written": False,
        "generation_id": generation_id,
        "artifacts": artifacts,
        "upstream_receipts": upstream,
        "source_step_ledger": _write_ledger_prefix_snapshot(
            run_dir,
            ledger,
            consumer_step=NETWORK_PREREQUISITE_STEP,
        ),
    }
    output = _safe_run_child(
        run_dir, "network_dependents_prerequisites.json"
    )
    write_json(output, value)
    try:
        network.validate_prerequisites(run_dir, output)
    except Exception as exc:
        raise ShadowError(
            f"first-pass network prerequisite is invalid: {exc}"
        ) from exc
    return output


def prepare_step_prerequisites(
    run_dir: Path,
    *,
    step_name: str,
    ledger: Mapping[str, Any],
    materialized: Mapping[str, Any],
) -> dict[str, Any]:
    prepared: dict[str, Any] = {}
    if step_name == PILLAR_PREREQUISITE_STEP:
        path = build_pillars_ranks_prerequisites(run_dir, ledger)
        prepared["pillars_ranks"] = _run_relative_receipt(path, run_dir)
    if step_name == NETWORK_PREREQUISITE_STEP:
        path = build_network_dependents_prerequisites(run_dir, ledger)
        prepared["network_dependents"] = _run_relative_receipt(path, run_dir)
    if step_name == SAY_DO_FIRST_STEP:
        transient = (
            (materialized.get("immutable") or {})
            .get("transient_inputs", {})
            .get("member_desk_bundle.json")
        )
        if not isinstance(transient, Mapping):
            raise ShadowError("say/do baseline member-desk input is unsealed")
        source = _safe_run_child(
            run_dir, str(transient.get("canonical_path") or "")
        )
        expected = str(transient.get("sha256") or "").casefold()
        destination = _safe_run_child(run_dir, "member_desk_bundle.json")
        copied = _copy_immutable(
            source, destination, expected_sha256=expected
        )
        prepared["transient_member_desk_bundle"] = {
            **copied,
            "path": "member_desk_bundle.json",
            "canonical_path": str(
                source.relative_to(run_dir)
            ).replace("\\", "/"),
        }
    return prepared


def release_step_transients(
    run_dir: Path, *, step_name: str, prepared: Mapping[str, Any]
) -> None:
    if (
        step_name == SAY_DO_FIRST_STEP
        and "transient_member_desk_bundle" in prepared
    ):
        path = _safe_run_child(run_dir, "member_desk_bundle.json")
        expected = str(
            prepared["transient_member_desk_bundle"].get(
                "materialized_sha256"
            )
            or ""
        ).casefold()
        if not path.is_file() or sha256_file(path) != expected:
            raise ShadowError(
                "say/do transient member-desk compatibility input changed"
            )
        path.unlink()


def format_values(
    run_dir: Path,
    materialized: dict[str, Any],
) -> dict[str, str]:
    return {
        "run_dir": str(run_dir),
        "workspace_root": str(ROOT),
        "input_root": str(materialized["input_root"]),
        "accepted_input_dir": str(
            materialized["input_root"] / "accepted_generation"
        ),
        "direction_input_dir": str(
            materialized["input_root"] / "direction"
        ),
        "base_network_input_dir": str(
            materialized["input_root"] / "base_network"
        ),
        "rollcall_input_dir": str(
            materialized["input_root"] / "rollcall"
        ),
    }


def format_command(
    tokens: Sequence[str],
    run_dir: Path,
    materialized: dict[str, Any],
) -> list[str]:
    values = format_values(run_dir, materialized)
    try:
        return [token.format(**values) for token in tokens]
    except KeyError as exc:
        raise ShadowError(f"unknown consumer-plan placeholder: {exc}") from exc


def format_environment(
    raw: dict[str, str],
    run_dir: Path,
    materialized: dict[str, Any],
) -> dict[str, str]:
    values = format_values(run_dir, materialized)
    try:
        return {
            key: value.format(**values)
            for key, value in raw.items()
        }
    except KeyError as exc:
        raise ShadowError(
            f"unknown consumer-plan environment placeholder: {exc}"
        ) from exc


def semantic_hash(path: Path) -> str:
    suffix = path.suffix.casefold()
    if suffix == ".json":
        value = json.loads(path.read_text(encoding="utf-8-sig"))
        encoded = json.dumps(
            value, sort_keys=True, ensure_ascii=False, separators=(",", ":")
        ).encode("utf-8")
        return sha256_bytes(encoded)
    if suffix in {".ndjson", ".jsonl"}:
        values = [
            json.loads(line)
            for line in path.read_text(encoding="utf-8-sig").splitlines()
            if line.strip()
        ]
        encoded = json.dumps(
            values, sort_keys=True, ensure_ascii=False, separators=(",", ":")
        ).encode("utf-8")
        return sha256_bytes(encoded)
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            values = list(csv.DictReader(handle))
        encoded = json.dumps(
            values, sort_keys=True, ensure_ascii=False, separators=(",", ":")
        ).encode("utf-8")
        return sha256_bytes(encoded)
    if suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise ShadowError(
                f"cannot value-hash workbook without openpyxl: {path}"
            ) from exc
        workbook = openpyxl.load_workbook(
            path, read_only=True, data_only=False
        )
        values = {
            sheet_name: [
                list(row)
                for row in workbook[sheet_name].iter_rows(values_only=True)
            ]
            for sheet_name in workbook.sheetnames
        }
        workbook.close()
        encoded = json.dumps(
            values,
            sort_keys=True,
            ensure_ascii=False,
            default=str,
            separators=(",", ":"),
        ).encode("utf-8")
        return sha256_bytes(encoded)
    return sha256_file(path)


def load_rank_map(path: Path, *, label: str) -> dict[str, int]:
    """Read the member -> rank mapping out of a ranks_python.json artifact."""

    try:
        document = read_json(path)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        raise ShadowError(f"rank_drift_unreadable:{label}:{exc}") from exc
    ranks = document.get("rank") if isinstance(document, Mapping) else None
    if not isinstance(ranks, Mapping) or not ranks:
        raise ShadowError(
            f"rank_drift_rank_map_absent:{label}:{path}"
        )
    resolved: dict[str, int] = {}
    for member, value in ranks.items():
        if isinstance(value, bool) or not isinstance(value, int):
            raise ShadowError(
                f"rank_drift_non_integer_rank:{label}:{member}"
            )
        resolved[str(member)] = int(value)
    return resolved


def compute_rank_drift(
    baseline_path: Path,
    recomputed_path: Path,
    *,
    max_movers: int,
    max_abs_delta: int,
    threshold_source: str,
) -> dict[str, Any]:
    """Measure how far the recomputation moved the member ranking.

    Reports the count of members whose rank changed, the largest absolute rank
    delta, and the ordered top movers.  Roster churn (a member appearing or
    disappearing) is drift too and is reported separately rather than being
    silently dropped from the comparison.
    """

    old = load_rank_map(baseline_path, label="baseline")
    new = load_rank_map(recomputed_path, label="recomputed")
    added = sorted(set(new) - set(old))
    removed = sorted(set(old) - set(new))
    shared = sorted(set(old) & set(new))
    deltas = {
        member: new[member] - old[member]
        for member in shared
        if new[member] != old[member]
    }
    movers = sorted(
        (
            {
                "member": member,
                "baseline_rank": old[member],
                "recomputed_rank": new[member],
                "delta": delta,
                "abs_delta": abs(delta),
            }
            for member, delta in deltas.items()
        ),
        key=lambda row: (-row["abs_delta"], row["member"]),
    )
    max_abs_observed = max((row["abs_delta"] for row in movers), default=0)
    total_abs = sum(row["abs_delta"] for row in movers)
    roster_stable = not added and not removed
    within_threshold = (
        len(movers) <= max_movers and max_abs_observed <= max_abs_delta
    )
    return {
        "schema": RANK_DRIFT_SCHEMA,
        "baseline_path": str(baseline_path),
        "recomputed_path": str(recomputed_path),
        "baseline_member_count": len(old),
        "recomputed_member_count": len(new),
        "members_compared": len(shared),
        "members_added": added,
        "members_removed": removed,
        "roster_stable": roster_stable,
        "members_moved": len(movers),
        "max_abs_rank_delta": max_abs_observed,
        "mean_abs_rank_delta": (
            round(total_abs / len(shared), 6) if shared else 0.0
        ),
        "top_movers": movers[:RANK_DRIFT_TOP_MOVER_LIMIT],
        "top_mover_limit": RANK_DRIFT_TOP_MOVER_LIMIT,
        "threshold": {
            "max_members_moved": max_movers,
            "max_abs_rank_delta": max_abs_delta,
            "source": threshold_source,
        },
        "within_threshold": within_threshold,
        "drift_measured": True,
        "gate_passed": bool(within_threshold and roster_stable),
    }


def attribution_authorship(attribution: Mapping[str, Any]) -> str:
    """Classify who wrote a delta rationale.

    Absent authorship is machine authorship: the consumer plan builder stamps
    an identical rationale onto every baseline role, so silence means nobody
    looked.  A rationale that still matches a known auto-fill template is
    machine-written no matter what the ``authored_by`` field claims.
    """

    rationale = str(attribution.get("rationale") or "").strip().casefold()
    if rationale in MACHINE_RATIONALE_TEMPLATES:
        return "machine_generated"
    author = str(attribution.get("authored_by") or "").strip().casefold()
    if author in HUMAN_ATTRIBUTION_AUTHORS:
        return "human"
    return "machine_generated"


def validate_attribution(
    output: dict[str, Any],
    path: Path,
    baseline: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    role = output.get("baseline_role")
    new = receipt(path)
    row = {
        "schema": "legislation-shadow-old-new-attribution-v1",
        "output_path": str(output["path"]),
        "baseline_role": role,
        "new_sha256": new["sha256"],
        "new_value_sha256": semantic_hash(path),
    }
    if not role:
        row.update(
            {
                "baseline_sha256": None,
                "changed": True,
                "attribution_status": "staged_intermediate",
                "attribution_authorship": "machine_generated",
                "attribution": {
                    "drivers": ["accepted_legislation"],
                    "rationale": "Declared isolated intermediate artifact.",
                    "authored_by": "machine",
                },
            }
        )
        return row
    old = baseline[str(role)]
    changed = new["sha256"] != old["sha256"]
    row.update(
        {
            "baseline_path": str(old["path"]),
            "baseline_sha256": old["sha256"],
            "changed": changed,
        }
    )
    if not changed:
        row["attribution_status"] = "unchanged"
        row["attribution_authorship"] = None
        row["attribution"] = None
        return row
    attribution = output.get("attribution")
    if not isinstance(attribution, dict):
        raise ShadowError(
            f"unattributed_old_new_delta:{role}: attribution object absent"
        )
    drivers = attribution.get("drivers")
    rationale = str(attribution.get("rationale") or "").strip()
    if (
        not isinstance(drivers, list)
        or not drivers
        or not all(driver in ALLOWED_ATTRIBUTION_DRIVERS for driver in drivers)
        or not rationale
    ):
        raise ShadowError(
            f"unattributed_old_new_delta:{role}: drivers/rationale invalid"
        )
    if (
        "approved_computation_repair" in drivers
        and attribution.get("approved") is not True
    ):
        raise ShadowError(
            f"unapproved_computation_repair_delta:{role}"
        )
    authorship = attribution_authorship(attribution)
    row["attribution_authorship"] = authorship
    # A machine-written rationale is recorded, not accepted.  This is the one
    # path on which validate_attribution can return a status outside the
    # accepted set, which is what makes the old_new_delta_attribution gate a
    # real measurement instead of a tautology.
    row["attribution_status"] = (
        "attributed"
        if authorship == "human"
        else "attributed_machine_generated"
    )
    row["attribution"] = attribution
    return row


def execute_run(
    run_dir: Path,
    steps: list[dict[str, Any]],
    coverage: dict[str, Any],
    producers: dict[str, Any],
    baseline: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    run_dir.mkdir()
    materialized = materialize_inputs(
        run_dir, coverage, producers, baseline
    )
    validate_baseline_unchanged(baseline, f"{run_dir.name}:materialization")
    accepted_manifest = producers["accepted_generation"]["manifest"]
    generation_id = str(
        accepted_manifest.get("canonical_ledger_sha256") or ""
    ).casefold()
    if not HEX_SHA_RE.fullmatch(generation_id):
        raise ShadowError(
            "accepted_generation_id_absent: canonical ledger SHA-256 is "
            "required for downstream step receipts"
        )
    step_ledger_path = run_dir / STEP_LEDGER_NAME
    step_ledger_path.parent.mkdir(parents=True, exist_ok=True)
    step_ledger = {
        "schema": STEP_LEDGER_SCHEMA,
        "status": "RUNNING",
        "staged_only": True,
        "production_written": False,
        "generation_id": generation_id,
        "accepted_generation": {
            "manifest_sha256": producers["accepted_generation"][
                "manifest_receipt"
            ]["sha256"],
            "canonical_ledger_sha256": generation_id,
        },
        "completed_steps": [],
    }
    write_json(step_ledger_path, step_ledger)
    outputs: dict[str, dict[str, Any]] = {}
    attribution_rows: list[dict[str, Any]] = []
    step_receipts: list[dict[str, Any]] = []
    log_dir = run_dir / "logs"
    log_dir.mkdir()
    base_environment = dict(os.environ)
    base_environment.update(
        {
            "LEGISLATION_SHADOW_MODE": "1",
            "LEGISLATION_SHADOW_RUN_DIR": str(run_dir),
            "LEGISLATION_SHADOW_INPUT_ROOT": str(materialized["input_root"]),
            "PYTHONDONTWRITEBYTECODE": "1",
        }
    )
    for index, step in enumerate(steps):
        name = str(step["name"])
        prepared = prepare_step_prerequisites(
            run_dir,
            step_name=name,
            ledger=step_ledger,
            materialized=materialized,
        )
        command = format_command(step["command"], run_dir, materialized)
        environment = dict(base_environment)
        environment.update(
            format_environment(step["environment"], run_dir, materialized)
        )
        try:
            completed = subprocess.run(
                command,
                # The working directory is part of the isolation boundary.
                # Plans invoking repository code use {workspace_root}.
                cwd=run_dir,
                env=environment,
                text=True,
                encoding="utf-8",
                errors="replace",
                capture_output=True,
                check=False,
                timeout=step.get(
                    "timeout_seconds", DEFAULT_STEP_TIMEOUT_SECONDS
                ),
            )
        except subprocess.TimeoutExpired as exc:
            for stream, text in (("stdout", exc.stdout), ("stderr", exc.stderr)):
                if isinstance(text, bytes):
                    text = text.decode("utf-8", errors="replace")
                (log_dir / f"{index:02d}_{name}.{stream}.txt").write_text(
                    text or "", encoding="utf-8"
                )
            raise ShadowError(f"isolated_consumer_timeout:{name}") from exc
        finally:
            release_step_transients(
                run_dir, step_name=name, prepared=prepared
            )
        stdout_path = log_dir / f"{index:02d}_{name}.stdout.txt"
        stderr_path = log_dir / f"{index:02d}_{name}.stderr.txt"
        stdout_path.write_text(completed.stdout, encoding="utf-8")
        stderr_path.write_text(completed.stderr, encoding="utf-8")
        if completed.returncode != 0:
            raise ShadowError(
                f"isolated_consumer_failed:{name}:exit={completed.returncode}"
            )
        declared = []
        for output in step["outputs"]:
            path = (run_dir / output["path"]).resolve()
            if not _is_relative_to(path, run_dir):
                raise ShadowError(f"consumer output escaped run: {name}/{path}")
            if not path.is_file():
                raise ShadowError(
                    f"isolated_consumer_missing_output:{name}:{output['path']}"
                )
            key = str(output["path"]).replace("\\", "/")
            if key in outputs:
                raise ShadowError(f"duplicate declared output path: {key}")
            item = {
                "step": name,
                "path": key,
                "byte_sha256": sha256_file(path),
                "value_sha256": semantic_hash(path),
                "size": path.stat().st_size,
                "baseline_role": output.get("baseline_role"),
            }
            outputs[key] = item
            declared.append(item)
            attribution_rows.append(
                validate_attribution(output, path, baseline)
            )
        validate_baseline_unchanged(baseline, f"{run_dir.name}:{name}")
        step_receipts.append(
            {
                "name": name,
                "command": command,
                "exit_code": completed.returncode,
                "outputs": declared,
                "stdout": receipt(stdout_path),
                "stderr": receipt(stderr_path),
                "prepared_prerequisites": prepared,
            }
        )
        producer_path = Path(str(step.get("producer") or ""))
        producer_sha256 = (
            sha256_file(producer_path)
            if producer_path.is_file()
            else None
        )
        step_ledger["completed_steps"].append(
            {
                "name": name,
                "producer_sha256": producer_sha256,
                "prepared_prerequisites": prepared,
                "outputs": {
                    item["path"]: {
                        "sha256": item["byte_sha256"],
                        "size": item["size"],
                    }
                    for item in declared
                },
            }
        )
        write_json(step_ledger_path, step_ledger)
    step_ledger["status"] = "COMPLETE"
    write_json(step_ledger_path, step_ledger)
    run_manifest = {
        "schema": "legislation-shadow-isolated-run-v1",
        "production_read_only": True,
        "materialization": materialized["receipt"],
        "generation_id": generation_id,
        "step_ledger": receipt(step_ledger_path),
        "steps": step_receipts,
        "outputs": outputs,
        "attribution_row_count": len(attribution_rows),
    }
    write_json(run_dir / "run_manifest.json", run_manifest)
    return {
        "run_dir": run_dir,
        "materialization": materialized,
        "steps": step_receipts,
        "outputs": outputs,
        "attribution": attribution_rows,
        "manifest": receipt(run_dir / "run_manifest.json"),
    }


def idempotence_audit(
    first: dict[str, Any], second: dict[str, Any]
) -> tuple[list[dict[str, Any]], bool, bool]:
    first_outputs = first["outputs"]
    second_outputs = second["outputs"]
    if set(first_outputs) != set(second_outputs):
        raise ShadowError("cross_run_output_set_mismatch")
    rows = []
    for path in sorted(first_outputs):
        left = first_outputs[path]
        right = second_outputs[path]
        byte_equal = left["byte_sha256"] == right["byte_sha256"]
        value_equal = left["value_sha256"] == right["value_sha256"]
        rows.append(
            {
                "schema": "legislation-shadow-idempotence-v1",
                "path": path,
                "baseline_role": left["baseline_role"],
                "run_1_byte_sha256": left["byte_sha256"],
                "run_2_byte_sha256": right["byte_sha256"],
                "run_1_value_sha256": left["value_sha256"],
                "run_2_value_sha256": right["value_sha256"],
                "byte_idempotent": byte_equal,
                "value_idempotent": value_equal,
            }
        )
    return (
        rows,
        bool(rows) and all(row["value_idempotent"] for row in rows),
        bool(rows) and all(row["byte_idempotent"] for row in rows),
    )


def public_coverage(coverage: dict[str, Any]) -> dict[str, Any]:
    return {
        key: value
        for key, value in coverage.items()
        if not key.startswith("_")
    }


def public_producers(producers: dict[str, Any]) -> dict[str, Any]:
    return {
        label: {
            "manifest": item["manifest_receipt"],
            "schema": item["manifest"].get("schema"),
            "acceptance_passed": explicit_pass(item["manifest"]),
            "outputs": item["outputs"],
        }
        for label, item in producers.items()
    }


def build_gate_provenance(gates: Mapping[str, bool]) -> dict[str, Any]:
    """Attach honest provenance to every acceptance gate.

    ``gates`` keeps its ``{name: bool}`` shape because downstream consumers
    evaluate it with ``all(value is True ...)``.  This adds the missing half:
    whether each True was measured this run or merely restates a precondition
    that already raised.
    """

    unknown = sorted(set(gates) - set(GATE_PROVENANCE))
    if unknown:
        raise ShadowError(
            "gate_provenance_missing:" + ",".join(unknown)
        )
    rows: dict[str, Any] = {}
    for name, value in gates.items():
        provenance, evidence = GATE_PROVENANCE[name]
        rows[name] = {
            "value": value,
            "provenance": provenance,
            "evidence": evidence,
            "independently_measured_this_run": provenance == "measured",
        }
    return {
        "gates": rows,
        "measured_gate_count": sum(
            1 for row in rows.values() if row["provenance"] == "measured"
        ),
        "enforced_by_precondition_gate_count": sum(
            1
            for row in rows.values()
            if row["provenance"] == "enforced_by_precondition"
        ),
    }


def measure_rank_drift(
    args: argparse.Namespace,
    run: Mapping[str, Any],
    baseline: Mapping[str, dict[str, Any]],
    out: Path,
) -> dict[str, Any]:
    """Locate the recomputed ranking and measure it against the baseline."""

    del out
    recomputed = [
        item
        for item in run["outputs"].values()
        if item.get("baseline_role") == RANK_DRIFT_ROLE
    ]
    if len(recomputed) != 1:
        raise ShadowError(
            "rank_drift_output_absent:"
            f"expected exactly one {RANK_DRIFT_ROLE} output, "
            f"found {len(recomputed)}"
        )
    baseline_entry = baseline.get(RANK_DRIFT_ROLE)
    if not baseline_entry:
        raise ShadowError(f"rank_drift_baseline_absent:{RANK_DRIFT_ROLE}")
    max_movers = getattr(args, "max_rank_movers", None)
    max_abs_delta = getattr(args, "max_abs_rank_delta", None)
    explicit = max_movers is not None or max_abs_delta is not None
    return compute_rank_drift(
        Path(baseline_entry["path"]),
        Path(run["run_dir"]) / recomputed[0]["path"],
        max_movers=(
            DEFAULT_MAX_RANK_MOVERS if max_movers is None else int(max_movers)
        ),
        max_abs_delta=(
            DEFAULT_MAX_ABS_RANK_DELTA
            if max_abs_delta is None
            else int(max_abs_delta)
        ),
        threshold_source=(
            "explicit_human_allowance" if explicit else "default_zero_movement"
        ),
    )


def orchestrate(args: argparse.Namespace) -> dict[str, Any]:
    out = validate_out(args.out)
    out.mkdir(parents=True, exist_ok=True)
    audit_path = out / "shadow_recompute_audit.json"
    receipt_path = out / "staged_acceptance_receipt.json"
    try:
        coverage = validate_coverage(args)
        producers = validate_producers(args, coverage)
        baseline = load_baseline(args.baseline_manifest, out)
        plan, steps = load_plan(args.consumer_plan, baseline)
        validate_baseline_unchanged(baseline, "pre_run")

        first = execute_run(
            out / "run_1", steps, coverage, producers, baseline
        )
        second = execute_run(
            out / "run_2", steps, coverage, producers, baseline
        )
        validate_baseline_unchanged(baseline, "post_run")
        idempotence_rows, value_pass, byte_pass = idempotence_audit(
            first, second
        )
        write_ndjson(out / "idempotence_audit.ndjson", idempotence_rows)
        write_ndjson(
            out / "old_vs_new_attribution.ndjson", first["attribution"]
        )
        attribution_pass = bool(first["attribution"]) and all(
            row["attribution_status"] in ACCEPTED_ATTRIBUTION_STATUSES
            for row in first["attribution"]
        )
        unexplained_deltas = sorted(
            str(row.get("baseline_role"))
            for row in first["attribution"]
            if row["attribution_status"] == "attributed_machine_generated"
        )

        drift = measure_rank_drift(args, first, baseline, out)
        write_json(out / RANK_DRIFT_ARTIFACT_NAME, drift)

        gates = {
            "strict_discovery": True,
            "final_merged_coverage": True,
            "membership_review": True,
            "relationships": True,
            "coverage_reconciliation": True,
            "accepted_generation_current": True,
            "direction_current": True,
            "base_network_current": True,
            "rollcall_current": True,
            "production_baseline_immutable": True,
            "all_required_consumers_isolated": True,
            "old_new_delta_attribution": attribution_pass,
            "rank_drift_within_threshold": bool(drift["gate_passed"]),
            "cross_run_value_idempotence": value_pass,
            "cross_run_byte_idempotence": byte_pass,
            "complete_baseline_role_output_closure": (
                set(first["outputs"]) == set(second["outputs"])
                and all(
                    sorted(
                        item["baseline_role"]
                        for item in run["outputs"].values()
                        if item["baseline_role"]
                    )
                    == sorted(REQUIRED_BASELINE_ROLES)
                    for run in (first, second)
                )
            ),
        }
        gate_provenance = build_gate_provenance(gates)
        acceptance_passed = all(gates.values())
        if not acceptance_passed:
            detail = ",".join(
                name for name, passed in gates.items() if not passed
            )
            if unexplained_deltas:
                detail += (
                    ";machine_generated_attribution_is_not_an_explanation:"
                    + ",".join(unexplained_deltas)
                )
            if not gates["rank_drift_within_threshold"]:
                detail += (
                    ";rank_drift:"
                    f"members_moved={drift['members_moved']}"
                    f",max_abs_rank_delta={drift['max_abs_rank_delta']}"
                    f",roster_stable={drift['roster_stable']}"
                )
            raise ShadowError("technical_acceptance_failed:" + detail)
        audit = {
            "schema": AUDIT_SCHEMA,
            "status": "PASS",
            "acceptance_passed": True,
            "production_read_only": True,
            "coverage": public_coverage(coverage),
            "producers": public_producers(producers),
            "baseline_manifest": receipt(args.baseline_manifest.resolve()),
            "consumer_plan": plan["receipt"],
            "gates": gates,
            "gate_provenance": gate_provenance["gates"],
            "measured_gate_count": gate_provenance["measured_gate_count"],
            "enforced_by_precondition_gate_count": gate_provenance[
                "enforced_by_precondition_gate_count"
            ],
            "rank_drift": drift,
            "rank_drift_receipt": receipt(out / RANK_DRIFT_ARTIFACT_NAME),
            "unexplained_delta_roles": unexplained_deltas,
            "run_1_output_count": len(first["outputs"]),
            "run_2_output_count": len(second["outputs"]),
            "run_1_manifest": first["manifest"],
            "run_2_manifest": second["manifest"],
            "idempotence_audit": receipt(out / "idempotence_audit.ndjson"),
            "old_new_attribution": receipt(
                out / "old_vs_new_attribution.ndjson"
            ),
        }
        write_json(audit_path, audit)
        staged_receipt = {
            "schema": RECEIPT_SCHEMA,
            "status": "PASS",
            "acceptance_passed": True,
            "idempotence_passed": value_pass and byte_pass,
            "value_idempotence_passed": value_pass,
            "byte_idempotence_passed": byte_pass,
            "old_new_delta_audit_passed": attribution_pass,
            "old_new_delta_human_explained": not unexplained_deltas,
            # Consumed by the staged integration gate's
            # recomputation.rank_drift_reviewed derivation.
            "rank_drift_reviewed": True,
            "rank_drift_within_threshold": bool(drift["gate_passed"]),
            "rank_drift_members_moved": drift["members_moved"],
            "rank_drift_max_abs_rank_delta": drift["max_abs_rank_delta"],
            "rank_drift_roster_stable": drift["roster_stable"],
            "rank_drift_threshold": drift["threshold"],
            "rank_drift_sha256": sha256_file(out / RANK_DRIFT_ARTIFACT_NAME),
            "measured_gate_count": gate_provenance["measured_gate_count"],
            "enforced_by_precondition_gate_count": gate_provenance[
                "enforced_by_precondition_gate_count"
            ],
            "production_baseline_immutable": True,
            "production_promotion_performed": False,
            "explicit_user_promotion_approval_still_required": True,
            "audit": receipt(audit_path),
            "coverage_manifest_hashes": {
                key: value["manifest"]["sha256"]
                for key, value in public_coverage(coverage).items()
                if isinstance(value, dict) and "manifest" in value
            },
            "producer_manifest_hashes": {
                label: item["manifest_receipt"]["sha256"]
                for label, item in producers.items()
            },
            "run_manifest_hashes": {
                "run_1": first["manifest"]["sha256"],
                "run_2": second["manifest"]["sha256"],
            },
            "baseline_manifest_sha256": sha256_file(
                args.baseline_manifest.resolve()
            ),
            "consumer_plan_sha256": sha256_file(args.consumer_plan.resolve()),
            "idempotence_audit_sha256": sha256_file(
                out / "idempotence_audit.ndjson"
            ),
            "old_new_attribution_sha256": sha256_file(
                out / "old_vs_new_attribution.ndjson"
            ),
        }
        write_json(receipt_path, staged_receipt)
        return staged_receipt
    except Exception as exc:
        # A failure audit is intentionally distinct from the PASS-only receipt.
        if receipt_path.exists():
            receipt_path.unlink()
        failure = {
            "schema": AUDIT_SCHEMA,
            "status": "FAIL",
            "acceptance_passed": False,
            "production_read_only": True,
            "blocking_issues": [str(exc)],
            "staged_acceptance_receipt_emitted": False,
        }
        write_json(audit_path, failure)
        return failure


def main(argv: Sequence[str] | None = None) -> int:
    args = parser().parse_args(argv)
    try:
        result = orchestrate(args)
    except ShadowError as exc:
        print(f"run_legislation_shadow_recompute: FAIL CLOSED: {exc}")
        return 2
    print(json.dumps(result, indent=2, sort_keys=True))
    return 0 if result.get("acceptance_passed") is True else 2


if __name__ == "__main__":
    raise SystemExit(main())
