# -*- coding: utf-8 -*-
"""Apply the canonical per-Congress decay to letter evidence everywhere it is materialized.

The archive scorer enforces the policy in ``bill_rules.evidence_decay``.  This repair also updates
the manual evidence ledger and the letter units already pasted into Legislative Score - Raw, then
re-materializes the affected normalized columns.  The workbook adjustment is delta-based:

    current cell - units at stored CSV weights + units at canonical bill weights

That makes a completed run idempotent while preserving non-letter activity in the same cells.
Backup-first, Excel-lock guarded, verified atomic workbook save.

Usage: python fec/legislation/repair_letter_decay.py [--dry-run]
"""
import collections
import csv
import datetime
import io
import os
import shutil
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import bill_rules as BR
from moneyball_core import save_workbook_atomic, write_text_atomic
from paths import BACKUPS, ROOT, WORKBOOK
import roster

MANUAL = os.path.join(ROOT, "manual_bill_evidence.csv")
RAW_SHEET = "Legislative Score - Raw"
NORM_SHEET = "Legislative Score - Normalized"
FIRST_DATA_ROW, LAST_DATA_ROW = roster.first_row(RAW_SHEET), roster.last_row(RAW_SHEET)

# (category, canonical role) -> Legislative Score - Raw column.
ROLE_COLS = {
    ("CDFI Fund", "sponsor"): 91,
    ("CDFI Fund", "original-cosponsor"): 93,
    ("CFPB", "sponsor"): 99,
    ("CFPB", "original-cosponsor"): 101,
    ("Credit Card Mandates", "sponsor"): 107,
    ("Credit Card Mandates", "original-cosponsor"): 109,
    ("CU Tax Status", "sponsor"): 153,
    ("CU Tax Status", "original-cosponsor"): 155,
}


def _stored_decay(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        # The historical manual-evidence producer treated blank as current/full weight.
        return 1.0


def _format_decay(value):
    number = float(value)
    return "%.1f" % number if number.is_integer() else "%.12g" % number


def _load_manual():
    with open(MANUAL, encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return list(reader), list(reader.fieldnames or [])


def _manual_text(rows, fields):
    out = io.StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=fields, lineterminator="\n")
    writer.writeheader()
    writer.writerows({field: row.get(field, "") for field in fields} for row in rows)
    return "\ufeff" + out.getvalue()


def _guard_not_locked():
    lock = os.path.join(os.path.dirname(WORKBOOK), "~$" + os.path.basename(WORKBOOK))
    if os.path.exists(lock):
        raise RuntimeError("Excel lock file present; close the workbook first: %s" % lock)


def main():
    dry_run = "--dry-run" in sys.argv
    _guard_not_locked()
    rows, fields = _load_manual()
    if "DecayWeight" not in fields:
        raise RuntimeError("manual evidence ledger has no DecayWeight column")

    old_units = collections.defaultdict(float)
    new_units = collections.defaultdict(float)
    letter_rows = changed_rows = 0
    unsupported = []
    for row in rows:
        if not BR.is_letter_evidence(row.get("Bill")):
            continue
        letter_rows += 1
        category = BR.canonical_category(row.get("Category"))
        role = BR.canonical_role(row.get("Role"))
        key = (row.get("Member", "").strip(), category, role)
        if (category, role) not in ROLE_COLS:
            unsupported.append((row.get("Bill"), category, role))
            continue
        old = _stored_decay(row.get("DecayWeight"))
        new = BR.congress_decay(row.get("Congress"))
        old_units[key] += old
        new_units[key] += new
        canonical = _format_decay(new)
        if str(row.get("DecayWeight", "")).strip() != canonical:
            changed_rows += 1
            row["DecayWeight"] = canonical
    if unsupported:
        raise RuntimeError("letter rows use unsupported category/role pairs: %r" % unsupported[:8])

    wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
    raw = wb[RAW_SHEET]
    norm = wb[NORM_SHEET]
    member_rows = {
        str(raw.cell(r, 1).value).strip(): r
        for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1)
        if raw.cell(r, 1).value
    }
    missing = sorted({member for member, _category, _role in old_units if member not in member_rows})
    if missing:
        wb.close()
        raise RuntimeError("letter members missing from workbook: %r" % missing[:8])

    changed_cells = []
    affected_cols = set()
    for key in sorted(old_units):
        member, category, role = key
        column = ROLE_COLS[(category, role)]
        row_number = member_rows[member]
        current = raw.cell(row_number, column).value
        if not isinstance(current, (int, float)):
            wb.close()
            raise RuntimeError("%s!%s is not numeric" % (RAW_SHEET, raw.cell(row_number, column).coordinate))
        updated = float(current) - old_units[key] + new_units[key]
        if updated < -1e-9:
            wb.close()
            raise RuntimeError("decay adjustment would make %s negative" % raw.cell(row_number, column).coordinate)
        updated = 0.0 if abs(updated) < 1e-12 else updated
        if abs(float(current) - updated) > 1e-12:
            changed_cells.append(
                (row_number, column, member, category, role, float(current), updated)
            )
        affected_cols.add(column)

    print("Letter-decay repair")
    print("  letter rows                :", letter_rows)
    print("  ledger rows changing       :", changed_rows)
    print("  workbook raw cells changing:", len(changed_cells))
    print("  affected columns           :", sorted(affected_cols))
    print("  canonical schedule         : 119=1.0, 118=0.8, 117=0.64")
    if changed_cells:
        print("  sample                     :", changed_cells[:5])
    if dry_run:
        wb.close()
        print("--dry-run: no files written.")
        return 0

    os.makedirs(BACKUPS, exist_ok=True)
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    workbook_backup = os.path.join(
        BACKUPS,
        os.path.basename(WORKBOOK).replace(".xlsx", ".PRE-LETTER-DECAY-%s.xlsx" % timestamp),
    )
    ledger_backup = os.path.join(BACKUPS, "manual_bill_evidence.PRE-LETTER-DECAY-%s.csv" % timestamp)
    shutil.copy2(WORKBOOK, workbook_backup)
    shutil.copy2(MANUAL, ledger_backup)
    print("  workbook backup            :", workbook_backup)
    print("  ledger backup              :", ledger_backup)

    for row_number, column, _member, _category, _role, _old, updated in changed_cells:
        raw.cell(row_number, column).value = updated

    # Re-materialize the same-column value/MAX normalization used by the legislative grid.
    for column in affected_cols:
        values = [
            float(raw.cell(row_number, column).value or 0.0)
            for row_number in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1)
        ]
        maximum = max([0.0] + values)
        for offset, row_number in enumerate(range(FIRST_DATA_ROW, LAST_DATA_ROW + 1)):
            norm.cell(row_number, column).value = 0.0 if maximum == 0 else values[offset] / maximum

    wb.calculation.fullCalcOnLoad = True
    save_workbook_atomic(wb, WORKBOOK)
    write_text_atomic(_manual_text(rows, fields), MANUAL, encoding="utf-8")
    print("SAVED: workbook and manual ledger now use bill-equivalent letter decay.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
