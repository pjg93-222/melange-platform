# -*- coding: utf-8 -*-
"""Build a deterministic, audit-only legislation cosponsorship network.

The sole scoring input is the canonical ledger written by
``build_accepted_legislation_generation.py``.  This producer deliberately has
no production default and refuses production-style output names.  It writes a
legacy-compatible two-sheet workbook plus contribution-, exclusion-, and hash
ledgers so every edge can be reconstructed from Congress-qualified measures.

Accepted-generation v1 input contract (NDJSON)
------------------------------------------------
Every nonblank line is one ``accepted-legislation-relationship-v1`` object:

``schema``
    Literal ``accepted-legislation-relationship-v1``.
``measure_key``
    Literal ``{congress}:{measure_type}:{measure_number}``.
``congress``
    Integer 116, 117, 118, or 119.
``measure_type``
    One of the ten canonical lowercase Congress.gov measure types.
``measure_number``
    Positive integer.
``category``
    Nonempty accepted category.
``member_bioguide_id`` / ``member_name``
    Congress member identity.
``role``
    Literal ``sponsor``, ``original_cosponsor``, or ``cosponsor``.
``relationship_status``
    Literal ``current`` or ``withdrawn``.
``membership_status``
    Literal ``included`` or ``excluded``.
``provenance``
    Nonempty object.  It is copied, not interpreted.

Category duplicates are expected: network relationships are deduplicated by
type-aware measure identity and member before edge construction.  Any schema
or identity mismatch is fatal; this producer does not guess field aliases.
"""
from __future__ import annotations

import os as _bootstrap_os, sys as _bootstrap_sys; _bootstrap_sys.path.insert(0, _bootstrap_os.path.dirname(_bootstrap_os.path.dirname(_bootstrap_os.path.abspath(__file__))))

import argparse
import csv
import datetime as dt
import hashlib
import json
import math
import os
import re
import tempfile
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Sequence

import networkx as nx
import openpyxl
from openpyxl import Workbook

try:
    from legislation_measure_key import measure_weight, measure_weight_policy
except ImportError:  # pragma: no cover - package import path
    from fec.legislation_measure_key import measure_weight, measure_weight_policy


HERE = Path(__file__).resolve().parent.parent
ROOT = HERE.parent
AUDIT_ROOT = (ROOT / "audits" / "legfix5").resolve()
REPO_TEMP_ROOT = (ROOT / "tmp").resolve()

ROW_SCHEMA = "accepted-legislation-relationship-v1"
MANIFEST_SCHEMA = "legislation-base-network-manifest-v1"
CONTRIBUTION_SCHEMA = "legislation-network-edge-contribution-v1"
EXCLUSION_SCHEMA = "legislation-network-zero-contribution-v1"

DECAY = {116: 0.16, 117: 0.33, 118: 0.66, 119: 1.0}
MEASURE_TYPES = frozenset(
    {
        "hr",
        "hres",
        "hjres",
        "hconres",
        "hamdt",
        "s",
        "sres",
        "sjres",
        "sconres",
        "samdt",
    }
)
ROLES = frozenset({"sponsor", "original_cosponsor", "cosponsor"})
RELATIONSHIP_STATUSES = frozenset({"current", "withdrawn"})
MEMBERSHIP_STATUSES = frozenset({"included", "excluded"})
BIOGUIDE_RE = re.compile(r"^[A-Z][0-9]{6}$")
PRODUCTION_WORKBOOK_RE = re.compile(
    r"^congress_cosponsor_network(?:_AUG)?_?.*\.xlsx$", re.IGNORECASE
)
FORBIDDEN_NAMES = frozenset(
    {
        "moneyball main version - blended.xlsx",
        "bipartisan.json",
        "behavior_panel.json",
        "behavioral_dispositions.json",
        "ranks_python.json",
    }
)
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
# openpyxl's save_workbook() overwrites properties.modified with the save-time clock,
# so write_workbook's pin cannot survive the save; core.xml is normalized on the way
# through _canonicalize_xlsx instead.
FIXED_DOC_TIME = (
    dt.datetime(*FIXED_ZIP_TIME).strftime("%Y-%m-%dT%H:%M:%SZ").encode("ascii")
)
CORE_PROPERTIES_MEMBER = "docProps/core.xml"
_CORE_MODIFIED = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")

EDGE_HEADERS = (
    "From",
    "To",
    "Weight",
    "NBills",
    "NOriginal",
    "NRegular",
    "From_label",
    "To_label",
)
CENTRALITY_HEADERS = (
    "node",
    "label",
    "InfluenceScore",
    "WeightedOutDegree",
    "WeightedInDegree",
    "OutDegree",
    "InDegree",
    "Closeness",
    "Betweenness",
    "Eigenvector",
    "PageRank",
)


class ContractError(ValueError):
    """The accepted ledger or requested output violates the staged contract."""


@dataclass(frozen=True)
class Relationship:
    row_number: int
    measure_key: str
    congress: int
    measure_type: str
    measure_number: int
    category: str
    member_bioguide_id: str
    member_name: str
    role: str
    relationship_status: str
    membership_status: str
    measure_weight: float
    effective_weight: float
    provenance: dict[str, Any]


@dataclass(frozen=True)
class OutputPaths:
    workbook: Path
    contributions: Path
    exclusions: Path
    manifest: Path
    legacy_comparison: Path


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a sealed legislation base network and Centrality workbook."
    )
    parser.add_argument("--accepted-ledger", required=True, type=Path)
    parser.add_argument(
        "--out",
        required=True,
        type=Path,
        help="Audit/temp output directory, or an explicit non-production .xlsx path.",
    )
    parser.add_argument(
        "--roster",
        type=Path,
        help=(
            "Optional .csv/.ndjson/.json/.xlsx roster. Required fields are "
            "bioguide_id/member_name (CSV/JSON) or node/label (XLSX Centrality)."
        ),
    )
    parser.add_argument(
        "--compare-old-workbook",
        type=Path,
        help="Optional read-only legacy workbook for an edge/Centrality delta audit.",
    )
    return parser.parse_args(argv)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def file_receipt(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "path": str(path.resolve()),
        "size": stat.st_size,
        "sha256": sha256(path),
    }


def _is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def allowed_output_roots() -> tuple[Path, ...]:
    candidates = {AUDIT_ROOT, REPO_TEMP_ROOT, Path(tempfile.gettempdir()).resolve()}
    for env_name in ("TEMP", "TMP", "TMPDIR"):
        value = os.environ.get(env_name)
        if value:
            candidates.add(Path(value).resolve())
    return tuple(sorted(candidates, key=str))


def output_paths(raw_out: Path) -> OutputPaths:
    out = raw_out.resolve()
    roots = allowed_output_roots()
    if not any(_is_relative_to(out, root) for root in roots):
        raise ContractError(
            "--out must resolve under audits/legfix5 or a temp directory; "
            f"received {out}"
        )
    if out.name.lower() in FORBIDDEN_NAMES or PRODUCTION_WORKBOOK_RE.fullmatch(
        out.name
    ):
        raise ContractError(f"production-style output name is prohibited: {out.name}")

    if out.suffix.lower() == ".xlsx":
        prefix = out.with_suffix("")
        paths = OutputPaths(
            workbook=out,
            contributions=Path(str(prefix) + ".edge_contributions.ndjson"),
            exclusions=Path(str(prefix) + ".zero_contribution_exclusions.ndjson"),
            manifest=Path(str(prefix) + ".manifest.json"),
            legacy_comparison=Path(str(prefix) + ".legacy_comparison.json"),
        )
    elif out.suffix:
        raise ContractError("--out must be a directory or have the .xlsx suffix")
    else:
        paths = OutputPaths(
            workbook=out / "legislation_base_network.xlsx",
            contributions=out / "edge_contributions.ndjson",
            exclusions=out / "zero_contribution_exclusions.ndjson",
            manifest=out / "manifest.json",
            legacy_comparison=out / "legacy_comparison.json",
        )
    for path in paths.__dict__.values():
        if path.name.lower() in FORBIDDEN_NAMES or PRODUCTION_WORKBOOK_RE.fullmatch(
            path.name
        ):
            raise ContractError(f"production-style output name is prohibited: {path.name}")
        if not any(_is_relative_to(path.resolve(), root) for root in roots):
            raise ContractError(f"derived output escaped the allowed roots: {path}")
    return paths


def _expect_type(
    row: dict[str, Any], field: str, expected: type, row_number: int
) -> Any:
    value = row.get(field)
    if expected is int and isinstance(value, bool):
        raise ContractError(f"row {row_number}: {field} must be an integer")
    if not isinstance(value, expected):
        raise ContractError(
            f"row {row_number}: {field} must be {expected.__name__}, "
            f"got {type(value).__name__}"
        )
    return value


def parse_relationship(raw: Any, row_number: int) -> Relationship:
    if not isinstance(raw, dict):
        raise ContractError(f"row {row_number}: each NDJSON value must be an object")
    if raw.get("schema") != ROW_SCHEMA:
        raise ContractError(
            f"row {row_number}: schema must be {ROW_SCHEMA!r}, "
            f"got {raw.get('schema')!r}"
        )
    congress = _expect_type(raw, "congress", int, row_number)
    if congress not in DECAY:
        raise ContractError(f"row {row_number}: unsupported Congress {congress}")
    measure_type = _expect_type(raw, "measure_type", str, row_number)
    if measure_type not in MEASURE_TYPES:
        raise ContractError(
            f"row {row_number}: noncanonical measure_type {measure_type!r}"
        )
    measure_number = _expect_type(raw, "measure_number", int, row_number)
    if measure_number <= 0:
        raise ContractError(f"row {row_number}: measure_number must be positive")
    measure_key = _expect_type(raw, "measure_key", str, row_number)
    expected_key = f"{congress}:{measure_type}:{measure_number}"
    if measure_key != expected_key:
        raise ContractError(
            f"row {row_number}: measure_key {measure_key!r} does not equal "
            f"type-aware identity {expected_key!r}"
        )
    category = _expect_type(raw, "category", str, row_number).strip()
    if not category:
        raise ContractError(f"row {row_number}: category must be nonempty")
    bioguide = _expect_type(raw, "member_bioguide_id", str, row_number).strip()
    if not BIOGUIDE_RE.fullmatch(bioguide):
        raise ContractError(
            f"row {row_number}: invalid member_bioguide_id {bioguide!r}"
        )
    member_name = _expect_type(raw, "member_name", str, row_number).strip()
    if not member_name:
        raise ContractError(f"row {row_number}: member_name must be nonempty")
    role = _expect_type(raw, "role", str, row_number)
    if role not in ROLES:
        raise ContractError(f"row {row_number}: invalid role {role!r}")
    relationship_status = _expect_type(
        raw, "relationship_status", str, row_number
    )
    if relationship_status not in RELATIONSHIP_STATUSES:
        raise ContractError(
            f"row {row_number}: invalid relationship_status "
            f"{relationship_status!r}"
        )
    membership_status = _expect_type(raw, "membership_status", str, row_number)
    if membership_status not in MEMBERSHIP_STATUSES:
        raise ContractError(
            f"row {row_number}: invalid membership_status {membership_status!r}"
        )
    expected_measure_weight = measure_weight(measure_key)
    expected_effective_weight = round(
        DECAY[congress] * expected_measure_weight, 12
    )
    declared_measure_weight = raw.get(
        "measure_weight", expected_measure_weight
    )
    declared_effective_weight = raw.get(
        "effective_weight", expected_effective_weight
    )
    if (
        isinstance(declared_measure_weight, bool)
        or not isinstance(declared_measure_weight, (int, float))
        or float(declared_measure_weight) != expected_measure_weight
    ):
        raise ContractError(
            f"row {row_number}: measure_weight does not match the "
            "original-method bill/amendment policy"
        )
    if (
        isinstance(declared_effective_weight, bool)
        or not isinstance(declared_effective_weight, (int, float))
        or float(declared_effective_weight) != expected_effective_weight
    ):
        raise ContractError(
            f"row {row_number}: effective_weight does not equal Congress "
            "decay times measure weight"
        )
    provenance = _expect_type(raw, "provenance", dict, row_number)
    if not provenance:
        raise ContractError(f"row {row_number}: provenance must be nonempty")
    return Relationship(
        row_number=row_number,
        measure_key=measure_key,
        congress=congress,
        measure_type=measure_type,
        measure_number=measure_number,
        category=category,
        member_bioguide_id=bioguide,
        member_name=member_name,
        role=role,
        relationship_status=relationship_status,
        membership_status=membership_status,
        measure_weight=expected_measure_weight,
        effective_weight=expected_effective_weight,
        provenance=provenance,
    )


def load_relationships(path: Path) -> list[Relationship]:
    if path.suffix.lower() not in {".ndjson", ".jsonl"}:
        raise ContractError("--accepted-ledger must be an .ndjson or .jsonl file")
    records: list[Relationship] = []
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                raw = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ContractError(
                    f"row {line_number}: invalid JSON: {exc.msg}"
                ) from exc
            records.append(parse_relationship(raw, line_number))
    if not records:
        raise ContractError("accepted ledger has no relationship records")
    return records


def _add_roster_name(
    roster: dict[str, str], bioguide: Any, name: Any, source: str
) -> None:
    if not isinstance(bioguide, str) or not BIOGUIDE_RE.fullmatch(bioguide.strip()):
        raise ContractError(f"{source}: invalid bioguide_id {bioguide!r}")
    if not isinstance(name, str) or not name.strip():
        raise ContractError(f"{source}: member_name must be nonempty")
    bioguide = bioguide.strip()
    name = name.strip()
    previous = roster.get(bioguide)
    if previous is not None and previous != name:
        raise ContractError(
            f"{source}: conflicting roster names for {bioguide}: "
            f"{previous!r} vs {name!r}"
        )
    roster[bioguide] = name


def load_roster(path: Path | None) -> dict[str, str]:
    if path is None:
        return {}
    suffix = path.suffix.lower()
    roster: dict[str, str] = {}
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames is None or not {
                "bioguide_id",
                "member_name",
            }.issubset(reader.fieldnames):
                raise ContractError(
                    "CSV roster requires bioguide_id and member_name columns"
                )
            for row_number, row in enumerate(reader, start=2):
                _add_roster_name(
                    roster,
                    row.get("bioguide_id"),
                    row.get("member_name"),
                    f"{path}: row {row_number}",
                )
    elif suffix in {".ndjson", ".jsonl"}:
        with path.open("r", encoding="utf-8") as handle:
            for row_number, line in enumerate(handle, start=1):
                if not line.strip():
                    continue
                try:
                    row = json.loads(line)
                except json.JSONDecodeError as exc:
                    raise ContractError(
                        f"{path}: row {row_number}: invalid JSON"
                    ) from exc
                if not isinstance(row, dict):
                    raise ContractError(
                        f"{path}: row {row_number}: roster value must be an object"
                    )
                _add_roster_name(
                    roster,
                    row.get("bioguide_id"),
                    row.get("member_name"),
                    f"{path}: row {row_number}",
                )
    elif suffix == ".json":
        raw = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(raw, list):
            raise ContractError("JSON roster must be a list of member objects")
        for index, row in enumerate(raw, start=1):
            if not isinstance(row, dict):
                raise ContractError(f"{path}: roster item {index} must be an object")
            _add_roster_name(
                roster,
                row.get("bioguide_id"),
                row.get("member_name"),
                f"{path}: item {index}",
            )
    elif suffix == ".xlsx":
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if "Centrality" not in workbook.sheetnames:
            raise ContractError("XLSX roster requires a Centrality sheet")
        sheet = workbook["Centrality"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, ())
        if "node" not in headers or "label" not in headers:
            raise ContractError("XLSX Centrality roster requires node and label columns")
        node_index = headers.index("node")
        label_index = headers.index("label")
        for row_number, row in enumerate(rows, start=2):
            if not row or row[node_index] is None:
                continue
            _add_roster_name(
                roster,
                row[node_index],
                row[label_index],
                f"{path}: Centrality row {row_number}",
            )
        workbook.close()
    else:
        raise ContractError("roster must be .csv, .ndjson/.jsonl, .json, or .xlsx")
    if not roster:
        raise ContractError("provided roster contains no members")
    return roster


def _json_clone(value: Any) -> Any:
    """Validate that copied provenance is JSON-compatible and detach it."""
    try:
        return json.loads(json.dumps(value, sort_keys=True, ensure_ascii=False))
    except (TypeError, ValueError) as exc:
        raise ContractError(f"provenance is not JSON-serializable: {exc}") from exc


def construct_edges(
    relationships: Iterable[Relationship],
    roster: dict[str, str] | None = None,
) -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
    list[dict[str, Any]],
    dict[str, str],
]:
    relationships = list(relationships)
    names = dict(roster or {})
    roster_is_authoritative = bool(roster)
    exclusions: list[dict[str, Any]] = []

    relationship_statuses: dict[tuple[str, str], set[str]] = defaultdict(set)
    for row in relationships:
        relationship_statuses[
            (row.measure_key, row.member_bioguide_id)
        ].add(row.relationship_status)
    conflicting_statuses = {
        key: sorted(statuses)
        for key, statuses in relationship_statuses.items()
        if len(statuses) > 1
    }
    if conflicting_statuses:
        first_key = sorted(conflicting_statuses)[0]
        raise ContractError(
            "category copies disagree on relationship status for "
            f"{first_key[0]}/{first_key[1]}: "
            f"{conflicting_statuses[first_key]!r}"
        )

    active: list[Relationship] = []
    for row in relationships:
        if not roster_is_authoritative:
            prior = names.get(row.member_bioguide_id)
            if prior is not None and prior != row.member_name:
                raise ContractError(
                    f"conflicting accepted-ledger names for "
                    f"{row.member_bioguide_id}: {prior!r} vs {row.member_name!r}"
                )
            names[row.member_bioguide_id] = row.member_name
        elif row.member_bioguide_id not in names:
            names[row.member_bioguide_id] = row.member_name

        reason: str | None = None
        if row.membership_status != "included":
            reason = "membership_excluded"
        elif row.relationship_status != "current":
            reason = "relationship_withdrawn"
        if reason:
            exclusions.append(
                {
                    "schema": EXCLUSION_SCHEMA,
                    "measure_key": row.measure_key,
                    "category": row.category,
                    "member_bioguide_id": row.member_bioguide_id,
                    "member_name": names[row.member_bioguide_id],
                    "role": row.role,
                    "relationship_status": row.relationship_status,
                    "membership_status": row.membership_status,
                    "reason": reason,
                    "network_contribution": 0.0,
                    "source_row_number": row.row_number,
                    "provenance": _json_clone(row.provenance),
                }
            )
        else:
            active.append(row)

    # Collapse accepted category memberships to one official relationship per
    # type-aware measure/member.  Repeated categories cannot multiply edges.
    grouped: dict[tuple[str, str], list[Relationship]] = defaultdict(list)
    for row in active:
        grouped[(row.measure_key, row.member_bioguide_id)].append(row)

    consolidated: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for (measure_key, member), rows in sorted(grouped.items()):
        roles = {row.role for row in rows}
        if "sponsor" in roles and len(roles) > 1:
            raise ContractError(
                f"{measure_key}/{member}: member cannot be sponsor and cosponsor"
            )
        if "original_cosponsor" in roles:
            role = "original_cosponsor"
        elif "cosponsor" in roles:
            role = "cosponsor"
        else:
            role = "sponsor"
        first = rows[0]
        consolidated[measure_key][member] = {
            "measure_key": measure_key,
            "congress": first.congress,
            "measure_type": first.measure_type,
            "measure_number": first.measure_number,
            "member_bioguide_id": member,
            "role": role,
            "measure_weight": first.measure_weight,
            "effective_weight": first.effective_weight,
            "categories": sorted({row.category for row in rows}),
            "source_row_numbers": sorted(row.row_number for row in rows),
            "provenance": [
                {
                    "category": row.category,
                    "source_row_number": row.row_number,
                    "value": _json_clone(row.provenance),
                }
                for row in sorted(rows, key=lambda item: (item.category, item.row_number))
            ],
        }

    contributions: list[dict[str, Any]] = []
    for measure_key, members in sorted(consolidated.items()):
        sponsors = sorted(
            member for member, row in members.items() if row["role"] == "sponsor"
        )
        cosponsors = sorted(
            member
            for member, row in members.items()
            if row["role"] in {"original_cosponsor", "cosponsor"}
        )
        if len(sponsors) != 1:
            reason = (
                "no_current_sponsor"
                if not sponsors
                else "multiple_current_sponsors"
            )
            for member in cosponsors:
                row = members[member]
                exclusions.append(
                    {
                        "schema": EXCLUSION_SCHEMA,
                        "measure_key": measure_key,
                        "category": row["categories"],
                        "member_bioguide_id": member,
                        "member_name": names[member],
                        "role": row["role"],
                        "relationship_status": "current",
                        "membership_status": "included",
                        "reason": reason,
                        "network_contribution": 0.0,
                        "source_row_number": row["source_row_numbers"],
                        "provenance": row["provenance"],
                    }
                )
            raise ContractError(
                f"{measure_key}: expected one current sponsor, "
                f"found {len(sponsors)} ({reason})"
            )
        sponsor = sponsors[0]
        count = len(cosponsors)
        if count == 0:
            continue
        congress = int(members[sponsor]["congress"])
        decay = DECAY[congress]
        factor = float(members[sponsor]["measure_weight"])
        effective = float(members[sponsor]["effective_weight"])
        denominator = math.sqrt(count)
        for member in cosponsors:
            row = members[member]
            role = row["role"]
            role_factor = 1.0 if role == "original_cosponsor" else 0.5
            contribution = role_factor * effective / denominator
            contributions.append(
                {
                    "schema": CONTRIBUTION_SCHEMA,
                    "measure_key": measure_key,
                    "congress": congress,
                    "measure_type": row["measure_type"],
                    "measure_number": row["measure_number"],
                    "sponsor_bioguide_id": sponsor,
                    "sponsor_name": names[sponsor],
                    "cosponsor_bioguide_id": member,
                    "cosponsor_name": names[member],
                    "cosponsor_role": role,
                    "current_cosponsor_count": count,
                    "role_factor": role_factor,
                    "congress_decay": decay,
                    "measure_weight": factor,
                    "effective_weight": effective,
                    "contribution": contribution,
                    "categories": sorted(
                        set(members[sponsor]["categories"]) | set(row["categories"])
                    ),
                    "sponsor_source_row_numbers": members[sponsor][
                        "source_row_numbers"
                    ],
                    "cosponsor_source_row_numbers": row["source_row_numbers"],
                    "sponsor_provenance": members[sponsor]["provenance"],
                    "cosponsor_provenance": row["provenance"],
                }
            )

    edge_buckets: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for contribution in contributions:
        edge_buckets[
            (
                contribution["sponsor_bioguide_id"],
                contribution["cosponsor_bioguide_id"],
            )
        ].append(contribution)
    edges: list[dict[str, Any]] = []
    for (source, target), items in sorted(edge_buckets.items()):
        measures = {item["measure_key"] for item in items}
        if len(measures) != len(items):
            raise AssertionError(f"duplicate measure contribution on {source}->{target}")
        edges.append(
            {
                "From": source,
                "To": target,
                "Weight": math.fsum(item["contribution"] for item in items),
                "NBills": len(measures),
                "NOriginal": sum(
                    item["cosponsor_role"] == "original_cosponsor" for item in items
                ),
                "NRegular": sum(
                    item["cosponsor_role"] == "cosponsor" for item in items
                ),
                "From_label": names[source],
                "To_label": names[target],
            }
        )
    contributions.sort(
        key=lambda row: (
            row["measure_key"],
            row["sponsor_bioguide_id"],
            row["cosponsor_bioguide_id"],
        )
    )
    exclusions.sort(
        key=lambda row: (
            row["measure_key"],
            str(row["member_bioguide_id"]),
            str(row["reason"]),
            str(row["category"]),
        )
    )
    return edges, contributions, exclusions, names


def compute_centrality(
    edges: Iterable[dict[str, Any]], names: dict[str, str]
) -> list[dict[str, Any]]:
    graph = nx.DiGraph()
    for node in sorted(names):
        graph.add_node(node)
    for edge in edges:
        weight = float(edge["Weight"])
        if not math.isfinite(weight) or weight <= 0:
            raise ContractError(f"nonpositive/nonfinite edge weight: {edge!r}")
        graph.add_edge(
            edge["From"],
            edge["To"],
            weight=weight,
            distance=1.0 / weight,
        )

    weighted_out = dict(graph.out_degree(weight="weight"))
    weighted_in = dict(graph.in_degree(weight="weight"))
    out_degree = dict(graph.out_degree())
    in_degree = dict(graph.in_degree())

    closeness: dict[str, float] = {}
    for node in sorted(graph):
        distances = nx.single_source_dijkstra_path_length(
            graph, node, weight="distance"
        )
        total = math.fsum(distances.values())
        closeness[node] = 1.0 / total if total > 0 else 0.0

    if graph.number_of_edges():
        betweenness = nx.betweenness_centrality(
            graph, normalized=False, weight="distance"
        )
        eigenvector = nx.eigenvector_centrality(
            graph, max_iter=1000, tol=1.0e-6, weight="weight"
        )
        max_eigenvector = max(eigenvector.values(), default=0.0)
        if max_eigenvector > 0:
            eigenvector = {
                node: value / max_eigenvector
                for node, value in eigenvector.items()
            }
        pagerank = nx.pagerank(graph, weight="weight")
    else:
        betweenness = {node: 0.0 for node in graph}
        eigenvector = {node: 0.0 for node in graph}
        pagerank = {
            node: (1.0 / graph.number_of_nodes() if graph.number_of_nodes() else 0.0)
            for node in graph
        }

    max_weighted_out = max(weighted_out.values(), default=0.0)
    max_betweenness = max(betweenness.values(), default=0.0)
    max_eigenvector = max(eigenvector.values(), default=0.0)
    raw: dict[str, float] = {}
    for node in graph:
        weighted_component = (
            weighted_out[node] / max_weighted_out if max_weighted_out else 0.0
        )
        between_component = (
            betweenness[node] / max_betweenness if max_betweenness else 0.0
        )
        eigen_component = (
            eigenvector[node] / max_eigenvector if max_eigenvector else 0.0
        )
        raw[node] = (
            0.60 * weighted_component
            + 0.25 * between_component
            + 0.15 * eigen_component
        )
    max_raw = max(raw.values(), default=0.0)
    rows = [
        {
            "node": node,
            "label": names[node],
            "InfluenceScore": (
                round(100.0 * raw[node] / max_raw, 1) if max_raw else 0.0
            ),
            "WeightedOutDegree": float(weighted_out[node]),
            "WeightedInDegree": float(weighted_in[node]),
            "OutDegree": int(out_degree[node]),
            "InDegree": int(in_degree[node]),
            "Closeness": float(closeness[node]),
            "Betweenness": float(betweenness[node]),
            "Eigenvector": float(eigenvector[node]),
            "PageRank": float(pagerank[node]),
        }
        for node in graph
    ]
    rows.sort(key=lambda row: (-row["InfluenceScore"], row["node"]))
    return rows


def write_ndjson(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    text = "".join(
        json.dumps(row, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
        + "\n"
        for row in rows
    )
    path.write_text(text, encoding="utf-8", newline="\n")


def _canonicalize_xlsx(path: Path) -> None:
    """Normalize ZIP member order/timestamps so identical inputs are byte-stable."""
    source = path.with_suffix(path.suffix + ".nondeterministic")
    path.replace(source)
    try:
        with zipfile.ZipFile(source, "r") as reader, zipfile.ZipFile(
            path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as writer:
            for name in sorted(reader.namelist()):
                data = reader.read(name)
                if name == CORE_PROPERTIES_MEMBER:
                    data = _CORE_MODIFIED.sub(
                        lambda match: match.group(1) + FIXED_DOC_TIME + match.group(2),
                        data,
                    )
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                writer.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED)
    finally:
        if source.exists():
            source.unlink()


def write_workbook(
    path: Path, edges: list[dict[str, Any]], centrality: list[dict[str, Any]]
) -> None:
    workbook = Workbook()
    edge_sheet = workbook.active
    edge_sheet.title = "Edges"
    edge_sheet.append(EDGE_HEADERS)
    for edge in edges:
        edge_sheet.append(tuple(edge[field] for field in EDGE_HEADERS))
    centrality_sheet = workbook.create_sheet("Centrality")
    centrality_sheet.append(CENTRALITY_HEADERS)
    for row in centrality:
        centrality_sheet.append(tuple(row[field] for field in CENTRALITY_HEADERS))
    workbook.properties.creator = "Melange staged legislation network producer"
    workbook.properties.lastModifiedBy = workbook.properties.creator
    workbook.properties.created = dt.datetime(1980, 1, 1)
    workbook.properties.modified = dt.datetime(1980, 1, 1)
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    workbook.save(path)
    _canonicalize_xlsx(path)


def _read_workbook_rows(
    path: Path,
) -> tuple[dict[tuple[str, str], dict[str, Any]], dict[str, dict[str, Any]]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Edges" not in workbook.sheetnames or "Centrality" not in workbook.sheetnames:
        raise ContractError(
            f"comparison workbook {path} requires Edges and Centrality sheets"
        )

    def keyed(sheet_name: str, keys: tuple[str, ...]) -> dict[Any, dict[str, Any]]:
        rows = workbook[sheet_name].iter_rows(values_only=True)
        headers = tuple(next(rows, ()))
        required = EDGE_HEADERS if sheet_name == "Edges" else CENTRALITY_HEADERS
        if not set(required).issubset(headers):
            raise ContractError(
                f"comparison workbook {sheet_name} headers are incomplete"
            )
        result: dict[Any, dict[str, Any]] = {}
        for row in rows:
            values = dict(zip(headers, row))
            key_parts = tuple(str(values[field]) for field in keys)
            key: Any = key_parts if len(key_parts) > 1 else key_parts[0]
            if key in result:
                raise ContractError(
                    f"comparison workbook duplicate {sheet_name} key {key!r}"
                )
            result[key] = values
        return result

    edges = keyed("Edges", ("From", "To"))
    centrality = keyed("Centrality", ("node",))
    workbook.close()
    return edges, centrality


def compare_legacy(
    new_edges: list[dict[str, Any]],
    new_centrality: list[dict[str, Any]],
    old_path: Path,
) -> dict[str, Any]:
    old_edges, old_centrality = _read_workbook_rows(old_path)
    edges = {(row["From"], row["To"]): row for row in new_edges}
    centrality = {row["node"]: row for row in new_centrality}
    shared_edges = sorted(set(edges) & set(old_edges))
    shared_nodes = sorted(set(centrality) & set(old_centrality))
    edge_numeric_fields = ("Weight", "NBills", "NOriginal", "NRegular")
    centrality_numeric_fields = CENTRALITY_HEADERS[2:]

    def numeric(value: Any) -> float:
        # Legacy workbooks leave metrics blank for some isolates.  The new
        # producer writes explicit zeros; they are computationally equivalent.
        return float(value or 0.0)

    edge_max_abs_delta = {
        field: max(
            (
                abs(
                    numeric(edges[key][field])
                    - numeric(old_edges[key][field])
                )
                for key in shared_edges
            ),
            default=0.0,
        )
        for field in edge_numeric_fields
    }
    centrality_max_abs_delta = {
        field: max(
            (
                abs(
                    numeric(centrality[key][field])
                    - numeric(old_centrality[key][field])
                )
                for key in shared_nodes
            ),
            default=0.0,
        )
        for field in centrality_numeric_fields
    }
    edge_label_mismatches = [
        list(key)
        for key in shared_edges
        if edges[key]["From_label"] != old_edges[key]["From_label"]
        or edges[key]["To_label"] != old_edges[key]["To_label"]
    ]
    node_label_mismatches = [
        key
        for key in shared_nodes
        if centrality[key]["label"] != old_centrality[key]["label"]
    ]
    exact_keys = (
        edges.keys() == old_edges.keys()
        and centrality.keys() == old_centrality.keys()
    )
    metric_tolerances = {
        "Weight": 1.0e-12,
        "NBills": 0.0,
        "NOriginal": 0.0,
        "NRegular": 0.0,
        "InfluenceScore": 0.0,
        "WeightedOutDegree": 1.0e-12,
        "WeightedInDegree": 1.0e-12,
        "OutDegree": 0.0,
        "InDegree": 0.0,
        "Closeness": 1.0e-12,
        "Betweenness": 1.0e-12,
        "Eigenvector": 1.0e-12,
        "PageRank": 1.0e-12,
    }
    metrics_within_tolerance = all(
        delta <= metric_tolerances[field]
        for field, delta in {
            **edge_max_abs_delta,
            **centrality_max_abs_delta,
        }.items()
    )
    return {
        "schema": "legislation-base-network-legacy-comparison-v1",
        "old_workbook": file_receipt(old_path),
        "new_edge_count": len(edges),
        "old_edge_count": len(old_edges),
        "shared_edge_count": len(shared_edges),
        "added_edges": [list(key) for key in sorted(set(edges) - set(old_edges))],
        "removed_edges": [
            list(key) for key in sorted(set(old_edges) - set(edges))
        ],
        "shared_edge_max_abs_delta": edge_max_abs_delta,
        "shared_edge_label_mismatches": edge_label_mismatches,
        "new_node_count": len(centrality),
        "old_node_count": len(old_centrality),
        "shared_node_count": len(shared_nodes),
        "added_nodes": sorted(set(centrality) - set(old_centrality)),
        "removed_nodes": sorted(set(old_centrality) - set(centrality)),
        "shared_centrality_max_abs_delta": centrality_max_abs_delta,
        "shared_node_label_mismatches": node_label_mismatches,
        "metric_tolerances": metric_tolerances,
        "exact_value_reproduction": bool(
            exact_keys
            and metrics_within_tolerance
            and not edge_label_mismatches
            and not node_label_mismatches
        ),
    }


def run(args: argparse.Namespace) -> dict[str, Any]:
    ledger_path = args.accepted_ledger.resolve()
    if not ledger_path.is_file():
        raise ContractError(f"accepted ledger does not exist: {ledger_path}")
    roster_path = args.roster.resolve() if args.roster else None
    if roster_path is not None and not roster_path.is_file():
        raise ContractError(f"roster does not exist: {roster_path}")
    old_path = (
        args.compare_old_workbook.resolve() if args.compare_old_workbook else None
    )
    if old_path is not None and not old_path.is_file():
        raise ContractError(f"comparison workbook does not exist: {old_path}")

    paths = output_paths(args.out)
    paths.workbook.parent.mkdir(parents=True, exist_ok=True)
    relationships = load_relationships(ledger_path)
    roster = load_roster(roster_path)
    edges, contributions, exclusions, names = construct_edges(
        relationships, roster
    )
    centrality = compute_centrality(edges, names)

    write_ndjson(paths.contributions, contributions)
    write_ndjson(paths.exclusions, exclusions)
    write_workbook(paths.workbook, edges, centrality)

    legacy_comparison: dict[str, Any] | None = None
    if old_path is not None:
        legacy_comparison = compare_legacy(edges, centrality, old_path)
        paths.legacy_comparison.write_text(
            json.dumps(
                legacy_comparison,
                indent=2,
                sort_keys=True,
                ensure_ascii=False,
            )
            + "\n",
            encoding="utf-8",
            newline="\n",
        )

    active_relationships = [
        row
        for row in relationships
        if row.membership_status == "included"
        and row.relationship_status == "current"
    ]
    zero_source_rows = {
        row.row_number
        for row in relationships
        if row.membership_status != "included"
        or row.relationship_status != "current"
    }
    contribution_source_rows = {
        int(row_number)
        for contribution in contributions
        for row_number in (
            contribution["sponsor_source_row_numbers"]
            + contribution["cosponsor_source_row_numbers"]
        )
    }
    invariants = {
        "all_roster_nodes_present": set(roster).issubset(
            {row["node"] for row in centrality}
        ),
        "no_withdrawn_or_excluded_contribution": not (
            zero_source_rows & contribution_source_rows
        ),
        "edge_counts_reconcile": all(
            row["NBills"] == row["NOriginal"] + row["NRegular"]
            for row in edges
        ),
        "positive_finite_edge_weights": all(
            math.isfinite(row["Weight"]) and row["Weight"] > 0
            for row in edges
        ),
        "unique_type_aware_measure_edge_contributions": len(contributions)
        == len(
            {
                (
                    row["measure_key"],
                    row["sponsor_bioguide_id"],
                    row["cosponsor_bioguide_id"],
                )
                for row in contributions
            }
        ),
    }
    manifest = {
        "schema": MANIFEST_SCHEMA,
        "status": "PASS" if all(invariants.values()) else "FAIL",
        "acceptance_passed": all(invariants.values()),
        "production_write_prohibited": True,
        "input_contract": ROW_SCHEMA,
        "inputs": {
            "accepted_ledger": file_receipt(ledger_path),
            "roster": file_receipt(roster_path) if roster_path else None,
            "comparison_workbook": file_receipt(old_path) if old_path else None,
        },
        "policy": {
            "measure_identity": "congress:measure_type:measure_number",
            "congress_decay": {str(key): value for key, value in DECAY.items()},
            "measure_weight": measure_weight_policy(),
            "later_cosponsor": (
                "0.5*congress_decay*measure_weight/"
                "sqrt(current_cosponsor_count)"
            ),
            "original_cosponsor": (
                "1.0*congress_decay*measure_weight/"
                "sqrt(current_cosponsor_count)"
            ),
            "withdrawn_or_excluded_contribution": 0.0,
            "influence_score": (
                "round(100*(.60*W+.25*B+.15*E)/"
                "max(.60*W+.25*B+.15*E),1)"
            ),
            "W": "WeightedOutDegree/max(WeightedOutDegree)",
            "B": "unnormalized inverse-weight Betweenness/max(Betweenness)",
            "E": "weighted directed Eigenvector/max(Eigenvector)",
        },
        "counts": {
            "input_relationship_rows": len(relationships),
            "active_input_relationship_rows": len(active_relationships),
            "zero_contribution_exclusion_rows": len(exclusions),
            "edge_contribution_rows": len(contributions),
            "edge_rows": len(edges),
            "centrality_nodes": len(centrality),
            "roster_nodes": len(roster),
            "original_contribution_rows": sum(
                row["cosponsor_role"] == "original_cosponsor"
                for row in contributions
            ),
            "later_contribution_rows": sum(
                row["cosponsor_role"] == "cosponsor" for row in contributions
            ),
        },
        "invariants": invariants,
        "outputs": {
            "workbook": file_receipt(paths.workbook),
            "edge_contributions": file_receipt(paths.contributions),
            "zero_contribution_exclusions": file_receipt(paths.exclusions),
            "legacy_comparison": (
                file_receipt(paths.legacy_comparison)
                if legacy_comparison is not None
                else None
            ),
        },
        "producer": file_receipt(Path(__file__).resolve()),
    }
    if not manifest["acceptance_passed"]:
        raise AssertionError(
            f"network invariants failed: {manifest['invariants']!r}"
        )
    paths.manifest.write_text(
        json.dumps(manifest, indent=2, sort_keys=True, ensure_ascii=False) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    print(
        "staged base network: "
        f"{len(contributions)} contributions -> {len(edges)} edges -> "
        f"{len(centrality)} nodes"
    )
    print(paths.manifest)
    return manifest


def main(argv: Sequence[str] | None = None) -> int:
    try:
        run(parse_args(argv))
    except ContractError as exc:
        print(f"ERROR: {exc}")
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
