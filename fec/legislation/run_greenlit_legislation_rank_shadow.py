# -*- coding: utf-8 -*-
"""Run the Greenlit legislation integration as an isolated paired rank shadow.

The treatment uses the existing workbook semantics:

* CFPB restraint/reform activity -> ``CFPB Restrain`` (CU:DB);
* CFPB expansion/support activity -> ``CFPB Reach (-)`` (EO:EV);
* bills have measure weight 1.0 and amendments 0.5; and
* direction is applied in ``archive_leg.json`` while workbook activity remains
  a nonnegative evidence magnitude.

Production files are hash-guarded and never opened for writing.  The script
creates a frozen-baseline control and an otherwise-identical treatment, runs
the existing category-stance, leg-stance, and Python rank-authority producers,
then emits complete edge coverage and rank-delta receipts.
"""
from __future__ import annotations

import os as _bootstrap_os, sys as _bootstrap_sys; _bootstrap_sys.path.insert(0, _bootstrap_os.path.dirname(_bootstrap_os.path.dirname(_bootstrap_os.path.abspath(__file__))))

import argparse
from collections import Counter, defaultdict
import csv
import datetime as dt
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import unicodedata
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook
import win32com.client as win32


ROOT = Path(__file__).resolve().parents[2]
AUDIT_ROOT = ROOT / "audits" / "legfix5"
BASELINE_ROOT = AUDIT_ROOT / "production_baseline_full_downstream"
BASELINE_FILES = BASELINE_ROOT / "files"
DEFAULT_INTEGRATION = (
    AUDIT_ROOT
    / "staged_integration"
    / "greenlight-20260728a"
)
DEFAULT_AMENDMENTS = (
    AUDIT_ROOT
    / "opus_adjudication"
    / "amendment_culling_opus5_20260728a"
    / "amendment_partition_3703.ndjson"
)
DEFAULT_OUT = (
    AUDIT_ROOT
    / "shadow_recompute"
    / "greenlight-20260728a"
    / "cfpb-route-corrected-v5"
)
WORKBOOK_NAME = "Moneyball Main Version - Blended.xlsx"
SCHEMA = "legislation-greenlit-paired-rank-shadow-v2"
EVENT_SCHEMA = "legislation-greenlit-shadow-event-v1"
EXPECTED_DECAY = {116: 0.16, 117: 0.33, 118: 0.66, 119: 1.0}
ADVANCED_PHRASES = (
    "agreed to",
    "passed house",
    "passed senate",
    "became law",
)
EXCEL_UP = -4162
EXCEL_AUTOMATIC = -4105

CATEGORY_LABELS = {
    "examination_supervision": "Examination & Supervision",
    "credit_union_capital_requirements": "CU Capital Requirements",
    "ncua_budget_share_insurance": "NCUA Budget & Insurance",
    "interchange": "Interchange",
    "payments": "Payments",
    "technology_emerging_issues": "Technology & Emerging",
    "fraud_prevention": "Fraud Prevention",
    "data_privacy_cybersecurity": "Data Privacy & Cyber",
    "charter_modernization_governance": "Charter Modernization",
    "defense_credit_union_issues": "Defense CU Issues",
    "field_of_membership": "Field of Membership",
    "cdfi_fund": "CDFI Fund",
    "cfpb_issues_structural_reform": "CFPB",
    "credit_card_consumer_fee_mandates": "Credit Card Mandates",
    "telephone_consumer_protection_act_issues": "TCPA",
    "member_business_lending": "Member Business Lending",
    "mortgage_lending_housing": "Mortgage & Housing",
    "credit_union_tax_status": "CU Tax Status",
}

SPONSOR_COLUMNS = {
    "Examination & Supervision": 3,  # C
    "CU Capital Requirements": 11,  # K
    "NCUA Budget & Insurance": 19,  # S
    "Interchange": 27,  # AA
    "Payments": 35,  # AI
    "Technology & Emerging": 43,  # AQ
    "Fraud Prevention": 51,  # AY
    "Data Privacy & Cyber": 59,  # BG
    "Charter Modernization": 67,  # BO
    "Defense CU Issues": 75,  # BW
    "Field of Membership": 83,  # CE
    "CDFI Fund": 91,  # CM
    "CFPB Restrain": 99,  # CU
    "Credit Card Mandates": 107,  # DC
    "TCPA": 115,  # DK
    "Member Business Lending": 123,  # DS
    "Mortgage & Housing": 131,  # EA
    "CFPB Reach (-)": 145,  # EO
    "CU Tax Status": 153,  # EW
}

PRODUCTION_GUARDS = (
    WORKBOOK_NAME,
    "archive_leg.json",
    "category_stance.json",
    "ranks_python.json",
    "manual_bill_evidence.csv",
    "data/legislation/bill_directions.csv",
)

BASELINE_INPUTS = (
    WORKBOOK_NAME,
    "archive_leg.json",
    "vote_category_lean.json",
    "data/engine/pillar_inputs_snapshot.json",
    "ranks_python.json",
    "say_do_gap.json",
    "legislation_routing_audit.json",
)

LIVE_SUPPORT_INPUTS = (
    "bill_repository_live.json",
    "data/engine/congestion.json",
    "data/members/formal_visibility_roles.json",
    "data/legislation/bill_overrides.csv",
    "manual_bill_evidence.csv",
)


PRODUCER_TIMEOUT_SECONDS = 3600


class ShadowError(RuntimeError):
    """A paired-shadow input, execution, or safety invariant failed."""


def canonical_bytes(value: Any) -> bytes:
    return json.dumps(
        value,
        sort_keys=True,
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def receipt(path: Path, records: int | None = None) -> dict[str, Any]:
    result = {
        "path": str(path.resolve()),
        "bytes": path.stat().st_size,
        "sha256": sha256_file(path),
    }
    if records is not None:
        result["records"] = records
    return result


def atomic_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    temporary.write_bytes(
        json.dumps(
            value,
            indent=2,
            ensure_ascii=False,
            sort_keys=True,
        ).encode("utf-8")
        + b"\n"
    )
    os.replace(temporary, path)


def atomic_ndjson(path: Path, rows: Iterable[Mapping[str, Any]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    count = 0
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(
                json.dumps(
                    row,
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                + "\n"
            )
            count += 1
    os.replace(temporary, path)
    return count


def read_object(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(value, dict):
        raise ShadowError(f"JSON root is not an object: {path}")
    return value


def read_ndjson(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open(encoding="utf-8-sig") as handle:
        for number, line in enumerate(handle, start=1):
            if not line.strip():
                raise ShadowError(f"{path}:{number}: blank NDJSON line")
            value = json.loads(line)
            if not isinstance(value, dict):
                raise ShadowError(f"{path}:{number}: row is not an object")
            rows.append(value)
    return rows


def normalize_name(value: str) -> str:
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(
        character
        for character in value
        if not unicodedata.combining(character)
    )
    value = value.casefold().replace("’", "'")
    value = re.sub(r"\b(jr|sr|ii|iii|iv)\b", " ", value)
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def parsed_sponsor_name(label: str) -> str:
    visible = str(label or "").split("[", 1)[0].strip()
    if "," not in visible:
        return visible
    last, first = visible.split(",", 1)
    return f"{first.strip()} {last.strip()}"


def display_measure(canonical_key: str) -> str:
    congress, kind, number = canonical_key.split(":")
    prefixes = {
        "hamdt": "H.Amdt.",
        "samdt": "S.Amdt.",
        "hr": "H.R.",
        "s": "S.",
        "hjres": "H.J.Res.",
        "sjres": "S.J.Res.",
        "hconres": "H.Con.Res.",
        "sconres": "S.Con.Res.",
        "hres": "H.Res.",
        "sres": "S.Res.",
    }
    if kind not in prefixes:
        raise ShadowError(f"unsupported canonical measure type: {kind}")
    return f"{prefixes[kind]}{int(number)}"


def status_from_latest_action(row: Mapping[str, Any]) -> str:
    action = str(row.get("latest_action") or "").casefold()
    if not any(phrase in action for phrase in ADVANCED_PHRASES):
        return "Introduced"
    kind = str(row.get("measure_type") or "")
    return "Passed House" if kind == "hamdt" else "Agreed to in Senate"


def production_hashes() -> dict[str, str | None]:
    result: dict[str, str | None] = {}
    for name in PRODUCTION_GUARDS:
        path = ROOT / name
        result[name] = sha256_file(path) if path.is_file() else None
    return result


def ensure_safe_output(path: Path) -> Path:
    resolved = path.resolve()
    allowed = (AUDIT_ROOT / "shadow_recompute").resolve()
    try:
        safe = resolved.is_relative_to(allowed)
    except AttributeError:  # pragma: no cover
        safe = str(resolved).casefold().startswith(
            str(allowed).casefold() + os.sep
        )
    if not safe or resolved == allowed:
        raise ShadowError(
            "output must be a child of audits/legfix5/shadow_recompute"
        )
    if resolved.exists():
        raise ShadowError(f"output already exists; choose a new run path: {resolved}")
    return resolved


def workbook_roster(path: Path) -> tuple[list[str], dict[str, int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rank = workbook["Rank & Allocation"]
        raw = workbook["Legislative Score - Raw"]
        members = [
            value[0]
            for value in rank.iter_rows(
                min_row=3,
                max_row=535,
                min_col=1,
                max_col=1,
                values_only=True,
            )
            if value[0]
        ]
        raw_rows = {
            str(value[0]): row_number
            for row_number, value in enumerate(
                raw.iter_rows(
                    min_row=3,
                    max_row=535,
                    min_col=1,
                    max_col=1,
                    values_only=True,
                ),
                start=3,
            )
            if value[0]
        }
    finally:
        workbook.close()
    if len(members) != 533 or set(members) != set(raw_rows):
        raise ShadowError("frozen workbook roster is not the expected 533-member set")
    return members, raw_rows


def map_sponsor_to_roster(
    sponsor_label: str,
    members: Sequence[str],
) -> tuple[str | None, str]:
    parsed = parsed_sponsor_name(sponsor_label)
    by_normalized = {normalize_name(member): member for member in members}
    direct = by_normalized.get(normalize_name(parsed))
    if direct:
        return direct, parsed
    parts = normalize_name(parsed).split()
    candidates = [
        member
        for member in members
        if parts
        and normalize_name(member).split()[:1] == parts[:1]
        and normalize_name(member).split()[-1:] == parts[-1:]
    ]
    if len(candidates) == 1:
        return candidates[0], parsed
    return None, parsed


def build_events(
    integration_root: Path,
    amendments_path: Path,
    members: Sequence[str],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    edges = read_ndjson(integration_root / "integrated_membership_edges.ndjson")
    decisions = read_ndjson(integration_root / "integrated_decisions.ndjson")
    scoring = [row for row in edges if row.get("scoring_eligible") is True]
    membership_only = [
        row for row in edges if row.get("scoring_eligible") is not True
    ]
    amendment_rows = read_ndjson(amendments_path)
    amendments = {
        str(row.get("canonical_key")): row
        for row in amendment_rows
    }
    reasons = {
        str(row.get("canonical_key")): str(row.get("reason") or "")
        for row in decisions
        if row.get("integration_state") == "included_scoring_ready"
    }
    if len(scoring) != 74 or len(membership_only) != 19:
        raise ShadowError("Greenlit 74-scoring/19-membership-only split drifted")
    keys = {str(row["canonical_key"]) for row in scoring}
    if len(keys) != 60 or any(
        key.split(":")[1] not in {"hamdt", "samdt"} for key in keys
    ):
        raise ShadowError("expected 60 scoring amendments")
    missing = sorted(keys - set(amendments))
    if missing:
        raise ShadowError(f"amendment evidence missing: {missing[:3]}")

    events: list[dict[str, Any]] = []
    for edge in sorted(
        scoring,
        key=lambda row: (
            str(row["canonical_key"]),
            str(row["category_id"]),
        ),
    ):
        canonical_key = str(edge["canonical_key"])
        congress = int(canonical_key.split(":")[0])
        source = amendments[canonical_key]
        member, parsed = map_sponsor_to_roster(
            str(source.get("sponsor") or ""),
            members,
        )
        category_id = str(edge["category_id"])
        if category_id not in CATEGORY_LABELS:
            raise ShadowError(f"unknown integrated category: {category_id}")
        category = CATEGORY_LABELS[category_id]
        direction = int(edge["ultimate_direction"])
        if direction not in {-1, 1}:
            raise ShadowError(f"scoring edge direction is not signed: {edge}")
        measure_weight = float(edge["measure_weight"])
        if not math.isclose(measure_weight, 0.5, abs_tol=1e-12):
            raise ShadowError("a scoring amendment does not have weight 0.5")
        activity = round(EXPECTED_DECAY[congress] * measure_weight, 12)
        if category == "CFPB":
            physical_block = (
                "CFPB Restrain" if direction == 1 else "CFPB Reach (-)"
            )
        else:
            physical_block = category
        events.append(
            {
                "schema": EVENT_SCHEMA,
                "canonical_key": canonical_key,
                "display_measure": display_measure(canonical_key),
                "congress": congress,
                "measure_type": canonical_key.split(":")[1],
                "category_id": category_id,
                "category": category,
                "direction": direction,
                "direction_meaning": (
                    "industry_aligned" if direction == 1 else "industry_adverse"
                ),
                "measure_weight": measure_weight,
                "congress_decay": EXPECTED_DECAY[congress],
                "effective_activity": activity,
                "archive_role_weight": 3.0,
                "archive_signed_delta": round(3.0 * activity * direction, 12),
                "raw_activity_is_nonnegative": True,
                "physical_block": physical_block,
                "raw_sponsor_column": SPONSOR_COLUMNS[physical_block],
                "sponsor_source": str(source.get("sponsor") or ""),
                "parsed_sponsor": parsed,
                "member": member,
                "roster_status": (
                    "mapped_current_roster"
                    if member is not None
                    else "historical_off_current_roster"
                ),
                "latest_action_source": source.get("latest_action"),
                "tracker_status": status_from_latest_action(source),
                "amends": source.get("amends"),
                "reason": reasons.get(canonical_key) or str(source.get("reason") or ""),
                "source_decision_id": edge.get("source_decision_id"),
                "integration_state": edge.get("integration_state"),
            }
        )
    return events, {
        "amendment_partition_records": len(amendment_rows),
        "integrated_category_edges": len(edges),
        "scoring_ready_category_edges": len(scoring),
        "membership_only_non_scoring_category_edges": len(membership_only),
        "scoring_ready_candidates": len(keys),
        "mapped_current_roster_edges": sum(
            row["member"] is not None for row in events
        ),
        "historical_off_roster_edges": sum(
            row["member"] is None for row in events
        ),
        "historical_off_roster_candidates": sorted(
            {row["canonical_key"] for row in events if row["member"] is None}
        ),
    }


def copy_root(target: Path) -> None:
    target.mkdir(parents=True, exist_ok=False)
    for name in BASELINE_INPUTS:
        source = BASELINE_FILES / name
        if not source.is_file():
            raise ShadowError(f"frozen baseline input is missing: {name}")
        destination = target / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, destination)
        os.chmod(destination, 0o666)
    for name in LIVE_SUPPORT_INPUTS:
        source = ROOT / name
        if not source.is_file():
            raise ShadowError(f"live support input is missing: {name}")
        destination = target / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, destination)
        os.chmod(destination, 0o666)


def replace_amendment_manual_rows(
    path: Path,
    events: Sequence[Mapping[str, Any]],
) -> dict[str, int]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = list(reader.fieldnames or [])
        existing = list(reader)
    required = {
        "Member",
        "Role",
        "Bill",
        "Title",
        "Sponsor",
        "Congress",
        "Category",
        "Direction",
        "DecayWeight",
        "LastAction",
        "DetailLink",
        "Summary",
        "Note",
    }
    if set(fieldnames) != required:
        raise ShadowError("manual_bill_evidence.csv schema drifted")

    def is_amendment(row: Mapping[str, str]) -> bool:
        return bool(
            re.search(r"(?i)\b[HS]\.?\s*Amdt\.?", row.get("Bill") or "")
            or "amendment evidence" in (row.get("Note") or "").casefold()
        )

    retained = [row for row in existing if not is_amendment(row)]
    removed = len(existing) - len(retained)
    rows = list(retained)
    for event in events:
        member = str(event["member"] or event["parsed_sponsor"])
        canonical_key = str(event["canonical_key"])
        congress = int(event["congress"])
        detail = (
            f"https://www.congress.gov/amendment/{congress}th-congress/"
            f"{'house' if event['measure_type'] == 'hamdt' else 'senate'}-"
            f"amendment/{int(canonical_key.split(':')[2])}"
        )
        rows.append(
            {
                "Member": member,
                "Role": "sponsor",
                "Bill": str(event["display_measure"]),
                "Title": f"Greenlit amendment to {event.get('amends') or 'measure'}",
                "Sponsor": str(event["sponsor_source"]),
                "Congress": str(congress),
                "Category": str(event["category"]),
                "Direction": f"{int(event['direction']):+d}",
                "DecayWeight": f"{float(event['effective_activity']):.12g}",
                "LastAction": str(event["tracker_status"]),
                "DetailLink": detail,
                "Summary": str(event["reason"]),
                "Note": (
                    "Greenlit culling integration; canonical_key "
                    f"{canonical_key}; amendment weight 0.5; existing "
                    f"{event['physical_block']} route; staged only"
                ),
            }
        )
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    with temporary.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    os.replace(temporary, path)
    return {
        "source_rows": len(existing),
        "removed_pre_adjudication_amendment_rows": removed,
        "retained_non_amendment_rows": len(retained),
        "greenlit_amendment_edge_rows": len(events),
        "integrated_rows": len(rows),
    }


def patch_archive(
    path: Path,
    events: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    archive = read_object(path)
    applied = 0
    deltas: defaultdict[tuple[str, str], float] = defaultdict(float)
    for event in events:
        member = event.get("member")
        if member is None:
            continue
        key = (str(member), str(event["category"]))
        deltas[key] += float(event["archive_signed_delta"])
        applied += 1
    for (member, category), delta in sorted(deltas.items()):
        archive.setdefault(member, {})
        archive[member][category] = (
            float(archive[member].get(category, 0.0)) + delta
        )
    atomic_json(path, archive)
    return {
        "applied_edges": applied,
        "member_category_cells_changed": len(deltas),
        "signed_delta_sum": round(sum(deltas.values()), 12),
    }


def append_sponsor_sheet_rows(
    workbook: Any,
    events: Sequence[Mapping[str, Any]],
) -> int:
    sheet = workbook.Worksheets("Sponsor Interactions")
    last_row = sheet.Cells(sheet.Rows.Count, 1).End(EXCEL_UP).Row
    rows = []
    for event in events:
        member = str(event["member"] or event["parsed_sponsor"])
        rows.append(
            (
                str(event["category"]),
                "sponsor",
                member,
                str(event["display_measure"]),
                f"Greenlit amendment to {event.get('amends') or 'measure'}",
                str(event["sponsor_source"]),
                0,
                "",
                str(event["tracker_status"]),
                int(event["congress"]),
                float(event["effective_activity"]),
                "",
                str(event["reason"]),
            )
        )
    if rows:
        target = sheet.Range(
            sheet.Cells(last_row + 1, 1),
            sheet.Cells(last_row + len(rows), 13),
        )
        target.Value = tuple(rows)
    return len(rows)


def mutate_treatment_workbook(
    path: Path,
    events: Sequence[Mapping[str, Any]],
    raw_rows: Mapping[str, int],
) -> dict[str, Any]:
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.AskToUpdateLinks = False
    workbook = None
    touched_columns: set[int] = set()
    applied = 0
    try:
        workbook = excel.Workbooks.Open(
            str(path.resolve()),
            UpdateLinks=0,
            ReadOnly=False,
        )
        raw = workbook.Worksheets("Legislative Score - Raw")
        normalized = workbook.Worksheets("Legislative Score - Normalized")
        for event in events:
            member = event.get("member")
            if member is None:
                continue
            row_number = raw_rows[str(member)]
            column = int(event["raw_sponsor_column"])
            current = raw.Cells(row_number, column).Value
            raw.Cells(row_number, column).Value = float(current or 0.0) + float(
                event["effective_activity"]
            )
            touched_columns.add(column)
            applied += 1

        normalization: dict[str, float] = {}
        for column in sorted(touched_columns):
            values = raw.Range(
                raw.Cells(3, column),
                raw.Cells(535, column),
            ).Value
            vector = [float(row[0] or 0.0) for row in values]
            maximum = max([value for value in vector if value > 0.0], default=0.0)
            normalized_values = tuple(
                (0.0 if maximum == 0.0 else value / maximum,)
                for value in vector
            )
            normalized.Range(
                normalized.Cells(3, column),
                normalized.Cells(535, column),
            ).Value = normalized_values
            normalization[str(column)] = maximum

        sponsor_rows = append_sponsor_sheet_rows(workbook, events)
        excel.CalculateFullRebuild()
        workbook.Save()
        workbook.Close(SaveChanges=True)
        workbook = None
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        excel.Quit()
    return {
        "raw_activity_edges_applied": applied,
        "historical_off_roster_edges_not_applied_to_grid": (
            len(events) - applied
        ),
        "touched_sponsor_columns": sorted(touched_columns),
        "normalization_positive_maxima": normalization,
        "sponsor_interaction_rows_appended": sponsor_rows,
        "cfpb_route": {
            "industry_aligned": "CFPB Restrain!CU",
            "industry_adverse": "CFPB Reach (-)!EO",
        },
    }


def run_producer(
    script_name: str,
    shadow_root: Path,
    logs: Path,
    *arguments: str,
) -> dict[str, Any]:
    script = ROOT / "fec" / script_name
    if not script.is_file():
        raise ShadowError(f"producer is missing: {script_name}")
    environment = dict(os.environ)
    environment["MONEYBALL_ROOT"] = str(shadow_root.resolve())
    environment["PILLAR_SNAPSHOT_FREEZE"] = "0"
    log_path = logs / f"{shadow_root.name}__{script.stem}.log"
    try:
        completed = subprocess.run(
            [sys.executable, str(script), *arguments],
            cwd=str(ROOT),
            env=environment,
            text=True,
            encoding="utf-8",
            errors="replace",
            capture_output=True,
            check=False,
            timeout=PRODUCER_TIMEOUT_SECONDS,
        )
    except subprocess.TimeoutExpired as exc:
        def _partial(value: Any) -> str:
            if value is None:
                return ""
            if isinstance(value, bytes):
                return value.decode("utf-8", errors="replace")
            return str(value)

        log_path.write_text(
            _partial(exc.stdout)
            + "\n[stderr]\n"
            + _partial(exc.stderr)
            + f"\n[timeout after {PRODUCER_TIMEOUT_SECONDS}s]\n",
            encoding="utf-8",
        )
        raise ShadowError(
            f"{script_name} timed out after {PRODUCER_TIMEOUT_SECONDS}s for "
            f"{shadow_root.name}; see {log_path}"
        ) from exc
    log_path.write_text(
        completed.stdout
        + ("\n[stderr]\n" + completed.stderr if completed.stderr else ""),
        encoding="utf-8",
    )
    if completed.returncode != 0:
        raise ShadowError(
            f"{script_name} failed for {shadow_root.name}; see {log_path}"
        )
    return {
        "script": receipt(script),
        "argv": [str(script), *arguments],
        "returncode": completed.returncode,
        "log": receipt(log_path),
    }


def execute_model(shadow_root: Path, logs: Path) -> list[dict[str, Any]]:
    return [
        run_producer("stance/per_category_stance.py", shadow_root, logs),
        run_producer("legislation/apply_leg_stance.py", shadow_root, logs),
        run_producer("engine/compute_pillars.py", shadow_root, logs, "--write"),
    ]


def rank_delta(
    baseline_path: Path,
    control_path: Path,
    treatment_path: Path,
) -> dict[str, Any]:
    baseline = read_object(baseline_path)
    control = read_object(control_path)
    treatment = read_object(treatment_path)
    names = sorted(baseline["rank"])
    if not (
        set(names) == set(control["rank"]) == set(treatment["rank"])
        and len(names) == 533
    ):
        raise ShadowError("baseline/control/treatment rank rosters differ")
    rows = []
    for member in names:
        control_rank = int(control["rank"][member])
        treatment_rank = int(treatment["rank"][member])
        control_score = float(control["score"][member])
        treatment_score = float(treatment["score"][member])
        rows.append(
            {
                "member": member,
                "baseline_rank": int(baseline["rank"][member]),
                "control_rank": control_rank,
                "treatment_rank": treatment_rank,
                "rank_delta_treatment_minus_control": (
                    treatment_rank - control_rank
                ),
                "rank_improvement": control_rank - treatment_rank,
                "baseline_score": float(baseline["score"][member]),
                "control_score": control_score,
                "treatment_score": treatment_score,
                "score_delta_treatment_minus_control": (
                    treatment_score - control_score
                ),
                "leg_percentile_control": float(
                    control["pcts"][member]["leg"]
                ),
                "leg_percentile_treatment": float(
                    treatment["pcts"][member]["leg"]
                ),
            }
        )
    moved = [row for row in rows if row["rank_delta_treatment_minus_control"]]
    absolute = [
        abs(int(row["rank_delta_treatment_minus_control"])) for row in moved
    ]
    replay_rank_mismatches = [
        member
        for member in names
        if int(control["rank"][member]) != int(baseline["rank"][member])
    ]
    replay_score_max = max(
        abs(
            float(control["score"][member])
            - float(baseline["score"][member])
        )
        for member in names
    )
    best = sorted(
        rows,
        key=lambda row: (
            -int(row["rank_improvement"]),
            row["member"],
        ),
    )[:20]
    worst = sorted(
        rows,
        key=lambda row: (
            int(row["rank_improvement"]),
            row["member"],
        ),
    )[:20]
    return {
        "schema": "legislation-greenlit-rank-delta-v1",
        "delta_convention": (
            "rank_delta=treatment-control; negative is an improvement"
        ),
        "control_replay": {
            "rank_mismatch_count_vs_frozen": len(replay_rank_mismatches),
            "rank_mismatch_examples": replay_rank_mismatches[:20],
            "max_abs_score_delta_vs_frozen": replay_score_max,
            "exact_rank_replay": not replay_rank_mismatches,
            "exact_score_replay": replay_score_max <= 1e-12,
        },
        "summary": {
            "member_count": len(rows),
            "members_with_rank_change": len(moved),
            "members_unchanged": len(rows) - len(moved),
            "max_abs_rank_delta": max(absolute, default=0),
            "mean_abs_rank_delta_among_moved": (
                sum(absolute) / len(absolute) if absolute else 0.0
            ),
            "members_with_score_change": sum(
                abs(float(row["score_delta_treatment_minus_control"])) > 1e-15
                for row in rows
            ),
            "max_abs_score_delta": max(
                abs(float(row["score_delta_treatment_minus_control"]))
                for row in rows
            ),
        },
        "largest_improvements": best,
        "largest_declines": worst,
        "members": rows,
    }


def workbook_verification(path: Path) -> dict[str, Any]:
    values = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True, data_only=False)
    try:
        rank_values = values["Rank & Allocation"]
        raw_values = values["Legislative Score - Raw"]
        raw_formulas = formulas["Legislative Score - Raw"]
        member_count = sum(
            row[0] is not None
            for row in rank_values.iter_rows(
                min_row=3,
                max_row=535,
                min_col=1,
                max_col=1,
                values_only=True,
            )
        )
        errors = 0
        for sheet_name, cell_range in (
            ("Blended Model", (3, 535, 1, 20)),
            ("Rank & Allocation", (3, 535, 1, 7)),
        ):
            sheet = values[sheet_name]
            min_row, max_row, min_col, max_col = cell_range
            for row in sheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
                values_only=True,
            ):
                errors += sum(
                    isinstance(value, str) and value.startswith("#")
                    for value in row
                )
        cfpb = {}
        for column in ("CU", "EO"):
            vector = [
                float(row[0] or 0.0)
                for row in raw_values.iter_rows(
                    min_row=3,
                    max_row=535,
                    min_col=SPONSOR_COLUMNS[
                        "CFPB Restrain"
                        if column == "CU"
                        else "CFPB Reach (-)"
                    ],
                    max_col=SPONSOR_COLUMNS[
                        "CFPB Restrain"
                        if column == "CU"
                        else "CFPB Reach (-)"
                    ],
                    values_only=True,
                )
            ]
            cfpb[column] = {
                "header": raw_formulas[f"{column}1"].value,
                "nonzero_rows": sum(value != 0.0 for value in vector),
                "sum": sum(vector),
                "max": max(vector),
            }
    finally:
        values.close()
        formulas.close()
    return {
        "member_count": member_count,
        "key_formula_error_count": errors,
        "cfpb_physical_blocks": cfpb,
    }


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description=__doc__)
    result.add_argument(
        "--integration-root",
        type=Path,
        default=DEFAULT_INTEGRATION,
    )
    result.add_argument(
        "--amendments",
        type=Path,
        default=DEFAULT_AMENDMENTS,
    )
    result.add_argument("--out", type=Path, default=DEFAULT_OUT)
    return result


def main(argv: Sequence[str] | None = None) -> int:
    args = parser().parse_args(argv)
    started = dt.datetime.now(dt.timezone.utc)
    out = ensure_safe_output(args.out)
    before = production_hashes()
    out.mkdir(parents=True, exist_ok=False)
    logs = out / "logs"
    logs.mkdir()
    control_root = out / "control"
    treatment_root = out / "treatment"
    try:
        integration_root = args.integration_root.resolve()
        integration_receipt = read_object(
            integration_root / "integration_receipt.json"
        )
        if integration_receipt.get("status") != (
            "PASS_STAGED_INTEGRATION_WITH_HELD_NONSCORING"
        ):
            raise ShadowError("integration receipt is not the Greenlit staged PASS")
        members, raw_rows = workbook_roster(BASELINE_FILES / WORKBOOK_NAME)
        events, event_counts = build_events(
            integration_root,
            args.amendments.resolve(),
            members,
        )
        if (
            args.amendments.resolve() == DEFAULT_AMENDMENTS.resolve()
            and event_counts["amendment_partition_records"] != 3703
        ):
            raise ShadowError("default amendment partition is not 3703 records")
        event_path = out / "shadow_event_ledger.ndjson"
        event_count = atomic_ndjson(event_path, events)
        if event_count != 74:
            raise ShadowError("shadow event ledger is not the full 74-edge set")

        copy_root(control_root)
        copy_root(treatment_root)
        control_manual = replace_amendment_manual_rows(
            control_root / "manual_bill_evidence.csv",
            [],
        )
        treatment_manual = replace_amendment_manual_rows(
            treatment_root / "manual_bill_evidence.csv",
            events,
        )
        if (
            control_manual["removed_pre_adjudication_amendment_rows"] != 160
            or treatment_manual["removed_pre_adjudication_amendment_rows"] != 160
        ):
            raise ShadowError(
                "expected to cull exactly 160 pre-adjudication amendment rows"
            )
        archive_audit = patch_archive(
            treatment_root / "archive_leg.json",
            events,
        )
        workbook_audit = mutate_treatment_workbook(
            treatment_root / WORKBOOK_NAME,
            events,
            raw_rows,
        )
        control_steps = execute_model(control_root, logs)
        treatment_steps = execute_model(treatment_root, logs)

        delta = rank_delta(
            BASELINE_FILES / "ranks_python.json",
            control_root / "ranks_python.json",
            treatment_root / "ranks_python.json",
        )
        delta_path = out / "rank_delta.json"
        atomic_json(delta_path, delta)
        workbook_check = workbook_verification(
            treatment_root / WORKBOOK_NAME
        )
        if not (
            workbook_check["member_count"] == 533
            and workbook_check["key_formula_error_count"] == 0
        ):
            raise ShadowError("treatment workbook key verification failed")

        after = production_hashes()
        if after != before:
            changed = [
                name for name in before if before.get(name) != after.get(name)
            ]
            raise ShadowError(f"production hash changed: {changed}")

        integration_counts = integration_receipt.get("counts") or {}
        result = {
            "schema": SCHEMA,
            "status": "PASS_ISOLATED_PAIRED_RANK_SHADOW",
            "started_at": started.isoformat(),
            "finished_at": dt.datetime.now(dt.timezone.utc).isoformat(),
            "integration_generation_id": integration_receipt.get(
                "generation_id"
            ),
            "method": {
                "control": (
                    "frozen production baseline with the 160 unadjudicated "
                    "amendment rows removed from manual evidence"
                ),
                "treatment": (
                    "same control plus all 74 scoring-ready Greenlit "
                    "amendment/category edges"
                ),
                "rank_chain": [
                    "stance/per_category_stance.py",
                    "legislation/apply_leg_stance.py",
                    "engine/compute_pillars.py --write",
                ],
                "measure_weights": {"bill": 1.0, "amendment": 0.5},
                "cfpb_direction": {
                    "diminish_restrain_reform": {
                        "direction": 1,
                        "physical_block": "CFPB Restrain",
                        "raw_range": "CU:DB",
                    },
                    "support_expand_restore": {
                        "direction": -1,
                        "physical_block": "CFPB Reach (-)",
                        "raw_range": "EO:EV",
                    },
                },
                "amendment_success_policy": (
                    "activity and authorship are integrated at weight 0.5; "
                    "only explicit advanced latest-action phrases receive an "
                    "advanced tracker status"
                ),
            },
            "counts": {
                **event_counts,
                "source_decision_units": integration_counts.get(
                    "source_decision_units"
                ),
                "held_visible_non_scoring": integration_counts.get(
                    "held_decision_units"
                ),
                "ratified_exclusions": integration_counts.get(
                    "integration_state_counts", {}
                ).get("excluded_false_positive"),
                "manual_control": control_manual,
                "manual_treatment": treatment_manual,
            },
            "coverage": {
                "all_74_scoring_edges_in_event_ledger": event_count == 74,
                "all_mapped_edges_applied_to_raw_grid": (
                    workbook_audit["raw_activity_edges_applied"]
                    == event_counts["mapped_current_roster_edges"]
                ),
                "all_mapped_edges_applied_to_archive": (
                    archive_audit["applied_edges"]
                    == event_counts["mapped_current_roster_edges"]
                ),
                "all_19_directionless_memberships_retained_non_scoring": (
                    event_counts[
                        "membership_only_non_scoring_category_edges"
                    ]
                    == 19
                ),
                "all_1184_held_visible_non_scoring": (
                    integration_counts.get("held_decision_units") == 1184
                ),
                "all_2908_exclusions_preserved": (
                    integration_counts.get(
                        "integration_state_counts", {}
                    ).get("excluded_false_positive")
                    == 2908
                ),
                "all_4180_source_decisions_accounted": (
                    integration_counts.get("source_decision_units") == 4180
                ),
            },
            "workbook_integration": workbook_audit,
            "archive_integration": archive_audit,
            "workbook_verification": workbook_check,
            "rank_delta_summary": delta["summary"],
            "control_replay": delta["control_replay"],
            "largest_improvements": delta["largest_improvements"][:10],
            "largest_declines": delta["largest_declines"][:10],
            "pipeline_steps": {
                "control": control_steps,
                "treatment": treatment_steps,
            },
            "inputs": {
                "baseline_manifest": receipt(
                    BASELINE_ROOT / "baseline_manifest.json"
                ),
                "frozen_workbook": receipt(
                    BASELINE_FILES / WORKBOOK_NAME
                ),
                "frozen_rank_authority": receipt(
                    BASELINE_FILES / "ranks_python.json"
                ),
                "integration_receipt": receipt(
                    integration_root / "integration_receipt.json"
                ),
                "integrated_membership_edges": receipt(
                    integration_root / "integrated_membership_edges.ndjson",
                    93,
                ),
                "amendment_partition": receipt(
                    args.amendments.resolve(),
                    event_counts["amendment_partition_records"],
                ),
            },
            "outputs": {
                "shadow_event_ledger": receipt(event_path, event_count),
                "rank_delta": receipt(delta_path),
                "control_rank_authority": receipt(
                    control_root / "ranks_python.json"
                ),
                "treatment_rank_authority": receipt(
                    treatment_root / "ranks_python.json"
                ),
                "treatment_workbook": receipt(
                    treatment_root / WORKBOOK_NAME
                ),
                "treatment_archive_leg": receipt(
                    treatment_root / "archive_leg.json"
                ),
                "treatment_category_stance": receipt(
                    treatment_root / "category_stance.json"
                ),
                "treatment_manual_evidence": receipt(
                    treatment_root / "manual_bill_evidence.csv"
                ),
            },
            "safety": {
                "isolated_shadow_only": True,
                "production_hashes_before": before,
                "production_hashes_after": after,
                "production_hashes_unchanged": before == after,
                "production_workbook_written": False,
                "production_rank_written": False,
                "promotion_performed": False,
                "external_sync_performed": False,
            },
        }
        if not all(result["coverage"].values()):
            raise ShadowError("one or more final coverage gates failed")
        result["record_content_sha256"] = hashlib.sha256(
            canonical_bytes(result)
        ).hexdigest()
        receipt_path = out / "paired_shadow_receipt.json"
        atomic_json(receipt_path, result)
        print(json.dumps(result, indent=2, ensure_ascii=False, sort_keys=True))
        return 0
    except Exception as exc:
        after = production_hashes()
        failure = {
            "schema": SCHEMA,
            "status": "FAIL_CLOSED",
            "started_at": started.isoformat(),
            "failed_at": dt.datetime.now(dt.timezone.utc).isoformat(),
            "error_type": type(exc).__name__,
            "error": str(exc),
            "safety": {
                "production_hashes_before": before,
                "production_hashes_after": after,
                "production_hashes_unchanged": before == after,
            },
        }
        atomic_json(out / "paired_shadow_failure.json", failure)
        print(
            f"Greenlit paired rank shadow: FAIL CLOSED: {exc}",
            file=sys.stderr,
        )
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
