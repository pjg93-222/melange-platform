# -*- coding: utf-8 -*-
"""Repair the 'Sub Layer - Raw' efficiency block (AG:AW) after the 2026-07-15 Legislative-Normalized
own-row repair shifted that sheet's member rows up by one and left every downstream reference behind
(independently verified 2026-07-17: SubRaw row-4 member Aaron Bean reads LegNorm H4 = Abraham
Hamadeh; row-536 Zoe Lofgren reads an EMPTY row; audit wf_73d7c183-627 HIGH findings, refuter-
confirmed; FIX_CAMPAIGN_LOG entries 010+).

--plan  (read-only): audits EVERY cell AG4:AW536 — parses the formula, resolves the member's OWN
        row in LegNorm by name, resolves the slot's CANONICAL numerator column (the category's own
        success-rate / co-sponsor column, exactly the pattern the majority of cells already use),
        and prints/dumps every deviation. Writes data/engine/repair_sublayer_plan.json. NO workbook writes.
--apply: backs the workbook up (PRE-SUBLAYER-REPAIR-<ts>), rewrites ONLY the deviating cells to the
        canonical formula shape (same shape as the majority pattern: numerator ref fixed to the
        member's own LegNorm row + canonical column; Table1 MIN/MAX basis matched to the same
        column), saves atomically. Excel recalc + chain + verification happen OUTSIDE this script.

Rank note: this is the correct-math repair for the leg pillar (30% of rank weight). Ranks are
EXPECTED to move when broken math is fixed; the movement is proven by the independent recompute
in the campaign log before it stands (charter 002).
"""
import os, sys, json, re, shutil, datetime, collections
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
from paths import ROOT as B
import roster

WB = B + r'\Moneyball Main Version - Blended.xlsx'
PLAN = B + r'\data\engine\repair_sublayer_plan.json'
REF = re.compile(r"('Legislative Score - Normalized'!)(\$?)([A-Z]{1,2})(\$?)(\d+)")
T1 = re.compile(r"Table1\[([^\]]+)\]")

def load():
    wb = openpyxl.load_workbook(WB)
    return wb, wb['Sub Layer - Raw'], wb['Legislative Score - Normalized']

def build_maps(sr, ln):
    # member -> LegNorm row (names in LN col A, members start row 3)
    ln_row = {}
    for r in range(2, ln.max_row + 1):
        v = ln.cell(r, 1).value
        if isinstance(v, str) and v.strip() and v.strip() != 'Weight':
            ln_row[v.strip()] = r
    # Canonical basis per slot = the MODE of the Table1[...] min/max basis names (the formulas'
    # own declared category basis — for the two known-broken slots the audit showed the numerator
    # column drifted while the Table1 basis stayed on the correct category, so the basis, NOT the
    # numerator mode, is the trustworthy anchor). The canonical NUMERATOR column is then the
    # LegNorm column whose row-1 header exactly matches that basis name.
    slot_t1 = collections.defaultdict(collections.Counter)
    for r in roster.rows('Sub Layer - Raw'):
        for c in range(33, 50):
            t = T1.search(str(sr.cell(r, c).value or ''))
            if t: slot_t1[c][t.group(1)] += 1
    canon_t1 = {c: slot_t1[c].most_common(1)[0][0] for c in slot_t1}
    hdr2col = {}
    for col in range(1, ln.max_column + 1):
        h = ln.cell(1, col).value
        if isinstance(h, str) and h.strip(): hdr2col[h.strip()] = col
    canon_col = {}
    for c, basis in canon_t1.items():
        col = hdr2col.get(basis.strip())
        if col is None:
            raise SystemExit('cannot resolve Table1 basis %r to a LegNorm header column (slot %s)' % (basis, get_column_letter(c)))
        canon_col[c] = get_column_letter(col)
    return ln_row, canon_col, canon_t1

def audit(sr, ln, ln_row, canon_col, canon_t1):
    devs = []
    for r in roster.rows('Sub Layer - Raw'):
        nm = sr.cell(r, 2).value
        if not nm: continue
        own = ln_row.get(str(nm).strip())
        for c in range(33, 50):
            f = str(sr.cell(r, c).value or '')
            m = REF.search(f)
            if not m:
                if f.strip():
                    devs.append({'cell': '%s%d' % (get_column_letter(c), r), 'member': nm,
                                 'kind': 'no-legnorm-ref', 'formula': f[:120]})
                continue
            col, refrow = m.group(3), int(m.group(5))
            t = T1.search(f)
            problems = []
            if own is None: problems.append('member-not-in-legnorm')
            elif refrow != own: problems.append('row %d != own %d' % (refrow, own))
            if col != canon_col.get(c): problems.append('col %s != canon %s' % (col, canon_col.get(c)))
            if t and t.group(1) != canon_t1.get(c): problems.append('t1basis %r != canon %r' % (t.group(1)[:40], canon_t1.get(c, '')[:40]))
            if problems:
                devs.append({'cell': '%s%d' % (get_column_letter(c), r), 'member': nm, 'row_should_be': own,
                             'col_should_be': canon_col.get(c), 'kind': '; '.join(problems), 'formula': f[:160]})
    return devs

def main():
    wb, sr, ln = load()
    ln_row, canon_col, canon_t1 = build_maps(sr, ln)
    print('canonical numerator columns:', {get_column_letter(c): v for c, v in sorted(canon_col.items())})
    devs = audit(sr, ln, ln_row, canon_col, canon_t1)
    kinds = collections.Counter(d['kind'].split(';')[0] for d in devs)
    print('deviating cells: %d of %d | by first-kind: %s' % (len(devs), roster.count() * 17, dict(kinds)))
    if '--plan' in sys.argv or '--apply' not in sys.argv:
        json.dump({'generated': datetime.datetime.now().isoformat(timespec='seconds'),
                   'canonical_cols': {get_column_letter(c): v for c, v in canon_col.items()},
                   'deviations': devs}, open(PLAN, 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
        print('plan written -> %s (read-only run; use --apply to repair)' % PLAN)
        wb.close(); return
    # ---- APPLY ----
    bk = WB.replace('.xlsx', '.PRE-SUBLAYER-REPAIR-%s.xlsx' % datetime.datetime.now().strftime('%Y%m%d_%H%M%S'))
    shutil.copy2(WB, bk); print('backup ->', os.path.basename(bk))
    fixed = 0
    for d in devs:
        if d['kind'].startswith('no-legnorm-ref') or d.get('row_should_be') is None:
            continue   # stray/orphan rows are reported, not rewritten (separate decision)
        cell = d['cell']; c = column_index_from_string(re.match(r'[A-Z]+', cell).group(0)); r = int(re.search(r'\d+', cell).group(0))
        f = str(sr.cell(r, c).value or '')
        f2 = REF.sub(lambda m: "%s%s%s%s%d" % (m.group(1), m.group(2), canon_col[c], m.group(4), d['row_should_be']), f)
        f2 = T1.sub('Table1[%s]' % canon_t1[c], f2)
        if f2 != f:
            sr.cell(r, c).value = f2; fixed += 1
    from moneyball_core import save_workbook_atomic
    save_workbook_atomic(wb, WB)
    print('APPLIED: %d cells rewritten (of %d deviations; no-ref/orphan rows untouched). Next: COM recalc -> chain -> verification.' % (fixed, len(devs)))

if __name__ == '__main__':
    main()
