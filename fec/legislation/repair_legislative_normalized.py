# -*- coding: utf-8 -*-
import os as _o, sys as _s
_s.path.insert(0, _o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))))
from paths import ROOT as MB_ROOT  # centralized root
"""Repair the ONE-ROW misalignment of 'Legislative Score - Normalized' (2026-07-15).

THE BUG (see tests/infra/test_legislative_alignment.py for the full write-up + evidence)
---------------------------------------------------------------------------------
Each member's normalized value row r held the normalization of the ALPHABETICAL PREDECESSOR
(raw row r-1) instead of the member's OWN raw row r.  The displacement is present in BOTH the
~24k pasted LITERAL cells (values equal to RAW[r-1]/colmax) and the ~53k FORMULA cells
(=IFERROR('Legislative Score - Raw'!<col><r-1>/MAX(...),0), whose relative reference is one row
too high).  Because openpyxl cannot evaluate Excel formulas, the only way to land -- and to prove
with openpyxl, no Excel and no pipeline run -- a coherent, fully own-row-aligned sheet is to
MATERIALIZE each value/MAX cell as the corrected literal number.  This mirrors the shipped
reference implementation (extract_workbook_audit_inputs.normalize_columns / build_audit_workbook,
which write literal normalized values) and the values-write + fullCalcOnLoad pattern of
sync_stance_g.py (Python owns the derived math; the workbook is a self-healing VIEW that recomputes
its remaining formulas when Paul next opens it, and Paul re-propagates the pipeline after review).

THE RULE (derived from the surviving aligned cells AND the formula text; no negatives exist, so
signed-magnitude reduces to value/MAX):
        NRM[r][c] = 0.0                     if colmax(c) == 0     (IFERROR(...,0))
                  = RAW[r][c] / colmax(c)   otherwise
        colmax(c) = max(0, max RAW over data rows 3..535)

LEFT UNCHANGED (reported): col 144 ('None', a MIN-MAX rule over a different raw column, already
own-row) and cols 146,148,150,151,152 (blank normalized + all-zero raw -> no derivable
normalization; a shifted blank column is still blank).  Everything outside the value region -- the
header row, the weight row, the name/district columns, every other sheet, and all their formulas --
is preserved untouched.

Backs up the workbook FIRST.  Excel/openpyxl only; NO pipeline run.  Idempotent (re-running writes
the same values).  Usage:  python fec/legislation/repair_legislative_normalized.py [--dry-run]
"""
import datetime
import os
import shutil
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from moneyball_core import save_workbook_atomic  # crash-safe verified workbook replace

B = MB_ROOT
WB = B + r"\Moneyball Main Version - Blended.xlsx"
BACKUP_DIR = os.path.join(B, "backups")
RAW_SHEET = "Legislative Score - Raw"
NRM_SHEET = "Legislative Score - Normalized"

FIRST_DATA_ROW, LAST_DATA_ROW = 3, 535
FIRST_VALUE_COL, LAST_VALUE_COL = 3, 160  # 160 since 2026-07-17 (CU Tax Status block 153-160)
WEIGHT_ROW = 2
EXCLUDED_COLS = frozenset({144, 146, 148, 150, 151, 152})


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def value_max_cols():
    return [c for c in range(FIRST_VALUE_COL, LAST_VALUE_COL + 1) if c not in EXCLUDED_COLS]


def guard_not_locked():
    lock = os.path.join(os.path.dirname(WB), "~$" + os.path.basename(WB))
    if os.path.exists(lock):
        sys.exit("REFUSING: Excel lock file present (%s) -- close the workbook first." % lock)


def main():
    dry = "--dry-run" in sys.argv
    guard_not_locked()

    # ---- read the RAW grid (all literal numbers) for the derivation -----------------------------
    rd = openpyxl.load_workbook(WB, data_only=False, read_only=True)
    raw = list(rd[RAW_SHEET].iter_rows(values_only=True))
    nrm_view = list(rd[NRM_SHEET].iter_rows(values_only=True))
    rd.close()

    def rawval(r, c):
        row = raw[r - 1]
        return num(row[c - 1]) if c - 1 < len(row) else 0.0

    def nrmview(r, c):
        row = nrm_view[r - 1]
        return row[c - 1] if c - 1 < len(row) else None

    cols = value_max_cols()
    colmax = {c: max([0.0] + [rawval(r, c) for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1)]) for c in cols}

    # ---- SAFETY: confirm rows-3..535 max reproduces the formula denominator (weight <= max) ------
    weight_violations = [c for c in cols if rawval(WEIGHT_ROW, c) > colmax[c] + 1e-12 and colmax[c] > 0]
    if weight_violations:
        sys.exit("REFUSING: weight row exceeds data max in cols %s -- denominator ambiguous; inspect before repair."
                 % weight_violations)

    # ---- plan: classify each target cell's current state, compute corrected value ----------------
    was_formula = was_literal_changed = was_literal_same = was_blank = 0
    for c in cols:
        mx = colmax[c]
        for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
            want = 0.0 if mx == 0 else rawval(r, c) / mx
            cur = nrmview(r, c)
            if cur is None:
                was_blank += 1
            elif isinstance(cur, str):
                was_formula += 1
            elif abs(num(cur) - want) > 1e-12:
                was_literal_changed += 1
            else:
                was_literal_same += 1
    total_cells = len(cols) * (LAST_DATA_ROW - FIRST_DATA_ROW + 1)

    print("Legislative-Normalized row-alignment repair")
    print("  workbook            :", WB)
    print("  value/MAX columns   : %d  (excluded, left unchanged: %s)" % (len(cols), sorted(EXCLUDED_COLS)))
    print("  target cells        : %d" % total_cells)
    print("  formula -> value    : %d  (were =IFERROR(..r-1..) formulas, now materialized own-row values)" % was_formula)
    print("  literal value fixed : %d  (pasted r-1 shift corrected to own row)" % was_literal_changed)
    print("  literal unchanged   : %d  (already own-row / zero)" % was_literal_same)
    print("  blank -> value      : %d" % was_blank)
    print("  cells whose numeric value changes: %d" % (was_formula + was_literal_changed + was_blank))

    if dry:
        print("\n--dry-run: no write performed.")
        return 0

    # ---- backup FIRST ---------------------------------------------------------------------------
    os.makedirs(BACKUP_DIR, exist_ok=True)
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = os.path.join(BACKUP_DIR, os.path.basename(WB).replace(".xlsx", ".PRE-LEGNORM-%s.xlsx" % ts))
    shutil.copy2(WB, backup)
    print("\n  backup -> %s" % backup)

    # ---- write corrected literal values (formulas preserved everywhere else) ---------------------
    wb = openpyxl.load_workbook(WB, data_only=False)   # formulas + tables + other sheets preserved
    ws = wb[NRM_SHEET]
    written = 0
    for c in cols:
        mx = colmax[c]
        for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
            ws.cell(r, c).value = 0.0 if mx == 0 else rawval(r, c) / mx
            written += 1
    wb.calculation.fullCalcOnLoad = True               # remaining formulas self-heal on next Excel open
    save_workbook_atomic(wb, WB)
    print("  wrote %d corrected value cells to '%s'; fullCalcOnLoad set; saved atomically." % (written, NRM_SHEET))
    print("SAVED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
