# -*- coding: utf-8 -*-
"""Add the 'Credit Union Tax Status' block (cols 153-160) to 'Legislative Score - Raw',
paste ALL 7 letters' Sponsor/Original-Co counts (CDFI 91/93, CFPB Restrain 99/101,
CCM 107/109, CU Tax 153/155), and re-materialize the affected 'Legislative Score -
Normalized' literal columns (house pattern: Python owns Raw->LegNorm math since the
2026-07-15 materialization; fullCalcOnLoad heals remaining formulas).

Paul-directed 2026-07-17 ("Add the block"). Backup FIRST; refuses if Excel lock present;
atomic save; Table1(A1:EM535 LegNorm)/Table7(A1:EV535 Raw) untouched (new cols outside both).
Usage: python add_cutax_block.py [--dry-run]
"""
import io,sys,csv,os,shutil,datetime
sys.stdout=io.TextIOWrapper(sys.stdout.buffer,encoding='utf-8')
import openpyxl
sys.path.insert(0,os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from paths import ROOT as MB_ROOT
import roster
from moneyball_core import save_workbook_atomic
import bill_rules as BR

B=MB_ROOT
WB=B+r'\Moneyball Main Version - Blended.xlsx'
DRY='--dry-run' in sys.argv

lock=os.path.join(os.path.dirname(WB),'~$'+os.path.basename(WB))
if os.path.exists(lock): sys.exit('REFUSING: Excel lock file present — close the workbook first.')

# ---- bill-equivalent decayed letter units per (member, category) from the live CSV ----
from collections import defaultdict
agg=defaultdict(lambda:[0,0])  # (member,cat) -> [sponsor, origco]
for r in csv.DictReader(open(B+r'\manual_bill_evidence.csv',encoding='utf-8-sig')):
    if not r['Bill'].startswith('LTR-'): continue
    k=(r['Member'],r['Category'])
    decay=BR.congress_decay(r['Congress'])
    if r['Role']=='sponsor': agg[k][0]+=decay
    elif r['Role']=='original-cosponsor': agg[k][1]+=decay

# category -> (sponsor_col, origco_col) in Legislative Score - Raw; CU Tax = the NEW block
CATCOLS={'CDFI Fund':(91,93),'CFPB':(99,101),'Credit Card Mandates':(107,109),'CU Tax Status':(153,155)}
NEWCOL0=153
NEWHDR=['Credit Union Tax Status - Sponsor','Credit Union Tax Status - Sponsor - Average',
        'Credit Union Tax Status - Original Co-Sponsor','Credit Union Tax Status - Original Co-Sponsor - Average',
        'Credit Union Tax Status - Co-Sponsor','Credit Union Tax Status - Success Rate',
        'Credit Union Tax Status - Sponsor - Success Rate','Credit Union Tax Status - Original Co-Sponsor - Success Rate']
NEWW=[2/7.0,0,1/7.0,0,0.5/7.0,0,0,0]   # mirrors every other block's row-2 weights

wb=openpyxl.load_workbook(WB)             # formulas preserved
ls=wb['Legislative Score - Raw']; ln=wb['Legislative Score - Normalized']

# sanity: expected headers before touching anything (fail loud on drift)
EXPECT={91:'CDFI Fund - Sponsor',93:'CDFI Fund - Original Co-Sponsorship',
        99:'CFPB Restrain - Sponsor',101:'CFPB Restrain - Original Co-Sponsor',
        107:'Credit Card & Consumer-Fee Mandates - Sponsor',109:'Credit Card & Consumer-Fee Mandates - Original Co-Sponsor'}
for c,h in EXPECT.items():
    got=str(ls.cell(1,c).value).strip()
    if got!=h: sys.exit(f'REFUSING: Raw col {c} header {got!r} != expected {h!r}')
for c in range(NEWCOL0,NEWCOL0+8):
    if ls.cell(1,c).value not in (None,''): sys.exit(f'REFUSING: Raw col {c} not empty ({ls.cell(1,c).value!r})')
    if ln.cell(1,c).value not in (None,''): sys.exit(f'REFUSING: LegNorm col {c} not empty')

# member rows (Raw + LegNorm both: the roster's member rows, same order)
members={}
for r in roster.rows('Legislative Score - Raw'):
    nm=ls.cell(r,1).value
    if nm: members[str(nm).strip()]=r
assert len(members)==roster.count(), f'member rows {len(members)}!={roster.count()}'
missing=[m for (m,c) in agg if m not in members]
if missing: sys.exit(f'REFUSING: letter members not in workbook rows: {sorted(set(missing))[:8]}')

# ---- 1. new CU Tax block on Raw (headers, weights, zeros) ----
for i,(h,w) in enumerate(zip(NEWHDR,NEWW)):
    ls.cell(1,NEWCOL0+i).value=h
    ls.cell(2,NEWCOL0+i).value=w
for r in roster.rows('Legislative Score - Raw'):
    for i in range(8): ls.cell(r,NEWCOL0+i).value=0

# ---- 2. paste letter counts (+=) ----
applied=[]
for (m,cat),(s,o) in sorted(agg.items()):
    sc,oc=CATCOLS[cat]; r=members[m]
    if s:
        cur=ls.cell(r,sc).value or 0; ls.cell(r,sc).value=cur+s; applied.append((m,cat,'sponsor',cur,cur+s))
    if o:
        cur=ls.cell(r,oc).value or 0; ls.cell(r,oc).value=cur+o; applied.append((m,cat,'origco',cur,cur+o))

# ---- 3. re-materialize LegNorm literals for every touched column ----
TOUCHED=[91,93,99,101,107,109]+list(range(NEWCOL0,NEWCOL0+8))
for i,h in enumerate(NEWHDR):                      # headers + weights on LegNorm too
    ln.cell(1,NEWCOL0+i).value=h
    ln.cell(2,NEWCOL0+i).value=NEWW[i]
recalc=0
for c in TOUCHED:
    vals={r:(ls.cell(r,c).value or 0) for r in roster.rows('Legislative Score - Raw')}
    mx=max([0.0]+[float(v) for v in vals.values() if isinstance(v,(int,float))])
    for r in roster.rows('Legislative Score - Raw'):
        v=float(vals[r]) if isinstance(vals[r],(int,float)) else 0.0
        ln.cell(r,c).value=0.0 if mx==0 else v/mx
        recalc+=1

print(f'letter count cells applied: {len(applied)} | LegNorm cells materialized: {recalc} (cols {TOUCHED})')
print('sample applied:',applied[:5]);print('...',applied[-3:])
if DRY:
    print('--dry-run: NO SAVE.'); sys.exit(0)

# ---- backup FIRST, then atomic save ----
os.makedirs(B+r'\backups',exist_ok=True)
ts=datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
bk=B+r'\backups\Moneyball Main Version - Blended.PRE-CUTAX-BLOCK-%s.xlsx'%ts
shutil.copy2(WB,bk); print('backup ->',bk)
wb.calculation.fullCalcOnLoad=True
save_workbook_atomic(wb,WB)
print('SAVED atomically. Block cols %d-%d added; counts landed; LegNorm literals updated.'%(NEWCOL0,NEWCOL0+7))
