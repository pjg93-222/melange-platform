# -*- coding: utf-8 -*-
"""Run the unchanged pillar/rank authority only from a verified shadow root.

This wrapper exists because ``compute_pillars.py`` intentionally tolerates a
dead workbook formula cache by reading ``pillar_inputs_snapshot.json``.  That
is useful in the live system but is not acceptable for a legislation
recomputation.  Here every Blended Model O/E/F cache must be present, must
match the generation-bound snapshot, and the snapshot is frozen while the
existing rank authority runs.

The wrapper never edits a workbook and refuses production roots.  Its receipt
binds the accepted legislation generation, verified workbook-recompute
receipt, exact prerequisite hashes, unchanged source contract, and output
hashes.  It does not claim that centrality affects ranks: the existing
authority still uses the workbook committee score as its influence pillar.
"""
from __future__ import annotations

import os as _bootstrap_os, sys as _bootstrap_sys; _bootstrap_sys.path.insert(0, _bootstrap_os.path.dirname(_bootstrap_os.path.dirname(_bootstrap_os.path.abspath(__file__))))

import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
from typing import Any, Iterable, Sequence

import openpyxl

from paths import ROOT as MB_ROOT

try:
    from lane_file_io import atomic_write_bytes
except ImportError:  # pragma: no cover - package import path
    from fec.lane_file_io import atomic_write_bytes


CODE_ROOT = Path(__file__).resolve().parents[2]
SCRIPT = CODE_ROOT / "fec" / "engine" / "compute_pillars.py"
PREREQUISITE_SCHEMA = "legislation-pillars-ranks-prerequisites-v1"
WORKBOOK_RECEIPT_SCHEMA = "legislation-workbook-recompute-receipt-v1"
RECEIPT_SCHEMA = "legislation-pillars-ranks-receipt-v1"
PREREQUISITE_NAME = "pillars_ranks_prerequisites.json"
RECEIPT_NAME = "pillars_ranks.receipt.json"
WORKBOOK_NAME = "Moneyball Main Version - Blended.xlsx"
SNAPSHOT_NAME = "data/engine/pillar_inputs_snapshot.json"
RANKS_NAME = "ranks_python.json"
ROLES_NAME = "data/members/formal_visibility_roles.json"
SNAPSHOT_FIRST = "intermediate/pillar_inputs_first_pass.json"
RANKS_FIRST = "intermediate/ranks_first_pass.json"
HEX_RE = re.compile(r"^[0-9a-f]{64}$")
REQUIRED_ARTIFACTS = frozenset(
    {WORKBOOK_NAME, SNAPSHOT_NAME, ROLES_NAME}
)
SOURCE_TOKENS = (
    "WRITE='--write' in sys.argv",
    "PILLAR_SNAPSHOT_FREEZE",
    "legp=_leg_percentiles[m]",
    "infp=E[m]/emax",
    "score[m]=AB[2]*legp+AB[3]*pressp+AB[4]*infp+AB[5]*demp",
    "B+r'\\ranks_python.json'",
)


class PillarError(RuntimeError):
    """A fail-closed staged rank computation error."""


def parser() -> argparse.ArgumentParser:
    root = Path(MB_ROOT)
    result = argparse.ArgumentParser(
        description="Run hash-bound legislation pillars/ranks in a shadow root."
    )
    result.add_argument("--run-dir", type=Path, default=root)
    result.add_argument("--prerequisites", type=Path)
    return result


def canonical_json(value: Any) -> str:
    return json.dumps(
        value, sort_keys=True, ensure_ascii=False, separators=(",", ":")
    )


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def receipt(path: Path, *, relative_to: Path | None = None) -> dict[str, Any]:
    shown = path.relative_to(relative_to) if relative_to else path.resolve()
    return {
        "path": str(shown).replace("\\", "/"),
        "size": path.stat().st_size,
        "sha256": sha256_file(path),
    }


def read_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise PillarError(f"required JSON does not exist: {path}")
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise PillarError(f"invalid JSON at {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise PillarError(f"JSON root must be an object: {path}")
    return value


def is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def validate_run_dir(raw: Path) -> Path:
    root = raw.expanduser().resolve()
    allowed = {
        (CODE_ROOT / "audits" / "legfix5").resolve(),
        (CODE_ROOT / "tmp").resolve(),
        Path(tempfile.gettempdir()).resolve(),
    }
    if root.suffix or any(root == parent for parent in allowed):
        raise PillarError("run directory must be a child staging directory")
    if not any(is_relative_to(root, parent) for parent in allowed):
        raise PillarError(
            "run directory must be under audits/legfix5, repo tmp, or OS temp"
        )
    if not root.is_dir():
        raise PillarError(f"run directory does not exist: {root}")
    return root


def normalize_artifact_path(run_dir: Path, raw: Any) -> Path:
    path = Path(str(raw or ""))
    if not str(raw or "") or path.is_absolute() or ".." in path.parts:
        raise PillarError(f"unsafe prerequisite artifact path: {raw!r}")
    resolved = (run_dir / path).resolve()
    if not is_relative_to(resolved, run_dir):
        raise PillarError(f"artifact escapes run directory: {raw!r}")
    return resolved


def validate_hash_entry(
    *,
    run_dir: Path,
    name: str,
    entry: Any,
) -> dict[str, Any]:
    if not isinstance(entry, dict):
        raise PillarError(f"artifact receipt is absent: {name}")
    path = normalize_artifact_path(run_dir, entry.get("path"))
    expected = str(entry.get("sha256") or "").casefold()
    if not HEX_RE.fullmatch(expected):
        raise PillarError(f"artifact lacks valid SHA-256: {name}")
    if not path.is_file():
        raise PillarError(f"artifact is missing: {path}")
    observed = sha256_file(path)
    if observed != expected:
        raise PillarError(
            f"artifact hash mismatch for {name}: expected {expected}, "
            f"observed {observed}"
        )
    if entry.get("size") is not None and path.stat().st_size != int(
        entry["size"]
    ):
        raise PillarError(f"artifact size mismatch: {name}")
    return {**receipt(path, relative_to=run_dir), "resolved_path": path}


def validate_workbook_receipt(
    *,
    run_dir: Path,
    item: Any,
    generation_id: str,
    workbook_sha256: str,
    snapshot_sha256: str,
) -> dict[str, Any]:
    if not isinstance(item, dict):
        raise PillarError("verified legislative-workbook receipt is absent")
    path = normalize_artifact_path(run_dir, item.get("path"))
    expected = str(item.get("sha256") or "").casefold()
    if not path.is_file() or not HEX_RE.fullmatch(expected):
        raise PillarError("legislative-workbook receipt path/hash is invalid")
    if sha256_file(path) != expected:
        raise PillarError("legislative-workbook receipt hash mismatch")
    value = read_json(path)
    outputs = value.get("outputs") or {}
    if not (
        value.get("schema") == WORKBOOK_RECEIPT_SCHEMA
        and value.get("status") == "PASS"
        and value.get("staged_only") is True
        and value.get("production_written") is False
        and value.get("generation_id") == generation_id
        and (outputs.get(WORKBOOK_NAME) or {}).get("sha256")
        == workbook_sha256
        and (outputs.get(SNAPSHOT_NAME) or {}).get("sha256")
        == snapshot_sha256
    ):
        raise PillarError(
            "legislative-workbook receipt is not PASS/hash-bound to this "
            "workbook, snapshot, and accepted generation"
        )
    return receipt(path, relative_to=run_dir)


def validate_source_contract() -> dict[str, Any]:
    if not SCRIPT.is_file():
        raise PillarError(f"rank authority script is missing: {SCRIPT}")
    text = SCRIPT.read_text(encoding="utf-8-sig", errors="replace")
    missing = [token for token in SOURCE_TOKENS if token not in text]
    if missing:
        raise PillarError(
            "compute_pillars source contract changed: " + ", ".join(missing)
        )
    return receipt(SCRIPT)


def workbook_cache(
    path: Path,
) -> tuple[list[str], dict[str, dict[str, float]]]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        required_sheets = {
            "Rank & Allocation",
            "Blended Model",
            "Press Release - Raw",
            "Press Release - Normalized",
            "Committee Data",
        }
        missing = sorted(required_sheets - set(workbook.sheetnames))
        if missing:
            raise PillarError(
                "rank workbook lacks required sheets: " + ", ".join(missing)
            )
        rank_sheet = workbook["Rank & Allocation"]
        blended = workbook["Blended Model"]
        members = [
            str(rank_sheet.cell(row, 1).value)
            for row in range(3, 536)
            if rank_sheet.cell(row, 1).value
        ]
        if not members or len(set(members)) != len(members):
            raise PillarError("rank workbook member roster is empty or duplicated")
        rows = {
            str(blended.cell(row, 1).value): row
            for row in range(3, 536)
            if blended.cell(row, 1).value
        }
        if set(rows) != set(members):
            raise PillarError(
                "Blended Model formula-cache roster differs from Rank & "
                "Allocation roster"
            )
        values: dict[str, dict[str, float]] = {"O": {}, "E": {}, "F": {}}
        columns = {"O": 15, "E": 5, "F": 6}
        for member in members:
            row = rows[member]
            for label, column in columns.items():
                raw = blended.cell(row, column).value
                if not isinstance(raw, (int, float)) or not math.isfinite(
                    float(raw)
                ):
                    raise PillarError(
                        f"dead/non-numeric {label} cache for {member!r}"
                    )
                values[label][member] = float(raw)
        return members, values
    finally:
        workbook.close()


def validate_snapshot(
    path: Path,
    members: Sequence[str],
    cached: dict[str, dict[str, float]],
) -> dict[str, Any]:
    value = read_json(path)
    for label in ("O", "E", "F"):
        observed = value.get(label)
        if not isinstance(observed, dict) or set(observed) != set(members):
            raise PillarError(
                f"snapshot {label} member set differs from workbook cache"
            )
        for member in members:
            try:
                number = float(observed[member])
            except (TypeError, ValueError) as exc:
                raise PillarError(
                    f"snapshot {label} value is invalid for {member!r}"
                ) from exc
            if not math.isclose(
                number, cached[label][member], rel_tol=0.0, abs_tol=1e-12
            ):
                raise PillarError(
                    f"snapshot {label} is stale for {member!r}"
                )
    return value


def file_state(root: Path) -> dict[str, str]:
    return {
        str(path.relative_to(root)).replace("\\", "/"): sha256_file(path)
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def validate_ranks(path: Path, members: Sequence[str]) -> dict[str, Any]:
    value = read_json(path)
    required_maps = ("score", "rank", "pcts", "ab_weights")
    if any(not isinstance(value.get(key), dict) for key in required_maps):
        raise PillarError("rank authority lacks required score/rank decomposition")
    if set(value["score"]) != set(members) or set(value["rank"]) != set(members):
        raise PillarError("rank authority member set differs from workbook")
    if any(
        not isinstance(value["score"][member], (int, float))
        or not math.isfinite(float(value["score"][member]))
        or not isinstance(value["rank"][member], int)
        or value["rank"][member] < 1
        for member in members
    ):
        raise PillarError("rank authority contains invalid score/rank values")
    return value


def write_receipt_idempotently(path: Path, value: dict[str, Any]) -> bool:
    body = json.dumps(
        value, indent=2, sort_keys=True, ensure_ascii=False
    ) + "\n"
    if path.exists():
        if path.read_text(encoding="utf-8-sig") != body:
            raise PillarError(
                "existing staged pillar/rank receipt differs for identical root"
            )
        return True
    path.write_text(body, encoding="utf-8", newline="\n")
    return False


def atomic_write(path: Path, body: bytes) -> None:
    # The scratch name used to be fixed, so two recomputes writing the same
    # output clobbered each other's staging file, and OneDrive syncs a stable
    # scratch name and mints conflict copies from it.  ``atomic_write_bytes``
    # is unique per process and per call, fsync'd before publication, and
    # retries the publish through the OneDrive/AV lock window.
    atomic_write_bytes(path, body)


def atomic_copy(source: Path, destination: Path) -> None:
    # Verification now compares the source against the exact bytes being
    # published, rather than against a separate scratch file that a sibling
    # producer could have replaced between the copy and the check.
    body = source.read_bytes()
    if sha256_bytes(body) != sha256_file(source):
        raise PillarError(f"hash-copy verification failed: {destination}")
    atomic_write(destination, body)


def run(
    *,
    run_dir: Path,
    prerequisites_path: Path | None,
) -> dict[str, Any]:
    root = validate_run_dir(run_dir)
    prerequisite_path = (
        prerequisites_path.expanduser().resolve()
        if prerequisites_path
        else root / PREREQUISITE_NAME
    )
    if not is_relative_to(prerequisite_path, root):
        raise PillarError("prerequisites must be inside the shadow run directory")
    prerequisite = read_json(prerequisite_path)
    generation_id = str(prerequisite.get("generation_id") or "")
    if not (
        prerequisite.get("schema") == PREREQUISITE_SCHEMA
        and prerequisite.get("status") == "PASS"
        and prerequisite.get("staged_only") is True
        and prerequisite.get("production_written") is False
        and HEX_RE.fullmatch(generation_id)
    ):
        raise PillarError(
            "pillar/rank prerequisites are not a PASS staged generation"
        )
    artifacts_raw = prerequisite.get("artifacts")
    if not isinstance(artifacts_raw, dict):
        raise PillarError("pillar/rank prerequisite artifacts are absent")
    missing = sorted(REQUIRED_ARTIFACTS - set(artifacts_raw))
    if missing:
        raise PillarError(
            "pillar/rank prerequisite artifacts missing: " + ", ".join(missing)
        )
    artifacts = {
        name: validate_hash_entry(
            run_dir=root, name=name, entry=artifacts_raw[name]
        )
        for name in sorted(REQUIRED_ARTIFACTS)
    }
    workbook_receipt = validate_workbook_receipt(
        run_dir=root,
        item=prerequisite.get("legislative_workbook_receipt"),
        generation_id=generation_id,
        workbook_sha256=artifacts[WORKBOOK_NAME]["sha256"],
        snapshot_sha256=artifacts[SNAPSHOT_NAME]["sha256"],
    )
    source = validate_source_contract()
    members, cached = workbook_cache(artifacts[WORKBOOK_NAME]["resolved_path"])
    validate_snapshot(
        artifacts[SNAPSHOT_NAME]["resolved_path"], members, cached
    )
    before = file_state(root)
    snapshot_before = artifacts[SNAPSHOT_NAME]["sha256"]
    workbook_before = artifacts[WORKBOOK_NAME]["sha256"]
    environment = os.environ.copy()
    environment.update(
        {
            "MONEYBALL_ROOT": str(root),
            "PILLAR_SNAPSHOT_FREEZE": "1",
        }
    )
    completed = subprocess.run(
        [sys.executable, str(SCRIPT), "--write"],
        cwd=str(CODE_ROOT),
        env=environment,
        text=True,
        capture_output=True,
        timeout=900,
        check=False,
    )
    if completed.returncode != 0:
        raise PillarError(
            "compute_pillars failed in shadow root: "
            + (completed.stderr or completed.stdout)[-2000:]
        )
    if sha256_file(artifacts[WORKBOOK_NAME]["resolved_path"]) != workbook_before:
        raise PillarError("compute_pillars modified the staged workbook")
    if sha256_file(artifacts[SNAPSHOT_NAME]["resolved_path"]) != snapshot_before:
        raise PillarError(
            "compute_pillars modified the frozen pillar snapshot"
        )
    ranks_path = root / RANKS_NAME
    validate_ranks(ranks_path, members)
    snapshot_first_path = root / SNAPSHOT_FIRST
    ranks_first_path = root / RANKS_FIRST
    atomic_copy(root / SNAPSHOT_NAME, snapshot_first_path)
    atomic_copy(ranks_path, ranks_first_path)
    after = file_state(root)
    changed = sorted(
        name
        for name in set(before) | set(after)
        if before.get(name) != after.get(name)
    )
    unexpected = [
        name
        for name in changed
        if name
        not in {RANKS_NAME, SNAPSHOT_FIRST, RANKS_FIRST, RECEIPT_NAME}
    ]
    if unexpected:
        raise PillarError(
            "rank authority wrote unexpected staged artifacts: "
            + ", ".join(unexpected)
        )
    outputs = {
        SNAPSHOT_NAME: receipt(root / SNAPSHOT_NAME, relative_to=root),
        RANKS_NAME: receipt(ranks_path, relative_to=root),
        SNAPSHOT_FIRST: receipt(snapshot_first_path, relative_to=root),
        RANKS_FIRST: receipt(ranks_first_path, relative_to=root),
    }
    basis = {
        "generation_id": generation_id,
        "prerequisites_sha256": sha256_file(prerequisite_path),
        "legislative_workbook_receipt_sha256": workbook_receipt["sha256"],
        "compute_pillars_sha256": source["sha256"],
        "outputs": outputs,
    }
    result = {
        "schema": RECEIPT_SCHEMA,
        "status": "PASS",
        "acceptance_passed": True,
        "staged_only": True,
        "production_written": False,
        "mainframe_written": False,
        "generation_id": generation_id,
        "stale_snapshot_fallback_used": False,
        "workbook_formula_cache_complete": True,
        "snapshot_matches_workbook_cache": True,
        "snapshot_frozen": True,
        "first_pass_copies_byte_identical": (
            outputs[SNAPSHOT_NAME]["sha256"]
            == outputs[SNAPSHOT_FIRST]["sha256"]
            and outputs[RANKS_NAME]["sha256"]
            == outputs[RANKS_FIRST]["sha256"]
        ),
        "centrality_rank_policy": (
            "unchanged live method: workbook committee score is the influence "
            "pillar; base/AUG centrality does not modify rank"
        ),
        "inputs": {
            "prerequisites": receipt(prerequisite_path, relative_to=root),
            "legislative_workbook_receipt": workbook_receipt,
            "artifacts": {
                name: {
                    key: value
                    for key, value in item.items()
                    if key != "resolved_path"
                }
                for name, item in artifacts.items()
            },
        },
        "producer": source,
        "member_count": len(members),
        "outputs": outputs,
        "generation_basis_sha256": sha256_bytes(
            canonical_json(basis).encode("utf-8")
        ),
        "subprocess": {
            "argv": ["{python}", "fec/engine/compute_pillars.py", "--write"],
            "moneyball_root": ".",
            "pillar_snapshot_freeze": True,
            "returncode": completed.returncode,
            "stdout_sha256": sha256_bytes(completed.stdout.encode("utf-8")),
            "stderr_sha256": sha256_bytes(completed.stderr.encode("utf-8")),
        },
    }
    reused = write_receipt_idempotently(root / RECEIPT_NAME, result)
    result["_idempotent_reuse"] = reused
    return result


def main(argv: Sequence[str] | None = None) -> int:
    args = parser().parse_args(argv)
    try:
        result = run(
            run_dir=args.run_dir,
            prerequisites_path=args.prerequisites,
        )
    except PillarError as exc:
        print(
            f"run_legislation_pillars_and_ranks: FAIL CLOSED: {exc}",
            file=sys.stderr,
        )
        return 2
    print(
        "legislation pillars/ranks: PASS | "
        f"members={result['member_count']} | "
        f"idempotent_reuse={result.pop('_idempotent_reuse', False)}"
    )
    print((args.run_dir.resolve() / RECEIPT_NAME))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
