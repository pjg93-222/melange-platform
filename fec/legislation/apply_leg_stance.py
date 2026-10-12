# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
"""Make the legislative pillar STANCE-AWARE: gate each member's per-topic cosponsorship activity (Sponsor +
Original-Co-Sponsor + Co-Sponsorship counts) by their per-category stance, weighted by issue importance.
Anti-CU members' activity now counts against them instead of inflating them. Writes the new leg raw to
Blended Model col4 (overriding the =R&A!D3 activity-only value) via COM + recalc -> Leg% -> Score -> Rank.
Backs up first.

2026-07-17 (Paul, letters integration, "whatever is most optimal"): Original Co-Sponsor count columns now
INCLUDED in the per-topic activity sum. The old filter zeroed them while counting weaker plain
co-sponsorship — incoherent with every other layer (bill_rules 3/2/1 role weights, the grid's own
4:2:1 row-2 weights, the say-do Champion gate's A-path). Revert switch: LEGSTANCE_ORIGCO=0 restores the
sponsor+plain-cosponsor-only sum."""
import openpyxl, json, io, sys, shutil, os, datetime, unicodedata
import win32com.client as win32
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
WB = MB_ROOT+r'\Moneyball Main Version - Blended.xlsx'
CS = json.load(io.open(MB_ROOT+r'\category_stance.json', encoding='utf-8'))['matrix']
# IMP anchored to the ABB26 handout (America's Credit Unions 2026 Advocacy Priorities, applied 2026-07-01):
# Protect flagships 2.0-3.0 (Tax Status, Credit Card Mandates, Interchange); named priorities 0.75-1.0
# (NCUA/regulatory regime, FOM/expand-access, Charter, Data Privacy, CFPB, Fraud, Tech/innovation, Payments);
# unlisted categories unchanged. NOTE: canonical dict = moneyball_core.IMP (imported here); legacy copies remain in
# sync_stance(.clean)/corrected_signals -- keep those in sync (build_dashboard_bundle now imports it).
from moneyball_core import IMP
LEGPRE = {'Examination & Supervision':'examination & supervision','CU Capital Requirements':'credit union capital requirements','NCUA Budget & Insurance':'ncua budget','Interchange':'interchange','Payments':'payments','Technology & Emerging':'technology','Fraud Prevention':'fraud prevention','Data Privacy & Cyber':'data privacy','Charter Modernization':'charter modernization','Defense CU Issues':'defense credit union','Field of Membership':'field of membership','CDFI Fund':'cdfi fund','CFPB':'cfpb','Credit Card Mandates':'credit card & consumer-fee','TCPA':'telephone consumer','Member Business Lending':'member business lending','Mortgage & Housing':'mortgage lending','CU Tax Status':'credit union tax status'}
SM = {'Champion':1.0,'Defender':1.0,'Ally':0.6,'Highly Persuadable':0.25,'Persuadable':0.1,'Lean Persuadable':0.0,'True Neutral':0.0,'Neutral':0.0,'Lean Skeptic':-0.3,'Skeptic':-0.6,'Highly Skeptical':-0.8,'Opponent':-1.0,'—':0.0}
def norm(s): return str(s).lower().replace(' and ',' & ').strip()
def nfc(s): return unicodedata.normalize('NFC', str(s)) if s is not None else s

# ---- compute stance-gated leg raw per member ----
wb = openpyxl.load_workbook(WB, read_only=True, data_only=True)
ls = wb['Legislative Score - Raw']; lh = [ls.cell(1, c).value for c in range(1, ls.max_column + 1)]
_ORIGCO = os.environ.get('LEGSTANCE_ORIGCO', '1') != '0'   # 2026-07-17: include Original Co-Sponsor counts (revert: LEGSTANCE_ORIGCO=0)
topcols = {t: [] for t in IMP}
for ci, h in enumerate(lh):
    if not h: continue
    hn = norm(h)
    if 'average' in hn or 'success' in hn or 'sponsor' not in hn: continue
    if 'original' in hn and not _ORIGCO: continue
    for t, pre in LEGPRE.items():
        if hn.startswith(norm(pre)): topcols[t].append(ci); break
legst = {}
for row in ls.iter_rows(min_row=2, values_only=True):
    nm = row[0]
    if not nm: continue
    cm = CS.get(nm, {}); s = 0.0
    for t in IMP:
        act = sum((row[ci] or 0) for ci in topcols[t] if isinstance(row[ci], (int, float)))
        s += IMP[t] * act * SM.get(cm.get(t, '—'), 0.0)
    legst[nfc(nm)] = round(s, 4)
wb.close()
print('computed stance-gated leg for %d members | range %.2f..%.2f' % (len(legst), min(legst.values()), max(legst.values())))

# ---- backup + COM write to Blended Model col4 + recalc ----
bk = WB.replace('.xlsx', '.PRE-LEGSTANCE-%s.xlsx' % datetime.datetime.now().strftime('%Y%m%d_%H%M%S'))
shutil.copy2(WB, bk); print('backup ->', os.path.basename(bk))
xl = win32.DispatchEx('Excel.Application'); xl.Visible = False; xl.DisplayAlerts = False
try:
    wb2 = xl.Workbooks.Open(WB); ws = wb2.Worksheets('Blended Model')
    n = 0; miss = 0
    for r in range(3, 536):
        nm = ws.Cells(r, 1).Value
        if not nm: continue
        key = nfc(nm)
        if key in legst: ws.Cells(r, 4).Value = legst[key]; n += 1
        else: miss += 1
    print('wrote leg-stance to %d rows (col4) | unmatched: %d' % (n, miss))
    xl.CalculateFull()
    wb2.Save(); wb2.Close(SaveChanges=True); print('SAVED + recalculated (Leg% / Score / Rank updated)')
finally:
    xl.Quit()
