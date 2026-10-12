# -*- coding: utf-8 -*-
"""Deterministic, offline recomputation of the 2026-07-24 legislative audit.

The program has two deliberately separate modes:

* ``--copy-phase`` is the only mode allowed to touch the named LIVE source
  directory.  It copies the fixed input inventory to the task scratch
  directory and writes a hash/size/mtime manifest.
* evidence mode verifies every scratch copy against that manifest and then
  reads only those copies.  It writes the JSON and Markdown evidence pack.

No network access, package installation, workbook save, or production-artifact
write is performed.  The workbook is opened data_only=True and read_only=False,
matching fec/engine/compute_pillars.py:32 and its documented random-access constraint.
"""
from __future__ import print_function

import argparse
import collections
import csv
import datetime
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import statistics
import sys
import tempfile
import unicodedata

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


WORKTREE_ROOT = Path(__file__).resolve().parent.parent.parent
APPROVED_OUT = (WORKTREE_ROOT / "audits" / "legfix2").resolve()
APPROVED_SCRATCH = Path(
    r"C:\Users\pgoff\AppData\Local\Temp\claude\legfix2-scratch"
).resolve()
APPROVED_LIVE = Path(r"C:\Users\pgoff\OneDrive\Desktop\Melange Systems").resolve()

INPUT_FILES = (
    "archive_leg.json",
    "data/legislation/bill_direction_queue_FILLED.csv",
    "data/legislation/bill_directions.csv",
    "data/legislation/bill_overrides.csv",
    "category_stance.json",
    "data/legislation/cfpb_bill_directions.csv",
    "data/engine/congestion.json",
    "congress_bills_archive_20260624_115822.xlsx",
    "congress_cosponsor_network_20260625_205554.xlsx",
    "congress_cosponsor_network_AUG20260719.xlsx",
    "data/legislation/cu_bill_directions.json",
    "data/stance/cu_valence_registry.json",
    "legislation_routing_audit.json",
    "manual_bill_evidence.csv",
    "Moneyball Main Version - Blended.xlsx",
    "data/engine/pillar_inputs_snapshot.json",
    "press_text_cache.jsonl",
    "ranks_python.json",
    "say_do_gap.json",
    "vote_category_lean.json",
)

MANIFEST_NAME = "source_manifest.json"
JSON_NAME = "evidence_recompute.json"
MD_NAME = "EVIDENCE_RECOMPUTE.md"
SCHEMA = "legfix2-evidence-recompute-v1"
MANIFEST_SCHEMA = "legfix2-source-manifest-v1"

REPORT_CITATION = "LEGISLATIVE_AUDIT_2026-07-24.md"
CURRENT_CONGRESS = 119
PERCENT_TOLERANCE_PP = 0.05
RHO_TOLERANCE = 0.005
FLOAT_TOLERANCE = 1e-9
DATA_INVALID = "DATA-INVALID"
E17_MANIFEST_EXCEPTION = (
    "E17 scans the running worktree's code \u2014 the one input outside the manifest"
)
E19C_PERCENT_NEAR_MISS_MAX = 0.5
E19C_RAW_NEAR_MISS_MAX = 0.05
EXCEL_ERROR_STRINGS = frozenset(
    {
        "#DIV/0!",
        "#N/A",
        "#NAME?",
        "#NULL!",
        "#NUM!",
        "#REF!",
        "#VALUE!",
    }
)
_DATA_INVALID_COUNTS = collections.Counter()
_EXPECTED_ABSENT_COUNTS = collections.Counter()

CATEGORIES = (
    "Examination & Supervision",
    "CU Capital Requirements",
    "NCUA Budget & Insurance",
    "Interchange",
    "Payments",
    "Technology & Emerging",
    "Fraud Prevention",
    "Data Privacy & Cyber",
    "Charter Modernization",
    "Defense CU Issues",
    "Field of Membership",
    "CU Tax Status",
    "CDFI Fund",
    "CFPB",
    "Credit Card Mandates",
    "TCPA",
    "Member Business Lending",
    "Mortgage & Housing",
)

LABEL_TO_CANONICAL = {
    "CFPB Issues & Structural Reform": "CFPB",
    "Charter Modernization & Governance": "Charter Modernization",
    "Credit Union Capital Requirements": "CU Capital Requirements",
    "Credit Card & Consumer-Fee Mandates": "Credit Card Mandates",
    "Credit Union Tax Status": "CU Tax Status",
    "Data Privacy & Cybersecurity": "Data Privacy & Cyber",
    "Defense Credit Union Issues": "Defense CU Issues",
    "Mortgage Lending & Housing": "Mortgage & Housing",
    "NCUA Budget & Share Insurance Fund": "NCUA Budget & Insurance",
    "Technology & Emerging Issues": "Technology & Emerging",
    "Telephone Consumer Protection Act Issues": "TCPA",
}

SECTION_KEYWORDS = (
    ("Examination & Supervision", "examination"),
    ("CU Capital Requirements", "capital"),
    ("NCUA Budget & Insurance", "ncua"),
    ("Interchange", "interchange"),
    ("Payments", "payments"),
    ("Technology & Emerging", "technology"),
    ("Fraud Prevention", "fraud"),
    ("Data Privacy & Cyber", "data privacy"),
    ("Charter Modernization", "charter"),
    ("Defense CU Issues", "defense"),
    ("Field of Membership", "field of membership"),
    ("CU Tax Status", "tax status"),
    ("CDFI Fund", "cdfi"),
    ("CFPB", "cfpb"),
    ("Credit Card Mandates", "credit card"),
    ("TCPA", "telephone"),
    ("Member Business Lending", "member business"),
    ("Mortgage & Housing", "mortgage"),
)

DISAPPROVAL_RE = re.compile(r"disapprov(?:al|ing)", re.I)
BILL_RE = re.compile(
    r"\s*([HS])\.?\s*(?:(R|J\.?\s*Res|Con\.?\s*Res|Res)\.?\s*)?(\d+)",
    re.I,
)
AMENDMENT_RE = re.compile(r"^\s*(?:H|S)\.?\s*Amdt\.?\s*\d+", re.I)

# Exact production constants copied from fec/legislation/apply_leg_stance.py:15-23.  That
# module is not imported because it executes workbook mutation at import time.
LEG_PREFIX = {
    "Examination & Supervision": "examination & supervision",
    "CU Capital Requirements": "credit union capital requirements",
    "NCUA Budget & Insurance": "ncua budget",
    "Interchange": "interchange",
    "Payments": "payments",
    "Technology & Emerging": "technology",
    "Fraud Prevention": "fraud prevention",
    "Data Privacy & Cyber": "data privacy",
    "Charter Modernization": "charter modernization",
    "Defense CU Issues": "defense credit union",
    "Field of Membership": "field of membership",
    "CDFI Fund": "cdfi fund",
    "CFPB": "cfpb",
    "Credit Card Mandates": "credit card & consumer-fee",
    "TCPA": "telephone consumer",
    "Member Business Lending": "member business lending",
    "Mortgage & Housing": "mortgage lending",
    "CU Tax Status": "credit union tax status",
}

# Exact Layer-1 mapping from AUDIT_SP/gapb.py:11-19.
L1_TO_CATEGORIES = {
    "Examination and Supervision": (
        "Examination & Supervision",
        "CU Capital Requirements",
        "NCUA Budget & Insurance",
    ),
    "Payments & Technology": (
        "Interchange",
        "Payments",
        "Technology & Emerging",
    ),
    "Information Security": ("Fraud Prevention", "Data Privacy & Cyber"),
    "Credit Union Powers & Charter": (
        "Charter Modernization",
        "Defense CU Issues",
        "Field of Membership",
    ),
    "Tax, Budget & Bank Attacks": ("CU Tax Status", "CDFI Fund"),
    "Consumer Protection": ("CFPB", "Credit Card Mandates", "TCPA"),
    "Lending": ("Member Business Lending", "Mortgage & Housing"),
}

# CACHED CONSTANT ONLY.  These are not represented by a per-letter offline
# input.  Citation and provenance are mandatory in E18's method note.
# AUDIT_SP/gapb.py:8-10.
ACU_L1_COUNTS_20260724 = {
    "Examination and Supervision": 176,
    "Payments & Technology": 92,
    "Consumer Protection": 87,
    "Lending": 85,
    "Tax, Budget & Bank Attacks": 82,
    "Credit Union Powers & Charter": 74,
    "Regulatory Relief": 23,
    "Information Security": 21,
    "Diversity, Equity & Inclusion": 2,
}

REPORTED_E18_MULTIPLIERS = {
    "Information Security": 4.04,
    "Lending": 2.77,
    "Tax, Budget & Bank Attacks": 0.44,
    "Credit Union Powers & Charter": 0.58,
    "Examination and Supervision": 0.60,
}

ADVANCE_WEIGHT = {
    "Introduced": 0.0,
    "Passed House": 1.0,
    "Passed Senate": 1.0,
    "Agreed to in Senate": 1.0,
    "Resolving Differences": 1.5,
    "Failed to pass over veto": 1.5,
    "Became Law": 2.0,
}
SAY_DO_ROLE_WEIGHT = {"sponsor": 1.0, "original": 0.5, "cosponsor": 0.25}

RAW_FIRST_VALUE_COL = 3
RAW_LAST_VALUE_COL = 160
RAW_FIRST_DATA_ROW = 3
RAW_LAST_DATA_ROW = 535
NORMALIZATION_EXCLUDED_COLS = frozenset({144, 146, 148, 150, 151, 152})

EXPECTED_BLOCKS = (
    ("C", "J", "Examination & Supervision"),
    ("K", "R", "CU Capital Requirements"),
    ("S", "Z", "NCUA Budget & Insurance"),
    ("AA", "AH", "Interchange"),
    ("AI", "AP", "Payments"),
    ("AQ", "AX", "Technology & Emerging"),
    ("AY", "BF", "Fraud Prevention"),
    ("BG", "BN", "Data Privacy & Cyber"),
    ("BO", "BV", "Charter Modernization"),
    ("BW", "CD", "Defense CU Issues"),
    ("CE", "CL", "Field of Membership"),
    ("CM", "CT", "CDFI Fund"),
    ("CU", "DB", "CFPB"),
    ("DC", "DJ", "Credit Card Mandates"),
    ("DK", "DR", "TCPA"),
    ("DS", "DZ", "Member Business Lending"),
    ("EA", "EH", "Mortgage & Housing"),
    ("EI", "EM", "institutional/reserved"),
    ("EN", "EN", "reserved"),
    ("EO", "EV", "legacy CFPB Reach"),
    ("EW", "FD", "CU Tax Status"),
)

INTERACTION_SHEETS = (
    "Sponsor Interactions",
    "Original Co-Sponsor Interaction",
    "Co-Sponsor Interactions",
)
SUCCESS_SHEETS = ("Success Sponsor", "Success Original Co-Sponsor")


class EvidenceError(RuntimeError):
    """A fail-closed input, manifest, or structural error."""


def _is_relative_to(path, root):
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


def require_paths(data_root, out):
    data_root = Path(os.path.realpath(os.path.abspath(str(data_root))))
    out = Path(os.path.realpath(os.path.abspath(str(out))))
    physical_worktree = Path(os.path.realpath(str(WORKTREE_ROOT)))
    if out != APPROVED_OUT:
        raise EvidenceError("output must be exactly %s; got %s" % (APPROVED_OUT, out))
    if not _is_relative_to(out, physical_worktree):
        raise EvidenceError("output physically escaped worktree %s" % physical_worktree)
    if not _is_relative_to(data_root, APPROVED_SCRATCH):
        raise EvidenceError(
            "data root must remain under approved scratch %s; got %s"
            % (APPROVED_SCRATCH, data_root)
        )
    return data_root, out


def sha256_file(path):
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
            size += len(block)
    return digest.hexdigest(), size


def _atomic_bytes(path, payload):
    path = Path(os.path.realpath(os.path.abspath(str(path))))
    physical_worktree = Path(os.path.realpath(str(WORKTREE_ROOT)))
    if not _is_relative_to(path, physical_worktree):
        raise EvidenceError("write path physically escaped worktree: %s" % path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(
        prefix="." + path.name + ".", suffix=".tmp", dir=str(path.parent)
    )
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(tmp_name, path)
    finally:
        if os.path.exists(tmp_name):
            os.unlink(tmp_name)


def atomic_text(path, text):
    _atomic_bytes(path, text.encode("utf-8"))


def json_text(value):
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False
    ) + "\n"


def atomic_json(path, value):
    atomic_text(path, json_text(value))


def iso_utc_from_ns(value):
    dt = datetime.datetime.fromtimestamp(value / 1_000_000_000, datetime.timezone.utc)
    return dt.isoformat(timespec="microseconds").replace("+00:00", "Z")


def copy_one_stable(source, destination):
    before = source.stat()
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(
        prefix="." + destination.name + ".",
        suffix=".copying",
        dir=str(destination.parent),
    )
    digest = hashlib.sha256()
    copied_bytes = 0
    try:
        with source.open("rb") as src, os.fdopen(fd, "wb") as dst:
            for block in iter(lambda: src.read(1024 * 1024), b""):
                dst.write(block)
                digest.update(block)
                copied_bytes += len(block)
            dst.flush()
            os.fsync(dst.fileno())
        after = source.stat()
        if (
            before.st_size != after.st_size
            or before.st_mtime_ns != after.st_mtime_ns
            or copied_bytes != before.st_size
        ):
            raise EvidenceError("source changed during copy: %s" % source)
        copied_hash, copied_size = sha256_file(Path(tmp_name))
        streamed_hash = digest.hexdigest()
        if copied_hash != streamed_hash or copied_size != copied_bytes:
            raise EvidenceError("scratch copy verification failed: %s" % source)
        os.replace(tmp_name, destination)
    except Exception:
        if os.path.exists(tmp_name):
            os.unlink(tmp_name)
        raise
    return {
        "bytes": copied_bytes,
        "sha256": digest.hexdigest(),
        "source_mtime": iso_utc_from_ns(before.st_mtime_ns),
        "source_mtime_ns": before.st_mtime_ns,
    }


def run_copy_phase(live_root, data_root, out):
    live_root = live_root.resolve()
    if live_root != APPROVED_LIVE:
        raise EvidenceError(
            "LIVE root must be exactly %s; got %s" % (APPROVED_LIVE, live_root)
        )
    if not live_root.is_dir():
        raise EvidenceError("LIVE root does not exist: %s" % live_root)
    missing = [name for name in INPUT_FILES if not (live_root / name).is_file()]
    if missing:
        raise EvidenceError("BLOCKED: missing LIVE inputs: %s" % ", ".join(missing))
    entries = []
    for name in INPUT_FILES:
        source = (live_root / name).resolve()
        try:
            source.relative_to(live_root)
        except ValueError:
            raise EvidenceError("unexpected source path resolution: %s" % source)
        destination = data_root / name
        observed = copy_one_stable(source, destination)
        entries.append(
            {
                "bytes": observed["bytes"],
                "sha256": observed["sha256"],
                "source_mtime": observed["source_mtime"],
                "source_mtime_ns": observed["source_mtime_ns"],
                "source_path": str(source),
                "scratch_name": name,
            }
        )
        print("COPIED %s (%d bytes)" % (name, observed["bytes"]))
    manifest = {
        "schema": MANIFEST_SCHEMA,
        "input_count": len(entries),
        "inputs": sorted(entries, key=lambda item: item["scratch_name"].casefold()),
    }
    atomic_json(out / MANIFEST_NAME, manifest)
    print("WROTE %s with %d entries" % (out / MANIFEST_NAME, len(entries)))
    return 0


def verified_scratch_paths(data_root, out):
    manifest_path = out / MANIFEST_NAME
    if not manifest_path.is_file():
        raise EvidenceError("manifest missing; run --copy-phase first: %s" % manifest_path)
    manifest = load_json(manifest_path)
    if manifest.get("schema") != MANIFEST_SCHEMA:
        raise EvidenceError("unexpected manifest schema")
    entries = manifest.get("inputs")
    if not isinstance(entries, list) or len(entries) != len(INPUT_FILES):
        raise EvidenceError("manifest must contain exactly %d inputs" % len(INPUT_FILES))
    by_name = {}
    for entry in entries:
        name = entry.get("scratch_name")
        if name in by_name:
            raise EvidenceError("duplicate manifest scratch_name: %s" % name)
        by_name[name] = entry
    if set(by_name) != set(INPUT_FILES):
        missing = sorted(set(INPUT_FILES) - set(by_name))
        extra = sorted(set(by_name) - set(INPUT_FILES))
        raise EvidenceError("manifest inventory mismatch missing=%r extra=%r" % (missing, extra))
    paths = {}
    for name in INPUT_FILES:
        entry = by_name[name]
        path = (data_root / name).resolve()
        try:
            path.relative_to(data_root.resolve())
        except ValueError:
            raise EvidenceError("scratch path escaped data root: %s" % path)
        if not path.is_file():
            raise EvidenceError("scratch input missing: %s" % path)
        observed_hash, observed_bytes = sha256_file(path)
        if observed_hash != entry.get("sha256") or observed_bytes != entry.get("bytes"):
            raise EvidenceError(
                "scratch input differs from manifest: %s (hash=%s bytes=%d)"
                % (name, observed_hash, observed_bytes)
            )
        paths[name] = path
        print("VERIFIED %s" % name)
    return manifest, paths


def load_json(path):
    with path.open("r", encoding="utf-8-sig") as handle:
        return json.load(handle)


def load_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def load_jsonl(path):
    rows = []
    with path.open("r", encoding="utf-8-sig") as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip():
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError:
                rows.append({"_invalid_json_line": line_number})
    return rows


def _poison_number(reason):
    _DATA_INVALID_COUNTS[reason] += 1
    return float("nan")


def _record_expected_absent(reason, count=1):
    _EXPECTED_ABSENT_COUNTS[reason] += count


def is_data_invalid(value):
    return (
        value == DATA_INVALID
        or (
            isinstance(value, float)
            and not math.isfinite(value)
        )
    )


def is_invalid_numeric_input(value):
    if is_data_invalid(value):
        return True
    if not isinstance(value, str):
        return False
    text = value.strip().upper()
    return text in EXCEL_ERROR_STRINGS or text in {
        "NAN",
        "INF",
        "+INF",
        "-INF",
        "INFINITY",
        "+INFINITY",
        "-INFINITY",
    }


def number(value, _legacy_default=None):
    """Return a finite float or a poisoned NaN that can never compare MATCH.

    The second argument is retained only so older internal call sites remain
    source-compatible.  It is deliberately ignored: generic missing, blank,
    non-numeric, Excel-error, NaN, and infinite values are data errors, not
    numeric zero/defaults.  The few production ``or 0.0`` mirrors are explicit
    at their cited call sites instead of flowing through this helper.
    """
    if value is None:
        return _poison_number("none")
    if isinstance(value, bool):
        return _poison_number("boolean")
    if isinstance(value, (int, float)):
        parsed = float(value)
        return parsed if math.isfinite(parsed) else _poison_number("non-finite")
    text = str(value).strip()
    if not text:
        return _poison_number("blank")
    if text.upper() in EXCEL_ERROR_STRINGS:
        return _poison_number("excel-error:%s" % text.upper())
    try:
        parsed = float(text)
    except (TypeError, ValueError):
        return _poison_number("non-numeric")
    return parsed if math.isfinite(parsed) else _poison_number("non-finite")


def cited_none_as_zero(value):
    """Coerce only None at production-semantic call sites that cite the rule."""
    return 0.0 if value is None else number(value)


def is_number(value):
    return (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and math.isfinite(float(value))
    )


def compact_number(value, digits=12):
    if value is None:
        return None
    if isinstance(value, float):
        if not math.isfinite(value):
            return DATA_INVALID
        return round(value, digits)
    return value


def pct(numerator, denominator):
    if is_data_invalid(numerator) or is_data_invalid(denominator):
        return float("nan")
    return 100.0 * numerator / denominator if denominator else None


def pct_delta(recomputed, reported):
    if recomputed is None:
        return None
    if is_data_invalid(recomputed) or is_data_invalid(reported):
        return float("nan")
    return recomputed - reported


def scalar_flag(recomputed, reported, tolerance):
    if recomputed is None:
        return "NOT-RECOMPUTABLE-OFFLINE"
    if is_data_invalid(recomputed) or is_data_invalid(reported):
        return DATA_INVALID
    return "MATCH" if abs(recomputed - reported) <= tolerance else "DRIFT"


def all_flags_verdict(flags, mismatch_word="DRIFT"):
    flags = list(flags)
    if any(flag == DATA_INVALID for flag in flags):
        return DATA_INVALID
    if any(flag == "DRIFT" for flag in flags):
        return mismatch_word
    if any(flag == "NOT-RECOMPUTABLE-OFFLINE" for flag in flags):
        return "NOT-RECOMPUTABLE-OFFLINE"
    return "MATCH" if all(flag == "MATCH" for flag in flags) else mismatch_word


def e16_weight_reconciliation(
    primary_weight,
    selected_total_weight,
    variant_attempts,
    reported_weight=1416.0,
    reported_pct=14.5,
):
    """Keep the p9 selected-only value primary; variants are hypotheses only."""
    primary_pct = pct(primary_weight, selected_total_weight)
    weight_flag = scalar_flag(primary_weight, reported_weight, 0.05)
    pct_flag = scalar_flag(
        primary_pct,
        reported_pct,
        PERCENT_TOLERANCE_PP,
    )
    hypothesis_names = sorted(
        name
        for name, value in variant_attempts.items()
        if not is_data_invalid(value)
        and abs(value - reported_weight) <= 0.05
    )
    hypothesis = hypothesis_names[0] if hypothesis_names else None
    invalid_attempts = sorted(
        name
        for name, value in variant_attempts.items()
        if is_data_invalid(value)
    )
    return {
        "primary_method": "selected_only_absolute_p9_exact",
        "primary_weight": compact_number(primary_weight),
        "primary_weight_pct": compact_number(primary_pct),
        "reported_weight": compact_number(reported_weight),
        "reported_weight_pct": compact_number(reported_pct),
        "weight_flag": weight_flag,
        "weight_pct_flag": pct_flag,
        "status": all_flags_verdict((weight_flag, pct_flag)),
        "reproduction_status": (
            DATA_INVALID
            if DATA_INVALID in (weight_flag, pct_flag)
            else (
                "REPRODUCED"
                if weight_flag == "MATCH" and pct_flag == "MATCH"
                else "NOT-REPRODUCED"
            )
        ),
        "attempts_metadata": {
            "attempts": {
                key: compact_number(value)
                for key, value in sorted(variant_attempts.items())
            },
            "hypothesis_for_reported_scope": hypothesis,
            "hypothesis_weight": (
                compact_number(variant_attempts[hypothesis])
                if hypothesis
                else None
            ),
            "data_invalid_attempts": invalid_attempts,
        },
    }


def e19c_unidentified_verdict(
    best_deviation,
    aggregate_like,
    percent_scale,
):
    """Classify close-but-nonzero candidate residuals as contradictions."""
    fallback = (
        "NOT-RECOMPUTABLE-OFFLINE"
        if aggregate_like
        else "UNIDENTIFIED-RAW-MEASURE"
    )
    if best_deviation is None or is_data_invalid(best_deviation):
        return fallback
    threshold = (
        E19C_PERCENT_NEAR_MISS_MAX
        if percent_scale
        else E19C_RAW_NEAR_MISS_MAX
    )
    if 0.0 < best_deviation <= threshold:
        return "NEAR-MISS-DIVERGENT"
    return fallback


def e19c_part_verdict(data_invalid_columns, contradictions, unresolved_columns):
    if data_invalid_columns:
        return DATA_INVALID
    if contradictions:
        return "MISMATCH"
    if unresolved_columns:
        return "PARTIALLY-IDENTIFIED"
    return "MATCH"


def e19d_classify_findings(
    cp1_observed,
    required_weight_findings,
    weight_mismatches,
    structural_anomalies,
    data_invalid_weight_cells,
):
    """Separate the four known E19d findings from newly observed anomalies."""
    known_weight_cells = frozenset({"P2", "X2", "DG2"})
    unexpected_weight_mismatches = [
        row
        for row in weight_mismatches
        if row.get("raw_cell") not in known_weight_cells
    ]
    unexpected_structural_anomalies = [
        row
        for row in structural_anomalies
        if row.get("finding_key") != "CP1"
    ]
    new_findings = [
        dict({"kind": "row-2-weight-mismatch"}, **row)
        for row in unexpected_weight_mismatches
    ] + [
        dict({"kind": "structural-anomaly"}, **row)
        for row in unexpected_structural_anomalies
    ]
    if data_invalid_weight_cells:
        verdict = DATA_INVALID
    elif new_findings:
        verdict = "NEW-FINDINGS"
    elif cp1_observed and all(
        row["observed_mismatch"] for row in required_weight_findings
    ):
        verdict = "FINDINGS-CONFIRMED"
    else:
        verdict = "MISMATCH"
    return {
        "verdict": verdict,
        "new_findings": new_findings,
        "unexpected_row_2_weight_mismatches": unexpected_weight_mismatches,
        "unexpected_structural_anomalies": unexpected_structural_anomalies,
    }


def e19_overall_verdict(nested_verdicts):
    if any(value == DATA_INVALID for value in nested_verdicts.values()):
        return DATA_INVALID
    if any(
        value in ("MISMATCH", "NEW-FINDINGS")
        for value in nested_verdicts.values()
    ):
        return "MISMATCH"
    if any(
        value in ("PARTIAL", "PARTIALLY-IDENTIFIED")
        for value in nested_verdicts.values()
    ):
        return "PARTIAL"
    if all(
        value in ("MATCH", "FINDINGS-CONFIRMED")
        for value in nested_verdicts.values()
    ):
        return "MATCH"
    return "PARTIAL"


def sanitize_evidence_value(value):
    if is_data_invalid(value):
        return DATA_INVALID
    if isinstance(value, dict):
        return {
            key: sanitize_evidence_value(item)
            for key, item in value.items()
        }
    if isinstance(value, (list, tuple)):
        return [sanitize_evidence_value(item) for item in value]
    return value


def evidence_entry(claim, reported, recomputed, delta, verdict, method_note):
    return sanitize_evidence_value({
        "claim": claim,
        "reported": reported,
        "recomputed": recomputed,
        "delta": delta,
        "verdict": verdict,
        "method_note": method_note,
    })


def norm_text(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def norm_header(value):
    return norm_text(value).casefold().replace(" and ", " & ")


def nfc(value):
    return unicodedata.normalize("NFC", str(value)) if value is not None else value


def canonical_category(value):
    raw = norm_text(value)
    if not raw:
        return None
    if raw in CATEGORIES:
        return raw
    if raw in LABEL_TO_CANONICAL:
        return LABEL_TO_CANONICAL[raw]
    low = raw.casefold()
    for category, keyword in SECTION_KEYWORDS:
        if keyword in low:
            return category
    return None


def canonical_role(value):
    role = norm_text(value).casefold()
    if "original" in role:
        return "original-cosponsor"
    if "cosponsor" in role or "co-spons" in role:
        return "cosponsor"
    if "sponsor" in role:
        return "sponsor"
    return "other"


def bill_key(value):
    match = BILL_RE.match(str(value or ""))
    return (match.group(1).upper(), int(match.group(3))) if match else None


def bill_type_key(value):
    match = BILL_RE.match(str(value or ""))
    if not match:
        return None
    token = re.sub(r"[^A-Za-z]", "", match.group(2) or "R").upper()
    return match.group(1).upper(), token, int(match.group(3))


def congress_text(value):
    try:
        return str(int(float(value)))
    except (TypeError, ValueError):
        return norm_text(value)


def congress_decay(value, current_congress=CURRENT_CONGRESS):
    try:
        source = int(float(value))
    except (TypeError, ValueError):
        raise EvidenceError("numeric Congress required for decay; got %r" % value)
    return round(0.8 ** max(0, int(current_congress) - source), 12)


def is_disapproval(value):
    return bool(DISAPPROVAL_RE.search(str(value or "")))


def workbook_rows(ws):
    headers = [norm_text(cell.value) for cell in ws[1]]
    result = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        result.append(
            {
                headers[index] if headers[index] else "Column%d" % (index + 1): value
                for index, value in enumerate(row[: len(headers)])
            }
        )
    return headers, result


def sheet_structure(ws):
    return {
        "dimension": ws.calculate_dimension(),
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "headers": [
            {"cell": "%s1" % get_column_letter(column), "value": ws.cell(1, column).value}
            for column in range(1, ws.max_column + 1)
        ],
    }


def selected_rows(routing):
    return [
        row
        for row in routing.get("rows", [])
        if row.get("disposition") == "selected"
    ]


def strict_rank(values):
    order = sorted(range(len(values)), key=lambda index: -values[index])
    ranks = [0] * len(values)
    for position, index in enumerate(order):
        ranks[index] = position + 1
    return ranks


def average_ranks(values):
    order = sorted(range(len(values)), key=lambda index: values[index])
    ranks = [0.0] * len(values)
    index = 0
    while index < len(order):
        end = index
        while (
            end + 1 < len(order)
            and values[order[end + 1]] == values[order[index]]
        ):
            end += 1
        rank = (index + end) / 2.0 + 1.0
        for cursor in range(index, end + 1):
            ranks[order[cursor]] = rank
        index = end + 1
    return ranks


def pearson(left, right):
    if len(left) != len(right) or not left:
        return None
    left_mean = sum(left) / len(left)
    right_mean = sum(right) / len(right)
    numerator = sum(
        (left[index] - left_mean) * (right[index] - right_mean)
        for index in range(len(left))
    )
    denominator = math.sqrt(
        sum((value - left_mean) ** 2 for value in left)
        * sum((value - right_mean) ** 2 for value in right)
    )
    return numerator / denominator if denominator else None


def spearman(left, right):
    if any(is_data_invalid(value) for value in list(left) + list(right)):
        return float("nan")
    return pearson(average_ranks(left), average_ranks(right))


def rank_disagreements(members, left, right, limit=10):
    if any(is_data_invalid(value) for value in list(left) + list(right)):
        return []
    left_rank = strict_rank(left)
    right_rank = strict_rank(right)
    order = sorted(
        range(len(members)),
        key=lambda index: -abs(left_rank[index] - right_rank[index]),
    )
    return [
        {
            "member": members[index],
            "left_rank": left_rank[index],
            "right_rank": right_rank[index],
            "abs_rank_delta": abs(left_rank[index] - right_rank[index]),
            "left_value": compact_number(left[index]),
            "right_value": compact_number(right[index]),
        }
        for index in order[:limit]
    ]


def percentrank(values, value):
    # Exact production rule, fec/engine/compute_pillars.py:44-48.
    if is_data_invalid(value) or any(is_data_invalid(item) for item in values):
        return float("nan")
    if len(values) <= 1:
        return 0.0
    below = sum(1 for item in values if item < value)
    return math.floor(below / (len(values) - 1) * 100000.0) / 100000.0


class EvidenceContext:
    def __init__(self, paths):
        self.paths = paths
        self.archive_leg = load_json(paths["archive_leg.json"])
        self.routing = load_json(paths["legislation_routing_audit.json"])
        self.selected = selected_rows(self.routing)
        self.snapshot = load_json(paths["data/engine/pillar_inputs_snapshot.json"])
        self.ranks = load_json(paths["ranks_python.json"])
        self.manual = load_csv(paths["manual_bill_evidence.csv"])
        self.bill_directions = load_csv(paths["data/legislation/bill_directions.csv"])
        self.cfpb_directions = load_csv(paths["data/legislation/cfpb_bill_directions.csv"])
        self.bill_overrides = load_csv(paths["data/legislation/bill_overrides.csv"])
        self.direction_queue = load_csv(paths["data/legislation/bill_direction_queue_FILLED.csv"])
        self.cu_directions = load_json(paths["data/legislation/cu_bill_directions.json"])
        self.valence_registry = load_json(paths["data/stance/cu_valence_registry.json"])
        self.press_cache = load_jsonl(paths["press_text_cache.jsonl"])
        self.say_do = load_json(paths["say_do_gap.json"])
        self.vote_lean = load_json(paths["vote_category_lean.json"])
        self.category_stance = load_json(paths["category_stance.json"])
        self.congestion = load_json(paths["data/engine/congestion.json"])
        self.bill_wb = openpyxl.load_workbook(
            paths["congress_bills_archive_20260624_115822.xlsx"],
            data_only=True,
            read_only=False,
        )
        self.model_wb = openpyxl.load_workbook(
            paths["Moneyball Main Version - Blended.xlsx"],
            data_only=True,
            read_only=False,
        )
        self.network_wb = openpyxl.load_workbook(
            paths["congress_cosponsor_network_AUG20260719.xlsx"],
            data_only=True,
            read_only=False,
        )
        self.base_network_wb = openpyxl.load_workbook(
            paths["congress_cosponsor_network_20260625_205554.xlsx"],
            data_only=True,
            read_only=False,
        )
        if "All Bills" not in self.bill_wb.sheetnames:
            raise EvidenceError("bill archive workbook lacks All Bills")
        self.bill_headers, self.bill_rows = workbook_rows(self.bill_wb["All Bills"])
        self.roster = [
            norm_text(row.get("member"))
            for row in self.congestion
            if norm_text(row.get("member"))
        ]
        self.roster_by_casefold = {member.casefold(): member for member in self.roster}

    def resolve_member(self, value):
        text = nfc(norm_text(value))
        return self.roster_by_casefold.get(text.casefold(), text if text in self.roster else None)

    def close(self):
        self.bill_wb.close()
        self.model_wb.close()
        self.network_wb.close()
        self.base_network_wb.close()


def claim_e1(ctx):
    totals = {
        member: sum(number(value) for value in categories.values())
        for member, categories in ctx.archive_leg.items()
    }
    o_values = ctx.snapshot.get("O", {})
    common = [member for member in o_values if member in totals]
    # AUDIT_SP/p14.py:25 explicitly coerces a missing O value with
    # ``float(O[m] or 0)``.  Preserve None->0 but poison corrupt non-finite,
    # blank, and Excel-error values instead of extending that fallback.
    left = [cited_none_as_zero(o_values[member]) for member in common]
    right = [totals[member] for member in common]
    rho = spearman(left, right)
    tail = rank_disagreements(common, left, right, 10)
    merkley = next((row for row in tail if row["member"] == "Jeff Merkley"), None)
    rho_flag = scalar_flag(rho, 0.912, RHO_TOLERANCE)
    merkley_flag = (
        "MATCH"
        if merkley
        and merkley["left_rank"] == 471
        and merkley["right_rank"] == 41
        else "DRIFT"
    )
    return evidence_entry(
        "Spearman rho between pillar snapshot O and archive_leg signed totals, plus the ten largest strict-rank tail disagreements",
        {
            "rho": 0.912,
            "tail_rows": 10,
            "required_tail_row": {
                "member": "Jeff Merkley",
                "O_rank": 471,
                "archive_rank": 41,
            },
            "citation": REPORT_CITATION + " §Q3",
        },
        {
            "common_members": len(common),
            "rho": compact_number(rho),
            "tail": tail,
            "merkley_in_top_10": merkley is not None,
        },
        {
            "rho": compact_number(None if rho is None else rho - 0.912),
            "rho_flag": rho_flag,
            "merkley_flag": merkley_flag,
        },
        all_flags_verdict((rho_flag, merkley_flag)),
        "Mirrors AUDIT_SP/p14.py:6-42 exactly: O insertion order, None->0.0, average-rank ties for Spearman/Pearson, strict descending position ranks for the tail. Archive values are signed category sums. Driver semantics are documented by AUDIT_SP/p15.py:3-19.",
    )


def parse_direction(value):
    text = norm_text(value)
    if not text:
        return None
    try:
        parsed = float(text)
    except (TypeError, ValueError):
        return 0
    return int(parsed) if parsed in (-1.0, 0.0, 1.0) else 0


def override_index(rows):
    result = {}
    for row in rows:
        key = bill_key(row.get("Bill"))
        congress = congress_text(row.get("Congress")) or "*"
        if key:
            result[(key, congress)] = row
    return result


def find_override(overrides, bill, congress):
    key = bill_key(bill)
    congress_key = congress_text(congress)
    if not key:
        return None
    return overrides.get((key, congress_key)) or overrides.get((key, "*"))


def selected_sign_source_rows(ctx):
    """Classify selected events by the exact build_archive sign precedence."""
    bdir = {}
    for row in ctx.bill_directions:
        key = bill_key(row.get("Bill"))
        if key:
            bdir[key] = norm_text(row.get("Direction"))
    overrides = override_index(ctx.bill_overrides)
    manual_direction = {
        (norm_text(row.get("Member")), norm_text(row.get("Bill")).casefold()): row.get(
            "Direction"
        )
        for row in ctx.manual
        if norm_text(row.get("Direction"))
    }
    classified = []
    for row in ctx.selected:
        key = bill_key(row.get("bill"))
        override = find_override(overrides, row.get("bill"), row.get("congress"))
        manual_key = (row.get("member"), norm_text(row.get("bill")).casefold())
        if override and norm_text(override.get("Direction")):
            source = "1 bill_overrides.csv Direction"
        elif manual_key in manual_direction:
            source = "2 manual ManualDirection"
        elif key in bdir:
            source = "3 bill_directions.csv registry"
        elif row.get("effective_category") == "CFPB":
            source = "4 CFPB code default (+1)"
        elif row.get("effective_category") == "Credit Card Mandates":
            source = "5 CreditCard title-keyword default"
        else:
            source = "6 CODE DEFAULT: -1 if NEG else +1"
        classified.append((row, source))
    return classified, bdir


def claim_e2(ctx):
    classified, bdir = selected_sign_source_rows(ctx)
    source_counts = collections.Counter()
    source_weight = collections.Counter()
    for row, source in classified:
        source_counts[source] += 1
        source_weight[source] += abs(number(row.get("contribution")))
    total_weight = sum(source_weight.values())
    blanket_default_weight = source_weight[
        "6 CODE DEFAULT: -1 if NEG else +1"
    ]
    code_signed_total_weight = sum(
        source_weight[key]
        for key in (
            "4 CFPB code default (+1)",
            "5 CreditCard title-keyword default",
            "6 CODE DEFAULT: -1 if NEG else +1",
        )
    )
    scored_bill_keys = {
        bill_key(row.get("bill")) for row in ctx.selected if bill_key(row.get("bill"))
    }
    adjudicated = len(scored_bill_keys & set(bdir))
    blanket_default_share = pct(blanket_default_weight, total_weight)
    code_signed_total_share = pct(code_signed_total_weight, total_weight)
    adjudicated_share = pct(adjudicated, len(scored_bill_keys))
    flags = (
        scalar_flag(blanket_default_share, 82.0, PERCENT_TOLERANCE_PP),
        "MATCH" if adjudicated == 54 else "DRIFT",
        "MATCH" if len(scored_bill_keys) == 436 else "DRIFT",
        scalar_flag(adjudicated_share, 12.4, PERCENT_TOLERANCE_PP),
    )
    return evidence_entry(
        "Share of stance weight signed by category defaults and coverage of adjudicated bill directions",
        {
            "blanket_default_pct": 82.0,
            "adjudicated_bills": 54,
            "scored_bills": 436,
            "adjudicated_pct": 12.4,
            "citation": REPORT_CITATION + " §Q3",
        },
        {
            "blanket_default_abs_weight": compact_number(blanket_default_weight),
            "blanket_default_pct": compact_number(blanket_default_share),
            "code_signed_total_abs_weight": compact_number(
                code_signed_total_weight
            ),
            "code_signed_total_pct": compact_number(code_signed_total_share),
            "all_selected_abs_weight": compact_number(total_weight),
            "adjudicated_bills": adjudicated,
            "scored_bills": len(scored_bill_keys),
            "adjudicated_pct": compact_number(adjudicated_share),
            "sign_source_rows": dict(sorted(source_counts.items())),
            "sign_source_abs_weight": {
                key: compact_number(value)
                for key, value in sorted(source_weight.items())
            },
            "cfpb_direction_rows": len(ctx.cfpb_directions),
            "valence_registry_entries": len(
                ctx.valence_registry.get("entries", ctx.valence_registry)
            ),
        },
        {
            "blanket_default_pp": compact_number(
                pct_delta(blanket_default_share, 82.0)
            ),
            "code_signed_total_pp_vs_reported_82_supplementary": compact_number(
                pct_delta(code_signed_total_share, 82.0)
            ),
            "adjudicated_bills": adjudicated - 54,
            "scored_bills": len(scored_bill_keys) - 436,
            "adjudicated_pp": compact_number(
                pct_delta(adjudicated_share, 12.4)
            ),
            "flags": list(flags),
        },
        all_flags_verdict(flags),
        "Sign-source precedence mirrors fec/legislation/build_archive_legislative.py:37-43,76-109 and AUDIT_SP/p16.py:23-47. The reported 82.0% is compared only to bucket 6, the blanket category default. A supplementary code-signed-total share also includes bucket 5's Credit Card title-keyword rule (and bucket 4 if present). Absolute selected-event contribution is the denominator.",
    )


def claim_e3(ctx):
    triples = collections.defaultdict(set)
    triple_rows = collections.defaultdict(list)
    for row in ctx.selected:
        key = (
            row["member"],
            norm_text(row.get("bill")).casefold(),
            congress_text(row.get("congress")),
        )
        triples[key].add(row["effective_category"])
        triple_rows[key].append(row)
    multi = {key: values for key, values in triples.items() if len(values) > 1}
    total_weight = sum(abs(number(row.get("contribution"))) for row in ctx.selected)
    extra_weight = 0.0
    for key in multi:
        rows = sorted(
            triple_rows[key],
            key=lambda row: -abs(number(row.get("contribution"))),
        )
        extra_weight += sum(abs(number(row.get("contribution"))) for row in rows[1:])
    multi_share = pct(len(multi), len(triples))
    extra_share = pct(extra_weight, total_weight)
    flags = (
        "MATCH" if len(multi) == 1463 else "DRIFT",
        "MATCH" if len(triples) == 5466 else "DRIFT",
        scalar_flag(multi_share, 26.8, PERCENT_TOLERANCE_PP),
        scalar_flag(extra_share, 28.9, PERCENT_TOLERANCE_PP),
    )
    return evidence_entry(
        "Cross-listed member/bill/Congress triples and the extra absolute weight carried by secondary categories",
        {
            "cross_listed_triples": 1463,
            "all_triples": 5466,
            "cross_listed_pct": 26.8,
            "extra_raw_weight_pct": 28.9,
            "citation": REPORT_CITATION + " §Q5",
        },
        {
            "cross_listed_triples": len(multi),
            "all_triples": len(triples),
            "cross_listed_pct": compact_number(multi_share),
            "extra_raw_weight": compact_number(extra_weight),
            "all_raw_weight": compact_number(total_weight),
            "extra_raw_weight_pct": compact_number(extra_share),
        },
        {
            "cross_listed_triples": len(multi) - 1463,
            "all_triples": len(triples) - 5466,
            "cross_listed_pp": compact_number(pct_delta(multi_share, 26.8)),
            "extra_raw_weight_pp": compact_number(pct_delta(extra_share, 28.9)),
            "flags": list(flags),
        },
        all_flags_verdict(flags),
        "Mirrors AUDIT_SP/p8.py:5-20 and production's event key in fec/legislation/build_archive_legislative.py:111-132. For each multi-category triple, the largest absolute contribution is retained and all remaining category contributions are counted as extra.",
    )


def claim_e4(ctx):
    records = collections.defaultdict(list)
    for row in ctx.bill_rows:
        records[(norm_text(row.get("Name")), congress_text(row.get("Congress")))].append(
            row
        )
    reported_total = 0
    captured_total = 0
    inconsistent_records = []
    for key, rows in records.items():
        counts = set()
        for row in rows:
            count = number(row.get("CosponsorCount"))
            if not is_data_invalid(count):
                counts.add(int(count))
        if len(counts) > 1:
            inconsistent_records.append({"bill": key[0], "congress": key[1], "counts": sorted(counts)})
        reported_total += max(counts) if counts else 0
        captured_total += len(
            {
                norm_text(row.get("Member"))
                for row in rows
                if canonical_role(row.get("Role")) in ("original-cosponsor", "cosponsor")
                and norm_text(row.get("Member"))
            }
        )
    capture_share = pct(captured_total, reported_total)
    flag = scalar_flag(capture_share, 83.2, PERCENT_TOLERANCE_PP)
    return evidence_entry(
        "Archive member rows capture the stated CosponsorCount reported by their own bill records",
        {
            "captured_pct": 83.2,
            "citation": REPORT_CITATION + " §Q2",
        },
        {
            "captured_cosponsor_pairs": captured_total,
            "reported_cosponsors": reported_total,
            "captured_pct": compact_number(capture_share),
            "bill_records": len(records),
            "records_with_inconsistent_CosponsorCount": inconsistent_records,
        },
        {
            "captured_pp": compact_number(pct_delta(capture_share, 83.2)),
            "flag": flag,
        },
        flag,
        "The hand-produced bill archive workbook is the sole specification for this claim. Rows are grouped by (Name,Congress); CosponsorCount is counted once per record and distinct original/plain cosponsor members are the captured numerator.",
    )


def corrected_act_title(value):
    text = unicodedata.normalize("NFKC", norm_text(value)).casefold()
    text = text.translate(
        str.maketrans(
            {
                "\u2018": "'",
                "\u2019": "'",
                "\u201b": "'",
                "\u02bc": "'",
                "\u2032": "'",
                "\uff07": "'",
            }
        )
    )
    text = "".join(
        " " if unicodedata.category(character).startswith("P") else character
        for character in text
    )
    return norm_text(text)


def act_category_disagreement(rows, title_normalizer, exclude_blank=False):
    acts = collections.defaultdict(dict)
    row_count_by_title = collections.Counter()
    blank_title_rows = 0
    for row in rows:
        title = title_normalizer(row.get("Title"))
        if not title and exclude_blank:
            blank_title_rows += 1
            continue
        instance = norm_text(row.get("Name")) + "/C" + congress_text(row.get("Congress"))
        category = canonical_category(row.get("Section"))
        acts[title].setdefault(instance, set()).add(category)
        row_count_by_title[title] += 1
    multi = {title: instances for title, instances in acts.items() if len(instances) > 1}
    inconsistent = {}
    for title, instances in multi.items():
        category_sets = list(instances.values())
        if any(value != category_sets[0] for value in category_sets[1:]):
            inconsistent[title] = instances
    rows_involved = sum(row_count_by_title[title] for title in inconsistent)
    act_share = pct(len(inconsistent), len(multi))
    row_share = pct(rows_involved, len(rows))
    examples = []
    for title in sorted(
        inconsistent, key=lambda value: (-row_count_by_title[value], value)
    )[:10]:
        examples.append(
            {
                "title": title,
                "rows": row_count_by_title[title],
                "instances": {
                    instance: sorted("" if item is None else item for item in categories)
                    for instance, categories in sorted(inconsistent[title].items())
                },
            }
        )
    return {
        "disagreeing_acts": len(inconsistent),
        "multi_number_acts": len(multi),
        "disagreeing_pct": compact_number(act_share),
        "rows_involved": rows_involved,
        "all_rows": len(rows),
        "rows_involved_pct": compact_number(row_share),
        "blank_title_rows_excluded": blank_title_rows,
        "worst_examples": examples,
    }


def claim_e5(ctx):
    audit_method = act_category_disagreement(
        ctx.bill_rows,
        lambda value: norm_text(value).casefold(),
    )
    corrected_method = act_category_disagreement(
        ctx.bill_rows,
        corrected_act_title,
        exclude_blank=True,
    )
    h_r_3323_titles = sorted(
        {
            norm_text(row.get("Title"))
            for row in ctx.bill_rows
            if norm_text(row.get("Name")).upper().replace(" ", "") == "H.R.3323"
        }
    )
    flags = (
        "MATCH" if audit_method["disagreeing_acts"] == 13 else "DRIFT",
        "MATCH" if audit_method["multi_number_acts"] == 112 else "DRIFT",
        scalar_flag(
            audit_method["disagreeing_pct"],
            11.6,
            PERCENT_TOLERANCE_PP,
        ),
        "MATCH" if audit_method["rows_involved"] == 912 else "DRIFT",
        scalar_flag(
            audit_method["rows_involved_pct"],
            8.3,
            PERCENT_TOLERANCE_PP,
        ),
    )
    return evidence_entry(
        "Multi-number acts whose bill instances disagree on canonical category sets",
        {
            "disagreeing_acts": 13,
            "multi_number_acts": 112,
            "disagreeing_pct": 11.6,
            "rows_involved": 912,
            "rows_involved_pct": 8.3,
            "citation": REPORT_CITATION + " §Q5",
        },
        {
            "audit_method": audit_method,
            "corrected_normalization": corrected_method,
            "normalization_artifact": {
                "bill": "H.R.3323",
                "titles_as_stored": h_r_3323_titles,
                "cause": "curly-versus-straight apostrophe split under the audit method",
            },
        },
        {
            "audit_method": {
                "disagreeing_acts": audit_method["disagreeing_acts"] - 13,
                "multi_number_acts": audit_method["multi_number_acts"] - 112,
                "disagreeing_pp": compact_number(
                    pct_delta(audit_method["disagreeing_pct"], 11.6)
                ),
                "rows_involved": audit_method["rows_involved"] - 912,
                "rows_involved_pp": compact_number(
                    pct_delta(audit_method["rows_involved_pct"], 8.3)
                ),
            },
            "corrected_minus_audit": {
                "disagreeing_acts": corrected_method["disagreeing_acts"]
                - audit_method["disagreeing_acts"],
                "multi_number_acts": corrected_method["multi_number_acts"]
                - audit_method["multi_number_acts"],
                "rows_involved": corrected_method["rows_involved"]
                - audit_method["rows_involved"],
            },
            "flags": list(flags),
        },
        all_flags_verdict(flags),
        "The flag mirrors AUDIT_SP/act_category.py:10-35 exactly: whitespace-normalized lowercase Title is the act key; Name/Congress is the instance; section labels use fec/bill_rules.py canonicalization. A second, non-flagged correction applies Unicode NFKC, folds apostrophe variants, strips punctuation, and excludes blank titles; it removes the manufactured H.R.3323 disagreement (curly versus straight apostrophe) and is expected to yield 12/112 acts and 899 rows.",
    )


def claim_e6(ctx):
    by_congress = collections.Counter()
    for row in ctx.selected:
        by_congress[congress_text(row.get("congress"))] += abs(
            number(row.get("contribution"))
        )
    total = sum(by_congress.values())
    share = pct(by_congress["119"], total)
    flag = scalar_flag(share, 41.7, PERCENT_TOLERANCE_PP)
    return evidence_entry(
        "119th Congress share of selected absolute legislative weight",
        {"weight_pct": 41.7, "citation": REPORT_CITATION + " §Q3"},
        {
            "weight_pct": compact_number(share),
            "weight": compact_number(by_congress["119"]),
            "all_weight": compact_number(total),
            "weight_by_congress": {
                key: compact_number(value) for key, value in sorted(by_congress.items())
            },
        },
        {"weight_pp": compact_number(pct_delta(share, 41.7)), "flag": flag},
        flag,
        "Mirrors AUDIT_SP/p3.py:8-20 on selected routing rows. Stored decay/contribution semantics originate at fec/legislation/build_archive_legislative.py:72-110 and fec/bill_rules.py:185-203.",
    )


def raw_header_columns(ws):
    result = {}
    for column in range(1, ws.max_column + 1):
        header = ws.cell(1, column).value
        if not header:
            continue
        normalized = norm_header(header)
        for category, prefix in LEG_PREFIX.items():
            if normalized.startswith(norm_header(prefix)):
                result.setdefault(category, []).append(column)
                break
    return result


def claim_e7(ctx):
    raw = ctx.model_wb["Legislative Score - Raw"]
    generic_success_columns = []
    for start, _end, classification in EXPECTED_BLOCKS:
        if classification in CATEGORIES:
            generic_success_columns.append(
                (classification, column_index_from_string(start) + 5)
            )
    all_success_labeled_columns = []
    for column in range(1, raw.max_column + 1):
        header = norm_header(raw.cell(1, column).value)
        if "success" in header:
            all_success_labeled_columns.append(column)

    generic_total_cells = 0
    generic_empty_cells = 0
    generic_blank_cells = 0
    generic_per_category = {}
    generic_column_numbers = {column for _category, column in generic_success_columns}
    for category, column in generic_success_columns:
        header = norm_text(raw.cell(1, column).value)
        values = [
            raw.cell(row, column).value
            for row in range(RAW_FIRST_DATA_ROW, RAW_LAST_DATA_ROW + 1)
        ]
        zero = sum(1 for value in values if value is None or number(value) == 0.0)
        blank = sum(value is None for value in values)
        generic_total_cells += len(values)
        generic_empty_cells += zero
        generic_blank_cells += blank
        generic_per_category[category] = {
            "column": get_column_letter(column),
            "header": header,
            "zero_or_blank": zero,
            "blank": blank,
            "cells": len(values),
            "constant_zero": zero == len(values),
        }
    generic_empty_share = pct(generic_empty_cells, generic_total_cells)
    constant_zero_generic = sorted(
        category
        for category, detail in generic_per_category.items()
        if detail["constant_zero"]
    )
    additional_columns = []
    all_success_empty_cells = 0
    all_success_blank_cells = 0
    for column in all_success_labeled_columns:
        values = [
            raw.cell(row, column).value
            for row in range(RAW_FIRST_DATA_ROW, RAW_LAST_DATA_ROW + 1)
        ]
        zero = sum(1 for value in values if value is None or number(value) == 0.0)
        blank = sum(value is None for value in values)
        all_success_empty_cells += zero
        all_success_blank_cells += blank
        if column not in generic_column_numbers and zero == len(values):
            additional_columns.append(
                {
                    "column": get_column_letter(column),
                    "header": raw.cell(1, column).value,
                    "zero_or_blank": zero,
                    "cells": len(values),
                }
            )
    reported_zero = sorted(
        (
            "Interchange",
            "Payments",
            "Field of Membership",
            "CDFI Fund",
            "CFPB",
            "Credit Card Mandates",
            "TCPA",
        )
    )
    flags = (
        scalar_flag(generic_empty_share, 90.7, PERCENT_TOLERANCE_PP),
        "MATCH" if constant_zero_generic == reported_zero else "DRIFT",
    )
    return evidence_entry(
        "Efficiency/success-rate slice sparsity and categories whose success rate is constant zero",
        {
            "empty_pct": 90.7,
            "constant_zero_categories": reported_zero,
            "citation": REPORT_CITATION + " §Q3",
        },
        {
            "generic_scope_empty_pct": compact_number(generic_empty_share),
            "generic_scope_cell_count": generic_total_cells,
            "generic_scope_empty_cells": generic_empty_cells,
            "generic_scope_column_count": len(generic_success_columns),
            "constant_zero_generic_success_categories": constant_zero_generic,
            "constant_zero_additional_columns": additional_columns,
            "audit_scope_generic_success_rate": {
                "columns": len(generic_success_columns),
                "cells": generic_total_cells,
                "empty_cells": generic_empty_cells,
                "blank_cells": generic_blank_cells,
                "empty_definition": "blank or numeric zero",
                "empty_pct": compact_number(generic_empty_share),
                "constant_zero_generic_success_categories": constant_zero_generic,
                "per_category": generic_per_category,
            },
            "supplementary_all_success_labeled_columns": {
                "columns": len(all_success_labeled_columns),
                "cells": len(all_success_labeled_columns)
                * (RAW_LAST_DATA_ROW - RAW_FIRST_DATA_ROW + 1),
                "empty_cells": all_success_empty_cells,
                "blank_cells": all_success_blank_cells,
                "empty_definition": "blank or numeric zero",
                "empty_pct": compact_number(
                    pct(
                        all_success_empty_cells,
                        len(all_success_labeled_columns)
                        * (RAW_LAST_DATA_ROW - RAW_FIRST_DATA_ROW + 1),
                    )
                ),
                "constant_zero_additional_columns": additional_columns,
            },
            "dependency_path": "Blended Model G -> K -> O; user dials AB10/AB11",
        },
        {
            "generic_empty_pp": compact_number(
                pct_delta(generic_empty_share, 90.7)
            ),
            "constant_zero_symmetric_difference": sorted(
                set(constant_zero_generic) ^ set(reported_zero)
            ),
            "flags": list(flags),
        },
        all_flags_verdict(flags),
        "The audit comparison uses exactly one generic Success Rate column per 18 category blocks (18x533 cells). All header-labeled success columns are separately reported as a broader supplementary scope. Constant-zero generic categories and additional role-specific columns are kept in parallel typed lists. The workbook is its own formula specification; apply_leg_stance.py:25-50 documents the activity-to-pillar gate.",
    )


def claim_e8(ctx):
    acu_rows = [
        row
        for row in ctx.press_cache
        if "americascreditunions.org" in str(row.get("url") or "")
    ]
    marker_count = 0
    lengths = []
    invalid = 0
    for row in acu_rows:
        if "_invalid_json_line" in row:
            invalid += 1
            continue
        text = str(row.get("text") or "")
        position = text.find("LOG IN FOR ACCESS")
        if position >= 0:
            marker_count += 1
            lengths.append(position)
        else:
            lengths.append(len(text))
    lengths.sort()
    median = lengths[len(lengths) // 2] if lengths else None
    share = pct(marker_count, len(acu_rows))
    flags = (
        scalar_flag(share, 97.7, PERCENT_TOLERANCE_PP),
        "MATCH" if median == 1927 else "DRIFT",
    )
    return evidence_entry(
        "ACU press cache paywall truncation and median characters before the cut",
        {
            "paywall_truncated_pct": 97.7,
            "median_chars": 1927,
            "citation": REPORT_CITATION + " §Q1",
        },
        {
            "acu_rows": len(acu_rows),
            "marker_rows": marker_count,
            "paywall_truncated_pct": compact_number(share),
            "median_chars": median,
            "invalid_json_lines_in_acu_subset": invalid,
            "definitional_sensitivity": {
                "median_1927_scope": (
                    "reproduces only under the audit's verify3.py definition"
                ),
                "independent_variant_range": [1894, 1940.5],
                "floor_claim": "definition-robust",
            },
        },
        {
            "paywall_truncated_pp": compact_number(pct_delta(share, 97.7)),
            "median_chars": None if median is None else median - 1927,
            "flags": list(flags),
        },
        all_flags_verdict(flags),
        "Mirrors AUDIT_SP/verify3.py:10-22: ACU-domain rows only; marker index when LOG IN FOR ACCESS occurs, otherwise full text length; sorted middle element is the median. Definitional sensitivity: the 1,927 median reproduces only under this verify3.py definition; independent variants span 1,894-1,940.5. The floor claim is definition-robust.",
    )


def claim_e9(ctx):
    roles = collections.Counter(norm_text(row.get("Role")) for row in ctx.bill_rows)
    records = {
        (norm_text(row.get("Name")), congress_text(row.get("Congress")))
        for row in ctx.bill_rows
    }
    amendments = {
        key for key in records if AMENDMENT_RE.match(key[0])
    }
    flags = (
        "MATCH" if len(roles) == 3 else "DRIFT",
        "MATCH" if len(ctx.bill_rows) == 10986 else "DRIFT",
        "MATCH" if len(records) == 468 else "DRIFT",
        "MATCH" if len(amendments) == 0 else "DRIFT",
    )
    return evidence_entry(
        "Role-column cardinality and amendment coverage in the bill archive",
        {
            "distinct_roles": 3,
            "rows": 10986,
            "bill_records": 468,
            "amendment_records": 0,
            "citation": REPORT_CITATION + " §Q1",
        },
        {
            "distinct_roles": len(roles),
            "role_counts": dict(sorted(roles.items())),
            "rows": len(ctx.bill_rows),
            "bill_records": len(records),
            "amendment_records": len(amendments),
            "amendments": [list(value) for value in sorted(amendments)],
        },
        {
            "distinct_roles": len(roles) - 3,
            "rows": len(ctx.bill_rows) - 10986,
            "bill_records": len(records) - 468,
            "amendment_records": len(amendments),
            "flags": list(flags),
        },
        all_flags_verdict(flags),
        "All Bills is enumerated directly. A record is (Name,Congress); amendments are formal H.Amdt/S.Amdt identifiers, not inferred from titles.",
    )


def activity_matrix(ctx):
    raw = ctx.model_wb["Legislative Score - Raw"]
    headers = [raw.cell(1, column).value for column in range(1, raw.max_column + 1)]
    columns = {category: [] for category in CATEGORIES}
    for index, header in enumerate(headers, 1):
        normalized = norm_header(header)
        if "average" in normalized or "success" in normalized or "sponsor" not in normalized:
            continue
        for category, prefix in LEG_PREFIX.items():
            if normalized.startswith(norm_header(prefix)):
                columns[category].append(index)
                break
    activity = {}
    for row in range(2, raw.max_row + 1):
        member = raw.cell(row, 1).value
        if not member:
            continue
        activity[nfc(member)] = {
            category: sum(
                number(raw.cell(row, column).value)
                for column in category_columns
                if is_number(raw.cell(row, column).value)
            )
            for category, category_columns in columns.items()
        }
    return activity, columns


def claim_e10(ctx):
    activity, columns = activity_matrix(ctx)
    name_of = ctx.vote_lean.get("name_of", {})
    name_to_bioguide = {name: bioguide for bioguide, name in name_of.items()}
    no_archive = []
    for member in ctx.roster:
        member_activity = activity.get(nfc(member), {})
        bioguide = name_to_bioguide.get(member, "")
        for category, detail in ctx.vote_lean.get("votes", {}).get(bioguide, {}).items():
            vote_n = number(detail.get("n"))
            if vote_n < 1:
                continue
            archive_categories = ctx.archive_leg.get(member, {})
            # Absence is a structural "no archive_leg cell" predicate, not a
            # numeric None->0 coercion.  A present invalid value is poisoned.
            if (
                category in archive_categories
                and abs(number(archive_categories[category])) > FLOAT_TOLERANCE
            ):
                continue
            act = number(member_activity.get(category))
            no_archive.append(
                {
                    "member": member,
                    "category": category,
                    "vote_n": int(vote_n),
                    "activity": compact_number(act),
                    "act_gated_contribution": 0.0 if act == 0.0 else None,
                }
            )
    audit_gate = [row for row in no_archive if row["vote_n"] >= 5]
    audit_gate_zero = [row for row in audit_gate if row["activity"] == 0.0]
    broad_true_vote_only = [row for row in no_archive if row["activity"] == 0.0]
    workbook_archive_residue = [row for row in no_archive if row["activity"] != 0.0]
    flags = (
        "MATCH" if len(audit_gate) == 148 else "DRIFT",
        "MATCH" if len(audit_gate_zero) == 148 else "DRIFT",
        "MATCH"
        if all(row["act_gated_contribution"] == 0.0 for row in broad_true_vote_only)
        else "DRIFT",
        "MATCH" if len(workbook_archive_residue) == 112 else "DRIFT",
    )
    return evidence_entry(
        "Vote-only member/category cells contribute zero through the production activity gate",
        {
            "audit_gate_n5_no_archive_cells": 148,
            "audit_gate_zero_contribution_cells": 148,
            "workbook_activity_residue_cells": 112,
            "citation": REPORT_CITATION + " §Q1",
        },
        {
            "tier_a_audit_definition": {
                "definition": "vote n>=5 and no archive_leg member/category value",
                "cells": len(audit_gate),
                "zero_activity_cells": len(audit_gate_zero),
                "zero_contribution_cells": len(audit_gate_zero),
                "nonzero_activity_examples": [
                    row for row in audit_gate if row["activity"] != 0.0
                ][:5],
            },
            "tier_b_broad_true_vote_only": {
                "definition": "vote n>=1, no archive_leg value, and zero workbook activity across sponsor/original/cosponsor columns",
                "cells": len(broad_true_vote_only),
                "all_contribute_exactly_zero": all(
                    row["act_gated_contribution"] == 0.0
                    for row in broad_true_vote_only
                ),
            },
            "tier_c_workbook_archive_desync": {
                "definition": "vote n>=1 and no archive_leg value, but nonzero workbook activity",
                "cells": len(workbook_archive_residue),
                "expected_by_design": True,
                "reason": "workbook/archive_leg exclusions and deduplication use different scopes",
                "examples": workbook_archive_residue[:5],
            },
            "broad_n1_no_archive_cells": len(no_archive),
            "activity_columns": {
                category: [get_column_letter(column) for column in value]
                for category, value in sorted(columns.items())
            },
        },
        {
            "audit_gate_cells": len(audit_gate) - 148,
            "audit_gate_zero_contribution_cells": len(audit_gate_zero) - 148,
            "workbook_activity_residue_cells": len(workbook_archive_residue) - 112,
            "flags": list(flags),
        },
        all_flags_verdict(flags),
        "Tier A mirrors vote_by_category's n>=5 gate and the no-archive test in AUDIT_SP/actcheck.py:31-48. Tier B broadens to n>=1 but requires ACT==0; fec/legislation/apply_leg_stance.py:25-50 then proves exact zero. Tier C isolates the nonzero-workbook/no-archive residue as an expected scope desynchronization caused by exclusions and deduplication, not a vote-only contribution.",
    )


def numbered_sheet_prefixes(sheetnames):
    values = []
    for name in sheetnames:
        match = re.match(r"^\s*(\d{2})\b", name)
        if match:
            values.append(int(match.group(1)))
    return sorted(values)


def find_category_sheet(workbook, category):
    for name in workbook.sheetnames:
        if canonical_category(re.sub(r"^\s*\d+\s*", "", name)) == category:
            return workbook[name]
    return None


def claim_e11(ctx):
    prefixes = numbered_sheet_prefixes(ctx.bill_wb.sheetnames)
    cdfi = find_category_sheet(ctx.bill_wb, "CDFI Fund")
    cdfi_rows = max(0, cdfi.max_row - 1) if cdfi else None
    manual_cutax = [
        row for row in ctx.manual if canonical_category(row.get("Category")) == "CU Tax Status"
    ]
    cutax_ids = collections.Counter(norm_text(row.get("Bill")) for row in manual_cutax)
    flags = (
        "MATCH" if prefixes == list(range(1, 18)) else "DRIFT",
        "MATCH" if cdfi_rows == 0 else "DRIFT",
        "MATCH"
        if len(manual_cutax) == 34
        and cutax_ids == {"LTR-119-CUTAX-CISCOMANI": 34}
        else "DRIFT",
    )
    return evidence_entry(
        "Archive category-sheet numbering, empty CDFI sheet, and provenance of all CU Tax Status manual rows",
        {
            "numbered_sheets": "01..17 contiguous",
            "cdfi_data_rows": 0,
            "cu_tax_rows": 34,
            "cu_tax_source_id": "LTR-119-CUTAX-CISCOMANI",
            "citation": REPORT_CITATION + " §Q2",
        },
        {
            "numbered_sheet_prefixes": prefixes,
            "cdfi_sheet": cdfi.title if cdfi else None,
            "cdfi_data_rows": cdfi_rows,
            "cu_tax_rows": len(manual_cutax),
            "cu_tax_source_ids": dict(sorted(cutax_ids.items())),
        },
        {
            "numbering_missing_or_extra": sorted(
                set(prefixes) ^ set(range(1, 18))
            ),
            "cdfi_data_rows": None if cdfi_rows is None else cdfi_rows,
            "cu_tax_rows": len(manual_cutax) - 34,
            "flags": list(flags),
        },
        all_flags_verdict(flags),
        "Workbook sheet names and data-row extents are read directly. CU Tax Status is absent from the 01..17 archive tabs and is independently counted from manual_bill_evidence.csv.",
    )


def claim_e12(ctx):
    sheet = find_category_sheet(ctx.bill_wb, "Interchange")
    if sheet is None:
        return evidence_entry(
            "Interchange archive footprint",
            {
                "rows": 36,
                "bills": 6,
                "act_families": 1,
                "citation": REPORT_CITATION + " §Q2",
            },
            {"missing_input": "Interchange category sheet"},
            None,
            "NOT-RECOMPUTABLE-OFFLINE",
            "The copied workbook did not contain a sheet canonicalizing to Interchange.",
        )
    _, rows = workbook_rows(sheet)
    bills = {
        (norm_text(row.get("Name")), congress_text(row.get("Congress")))
        for row in rows
        if norm_text(row.get("Name"))
    }
    acts = {
        norm_text(row.get("Title")).casefold()
        for row in rows
        if norm_text(row.get("Title"))
    }
    act_families = {
        re.sub(r"\s+of\s+\d{4}\s*$", "", title).strip() for title in acts
    }
    flags = (
        "MATCH" if len(rows) == 36 else "DRIFT",
        "MATCH" if len(bills) == 6 else "DRIFT",
        "MATCH" if len(act_families) == 1 else "DRIFT",
    )
    return evidence_entry(
        "Interchange archive footprint",
        {
            "rows": 36,
            "bills": 6,
            "act_families": 1,
            "citation": REPORT_CITATION + " §Q2",
        },
        {
            "sheet": sheet.title,
            "rows": len(rows),
            "bills": len(bills),
            "bill_instances": [list(value) for value in sorted(bills)],
            "act_families": len(act_families),
            "normalized_act_families": sorted(act_families),
            "distinct_title_vintages": len(acts),
            "normalized_titles": sorted(acts),
        },
        {
            "rows": len(rows) - 36,
            "bills": len(bills) - 6,
            "act_families": len(act_families) - 1,
            "flags": list(flags),
        },
        all_flags_verdict(flags),
        "Counts the copied Interchange sheet directly. Act family strips only a trailing 'of <year>' from the whitespace-normalized lowercase title, so the 2022/2023/2026 title vintages remain visible while reconciling to one Credit Card Competition Act family.",
    )


def claim_e13(ctx):
    source_rows = [
        {
            "title": row.get("Title"),
            "congress": row.get("Congress"),
            "source": "bill archive",
        }
        for row in ctx.bill_rows
    ] + [
        {
            "title": row.get("Title"),
            "congress": row.get("Congress"),
            "source": "manual evidence",
        }
        for row in ctx.manual
    ]
    title_literal_hits_all_rows = [
        row for row in source_rows if "disapproval" in str(row.get("title") or "").casefold()
    ]
    classified, _bdir = selected_sign_source_rows(ctx)
    ccm_title_branch_rows = [
        row for row, source in classified if source.startswith("5 ")
    ]
    ccm_branch_literal_firings = [
        row
        for row in ccm_title_branch_rows
        if "disapproval" in str(row.get("title") or "").casefold()
    ]
    current_candidates = [
        row
        for row in source_rows
        if congress_text(row.get("congress")) == str(CURRENT_CONGRESS)
        and DISAPPROVAL_RE.search(str(row.get("title") or ""))
    ]
    new_hits = [row for row in current_candidates if is_disapproval(row.get("title"))]
    literal_absent = [
        row
        for row in current_candidates
        if "disapproval" not in str(row.get("title") or "").casefold()
    ]
    flags = (
        "MATCH" if len(source_rows) == 11283 else "DRIFT",
        "MATCH" if len(ccm_branch_literal_firings) == 0 else "DRIFT",
        "MATCH" if len(literal_absent) == 123 else "DRIFT",
        "MATCH" if len(current_candidates) == 130 else "DRIFT",
    )
    return evidence_entry(
        "Literal disapproval branch versus the current disapprov(al|ing) rule",
        {
            "ccm_branch_literal_firings": 0,
            "all_source_rows": 11283,
            "mismatch_rows": 123,
            "current_disapproval_rows": 130,
            "citation": REPORT_CITATION + " §Q5",
        },
        {
            "ccm_branch_literal_firings": len(ccm_branch_literal_firings),
            "ccm_title_keyword_branch_rows": len(ccm_title_branch_rows),
            "title_literal_hits_all_rows": len(title_literal_hits_all_rows),
            "first_class_finding": (
                {
                    "status": "UNEXPECTED CCM LITERAL BRANCH FIRING",
                    "count": len(ccm_branch_literal_firings),
                    "examples": ccm_branch_literal_firings[:10],
                }
                if ccm_branch_literal_firings
                else {
                    "status": "NONE",
                    "count": 0,
                    "examples": [],
                }
            ),
            "all_source_rows": len(source_rows),
            "new_is_disapproval_current_congress_hits": len(new_hits),
            "current_disapproval_rows": len(current_candidates),
            "current_rows_without_literal_disapproval": len(literal_absent),
        },
        {
            "ccm_branch_literal_firings": len(ccm_branch_literal_firings),
            "title_literal_hits_all_rows_supplementary": len(
                title_literal_hits_all_rows
            ),
            "all_source_rows": len(source_rows) - 11283,
            "mismatch_rows": len(literal_absent) - 123,
            "current_disapproval_rows": len(current_candidates) - 130,
            "flags": list(flags),
        },
        all_flags_verdict(flags),
        "The reported zero is a branch-firing count, not a corpus-wide title search. It applies the full fec/legislation/build_archive_legislative.py:76-109 precedence and counts literal 'disapproval' only among rows that reach bucket 5, the Credit Card Mandates title-keyword branch. Corpus-wide literal title hits are supplementary. The replacement mirrors fec/bill_rules.py:14,105-107; current Congress is 119.",
    )


def claim_e14(ctx):
    sponsor_substring = [
        row for row in ctx.manual if "sponsor" in norm_text(row.get("Role")).casefold()
    ]
    ltr_rows_any_role = [
        row
        for row in ctx.manual
        if norm_text(row.get("Bill")).upper().startswith("LTR-")
    ]
    non_ltr_rows = [
        row
        for row in ctx.manual
        if not norm_text(row.get("Bill")).upper().startswith("LTR-")
    ]
    matrix = ctx.category_stance.get("matrix", {})
    corrob = ctx.category_stance.get("gate", {}).get("corrob", {})
    letter_flags = collections.defaultdict(lambda: collections.defaultdict(list))
    for row in ctx.selected:
        letter_flags[row["member"]][row["effective_category"]].append(
            norm_text(row.get("bill")).upper().startswith("LTR-")
        )
    aligned_vote = collections.defaultdict(set)
    name_of = ctx.vote_lean.get("name_of", {})
    for bioguide, categories in ctx.vote_lean.get("votes", {}).items():
        member = name_of.get(bioguide)
        if not member:
            continue
        for category, detail in categories.items():
            if number(detail.get("lean")) >= 0.5 and number(detail.get("n")) >= 1:
                aligned_vote[member].add(category)
    resting = []
    for member, categories in matrix.items():
        for category, tier in categories.items():
            if tier not in ("Champion", "Defender"):
                continue
            if corrob.get(member, {}).get(category) != "A":
                continue
            flags = letter_flags.get(member, {}).get(category)
            if flags and all(flags) and category not in aligned_vote.get(member, set()):
                resting.append({"member": member, "category": category})
    results = (
        "MATCH" if len(ltr_rows_any_role) == 296 else "DRIFT",
        "MATCH" if len(sponsor_substring) == 297 else "DRIFT",
        "MATCH" if len(ctx.manual) == 297 else "DRIFT",
        "MATCH" if len(resting) == 187 else "DRIFT",
    )
    return evidence_entry(
        "Manual letter rows admitted by the sponsor-substring gate and top-tier cells resting on letter-only Rule A",
        {
            "ltr_rows_any_role": 296,
            "total_sponsor_substring_pass": 297,
            "manual_rows": 297,
            "letter_only_top_tier_cells": 187,
            "citation": REPORT_CITATION + " §Q1",
        },
        {
            "ltr_rows_any_role": len(ltr_rows_any_role),
            "ltr_role_breakdown": dict(
                sorted(
                    collections.Counter(
                        norm_text(row.get("Role")) for row in ltr_rows_any_role
                    ).items()
                )
            ),
            "total_sponsor_substring_pass": len(sponsor_substring),
            "manual_rows": len(ctx.manual),
            "non_ltr_rows": [
                {
                    key: row.get(key)
                    for key in (
                        "Member",
                        "Role",
                        "Bill",
                        "Title",
                        "Congress",
                        "Category",
                    )
                }
                for row in non_ltr_rows
            ],
            "roles": dict(
                sorted(collections.Counter(row.get("Role") for row in ctx.manual).items())
            ),
            "letter_only_top_tier_cells": len(resting),
            "by_category": dict(
                sorted(collections.Counter(row["category"] for row in resting).items())
            ),
            "cells": sorted(resting, key=lambda row: (row["member"], row["category"])),
        },
        {
            "ltr_rows_any_role": len(ltr_rows_any_role) - 296,
            "total_sponsor_substring_pass": len(sponsor_substring) - 297,
            "manual_rows": len(ctx.manual) - 297,
            "letter_only_top_tier_cells": len(resting) - 187,
            "flags": list(results),
        },
        all_flags_verdict(results),
        "The reported 296 is the LTR-* population across every role, not only original cosponsors. Its role breakdown and every non-LTR identity are explicit. The literal moneyball_core.py:355-369 sponsor-substring gate passes all 297 rows. The 187-cell counterfactual mirrors AUDIT_SP/p13.py:15-34.",
    )


def claim_e15(ctx):
    targets = ("S.427", "H.R.3380", "H.R.1", "S.3940", "S.1880")
    archive_keys = collections.Counter(
        bill_key(row.get("Name")) for row in ctx.bill_rows if bill_key(row.get("Name"))
    )
    manual_keys = collections.Counter(
        bill_key(row.get("Bill")) for row in ctx.manual if bill_key(row.get("Bill"))
    )
    zero_rows = {
        target: archive_keys[bill_key(target)] + manual_keys[bill_key(target)]
        for target in targets
    }
    s2704 = [
        {
            "section": row.get("Section"),
            "role": row.get("Role"),
            "member": row.get("Member"),
            "name": row.get("Name"),
            "title": row.get("Title"),
            "congress": row.get("Congress"),
        }
        for row in ctx.bill_rows
        if bill_key(row.get("Name")) == bill_key("S.2704")
    ]
    observed_titles = sorted({norm_text(row["title"]) for row in s2704})
    observed_congresses = sorted({congress_text(row["congress"]) for row in s2704})
    different = (
        bool(s2704)
        and observed_congresses == ["117"]
        and any("small business surcharge relief" in title.casefold() for title in observed_titles)
        and not any("cdfi fund transparency" in title.casefold() for title in observed_titles)
    )
    flags = (
        "MATCH" if all(value == 0 for value in zero_rows.values()) else "DRIFT",
        "MATCH" if different else "DRIFT",
    )
    return evidence_entry(
        "Five outline bills absent from the corpus and S.2704 identifying a different act",
        {
            "zero_row_bills": list(targets),
            "s2704_actual": "117th Small Business Surcharge Relief",
            "s2704_not": "CDFI Fund Transparency",
            "citation": REPORT_CITATION + " §Q4",
        },
        {
            "target_row_counts": zero_rows,
            "s2704_rows": len(s2704),
            "s2704_titles": observed_titles,
            "s2704_congresses": observed_congresses,
            "is_different_bill": different,
        },
        {
            "nonzero_target_bills": [
                target for target, value in zero_rows.items() if value != 0
            ],
            "s2704_flag": flags[1],
        },
        all_flags_verdict(flags),
        "Outline membership mirrors AUDIT_SP/outline.py:41-70 using bill_key matching over archive plus manual rows. The S.2704 identity check mirrors AUDIT_SP/vf_s2704.py:2-8.",
    )


def network_pairs(ctx, sheet_name, symmetrize, layer_filter=None, workbook=None):
    workbook = workbook or ctx.network_wb
    if sheet_name not in workbook.sheetnames:
        raise EvidenceError("network workbook lacks %s sheet" % sheet_name)
    ws = workbook[sheet_name]
    headers, rows = workbook_rows(ws)
    lower = {norm_header(header): header for header in headers}
    source_key = lower.get("source_bioguide") or headers[0]
    target_key = lower.get("target_bioguide") or headers[1]
    layer_key = lower.get("layer")
    pairs = set()
    for row in rows:
        if (
            layer_filter
            and layer_key
            and layer_filter not in norm_text(row.get(layer_key)).casefold()
        ):
            continue
        source = norm_text(row.get(source_key)).upper()
        target = norm_text(row.get(target_key)).upper()
        if source and target:
            pairs.add((source, target))
            if symmetrize:
                pairs.add((target, source))
    return pairs, ws.title


def claim_e16(ctx):
    # AUDIT_SP/p9.py:17-28 exact footprint predicate and selected-event weight.
    def is_safe_routing(row):
        return (
            "safe" in str(row.get("title") or "").casefold()
            and "banking act" in str(row.get("title") or "").casefold()
        )

    safe_selected = [
        row for row in ctx.selected if is_safe_routing(row)
    ]
    safe_weight = sum(abs(number(row.get("contribution"))) for row in safe_selected)
    all_weight = sum(abs(number(row.get("contribution"))) for row in ctx.selected)
    safe_share = pct(safe_weight, all_weight)
    safe_all_dispositions = [
        row for row in ctx.routing.get("rows", []) if is_safe_routing(row)
    ]
    all_disposition_contributions = []
    for row in safe_all_dispositions:
        contribution = row.get("contribution")
        if contribution is None and row.get("disposition") != "selected":
            _record_expected_absent(
                "E16.non_selected_routing_contribution", 1
            )
            continue
        all_disposition_contributions.append(number(contribution))
    variant_weight_attempts = {
        "selected_only_signed": sum(
            number(row.get("contribution")) for row in safe_selected
        ),
        "all_dispositions_absolute": sum(
            abs(value) for value in all_disposition_contributions
        ),
        "all_dispositions_signed": sum(all_disposition_contributions),
    }
    weight_reconciliation = e16_weight_reconciliation(
        safe_weight,
        all_weight,
        variant_weight_attempts,
    )
    safe_bills = {
        (norm_text(row.get("bill")), congress_text(row.get("congress")))
        for row in safe_selected
    }
    name_to_bioguide = {
        norm_text(row.get("member")): norm_text(row.get("bioguide")).upper()
        for row in ctx.congestion
        if norm_text(row.get("member")) and norm_text(row.get("bioguide"))
    }
    safe_archive_groups = collections.defaultdict(list)
    for row in ctx.bill_rows:
        title = str(row.get("Title") or "").casefold()
        if "safe" in title and "banking act" in title:
            safe_archive_groups[
                (norm_text(row.get("Name")), congress_text(row.get("Congress")))
            ].append(row)

    # Production-semantic comparison: moneyball_core.py:145-164 selects AUG
    # Edges and symmetrizes them.  Pair links are kept bill-specific so the
    # same member tie on two SAFE bills remains two evidence links.
    required_unique_bill_pairs = set()
    unmapped_members = set()
    for bill, rows in safe_archive_groups.items():
        sponsors = {
            norm_text(row.get("Member"))
            for row in rows
            if canonical_role(row.get("Role")) == "sponsor"
        }
        cosponsors = {
            norm_text(row.get("Member"))
            for row in rows
            if canonical_role(row.get("Role")) in ("original-cosponsor", "cosponsor")
        }
        for sponsor in sponsors:
            for cosponsor in cosponsors:
                source = name_to_bioguide.get(sponsor)
                target = name_to_bioguide.get(cosponsor)
                if source and target:
                    if source != target:
                        required_unique_bill_pairs.add(
                            (bill[0], bill[1], source, target)
                        )
                else:
                    if not source:
                        unmapped_members.add(sponsor)
                    if not target:
                        unmapped_members.add(cosponsor)

    production_edges, production_sheet = network_pairs(
        ctx, "Edges", symmetrize=True
    )
    production_missing = sorted(
        required_unique_bill_pairs,
        key=lambda item: (item[0], item[1], item[2], item[3]),
    )
    production_missing = [
        item
        for item in production_missing
        if (item[2], item[3]) not in production_edges
    ]

    # Preserve the first-run method as its own labeled scope: cosponsor-only
    # provenance, mean-symmetrized, unique bill/member links.
    provenance_edges, provenance_sheet = network_pairs(
        ctx, "Provenance", symmetrize=True, layer_filter="cosponsor"
    )
    provenance_missing = sorted(
        (
            item
            for item in required_unique_bill_pairs
            if (item[2], item[3]) not in provenance_edges
        ),
        key=lambda item: (item[0], item[1], item[2], item[3]),
    )

    # AUDIT_SP/net_gap2.py:18-59 semantics: keep cosponsor row
    # multiplicity, use sponsor->cosponsor direction only, skip mapping
    # failures, and count each row/sponsor combination.  This replay uses the
    # exact June-25 base workbook named at net_gap2.py:10.
    audit_direct_edges, audit_replay_sheet = network_pairs(
        ctx,
        "Edges",
        symmetrize=False,
        workbook=ctx.base_network_wb,
    )
    audit_required_instances = 0
    audit_hits = 0
    audit_misses = 0
    audit_mapping_skips = 0
    audit_hits_by_bill = collections.Counter()
    audit_misses_by_bill = collections.Counter()
    audit_mapping_skips_by_bill = collections.Counter()
    audit_missing_examples = []
    for bill, rows in sorted(safe_archive_groups.items()):
        sponsors = {
            norm_text(row.get("Member"))
            for row in rows
            if canonical_role(row.get("Role")) == "sponsor"
        }
        for row in rows:
            if canonical_role(row.get("Role")) not in (
                "original-cosponsor",
                "cosponsor",
            ):
                continue
            cosponsor = norm_text(row.get("Member"))
            cosponsor_bio = name_to_bioguide.get(cosponsor)
            if not cosponsor_bio:
                audit_mapping_skips += 1
                audit_mapping_skips_by_bill[bill] += 1
                continue
            for sponsor in sorted(sponsors):
                sponsor_bio = name_to_bioguide.get(sponsor)
                if not sponsor_bio or sponsor_bio == cosponsor_bio:
                    audit_mapping_skips += 1
                    audit_mapping_skips_by_bill[bill] += 1
                    continue
                audit_required_instances += 1
                if (sponsor_bio, cosponsor_bio) in audit_direct_edges:
                    audit_hits += 1
                    audit_hits_by_bill[bill] += 1
                else:
                    audit_misses += 1
                    audit_misses_by_bill[bill] += 1
                    if len(audit_missing_examples) < 20:
                        audit_missing_examples.append(
                            {
                                "bill": bill[0],
                                "congress": bill[1],
                                "sponsor_bioguide": sponsor_bio,
                                "cosponsor_bioguide": cosponsor_bio,
                            }
                        )

    # Full net_gap2.py:18-59 replay across the entire archive, retained as a
    # cross-check on Claude's verbatim 2026-07-25 run.
    archive_groups = collections.defaultdict(list)
    for row in ctx.bill_rows:
        archive_groups[
            (norm_text(row.get("Name")), congress_text(row.get("Congress")))
        ].append(row)
    archive_wide_hits = 0
    archive_wide_misses = 0
    archive_wide_mapping_skips = 0
    for _bill, rows in sorted(archive_groups.items()):
        sponsors = {
            norm_text(row.get("Member"))
            for row in rows
            if canonical_role(row.get("Role")) == "sponsor"
        }
        if not any(name_to_bioguide.get(sponsor) for sponsor in sponsors):
            continue
        for row in rows:
            if canonical_role(row.get("Role")) not in (
                "original-cosponsor",
                "cosponsor",
            ):
                continue
            cosponsor_bio = name_to_bioguide.get(
                norm_text(row.get("Member"))
            )
            if not cosponsor_bio:
                archive_wide_mapping_skips += 1
                continue
            for sponsor in sponsors:
                sponsor_bio = name_to_bioguide.get(sponsor)
                if not sponsor_bio or sponsor_bio == cosponsor_bio:
                    archive_wide_mapping_skips += 1
                    continue
                if (sponsor_bio, cosponsor_bio) in audit_direct_edges:
                    archive_wide_hits += 1
                else:
                    archive_wide_misses += 1

    reported_bill = ("H.R.2891", "118")
    reported_bill_missing = audit_misses_by_bill[reported_bill]
    archive_wide_pairs = archive_wide_hits + archive_wide_misses
    flags = {
        "reported_missing_cosponsor_pairs": (
            "MATCH" if reported_bill_missing == 448 else "DRIFT"
        ),
        "reported_bill_count": (
            "MATCH" if len(safe_bills) == 5 else "DRIFT"
        ),
        "reported_weight": weight_reconciliation["weight_flag"],
        "reported_weight_pct": weight_reconciliation["weight_pct_flag"],
    }
    misscope_replay_failures = []
    if reported_bill_missing != 448:
        misscope_replay_failures.append("single-bill expected 448 got %d"
                                        % reported_bill_missing)
    if audit_misses != 472:
        misscope_replay_failures.append("family-total expected 472 got %d"
                                        % audit_misses)
    verdict = (
        "DRIFT: MISSCOPING-NOT-REPRODUCED (%s)"
        % "; ".join(misscope_replay_failures)
        if misscope_replay_failures
        else "REPORTED-VALUE-MISSCOPED"
    )
    supplementary_replay_failures = []
    if len(production_missing) != 108:
        supplementary_replay_failures.append(
            "production-AUG expected 108 got %d" % len(production_missing))
    if archive_wide_misses != 2186 or archive_wide_pairs != 8518:
        supplementary_replay_failures.append(
            "archive-wide expected 2186/8518 got %d/%d"
            % (archive_wide_misses, archive_wide_pairs))
    if len(safe_bills) != 5:
        supplementary_replay_failures.append(
            "SAFE-bills expected 5 got %d" % len(safe_bills))
    if supplementary_replay_failures:
        verdict += "; replay DRIFT: %s" % ", ".join(
            supplementary_replay_failures)
    return evidence_entry(
        "SAFE/SAFER Banking footprint and bill-specific cosponsor pairs absent from the selected network artifact",
        {
            "weight_pct": 14.5,
            "bills": 5,
            "weight": 1416,
            "missing_cosponsor_pairs": 448,
            "citation": REPORT_CITATION + " §Q3",
        },
        {
            "weight": compact_number(safe_weight),
            "weight_pct": compact_number(safe_share),
            "weight_flag": weight_reconciliation["weight_flag"],
            "weight_pct_flag": weight_reconciliation["weight_pct_flag"],
            "weight_status": weight_reconciliation["reproduction_status"],
            "weight_method_attempts": {
                "selected_events": len(safe_selected),
                "all_disposition_rows_with_safe_title": len(
                    safe_all_dispositions
                ),
                "classification": (
                    "attempts metadata only; no variant can reproduce the "
                    "primary p9-mirror result"
                ),
                "primary": weight_reconciliation,
                "bills": len(safe_bills),
                "bill_instances": [list(item) for item in sorted(safe_bills)],
            },
            "audit_method_vs_base_network": {
                "reported_448_actual_scope": "H.R.2891 (118th) alone, per net_gap2.py TOP-ZERO line",
                "reported_bill_missing_pairs": reported_bill_missing,
                "corrected_family_total_vs_base_network": audit_misses,
                "required_pair_instances": audit_required_instances,
                "directed_hits": audit_hits,
                "mapping_skips": audit_mapping_skips,
                "pair_multiplicity": "archive cosponsor rows x roster-mapped sponsor rows",
                "direction": "sponsor->cosponsor only",
                "network_artifact": "congress_cosponsor_network_20260625_205554.xlsx",
                "network_sheet": audit_replay_sheet,
                "network_in_manifest": True,
                "safe_bill_missing_pairs": {
                    "%s C%s" % bill: {
                        "missing_pairs": audit_misses_by_bill[bill],
                        "mapping_skips": audit_mapping_skips_by_bill[bill],
                        "status": (
                            "MISSING-PAIRS"
                            if audit_misses_by_bill[bill]
                            else (
                                "ZERO-MISSING-WITH-MAPPING-SKIPS"
                                if audit_mapping_skips_by_bill[bill]
                                else "ZERO-MISSING-NO-MAPPING-SKIPS"
                            )
                        ),
                    }
                    for bill in sorted(safe_archive_groups)
                },
                "missing_pair_sample": audit_missing_examples,
            },
            "production_semantics_vs_newest_AUG": {
                "network_artifact": "congress_cosponsor_network_AUG20260719.xlsx",
                "network_sheet": production_sheet,
                "symmetrized": True,
                "required_unique_bill_pairs": len(required_unique_bill_pairs),
                "missing_unique_bill_pairs": len(production_missing),
                "unmapped_members": sorted(unmapped_members),
                "matches_moneyball_core_selection_semantics": True,
                "unit_caveat_vs_reported_448": (
                    "108 is a count of unique symmetrized bill/member pairs; "
                    "448 is a count of directed multiplicity instances. The "
                    "two values are supplementary and not commensurate."
                ),
            },
            "archive_wide_base_network_replay": {
                "archive_wide_missing": archive_wide_misses,
                "archive_wide_pairs": archive_wide_pairs,
                "archive_wide_hits": archive_wide_hits,
                "mapping_skips": archive_wide_mapping_skips,
                "independent_replay_citation": "Claude verbatim replay of AUDIT_SP/net_gap2.py, 2026-07-25",
                "independent_replay_reported": {
                    "missing": 2186,
                    "pairs": 8518,
                },
            },
            "first_run_provenance_deduplicated_method": {
                "network_sheet": provenance_sheet,
                "layer_filter": "cosponsor",
                "symmetrized": True,
                "required_unique_bill_pairs": len(required_unique_bill_pairs),
                "missing_unique_bill_pairs": len(provenance_missing),
                "matches_moneyball_core_selection_semantics": False,
            },
            "method_scope_difference": {
                "audit_base": "June-25 base Edges; directed sponsor->cosponsor row multiplicity; mapping failures skipped",
                "reported_448": "single H.R.2891/C118 TOP-ZERO row transcribed as though it covered the five-bill family",
                "corrected_family": "all five SAFE/SAFER bill instances under the base-network method",
                "first_run": "cosponsor-provenance-only, symmetrized, unique bill/member links",
                "production_AUG": "July-19 AUG all-layer Edges, symmetrized, unique bill/member links for this comparison",
            },
        },
        {
            "weight_pp": compact_number(
                pct_delta(safe_share, 14.5)
            ),
            "bills": len(safe_bills) - 5,
            "weight": compact_number(safe_weight - 1416.0),
            "hypothesis_for_reported_scope": weight_reconciliation[
                "attempts_metadata"
            ]["hypothesis_for_reported_scope"],
            "reported_bill_missing_pairs": reported_bill_missing - 448,
            "recomputed_only": {
                "corrected_family_total_vs_base_network": audit_misses,
                "archive_wide_missing": archive_wide_misses,
                "archive_wide_pairs": archive_wide_pairs,
                "production_semantics_missing_unique_pairs": len(
                    production_missing
                ),
                "first_run_missing_unique_pairs": len(provenance_missing),
            },
            "flags": flags,
        },
        verdict,
        "Arbitration: the report's 448 is the H.R.2891/C118 TOP-ZERO count, not the five-bill SAFE/SAFER family total. audit_method_vs_base_network mirrors AUDIT_SP/net_gap2.py:10,18-59 against the copied June-25 BASE and yields the corrected family total; archive-wide results are independently cited to 'Claude verbatim replay of AUDIT_SP/net_gap2.py, 2026-07-25'. production_semantics_vs_newest_AUG uses the copied July-19 AUG and fec/moneyball_core.py:145-164 symmetrization. Its 108 missing unique symmetrized bill/member pairs are a supplementary unit, not a comparison with the 448 directed multiplicity instances. The p9-mirror selected-only absolute weight is the primary recomputed value and is independently flagged against both 1,416 and 14.5%; its reproduction status is NOT-REPRODUCED. Alternate scopes are attempts metadata only, and a close alternate is named solely as hypothesis_for_reported_scope.",
    )


def static_consumers():
    needles = ("data/legislation/cu_bill_directions.json", "data/legislation/bill_direction_queue_FILLED.csv")
    excluded_tooling_modules = frozenset(
        {
            "fec/legislation/leg_evidence_recompute.py",
            "fec/legislation/bill_archive_provenance.py",
            "fec/legislation/make_bill_directions_congress_candidate.py",
        }
    )
    guard_modules = frozenset({"fec/legislation/bill_coverage_audit.py"})
    scoring_hits = []
    guard_hits = []
    scanned_files = []
    read_errors = []
    for path in sorted((WORKTREE_ROOT / "fec").rglob("*"), key=lambda p: str(p).casefold()):
        if not path.is_file() or path.suffix.casefold() not in (".py", ".js"):
            continue
        if "__pycache__" in path.parts or "attic" in path.parts:
            continue
        relative_path = str(path.relative_to(WORKTREE_ROOT)).replace("\\", "/")
        try:
            payload = path.read_bytes()
        except OSError as exc:
            read_errors.append(
                {
                    "path": relative_path,
                    "error": exc.__class__.__name__,
                }
            )
            continue
        scanned_files.append(
            [relative_path, hashlib.sha256(payload).hexdigest()]
        )
        text = payload.decode("utf-8", errors="replace")
        for needle in needles:
            if needle not in text:
                continue
            hit = {"path": relative_path, "needle": needle}
            if relative_path in guard_modules:
                guard_hits.append(hit)
            elif (
                relative_path not in excluded_tooling_modules
                and not relative_path.startswith("fec/tests/")
            ):
                scoring_hits.append(hit)
    code_scan_fingerprint = hashlib.sha256(
        json.dumps(
            sorted(scanned_files),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    return {
        "scoring_consumers": scoring_hits,
        "guard_readers": guard_hits,
        "scanned_file_count": len(scanned_files),
        "code_scan_fingerprint": code_scan_fingerprint,
        "scan_read_errors": read_errors,
    }


def claim_e17(ctx):
    bills = ctx.cu_directions.get("bills", ctx.cu_directions)
    positions = collections.Counter()
    blank_position_bills = []
    for value in bills.values():
        position = value.get("position") if isinstance(value, dict) else value
        positions[norm_text(position).upper() or "(BLANK)"] += 1
    for bill, value in bills.items():
        position = value.get("position") if isinstance(value, dict) else value
        if not norm_text(position):
            blank_position_bills.append(bill)
    blank_position_bills.sort()
    consumer_scan = static_consumers()
    scoring_consumer_hits = consumer_scan["scoring_consumers"]
    guard_reader_hits = consumer_scan["guard_readers"]
    expected_guard_readers = ["fec/legislation/bill_coverage_audit.py"]
    observed_guard_readers = sorted(
        {hit["path"] for hit in guard_reader_hits}
    )
    queue_filled = sum(
        bool(norm_text(row.get("direction_+1_0_-1")))
        for row in ctx.direction_queue
    )
    reported_positions = {"SUPPORT": 54, "OPPOSE": 6, "MONITOR": 1}
    flags = {
        "cu_direction_bills": "MATCH" if len(bills) == 63 else "DRIFT",
        "reported_positions": (
            "MATCH"
            if all(
                positions.get(key, 0) == value
                for key, value in reported_positions.items()
            )
            else "DRIFT"
        ),
        "queue_filled_rows": "MATCH" if queue_filled == 57 else "DRIFT",
        "blank_position_bills": (
            "MATCH" if len(blank_position_bills) == 2 else "DRIFT"
        ),
        "scoring_consumers": (
            DATA_INVALID
            if consumer_scan["scan_read_errors"]
            else (
                "MATCH"
                if len(scoring_consumer_hits) == 0
                else "DRIFT"
            )
        ),
        "guard_readers": (
            DATA_INVALID
            if consumer_scan["scan_read_errors"]
            else (
                "MATCH"
                if observed_guard_readers == expected_guard_readers
                else "DRIFT"
            )
        ),
    }
    return evidence_entry(
        "Adjudicated CU bill directions, filled queue, scoring consumers, and guard readers",
        {
            "cu_direction_bills": 63,
            "positions": reported_positions,
            "queue_filled_rows": 57,
            "blank_position_bills": 2,
            "scoring_consumers": 0,
            "guard_readers": expected_guard_readers,
            "citation": REPORT_CITATION + " §Tier 1",
        },
        {
            "cu_direction_bills": len(bills),
            "positions": dict(sorted(positions.items())),
            "blank_position_bills": blank_position_bills,
            "queue_rows": len(ctx.direction_queue),
            "queue_filled_rows": queue_filled,
            "scoring_consumers": len(scoring_consumer_hits),
            "scoring_consumer_hits": scoring_consumer_hits,
            "guard_readers": observed_guard_readers,
            "guard_reader_hits": guard_reader_hits,
            "manifest_exception": E17_MANIFEST_EXCEPTION,
            "code_scan_fingerprint": consumer_scan[
                "code_scan_fingerprint"
            ],
            "code_tree_scan": {
                "scanned_file_count": consumer_scan["scanned_file_count"],
                "code_scan_fingerprint": consumer_scan[
                    "code_scan_fingerprint"
                ],
                "read_errors": consumer_scan["scan_read_errors"],
                "manifest_verified_input": False,
            },
        },
        {
            "cu_direction_bills": len(bills) - 63,
            "position_deltas": {
                key: positions.get(key, 0) - value
                for key, value in reported_positions.items()
            },
            "queue_filled_rows": queue_filled - 57,
            "queue_rows_total_supplementary": len(ctx.direction_queue),
            "blank_position_bills": len(blank_position_bills) - 2,
            "scoring_consumers": len(scoring_consumer_hits),
            "guard_reader_delta": sorted(
                set(observed_guard_readers).symmetric_difference(
                    expected_guard_readers
                )
            ),
            "code_tree_scan_read_errors": len(
                consumer_scan["scan_read_errors"]
            ),
            "flags": flags,
        },
        all_flags_verdict(flags.values()),
        "The reported 57 is the count of rows carrying a human direction, not the queue's 216 total rows. cu_bill_directions.json contains 63 bill objects but the reported 54/6/1 positions sum to 61, so the two blank-position bill IDs are elevated as a first-class Tier-1 wiring finding. %s. The scan covers fec/**/*.py and fec/**/*.js and records a deterministic combined sha256 over the sorted (path, file-sha256) list. leg_evidence_recompute.py, bill_archive_provenance.py, make_bill_directions_congress_candidate.py, and fec/tests/* are excluded from consumer classification. bill_coverage_audit.py is a guard_reader because C14 reads directions for guard purposes, not scoring; any hit in another fec module is a scoring_consumer and is compared with the reported zero."
        % E17_MANIFEST_EXCEPTION,
    )


def claim_e18(ctx):
    category_weight = collections.Counter()
    for row in ctx.selected:
        category_weight[row.get("effective_category")] += abs(
            number(row.get("contribution"))
        )
    pillar_weight = {
        layer: sum(category_weight[category] for category in categories)
        for layer, categories in L1_TO_CATEGORIES.items()
    }
    pillar_total = sum(pillar_weight.values())
    pillar_share = {
        layer: value / pillar_total if pillar_total else None
        for layer, value in pillar_weight.items()
    }
    acu_mapped_total = sum(
        ACU_L1_COUNTS_20260724[layer] for layer in L1_TO_CATEGORIES
    )
    acu_share = {
        layer: ACU_L1_COUNTS_20260724[layer] / acu_mapped_total
        for layer in L1_TO_CATEGORIES
    }
    multipliers = {}
    flags = {}
    deltas = {}
    for layer, reported in REPORTED_E18_MULTIPLIERS.items():
        recomputed = (
            pillar_share[layer] / acu_share[layer]
            if pillar_share[layer] is not None and acu_share[layer]
            else None
        )
        multipliers[layer] = compact_number(recomputed)
        deltas[layer] = compact_number(
            None if recomputed is None else recomputed - reported
        )
        flags[layer] = scalar_flag(recomputed, reported, RHO_TOLERANCE)
    rendered_flags = {
        layer: {
            "flag": flag,
            **(
                {"caveat": "proves pillar-side stability only"}
                if flag == "MATCH"
                else {}
            ),
        }
        for layer, flag in sorted(flags.items())
    }
    drift_layers = sorted(
        layer for layer, flag in flags.items() if flag == "DRIFT"
    )
    invalid_layers = sorted(
        layer for layer, flag in flags.items() if flag == DATA_INVALID
    )
    verdict = "SPLIT: pillar-side recomputed, ACU-side cached-constant"
    if invalid_layers:
        verdict += " (DATA-INVALID: %s)" % ", ".join(invalid_layers)
    elif drift_layers:
        verdict += " (DRIFT: %s)" % ", ".join(drift_layers)
    return evidence_entry(
        "Layer-1 advocacy-to-pillar gap multipliers with recomputed pillar side and explicitly cached ACU side",
        {
            "multipliers": REPORTED_E18_MULTIPLIERS,
            "citation": REPORT_CITATION + " §Q2 and AUDIT_SP/gapb.py:8-19",
        },
        {
            "E18a_PILLAR_SIDE_RECOMPUTED": {
                "absolute_weight": {
                    key: compact_number(value)
                    for key, value in sorted(pillar_weight.items())
                },
                "shares": {
                    key: compact_number(value)
                    for key, value in sorted(pillar_share.items())
                },
                "mapped_total_weight": compact_number(pillar_total),
            },
            "E18b_ACU_SIDE_CACHED_CONSTANT": {
                "counts_all_nine_layers": ACU_L1_COUNTS_20260724,
                "counts_all_nine_layers_sum": sum(
                    ACU_L1_COUNTS_20260724.values()
                ),
                "letter_count": 533,
                "multi_tag_note": (
                    "Nine-layer counts sum to 642 over 533 letters because "
                    "letters can carry multiple tags."
                ),
                "mapped_seven_layer_denominator": acu_mapped_total,
                "mapped_shares": {
                    key: compact_number(value)
                    for key, value in sorted(acu_share.items())
                },
                "fresh_recompute": False,
            },
            "multipliers": multipliers,
        },
        {"multipliers": deltas, "per_multiplier_flags": rendered_flags},
        verdict,
        "E18 is intentionally not a full fresh recompute. E18a mirrors AUDIT_SP/gapb.py:20-29 from selected routing rows and the cited L1 mapping. E18b embeds AUDIT_SP/gapb.py:8-10 with provenance: 2026-07-24 live GraphQL pull (fieldAdvocacyIssueTypes, 533 letters); re-fetch requires network and is user-gated. Nine-layer counts sum to 642 over 533 letters because letters can carry multiple tags. A per-multiplier MATCH proves pillar-side stability only; multiplier deltas can therefore be attributable ONLY to pillar-side drift.",
    )


def compare_member_maps(left, right, members, tolerance=FLOAT_TOLERANCE):
    compared = 0
    missing_left = []
    missing_right = []
    invalid_values = []
    max_deviation = 0.0
    worst = None
    for member in members:
        if member not in left:
            missing_left.append(member)
            continue
        if member not in right:
            missing_right.append(member)
            continue
        left_value = number(left[member])
        right_value = number(right[member])
        if is_data_invalid(left_value) or is_data_invalid(right_value):
            invalid_values.append(
                {
                    "member": member,
                    "left": compact_number(left_value),
                    "right": compact_number(right_value),
                }
            )
            continue
        deviation = abs(left_value - right_value)
        compared += 1
        if deviation > max_deviation:
            max_deviation = deviation
            worst = {
                "member": member,
                "left": compact_number(left_value),
                "right": compact_number(right_value),
                "abs_deviation": compact_number(deviation),
            }
    return {
        "compared": compared,
        "missing_left": missing_left,
        "missing_right": missing_right,
        "data_invalid": invalid_values,
        "max_abs_deviation": compact_number(max_deviation),
        "worst": worst,
        "verdict": (
            DATA_INVALID
            if invalid_values
            else (
                "MATCH"
                if not missing_left
                and not missing_right
                and max_deviation <= tolerance
                else "MISMATCH"
            )
        ),
    }


def e19a(ctx):
    wb = ctx.model_wb
    blended = wb["Blended Model"]
    rank_allocation = wb["Rank & Allocation"]
    members = [
        rank_allocation.cell(row, 1).value
        for row in range(3, 536)
        if rank_allocation.cell(row, 1).value
    ]
    bm_rows = {member: index + 3 for index, member in enumerate(members)}
    names_alive = sum(
        blended.cell(bm_rows[member], 1).value == member for member in members
    )
    cached_o = {member: blended.cell(bm_rows[member], 15).value for member in members}
    cached_e = {member: blended.cell(bm_rows[member], 5).value for member in members}
    cached_f = {member: blended.cell(bm_rows[member], 6).value for member in members}
    cache_alive_count = sum(value is not None for value in cached_o.values())
    cache_alive = cache_alive_count >= len(members) * 0.9 and names_alive >= len(members) * 0.9
    snapshot_o = ctx.snapshot.get("O", {})
    snapshot_e = ctx.snapshot.get("E", {})
    snapshot_f = ctx.snapshot.get("F", {})
    # fec/engine/compute_pillars.py:54-59 is the authoritative _read_OEF mirror:
    # formula-cache None values become 0.0 only when that cache is alive.
    source_o = (
        {
            member: cited_none_as_zero(cached_o[member])
            for member in members
        }
        if cache_alive
        else snapshot_o
    )
    rank_leg = {
        member: values.get("leg")
        for member, values in ctx.ranks.get("pcts", {}).items()
        if isinstance(values, dict) and values.get("leg") is not None
    }
    source_values = [number(source_o.get(member)) for member in members]
    recomputed_pct = {
        member: percentrank(source_values, number(source_o.get(member)))
        for member in members
    }
    cache_vs_snapshot = (
        compare_member_maps(
            # fec/engine/compute_pillars.py:54-59 _read_OEF ``or 0.0`` mirror.
            {
                member: cited_none_as_zero(cached_o[member])
                for member in members
            },
            snapshot_o,
            members,
        )
        if cache_alive
        else {
            "compared": 0,
            "max_abs_deviation": None,
            "verdict": "EXPECTED-CACHE-DEAD",
        }
    )
    cache_vs_snapshot_e = (
        compare_member_maps(
            # fec/engine/compute_pillars.py:54-59 _read_OEF ``or 0.0`` mirror.
            {
                member: cited_none_as_zero(cached_e[member])
                for member in members
            },
            snapshot_e,
            members,
        )
        if cache_alive
        else {
            "compared": 0,
            "max_abs_deviation": None,
            "verdict": "EXPECTED-CACHE-DEAD",
        }
    )
    cache_vs_snapshot_f = (
        compare_member_maps(
            # fec/engine/compute_pillars.py:54-59 _read_OEF ``or 0.0`` mirror.
            {
                member: cited_none_as_zero(cached_f[member])
                for member in members
            },
            snapshot_f,
            members,
        )
        if cache_alive
        else {
            "compared": 0,
            "max_abs_deviation": None,
            "verdict": "EXPECTED-CACHE-DEAD",
        }
    )
    ranks_check = compare_member_maps(recomputed_pct, rank_leg, members, tolerance=1e-9)
    common = [member for member in members if member in rank_leg]
    left = [number(source_o.get(member)) for member in common]
    right = [number(rank_leg.get(member)) for member in common]
    rho = spearman(left, right)
    tail = rank_disagreements(common, left, right, 10)
    return {
        "member_count": len(members),
        "rank_allocation_names_alive": names_alive,
        "cache_non_none_O": cache_alive_count,
        "cache_none_fraction": compact_number(
            (len(members) - cache_alive_count) / len(members) if members else None
        ),
        "cache_alive": cache_alive,
        "production_O_source": "Blended Model cache" if cache_alive else "data/engine/pillar_inputs_snapshot.json",
        "cache_vs_snapshot_O": cache_vs_snapshot,
        "cache_vs_snapshot_E": cache_vs_snapshot_e,
        "cache_vs_snapshot_F": cache_vs_snapshot_f,
        "snapshot_keys": {
            key: len(ctx.snapshot.get(key, {})) for key in ("O", "E", "F")
        },
        "percentrank_vs_ranks_python_pcts_leg": ranks_check,
        "O_source_vs_ranks_leg_rho": compact_number(rho),
        "O_source_vs_ranks_leg_worst_tail": tail,
        "Blended_Model_headers": sheet_structure(blended),
        "Rank_Allocation_headers": sheet_structure(rank_allocation),
        "OEF_columns": {"O": "O", "E": "E", "F": "F"},
        "citation": "fec/engine/compute_pillars.py:32-61",
    }


def e19b(ctx):
    raw = ctx.model_wb["Legislative Score - Raw"]
    normalized = ctx.model_wb["Legislative Score - Normalized"]
    columns = []
    all_match = True
    for column in range(RAW_FIRST_VALUE_COL, RAW_LAST_VALUE_COL + 1):
        letter = get_column_letter(column)
        excluded = column in NORMALIZATION_EXCLUDED_COLS
        if excluded:
            raw_values = []
            for row in range(RAW_FIRST_DATA_ROW, RAW_LAST_DATA_ROW + 1):
                raw_value = raw.cell(row, column).value
                if raw_value is None:
                    _record_expected_absent(
                        "E19b.production_excluded_raw_cell", 1
                    )
                elif (
                    isinstance(raw_value, (int, float))
                    and not isinstance(raw_value, bool)
                    and math.isfinite(float(raw_value))
                ):
                    raw_values.append(float(raw_value))
            invalid_cells = []
        else:
            raw_values = [
                number(raw.cell(row, column).value)
                for row in range(RAW_FIRST_DATA_ROW, RAW_LAST_DATA_ROW + 1)
            ]
            invalid_cells = [
                "%s%d" % (letter, row)
                for row, value in zip(
                    range(RAW_FIRST_DATA_ROW, RAW_LAST_DATA_ROW + 1),
                    raw_values,
                )
                if is_data_invalid(value)
            ]
        maximum = max(
            [0.0]
            + [value for value in raw_values if not is_data_invalid(value)]
        )
        deviations = []
        mismatch_cells = []
        if not excluded:
            for offset, row in enumerate(
                range(RAW_FIRST_DATA_ROW, RAW_LAST_DATA_ROW + 1)
            ):
                if is_data_invalid(raw_values[offset]):
                    continue
                expected = raw_values[offset] / maximum if maximum else 0.0
                actual = number(normalized.cell(row, column).value)
                if is_data_invalid(actual):
                    invalid_cells.append("%s%d" % (letter, row))
                    continue
                deviation = abs(actual - expected)
                deviations.append(deviation)
                if deviation > FLOAT_TOLERANCE and len(mismatch_cells) < 10:
                    mismatch_cells.append(
                        {
                            "cell": "%s%d" % (letter, row),
                            "raw_cell": "%s%d" % (letter, row),
                            "raw": compact_number(raw_values[offset]),
                            "expected": compact_number(expected),
                            "actual": compact_number(actual),
                            "abs_deviation": compact_number(deviation),
                        }
                    )
            verdict = (
                DATA_INVALID
                if invalid_cells
                else (
                    "MATCH"
                    if max(deviations or [0.0]) <= FLOAT_TOLERANCE
                    else "MISMATCH"
                )
            )
            all_match = all_match and verdict == "MATCH"
        else:
            verdict = "EXCLUDED-BY-PRODUCTION-RULE"
        columns.append(
            {
                "column": letter,
                "raw_header": raw.cell(1, column).value,
                "normalized_header": normalized.cell(1, column).value,
                "raw_max": compact_number(maximum),
                "excluded": excluded,
                "verdict": verdict,
                "data_invalid_cells": sorted(set(invalid_cells))[:50],
                "max_abs_deviation": (
                    None if excluded else compact_number(max(deviations or [0.0]))
                ),
                "mismatch_sample": mismatch_cells,
            }
        )
    return {
        "formula": "normalized[row,col] = raw[row,col] / max(0,max(raw[3:535,col])); zero when max is zero",
        "sampled_rows": "%d..%d" % (RAW_FIRST_DATA_ROW, RAW_LAST_DATA_ROW),
        "sampled_columns": "C:FD",
        "sampled_cells": (RAW_LAST_DATA_ROW - RAW_FIRST_DATA_ROW + 1)
        * (RAW_LAST_VALUE_COL - RAW_FIRST_VALUE_COL + 1),
        "normalization_excluded_columns": [
            get_column_letter(value) for value in sorted(NORMALIZATION_EXCLUDED_COLS)
        ],
        "columns": columns,
        "verdict": "MATCH" if all_match else "MISMATCH",
        "citation": "fec/legislation/repair_legislative_normalized.py:21-30,52-67",
    }


def interaction_records(ctx):
    result = {}
    for sheet_name in INTERACTION_SHEETS + SUCCESS_SHEETS:
        ws = ctx.model_wb[sheet_name]
        headers, rows = workbook_rows(ws)
        result[sheet_name] = {"headers": headers, "rows": rows, "ws": ws}
    return result


def row_signature(row, headers):
    values = []
    for header in headers:
        value = row.get(header)
        if isinstance(value, float):
            values.append(round(value, 12))
        else:
            values.append(norm_text(value) if value is not None else None)
    return tuple(values)


def candidate_vectors(ctx, records, members):
    per_role = {
        "sponsor": "Sponsor Interactions",
        "original": "Original Co-Sponsor Interaction",
        "cosponsor": "Co-Sponsor Interactions",
    }
    cfpb_direction = {}
    for row in ctx.cfpb_directions:
        key = bill_key(row.get("Bill"))
        label = norm_text(row.get("Category")).casefold()
        if key:
            cfpb_direction[key] = (
                "restrain"
                if label == "reduce"
                else ("reach" if label == "expand" else "neutral")
            )
    cfpb_restrain_reroute_targets = (
        ("SJRES18", bill_key("S.J.Res.18"), "119"),
        ("HJRES59", bill_key("H.J.Res.59"), "119"),
        ("HR6955", bill_key("H.R.6955"), "119"),
    )
    data = collections.defaultdict(lambda: collections.defaultdict(lambda: collections.defaultdict(list)))
    cfpb_split_data = collections.defaultdict(
        lambda: collections.defaultdict(lambda: collections.defaultdict(list))
    )
    for role, sheet_name in per_role.items():
        for row in records[sheet_name]["rows"]:
            member = ctx.resolve_member(row.get("Member"))
            category = canonical_category(row.get("Section"))
            if member and category:
                data[member][category][role].append(row)
                if category == "CFPB":
                    split = cfpb_direction.get(
                        bill_key(row.get("Name")), "neutral"
                    )
                    cfpb_split_data[member][split][role].append(row)
    success_data = collections.defaultdict(
        lambda: collections.defaultdict(lambda: collections.defaultdict(list))
    )
    cfpb_split_success = collections.defaultdict(
        lambda: collections.defaultdict(lambda: collections.defaultdict(list))
    )
    for role, sheet_name in (
        ("sponsor", "Success Sponsor"),
        ("original", "Success Original Co-Sponsor"),
    ):
        for row in records[sheet_name]["rows"]:
            member = ctx.resolve_member(row.get("Member"))
            category = canonical_category(row.get("Section"))
            if member and category:
                success_data[member][category][role].append(row)
                if category == "CFPB":
                    split = cfpb_direction.get(
                        bill_key(row.get("Name")), "neutral"
                    )
                    cfpb_split_success[member][split][role].append(row)

    # add_cutax_block.py:27-44,76-97 adds decayed LTR units only to
    # sponsor/original activity in these four blocks.  Retain both the base
    # interaction candidates and letter-augmented alternatives.
    letter_categories = {
        "CDFI Fund",
        "CU Tax Status",
        "CFPB",
        "Credit Card Mandates",
    }
    letter_units = collections.defaultdict(
        lambda: collections.defaultdict(lambda: collections.defaultdict(list))
    )
    for row in ctx.manual:
        if not norm_text(row.get("Bill")).upper().startswith("LTR-"):
            continue
        category = canonical_category(row.get("Category"))
        if category not in letter_categories:
            continue
        member = ctx.resolve_member(row.get("Member"))
        role = norm_text(row.get("Role")).casefold()
        if role == "sponsor":
            role_key = "sponsor"
        elif role == "original-cosponsor":
            role_key = "original"
        else:
            continue
        if member:
            letter_units[member][category][role_key].append(
                congress_decay(row.get("Congress"))
            )

    def metrics(rows, extra_letter_weights=()):
        weights = [number(row.get("DecayWeight"), 1.0) for row in rows]
        cosponsors = [number(row.get("CosponsorCount")) for row in rows]
        extra_letter_weights = list(extra_letter_weights)
        augmented_weights = weights + extra_letter_weights
        augmented_cosponsors = cosponsors + [0.0] * len(extra_letter_weights)
        return {
            "count": float(len(rows)),
            "sum_decay": sum(weights),
            "average_decay": statistics.fmean(weights) if weights else 0.0,
            "average_cosponsor_count": statistics.fmean(cosponsors)
            if cosponsors
            else 0.0,
            "decay_weighted_average_cosponsor_count": (
                sum(value * weight for value, weight in zip(cosponsors, weights))
                / sum(weights)
                if sum(weights)
                else 0.0
            ),
            "letter_augmented_count": float(len(rows) + len(extra_letter_weights)),
            "letter_augmented_sum_decay": sum(augmented_weights),
            "letter_augmented_average_decay": (
                statistics.fmean(augmented_weights)
                if augmented_weights
                else 0.0
            ),
            "letter_augmented_average_cosponsor_count": (
                statistics.fmean(augmented_cosponsors)
                if augmented_cosponsors
                else 0.0
            ),
            "letter_augmented_decay_weighted_average_cosponsor_count": (
                sum(
                    value * weight
                    for value, weight in zip(
                        augmented_cosponsors, augmented_weights
                    )
                )
                / sum(augmented_weights)
                if sum(augmented_weights)
                else 0.0
            ),
            "manual_letter_rows": float(len(extra_letter_weights)),
            "manual_letter_decay_units": sum(extra_letter_weights),
        }

    def add_candidate_with_rounding(candidate, name, value):
        candidate[name] = value
        candidate[name + "_rounded_1dp"] = round(value, 1)
        candidate[name + "_rounded_2dp"] = round(value, 2)

    def add_rate_candidates(
        candidate,
        prefix,
        success_count,
        eligible_count,
        success_weight,
        eligible_weight,
    ):
        fraction_count = success_count / eligible_count if eligible_count else 0.0
        fraction_decay = (
            success_weight / eligible_weight if eligible_weight else 0.0
        )
        values = {
            prefix + "_count_fraction": fraction_count,
            prefix + "_decay_weighted_fraction": fraction_decay,
            prefix + "_count_pct": 100.0 * fraction_count,
            prefix + "_decay_weighted_pct": 100.0 * fraction_decay,
        }
        for name, value in values.items():
            add_candidate_with_rounding(candidate, name, value)

    vectors = collections.defaultdict(dict)
    for member in members:
        for category in CATEGORIES:
            role_metrics = {
                role: metrics(
                    data[member][category][role],
                    letter_units[member][category][role]
                    if role in ("sponsor", "original")
                    else (),
                )
                for role in ("sponsor", "original", "cosponsor")
            }
            success_metrics = {
                role: metrics(success_data[member][category][role])
                for role in ("sponsor", "original")
            }
            candidate = {}
            for role in ("sponsor", "original", "cosponsor"):
                for name, value in role_metrics[role].items():
                    candidate["%s_%s" % (role, name)] = value
            # add_cutax_block.py:76-83 pastes a SUM of sponsor letter decay
            # units.  Keep an explicitly named producer-semantic candidate;
            # an average can coincide only while each member has <=1 letter.
            candidate["sponsor_letter_augmented_SUM_decay"] = role_metrics[
                "sponsor"
            ]["letter_augmented_sum_decay"]
            for role in ("sponsor", "original"):
                for name, value in success_metrics[role].items():
                    candidate["%s_success_%s" % (role, name)] = value

            for scope, roles in (
                ("sponsor_success_rate", ("sponsor",)),
                ("original_success_rate", ("original",)),
                ("sponsor_plus_original_success_rate", ("sponsor", "original")),
            ):
                success_count = sum(
                    success_metrics[role]["count"] for role in roles
                )
                success_weight = sum(
                    success_metrics[role]["sum_decay"] for role in roles
                )
                eligible_count = sum(
                    role_metrics[role]["count"] for role in roles
                )
                eligible_weight = sum(
                    role_metrics[role]["sum_decay"] for role in roles
                )
                add_rate_candidates(
                    candidate,
                    scope,
                    success_count,
                    eligible_count,
                    success_weight,
                    eligible_weight,
                )
                augmented_eligible_count = sum(
                    role_metrics[role]["letter_augmented_count"]
                    for role in roles
                )
                augmented_eligible_weight = sum(
                    role_metrics[role]["letter_augmented_sum_decay"]
                    for role in roles
                )
                add_rate_candidates(
                    candidate,
                    scope + "_letter_augmented",
                    success_count,
                    augmented_eligible_count,
                    success_weight,
                    augmented_eligible_weight,
                )

            # Generic category Success Rate columns can be the arithmetic mean
            # of the sponsor and original role rates rather than the pooled
            # numerator/denominator.  Test both empty-role conventions under
            # count and decay weighting, with and without manual-letter units.
            for augmented, suffix in (
                (False, ""),
                (True, "_letter_augmented"),
            ):
                for weighting, eligible_key, success_key in (
                    ("count", "count", "count"),
                    ("decay_weighted", "sum_decay", "sum_decay"),
                ):
                    if augmented:
                        eligible_key = "letter_augmented_" + eligible_key
                    role_rates = []
                    eligible_rates = []
                    for role in ("sponsor", "original"):
                        eligible = role_metrics[role][eligible_key]
                        successful = success_metrics[role][success_key]
                        rate = 100.0 * successful / eligible if eligible else 0.0
                        role_rates.append(rate)
                        if eligible:
                            eligible_rates.append(rate)
                    include_zero_mean = sum(role_rates) / 2.0
                    exclude_empty_mean = (
                        statistics.fmean(eligible_rates)
                        if eligible_rates
                        else 0.0
                    )
                    base_name = (
                        "mean_role_success_rate_%s_pct%s"
                        % (weighting, suffix)
                    )
                    add_candidate_with_rounding(
                        candidate,
                        base_name + "_empty_role_zero",
                        include_zero_mean,
                    )
                    add_candidate_with_rounding(
                        candidate,
                        base_name + "_empty_role_excluded",
                        exclude_empty_mean,
                    )

            if category == "CFPB":
                for block in ("restrain", "reach"):
                    for neutral_mode in ("both", "neither"):
                        split_role_metrics = {}
                        split_success_metrics = {}
                        for role in ("sponsor", "original", "cosponsor"):
                            rows = list(cfpb_split_data[member][block][role])
                            if neutral_mode == "both":
                                rows.extend(
                                    cfpb_split_data[member]["neutral"][role]
                                )
                            extra_letters = (
                                letter_units[member]["CFPB"][role]
                                if neutral_mode == "both"
                                and role in ("sponsor", "original")
                                else ()
                            )
                            split_role_metrics[role] = metrics(
                                rows, extra_letters
                            )
                        for role in ("sponsor", "original"):
                            rows = list(cfpb_split_success[member][block][role])
                            if neutral_mode == "both":
                                rows.extend(
                                    cfpb_split_success[member]["neutral"][role]
                                )
                            split_success_metrics[role] = metrics(rows)
                        split_prefix = "cfpb_%s_neutral_%s" % (
                            block,
                            neutral_mode,
                        )
                        for role in ("sponsor", "original", "cosponsor"):
                            for name, value in split_role_metrics[role].items():
                                candidate[
                                    "%s_%s_%s"
                                    % (split_prefix, role, name)
                                ] = value
                        for role in ("sponsor", "original"):
                            for name, value in split_success_metrics[role].items():
                                candidate[
                                    "%s_%s_success_%s"
                                    % (split_prefix, role, name)
                                ] = value
                        for scope, roles in (
                            ("sponsor_success_rate", ("sponsor",)),
                            ("original_success_rate", ("original",)),
                            (
                                "sponsor_plus_original_success_rate",
                                ("sponsor", "original"),
                            ),
                        ):
                            success_count = sum(
                                split_success_metrics[role]["count"]
                                for role in roles
                            )
                            success_weight = sum(
                                split_success_metrics[role]["sum_decay"]
                                for role in roles
                            )
                            eligible_count = sum(
                                split_role_metrics[role]["count"]
                                for role in roles
                            )
                            eligible_weight = sum(
                                split_role_metrics[role]["sum_decay"]
                                for role in roles
                            )
                            add_rate_candidates(
                                candidate,
                                split_prefix + "_" + scope,
                                success_count,
                                eligible_count,
                                success_weight,
                                eligible_weight,
                            )
                            add_rate_candidates(
                                candidate,
                                split_prefix
                                + "_"
                                + scope
                                + "_letter_augmented",
                                success_count,
                                sum(
                                    split_role_metrics[role][
                                        "letter_augmented_count"
                                    ]
                                    for role in roles
                                ),
                                success_weight,
                                sum(
                                    split_role_metrics[role][
                                        "letter_augmented_sum_decay"
                                    ]
                                    for role in roles
                                ),
                            )

                # The audit identified three 119th-Congress bills whose
                # hand-built Restrain-block membership may differ from the
                # direction file.  Test every inclusion combination so CU/CW
                # can name the exact workbook membership when one reproduces
                # the sheet; otherwise the closest residual remains explicit.
                for mask in range(1 << len(cfpb_restrain_reroute_targets)):
                    included = {
                        (target_key, target_congress)
                        for index, (_label, target_key, target_congress) in enumerate(
                            cfpb_restrain_reroute_targets
                        )
                        if mask & (1 << index)
                    }
                    included_labels = [
                        label
                        for index, (label, _target_key, _target_congress) in enumerate(
                            cfpb_restrain_reroute_targets
                        )
                        if mask & (1 << index)
                    ]
                    excluded_labels = [
                        label
                        for index, (label, _target_key, _target_congress) in enumerate(
                            cfpb_restrain_reroute_targets
                        )
                        if not mask & (1 << index)
                    ]
                    for neutral_mode in ("both", "neither"):
                        toggle_role_metrics = {}
                        toggle_success_metrics = {}
                        for role in ("sponsor", "original", "cosponsor"):
                            rows = []
                            for row in data[member]["CFPB"][role]:
                                row_identity = (
                                    bill_key(row.get("Name")),
                                    congress_text(row.get("Congress")),
                                )
                                if row_identity in {
                                    (target_key, target_congress)
                                    for _label, target_key, target_congress
                                    in cfpb_restrain_reroute_targets
                                }:
                                    keep = row_identity in included
                                else:
                                    split = cfpb_direction.get(
                                        bill_key(row.get("Name")), "neutral"
                                    )
                                    keep = split == "restrain" or (
                                        neutral_mode == "both"
                                        and split == "neutral"
                                    )
                                if keep:
                                    rows.append(row)
                            extra_letters = (
                                letter_units[member]["CFPB"][role]
                                if neutral_mode == "both"
                                and role in ("sponsor", "original")
                                else ()
                            )
                            toggle_role_metrics[role] = metrics(
                                rows, extra_letters
                            )
                        for role in ("sponsor", "original"):
                            rows = []
                            for row in success_data[member]["CFPB"][role]:
                                row_identity = (
                                    bill_key(row.get("Name")),
                                    congress_text(row.get("Congress")),
                                )
                                if row_identity in {
                                    (target_key, target_congress)
                                    for _label, target_key, target_congress
                                    in cfpb_restrain_reroute_targets
                                }:
                                    keep = row_identity in included
                                else:
                                    split = cfpb_direction.get(
                                        bill_key(row.get("Name")), "neutral"
                                    )
                                    keep = split == "restrain" or (
                                        neutral_mode == "both"
                                        and split == "neutral"
                                    )
                                if keep:
                                    rows.append(row)
                            toggle_success_metrics[role] = metrics(rows)
                        toggle_prefix = (
                            "cfpb_restrain_toggle_in_%s_out_%s_neutral_%s"
                            % (
                                "-".join(included_labels) or "NONE",
                                "-".join(excluded_labels) or "NONE",
                                neutral_mode,
                            )
                        )
                        for role in ("sponsor", "original", "cosponsor"):
                            for name, value in toggle_role_metrics[role].items():
                                candidate[
                                    "%s_%s_%s"
                                    % (toggle_prefix, role, name)
                                ] = value
                        for role in ("sponsor", "original"):
                            for name, value in toggle_success_metrics[role].items():
                                candidate[
                                    "%s_%s_success_%s"
                                    % (toggle_prefix, role, name)
                                ] = value
                        for scope, roles in (
                            ("sponsor_success_rate", ("sponsor",)),
                            ("original_success_rate", ("original",)),
                            (
                                "sponsor_plus_original_success_rate",
                                ("sponsor", "original"),
                            ),
                        ):
                            success_count = sum(
                                toggle_success_metrics[role]["count"]
                                for role in roles
                            )
                            success_weight = sum(
                                toggle_success_metrics[role]["sum_decay"]
                                for role in roles
                            )
                            eligible_count = sum(
                                toggle_role_metrics[role]["count"]
                                for role in roles
                            )
                            eligible_weight = sum(
                                toggle_role_metrics[role]["sum_decay"]
                                for role in roles
                            )
                            add_rate_candidates(
                                candidate,
                                toggle_prefix + "_" + scope,
                                success_count,
                                eligible_count,
                                success_weight,
                                eligible_weight,
                            )
                            add_rate_candidates(
                                candidate,
                                toggle_prefix
                                + "_"
                                + scope
                                + "_letter_augmented",
                                success_count,
                                sum(
                                    toggle_role_metrics[role][
                                        "letter_augmented_count"
                                    ]
                                    for role in roles
                                ),
                                success_weight,
                                sum(
                                    toggle_role_metrics[role][
                                        "letter_augmented_sum_decay"
                                    ]
                                    for role in roles
                                ),
                            )

            # Retain the original combined fraction names so prior evidence
            # can be diffed directly.
            candidate["success_rate_count"] = candidate[
                "sponsor_plus_original_success_rate_count_fraction"
            ]
            candidate["success_rate_decay_weighted"] = candidate[
                "sponsor_plus_original_success_rate_decay_weighted_fraction"
            ]
            vectors[(member, category)] = candidate
    return vectors


def identify_aggregate_columns(ctx, records):
    raw = ctx.model_wb["Legislative Score - Raw"]
    member_rows = {
        raw.cell(row, 1).value: row
        for row in range(RAW_FIRST_DATA_ROW, RAW_LAST_DATA_ROW + 1)
        if raw.cell(row, 1).value
    }
    members = list(member_rows)
    vectors = candidate_vectors(ctx, records, members)
    generic_success_columns = {
        column_index_from_string(start) + 5: classification
        for start, _end, classification in EXPECTED_BLOCKS
        if classification in CATEGORIES
    }
    columns = []
    for column in range(RAW_FIRST_VALUE_COL, RAW_LAST_VALUE_COL + 1):
        header = norm_text(raw.cell(1, column).value)
        category = canonical_category(header)
        if not category:
            continue
        observed = {
            raw.cell(row, 1).value: number(raw.cell(row, column).value)
            for row in range(RAW_FIRST_DATA_ROW, RAW_LAST_DATA_ROW + 1)
            if raw.cell(row, 1).value
        }
        observed_invalid_members = sorted(
            member
            for member, value in observed.items()
            if is_data_invalid(value)
        )
        sheet_candidate_vectors = {}
        if column in generic_success_columns:
            sibling_zero = {}
            sibling_excluded = {}
            for member in members:
                row = member_rows[member]
                sponsor_cell = raw.cell(row, column + 1).value
                original_cell = raw.cell(row, column + 2).value
                # This candidate deliberately mirrors the documented
                # "empty role treated as zero" formula alternative.  It is
                # not the generic numeric parser's behavior.
                sponsor_value = (
                    0.0
                    if sponsor_cell is None or not norm_text(sponsor_cell)
                    else number(sponsor_cell)
                )
                original_value = (
                    0.0
                    if original_cell is None or not norm_text(original_cell)
                    else number(original_cell)
                )
                sibling_zero[member] = (sponsor_value + original_value) / 2.0
                nonempty = [
                    number(value)
                    for value in (sponsor_cell, original_cell)
                    if value is not None and norm_text(value)
                ]
                sibling_excluded[member] = (
                    statistics.fmean(nonempty) if nonempty else 0.0
                )
            for name, values in (
                (
                    "sheet_sibling_mean_success_rate_empty_role_zero",
                    sibling_zero,
                ),
                (
                    "sheet_sibling_mean_success_rate_empty_role_excluded",
                    sibling_excluded,
                ),
            ):
                sheet_candidate_vectors[name] = values
                sheet_candidate_vectors[name + "_rounded_1dp"] = {
                    member: round(value, 1)
                    for member, value in values.items()
                }
                sheet_candidate_vectors[name + "_rounded_2dp"] = {
                    member: round(value, 2)
                    for member, value in values.items()
                }
        candidate_names = sorted(
            (
                {
                    name
                    for member in members
                    for name in vectors[(member, category)].keys()
                }
                if vectors
                else set()
            )
            | set(sheet_candidate_vectors)
        )
        scores = []
        for candidate in candidate_names:
            candidate_values = (
                sheet_candidate_vectors[candidate]
                if candidate in sheet_candidate_vectors
                else {
                    member: number(
                        vectors[(member, category)].get(candidate)
                    )
                    for member in members
                }
            )
            invalid_members = sorted(
                member
                for member in members
                if is_data_invalid(observed[member])
                or is_data_invalid(candidate_values[member])
            )
            valid_members = [
                member for member in members if member not in set(invalid_members)
            ]
            deviations = [
                abs(observed[member] - candidate_values[member])
                for member in valid_members
            ]
            rounded_2dp_deviations = [
                abs(observed[member] - round(candidate_values[member], 2))
                for member in valid_members
            ]
            rounded_1dp_deviations = [
                abs(observed[member] - round(candidate_values[member], 1))
                for member in valid_members
            ]
            max_deviation = max(deviations or [0.0])
            rounded_2dp_max_deviation = max(
                rounded_2dp_deviations or [0.0]
            )
            rounded_1dp_max_deviation = max(
                rounded_1dp_deviations or [0.0]
            )
            is_explicit_rounding = candidate.endswith(
                ("_rounded_1dp", "_rounded_2dp")
            )
            scores.append(
                {
                    "candidate": candidate,
                    "data_invalid_members": invalid_members,
                    "max_abs_deviation": max_deviation,
                    "sum_abs_deviation": sum(deviations),
                    "rounded_2dp_max_abs_deviation": rounded_2dp_max_deviation,
                    "rounded_1dp_max_abs_deviation": rounded_1dp_max_deviation,
                    "exact_match": (
                        not invalid_members
                        and max_deviation <= FLOAT_TOLERANCE
                    ),
                    "rounded_2dp_match": (
                        not invalid_members
                        and not is_explicit_rounding
                        and max_deviation <= 0.005 + 1e-12
                        and rounded_2dp_max_deviation <= FLOAT_TOLERANCE
                    ),
                    "rounded_1dp_match": (
                        not invalid_members
                        and not is_explicit_rounding
                        and "success_rate" in candidate
                        and max_deviation <= 0.05 + 1e-12
                        and rounded_1dp_max_deviation <= FLOAT_TOLERANCE
                    ),
                }
            )
        numeric_scores = sorted(
            (item for item in scores if not item["data_invalid_members"]),
            key=lambda item: (
                item["max_abs_deviation"],
                item["sum_abs_deviation"],
                item["candidate"],
            )
        )
        best = numeric_scores[0] if numeric_scores else None
        exact_matches = sorted(
            (
                item
                for item in scores
                if item["exact_match"]
                and not item["candidate"].endswith(
                    ("_rounded_1dp", "_rounded_2dp")
                )
            ),
            key=lambda item: (item["sum_abs_deviation"], item["candidate"]),
        )
        rounded_matches = sorted(
            (item for item in scores if item["rounded_2dp_match"]),
            key=lambda item: (
                item["max_abs_deviation"],
                item["sum_abs_deviation"],
                item["candidate"],
            ),
        )
        rounded_1dp_matches = sorted(
            (item for item in scores if item["rounded_1dp_match"]),
            key=lambda item: (
                item["max_abs_deviation"],
                item["sum_abs_deviation"],
                item["candidate"],
            ),
        )
        explicit_rounded_2dp_matches = sorted(
            (
                item
                for item in scores
                if item["exact_match"]
                and item["candidate"].endswith("_rounded_2dp")
            ),
            key=lambda item: (item["sum_abs_deviation"], item["candidate"]),
        )
        explicit_rounded_1dp_matches = sorted(
            (
                item
                for item in scores
                if item["exact_match"]
                and item["candidate"].endswith("_rounded_1dp")
            ),
            key=lambda item: (item["sum_abs_deviation"], item["candidate"]),
        )
        column_letter = get_column_letter(column)
        preferred_producer_candidate = {
            "EW": "sponsor_letter_augmented_SUM_decay",
        }.get(column_letter)
        preferred_producer_match = next(
            (
                item
                for item in scores
                if item["candidate"] == preferred_producer_candidate
                and item["exact_match"]
            ),
            None,
        )
        cfpb_target_role_fragment = {
            "CU": "_sponsor_",
            "CW": "_original_",
        }.get(column_letter)
        cfpb_toggle_scores = sorted(
            (
                item
                for item in scores
                if item["candidate"].startswith("cfpb_restrain_toggle_")
                and (
                    cfpb_target_role_fragment is None
                    or (
                        cfpb_target_role_fragment in item["candidate"]
                        and "_success_" not in item["candidate"]
                        and "success_rate" not in item["candidate"]
                    )
                )
            ),
            key=lambda item: (
                item["max_abs_deviation"],
                item["sum_abs_deviation"],
                item["candidate"],
            ),
        )
        cfpb_toggle_exact_matches = [
            item for item in cfpb_toggle_scores if item["exact_match"]
        ]
        if observed_invalid_members:
            identified = None
            identified_verdict = DATA_INVALID
        elif preferred_producer_match:
            identified = preferred_producer_match
            identified_verdict = "MATCH"
        elif column_letter in ("CU", "CW"):
            if cfpb_toggle_exact_matches:
                identified = cfpb_toggle_exact_matches[0]
                identified_verdict = "MATCH"
            else:
                identified = None
                identified_verdict = None
        elif exact_matches:
            identified = exact_matches[0]
            identified_verdict = "MATCH"
        elif rounded_matches:
            identified = rounded_matches[0]
            identified_verdict = "MATCH-ROUNDED-2DP"
        elif rounded_1dp_matches:
            identified = rounded_1dp_matches[0]
            identified_verdict = "MATCH-ROUNDED-1DP"
        elif explicit_rounded_2dp_matches:
            identified = explicit_rounded_2dp_matches[0]
            identified_verdict = "MATCH-ROUNDED-2DP"
        elif explicit_rounded_1dp_matches:
            identified = explicit_rounded_1dp_matches[0]
            identified_verdict = "MATCH-ROUNDED-1DP"
        else:
            identified = None
            identified_verdict = None
        best_candidate_name = best["candidate"] if best else ""
        best_candidate_percent_scale = (
            "_pct" in best_candidate_name
            or best_candidate_name.startswith(
                "sheet_sibling_mean_success_rate"
            )
        )
        unidentified_verdict = e19c_unidentified_verdict(
            best["max_abs_deviation"] if best else None,
            aggregate_like=(
                "average" in header.casefold()
                or "success" in header.casefold()
            ),
            percent_scale=best_candidate_percent_scale,
        )
        cfpb_reroute_arbitration = None
        if column_letter in ("CU", "CW"):
            closest_toggle = (
                cfpb_toggle_exact_matches[0]
                if cfpb_toggle_exact_matches
                else (cfpb_toggle_scores[0] if cfpb_toggle_scores else None)
            )
            residual_rows = []
            if closest_toggle:
                closest_name = closest_toggle["candidate"]
                closest_values = (
                    sheet_candidate_vectors[closest_name]
                    if closest_name in sheet_candidate_vectors
                    else {
                        member: number(
                            vectors[(member, category)].get(closest_name)
                        )
                        for member in members
                    }
                )
                for member in sorted(members):
                    residual = observed[member] - closest_values[member]
                    if abs(residual) > FLOAT_TOLERANCE:
                        residual_rows.append(
                            {
                                "member": member,
                                "sheet_minus_candidate": compact_number(
                                    residual
                                ),
                                "residual_in_119th_bill_units": compact_number(
                                    residual / congress_decay("119")
                                ),
                            }
                        )
            cfpb_reroute_arbitration = {
                "status": (
                    "EXACT-TOGGLE-COMBINATION"
                    if cfpb_toggle_exact_matches
                    else "NO-EXACT-TOGGLE-COMBINATION"
                ),
                "identified_combination": (
                    cfpb_toggle_exact_matches[0]["candidate"]
                    if cfpb_toggle_exact_matches
                    else None
                ),
                "closest_toggle_candidate": (
                    closest_toggle["candidate"] if closest_toggle else None
                ),
                "known_reroute_bills": [
                    "S.J.Res.18 C119",
                    "H.J.Res.59 C119",
                    "H.R.6955 C119",
                ],
                "tested_combinations": 8,
                "neutral_membership_modes_per_combination": [
                    "both",
                    "neither",
                ],
                "current_congress_bill_unit": compact_number(
                    congress_decay("119")
                ),
                "nonzero_residual_cells": len(residual_rows),
                "max_abs_residual": compact_number(
                    max(
                        [
                            abs(number(row["sheet_minus_candidate"]))
                            for row in residual_rows
                        ]
                        or [0.0]
                    )
                ),
                "distinct_signed_residuals": sorted(
                    {
                        row["sheet_minus_candidate"]
                        for row in residual_rows
                    },
                    key=number,
                ),
                "residual_sample": residual_rows[:10],
            }
        header_low = header.casefold()
        aggregate_like = "average" in header_low or "success" in header_low
        letter_counts_by_role = {
            role: {
                member: vectors[(member, category)].get(
                    "%s_manual_letter_rows" % role,
                    0.0,
                )
                for member in members
            }
            for role in ("sponsor", "original")
        }
        manual_letter_rows_for_category = int(
            sum(
                sum(values.values())
                for values in letter_counts_by_role.values()
            )
        )
        identified_name = identified["candidate"] if identified else None
        if identified_name and identified_name.startswith("sponsor_"):
            identified_letter_roles = ("sponsor",)
        elif identified_name and identified_name.startswith("original_"):
            identified_letter_roles = ("original",)
        else:
            identified_letter_roles = ("sponsor", "original")
        identified_letter_counts = {
            member: sum(
                letter_counts_by_role[role][member]
                for role in identified_letter_roles
            )
            for member in members
        }
        identified_candidate_letter_rows = int(
            sum(identified_letter_counts.values())
        )
        max_letter_rows_per_member = int(
            max(list(identified_letter_counts.values()) or [0.0])
        )
        if identified_name and "letter_augmented" in identified_name:
            if identified_candidate_letter_rows == 0:
                letter_augmentation_note = (
                    "VACUOUS qualifier: no manual letter units exist for "
                    "this category, so letter-augmented equals base."
                )
            elif max_letter_rows_per_member <= 1:
                letter_augmentation_note = (
                    "SUM/average degeneracy: no member has a second letter; "
                    "a second letter would split the candidates."
                )
            else:
                letter_augmentation_note = (
                    "Non-degenerate letter augmentation: at least one member "
                    "has multiple letter units."
                )
        else:
            letter_augmentation_note = None
        columns.append(
            {
                "column": column_letter,
                "header": header,
                "category": category,
                "aggregate_like": aggregate_like,
                "identified_source": (
                    identified_name
                ),
                "letter_augmentation_semantics": {
                    "manual_letter_rows_for_category": manual_letter_rows_for_category,
                    "identified_candidate_roles": list(
                        identified_letter_roles
                    ),
                    "identified_candidate_letter_rows": (
                        identified_candidate_letter_rows
                    ),
                    "identified_candidate_max_letter_rows_per_member": (
                        max_letter_rows_per_member
                    ),
                    "note": letter_augmentation_note,
                },
                "identification_max_abs_deviation": (
                    compact_number(identified["max_abs_deviation"])
                    if identified
                    else None
                ),
                "identification_rounded_2dp_max_abs_deviation": (
                    compact_number(
                        identified["rounded_2dp_max_abs_deviation"]
                    )
                    if identified
                    else None
                ),
                "identification_rounded_1dp_max_abs_deviation": (
                    compact_number(
                        identified["rounded_1dp_max_abs_deviation"]
                    )
                    if identified
                    else None
                ),
                "best_attempt": (
                    {
                        "candidate": best["candidate"],
                        "scale": (
                            "percent"
                            if best_candidate_percent_scale
                            else "raw"
                        ),
                        "near_miss_max_deviation": (
                            E19C_PERCENT_NEAR_MISS_MAX
                            if best_candidate_percent_scale
                            else E19C_RAW_NEAR_MISS_MAX
                        ),
                        "max_abs_deviation": compact_number(
                            best["max_abs_deviation"]
                        ),
                        "sum_abs_deviation": compact_number(
                            best["sum_abs_deviation"]
                        ),
                    }
                    if best
                    else None
                ),
                "verdict": (
                    identified_verdict
                    if identified or identified_verdict == DATA_INVALID
                    else unidentified_verdict
                ),
                "attempted_candidates": candidate_names,
                "data_invalid_observed_members": observed_invalid_members,
                "data_invalid_candidates": [
                    {
                        "candidate": item["candidate"],
                        "members": item["data_invalid_members"],
                    }
                    for item in scores
                    if item["data_invalid_members"]
                ],
                "cfpb_reroute_arbitration": cfpb_reroute_arbitration,
            }
        )
    return columns


def e19c(ctx):
    records = interaction_records(ctx)
    subset_checks = {}
    source_for_success = {
        "Success Sponsor": "Sponsor Interactions",
        "Success Original Co-Sponsor": "Original Co-Sponsor Interaction",
    }
    for success_name, source_name in source_for_success.items():
        success = records[success_name]
        source = records[source_name]
        source_advanced = [
            row
            for row in source["rows"]
            if ADVANCE_WEIGHT.get(norm_text(row.get("LastAction")), 0.0) > 0
        ]
        headers = success["headers"]
        actual_counter = collections.Counter(
            row_signature(row, headers) for row in success["rows"]
        )
        source_counter = collections.Counter(
            row_signature(row, headers) for row in source_advanced
        )
        subset_checks[success_name] = {
            "source_sheet": source_name,
            "actual_rows": len(success["rows"]),
            "rederived_advanced_rows": len(source_advanced),
            "row_multiset_match": actual_counter == source_counter,
            "missing_from_success": sum((source_counter - actual_counter).values()),
            "extra_in_success": sum((actual_counter - source_counter).values()),
            "verdict": "MATCH" if actual_counter == source_counter else "MISMATCH",
        }
    aggregate_columns = identify_aggregate_columns(ctx, records)
    aggregate_like = [row for row in aggregate_columns if row["aggregate_like"]]
    match_family = ("MATCH", "MATCH-ROUNDED-1DP", "MATCH-ROUNDED-2DP")
    unresolved_family = (
        "UNIDENTIFIED-RAW-MEASURE",
        "NOT-RECOMPUTABLE-OFFLINE",
    )
    source_contradictions = [
        {
            "source": name,
            "verdict": value["verdict"],
            "reason": "success-sheet row multiset diverges from its identified interaction-sheet source",
        }
        for name, value in sorted(subset_checks.items())
        if value["verdict"] == "MISMATCH"
    ]
    column_contradictions = [
        {
            "column": value["column"],
            "header": value["header"],
            "verdict": value["verdict"],
        }
        for value in aggregate_columns
        if value["verdict"] not in match_family + unresolved_family
        and value["verdict"] != DATA_INVALID
    ]
    contradictions = source_contradictions + column_contradictions
    identified_columns = [
        value for value in aggregate_columns if value["verdict"] in match_family
    ]
    unresolved_columns = [
        value
        for value in aggregate_columns
        if value["verdict"] in unresolved_family
    ]
    data_invalid_columns = [
        value
        for value in aggregate_columns
        if value["verdict"] == DATA_INVALID
    ]
    part_verdict = e19c_part_verdict(
        data_invalid_columns,
        contradictions,
        unresolved_columns,
    )
    return {
        "sheet_structures": {
            name: sheet_structure(records[name]["ws"])
            for name in INTERACTION_SHEETS + SUCCESS_SHEETS
        },
        "success_sheet_rederivation": subset_checks,
        "raw_aggregate_columns": aggregate_like,
        "all_raw_measure_columns": aggregate_columns,
        "identification_counts": {
            "total_measure_columns": len(aggregate_columns),
            "identified": len(identified_columns),
            "unresolved": len(unresolved_columns),
            "contradictions": len(contradictions),
            "data_invalid": len(data_invalid_columns),
            "identified_by_verdict": dict(
                sorted(
                    collections.Counter(
                        value["verdict"] for value in identified_columns
                    ).items()
                )
            ),
            "unresolved_by_verdict": dict(
                sorted(
                    collections.Counter(
                        value["verdict"] for value in unresolved_columns
                    ).items()
                )
            ),
        },
        "contradictions": contradictions,
        "verdict": part_verdict,
        "method": "Candidate formulas test role count/sum decay/average decay/average CosponsorCount/decay-weighted average CosponsorCount; sponsor/original/combined and mean-of-role success rates under count/decay weighting with both empty-role conventions on 0..1 and 0..100 scales; generic Success Rate means from the sheet's own adjacent role-rate cells; exact, 1dp, and 2dp materialization; base versus add_cutax letter augmentation; CFPB Reduce/Expand/Neutral direction splits; and every in/out combination of the three known 119th-Congress Restrain re-route bills. Column EW uses the producer-semantic sponsor_letter_augmented_SUM_decay candidate because fec/legislation/add_cutax_block.py:76-83 pastes a SUM; it currently equals the average only because no member has a second letter, and a second letter would split them. Letter-augmented matches in categories with no letter rows are labeled vacuous rather than treated as independent evidence. A best nonzero deviation up to 0.5 on percent-scale candidates or 0.05 on raw-scale candidates is NEAR-MISS-DIVERGENT and is a contradiction; larger residuals, including today's CU=0.64 and CW=0.8, remain unresolved. Unidentified local measures remain explicit and make this part PARTIALLY-IDENTIFIED.",
        "citations": [
            "fec/legislation/add_cutax_block.py:27-44,76-97",
            "fec/stance/say_do_gap.py:36-52",
        ],
    }


def e19d(ctx):
    raw = ctx.model_wb["Legislative Score - Raw"]
    normalized = ctx.model_wb["Legislative Score - Normalized"]
    raw_shape = (raw.max_row, raw.max_column)
    normalized_shape = (normalized.max_row, normalized.max_column)
    compared_columns = max(raw_shape[1], normalized_shape[1])
    weight_mismatches = []
    data_invalid_weight_cells = []
    for column in range(1, compared_columns + 1):
        raw_value = raw.cell(2, column).value
        normalized_value = normalized.cell(2, column).value
        if is_invalid_numeric_input(raw_value):
            data_invalid_weight_cells.append(
                "%s2" % get_column_letter(column)
            )
        if is_invalid_numeric_input(normalized_value):
            data_invalid_weight_cells.append(
                "%s2" % get_column_letter(column)
            )
        if raw_value != normalized_value:
            weight_mismatches.append(
                {
                    "raw_cell": "%s2" % get_column_letter(column),
                    "normalized_cell": "%s2" % get_column_letter(column),
                    "raw": compact_number(raw_value),
                    "normalized": compact_number(normalized_value),
                }
            )
    required_weight_findings = []
    for cell in ("P2", "X2", "DG2"):
        required_weight_findings.append(
            {
                "cell": cell,
                "raw": compact_number(raw[cell].value),
                "normalized": compact_number(normalized[cell].value),
                "observed_mismatch": raw[cell].value != normalized[cell].value,
            }
        )
    cp1 = {
        "raw_cell": "CP1",
        "raw_value": raw["CP1"].value,
        "normalized_cell": "CP1",
        "normalized_value": normalized["CP1"].value,
        "observed_anomaly": (
            raw["CP1"].value == "Column17"
            and normalized["CP1"].value
            == "CDFI Fund - Original Co-Sponsorship - Average"
        ),
    }
    structural_anomalies = []
    if raw_shape != normalized_shape:
        structural_anomalies.append({
            "finding_key": "SHEET-DIMENSIONS",
            "raw_shape": list(raw_shape),
            "normalized_shape": list(normalized_shape),
        })
    for column in range(1, compared_columns + 1):
        raw_header = raw.cell(1, column).value
        normalized_header = normalized.cell(1, column).value
        if raw_header != normalized_header:
            cell = "%s1" % get_column_letter(column)
            structural_anomalies.append({
                "finding_key": cell,
                "raw_cell": cell,
                "normalized_cell": cell,
                "raw": raw_header,
                "normalized": normalized_header,
            })
    finding_classification = e19d_classify_findings(
        cp1["observed_anomaly"],
        required_weight_findings,
        weight_mismatches,
        structural_anomalies,
        data_invalid_weight_cells,
    )
    reserved_legacy = {
        "EI:EM": {
            "classification": "institutional columns",
            "headers": [
                {
                    "cell": "%s1" % get_column_letter(column),
                    "raw": raw.cell(1, column).value,
                    "normalized": normalized.cell(1, column).value,
                }
                for column in range(139, 144)
            ],
        },
        "EN": {
            "classification": "reserved; different MIN-MAX rule and excluded from value/MAX repair",
            "headers": {
                "raw": raw["EN1"].value,
                "normalized": normalized["EN1"].value,
            },
        },
        "EO:EV": {
            "classification": "legacy CFPB Reach block",
            "headers": [
                {
                    "cell": "%s1" % get_column_letter(column),
                    "raw": raw.cell(1, column).value,
                    "normalized": normalized.cell(1, column).value,
                }
                for column in range(145, 153)
            ],
        },
        "normalization_exceptions": [
            {
                "column": get_column_letter(column),
                "column_number": column,
                "raw_header": raw.cell(1, column).value,
                "normalized_header": normalized.cell(1, column).value,
            }
            for column in sorted(NORMALIZATION_EXCLUDED_COLS)
        ],
    }
    sheet_totals = {}
    for name in INTERACTION_SHEETS + SUCCESS_SHEETS:
        ws = ctx.model_wb[name]
        headers, rows = workbook_rows(ws)
        sheet_totals[name] = {
            "data_rows": len(rows),
            "sum_DecayWeight": compact_number(
                sum(number(row.get("DecayWeight")) for row in rows)
            ),
            "distinct_members": len(
                {norm_text(row.get("Member")) for row in rows if row.get("Member")}
            ),
            "distinct_bill_instances": len(
                {
                    (norm_text(row.get("Name")), congress_text(row.get("Congress")))
                    for row in rows
                    if row.get("Name")
                }
            ),
            "headers": headers,
        }
    return {
        "sheet_structures": {
            "Legislative Score - Raw": sheet_structure(raw),
            "Legislative Score - Normalized": sheet_structure(normalized),
        },
        "expected_block_map": [
            {"start": start, "end": end, "classification": classification}
            for start, end, classification in EXPECTED_BLOCKS
        ],
        "first_class_findings": {
            "raw_CP1_Column17": cp1,
            "row_2_weight_drift_required_cells": required_weight_findings,
            "all_row_2_weight_mismatches": weight_mismatches,
            "all_structural_anomalies": structural_anomalies,
            "known_finding_cells": ["CP1", "P2", "X2", "DG2"],
            "new_findings": finding_classification["new_findings"],
            "unexpected_row_2_weight_mismatches": finding_classification[
                "unexpected_row_2_weight_mismatches"
            ],
            "unexpected_structural_anomalies": finding_classification[
                "unexpected_structural_anomalies"
            ],
            "reserved_and_legacy_columns": reserved_legacy,
            "data_invalid_weight_cells": sorted(
                set(data_invalid_weight_cells)
            ),
        },
        "interaction_and_success_totals": sheet_totals,
        "verdict": finding_classification["verdict"],
        "citation": "fec/legislation/repair_legislative_normalized.py:21-30,52-67",
    }


def claim_e19(ctx):
    part_a = e19a(ctx)
    part_b = e19b(ctx)
    part_c = e19c(ctx)
    part_d = e19d(ctx)
    e19a_verdicts = [
        part_a["percentrank_vs_ranks_python_pcts_leg"]["verdict"],
    ] + [
        part_a[key]["verdict"]
        for key in (
            "cache_vs_snapshot_O",
            "cache_vs_snapshot_E",
            "cache_vs_snapshot_F",
        )
    ]
    nested_verdicts = {
        "E19a": (
            DATA_INVALID
            if DATA_INVALID in e19a_verdicts
            else (
                "MATCH"
                if e19a_verdicts[0] == "MATCH"
                and all(
                    value in ("MATCH", "EXPECTED-CACHE-DEAD")
                    for value in e19a_verdicts[1:]
                )
                else "MISMATCH"
            )
        ),
        "E19b": part_b["verdict"],
        "E19c": part_c["verdict"],
        "E19d": part_d["verdict"],
    }
    overall = e19_overall_verdict(nested_verdicts)
    return evidence_entry(
        "Workbook legislative aggregate reconciliation: O/E/F authority, raw normalization, success/average derivations, and structural findings",
        {
            "requirements": [
                "E19a three-way O consistency and exact percentrank",
                "E19b all-member C:FD raw-to-normalized reconciliation",
                "E19c success sheets plus every identifiable rate/average",
                "E19d all other aggregates and structural anomalies",
            ],
            "known_findings": [
                "Raw CP1 is Column17 while Normalized CP1 names the CDFI original-cosponsor average",
                "row-2 weight drift at P2, X2, DG2",
                "EI:EM institutional, EN reserved, EO:EV legacy; production normalization exceptions EN/EP/ER/ET/EU/EV",
            ],
            "citation": REPORT_CITATION + " §Q3",
        },
        {"E19a": part_a, "E19b": part_b, "E19c": part_c, "E19d": part_d},
        {
            "part_verdicts": nested_verdicts,
            "nonmatching_parts": [
                key
                for key, value in nested_verdicts.items()
                if value not in ("MATCH", "FINDINGS-CONFIRMED")
            ],
        },
        overall,
        "All workbook reads use data_only=True, read_only=False on the verified scratch copy. E19a mirrors fec/engine/compute_pillars.py:32-61 and exact percentrank lines 44-48. E19b mirrors fec/legislation/repair_legislative_normalized.py:21-30,52-67. E19c identifies producers by exhaustive local candidate comparison; MATCH-family means identified and reproduced, PARTIALLY-IDENTIFIED means at least one local measure remains honestly unresolved with no contradiction, and NEAR-MISS-DIVERGENT is a contradiction. E19d treats CP1 and P2/X2/DG2 as the known finding set; any additional row-2 weight mismatch or Raw/Normalized structural anomaly is listed under new_findings, returns NEW-FINDINGS for the part, and is non-matching for E19 overall.",
    )


def excluded_title(title, summary=""):
    phrases = (
        "george floyd",
        "save money on auto repair",
        "auto repair transportation",
        "accessible, affordable internet",
        "affordable internet for all",
        "health and location data",
        "fourth amendment is not for sale",
        "app store accountability",
        "safer gaming",
        "protecting military servicemembers",
        "data from foreign adversaries",
        "cease orwellian surveillance",
    )
    text = (str(title or "") + " " + str(summary or "")).casefold()
    return next((phrase for phrase in phrases if phrase in text), None)


def override_included(row):
    if row is None:
        return True
    value = norm_text(row.get("Include")).casefold()
    return value not in ("0", "false", "no", "exclude", "excluded")


def e20_bills(ctx):
    overrides = override_index(ctx.bill_overrides)
    role_map = {
        "Sponsor Interactions": "sponsor",
        "Original Co-Sponsor Interaction": "original",
        "Co-Sponsor Interactions": "cosponsor",
    }
    bills = collections.defaultdict(list)
    unmatched = set()
    excluded = 0
    for sheet_name, role in role_map.items():
        _, rows = workbook_rows(ctx.model_wb[sheet_name])
        for row in rows:
            member = ctx.resolve_member(row.get("Member"))
            if not member:
                unmatched.add(norm_text(row.get("Member")))
                continue
            title = norm_text(row.get("Title"))
            summary = norm_text(row.get("bill_summary"))
            override = find_override(
                overrides, row.get("Name"), row.get("Congress")
            )
            if excluded_title(title, summary) or not override_included(override):
                excluded += 1
                continue
            category = canonical_category(row.get("Section")) or norm_text(
                row.get("Section")
            )
            if override and override.get("Category"):
                category = canonical_category(override.get("Category")) or category
            bills[member].append(
                {
                    "role": role,
                    "area": category,
                    "bill": norm_text(row.get("Name")),
                    "congress": congress_text(row.get("Congress")),
                    "action": norm_text(row.get("LastAction")),
                }
            )
    for row in ctx.manual:
        member = ctx.resolve_member(row.get("Member"))
        if not member:
            unmatched.add(norm_text(row.get("Member")))
            continue
        if excluded_title(row.get("Title"), row.get("Summary")):
            excluded += 1
            continue
        raw_role = norm_text(row.get("Role")).casefold()
        role = (
            "sponsor"
            if raw_role == "sponsor"
            else ("original" if "original" in raw_role else "cosponsor")
        )
        bills[member].append(
            {
                "role": role,
                "area": canonical_category(row.get("Category"))
                or norm_text(row.get("Category")),
                "bill": norm_text(row.get("Bill")),
                "congress": congress_text(row.get("Congress")),
                "action": norm_text(row.get("LastAction")),
            }
        )
    order = {"sponsor": 0, "original": 1, "cosponsor": 2}
    collapsed = {}
    for member, rows in bills.items():
        groups = {}
        for row in rows:
            key = (row["bill"], str(row["congress"]))
            current = groups.get(key)
            if current is None:
                current = dict(row)
                current["areas"] = []
                groups[key] = current
            if order.get(row["role"], 9) < order.get(current["role"], 9):
                current["role"] = row["role"]
            if row["area"] and row["area"] not in current["areas"]:
                current["areas"].append(row["area"])
        for row in groups.values():
            row["areas"].sort()
        collapsed[member] = sorted(
            groups.values(),
            key=lambda row: (
                order.get(row["role"], 9),
                str(row["congress"]),
                row["areas"][0] if row["areas"] else "",
            ),
        )
    return collapsed, sorted(unmatched), excluded


def claim_e20(ctx):
    bills, unmatched, excluded = e20_bills(ctx)
    rederived = {}
    for member in ctx.roster:
        member_bills = bills.get(member, [])
        do_score = sum(
            SAY_DO_ROLE_WEIGHT.get(row.get("role"), 0.25)
            * ADVANCE_WEIGHT.get(row.get("action"), 0.0)
            for row in member_bills
        )
        sponsored = [row for row in member_bills if row.get("role") == "sponsor"]
        advanced = [
            row
            for row in sponsored
            if ADVANCE_WEIGHT.get(row.get("action"), 0.0) > 0
        ]
        rederived[member] = {
            "do_score": round(do_score, 3),
            "sponsored_cu": len(sponsored),
            "advanced_cu": len(advanced),
            "adv_rate": (
                round(len(advanced) / len(sponsored), 3)
                if len(sponsored) >= 3
                else None
            ),
        }
    artifact = ctx.say_do.get("members", {})
    fields = ("do_score", "sponsored_cu", "advanced_cu", "adv_rate")
    mismatches = []
    invalid_values = []
    compared = 0
    per_field_max = {field: 0.0 for field in fields}
    for member in sorted(set(rederived) | set(artifact)):
        expected = rederived.get(member)
        observed = artifact.get(member)
        if expected is None or observed is None:
            mismatches.append(
                {
                    "member": member,
                    "kind": "missing_member",
                    "rederived_present": expected is not None,
                    "artifact_present": observed is not None,
                }
            )
            continue
        compared += 1
        for field in fields:
            left = expected.get(field)
            right = observed.get(field)
            if left is None or right is None:
                match = left is None and right is None
                deviation = None
            else:
                left_number = number(left)
                right_number = number(right)
                if is_data_invalid(left_number) or is_data_invalid(right_number):
                    invalid_values.append(
                        {
                            "member": member,
                            "field": field,
                            "rederived": compact_number(left_number),
                            "artifact": compact_number(right_number),
                        }
                    )
                    match = False
                    deviation = float("nan")
                else:
                    deviation = abs(left_number - right_number)
                if not is_data_invalid(deviation):
                    per_field_max[field] = max(per_field_max[field], deviation)
                    match = deviation <= FLOAT_TOLERANCE
            if not match:
                mismatches.append(
                    {
                        "member": member,
                        "field": field,
                        "rederived": left,
                        "artifact": right,
                        "abs_deviation": compact_number(deviation),
                    }
                )
    verdict = (
        DATA_INVALID
        if invalid_values
        else ("MATCH" if not mismatches else "DRIFT")
    )
    return evidence_entry(
        "say_do_gap.json advancement and success-rate internal consistency",
        {
            "expectation": "all per-member do_score, sponsored_cu, advanced_cu, and adv_rate values reproduce",
            "citation": REPORT_CITATION + " §Q3",
        },
        {
            "compared_members": compared,
            "rederived_members": len(rederived),
            "artifact_members": len(artifact),
            "mismatch_count": len(mismatches),
            "mismatch_sample": mismatches[:50],
            "data_invalid_values": invalid_values[:50],
            "per_field_max_abs_deviation": {
                key: compact_number(value) for key, value in per_field_max.items()
            },
            "unmatched_source_members": unmatched,
            "excluded_source_rows": excluded,
        },
        {
            "mismatch_count": len(mismatches),
            "member_count_delta": len(rederived) - len(artifact),
        },
        verdict,
        "Rebuilds CU desk bills from the three interaction sheets plus manual evidence, then collapses by member/(bill,Congress) with sponsor>original>cosponsor exactly as fec/surfaces/make_member_desk.py:907-973. Applies ROLE_W/ADV_W and the >=3 sponsor rate floor from fec/stance/say_do_gap.py:36-52.",
    )


def build_claims(ctx):
    builders = (
        ("E1", claim_e1),
        ("E2", claim_e2),
        ("E3", claim_e3),
        ("E4", claim_e4),
        ("E5", claim_e5),
        ("E6", claim_e6),
        ("E7", claim_e7),
        ("E8", claim_e8),
        ("E9", claim_e9),
        ("E10", claim_e10),
        ("E11", claim_e11),
        ("E12", claim_e12),
        ("E13", claim_e13),
        ("E14", claim_e14),
        ("E15", claim_e15),
        ("E16", claim_e16),
        ("E17", claim_e17),
        ("E18", claim_e18),
        ("E19", claim_e19),
        ("E20", claim_e20),
    )
    claims = {}
    for identifier, builder in builders:
        print("RECOMPUTE %s" % identifier)
        invalid_before = _DATA_INVALID_COUNTS.copy()
        expected_absent_before = _EXPECTED_ABSENT_COUNTS.copy()
        entry = builder(ctx)
        invalid_after = _DATA_INVALID_COUNTS.copy()
        expected_absent_after = _EXPECTED_ABSENT_COUNTS.copy()
        expected_absent_delta = {
            reason: expected_absent_after[reason] - expected_absent_before[reason]
            for reason in sorted(expected_absent_after)
            if expected_absent_after[reason] > expected_absent_before[reason]
        }
        if expected_absent_delta:
            recomputed = dict(entry["recomputed"])
            recomputed["expected_absent_skipped"] = expected_absent_delta
            entry["recomputed"] = recomputed
        invalid_delta = {
            reason: invalid_after[reason] - invalid_before[reason]
            for reason in sorted(invalid_after)
            if invalid_after[reason] > invalid_before[reason]
        }
        if invalid_delta:
            recomputed = dict(entry["recomputed"])
            recomputed["data_invalid_numeric_inputs"] = invalid_delta
            delta = dict(entry["delta"])
            delta["data_invalid_flag"] = DATA_INVALID
            entry["recomputed"] = recomputed
            entry["delta"] = delta
            entry["verdict"] = DATA_INVALID
            entry["method_note"] += (
                " At least one numeric input was missing, non-finite, an "
                "Excel error, or otherwise non-numeric; poisoned values can "
                "never compare as MATCH."
            )
            entry = sanitize_evidence_value(entry)
        if set(entry) != {
            "claim",
            "reported",
            "recomputed",
            "delta",
            "verdict",
            "method_note",
        }:
            raise EvidenceError("%s entry does not have exactly the six required fields" % identifier)
        claims[identifier] = entry
    return claims


def md_cell(value, limit=180):
    if isinstance(value, (dict, list)):
        text = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    elif value is None:
        text = "null"
    else:
        text = str(value)
    text = text.replace("|", "\\|").replace("\r", " ").replace("\n", " ")
    return text if len(text) <= limit else text[: limit - 1] + "…"


def render_markdown(pack):
    claims = pack["claims"]
    lines = [
        "# Legislative Evidence Recompute — D1",
        "",
        "Deterministic offline evidence pack for the quantitative claims in the 2026-07-24 legislative data-plane audit. Inputs were read only from hash-verified scratch copies.",
        "",
        "Manifest exception: %s." % pack["manifest_exception"],
        "",
        "Code scan fingerprint (combined sha256 over sorted path/file-sha256 pairs): `%s`."
        % pack["code_scan_fingerprint"],
        "",
        "## Summary delta table",
        "",
        "| ID | Claim | Reported | Recomputed | Delta | Verdict |",
        "|---|---|---|---|---|---|",
    ]
    for identifier in sorted(claims, key=lambda value: int(value[1:])):
        entry = claims[identifier]
        lines.append(
            "| %s | %s | %s | %s | %s | %s |"
            % (
                identifier,
                md_cell(entry["claim"]),
                md_cell(entry["reported"]),
                md_cell(entry["recomputed"]),
                md_cell(entry["delta"]),
                md_cell(entry["verdict"]),
            )
        )
    lines.extend(
        [
            "",
            "Verdict rules: integer claims require exact equality. Percentages match within ±0.05 percentage points; rho matches within ±0.005. DRIFT is a finding, not a runtime error. E18 is intentionally split-verdict evidence, not a full fresh recompute.",
            "",
        ]
    )
    for identifier in sorted(claims, key=lambda value: int(value[1:])):
        entry = claims[identifier]
        lines.extend(
            [
                "## %s — %s" % (identifier, entry["claim"]),
                "",
                "- Verdict: `%s`" % entry["verdict"],
                "- Method: %s" % entry["method_note"],
                "",
                "### Reported",
                "",
                "```json",
                json.dumps(
                    entry["reported"],
                    ensure_ascii=False,
                    sort_keys=True,
                    indent=2,
                    allow_nan=False,
                ),
                "```",
                "",
            ]
        )
        if identifier == "E18":
            lines.extend(
                [
                    "### Recomputed (pillar side)",
                    "",
                    "```json",
                    json.dumps(
                        {
                            "E18a_PILLAR_SIDE_RECOMPUTED": entry[
                                "recomputed"
                            ]["E18a_PILLAR_SIDE_RECOMPUTED"],
                            "multipliers": entry["recomputed"]["multipliers"],
                        },
                        ensure_ascii=False,
                        sort_keys=True,
                        indent=2,
                        allow_nan=False,
                    ),
                    "```",
                    "",
                    "### Cached constant (ACU side, 2026-07-24 pull)",
                    "",
                    "```json",
                    json.dumps(
                        entry["recomputed"][
                            "E18b_ACU_SIDE_CACHED_CONSTANT"
                        ],
                        ensure_ascii=False,
                        sort_keys=True,
                        indent=2,
                        allow_nan=False,
                    ),
                    "```",
                    "",
                ]
            )
        else:
            lines.extend(
                [
                    "### Recomputed",
                    "",
                    "```json",
                    json.dumps(
                        entry["recomputed"],
                        ensure_ascii=False,
                        sort_keys=True,
                        indent=2,
                        allow_nan=False,
                    ),
                    "```",
                    "",
                ]
            )
        lines.extend(
            [
                "### Delta",
                "",
                "```json",
                json.dumps(
                    entry["delta"],
                    ensure_ascii=False,
                    sort_keys=True,
                    indent=2,
                    allow_nan=False,
                ),
                "```",
                "",
            ]
        )
    return "\n".join(lines).rstrip() + "\n"


def run_evidence(data_root, out):
    manifest, paths = verified_scratch_paths(data_root, out)
    ctx = EvidenceContext(paths)
    try:
        claims = build_claims(ctx)
    finally:
        ctx.close()
    verdict_counts = dict(
        sorted(collections.Counter(entry["verdict"] for entry in claims.values()).items())
    )
    pack = {
        "schema": SCHEMA,
        "audit_report": "2026-07-24",
        "input_manifest": MANIFEST_NAME,
        "input_count": manifest["input_count"],
        "manifest_exception": E17_MANIFEST_EXCEPTION,
        "code_scan_fingerprint": claims["E17"]["recomputed"][
            "code_scan_fingerprint"
        ],
        "tolerances": {
            "integer": "exact",
            "percentage_points": PERCENT_TOLERANCE_PP,
            "rho": RHO_TOLERANCE,
        },
        "summary": {"claim_count": len(claims), "verdict_counts": verdict_counts},
        "claims": claims,
    }
    atomic_json(out / JSON_NAME, pack)
    atomic_text(out / MD_NAME, render_markdown(pack))
    print("WROTE %s" % (out / JSON_NAME))
    print("WROTE %s" % (out / MD_NAME))
    print("CLAIMS %d; VERDICTS %s" % (len(claims), verdict_counts))
    return 0


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description="Copy and recompute the D1 legislative evidence pack."
    )
    parser.add_argument(
        "--copy-phase",
        action="store_true",
        help="copy the fixed LIVE input inventory into scratch and write its manifest",
    )
    parser.add_argument("--live-root", help="LIVE copy-source root; valid only with --copy-phase")
    parser.add_argument("--data-root", required=True, help="approved scratch data directory")
    parser.add_argument("--out", required=True, help="must resolve to audits/legfix2")
    return parser.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    data_root, out = require_paths(Path(args.data_root), Path(args.out))
    if args.copy_phase:
        if not args.live_root:
            raise EvidenceError("--live-root is required with --copy-phase")
        return run_copy_phase(Path(args.live_root), data_root, out)
    if args.live_root:
        raise EvidenceError("--live-root is prohibited outside --copy-phase")
    return run_evidence(data_root, out)


if __name__ == "__main__":
    try:
        sys.exit(main())
    except EvidenceError as exc:
        print("BLOCKED: %s" % exc, file=sys.stderr)
        sys.exit(2)
