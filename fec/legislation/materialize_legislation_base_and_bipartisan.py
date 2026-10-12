# -*- coding: utf-8 -*-
"""Stage a verified base-network workbook and pre-behavior bipartisanship.

The adapter has two deliberately separate responsibilities:

* copy the PASS output of ``build_legislation_base_network.py`` byte-for-byte
  to the isolated run-root filename expected by legacy consumers; and
* reproduce the pre-behavior ``bipartisan.json`` calculation currently defined
  in ``build_dashboard_bundle.py``.

The bipartisan calculation is not inferred from names.  Network nodes and
party joins use Bioguide IDs.  A complete current-member roster and Voteview
party sources for Congresses 119, 118, and 117 are hash-bound inputs.  Missing
or same-Congress-ambiguous party evidence produces a staged BLOCKED manifest,
not a fabricated bipartisan score.

No workbook cell is edited.  No production or mainframe path is accepted.
"""
from __future__ import annotations

import os as _bootstrap_os, sys as _bootstrap_sys; _bootstrap_sys.path.insert(0, _bootstrap_os.path.dirname(_bootstrap_os.path.dirname(_bootstrap_os.path.abspath(__file__))))

import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
from typing import Any, Iterable, Sequence

import openpyxl

from paths import ROOT as MB_ROOT


CODE_ROOT = Path(__file__).resolve().parents[2]
SCHEMA = "legislation-base-bipartisan-materialization-v1"
MANIFEST_NAME = "base_bipartisan_manifest.json"
WORKBOOK_NAME = "congress_cosponsor_network_shadow.xlsx"
SUCCESS_OUTPUTS = (
    WORKBOOK_NAME,
    "bipartisan.json",
    "bipartisan_by_bioguide.json",
    "bipartisan_member_audit.ndjson",
    "bipartisan_edge_exclusions.ndjson",
    "base_bipartisan_consumer_compatibility.json",
)
BLOCKED_OUTPUTS = (
    WORKBOOK_NAME,
    "bipartisan_member_audit.ndjson",
    "bipartisan_edge_exclusions.ndjson",
    "base_bipartisan_consumer_compatibility.json",
)
REQUIRED_BASE_OUTPUTS = {
    "workbook",
    "edge_contributions",
    "zero_contribution_exclusions",
}
EXPECTED_DECAY = {"116": 0.16, "117": 0.33, "118": 0.66, "119": 1.0}
PARTY_CONGRESSES = (119, 118, 117)
ALPHA = 6.0
FALLBACK_PRIOR = 0.12
BIOGUIDE_RE = re.compile(r"^[A-Z]\d{6}$")
HEX_SHA_RE = re.compile(r"^[0-9a-f]{64}$")
EDGE_HEADERS = (
    "From",
    "To",
    "Weight",
    "NBills",
    "NOriginal",
    "NRegular",
    "From_label",
    "To_label",
)
CENTRALITY_HEADERS = (
    "node",
    "label",
    "InfluenceScore",
    "WeightedOutDegree",
    "WeightedInDegree",
    "OutDegree",
    "InDegree",
    "Closeness",
    "Betweenness",
    "Eigenvector",
    "PageRank",
)


class AdapterError(RuntimeError):
    """A fail-closed input or output contract violation."""


def parser() -> argparse.ArgumentParser:
    root = Path(MB_ROOT)
    result = argparse.ArgumentParser(
        description=(
            "Copy a PASS staged base network byte-for-byte and materialize "
            "the canonical pre-behavior bipartisan percentage."
        )
    )
    result.add_argument(
        "--base-network",
        type=Path,
        default=root / "inputs" / "base_network",
        help="PASS base-network directory or manifest",
    )
    result.add_argument(
        "--member-roster",
        type=Path,
        help=(
            "Current member roster (JSON, CSV, or XLSX). If omitted, only "
            "manifest-adjacent/run-root staged candidates are considered."
        ),
    )
    result.add_argument(
        "--party-roster",
        type=Path,
        action="append",
        default=None,
        help=(
            "Voteview member CSV; repeat as needed. The combined inputs must "
            "contain Congresses 119, 118, and 117."
        ),
    )
    result.add_argument("--out", type=Path, default=root)
    return result


def canonical_json(value: Any) -> str:
    return json.dumps(
        value, sort_keys=True, ensure_ascii=False, separators=(",", ":")
    )


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def file_receipt(path: Path, *, relative_to: Path | None = None) -> dict[str, Any]:
    shown = path.relative_to(relative_to) if relative_to else path.resolve()
    return {
        "path": str(shown).replace("\\", "/"),
        "size": path.stat().st_size,
        "sha256": sha256_file(path),
    }


def write_json(path: Path, value: Any) -> None:
    path.write_text(
        json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False) + "\n",
        encoding="utf-8",
        newline="\n",
    )


def write_ndjson(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.write_text(
        "".join(canonical_json(row) + "\n" for row in rows),
        encoding="utf-8",
        newline="\n",
    )


def read_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        raise AdapterError(f"invalid JSON at {path}: {exc}") from exc


def resolve_manifest(value: Path) -> tuple[Path, dict[str, Any]]:
    path = value.expanduser().resolve()
    if path.is_dir():
        path = path / "manifest.json"
    if not path.is_file():
        raise AdapterError(f"base-network manifest does not exist: {path}")
    manifest = read_json(path)
    if not isinstance(manifest, dict):
        raise AdapterError("base-network manifest root must be an object")
    return path, manifest


def receipt_candidates(manifest_path: Path, item: dict[str, Any]) -> list[Path]:
    raw = item.get("path")
    if not isinstance(raw, str) or not raw.strip():
        raise AdapterError(f"receipt lacks a path in {manifest_path}")
    declared = Path(raw)
    declared = (
        declared.resolve()
        if declared.is_absolute()
        else (manifest_path.parent / declared).resolve()
    )
    local = (manifest_path.parent / Path(raw).name).resolve()
    return [local, declared] if local != declared else [declared]


def verify_receipt(
    manifest_path: Path, item: dict[str, Any], label: str
) -> dict[str, Any]:
    sha = str(item.get("sha256") or "").casefold()
    if not HEX_SHA_RE.fullmatch(sha):
        raise AdapterError(f"{label} lacks a valid SHA-256")
    candidates = receipt_candidates(manifest_path, item)
    existing = [path for path in candidates if path.is_file()]
    if not existing:
        raise AdapterError(
            f"{label} is missing: " + ", ".join(str(path) for path in candidates)
        )
    # Prefer a copied, manifest-adjacent isolated input. If it is corrupt, do
    # not silently escape to the producer's original absolute path.
    path = existing[0]
    observed = sha256_file(path)
    if observed != sha:
        raise AdapterError(
            f"{label} hash mismatch: expected {sha}, observed {observed}"
        )
    if item.get("size") is not None and path.stat().st_size != int(item["size"]):
        raise AdapterError(f"{label} size mismatch: {path}")
    return {**file_receipt(path), "resolved_path": path}


def validate_base_generation(value: Path) -> dict[str, Any]:
    manifest_path, manifest = resolve_manifest(value)
    policy = manifest.get("policy") or {}
    invariants = manifest.get("invariants") or {}
    if not (
        manifest.get("schema") == "legislation-base-network-manifest-v1"
        and manifest.get("status") == "PASS"
        and manifest.get("acceptance_passed") is True
        and manifest.get("production_write_prohibited") is True
        and policy.get("congress_decay") == EXPECTED_DECAY
        and policy.get("measure_identity")
        == "congress:measure_type:measure_number"
        and invariants.get("unique_type_aware_measure_edge_contributions") is True
        and invariants.get("no_withdrawn_or_excluded_contribution") is True
        and invariants.get("positive_finite_edge_weights") is True
    ):
        raise AdapterError(
            "base-network generation is not PASS under the type-aware "
            "1/.66/.33/.16 staged policy"
        )
    raw_outputs = manifest.get("outputs")
    if not isinstance(raw_outputs, dict):
        raise AdapterError("base-network manifest has no output receipts")
    outputs = {}
    for name, item in sorted(raw_outputs.items()):
        if item is None:
            continue
        if not isinstance(item, dict):
            raise AdapterError(f"invalid base-network output receipt: {name}")
        outputs[str(name)] = verify_receipt(
            manifest_path, item, f"base-network/{name}"
        )
    missing = sorted(REQUIRED_BASE_OUTPUTS - set(outputs))
    if missing:
        raise AdapterError(
            "base-network generation lacks required outputs: "
            + ", ".join(missing)
        )
    accepted_input = (manifest.get("inputs") or {}).get("accepted_ledger")
    if not isinstance(accepted_input, dict) or not HEX_SHA_RE.fullmatch(
        str(accepted_input.get("sha256") or "").casefold()
    ):
        raise AdapterError(
            "base-network manifest is not hash-bound to an accepted ledger"
        )
    return {
        "manifest_path": manifest_path,
        "manifest": manifest,
        "manifest_receipt": file_receipt(manifest_path),
        "outputs": outputs,
        "declared_roster_receipt": (manifest.get("inputs") or {}).get("roster"),
        "accepted_ledger_sha256": str(accepted_input["sha256"]).casefold(),
    }


def normalize_name(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def roster_rows(path: Path) -> list[tuple[Any, Any]]:
    suffix = path.suffix.casefold()
    if suffix == ".json":
        payload = read_json(path)
        if isinstance(payload, dict):
            payload = (
                payload.get("members")
                or payload.get("records")
                or payload.get("roster")
            )
        if not isinstance(payload, list):
            raise AdapterError(f"member roster JSON must contain a list: {path}")
        return [
            (
                row.get("bioguide_id")
                or row.get("bioguide")
                or row.get("Bioguide"),
                row.get("member_name")
                or row.get("member")
                or row.get("name")
                or row.get("label"),
            )
            for row in payload
            if isinstance(row, dict)
        ]
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return [
                (
                    row.get("bioguide_id")
                    or row.get("bioguide")
                    or row.get("Bioguide"),
                    row.get("member_name")
                    or row.get("member")
                    or row.get("name")
                    or row.get("label"),
                )
                for row in csv.DictReader(handle)
            ]
    if suffix == ".xlsx":
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if "Centrality" not in workbook.sheetnames:
                raise AdapterError("XLSX member roster requires Centrality sheet")
            rows = workbook["Centrality"].iter_rows(values_only=True)
            headers = tuple(next(rows, ()))
            if "node" not in headers or "label" not in headers:
                raise AdapterError("Centrality roster lacks node/label")
            node_index, label_index = headers.index("node"), headers.index("label")
            return [(row[node_index], row[label_index]) for row in rows]
        finally:
            workbook.close()
    raise AdapterError(f"unsupported member roster format: {path}")


def load_member_roster(path: Path) -> dict[str, str]:
    names: dict[str, str] = {}
    names_to_bio: dict[str, str] = {}
    rows = roster_rows(path)
    if not rows:
        raise AdapterError("member roster is empty")
    for row_number, (raw_bio, raw_name) in enumerate(rows, 2):
        bioguide = normalize_name(raw_bio).upper()
        name = normalize_name(raw_name)
        if not BIOGUIDE_RE.fullmatch(bioguide) or not name:
            raise AdapterError(
                f"invalid member roster identity at {path}:{row_number}"
            )
        if bioguide in names and names[bioguide] != name:
            raise AdapterError(
                f"member roster name conflict for {bioguide}: "
                f"{names[bioguide]!r} vs {name!r}"
            )
        folded = name.casefold()
        if folded in names_to_bio and names_to_bio[folded] != bioguide:
            raise AdapterError(
                f"member roster name maps to multiple Bioguides: {name!r}"
            )
        names[bioguide] = name
        names_to_bio[folded] = bioguide
    return names


def discover_member_roster(
    *,
    explicit: Path | None,
    root: Path,
    base: dict[str, Any],
) -> tuple[Path | None, list[str]]:
    blockers: list[str] = []
    manifest_dir = base["manifest_path"].parent
    declared = base.get("declared_roster_receipt")
    candidates: list[Path] = []
    if explicit is not None:
        candidates.append(explicit.expanduser().resolve())
    else:
        if isinstance(declared, dict) and declared.get("path"):
            candidates.append(
                (manifest_dir / Path(str(declared["path"])).name).resolve()
            )
        candidates.extend(
            [
                (root / "data" / "engine" / "congestion.json").resolve(),
                (manifest_dir / "data" / "engine" / "congestion.json").resolve(),
                (manifest_dir / "member_roster.json").resolve(),
                (manifest_dir / "member_roster.csv").resolve(),
            ]
        )
    existing = []
    seen = set()
    for path in candidates:
        if path not in seen and path.is_file():
            seen.add(path)
            existing.append(path)
    if not existing:
        blockers.append("current_member_roster_missing_from_isolated_inputs")
        return None, blockers
    if explicit is None and len(existing) > 1:
        hashes = {sha256_file(path) for path in existing}
        if len(hashes) > 1:
            blockers.append(
                "multiple_nonidentical_current_member_roster_candidates"
            )
            return None, blockers
    selected = existing[0]
    if isinstance(declared, dict) and declared.get("sha256"):
        declared_sha = str(declared["sha256"]).casefold()
        if not HEX_SHA_RE.fullmatch(declared_sha):
            blockers.append("base_manifest_roster_receipt_invalid")
        elif sha256_file(selected) != declared_sha:
            blockers.append("current_member_roster_not_bound_to_base_manifest")
    return selected, blockers


def discover_party_rosters(
    *,
    explicit: Sequence[Path] | None,
    root: Path,
    base_dir: Path,
) -> list[Path]:
    if explicit:
        return [path.expanduser().resolve() for path in explicit]
    candidates = []
    for congress in PARTY_CONGRESSES:
        locations = (
            root / "votes" / f"HS{congress}_members.csv",
            base_dir / f"HS{congress}_members.csv",
            base_dir / "party_roster" / f"HS{congress}_members.csv",
        )
        existing = [path.resolve() for path in locations if path.is_file()]
        if existing:
            candidates.append(existing[0])
    return candidates


def load_party_evidence(
    paths: Sequence[Path],
    roster: dict[str, str],
) -> tuple[dict[str, str | None], list[dict[str, Any]], dict[str, Any], list[str]]:
    evidence: dict[tuple[int, str], set[str]] = defaultdict(set)
    source_rows: dict[tuple[int, str], list[dict[str, Any]]] = defaultdict(list)
    source_audit = []
    blockers: list[str] = []
    congresses_seen = set()
    for path in paths:
        if not path.is_file():
            blockers.append(f"party_roster_missing:{path}")
            continue
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            required = {"congress", "bioguide_id", "party_code"}
            if not reader.fieldnames or not required.issubset(reader.fieldnames):
                blockers.append(f"party_roster_headers_invalid:{path}")
                continue
            row_count = 0
            included_identity_rows = 0
            missing_bioguide_rows = 0
            for line_number, row in enumerate(reader, 2):
                row_count += 1
                try:
                    congress = int(str(row.get("congress") or "").strip())
                except ValueError:
                    blockers.append(
                        f"party_roster_congress_invalid:{path}:{line_number}"
                    )
                    continue
                if congress not in PARTY_CONGRESSES:
                    continue
                congresses_seen.add(congress)
                bioguide = normalize_name(row.get("bioguide_id")).upper()
                if not bioguide:
                    missing_bioguide_rows += 1
                    continue
                if not BIOGUIDE_RE.fullmatch(bioguide):
                    blockers.append(
                        f"party_roster_bioguide_invalid:{path}:{line_number}"
                    )
                    continue
                code = str(row.get("party_code") or "").strip()
                party = "D" if code == "100" else "R" if code == "200" else "I"
                evidence[(congress, bioguide)].add(party)
                source_rows[(congress, bioguide)].append(
                    {
                        "path": str(path.resolve()),
                        "line": line_number,
                        "party_code": code,
                        "mapped_party": party,
                    }
                )
                included_identity_rows += 1
            source_audit.append(
                {
                    **file_receipt(path),
                    "row_count": row_count,
                    "included_identity_rows": included_identity_rows,
                    "missing_bioguide_rows_excluded": missing_bioguide_rows,
                }
            )
    missing_congresses = sorted(set(PARTY_CONGRESSES) - congresses_seen)
    if missing_congresses:
        blockers.append(
            "required_party_congresses_missing:"
            + ",".join(str(value) for value in missing_congresses)
        )

    selected: dict[str, str | None] = {}
    member_rows = []
    ambiguous_count = 0
    unknown_count = 0
    for bioguide, name in sorted(roster.items()):
        selection = None
        selection_congress = None
        status = "unknown_no_voteview_party"
        history = []
        ambiguous = False
        for congress in PARTY_CONGRESSES:
            parties = sorted(evidence.get((congress, bioguide), set()))
            if parties:
                history.append(
                    {
                        "congress": congress,
                        "parties": parties,
                        "sources": source_rows[(congress, bioguide)],
                    }
                )
            if selection_congress is not None or not parties:
                continue
            if len(parties) > 1:
                ambiguous = True
                selection_congress = congress
                status = "ambiguous_same_congress_party"
                ambiguous_count += 1
                blockers.append(
                    f"ambiguous_same_congress_party:{congress}:{bioguide}"
                )
            else:
                selection = parties[0]
                selection_congress = congress
                status = (
                    "included_major_party"
                    if selection in {"D", "R"}
                    else "included_nonmajor_party_prior_only"
                )
        if selection_congress is None:
            unknown_count += 1
        selected[bioguide] = None if ambiguous else selection
        member_rows.append(
            {
                "schema": SCHEMA,
                "bioguide_id": bioguide,
                "member_name": name,
                "selected_party": None if ambiguous else selection,
                "selected_congress": selection_congress,
                "party_status": status,
                "party_history": history,
                "network_evidence_eligible": (
                    not ambiguous and selection in {"D", "R"}
                ),
            }
        )
    summary = {
        "required_congresses": list(PARTY_CONGRESSES),
        "congresses_seen": sorted(congresses_seen, reverse=True),
        "sources": source_audit,
        "member_party_counts": dict(
            sorted(
                Counter(
                    party if party is not None else "unknown_or_ambiguous"
                    for party in selected.values()
                ).items()
            )
        ),
        "ambiguous_same_congress_member_count": ambiguous_count,
        "unknown_member_count": unknown_count,
    }
    return selected, member_rows, summary, blockers


def read_base_workbook(
    path: Path,
) -> tuple[list[dict[str, Any]], dict[str, str]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "Edges" not in workbook.sheetnames or "Centrality" not in workbook.sheetnames:
            raise AdapterError("base workbook requires Edges and Centrality sheets")
        edge_values = workbook["Edges"].iter_rows(values_only=True)
        edge_headers = tuple(next(edge_values, ()))
        if edge_headers != EDGE_HEADERS:
            raise AdapterError(
                f"base workbook Edges headers differ: {edge_headers!r}"
            )
        edges = []
        directed_keys = set()
        for row_number, values in enumerate(edge_values, 2):
            row = dict(zip(edge_headers, values))
            source = normalize_name(row["From"]).upper()
            target = normalize_name(row["To"]).upper()
            if not BIOGUIDE_RE.fullmatch(source) or not BIOGUIDE_RE.fullmatch(target):
                raise AdapterError(
                    f"invalid Bioguide edge endpoint at workbook row {row_number}"
                )
            try:
                weight = float(row["Weight"])
            except (TypeError, ValueError) as exc:
                raise AdapterError(
                    f"invalid edge weight at workbook row {row_number}"
                ) from exc
            if not math.isfinite(weight) or weight <= 0:
                raise AdapterError(
                    f"nonpositive/nonfinite edge at workbook row {row_number}"
                )
            key = (source, target)
            if key in directed_keys:
                raise AdapterError(f"duplicate directed base edge: {source}->{target}")
            directed_keys.add(key)
            edges.append(
                {
                    "source": source,
                    "target": target,
                    "weight": weight,
                    "source_row": row_number,
                    "from_label": row["From_label"],
                    "to_label": row["To_label"],
                }
            )

        centrality_values = workbook["Centrality"].iter_rows(values_only=True)
        centrality_headers = tuple(next(centrality_values, ()))
        if centrality_headers != CENTRALITY_HEADERS:
            raise AdapterError(
                f"base workbook Centrality headers differ: {centrality_headers!r}"
            )
        centrality = {}
        for row_number, values in enumerate(centrality_values, 2):
            row = dict(zip(centrality_headers, values))
            bioguide = normalize_name(row["node"]).upper()
            name = normalize_name(row["label"])
            if not BIOGUIDE_RE.fullmatch(bioguide) or not name:
                raise AdapterError(
                    f"invalid Centrality identity at workbook row {row_number}"
                )
            if bioguide in centrality and centrality[bioguide] != name:
                raise AdapterError(f"conflicting Centrality label for {bioguide}")
            centrality[bioguide] = name
        return edges, centrality
    finally:
        workbook.close()


def compute_bipartisan(
    *,
    edges: list[dict[str, Any]],
    centrality: dict[str, str],
    roster: dict[str, str],
    party: dict[str, str | None],
    party_audit: list[dict[str, Any]],
) -> tuple[
    dict[str, int],
    dict[str, int],
    list[dict[str, Any]],
    list[dict[str, Any]],
    dict[str, Any],
    list[str],
]:
    blockers: list[str] = []
    missing_centrality = sorted(set(roster) - set(centrality))
    if missing_centrality:
        blockers.append(
            f"current_roster_missing_from_base_centrality:{len(missing_centrality)}"
        )
    name_conflicts = sorted(
        bioguide
        for bioguide in set(roster) & set(centrality)
        if roster[bioguide] != centrality[bioguide]
    )
    if name_conflicts:
        blockers.append(
            f"base_centrality_member_name_conflicts:{len(name_conflicts)}"
        )

    pair_rows: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    exclusions: list[dict[str, Any]] = []
    for edge in edges:
        source, target = edge["source"], edge["target"]
        if source == target:
            exclusions.append(
                {
                    "schema": SCHEMA,
                    **edge,
                    "disposition": "excluded",
                    "reason": "self_edge",
                }
            )
            continue
        outside = sorted({source, target} - set(roster))
        if outside:
            exclusions.append(
                {
                    "schema": SCHEMA,
                    **edge,
                    "disposition": "excluded",
                    "reason": "endpoint_outside_current_member_roster",
                    "outside_roster_bioguides": outside,
                }
            )
            continue
        pair_rows[tuple(sorted((source, target)))].append(edge)

    undirected = []
    for pair, rows in sorted(pair_rows.items()):
        if len(rows) > 2:
            blockers.append(f"more_than_two_directed_rows_for_pair:{pair[0]}:{pair[1]}")
            continue
        if len(rows) == 2:
            directions = {(row["source"], row["target"]) for row in rows}
            if len(directions) != 2:
                blockers.append(
                    f"duplicate_direction_in_pair:{pair[0]}:{pair[1]}"
                )
                continue
            weight = math.fsum(row["weight"] for row in rows) / 2.0
            symmetrization = "mean_of_reverse_directed_edges"
        else:
            weight = rows[0]["weight"]
            symmetrization = "single_directed_edge_used_as_undirected_weight"
        undirected.append(
            {
                "a": pair[0],
                "b": pair[1],
                "weight": weight,
                "symmetrization": symmetrization,
                "source_rows": [row["source_row"] for row in rows],
                "source_directed_weights": [row["weight"] for row in rows],
            }
        )

    same_cross = {bioguide: [0.0, 0.0] for bioguide in roster}
    included_pairs = 0
    for edge in undirected:
        pa, pb = party.get(edge["a"]), party.get(edge["b"])
        if pa not in {"D", "R"} or pb not in {"D", "R"}:
            exclusions.append(
                {
                    "schema": SCHEMA,
                    **edge,
                    "party_a": pa,
                    "party_b": pb,
                    "disposition": "excluded",
                    "reason": "nonmajor_or_unknown_party_endpoint",
                }
            )
            continue
        included_pairs += 1
        cross = pa != pb
        for node in (edge["a"], edge["b"]):
            same_cross[node][1 if cross else 0] += edge["weight"]

    global_same = math.fsum(value[0] for value in same_cross.values())
    global_cross = math.fsum(value[1] for value in same_cross.values())
    prior = (
        global_cross / (global_same + global_cross)
        if global_same + global_cross > 0
        else FALLBACK_PRIOR
    )
    by_name: dict[str, int] = {}
    by_bioguide: dict[str, int] = {}
    audit_by_bio = {row["bioguide_id"]: row for row in party_audit}
    member_audit = []
    for bioguide, name in sorted(roster.items(), key=lambda item: item[1].casefold()):
        same, cross = same_cross[bioguide]
        share = (cross + ALPHA * prior) / (same + cross + ALPHA)
        percent = round(share * 100)
        if name in by_name:
            blockers.append(f"duplicate_legacy_member_name:{name}")
        by_name[name] = percent
        by_bioguide[bioguide] = percent
        party_row = audit_by_bio.get(bioguide) or {}
        selected_party = party.get(bioguide)
        evidence_disposition = (
            "included_weighted_network_evidence"
            if selected_party in {"D", "R"}
            else "included_prior_only_no_major_party_evidence"
        )
        member_audit.append(
            {
                **party_row,
                "same_party_weight": round(same, 12),
                "cross_party_weight": round(cross, 12),
                "eligible_weight": round(same + cross, 12),
                "global_same_party_weight": round(global_same, 12),
                "global_cross_party_weight": round(global_cross, 12),
                "global_cross_party_prior": round(prior, 12),
                "shrinkage_alpha": ALPHA,
                "shrunk_cross_party_share": round(share, 12),
                "legacy_bipartisan_percent": percent,
                "disposition": evidence_disposition,
                "reason": (
                    "weighted_cross_party_share_with_global_prior"
                    if selected_party in {"D", "R"}
                    else "canonical_prior_only_value_for_nonmajor_or_unknown_party"
                ),
            }
        )
    exclusions.sort(
        key=lambda row: (
            str(row.get("reason") or ""),
            str(row.get("a") or row.get("source") or ""),
            str(row.get("b") or row.get("target") or ""),
            int(row.get("source_row") or 0),
        )
    )
    summary = {
        "directed_edge_rows": len(edges),
        "undirected_current_roster_pairs": len(undirected),
        "major_party_eligible_pairs": included_pairs,
        "excluded_edge_rows_or_pairs": len(exclusions),
        "global_same_party_weight": round(global_same, 12),
        "global_cross_party_weight": round(global_cross, 12),
        "global_cross_party_prior": round(prior, 12),
        "fallback_prior_used": global_same + global_cross == 0,
        "member_count": len(roster),
        "major_party_member_count": sum(
            value in {"D", "R"} for value in party.values()
        ),
        "prior_only_member_count": sum(
            value not in {"D", "R"} for value in party.values()
        ),
        "extra_centrality_nodes_excluded_from_bipartisan": len(
            set(centrality) - set(roster)
        ),
    }
    return (
        dict(sorted(by_name.items())),
        dict(sorted(by_bioguide.items())),
        member_audit,
        exclusions,
        summary,
        blockers,
    )


def audit_consumers() -> tuple[dict[str, Any], list[str]]:
    specifications = {
        "canonical_formula": {
            "path": CODE_ROOT / "fec" / "surfaces" / "build_dashboard_bundle.py",
            "tokens": (
                "prior=gc/(gs+gc) if (gs+gc)>0 else 0.12; ALPHA=6.0",
                "(raw[n][1]+ALPHA*prior)/(raw[n][0]+raw[n][1]+ALPHA)",
                "BP.append(round(bp[n]*100))",
            ),
        },
        "behavior_panel": {
            "path": CODE_ROOT / "fec" / "members" / "behavior_panel.py",
            "tokens": (
                "bipartisan.json",
                "P[m]['bipartisanship'] = round(float(v), 4)",
            ),
        },
        "moneyball_core": {
            "path": CODE_ROOT / "fec" / "moneyball_core.py",
            "tokens": ("bipartisan.json", "BP.get(m,24)/100.0"),
        },
        "cooperation_v2": {
            "path": CODE_ROOT / "fec" / "members" / "coop_v2_candidate.py",
            "tokens": ("c1_bip", "z1, z2, z3 = zscore(c1)"),
        },
    }
    sources = {}
    blockers = []
    for name, spec in specifications.items():
        path = spec["path"]
        if not path.is_file():
            sources[name] = {"path": str(path), "exists": False}
            blockers.append(f"consumer_source_missing:{name}")
            continue
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        missing = [token for token in spec["tokens"] if token not in text]
        if missing:
            blockers.append(f"consumer_contract_changed:{name}")
        sources[name] = {
            **file_receipt(path),
            "expected_contract_tokens_present": not missing,
            "missing_contract_tokens": missing,
        }
    report = {
        "schema": "base-bipartisan-consumer-compatibility-v1",
        "status": "PASS" if not blockers else "BLOCKED",
        "sources": sources,
        "canonical_semantics": {
            "graph": (
                "undirected pure-cosponsorship graph; reverse directed edge "
                "weights are arithmetic-mean symmetrized"
            ),
            "global_prior": "global_cross/(global_same+global_cross), else 0.12",
            "member_share": "(cross + 6*prior)/(same + cross + 6)",
            "legacy_output": "round(member_share*100) integer keyed by member name",
            "identity": "Bioguide joins; names are compatibility output keys",
        },
        "consumers": {
            "behavior_panel": (
                "loads the 0-100 value as bipartisanship; later z-scoring is "
                "scale invariant"
            ),
            "moneyball_core.bpf": (
                "divides the 0-100 value by 100; missing member fallback is 24%"
            ),
            "coop_v2_candidate": (
                "comments call c1 a ratio but z-scores the 0-100 artifact, "
                "which is scale-equivalent"
            ),
        },
        "known_limitations": [
            "legacy bipartisan.json is keyed by member name, not Bioguide",
            "nonmajor/unknown-party members receive the canonical prior-only value",
            "edges touching nonmajor/unknown-party members do not enter same/cross totals",
            "party evidence uses recent-Congress priority 119 then 118 then 117",
        ],
        "blocking_issues": blockers,
    }
    return report, blockers


def is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def validate_output_path(value: Path) -> Path:
    path = value.expanduser().resolve()
    audit_root = (CODE_ROOT / "audits" / "legfix5").resolve()
    temp_root = Path(tempfile.gettempdir()).resolve()
    if path in {audit_root, temp_root}:
        raise AdapterError("output must be a child staging directory")
    if not (
        is_relative_to(path, audit_root) or is_relative_to(path, temp_root)
    ):
        raise AdapterError(
            "output is restricted to audits/legfix5 or a temporary directory"
        )
    return path


def install_idempotently(temp: Path, out: Path, names: Sequence[str]) -> bool:
    targets = [out / name for name in names]
    existing = [path.exists() for path in targets]
    if any(existing):
        if not all(existing):
            raise AdapterError(
                "partial prior adapter output exists; refusing a mixed generation"
            )
        mismatched = [
            path.name
            for path in targets
            if path.read_bytes() != (temp / path.name).read_bytes()
        ]
        if mismatched:
            raise AdapterError(
                "immutable staged output differs: " + ", ".join(mismatched)
            )
        return True
    out.mkdir(parents=True, exist_ok=True)
    for name in names:
        os.replace(temp / name, out / name)
    return False


def materialize(
    *,
    base_network: Path,
    member_roster: Path | None,
    party_rosters: Sequence[Path] | None,
    out: Path,
) -> dict[str, Any]:
    out = validate_output_path(out)
    root = Path(MB_ROOT).resolve()
    base = validate_base_generation(base_network)
    roster_path, discovery_blockers = discover_member_roster(
        explicit=member_roster,
        root=root,
        base=base,
    )
    blockers = list(discovery_blockers)
    roster: dict[str, str] = {}
    roster_receipt = None
    if roster_path is not None:
        try:
            roster = load_member_roster(roster_path)
            roster_receipt = file_receipt(roster_path)
        except AdapterError as exc:
            blockers.append(f"current_member_roster_invalid:{exc}")

    party_paths = discover_party_rosters(
        explicit=party_rosters,
        root=root,
        base_dir=base["manifest_path"].parent,
    )
    party = {bioguide: None for bioguide in roster}
    party_audit = [
        {
            "schema": SCHEMA,
            "bioguide_id": bioguide,
            "member_name": name,
            "selected_party": None,
            "selected_congress": None,
            "party_status": "party_evidence_unavailable",
            "party_history": [],
            "network_evidence_eligible": False,
        }
        for bioguide, name in sorted(roster.items())
    ]
    party_summary = {
        "required_congresses": list(PARTY_CONGRESSES),
        "congresses_seen": [],
        "sources": [],
        "member_party_counts": {"unknown_or_ambiguous": len(roster)},
        "ambiguous_same_congress_member_count": 0,
        "unknown_member_count": len(roster),
    }
    if roster:
        party, party_audit, party_summary, party_blockers = load_party_evidence(
            party_paths, roster
        )
        blockers.extend(party_blockers)
    else:
        blockers.append("current_member_roster_unavailable_for_party_join")

    workbook_source = base["outputs"]["workbook"]["resolved_path"]
    edges, centrality = read_base_workbook(workbook_source)
    (
        bipartisan,
        bipartisan_by_bio,
        member_audit,
        exclusions,
        formula_summary,
        formula_blockers,
    ) = compute_bipartisan(
        edges=edges,
        centrality=centrality,
        roster=roster,
        party=party,
        party_audit=party_audit,
    )
    blockers.extend(formula_blockers)
    compatibility, consumer_blockers = audit_consumers()
    blockers.extend(consumer_blockers)
    unique_blockers = sorted(set(blockers))

    out.parent.mkdir(parents=True, exist_ok=True)
    temp = Path(
        tempfile.mkdtemp(prefix=f".{out.name}.base-bip-", dir=str(out.parent))
    )
    try:
        shutil.copyfile(workbook_source, temp / WORKBOOK_NAME)
        write_ndjson(temp / "bipartisan_member_audit.ndjson", member_audit)
        write_ndjson(temp / "bipartisan_edge_exclusions.ndjson", exclusions)
        write_json(
            temp / "base_bipartisan_consumer_compatibility.json",
            compatibility,
        )
        output_names = list(BLOCKED_OUTPUTS)
        if not unique_blockers:
            write_json(temp / "bipartisan.json", bipartisan)
            write_json(
                temp / "bipartisan_by_bioguide.json", bipartisan_by_bio
            )
            output_names = list(SUCCESS_OUTPUTS)
        outputs = {
            name: file_receipt(temp / name, relative_to=temp)
            for name in output_names
        }
        base_workbook_hash = base["outputs"]["workbook"]["sha256"]
        generation_basis = {
            "base_manifest_sha256": base["manifest_receipt"]["sha256"],
            "base_workbook_sha256": base_workbook_hash,
            "accepted_ledger_sha256": base["accepted_ledger_sha256"],
            "member_roster_sha256": (
                roster_receipt["sha256"] if roster_receipt else None
            ),
            "party_roster_sha256": [
                item["sha256"] for item in party_summary["sources"]
            ],
            "formula_source_sha256": {
                name: item.get("sha256")
                for name, item in compatibility["sources"].items()
            },
            "status": "PASS" if not unique_blockers else "BLOCKED",
            "blocking_issues": unique_blockers,
            "outputs": outputs,
        }
        manifest = {
            "schema": SCHEMA,
            "status": "PASS" if not unique_blockers else "BLOCKED",
            "acceptance_passed": not unique_blockers,
            "generated_at": None,
            "generation_id": sha256_bytes(
                canonical_json(generation_basis).encode("utf-8")
            ),
            "staging_only": True,
            "production_written": False,
            "mainframe_written": False,
            "workbook_cells_edited": False,
            "base_workbook_byte_copy": {
                "source_sha256": base_workbook_hash,
                "output_sha256": outputs[WORKBOOK_NAME]["sha256"],
                "byte_identical": (
                    base_workbook_hash == outputs[WORKBOOK_NAME]["sha256"]
                ),
                "output_name": WORKBOOK_NAME,
            },
            "upstream_policy": {
                "type_aware_measure_identity": (
                    "congress:measure_type:measure_number"
                ),
                "congress_decay_already_embedded": EXPECTED_DECAY,
                "adapter_applies_no_additional_decay": True,
            },
            "bipartisan_policy": {
                "canonical_source": "fec/surfaces/build_dashboard_bundle.py",
                "identity": "Bioguide",
                "graph": "undirected pure base cosponsorship",
                "reverse_edge_symmetrization": "arithmetic mean",
                "party_priority": list(PARTY_CONGRESSES),
                "major_parties": ["D", "R"],
                "global_prior": (
                    "global_cross_weight/(global_same_weight+"
                    "global_cross_weight), fallback 0.12"
                ),
                "alpha": ALPHA,
                "member_share": "(cross + alpha*prior)/(same+cross+alpha)",
                "legacy_output": "round(member_share*100)",
                "legacy_scale": "integer percentage 0..100",
                "unknown_or_nonmajor_party": (
                    "prior-only output; incident edges excluded from totals"
                ),
            },
            "inputs": {
                "base_network_manifest": base["manifest_receipt"],
                "base_network_outputs": {
                    name: {
                        key: value
                        for key, value in receipt.items()
                        if key != "resolved_path"
                    }
                    for name, receipt in base["outputs"].items()
                },
                "accepted_ledger_sha256": base["accepted_ledger_sha256"],
                "current_member_roster": roster_receipt,
                "current_member_roster_declared_by_base_manifest": (
                    base["declared_roster_receipt"]
                ),
                "party_rosters": party_summary["sources"],
                "hash_binding_passed": not any(
                    issue
                    in {
                        "base_manifest_roster_receipt_invalid",
                        "current_member_roster_not_bound_to_base_manifest",
                    }
                    for issue in unique_blockers
                ),
            },
            "party_evidence": party_summary,
            "formula_summary": formula_summary,
            "exclusion_policy": {
                "all_exclusions_visible": True,
                "self_edges": "excluded",
                "endpoint_outside_current_roster": "excluded",
                "nonmajor_or_unknown_party_endpoint": "excluded",
                "unknown_or_nonmajor_members": "published at prior-only value",
            },
            "consumer_compatibility": {
                "audit": outputs[
                    "base_bipartisan_consumer_compatibility.json"
                ],
                "behavior_panel_compatible": not unique_blockers,
                "moneyball_core_bpf_compatible": not unique_blockers,
                "cooperation_v2_scale_equivalent": not unique_blockers,
            },
            "counts": {
                **formula_summary,
                "member_audit_rows": len(member_audit),
                "edge_exclusion_rows": len(exclusions),
                "bipartisan_name_rows": (
                    len(bipartisan) if not unique_blockers else 0
                ),
                "bipartisan_bioguide_rows": (
                    len(bipartisan_by_bio) if not unique_blockers else 0
                ),
            },
            "blocking_issue_count": len(unique_blockers),
            "blocking_issues": unique_blockers,
            "outputs": outputs,
        }
        write_json(temp / MANIFEST_NAME, manifest)
        installed_names = (*output_names, MANIFEST_NAME)
        reused = install_idempotently(temp, out, installed_names)
        result = read_json(out / MANIFEST_NAME)
        result["_idempotent_reuse"] = reused
        return result
    finally:
        shutil.rmtree(temp, ignore_errors=True)


def main(argv: Sequence[str] | None = None) -> int:
    args = parser().parse_args(argv)
    try:
        manifest = materialize(
            base_network=args.base_network,
            member_roster=args.member_roster,
            party_rosters=args.party_roster,
            out=args.out,
        )
    except AdapterError as exc:
        print(
            f"materialize_legislation_base_and_bipartisan: FAIL CLOSED: {exc}",
            file=sys.stderr,
        )
        return 2
    print(
        "base network + bipartisan: "
        f"{manifest['status']} | members={manifest['counts']['member_count']} | "
        f"prior={manifest['formula_summary']['global_cross_party_prior']} | "
        f"idempotent_reuse={manifest.pop('_idempotent_reuse', False)}"
    )
    print(args.out.resolve() / MANIFEST_NAME)
    return 0 if manifest["status"] == "PASS" else 2


if __name__ == "__main__":
    raise SystemExit(main())
