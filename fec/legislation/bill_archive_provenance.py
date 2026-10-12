# -*- coding: utf-8 -*-
r"""Recover Congress.gov API-browser query provenance from the bill archive.

The archive has no producer script.  Its executed hand-built search queries
survive in the ``hl=`` parameter of ``DetailLink`` values.  This module
reconstructs that query specification without network access and without
altering or inventing any bill row.

Evidence mode is fail-closed: the workbook must be the fixed, hash-verified
copy under ``legfix2-scratch\data`` and must match the committed D1 manifest.
The module is import-safe; openpyxl is imported only inside workbook analysis.
"""
import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
import unicodedata
from pathlib import Path
from urllib.parse import unquote_plus, urlsplit


sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import bill_rules as BR  # noqa: E402  (pure module: constants + defs only)

try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')
except AttributeError:
    pass


WORKTREE_ROOT = Path(__file__).resolve().parent.parent.parent
APPROVED_DATA_ROOT = Path(
    r'C:\Users\pgoff\AppData\Local\Temp\claude\legfix2-scratch\data'
).resolve()
APPROVED_OUT = (WORKTREE_ROOT / 'audits' / 'legfix2').resolve()
ARCHIVE_NAME = 'congress_bills_archive_20260624_115822.xlsx'
MANIFEST_NAME = 'source_manifest.json'
MANIFEST_SCHEMA = 'legfix2-source-manifest-v1'
OUTPUT_SCHEMA = 'legfix2-bill-archive-provenance-v1'
JSON_NAME = 'bill_archive_provenance.json'
MD_NAME = 'BILL_ARCHIVE_PROVENANCE.md'
AUDIT_REPORTED_NONEMPTY_HL = 8520
QUERY_SET_ZERO_ALLOWLIST = frozenset({'CDFI Fund', 'CU Tax Status'})
MAX_DECODED_HL_CHARS = 10000
MAX_DISTINCT_QUERIES_PER_SHEET = 1000
TRUNCATED_MARKER = 'TRUNCATED'

BAD_PERCENT_RE = re.compile(r'%(?![0-9A-Fa-f]{2})')
SHEET_NUMBER_RE = re.compile(r'^\s*(\d{2})\s+(.+?)\s*$')
OPERATOR_RE = re.compile(r'^(?:AND|OR|NOT|NEAR(?:/\d+)?)$', re.I)


class ProvenanceError(RuntimeError):
    """A fail-closed manifest, path, workbook, or query-structure error."""


def _is_relative_to(path, root):
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


def approved_paths(data_root, out):
    """Resolve and enforce the D2 read/write envelope."""
    data_root = Path(os.path.realpath(os.path.abspath(str(data_root))))
    out = Path(os.path.realpath(os.path.abspath(str(out))))
    physical_worktree = Path(os.path.realpath(str(WORKTREE_ROOT)))
    if data_root != APPROVED_DATA_ROOT:
        raise ProvenanceError(
            'data root must be the approved scratch copy: %s' % APPROVED_DATA_ROOT
        )
    if out != APPROVED_OUT:
        raise ProvenanceError(
            'output directory must be the worktree evidence directory: %s'
            % APPROVED_OUT
        )
    if not _is_relative_to(out, physical_worktree):
        raise ProvenanceError('output directory physically escaped the worktree')
    return data_root, out


def sha256_file(path):
    digest = hashlib.sha256()
    size = 0
    with open(path, 'rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
            size += len(chunk)
    return digest.hexdigest(), size


def load_verified_archive(data_root, manifest_path):
    """Return the archive scratch path and manifest entry after hash verification."""
    data_root = Path(data_root).resolve()
    manifest_path = Path(manifest_path).resolve()
    if data_root != APPROVED_DATA_ROOT:
        raise ProvenanceError(
            'C13/provenance requires the approved scratch data root: %s'
            % APPROVED_DATA_ROOT
        )
    expected_manifest = (APPROVED_OUT / MANIFEST_NAME).resolve()
    if manifest_path != expected_manifest:
        raise ProvenanceError(
            'manifest must be the committed D1 manifest: %s' % expected_manifest
        )
    if not manifest_path.is_file():
        raise ProvenanceError('manifest missing: %s' % manifest_path)
    try:
        with open(manifest_path, encoding='utf-8') as handle:
            manifest = json.load(handle)
    except (OSError, ValueError) as exc:
        raise ProvenanceError('manifest unreadable: %s' % exc)
    if manifest.get('schema') != MANIFEST_SCHEMA:
        raise ProvenanceError(
            'unexpected manifest schema: %r' % manifest.get('schema')
        )
    inputs = manifest.get('inputs')
    if not isinstance(inputs, list):
        raise ProvenanceError('manifest inputs must be a list')
    if manifest.get('input_count') != len(inputs):
        raise ProvenanceError('manifest input_count does not match its input list')
    matches = [
        item for item in inputs
        if item.get('scratch_name') == ARCHIVE_NAME
    ]
    if len(matches) != 1:
        raise ProvenanceError(
            'manifest must contain exactly one %s entry' % ARCHIVE_NAME
        )
    archive_path = (data_root / ARCHIVE_NAME).resolve()
    if archive_path.parent != data_root:
        raise ProvenanceError('archive path escaped the scratch data root')
    if not archive_path.is_file():
        raise ProvenanceError('scratch archive missing: %s' % archive_path)
    observed_hash, observed_bytes = sha256_file(archive_path)
    entry = matches[0]
    if (
        observed_hash != entry.get('sha256')
        or observed_bytes != entry.get('bytes')
    ):
        raise ProvenanceError(
            'scratch archive differs from manifest '
            '(hash=%s bytes=%d)' % (observed_hash, observed_bytes)
        )
    return archive_path, dict(entry)


def optional_manifest_cross_check(archive_path, manifest_path):
    """Compare an archive with a covering manifest entry when one is available."""
    archive_path = Path(archive_path).resolve()
    manifest_path = Path(manifest_path).resolve()
    unavailable = 'manifest cross-check unavailable — direct data-root analysis'
    result = {
        'status': 'UNAVAILABLE',
        'detail': unavailable,
        'manifest': MANIFEST_NAME,
    }
    if not manifest_path.is_file():
        return result
    try:
        with open(manifest_path, encoding='utf-8') as handle:
            manifest = json.load(handle)
    except (OSError, ValueError) as exc:
        result['detail'] = '%s (%s)' % (unavailable, exc)
        return result
    inputs = manifest.get('inputs')
    if not isinstance(inputs, list):
        result['detail'] = '%s (manifest inputs are not a list)' % unavailable
        return result
    matches = [
        item for item in inputs
        if item.get('scratch_name') == archive_path.name
    ]
    if len(matches) != 1:
        result['detail'] = (
            '%s (no unique entry for %s)' % (unavailable, archive_path.name)
        )
        return result
    observed_hash, observed_bytes = sha256_file(archive_path)
    entry = matches[0]
    result.update({
        'status': 'PASS',
        'detail': '%d bytes; sha256=%s' % (observed_bytes, observed_hash),
        'scratch_name': entry.get('scratch_name'),
        'sha256': observed_hash,
        'bytes': observed_bytes,
    })
    if (
        observed_hash != entry.get('sha256')
        or observed_bytes != entry.get('bytes')
    ):
        result.update({
            'status': 'MISMATCH',
            'detail': (
                'optional manifest cross-check mismatch '
                '(observed sha256=%s bytes=%d; expected sha256=%s bytes=%r)'
                % (observed_hash, observed_bytes, entry.get('sha256'),
                   entry.get('bytes'))
            ),
            'expected_sha256': entry.get('sha256'),
            'expected_bytes': entry.get('bytes'),
        })
    return result


def _atomic_bytes(path, payload):
    path = Path(os.path.realpath(os.path.abspath(str(path))))
    physical_worktree = Path(os.path.realpath(str(WORKTREE_ROOT)))
    if not _is_relative_to(path, physical_worktree):
        raise ProvenanceError(
            'write path physically escaped worktree: %s' % path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix='.' + path.name + '.', suffix='.tmp', dir=str(path.parent)
    )
    try:
        with os.fdopen(fd, 'wb') as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, str(path))
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)


def atomic_text(path, text):
    _atomic_bytes(path, text.encode('utf-8'))


def json_text(value):
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        indent=2,
        allow_nan=False,
    ) + '\n'


def atomic_json(path, value):
    atomic_text(path, json_text(value))


def _catalog_text(value, limit=MAX_DECODED_HL_CHARS):
    text = str(value or '')
    if len(text) <= limit:
        return text
    marker = ' [%s]' % TRUNCATED_MARKER
    return text[:max(0, limit - len(marker))] + marker


def decode_hl(detail_link):
    """Decode one DetailLink ``hl`` parameter without accepting malformed data.

    ``MISSING`` covers an absent/blank DetailLink or a link with no hl
    parameter. ``EMPTY`` covers an explicit ``hl=``. ``MALFORMED`` covers
    duplicate parameters, bad percent escapes, bad UTF-8, or URL parsing
    errors. ``TRUNCATED`` is still decodable but exceeds the bounded catalog
    length. Repetition of the same valid query across *different rows* is
    deliberately not an anomaly.
    """
    text = str(detail_link or '').strip()
    if not text:
        return {
            'status': 'MISSING',
            'raw_hl': None,
            'decoded_hl': None,
            'error': 'DetailLink is blank',
        }
    try:
        query = urlsplit(text).query
    except ValueError as exc:
        return {
            'status': 'MALFORMED',
            'raw_hl': None,
            'decoded_hl': None,
            'error': 'URL parse error: %s' % exc,
        }
    matches = []
    for component in query.split('&'):
        name, separator, raw_value = component.partition('=')
        if BAD_PERCENT_RE.search(name):
            return {
                'status': 'MALFORMED',
                'raw_hl': None,
                'decoded_hl': None,
                'error': 'parameter name contains an invalid percent escape',
            }
        try:
            decoded_name = unquote_plus(name, encoding='utf-8', errors='strict')
        except (UnicodeDecodeError, ValueError) as exc:
            return {
                'status': 'MALFORMED',
                'raw_hl': None,
                'decoded_hl': None,
                'error': 'parameter name decoding failed: %s' % exc,
            }
        if decoded_name.casefold() == 'hl':
            matches.append(raw_value if separator else '')
    if not matches:
        return {
            'status': 'MISSING',
            'raw_hl': None,
            'decoded_hl': None,
            'error': 'DetailLink has no hl parameter',
        }
    if len(matches) != 1:
        return {
            'status': 'MALFORMED',
            'raw_hl': None,
            'decoded_hl': None,
            'error': 'DetailLink has %d hl parameters' % len(matches),
        }
    raw_hl = matches[0]
    if raw_hl == '':
        return {
            'status': 'EMPTY',
            'raw_hl': '',
            'decoded_hl': '',
            'error': 'hl parameter is empty',
        }
    if BAD_PERCENT_RE.search(raw_hl):
        return {
            'status': 'MALFORMED',
            'raw_hl': _catalog_text(raw_hl),
            'decoded_hl': None,
            'error': 'hl parameter contains an invalid percent escape',
        }
    try:
        decoded = unquote_plus(raw_hl, encoding='utf-8', errors='strict')
    except (UnicodeDecodeError, ValueError) as exc:
        return {
            'status': 'MALFORMED',
            'raw_hl': _catalog_text(raw_hl),
            'decoded_hl': None,
            'error': 'hl decoding failed: %s' % exc,
        }
    if len(decoded) > MAX_DECODED_HL_CHARS:
        return {
            'status': 'TRUNCATED',
            'raw_hl': _catalog_text(raw_hl),
            'decoded_hl': decoded[:MAX_DECODED_HL_CHARS],
            'decoded_hl_original_length': len(decoded),
            'catalog_marker': TRUNCATED_MARKER,
            'error': (
                'decoded hl exceeded %d characters and was truncated'
                % MAX_DECODED_HL_CHARS
            ),
        }
    return {
        'status': 'OK',
        'raw_hl': _catalog_text(raw_hl),
        'decoded_hl': decoded,
        'decoded_hl_original_length': len(decoded),
        'catalog_marker': None,
        'error': None,
    }


def _unescape_query_token(value):
    output = []
    index = 0
    while index < len(value):
        if value[index] == '\\' and index + 1 < len(value):
            output.append(value[index + 1])
            index += 2
        else:
            output.append(value[index])
            index += 1
    return ''.join(output)


def parse_query_atoms(decoded_query):
    """Lex a decoded Congress.gov query into quoted, bare, and operator atoms."""
    source = str(decoded_query or '')
    atoms = []
    errors = []
    index = 0
    while index < len(source):
        if source[index].isspace():
            index += 1
            continue
        if source[index] == '"':
            start = index
            index += 1
            content = []
            closed = False
            while index < len(source):
                char = source[index]
                if char == '\\' and index + 1 < len(source):
                    content.append(char)
                    content.append(source[index + 1])
                    index += 2
                    continue
                if char == '"':
                    index += 1
                    closed = True
                    break
                content.append(char)
                index += 1
            raw = source[start:index]
            text = _unescape_query_token(''.join(content))
            atoms.append({
                'kind': 'quoted_phrase',
                'raw': raw,
                'text': text,
            })
            if not closed:
                errors.append('unclosed quoted phrase at offset %d' % start)
            continue
        start = index
        while index < len(source) and not source[index].isspace():
            index += 1
        raw = source[start:index]
        text = _unescape_query_token(raw)
        kind = 'operator' if OPERATOR_RE.match(text) else 'bare_term'
        atoms.append({'kind': kind, 'raw': raw, 'text': text})
    quoted = [atom['text'] for atom in atoms if atom['kind'] == 'quoted_phrase']
    bare = [atom['text'] for atom in atoms if atom['kind'] == 'bare_term']
    operators = [atom['text'] for atom in atoms if atom['kind'] == 'operator']
    return {
        'atoms': atoms,
        'query_atom_count': len(atoms),
        'whitespace_token_count': len(source.split()),
        'quoted_phrases': quoted,
        'bare_terms': bare,
        'operators': operators,
        'parse_errors': errors,
    }


def normalize_search_text(value):
    text = unicodedata.normalize('NFKC', str(value or '')).casefold()
    return re.sub(r'\s+', ' ', text).strip()


def sheet_identity(sheet_name):
    name = str(sheet_name or '').strip()
    if name == 'All Bills':
        return {
            'kind': 'all_bills',
            'sheet_number': None,
            'category': None,
        }
    match = SHEET_NUMBER_RE.match(name)
    if not match:
        return {
            'kind': 'other',
            'sheet_number': None,
            'category': BR.canonical_category(name),
        }
    return {
        'kind': 'numbered',
        'sheet_number': int(match.group(1)),
        'category': BR.canonical_category(name),
    }


def analyze_sheet_rows(sheet_name, rows):
    """Analyze synthetic or workbook-adapted row dictionaries for one sheet."""
    identity = sheet_identity(sheet_name)
    normalized_rows = []
    groups = {}
    missing_parameter_rows = []
    empty_rows = []
    malformed_rows = []
    truncated_hl_rows = []
    query_catalog_overflow_rows = 0
    query_catalog_overflow_row_sample = []
    nonempty_hl = 0
    decodable_hl = 0
    for ordinal, source_row in enumerate(rows, start=2):
        row = dict(source_row)
        excel_row = int(row.pop('_excel_row', ordinal))
        title = str(row.get('Title') or '')
        summary = str(row.get('bill_summary') or row.get('Summary') or '')
        search_text = normalize_search_text(title + ' ' + summary)
        decoded = decode_hl(row.get('DetailLink'))
        normalized_rows.append({
            'excel_row': excel_row,
            'search_text': search_text,
        })
        if decoded['status'] == 'MISSING':
            missing_parameter_rows.append(excel_row)
            continue
        if decoded['status'] == 'EMPTY':
            empty_rows.append(excel_row)
            continue
        if decoded.get('raw_hl'):
            nonempty_hl += 1
        if decoded['status'] == 'MALFORMED':
            malformed_rows.append({
                'row': excel_row,
                'raw_hl': decoded.get('raw_hl'),
                'error': decoded.get('error'),
            })
            continue
        if decoded['status'] == 'TRUNCATED':
            truncated_hl_rows.append({
                'row': excel_row,
                'decoded_hl_original_length': (
                    decoded['decoded_hl_original_length']),
                'catalog_marker': TRUNCATED_MARKER,
            })
        decodable_hl += 1
        query = decoded['decoded_hl']
        if (
            query not in groups
            and len(groups) >= MAX_DISTINCT_QUERIES_PER_SHEET
        ):
            query_catalog_overflow_rows += 1
            if len(query_catalog_overflow_row_sample) < 100:
                query_catalog_overflow_row_sample.append(excel_row)
            continue
        group = groups.setdefault(query, {
            'rows': [],
            'raw_hl_variants': set(),
            'detail_link_samples': set(),
            'truncated_rows': [],
            'decoded_hl_original_lengths': set(),
        })
        group['rows'].append(excel_row)
        group['raw_hl_variants'].add(decoded['raw_hl'])
        if decoded['status'] == 'TRUNCATED':
            group['truncated_rows'].append(excel_row)
            group['decoded_hl_original_lengths'].add(
                decoded['decoded_hl_original_length'])
        link = str(row.get('DetailLink') or '')
        if link:
            group['detail_link_samples'].add(_catalog_text(link))

    query_list = []
    sheet_suspects = []
    for decoded_query in sorted(groups):
        group = groups[decoded_query]
        structure = parse_query_atoms(decoded_query)
        phrase_yields = []
        for phrase in sorted(set(structure['quoted_phrases']), key=str.casefold):
            needle = normalize_search_text(phrase)
            hit_rows = sorted(
                item['excel_row']
                for item in normalized_rows
                if needle and needle in item['search_text']
            )
            result = {
                'phrase': phrase,
                'artifact_row_hits': len(hit_rows),
                'hit_rows': hit_rows,
                'search_scope': 'Title + bill_summary within this sheet',
                'status': (
                    'HIT-IN-ARTIFACT'
                    if hit_rows
                    else 'SUSPECT-ZERO-YIELD'
                ),
            }
            phrase_yields.append(result)
            if not hit_rows:
                sheet_suspects.append({
                    'phrase': phrase,
                    'decoded_query': decoded_query,
                    'artifact_row_hits': 0,
                    'search_scope': 'Title + bill_summary within this sheet',
                    'status': 'SUSPECT-ZERO-YIELD',
                })
        query_list.append({
            'decoded_query': decoded_query,
            'catalog_status': (
                TRUNCATED_MARKER if group['truncated_rows'] else 'COMPLETE'
            ),
            'truncated_rows': sorted(group['truncated_rows']),
            'decoded_hl_original_lengths': sorted(
                group['decoded_hl_original_lengths']),
            'raw_hl_variants': sorted(group['raw_hl_variants']),
            'rows': sorted(group['rows']),
            'row_count': len(group['rows']),
            'detail_link_samples': sorted(group['detail_link_samples'])[:3],
            'query_atom_count': structure['query_atom_count'],
            'whitespace_token_count': structure['whitespace_token_count'],
            'quoted_phrases': structure['quoted_phrases'],
            'bare_terms': structure['bare_terms'],
            'operators': structure['operators'],
            'atoms': structure['atoms'],
            'parse_errors': structure['parse_errors'],
            'quoted_phrase_yields': phrase_yields,
        })

    atom_counts = sorted(
        set(query['query_atom_count'] for query in query_list)
    )
    whitespace_counts = sorted(
        set(query['whitespace_token_count'] for query in query_list)
    )
    if identity['category'] == 'Interchange' and query_list:
        audit_comparison = {
            'reported_measure': 'query_atom_count',
            'reported': 1,
            'recomputed': atom_counts,
            'verdict': 'MATCH' if atom_counts == [1] else 'DRIFT',
        }
    elif identity['category'] and query_list:
        audit_comparison = {
            'reported_measure': 'whitespace_token_count within one full query',
            'reported_range': [17, 72],
            'recomputed': whitespace_counts,
            'verdict': (
                'MATCH'
                if whitespace_counts
                and all(17 <= value <= 72 for value in whitespace_counts)
                else 'DRIFT'
            ),
        }
    else:
        audit_comparison = {
            'reported_measure': None,
            'reported': None,
            'recomputed': [],
            'verdict': 'NOT-APPLICABLE',
        }
    missing_hl_rows = sorted(missing_parameter_rows + empty_rows)
    return {
        'sheet_name': str(sheet_name),
        'sheet_kind': identity['kind'],
        'sheet_number': identity['sheet_number'],
        'category': identity['category'],
        'rows_total': len(normalized_rows),
        'rows_nonempty_hl': nonempty_hl,
        'rows_decodable_hl': decodable_hl,
        'rows_missing_hl': len(missing_hl_rows),
        'rows_missing_parameter_hl': len(missing_parameter_rows),
        'rows_empty_hl': len(empty_rows),
        'rows_malformed_hl': len(malformed_rows),
        'rows_truncated_hl': len(truncated_hl_rows),
        'truncated_hl_rows': sorted(
            truncated_hl_rows, key=lambda item: item['row']),
        'missing_hl_row_numbers': missing_hl_rows,
        'missing_parameter_row_numbers': sorted(missing_parameter_rows),
        'empty_hl_row_numbers': sorted(empty_rows),
        'malformed_hl_rows': sorted(
            malformed_rows, key=lambda item: item['row']
        ),
        'distinct_queries': len(query_list),
        'distinct_full_hl_query_count': len(query_list),
        'distinct_full_hl_query_count_is_lower_bound': bool(
            truncated_hl_rows or query_catalog_overflow_rows),
        'query_catalog_status': (
            TRUNCATED_MARKER
            if query_catalog_overflow_rows
            else 'COMPLETE'
        ),
        'query_catalog_limit': MAX_DISTINCT_QUERIES_PER_SHEET,
        'query_catalog_overflow_rows': query_catalog_overflow_rows,
        'query_catalog_overflow_row_sample': sorted(
            query_catalog_overflow_row_sample),
        'query_atom_counts': atom_counts,
        'whitespace_token_counts': whitespace_counts,
        'query_list': query_list,
        'suspect_zero_yield': sorted(
            sheet_suspects,
            key=lambda item: (
                item['phrase'].casefold(),
                item['decoded_query'],
            ),
        ),
        'audit_term_count_comparison': audit_comparison,
    }


def _sheet_rows(workbook_sheet):
    iterator = workbook_sheet.iter_rows(values_only=True)
    try:
        raw_header = next(iterator)
    except StopIteration:
        return []
    header = [str(value or '').strip() for value in raw_header]
    rows = []
    for excel_row, values in enumerate(iterator, start=2):
        if all(
            value is None
            or (isinstance(value, str) and not value.strip())
            for value in values
        ):
            continue
        row = dict(zip(header, values))
        row['_excel_row'] = excel_row
        rows.append(row)
    return rows


def analyze_workbook(path):
    """Analyze the verified workbook; caller must verify the path first."""
    import openpyxl

    book = openpyxl.load_workbook(
        str(path),
        data_only=True,
        read_only=False,
    )
    try:
        if 'All Bills' not in book.sheetnames:
            raise ProvenanceError('archive workbook lacks All Bills')
        sheets = {}
        sheet_order = []
        for name in book.sheetnames:
            result = analyze_sheet_rows(name, _sheet_rows(book[name]))
            sheets[name] = result
            sheet_order.append(name)
    finally:
        book.close()

    numbered = [
        sheets[name] for name in sheet_order
        if sheets[name]['sheet_kind'] == 'numbered'
    ]
    numbers = [item['sheet_number'] for item in numbered]
    counts = {}
    for value in numbers:
        counts[value] = counts.get(value, 0) + 1
    observed = sorted(counts)
    expected = list(range(1, 18))
    numbering = {
        'expected': expected,
        'observed': observed,
        'missing': sorted(set(expected) - set(observed)),
        'unexpected': sorted(set(observed) - set(expected)),
        'duplicates': sorted(
            value for value, count in counts.items() if count > 1
        ),
        'contiguous_01_through_17': (
            observed == expected
            and all(counts[value] == 1 for value in expected)
        ),
    }
    all_bills = sheets['All Bills']
    numbered_rows = sum(item['rows_total'] for item in numbered)
    numbered_nonempty = sum(item['rows_nonempty_hl'] for item in numbered)
    numbered_decodable = sum(item['rows_decodable_hl'] for item in numbered)
    numbered_missing = sum(item['rows_missing_hl'] for item in numbered)
    numbered_malformed = sum(item['rows_malformed_hl'] for item in numbered)
    numbered_truncated = sum(item['rows_truncated_hl'] for item in numbered)
    numbered_query_overflow = sum(
        item['query_catalog_overflow_rows'] for item in numbered)
    totals = {
        'all_bills_rows_total': all_bills['rows_total'],
        'all_bills_rows_nonempty_hl': all_bills['rows_nonempty_hl'],
        'all_bills_rows_decodable_hl': all_bills['rows_decodable_hl'],
        'all_bills_rows_missing_hl': all_bills['rows_missing_hl'],
        'all_bills_rows_malformed_hl': all_bills['rows_malformed_hl'],
        'all_bills_rows_truncated_hl': all_bills['rows_truncated_hl'],
        'all_bills_query_catalog_overflow_rows': (
            all_bills['query_catalog_overflow_rows']),
        'all_bills_distinct_full_hl_queries': (
            all_bills['distinct_full_hl_query_count']
        ),
        'numbered_sheets_rows_total': numbered_rows,
        'numbered_sheets_rows_nonempty_hl': numbered_nonempty,
        'numbered_sheets_rows_decodable_hl': numbered_decodable,
        'numbered_sheets_rows_missing_hl': numbered_missing,
        'numbered_sheets_rows_malformed_hl': numbered_malformed,
        'numbered_sheets_rows_truncated_hl': numbered_truncated,
        'numbered_sheets_query_catalog_overflow_rows': (
            numbered_query_overflow),
        'audit_reported_nonempty_hl': AUDIT_REPORTED_NONEMPTY_HL,
        'all_bills_nonempty_hl_delta': (
            all_bills['rows_nonempty_hl'] - AUDIT_REPORTED_NONEMPTY_HL
        ),
        'numbered_sheets_nonempty_hl_delta': (
            numbered_nonempty - AUDIT_REPORTED_NONEMPTY_HL
        ),
        'physical_nonempty_hl_cells_across_duplicate_views': (
            all_bills['rows_nonempty_hl'] + numbered_nonempty
        ),
        'counting_note': (
            'All Bills and the numbered sheets are duplicate workbook views; '
            'the audit comparison is 8,520 in either view, not their sum.'
        ),
    }
    return {
        'sheet_order': sheet_order,
        'sheets': sheets,
        'sheet_numbering': numbering,
        'totals': totals,
    }


def evaluate_c13(
    analysis,
    canonical_categories=None,
    known_missing=None,
):
    """Return pure C13 decisions; repeated full queries are metadata only."""
    canonical_categories = tuple(
        canonical_categories if canonical_categories is not None
        else BR.CATEGORIES
    )
    known_missing = frozenset(
        known_missing if known_missing is not None
        else QUERY_SET_ZERO_ALLOWLIST
    )
    events = []
    coverage = {}
    for name in analysis['sheet_order']:
        sheet = analysis['sheets'][name]
        total = sheet['rows_total']
        missing = sheet['rows_missing_parameter_hl']
        empty = sheet['rows_empty_hl']
        malformed = sheet['rows_malformed_hl']
        if total == 0:
            status = 'WARN'
            detail = '0 data rows; hl coverage unavailable'
        elif (
            missing == 0
            and empty == 0
            and malformed == 0
            and sheet['rows_decodable_hl'] == total
        ):
            status = 'PASS'
            detail = '%d/%d rows carry decodable hl' % (total, total)
        else:
            status = 'WARN'
            detail = (
                '%d/%d decodable; missing=%d empty=%d malformed=%d'
                % (
                    sheet['rows_decodable_hl'],
                    total,
                    missing,
                    empty,
                    malformed,
                )
            )
        coverage[name] = {
            'status': status,
            'rows_total': total,
            'rows_decodable_hl': sheet['rows_decodable_hl'],
            'rows_missing_parameter_hl': missing,
            'rows_empty_hl': empty,
            'rows_malformed_hl': malformed,
            'distinct_full_hl_query_count': (
                sheet['distinct_full_hl_query_count']
            ),
        }
        events.append({
            'status': status,
            'name': 'C13 hl coverage: %s' % name,
            'detail': detail,
        })

    numbering = analysis['sheet_numbering']
    if numbering['contiguous_01_through_17']:
        events.append({
            'status': 'PASS',
            'name': 'C13 numbered archive sheets are contiguous 01..17',
            'detail': '17 unique numbered sheets',
        })
    else:
        events.append({
            'status': 'FAIL',
            'name': 'C13 numbered archive sheet sequence is broken',
            'detail': 'missing=%s unexpected=%s duplicates=%s' % (
                numbering['missing'],
                numbering['unexpected'],
                numbering['duplicates'],
            ),
        })

    single_atom = []
    configured_categories = set()
    for name in analysis['sheet_order']:
        sheet = analysis['sheets'][name]
        if sheet['sheet_kind'] != 'numbered':
            continue
        if sheet['category'] and sheet['distinct_full_hl_query_count'] > 0:
            configured_categories.add(sheet['category'])
        for query in sheet['query_list']:
            if query['query_atom_count'] == 1:
                single_atom.append({
                    'sheet': name,
                    'category': sheet['category'],
                    'decoded_query': query['decoded_query'],
                    'row_count': query['row_count'],
                })
    single_atom = sorted(
        single_atom,
        key=lambda item: (
            item['sheet'],
            item['decoded_query'],
        ),
    )
    if single_atom:
        events.append({
            'status': 'WARN',
            'name': 'C13 single-atom API-browser query sets',
            'detail': ', '.join(
                '%s (%d rows)' % (item['sheet'], item['row_count'])
                for item in single_atom
            ),
        })
    else:
        events.append({
            'status': 'PASS',
            'name': 'C13 no single-atom API-browser query sets',
            'detail': '',
        })

    canonical_set = set(canonical_categories)
    missing_sets = canonical_set - configured_categories
    documented = sorted(missing_sets & known_missing)
    unexpected = sorted(missing_sets - known_missing)
    stale = sorted(configured_categories & known_missing)
    if documented:
        events.append({
            'status': 'WARN',
            'name': 'C13 documented canonical categories with no query set',
            'detail': ', '.join(documented),
        })
    if unexpected:
        events.append({
            'status': 'FAIL',
            'name': 'C13 canonical categories with undocumented query gaps',
            'detail': ', '.join(unexpected),
        })
    else:
        events.append({
            'status': 'PASS',
            'name': 'C13 no undocumented canonical query-set gaps',
            'detail': (
                'documented known: %s'
                % (', '.join(documented) or 'none')
            ),
        })
    if stale:
        events.append({
            'status': 'WARN',
            'name': 'C13 query-set zero allowlist entries are stale',
            'detail': ', '.join(stale),
        })
    return {
        'events': events,
        'coverage_by_sheet': coverage,
        'sheet_numbering': numbering,
        'single_atom_queries': single_atom,
        'query_sets': {
            'configured': sorted(configured_categories),
            'documented_missing': documented,
            'unexpected_missing': unexpected,
            'stale_allowlist': stale,
        },
    }


def build_document(analysis, manifest_entry):
    c13 = evaluate_c13(analysis)
    return {
        'schema': OUTPUT_SCHEMA,
        'producer_provenance': {
            'producer': 'manual Congress.gov API-browser',
            'statement_by': 'Paul',
            'statement_date': '2026-07-25',
            'producer_script': None,
            'recovery_scope': 'query specification only',
            'recovery_boundary': (
                'No data row was invented, re-fetched, or modified.'
            ),
            'api_browser_signature': (
                'one hand-built query per category, pasted across its result rows'
            ),
        },
        'source': {
            'scratch_name': manifest_entry.get('scratch_name'),
            'source_path': manifest_entry.get('source_path'),
            'sha256': manifest_entry.get('sha256'),
            'bytes': manifest_entry.get('bytes'),
            'source_mtime': manifest_entry.get('source_mtime'),
            'manifest': MANIFEST_NAME,
            'manifest_schema': MANIFEST_SCHEMA,
        },
        'audit_reported': {
            'nonempty_hl': AUDIT_REPORTED_NONEMPTY_HL,
            'other_category_whitespace_token_range': [17, 72],
            'interchange_query_atom_count': 1,
            'citation': (
                'LEGISLATIVE_AUDIT_2026-07-24.md §Q2 and Tier 3 item 10'
            ),
        },
        'count_definitions': {
            'distinct_full_hl_query_count': (
                'unique complete URL-decoded hl strings within a sheet'
            ),
            'query_atom_count': (
                'quoted phrase, bare term, or recognized operator; '
                'each counts once'
            ),
            'whitespace_token_count': (
                'len(decoded_query.split()); historical audit 17-72 measure'
            ),
            'truncation_policy': (
                'decoded hl values are capped at 10,000 characters and '
                'distinct-query catalogs at 1,000 entries per sheet; '
                'affected rows carry explicit TRUNCATED status/count fields'
            ),
        },
        'totals': analysis['totals'],
        'sheet_order': analysis['sheet_order'],
        'sheet_numbering': analysis['sheet_numbering'],
        'query_sets': c13['query_sets'],
        'dictionary_check': {
            'status': 'OMITTED',
            'reason': (
                'No deterministic dictionary or approved known-vocabulary '
                'source exists in the offline input set.'
            ),
        },
        'c13_assessment': c13,
        'sheets': analysis['sheets'],
        'method_citations': [
            'AUDIT_SP/sections.py:14-32 (All Bills and per-sheet traversal semantics)',
            'fec/bill_rules.py:16-58,110-122 (canonical category semantics)',
            'fec/legislation/leg_evidence_recompute.py:300-327,415-450 '
            '(atomic output and manifest-verification house style)',
        ],
    }


def _md_escape(value):
    return str(value).replace('|', '\\|').replace('\r', ' ').replace('\n', ' ')


def _md_query_text(value):
    """Make decoded query text inert and single-line in generated Markdown."""
    rendered = []
    for character in str(value):
        codepoint = ord(character)
        if character == '`':
            rendered.append('\\x60')
        elif character == '\n':
            rendered.append('\\n')
        elif character == '\r':
            rendered.append('\\r')
        elif character == '\t':
            rendered.append('\\t')
        elif codepoint < 32 or codepoint == 127 or 128 <= codepoint <= 159:
            rendered.append('\\x%02X' % codepoint)
        else:
            rendered.append(character)
    return ''.join(rendered)


def _row_ranges(values):
    values = sorted(set(int(value) for value in values))
    if not values:
        return 'none'
    groups = []
    start = previous = values[0]
    for value in values[1:]:
        if value == previous + 1:
            previous = value
            continue
        groups.append(
            str(start) if start == previous else '%d-%d' % (start, previous)
        )
        start = previous = value
    groups.append(
        str(start) if start == previous else '%d-%d' % (start, previous)
    )
    return ', '.join(groups)


def render_markdown(document):
    totals = document['totals']
    lines = [
        '# Bill Archive Congress.gov Query Provenance',
        '',
        '## Provenance statement',
        '',
        '- Producer: **manual Congress.gov API-browser**.',
        '- Statement: Paul, 2026-07-25.',
        '- There is not and never was a producer script.',
        '- Recovery scope: query specification only. No data row was invented, '
        're-fetched, or modified.',
        '- API-browser signature: **one hand-built query per category, pasted '
        'across its result rows**.',
        '',
        'Source scratch copy: `%s`  ' % document['source']['scratch_name'],
        'SHA-256: `%s`  ' % document['source']['sha256'],
        'Bytes: `%s`' % document['source']['bytes'],
        '',
        '## Recovery method and count definitions',
        '',
        'Every nonblank `DetailLink` is inspected for an `hl=` parameter. The '
        'parameter is validated, URL-decoded, and lexed into quoted phrases, '
        'bare terms, and operators. Repeated full queries across rows are '
        'provenance, not duplication errors.',
        '',
        '- **Distinct full queries:** unique complete decoded `hl` strings.',
        '- **Query atoms:** each quoted phrase, bare term, or operator counts once.',
        '- **Whitespace tokens:** `len(decoded_query.split())`; this is the '
        'audit’s historical 17–72 measure.',
        '',
        'Quoted phrases are checked only against `Title + bill_summary` rows in '
        'their own sheet. `SUSPECT-ZERO-YIELD` means zero yield in this artifact, '
        'not a claim about current Congress.gov results.',
        '',
        '## 8,520 reconciliation',
        '',
        '| View | Rows | Non-empty hl | Decodable hl | Missing/empty | Malformed | Delta vs 8,520 |',
        '|---|---:|---:|---:|---:|---:|---:|',
        '| All Bills | %d | %d | %d | %d | %d | %+d |' % (
            totals['all_bills_rows_total'],
            totals['all_bills_rows_nonempty_hl'],
            totals['all_bills_rows_decodable_hl'],
            totals['all_bills_rows_missing_hl'],
            totals['all_bills_rows_malformed_hl'],
            totals['all_bills_nonempty_hl_delta'],
        ),
        '| Numbered sheets | %d | %d | %d | %d | %d | %+d |' % (
            totals['numbered_sheets_rows_total'],
            totals['numbered_sheets_rows_nonempty_hl'],
            totals['numbered_sheets_rows_decodable_hl'],
            totals['numbered_sheets_rows_missing_hl'],
            totals['numbered_sheets_rows_malformed_hl'],
            totals['numbered_sheets_nonempty_hl_delta'],
        ),
        '',
        'The two rows above are duplicate workbook views of the same archive '
        'records. Their physical-cell sum is not the audit denominator.',
        '',
        '## Per-sheet summary',
        '',
        '| Sheet | Category | Rows | Decodable | Missing/empty | Malformed | '
        'Distinct full queries | Atom counts | Whitespace-token counts | Suspects |',
        '|---|---|---:|---:|---:|---:|---:|---|---|---:|',
    ]
    for name in document['sheet_order']:
        sheet = document['sheets'][name]
        atom_counts = _md_escape(sheet['query_atom_counts'])
        whitespace_counts = _md_escape(sheet['whitespace_token_counts'])
        if name == 'All Bills':
            atom_counts = 'deduplicated set %s' % atom_counts
            whitespace_counts = 'deduplicated set %s' % whitespace_counts
        lines.append(
            '| %s | %s | %d | %d | %d | %d | %d | %s | %s | %d |'
            % (
                _md_escape(name),
                _md_escape(sheet['category'] or 'mixed / none'),
                sheet['rows_total'],
                sheet['rows_decodable_hl'],
                sheet['rows_missing_hl'],
                sheet['rows_malformed_hl'],
                sheet['distinct_full_hl_query_count'],
                atom_counts,
                whitespace_counts,
                len(sheet['suspect_zero_yield']),
            )
        )
    lines.extend([
        '',
        '## Never-configured category evidence',
        '',
        '- Numbered archive sheets are `%s`.' % (
            'contiguous 01..17'
            if document['sheet_numbering']['contiguous_01_through_17']
            else 'not contiguous'
        ),
        '- `12 CDFI Fund` has zero data rows and no query set.',
        '- No `CU Tax Status` numbered sheet or query set exists.',
        '- Documented-known missing query sets: %s.' % (
            ', '.join(document['query_sets']['documented_missing']) or 'none'
        ),
        '',
        '## Per-sheet query catalog',
        '',
    ])
    for name in document['sheet_order']:
        sheet = document['sheets'][name]
        lines.extend([
            '### %s' % name,
            '',
            'Rows: %d; distinct full queries cataloged: %d; '
            'missing/empty hl rows: %d; malformed hl rows: %d; '
            'decoded-hl truncations: %d; query-catalog status: %s; '
            'overflow rows: %d.'
            % (
                sheet['rows_total'],
                sheet['distinct_full_hl_query_count'],
                sheet['rows_missing_hl'],
                sheet['rows_malformed_hl'],
                sheet['rows_truncated_hl'],
                sheet['query_catalog_status'],
                sheet['query_catalog_overflow_rows'],
            ),
            '',
        ])
        if not sheet['query_list']:
            lines.extend(['No recovered query set.', ''])
            continue
        for index, query in enumerate(sheet['query_list'], start=1):
            lines.extend([
                '**Query %d** — %d rows; Excel rows %s' % (
                    index,
                    query['row_count'],
                    _row_ranges(query['rows']),
                ),
                '',
                '```text',
                _md_query_text(query['decoded_query']),
                '```',
                '',
                '- Query atoms: %d' % query['query_atom_count'],
                '- Whitespace tokens: %d' % query['whitespace_token_count'],
                '- Catalog status: %s%s' % (
                    query['catalog_status'],
                    (
                        ' (decoded hl truncated at %d characters)'
                        % MAX_DECODED_HL_CHARS
                        if query['catalog_status'] == TRUNCATED_MARKER
                        else ''
                    ),
                ),
                '- Quoted phrases: %s' % (
                    ', '.join(
                        '`%s`' % _md_query_text(value)
                        for value in query['quoted_phrases']
                    )
                    or 'none'
                ),
                '- Bare terms: %s' % (
                    ', '.join(
                        '`%s`' % _md_query_text(value)
                        for value in query['bare_terms']
                    )
                    or 'none'
                ),
                '- Operators: %s' % (
                    ', '.join(
                        '`%s`' % _md_query_text(value)
                        for value in query['operators']
                    )
                    or 'none'
                ),
                '',
            ])
    suspects = []
    for name in document['sheet_order']:
        if name != 'All Bills':
            continue
        query_identifiers = {
            query['decoded_query']: 'Query %d' % index
            for index, query in enumerate(
                document['sheets'][name]['query_list'], start=1
            )
        }
        for item in document['sheets'][name]['suspect_zero_yield']:
            suspects.append(
                (name, query_identifiers[item['decoded_query']], item)
            )
    distinct_suspect_phrases = len(
        {item['phrase'] for _name, _query_identifier, item in suspects}
    )
    lines.extend([
        '## SUSPECT-ZERO-YIELD quoted phrases',
        '',
        'Counts are per query-occurrence: %d occurrences / %d distinct phrases.'
        % (len(suspects), distinct_suspect_phrases),
        '',
        '| Sheet | Query identifier | Phrase | Artifact hits | Scope |',
        '|---|---|---|---:|---|',
    ])
    if suspects:
        for name, query_identifier, item in suspects:
            lines.append(
                '| %s | %s | `%s` | %d | %s |'
                % (
                    _md_escape(name),
                    _md_escape(query_identifier),
                    _md_escape(_md_query_text(item['phrase'])),
                    item['artifact_row_hits'],
                    _md_escape(item['search_scope']),
                )
            )
    else:
        lines.append('| none | — | — | 0 | — |')
    lines.extend([
        '',
        'No network claim is made by this list.',
        '',
        '## Dictionary check',
        '',
        '**OMITTED.** %s' % document['dictionary_check']['reason'],
        '',
        '## C13 decision summary',
        '',
        '| Status | Check | Detail |',
        '|---|---|---|',
    ])
    for event in document['c13_assessment']['events']:
        lines.append(
            '| %s | %s | %s |'
            % (
                event['status'],
                _md_escape(event['name']),
                _md_escape(event['detail']),
            )
        )
    lines.extend([
        '',
        '## Method citations',
        '',
    ])
    for citation in document['method_citations']:
        lines.append('- `%s`' % citation)
    lines.append('')
    return '\n'.join(lines)


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description='Recover query provenance from the bill archive scratch copy.'
    )
    parser.add_argument(
        '--data-root',
        required=True,
        help='approved legfix2 scratch data directory',
    )
    parser.add_argument(
        '--out',
        required=True,
        help='must resolve to audits/legfix2 in this worktree',
    )
    return parser.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    try:
        data_root, out = approved_paths(args.data_root, args.out)
        manifest_path = out / MANIFEST_NAME
        archive_path, manifest_entry = load_verified_archive(
            data_root, manifest_path
        )
        print('VERIFIED %s' % archive_path.name)
        analysis = analyze_workbook(archive_path)
        document = build_document(analysis, manifest_entry)
        atomic_json(out / JSON_NAME, document)
        atomic_text(out / MD_NAME, render_markdown(document))
        print('WROTE %s' % (out / JSON_NAME))
        print('WROTE %s' % (out / MD_NAME))
        print(
            'HL all_bills=%d numbered=%d delta=%d'
            % (
                document['totals']['all_bills_rows_nonempty_hl'],
                document['totals']['numbered_sheets_rows_nonempty_hl'],
                document['totals']['all_bills_nonempty_hl_delta'],
            )
        )
        return 0
    except Exception as exc:
        print('BLOCKED: %s' % exc, file=sys.stderr)
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
