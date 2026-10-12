"""Committee score is a chain producer, not a hand-maintained column.

Pins the scoring rule (engine/apply_committee_reweight.py's 2026-06-30 port), the two
deliberate departures from the lane scripts -- the deleted 'Small Business' override and the
now-fatal bracketless token -- the check/write contract, and a live-data control that the
shipped workbook already agrees with its own seat text.
"""
from __future__ import annotations

import hashlib
import io
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest

import openpyxl

FEC = Path(__file__).resolve().parents[2]
ROOT = FEC.parent
if str(FEC) not in sys.path:
    sys.path.insert(0, str(FEC))

import engine.committee_scores as cs

WORKBOOK = ROOT / 'Moneyball Main Version - Blended.xlsx'


def build_workbook(path, rows):
    """A minimal stand-in for the live sheet: the real header, the given rows, and a second
    sheet carrying a FORMULA so every test proves the save did not flatten formulas."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = cs.SHEET
    for column, label in enumerate(cs.EXPECTED_HEADER, start=1):
        sheet.cell(1, column).value = label
    for index, (name, score, membership, leadership, seats) in enumerate(rows, start=2):
        sheet.cell(index, 1).value = name
        sheet.cell(index, 2).value = score
        sheet.cell(index, 3).value = membership
        sheet.cell(index, 4).value = leadership
        sheet.cell(index, 5).value = seats
    other = book.create_sheet('Untouched')
    other['A1'] = 'keep me'
    other['A2'] = '=SUM(1,2)'
    book.save(path)
    book.close()


def priced(seats):
    """(membership, leadership, score) -- the module's own answer for one row's seat text."""
    return cs.compute(seats)


def row_of(path, index):
    book = openpyxl.load_workbook(path, data_only=False)
    sheet = book[cs.SHEET]
    values = tuple(sheet.cell(index, column).value for column in (2, 3, 4, 5))
    book.close()
    return values


def call(path, **kwargs):
    stream = io.StringIO()
    code = cs.run(str(path), stream=stream, **kwargs)
    return code, stream.getvalue()


class ScoringRuleTests(unittest.TestCase):
    """The weights that engine/apply_committee_reweight.py ported from Excel on 2026-06-30."""

    def test_three_markers_carry_the_documented_full_committee_weights(self):
        for marker, weight in (('Highly', 3.0), ('Moderately', 1.8), ('Slightly', 0.9)):
            membership, leadership = cs.seat_weights('Banking (full) [%s]' % marker)
            self.assertEqual(membership, weight, marker)
            self.assertEqual(leadership, 0.0, marker)
            self.assertEqual(cs.compute('Banking (full) [%s]' % marker),
                             (weight, 0.0, round(weight / 2, 6)))

    def test_a_subcommittee_seat_weighs_one_sixth_of_its_parent(self):
        for marker, weight in (('Highly', 3.0), ('Moderately', 1.8), ('Slightly', 0.9)):
            membership, _ = cs.seat_weights('Banking - Digital Assets [%s]' % marker)
            self.assertAlmostEqual(membership, weight / 6.0, places=12, msg=marker)
        # ' - ' is the ONLY subcommittee signal; a hyphen without spaces is a full committee.
        self.assertEqual(cs.seat_weights('Ways & Means (full) [Highly]')[0], 3.0)
        self.assertEqual(cs.seat_weights('Health-Care Reform (full) [Highly]')[0], 3.0)

    def test_chair_and_ranking_multiply_the_seats_own_membership_weight(self):
        # full committee: Chair x3.0, Ranking x2.0, Vice x1.0
        self.assertEqual(cs.seat_weights('Banking (full) [Highly/Chair]'), (3.0, 9.0))
        self.assertEqual(cs.seat_weights('Banking (full) [Highly/Ranking]'), (3.0, 6.0))
        self.assertEqual(cs.seat_weights('Banking (full) [Highly/Vice]'), (3.0, 3.0))
        self.assertEqual(cs.seat_weights('Banking (full) [Slightly/Chair]'), (0.9, 2.7))
        # subcommittee: Chair x7.5, Ranking x4.5, Vice x2.0, applied to base/6
        membership, leadership = cs.seat_weights('Banking - Digital Assets [Highly/Chair]')
        self.assertAlmostEqual(membership, 0.5, places=12)
        self.assertAlmostEqual(leadership, 3.75, places=12)
        membership, leadership = cs.seat_weights('Banking - Digital Assets [Highly/Ranking]')
        self.assertAlmostEqual(leadership, 2.25, places=12)
        membership, leadership = cs.seat_weights('Banking - Digital Assets [Highly/Vice]')
        self.assertAlmostEqual(leadership, 1.0, places=12)

    def test_committee_score_is_the_mean_of_membership_and_leadership(self):
        seats = ('Banking (full) [Highly] | Banking - Digital Assets [Moderately] | '
                 "Veterans' Affairs (full) [Slightly/Ranking]")
        membership, leadership, score = cs.compute(seats)
        self.assertAlmostEqual(membership, 3.0 + 1.8 / 6 + 0.9, places=12)
        self.assertAlmostEqual(leadership, 0.9 * 2.0, places=12)
        self.assertAlmostEqual(score, (membership + leadership) / 2, places=12)

    def test_a_member_with_no_seats_scores_zero(self):
        self.assertEqual(cs.compute(''), (0.0, 0.0, 0.0))
        self.assertEqual(cs.compute(None), (0.0, 0.0, 0.0))


class SmallBusinessOverrideTests(unittest.TestCase):
    """The 2026-09-18 ruling: the DISPLAYED marker is authoritative.

    The lane scripts force every 'Small Business' seat to [Highly] whatever the text says.
    It was proved inert before deletion -- all 74 live Small Business tokens already read
    [Highly] -- so these pins are a negative control against the override coming back."""

    def test_small_business_highly_is_unchanged_by_the_deletion(self):
        self.assertEqual(cs.seat_weights('Small Business (full) [Highly]'), (3.0, 0.0))
        self.assertEqual(cs.seat_weights('Small Business (full) [Highly/Chair]'), (3.0, 9.0))
        membership, _ = cs.seat_weights(
            'Small Business - Rural Development, Energy & Supply Chains [Highly]')
        self.assertAlmostEqual(membership, 0.5, places=12)

    def test_a_retiered_small_business_seat_is_now_priced_as_written(self):
        # Under the deleted override both of these would still have scored 3.0.
        self.assertEqual(cs.seat_weights('Small Business (full) [Moderately]'), (1.8, 0.0))
        self.assertEqual(cs.seat_weights('Small Business (full) [Slightly]'), (0.9, 0.0))
        self.assertNotIn('Small Business',
                         (FEC / 'engine' / 'committee_scores.py').read_text(
                             encoding='utf-8').split('SCOPE AND SAFETY')[-1],
                         'the override must not be reintroduced below the docstring')


class MalformedSeatTextTests(unittest.TestCase):
    """The lane scripts skip these silently, which scores the seat as zero. This one refuses."""

    def test_a_bracketless_token_raises_instead_of_scoring_zero(self):
        with self.assertRaises(cs.CommitteeSeatError) as caught:
            cs.compute('Banking (full) [Highly] | Judiciary (full) Slightly')
        self.assertIn('no [marker]', str(caught.exception))

    def test_an_unknown_marker_or_role_raises(self):
        with self.assertRaises(cs.CommitteeSeatError):
            cs.compute('Banking (full) [Somewhat]')
        with self.assertRaises(cs.CommitteeSeatError):
            cs.compute('Banking (full) [Highly/Secretary]')

    def test_a_bracketless_token_names_its_row_and_member(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'wb.xlsx')
            build_workbook(path, [('Alice Example', 3.0, 3.0, 0.0, 'Banking (full) Highly')])
            with self.assertRaises(cs.CommitteeSeatError) as caught:
                cs.run(path, write=True, stream=io.StringIO())
            message = str(caught.exception)
            self.assertIn('row 2', message)
            self.assertIn('Alice Example', message)


class CheckAndWriteTests(unittest.TestCase):
    """The defect, reproduced: a marker edited in E with B/C/D left behind."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self._tmp.name, 'wb.xlsx')
        self.addCleanup(self._tmp.cleanup)
        # Row 2 is correct. Row 3 carries the [Moderately] weights its text no longer claims --
        # exactly what retier_committee.py left behind on 44 rows on 2026-09-18.
        stale_membership, stale_leadership, stale_score = cs.compute(
            "Banking (full) [Highly] | Veterans' Affairs (full) [Moderately]")
        self.edited_seats = "Banking (full) [Highly] | Veterans' Affairs (full) [Slightly]"
        clean = cs.compute('Finance (full) [Highly/Chair]')
        build_workbook(self.path, [
            ('Clean Member', clean[2], clean[0], clean[1], 'Finance (full) [Highly/Chair]'),
            ('Retiered Member', stale_score, stale_membership, stale_leadership,
             self.edited_seats),
        ])
        self.truth = cs.compute(self.edited_seats)

    def test_check_flags_the_edited_row_writes_nothing_and_exits_nonzero(self):
        before = Path(self.path).read_bytes()
        code, output = call(self.path)
        self.assertEqual(code, cs.EXIT_DISAGREE)
        self.assertIn('Retiered Member', output)
        self.assertIn('1 disagreements', output)
        self.assertNotIn('Clean Member', output)
        self.assertEqual(Path(self.path).read_bytes(), before, '--check must not write')

    def test_write_repairs_the_edited_row_and_leaves_the_clean_row_alone(self):
        clean_before = row_of(self.path, 2)
        code, output = call(self.path, write=True)
        self.assertEqual(code, cs.EXIT_OK)
        self.assertIn('1 repaired', output)
        score, membership, leadership, seats = row_of(self.path, 3)
        self.assertEqual((membership, leadership, score),
                         (self.truth[0], self.truth[1], self.truth[2]))
        self.assertEqual(seats, self.edited_seats, 'column E must never be rewritten')
        self.assertEqual(row_of(self.path, 2), clean_before)
        self.assertEqual(call(self.path)[0], cs.EXIT_OK, '--check must pass after --write')

    def test_write_is_idempotent_and_a_clean_tree_is_not_rewritten(self):
        self.assertEqual(call(self.path, write=True)[0], cs.EXIT_OK)
        repaired = Path(self.path).read_bytes()
        stat = os.stat(self.path)
        code, output = call(self.path, write=True)
        self.assertEqual(code, cs.EXIT_OK)
        self.assertIn('0 repaired', output)
        self.assertIn('not rewritten', output)
        self.assertEqual(Path(self.path).read_bytes(), repaired,
                         'a second --write must not move one byte')
        self.assertEqual(os.stat(self.path).st_mtime_ns, stat.st_mtime_ns,
                         'a clean --write must not even touch the file')

    def test_a_repairing_write_keeps_the_other_sheet_and_its_formula(self):
        call(self.path, write=True)
        book = openpyxl.load_workbook(self.path, data_only=False)
        self.assertIn('Untouched', book.sheetnames)
        self.assertEqual(book['Untouched']['A1'].value, 'keep me')
        self.assertEqual(book['Untouched']['A2'].value, '=SUM(1,2)')
        book.close()

    def test_a_formula_in_the_cached_columns_is_refused_not_overwritten(self):
        book = openpyxl.load_workbook(self.path, data_only=False)
        book[cs.SHEET].cell(3, 2).value = '=C3+D3'
        book.save(self.path)
        book.close()
        before = Path(self.path).read_bytes()
        self.assertEqual(call(self.path, write=True)[0], cs.EXIT_REFUSED)
        self.assertEqual(Path(self.path).read_bytes(), before)

    def test_an_excel_lock_file_refuses_both_modes(self):
        lock = os.path.join(os.path.dirname(self.path), '~$' + os.path.basename(self.path))
        open(lock, 'wb').close()
        self.addCleanup(lambda: os.path.exists(lock) and os.remove(lock))
        self.assertEqual(call(self.path)[0], cs.EXIT_REFUSED)
        self.assertEqual(call(self.path, write=True)[0], cs.EXIT_REFUSED)


class ChainWiringTests(unittest.TestCase):
    def test_the_producer_runs_immediately_before_compute_pillars_write(self):
        chain = (FEC / 'run_pipeline.py').read_text(encoding='utf-8')
        producer = "('committee_scores.py',['--write']),"
        consumer = "('compute_pillars.py',['--write']),"
        self.assertIn(producer, chain)
        self.assertIn(consumer, chain)
        between = chain.split(producer)[1].split(consumer)[0]
        self.assertNotIn("('", between,
                         'no chain step may sit between the producer and compute_pillars')


class LiveWorkbookControlTests(unittest.TestCase):
    """The shipped workbook must already agree with its own seat text.

    This is the gate the producer exists to hold: if it ever fails, some row's published
    Influence contradicts the seats printed beside it on the member desk."""

    def setUp(self):
        if not WORKBOOK.is_file():
            self.skipTest('live workbook absent (%s)' % WORKBOOK)

    def test_every_live_row_already_agrees(self):
        stream = io.StringIO()
        code = cs.run(str(WORKBOOK), stream=stream)
        output = stream.getvalue()
        self.assertEqual(code, cs.EXIT_OK, output)
        self.assertIn('0 disagreements', output)

    def test_check_does_not_touch_the_workbook_it_reads(self):
        # Deliberately a COPY, not the live file. The property under test belongs to --check;
        # asserting it against a shared artifact that another session may be writing would let
        # an unrelated concurrent save fail this suite for a reason that is not about --check.
        with tempfile.TemporaryDirectory() as tmp:
            copy = os.path.join(tmp, WORKBOOK.name)
            shutil.copy2(str(WORKBOOK), copy)
            before = os.stat(copy)
            digest = hashlib.md5(Path(copy).read_bytes()).hexdigest()
            self.assertEqual(cs.run(copy, stream=io.StringIO()), cs.EXIT_OK)
            after = os.stat(copy)
            self.assertEqual((before.st_size, before.st_mtime_ns),
                             (after.st_size, after.st_mtime_ns),
                             '--check must not touch the workbook')
            self.assertEqual(hashlib.md5(Path(copy).read_bytes()).hexdigest(), digest)

    def test_every_live_seat_token_is_priceable_and_small_business_reads_highly(self):
        book = openpyxl.load_workbook(str(WORKBOOK), data_only=False)
        sheet = book[cs.SHEET]
        tokens = 0
        small_business = 0
        for index in range(2, sheet.max_row + 1):
            if not sheet.cell(index, 1).value:
                continue
            for token in str(sheet.cell(index, 5).value or '').split(cs.SEPARATOR):
                if not token.strip():
                    continue
                tokens += 1
                cs.seat_weights(token)          # raises on anything unpriceable
                if 'Small Business' in token:
                    small_business += 1
                    self.assertIn('[Highly', token,
                                  'a Small Business seat is no longer [Highly]; the deleted '
                                  'override was load-bearing after all -- tell Paul before '
                                  'the next --write')
        book.close()
        self.assertGreater(tokens, 0)
        self.assertGreater(small_business, 0)


if __name__ == '__main__':
    unittest.main()
