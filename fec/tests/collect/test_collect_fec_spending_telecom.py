# -*- coding: utf-8 -*-
"""Telecom's compact cycle-sheet schema through the FEC spending collector.

Synthetic workbook only: it reproduces the real Telecom workbook's shape (12
column cycle sheets with header on row 4, no Cycle or Committee ID column, Fact
Type Master/Employee/Recipient, Recipient Key FEC:C########) and never reads a
Mentat file.
"""
import os
import sys
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from collect import collect_fec_spending as C  # noqa: E402

CYCLES = C.DETAIL_SHEETS
COMPACT_HEADER = [
    "Fact Type", "Company", "Employee Tier", "Amount", "Analysis Amount",
    "Party", "Recipient Category", "Recipient Key", "Gift Count",
    "Include in Analysis", "Person Key", "Logic Check",
]
FULL_HEADER = [
    "Source ID", "Cycle", "Company", "Contributor", "Employee Tier", "Amount",
    "Analysis Amount", "Party", "Recipient Category", "Recipient",
    "Committee ID", "Recipient Key", "Date", "Include in Analysis",
]
# Recipient-level facts per cycle: (category, key, amount, company)
RECIPIENT_FACTS = {
    "2021-2022": [
        ("Candidate", "FEC:C00000001", "100.00", "AT&T"),
        ("Leadership PAC", "FEC:C00000002", "50.00", "Verizon"),
        ("Super PAC", "FEC:C00000003", "25.00", "Comcast"),
    ],
    "2023-2024": [("Candidate", "FEC:C00000001", "200.00", "AT&T")],
    "2025-2026": [("Candidate", "FEC:C00000001", "300.00", "T-Mobile")],
}


def _compact_sheet(wb, title, facts, exception=False):
    ws = wb.create_sheet(title)
    ws.append([f"{title} Telecom Formula Facts"])
    ws.append(["subtitle"])
    ws.append(["Source DB SHA-256: synthetic"])
    ws.append(COMPACT_HEADER)
    for category, key, amount, company in facts:
        if exception:
            ws.append(["Exception", company, None, float(amount), 0, "Other / Unclassified",
                       None, key, 1, "No", None, "OK"])
            continue
        for fact_type, recipient_key in (("Master", None), ("Employee", None), ("Recipient", key)):
            ws.append([fact_type, company, "Employee", float(amount), float(amount), "D",
                       category, recipient_key, 3, "Yes",
                       "P|X" if fact_type == "Employee" else None, "OK"])
    return ws


def _with_header_row_5(wb, title, header, rows):
    ws = wb.create_sheet(title)
    for _ in range(4):
        ws.append([title])
    ws.append(header)
    for row in rows:
        ws.append(row)
    return ws


def _build_workbook(path, recipient_facts=None, exception_rows=2):
    recipient_facts = recipient_facts or RECIPIENT_FACTS
    wb = Workbook()
    wb.remove(wb.active)
    totals = {}
    for cycle in CYCLES:
        _compact_sheet(wb, cycle, recipient_facts[cycle])
        totals[cycle] = sum(float(amount) for _c, _k, amount, _co in recipient_facts[cycle])
    _compact_sheet(
        wb, C.EXCEPTIONS_SHEET,
        [("", "FEC:C00000009", "5.00", "AT&T")] * exception_rows, exception=True,
    )
    map_header = ["Cycle", "Recipient Committee", "Committee ID", "Lookup Key", "Recipient Category",
                  "Direct Candidate / Sponsor", "Candidate ID(s)", "Relationship"]
    map_rows = []
    for cycle in CYCLES:
        map_rows.append([cycle, "FRIENDS OF ONE", "C00000001", f"{cycle}|C00000001", "Candidate",
                         "ONE, ADA", "H0AA00001", "Principal campaign committee"])
        map_rows.append([cycle, "ONE LEADERSHIP", "C00000002", f"{cycle}|C00000002", "Leadership PAC",
                         "ONE, ADA", "H0AA00001", "Leadership PAC sponsor"])
    _with_header_row_5(wb, C.MAP_SHEET, map_header, map_rows)
    _with_header_row_5(wb, C.COMPANY_SHEET, ["Cycle", "Rank", "Company"],
                       [[cycle, 1, "AT&T"] for cycle in CYCLES] + [["2021-2022", 2, "Verizon"]])
    assoc_header = ["Cycle", "View", "Company", "Candidate / Officeholder", "Candidate ID", "Party",
                    "Office", "State", "District", "Associated Telecom Funding"]
    _with_header_row_5(
        wb, C.SUMMARY_SHEET, assoc_header,
        [["2021-2022", "Industry", "", "ONE, ADA", "H0AA00001", "D", "House", "AA", "1", 40.0]],
    )
    bridge_header = ["Cycle", "View", "Source Committee ID", "Candidate ID", "Directions",
                     "Activity Relationships", "Associated Telecom Funding"]
    _with_header_row_5(
        wb, C.BRIDGE_SHEET, bridge_header,
        [["2021-2022", "Industry", "C00000003", "H0AA00001", "Support", "Independent expenditure - support", 40.0]],
    )
    _with_header_row_5(
        wb, "Spending Overview", ["Cycle", "Scope", "Total Giving"],
        [[cycle, "Total Giving", totals[cycle]] for cycle in CYCLES],
    )
    wb.save(path)
    return totals


@pytest.fixture()
def workbook_path(tmp_path):
    path = str(tmp_path / "Telecom Industry Spending Overview.xlsx")
    _build_workbook(path)
    return path


def test_telecom_is_registered_as_an_opposition_workbook():
    assert C.WORKBOOKS["telecom"] == {
        "filename": "Telecom Industry Spending Overview.xlsx",
        "side": "opposition",
    }


def test_compact_schema_derives_cycle_and_committee_id(workbook_path):
    fec_to_bioguide = {"H0AA00001": "A000001"}
    manifest, members, totals, map_rows, assoc_rows, receipt = C._process_workbook(
        workbook_path, "telecom", "opposition", fec_to_bioguide
    )
    # Only the Recipient fact is consumed (Master and Employee restate its dollars).
    assert manifest["detail_row_counts"] == {"2021-2022": 3, "2023-2024": 1, "2025-2026": 1}
    assert manifest["view_rows_skipped"] == 2 * 5
    assert manifest["view_row_counts"] == {"Employee": 5, "Master": 5, "Recipient": 5}
    assert manifest["date_status"] == "cycle_only"
    assert manifest["detail_schema"]["kind"] == "compact_cycle_sheet"
    assert manifest["detail_schema"]["compact_sheets"] == sorted(CYCLES)
    # The overview's Total Giving reconciles to the Recipient-fact sum, per cycle.
    assert {c: r["ok"] for c, r in manifest["reconciliation"].items()} == {c: True for c in CYCLES}
    assert manifest["reconciliation"]["2021-2022"]["computed_total"] == 175.0
    # Cycle came from the sheet name and Committee ID from the Recipient Key, so
    # the (cycle, committee) map join attributes dollars to the bioguide.
    direct = members[("A000001", "telecom", "2021-2022", "")]
    assert direct["direct_dollars"] == Decimal("100.00")
    assert direct["sponsor_dollars"] == Decimal("50.00")
    assert direct["candidate_ids"] == {"H0AA00001"}
    assert members[("A000001", "telecom", "2025-2026", "")]["direct_dollars"] == Decimal("300.00")
    # The Super PAC recipient has no candidate, so it stays unattributed.
    assert receipt["unattributed_dollars"] == Decimal("25.00")
    assert receipt["direct_dollars"] == Decimal("600.00")
    assert receipt["committees"] == {"C00000001", "C00000002", "C00000003"}
    assert totals[("telecom", "2021-2022", "Candidate")]["total_dollars"] == Decimal("100.00")
    assert all(row["industry"] == "telecom" for row in map_rows)
    assert assoc_rows and assoc_rows[0]["side"] == "opposition"
    # All Exception rows are counted (a single unstacked view).
    assert receipt["excluded_exception_rows"] == 2
    assert manifest["excluded_exception_rows_all_views"] == 2


def test_compact_recipient_key_must_be_an_fec_key(tmp_path):
    bad = {cycle: list(facts) for cycle, facts in RECIPIENT_FACTS.items()}
    bad["2021-2022"][0] = ("Candidate", "ACTBLUE:C00000001", "100.00", "AT&T")
    path = str(tmp_path / "bad.xlsx")
    _build_workbook(path, recipient_facts=bad)
    with pytest.raises(C.IngestError, match="is not an FEC:<committee id> key"):
        C._process_workbook(path, "telecom", "opposition", {"H0AA00001": "A000001"})


def test_full_schema_sheets_are_read_exactly_as_before(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "2021-2022"
    for _ in range(3):
        ws.append(["title"])
    ws.append(FULL_HEADER)
    path = str(tmp_path / "full.xlsx")
    wb.save(path)
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True)["2021-2022"]
    positions, compact = C._detail_headers(ws, 4, "full.xlsx")
    assert compact is False
    assert positions == C._headers(ws, 4, C.REQUIRED_DETAIL_COLUMNS, "full.xlsx")
    assert positions["Cycle"] == FULL_HEADER.index("Cycle")
    assert positions["Committee ID"] == FULL_HEADER.index("Committee ID")


def test_a_sheet_missing_both_schemas_still_fails_closed(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "2021-2022"
    for _ in range(3):
        ws.append(["title"])
    ws.append(["Company", "Amount"])
    path = str(tmp_path / "neither.xlsx")
    wb.save(path)
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True)["2021-2022"]
    with pytest.raises(C.IngestError, match="required column"):
        C._detail_headers(ws, 4, "neither.xlsx")
