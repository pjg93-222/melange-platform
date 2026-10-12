# -*- coding: utf-8 -*-
"""COR-7 (determinism: rebuild == shipped, canonical row-hash) + COR-3 (parity-mode dominance
vs the live CU base) for the holistic v1 BASE. Subprocess builder runs ride the bill cache, so
each is assemble+write only. Writes cor3_ledger.json beside this script's eval record."""
import os, sys, json, hashlib, subprocess, tempfile
HERE = os.path.dirname(os.path.abspath(__file__))
WT = os.path.normpath(os.path.join(HERE, '..', '..', '..'))
BUILDER = os.path.join(WT, 'fec', 'network', 'build_holistic_base.py')
SHIPPED = os.path.join(WT, 'holistic_network_BASE_20260817.xlsx')
CU_BASE = os.path.join(WT, 'congress_cosponsor_network_20260625_205554.xlsx')
REC = os.path.join(WT, 'audits', 'holistic_v1_eval')
import openpyxl

def rowhash(path):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)['Edges']
    h = hashlib.sha256()
    for r in ws.iter_rows(values_only=True):
        h.update(repr(r).encode('utf-8'))
    return h.hexdigest()

env0 = {k: v for k, v in os.environ.items() if not k.startswith('HOLISTIC_') and k != 'NET_CONFIG'}
tmp = tempfile.mkdtemp(prefix='holistic_cor37_')

print('[COR-7] rebuilding with identical dials...')
o1 = os.path.join(tmp, 'holistic_network_BASE_rebuild.xlsx')
r = subprocess.run([sys.executable, BUILDER], env=dict(env0, HOLISTIC_BASE_OUT=o1),
                   capture_output=True, text=True, encoding='utf-8', errors='replace')
if r.returncode != 0: print(r.stdout[-500:], r.stderr[-500:]); sys.exit(1)
h_ship, h_re = rowhash(SHIPPED), rowhash(o1)
print('COR-7 determinism: %s (shipped %s / rebuild %s)' %
      ('PASS — canonical row-hash identical' if h_ship == h_re else 'FAIL', h_ship[:12], h_re[:12]))

print('[COR-3] building parity-mode BASE (HOLISTIC_REGULAR_W=0.5)...')
o2 = os.path.join(tmp, 'holistic_network_BASE_parity05.xlsx')
r = subprocess.run([sys.executable, BUILDER], env=dict(env0, HOLISTIC_BASE_OUT=o2, HOLISTIC_REGULAR_W='0.5'),
                   capture_output=True, text=True, encoding='utf-8', errors='replace')
if r.returncode != 0: print(r.stdout[-500:], r.stderr[-500:]); sys.exit(1)
par = {}
for row in openpyxl.load_workbook(o2, read_only=True, data_only=True)['Edges'].iter_rows(min_row=2, values_only=True):
    par[(row[0], row[1])] = row[2]
roster = {m['bioguide'] for m in json.load(open(os.path.join(WT, 'data', 'engine', 'congestion.json'), encoding='utf-8')) if m.get('bioguide')}
n = off = missing = dom_viol = 0
viol_samples = []
for row in openpyxl.load_workbook(CU_BASE, read_only=True, data_only=True)['Edges'].iter_rows(min_row=2, values_only=True):
    a, b, w = row[0], row[1], row[2]
    if not isinstance(w, (int, float)): continue
    if a not in roster or b not in roster: off += 1; continue
    n += 1
    pw = par.get((a, b))
    if pw is None:
        missing += 1
        if len(viol_samples) < 10: viol_samples.append({'edge': [a, b], 'cu_w': w, 'parity_w': None, 'class': 'missing-pair'})
    elif pw < w - 1e-9:
        dom_viol += 1
        if len(viol_samples) < 10: viol_samples.append({'edge': [a, b], 'cu_w': w, 'parity_w': pw, 'class': 'dominance'})
strict = n - missing - dom_viol
ledger = {'cu_edges_roster_internal': n, 'cu_edges_off_roster_skipped': off,
          'strict_dominant_or_equal': strict, 'strict_rate': round(strict / n, 6) if n else None,
          'missing_pairs': missing, 'dominance_violations': dom_viol, 'samples': viol_samples,
          'pin': '>=99.5% strict; 100% explained (freeze addendum 2026-08-17)',
          'note': 'CU base vintage 2026-06-25 scrape vs holistic BILLSTATUS 2026-08-01 — source-discrepancy classes for any residue'}
json.dump(ledger, open(os.path.join(REC, 'cor3_ledger.json'), 'w', encoding='utf-8'), indent=1)
print('COR-3 dominance: %d roster-internal CU edges | strict %.4f%% | missing %d | dominance-violations %d | off-roster skipped %d' %
      (n, 100.0 * strict / n if n else 0, missing, dom_viol, off))
print('COR-3 verdict vs pin: %s' % ('PASS' if n and strict / n >= 0.995 and (missing + dom_viol) == 0
      else ('PASS-with-ledger' if n and strict / n >= 0.995 else 'FAIL — below strict floor')))
