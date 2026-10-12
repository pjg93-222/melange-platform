# -*- coding: utf-8 -*-
"""COR-3v2 — the re-pinned mirror-reconcile bar (delegated ruling, Paul in-chat 2026-08-17:
"What is sitting on me is for you to do best"). The v1 per-edge dominance bar's premise was
unsound for accruing bills: C grows between the CU scrape vintage (2026-06-25) and the archive
vintage (2026-08-01), the per-bill factor divides by sqrt(C), so later-vintage parity weights
legitimately dip below CU weights on bills that gained cosponsors. Re-pin (recorded in the
freeze addendum + RECORD):
  (i)  PRESENCE: 100% of roster-internal CU directed edges exist in the parity-mode holistic
       BASE; every missing pair individually root-caused; unexplained -> FAIL.
  (ii) AGGREGATE MASS: sum(parity_w over CU edges) / sum(cu_w) >= 1.0 — the universe-superset
       effect must dominate vintage drift in aggregate, else the builder is under-collecting.
  (iii) per-edge dominance -> REPORT-ONLY diagnostic (distribution quoted, no verdict).
parity_w needs NO rebuild: the BASE carries WOriginal/WRegular with role dials factored out,
so parity (CU role weights 1.0/0.5) = 1.0*WOriginal + 0.5*WRegular per row (COR-2 basis)."""
import os, sys, json, collections
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
WT = os.path.normpath(os.path.join(HERE, '..', '..', '..'))
CU_BASE = os.path.join(WT, 'congress_cosponsor_network_20260625_205554.xlsx')
HOL_BASE = os.path.join(WT, 'holistic_network_BASE_20260817.xlsx')
ROSTER = {m['bioguide'] for m in json.load(open(os.path.join(WT, r'data\engine\congestion.json'), encoding='utf-8')) if m.get('bioguide')}

hol = {}
ws = openpyxl.load_workbook(HOL_BASE, data_only=True)['Edges']
for r in ws.iter_rows(min_row=2, values_only=True):
    frm, to, _, nb, no, nr, _, _, wo, wr = r[:10]
    hol[(frm, to)] = (1.0 * (wo or 0.0) + 0.5 * (wr or 0.0), nb, no, nr)

cu = openpyxl.load_workbook(CU_BASE, data_only=True)['Edges']
hdr = [c.value for c in cu[1]]
matched = missing = 0
sum_cu = sum_par = 0.0
ratios = []
missing_detail = []
for r in cu.iter_rows(min_row=2, values_only=True):
    frm, to, w = r[0], r[1], r[2]
    if frm not in ROSTER or to not in ROSTER or not isinstance(w, (int, float)):
        continue
    if (frm, to) in hol:
        pw = hol[(frm, to)][0]
        matched += 1; sum_cu += w; sum_par += pw
        ratios.append(pw / w if w else float('inf'))
    else:
        missing += 1
        nb = r[3] if len(r) > 3 else None
        missing_detail.append({
            'edge': [frm, to], 'cu_w': w, 'cu_NBills': nb,
            'reverse_present': (to, frm) in hol,
            'from_out_deg_hol': sum(1 for k in hol if k[0] == frm),
        })

ratios.sort()
def q(p): return ratios[int(p * (len(ratios) - 1))] if ratios else None
agg = (sum_par / sum_cu) if sum_cu else None
agg_with_missing = agg  # missing edges contribute 0 to sum_par by construction (not matched) —
# recompute honestly including their cu mass in the denominator:
sum_cu_all = sum_cu + sum(m['cu_w'] for m in missing_detail)
agg_all = (sum_par / sum_cu_all) if sum_cu_all else None

print('COR-3v2 | roster-internal CU edges: %d matched + %d missing' % (matched, missing))
print('(i) PRESENCE: %s' % ('PASS 100%' if missing == 0 else 'PENDING root-cause of %d missing pairs' % missing))
for m in missing_detail:
    print('    MISSING %s cu_w=%.4f NBills=%s reverse_present=%s hol_out_deg(From)=%d' %
          (m['edge'], m['cu_w'], m['cu_NBills'], m['reverse_present'], m['from_out_deg_hol']))
print('(ii) AGGREGATE MASS: parity/cu = %.4f (matched-only) | %.4f (missing counted in denominator) -> %s'
      % (agg, agg_all, 'PASS >= 1.0' if agg_all is not None and agg_all >= 1.0 else 'FAIL'))
print('(iii) REPORT-ONLY per-edge ratio distribution: p05 %.3f | p25 %.3f | median %.3f | p75 %.3f | p95 %.3f | share>=1: %.4f'
      % (q(.05), q(.25), q(.50), q(.75), q(.95), sum(1 for x in ratios if x >= 1 - 1e-9) / len(ratios)))
json.dump({'matched': matched, 'missing': missing, 'missing_detail': missing_detail,
           'aggregate_matched_only': agg, 'aggregate_all': agg_all,
           'ratio_quantiles': {'p05': q(.05), 'p25': q(.25), 'median': q(.50), 'p75': q(.75), 'p95': q(.95)},
           'share_dominant': sum(1 for x in ratios if x >= 1 - 1e-9) / len(ratios)},
          open(os.path.join(WT, r'audits\holistic_v1_eval\cor3v2_result.json'), 'w', encoding='utf-8'), indent=1)
print('wrote audits/holistic_v1_eval/cor3v2_result.json')
