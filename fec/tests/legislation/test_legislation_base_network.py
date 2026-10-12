# -*- coding: utf-8 -*-
"""Focused contract tests for the staged legislation base-network producer."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
import tempfile
from pathlib import Path

import openpyxl


FEC = Path(__file__).resolve().parents[2]
ROOT = FEC.parent
sys.path.insert(0, str(FEC))

from legislation import build_legislation_base_network as NET  # noqa: E402


class raises:
    """Tiny dependency-free equivalent for the focused exception assertions."""

    def __init__(self, exception: type[BaseException], match: str):
        self.exception = exception
        self.pattern = re.compile(match)

    def __enter__(self):
        return self

    def __exit__(self, kind, value, traceback):
        assert kind is not None, f"expected {self.exception.__name__}"
        assert issubclass(kind, self.exception), value
        assert self.pattern.search(str(value)), (
            f"{value!s} does not match {self.pattern.pattern!r}"
        )
        return True


def relationship(
    *,
    congress: int = 119,
    measure_type: str = "hr",
    measure_number: int = 1,
    category: str = "Payments",
    bioguide: str,
    name: str,
    role: str,
    relationship_status: str = "current",
    membership_status: str = "included",
) -> dict:
    return {
        "schema": NET.ROW_SCHEMA,
        "measure_key": f"{congress}:{measure_type}:{measure_number}",
        "congress": congress,
        "measure_type": measure_type,
        "measure_number": measure_number,
        "category": category,
        "member_bioguide_id": bioguide,
        "member_name": name,
        "role": role,
        "relationship_status": relationship_status,
        "membership_status": membership_status,
        "provenance": {
            "official_url": (
                f"https://www.congress.gov/bill/{congress}th-congress/"
                f"{measure_type}/{measure_number}"
            ),
            "retrieved_at": "2026-07-25T00:00:00Z",
        },
    }


def parsed(**kwargs) -> NET.Relationship:
    return NET.parse_relationship(relationship(**kwargs), 1)


def write_ledger(path: Path, rows: list[dict]) -> None:
    path.write_text(
        "".join(json.dumps(row, sort_keys=True) + "\n" for row in rows),
        encoding="utf-8",
    )


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_formula_dedup_withdrawn_excluded_and_roster_isolate():
    rows = [
        parsed(bioguide="A000001", name="Sponsor", role="sponsor"),
        parsed(
            bioguide="B000001",
            name="Original",
            role="original_cosponsor",
        ),
        parsed(bioguide="C000001", name="Later", role="cosponsor"),
        # A second category is the same official relationship, not a new bill.
        parsed(
            category="Data Privacy & Cyber",
            bioguide="B000001",
            name="Original",
            role="cosponsor",
        ),
        parsed(
            bioguide="D000001",
            name="Withdrawn",
            role="cosponsor",
            relationship_status="withdrawn",
        ),
        parsed(
            bioguide="E000001",
            name="Excluded",
            role="cosponsor",
            membership_status="excluded",
        ),
        parsed(
            congress=118,
            measure_number=2,
            bioguide="A000001",
            name="Sponsor",
            role="sponsor",
        ),
        parsed(
            congress=118,
            measure_number=2,
            bioguide="B000001",
            name="Original",
            role="original_cosponsor",
        ),
    ]
    roster = {
        "A000001": "Sponsor",
        "B000001": "Original",
        "C000001": "Later",
        "D000001": "Withdrawn",
        "E000001": "Excluded",
        "F000001": "Roster Isolate",
    }
    edges, contributions, exclusions, names = NET.construct_edges(rows, roster)
    edge_by_pair = {(row["From"], row["To"]): row for row in edges}
    assert len(contributions) == 3
    assert set(edge_by_pair) == {
        ("A000001", "B000001"),
        ("A000001", "C000001"),
    }
    assert math.isclose(
        edge_by_pair[("A000001", "B000001")]["Weight"],
        1.0 / math.sqrt(2) + 0.66,
        rel_tol=0,
        abs_tol=1.0e-15,
    )
    assert edge_by_pair[("A000001", "B000001")]["NBills"] == 2
    assert edge_by_pair[("A000001", "B000001")]["NOriginal"] == 2
    assert edge_by_pair[("A000001", "B000001")]["NRegular"] == 0
    assert math.isclose(
        edge_by_pair[("A000001", "C000001")]["Weight"],
        0.5 / math.sqrt(2),
        rel_tol=0,
        abs_tol=1.0e-15,
    )
    assert {row["reason"] for row in exclusions} == {
        "relationship_withdrawn",
        "membership_excluded",
    }
    centrality = NET.compute_centrality(edges, names)
    centrality_by_node = {row["node"]: row for row in centrality}
    assert set(centrality_by_node) == set(roster)
    assert centrality_by_node["F000001"]["WeightedOutDegree"] == 0
    assert centrality_by_node["F000001"]["InfluenceScore"] == 0
    assert centrality_by_node["A000001"]["InfluenceScore"] == 100.0


def test_type_aware_identity_keeps_same_number_as_two_bills():
    rows = [
        parsed(
            measure_type=measure_type,
            bioguide=bioguide,
            name=name,
            role=role,
        )
        for measure_type in ("hr", "hres")
        for bioguide, name, role in (
            ("A000001", "Sponsor", "sponsor"),
            ("B000001", "Cosponsor", "cosponsor"),
        )
    ]
    edges, contributions, exclusions, _ = NET.construct_edges(rows)
    assert not exclusions
    assert len(contributions) == 2
    assert {row["measure_key"] for row in contributions} == {
        "119:hr:1",
        "119:hres:1",
    }
    assert edges[0]["NBills"] == 2
    assert edges[0]["NRegular"] == 2


def test_amendments_are_half_weighted_and_forged_weights_fail_closed():
    bill_rows = [
        parsed(
            measure_type="hr",
            bioguide="A000001",
            name="Sponsor",
            role="sponsor",
        ),
        parsed(
            measure_type="hr",
            bioguide="B000001",
            name="Cosponsor",
            role="original_cosponsor",
        ),
    ]
    amendment_rows = [
        parsed(
            measure_type="hamdt",
            measure_number=2,
            bioguide="A000001",
            name="Sponsor",
            role="sponsor",
        ),
        parsed(
            measure_type="hamdt",
            measure_number=2,
            bioguide="B000001",
            name="Cosponsor",
            role="original_cosponsor",
        ),
    ]
    _, contributions, _, _ = NET.construct_edges(bill_rows + amendment_rows)
    by_measure = {
        row["measure_key"]: row for row in contributions
    }
    assert by_measure["119:hr:1"]["contribution"] == 1.0
    assert by_measure["119:hamdt:2"]["contribution"] == 0.5
    assert by_measure["119:hamdt:2"]["measure_weight"] == 0.5
    assert by_measure["119:hamdt:2"]["effective_weight"] == 0.5

    forged = relationship(
        measure_type="hamdt",
        bioguide="A000001",
        name="Sponsor",
        role="sponsor",
    )
    forged["measure_weight"] = 1.0
    forged["effective_weight"] = 1.0
    with raises(NET.ContractError, match="measure_weight"):
        NET.parse_relationship(forged, 10)


def test_strict_schema_and_identity_fail_loudly():
    bad_schema = relationship(
        bioguide="A000001", name="Sponsor", role="sponsor"
    )
    bad_schema["schema"] = "something-else"
    with raises(NET.ContractError, match="schema must be"):
        NET.parse_relationship(bad_schema, 7)

    bad_key = relationship(
        bioguide="A000001", name="Sponsor", role="sponsor"
    )
    bad_key["measure_key"] = "119:hres:1"
    with raises(NET.ContractError, match="type-aware identity"):
        NET.parse_relationship(bad_key, 8)

    bad_type = relationship(
        bioguide="A000001", name="Sponsor", role="sponsor"
    )
    bad_type["measure_type"] = "H.R."
    bad_type["measure_key"] = "119:H.R.:1"
    with raises(NET.ContractError, match="noncanonical measure_type"):
        NET.parse_relationship(bad_type, 9)

    current = parsed(
        bioguide="A000001", name="Sponsor", role="sponsor"
    )
    withdrawn_copy = NET.Relationship(
        **{
            **current.__dict__,
            "row_number": 2,
            "category": "Data Privacy & Cyber",
            "relationship_status": "withdrawn",
        }
    )
    with raises(NET.ContractError, match="disagree on relationship status"):
        NET.construct_edges([current, withdrawn_copy])

    with raises(NET.ContractError, match="expected one current sponsor"):
        NET.construct_edges(
            [parsed(bioguide="B000001", name="Later", role="cosponsor")]
        )


def test_output_guard_rejects_workspace_and_production_style_names():
    with raises(NET.ContractError, match="must resolve under"):
        NET.output_paths(ROOT / "new_network.xlsx")
    with tempfile.TemporaryDirectory() as directory:
        tmp_path = Path(directory)
        with raises(NET.ContractError, match="production-style"):
            NET.output_paths(
                tmp_path / "congress_cosponsor_network_20260725.xlsx"
            )
        accepted = NET.output_paths(tmp_path / "shadow_network.xlsx")
        assert accepted.workbook == (tmp_path / "shadow_network.xlsx").resolve()


def test_end_to_end_outputs_hashes_and_byte_idempotence():
    with tempfile.TemporaryDirectory() as directory:
        tmp_path = Path(directory)
        ledger = tmp_path / "accepted_relationships.ndjson"
        roster = tmp_path / "roster.csv"
        out = tmp_path / "stage"
        write_ledger(
            ledger,
            [
                relationship(
                    bioguide="A000001", name="Sponsor", role="sponsor"
                ),
                relationship(
                    bioguide="B000001",
                    name="Original",
                    role="original_cosponsor",
                ),
                relationship(
                    bioguide="C000001", name="Later", role="cosponsor"
                ),
                relationship(
                    bioguide="D000001",
                    name="Withdrawn",
                    role="cosponsor",
                    relationship_status="withdrawn",
                ),
            ],
        )
        roster.write_text(
            "bioguide_id,member_name\n"
            "A000001,Sponsor\n"
            "B000001,Original\n"
            "C000001,Later\n"
            "D000001,Withdrawn\n"
            "E000001,Isolate\n",
            encoding="utf-8",
        )
        args = argparse.Namespace(
            accepted_ledger=ledger,
            out=out,
            roster=roster,
            compare_old_workbook=None,
        )
        first = NET.run(args)
        paths = NET.output_paths(out)
        first_hashes = {
            name: digest(path)
            for name, path in {
                "workbook": paths.workbook,
                "contributions": paths.contributions,
                "exclusions": paths.exclusions,
            }.items()
        }
        second = NET.run(args)
        second_hashes = {
            name: digest(path)
            for name, path in {
                "workbook": paths.workbook,
                "contributions": paths.contributions,
                "exclusions": paths.exclusions,
            }.items()
        }
        assert first_hashes == second_hashes
        assert first["counts"] == second["counts"]
        assert first["outputs"]["workbook"]["sha256"] == first_hashes["workbook"]
        assert first["counts"]["edge_contribution_rows"] == 2
        assert first["counts"]["zero_contribution_exclusion_rows"] == 1

        workbook = openpyxl.load_workbook(
            paths.workbook, read_only=True, data_only=True
        )
        assert workbook.sheetnames == ["Edges", "Centrality"]
        assert tuple(next(workbook["Edges"].values)) == NET.EDGE_HEADERS
        assert (
            tuple(next(workbook["Centrality"].values))
            == NET.CENTRALITY_HEADERS
        )
        assert workbook["Centrality"].max_row == 6
        workbook_edges = [
            dict(zip(NET.EDGE_HEADERS, row))
            for row in list(workbook["Edges"].values)[1:]
        ]
        workbook_centrality = [
            dict(zip(NET.CENTRALITY_HEADERS, row))
            for row in list(workbook["Centrality"].values)[1:]
        ]
        workbook.close()

        comparison = NET.compare_legacy(
            workbook_edges,
            workbook_centrality,
            paths.workbook,
        )
        assert comparison["exact_value_reproduction"] is True


if __name__ == "__main__":
    for name, value in sorted(globals().copy().items()):
        if name.startswith("test_") and callable(value):
            value()
            print(f"PASS {name}")
