# -*- coding: utf-8 -*-
"""Focused tests for the fail-closed staged pillar/rank wrapper."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile

from openpyxl import Workbook


FEC = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(FEC))

from legislation import run_legislation_pillars_and_ranks as PR  # noqa: E402


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def artifact(path: Path, root: Path) -> dict[str, object]:
    return {
        "path": str(path.relative_to(root)).replace("\\", "/"),
        "size": path.stat().st_size,
        "sha256": sha(path),
    }


def fixture(root: Path, *, stale_snapshot: bool = False) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    rank_sheet = workbook.create_sheet("Rank & Allocation")
    blended = workbook.create_sheet("Blended Model")
    workbook.create_sheet("Press Release - Raw")
    workbook.create_sheet("Press Release - Normalized")
    workbook.create_sheet("Committee Data")
    members = ["Member One", "Member Two"]
    cached = {
        "O": {"Member One": 2.0, "Member Two": 4.0},
        "E": {"Member One": 3.0, "Member Two": 6.0},
        "F": {"Member One": 5.0, "Member Two": 10.0},
    }
    for offset, member in enumerate(members, 3):
        rank_sheet.cell(offset, 1).value = member
        blended.cell(offset, 1).value = member
        blended.cell(offset, 15).value = cached["O"][member]
        blended.cell(offset, 5).value = cached["E"][member]
        blended.cell(offset, 6).value = cached["F"][member]
    workbook_path = root / PR.WORKBOOK_NAME
    workbook.save(workbook_path)
    snapshot = {
        **cached,
        "stamped": "2026-07-25T00:00:00-05:00",
    }
    if stale_snapshot:
        snapshot["O"]["Member One"] = 999.0
    snapshot_path = root / PR.SNAPSHOT_NAME
    snapshot_path.parent.mkdir(parents=True, exist_ok=True)
    write_json(snapshot_path, snapshot)
    roles_path = root / PR.ROLES_NAME
    write_json(
        roles_path,
        {
            "schema": "party-elite-influence-roles-v3",
            "verified_on": "2026-07-25",
            "valid_until": "2027-01-03",
            "members": {},
        },
    )
    generation_id = "a" * 64
    workbook_receipt_path = root / "legislation_workbook.receipt.json"
    workbook_receipt = {
        "schema": PR.WORKBOOK_RECEIPT_SCHEMA,
        "status": "PASS",
        "staged_only": True,
        "production_written": False,
        "generation_id": generation_id,
        "outputs": {
            PR.WORKBOOK_NAME: artifact(workbook_path, root),
            PR.SNAPSHOT_NAME: artifact(snapshot_path, root),
        },
    }
    write_json(workbook_receipt_path, workbook_receipt)
    prerequisite = {
        "schema": PR.PREREQUISITE_SCHEMA,
        "status": "PASS",
        "staged_only": True,
        "production_written": False,
        "generation_id": generation_id,
        "artifacts": {
            name: artifact(root / name, root)
            for name in PR.REQUIRED_ARTIFACTS
        },
        "legislative_workbook_receipt": artifact(
            workbook_receipt_path, root
        ),
    }
    prerequisite_path = root / PR.PREREQUISITE_NAME
    write_json(prerequisite_path, prerequisite)
    return prerequisite_path


def fake_compute(
    args: list[str],
    *,
    cwd: str,
    env: dict[str, str],
    text: bool,
    capture_output: bool,
    timeout: int,
    check: bool,
) -> subprocess.CompletedProcess[str]:
    del args, cwd, text, capture_output, timeout, check
    root = Path(env["MONEYBALL_ROOT"])
    snapshot = json.loads(
        (root / PR.SNAPSHOT_NAME).read_text(encoding="utf-8")
    )
    members = list(snapshot["O"])
    write_json(
        root / PR.RANKS_NAME,
        {
            "score": {member: float(index) for index, member in enumerate(members)},
            "rank": {member: index + 1 for index, member in enumerate(members)},
            "pcts": {member: {} for member in members},
            "ab_weights": {"2": 0.45, "3": 0.1, "4": 0.4, "5": 0.05},
        },
    )
    return subprocess.CompletedProcess([], 0, "ok\n", "")


def test_pass_is_hash_bound_and_idempotent(monkeypatch):
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory) / "shadow"
        root.mkdir()
        prerequisite = fixture(root)
        monkeypatch.setattr(PR.subprocess, "run", fake_compute)
        first = PR.run(run_dir=root, prerequisites_path=prerequisite)
        assert first["status"] == "PASS"
        assert first["stale_snapshot_fallback_used"] is False
        assert first["snapshot_matches_workbook_cache"] is True
        assert first["first_pass_copies_byte_identical"] is True
        assert (
            sha(root / PR.SNAPSHOT_NAME)
            == sha(root / PR.SNAPSHOT_FIRST)
        )
        assert sha(root / PR.RANKS_NAME) == sha(root / PR.RANKS_FIRST)
        receipt_bytes = (root / PR.RECEIPT_NAME).read_bytes()
        second = PR.run(run_dir=root, prerequisites_path=prerequisite)
        assert second["_idempotent_reuse"] is True
        assert (root / PR.RECEIPT_NAME).read_bytes() == receipt_bytes


def test_stale_snapshot_blocks_before_compute(monkeypatch):
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory) / "shadow"
        root.mkdir()
        prerequisite = fixture(root, stale_snapshot=True)
        called = []
        monkeypatch.setattr(
            PR.subprocess,
            "run",
            lambda *args, **kwargs: called.append((args, kwargs)),
        )
        try:
            PR.run(run_dir=root, prerequisites_path=prerequisite)
        except PR.PillarError as exc:
            assert "snapshot O is stale" in str(exc)
        else:
            raise AssertionError("stale snapshot unexpectedly passed")
        assert not called
        assert not (root / PR.RANKS_NAME).exists()


def test_receipt_generation_mismatch_blocks(monkeypatch):
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory) / "shadow"
        root.mkdir()
        prerequisite = fixture(root)
        value = json.loads(prerequisite.read_text(encoding="utf-8"))
        receipt_path = root / value["legislative_workbook_receipt"]["path"]
        receipt_value = json.loads(receipt_path.read_text(encoding="utf-8"))
        receipt_value["generation_id"] = "b" * 64
        write_json(receipt_path, receipt_value)
        value["legislative_workbook_receipt"] = artifact(receipt_path, root)
        write_json(prerequisite, value)
        called = []
        monkeypatch.setattr(
            PR.subprocess,
            "run",
            lambda *args, **kwargs: called.append((args, kwargs)),
        )
        try:
            PR.run(run_dir=root, prerequisites_path=prerequisite)
        except PR.PillarError as exc:
            assert "not PASS/hash-bound" in str(exc)
        else:
            raise AssertionError("generation mismatch unexpectedly passed")
        assert not called


if __name__ == "__main__":
    # Lightweight fallback when pytest is unavailable.
    class Monkey:
        def setattr(self, obj, name, value):
            setattr(obj, name, value)

    monkey = Monkey()
    for name, value in sorted(globals().copy().items()):
        if name.startswith("test_") and callable(value):
            value(monkey)
            print(f"PASS {name}")
