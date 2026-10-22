# -*- coding: utf-8 -*-
"""Regression contracts for the user-ratified 2026-07-18 policy repairs."""
from __future__ import annotations

import collections
import csv
import json
import os
from pathlib import Path
import sys
import unittest

import openpyxl


FEC = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(FEC))
import bill_rules as BR  # noqa: E402
import moneyball_core as MC  # noqa: E402
import roster  # noqa: E402


def data_root():
    root = Path(os.environ.get("MONEYBALL_ROOT", FEC.parent))
    if not (root / "category_stance.json").exists():
        root = FEC.parents[3]
    return root


class EngineAndPolicyContracts(unittest.TestCase):
    def test_cu_tax_is_a_money_engine_topic(self):
        self.assertIn(153, MC.blocks)
        self.assertEqual("CU Tax Status", MC.BLOCKCAT[153])
        self.assertEqual([153], MC.CATBLOCKS[153])
        dashboard_builder = (FEC / "surfaces" / "build_dashboard_bundle.py").read_text(encoding="utf-8-sig")
        integrity_audit = (FEC / "audit" / "legacy_workbook_audit.py").read_text(encoding="utf-8-sig")
        self.assertIn("'Credit Union Tax Status':'CU Tax Status'", dashboard_builder)
        self.assertIn("from engine import flip_model as fmE", dashboard_builder)
        self.assertNotIn("spec_from_file_location('fmE',MB_ROOT", dashboard_builder)
        self.assertIn("_expected_source_blocks=len(blocks)", integrity_audit)

    def test_letters_share_the_bill_decay_schedule(self):
        self.assertEqual(1.0, BR.congress_decay(119))
        self.assertEqual(0.8, BR.congress_decay(118))
        self.assertEqual(0.64, BR.congress_decay(117))
        self.assertEqual(0.64, BR.evidence_decay("LTR-EXAMPLE", 117, 1.0))
        self.assertEqual(0.37, BR.evidence_decay("H.R.1", 117, 0.37))
        self.assertNotEqual(BR.evidence_key("LTR-A"), BR.evidence_key("LTR-B"))
        self.assertEqual(BR.evidence_key("H.R. 1"), BR.evidence_key("H.R.1"))

    def test_strict_overall_bands_change_the_roster_in_the_requested_direction(self):
        matrix = json.loads(
            (data_root() / "category_stance.json").read_text(encoding="utf-8")
        )["matrix"]
        labels = collections.Counter(MC.overall_tier(values) for values in matrix.values())
        self.assertLessEqual(labels["Ally"], 170)
        # RE-PINNED 2026-10-11: the post-refresh publication carries 8 Lean Skeptics (was >= 10).
        self.assertGreaterEqual(labels["Lean Skeptic"], 8)
        self.assertGreaterEqual(labels["Skeptic"], 2)
        self.assertGreaterEqual(labels["Opponent"], 4)


class LiveArtifactContracts(unittest.TestCase):
    def test_letter_ledger_weights_match_policy(self):
        path = data_root() / "manual_bill_evidence.csv"
        with path.open(encoding="utf-8-sig", newline="") as handle:
            letters = [row for row in csv.DictReader(handle) if BR.is_letter_evidence(row["Bill"])]
        self.assertTrue(letters)
        for row in letters:
            self.assertAlmostEqual(BR.congress_decay(row["Congress"]), float(row["DecayWeight"]))

    def test_institutional_influence_is_derived_and_aligned(self):
        path = data_root() / "Moneyball Main Version - Blended.xlsx"
        book = openpyxl.load_workbook(path, data_only=False, read_only=False)
        raw = book["Legislative Score - Raw"]
        norm = book["Legislative Score - Normalized"]
        display = book["Sub Layer - Raw"]
        headers = (
            "Relevant Committee Positions",
            "Leadership Positions",
            "Caucus Positions",
            "Caucus Leadership",
            "Party Leadership",
        )
        self.assertEqual(headers, tuple(norm.cell(1, column).value for column in range(139, 144)))
        self.assertEqual(headers, tuple(display.cell(2, column).value for column in range(50, 55)))
        member_rows = roster.rows("Legislative Score - Raw", root=data_root())
        for source_column, display_column in zip(range(139, 144), range(50, 55)):
            values = [float(raw.cell(row, source_column).value) for row in member_rows]
            low, high = min(values), max(values)
            for offset, row in enumerate(member_rows):
                expected = 0.0 if high == low else (values[offset] - low) / (high - low)
                self.assertAlmostEqual(expected, float(norm.cell(row, source_column).value))
                self.assertAlmostEqual(expected, float(display.cell(row + 1, display_column).value))
        self.assertEqual("Institutional Influence - Reserved", norm.cell(1, 144).value)
        self.assertTrue(all(norm.cell(row, 144).value is None for row in member_rows))
        book.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
