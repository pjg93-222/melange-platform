"""fec/workbook_rows.py: row insert/delete with Excel's reference semantics, on a synthetic two-sheet workbook."""
import os
import sys

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

FEC = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, FEC)
import workbook_rows as W  # noqa: E402


def _book(tmp_path):
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = 'Rank & Allocation'
    a['A1'] = 'Names'
    for i, n in enumerate(['Ann', 'Bob', 'Cid', 'Dee'], start=2):
        a.cell(i, 1, n)
        a.cell(i, 2, i * 10)
        a.cell(i, 3, '=B%d*2' % i)
    a['E1'] = '=SUM(B2:B5)'
    a['E2'] = '=B5'
    a.add_table(Table(displayName='T1', ref='A1:C5'))
    b = wb.create_sheet('Blended Model')
    for i in range(2, 6):
        b.cell(i, 1, "='Rank & Allocation'!A%d" % i)
    b['C1'] = "=MAX('Rank & Allocation'!$B$2:$B$5)"
    b['C2'] = "=COUNTA('Rank & Allocation'!A:A)"
    b['C3'] = "='Rank & Allocation'!$B$6"            # the footer row under the block
    a['B6'] = 99
    p = tmp_path / 'b.xlsx'
    wb.save(p)
    return p


def _load(p):
    return openpyxl.load_workbook(p)


def test_insert_mid_block_grows_ranges_and_moves_refs(tmp_path):
    p = _book(tmp_path)
    bk = W.Book(str(p))
    bk.apply_group([W.Op('Rank & Allocation', 'insert', 4), W.Op('Blended Model', 'insert', 4)])
    bk.clone_row('Rank & Allocation', 3, 4)
    bk.clone_row('Blended Model', 3, 4)
    bk.set_value('Rank & Allocation', 'A4', 'Bea')
    bk.set_value('Rank & Allocation', 'B4', 7)
    out = tmp_path / 'o.xlsx'
    bk.save(str(out))
    wb = _load(out)
    a, b = wb['Rank & Allocation'], wb['Blended Model']
    assert [a.cell(r, 1).value for r in range(2, 7)] == ['Ann', 'Bob', 'Bea', 'Cid', 'Dee']
    assert a['C4'].value == '=B4*2' and a['C5'].value == '=B5*2'           # cloned + moved formulas
    assert a['E1'].value == '=SUM(B2:B6)' and a['E2'].value == '=B6'
    assert b['C1'].value == "=MAX('Rank & Allocation'!$B$2:$B$6)"
    assert b['C2'].value == "=COUNTA('Rank & Allocation'!A:A)"
    assert b['C3'].value == "='Rank & Allocation'!$B$7"
    assert [b.cell(r, 1).value for r in range(2, 7)] == ["='Rank & Allocation'!A%d" % r for r in range(2, 7)]
    assert dict(a.tables.items())['T1'] == 'A1:C6'
    assert a['B7'].value == 99


def test_append_at_block_end_needs_grow_end(tmp_path):
    p = _book(tmp_path)
    for grow, want in ((False, '=SUM(B2:B5)'), (True, '=SUM(B2:B6)')):
        bk = W.Book(str(p))
        bk.apply_group([W.Op('Rank & Allocation', 'insert', 6, 1, grow_end=grow)])
        out = tmp_path / ('g%d.xlsx' % grow)
        bk.save(str(out))
        a = _load(out)['Rank & Allocation']
        assert a['E1'].value == want
        assert a['E2'].value == '=B5'             # a single cell on the last member row never moves
        assert a['B7'].value == 99


def test_group_delete_across_mutually_referencing_sheets(tmp_path):
    p = _book(tmp_path)
    bk = W.Book(str(p))
    bk.apply_group([W.Op('Rank & Allocation', 'delete', 3), W.Op('Blended Model', 'delete', 3)])
    out = tmp_path / 'd.xlsx'
    bk.save(str(out))
    wb = _load(out)
    a, b = wb['Rank & Allocation'], wb['Blended Model']
    assert [a.cell(r, 1).value for r in range(2, 5)] == ['Ann', 'Cid', 'Dee']
    assert a['E1'].value == '=SUM(B2:B4)' and a['C3'].value == '=B3*2'
    assert [b.cell(r, 1).value for r in range(2, 5)] == ["='Rank & Allocation'!A%d" % r for r in range(2, 5)]
    assert b['C3'] is not None and b['C1'].value == "=MAX('Rank & Allocation'!$B$2:$B$4)"
    assert dict(a.tables.items())['T1'] == 'A1:C4'


def test_delete_refuses_a_surviving_reference_into_the_deleted_row(tmp_path):
    p = _book(tmp_path)
    bk = W.Book(str(p))
    with pytest.raises(W.RowOpError):
        bk.apply_group([W.Op('Rank & Allocation', 'delete', 5)])     # E2 = B5 survives and points at it


def test_untouched_parts_are_byte_identical(tmp_path):
    import zipfile
    p = _book(tmp_path)
    bk = W.Book(str(p))
    bk.apply_group([W.Op('Blended Model', 'insert', 4)])
    out = tmp_path / 'u.xlsx'
    bk.save(str(out))
    with zipfile.ZipFile(p) as x, zipfile.ZipFile(out) as y:
        assert [i.filename for i in x.infolist()] == [i.filename for i in y.infolist()]
        changed = [n for n in x.namelist() if x.read(n) != y.read(n)]
    assert set(changed) <= {'xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml'}
