"""Recognition packaging must retain existing analytical model cells and caches."""
import sys
import ast
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from stance.honors_workbook import merge_views, workbook_views, publish, prepare_model_update, sha, SHEETS
import openpyxl


class HonorsWorkbookTests(unittest.TestCase):
    def bound_workbook(self, root):
        workbook = root / 'model.xlsx'
        workbook.write_bytes(b'last published model package')
        backup = root / 'data/stance/rank_honors/model.before.xlsx'
        backup.parent.mkdir(parents=True)
        backup.write_bytes(b'original analytical model package')
        binding_path = backup.parent / 'workbook_publications.json'
        binding = {'generation_id': 'previous', 'before_sha256': sha(backup),
                   'after_sha256': sha(workbook), 'original_path': backup.relative_to(root).as_posix()}
        binding_path.write_text(json.dumps({str(workbook.resolve()): binding}), encoding='utf-8')
        return workbook, backup, binding_path

    def test_intervening_model_edits_block_republication_and_rollback(self):
        for honors in ({'generation_id': 'previous'}, {'generation_id': 'next'}, None):
            with self.subTest(honors=honors), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                workbook, backup, binding_path = self.bound_workbook(root)
                workbook.write_bytes(b'legitimate intervening model edits')
                before = {path: path.read_bytes() for path in (workbook, backup, binding_path)}
                with patch('stance.honors_views.load_honors', return_value=honors), \
                     patch('stance.honors_workbook.subprocess.run') as author:
                    with self.assertRaisesRegex(RuntimeError, 'intervening edits'):
                        publish(workbook, root)
                author.assert_not_called()
                self.assertEqual(before, {path: path.read_bytes() for path in before})

    def test_unchanged_binding_skips_export_and_allows_exact_rollback(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook, backup, binding_path = self.bound_workbook(root)
            with patch('stance.honors_views.load_honors', return_value={'generation_id': 'previous'}), \
                 patch('stance.honors_workbook.subprocess.run') as author:
                self.assertEqual(publish(workbook, root)['status'], 'UNCHANGED')
            author.assert_not_called()
            with patch('stance.honors_views.load_honors', return_value=None):
                self.assertEqual(publish(workbook, root)['status'], 'ROLLED_BACK')
            self.assertEqual(workbook.read_bytes(), backup.read_bytes())
            self.assertEqual(json.loads(binding_path.read_text(encoding='utf-8')), {})

    def test_prepare_model_update_is_independent_of_the_active_honors_pointer(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook, backup, binding_path = self.bound_workbook(root)
            with patch('stance.honors_views.load_honors', side_effect=AssertionError('must not load active pointer')):
                self.assertEqual(prepare_model_update(workbook, root)['status'], 'ROLLED_BACK')
                self.assertEqual(prepare_model_update(workbook, root)['status'], 'NO_PRIOR_PUBLICATION')
            self.assertEqual(workbook.read_bytes(), backup.read_bytes())
            self.assertEqual(json.loads(binding_path.read_text(encoding='utf-8')), {})

    def test_prepare_model_update_rejects_intervening_edits(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook, backup, binding_path = self.bound_workbook(root)
            workbook.write_bytes(b'intervening model update')
            before = {path: path.read_bytes() for path in (workbook, backup, binding_path)}
            with self.assertRaisesRegex(RuntimeError, 'intervening edits'):
                prepare_model_update(workbook, root)
            self.assertEqual(before, {path: path.read_bytes() for path in before})

    def test_recurring_chain_prepares_before_media_and_retains_both_model_updates(self):
        """Exercise the actual chain order against a bound publication fixture."""
        source = Path(__file__).resolve().parents[2] / 'run_pipeline.py'
        tree = ast.parse(source.read_text(encoding='utf-8'))
        chain = next(ast.literal_eval(node.value) for node in tree.body
                     if isinstance(node, ast.Assign)
                     and any(isinstance(target, ast.Name) and target.id == 'CHAIN'
                             for target in node.targets))
        first = next(i for i, (name, _) in enumerate(chain)
                     if name == 'verify_collector_roster_updates.py')
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook, backup, binding_path = self.bound_workbook(root)
            original = backup.read_bytes()
            events = []
            for name, arguments in chain[first:]:
                if name == 'honors_workbook.py' and arguments == ['--prepare-model-update']:
                    self.assertEqual(prepare_model_update(workbook, root)['status'], 'ROLLED_BACK')
                    events.append('prepare')
                elif name == 'build_media_blend.py':
                    workbook.write_bytes(workbook.read_bytes() + b' + refreshed media inputs')
                    events.append('media')
                elif name == 'sync_stance_g.py':
                    workbook.write_bytes(workbook.read_bytes() + b' + refreshed stance inputs')
                    events.append('stance')
                elif name == 'honors_workbook.py' and arguments == []:
                    self.assertEqual(events, ['prepare', 'media', 'stance'])
                    self.assertEqual(workbook.read_bytes(), original + b' + refreshed media inputs + refreshed stance inputs')
                    self.assertEqual(json.loads(binding_path.read_text(encoding='utf-8')), {})
                    events.append('publish')
            self.assertEqual(events, ['prepare', 'media', 'stance', 'publish'])

    def test_fresh_generated_package_uses_a_new_rollback_baseline(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook, old_backup, binding_path = self.bound_workbook(root)
            generated = openpyxl.Workbook()
            generated.active.title = 'Category Stance'
            generated.active['A1'] = 'legitimate refreshed model'
            generated.save(workbook)
            fresh_hash = sha(workbook)
            old_backup_bytes = old_backup.read_bytes()
            runtime = root / 'data/engine/artifact_tool_runtime.json'
            runtime.parent.mkdir(parents=True)
            runtime.write_text(json.dumps({'node': 'mock-author', 'node_modules': 'mock-libraries'}))
            views = [{'name': name, 'values': [['fixture recognition']]} for name in SHEETS]

            def fixture_author(arguments, **kwargs):
                helper, plan_path = Path(arguments[1]), Path(arguments[2])
                plan = json.loads(plan_path.read_text(encoding='utf-8'))
                authored = Path(plan['output'])
                book = openpyxl.Workbook()
                book.remove(book.active)
                for view in views:
                    sheet = book.create_sheet(view['name'])
                    sheet.append(view['values'][0])
                book.save(authored)
                Path(str(authored)+'.receipt.json').write_text(json.dumps({
                    'status': 'PASS', 'plan_sha256': sha(plan_path),
                    'helper_sha256': sha(helper), 'output_sha256': sha(authored)}))
                return SimpleNamespace(returncode=0)

            with patch('stance.honors_views.load_honors', return_value={'generation_id': 'next', 'members': {}}), \
                 patch('stance.honors_workbook.workbook_views', return_value=views), \
                 patch('stance.honors_workbook.subprocess.run', side_effect=fixture_author):
                self.assertEqual(publish(workbook, root, replace_generated=True)['status'], 'PASS')
            binding = json.loads(binding_path.read_text(encoding='utf-8'))[str(workbook.resolve())]
            self.assertEqual(binding['before_sha256'], fresh_hash)
            self.assertEqual(sha(root / binding['original_path']), fresh_hash)
            self.assertNotEqual(root / binding['original_path'], old_backup)
            self.assertEqual(old_backup.read_bytes(), old_backup_bytes)
            prepare_model_update(workbook, root)
            self.assertEqual(sha(workbook), fresh_hash)

    def test_generated_override_rejects_either_managed_honors_tab(self):
        for managed in SHEETS:
            with self.subTest(managed=managed), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                workbook, backup, binding_path = self.bound_workbook(root)
                book = openpyxl.Workbook()
                book.active.title = managed
                book.save(workbook)
                before = {path: path.read_bytes() for path in (workbook, backup, binding_path)}
                with self.assertRaisesRegex(RuntimeError, 'managed honors views'):
                    publish(workbook, root, replace_generated=True)
                self.assertEqual(before, {path: path.read_bytes() for path in before})

    def test_append_preserves_existing_worksheet_bytes_and_replaces_views(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source, authored, output, repeated = [root / n for n in ('source.xlsx', 'authored.xlsx', 'output.xlsx', 'repeated.xlsx')]
            original = openpyxl.Workbook()
            original.active.title = 'Category Stance'
            original.active['A1'] = 'Legacy model input'
            original.active['B1'] = '=2*3'
            model = original.create_sheet('Rank & Allocation')
            model['G1'] = 'General Stance'
            model['G2'] = 'Ally'
            model['H2'] = '=2*3'
            model.column_dimensions.group('F', 'H', hidden=False)
            original.create_sheet('Coalition Profiles')['A1'] = 'Legacy analytical ladder'
            agenda = original.create_sheet('Legislative Agenda')
            agenda['C1'] = 'General Stance'
            agenda['C2'] = 'Ally'
            original.save(source)
            views = openpyxl.Workbook()
            views.active.title = SHEETS[0]
            views.active['A1'] = 'Champion'
            views.create_sheet(SHEETS[1])['A1'] = 'Tied 5–7'
            views[SHEETS[1]]['B2'] = 0
            views[SHEETS[1]]['B2'].number_format = '0.000'
            views.create_sheet(SHEETS[2])['A1'] = 'Overall standing'
            views.save(authored)
            merge_views(source, authored, output)
            merge_views(output, authored, repeated)
            with ZipFile(output) as first, ZipFile(repeated) as second:
                for tag in ('fonts', 'fills', 'borders', 'cellXfs'):
                    q = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                    self.assertEqual(len(ET.fromstring(first.read('xl/styles.xml')).find(q+tag)),
                                     len(ET.fromstring(second.read('xl/styles.xml')).find(q+tag)))
            for result in (output, repeated):
                with ZipFile(source) as before, ZipFile(result) as after:
                    self.assertEqual(before.read('xl/worksheets/sheet1.xml'), after.read('xl/worksheets/sheet1.xml'))
                    q = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                    self.assertEqual(ET.tostring(ET.fromstring(before.read('xl/worksheets/sheet2.xml')).find(q+'sheetData')),
                                     ET.tostring(ET.fromstring(after.read('xl/worksheets/sheet2.xml')).find(q+'sheetData')))
                book = openpyxl.load_workbook(result, data_only=False)
                self.assertEqual(book.sheetnames[:len(SHEETS)], list(SHEETS))
                self.assertEqual(book[SHEETS[0]]['A1'].value, 'Champion')
                self.assertEqual(book[SHEETS[1]]['A1'].value, 'Tied 5–7')
                self.assertEqual(book[SHEETS[1]]['B2'].value, 0)
                self.assertEqual(book[SHEETS[1]]['B2'].number_format, '0.000')
                self.assertEqual(book['Category Stance']['B1'].value, '=2*3')
                self.assertEqual(book['Category Stance'].sheet_state, 'hidden')
                self.assertEqual(book['Coalition Profiles'].sheet_state, 'hidden')
                self.assertTrue(book['Rank & Allocation'].column_dimensions['G'].hidden)
                self.assertFalse(book['Rank & Allocation'].column_dimensions['F'].hidden)
                self.assertFalse(book['Rank & Allocation'].column_dimensions['H'].hidden)
                self.assertEqual(book['Rank & Allocation']['G2'].value, 'Ally')
                self.assertEqual(book['Rank & Allocation']['H2'].value, '=2*3')
                self.assertTrue(book['Legislative Agenda'].column_dimensions['C'].hidden)
                self.assertEqual(book['Legislative Agenda']['C2'].value, 'Ally')
                book.close()

    def test_view_keeps_absence_separate_from_zero_and_shared_rank(self):
        cell = dict(evidence_status='observed', action_count=2, raw_field_score='0', tensor_tiebreak=None,
                    rank_start=3, rank_end=5, cohort_size=18, recognition=None,
                    cutoffs=dict(champion=0, ally=2, highly_persuadable=3),
                    reason='nonpositive_record', alignment='Mixed actions')
        doc = dict(cats=['Tax'], generation_id='fixture', as_of='2026-09-12',
                   policy={'champion_pct':5,'ally_pct':15,'highly_persuadable_pct':20},
                   members={'Example': {'overall': dict(recognition=None, alignment='Mixed actions',
                       priority_tier=2, rank=14, stance='True Neutral', stance_rank_start=3, stance_rank_end=5,
                       stance_cohort_size=18, record_pattern='balanced', raw_overall_score='0', action_count=2,
                       positive_action_count=1, negative_action_count=1, neutral_action_count=0,
                       stance_cutoffs=dict(ally=2, highly_persuadable=3, persuadable=9, lean_persuadable=13),
                       opposition_evidence=dict(measures=[],lead_measures=[],congresses=[],label=None),
                       opposition_rank_start=None, opposition_rank_end=None, opposition_cohort_size=0,
                       opposition_cutoffs=dict(opponent=0, highly_skeptical=0, skeptic=0),
                       tensor_tiebreak=None, tensor_status='incomplete', stance_reason='balanced_record'),
                       'fields': {'Tax': cell}}})
        views = workbook_views(doc)
        detail = views[1]['values'][4]
        self.assertEqual(detail[4:7], [3, 5, 18])
        self.assertEqual(detail[8], 0)
        self.assertIsNone(detail[9])
        self.assertEqual(views[0]['values'][4][:6], ['Example', 'Tier 2', 14, 'True Neutral', 'Primacy', 'Balanced actions'])
        self.assertEqual(views[0]['values'][4][6], 'Below threshold')
        self.assertEqual(views[2]['values'][4][6], 0)
        self.assertIsNone(views[2]['values'][4][21])
        self.assertTrue(all(len(row) == 24 for row in views[2]['values']))
        self.assertTrue(all(len(row) == 14 for row in views[1]['values']))
        doc['policy']['champion_pct'] = 3
        self.assertIn('Champion top 3%', workbook_views(doc)[0]['values'][1][0])


if __name__ == '__main__':
    unittest.main()
