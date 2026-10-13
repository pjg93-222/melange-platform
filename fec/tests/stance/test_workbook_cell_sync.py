from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
import openpyxl
import pytest
from openpyxl.worksheet.table import Table
from stance.workbook_cell_sync import merge_cells,Q
from stance.workbook_stance_views import sync_category_stance_sheet,sync_coalition_profiles_sheet

def package(path,sheet,extra=None):
    with ZipFile(path,'w') as z:
        z.writestr('xl/workbook.xml','<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Main" sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr('xl/_rels/workbook.xml.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        z.writestr('xl/worksheets/sheet1.xml',sheet)
        for name,value in (extra or {}).items():z.writestr(name,value)

def test_changed_cells_preserve_native_parts_styles_and_untargeted_formulas(tmp_path):
    start='<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><dimension ref="A1:C1"/><sheetData>'
    end='</sheetData><sheetProtection sheet="1"/></worksheet>'
    source=tmp_path/'source.xlsx';authored=tmp_path/'authored.xlsx';out=tmp_path/'out.xlsx'
    package(source,start+'<row r="1"><c r="A1" s="9"><v>3</v></c><c r="B1" s="2"><f>A1*2</f><v>6</v></c><c r="C1"><v>7</v></c></row>'+end,
            {'xl/styles.xml':'native styles','xl/charts/chart1.xml':'native chart'})
    package(authored,start+'<row r="1"><c r="A1" s="1" t="s"><v>0</v></c><c r="B1"><v>999</v></c></row><row r="2"><c r="D2"><v>4</v></c></row>'+end,
            {'xl/sharedStrings.xml':'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>Champion</t></si></sst>'})
    merge_cells(source,authored,[{'sheet':'Main','cell':a} for a in ('A1','C1','D2')],out)
    with ZipFile(out) as z:
        assert z.read('xl/charts/chart1.xml')==b'native chart'
        assert z.read('xl/styles.xml')==b'native styles'
        tree=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        cells={c.get('r'):c for c in tree.iter(Q+'c')}
        assert cells['A1'].get('s')=='9'
        assert cells['A1'].find(Q+'is').find(Q+'t').text=='Champion'
        assert cells['B1'].find(Q+'f').text=='A1*2'
        assert cells['B1'].find(Q+'v').text=='6'
        assert cells['C1'].find(Q+'v') is None
        assert cells['D2'].find(Q+'v').text=='4'
        assert tree.find(Q+'dimension').get('ref')=='A1:D2'
        assert tree.find(Q+'sheetProtection').get('sheet')=='1'

MAIN='xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
TABLE_REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships/table'

def table_xml(number,ref,names):
    columns=''.join('<tableColumn id="%d" name="%s"%s/>'%(i,n,' dataDxfId="7"' if i==2 else '') for i,n in enumerate(names,1))
    return '<table %s id="%d" name="T%d" displayName="T%d" ref="%s"><autoFilter ref="%s"/><tableColumns count="%d">%s</tableColumns></table>'%(
        MAIN,number,number,number,ref,ref,len(names),columns)

def header_cells(sheet):
    return '<row r="1"><c r="A1" s="3" t="inlineStr"><is><t>Old A</t></is></c><c r="B1" t="inlineStr"><is><t>Keep &amp; Co</t></is></c>'            '<c r="E1" t="inlineStr"><is><t>X</t></is></c><c r="F1" t="inlineStr"><is><t>Y</t></is></c></row>'

def table_package(tmp_path,new_text):
    start='<worksheet %s><dimension ref="A1:F2"/><sheetData>'%MAIN
    end='</sheetData><tableParts count="2"/></worksheet>'
    rels=('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
          '<Relationship Id="rId1" Type="%s" Target="/xl/tables/table1.xml"/>'
          '<Relationship Id="rId2" Type="%s" Target="../tables/table2.xml"/></Relationships>')%(TABLE_REL,TABLE_REL)
    extra={'xl/worksheets/_rels/sheet1.xml.rels':rels,'xl/charts/chart1.xml':'native chart',
           'xl/tables/table1.xml':table_xml(1,'A1:C2',['Old A','Keep &amp; Co','Column2']),
           'xl/tables/table2.xml':table_xml(2,'E1:F2',['X','Y'])}
    source=tmp_path/'source.xlsx';authored=tmp_path/'authored.xlsx';out=tmp_path/'out.xlsx'
    package(source,start+header_cells(None)+'<row r="2"><c r="A2"><v>1</v></c></row>'+end,extra)
    package(authored,start+'<row r="1"><c r="A1" t="s"><v>0</v></c></row>'+end,
            {'xl/sharedStrings.xml':'<sst %s><si><t>%s</t></si></sst>'%(MAIN,new_text)})
    return source,authored,out

def test_table_columns_follow_their_header_cells_and_a_blank_header_is_filled(tmp_path):
    source,authored,out=table_package(tmp_path,'New A')
    parts=merge_cells(source,authored,[{'sheet':'Main','cell':'A1'}],out)
    assert parts==['xl/tables/table1.xml','xl/worksheets/sheet1.xml']      # table2 already agreed with its cells: untouched
    with ZipFile(source) as before,ZipFile(out) as z:
        names=[c.get('name') for c in ET.fromstring(z.read('xl/tables/table1.xml')).find(Q+'tableColumns')]
        assert names==['New A','Keep & Co','Column2']
        assert z.read('xl/tables/table1.xml').decode().count('dataDxfId="7"')==1       # every other attribute is kept
        assert z.read('xl/tables/table2.xml')==before.read('xl/tables/table2.xml')
        assert z.read('xl/charts/chart1.xml')==b'native chart'
        cells={c.get('r'):c for c in ET.fromstring(z.read('xl/worksheets/sheet1.xml')).iter(Q+'c')}
        assert cells['A1'].get('s')=='3' and cells['A1'].find(Q+'is').find(Q+'t').text=='New A'
        assert cells['C1'].find(Q+'is').find(Q+'t').text=='Column2'                  # a missing header cell now exists

def test_a_header_edit_that_duplicates_another_column_fails_closed(tmp_path):
    source,authored,out=table_package(tmp_path,'keep &amp; co')
    with pytest.raises(ValueError,match='duplicate column names'):
        merge_cells(source,authored,[{'sheet':'Main','cell':'A1'}],out)
    assert not out.exists()

def table_defects(path):
    """(sheet, header cell) pairs whose tableColumn name differs from the cell, or whose cell is empty."""
    bad=[]
    for ws in openpyxl.load_workbook(path):
        for table in ws.tables.values():
            for cell,column in zip(ws[table.ref][0],table.tableColumns):
                if cell.value!=column.name:bad.append((ws.title,cell.coordinate))
    return bad

def test_the_stance_view_writers_leave_every_table_matching_its_header_cells(tmp_path):
    old=['Cat %02d'%i for i in range(18)]
    new=old[1:]+old[:1]                                  # a re-ordered category list, as the live Category Stance sheet had
    wb=openpyxl.Workbook();wb.active.title='Rank & Allocation';wb.create_sheet('Blended Model')
    cs=wb.create_sheet('Category Stance')
    cs.append(['Member','General Stance']+old);cs.append(['Ada','Ally']+['Ally']*18)
    cs.add_table(Table(displayName='Table21',ref='A1:T2'))
    cp=wb.create_sheet('Coalition Profiles')
    cp.append([]);cp.append([])
    cp.append(['Coalition','Anchors','Members (n)','Mean Score','Mean Rank','Best Rank','Funded (n)','Total Allocation','% of Pot',
               'Champion','Ally','Neutral','Unengaged','Skeptic','Opponent','Dominant Stance'])
    for row in range(4,14):cp.cell(row,1).value='Bloc %d'%row
    cp.cell(16,1).value='BY CU-STANCE TIER'
    for col in range(2,7):cp.cell(16,col).value='Column%d'%(col-1)
    cp.add_table(Table(displayName='Table24',ref='A3:P13'));cp.add_table(Table(displayName='Table25',ref='A16:F24'))
    source=tmp_path/'source.xlsx';authored=tmp_path/'authored.xlsx';out=tmp_path/'out.xlsx'
    wb.save(source)
    assert table_defects(source)==[]
    wb=openpyxl.load_workbook(source)
    before={(ws.title,c.coordinate):c.value for ws in wb for row in ws for c in row if c.value is not None}
    sync_category_stance_sheet(wb,{'cats':new,'matrix':{'Ada':{c:'Ally' for c in new}}},['Ada'],{'Ada':'Ally'})
    sync_coalition_profiles_sheet(wb)
    changes=[{'sheet':ws.title,'cell':c.coordinate,'value':c.value} for ws in wb for row in ws for c in row
             if before.get((ws.title,c.coordinate))!=c.value]
    wb.save(authored)
    merge_cells(source,authored,changes,out)
    assert table_defects(out)==[]
    result=openpyxl.load_workbook(out)
    assert [c.name for c in result['Category Stance'].tables['Table21'].tableColumns][2:]==new
    assert result['Coalition Profiles']['A16'].value=='BY CURRENT CU-STANCE TIER'
    assert [result['Coalition Profiles'].cell(16,col).value for col in range(2,7)]==['Column%d'%i for i in range(1,6)]
