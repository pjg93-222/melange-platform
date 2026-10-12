from __future__ import annotations

import sys
from pathlib import Path
import unittest

import openpyxl


FEC = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(FEC))

from stance.workbook_stance_views import (  # noqa: E402
    GENERAL_STANCE_TIERS,
    stance_view_contract,
    sync_category_stance_sheet,
    sync_coalition_profiles_sheet,
)


class WorkbookStanceViewsTests(unittest.TestCase):
    def workbook(self):
        wb = openpyxl.Workbook()
        wb.active.title = "Rank & Allocation"
        wb.create_sheet("Blended Model")
        cp = wb.create_sheet("Coalition Profiles")
        for row, bloc in enumerate(
            ("Bloc 1", "Bloc 2", "Bloc 3", "Bloc 4", "Bloc 5",
             "Bloc 6", "Bloc 7", "Bloc 8", "Unnetworked"),
            start=4,
        ):
            cp.cell(row, 1).value = bloc
        return wb

    def test_contract_requires_champions_to_equal_tier_one_top_13(self):
        members = [f"Member {index:02d}" for index in range(1, 15)]
        category = {"cats": [f"Category {i}" for i in range(18)],
                    "matrix": {m: {} for m in members}}
        overall = {m: ("Champion" if i < 13 else "Ally") for i, m in enumerate(members)}
        tiers = {m: (1 if i < 13 else 0) for i, m in enumerate(members)}
        receipt = stance_view_contract(category, members, overall, tiers)
        self.assertEqual(13, len(receipt["champions"]))
        overall[members[-1]] = "Champion"
        with self.assertRaisesRegex(ValueError, "top 13"):
            stance_view_contract(category, members, overall, tiers)

    def test_category_stance_is_written_in_live_category_order(self):
        wb = self.workbook()
        members = ["Ada", "Grace"]
        cats = [f"Category {i}" for i in range(18)]
        matrix = {
            "Ada": {cat: "Ally" for cat in cats},
            "Grace": {cat: "Persuadable" for cat in cats},
        }
        sync_category_stance_sheet(
            wb, {"cats": cats, "matrix": matrix}, members,
            {"Ada": "Champion", "Grace": "True Neutral"},
        )
        ws = wb["Category Stance"]
        self.assertEqual(cats, [ws.cell(1, c).value for c in range(3, 21)])
        self.assertEqual("Champion", ws["B2"].value)
        self.assertEqual("Persuadable", ws["T3"].value)

    def test_coalition_profiles_counts_every_current_tier(self):
        wb = self.workbook()
        sync_coalition_profiles_sheet(wb)
        ws = wb["Coalition Profiles"]
        self.assertEqual(
            GENERAL_STANCE_TIERS,
            tuple(ws.cell(row, 1).value for row in range(18, 29)),
        )
        self.assertEqual(
            GENERAL_STANCE_TIERS,
            tuple(ws.cell(3, col).value for col in range(10, 21)),
        )
        self.assertEqual("=SUM(B18:B28)", ws["B29"].value)
        self.assertIn("'Rank & Allocation'!$G$3:$G$535", ws["B18"].value)
        self.assertIn("$J$3:$T$3", ws["U4"].value)


if __name__ == "__main__":
    unittest.main(verbosity=2)
