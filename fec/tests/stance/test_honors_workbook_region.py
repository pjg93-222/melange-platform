# -*- coding: utf-8 -*-
"""The honors workbook binding owns only the sheets a publication writes.

Until 2026-09-18 the binding was a hash of the WHOLE workbook file, so an edit to
any sheet the publication never touches -- Committee Data, which is Paul's to
edit -- made `--prepare-model-update` raise 'Workbook changed since honors
publication' and blocked the chain. It cost this lane two manual rebases
(2026-09-17 and 2026-09-18; audits/members_section_20260914/evidence/
binding-rebase-20260918.md).

These suites pin the scoped replacement and its revert switch on synthetic
workbooks: an edit OUTSIDE the region rolls back and survives, an edit INSIDE an
honors view still refuses by name, and HONORS_WORKBOOK_SCOPE=legacy reproduces
the old whole-file refusal. The acceptance proofs against a real 533-member
publication live in the lane packet; these are the registered regression.
"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from stance import honors_workbook as H
from stance.workbook_cell_sync import sheet_paths
import openpyxl

# Deliberately not rectangular and deliberately mixed-typed: trailing blanks are
# packaging rather than content, and an integer read back out of XLSX where a
# float was authored must not read as an edit to the honors output.
VIEWS = [{'name': 'Field Recognition',
          'values': [['Melange — Field Recognition', None, None],
                     ['Member', 'Overall priority tier', 'Tax'],
                     ['Example Member', 'Tier 1 · Champion', 'Champion']]},
         {'name': 'Recognition Detail',
          'values': [['Member', 'Rank from', 'Raw action score'],
                     ['Example Member', 3, 0.0]]},
         {'name': 'Overall Standing',
          'values': [['Member', 'Actions'], ['Example Member', 2]]}]
HONORS = {'generation_id': 'fixture-generation', 'members': {'Example Member': None}}


class RegionScopedBindingTests(unittest.TestCase):
    """Every case publishes a real merge_views package onto a real model workbook."""

    def model(self, root):
        """A model workbook shaped like the live one: both label-column sheets, both
        tabs a publication hides, formulas, and a sheet that is Paul's to edit."""
        workbook = root / 'model.xlsx'
        book = openpyxl.Workbook()
        book.active.title = 'Rank & Allocation'
        allocation = book.active
        allocation['G1'], allocation['G2'] = 'General Stance', 'Ally'
        allocation['H2'] = '=2*3'
        allocation.column_dimensions.group('F', 'H', hidden=False)
        agenda = book.create_sheet('Legislative Agenda')
        agenda['C1'], agenda['C2'] = 'General Stance', 'Ally'
        book.create_sheet('Category Stance')['A1'] = 'Legacy model input'
        book.create_sheet('Coalition Profiles')['A1'] = 'Legacy analytical ladder'
        committee = book.create_sheet('Committee Data')
        committee['A1'], committee['E6'] = 'Member (model)', 'Banking (full) [Highly]'
        committee['B6'] = '=1+1'
        book.save(workbook)
        book.close()
        (root / 'data/engine').mkdir(parents=True)
        (root / 'data/engine/artifact_tool_runtime.json').write_text(
            json.dumps({'node': 'fixture-node', 'node_modules': 'fixture-modules'}), encoding='utf-8')
        return workbook

    def author(self, arguments, **kwargs):
        helper, plan_path = Path(arguments[1]), Path(arguments[2])
        plan = json.loads(plan_path.read_text(encoding='utf-8'))
        authored = Path(plan['output'])
        book = openpyxl.Workbook()
        book.remove(book.active)
        for view in plan['views']:
            sheet = book.create_sheet(view['name'])
            for row in view['values']:
                sheet.append(row)
        book.save(authored)
        book.close()
        Path(str(authored) + '.receipt.json').write_text(json.dumps(
            {'status': 'PASS', 'plan_sha256': H.sha(plan_path), 'helper_sha256': H.sha(helper),
             'output_sha256': H.sha(authored)}), encoding='utf-8')
        return SimpleNamespace(returncode=0)

    def publish(self, workbook, root, generation=HONORS['generation_id']):
        honors = dict(HONORS, generation_id=generation)
        with patch('stance.honors_views.load_honors', return_value=honors), \
             patch('stance.honors_workbook.workbook_views', return_value=VIEWS), \
             patch('stance.honors_workbook.subprocess.run', side_effect=self.author):
            return H.publish(workbook, root)

    def published(self, tmp):
        """A bound publication over a real model workbook, ready to be edited."""
        root = Path(tmp)
        workbook = self.model(root)
        self.assertEqual(self.publish(workbook, root)['status'], 'PASS')
        return root, workbook, root / 'data/stance/rank_honors/workbook_publications.json'

    def edit(self, workbook, sheet, address, value):
        """The realistic intervening edit: an openpyxl rewrite, exactly what this
        lane's own retier and recompute scripts do to Committee Data."""
        book = openpyxl.load_workbook(workbook)
        book[sheet][address] = value
        book.save(workbook)
        book.close()

    def value(self, workbook, sheet, address):
        book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
        try:
            return book[sheet][address].value
        finally:
            book.close()

    def names(self, workbook):
        book = openpyxl.load_workbook(workbook, read_only=True)
        try:
            return list(book.sheetnames)
        finally:
            book.close()

    # ---------------------------------------------- the defect this lane hit twice
    def test_an_edit_outside_the_honors_region_no_longer_blocks_the_model_update(self):
        with tempfile.TemporaryDirectory() as tmp:
            root, workbook, binding_path = self.published(tmp)
            self.edit(workbook, 'Committee Data', 'E6', 'Banking (full) [Slightly]')
            outcome = H.prepare_model_update(workbook, root)
            self.assertEqual(outcome['status'], 'ROLLED_BACK')
            self.assertEqual(outcome['scope'], 'region')
            # The edit survives, the views are gone, and the binding is released.
            self.assertEqual(self.value(workbook, 'Committee Data', 'E6'), 'Banking (full) [Slightly]')
            self.assertFalse(set(H.SHEETS) & set(self.names(workbook)))
            self.assertEqual(json.loads(binding_path.read_text(encoding='utf-8')), {})
            # Everything the publication flipped is back at its pre-publication value.
            region = H.honors_region(workbook)
            self.assertIs(region['columns']['Rank & Allocation!7'], False)
            self.assertIs(region['columns']['Legislative Agenda!3'], False)
            self.assertEqual(region['tabs'], {'Category Stance': 'visible', 'Coalition Profiles': 'visible'})
            # ... and nothing the publication never wrote moved at all.
            self.assertEqual(self.value(workbook, 'Rank & Allocation', 'G2'), 'Ally')
            book = openpyxl.load_workbook(workbook)
            self.assertEqual(book['Rank & Allocation']['H2'].value, '=2*3')
            self.assertEqual(book['Committee Data']['B6'].value, '=1+1')
            book.close()

    def test_the_scoped_rollback_removes_the_honors_parts_and_nothing_else(self):
        with tempfile.TemporaryDirectory() as tmp:
            root, workbook, _ = self.published(tmp)
            self.edit(workbook, 'Committee Data', 'E6', 'edited between publications')
            with ZipFile(workbook) as package:
                before = {name: package.read(name) for name in package.namelist()}
                honors_parts = {sheet_paths(package)[name] for name in H.SHEETS}
            reverted = set(H.prepare_model_update(workbook, root)['reverted_parts'])
            with ZipFile(workbook) as package:
                after = {name: package.read(name) for name in package.namelist()}
                self.assertIsNone(package.testzip())
            self.assertEqual(set(before) - set(after), honors_parts)
            self.assertEqual(set(after) - set(before), set())
            self.assertLessEqual(honors_parts | {'xl/workbook.xml'}, reverted)
            # Only parts the publication wrote may differ. xl/styles.xml keeps its
            # appended entries on purpose: an intervening edit may reference a style
            # index appended after the publication, and restoring the old styles
            # would repoint it.
            self.assertEqual({name for name in after if after[name] != before[name]} - reverted, set())
            self.assertEqual(after['xl/styles.xml'], before['xl/styles.xml'])

    # ------------------------------------------ the protection that must still hold
    def test_an_edit_inside_an_honors_view_still_refuses_and_names_the_sheet(self):
        for sheet in H.SHEETS:
            with self.subTest(sheet=sheet), tempfile.TemporaryDirectory() as tmp:
                root, workbook, binding_path = self.published(tmp)
                self.edit(workbook, sheet, 'A3', 'a hand edit of the honors output')
                frozen = {path: path.read_bytes() for path in (workbook, binding_path)}
                with self.assertRaises(RuntimeError) as refusal:
                    H.prepare_model_update(workbook, root)
                self.assertIn(sheet, str(refusal.exception))
                self.assertIn('honors output', str(refusal.exception))
                self.assertEqual(frozen, {path: path.read_bytes() for path in frozen})

    def test_a_removed_honors_view_refuses_rather_than_rolling_back_silently(self):
        with tempfile.TemporaryDirectory() as tmp:
            root, workbook, _ = self.published(tmp)
            book = openpyxl.load_workbook(workbook)
            book.remove(book['Overall Standing'])
            book.save(workbook)
            book.close()
            with self.assertRaises(RuntimeError) as refusal:
                H.prepare_model_update(workbook, root)
            self.assertIn('Overall Standing', str(refusal.exception))
            self.assertIn('was removed', str(refusal.exception))

    def test_republication_refuses_an_honors_edit_on_the_same_terms(self):
        """The publish-side guard reads the same binding and must scope identically."""
        with tempfile.TemporaryDirectory() as tmp:
            root, workbook, _ = self.published(tmp)
            self.edit(workbook, 'Field Recognition', 'A3', 'a hand edit of the honors output')
            with self.assertRaises(RuntimeError) as refusal:
                self.publish(workbook, root)
            self.assertIn('Field Recognition', str(refusal.exception))

    # -------------------------------------------------------------------- the dial
    def test_the_legacy_dial_restores_the_whole_file_refusal(self):
        with tempfile.TemporaryDirectory() as tmp:
            root, workbook, binding_path = self.published(tmp)
            self.edit(workbook, 'Committee Data', 'E6', 'Banking (full) [Slightly]')
            frozen = {path: path.read_bytes() for path in (workbook, binding_path)}
            with patch('stance.honors_workbook.HONORS_WORKBOOK_SCOPE', 'legacy'):
                with self.assertRaises(RuntimeError) as refusal:
                    H.prepare_model_update(workbook, root)
                with self.assertRaisesRegex(RuntimeError, 'intervening edits'):
                    self.publish(workbook, root)
            self.assertEqual(str(refusal.exception),
                             'Workbook changed since honors publication; refusing rollback over intervening edits')
            self.assertEqual(frozen, {path: path.read_bytes() for path in frozen})
            # The same state rolls back under the shipped scope: the dial is the only
            # difference between a blocked chain and a running one.
            self.assertEqual(H.prepare_model_update(workbook, root)['scope'], 'region')

    def test_the_dial_is_read_from_the_environment_with_region_as_the_default(self):
        source = Path(H.__file__).read_text(encoding='utf-8')
        self.assertIn("HONORS_WORKBOOK_SCOPE = os.environ.get('HONORS_WORKBOOK_SCOPE', 'region')", source)
        self.assertIn('HONORS_WORKBOOK_SCOPE=legacy', H.__doc__)
        self.assertEqual(H.HONORS_WORKBOOK_SCOPE, os.environ.get('HONORS_WORKBOOK_SCOPE', 'region'))

    # ------------------------------------------------------- the lossless fast path
    def test_an_untouched_publication_still_rolls_back_byte_for_byte(self):
        """Nothing else moved, so the recorded package is restorable exactly. The
        scoped path must not cost the chain that losslessness."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = self.model(root)
            original = workbook.read_bytes()
            self.publish(workbook, root)
            self.assertNotEqual(workbook.read_bytes(), original)
            outcome = H.prepare_model_update(workbook, root)
            self.assertEqual(outcome['scope'], 'whole-file')
            self.assertEqual(workbook.read_bytes(), original)

    def test_a_publication_over_intervening_edits_may_only_be_reverted_by_region(self):
        """A publication that lands on top of someone's edits inherits the PREVIOUS
        publication's rollback copy, which predates them. The byte-exact whole-file
        fast path must close for that binding, or the next cycle silently discards
        their work -- the very loss the old blanket refusal was protecting."""
        with tempfile.TemporaryDirectory() as tmp:
            root, workbook, binding_path = self.published(tmp)

            def binding():
                return json.loads(binding_path.read_text(encoding='utf-8'))[str(workbook.resolve())]

            first = binding()['original_path']
            self.assertIs(binding()['original_exact'], True)
            self.edit(workbook, 'Committee Data', 'E6', 'Banking (full) [Slightly]')
            self.assertEqual(self.publish(workbook, root, 'second-generation')['status'], 'PASS')
            # Same copy, which no longer restores the whole file losslessly.
            self.assertEqual(binding()['original_path'], first)
            self.assertIs(binding()['original_exact'], False)
            # Sticky: a later clean publication must not hand exactness back.
            self.assertEqual(self.publish(workbook, root, 'third-generation')['status'], 'PASS')
            self.assertIs(binding()['original_exact'], False)
            self.assertEqual(H.prepare_model_update(workbook, root)['scope'], 'region')
            self.assertEqual(self.value(workbook, 'Committee Data', 'E6'), 'Banking (full) [Slightly]')
            # The next publication starts a fresh, exact copy from the kept state.
            self.assertEqual(self.publish(workbook, root, 'fourth-generation')['status'], 'PASS')
            self.assertIs(binding()['original_exact'], True)
            self.assertNotEqual(binding()['original_path'], first)

    # --------------------------------------------------------------- the binding
    def test_the_binding_records_the_region_and_keeps_the_whole_file_hashes(self):
        with tempfile.TemporaryDirectory() as tmp:
            root, workbook, binding_path = self.published(tmp)
            binding = json.loads(binding_path.read_text(encoding='utf-8'))[str(workbook.resolve())]
            self.assertEqual(binding['after_sha256'], H.sha(workbook))
            self.assertEqual(H.sha(root / binding['original_path']), binding['before_sha256'])
            region = binding['honors_region']
            self.assertEqual(region['schema'], H.REGION_SCHEMA)
            self.assertEqual(sorted(region['sheets']), sorted(H.SHEETS))
            self.assertTrue(all(region['sheets'].values()))
            # The recorded digest is the digest of what is actually on disk.
            self.assertEqual(H.honors_region(workbook)['sheets'], region['sheets'])
            self.assertEqual(region['columns'], {'Rank & Allocation!7': True, 'Legislative Agenda!3': True})
            self.assertEqual(region['tabs'], {'Category Stance': 'hidden', 'Coalition Profiles': 'hidden'})
            self.assertIs(binding['original_exact'], True)

    def test_a_pre_scope_binding_is_rebuilt_from_its_own_publication_plan(self):
        """The binding sitting in the live tree today carries no honors_region field."""
        with tempfile.TemporaryDirectory() as tmp:
            root, workbook, binding_path = self.published(tmp)
            bindings = json.loads(binding_path.read_text(encoding='utf-8'))
            key = str(workbook.resolve())
            recorded = bindings[key].pop('honors_region')
            binding_path.write_text(json.dumps(bindings), encoding='utf-8')
            self.assertEqual(H._published_region(bindings[key], root)['sheets'], recorded['sheets'])
            self.edit(workbook, 'Committee Data', 'E6', 'Banking (full) [Slightly]')
            self.assertEqual(H.prepare_model_update(workbook, root)['scope'], 'region')
            self.assertEqual(self.value(workbook, 'Committee Data', 'E6'), 'Banking (full) [Slightly]')

    def test_a_plan_that_is_not_this_publications_is_refused_not_trusted(self):
        with tempfile.TemporaryDirectory() as tmp:
            root, workbook, binding_path = self.published(tmp)
            bindings = json.loads(binding_path.read_text(encoding='utf-8'))
            key = str(workbook.resolve())
            for field, value in (('after_sha256', '0' * 64), ('generation_id', 'another-generation')):
                previous = dict(bindings[key])
                previous.pop('honors_region')
                previous[field] = value
                with self.subTest(field=field):
                    self.assertIsNone(H._published_region(previous, root))
            # ... and a binding that cannot be scoped fails CLOSED, never open.
            previous = dict(bindings[key])
            previous.pop('honors_region')
            previous['after_sha256'] = '0' * 64
            bindings[key] = previous
            binding_path.write_text(json.dumps(bindings), encoding='utf-8')
            self.edit(workbook, 'Committee Data', 'E6', 'Banking (full) [Slightly]')
            with self.assertRaisesRegex(RuntimeError, 'intervening edits'):
                H.prepare_model_update(workbook, root)

    # ------------------------------------------------------------------ robustness
    def test_the_region_digest_is_value_semantic_not_byte_level(self):
        """A save rewrites part names, namespace prefixes and string encodings. A
        byte digest would read that as an edit to the honors output and refuse."""
        with tempfile.TemporaryDirectory() as tmp:
            root, workbook, _ = self.published(tmp)
            with ZipFile(workbook) as package:
                before = {name: (sheet_paths(package)[name], package.read(sheet_paths(package)[name]))
                          for name in H.SHEETS}
            published = H.honors_region(workbook)['sheets']
            self.edit(workbook, 'Category Stance', 'A1', 'a save that rewrites every part')
            with ZipFile(workbook) as package:
                after = {name: (sheet_paths(package)[name], package.read(sheet_paths(package)[name]))
                         for name in H.SHEETS}
            self.assertNotEqual(before, after)
            self.assertNotEqual({path for path, _ in before.values()}, {path for path, _ in after.values()})
            self.assertEqual(H.honors_region(workbook)['sheets'], published)
            self.assertEqual(H.prepare_model_update(workbook, root)['scope'], 'region')

    def test_a_package_that_cannot_be_opened_is_owned_whole(self):
        """Fail-safe, not fail-open: no resolvable region means the legacy binding."""
        with tempfile.TemporaryDirectory() as tmp:
            workbook = Path(tmp) / 'model.xlsx'
            workbook.write_bytes(b'not an OPC package')
            region = H.honors_region(workbook)
            self.assertFalse(region['package'])
            self.assertEqual(region['whole_file_sha256'], H.sha(workbook))
            self.assertEqual(H._guard(region)['whole_file_sha256'], H.sha(workbook))


if __name__ == '__main__':
    unittest.main()
