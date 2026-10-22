"""Temp-only regression proposals. No server, process enumeration, or live data.

Workbook functions are compiled from their actual source AST to avoid importing
the workbook/model dependency graph. Glass Box modules only define objects at
import; tests instantiate the in-memory store and replace process lookup.
"""
import ast
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import types
import unittest
from unittest import mock

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import roster  # noqa: E402
ROOT = Path(os.environ.get('FEC_SURFACE_SOURCE_ROOT', Path(__file__).resolve().parents[3])) / 'fec/surfaces'


def module(name):
    spec = importlib.util.spec_from_file_location('surface_fixture_' + name, ROOT / (name + '.py'))
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


glass = module('glassbox_app_server')
emitter = module('glassbox_emit')


def extraction_namespace(loader):
    tree = ast.parse((ROOT / 'extract_workbook_audit_inputs.py').read_bytes())
    selected = ast.Module(body=[n for n in tree.body if isinstance(n, ast.FunctionDef)
                               and n.name in {'sha256_file', 'number', 'normalize_columns', 'extract'}], type_ignores=[])
    env = {'hashlib': hashlib, 'os': os, 'io': io, 'tempfile': tempfile, 'roster': roster,
           'openpyxl': types.SimpleNamespace(load_workbook=loader)}
    exec(compile(selected, str(ROOT / 'extract_workbook_audit_inputs.py'), 'exec'), env)
    return env


class FakeWorkbook:
    def __init__(self, missing=False):
        self.closed = False
        self.missing = missing

    def __getitem__(self, name):
        if self.missing:
            raise KeyError('synthetic missing sheet')
        rows = ([tuple(['header', 'type'] + [0] * 158)] * 2
                + [tuple(['member', 'type'] + [0] * 158)]) if name.endswith('Raw') else (
                    [('member', None, None, 0, None, None, 'fixture')] if name == 'Rank & Allocation'
                    else [tuple(['member', 'type'] + [0] * 158)])
        return types.SimpleNamespace(iter_rows=lambda **kwargs: iter(rows))

    def close(self):
        self.closed = True


class SurfaceSnapshotTests(unittest.TestCase):
    def test_workbook_hash_binds_same_bytes_as_extraction(self):
        with tempfile.TemporaryDirectory(prefix='surface-workbook-') as temporary:
            path = Path(temporary) / 'synthetic.xlsx'
            path.write_bytes(b'first workbook snapshot')
            consumed = []
            book = FakeWorkbook()

            def loader(source, **kwargs):
                consumed.append(source.read() if hasattr(source, 'read') else Path(source).read_bytes())
                path.write_bytes(b'replacement workbook snapshot')
                return book

            result = extraction_namespace(loader)['extract'](path)
            self.assertEqual(result['source_workbook_sha256'], hashlib.sha256(consumed[0]).hexdigest())
            self.assertTrue(book.closed)

    def test_missing_sheet_closes_extraction_handle(self):
        with tempfile.TemporaryDirectory(prefix='surface-workbook-') as temporary:
            path = Path(temporary) / 'synthetic.xlsx'
            path.write_bytes(b'fixture')
            book = FakeWorkbook(missing=True)
            with self.assertRaises(KeyError):
                extraction_namespace(lambda *a, **k: book)['extract'](path)
            self.assertTrue(book.closed)

    def test_event_id_exposes_actual_pid_to_existing_parser(self):
        with tempfile.TemporaryDirectory(prefix='surface-events-') as temporary, mock.patch.object(emitter, 'EVENTS_DIR', Path(temporary)):
            ident = emitter.emit('fixture', 'running', 'synthetic event')
            self.assertIsNotNone(ident)
            match = glass.PID_RE.search(ident + '.json')
            self.assertIsNotNone(match)
            self.assertEqual(int(match.group(1)), os.getpid())
            self.assertEqual(len(list(Path(temporary).iterdir())), 1)

    def test_nonfinite_numbers_are_unavailable(self):
        for value in (float('nan'), float('inf'), -float('inf'), 'NaN', 'Infinity', '-inf', 10 ** 400):
            with self.subTest(value_type=type(value).__name__):
                self.assertIsNone(glass.as_number(value))
        for value in (0, -3.5, '12.5'):
            self.assertEqual(glass.as_number(value), float(value))

    def test_event_numeric_fields_keep_all_views_json_serializable(self):
        store = glass.EventStore('not-accessed')
        store.live_pids = lambda *args: {os.getpid()}
        store._ingest({'lane': 'fixture', 'source': 'fixture', 'run_id': 'fixture',
                       'status': 'running', 'rate': float('nan'), 'current': 'inf',
                       'total': 10, 'workers': float('inf')},
                      'evt-fixture-py-%d-tail.json' % os.getpid(), 1.0)
        rows, now = store.snapshot_rows()
        for view in (rows, store.feed(), store.history('fixture|fixture|fixture'), store.summary(rows, now)):
            json.dumps(view, allow_nan=False)

    def test_nonfinite_extra_alone_does_not_break_rows_or_history(self):
        store = glass.EventStore('not-accessed')
        store.live_pids = lambda *args: {os.getpid()}
        store._ingest({'lane': 'fixture', 'source': 'fixture', 'run_id': 'fixture',
                       'status': 'running', 'workers': float('inf')},
                      'evt-fixture-py-%d-tail.json' % os.getpid(), 1.0)
        rows, _ = store.snapshot_rows()
        for view in (rows, store.history('fixture|fixture|fixture')):
            json.dumps(view, allow_nan=False)

    def test_finite_aggregate_overflow_is_unavailable_without_losing_rows(self):
        store = glass.EventStore('not-accessed')
        store.live_pids = lambda *args: {os.getpid()}
        for index in range(2):
            store._ingest({'lane': 'fixture', 'source': 'fixture', 'run_id': str(index),
                           'status': 'running', 'current': 1e308, 'total': 1e308},
                          'evt-fixture-py-%d-%d.json' % (os.getpid(), index), 1.0)
        rows, now = store.snapshot_rows()
        summary = store.summary(rows, now)
        json.dumps(summary, allow_nan=False)
        self.assertEqual(summary['rows_total'], 2)
        self.assertTrue(all(row['current'] == 1e308 for row in rows))
        self.assertIsNone(summary['overall']['current'])
        self.assertIsNone(summary['overall']['total'])
        self.assertIsNone(summary['overall']['pct'])

    def test_valid_numeric_progress_and_extra_remain_unchanged(self):
        store = glass.EventStore('not-accessed')
        store.live_pids = lambda *args: {os.getpid()}
        store._ingest({'lane': 'fixture', 'source': 'fixture', 'run_id': 'fixture',
                       'status': 'running', 'current': 2, 'total': 8, 'workers': 3,
                       'rate': '1.5', 'transport': 'fixture'},
                      'evt-fixture-py-%d-tail.json' % os.getpid(), 1.0)
        rows, now = store.snapshot_rows()
        self.assertEqual(rows[0]['extra'], {'workers': 3, 'transport': 'fixture'})
        self.assertEqual(rows[0]['rate'], 1.5)
        self.assertEqual(store.summary(rows, now)['overall'], {'current': 2.0, 'total': 8.0, 'pct': 25.0, 'runs': 1})

    def test_real_workbook_snapshot_extracts_expected_cells(self):
        import openpyxl
        with tempfile.TemporaryDirectory(prefix='surface-real-workbook-') as temporary:
            path = Path(temporary) / 'synthetic.xlsx'
            book = openpyxl.Workbook()
            book.remove(book.active)
            for name in ('Legislative Score - Raw', 'Rank & Allocation', 'Legislative Score - Normalized'):
                sheet = book.create_sheet(name)
                sheet.cell(3, 1, 'Fixture Member')
            book['Legislative Score - Raw'].cell(1, 160, 'header')
            book['Rank & Allocation'].cell(3, 4, 7)
            book['Rank & Allocation'].cell(3, 7, 'fixture label')
            book.save(path)
            book.close()
            expected = hashlib.sha256(path.read_bytes()).hexdigest()
            result = extraction_namespace(openpyxl.load_workbook)['extract'](path)
            self.assertEqual(result['source_workbook_sha256'], expected)
            self.assertEqual(result['rank_rows'][0], ['Fixture Member', 7.0, 'fixture label'])


if __name__ == '__main__':
    unittest.main(verbosity=2)
