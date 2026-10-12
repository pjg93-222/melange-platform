# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
"""QUORUM MASTER PIPELINE (user spec 2026-07-01, confirmed step by step):
  1. CU-domain RELEVANCE filter (same CU_TERMS the collector uses).
  2. TRUE-duplicate removal ONLY: key = (Source, normalized Content, Date).  Retweets are KEPT —
     RT overlap across sources is engagement signal (Fink layer), never deduped.
  3. AUTHOR SEGREGATION: current members (incl. their campaign-tagged posts and pre-office
     candidate posts — "even when they were running") = main corpus; staffers/aides KEPT + LABELED
     with the principal member office where parsable (user's aside asset); former members,
     non-member candidates, executive-branch, committees, orgs segregated with labels.
  4. Tweet-form stance scoring (stance_clf_composed.pkl, the golden-label-trained LinearSVM) on the
     member+relevant corpus, with decision-margin confidence and RT flag.
Outputs: 'Quorum Processed.xlsx' (sheets: Members, Staffers, Excluded, Summary) +
quorum_member_tweets.csv for downstream pipeline use."""
import re,json,sys,collections,unicodedata,pickle
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
import pandas as pd
M=MB_ROOT
# full vocabulary of the corrected Quorum query (quorum_query.txt) — keep in sync with it
CU_TERMS=re.compile(r'credit union|\bncua\b|national credit union|share insurance|ncusif|central liquidity|'
    r'tax.?exempt|tax exemption|tax status|field of membership|common bond|federal charter|charter modernization|'
    r'low.?income designation|\bcdfi\b|community development financial|minority depository|'
    r'\bcfpb\b|consumer financial protection|section 1071|small business lending rule|section 1033|open banking|'
    r'interchange|swipe fee|credit card competition|\bccca\b|durbin|routing mandate|debit card|credit card fee|'
    r'late fee|junk fee|overdraft|nsf fee|insufficient funds fee|'
    r'member business lending|mbl cap|trigger lead|\bhmda\b|small.?dollar loan|payday|'
    r'community bank|big bank|megabank|wall street bank|bank consolidation|de novo charter|industrial loan company|'
    r'fednow|real.?time payments|faster payments|payments system|stablecoin|digital dollar|'
    r'central bank digital currency|\bcbdc\b|earned wage access|'
    r'check fraud|elder financial|elder fraud|p2p fraud|peer.?to.?peer payment|regulation e\b|'
    r'data breach notification|data security standards|gramm.?leach|\btcpa\b|telephone consumer protection|robocall|'
    r'military lending act|defense credit union|military banking|banks on military|'
    r'federal home loan bank|gse reform|mortgage servicing|loan participations|credit invisible|'
    r'financial inclusion|de.?risking|brokered deposits',re.I)
STAFF_TITLES=re.compile(r'\b(Director|Chief of Staff|Press Secretary|Secretary|Advisor|Adviser|Counsel|Coordinator|'
    r'Assistant|Manager|Scheduler|Liaison|Aide|Spokes|Staff|Caseworker|Constituent Services|Field Representative|'
    r'Regional Representative|Legislative|Outreach|Deputy|Intern)\b',re.I)
EXEC_PAT=re.compile(r'White House|Office of the (President|Vice President)|Special Envoy|Ambassador|'
    r'\bSecretary of\b|Administrator|Department of',re.I)
cong={m['member'] for m in json.load(open(M+r'\data\engine\congestion.json'))}
def nfc(s): return unicodedata.normalize('NFC',s or '')
def fold(s):   # accent-fold for matching (Quorum ships ASCII 'Velazquez'; congestion has 'Velázquez')
    return ''.join(c for c in unicodedata.normalize('NFD',s or '') if not unicodedata.combining(c)).lower()
byfold={fold(m):m for m in cong}
sur2mem=collections.defaultdict(list)
for m in cong: sur2mem[fold(m.split()[-1])].append(m)
NICK={
    'buddy carter':'Earl Carter','bobby scott':'Robert Scott',
    'dick durbin':'Richard Durbin','zach nunn':'Zachary Nunn',
    'jan schakowsky':'Janice Schakowsky','bernie sanders':'Bernard Sanders',
    'jerry nadler':'Jerrold Nadler','jen kiggans':'Jennifer Kiggans',
    'eli crane':'Elijah Crane','hank johnson':'Henry Johnson',
    'brad schneider':'Bradley Schneider','bill keating':'William Keating',
    'lou correa':'Luis Correa',
}   # Quorum uses nicknames; congestion is canonical
FIRST_EQUIV=[
    {'richard','dick','rick','rich'}, {'charles','chuck'}, {'william','bill','billy'},
    {'michael','mike'}, {'robert','bob','bobby','rob'}, {'edward','ed','eddie'},
    {'james','jim','jimmy'}, {'thomas','tom','tommy'}, {'john','jack','johnny'},
    {'joseph','joe'}, {'gerald','jerry'}, {'jerrold','jerry'}, {'daniel','dan'},
    {'donald','don'}, {'ronald','ron'}, {'andrew','andy'}, {'kenneth','ken'},
    {'benjamin','ben'}, {'nicholas','nick'}, {'christopher','chris'},
    {'stephen','steve'}, {'timothy','tim'}, {'matthew','matt'}, {'david','dave'},
    {'samuel','sam'}, {'patrick','pat'}, {'anthony','tony'}, {'alexander','alex'},
    {'katherine','katie','kate'}, {'elizabeth','liz','beth'}, {'margaret','maggie'},
    {'cynthia','cindy'}, {'earl','buddy'}, {'addison','mitch'}, {'rafael','ted'},
    {'louis','lou'}, {'luis','lou'}, {'francis','frank'}, {'rebecca','becca'},
    {'abraham','abe'}, {'gregory','greg'}, {'douglas','doug'}, {'raymond','ray'},
    {'harold','hal'}, {'zachary','zach'}, {'janice','jan'}, {'bernard','bernie'},
    {'jennifer','jen'}, {'elijah','eli'}, {'henry','hank'}, {'bradley','brad'},
]
FIRST_CANON={n:i for i,g in enumerate(FIRST_EQUIV) for n in g}
def same_first(a,b):
    a,b=fold(a),fold(b)
    return a==b or (a in FIRST_CANON and FIRST_CANON.get(a)==FIRST_CANON.get(b))
def matchname(nm):
    nm=nfc(nm.strip())
    nm=NICK.get(fold(nm),nm)
    if nm in cong: return nm
    if fold(nm) in byfold: return byfold[fold(nm)]
    ts=nm.split()
    if len(ts)>=2:
        cands=sur2mem.get(fold(ts[-1]),[])
        # Never infer identity from surname alone.  Quorum contains unrelated people and
        # candidates who share a legislator's surname (Rick/Susie Lee, Randy/Yvette Clarke,
        # Sherrod/Shontel Brown, Dan/Sanford Bishop, etc.).  Exact or explicit nickname-
        # equivalent first+last names are required; ambiguity is segregated for review.
        hits=[c for c in cands if same_first(c.split()[0],ts[0])]
        if len(hits)==1: return hits[0]
    return None
def classify(src):
    """-> (bucket, member_name_or_None, principal_member_or_None)"""
    s=nfc(str(src or '')); pre=s.split('|')[0].strip()
    # staffer: 'Name - Title...' with staff-title keywords
    if ' - ' in pre and STAFF_TITLES.search(pre.split(' - ',1)[1]):
        principal=None
        mm=re.search(r'\((?:Rep|Sen|Del)\.?\s+([^)]+)\)',s) or re.search(r'Office of (?:Rep|Sen)[a-z.]*\s+([^|)]+)',s)
        if mm: principal=matchname(re.sub(r'\((R|D|I)[-)][^)]*\)?','',mm.group(1)).strip())
        return ('staffer',None,principal)
    if re.match(r'^\s*Former\b',pre): return ('former_member',None,None)
    m=re.match(r'^\s*(?:Rep|Sen|Del|Resident Commissioner|Whip|Leader|Speaker)\.?\s+(.+)$',pre)
    if m:
        nm=re.sub(r'\((R|D|I)[^)]*\)\s*$','',m.group(1)).strip()
        hit=matchname(nm)
        return ('member',hit,None) if hit else ('member_unmatched',nm,None)
    if re.search(r'\bCandidate\b',pre,re.I):
        nm=re.sub(r'^.*\)\s*','',pre).strip() or pre.split()[-2]+' '+pre.split()[-1]
        hit=matchname(re.sub(r'\((R|D|I)[^)]*\)\s*$','',nm).strip())
        return ('member',hit,None) if hit else ('candidate_nonmember',None,None)   # current member running = main corpus
    if EXEC_PAT.search(s): return ('executive',None,None)
    if re.search(r'\bCommittee\b',s,re.I): return ('committee',None,None)
    # bare-name member forms (no honorific): try a full-name match before giving up
    nm=re.sub(r'\((R|D|I)[^)]*\)\s*$','',pre).strip()
    hit=matchname(nm) if re.match(r'^[A-Z][\w.\'-]+( [A-Z][\w.\'-]+)+$',nm) else None
    if hit: return ('member',hit,None)
    return ('organization',None,None)
df=pd.read_excel(M+r'\Quorum Master.xlsx',dtype=str)
df['dt']=pd.to_datetime(df['Date'],errors='coerce',format='mixed')
# --- step 2: TRUE duplicates (same source, same content, same date) — retweets across sources survive
def normc(t): return re.sub(r'\s+',' ',str(t or '')).strip().lower()
df['_k']=df['Source'].fillna('')+'\x1f'+df['Content'].map(normc)+'\x1f'+df['dt'].dt.date.astype(str)
before=len(df); df=df.drop_duplicates(subset='_k',keep='first').drop(columns='_k'); truedup=before-len(df)
# --- step 3: segregation
cl=df['Source'].map(classify)
df['bucket']=[c[0] for c in cl]; df['member']=[c[1] for c in cl]; df['principal_member']=[c[2] for c in cl]
# --- step 1: relevance
df['cu_relevant']=[1 if CU_TERMS.search((t or '')+' '+(c or '')) else 0
                   for t,c in zip(df['Title'].fillna(''),df['Content'].fillna(''))]
df['is_rt']=df['Content'].fillna('').str.match(r'\s*RT @').astype(int)
# --- step 4: stance on the member+relevant corpus
from lane_file_io import safe_to_unpickle
if not safe_to_unpickle(M+r'\data\stance\stance_clf_composed.pkl'): sys.exit('refusing to unpickle stance_clf_composed.pkl: it fails its integrity sidecar')
clf=pickle.load(open(M+r'\data\stance\stance_clf_composed.pkl','rb'))
mem=df[(df['bucket']=='member')]
score_ix=mem[mem['cu_relevant']==1].index
X=['target: Credit Unions (direct). speaker: member. '+str(df.loc[i,'Content'] or df.loc[i,'Title'] or '')[:1500] for i in score_ix]
df['stance']=''; df['stance_margin']=''
if len(X):
    pred=clf.predict(X)
    dec=clf.decision_function(X)
    import numpy as np
    marg=(np.sort(dec,axis=1)[:,-1]-np.sort(dec,axis=1)[:,-2]) if dec.ndim>1 else abs(dec)
    df.loc[score_ix,'stance']=pred
    df.loc[score_ix,'stance_margin']=['%.3f'%float(x) for x in marg]
# --- outputs
cols=['Quorum ID','Source','member','principal_member','bucket','dt','Document Type','Region','Title','Content',
      'Issues','cu_relevant','is_rt','stance','stance_margin','Source URL']
df['dt']=df['dt'].dt.strftime('%Y-%m-%d')
members=df[df['bucket']=='member'][cols]
staff=df[df['bucket']=='staffer'][cols]
excl=df[~df['bucket'].isin(['member','staffer'])][cols]
memtw=members[(members['cu_relevant']==1)]
memtw.to_csv(M+r'\data\stance\quorum_member_tweets.csv',index=False,encoding='utf-8-sig')
with pd.ExcelWriter(M+r'\Quorum Processed.xlsx',engine='openpyxl') as w:
    members.to_excel(w,index=False,sheet_name='Members')
    staff.to_excel(w,index=False,sheet_name='Staffers')
    excl.to_excel(w,index=False,sheet_name='Excluded')
    summ=pd.DataFrame([{'rows_in':before,'true_duplicates_removed':truedup,'rows':len(df)},])
    summ.to_excel(w,index=False,sheet_name='Summary')
from openpyxl import load_workbook
from openpyxl.styles import Font
wb=load_workbook(M+r'\Quorum Processed.xlsx')
for sn in wb.sheetnames:
    ws=wb[sn]
    for c in ws[1]: c.font=Font(name='Arial',bold=True)
    ws.freeze_panes='A2'
wb.save(M+r'\Quorum Processed.xlsx')
print('true duplicates removed (Source+Content+Date): %d  (retweets preserved: %d RTs in corpus)'%(truedup,int(df['is_rt'].sum())))
print('buckets:',dict(collections.Counter(df['bucket'])))
print('member rows: %d (distinct members %d) | CU-relevant member rows: %d | scored: %d'%(
    len(members),members['member'].nunique(),int(members['cu_relevant'].sum()),len(score_ix)))
print('staffers: %d rows, %d with principal member attributed'%(len(staff),staff['principal_member'].notna().sum()))
st=collections.Counter(df.loc[score_ix,'stance'])
print('stance distribution (member, CU-relevant):',dict(st))
um=df[df['bucket']=='member_unmatched']['Source'].value_counts()
print('member-shaped but UNMATCHED sources (review): %d rows, top: %s'%(len(df[df['bucket']=='member_unmatched']),um.head(4).to_dict()))
print('wrote Quorum Processed.xlsx + quorum_member_tweets.csv')
