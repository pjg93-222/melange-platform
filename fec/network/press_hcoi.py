# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT,DRAFT as MB_DRAFT  # centralized root (arch review 2026-07-02)
"""PRESS HCoI — Malaterre & Lareau (2024, Data & Policy) "hidden communities of interest" mapped onto
the CU policy domain.  Their pipeline: actor topic profiles (from an LDA over actor texts) -> Pearson
correlation between profiles -> threshold 0.6 -> Louvain -> communities with directly interpretable
topic profiles + diachronic genealogies (Hellinger distance between period community profiles).

Ours: members' press already carries a curated 19-category topical labeling (better than an unsupervised
LDA for this corpus), so the member "topic profile" = his/her distribution of press attention across the
19 categories (recency-weighted, multi-label rows split evenly).  The resulting HCoI network answers
who-talks-about-the-same-CU-issues-as-who — a relationship the cosponsor graph CANNOT see, because
LEADERSHIP DOESN'T COSPONSOR (finding D4: Hoyer, Clyburn, Scalise, Jordan, Jason Smith + freshmen are
isolated in the cosponsor layer but publish press constantly).  [Their case: philosophy-of-science
authors 1930-2017; ours: 385 press-active members 2025-2026, genealogy across the two half-periods.]
Writes press_hcoi.json: communities, profiles, correlation edges (candidate 2nd graph layer), genealogy."""
import csv,json,sys,math,collections,datetime,glob,os
from press_stream import load_tone_rows
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
import networkx as nx
B=MB_ROOT
SRC=MB_DRAFT+r'\cu_press_stance_tone.csv'
CORR_T=float(os.environ.get('HCOI_CORR','0.6'))     # their correlation threshold
MIN_W=float(os.environ.get('HCOI_MINW','3.0'))      # their transient-author filter (weighted pubs > 2)
rows=load_tone_rows()
def cats_of(r):
    c=(r.get('categories') or '').strip()
    return [p.strip() for p in c.replace('|',';').split(';') if p.strip()]
def pdate(r):
    d=(r.get('date') or '').strip()
    if not d or d.upper()=='NA': return None
    try: return datetime.date.fromisoformat(d[:10])
    except Exception: pass
    try: mm,dd,yy=d.split('/'); return datetime.date(int(yy),int(mm),int(dd))
    except Exception: return None
CATS=sorted({c for r in rows for c in cats_of(r)})
CIX={c:i for i,c in enumerate(CATS)}
def profiles(sub):
    """member -> (weight, prob-dist over CATS); multi-label split evenly, recency-weighted."""
    acc=collections.defaultdict(lambda:[0.0]*len(CATS)); wt=collections.Counter()
    for r in sub:
        cs=cats_of(r)
        if not cs: continue
        try: w=float(r.get('recency_weight') or 1.0)
        except ValueError: w=1.0
        for c in cs: acc[r['member']][CIX[c]]+=w/len(cs)
        wt[r['member']]+=w
    out={}
    for m,v in acc.items():
        s=sum(v)
        if wt[m]>=MIN_W and s>0: out[m]=(wt[m],[x/s for x in v])
    return out
def pearson(a,b):
    n=len(a); ma=sum(a)/n; mb=sum(b)/n
    num=sum((x-ma)*(y-mb) for x,y in zip(a,b))
    da=math.sqrt(sum((x-ma)**2 for x in a)); db=math.sqrt(sum((y-mb)**2 for y in b))
    return num/(da*db) if da>0 and db>0 else 0.0
def hcoi(sub,label):
    prof=profiles(sub); ms=sorted(prof)
    G=nx.Graph(); G.add_nodes_from(ms); edges=[]
    for i,m1 in enumerate(ms):
        for m2 in ms[i+1:]:
            r=pearson(prof[m1][1],prof[m2][1])
            if r>=CORR_T: G.add_edge(m1,m2,weight=r); edges.append((m1,m2,round(r,3)))
    comms=[sorted(c) for c in nx.community.louvain_communities(G,weight='weight',seed=7) if len(c)>=3]
    comms.sort(key=len,reverse=True)
    out=[]
    for k,c in enumerate(comms):
        tp=[0.0]*len(CATS); tw=0.0
        for m in c:
            w,pv=prof[m]; tw+=w
            for i,x in enumerate(pv): tp[i]+=w*x
        tp=[x/tw for x in tp]
        top=sorted(zip(CATS,tp),key=lambda t:-t[1])[:3]
        out.append({'id':'%s-%d'%(label,k+1),'label':' + '.join(t[0] for t in top[:2]),
                    'members':c,'n':len(c),'profile':{c2:round(v,4) for c2,v in zip(CATS,tp) if v>=0.01},
                    'top':[(t[0],round(t[1],3)) for t in top]})
    return prof,edges,out
def hellinger(p,q):
    return math.sqrt(0.5*sum((math.sqrt(p.get(c,0.0))-math.sqrt(q.get(c,0.0)))**2 for c in CATS))
# full-period HCoI
prof,edges,comms=hcoi(rows,'H')
print('PRESS HCoI (full period): %d members profiled (>=%.0f weighted items), %d correlation edges >= %.2f'%(len(prof),MIN_W,len(edges),CORR_T))
for c in comms:
    print('  %-5s n=%-3d %-52s top: %s'%(c['id'],c['n'],c['label'][:52],', '.join('%s %.2f'%t for t in c['top'])))
# D4 check: cosponsor-isolated members recovered by the HCoI layer
try:
    import openpyxl
    C={m['member']:m for m in json.load(open(B+r'\data\engine\congestion.json'))}
    bio2name={m['bioguide']:m['member'] for m in C.values() if m['bioguide']}
    net=max(glob.glob(B+r'\congress_cosponsor_network_*.xlsx'),key=os.path.getmtime)
    ne=openpyxl.load_workbook(net,data_only=True)['Edges']; deg=collections.Counter()
    for r in range(2,ne.max_row+1):
        a,b=bio2name.get(ne.cell(r,1).value),bio2name.get(ne.cell(r,2).value)
        if a in C and b in C and a!=b: deg[a]+=1; deg[b]+=1
    isolated=[m for m in C if deg[m]==0]
    incomm={m for c in comms for m in c['members']}
    rec=[m for m in isolated if m in incomm]
    print('D4 RECOVERY: %d/%d cosponsor-isolated members are placed in an HCoI community: %s%s'%(
        len(rec),len(isolated),', '.join(rec[:8]),'...' if len(rec)>8 else ''))
except Exception as e:
    isolated=[]; rec=[]; print('(D4 check skipped: %s)'%e)
# diachrony: split at median date -> genealogy (their Fig. 9)
dated=[(pdate(r),r) for r in rows]; dated=[(d,r) for d,r in dated if d]
dated.sort(key=lambda t:t[0]); mid=dated[len(dated)//2][0]
p1=[r for d,r in dated if d<mid]; p2=[r for d,r in dated if d>=mid]
_,_,c1=hcoi(p1,'P1'); _,_,c2=hcoi(p2,'P2')
gen=[]
print('\nGENEALOGY (split %s): period1 %d communities -> period2 %d'%(mid.isoformat(),len(c1),len(c2)))
for b in c2:
    best=min(((hellinger(dict(zip(CATS,[a['profile'].get(c,0) for c in CATS])),
                         dict(zip(CATS,[b['profile'].get(c,0) for c in CATS]))),a) for a in c1),
             default=(None,None))
    if best[1] is not None:
        d,a=best; gen.append({'from':a['id'],'from_label':a['label'],'to':b['id'],'to_label':b['label'],'hellinger':round(d,3)})
        tag='PERSISTS' if d<0.25 else ('drifts' if d<0.45 else 'NOVEL')
        print('  %-42s -> %-42s H=%.2f %s'%(a['label'][:42],b['label'][:42],d,tag))
json.dump({'threshold':CORR_T,'min_weight':MIN_W,'cats':CATS,
           'members_profiled':len(prof),'communities':comms,
           'edges':[{'a':a,'b':b,'r':r} for a,b,r in edges],
           'd4_isolated':isolated,'d4_recovered':rec,
           'periods':{'split':mid.isoformat(),'p1':c1,'p2':c2,'genealogy':gen}},
          open(B+r'\data\stance\press_hcoi.json','w'),indent=1)
print('\nwrote press_hcoi.json (%d communities, %d edges — candidate 2nd graph layer for the cascade)'%(len(comms),len(edges)))
