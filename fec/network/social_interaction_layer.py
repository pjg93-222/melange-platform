# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
"""SOCIAL INTERACTION LAYER — member<->member reposts & mentions as NETWORK data (user 2026-07-02:
"everything we designed, including the posts and re-posts, incorporated as both edge weights and
network linkages").

Events:
  * X/Quorum reposts: rows flagged is_rt whose text opens "RT @handle" where the handle maps to a
    CURRENT member (members.twitter in social_media.db). A repost is the strongest public social
    endorsement between offices (weight 1.0).
  * X/Quorum mentions: @handle tokens in ORIGINAL posts (RT rows excluded — the RT already counts
    once; its embedded mentions belong to the original author). Weight 0.4 (a mention is contact,
    not endorsement).
  * Bluesky mentions: interactions table (bioguide->bioguide), weight 0.4.
Normalization (Fink et al. 2023, same discipline as press_engagement_layer): a directed engagement
probability p(a->b) = weighted-events(a->b) / max(VOLFLOOR, posts(a)) — heavy posters aren't credited
for volume alone; VOLFLOOR=25 (TW_MIN_POSTS) stops one lucky RT from a near-silent account becoming
a strong tie. Pair probability = mean of both directions, capped 0.25 (a social tie is evidence of a
channel, not a whip relationship).
Output: social_interactions.json {'symmetric':[{a,b,w,rt,men}...]} consumed by build_multilayer_network."""
import sys,os,re,json,sqlite3,collections
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
import openpyxl
B=MB_ROOT
VOLFLOOR=int(os.environ.get('SOC_VOLFLOOR','25')); MEN_W=float(os.environ.get('SOC_MEN_W','0.4')); PCAP=0.25
C={m['member']:m for m in json.load(open(B+r'\data\engine\congestion.json'))}
bio2name={C[m]['bioguide']:m for m in C if C[m].get('bioguide')}
con=sqlite3.connect(B+r'\social_media.db')
h2m={}
for bio,tw in con.execute("SELECT bioguide,twitter FROM members WHERE twitter IS NOT NULL AND twitter!=''"):
    if bio in bio2name: h2m[tw.lower().lstrip('@')]=bio2name[bio]
print('handle map: %d X handles -> members'%len(h2m))
RTP=re.compile(r'^RT @(\w+)'); MEN=re.compile(r'@(\w+)')
ws=openpyxl.load_workbook(B+r'\Quorum Processed.xlsx',read_only=True)['Members']
hdr=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
im,ic,irt=hdr.index('member'),hdr.index('Content'),hdr.index('is_rt')
ev=collections.Counter(); rtc=collections.Counter(); menc=collections.Counter(); vol=collections.Counter()
for row in ws.iter_rows(min_row=2,values_only=True):
    a=row[im]; t=str(row[ic] or '')
    if a not in C: continue
    vol[a]+=1
    if str(row[irt]).lower() in ('1','true'):
        mt=RTP.match(t); b=h2m.get(mt.group(1).lower()) if mt else None
        if b and b!=a: ev[(a,b)]+=1.0; rtc[(a,b)]+=1
    else:
        for h in MEN.findall(t):
            b=h2m.get(h.lower())
            if b and b!=a: ev[(a,b)]+=MEN_W; menc[(a,b)]+=1
nbsky=0
for src,dst in con.execute("SELECT src_bioguide,dst_bioguide FROM interactions WHERE kind='mention'"):
    a,b=bio2name.get(src),bio2name.get(dst)
    if a and b and a!=b: ev[(a,b)]+=MEN_W; menc[(a,b)]+=1; nbsky+=1
pairs={}
for (a,b),s in ev.items():
    k=(min(a,b),max(a,b))
    pab=min(PCAP,s/max(VOLFLOOR,vol[a]))
    pairs.setdefault(k,[0.0,0.0])
    pairs[k][0 if (a,b)==k else 1]=pab
sym=[]
for (a,b),(p1,p2) in pairs.items():
    p=(p1+p2)/2.0
    if p>=0.002:
        sym.append({'a':a,'b':b,'w':round(p,4),
                    'rt':rtc[(a,b)]+rtc[(b,a)],'men':menc[(a,b)]+menc[(b,a)]})
sym.sort(key=lambda e:-e['w'])
from moneyball_core import dump_atomic   # atomic write: this runs UNATTENDED in weekly_refresh
dump_atomic({'symmetric':sym,'volfloor':VOLFLOOR,'mention_weight':MEN_W,
           'provenance':'X reposts(1.0)+mentions(%.1f)+bluesky mentions; Fink author-volume normalization; pair=mean dir, cap %.2f'%(MEN_W,PCAP)},
          B+r'\data\lobbying\social_interactions.json',indent=1)
print('social interaction pairs: %d (from %d RT events, %d mention events incl. %d bluesky)'%(
    len(sym),sum(rtc.values()),sum(menc.values()),nbsky))
print('top ties:',', '.join('%s<->%s %.3f (rt%d)'%(e['a'].split()[-1],e['b'].split()[-1],e['w'],e['rt']) for e in sym[:6]))
