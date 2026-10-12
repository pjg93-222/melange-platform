# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
"""SOCIAL MASTER — the bow.  Unifies every social channel into one scored corpus + per-member rollup:
  * Quorum (licensed): X tweets, Facebook posts, Truth Social posts — from quorum_member_tweets.csv
    (already member-attributed, CU-relevant, stance-scored by process_quorum.py).
  * Direct Bluesky collection: social_media.db posts, member-attributed via the handle registry.
    SAFETY: posts fetched from pattern-matched UNVERIFIED handles are included but flagged
    (verified_handle=0) and EXCLUDED from the member rollup until the user confirms the handle —
    a look-alike @bradsherman.bsky.social must not speak for Rep. Sherman.
Outputs: 'Social Master.xlsx' (Corpus + MemberRollup + Coverage sheets) + social_member_summary.json
(the per-member social stance aggregate: a NEW provenance-tagged evidence channel; pillar adoption
remains the user's decision)."""
import json,sys,sqlite3,re,collections,pickle
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
import pandas as pd
M=MB_ROOT
sys.path.insert(0,M+r'\fec')
from network.process_quorum import CU_TERMS   # single relevance vocabulary, kept in sync with quorum_query.txt
from collect.tweet_collector import bluesky_post_url, verified_bluesky_handle
DOCMAP={'Tweet':'x','Facebook Post':'facebook','Truth Social Post':'truthsocial'}
q=pd.read_csv(M+r'\data\stance\quorum_member_tweets.csv',dtype=str)
q['platform']=q['Document Type'].map(DOCMAP).fillna('other')
qc=pd.DataFrame({'platform':q['platform'],'member':q['member'],'date':q['dt'],'text':q['Content'].fillna(q['Title']),
                 'is_rt':q['is_rt'],'stance':q['stance'],'stance_margin':q['stance_margin'],
                 'source':'quorum','verified_handle':1,'url':q['Source URL'],
                 'post_id':'','repost_of':'','original_post_uri':''})
# direct bluesky
con=sqlite3.connect(M+r'\social_media.db')
handles={r[0]:(r[1] or '') for r in con.execute('SELECT bioguide,bluesky FROM members')}
bio2name={m['bioguide']:m['member'] for m in json.load(open(M+r'\data\engine\congestion.json')) if m.get('bioguide')}
# backfill member-member mention edges from post TEXT (facet features carry DIDs, not handles)
h2bg={h.lower():bg for bg,h in handles.items() if verified_bluesky_handle(h)}
n_ment=0
for bg,pid,dt,tx in con.execute("SELECT bioguide,post_id,created_at,text FROM posts WHERE platform='bluesky' AND COALESCE(repost_of,'')=''"):
    if not verified_bluesky_handle(handles.get(bg)): continue
    for hm in re.findall(r'@([a-z0-9.-]+)',(tx or '').lower()):
        tgt=h2bg.get(hm)
        if tgt and tgt!=bg:
            con.execute('INSERT OR IGNORE INTO interactions VALUES(?,?,?,?,?,?)',('bluesky','mention',bg,tgt,pid,dt)); n_ment+=1
con.commit()
print('mention backfill: %d member-member mention events extracted from post text'%n_ment)
rows=con.execute("SELECT bioguide,created_at,text,repost_of,post_id FROM posts WHERE platform='bluesky'").fetchall()
from lane_file_io import safe_to_unpickle
if not safe_to_unpickle(M+r'\data\stance\stance_clf_composed.pkl'): sys.exit('refusing to unpickle stance_clf_composed.pkl: it fails its integrity sidecar')
clf=pickle.load(open(M+r'\data\stance\stance_clf_composed.pkl','rb'))
brec=[]
for bg,dt,tx,rp,pid in rows:
    nm=bio2name.get(bg)
    if not nm: continue
    rel=1 if CU_TERMS.search(tx or '') else 0
    brec.append({'platform':'bluesky','member':nm,'date':(dt or '')[:10],'text':tx or '','is_rt':1 if rp else 0,
                 'rel':rel,'source':'direct','verified_handle':int(verified_bluesky_handle(handles.get(bg))),
                 'url':bluesky_post_url(rp or pid),'post_id':pid,'repost_of':rp or '',
                 'original_post_uri':rp or pid})
b=pd.DataFrame(brec)
if len(b):
    ri=b[b['rel']==1].index
    if len(ri):
        pred=clf.predict(['target: Credit Unions (direct). speaker: member. '+t[:1500] for t in b.loc[ri,'text']])
        import numpy as np
        dec=clf.decision_function(['target: Credit Unions (direct). speaker: member. '+t[:1500] for t in b.loc[ri,'text']])
        marg=(np.sort(dec,axis=1)[:,-1]-np.sort(dec,axis=1)[:,-2]) if getattr(dec,'ndim',1)>1 else abs(dec)
        b['stance']='';b['stance_margin']=''
        b.loc[ri,'stance']=pred; b.loc[ri,'stance_margin']=['%.3f'%float(x) for x in marg]
    bc=b[b['rel']==1][list(qc.columns)]
else:
    bc=pd.DataFrame(columns=qc.columns)
corpus=pd.concat([qc,bc],ignore_index=True)
# per-member rollup (verified handles + quorum only)
ok=corpus[(corpus['verified_handle']==1)&(corpus['stance']!='')]
roll=[]
for mname,g in ok.groupby('member'):
    n=len(g); s=(g['stance']=='support').sum(); o=(g['stance']=='oppose').sum()
    roll.append({'member':mname,'n_posts':int(n),'support':int(s),'oppose':int(o),
                 'neutral':int(n-s-o),'net_share':round((s-o)/n,3),
                 'platforms':','.join(sorted(g['platform'].unique()))})
roll.sort(key=lambda r:-r['n_posts'])
cov=corpus.groupby(['platform','source']).size().reset_index(name='posts')
with pd.ExcelWriter(M+r'\Social Master.xlsx',engine='openpyxl') as w:
    corpus.to_excel(w,index=False,sheet_name='Corpus')
    pd.DataFrame(roll).to_excel(w,index=False,sheet_name='MemberRollup')
    cov.to_excel(w,index=False,sheet_name='Coverage')
from openpyxl import load_workbook
from openpyxl.styles import Font
wb=load_workbook(M+r'\Social Master.xlsx')
for sn in wb.sheetnames:
    ws=wb[sn]
    for c in ws[1]: c.font=Font(name='Arial',bold=True)
    ws.freeze_panes='A2'
wb.save(M+r'\Social Master.xlsx')
json.dump({'rollup':roll,'coverage':cov.to_dict('records'),
           'provenance':'social channel — Quorum licensed + direct Bluesky (verified handles only in rollup); adoption into pillars = user decision'},
          open(M+r'\data\members\social_member_summary.json','w'),indent=1)
print('SOCIAL MASTER: %d posts (%s)'%(len(corpus),dict(collections.Counter(corpus['platform']))))
print('  member rollup: %d members with scored posts | bluesky rows: %d (%d from UNVERIFIED handles, excluded from rollup)'%(
    len(roll),len(bc),int((bc['verified_handle']==0).sum()) if len(bc) else 0))
print('wrote Social Master.xlsx + social_member_summary.json')
