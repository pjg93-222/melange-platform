# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
"""Rebuild corrected per-member PRESS and LEGISLATIVE leans from ONE consistent, direction-aware source —
the Press Release Articles sheet (real articles + the billstatus bill-cosponsorship entries, each carrying a
CU-direction = sign of its score). Fixes: (A) press normalization/mislabels, (B) per-bill legislative direction
(Interchange/Credit-Card/CFPB bills signed by the bill, not assumed anti), (C) consistency (one source, not the
disjoint Legislative-Score sheet). Writes corrected_signals.json {name:{press,leg,stance}}. Inspection mode prints
distributions + key members; --write also writes the JSON."""
import openpyxl, sys, io, json, collections, statistics
import roster
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
WB = MB_ROOT+r'\Moneyball Main Version - Blended.xlsx'
def fs(x):
    try: x = float(x)
    except (TypeError, ValueError): return 0
    return 1 if x > 0 else (-1 if x < 0 else 0)
# canonical area normalization (article/bill labels -> 17 canonical) + area importance (|legislative block weight|)
LBL2CAT = {'CFPB Issues & Structural Reform':'CFPB','Charter Modernization & Governance':'Charter Modernization',
 'Credit Union Capital Requirements':'CU Capital Requirements','Credit Card & Consumer-Fee Mandates':'Credit Card Mandates',
 'Data Privacy & Cybersecurity':'Data Privacy & Cyber','Defense Credit Union Issues':'Defense CU Issues',
 'Mortgage Lending & Housing':'Mortgage & Housing','NCUA Budget & Share Insurance Fund':'NCUA Budget & Insurance',
 'Technology & Emerging Issues':'Technology & Emerging','Telephone Consumer Protection Act':'TCPA',
 'Telephone Consumer Protection Act Issues':'TCPA'}
from moneyball_core import IMP as AREA_IMP
def norm(c):
    c=(c or '').strip(); return LBL2CAT.get(c,c)

wb = openpyxl.load_workbook(WB, read_only=True, data_only=True)
ra = wb['Rank & Allocation']; L={r:ra.cell(r,12).value for r in range(3,10)}
members={}; oldleg={}
for r in roster.rows('Rank & Allocation'):
    n=ra.cell(r,1).value
    if n: members[n]=ra.cell(r,7).value; oldleg[n]=ra.cell(r,4).value or 0
ws=wb['Press Release Articles']; H={h:i for i,h in enumerate(next(ws.iter_rows(min_row=1,max_row=1,values_only=True)))}
press=collections.defaultdict(lambda:[0.0,0.0]); leg=collections.defaultdict(float)
for r in ws.iter_rows(min_row=2,values_only=True):
    nm=r[H['member']]
    if nm not in members: continue
    src=r[H['stance_src']]; sc=r[H['score']]; fsign=fs(sc) if src=='bill' else fs(r[H['stance_sign']])
    rw=r[H['recency_weight']] or 1
    is_bill=str(r[H['url']] or '').startswith('billstatus:')
    if is_bill:   # legislative: bill cosponsorship, signed by the bill's CU-direction (sign of score)
        for c in str(r[H['categories']] or '').split(';'):
            a=norm(c)
            if a in AREA_IMP: leg[nm]+= fsign*AREA_IMP[a]*rw
    else:         # press: corrected per-article sentiment, recency-weighted mean
        mag=abs(r[H['score_graded']] or 0); grd=0.0 if fsign==0 else fsign*(mag if mag>0 else 0.2)
        press[nm][0]+=grd*rw; press[nm][1]+=rw
wb.close()

P={n:(press[n][0]/press[n][1] if press[n][1]>0 else 0.0) for n in members}      # corrected press lean
Lg_raw={n:leg.get(n,0.0) for n in members}
# calibrate legislative to ~ the old leg scale via a robust divisor (90th pct |value|) so dials L3/L7 still apply
sc=sorted(abs(v) for v in Lg_raw.values() if v); div=sc[int(len(sc)*0.90)] if sc else 1.0
Lg={n: round(Lg_raw[n]/div*0.4,4) for n in members}   # 0.4 ~ old leg p90 magnitude
L3,L4,L5,L6,L7,L8,L9=(L[i] for i in range(3,10))
def stance(lg,pr):
    if lg<=L3 or pr<=L4: return 'Opponent'
    if lg<=L5 or pr<=L6: return 'Skeptic'
    if lg>=L7: return 'Champion'
    if lg>=L8: return 'Ally'
    if lg>L9: return 'Neutral'
    return 'Unengaged'
new={n:stance(Lg[n],P[n]) for n in members}

print('press lean: min %.2f med %.2f max %.2f | leg(calib): min %.2f med %.2f max %.2f'%(
    min(P.values()),statistics.median(P.values()),max(P.values()),min(Lg.values()),statistics.median(Lg.values()),max(Lg.values())))
co=collections.Counter(members[n] for n in members); cn=collections.Counter(new.values())
print('\nTier        OLD(workbook) -> NEW(rebuilt)')
for t in ['Champion','Ally','Neutral','Skeptic','Opponent','Unengaged']:
    print('  %-10s %4d -> %4d'%(t,co.get(t,0),cn.get(t,0)))
print('\nkey members (old -> new | press, leg_calib, leg_old):')
for n in ['Bill Foster','Thom Tillis','Sean Casten','Tina Smith','Roger Wicker','Bernie Moreno','Richard Hudson','Jacky Rosen','John Garamendi','Brad Sherman']:
    if n in members: print('  %-18s %-9s -> %-9s | press %+.2f leg %+.2f (old %+.2f)'%(n,members[n],new[n],P[n],Lg[n],oldleg[n]))
if '--write' in sys.argv:
    json.dump({n:{'press':round(P[n],4),'leg':Lg[n],'stance':new[n]} for n in members},
              io.open(MB_ROOT+r'\corrected_signals.json','w',encoding='utf-8'))
    print('\nwrote corrected_signals.json')
