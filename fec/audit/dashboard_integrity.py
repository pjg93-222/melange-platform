# -*- coding: utf-8 -*-
"""Fail-closed identity/category checks for generated dashboard artifacts."""
import hashlib, json, os, re
import openpyxl
from paths import ROOT, WORKBOOK, p
import moneyball_core as MC
import roster as roster_rows   # not `roster`: the checkers below bind a local `roster` list
import stance.stance_rules as SR
import district.demographic_potential as DP
from engine.pillar_policy import PILLAR_WEIGHTS


class IntegrityError(RuntimeError):
    pass


OPTIONAL_DISTRICT_SOURCES = (
    'data/lobbying/district_industry.json',
    'data/lobbying/district_industry_4digit.json',
    'data/district/district_community.json',
    'data/district/district_employers.json',
    'data/district/district_financial.json',
    'data/district/district_hmda.json',
)


def _sha(obj):
    raw=json.dumps(obj,sort_keys=True,separators=(',',':'),ensure_ascii=False).encode('utf-8')
    return hashlib.sha256(raw).hexdigest()


def _file_sha(path):
    """Hash file contents, not mtimes, so OneDrive/Excel touches are not data changes."""
    if not os.path.exists(path):
        return None
    h=hashlib.sha256()
    with open(path,'rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''):
            h.update(chunk)
    return h.hexdigest()


def _check(condition, message, errors):
    if not condition: errors.append(message)


def preflight(member_names, category_stance, dashboard_data, archive_leg, analytic_inputs_only=False):
    """Validate the identity and category joins before any desk artifact is written.

    analytic_inputs_only (make_member_desk --analytic-inputs-only): that stage runs before dashboard_data.json is rebuilt for the
    current roster, so dashboard_data roster drift is a warning there and a hard error everywhere else."""
    errors=[]; warnings=[]; roster=MC.members(); rset=set(roster); cats=category_stance.get('cats') or []
    _check(len(roster)==len(rset), 'canonical roster contains duplicate names', errors)
    _check(len(member_names)==len(set(member_names)), 'member desk source contains duplicate names', errors)
    _check(set(member_names)==rset, 'member desk source roster differs from canonical roster', errors)
    snames=dashboard_data.get('name') or []
    _check(len(snames)==len(set(snames)), 'dashboard_data contains duplicate names', errors)
    if analytic_inputs_only and set(snames)!=rset:
        warnings.append('dashboard_data roster differs from canonical roster (analytic-inputs stage; rebuilt downstream): '
                        'missing=%d extras=%d'%(len(rset-set(snames)),len(set(snames)-rset)))
    else:
        _check(set(snames)==rset, 'dashboard_data roster differs from canonical roster', errors)
    _check(set(category_stance.get('matrix') or {})==rset, 'stance matrix roster differs from canonical roster', errors)
    _check(set(category_stance.get('scores') or {})==rset, 'stance score roster differs from canonical roster', errors)
    _check(set(archive_leg).issubset(rset), 'legislation archive contains unknown/current-roster-mismatched names', errors)
    for name in roster:
        _check(set((category_stance.get('matrix') or {}).get(name,{}) or {})==set(cats),
               'stance categories missing/misaligned for '+name, errors)
        _check(set((category_stance.get('scores') or {}).get(name,{}) or {})==set(cats),
               'stance score categories missing/misaligned for '+name, errors)
    issues=dashboard_data.get('issues') or []
    stance_rows=dashboard_data.get('stance') or []
    _check(len(stance_rows)==len(snames), 'dashboard stance row count does not match dashboard names', errors)
    _check(all(len(s)==len(issues) for s in stance_rows), 'dashboard stance strings do not align to issue count', errors)

    wb=openpyxl.load_workbook(WORKBOOK,read_only=True,data_only=True)
    ra=wb['Rank & Allocation']; wn=[ra.cell(r,1).value for r in roster_rows.rows('Rank & Allocation') if ra.cell(r,1).value]
    # Press Raw: the member rows (roster_rows.rows) are followed by the literal "Weights" footer (roster_rows.footer_row).
    rw=wb['Press Release - Raw']; mn=[rw.cell(r,1).value for r in roster_rows.rows('Press Release - Raw') if rw.cell(r,1).value]
    wb.close()
    _check(len(wn)==len(set(wn)) and set(wn)==rset, 'Rank & Allocation roster is duplicated or misaligned', errors)
    _check(len(mn)==len(set(mn)) and set(mn)==rset, 'Press Release - Raw roster is duplicated or misaligned', errors)

    # General CU must be conserved exactly when distributed.
    alloc=((category_stance.get('media') or {}).get('general_credit_union') or {})
    bad_conservation=[]
    for name,a in alloc.items():
        if name not in rset or abs(float(a.get('general',0))-sum(float(x) for x in (a.get('to') or {}).values()))>1e-8:
            bad_conservation.append(name)
    _check(not bad_conservation, 'General CU allocation is not identity-safe/conserved: '+', '.join(bad_conservation[:5]), errors)
    if not alloc: warnings.append('no nonzero General Credit Union values were available to disperse')

    report={'status':'pass' if not errors else 'fail','errors':errors,'warnings':warnings,
            'counts':{'roster':len(roster),'categories':len(cats),'general_cu_members':len(alloc)},
            'roster_order_sha256':_sha(roster),'category_order_sha256':_sha(cats)}
    if errors: raise IntegrityError('; '.join(errors[:8]))
    return report


# --- DEMO_POTENTIAL-aware demographic-potential contract (D-034) -------------
# The producer (fec/district/demographic_potential.py) reads DEMO_POTENTIAL ONCE
# at import and derives SCORE_WEIGHTS / SCORE_VERSION / the qualification gate
# from it.  make_member_desk and this module share that single module object, so
# reading DP.<attr> at CALL time -- never caching the mode in a constant here --
# is what guarantees the postflight asserts the shape the bundle was actually
# built with.  audit_full.py was made dial-aware the same way.
FOOTPRINT_UNQUALIFIED_STUB = {
    'score': None,
    'rank': None,
    'market_count': None,
    'percentile': None,
    'tier': None,
    'components': {},
    'raw': {},
    'available': {},
    'evidence_coverage': 0.0,
    'signals': [],
    'sensitivity': {},
}


def _tier_label(percentile):
    """Mirror demographic_potential._tier: first threshold at or below wins."""
    for threshold,label in DP.TIER_THRESHOLDS:
        if percentile>=threshold: return label
    return DP.TIER_THRESHOLDS[-1][1]


def _tier_labels_for(percentile):
    """Every tier label admissible for a percentile the producer rounded to 3dp.

    score_members emits round(percentile,3) but computes the tier from the
    UNROUNDED value, so the true percentile lies in [p-0.0005, p+0.0005].  That
    band is 0.001 wide and the cuts are 0.90/0.70/0.40, so it spans at most one
    cut and the two endpoints enumerate the admissible labels.  Asserting a
    single label would manufacture a failure for a member sitting exactly on a
    cut -- a false alarm 3831 seconds into the surfaces step.
    """
    return {_tier_label(percentile-0.0005),_tier_label(percentile+0.0005)}


def _footprint_stub_ok(potential):
    """An unqualified footprint sticker must be exactly the producer's stub.

    score_members resets every scored field on EVERY member before the
    qualification branch, so an unqualified member carries the stub plus the
    qualification verdict (label/signals/reason) that evaluate() returned.
    """
    if any(potential.get(key)!=value
           for key,value in FOOTPRINT_UNQUALIFIED_STUB.items()): return False
    structural=potential.get('structural_adjustment') or {}
    reason=potential.get('reason')
    return (structural.get('available') is False and
            structural.get('version')==DP.STRUCTURAL_ADJUSTMENT_VERSION and
            potential.get('score_version')==DP.SCORE_VERSION_FOOTPRINT and
            isinstance(reason,str) and bool(reason.strip()))


def check_demographic_potential(bundle, members, errors):
    """Validate demographic-potential metadata and stickers for the ACTIVE mode.

    growth_room keeps every assertion this module has always made.  Under
    footprint (ruling R7, the default since design-d2-d5.md D3) the producer:

      * scores four CU-presence components and emits NO sector ``signals`` at
        all -- _evaluate_footprint returns ``'signals': []`` unconditionally;
      * runs the sensitivity audit over its own four weights, so scenario_count
        is 1 equal-weight + 4 leave-one-out = 5, not growth_room's 1+8 = 9;
      * returns the reach-residual structural companion as
        ``not_applicable_for_mode`` with adjusted_markets 0, because neither
        large_employer_reach nor small_business_reach is a scored component.

    Those three growth-room-shaped assertions -- plus the CBP large-employer
    availability rule, which cannot arise when the component is not scored --
    are REPLACED by a footprint contract, not merely skipped.

    Returns ``(qualified_member_count, markets_by_(kind, geo))``.
    """
    footprint=getattr(DP,'DEMO_POTENTIAL','growth_room')=='footprint'
    potential_meta=bundle.get('demographic_potential_meta') or {}
    potential_codes={row.get('sector') for row in potential_meta.get('eligible_sectors') or []
                     if isinstance(row,dict)}
    potential_lq_min=float(potential_meta.get('lq_min',1.25))
    potential_count=0
    potential_markets={}
    score_weights=potential_meta.get('score_weights') or {}
    _check(score_weights==DP.SCORE_WEIGHTS,
           'demographic-potential weights differ from the pinned formula',
           errors)
    _check(abs(sum(float(value) for value in score_weights.values())-1.0)<1e-9,
           'demographic-potential weights do not sum to 100%',errors)
    _check(potential_meta.get('score_version')==DP.SCORE_VERSION,
           'demographic-potential score version is not current',errors)
    _check(potential_meta.get('weight_basis')==DP.WEIGHT_BASIS and
           potential_meta.get('not_empirically_optimized') is True,
           'demographic-potential weight provenance is missing or changed',
           errors)
    component_definitions=potential_meta.get('component_definitions') or []
    _check([row.get('key') for row in component_definitions
            if isinstance(row,dict)]==list(DP.SCORE_WEIGHTS),
           'demographic-potential component formula catalog is incomplete',
           errors)
    source_catalog=potential_meta.get('source_catalog') or {}
    _check(source_catalog==DP.SOURCE_CATALOG and all(
        str((row or {}).get('official_url') or '').startswith('https://')
        for row in source_catalog.values()),
        'demographic-potential official source catalog is incomplete',errors)
    sensitivity_audit=potential_meta.get('sensitivity_audit') or {}
    # _sensitivity_audit builds one equal-weight scenario plus one leave-one-out
    # scenario per SCORED component, so the historical literal 9 is growth_room's
    # 1+8.  Footprint scores four components and legitimately reports 5.
    expected_scenarios=(1+len(DP.SCORE_WEIGHTS)) if footprint else 9
    _check(sensitivity_audit.get('version')==DP.SENSITIVITY_VERSION and
           sensitivity_audit.get('scenario_count')==expected_scenarios,
           'demographic-potential sensitivity audit is missing',errors)
    if footprint:
        _check(len(sensitivity_audit.get('leave_one_out') or [])==
               len(DP.SCORE_WEIGHTS),
               'demographic-potential sensitivity audit does not omit every '
               'footprint component',errors)
        _check(set(score_weights)==set(DP.SCORE_WEIGHTS_FOOTPRINT) and
               potential_meta.get('score_version')==DP.SCORE_VERSION_FOOTPRINT and
               potential_meta.get('footprint_employment_lq_min')==
               DP.FOOTPRINT_EMPLOYMENT_LQ_MIN,
               'demographic-potential metadata is not the footprint contract',
               errors)
    community_meta=bundle.get('community_meta') or {}
    _check(
        community_meta.get('schema')=='district-community-acs-v1' and
        community_meta.get('vintage')==2024 and
        community_meta.get('survey')=='ACS 1-year',
        'ACS community metadata is missing or not the pinned 2024 vintage',
        errors,
    )
    _check(
        set(community_meta.get('profile_families') or {})==
        {'DP03','DP04','DP05'} and
        all(str(url).startswith('https://')
            for url in (community_meta.get('official_urls') or {}).values()),
        'ACS community table/source catalog is incomplete',
        errors,
    )
    structural_meta=potential_meta.get('structural_adjustment') or {}
    if footprint:
        _check(
            structural_meta.get('version')==
            DP.STRUCTURAL_ADJUSTMENT_VERSION and
            structural_meta.get('status')=='not_applicable_for_mode' and
            structural_meta.get('adjusted_markets')==0,
            'demographic-potential structural adjustment is not the '
            'not-applicable-for-mode companion footprint requires',
            errors,
        )
    else:
        _check(
            structural_meta.get('version')==
            DP.STRUCTURAL_ADJUSTMENT_VERSION and
            structural_meta.get('status')==
            'descriptive_companion_not_causal',
            'demographic-potential structural adjustment is missing',
            errors,
        )
    structurally_adjusted_markets={}
    for m in members:
        potential=m.get('demographic_potential') or {}
        if not potential.get('qualified'):
            _check(potential.get('score') is None,
                   'unqualified member has a demographic-potential score for '+
                   str(m.get('name')),errors)
            if footprint:
                _check(_footprint_stub_ok(potential),
                       'unqualified member does not carry the documented '
                       'footprint stub for '+str(m.get('name')),errors)
            continue
        potential_count+=1
        industry=m.get('industry') or {}
        market=(industry.get('kind'),industry.get('geo'))
        prior=potential_markets.get(market)
        if prior is not None:
            _check(
                (prior.get('score'),prior.get('rank'))==
                (potential.get('score'),potential.get('rank')),
                'members in the same demographic market have different '
                'scores/ranks for '+str(market),errors)
            _check(
                (
                    (prior.get('structural_adjustment') or {}).get('score'),
                    (prior.get('structural_adjustment') or {}).get('rank'),
                ) == (
                    (potential.get('structural_adjustment') or {}).get(
                        'score'),
                    (potential.get('structural_adjustment') or {}).get(
                        'rank'),
                ),
                'members in the same demographic market have different '
                'structurally adjusted scores/ranks for '+str(market),
                errors,
            )
        else:
            potential_markets[market]=potential
        if footprint:
            # No sector concept exists under footprint; a populated list would
            # mean a growth-room sticker leaked into a footprint bundle.
            _check(potential.get('signals')==[],
                   'footprint demographic potential carries sector signals for '+
                   str(m.get('name')),errors)
        else:
            signals=potential.get('signals') or []
            _check(bool(signals),'qualified demographic potential has no signals for '+str(m.get('name')),errors)
            for signal in signals:
                code=signal.get('sector')
                lq=signal.get('lq')
                shift=signal.get('competitive_shift')
                _check(code in potential_codes,
                       'unknown demographic-potential sector for '+str(m.get('name')),errors)
                _check(isinstance(lq,(int,float)) and not isinstance(lq,bool) and lq>=potential_lq_min,
                       'demographic-potential LQ below threshold for '+str(m.get('name')),errors)
                _check(isinstance(shift,(int,float)) and not isinstance(shift,bool) and shift>0,
                       'demographic-potential shift is not positive for '+str(m.get('name')),errors)
                _check(signal.get('crosscheck') in ('passed','not_covered'),
                       'invalid demographic-potential cross-check for '+str(m.get('name')),errors)
        score=potential.get('score')
        components=potential.get('components') or {}
        available=potential.get('available') or {}
        coverage=potential.get('evidence_coverage')
        sensitivity=potential.get('sensitivity') or {}
        _check(isinstance(score,(int,float)) and not isinstance(score,bool) and 0<=score<=100,
               'invalid demographic-potential score for '+str(m.get('name')),errors)
        _check(set(components)==set(score_weights),
               'demographic-potential components do not match weights for '+
               str(m.get('name')),errors)
        _check(set(available)==set(score_weights),
               'demographic-potential availability does not match weights for '+
               str(m.get('name')),errors)
        _check(all(isinstance(value,(int,float)) and not isinstance(value,bool)
                   and 0<=value<=1 for value in components.values()),
               'demographic-potential component is outside [0,1] for '+
               str(m.get('name')),errors)
        if set(components)==set(score_weights) and isinstance(score,(int,float)):
            # Both sides are rounded to 1dp exactly as _weighted_scores does; the
            # emitted components are themselves rounded to 6dp, so a raw 1e-6
            # tolerance against the unrounded product would false-fail.
            reconstructed=round(100*sum(
                float(score_weights[key])*float(components[key])
                for key in score_weights),1)
            _check(abs(reconstructed-float(score))<=0.01,
                   'demographic-potential score does not reconstruct for '+
                   str(m.get('name')),errors)
        _check(isinstance(coverage,(int,float)) and not isinstance(coverage,bool)
               and 0<=coverage<=1,
               'invalid demographic-potential evidence coverage for '+
               str(m.get('name')),errors)
        _check(isinstance(potential.get('rank'),int) and
               potential.get('rank')>=1,
               'invalid demographic-potential rank for '+str(m.get('name')),
               errors)
        _check(
            isinstance(sensitivity.get('rank_min'),int) and
            isinstance(sensitivity.get('rank_max'),int) and
            sensitivity['rank_min']<=potential.get('rank')<=
            sensitivity['rank_max'],
            'invalid demographic-potential sensitivity range for '+
            str(m.get('name')),errors)
        if footprint:
            _check('large_employer_reach' not in available and
                   'small_business_reach' not in available,
                   'footprint demographic potential scores a reach component '
                   'for '+str(m.get('name')),errors)
            _check(potential.get('score_version')==DP.SCORE_VERSION_FOOTPRINT,
                   'demographic-potential sticker is not the footprint score '
                   'version for '+str(m.get('name')),errors)
            percentile=potential.get('percentile')
            percentile_ok=(isinstance(percentile,(int,float)) and
                           not isinstance(percentile,bool) and
                           0<=percentile<=1)
            _check(percentile_ok,
                   'invalid demographic-potential percentile for '+
                   str(m.get('name')),errors)
            if percentile_ok:
                _check(potential.get('tier') in _tier_labels_for(percentile),
                       'demographic-potential tier does not follow the '
                       'percentile cuts for '+str(m.get('name')),errors)
        else:
            large=((m.get('employers') or {}).get('large') or {})
            if large.get('complete') is False:
                _check(
                    available.get('large_employer_reach') is False,
                    'incomplete CBP large-employer evidence was scored as '
                    'available for '+str(m.get('name')),errors)
        _check(potential.get('tier') in {
            label for _,label in DP.TIER_THRESHOLDS},
            'invalid demographic-potential tier for '+str(m.get('name')),
            errors)
        structural=potential.get('structural_adjustment') or {}
        if structural.get('available'):
            structurally_adjusted_markets.setdefault(market,structural)
            community=m.get('community') or {}
            _check(
                community.get('complete') is True and
                community.get('vintage')==2024,
                'structurally adjusted member lacks complete 2024 ACS '
                'context for '+str(m.get('name')),errors)
            _check(
                structural.get('version')==
                DP.STRUCTURAL_ADJUSTMENT_VERSION and
                isinstance(structural.get('score'),(int,float)) and
                0<=structural.get('score')<=100 and
                isinstance(structural.get('rank'),int) and
                structural.get('rank')>=1,
                'invalid structurally adjusted score for '+
                str(m.get('name')),errors)
            controls=structural.get('controls') or {}
            _check(
                set(controls)==set(DP.STRUCTURAL_NUMERIC_CONTROLS)|{
                    'region'},
                'structural control vector is incomplete for '+
                str(m.get('name')),errors)
            adjusted_components=dict(components)
            reach_details=structural.get('components') or {}
            _check(
                set(reach_details)==set(DP.STRUCTURAL_TARGETS),
                'structural reach-component audit is incomplete for '+
                str(m.get('name')),errors)
            for target in DP.STRUCTURAL_TARGETS:
                adjusted=(reach_details.get(target) or {}).get(
                    'adjusted_percentile')
                if isinstance(adjusted,(int,float)):
                    adjusted_components[target]=adjusted
            if (
                set(adjusted_components)==set(score_weights) and
                isinstance(structural.get('score'),(int,float))
            ):
                reconstructed=round(100*sum(
                    float(score_weights[key])*
                    float(adjusted_components[key])
                    for key in score_weights
                ),1)
                _check(
                    abs(reconstructed-float(structural['score']))<=0.01,
                    'structurally adjusted score does not reconstruct for '+
                    str(m.get('name')),errors)
    # Chip-vs-artifact reconciliation: the displayed sticker count, the market
    # count every sticker carries, and the metadata block must agree.  Mode
    # independent -- footprint changes what a sticker contains, not how many.
    _check(potential_count==int(potential_meta.get('qualified_members',-1)),
           'demographic-potential metadata count does not match member stickers',errors)
    _check(len(potential_markets)==int(
        potential_meta.get('qualified_markets',-1)),
        'demographic-potential metadata count does not match qualifying markets',
        errors)
    _check(all(potential.get('market_count')==len(potential_markets)
               for potential in potential_markets.values()),
           'demographic-potential member market counts are inconsistent',errors)
    _check(
        len(structurally_adjusted_markets)==int(
            structural_meta.get('adjusted_markets',-1)
        ),
        'structural-adjustment metadata count does not match markets',
        errors,
    )
    _check(all(
        structural.get('market_count')==len(
            structurally_adjusted_markets)
        for structural in structurally_adjusted_markets.values()
    ),'structurally adjusted market counts are inconsistent',errors)
    return potential_count,potential_markets


def district_brief_errors(members, meta, roster):
    """District-brief checks for the desk bundle: (errors, warnings, attached_count).

    The desk WITHHOLDS the brief of a roster arrival the ratified corpus predates (meta['coverage']['withheld']): that member must
    carry no brief, every other member a valid one, and the corpus accounts for exactly the roster less the withheld plus the
    departed rows it still holds (meta['coverage']['dropped']).  With nothing withheld or dropped this is the original full-roster check."""
    errors=[]; warnings=[]
    coverage=meta.get('coverage') or {}
    withheld=set(coverage.get('withheld') or []); dropped=set(coverage.get('dropped') or [])
    _check(withheld.issubset(set(roster)),
           'district briefs withheld for members outside the roster: '+', '.join(sorted(withheld-set(roster))),errors)
    attached=0
    for m in members:
        name=m.get('name'); brief=m.get('district_brief')
        if name in withheld:
            _check(brief is None,'district brief attached to withheld member '+str(name),errors)
            continue
        brief=brief or {}
        paragraphs=[brief.get('p%d'%index) for index in range(1,4)]
        brief_ok=(all(isinstance(value,str) and value.strip() for value in paragraphs)
                  and type(brief.get('words')) is int
                  and 200<=brief.get('words')<=250)
        _check(brief_ok,'district brief is missing or malformed for '+str(name),errors)
        if brief_ok: attached+=1
    counts=meta.get('counts') or {}
    corpus_rows=len(roster)-len(withheld)+len(dropped)
    _check(
        meta.get('schema')=='district-demographic-briefs-v1'
        and meta.get('status')=='RATIFIED_FOR_DASHBOARD'
        and counts.get('roster')==corpus_rows
        and counts.get('accepted')==corpus_rows
        and counts.get('missing')==0
        and attached==len(roster)-len(withheld),
        'ratified district brief corpus is absent or incomplete',errors)
    if withheld:
        warnings.append('%d district briefs WITHHELD for roster arrivals the corpus predates: %s'%(
            len(withheld),', '.join(sorted(withheld))))
    return errors,warnings,attached


def postflight(bundle, preflight_report):
    """Validate the completed desk bundle and enforce deterministic regeneration."""
    errors=[]; warnings=[]; roster=MC.members(); rset=set(roster); cats=bundle.get('cats') or []
    members=bundle.get('members') or []; names=[m.get('name') for m in members]
    _check(len(names)==len(set(names)) and set(names)==rset, 'desk bundle roster is duplicated or misaligned', errors)
    allowed=set(cats)|set(bundle.get('art_extra_cats') or [])
    expected_weights={'leg':PILLAR_WEIGHTS[2], 'press':PILLAR_WEIGHTS[3],
                      'comm':PILLAR_WEIGHTS[4], 'demo':PILLAR_WEIGHTS[5]}
    _check(all((bundle.get('weights') or {}).get(key)==value
               for key,value in expected_weights.items()),
           'desk bundle does not carry the ratified diminished-sentiment pillar weights',errors)
    n2b=MC.name2bio(); social_members=set(); article_total=bill_total=0; missing_bio=[]
    media_count_mismatches=[]; unknown_media_sources=[]
    for m in members:
        name=m.get('name'); article_total+=len(m.get('articles') or []); bill_total+=len(m.get('bills') or [])
        _check(len(m.get('stances') or '')==len(cats), 'desk stance string misaligned for '+str(name), errors)
        expected=n2b.get(name); actual=(m.get('bio') or {}).get('bioguide')
        if expected and actual: _check(actual==expected, 'bioguide mismatch for '+str(name), errors)
        elif expected and not actual: missing_bio.append(name)
        member_articles=m.get('articles') or []
        for a in member_articles:
            _check(set(a.get('cats') or []).issubset(allowed), 'unknown media category for '+str(name), errors)
            if a.get('source')=='social': social_members.add(name)
            if a.get('source') not in {'press','social','bluesky','interview','hearing'}:
                unknown_media_sources.append((name,a.get('source')))
        totals=m.get('media_totals') or {}
        for key,source in (('iv','interview'),('hr','hearing'),('bs','bluesky')):
            pair=totals.get(key)
            if pair is None: continue
            actual=sum(a.get('source')==source for a in member_articles)
            if pair != [actual,actual]: media_count_mismatches.append((name,key,pair,actual))
        if totals.get('iv_sources') and sum(totals['iv_sources'].values()) != sum(
                a.get('source')=='interview' for a in member_articles):
            media_count_mismatches.append((name,'iv_sources',totals.get('iv_sources'),None))
        for b in m.get('bills') or []:
            _check(set(b.get('areas') or []).issubset(allowed), 'unknown bill category for '+str(name), errors)
    _check(not unknown_media_sources,
           'desk bundle contains unaccounted media source types: '+str(unknown_media_sources[:5]),errors)
    _check(not media_count_mismatches,
           'desk media totals show producer-side omissions/caps: '+str(media_count_mismatches[:5]),errors)
    district_briefs_meta=bundle.get('district_briefs_meta') or {}
    brief_errors,brief_warnings,district_brief_count=district_brief_errors(members,district_briefs_meta,roster)
    errors.extend(brief_errors); warnings.extend(brief_warnings)
    _check(
        bool(re.fullmatch(r'[0-9a-f]{64}',str(
            district_briefs_meta.get('evidence_sha256') or ''))),
        'district brief evidence fingerprint is missing or malformed',errors)
    potential_count,potential_markets=check_demographic_potential(
        bundle,members,errors)
    source_display=((bundle.get('integrity') or {}).get('source_display') or {})
    # DESK_OUT preview runs (make_member_desk no-clobber redirect) must not overwrite the LIVE
    # integrity state/report/reconciliation artifacts — route this module's writes (and the
    # regeneration-state read) to the same directory.
    _outdir=os.environ.get('DESK_OUT')
    _op=(lambda rel: os.path.join(_outdir, rel)) if _outdir else p
    _check(source_display.get('social_displayed',0)>0, 'social source loaded zero visible items', errors)
    def _recon(source,parts):
        accounted=sum(int(source_display.get(k,0) or 0) for k in parts)
        return {'source_rows':int(source),'accounted_rows':accounted,
                'difference':int(source)-accounted,'components':{k:int(source_display.get(k,0) or 0) for k in parts},
                'status':'pass' if int(source)==accounted else 'fail'}
    reconciliation={
        'press':_recon(source_display.get('press_source_rows',0),(
            'press_displayed','press_unknown_member_excluded','press_billstatus_rows_routed_to_legislation')),
        'social':_recon(source_display.get('social_source_rows',0),(
            'social_displayed','social_unverified_handle_excluded','social_unknown_member_excluded',
            'social_unrelated_excluded','social_offtopic_excluded')),
        # 2026-08-29: the verified collector roster sidecar is the FOURTH declared legislation
        # source (verify_collector_roster_updates.py -> make_member_desk). Its source term counts
        # APPENDED LEDGER ROWS plus explicitly skipped lines, so multi-area rows reconcile exactly.
        'legislation':_recon((source_display.get('bill_source_rows',0) or 0)+
                              (source_display.get('bill_manual_rows',0) or 0)+
                              (source_display.get('bill_primacy_membership_rows',0) or 0)+
                              (source_display.get('bill_collector_roster_rows',0) or 0),(
            'bills_displayed_unique','bill_duplicate_rows_collapsed','bill_unknown_member_excluded',
            'bill_rules_excluded','bill_override_excluded','bill_manual_excluded',
            'bill_collector_roster_excluded',
            # 2026-08-29: an adjudicated per-member withdrawal removes PRE-COLLAPSE
            # ledger rows (data/legislation/membership_withdrawals.csv). Those rows were
            # already counted in a declared source term, so without this exclusion
            # component the legislation reconciliation would show a deficit exactly equal
            # to the rows removed. The term generalizes to multi-area and duplicate-role
            # memberships: removing the N pre-collapse rows of one membership drops
            # bills_displayed_unique by 1 and bill_duplicate_rows_collapsed by N-1, and
            # this +N restores the balance. It is 0 when no withdrawal file exists.
            'bill_withdrawn_rows_removed')),
        # media footprint (2026-07-16): every interview episode / hearing appearance / bluesky post is
        # displayed, summarized beyond the display cap, or carries one explicit exclusion reason.
        'interview_footprint':_recon(source_display.get('interview_episode_rows',0),(
            'interview_appearances_displayed','interview_beyond_display_cap',
            'interview_unknown_member_excluded')),
        # hearing footprint draws on TWO declared sources (2026-07-19): the hearing_antag appearance
        # rows, plus graded CU passages whose hearing has no antag row at all (hearing_antag records
        # the antagonism pass, not every speaker, so a committee chair can be graded while owning no
        # antag row). Both are displayed or accounted; DESK_MEDIA_STANCE=0 makes the second term 0.
        'hearing_footprint':_recon((source_display.get('hearing_appearance_rows',0) or 0)+
                                   (source_display.get('hearing_graded_rows_added',0) or 0),(
            'hearing_appearances_displayed','hearing_beyond_display_cap',
            'hearing_unknown_member_excluded')),
        'bluesky_footprint':_recon(source_display.get('bluesky_post_rows',0),(
            'bluesky_footprint_displayed','bluesky_beyond_display_cap',
            'bluesky_unverified_handle_posts_excluded','bluesky_unknown_member_excluded',
            'bluesky_empty_text_excluded','bluesky_scored_rows_routed_to_social')),
    }
    for source,row in reconciliation.items():
        _check(row['status']=='pass','%s source rows are not fully reconciled (%+d)'%(source,row['difference']),errors)
    _check(article_total==(source_display.get('press_displayed',0) or 0)+(source_display.get('social_displayed',0) or 0)
           +(source_display.get('interview_appearances_displayed',0) or 0)
           +(source_display.get('hearing_appearances_displayed',0) or 0)
           +(source_display.get('bluesky_footprint_displayed',0) or 0),
           'desk article total does not match reconciled press + social + footprint display rows',errors)
    _check(bill_total==(source_display.get('bills_displayed_unique',0) or 0),
           'desk bill total does not match reconciled unique legislation rows',errors)
    MC.dump_atomic({'status':'pass' if all(v['status']=='pass' for v in reconciliation.values()) else 'fail',
                    'rule':'Every source row must be displayed or assigned one explicit exclusion/routing reason.',
                    'sources':reconciliation,'exclusion_samples':source_display.get('exclusion_samples') or {}},
                   _op('dashboard_exclusion_reconciliation.json'),indent=2,ensure_ascii=False)
    if missing_bio: warnings.append('%d members lack desk bioguide enrichment (identity not guessed)'%len(missing_bio))

    # Strip self-referential validation and volatile freshness metadata before hashing the semantic
    # product.  ``generated_at`` and source mtimes are provenance/display metadata, not member data;
    # including them makes every same-input rebuild fail merely because it ran in a later second.
    projection={k:v for k,v in bundle.items() if k not in ('integrity','freshness')}
    output_hash=_sha(projection)
    # Fingerprint every source consumed by make_member_desk plus its routing code
    # and templates.  Content hashes prevent timestamp-only Excel/OneDrive touches
    # from defeating the same-input/same-output corruption check.
    source_files=[
        'Moneyball Main Version - Blended.xlsx','Social Master.xlsx','data/stance/tweet_stance_final.csv',
        'manual_bill_evidence.csv','data/legislation/bill_overrides.csv','archive_leg.json',
        # 2026-08-29: adjudicated per-member withdrawals are a desk INPUT that changes
        # the bills ledger. Registering it here is the same lesson the bill-lobbying
        # sidecar taught on 08-06 — an unregistered dependency makes the first
        # withdrawal-only change look like same-input/different-bundle corruption.
        'data/legislation/membership_withdrawals.csv',
        'bill_beneficiary_tags.ndjson','data/members/primacy_act_memberships.ndjson',
        'data/members/registry_anchor_memberships.ndjson',
        'data/legislation/registry_anchor_reconciliation.json',
        'data/engine/primacy_factors.json',
        'grievance_ledger.ndjson','delivery_ledger.ndjson',
        'tier_overrides.json','legislation_culling_receipt.json',
        'data/legislation/legislation_culling_status.json',
        'legislation_membership_edges.ndjson',
        'behavioral_dispositions.json','data/members/behavioral_roles_v2_candidate.json',
        'data/members/big5_blue_emb.json','data/members/big5_blue_validated.json','data/members/big5_llm_opus.json',
        'data/stance/big5_senate_expert_gold.json','data/stance/big5_external_gold_validation.json','data/stance/big5_e_gold_anchored.json',
        'data/stance/big5_gold_anchored.json',
        'data/members/big5_traits.json','data/members/big5_v2.json','category_stance.json','communication_style.json',
        'coop_v2_candidate.json','dashboard_data.json','data/money/electoral_history.json',
        'data/money/electoral_plays.json','data/stance/interview_antagonism.json','data/network/kb_benchmark.json',
        # COMPET_MULT (Paul 2026-10-05): the competitiveness multiplier artifact (promoted path and .candidate) shapes the desk ROI when the dial is ON;
        # hashed under either dial position, like the other declared inputs (an undeclared one that IS read breaks the same-input check).
        'data/money/competitiveness_c.json','data/money/competitiveness_c.json.candidate',
        'bill_repository_live.json',
        # Verified sentinel membership rows and their durable Excel ledger are
        # direct legislation inputs.  Without these hashes, a legitimate new
        # join looks like same-input/different-output corruption.
        'data/members/collector_roster_updates.ndjson',
        'data/members/collector_roster_pending.ndjson',
        'data/engine/sentinel_bill_workbook_receipt.json',
        'Sentinel Bill Ledger.xlsx',
        'data/members/lep_criterion.json','member_desk_manual.json','data/members/member_desk_src.json',
        'member_narratives_v2.json','data/members/member_personal_texture.json',
        # Profile acceptance changes when its evidence or receipt changes, even if
        # the prose is unchanged. Register those inputs and the descriptive producer.
        'data/members/member_narrative_evidence_v4.json',
        'data/members/member_narratives_v2_semantic_receipt.json',
        'data/members/member_narratives_v3_descriptive_receipt.json',
        os.path.join('fec','members','validate_descriptive_profiles.py'),
        os.path.join('fec','members','regenerate_descriptive_profiles.py'),
        os.path.join('fec','members','descriptive_profile_overrides.py'),
        os.path.join('fec','members','descriptive_public_accounts.py'),
        'data/members/member_wikipedia_summaries.json','member_receptivity.json',
        'member_references.json','model_bio.json','optimize_result.json',
        'data/engine/pillar_inputs_snapshot.json','data/stance/press_date_overrides.json','data/stance/press_stance_overrides.json',
# N1a 2026-09-19: under DESK_PRESS_SOURCE=engine the desk's press rows are the union of
# 'Press Release Articles' and the tone stream fec/press_stream.py returns, so the refresh tone
# file and the receipt that hash-binds it are desk INPUTS. Registering them is the same lesson
# the bill-lobbying sidecar taught on 08-06: an unregistered dependency makes the first
# refresh-only change look like same-input/different-bundle corruption. They are hashed under
# either dial position — a declared input that is not read can only relax this check, while an
# undeclared one that IS read breaks it.
        'data/stance/press_refresh_tone.csv','data/stance/press_refresh_feed_receipt.json',
        os.path.join('fec','press_stream.py'),
        os.path.join('fec','surfaces','desk_press_source.py'),
        'data/members/profile_basis_v2.json','ranks_python.json','data/lobbying/signed_coalitions.json',
        'data/district/district_briefs_v1.json',
# 2026-08-06: the bill-lobbying sidecar became a desk input in the 08-05 five-lane
# pass but was never fingerprinted; the first sidecar-only change tripped the
# same-input/different-bundle alarm (correctly - unregistered dependency).
'data/lobbying/bill_lobbying_sidecar.json',
        'data/members/staff_directory.json','staff_interactions.json','staff_state.json','data/members/formal_visibility_roles.json',
        os.path.join('fec','legislators.csv'),os.path.join('fec','surfaces','make_member_desk.py'),
        os.path.join('fec','stance','media_rules.py'),os.path.join('fec','bill_rules.py'),
        os.path.join('fec','members','member_identity.py'),os.path.join('fec','audit','dashboard_integrity.py'),
        os.path.join('fec','legislation','build_registry_anchor_backfill.py'),
        os.path.join('fec','district','demographic_potential.py'),
        os.path.join('fec','district','demographic_normalization.py'),
        os.path.join('fec','district','district_cu_metrics.py'),
        os.path.join('fec','legislation','congress_bill_repository.py'),
        os.path.join('fec','legislation','dashboard_legislative_workbench.py'),
        os.path.join('fec','legislation','update_bill_repository.py'),
        os.path.join('fec','refresh_bill_repository_headless.mjs'),
        'DEMOGRAPHIC_POTENTIAL_METHOD.md',
        os.path.join('fec','moneyball_core.py'),os.path.join('fec','stance','stance_rules.py'),
        os.path.join('fec','tier_util.py'),os.path.join('fec','engine','primacy_factors.py'),
        os.path.join('fec','legislation','legislation_culling_status.py'),
        os.path.join('fec','stance','per_category_stance.py'),os.path.join('fec','engine','compute_pillars.py'),
        os.path.join('fec','engine','influence_roles.py'),
        os.path.join('fec','surfaces','templates','make_member_desk_00.tpl'),
        # media-footprint source (2026-07-16): the DB plus its WAL sidecar — recent commits live in
        # the -wal until checkpoint, and hashing the main file alone would call changed data "same".
        'social_media.db','social_media.db-wal',
    ]
    source_state={rel:_file_sha(p(rel)) for rel in source_files}
    # model_bio follows make_member_desk.L(): the root path wins and the fec/ copy is the
    # maintained fallback.  The live tree currently has only fec/model_bio.json; hashing the
    # nonexistent root alone made biography input drift invisible to this integrity state.
    _model_bio_root=p('model_bio.json')
    _model_bio_fallback=p(os.path.join('fec','model_bio.json'))
    source_state['model_bio.json']=_file_sha(
        _model_bio_root if os.path.exists(_model_bio_root) else _model_bio_fallback)
    if os.environ.get('CATEGORY_STANCE_PATH'):
        source_state['category_stance.json']=_file_sha(os.environ['CATEGORY_STANCE_PATH'])
    if os.environ.get('DASHBOARD_DATA_PATH'):
        source_state['dashboard_data.json']=_file_sha(os.environ['DASHBOARD_DATA_PATH'])
    # Match make_member_desk.L(): root copy wins, then fec/ fallback; if
    # neither exists, _file_sha records None.  These optional district layers
    # all change the semantic desk bundle and therefore must participate in
    # the same-input/same-output fingerprint.
    for rel in OPTIONAL_DISTRICT_SOURCES:
        root_path=os.path.join(ROOT,rel)
        fallback_path=os.path.join(ROOT,'fec',rel)
        selected_path=root_path if os.path.exists(root_path) else fallback_path
        source_state[rel]=_file_sha(selected_path)
    # The HISTORICAL half of the desk's press stream lives OUTSIDE the project root (the R
    # pipeline's Draft Systems tree), so it cannot be expressed as a root-relative entry like
    # every other source above. Hash it by the constant press_stream itself resolves, which is
    # also what MONEYBALL_DRAFT relocates. press_stream has no top-level I/O.
    import press_stream as PS
    source_state['press_stream:BASE_TONE_PATH']=_file_sha(PS.BASE_TONE_PATH)
    # the footprint display dials shape the bundle, so they are inputs to the same-input check
    source_state['env:DESK_FOOTPRINT']=os.environ.get('DESK_FOOTPRINT','1')
    source_state['env:DESK_MEDIA_CAP']=os.environ.get('DESK_MEDIA_CAP','0')
    source_state['env:DESK_MEDIA_STANCE']=os.environ.get('DESK_MEDIA_STANCE','1')
    # N1a: which press corpus the desk read is the largest single determinant of its press rows.
    source_state['env:DESK_PRESS_SOURCE']=os.environ.get('DESK_PRESS_SOURCE','engine')
    # COMPET_MULT: the dial and its artifact override are inputs to the bundle's ROI column (unset or '' == ON, the default since 2026-10-06).
    source_state['env:COMPET_MULT']=os.environ.get('COMPET_MULT','1') or '1'
    source_state['env:COMPET_SRC']=os.environ.get('COMPET_SRC','')
    if os.environ.get('OPTIMIZE_RESULT_PATH'):    # desk plan source override (make_member_desk.py): the allocation columns come from THIS file when set
        source_state['optimize_result.json']=_file_sha(os.environ['OPTIMIZE_RESULT_PATH'])
    input_hash=_sha(source_state); state_path=_op('data/engine/dashboard_integrity_state.json')
    os.makedirs(os.path.dirname(state_path), exist_ok=True)
    prior=json.load(open(state_path,encoding='utf-8')) if os.path.exists(state_path) else None
    if prior and prior.get('input_fingerprint')==input_hash:
        _check(prior.get('output_semantic_sha256')==output_hash,
               'same dashboard inputs produced a different semantic bundle', errors)
    report={'status':'pass' if not errors else 'fail','errors':errors,'warnings':warnings,
            'preflight':preflight_report,'counts':{'members':len(members),'articles':article_total,
                'bills':bill_total,'members_with_social':len(social_members),
                'district_briefs':district_brief_count,
                'demographic_potential':potential_count,
                'demographic_potential_markets':len(potential_markets)},
            'input_fingerprint':input_hash,'output_semantic_sha256':output_hash,
            'regeneration_match': None if not prior or prior.get('input_fingerprint')!=input_hash else not errors,
            'missing_bioguide_enrichment':missing_bio,'exclusion_reconciliation':reconciliation}
    if errors: raise IntegrityError('; '.join(errors[:8]))
    MC.dump_atomic(report,_op('data/engine/dashboard_integrity_report.json'),indent=2,ensure_ascii=False)
    MC.dump_atomic({'input_fingerprint':input_hash,'output_semantic_sha256':output_hash},state_path,indent=2)
    return report
