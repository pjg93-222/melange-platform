"""Reconcile scoring arithmetic, adjudicated labels and published consumers."""
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
import hashlib
import json
import math
import sys
import unicodedata

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import bill_rules as BR
import moneyball_core as MC
import tier_util as TU
from engine import primacy_factors as PF
from engine.live_legislative_pillar import measure, percentranks
from engine.pillar_policy import PILLAR_WEIGHTS
from legislation import sentinel_admissions as SA
from stance import stance_rules as SR
from stance.adjudicated_corroboration import from_routing
from stance.honors_views import load_honors

def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def nfc(value):
    return unicodedata.normalize('NFC', str(value))

def keyed(mapping):
    result = {nfc(k): v for k, v in mapping.items()}
    if len(result) != len(mapping):
        raise ValueError('Duplicate normalized member identity')
    return result

def active_overall_board(root):
    """The crown authority the surfaces publish from (D-030/D-031).

    The overall crown is the active honors board's top 13 with a qualifying action, NOT
    the retired 'Legislative+Press Tier 1 -> Champion' overlay. This reads the same
    fail-closed generation the producers read (surfaces/build_dashboard_bundle.py:96),
    so the guard and the published surfaces cannot drift apart. None means no active
    generation is in force, which is exactly the fallback the producers take.
    """
    if not (Path(root)/'data/stance/rank_honors/active.json').is_file():
        return None
    return load_honors(root)

def verify(root):
    root = Path(root)
    errors, counts = [], defaultdict(int)
    def check(condition, message, group):
        counts[group] += 1
        if not condition:
            errors.append(message)
    def near(a, b, tolerance=1e-6):
        return (isinstance(a, (int, float)) and isinstance(b, (int, float))
                and math.isfinite(a) and math.isfinite(b) and abs(a-b) <= tolerance)
    inputs = {}
    try:
        paths = ('ranks_python.json','archive_leg.json','legislation_routing_audit.json',
                 'bill_repository_live.json','category_stance.json','member_desk_bundle.json',
                 'dashboard_data.json','data/engine/primacy_factors.json',
                 'data/lobbying/industry_tiers.json','vote_category_lean.json')
        docs = {p: read(root/p) for p in paths}
        inputs = {p: sha(root/p) for p in paths}
        ranks = docs['ranks_python.json']; names = list(ranks['rank'])
        normalized_names = {nfc(n) for n in names}
        routing = docs['legislation_routing_audit.json']['rows']
        repo = docs['bill_repository_live.json']['bills']
        live = ranks['legislative_pillar']
        current = measure(names, routing, repo, live['topic_weights'], live['efficiency_weights'], live['blend'])
        for field in ('agenda','efficiency','pillar','agenda_percentile','efficiency_percentile'):
            check(all(near(current[field][n],live[field][n],1e-9) for n in names),
                  'Live legislative '+field+' does not match adjudicated rows','legislative_math')
        archive = docs['archive_leg.json']; totals = defaultdict(lambda: defaultdict(float)); seen = set()
        for row in routing:
            if row.get('disposition') != 'selected': continue
            identity = (row['member'],row['bill'],BR._congress(row['congress']),row['effective_category'])
            check(identity not in seen, 'Duplicate selected row: '+str(identity),'routing')
            seen.add(identity)
            expected = row['sign'] * BR.role_weight(BR.canonical_role(row['effective_role'])) * row['decay_weight']
            check(near(row['contribution'],expected,1e-9),'Signed role/decay arithmetic: '+str(identity),'routing')
            totals[row['member']][row['effective_category']] += row['contribution']
        for n in names:
            for cat in set(totals[n]) | set(archive.get(n,{})):
                check(near(totals[n][cat],archive.get(n,{}).get(cat,0),1e-9),
                      'Archive/routing mismatch: '+n+' '+cat,'routing')
        weights = {'leg':PILLAR_WEIGHTS[2], 'press':PILLAR_WEIGHTS[3],
                   'inf':PILLAR_WEIGHTS[4], 'dem':PILLAR_WEIGHTS[5]}
        check(all(near(ranks['ab_weights'][str(k)],v,1e-12) for k,v in PILLAR_WEIGHTS.items()),
              'Headline weights differ from policy','headline_math')
        legislative = percentranks(current['pillar'])
        for n in names:
            check(near(ranks['pcts'][n]['leg'],legislative[n],1e-9),'Legislative percentile: '+n,'headline_math')
            score = sum(w*ranks['pcts'][n][k] for k,w in weights.items())
            check(near(score,ranks['score'][n],1e-9),'Weighted headline score: '+n,'headline_math')
            expected_rank = 1+sum(s>ranks['score'][n] for s in ranks['score'].values())
            check(expected_rank==ranks['rank'][n],'Competition rank/tie arithmetic: '+n,'headline_math')

        desk = docs['member_desk_bundle.json']; dashboard = docs['dashboard_data.json']
        by_name = keyed({m['name']:m for m in desk['members']})
        check(set(by_name)==normalized_names,'Desk roster mismatch','surfaces')
        check(len(dashboard['name'])==len(set(map(nfc,dashboard['name'])))==len(names),
              'Dashboard duplicated or incomplete roster','surfaces')
        factors = docs['data/engine/primacy_factors.json']
        tags = SA.effective_tags(PF.load_tags(root/'bill_beneficiary_tags.ndjson'),SA.load(root))
        credits = PF.load_adjudicated_credits(root/PF.ADJUDICATION_LEDGER)
        cells = PF.load_cu_tensor_cells(root/PF.CU_TENSOR)
        for member in desk['members']:
            keys = [(b['bill'],str(b['congress'])) for b in member['bills']]
            check(len(keys)==len(set(keys)),'Duplicate primacy act: '+member['name'],'primacy_math')
        recomputed, baseline = PF.compute(desk,tags,credits,cu_cells=cells)
        check(near(baseline,factors['_constants']['roster_baseline']), 'Primacy baseline mismatch','primacy_math')
        factor_fields = ('factor_raw','agenda_credit','factor_shrunk','factor_ranked','lambda')
        if PF.PRIMACY_T1_MIN_AFFIRMATIVE > 0:
            factor_fields += (PF.AFFIRMATIVE_PRIMARY_FIELD,)
        for n, values in recomputed.items():
            for field in factor_fields:
                check(near(values[field],factors['members'][n][field]),'Primacy '+field+': '+n,'primacy_math')
        tiers, positions = TU.legislative_press_tiers(names,root=str(root))
        tiers, positions = keyed(tiers), keyed(positions)
        labels = docs['category_stance.json']; matrix = keyed(labels['matrix'])
        scores = keyed(labels['scores'])
        check(set(matrix)==normalized_names,'Label roster mismatch','labels')
        binding = labels.get('source_contract',{})
        check(binding.get('schema')=='adjudicated-label-inputs-v1','Labels lack adjudicated source binding','labels')
        for path in ('archive_leg.json','legislation_routing_audit.json','vote_category_lean.json','bill_repository_live.json'):
            check(binding.get('inputs_sha256',{}).get(path)==sha(root/path),'Stale label source: '+path,'labels')
        signal = labels.get('signal')
        check(signal is not None,'Labels lack explicit evidence-presence state','labels')
        corr = keyed(from_routing(routing,names,repo))
        votes = docs['vote_category_lean.json']; vote_corr = defaultdict(set)
        for bio, cats in votes['votes'].items():
            name = nfc(votes['name_of'].get(bio,''))
            for cat, v in cats.items():
                if v.get('lean',0)>=0.5 and v.get('n',0)>=1: vote_corr[name].add(cat)
        if signal is not None:
            signal = keyed(signal)
            curved = SR.curve_label_scores(scores,signal,SR.CURVED_CATEGORIES)
            for n in matrix:
                for cat in labels['cats']:
                    raw = labels['raw_scores'][n][cat]
                    rounded = SR.clip(raw+labels['rounding']['factors'][n][cat])
                    check(near(scores[n][cat],rounded,0.000501),'Label rounding: '+n+' '+cat,'labels')
                    expected = SR.classify(curved[n][cat],signal[n][cat],cat in labels['defensive'])
                    authored = cat in corr.get(n,{}).get('sponsored',set())
                    voted = cat in vote_corr[n]
                    if expected in ('Champion','Defender') and not (authored or voted): expected='Ally'
                    check(matrix[n][cat]==expected,'Label assignment: '+n+' '+cat,'labels')

        honors = active_overall_board(root)
        overall_caps = (labels.get('gate') or {}).get('overall_capped') or {}
        board_top13 = None
        if honors is not None:
            board_top13 = {nfc(n) for n in ((honors.get('overall_board') or {}).get('top13') or ())} or None
            check(board_top13 is not None and len(board_top13)==13,
                  'Active overall board does not publish a top-13 crown','crown')
            check(board_top13 is None or board_top13<=normalized_names,
                  'Overall crown names lie outside the published roster','crown')
            honors_members = keyed(honors.get('members') or {})
            for n in sorted(board_top13 or ()):
                row = (honors_members.get(n) or {}).get('overall') or {}
                # The producer refuses a crown without a recorded action
                # (stance/overall_stance_policy.py:171); verify the published row carries it.
                check(row.get('recognition')=='Champion'
                      and row.get('reason')=='active_overall_board_top13'
                      and row.get('evidence_status')=='observed'
                      and (row.get('action_count') or 0)>=1,
                      'Overall crown lacks a qualifying action: '+n,'crown')
            for relative in ('data/stance/rank_honors/active.json',
                             'data/stance/rank_honors/%s/rank_honors.json'%honors['generation_id']):
                inputs[relative] = sha(root/relative)
        letters = {v:k for k,v in desk['tiername'].items()}
        normalized_rank = keyed(ranks['rank']); normalized_score = keyed(ranks['score'])
        pcts = keyed(ranks['pcts']); raw_leg = keyed(live['pillar'])
        for i, name in enumerate(dashboard['name']):
            n=nfc(name); member=by_name[n]
            expected_overall = letters[MC.overall_display_tier(matrix[n],tiers[n],n,overall_caps,board_top13)]
            for field,expected in (('rank',normalized_rank[n]),('tier',tiers[n]),('tierrank',positions[n])):
                check(dashboard[field][i]==member[field]==expected,'Rank/tier surface mismatch: '+n+' '+field,'surfaces')
            check(near(member['score'],normalized_score[n],1e-9),'Desk headline score: '+n,'surfaces')
            check(near(member['leg_raw'],raw_leg[n],1e-9),'Desk stale legislative raw value: '+n,'surfaces')
            for dest,source in (('leg_n','leg'),('press_n','press'),('comm_n','inf'),('demo_n','dem')):
                check(near(member[dest],pcts[n][source],1e-9),'Desk pillar mismatch: '+n+' '+dest,'surfaces')
            check(dashboard['gen'][i]==member['overall_stance']==expected_overall,'Overall label mismatch: '+n,'surfaces')
            check(member['stances']==''.join(letters[matrix[n][cat]] for cat in desk['cats']),
                  'Desk issue labels: '+n,'surfaces')
            engine=dashboard['engine']
            check(near(engine['tierp'][i],MC.primacy_engine_value(tiers[n])),'Engine primacy tier: '+n,'engine_inputs')
            for j,cat in enumerate(engine['categories']):
                expected=letters[matrix[n][cat]] if cat else expected_overall
                check(dashboard['stance'][i][j]==expected,'Dashboard issue label: '+n+' '+str(cat),'surfaces')
                check(near(engine['priority'][j][i],MC.blended_engine_priority(tiers[n],scores[n],cat,'flip')),
                      'Engine priority uses stale labels/tier: '+n+' '+str(cat),'engine_inputs')
                check(near(engine['orientation'][j][i],MC.engine_orientation(scores[n],cat)),
                      'Engine orientation uses stale category scores: '+n+' '+str(cat),'engine_inputs')
        industry=docs['data/lobbying/industry_tiers.json']['industries']['credit_union']['rows']
        check(len(industry)==len(names) and {nfc(row['member']) for row in industry}==normalized_names,
              'Industry CU board duplicated or incomplete roster','surfaces')
        for row in industry:
            n=nfc(row['member'])
            check(n in tiers and (row['tier'],row['tierrank'])==(tiers.get(n),positions.get(n)),
                  'Industry CU board drift: '+n,'surfaces')
        workbook_path=root/'Moneyball Main Version - Blended.xlsx'
        if workbook_path.exists():
            import openpyxl
            from stance.workbook_stance_views import GENERAL_STANCE_TIERS
            inputs[workbook_path.name]=sha(workbook_path)
            wb=openpyxl.load_workbook(workbook_path,read_only=True,data_only=False)
            ra=wb['Rank & Allocation'];bm=wb['Blended Model'];cs=wb['Category Stance'];cp=wb['Coalition Profiles']
            check([c.value for c in next(cs.iter_rows(min_row=1,max_row=1,min_col=3,max_col=20))]==labels['cats'],
                  'Workbook category order is stale','workbook')
            book_labels={nfc(row[0]):row[1:20] for row in cs.iter_rows(min_row=2,max_row=534,values_only=True) if row[0]}
            check(set(book_labels)==normalized_names,'Workbook label roster mismatch','workbook')
            bm_rows=list(bm.iter_rows(min_row=3,max_row=535,values_only=True))
            for index,row in enumerate(ra.iter_rows(min_row=3,max_row=535,values_only=True)):
                n=nfc(row[0]);r=index+3;values=bm_rows[index]
                # Byte-for-byte the call stance/sync_stance_g.py:agg makes, so the
                # workbook is checked against the one crown source and not a third
                # rule: board seats first, then the capped analytic band.
                expected=MC.overall_display_tier(matrix[n],tiers[n],n,overall_caps,board_top13)
                check(row[6]==expected and book_labels[n][0]==expected,'Workbook overall label: '+n,'workbook')
                check(list(book_labels[n][1:])==[matrix[n][cat] for cat in labels['cats']],
                      'Workbook issue labels: '+n,'workbook')
                check(near(values[3],live['agenda'][row[0]],1e-9) and near(values[6],live['efficiency'][row[0]],1e-9),
                      'Workbook legislative inputs stale: '+n,'workbook')
                check(row[3]==f"='Blended Model'!D{r}",'Workbook detached legislative input: '+n,'workbook')
                check(values[14]==f'=$AB$10*J{r}+$AB$11*K{r}' and values[18]==f'=$AB$2*P{r}+$AB$3*I{r}+$AB$4*R{r}+$AB$5*M{r}',
                      'Workbook scoring formula drift: '+n,'workbook')
            check(tuple(row[0] for row in cp.iter_rows(min_row=18,max_row=28,min_col=1,max_col=1,values_only=True))==GENERAL_STANCE_TIERS,
                  'Workbook coalition taxonomy stale','workbook')
            wb.close()
    except Exception as exc:
        errors.append(type(exc).__name__+': '+str(exc))
    return {'schema':'measurement-layer-reconciliation-v1','status':'FAIL' if errors else 'PASS',
            'checked_at':datetime.now(timezone.utc).isoformat(),'checks':dict(counts),
            'errors':errors,'inputs_sha256':inputs}

if __name__=='__main__':
    from paths import ROOT
    result=verify(ROOT)
    MC.dump_atomic(result,str(Path(ROOT)/'data/engine/measurement_layer_receipt.json'))
    print(json.dumps({'status':result['status'],'checks':result['checks'],'error_count':len(result['errors']),
                      'errors':result['errors'][:15]},ensure_ascii=True))
    raise SystemExit(0 if result['status']=='PASS' else 1)
