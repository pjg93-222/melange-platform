# -*- coding: utf-8 -*-
"""Insert and delete whole rows in an .xlsx package with Excel's reference semantics (roster lane, 2026-10-20).

The roster refresh grows and shrinks the member blocks of the Melange workbook ('Rank & Allocation', 'Blended Model', the
Press / Legislative / Sub Layer / Demographic sheets ...).  openpyxl's insert_rows moves cells but rewrites no formula, and a
LibreOffice or Excel round trip rewrites the whole package (D-174).  This module edits the package directly:

  * every formula in every sheet (and table calculatedColumnFormula, conditional-format formula) that points at the target
    sheet is re-addressed the way Excel does it on Insert/Delete: a reference at or below the cut moves; a range that spans
    the cut grows or shrinks.  `grow_end=True` on an insert ALSO grows ranges that end exactly on the row above the cut, which
    is what appending a member at the bottom of a block needs (Excel would leave those ranges short);
  * the target sheet's <row>/<c> addresses, dimension, merged cells, conditional-format sqref, hyperlinks and the refs of its
    tables (and their autoFilter) follow;
  * an inserted row is cloned from a template row: its styles and its formulas (translated to the new row, relative parts
    moving as on fill-down); its literal values are NOT copied.  The caller supplies the new row's values.
A reference INTO a deleted row from a cell that survives raises RowOpError (Excel would write #REF!).

Package level: every part that is not rewritten is carried over byte for byte (entry order, names, dates, compression).
Stdlib + openpyxl's formula tokenizer and translator only.  Never writes the source; the caller puts dst in place.
"""
import re
import zipfile
import posixpath
from xml.etree import ElementTree as ET

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import column_index_from_string, get_column_letter

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
Q = '{%s}' % NS
for _p, _u in (('', NS), ('r', RNS), ('mc', 'http://schemas.openxmlformats.org/markup-compatibility/2006'),
               ('x14ac', 'http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac')):
    ET.register_namespace(_p, _u)

CELL = re.compile(r'(\$?)([A-Z]{1,3})(\$?)([0-9]+)')
AREA = re.compile(r'(\$?)([A-Z]{1,3})(\$?)([0-9]+)(?::(\$?)([A-Z]{1,3})(\$?)([0-9]+))?')
ROWS = re.compile(r'(\$?)([0-9]+):(\$?)([0-9]+)')
SHEETED = re.compile(r"^((?:'(?:[^']|'')+'|[A-Za-z0-9_.]+)!)(.*)$")


class RowOpError(ValueError):
    """The operation would leave a broken reference or does not fit the workbook."""


def _sheet_name(prefix):
    s = prefix[:-1]
    return s[1:-1].replace("''", "'") if s.startswith("'") else s


class Op:
    """kind 'insert': n new rows at row `at` (old row `at` moves to at+n).  kind 'delete': rows at .. at+n-1 go."""

    def __init__(self, sheet, kind, at, n=1, grow_end=False):
        if kind not in ('insert', 'delete') or at < 1 or n < 1:
            raise RowOpError('bad op %r %r %r %r' % (sheet, kind, at, n))
        self.sheet, self.kind, self.at, self.n, self.grow_end = sheet, kind, at, n, grow_end

    def __repr__(self):
        return 'Op(%r, %r, %d, %d%s)' % (self.sheet, self.kind, self.at, self.n, ', grow_end' if self.grow_end else '')

    # ---- row arithmetic -------------------------------------------------------------------------------------------
    def cell_row(self, r):
        """New row of a single reference, or None when it pointed into a deleted row."""
        if self.kind == 'insert':
            return r + self.n if r >= self.at else r
        if r < self.at:
            return r
        return None if r < self.at + self.n else r - self.n

    def span(self, r1, r2):
        """New (r1, r2) of a row span, or None when the whole span was deleted."""
        lo, hi = min(r1, r2), max(r1, r2)
        if self.kind == 'insert':
            if lo >= self.at:
                return lo + self.n, hi + self.n
            if hi >= self.at or (self.grow_end and hi == self.at - 1):
                return lo, hi + self.n
            return lo, hi
        last = self.at + self.n - 1
        nlo = lo if lo < self.at else (self.at if lo <= last else lo - self.n)
        nhi = hi if hi < self.at else (self.at - 1 if hi <= last else hi - self.n)
        return None if nhi < nlo else (nlo, nhi)


def _ref(op, text, where):
    """Re-address one A1 reference text (no sheet prefix).  Returns the new text."""
    m = AREA.fullmatch(text)
    if m:
        c1, a, c2, r1 = m.group(1, 2, 3, 4)
        if m.group(6) is None:
            nr = op.cell_row(int(r1))
            if nr is None:
                raise RowOpError('%s refers to %s, a deleted row of %s' % (where, text, op.sheet))
            return '%s%s%s%d' % (c1, a, c2, nr)
        d1, b, d2, r2 = m.group(5, 6, 7, 8)
        s = op.span(int(r1), int(r2))
        if s is None:
            raise RowOpError('%s refers to %s, wholly inside deleted rows of %s' % (where, text, op.sheet))
        lo, hi = s if int(r1) <= int(r2) else (s[1], s[0])
        return '%s%s%s%d:%s%s%s%d' % (c1, a, c2, lo, d1, b, d2, hi)
    m = ROWS.fullmatch(text)
    if m:
        s = op.span(int(m.group(2)), int(m.group(4)))
        if s is None:
            raise RowOpError('%s refers to rows %s, all deleted from %s' % (where, text, op.sheet))
        return '%s%d:%s%d' % (m.group(1), s[0], m.group(3), s[1])
    return text            # whole columns (A:A), structured references, LET names, errors


def shift_formula(formula, home_sheet, op, where=''):
    """formula without the leading '='; home_sheet = the sheet the formula lives on (for unprefixed references).
    op is one Op, or {sheet: Op} for a group (one op per sheet: their re-addressings touch disjoint references)."""
    ops = op if isinstance(op, dict) else {op.sheet: op}
    if home_sheet not in ops and not any(s.replace("'", "''") in formula for s in ops):
        return formula
    tok = Tokenizer('=' + formula)
    out = []
    for t in tok.items:
        v = t.value
        if t.type == 'OPERAND' and t.subtype == 'RANGE' and '[' not in v:
            m = SHEETED.match(v)
            if m:
                o = ops.get(_sheet_name(m.group(1)))
                if o is not None:
                    v = m.group(1) + _ref(o, m.group(2), where)
            elif home_sheet in ops:
                v = _ref(ops[home_sheet], v, where)
        out.append(v)
    return ''.join(out)


def _sqref(op, text, where, tolerate_loss=False):
    parts = []
    for p in text.split():
        try:
            parts.append(_ref(op, p, where))
        except RowOpError:
            if not tolerate_loss:
                raise
    return ' '.join(parts)


# ---------------------------------------------------------------------------------------------------------- package

def _paths(z):
    rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    tgt = {}
    for r in rels:
        t = r.get('Target')
        tgt[r.get('Id')] = t.lstrip('/') if t.startswith('/') else posixpath.normpath(posixpath.join('xl', t))
    book = ET.fromstring(z.read('xl/workbook.xml'))
    sheets = {s.get('name'): tgt[s.get('{%s}id' % RNS)] for s in book.find(Q + 'sheets')}
    tables = {}            # sheet name -> [table part]
    for name, part in sheets.items():
        rp = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
        if rp in z.namelist():
            for r in ET.fromstring(z.read(rp)):
                if r.get('Type', '').endswith('/table'):
                    t = r.get('Target')
                    tables.setdefault(name, []).append(
                        t.lstrip('/') if t.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(part), t)))
    return sheets, tables


def _rownum(addr):
    return int(CELL.fullmatch(addr).group(4))


def _col(addr):
    return CELL.fullmatch(addr).group(2)


class Book:
    """The parsed parts a row operation touches.  apply() runs ops in order; save() writes the package."""

    def __init__(self, src):
        self.src = src
        with zipfile.ZipFile(src) as z:
            self.sheets, self.tables = _paths(z)
            self.trees = {n: ET.fromstring(z.read(p)) for n, p in self.sheets.items()}
            self.ttrees = {p: ET.fromstring(z.read(p)) for ps in self.tables.values() for p in ps}
        self.dirty = set()

    # every formula-bearing element of a sheet: (element, where)
    def _formulas(self, name):
        tree = self.trees[name]
        for c in tree.iter(Q + 'c'):
            f = c.find(Q + 'f')
            if f is not None and f.text:
                if f.get('t') in ('shared', 'array'):
                    raise RowOpError('%s!%s holds a %s formula; not supported' % (name, c.get('r'), f.get('t')))
                yield c, f
        for cf in tree.iter(Q + 'conditionalFormatting'):
            for fe in cf.iter(Q + 'formula'):
                if fe.text:
                    yield None, fe

    def apply_group(self, ops):
        """Run ops (one per sheet) as one edit: a cell that a LATER delete in the group removes is not re-addressed by an
        earlier op, so the same member row can leave several sheets that point at each other ('Blended Model'!A301 =
        'Rank & Allocation'!A301) without a false broken reference."""
        if len({o.sheet for o in ops}) != len(ops):
            raise RowOpError('a group holds at most one op per sheet')
        for o in ops:
            if o.sheet not in self.trees:
                raise RowOpError('no sheet %r' % o.sheet)
        by = {o.sheet: o for o in ops}
        skip = {o.sheet: set(range(o.at, o.at + o.n)) for o in ops if o.kind == 'delete'}
        # 1. formulas everywhere, one tokenizer pass per formula for the whole group
        for name in self.trees:
            gone = skip.get(name, ())
            for c, f in self._formulas(name):
                if c is not None and _rownum(c.get('r')) in gone:
                    continue
                where = '%s!%s' % (name, c.get('r') if c is not None else 'conditional format')
                new = shift_formula(f.text, name, by, where)
                if new != f.text:
                    f.text = new
                    self.dirty.add(name)
        for name, parts in self.tables.items():
            for p in parts:
                for fe in self.ttrees[p].iter(Q + 'calculatedColumnFormula'):
                    if fe.text:
                        new = shift_formula(fe.text, name, by, '%s table formula' % name)
                        if new != fe.text:
                            fe.text = new
                            self.dirty.add(p)
        # 2. each target sheet's own geometry
        for op in ops:
            self._geometry(op)

    def apply(self, op):
        self.apply_group([op])

    def _geometry(self, op):
        dying = set(range(op.at, op.at + op.n)) if op.kind == 'delete' else set()
        tree = self.trees[op.sheet]
        data = tree.find(Q + 'sheetData')
        for row in list(data):
            r = int(row.get('r'))
            if r in dying:
                data.remove(row)
                continue
            nr = op.cell_row(r)
            if nr != r:
                row.set('r', str(nr))
                for c in row.iter(Q + 'c'):
                    c.set('r', _col(c.get('r')) + str(nr))
        where = op.sheet + ' geometry'
        dim = tree.find(Q + 'dimension')
        if dim is not None:
            dim.set('ref', _sqref(op, dim.get('ref'), where, tolerate_loss=True) or 'A1')
        for mc in list(tree.iter(Q + 'mergeCell')):
            mc.set('ref', _ref(op, mc.get('ref'), where))
        for cf in tree.iter(Q + 'conditionalFormatting'):
            cf.set('sqref', _sqref(op, cf.get('sqref'), where))
        for h in tree.iter(Q + 'hyperlink'):
            h.set('ref', _ref(op, h.get('ref'), where))
        for sel in tree.iter(Q + 'selection'):
            for k in ('activeCell', 'sqref'):
                if sel.get(k):
                    sel.set(k, _sqref(op, sel.get(k), where, tolerate_loss=True) or 'A1')
        for p in self.tables.get(op.sheet, []):
            t = self.ttrees[p]
            t.set('ref', _ref(op, t.get('ref'), where))
            for af in t.iter(Q + 'autoFilter'):
                af.set('ref', _ref(op, af.get('ref'), where))
            self.dirty.add(p)
        self.dirty.add(op.sheet)

    def _row(self, sheet, r):
        for row in self.trees[sheet].find(Q + 'sheetData'):
            if int(row.get('r')) == r:
                return row
        return None

    def clone_row(self, sheet, template, new):
        """Write row `new` (must be empty) as a copy of row `template`: attributes, styles and formulas (translated);
        no literal values."""
        data = self.trees[sheet].find(Q + 'sheetData')
        src = self._row(sheet, template)
        if src is None:
            raise RowOpError('%s has no template row %d' % (sheet, template))
        if self._row(sheet, new) is not None:
            raise RowOpError('%s row %d is not empty' % (sheet, new))
        row = ET.Element(Q + 'row', {k: v for k, v in src.attrib.items()})
        row.set('r', str(new))
        for c in src.findall(Q + 'c'):
            col = _col(c.get('r'))
            nc = ET.SubElement(row, Q + 'c', {'r': col + str(new)})
            if c.get('s') is not None:
                nc.set('s', c.get('s'))
            f = c.find(Q + 'f')
            if f is not None and f.text:
                nf = ET.SubElement(nc, Q + 'f')
                nf.text = Translator('=' + f.text, origin=c.get('r')).translate_formula(col + str(new))[1:]
                ET.SubElement(nc, Q + 'v')
        rows = list(data)
        idx = next((i for i, x in enumerate(rows) if int(x.get('r')) > new), len(rows))
        data.insert(idx, row)
        self.dirty.add(sheet)
        return row

    def set_value(self, sheet, addr, value):
        """Literal value into a cell (str -> inline string, int/float -> number, None -> empty); refuses formula cells."""
        r = _rownum(addr)
        row = self._row(sheet, r)
        if row is None:
            raise RowOpError('%s row %d does not exist' % (sheet, r))
        cells = row.findall(Q + 'c')
        c = next((x for x in cells if x.get('r') == addr), None)
        if c is None:
            c = ET.Element(Q + 'c', {'r': addr})
            ci = column_index_from_string(_col(addr))
            i = next((k for k, x in enumerate(cells) if column_index_from_string(_col(x.get('r'))) > ci), len(cells))
            row.insert(i, c)
        if c.find(Q + 'f') is not None:
            raise RowOpError('%s!%s holds a formula; not overwriting it' % (sheet, addr))
        for ch in list(c):
            c.remove(ch)
        c.attrib.pop('t', None)
        if isinstance(value, str):
            c.set('t', 'inlineStr')
            t = ET.SubElement(ET.SubElement(c, Q + 'is'), Q + 't')
            t.text = value
            if value != value.strip():
                t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        elif isinstance(value, bool):
            c.set('t', 'b')
            ET.SubElement(c, Q + 'v').text = '1' if value else '0'
        elif isinstance(value, (int, float)):
            ET.SubElement(c, Q + 'v').text = repr(value) if isinstance(value, float) else str(value)
        elif value is not None:
            raise RowOpError('unsupported value type %r' % type(value))
        self.dirty.add(sheet)

    def save(self, dst):
        with zipfile.ZipFile(self.src) as zin, zipfile.ZipFile(dst, 'w') as zout:
            zout.comment = zin.comment
            parts = {p: n for n, p in self.sheets.items()}
            for info in zin.infolist():
                zi = zipfile.ZipInfo(info.filename, info.date_time)
                zi.compress_type, zi.external_attr, zi.create_system, zi.comment = (
                    info.compress_type, info.external_attr, info.create_system, info.comment)
                name = parts.get(info.filename)
                if name is not None and name in self.dirty:
                    body = _xml(self.trees[name])
                elif info.filename in self.ttrees and info.filename in self.dirty:
                    body = _xml(self.ttrees[info.filename])
                else:
                    body = zin.read(info)
                zout.writestr(zi, body, compress_type=info.compress_type)


def _xml(tree):
    return b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + ET.tostring(tree, encoding='utf-8', xml_declaration=False)


def column_letter(i):
    return get_column_letter(i)
