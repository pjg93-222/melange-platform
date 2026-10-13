# -*- coding: utf-8 -*-
"""Committee score is a DERIVED column, and this is its producer.

THE DEFECT THIS CLOSES (2026-09-18)
    'Committee Data' holds, per member row: A member name, B Committee score, C Membership,
    D Leadership, E Seats held. E is the human-authored text -- pipe-joined
    "Name [Highly|Moderately|Slightly]" tokens, optionally "/Chair", "/Ranking" or "/Vice".
    B, C and D are CACHED VALUES, not formulas (verified: 0 of 1170 B/C/D cells hold a
    formula), and until this module existed NO chain step recomputed them -- `committee` had
    zero matches in run_pipeline.py.

    So editing a tier marker in column E -- Paul's prerogative, and the whole point of the
    column being human-readable -- silently left B/C/D carrying the OLD weight. That is not
    a cosmetic drift: engine/live_legislative_pillar.nonlegislative_inputs() reads column B
    as the Influence pillar, and engine/compute_pillars.py weights Influence at 40% of the
    composite score. A stale column B publishes a rank that disagrees with the seat text
    printed beside it on the member desk.

    It happened: retier_committee.py moved 44 "Veterans' Affairs (full)" markers from
    [Moderately] to [Slightly] and the scores did not follow, and the lane repaired 92 cells
    by hand (audits/members_section_20260914/evidence/committee-score-recompute.md). This
    module makes that repair a chain step so the next marker edit cannot go unpriced.

THE SCORING RULE
    A byte-consistent mirror of engine/apply_committee_reweight.py's compute(), which is the
    reproducible port of the 2026-06-30 manual Excel edit and is itself re-used verbatim by
    engine/add_judiciary_seats.py and by both lane scripts
    (audits/members_section_20260914/scripts/{apply_committee_rows,recompute_committee_scores}.py):

        BASE  = Highly 3.0 | Moderately 1.8 | Slightly 0.9
        SUBF  = 1/6                a subcommittee seat weighs its parent's base / 6
        KLEAD = Chair   full 3.0 / sub 7.5
                Ranking full 2.0 / sub 4.5
                Vice    full 1.0 / sub 2.0     (multiplies the seat's OWN membership weight)
        Membership      = sum of seat weights
        Leadership      = sum of leadership contributions
        Committee score = (Membership + Leadership) / 2
        A seat is a SUBcommittee iff its name contains ' - ' (space-hyphen-space).

    The rule is copied rather than imported because apply_committee_reweight.py does its
    workbook IO at module scope -- importing it would read the workbook and, with --apply on
    the command line, drive Excel COM.

TWO DELIBERATE DEPARTURES FROM THE SCRIPTS THIS MIRRORS
    (1) THE 'Small Business' OVERRIDE IS GONE. Those scripts force any seat whose name
        contains 'Small Business' to score as [Highly] regardless of the marker displayed in
        column E -- a 2026-06-30 sheet-wide edit frozen into code. That makes the displayed
        marker a lie for exactly one committee, and it silently defeats the edit this module
        exists to price: retiering a Small Business seat would change the text and nothing
        else. Removing it was gated on proving it inert first, and it is: all 74 Small
        Business seat tokens on the live sheet (4 distinct full/sub names) already read
        [Highly], so dropping the override reproduces B/C/D on 390 of 390 rows, bit for bit.
        The displayed marker is now authoritative, full stop. Evidence:
        audits/members_section_20260914/evidence/fix-committee-producer.md.
        If a Small Business seat is ever retiered to [Moderately] or [Slightly], this module
        will price it as written -- which is the intended behaviour, not a regression.

    (2) A BRACKETLESS SEAT TOKEN IS FATAL. Those scripts `continue` past any token with no
        trailing [marker], so a typo like "Judiciary (full) Slightly" quietly contributes
        ZERO weight and the member's influence silently drops. This module raises
        CommitteeSeatError naming the row, member and token. Today that changes nothing:
        0 of the live sheet's 919 seat tokens are bracketless. It only ever fires on a
        malformed edit, which is precisely when silence is worst.

MODES
    --check   recompute every row, report disagreements, write NOTHING, exit 1 if any.
    --write   repair the disagreeing rows, exit 0. Idempotent.

    --write is a NO-OP ON A CLEAN TREE in the literal sense: when no cell needs to change the
    workbook is never saved and not one byte moves. That matters more here than usual --
    openpyxl re-serializes the whole 33-sheet file, and an unconditional save on every chain
    run would churn 11 MB and re-stale every formula cache for no reason. Saving only on a
    real repair keeps this step invisible until the day it is needed.

SCOPE AND SAFETY
    * Sheet 'Committee Data' only. Columns B, C and D only. Column A (name) and column E
      (seat text) are READ, never written -- this module derives FROM E, it never edits it.
    * Loaded data_only=False so every other sheet's formulas survive the save verbatim. The
      formula CACHES go stale, exactly as add_judiciary_seats.py:71 documents as the standing
      workbook contract; the engine reads typed cells, and the workbook is a recalculating
      view artifact (run_pipeline.py's module docstring).
    * Saves through moneyball_core.save_workbook_atomic -- verified sibling ZIP then
      os.replace -- so an interrupted save cannot truncate the workbook.
    * Refuses if an Excel lock file (~$...) sits beside the workbook, and refuses if the
      workbook's size or mtime moved between load and save (someone else wrote it while we
      held it in memory).
    * Refuses if any B/C/D cell holds a FORMULA: this module repairs cached values and must
      never overwrite somebody's formula.
    * No backup is taken, deliberately. B/C/D are a pure function of column E and column E is
      never touched, so any repair is reproducible and reversible by re-running against the
      corrected text. --backup is available for a cautious manual run.

CHAIN POSITION
    Immediately before compute_pillars.py --write, the first consumer of column B. Every
    later consumer (build_dashboard_bundle, make_member_desk, audit_full) therefore reads the
    repaired values in the same cycle. No earlier chain step writes this sheet.
"""
import argparse
import datetime
import hashlib
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from moneyball_core import save_workbook_atomic
from paths import WORKBOOK

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

SHEET = 'Committee Data'
EXPECTED_HEADER = ['Member (model)', 'Committee score', 'Membership', 'Leadership', 'Seats held']
COL_NAME = 1
COL_SCORE = 2
COL_MEMBERSHIP = 3
COL_LEADERSHIP = 4
COL_SEATS = 5
COL_LABEL = {COL_SCORE: 'Committee score', COL_MEMBERSHIP: 'Membership', COL_LEADERSHIP: 'Leadership'}

# ---- the scoring rule (mirror of engine/apply_committee_reweight.py; keep byte-consistent) ----
BASE = {'Highly': 3.0, 'Moderately': 1.8, 'Slightly': 0.9}
SUBF = 1.0 / 6.0
KLEAD = {'Chair': {'full': 3.0, 'sub': 7.5},
         'Ranking': {'full': 2.0, 'sub': 4.5},
         'Vice': {'full': 1.0, 'sub': 2.0}}
SEPARATOR = ' | '
SUBCOMMITTEE_MARK = ' - '
BRK = re.compile(r'\[([^\]]+)\]\s*$')

EXIT_OK = 0
EXIT_DISAGREE = 1
EXIT_REFUSED = 2


class CommitteeSeatError(ValueError):
    """A seat token this module refuses to price rather than silently score as zero."""


def seat_weights(token):
    """(membership, leadership) for ONE seat token, e.g. 'Banking - Digital Assets [Highly/Chair]'.

    Raises CommitteeSeatError on a token with no trailing [marker], on an unknown tier marker,
    and on an unknown leadership role. The scripts this mirrors skip all three."""
    text = token.strip()
    bracket = BRK.search(text)
    if not bracket:
        raise CommitteeSeatError(
            'seat token has no [marker] and cannot be priced: %r' % token)
    name = text[:bracket.start()].rstrip()
    parts = [part.strip() for part in bracket.group(1).split('/')]
    level = 'sub' if SUBCOMMITTEE_MARK in name else 'full'
    tier = parts[0]
    if tier not in BASE:
        raise CommitteeSeatError(
            'unknown tier marker %r in seat token %r (expected one of %s)'
            % (tier, token, ', '.join(sorted(BASE))))
    membership = BASE[tier] * (1.0 if level == 'full' else SUBF)
    leadership = 0.0
    if len(parts) > 1:
        role = parts[1]
        if role not in KLEAD:
            raise CommitteeSeatError(
                'unknown leadership role %r in seat token %r (expected one of %s)'
                % (role, token, ', '.join(sorted(KLEAD))))
        leadership = membership * KLEAD[role][level]
    return membership, leadership


def compute(seats_text):
    """(Membership, Leadership, Committee score) for one row's 'Seats held' text.

    Empty text scores (0, 0, 0): a member with no seats has no committee influence. Rows whose
    seat text is empty are reported separately by scan() so an accidental blanking is visible."""
    membership = leadership = 0.0
    for token in str(seats_text or '').split(SEPARATOR):
        if not token.strip():
            continue
        seat_membership, seat_leadership = seat_weights(token)
        membership += seat_membership
        leadership += seat_leadership
    return round(membership, 6), round(leadership, 6), round((membership + leadership) / 2, 6)


def _same(cached, recomputed):
    """Exact equality, tolerating an int/float spelling of the same number.

    Exact rather than tolerant on purpose. round(x, 6) yields a double, and writing that double
    through openpyxl and reading it back returns the identical double, so one repair makes a row
    permanently stable -- --write cannot enter a rewrite-every-run loop. A tolerance would instead
    let a real (small) disagreement survive as a published number that contradicts the seat text."""
    if isinstance(cached, (int, float)) and not isinstance(cached, bool):
        return float(cached) == float(recomputed)
    if cached is None:
        return float(recomputed) == 0.0
    return False


def scan(sheet):
    """Recompute every data row. Returns (rows, formula_cells).

    `rows` carries one dict per row with a member name: its cached B/C/D, its recomputed
    B/C/D, whether they disagree, and whether its seat text is empty. `formula_cells` lists
    any B/C/D coordinate holding a formula -- the caller refuses on a non-empty list."""
    header = [sheet.cell(1, column).value for column in range(1, 6)]
    if header != EXPECTED_HEADER:
        raise CommitteeSeatError(
            'header mismatch on %r, refusing to guess columns.\n  got:      %s\n  expected: %s'
            % (SHEET, header, EXPECTED_HEADER))
    rows = []
    formula_cells = []
    for index in range(2, sheet.max_row + 1):
        name = sheet.cell(index, COL_NAME).value
        if not name:
            continue
        cached = {column: sheet.cell(index, column).value
                  for column in (COL_SCORE, COL_MEMBERSHIP, COL_LEADERSHIP)}
        for column, value in cached.items():
            if isinstance(value, str) and value.startswith('='):
                formula_cells.append(sheet.cell(index, column).coordinate)
        seats = str(sheet.cell(index, COL_SEATS).value or '')
        try:
            membership, leadership, score = compute(seats)
        except CommitteeSeatError as error:
            raise CommitteeSeatError(
                'row %d (%s): %s\n  seats held: %s' % (index, str(name).strip(), error, seats))
        fresh = {COL_SCORE: score, COL_MEMBERSHIP: membership, COL_LEADERSHIP: leadership}
        rows.append({
            'row': index,
            'member': str(name).strip(),
            'seats': seats,
            'seats_empty': not seats.strip(),
            'cached': cached,
            'fresh': fresh,
            'disagrees': [column for column in (COL_SCORE, COL_MEMBERSHIP, COL_LEADERSHIP)
                          if not _same(cached[column], fresh[column])],
        })
    return rows, formula_cells


def _md5(path):
    digest = hashlib.md5()
    with open(path, 'rb') as handle:
        for block in iter(lambda: handle.read(1 << 20), b''):
            digest.update(block)
    return digest.hexdigest().upper()


def _fingerprint(path):
    """(size, mtime_ns) -- re-checked before the save so a concurrent writer cannot be clobbered."""
    stat = os.stat(path)
    return stat.st_size, stat.st_mtime_ns


def run(workbook_path, write=False, receipt_path=None, backup_path=None, stream=sys.stdout):
    """Check or repair B/C/D. Returns the process exit code."""
    if not os.path.isfile(workbook_path):
        print('ERROR: workbook not found: %s' % workbook_path, file=sys.stderr)
        return EXIT_REFUSED
    lock = os.path.join(os.path.dirname(workbook_path), '~$' + os.path.basename(workbook_path))
    if os.path.exists(lock):
        print('ERROR: Excel lock file present (%s) -- the workbook is open. Refusing to touch it.'
              % lock, file=sys.stderr)
        return EXIT_REFUSED

    fingerprint_before = _fingerprint(workbook_path)
    # data_only=False is load-bearing: it keeps every other sheet's FORMULAS, which a
    # data_only=True load would replace with their cached values and the save would then
    # freeze into the file permanently.
    book = openpyxl.load_workbook(workbook_path, data_only=False, keep_vba=True)
    try:
        if SHEET not in book.sheetnames:
            print('ERROR: %r sheet not found in %s' % (SHEET, workbook_path), file=sys.stderr)
            return EXIT_REFUSED
        sheet = book[SHEET]
        rows, formula_cells = scan(sheet)

        if formula_cells:
            print('ERROR: %d B/C/D cell(s) hold a FORMULA (e.g. %s). This producer repairs CACHED '
                  'values and refuses to overwrite formulas.' % (len(formula_cells), formula_cells[:5]),
                  file=sys.stderr)
            return EXIT_REFUSED

        stale = [row for row in rows if row['disagrees']]
        blank = [row for row in rows if row['seats_empty']]
        top = max((float(row['fresh'][COL_SCORE]) for row in rows), default=0.0)
        top_member = next((row['member'] for row in rows
                           if float(row['fresh'][COL_SCORE]) == top), '')

        print('workbook : %s' % workbook_path, file=stream)
        print('sheet    : %s   rows: %d   seat tokens: %d'
              % (SHEET, len(rows),
                 sum(len([t for t in row['seats'].split(SEPARATOR) if t.strip()]) for row in rows)),
              file=stream)
        if blank:
            print('note     : %d row(s) have EMPTY seat text and therefore score 0: %s'
                  % (len(blank), ', '.join(row['member'] for row in blank[:10])), file=stream)
        for row in stale:
            for column in row['disagrees']:
                print('  %-6s %-28s %-16s cached=%s  recomputed=%s'
                      % (sheet.cell(row['row'], column).coordinate, row['member'],
                         COL_LABEL[column], row['cached'][column], row['fresh'][column]),
                      file=stream)
                print('         seats: %s' % row['seats'], file=stream)

        repaired = 0
        saved = False
        if write and stale:
            for row in stale:
                for column in row['disagrees']:
                    sheet.cell(row['row'], column).value = row['fresh'][column]
                repaired += 1
            if _fingerprint(workbook_path) != fingerprint_before:
                print('ERROR: the workbook changed on disk while this run held it in memory '
                      '(size/mtime moved). Nothing written -- re-run.', file=sys.stderr)
                return EXIT_REFUSED
            if backup_path:
                import shutil
                os.makedirs(os.path.dirname(backup_path) or '.', exist_ok=True)
                shutil.copy2(workbook_path, backup_path)
                print('backup   : %s' % backup_path, file=stream)
            save_workbook_atomic(book, workbook_path)
            saved = True
    finally:
        book.close()

    if receipt_path:
        receipt = {
            'producer': 'fec/engine/committee_scores.py',
            'generated_utc': datetime.datetime.now(datetime.timezone.utc).isoformat(),
            'workbook': workbook_path,
            'sheet': SHEET,
            'mode': 'write' if write else 'check',
            'rows_scanned': len(rows),
            'rows_disagreeing': len(stale),
            'rows_repaired': repaired,
            'workbook_rewritten': saved,
            'rows_with_empty_seats': [row['member'] for row in blank],
            'max_committee_score': {'value': top, 'member': top_member},
            'md5_after': _md5(workbook_path),
            'small_business_override': 'deleted 2026-09-18; displayed marker is authoritative',
            'disagreements': [
                {'row': row['row'], 'member': row['member'], 'seats': row['seats'],
                 'cells': {COL_LABEL[column]: {'cached': row['cached'][column],
                                               'recomputed': row['fresh'][column]}
                           for column in row['disagrees']}}
                for row in stale],
        }
        os.makedirs(os.path.dirname(receipt_path) or '.', exist_ok=True)
        with open(receipt_path, 'w', encoding='utf-8', newline='\r\n') as handle:
            json.dump(receipt, handle, indent=2)
        print('receipt  : %s' % receipt_path, file=stream)

    # run_pipeline keeps only this LAST line, truncated to 90 characters.
    if write:
        print('committee scores: %d scanned, %d repaired, max %g%s'
              % (len(rows), repaired, top, '' if saved else ' (clean; not rewritten)'), file=stream)
        return EXIT_OK
    print('committee scores: %d scanned, %d disagreements, max %g'
          % (len(rows), len(stale), top), file=stream)
    return EXIT_DISAGREE if stale else EXIT_OK


def main(argv=None):
    parser = argparse.ArgumentParser(
        description='Recompute the Committee Data B/C/D cache from column E (Seats held).')
    parser.add_argument('--workbook', default=WORKBOOK, help='workbook to check or repair')
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument('--check', action='store_true',
                      help='report disagreements, write nothing, exit 1 if any (default)')
    mode.add_argument('--write', action='store_true',
                      help='repair disagreeing B/C/D cells; idempotent; exit 0')
    parser.add_argument('--receipt', default=None, help='optional JSON receipt path')
    parser.add_argument('--backup', default=None,
                        help='optional copy of the workbook taken immediately before a --write save')
    args = parser.parse_args(argv)
    try:
        return run(args.workbook, write=args.write, receipt_path=args.receipt,
                   backup_path=args.backup)
    except CommitteeSeatError as error:
        print('ERROR: %s' % error, file=sys.stderr)
        print('committee scores: REFUSED -- unpriceable seat text', file=sys.stdout)
        return EXIT_REFUSED


if __name__ == '__main__':
    sys.exit(main())
