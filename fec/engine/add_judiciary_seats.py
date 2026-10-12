# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
"""ADD JUDICIARY TO THE TRACKED COMMITTEE SET (user-directed 2026-07-05: committee coverage is on
purpose "unless there are committees where tracked legislation is introduced that I missed. If so,
add those committees and adjust").

Evidence (congress_bills_archive, 481 distinct tracked bills, primary referral = introduction venue):
  tracked committees carry every venue with >=7 primary referrals EXCEPT Judiciary: 10 primary / 19
  total (interchange-antitrust + data/TCPA adjacency). Next untracked: HSGAC 4, Ag 3, Armed Svcs 3
  (defense_cu, IMP 0.2) — noise floor, deliberately excluded per the user. Tier = [Slightly] by the
  sheet's own precedent: Energy & Commerce holds 30 primary referrals at [Slightly]; ranking
  Judiciary (10) above that would contradict the sheet's evidence ordering. One-tier bump is a
  single text edit + apply_committee_reweight.py --apply if the user prefers.

Mechanics (one-shot, reproducible-port pattern of apply_committee_reweight.py):
  1. Backup the workbook (.PRE-JUDICIARY-<ts>.xlsx).
  2. Committee Data: append ' | Judiciary (full) [Slightly(/Role)]' to the 26 members who already
     hold tracked seats; add NEW rows (A..E) for the 38 who don't (Durbin enters here, Ranking).
     B/C/D recomputed with the EXACT apply_committee_reweight formula (mirrored below).
  3. Rank & Allocation col E INDEX/MATCH ranges extended $2:$320 -> $2:$360 (326 rows now exceed
     the old cap; new cap leaves headroom).
  4. Caches: openpyxl save leaves formula caches stale by contract (view self-heals) — run
     `python fec/engine/apply_committee_reweight.py` (expect all-rows reproduced) then `--apply` for the
     COM CalculateFull that refreshes the caches compute_pillars reads, then
     `python fec/run_pipeline.py --from compute_pillars`. Rank movement is EXPECTED and reported there.
Roster source: unitedstates/congress-legislators committee-membership-current (HSJU 42 + SSJU 22),
bioguide-matched to the model roster (64/64 matched, 0 unmatched)."""
import json, re, sys, os, shutil, datetime
import openpyxl
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
B = MB_ROOT
WB = B + r"\Moneyball Main Version - Blended.xlsx"
ROSTER = os.environ.get('CMTE_ROSTER',os.path.join(B,'source_data','committee_membership.json'))
TIER = 'Slightly'
# ---- formula mirror (apply_committee_reweight.py — keep byte-consistent) ----
BASE = {'Highly': 3.0, 'Moderately': 1.8, 'Slightly': 0.9}; SUBF = 1.0 / 6.0
KLEAD = {'Chair': {'full': 3.0, 'sub': 7.5}, 'Ranking': {'full': 2.0, 'sub': 4.5}, 'Vice': {'full': 1.0, 'sub': 2.0}}
BRK = re.compile(r'\[([^\]]+)\]\s*$')
def compute(seats_text):
    mem = lead = 0.0
    for seat in str(seats_text).split(' | '):
        s = seat.strip()
        if not s: continue
        m = BRK.search(s)
        if not m: continue
        name = s[:m.start()].rstrip(); parts = [p.strip() for p in m.group(1).split('/')]
        level = 'sub' if ' - ' in name else 'full'
        tier = 'Highly' if 'Small Business' in name else parts[0]
        mw = BASE[tier] * (1.0 if level == 'full' else SUBF); mem += mw
        if len(parts) > 1: lead += mw * KLEAD[parts[1]][level]
    return round(mem, 6), round(lead, 6), round((mem + lead) / 2, 6)

def role_of(title):
    t = (title or '').lower()
    if 'vice' in t: return 'Vice'
    if 'chair' in t: return 'Chair'
    if 'ranking' in t: return 'Ranking'
    return None

cm = json.load(open(ROSTER))
C = {m['bioguide']: m['member'] for m in json.load(open(B + r'\data\engine\congestion.json')) if m.get('bioguide')}
seats = {}
for code in ('HSJU', 'SSJU'):
    for m in cm.get(code, []):
        nm = C.get(m.get('bioguide', ''))
        if nm: seats[nm] = role_of(m.get('title'))
print('Judiciary roster mapped: %d members' % len(seats))

bk = WB.replace('.xlsx', '.PRE-JUDICIARY-%s.xlsx' % datetime.datetime.now().strftime('%Y%m%d_%H%M%S'))
shutil.copy2(WB, bk); print('backup ->', os.path.basename(bk))
wb = openpyxl.load_workbook(WB)          # formulas preserved; caches go stale by contract
cd = wb['Committee Data']
last = 1
for r in range(2, cd.max_row + 1):
    if cd.cell(r, 1).value: last = r
byname = {}
for r in range(2, last + 1):
    nm = cd.cell(r, 1).value
    if nm: byname[str(nm).strip()] = r
edits = adds = 0
for nm, role in sorted(seats.items()):
    seat = 'Judiciary (full) [%s%s]' % (TIER, ('/' + role) if role else '')
    if nm in byname:
        r = byname[nm]; ev = str(cd.cell(r, 5).value or '')
        if 'Judiciary' in ev: continue                      # idempotent
        ev = (ev + ' | ' + seat) if ev else seat
        mem, lead, sc = compute(ev)
        cd.cell(r, 5).value = ev; cd.cell(r, 2).value = sc; cd.cell(r, 3).value = mem; cd.cell(r, 4).value = lead
        edits += 1
    else:
        last += 1; mem, lead, sc = compute(seat)
        cd.cell(last, 1).value = nm; cd.cell(last, 2).value = sc; cd.cell(last, 3).value = mem
        cd.cell(last, 4).value = lead; cd.cell(last, 5).value = seat
        adds += 1
print('Committee Data: %d rows edited, %d rows added (last row now %d)' % (edits, adds, last))
# ---- extend the R&A INDEX/MATCH cap (320 -> 360) so the new rows are inside the lookup range ----
ra = wb['Rank & Allocation']
ext = 0
for r in range(2, ra.max_row + 1):
    f = ra.cell(r, 5).value
    if isinstance(f, str) and '$320' in f and 'Committee Data' in f:
        ra.cell(r, 5).value = f.replace('$B$2:$B$320', '$B$2:$B$360').replace('$A$2:$A$320', '$A$2:$A$360')
        ext += 1
print('R&A committee INDEX ranges extended 320->360 on %d formula rows' % ext)
wb.save(WB); wb.close()
print('SAVED. Next: python fec/engine/apply_committee_reweight.py  (validate; expect all reproduced)')
print('      then --apply (COM CalculateFull refreshes the caches compute_pillars reads)')
print('      then python fec/run_pipeline.py --from compute_pillars   (rank movement EXPECTED)')
