# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
"""Committee re-weight (reproducible port of the 2026-06-30 manual Excel edit). Recomputes Committee Data
Membership (C), Leadership (D) and Committee score (B) from the Seats-held text (E), applying:
  * ALL 'Small Business' seats -> [Highly]
  * leadership bonuses x1.5:  Chair full 3.0 / sub 7.5,  Ranking full 2.0 / sub 4.5,  Vice full 1.0 / sub 2.0
  * membership tier weights:  Highly 3.0 / Moderately 1.8 / Slightly 0.9   (subcommittee = base/6)
  * committee score = (Membership + Leadership) / 2
Default = VALIDATE (read-only): recompute from the live workbook and confirm it reproduces B/C/D for every row.
--apply = also rewrite E (SB->Highly) + B/C/D via Excel COM + full recalc (idempotent on the current workbook)."""
import openpyxl, re, sys, shutil, os, datetime, time
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
B = MB_ROOT
WB = B + r"\Moneyball Main Version - Blended.xlsx"
BASE = {'Highly': 3.0, 'Moderately': 1.8, 'Slightly': 0.9}; SUBF = 1.0 / 6.0
KLEAD = {'Chair': {'full': 3.0, 'sub': 7.5}, 'Ranking': {'full': 2.0, 'sub': 4.5}, 'Vice': {'full': 1.0, 'sub': 2.0}}
BRK = re.compile(r'\[([^\]]+)\]\s*$')

def parse(seat):
    s = seat.strip(); m = BRK.search(s)
    if not m: return None
    inside = m.group(1); name = s[:m.start()].rstrip()
    level = 'sub' if ' - ' in name else 'full'; parts = [p.strip() for p in inside.split('/')]
    return name, level, parts[0], (parts[1] if len(parts) > 1 else None), m.start()

def compute(seats_text):
    mem = lead = 0.0; out = []
    for seat in str(seats_text).split(' | '):
        if not seat.strip(): continue
        p = parse(seat)
        if not p: out.append(seat); continue
        name, level, tier, role, brk = p
        is_sb = 'Small Business' in name
        mw = BASE['Highly' if is_sb else tier] * (1.0 if level == 'full' else SUBF); mem += mw
        if role: lead += mw * KLEAD[role][level]
        out.append(seat[:brk].rstrip() + ' [' + ('Highly' + (('/' + role) if role else '')) + ']' if is_sb else seat)
    return round(mem, 6), round(lead, 6), round((mem + lead) / 2, 6), ' | '.join(out)

wb = openpyxl.load_workbook(WB, data_only=True); cd = wb["Committee Data"]
last = 1
for r in range(2, cd.max_row + 1):
    if cd.cell(r, 1).value: last = r
recs = [(r, cd.cell(r, 1).value, cd.cell(r, 2).value, cd.cell(r, 3).value, cd.cell(r, 4).value, cd.cell(r, 5).value)
        for r in range(2, last + 1) if cd.cell(r, 1).value]
wb.close()

ok = 0; miss = []; out_rows = {}
for r, nm, Bv, Cv, Dv, Ev in recs:
    if Ev is None: out_rows[r] = (Bv, Cv, Dv, Ev); continue
    nmem, nlead, nB, nE = compute(Ev); out_rows[r] = (nB, nmem, nlead, nE)
    if abs(nB - (Bv or 0)) < 1e-6 and abs(nmem - (Cv or 0)) < 1e-6 and abs(nlead - (Dv or 0)) < 1e-6:
        ok += 1
    else:
        miss.append((nm, (Cv, Dv, Bv), (nmem, nlead, nB)))
n = sum(1 for r, nm, Bv, Cv, Dv, Ev in recs if Ev is not None)
print("VALIDATE: recomputed committee score reproduces %d/%d rows" % (ok, n))
for nm, old, new in miss[:20]:
    print("  MISS %-24s stored(mem,lead,score)=%s  recomputed(mem,lead,score)=%s" % (nm, old, new))

if '--apply' in sys.argv:
    if miss: print("apply ABORTED: %d mismatches -- investigate before writing" % len(miss)); sys.exit(1)
    import win32com.client as win32
    bk = WB.replace('.xlsx', '.PRE-CMTEREWEIGHT-%s.xlsx' % datetime.datetime.now().strftime('%Y%m%d_%H%M%S'))
    shutil.copy2(WB, bk); print('backup ->', os.path.basename(bk))
    arr = tuple(tuple(float(v) if isinstance(v, (int, float)) else v for v in out_rows[r]) for r in range(2, last + 1))
    xl = win32.DispatchEx('Excel.Application'); xl.Visible = False; xl.DisplayAlerts = False
    try:
        w = xl.Workbooks.Open(WB); ws = w.Worksheets('Committee Data')
        ws.Range("B2:E%d" % last).Value = arr; xl.CalculateFull()
        saved = False
        for a in range(6):
            try: w.Save(); saved = True; break
            except Exception as e: print('save retry %d: %s' % (a, str(e)[:70])); time.sleep(3)
        w.Close(SaveChanges=False); print('APPLIED + recalculated' if saved else 'NOT SAVED')
    finally:
        xl.Quit()
else:
    print("(validate-only; pass --apply to write via COM)")
