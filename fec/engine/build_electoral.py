import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
from moneyball_core import dump_atomic   # S2-03: crash-safe artifact write
import roster
import openpyxl,json,sys,glob,os,collections
from members.member_identity import Resolver
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
B=MB_ROOT
LELEV=float(os.environ.get('LELEV','0.8'))   # electoral-leverage strength (competitive seats are cheaper to move)
members=[m['member'] for m in json.load(open(os.path.join(B, 'data', 'engine', 'congestion.json')))]
_resolver=Resolver(members)
def match(nm): return _resolver.resolve(nm)
_bpc=glob.glob(os.path.join(B, 'Ballotpedia_119th_CU_Influence_*.xlsx'))+glob.glob(os.path.join(B, 'archive', 'Ballotpedia_119th_CU_Influence_*.xlsx'))
bp=sorted(_bpc,key=os.path.basename)[-1]   # newest by filename timestamp (file may live in archive/ after cleanup)
ws=openpyxl.load_workbook(bp,data_only=True,read_only=True)['Members']
hdr={ws.cell(1,c).value:c for c in range(1,ws.max_column+1)}
ec=hdr['Election_Average']; mc=hdr['Member']
elavg={}
for r in range(2,ws.max_row+1):
    nm=ws.cell(r,mc).value; v=ws.cell(r,ec).value
    if nm is None: continue
    mm=match(nm)
    if mm and isinstance(v,(int,float)): elavg[mm]=v
# security percentile (higher Election_Average -> safer -> lower leverage)
vals=sorted(elavg.values()); n=len(vals)
def pct(v):
    lo=sum(1 for x in vals if x<v); return lo/max(1,n-1)
elec={}
for m in members:
    if m in elavg:
        sec=pct(elavg[m]); lev=round(1+LELEV*(1-sec),3); elec[m]={'elavg':round(elavg[m],1),'security':round(sec,3),'leverage':lev}
    else:
        elec[m]={'elavg':None,'security':0.5,'leverage':round(1+LELEV*0.5,3)}   # unmatched -> neutral
dump_atomic(elec,os.path.join(B, 'data', 'money', 'electoral.json'))
matched=sum(1 for m in members if elec[m]['elavg'] is not None)
print('electoral.json: %d/%d matched, LELEV=%.2f'%(matched,roster.count(),LELEV))
hi=sorted([m for m in members if elec[m]['elavg'] is not None],key=lambda m:-elec[m]['leverage'])[:6]
lo=sorted([m for m in members if elec[m]['elavg'] is not None],key=lambda m:elec[m]['leverage'])[:6]
print('MOST leverage (competitive):',[(m.split()[-1],elec[m]['elavg'],'x%.2f'%elec[m]['leverage']) for m in hi])
print('LEAST leverage (safe):',[(m.split()[-1],elec[m]['elavg'],'x%.2f'%elec[m]['leverage']) for m in lo])
