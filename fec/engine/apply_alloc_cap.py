# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
"""Capped, graduated allocator: distribute the Pot ($AB$6) across the top-N (=$AB$7) ranked members,
proportional to Score but with a hard per-member Cap ($AB$8). Water-fill: alloc = MIN(cap, score*lambda),
lambda solved so the funded allocations sum to the pot (capped members give their excess to the rest).
Adds the Cap + fill-multiplier as labelled parameter cells and rewrites the Allocation (col U) formula, so
it stays live; re-run this if Score or N/Cap change (it re-solves lambda). COM + recalc. Backs up first."""
import openpyxl, sys, shutil, os, datetime
import win32com.client as win32
import roster
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
B  = MB_ROOT
WB = B + r'\Moneyball Main Version - Blended.xlsx'
POT = 2_000_000          # $AB$6 Pot
CAP = 10_000             # $AB$8 Cap ($) per member   <-- edit here
N   = 250                # $AB$7 Fund top N           <-- edit here

# ---- read top-N scores (col19=Score, col20=Rank) ----
wv = openpyxl.load_workbook(WB, data_only=True); bm = wv['Blended Model']
rows = []
for r in roster.rows('Blended Model'):
    nm = bm.cell(r, 1).value; sc = bm.cell(r, 19).value; rk = bm.cell(r, 20).value
    if nm and isinstance(sc, (int, float)) and isinstance(rk, (int, float)):
        rows.append((rk, sc))
wv.close()
rows.sort()
top = [sc for rk, sc in rows if rk <= N]
assert POT <= CAP * len(top) + 1, 'pot exceeds N*cap; raise N or Cap'

# ---- solve lambda: sum(min(CAP, s*lam)) = POT ----
lo, hi = 0.0, 1e7
for _ in range(200):
    mid = (lo + hi) / 2
    if sum(min(CAP, s * mid) for s in top) < POT: lo = mid
    else: hi = mid
lam = (lo + hi) / 2
alloc = [min(CAP, s * lam) for s in top]
ncap = sum(1 for a in alloc if a >= CAP - 1e-6)
print('lambda=%.6f  funded=%d  total=$%.2f  capped@$%d=%d  min=$%.0f  max=$%.0f'
      % (lam, len(top), sum(alloc), CAP, ncap, min(alloc), max(alloc)))

# ---- backup + COM write ----
bk = B + r'\backups\Moneyball Main Version - Blended.PRE-ALLOCCAP-%s.xlsx' % datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
shutil.copy2(WB, bk); print('backup ->', os.path.basename(bk))
xl = win32.DispatchEx('Excel.Application'); xl.Visible = False; xl.DisplayAlerts = False
try:
    try: xl.Calculation = -4105   # xlCalculationAutomatic
    except Exception: pass
    wb = xl.Workbooks.Open(WB); ws = wb.Worksheets('Blended Model')
    ws.Cells(7, 28).Value = N                                   # AB7  Fund top N
    ws.Cells(8, 27).Value = 'Cap ($)';        ws.Cells(8, 28).Value = CAP    # AA8 / AB8
    ws.Cells(14, 27).Value = 'Alloc fill mult'; ws.Cells(14, 28).Value = lam  # AA14 / AB14
    ws.Range(roster.col_range('U', 'Blended Model', absolute=False)).Formula = '=IF(T3<=$AB$7,MIN($AB$8,S3*$AB$14),0)'
    try:
        tw = wb.Worksheets('Topic Weights'); tw.Cells(5,3).Value = 0.3   # cosmetic: TOP-tier 'Press wt' was 0, inconsistent with its own detailed dials (0.3). Model uses hardcoded 0.70/0.35 (per_category_stance.py), NOT this sheet.
        print('Topic Weights C5 (TOP press wt) 0 -> 0.3 (cosmetic; sheet is reference-only)')
    except Exception as e:
        print('Topic Weights fix skipped:', e)
    xl.CalculateFull()
    wb.Save(); wb.Close(SaveChanges=True); print('SAVED + recalculated.')
finally:
    xl.Quit()

# ---- verify from the saved workbook ----
wv = openpyxl.load_workbook(WB, data_only=True); bm = wv['Blended Model']
av = sorted([bm.cell(r, 21).value for r in roster.rows('Blended Model')
             if isinstance(bm.cell(r, 21).value, (int, float)) and bm.cell(r, 21).value > 0], reverse=True)
print('VERIFY: funded=%d  total=$%.0f  max=$%.0f  min=$%.0f  at_cap=%d'
      % (len(av), sum(av), av[0], av[-1], sum(1 for a in av if a >= CAP - 0.5)))
wv.close()
