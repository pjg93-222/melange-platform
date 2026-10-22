import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
import openpyxl, sys, json, collections, math, unicodedata, hashlib, os
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
M=os.path.join(MB_ROOT, 'Moneyball Main Version - Blended.xlsx')
wb=openpyxl.load_workbook(M,data_only=True)
ra=wb['Rank & Allocation']
import moneyball_core as MC
import roster
import stance.stance_rules as SR
from engine.pillar_policy import LABEL_EVIDENCE_WEIGHTS
_normM,_w535,_=MC.press_matrix()   # PHASE-2: Python-owned press normalization (raw values + press footer-row dial literals; no recalc dependency)
# S2-04 (audit run-2, user-ratified): col G is NOT read here anymore. It was the previous run's
# sync_stance_g output, which made stance_N+1=F(stance_N) — a cross-run recurrence that orbited a
# period-2 limit cycle (~39 persuadable-middle cells flipping per run). The overall-tier feedback
# is now resolved IN-RUN from a canonical start (see the gen iteration below); col G stays a view.
def nfc(value):
    return unicodedata.normalize('NFC',str(value)) if value is not None else value
members=[nfc(ra.cell(r,1).value) for r in roster.rows('Rank & Allocation') if ra.cell(r,1).value]
with open(os.path.join(MB_ROOT, 'archive_leg.json'),encoding='utf-8') as _handle:
    LEG={nfc(member):values for member,values in json.load(_handle).items()}
with open(os.path.join(MB_ROOT, 'legislation_routing_audit.json'),encoding='utf-8') as _handle:
    _routing=json.load(_handle)
# Equal positive and negative contributions are measured zero, not absent evidence.
LEG_SIGNAL={(nfc(row['member']),row['effective_category']) for row in _routing['rows']
            if row.get('disposition')=='selected'}
with open(os.path.join(MB_ROOT, 'vote_category_lean.json'),encoding='utf-8') as _handle:
    VC=json.load(_handle)
n2b={nfc(v):k for k,v in VC['name_of'].items()}; votes=VC['votes']
cats=[('Examination & Supervision',5),('CU Capital Requirements',6),('NCUA Budget & Insurance',7),('Interchange',8),
('Payments',9),('Technology & Emerging',10),('Fraud Prevention',11),('Data Privacy & Cyber',12),('Charter Modernization',13),
('Defense CU Issues',14),('Field of Membership',15),('CU Tax Status',16),('CDFI Fund',17),('CFPB',19),('Credit Card Mandates',20),
('TCPA',21),('Member Business Lending',22),('Mortgage & Housing',23)]
# deep-audit fix 2026-07-02: 'CU Tax Status' (Normalized col 16) was added to category_stance.json by hand
# on 2026-07-01 but never to THIS producer list — the next regeneration silently dropped the 18th category
# while the ABB26 IMP dicts (weight 3.0!) kept expecting it. Data now exists (190 press rows + social).
# Paul 2026-07-15: CU Tax Status IS a defend-type fight -> added to DEFENSIVE (weighting is the user's call, now made).
DEFENSIVE={'Interchange','Credit Card Mandates','CU Tax Status'}
def weights(cat):
    policy = LABEL_EVIDENCE_WEIGHTS[
        'defensive' if cat in DEFENSIVE else 'standard'
    ]
    return policy['legislation'], policy['media'], policy['vote']
pw=[_w535[pc] for _,pc in cats]
def clip(x): return max(-1.0,min(1.0,x))
_normM={column:{nfc(member):value for member,value in values.items()}
        for column,values in _normM.items()}
press={(m,lab):_normM[pc].get(m,0.0) for m in members for lab,pc in cats}
opl={m:sum(press[(m,lab)]*pw[i] for i,(lab,pc) in enumerate(cats)) for m in members}
neg=sorted(v for v in opl.values() if v<0); ascale=abs(neg[int(0.08*len(neg))]) if neg else 1.0
ANTI_STANCE={'Opponent':1.0,'Skeptic':0.5}   # floors keyed on the OVERALL ladder — resolved in-run below (S2-04)
anti0={m:clip(-opl[m]/(ascale or 1)) for m in members}   # data-driven base; overall-tier floors join inside the gen iteration
PRIOR=0.70
def classify(anti):
    """Score cells, then curve labels only in the declared saturated domains."""
    raw_scores=collections.defaultdict(dict); signal=collections.defaultdict(dict)
    for lab,pc in cats:
        WL,WP,WV=weights(lab)
        legraw={m:LEG.get(m,{}).get(lab,0.0) for m in members}; prraw={m:press[(m,lab)] for m in members}
        legc={m:math.copysign(math.sqrt(abs(legraw[m])),legraw[m]) for m in members}   # sqrt-compress heavy tail -> reduce ceiling saturation
        for m in members:
            has_leg=(m,lab) in LEG_SIGNAL; has_press=abs(prraw[m])>1e-9
            ln=clip(legc[m]/SR.LEGISLATION_SQRT_SCALE); pn=clip(prraw[m]/SR.MEDIA_SCALE)
            vinfo=votes.get(n2b.get(m,''),{}).get(lab); vn=vinfo['n'] if vinfo else 0; hv=bool(vinfo) and vn>=5
            c=SR.combine_evidence(
                {'legislation':ln if has_leg else None,'media':pn if has_press else None,
                 'vote':clip(vinfo['lean']) if hv else None},
                {'legislation':WL,'media':WP,'vote':WV})
            if c>0 and prraw[m]<=0.02 and anti[m]>0: c*=(1-PRIOR*anti[m])
            raw_scores[m][lab]=clip(c)
            signal[m][lab]=has_leg or has_press or hv
    matrix=collections.defaultdict(dict); scores=collections.defaultdict(dict)
    raw_out=collections.defaultdict(dict); rounding=collections.defaultdict(dict)
    for m in members:
        adjusted,factors=SR.exceptional_rounding(raw_scores[m],signal[m])
        for lab,_ in cats:
            raw_out[m][lab]=round(raw_scores[m][lab],6)
            rounding[m][lab]=round(factors[lab],6)
            scores[m][lab]=round(adjusted[lab],3)
    label_scores=SR.curve_label_scores(
        scores,signal,SR.CURVED_CATEGORIES)
    for m in members:
        for lab,_ in cats:
            matrix[m][lab]=SR.classify(
                label_scores[m][lab],signal[m][lab],lab in DEFENSIVE)
    return matrix,scores,raw_out,rounding,label_scores,signal
# ---- S2-04 gen iteration (user-ratified "whatever optimizes it the most"): resolve the overall-tier
# feedback IN-RUN from a canonical start (gen={}), instead of consuming last run's col G. The overall
# ladder is MC.overall_tier — the exact law sync_stance_g writes to col G afterwards. Fixed point ->
# self-consistent matrix; detected period-2 cycle -> UNION anti floor over both phases + one final
# pass (deterministic damping). Same data therefore ALWAYS yields the same matrix: pipeline idempotent.
gen_in={}; prev_in=None
for _it in range(1,9):
    anti={m:max(anti0[m],ANTI_STANCE.get(gen_in.get(m),0.0)) for m in members}
    matrix,scores,raw_scores,rounding,label_scores,signal=classify(anti)
    gen_out={m:MC.overall_tier(matrix[m]) for m in members}
    if gen_out==gen_in:
        print('gen iteration: self-consistent fixed point at iteration %d'%_it); break
    if prev_in is not None and gen_out==prev_in:
        anti={m:max(anti0[m],ANTI_STANCE.get(gen_in.get(m),0.0),ANTI_STANCE.get(gen_out.get(m),0.0)) for m in members}
        matrix,scores,raw_scores,rounding,label_scores,signal=classify(anti)
        print('gen iteration: period-2 cycle at iteration %d -> union-floor damped final pass (S2-04)'%_it); break
    prev_in,gen_in=gen_in,gen_out
else:
    anti={m:max(anti0[m],ANTI_STANCE.get(gen_in.get(m),0.0),ANTI_STANCE.get(prev_in.get(m),0.0)) for m in members}
    matrix,scores,raw_scores,rounding,label_scores,signal=classify(anti)
    print('gen iteration: no fixed point in 8 iterations -> union-floor damped final pass (S2-04)')
genF={m:MC.overall_tier(matrix[m]) for m in members}
# ---- SAY-DO corroboration gate: recorded action (A/V) + Rule C 'proven' badge,
# applied AFTER the S2-04 fixed point converges, as a pure OUTPUT transform. Safe by construction:
# the cell cap only lowers Champion/Defender toward Ally (+2), which can never create the Opponent/
# Skeptic overalls the ANTI_STANCE floors key on, so convergence and idempotence are untouched.
# Raw tiers preserved in the 'gate' block; single law lives in MC.cap_uncorroborated.
raw_matrix={m:dict(matrix[m]) for m in members}
gate_capped={}; gate_proven={}; gate_corrob={}
for m in members:
    matrix[m],_cap,_prov,_cor=MC.cap_uncorroborated(m,matrix[m])
    if _cap: gate_capped[m]=_cap
    if _prov: gate_proven[m]=_prov
    if _cor: gate_corrob[m]=_cor
gate_overall={}
for m in members:
    _final,_detail=MC.overall_say_do_gate(m,matrix[m],raw_matrix[m],gate_corrob.get(m,{}))
    if _detail['raw']!=_final: gate_overall[m]=_detail
_cbp=collections.Counter(pth for cc in gate_corrob.values() for pth in cc.values())
print('say-do gate: %d cells capped across %d members | %d high-score crown candidates capped | '
      '%d proven badges | corroborated A%d V%d'%(
      sum(len(v) for v in gate_capped.values()),len(gate_capped),len(gate_overall),
      len(gate_proven),_cbp['A'],_cbp['V']))
cnt=collections.Counter(matrix[m][lab] for m in members for lab,_ in cats)
_OUT = _o.environ.get('CATEGORY_STANCE_OUT') or os.path.join(MB_ROOT, 'category_stance.json')
_source_paths=('archive_leg.json','legislation_routing_audit.json','vote_category_lean.json',
               'bill_repository_live.json')
_source_hashes={path:hashlib.sha256(open(os.path.join(MB_ROOT,path),'rb').read()).hexdigest()
                for path in _source_paths}
MC.dump_atomic({'cats':[c[0] for c in cats],'defensive':sorted(DEFENSIVE),'matrix':{m:matrix[m] for m in members},
           'source_contract':{'schema':'adjudicated-label-inputs-v1','inputs_sha256':_source_hashes,
                              'corroboration':'positive selected sponsor/original rows; current official advancement',
                              'zero_legislation':'observed cancellation retains evidence weight'},
           'scores':{m:scores[m] for m in members},
           'signal':{m:signal[m] for m in members},
           'raw_scores':{m:raw_scores[m] for m in members},
           'curve':{
               'categories':sorted(SR.CURVED_CATEGORIES),
               'method':SR.thresholds()['domain_label_curves']['method'],
               'label_scores':{
                   m:{cat:round(label_scores[m][cat],6)
                      for cat in sorted(SR.CURVED_CATEGORIES)}
                   for m in members},
               'raw_and_adjusted_scores_unchanged':True,
               'empty_domains_inherit':False},
           'rounding':{'rule':SR.thresholds()['exceptional_rounding'],
                       'factors':{m:rounding[m] for m in members},
                       'overall_factor':{m:round(MC.overall_rounding(matrix[m]),6) for m in members}},
           'thresholds':{**SR.thresholds(),'overall':MC.overall_thresholds(),
                         'evidence_weights':LABEL_EVIDENCE_WEIGHTS,
                         'press_policy':'active_diminished; may affect issue and overall labels'},
            'media':{'general_credit_union_exclusions':MC.press_general_exclusions()},
           'gate':{'rule':'A Champion/Defender cell keeps its tier only if the member (A) authored/orig-'
                          'cosponsored an in-category bill or (V) has a recorded CU-aligned final-passage '
                          'vote in it; committee office is not action. A legacy high-score candidate for '
                          'the aggregate say-do provenance audit requires A/V evidence in at least '
                          'two-thirds of its raw top-tier cells and '
                          'at least six raw top-tier domains; it also caps a 10+ direct-sponsorship paper '
                          'trail when no more than 20% advanced, unless 20+ direct sponsored bills '
                          'corroborate every raw top-tier domain. Cells '
                          'and legacy high-score candidates otherwise cap to Ally; raw values are preserved '
                          'here. The overall Champion display crown is independent: every Legislative+Press '
                          'Tier-1 member receives it. proven = authored bill in category advanced past introduction.',
                   'overall_delivery_threshold':round(MC.SAY_DO_OVERALL_RATIO,6),
                   'overall_minimum_top_tier_cells':MC.SAY_DO_MIN_TOP_CELLS,
                   'high_volume_followthrough':{'direct_sponsored_bills_at_least':MC.SAY_DO_DIRECT_VOLUME_MIN,
                                                'advancement_rate_at_most':MC.SAY_DO_MAX_ADVANCEMENT_RATE,
                                                'exceptional_authorship_bills_at_least':MC.SAY_DO_EXCEPTIONAL_AUTHORSHIP_MIN,
                                                'exception_requires_full_top_domain_coverage':True},
                   'capped':gate_capped,'overall_capped':gate_overall,
                   'proven':gate_proven,'corrob':gate_corrob}},_OUT)
print('tier distribution:',dict(cnt))
persu=sum(cnt[t] for t in ('Highly Persuadable','Persuadable','Lean Persuadable','True Neutral','Lean Skeptic','Skeptic','Highly Skeptical'))
print('middle tiers: HP %d | P %d | LP %d | TN %d | LS %d | S %d | HS %d (total %d) | Champions+Defenders %d'%(
      cnt['Highly Persuadable'],cnt['Persuadable'],cnt['Lean Persuadable'],cnt['True Neutral'],
      cnt['Lean Skeptic'],cnt['Skeptic'],cnt['Highly Skeptical'],persu,cnt['Champion']+cnt['Defender']))
print('\nLABEL thresholds and declared curved domains:',SR.thresholds())
def prof(m): return ([c for c,_ in cats if matrix[m][c] in ('Champion','Defender')],[c for c,_ in cats if matrix[m][c]=='Opponent'])
print('\nVALIDATION:')
for who in ['Richard Durbin','Tim Scott','Andy Barr','Maxine Waters']:
    if who in matrix:
        ch,op=prof(who); print('   %-15s [%s] champ/def=%s opp=%s'%(who,genF.get(who),[c.split(' &')[0][:8] for c in ch],[c.split(' &')[0] for c in op]))
print('Interchange DEFENDERS:',[m.split()[-1] for m in members if matrix[m]['Interchange']=='Defender'])
