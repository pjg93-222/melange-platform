"""Artifact-authored recognition views, packaged without rewriting model cells.

HONORS_WORKBOOK_SCOPE=legacy restores the pre-2026-09-18 behaviour: the
publication binding is a hash of the WHOLE workbook file and
--prepare-model-update replaces the whole file with the recorded
pre-publication package, so an edit to any sheet -- Committee Data, for
instance -- refuses the rollback and blocks the chain. The default,
HONORS_WORKBOOK_SCOPE=region, binds and reverts only the region a
publication actually writes: the SHEETS recognition views, the
LABEL_COLUMNS label columns it hides and the HIDDEN_SHEETS tabs it
collapses. Edits to the honors views themselves still refuse, by name.

A publication may then land on top of unrelated edits, so its rollback copy no
longer restores the whole file losslessly; the binding records that as
original_exact and such a binding is only ever reverted by region.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
from pathlib import Path
import posixpath
import re
import subprocess
import sys
import tempfile
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from paths import ROOT
from stance.workbook_cell_sync import sheet_paths
from xlsx_openpyxl_authors import AUTHOR_MODULE, author_honors_views, xlsx_writer

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
Q = '{' + NS + '}'
SHEETS = ('Field Recognition', 'Recognition Detail', 'Overall Standing')
# The only existing cells a publication writes: it hides the analytical label
# column (1-based) on these sheets without touching a value, formula or cache.
LABEL_COLUMNS = (('Rank & Allocation', 7), ('Legislative Agenda', 3))
# The analytical tabs a publication collapses behind the public views.
HIDDEN_SHEETS = ('Category Stance', 'Coalition Profiles')
REGION_SCHEMA = 'honors-region-v1'
# 'region' binds only what a publication writes; 'legacy' binds the whole file.
HONORS_WORKBOOK_SCOPE = os.environ.get('HONORS_WORKBOOK_SCOPE', 'region')


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def workbook_views(honors):
    """Rectangular public views; all recognition math belongs to the producer."""
    cats = honors['cats']
    policy = honors['policy']
    note = ('Recorded field actions through %s. Champion top %d%%; Ally top %d%%; '
            'Highly Persuadable top %d%%. Equal merit stays together.' % (
                honors['as_of'], policy['champion_pct'], policy['ally_pct'], policy['highly_persuadable_pct']))
    overview = [['Melange — Field Recognition'], [note],
                ['Generation', honors['generation_id']],
                ['Member', 'Overall priority tier', 'Primacy rank', 'Stance',
                 'Rank basis', 'Observed action record'] + cats]
    detail = [['Melange — Recognition Detail'], [note],
              ['Generation', honors['generation_id']],
              ['Member', 'Field', 'Recognition', 'Action record', 'Rank from', 'Rank to',
               'Eligible members', 'Actions', 'Raw action score', 'Tensor tie-break',
               'Champion cutoff', 'Ally cutoff', 'Highly Persuadable cutoff', 'Reason']]
    overall_detail = [['Melange — Overall Standing'],
        ['Labels follow primacy. Repeated opposing leadership is not canceled by unrelated positive activity.'],
        ['Generation', honors['generation_id']],
        ['Member', 'Overall priority tier', 'Primacy rank', 'Stance', 'Rank basis',
         'Observed action record', 'Raw action score', 'Actions', 'Supportive actions', 'Opposing actions',
         'Neutral actions', 'Observed members', 'Ally cutoff', 'Highly Persuadable cutoff', 'Persuadable cutoff',
         'Opposing measures', 'Opposing lead measures', 'Members with opposing acts', 'Leadership Congresses',
         'Lean Persuadable cutoff', 'Opposition classification', 'Tensor diagnostic', 'Tensor coverage', 'Reason']]
    for name, record in sorted(honors['members'].items(), key=lambda item: item[1]['overall']['rank'] or 9999):
        overall = record['overall']
        priority = ('Tier 1 · Champion' if overall['priority_tier'] == 1 else
                    'Tier %d' % overall['priority_tier'] if overall['priority_tier'] else 'Outside priority tiers')
        start, end = overall['stance_rank_start'], overall['stance_rank_end']
        action_rank = ('—' if start is None else '%s of %s' % (
            str(start) if start == end else '%d–%d' % (start, end), overall['stance_cohort_size']))
        action_pattern = 'Balanced actions' if overall['record_pattern'] == 'balanced' else overall['alignment']
        row = [name, priority, overall['rank'], overall['stance'], 'Primacy', action_pattern]
        adverse = overall['opposition_evidence']
        cuts, opposition = overall['stance_cutoffs'], overall['opposition_cutoffs']
        overall_detail.append(row[:] + [float(overall['raw_overall_score']) if overall['raw_overall_score'] is not None else None,
            overall['action_count'], overall['positive_action_count'], overall['negative_action_count'],
            overall['neutral_action_count'], overall['stance_cohort_size'], cuts['ally'], cuts['highly_persuadable'],
            cuts['persuadable'], len(adverse['measures']), len(adverse['lead_measures']),
            overall['opposition_cohort_size'], len(adverse['congresses']), cuts['lean_persuadable'], adverse['label'],
            float(overall['tensor_tiebreak']) if overall['tensor_tiebreak'] is not None else None,
            overall['tensor_status'], overall['stance_reason']])
        for cat in cats:
            cell = record['fields'][cat]
            label = cell.get('recognition') or (
                'No recorded action' if cell['evidence_status'] != 'observed' else 'Below threshold')
            row.append(label)
            cuts = cell['cutoffs']
            detail.append([name, cat, cell.get('recognition') or '—', cell['alignment'],
                           cell.get('rank_start'), cell.get('rank_end'), cell['cohort_size'],
                           cell['action_count'], float(cell['raw_field_score']) if cell.get('raw_field_score') is not None else None,
                           float(cell['tensor_tiebreak']) if cell.get('tensor_tiebreak') is not None else None,
                           cuts['champion'], cuts['ally'], cuts['highly_persuadable'], cell['reason']])
        overview.append(row)
    def rectangular(rows):
        width = max(map(len, rows))
        return [row + [None] * (width - len(row)) for row in rows]
    return [{'name': SHEETS[0], 'values': rectangular(overview)},
            {'name': SHEETS[1], 'values': rectangular(detail)},
            {'name': SHEETS[2], 'values': rectangular(overall_detail)}]


def merge_views(original, authored, destination):
    """Append only authored views/styles; keep all existing model worksheet bytes."""
    with ZipFile(original) as src, ZipFile(authored) as made:
        original_paths = sheet_paths(src)
        authored_paths = sheet_paths(made)
        if not set(SHEETS) <= set(authored_paths):
            raise ValueError('Authored workbook is missing recognition views')
        book = ET.fromstring(src.read('xl/workbook.xml'))
        sheets = book.find(Q + 'sheets')
        rels = ET.fromstring(src.read('xl/_rels/workbook.xml.rels'))
        types = ET.fromstring(src.read('[Content_Types].xml'))
        styles = ET.fromstring(src.read('xl/styles.xml'))
        extra_styles = ET.fromstring(made.read('xl/styles.xml'))
        strings = list(ET.fromstring(made.read('xl/sharedStrings.xml'))) if 'xl/sharedStrings.xml' in made.namelist() else []
        # Keep every original style index valid. Append the artifact's style
        # components and map only the newly authored cells to appended indexes.
        def append_unique(parent, entry):
            key = ET.tostring(entry)
            for index, existing in enumerate(parent):
                if ET.tostring(existing) == key:
                    return index
            parent.append(entry)
            return len(parent) - 1
        component_ids = {}
        for tag in ('fonts', 'fills', 'borders'):
            old = styles.find(Q + tag)
            new = extra_styles.find(Q + tag)
            component_ids[tag] = {}
            if new is not None:
                for index, entry in enumerate(new):
                    component_ids[tag][index] = append_unique(old, copy.deepcopy(entry))
            old.set('count', str(len(old)))
        old_formats = styles.find(Q + 'numFmts')
        if old_formats is None:
            old_formats = ET.Element(Q + 'numFmts', count='0')
            styles.insert(0, old_formats)
        next_format = max([163] + [int(x.get('numFmtId')) for x in old_formats]) + 1
        fmt_map = {}
        authored_formats = extra_styles.find(Q + 'numFmts')
        for fmt in authored_formats if authored_formats is not None else []:
            copied = copy.deepcopy(fmt)
            existing = next((f for f in old_formats if f.get('formatCode') == fmt.get('formatCode')), None)
            if existing is not None:
                fmt_map[int(fmt.get('numFmtId'))] = int(existing.get('numFmtId'))
                continue
            fmt_map[int(fmt.get('numFmtId'))] = next_format
            copied.set('numFmtId', str(next_format))
            old_formats.append(copied)
            next_format += 1
        old_formats.set('count', str(len(old_formats)))
        xfs = styles.find(Q + 'cellXfs')
        xf_ids = {}
        for index, xf in enumerate(extra_styles.find(Q + 'cellXfs')):
            copied = copy.deepcopy(xf)
            for attr, group in (('fontId', 'fonts'), ('fillId', 'fills'), ('borderId', 'borders')):
                copied.set(attr, str(component_ids[group][int(copied.get(attr, '0'))]))
            copied.set('xfId', '0')
            number_format = int(copied.get('numFmtId', '0'))
            if number_format in fmt_map:
                copied.set('numFmtId', str(fmt_map[number_format]))
            xf_ids[index] = append_unique(xfs, copied)
        xfs.set('count', str(len(xfs)))
        replacements = {'xl/styles.xml': ET.tostring(styles, encoding='utf-8')}
        next_sheet = max(int(s.get('sheetId')) for s in sheets) + 1
        used_relationships = {r.get('Id') for r in rels}
        for index, name in enumerate(SHEETS):
            existing = next((s for s in sheets if s.get('name') == name), None)
            if existing is not None:
                target = original_paths[name]
            else:
                target = 'xl/worksheets/melange_recognition_%d.xml' % index
                if target in src.namelist():
                    raise ValueError('Recognition worksheet target already belongs to another sheet')
                rid = 'rIdMelangeRecognition%d' % index
                if rid in used_relationships:
                    raise ValueError('Recognition relationship ID collision')
                ET.SubElement(rels, '{' + PKG + '}Relationship',
                              Id=rid, Type=REL + '/worksheet', Target=posixpath.relpath(target, 'xl'))
                existing = ET.Element(Q + 'sheet', name=name, sheetId=str(next_sheet))
                existing.set('{' + REL + '}id', rid)
                sheets.insert(index, existing)
                # A sheet-local defined name addresses its sheet by POSITION (see
                # remove_views): shift every one at or after the insertion point.
                defined = book.find(Q + 'definedNames')
                for entry in list(defined) if defined is not None else []:
                    local = entry.get('localSheetId')
                    if local is not None and int(local) >= index:
                        entry.set('localSheetId', str(int(local) + 1))
                next_sheet += 1
                ET.SubElement(types, '{' + CT + '}Override', PartName='/' + target,
                              ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
            tree = ET.fromstring(made.read(authored_paths[name]))
            # Authoring uses values and styles only; external sheet resources
            # must never be silently lost at this packaging boundary.
            if any('{' + REL + '}id' in node.attrib for node in tree.iter()):
                raise ValueError('Recognition sheet contains an unsupported relationship')
            for cell in tree.iter(Q + 'c'):
                cell.set('s', str(xf_ids[int(cell.get('s', '0'))]))
                if cell.get('t') == 's':
                    value = cell.find(Q + 'v')
                    item = strings[int(value.text)]
                    cell.remove(value)
                    inline = ET.SubElement(cell, Q + 'is')
                    for child in item:
                        inline.append(copy.deepcopy(child))
                    cell.set('t', 'inlineStr')
            replacements[target] = ET.tostring(tree, encoding='utf-8')
        # The analytical matrix remains available to workbook formulas. Its tab
        # is internal; the public recognition sheets are placed first.
        for sheet in sheets:
            if sheet.get('name') in HIDDEN_SHEETS:
                sheet.set('state', 'hidden')
        for analytical_sheet, label_column in LABEL_COLUMNS:
            if analytical_sheet not in original_paths:
                continue
            # Hide the old analytical label column without changing its values,
            # formulas, cached results or any other cell in the model sheet.
            rank_path = original_paths[analytical_sheet]
            raw = src.read(rank_path)
            rank_tree = ET.fromstring(raw)
            old_cols = rank_tree.find(Q + 'cols')
            cols = ET.Element(Q + 'cols')
            found = False
            for col in old_cols if old_cols is not None else []:
                lo, hi = int(col.get('min')), int(col.get('max'))
                if not lo <= label_column <= hi:
                    cols.append(copy.deepcopy(col))
                    continue
                found = True
                if lo < label_column:
                    left = copy.deepcopy(col); left.set('max', str(label_column-1)); cols.append(left)
                middle = copy.deepcopy(col); middle.set('min', str(label_column)); middle.set('max', str(label_column)); middle.set('hidden', '1'); cols.append(middle)
                if hi > label_column:
                    right = copy.deepcopy(col); right.set('min', str(label_column+1)); cols.append(right)
            if not found:
                ET.SubElement(cols, Q + 'col', min=str(label_column), max=str(label_column), hidden='1')
            cols[:] = sorted(cols, key=lambda c: int(c.get('min')))
            cols_xml = ET.tostring(cols, encoding='utf-8')
            if old_cols is not None:
                raw = re.sub(rb'<(?:\w+:)?cols\b[^>]*>.*?</(?:\w+:)?cols>', lambda _: cols_xml, raw, count=1, flags=re.S)
            else:
                raw = re.sub(rb'<(?:\w+:)?sheetData\b', lambda m: cols_xml + m.group(0), raw, count=1)
            replacements[rank_path] = raw
        views = book.find(Q + 'bookViews')
        if views is not None:
            for view in views:
                view.set('activeTab', '0')
        replacements['xl/workbook.xml'] = ET.tostring(book, encoding='utf-8')
        replacements['xl/_rels/workbook.xml.rels'] = ET.tostring(rels, encoding='utf-8')
        replacements['[Content_Types].xml'] = ET.tostring(types, encoding='utf-8')
        with ZipFile(destination, 'w') as output:
            for info in src.infolist():
                output.writestr(copy.copy(info), replacements.get(info.filename, src.read(info.filename)))
            for name, content in replacements.items():
                if name not in src.namelist():
                    output.writestr(name, content)
        with ZipFile(destination) as output:
            for name in src.namelist():
                if name not in replacements and src.read(name) != output.read(name):
                    raise ValueError('Untargeted workbook part changed: ' + name)
        return sorted(replacements)


def _cell(value):
    """One canonical form per published cell, so an Excel round-trip is not an edit."""
    if value is None or isinstance(value, (str, bool)):
        return value
    if isinstance(value, (int, float)):
        return float(value)
    return str(value)


def _rows(values):
    """Trailing blanks are packaging, not content; both sides normalize identically."""
    rows = []
    for row in values:
        row = [_cell(value) for value in row]
        while row and row[-1] is None:
            row.pop()
        rows.append(row)
    while rows and not rows[-1]:
        rows.pop()
    return rows


def _digest(values):
    return hashlib.sha256(json.dumps(_rows(values), ensure_ascii=False).encode('utf-8')).hexdigest()


def _column_hidden(tree, column):
    cols = tree.find(Q + 'cols')
    for col in cols if cols is not None else []:
        if int(col.get('min')) <= column <= int(col.get('max')):
            # OOXML booleans are 1/0/true/false; Excel and openpyxl do not agree.
            return (col.get('hidden') or '0').lower() in ('1', 'true')
    return False


def _structure(package):
    """The non-cell facts a publication sets: hidden label columns, tab states, active tab."""
    paths = sheet_paths(package)
    book = ET.fromstring(package.read('xl/workbook.xml'))
    states = {s.get('name'): (s.get('state') or 'visible') for s in book.find(Q + 'sheets')}
    views = book.find(Q + 'bookViews')
    return {'schema': REGION_SCHEMA, 'package': True,
            'columns': {'%s!%d' % (name, column):
                        (_column_hidden(ET.fromstring(package.read(paths[name])), column)
                         if name in paths else None) for name, column in LABEL_COLUMNS},
            'tabs': {name: states.get(name) for name in HIDDEN_SHEETS},
            'active_tab': views[0].get('activeTab') if views is not None and len(views) else None}


def honors_region(workbook):
    """Exactly what a publication owns in a package, and nothing else.

    A file that is not an OPC package has no resolvable region, so it is owned
    whole: that is the legacy binding, reached only as a fail-safe.
    """
    workbook = Path(workbook)
    try:
        package = ZipFile(workbook)
    except BadZipFile:
        return {'schema': REGION_SCHEMA, 'package': False, 'columns': {}, 'tabs': {},
                'active_tab': None, 'sheets': {name: None for name in SHEETS},
                'whole_file_sha256': sha(workbook)}
    with package:
        region = _structure(package)
        present = [name for name in SHEETS if name in sheet_paths(package)]
    region['sheets'] = {name: None for name in SHEETS}
    if present:
        import openpyxl
        read = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
        try:
            for name in present:
                region['sheets'][name] = _digest(read[name].iter_rows(values_only=True))
        finally:
            read.close()
    return region


def _guard(region):
    """The comparable part of a region: the views, plus the file itself when it is not a package."""
    return {'sheets': region['sheets'], 'whole_file_sha256': region.get('whole_file_sha256')}


def _published_region(previous, root):
    """The region as published, from the binding -- or rebuilt from that publication's
    own plan when the binding predates region scoping. Returns None when neither exists."""
    recorded = previous.get('honors_region')
    if recorded and recorded.get('schema') == REGION_SCHEMA:
        return recorded
    work = Path(root) / 'data/stance/rank_honors' / previous['generation_id'] / 'workbook'
    plan, proof = work / 'plan.json', work / 'publish_receipt.json'
    if not (plan.exists() and proof.exists()):
        return None
    receipt = json.loads(proof.read_text(encoding='utf-8'))
    # The receipt ties that plan to these exact published bytes; without both
    # matching the binding, the plan is some other publication's.
    if (receipt.get('generation_id') != previous['generation_id']
            or receipt.get('after_sha256') != previous['after_sha256']):
        return None
    authored = json.loads(plan.read_text(encoding='utf-8'))
    if authored.get('generation_id') != previous['generation_id']:
        return None
    views = {view['name']: _digest(view['values']) for view in authored['views']}
    return {'schema': REGION_SCHEMA, 'package': True, 'columns': None, 'tabs': None,
            'active_tab': None, 'sheets': {name: views.get(name) for name in SHEETS}}


def _refuse_on_honors_edits(previous, workbook, root, refusal):
    """Fail closed when the honors OUTPUT changed; pass when only other sheets did."""
    published = _published_region(previous, root)
    if published is None:
        raise RuntimeError(refusal + ' (this binding predates region scoping and its'
                           ' publication plan is unavailable; HONORS_WORKBOOK_SCOPE'
                           ' cannot narrow it)')
    current = honors_region(workbook)
    if not current['package']:
        raise RuntimeError(refusal + ' (the workbook is not a readable package, so no'
                           ' honors region can be resolved)')
    changed = [name for name in SHEETS if _guard(current)['sheets'][name] != published['sheets'][name]]
    if changed:
        raise RuntimeError("Honors sheet '%s' %s since publication; refusing to roll back over"
                           ' edits to the honors output itself'
                           % (changed[0], 'was removed' if current['sheets'][changed[0]] is None else 'changed'))
    return current


def _rels_path(part):
    return posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')


def _set_column_hidden(raw, column, hidden):
    """Restore one column's hidden flag, splitting the ranges around it and leaving
    every other column property -- width, style, outline -- exactly as found."""
    old = ET.fromstring(raw).find(Q + 'cols')
    if old is None:
        return raw
    cols, found = ET.Element(Q + 'cols'), False
    for col in old:
        lo, hi = int(col.get('min')), int(col.get('max'))
        if not lo <= column <= hi:
            cols.append(copy.deepcopy(col))
            continue
        found = True
        if lo < column:
            left = copy.deepcopy(col); left.set('max', str(column-1)); cols.append(left)
        middle = copy.deepcopy(col); middle.set('min', str(column)); middle.set('max', str(column))
        if hidden:
            middle.set('hidden', '1')
        else:
            middle.attrib.pop('hidden', None)
        cols.append(middle)
        if hi > column:
            right = copy.deepcopy(col); right.set('min', str(column+1)); cols.append(right)
    if not found:
        return raw
    cols[:] = sorted(cols, key=lambda c: int(c.get('min')))
    cols_xml = ET.tostring(cols, encoding='utf-8')
    return re.sub(rb'<(?:\w+:)?cols\b[^>]*>.*?</(?:\w+:)?cols>', lambda _: cols_xml, raw, count=1, flags=re.S)


def remove_views(current, original, destination):
    """Undo exactly the region merge_views wrote, on top of the workbook as it stands.

    Every other ZIP entry is carried across byte for byte and verified -- including
    xl/styles.xml, whose appended entries are inert because merge_views keeps every
    original style index valid, and whose pre-publication copy must NOT be restored:
    an intervening edit may legitimately reference a style appended after it.
    """
    with ZipFile(current) as live, ZipFile(original) as before:
        live_paths, before_paths = sheet_paths(live), sheet_paths(before)
        stale = [name for name in SHEETS if name in before_paths]
        if stale:
            raise RuntimeError('Recorded rollback copy already contains the managed honors view '
                               + repr(stale[0]) + '; use HONORS_WORKBOOK_SCOPE=legacy')
        book = ET.fromstring(live.read('xl/workbook.xml'))
        sheets = book.find(Q + 'sheets')
        rels = ET.fromstring(live.read('xl/_rels/workbook.xml.rels'))
        types = ET.fromstring(live.read('[Content_Types].xml'))
        dropped, relationship_ids, positions = set(), set(), []
        for position, entry in enumerate(list(sheets)):
            if entry.get('name') not in SHEETS:
                continue
            positions.append(position)
            relationship_ids.add(entry.get('{' + REL + '}id'))
            dropped.add(live_paths[entry.get('name')])
            sheets.remove(entry)
        for relationship in list(rels):
            if relationship.get('Id') in relationship_ids:
                rels.remove(relationship)
        for override in list(types):
            if (override.get('PartName') or '').lstrip('/') in dropped:
                types.remove(override)
        # A sheet-local defined name addresses its sheet by POSITION. Removing sheets
        # from the front of the book would silently repoint every one that follows.
        defined = book.find(Q + 'definedNames')
        for entry in list(defined) if defined is not None else []:
            local = entry.get('localSheetId')
            if local is None:
                continue
            local = int(local)
            if local in positions:
                defined.remove(entry)
            else:
                entry.set('localSheetId', str(local - sum(1 for p in positions if p < local)))
        replacements = {}
        for name, column in LABEL_COLUMNS:
            if name not in live_paths or name not in before_paths:
                continue
            raw = live.read(live_paths[name])
            restored = _set_column_hidden(
                raw, column, _column_hidden(ET.fromstring(before.read(before_paths[name])), column))
            if restored != raw:
                replacements[live_paths[name]] = restored
        prior = ET.fromstring(before.read('xl/workbook.xml'))
        states = {s.get('name'): s.get('state') for s in prior.find(Q + 'sheets')}
        for entry in sheets:
            if entry.get('name') in HIDDEN_SHEETS and entry.get('name') in states:
                if states[entry.get('name')] is None:
                    entry.attrib.pop('state', None)
                else:
                    entry.set('state', states[entry.get('name')])
        views, prior_views = book.find(Q + 'bookViews'), prior.find(Q + 'bookViews')
        if views is not None and prior_views is not None and len(views) and len(prior_views):
            if prior_views[0].get('activeTab') is None:
                views[0].attrib.pop('activeTab', None)
            else:
                views[0].set('activeTab', prior_views[0].get('activeTab'))
        replacements['xl/workbook.xml'] = ET.tostring(book, encoding='utf-8')
        replacements['xl/_rels/workbook.xml.rels'] = ET.tostring(rels, encoding='utf-8')
        replacements['[Content_Types].xml'] = ET.tostring(types, encoding='utf-8')
        dropped |= {_rels_path(part) for part in dropped} & set(live.namelist())
        with ZipFile(destination, 'w') as output:
            for info in live.infolist():
                if info.filename in dropped:
                    continue
                output.writestr(copy.copy(info), replacements.get(info.filename, live.read(info.filename)))
        with ZipFile(destination) as output:
            if set(output.namelist()) != set(live.namelist()) - dropped:
                raise ValueError('Rollback changed the workbook part inventory')
            for name in output.namelist():
                if name not in replacements and live.read(name) != output.read(name):
                    raise ValueError('Untargeted workbook part changed: ' + name)
            if set(SHEETS) & set(sheet_paths(output)):
                raise ValueError('Rollback left a managed honors view in the workbook')
        return sorted(dropped | set(replacements))


def _write_bindings(path, bindings):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix='bindings-', suffix='.json', dir=path.parent)
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as handle:
            json.dump(bindings, handle, indent=2)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def prepare_model_update(workbook, root=ROOT):
    """Restore the bound baseline before an authorized model writer runs.

    This operation deliberately does not read the active honors pointer: preparing
    model inputs must also work before a new generation can be constructed.

    When the package is still byte-for-byte as published AND its recorded copy is
    still an exact whole-file restore, that whole file is restored -- the legacy
    behaviour, and the only lossless answer. Otherwise only the honors region is
    reverted and every other edit is kept; an edit INSIDE an honors view still
    refuses, naming the sheet.
    """
    root, workbook = Path(root), Path(workbook)
    binding_path = root / 'data/stance/rank_honors/workbook_publications.json'
    bindings = json.loads(binding_path.read_text(encoding='utf-8')) if binding_path.exists() else {}
    key = str(workbook.resolve())
    previous = bindings.get(key)
    if previous is None:
        return {'status': 'NO_PRIOR_PUBLICATION'}
    observed = sha(workbook)
    as_published = observed == previous['after_sha256']
    # A publication that landed on top of unrelated edits inherits the PREVIOUS
    # publication's rollback copy, which predates them. Restoring that copy whole
    # would discard them, so such a binding may only ever be reverted by region.
    exact = previous.get('original_exact', True)
    if HONORS_WORKBOOK_SCOPE == 'legacy' and not exact:
        # Today's code cannot write such a binding, so refusing here changes no
        # legacy behaviour -- it only stops the dial from discarding kept work.
        raise RuntimeError('Recorded rollback copy predates intervening model edits; the legacy'
                           ' whole-file scope would discard them. Clear HONORS_WORKBOOK_SCOPE'
                           ' to revert by region.')
    scoped = HONORS_WORKBOOK_SCOPE != 'legacy' and not (as_published and exact)
    if scoped and not as_published:
        _refuse_on_honors_edits(previous, workbook, root,
                                'Workbook changed since honors publication; refusing rollback over intervening edits')
    elif not scoped and not as_published:
        raise RuntimeError('Workbook changed since honors publication; refusing rollback over intervening edits')
    original = root / previous['original_path']
    original_bytes = original.read_bytes()
    if hashlib.sha256(original_bytes).hexdigest() != previous['before_sha256']:
        raise RuntimeError('Invalid original workbook rollback copy')
    fd, temporary = tempfile.mkstemp(prefix='honors-rollback-', suffix='.xlsx', dir=workbook.parent)
    reverted = None
    try:
        if scoped:
            os.close(fd)
            reverted = remove_views(workbook, original, temporary)
        else:
            with os.fdopen(fd, 'wb') as handle:
                handle.write(original_bytes)
        if sha(workbook) != observed:
            raise RuntimeError('Workbook changed during rollback; refusing to overwrite intervening edits')
        os.replace(temporary, workbook)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    del bindings[key]
    _write_bindings(binding_path, bindings)
    result = {'status': 'ROLLED_BACK', 'scope': 'region' if scoped else 'whole-file',
              'after_sha256': sha(workbook)}
    if reverted is not None:
        result['reverted_parts'] = reverted
    return result


def publish(workbook, root=ROOT, *, replace_generated=False):
    from stance.honors_views import load_honors
    root, workbook = Path(root), Path(workbook)
    binding_path = root / 'data/stance/rank_honors/workbook_publications.json'
    bindings = json.loads(binding_path.read_text(encoding='utf-8')) if binding_path.exists() else {}
    key = str(workbook.resolve())
    previous = bindings.get(key)
    if replace_generated:
        # Only a producer that has just recreated a clean analytical workbook
        # may replace its binding. A managed published package is never fresh.
        with ZipFile(workbook) as package:
            if any(name in sheet_paths(package) for name in SHEETS):
                raise RuntimeError('Fresh generated workbook still contains managed honors views')
        bindings.pop(key, None)
        previous = None
    # A saved pre-honors backup is valid only for the exact package we last
    # published: under the legacy whole-file scope, reusing it after a model/user
    # edit would let a later rollback erase that intervening work. Under the
    # region scope a rollback removes only the views this publication wrote, so
    # unrelated edits are safe -- but an edit to an honors view still refuses.
    observed = sha(workbook)
    scoped = HONORS_WORKBOOK_SCOPE != 'legacy' and bool(previous) and observed != previous['after_sha256']
    if scoped:
        _refuse_on_honors_edits(previous, workbook, root,
                                'Workbook changed since honors publication; refusing to reuse an earlier'
                                ' rollback copy over intervening edits')
    elif previous and observed != previous['after_sha256']:
        raise RuntimeError('Workbook changed since honors publication; refusing to reuse an earlier rollback copy over intervening edits')
    honors = load_honors(root)
    if honors is None:
        if replace_generated:
            _write_bindings(binding_path, bindings)
            return {'status': 'SKIPPED_LEGACY'}
        if previous is None:
            return {'status': 'SKIPPED_LEGACY'}
        return prepare_model_update(workbook, root)
    # Same generation, and the honors region on disk is the one this binding
    # published: republishing would rewrite identical views over unrelated edits.
    if previous and previous['generation_id'] == honors['generation_id'] and (
            scoped or observed == previous['after_sha256']):
        return {'status': 'UNCHANGED', 'generation_id': honors['generation_id'], 'after_sha256': observed}
    work = root / 'data/stance/rank_honors' / honors['generation_id'] / 'workbook'
    work.mkdir(parents=True, exist_ok=True)
    plan = work / 'plan.json'
    authored = work / 'recognition_views.xlsx'
    plan.write_text(json.dumps({'generation_id': honors['generation_id'], 'output': str(authored),
                                'views': workbook_views(honors)}, ensure_ascii=False), encoding='utf-8')
    # MELANGE_XLSX_WRITER=openpyxl authors the views in process (D-178); the default runs the Node helper.
    writer = xlsx_writer()
    if writer == 'artifact':
        runtime = json.loads((root / 'data/engine/artifact_tool_runtime.json').read_text(encoding='utf-8'))
        helper = Path(__file__).with_suffix('.cjs')
        env = {**os.environ, 'ARTIFACT_TOOL_NODE_MODULES': runtime['node_modules']}
    else:
        helper = AUTHOR_MODULE
    # The hash of the bytes about to be modified -- equal to the binding's
    # after_sha256 whenever the legacy scope would have allowed this publication.
    before = observed
    # Keep the pre-recognition package for an exact rollback, scoped to this
    # specific target. Never overwrite it with a later already-modified copy.
    backup_name = hashlib.sha256(key.encode('utf-8')).hexdigest()[:16] + '.' + before[:16] + '.before.xlsx'
    original_path = root / 'data/stance/rank_honors' / backup_name
    if previous is None:
        original_path.parent.mkdir(parents=True, exist_ok=True)
        source_bytes = workbook.read_bytes()
        if hashlib.sha256(source_bytes).hexdigest() != before:
            raise RuntimeError('Workbook changed while preparing its rollback copy')
        if original_path.exists() and sha(original_path) != before:
            raise RuntimeError('Invalid existing workbook rollback copy')
        if not original_path.exists():
            original_path.write_bytes(source_bytes)
        original_hash = before
    else:
        original_path = root / previous['original_path']
        original_hash = previous['before_sha256']
    # Exactness is sticky and can only be lost: a reused copy that already predates
    # someone's edits does not become a whole-file restore again at the next cycle.
    original_exact = True if previous is None else (previous.get('original_exact', True) and not scoped)
    if writer == 'artifact':
        author_exit = subprocess.run([runtime['node'], str(helper), str(plan)], env=env, timeout=600).returncode
    else:
        author_exit = None
        author_honors_views(plan)
    proof_path = Path(str(authored) + '.receipt.json')
    if not proof_path.exists():
        raise RuntimeError('No artifact-authored recognition workbook receipt; exit %s' % author_exit)
    proof = json.loads(proof_path.read_text(encoding='utf-8'))
    if (proof.get('status') != 'PASS' or proof.get('plan_sha256') != sha(plan)
            or proof.get('helper_sha256') != sha(helper) or proof.get('output_sha256') != sha(authored)):
        raise RuntimeError('Invalid artifact-authored recognition workbook receipt')
    fd, temporary = tempfile.mkstemp(suffix='.xlsx', prefix='honors-', dir=workbook.parent)
    os.close(fd)
    try:
        parts = merge_views(workbook, authored, temporary)
        # Read independently after packaging, before replacing the live file.
        import openpyxl
        check = openpyxl.load_workbook(temporary, read_only=True, data_only=True)
        try:
            for view in workbook_views(honors):
                actual = list(check[view['name']].iter_rows(values_only=True))
                if [list(row) for row in actual] != view['values']:
                    raise ValueError('Authored workbook cell parity failed: ' + view['name'])
        finally:
            check.close()
        # The parity check above proves these views ARE the packaged bytes, so the
        # bound region is recorded from them rather than re-read from the file.
        with ZipFile(temporary) as packaged:
            region = _structure(packaged)
        region['sheets'] = {view['name']: _digest(view['values']) for view in workbook_views(honors)}
        if sha(workbook) != before:
            raise RuntimeError('Workbook changed during recognition synchronization')
        os.replace(temporary, workbook)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    receipt = {'schema': 'honors-workbook-v1', 'status': 'PASS',
               'generation_id': honors['generation_id'], 'before_sha256': before,
               'after_sha256': sha(workbook), 'changed_parts': parts,
               'author_exit_code': author_exit, 'verified_members': len(honors['members']),
               'preservation': 'All existing model cell values, formulas and caches retained; analytical label column hidden.'}
    (work / 'publish_receipt.json').write_text(json.dumps(receipt, indent=2), encoding='utf-8')
    # before_sha256/after_sha256 stay as whole-file provenance; honors_region is
    # what the guard and the rollback are scoped to.
    bindings[key] = {'generation_id': honors['generation_id'], 'before_sha256': original_hash,
                     'after_sha256': receipt['after_sha256'], 'original_path': original_path.relative_to(root).as_posix(),
                     'original_exact': original_exact, 'honors_region': region}
    _write_bindings(binding_path, bindings)
    return receipt


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', default=ROOT)
    parser.add_argument('--workbook')
    parser.add_argument('--prepare-model-update', action='store_true',
                        help='Revert the honors region before running analytical workbook writers'
                             ' (HONORS_WORKBOOK_SCOPE=legacy reverts the whole file)')
    args = parser.parse_args()
    operation = prepare_model_update if args.prepare_model_update else publish
    print(json.dumps(operation(args.workbook or Path(args.root) / 'Moneyball Main Version - Blended.xlsx', args.root)))


if __name__ == '__main__':
    main()
