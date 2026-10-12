"""Snapshot accepted legislative rows and independently mapped roll-call events.

This adapter never imports executable archive/vote producers. The retained vote
rules are read as literals and the same typed measure/category mapping is
reconstructed from their receipt-bound inputs. Aggregate lean is never action.
"""
from __future__ import annotations

import ast
from collections import Counter, defaultdict
import csv
from datetime import date
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import os
from pathlib import Path, PureWindowsPath
import re

from stance.rank_honors_policy import POLICY, CONVERSION_TOLERANCE, literal, number
import bill_rules as BR


def integer(value):
    if isinstance(value, bool) or value is None:
        raise ValueError('Expected an integral source value')
    try:
        raw = Decimal(str(value))
        if not raw.is_finite() or raw != raw.to_integral_value():
            raise ValueError('Nonintegral source value: ' + repr(value))
        return int(raw)
    except InvalidOperation as exc:
        raise ValueError('Invalid integral source value') from exc


def canonical_measure(value, congress):
    term = integer(congress)
    if not 1 <= term <= 999:
        raise ValueError('Congress outside valid identity range')
    token = re.sub(r'[.\s]', '', str(value)).lower()
    if re.fullmatch(r'(hr|s|hjres|sjres|hconres|sconres|hres|sres|hamdt|samdt)[1-9]\d*', token):
        kind, ident = re.fullmatch(r'([a-z]+)(\d+)', token).groups()
        return f'{term}:{kind}:{int(ident)}'
    if re.fullmatch(r'ltr-[a-z0-9-]+', token):
        embedded = re.match(r'ltr-(\d+)-', token)
        if embedded and int(embedded[1]) != term:
            raise ValueError('Letter Congress conflicts with typed identity')
        return f'{term}:letter:{token.upper()}'
    raise ValueError('Unresolved typed measure: ' + str(value))


class Snapshot:
    """All calculations use captured bytes; verify original inputs before publish."""
    def __init__(self, root):
        self.root = Path(root).resolve()
        self.data = {}
        self.hashes = {}

    def path(self, value):
        p = Path(value)
        if p.is_absolute():
            return p.resolve()
        # Windows references in receipts must also work in a copied test root.
        parts = PureWindowsPath(str(value)).parts
        if PureWindowsPath(str(value)).is_absolute():
            if self.root.name in parts:
                return self.root.joinpath(*parts[parts.index(self.root.name) + 1:])
            return Path(value).resolve()
        return (self.root / p).resolve()

    def read(self, value):
        path = self.path(value)
        if path not in self.data:
            self.data[path] = path.read_bytes()
            self.hashes[path] = hashlib.sha256(self.data[path]).hexdigest()
        return self.data[path]

    def digest(self, value):
        self.read(value)
        return self.hashes[self.path(value)]

    def json(self, value):
        return json.loads(self.read(value).decode('utf-8-sig'))

    def rows(self, value):
        return list(csv.DictReader(io.StringIO(self.read(value).decode('utf-8-sig'))))

    def ndjson(self, value):
        return [json.loads(line) for line in self.read(value).decode('utf-8-sig').splitlines() if line.strip()]

    def verify(self):
        for path, digest in self.hashes.items():
            if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
                raise ValueError('Input changed during honors build: ' + str(path))

    def manifest(self):
        result = {}
        for path, digest in self.hashes.items():
            try:
                key = path.relative_to(self.root).as_posix()
            except ValueError:
                key = str(path)
            result[key] = digest
        return dict(sorted(result.items()))


def _require_hash(snapshot, path, expected):
    if snapshot.digest(path) != expected:
        raise ValueError('Accepted source hash mismatch: ' + str(path))


def load_roster(snapshot):
    rows = snapshot.json('data/engine/congestion.json')
    roster, seen = {}, set()
    for row in rows:
        name, bio = row['member'], row.get('bioguide')
        if not name or name in roster or not re.fullmatch(r'[A-Z]\d{6}', str(bio)) or bio in seen:
            raise ValueError('Roster has duplicate or unresolved stable identity')
        if row['chamber'] not in ('House', 'Senate'):
            raise ValueError('Unknown roster chamber')
        seen.add(bio)
        roster[name] = {'member_id': bio, 'chamber': row['chamber']}
    if not roster:
        raise ValueError('Empty roster')
    return dict(sorted(roster.items()))


def _admissions(snapshot):
    doc = snapshot.json('data/legislation/sentinel_admissions.json')
    if doc.get('schema') != 'sentinel-admissions-v1' or doc.get('status') != 'PASS':
        raise ValueError('Invalid official admissions')
    for record in doc['sources']:
        _require_hash(snapshot, record['path'], record['sha256'])
    for key, bill in doc['bills'].items():
        decision = doc['decisions'][key]
        if bill['key'] != key or decision.get('canonical_key') != key or decision.get('disposition') != 'include' or decision.get('confidence') != 'high':
            raise ValueError('Unaccepted official bill decision')
        for row in decision['categories']:
            if type(row['direction']) is not int or row['direction'] not in (-1, 0, 1) or not row.get('quote') or not row.get('rationale'):
                raise ValueError('Official category direction lacks provenance')
    return doc


def load_tensor(snapshot):
    result, seen = {}, set()
    for row in snapshot.ndjson('data/lobbying/industry_axes.ndjson'):
        if row.get('industry') != 'credit_union':
            continue
        key = canonical_measure(row['bill'], row['congress'])
        if key in seen:
            raise ValueError('Duplicate CU tensor measure cell: ' + key)
        seen.add(key)
        if row.get('status') not in ('active', 'inactive'):
            raise ValueError('Unknown CU tensor status')
        relevance, valence, confidence = [number(row[k]) for k in ('relevance', 'valence', 'confidence')]
        if not (0 <= relevance <= 1 and -1 <= valence <= 1 and 0 < confidence <= 1):
            raise ValueError('Invalid CU tensor value')
        if not all(row.get(k) for k in ('basis', 'rationale', 'adjudicated_by', 'adjudicated')):
            raise ValueError('Unadjudicated tensor cell')
        if row['status'] == 'active':
            result[key] = literal(relevance * valence)
    return result


def _action(member_id, category, measure, kind, event, alignment, role, decay, **extra):
    if type(alignment) is not int or alignment not in (-1, 0, 1):
        raise ValueError('Invalid accepted action direction')
    weight = POLICY['vote_weight'] if kind == 'vote' else POLICY['role_weights'].get(role)
    if weight is None:
        raise ValueError('Unsupported legislative action role: ' + str(role))
    base = f'{kind}|{member_id}|{event}'
    ident = base + '|' + category
    return {'evidence_id': ident, 'overall_event_id': base, 'member_id': member_id,
            'category_id': category, 'measure_id': measure, 'kind': kind,
            'alignment': alignment, 'weight': weight, 'decay': literal(decay),
            'role': role, 'position_on_measure': 1,
            'tensor_position_scoped': kind != 'letter', **extra}


def _deduplicate(actions, counts):
    out = {}
    for row in actions:
        key = row['evidence_id']
        prior = out.get(key)
        if prior:
            if any(prior[k] != row[k] for k in ('member_id', 'category_id', 'measure_id', 'alignment', 'decay', 'position_on_measure')):
                raise ValueError('Conflicting duplicate action: ' + key)
            counts['duplicate_imports'] += 1
            if number(prior['weight']) >= number(row['weight']):
                continue
        out[key] = row
    return sorted(out.values(), key=lambda r: r['evidence_id'])


def legislative_actions(snapshot, roster, categories, admissions, counts, titles=None):
    doc = snapshot.json('legislation_routing_audit.json')
    archive = snapshot.json('archive_leg.json')
    if doc.get('schema') != 'legislation-routing-audit-v1' or not isinstance(doc.get('rows'), list):
        raise ValueError('Invalid legislative routing audit')
    registry = {}
    for row in snapshot.rows('data/legislation/bill_directions.csv'):
        if row.get('Congress'):
            key = canonical_measure(row['Bill'], row['Congress'])
            value = integer(row['Direction'])
            if value not in (-1, 0, 1) or (key in registry and registry[key] != value):
                raise ValueError('Conflicting direction registry')
            registry[key] = value
    manual = {}
    for row in snapshot.rows('manual_bill_evidence.csv'):
        if str(row['Bill']).upper().startswith('LTR-'):
            key = (row['Member'], canonical_measure(row['Bill'], row['Congress']), BR.canonical_category(row['Category']))
            manual[key] = row
    movements = snapshot.json('data/legislation/official_membership_movements.json')
    if movements.get('schema') != 'official-membership-movements-v1' or movements.get('status') != 'PASS':
        raise ValueError('Invalid official membership movements')
    _require_hash(snapshot, 'bill_repository_live.json', movements['repository_sha256'])
    facts = {(r['repository_key'], r['bioguide']): r for r in movements['facts']}
    sums = defaultdict(Decimal)
    actions = []
    for row in doc['rows']:
        disposition = row['disposition']
        counts['legislative_disposition:' + disposition] += 1
        if disposition not in ('selected', 'neutral-registry', 'neutral-official-evidence'):
            continue
        name, category = row.get('member'), row.get('effective_category')
        if name not in roster or category not in categories:
            raise ValueError('Accepted row has unresolved member/category')
        measure = canonical_measure(row['bill'], row['congress'])
        term = int(measure.split(':')[0])
        if titles is not None and row.get('title') and measure not in titles:
            titles[measure] = str(row['title'])
        role = BR.canonical_role(row.get('effective_role') or row['source_role'])
        weight = POLICY['role_weights'].get(role)
        if weight is None:
            raise ValueError('Unsupported accepted role')
        sign = integer(row['sign'])
        if sign not in (-1, 0, 1):
            raise ValueError('Invalid accepted sign')
        bio = roster[name]['member_id']
        fact = facts.get((measure, bio))
        if fact and (fact.get('withdrawn') or BR.canonical_role(fact['role']) != role):
            raise ValueError('Accepted routing row conflicts with current official membership')
        is_letter = ':letter:' in measure
        if disposition != 'selected':
            # Neutral dispositions were emitted only after member/category and
            # official-current-membership routing. Require the exact direction
            # source as additional provenance, not a disposition-string guess.
            if sign != 0 or not row.get('source_row') or not row.get('title') or not row.get('member_source'):
                raise ValueError('Neutral membership lacks explicit provenance')
            if disposition == 'neutral-registry':
                if row.get('direction_mechanism') != 'registry-exact' or registry.get(measure) != 0:
                    raise ValueError('Neutral registry evidence is not exact and explicit')
            else:
                candidates = admissions['decisions'].get(measure, {}).get('categories', [])
                if not any(BR.canonical_category(r['category']) == category and r['direction'] == 0 for r in candidates):
                    raise ValueError('Neutral official direction is not admitted')
            decay = Decimal('0.8') ** max(0, 119 - term)
        else:
            decay = number(row['decay_weight'])
        contribution = number(weight) * decay * sign
        if abs(contribution - number(row['contribution'])) > CONVERSION_TOLERANCE:
            raise ValueError('Signed row does not reconcile to its role and decay')
        if disposition == 'selected':
            sums[(name, category)] += contribution
        if term not in POLICY['congresses']:
            counts['legislative_outside_congress_window'] += 1
            continue
        if is_letter:
            source = manual.get((name, measure, category))
            if not source or not source.get('Note') or integer(source['Direction']) != sign or BR.canonical_role(source['Role']) != role:
                raise ValueError('Letter lacks accepted signed ledger provenance')
            if decay != number(POLICY['congress_decay'][str(term)]):
                raise ValueError('Letter uses stale Congress decay')
        actions.append(_action(bio, category, measure, 'letter' if is_letter else 'legislative', measure,
            sign, role, decay, congress=term, source='legislation_routing_audit.json',
            source_row=row['source_row'], direction_mechanism=row.get('direction_mechanism'),
            source_disposition=disposition, source_url=(source.get('DetailLink') or None) if is_letter else
            f'https://www.congress.gov/{"amendment" if "amdt" in measure else "bill"}/{term}th-congress/'
            + {'hr': 'house-bill', 's': 'senate-bill', 'hjres': 'house-joint-resolution',
               'sjres': 'senate-joint-resolution', 'hres': 'house-resolution', 'sres': 'senate-resolution',
               'hconres': 'house-concurrent-resolution', 'sconres': 'senate-concurrent-resolution',
               'hamdt': 'house-amendment', 'samdt': 'senate-amendment'}[measure.split(':')[1]] + '/' + measure.split(':')[2]))
    expected = {(m, c): number(v) for m, cats in archive.items() for c, v in cats.items()}
    if any(abs(sums.get(key, Decimal(0)) - expected.get(key, Decimal(0))) > CONVERSION_TOLERANCE
           for key in set(sums) | set(expected)):
        raise ValueError('Selected legislative subtotal differs from archive_leg.json')
    counts['archive_reconciliation_cells'] = len(expected)
    return _deduplicate(actions, counts)


def _vote_mapping(snapshot, receipt, admissions):
    """Mirror build_vote_category's retained mapping using its bound sources."""
    bound = {}
    for path, digest in receipt['inputs_sha256'].items():
        bound[PureWindowsPath(path).name] = (path, digest)
    needed = ('congress_bills_archive_', 'historical_bill_exclusions.csv',
              'cfpb_bill_directions.csv', 'legislation_exclusion_ledger.ndjson',
              'legislation_hold_ledger.ndjson', 'vote_by_category.py', 'sentinel_admissions.json')
    if any(not any(name.startswith(prefix) for name in bound) for prefix in needed):
        raise ValueError('Vote receipt omits a required category-mapping source hash')
    for name, (path, digest) in bound.items():
        if any(name.startswith(prefix) for prefix in needed):
            _require_hash(snapshot, path, digest)
    rules = {}
    tree = ast.parse(snapshot.read('fec/legislation/vote_by_category.py').decode('utf-8-sig'))
    for node in tree.body:
        if isinstance(node, ast.Assign):
            for target in node.targets:
                if isinstance(target, ast.Name) and target.id in {'CATLAB', 'ANTI', 'SECKW'}:
                    rules[target.id] = ast.literal_eval(node.value)
    if set(rules) != {'CATLAB', 'ANTI', 'SECKW'}:
        raise ValueError('Missing retained vote mapping rules')
    labels, anti, keywords = rules['CATLAB'], rules['ANTI'], rules['SECKW']
    if {str(k): v for k, v in labels.items()} != receipt['retained_model_rules']['CATLAB'] or sorted(anti) != receipt['retained_model_rules']['ANTI']:
        raise ValueError('Vote mapping policy differs from accepted receipt')
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(snapshot.read(receipt['archive'])), read_only=True, data_only=True)
    sheet = workbook['All Bills']
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    mapping = defaultdict(set)
    for values in iterator:
        row = dict(zip(headers, values))
        try:
            key = canonical_measure(row['Name'], row['Congress'])
        except ValueError:
            continue
        category = next((i for i, word in keywords if word in str(row['Section']).lower()), None)
        if category:
            mapping[key].add(category)
    workbook.close()
    historical_path = bound['historical_bill_exclusions.csv'][0]
    excluded = set()
    for row in snapshot.rows(historical_path):
        try:
            excluded.add(canonical_measure(row['bill'], row['congress']))
        except ValueError:
            continue
    category_ids = {label: i for i, label in labels.items()}
    scoped_exclusions = set()
    for path in ('legislation_exclusion_ledger.ndjson', 'legislation_hold_ledger.ndjson'):
        for row in snapshot.ndjson(path):
            if row.get('decision_scope') == 'candidate':
                excluded.add(row['canonical_key'])
            else:
                for value in row.get('category_ids') or [row.get('query_id')]:
                    category = category_ids.get(BR.canonical_category(value))
                    if category:
                        scoped_exclusions.add((row['canonical_key'], category))
    for key in excluded:
        mapping.pop(key, None)
    admitted = {}
    for key in admissions['bills']:
        if key in excluded:
            raise ValueError('Official admission conflicts with historical exclusion')
        for row in admissions['decisions'][key]['categories']:
            category = category_ids.get(BR.canonical_category(row['category']))
            if category is None:
                raise ValueError('Official category has no retained vote mapping')
            mapping[key].add(category)
            admitted[(key, category)] = row['direction']
    for key, category in scoped_exclusions:
        mapping[key].discard(category)
    cfpb = {}
    for row in snapshot.rows('data/legislation/cfpb_bill_directions.csv'):
        token = canonical_measure(row['Bill'], 119).split(':', 1)[1]
        candidates = [key for key, cats in mapping.items() if 13 in cats and key.split(':', 1)[1] == token]
        if len(candidates) == 1 and row['Direction'].strip() in ('1', '-1', '0'):
            value = int(row['Direction'])
            if candidates[0] in cfpb and cfpb[candidates[0]] != value:
                raise ValueError('Conflicting CFPB vote direction')
            cfpb[candidates[0]] = value
    result = {}
    for key, cats in mapping.items():
        for category in cats:
            direction = vote_direction(key, category, admitted, cfpb, anti)
            if direction is not None:
                result[(key, labels[category])] = direction
    return result


LEGACY_CAST_CODES = ({1}, {4})


def vote_cast_codes(snapshot=None):
    """Yea/Nay cast codes, lifted from the analytic producer's own mapping.

    VOTE_DIRECTION=legacy keeps the historical (1, 4) admission gate, which
    silently dropped every Nay because the accepted roll-call CSV writes cast
    code 6 for Nay and never writes 4. The default reads build_vote_category's
    literal sets so the two lanes cannot drift apart again.
    """
    dial = os.environ.get('VOTE_DIRECTION', 'honors').strip().lower()
    if dial not in ('honors', 'legacy'):
        raise ValueError('Unknown VOTE_DIRECTION: ' + dial)
    if dial == 'legacy':
        return LEGACY_CAST_CODES
    if snapshot is None:
        raise ValueError('Unified vote direction needs its bound producer source')
    tree = ast.parse(snapshot.read('fec/legislation/build_vote_category.py').decode('utf-8-sig'))
    for node in ast.walk(tree):
        if not (isinstance(node, ast.Assign) and isinstance(node.value, ast.IfExp)
                and any(isinstance(t, ast.Name) and t.id == 'vote' for t in node.targets)):
            continue
        mapping, branch = {}, node.value
        while isinstance(branch, ast.IfExp):
            test = branch.test
            if (isinstance(test, ast.Compare) and len(test.ops) == 1 and isinstance(test.ops[0], ast.In)
                    and isinstance(test.left, ast.Name) and test.left.id == 'code'
                    and isinstance(branch.body, ast.Constant)):
                mapping[branch.body.value] = frozenset(ast.literal_eval(test.comparators[0]))
            branch = branch.orelse
        if 'Yea' in mapping and 'Nay' in mapping:
            yea, nay = mapping['Yea'], mapping['Nay']
            if not yea or not nay or yea & nay or any(type(c) is not int for c in yea | nay):
                raise ValueError('Retained cast-code mapping is not a disjoint integer partition')
            return yea, nay
    raise ValueError('Retained roll-call cast-code mapping not found')


def vote_direction(key, category, admitted, cfpb, anti):
    """An explicit neutral direction is known evidence; missing CFPB is not."""
    if (key, category) in admitted:
        return admitted[(key, category)]
    if category == 13:
        return cfpb.get(key)
    return -1 if category in anti else 1


def vote_actions(snapshot, roster, categories, admissions, as_of, counts):
    receipt = snapshot.json('data/legislation/vote_refresh_receipt.json')
    if receipt.get('schema') != 'vote-category-refresh-v1' or receipt.get('status') != 'PASS':
        raise ValueError('Vote events have no valid accepted refresh receipt')
    _require_hash(snapshot, 'votes/cu_rollcall_votes_clean.csv', receipt['outputs_sha256']['cu_rollcall_votes_clean.csv'])
    _require_hash(snapshot, 'vote_category_lean.json', receipt['outputs_sha256']['vote_category_lean.json'])
    mapping = _vote_mapping(snapshot, receipt, admissions)
    yea, nay = vote_cast_codes(snapshot)
    by_measure = defaultdict(list)
    for (measure, category), direction in mapping.items():
        if category not in categories:
            raise ValueError('Vote category outside canonical fields')
        by_measure[measure].append((category, direction))
    bios = {r['member_id'] for r in roster.values()}
    actions, seen = [], {}
    for row in snapshot.rows('votes/cu_rollcall_votes_clean.csv'):
        term = int(row['congress'])
        if term not in POLICY['congresses']:
            counts['vote_outside_congress_window'] += 1
            continue
        if not re.fullmatch(r'[A-Z]\d{6}', row['bioguide']):
            raise ValueError('Unresolved roll-call member identity')
        if row['bioguide'] not in bios:
            counts['vote_inactive_member'] += 1
            continue
        if date.fromisoformat(row['date']) > as_of:
            counts['vote_after_as_of'] += 1
            continue
        question = row['vote_question']
        if not re.search(r'passage|suspend the rules and pass|concur', question, re.I) or re.search(r'recommit|amendment', question, re.I):
            counts['vote_not_substantive_final_passage'] += 1
            continue
        measure = canonical_measure(row['bill_number'], term)
        if measure != row['canonical_key'] or row['chamber'] not in ('House', 'Senate') or not re.fullmatch(r'[1-9]\d*', row['rollnumber']):
            raise ValueError('Invalid typed roll-call identity')
        event = f'{term}:{row["chamber"]}:{row["rollnumber"]}'
        key = (row['bioguide'], event)
        code = int(row['cast_code'])
        value = (measure, code, row['date'], row['vote'])
        if key in seen:
            if seen[key] != value:
                raise ValueError('Conflicting duplicate roll-call vote')
            counts['vote_duplicate_import'] += 1
            continue
        seen[key] = value
        if code not in yea and code not in nay:
            counts['vote_non_actual_cast:' + str(code)] += 1
            continue
        position = 1 if code in yea else -1
        if row['vote'] != ('Yea' if code in yea else 'Nay'):
            raise ValueError('Cast code contradicts vote text')
        counts['vote_cast_position:' + ('yea' if position == 1 else 'nay')] += 1
        for category, direction in sorted(by_measure[measure]):
            actions.append(_action(row['bioguide'], category, measure, 'vote', event,
                position * direction, 'vote', POLICY['congress_decay'][str(term)],
                congress=term, position_on_measure=position, cast_code=code,
                rollcall_id=event, date=row['date'], source='votes/cu_rollcall_votes_clean.csv',
                direction_mechanism='accepted_vote_category_mapping',
                source_url='https://voteview.com/rollcall/' + f'{row["chamber"][0]}{term}{int(row["rollnumber"]):04d}'))
        if not by_measure[measure]:
            counts['vote_no_accepted_category_direction'] += 1
    return _deduplicate(actions, counts)


def build_evidence(snapshot, as_of):
    as_of = date.fromisoformat(str(as_of))
    roster = load_roster(snapshot)
    category = snapshot.json('category_stance.json')
    categories = list(BR.CATEGORIES)
    if category['cats'] != categories or set(category['matrix']) != set(roster):
        raise ValueError('Canonical roster/category mismatch with analytic input')
    source_contract = category.get('source_contract', {})
    if source_contract.get('schema') != 'adjudicated-label-inputs-v1' or set(source_contract.get('inputs_sha256', {})) != {
            'archive_leg.json', 'legislation_routing_audit.json', 'vote_category_lean.json', 'bill_repository_live.json'}:
        raise ValueError('Analytic input omits required accepted action-source hashes')
    for path, digest in source_contract['inputs_sha256'].items():
        _require_hash(snapshot, path, digest)
    admissions = _admissions(snapshot)
    if date.fromisoformat(admissions['checked_at'][:10]) > as_of:
        raise ValueError('As-of precedes accepted source snapshot; historical backtesting unsupported')
    counts = Counter()
    titles = {}
    legislation = legislative_actions(snapshot, roster, categories, admissions, counts, titles)
    votes = vote_actions(snapshot, roster, categories, admissions, as_of, counts)
    actions = sorted(legislation + votes, key=lambda r: r['evidence_id'])
    tensor = load_tensor(snapshot)
    counts.update({'accepted_actions': len(actions), 'accepted_legislative_actions': len(legislation),
                   'accepted_vote_actions': len(votes), 'active_tensor_cells': len(tensor)})
    measures = {a['measure_id'] for a in actions}
    return {'roster': roster, 'cats': categories, 'actions': actions, 'tensor': tensor,
            'measure_titles': {k: v for k, v in sorted(titles.items()) if k in measures},
            'counts': dict(sorted(counts.items())), 'analytic': category,
            'evidence_window': {'congresses': list(POLICY['congresses']), 'as_of': str(as_of),
                'basis': 'accepted_source_snapshot; event dates enforced where available',
                'vote_source_generated_at': snapshot.json('data/legislation/vote_refresh_receipt.json')['generated_at'],
                'legislative_source': snapshot.json('legislation_routing_audit.json')['source']}}
