# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root
"""REBUILD press_raw_original.json — the media-blend idempotence cache (audit 2026-07-02 S2-05).

build_media_blend.py NEEDS this cache to separate the PRESS component from the SOCIAL component in
the 'Press Release - Raw' cells: after the first blend those cells hold press+social, so the press
half must always come from a cached snapshot of the PRE-BLEND raw values (never re-read live).

The cache is created ONCE from a pre-blend workbook and is gitignored → if lost it must be rebuilt
from a backup taken BEFORE any media blend ran. This script does that, safely and atomically.

Usage:
  python rebuild_press_cache.py --from "backups/Moneyball Main Version - Blended.PRE-<...>.xlsx"
  python rebuild_press_cache.py --list          # show candidate pre-blend backups
A pre-blend backup is one whose 'Press Release - Raw' footer row (the row right under the member block) label cell says 'Weights' AND whose
raw cells are the ORIGINAL press values. The safest source is a PRE-* backup from before 2026-07-02
(the media-blend date); the script refuses a source whose Z1 note already says 'MEDIA BLEND'."""
import sys, os, glob, argparse
import openpyxl
import roster
from moneyball_core import dump_atomic
M=MB_ROOT
def candidates():
    return sorted(glob.glob(M+r'\backups\Moneyball Main Version - Blended.*.xlsx'))
a=argparse.ArgumentParser(); a.add_argument('--from',dest='src'); a.add_argument('--list',action='store_true')
args=a.parse_args()
if args.list or not args.src:
    print('candidate backups (pick one taken BEFORE the first media blend):')
    for c in candidates(): print('  ',os.path.relpath(c,M))
    print('\nre-run with: python rebuild_press_cache.py --from "<one of the above>"'); sys.exit(0)
src=args.src if os.path.isabs(args.src) else os.path.join(M,args.src)
if not os.path.exists(src): print('!! not found:',src); sys.exit(1)
wb=openpyxl.load_workbook(src,data_only=True)
rw=wb['Press Release - Raw']
z1=str(rw['Z1'].value or '')
if 'MEDIA BLEND' in z1:
    print('!! REFUSING: %s is a POST-blend workbook (Z1=%r). Its raw cells already contain press+social; '
          'pick a PRE-blend backup.'%(os.path.basename(src),z1[:60])); sys.exit(2)
hdr={rw.cell(1,c).value:c for c in range(1,30) if rw.cell(1,c).value}
PRESS_ROWS=roster.rows('Press Release - Raw')   # member rows only; the 'Weights' footer below them was never a member
members=[rw.cell(r,1).value for r in PRESS_ROWS if rw.cell(r,1).value]
cache={}
for h,c in hdr.items():
    if h in ('Names','District/State') or not isinstance(h,str): continue
    cache[h]={m:(rw.cell(r,c).value or 0.0) for r,m in ((r,rw.cell(r,1).value) for r in PRESS_ROWS) if m}
out=M+r'\data\stance\press_raw_original.json'
if os.path.exists(out):
    bak=out+'.replaced-%s'%__import__('datetime').datetime.now().strftime('%Y%m%d-%H%M%S')
    __import__('shutil').copy2(out,bak); print('existing cache backed up to',os.path.basename(bak))
dump_atomic(cache,out)
print('rebuilt press_raw_original.json from %s: %d categories x %d members'%(os.path.basename(src),len(cache),len(members)))
