# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT,DRAFT as MB_DRAFT  # centralized root (arch review 2026-07-02)
"""S1 v2 — BALANCED gold-label queue (user spec 2026-07-01: "a well balanced mix from many members
across all issues, covering all stances"). Stratification:
  * per-CATEGORY quota across ALL 19 press category labels (incl. Credit Union Tax Status and
    CU-Bank Competitive Dynamics),
  * within each category, round-robin across stance classes (oppose first -- it's the rare class),
  * MEMBER CAP <= 3 rows per member (385 distinct members in the source),
  * the known confusion articles always included (permanent regression suite).
Output gold_label_queue.csv carries the richer Burnham-factorization schema; the api_* columns are
filled by fec/stance/api_prelabel.py (Claude API pre-label pass) and the gold_* columns are the USER'S to
adjudicate (pre-seeded from the API pass for correction-in-place)."""
import csv, io, sys, hashlib, collections, random
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
random.seed(7)
B = MB_ROOT
SRC = MB_DRAFT+r'\cu_press_stance_tone.csv'
rows = list(csv.DictReader(open(SRC, encoding='utf-8-sig')))
print('source rows:', len(rows))

PER_CATEGORY = 32     # target per category label (19 labels -> ~600 before dedup/overlap)
MEMBER_CAP = 3

def cats_of(r):
    c = (r.get('categories') or '').strip()
    return [p.strip() for p in c.replace('|', ';').split(';') if p.strip()] or ['(uncategorized)']
def channel(r):
    url = (r.get('url') or '').lower()
    return ('acu' if 'americascreditunions' in url else
            ('own-domain' if ('.senate.gov' in url or '.house.gov' in url) else 'third-party'))
def key(r):
    return hashlib.md5(((r.get('url') or '') + '|' + (r.get('member') or '')).encode()).hexdigest()[:12]

# regression suite: known-confusion articles
regset = set()
try:
    import openpyxl
    wb = openpyxl.load_workbook(B + r'\CU_Confusion_Articles.xlsx', read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    hdr = [str(c.value or '').lower() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    uc = hdr.index('url') if 'url' in hdr else None
    tc = hdr.index('title') if 'title' in hdr else None
    for rr in ws.iter_rows(min_row=2, values_only=True):
        if uc is not None and rr[uc]: regset.add(str(rr[uc]).strip())
        elif tc is not None and rr[tc]: regset.add(str(rr[tc]).strip()[:60])
    wb.close()
except Exception as e:
    print('(confusion workbook unreadable: %s)' % e)
def isreg(r):
    return (r.get('url') or '').strip() in regset or (r.get('title') or '').strip()[:60] in regset

picked = {}                       # key -> (row, set(reasons))
percap = collections.Counter()    # member -> picks
def take(r, why, force=False):
    k = key(r)
    if k in picked:
        picked[k][1].add(why); return True
    m = r.get('member') or ''
    if not force and percap[m] >= MEMBER_CAP: return False
    picked[k] = (r, {why}); percap[m] += 1
    return True

# 1) regression cases always in (exempt from member cap)
for r in rows:
    if isreg(r): take(r, 'regression', force=True)

# 2) per-category x stance round-robin under the member cap
bycat = collections.defaultdict(lambda: {'oppose': [], 'support': [], 'neutral': []})
for r in rows:
    st = (r.get('stance') or 'neutral').strip() or 'neutral'
    if st not in ('oppose', 'support', 'neutral'): st = 'neutral'
    for c in cats_of(r):
        bycat[c][st].append(r)
for c in bycat:
    for st in bycat[c]: random.shuffle(bycat[c][st])

for c in sorted(bycat):
    got = sum(1 for k, (r, w) in picked.items() if c in cats_of(r))
    order = ['oppose', 'support', 'neutral']   # rare class first
    idx = {st: 0 for st in order}
    while got < PER_CATEGORY:
        progress = False
        for st in order:
            pool = bycat[c][st]
            while idx[st] < len(pool):
                r = pool[idx[st]]; idx[st] += 1
                if key(r) in picked: continue
                if take(r, 'cat:' + c):
                    got += 1; progress = True; break
            if got >= PER_CATEGORY: break
        if not progress: break     # category exhausted under the member cap

OUT = B + r'\data\stance\gold_label_queue.csv'
cols = ['id', 'member', 'party', 'date', 'title', 'url', 'categories', 'source_channel', 'slices',
        'model_stance', 'model_stance_src', 'model_conf', 'model_tone(lsd_net_norm)',
        'api_target', 'api_stance_toward_target', 'api_target_cu_valence', 'api_cu_alignment',
        'api_attribution', 'api_confidence', 'api_rationale', 'debate_vs_api',
        'gold_target', 'gold_stance_toward_target', 'gold_target_cu_valence', 'gold_cu_alignment',
        'gold_attribution', 'gold_notes']
with io.open(OUT, 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f); w.writerow(cols)
    for k, (r, whys) in sorted(picked.items(), key=lambda kv: (kv[1][0].get('member') or '', kv[1][0].get('date') or '')):
        w.writerow([k, r.get('member'), r.get('party'), r.get('date'), r.get('title'), r.get('url'),
                    r.get('categories'), channel(r), ';'.join(sorted(whys)),
                    r.get('stance'), r.get('stance_src'), r.get('stance_conf'), r.get('lsd_net_norm'),
                    '', '', '', '', '', '', '', '', '', '', '', '', '', ''])
cc = collections.Counter(pr[0].get('stance') for pr in picked.values())
mm = collections.Counter(pr[0].get('member') for pr in picked.values())
ch = collections.Counter(channel(pr[0]) for pr in picked.values())
catcov = collections.Counter(c for pr in picked.values() for c in cats_of(pr[0]))
print('QUEUE: %d items | distinct members: %d (max/member %d) | channels %s' % (
    len(picked), len(mm), max(mm.values()), dict(ch)))
print('stance mix:', dict(cc))
print('category coverage:')
for c, n in catcov.most_common(): print('   %-42s %d' % (c, n))
