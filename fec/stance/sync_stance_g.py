# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root
"""Synchronize workbook views with Python's live labels and legislative inputs.

Python prepares cell changes; artifact-tool authors the changed XLSX cells.
The package merger preserves all unrelated worksheet content and native objects.
Run after compute_pillars and primacy generation (run_pipeline owns this order).
"""
import json,sys,collections
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
import openpyxl
import roster
from moneyball_core import overall_display_tier
from stance.honors_views import load_honors
from stance.workbook_cell_sync import publish_cells
from tier_util import legislative_press_tiers
from stance.workbook_stance_views import (
    stance_view_contract, sync_category_stance_sheet, sync_coalition_profiles_sheet,
)
B=MB_ROOT
RA_ROWS=roster.rows('Rank & Allocation')
CS_FIRST,CS_LAST=roster.first_row('Category Stance'),roster.last_row('Category Stance')
WB=_o.path.join(B,'Moneyball Main Version - Blended.xlsx')
CS=json.load(open(_o.path.join(B,'category_stance.json'),encoding='utf-8'))
M=CS['matrix']; OVERALL_CAPS=(CS.get('gate') or {}).get('overall_capped') or {}
_tier,_trank=legislative_press_tiers(list(M))
# ONE CROWN SOURCE (D-030/D-031). Column G is a published surface, so the crown it
# shows is the active honors generation's overall_board.top13 -- the same board the
# member desk, the dashboard bundle and the publication guard read -- and is no longer
# re-derived here from the Legislative+Press tier. run_pipeline.py:153-155 publishes
# that generation two steps before this one, and load_honors refuses a generation
# seated on a stale board basis, so a mid-chain reseat can no longer split the
# workbook from the pages. Value-identical by construction: honors_views rejects any
# generation whose priority tiers are not the board's 13/37/85 buckets, so tier 1 and
# top13 are the same thirteen names in every generation that loads.
_BOARD13=set((load_honors(B).get('overall_board') or {}).get('top13') or ())
if len(_BOARD13)!=13: raise ValueError('Active honors board must seat exactly 13 crowns')
def agg(n): return overall_display_tier(M.get(n,{}),_tier.get(n,0),n,OVERALL_CAPS,_BOARD13)
wb=openpyxl.load_workbook(WB)                       # formulas preserved
before={(ws.title,c.coordinate):c.value for ws in wb for row in ws for c in row if c.value is not None}
style_counts=collections.defaultdict(collections.Counter);style_refs={}
for row in wb['Category Stance'].iter_rows(min_row=CS_FIRST,max_row=CS_LAST,min_col=3,max_col=20):
    for c in row:
        style_counts[c.value][c.style_id]+=1
        style_refs[(c.value,c.style_id)]=c.coordinate
styles={label:(counts.most_common(1)[0][0],style_refs[(label,counts.most_common(1)[0][0])]) for label,counts in style_counts.items()}
for label,base in {'Persuadable':'Highly Persuadable','Lean Persuadable':'Highly Persuadable','Highly Skeptical':'Skeptic'}.items():
    styles.setdefault(label,styles[base])
if '—' not in styles:styles['—']=styles[max(style_counts,key=lambda label:sum(style_counts[label].values()))]
ws=wb['Rank & Allocation']; cnt=collections.Counter(); n=0
members=[ws.cell(r,1).value for r in RA_ROWS if ws.cell(r,1).value]
overall={name:agg(name) for name in members}
stance_view_contract(CS,members,overall,_tier)
for r in RA_ROWS:
    nm=ws.cell(r,1).value
    if not nm: continue
    s=agg(nm); ws.cell(r,7).value=s; cnt[s]+=1; n+=1
sync_category_stance_sheet(wb,CS,members,overall)
sync_coalition_profiles_sheet(wb)
ranks=json.load(open(_o.path.join(B,'ranks_python.json'),encoding='utf-8'))
live=ranks.get('legislative_pillar')
if live:
    if set(live['agenda'])!=set(members):raise ValueError('live legislative roster mismatch')
    bm=wb['Blended Model']
    for r,name in enumerate(members,start=3):
        bm.cell(r,4).value=live['agenda'][name]
        bm.cell(r,7).value=live['efficiency'][name]
        ws.cell(r,4).value=f"='Blended Model'!D{r}"
changes=[]
for sheet in wb:
    for row in sheet:
        for cell in row:
            style=styles.get(cell.value) if sheet.title=='Category Stance' and CS_FIRST<=cell.row<=CS_LAST and 3<=cell.column<=20 else None
            if before.get((sheet.title,cell.coordinate))!=cell.value or (style and cell.style_id!=style[0]):
                change={'sheet':sheet.title,'cell':cell.coordinate,'value':cell.value}
                if style:change.update({'style_id':style[0],'style_reference':style[1]})
                changes.append(change)
publish_cells(WB,changes,B)
print('wrote %d overall stance cells and reconciled live legislative inputs'%n)
print('distribution:',dict(cnt))
print('Category Stance and Coalition Profiles synchronized with the same live labels.')
print('SAVED.')
