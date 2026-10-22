import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
import shutil,os,json,sys,subprocess
import roster
import win32com.client as w
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
src=MB_ROOT+r'\Moneyball Main Version - Blended.xlsx'
bak=MB_ROOT+r'\Moneyball Main Version - Blended.PRE-IM-20260626.xlsx'
tmp=r'C:\Temp\mb_im.xlsx'
# M4: regenerate im_allocation.json from the LIVE per-topic model before writing, so the sheet can never
# fall back to the stale legacy im_allocate2.py output (old taxonomy, Dan Sullivan/Shaheen top picks).
# build_im_allocation.py sums per_topic_alloc.json (FLIP) + per_topic_coalition.json (COALITION) per member.
subprocess.run([sys.executable,os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),'engine','build_im_allocation.py')],check=True)
shutil.copy(src,bak); os.makedirs(r'C:\Temp',exist_ok=True); shutil.copy(src,tmp)
IM=json.load(open(MB_ROOT+r'\im_allocation.json'))
PT=json.load(open(MB_ROOT+r'\per_topic_alloc.json'))
flip=IM['flip']; coal=IM['coalition']; topics=list(PT.keys())
import openpyxl
bm=openpyxl.load_workbook(src,data_only=True)['Blended Model']
members=[bm.cell(r,1).value for r in roster.rows('Blended Model') if bm.cell(r,1).value]
xl=w.DispatchEx('Excel.Application'); xl.Visible=False; xl.DisplayAlerts=False
wb=xl.Workbooks.Open(tmp)
for s in list(wb.Worksheets):
    if s.Name=='IM Allocation': s.Delete()
ws=wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count)); ws.Name='IM Allocation'
ws.Cells(1,1).Value='MODE'
ws.Cells(1,2).Value='FLIP'                          # <- the switch: type FLIP or COALITION
ws.Cells(1,4).Value='Set B1 = FLIP or COALITION; "Active $" follows. FLIP = competitive Linear Threshold (Borodin, contest persuadables). COALITION = competitive Independent Cascade (Bharathi-Kempe, grow the bloc). Topic columns = per-sub-category flip allocation ($2M each, independent).'
hdr=['Member','Active $','Flip $','Coalition $']+topics
for c,h in enumerate(hdr,1): ws.Cells(2,c).Value=h
data=[]
for m in members:
    row=[m,None,round(flip.get(m,0)),round(coal.get(m,0))]+[round(PT[t].get(m,0)) for t in topics]
    data.append(row)
ws.Range(ws.Cells(3,1),ws.Cells(2+len(data),len(hdr))).Value=data
# Active $ formula (col B) = flip or coalition by mode
for r in range(3,3+len(data)):
    ws.Cells(r,2).Formula='=IF($B$1="COALITION",D%d,C%d)'%(r,r)
ws.Rows(2).Font.Bold=True; ws.Cells(1,1).Font.Bold=True; ws.Cells(1,2).Font.Bold=True
wb.Save(); wb.Close(False); xl.Quit(); shutil.copy(tmp,src)
print('IM Allocation sheet written: %d members x %d cols (flip,coal,%d topics) | mode switch in B1 | backup %s'%(len(data),len(hdr),len(topics),os.path.basename(bak)))
