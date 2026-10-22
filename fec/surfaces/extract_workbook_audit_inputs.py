# -*- coding: utf-8 -*-
"""Read-only extraction for the compact Moneyball audit workbook.

This script never authors an XLSX.  It converts the legacy source workbook's relevant cells into a
small JSON handoff; ``build_audit_workbook.mjs`` performs all spreadsheet creation and formatting.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import tempfile

import openpyxl

sys.path.insert(0,os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import roster
from moneyball_core import dump_atomic
from paths import ROOT,WORKBOOK


def sha256_file(path):
    h=hashlib.sha256()
    with open(path,'rb') as handle:
        for chunk in iter(lambda:handle.read(1024*1024),b''):
            h.update(chunk)
    return h.hexdigest()


def number(value):
    return float(value) if isinstance(value,(int,float)) and not isinstance(value,bool) else 0.0


def normalize_columns(rows):
    out=[list(rows[0]),list(rows[1])]
    data=[list(row) for row in rows[2:]]
    for col in range(2,len(rows[0])):
        values=[number(row[col]) for row in data]
        pos=max([0.0]+values); neg=abs(min([0.0]+values))
        for row,value in zip(data,values):
            row[col]=0.0 if value==0 else value/pos if value>0 and pos else value/neg if neg else 0.0
    out.extend(data)
    return out


def extract(workbook_path):
    # Bind the receipt to the exact parsed bytes, even if Excel replaces the source.
    # Large workbooks spill to a private temporary file instead of unbounded RAM.
    digest=hashlib.sha256()
    with tempfile.SpooledTemporaryFile(max_size=8*1024*1024,mode='w+b') as snapshot:
        with open(workbook_path,'rb') as handle:
            for block in iter(lambda:handle.read(1024*1024),b''):
                digest.update(block)
                snapshot.write(block)
        snapshot.seek(0)
        values=openpyxl.load_workbook(snapshot,read_only=True,data_only=True)
        try:
            raw=list(values['Legislative Score - Raw'].iter_rows(min_row=1,max_row=roster.last_row('Legislative Score - Raw'),min_col=1,max_col=160,values_only=True))  # 160: CU Tax Status block cols 153-160 added 2026-07-17
            rank_rows=list(values['Rank & Allocation'].iter_rows(min_row=roster.first_row('Rank & Allocation'),max_row=roster.last_row('Rank & Allocation'),min_col=1,max_col=7,values_only=True))
            current_norm=list(values['Legislative Score - Normalized'].iter_rows(min_row=roster.first_row('Legislative Score - Normalized'),max_row=roster.last_row('Legislative Score - Normalized'),min_col=1,max_col=160,values_only=True))
        finally:
            values.close()
    expected=normalize_columns(raw)
    mismatches=0
    for old,new in zip(current_norm,expected[2:]):
        for col in range(2,160):
            if abs(number(old[col])-number(new[col]))>1e-10:
                mismatches+=1
    return {
        'schema':'moneyball-workbook-audit-inputs-v1',
        'source_workbook':os.path.basename(workbook_path),
        'source_workbook_sha256':digest.hexdigest(),
        'legacy_normalized_cell_mismatches':mismatches,
        'rank_rows':[[row[0],number(row[3]),row[6]] for row in rank_rows],
        'legislative_raw':[list(row) for row in raw],
        'legislative_normalized':expected,
    }


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--workbook',default=WORKBOOK)
    parser.add_argument('--out',required=True)
    args=parser.parse_args()
    doc=extract(os.path.abspath(args.workbook))
    dump_atomic(doc,os.path.abspath(args.out),indent=1)
    print('extracted workbook audit inputs: %d members | %d displaced normalized cells'%(
          len(doc['rank_rows']),doc['legacy_normalized_cell_mismatches']))


if __name__=='__main__':
    main()
